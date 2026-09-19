#!/usr/bin/env python3
"""Deterministically apply format-spec.json to a DOCX and verify attributes."""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from artifact_io import atomic_write_text, commit_files, sibling_temp
from docx_semantics import (
    all_body_paragraphs,
    all_story_paragraphs,
    all_table_paragraphs,
    dominant_structural_style,
    fixed_text_sha256,
    has_drawing,
    has_visible_or_object_content,
    is_abstract_title_zh,
    is_display_equation,
    has_toc_field,
    matches_structural_detector,
    iter_document_nodes,
)
from format_spec_validation import load_and_validate
from compliance import annotate_satisfied_inputs, finalize_records, report as compliance_report
from role_registry import find_existing_style, generated_style, role_config, role_names, structural_detector, style_aliases
from resource_registry import resource_items
from submission_audit import audit_docx as audit_submission_docx
from section_executor import (
    add_page_field as _section_add_page_field,
    all_story_paragraphs,
    execute_section_plan,
    normalize_section_property_order,
    remove_page_fields,
    unlink_story_preserving_content,
)
from section_model import audit_plan_against_docx, compile_section_plan
from property_receipts import (
    audit_property_receipts,
    build_property_receipts,
    expected_receipt_ids,
    flatten,
)

ROLE_STYLES = {role: style_aliases(role) for role in role_names()}
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
         "distributed": WD_ALIGN_PARAGRAPH.DISTRIBUTE}
LINE = {"exact": WD_LINE_SPACING.EXACTLY, "at_least": WD_LINE_SPACING.AT_LEAST,
        "single": WD_LINE_SPACING.SINGLE, "one_point_five": WD_LINE_SPACING.ONE_POINT_FIVE,
        "double": WD_LINE_SPACING.DOUBLE, "multiple": WD_LINE_SPACING.MULTIPLE}

# A missing thesis section is an input gap, not a reason to abandon the
# formatting run.  Keep the placeholder deliberately neutral and make the
# pending state explicit in validation-report.json instead of treating it as
# real thesis content.
NEUTRAL_CONTENT_PLACEHOLDER = "——"
PLACEHOLDER_CONTENT_ROLES = {
    "heading_1", "heading_2", "heading_3", "heading_4",
    "heading_acknowledgments", "heading_appendix", "heading_conclusion",
    "heading_publications", "heading_references",
    "body_text", "abstract_body_zh", "abstract_body_en",
    "keywords_zh", "keywords_en", "bibliography_heading", "bibliography_entry",
}
PLACEHOLDER_CONTENT_LABELS = {
    "heading_1": "一级章节标题",
    "heading_2": "二级章节标题",
    "heading_3": "三级章节标题",
    "heading_4": "四级章节标题",
    "heading_acknowledgments": "致谢/后记章节标题",
    "heading_appendix": "附录章节标题",
    "heading_conclusion": "结论章节标题",
    "heading_publications": "学术成果章节标题",
    "heading_references": "参考文献章节标题",
    "body_text": "论文正文",
    "abstract_body_zh": "中文摘要正文",
    "abstract_body_en": "英文摘要正文",
    "keywords_zh": "中文关键词",
    "keywords_en": "英文关键词",
    "bibliography_heading": "参考文献标题",
    "bibliography_entry": "参考文献条目",
}


def load_json(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict): raise ValueError(f"{path} must contain an object")
    return data


def format_spec_blockers(spec: dict[str, Any], compliance_mode: str = "full") -> list[str]:
    """Identify concrete apply-time blockers for the selected compliance mode."""
    completeness = spec.get("completeness") if isinstance(spec.get("completeness"), dict) else {}
    blockers = []
    if spec.get("status") == "needs_clarification": blockers.append("status_needs_clarification")
    if completeness.get("unresolved_clause_ids"): blockers.append("unresolved_clauses")
    if completeness.get("missing_clause_ids"): blockers.append("missing_clauses")
    if spec.get("blocking_errors"): blockers.append("blocking_errors")
    if compliance_mode == "full" and completeness.get("unsupported_items"):
        blockers.append("unsupported_items")
    records = spec.get("clause_compliance") if isinstance(spec.get("clause_compliance"), list) else []
    if compliance_mode == "full":
        if not records: blockers.append("missing_clause_compliance")
        elif not compliance_report(records, "full", "analysis")["execution_ready"]:
            blockers.append("full_compliance_analysis_failed")
    return blockers
def canonicalize_docx_zip(path: Path) -> None:
    """Remove wall-clock ZIP metadata so identical OOXML yields identical DOCX bytes."""
    temporary = path.with_suffix(path.suffix + ".canonical")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w") as target:
        for name in source.namelist():
            original = source.getinfo(name)
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = original.compress_type
            info.comment = original.comment
            info.extra = b""
            info.internal_attr = original.internal_attr
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_names(doc: Document) -> set[str]:
    return {s.name for s in doc.styles}


def ensure_paragraph_style(doc: Document, name: str, base: str = "Normal"):
    if name in style_names(doc):
        return doc.styles[name]
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base in style_names(doc): style.base_style = doc.styles[base]
    return style


def _insert_paragraph_before(anchor: Paragraph, text: str, style_name: str) -> Paragraph:
    paragraph = anchor._parent.add_paragraph(text, style=style_name)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def _insert_paragraph_after(anchor: Paragraph, text: str, style_name: str) -> Paragraph:
    """Insert a paragraph after an XML paragraph, including inside table cells."""
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style_name
    paragraph.add_run(text)
    return paragraph


def insert_missing_content_placeholders(doc: Document, spec: dict[str, Any],
                                        mappings: dict[str, Any]) -> list[dict[str, Any]]:
    """Fill missing executable content roles with neutral, auditable placeholders.

    Structural mapping failures remain failures: if a table cell or semantic
    object exists but was assigned to the wrong role, silently adding a dummy
    paragraph would hide a formatting bug.  Only genuine ``content_missing``
    coverage gaps receive ``——``.  Existing placeholders are reported too, so
    a second run cannot accidentally turn a pending document into a ready one.
    """
    body = list(all_body_paragraphs(doc))
    pending: dict[str, dict[str, Any]] = {}

    for role in PLACEHOLDER_CONTENT_ROLES:
        mapping = mappings.get(role)
        if not mapping:
            continue
        count = sum(1 for paragraph in body
                    if paragraph.style.name == mapping["style_name"]
                    and paragraph.text.strip() == NEUTRAL_CONTENT_PLACEHOLDER)
        if count:
            pending[role] = {
                "role": role,
                "label": PLACEHOLDER_CONTENT_LABELS.get(role, role),
                "category": "thesis_content",
                "status": "pending_user_content",
                "placeholder": NEUTRAL_CONTENT_PLACEHOLDER,
                "style": mapping["style_name"],
                "count": count,
                "origin": "existing_placeholder",
            }

    coverage = role_coverage(doc, spec.get("roles", {}), mappings)
    for item in coverage:
        role = item["role"]
        if role not in PLACEHOLDER_CONTENT_ROLES:
            continue
        if item.get("expectation") != "required" or item.get("status") != "missing":
            continue
        # A structural match with the wrong style is a role-migration problem,
        # not absent thesis content.  The semantic-priority pass handles the
        # known heading/keyword case; other mapping failures must stay visible.
        if item.get("failure_type") != "content_missing":
            continue
        if role in pending:
            continue
        mapping = mappings.get(role)
        if not mapping:
            continue
        paragraph = doc.add_paragraph(NEUTRAL_CONTENT_PLACEHOLDER, mapping["style_name"])
        pending[role] = {
            "role": role,
            "label": PLACEHOLDER_CONTENT_LABELS.get(role, role),
            "category": "thesis_content",
            "status": "pending_user_content",
            "placeholder": NEUTRAL_CONTENT_PLACEHOLDER,
            "style": mapping["style_name"],
            "count": 1,
            "origin": "inserted_document_end",
            "reason": "required content role had no source content mapped",
            "source_structural_count": item.get("structural_count", 0),
            "paragraph_text": paragraph.text,
        }
        body.append(paragraph)
    return [pending[role] for role in sorted(pending)]


def _abstract_anchor(doc: Document) -> Paragraph | None:
    return next((p for p in doc.paragraphs if is_abstract_title_zh(p)), None)


def _document_start_anchor(doc: Document) -> tuple[Paragraph, bool]:
    """Return a body anchor at the absolute document start.

    A cover is a document-level front-matter region, not content that belongs
    next to the first abstract.  The synthetic-anchor path also handles source
    fragments whose first body item is a table rather than a paragraph.
    """
    body = doc._body._element
    first = next((child for child in body if child.tag != qn("w:sectPr")), None)
    if first is not None and first.tag == qn("w:p"):
        return Paragraph(first, doc._body), False
    anchor = doc.add_paragraph()
    body.remove(anchor._p)
    if first is None:
        body.insert(0, anchor._p)
    else:
        first.addprevious(anchor._p)
    return anchor, True


def _paragraph_has_page_break(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))


def _remove_generated_paragraphs(doc: Document, style_names_to_remove: set[str]) -> int:
    removed = 0
    for paragraph in list(doc.paragraphs):
        if paragraph.style.name in style_names_to_remove:
            paragraph._element.getparent().remove(paragraph._element); removed += 1
    return removed


def _generated_cover_block(doc: Document, institution: str) -> list[Paragraph]:
    """Return the bounded front-matter block produced by the cover compiler."""
    paragraphs = list(doc.paragraphs)
    institution_styles = {"Thesis Cover Institution"}
    start = next(
        (index for index, paragraph in enumerate(paragraphs)
         if paragraph.style.name in institution_styles
         and paragraph.text.strip() == str(institution or "").strip()),
        None,
    )
    if start is None:
        return []
    end = next(
        (index for index in range(start, len(paragraphs))
         if _paragraph_has_page_break(paragraphs[index])),
        None,
    )
    if end is None:
        return []
    return paragraphs[start:end + 1]


def _remove_generated_cover_block(doc: Document, style_names_to_remove: set[str],
                                  institution: str) -> int:
    """Remove only a previously generated cover block.

    ``Title`` and ``English Title`` are legitimate source-document styles, so
    removing every paragraph carrying a mapped title style would delete the
    thesis titles before the cover compiler runs.  A generated cover has a
    stable institution anchor and an explicit page break; use that bounded
    block as the idempotent cleanup scope instead of treating a semantic style
    name as an ownership marker.
    """
    block = _generated_cover_block(doc, institution)
    if not block:
        # A cover without its terminal page break is not a verified generated
        # block; leave it intact so the missing structural invariant remains
        # visible to validation rather than being silently rewritten.
        return 0
    removed = 0
    for paragraph in block:
        if paragraph.style.name in style_names_to_remove:
            paragraph._element.getparent().remove(paragraph._element)
            removed += 1
    return removed


def _metadata_value(metadata: dict[str, Any], field_id: str) -> str:
    value = metadata.get(field_id)
    if field_id == "co_supervisors" and isinstance(value, list):
        return "、".join(str(item.get("name", "")) for item in value if isinstance(item, dict))
    return "" if value is None else str(value)


def _trusted_cover_metadata(profile: dict[str, Any]) -> tuple[dict[str, Any], bool]:
    metadata = profile.get("cover_metadata", {})
    if not isinstance(metadata, dict):
        return {}, False
    trust = metadata.get("trust", {})
    trusted = (isinstance(trust, dict) and trust.get("confirmed") is True and
               trust.get("source") in {"user_confirmed", "source_document", "authoritative_record"})
    return (metadata if trusted else {}), trusted


def _cover_field_value(cover: dict[str, Any], metadata: dict[str, Any], field: dict[str, Any]) -> tuple[str, bool]:
    value = _metadata_value(metadata, field["id"])
    if value.strip():
        return value, False
    if field.get("display_policy") == "required":
        return str(cover.get("missing_value_placeholder") or "——"), True
    return "", False


def _compile_non_public_administration(cover: dict[str, Any], profile: dict[str, Any],
                                       metadata: dict[str, Any], trusted: bool) -> dict[str, Any] | None:
    administration = cover.get("non_public_administration")
    if not isinstance(administration, dict):
        return None
    security_level = profile.get("security_level")
    fields: list[dict[str, Any]] = []
    for field in sorted(administration.get("fields", []), key=lambda item: item.get("order", 0)):
        value = _metadata_value(metadata, field.get("id", "")) if trusted else ""
        fields.append({
            "id": field.get("id"), "label": field.get("label"), "order": field.get("order"),
            "value": value if value.strip() else "",
            "value_kind": "trusted" if value.strip() else "omitted",
            "source": field.get("value_from"),
            "display_policy": field.get("display_policy"),
        })
    if security_level == "public":
        status = "blank_public"
    elif security_level in {"restricted", "classified"} and all(item["value_kind"] == "trusted" for item in fields):
        status = "provided_unverified"
    else:
        status = "pending_external_approval"
    return {
        "applicability": administration.get("applicability"),
        "public_policy": administration.get("public_policy"),
        "source_region": administration.get("source_region"),
        "security_level": security_level,
        "status": status,
        "fields": fields,
    }


def compile_cover_contract(cover: dict[str, Any], profile: dict[str, Any]) -> dict[str, Any]:
    """Compile written requirements and trusted instance data into an executable contract."""
    metadata, trusted = _trusted_cover_metadata(profile)
    fields = []
    for field in sorted(cover.get("fields", []), key=lambda item: item["order"]):
        value, placeholder = _cover_field_value(cover, metadata, field)
        fields.append({
            "id": field["id"], "label": field["label"], "order": field["order"],
            "required": field.get("display_policy") == "required", "value": value,
            "value_kind": "placeholder" if placeholder else ("trusted" if value else "omitted"),
            "source": field.get("value_from"),
            "style_role": field["id"] if field["id"] in {"title_zh", "title_en"} else "cover_field_value",
        })
    contract = {
        "schema_version": "1.0", "institution": cover.get("institution"),
        # Normalize the legacy abstract-relative declaration to the actual
        # executable contract: every generated cover starts the document.
        "before_role": "document_start",
        "metadata_status": "trusted" if trusted else "pending",
        "missing_value_policy": cover.get("missing_value_policy", "placeholder"),
        "missing_value_placeholder": cover.get("missing_value_placeholder", "——"),
        "layout_id": cover.get("layout_id", "linear"),
        "fields": fields,
    }
    administration = _compile_non_public_administration(cover, profile, metadata, trusted)
    if administration is not None:
        contract["non_public_administration"] = administration
    return contract


def apply_cover(doc: Document, cover: dict[str, Any], profile: dict[str, Any],
                contract: dict[str, Any] | None = None,
                mappings: dict[str, Any] | None = None) -> dict[str, Any]:
    counts: dict[str, Any] = {"removed_generated_paragraphs": 0, "fields_written": 0,
                              "trusted_fields_written": 0, "placeholder_fields_written": 0,
                              "metadata_status": "not_applicable", "metadata_pending_fields": []}
    if not cover: return counts
    contract = contract or compile_cover_contract(cover, profile)
    counts["metadata_status"] = contract["metadata_status"]
    mappings = mappings or {}
    styles = {
        "institution": "Thesis Cover Institution",
        "title_zh": mappings.get("thesis_title_zh", {}).get("style_name") or generated_style("thesis_title_zh") or "Thesis Cover Title ZH",
        "title_en": mappings.get("thesis_title_en", {}).get("style_name") or generated_style("thesis_title_en") or "Thesis Cover Title EN",
        "label": generated_style("cover_field_label") or "Thesis Cover Field Label",
        "value": generated_style("cover_field_value") or "Thesis Cover Field Value",
    }
    counts["removed_generated_paragraphs"] = _remove_generated_cover_block(
        doc, set(styles.values()), cover.get("institution", "")
    )
    for name in styles.values(): ensure_paragraph_style(doc, name)
    anchor, synthetic_anchor = _document_start_anchor(doc)
    if cover.get("layout_id") == "official_template_regions":
        raise ValueError(
            "cover layout official_template_regions must be assembled from a verified official "
            "DOCX template before style application; synthetic cover drawing is prohibited"
        )
    institution = _insert_paragraph_before(anchor, cover["institution"], styles["institution"])
    institution.alignment = WD_ALIGN_PARAGRAPH.CENTER
    last_cover_paragraph = institution
    for field in contract["fields"]:
        value = field["value"]; placeholder = field["value_kind"] == "placeholder"
        if not value: continue
        style = styles["value"]
        if field["id"] == "title_zh": style = styles["title_zh"]
        elif field["id"] == "title_en": style = styles["title_en"]
        paragraph = _insert_paragraph_before(anchor, value if field["id"] in {"title_zh", "title_en"} else f"{field['label']}：{value}", style)
        if field["id"] == "title_zh":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs: run.font.name = "SimHei"; run.font.size = Pt(22)
        elif field["id"] == "title_en":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs: run.font.name = "Times New Roman"; run.font.size = Pt(16); run.bold = True
        last_cover_paragraph = paragraph
        counts["fields_written"] += 1
        if placeholder:
            counts["placeholder_fields_written"] += 1
            counts["metadata_pending_fields"].append(field["id"])
        else:
            counts["trusted_fields_written"] += 1
    # A generated cover is an independent front-matter page even when the
    # input is only a chapter/body fragment and contains no abstract at all.
    last_cover_paragraph.add_run().add_break(WD_BREAK.PAGE)
    if synthetic_anchor:
        anchor._element.getparent().remove(anchor._element)
    counts["placement"] = "document_start"
    counts["page_break_written"] = 1
    administration = contract.get("non_public_administration")
    if isinstance(administration, dict):
        counts["non_public_administration_status"] = administration.get("status")
        counts["non_public_administration_pending_fields"] = [
            item.get("id") for item in administration.get("fields", [])
            if isinstance(item, dict) and item.get("value_kind") != "trusted"
        ]
        # Public theses deliberately keep this region blank.  Non-public
        # administration is not drawn into a generic linear cover: without a
        # verified official region, writing it here would recreate the
        # original R00138 semantic/layout error.
    return counts


def _declaration_anchor(doc: Document, before_role: str) -> tuple[Paragraph | None, bool]:
    """Resolve the current run's declaration insertion anchor by semantic role."""
    if before_role == "document_start":
        return _document_start_anchor(doc)
    detector = structural_detector(before_role)
    if detector:
        return next((p for p in doc.paragraphs if matches_structural_detector(p, detector)), None), False
    aliases = set(style_aliases(before_role))
    return next((p for p in doc.paragraphs if p.style.name in aliases), None), False


def _declaration_styles() -> dict[str, str]:
    roles = {
        "heading": "declaration_heading",
        "body": "declaration_body",
        "author": "signature_line_author",
        "supervisor": "signature_line_supervisor",
        "date": "signature_date_line",
    }
    return {
        key: generated_style(role) or f"Thesis {role.replace('_', ' ').title()}"
        for key, role in roles.items()
    }


def _resource_body_parts(resource: dict[str, Any]) -> list[str]:
    parts = resource.get("body_parts")
    if isinstance(parts, list):
        return [part for part in parts if isinstance(part, str) and part.strip()]
    body = resource.get("body")
    return [body] if isinstance(body, str) and body.strip() else []


def apply_declarations(doc: Document, declarations: dict[str, Any],
                       resources: dict[str, dict[str, Any]] | None = None) -> dict[str, int]:
    counts = {"removed_generated_paragraphs": 0, "items_written": 0, "placeholders_written": 0}
    if not declarations: return counts
    anchor, synthetic_anchor = _declaration_anchor(doc, str(declarations.get("before_role", "document_start")))
    if anchor is None: return counts
    styles = _declaration_styles()
    generated_styles = set(styles.values())
    counts["removed_generated_paragraphs"] = _remove_generated_paragraphs(doc, generated_styles)
    for name in generated_styles: ensure_paragraph_style(doc, name)
    resources = resources or {}
    for item in declarations.get("items", []):
        resource_id = item.get("resource_id")
        resource = resources.get(resource_id) if isinstance(resource_id, str) else None
        if not isinstance(resource, dict):
            raise ValueError(f"declaration resource {resource_id!r} is not registered in this run")
        body_parts = _resource_body_parts(resource)
        if not body_parts:
            raise ValueError(f"declaration resource {resource_id!r} has no fixed body text")
        heading = _insert_paragraph_before(anchor, resource["heading"], styles["heading"])
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for body in body_parts:
            _insert_paragraph_before(anchor, body, styles["body"])
        counts["items_written"] += 1
        for placeholder in item.get("signature_placeholders", []):
            role = placeholder["role"]
            style = styles[role]
            _insert_paragraph_before(anchor, f"{placeholder['label']}：________________", style)
            counts["placeholders_written"] += 1
    if synthetic_anchor:
        anchor._element.getparent().remove(anchor._element)
    return counts


def audit_cover(doc: Document, cover: dict[str, Any], profile: dict[str, Any],
                mappings: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    if not cover: return []
    findings: list[dict[str, Any]] = []
    metadata, _trusted = _trusted_cover_metadata(profile); mappings = mappings or {}
    paragraphs = doc.paragraphs
    nonempty = [(i, p) for i, p in enumerate(paragraphs) if p.text.strip()]
    first_content_index = nonempty[0][0] if nonempty else None
    if cover.get("layout_id") == "official_template_regions":
        # The style applier cannot verify protected-package fidelity; the
        # assembly executor must supply fixed-region hash evidence.
        findings.append({"role": "cover", "property": "assembly_evidence",
                         "template_value": "not supplied to style applier",
                         "required_value": "verified official-template fixed-region hashes"})
        return findings
    cover_block = _generated_cover_block(doc, cover.get("institution", ""))
    institution_matches = [p for p in cover_block
                           if p.text.strip() == cover.get("institution", "").strip()]
    if len(institution_matches) != 1:
        findings.append({"role": "cover", "property": "institution", "template_value": len(institution_matches),
                         "required_value": "exactly one institution at document start"})
    elif (institution_matches[0]._p.getroottree().getpath(institution_matches[0]._p)
          != nonempty[0][1]._p.getroottree().getpath(nonempty[0][1]._p)):
        findings.append({"role": "cover", "property": "placement", "template_value": "not_first_content",
                         "required_value": "first content-bearing paragraph in the document"})
    generated_cover_styles = {
        "Thesis Cover Institution",
        mappings.get("thesis_title_zh", {}).get("style_name") or generated_style("thesis_title_zh") or "Thesis Cover Title ZH",
        mappings.get("thesis_title_en", {}).get("style_name") or generated_style("thesis_title_en") or "Thesis Cover Title EN",
        generated_style("cover_field_value") or "Thesis Cover Field Value",
    }
    generated_cover = [p for p in cover_block if p.style.name in generated_cover_styles]
    if not generated_cover or not _paragraph_has_page_break(generated_cover[-1]):
        findings.append({"role": "cover", "property": "page_break", "template_value": False,
                         "required_value": "explicit page break after generated cover"})
    for field in cover.get("fields", []):
        value, placeholder = _cover_field_value(cover, metadata, field)
        if not value: continue
        expected = value if field["id"] in {"title_zh", "title_en"} else f"{field['label']}：{value}"
        expected_style = {"title_zh": mappings.get("thesis_title_zh", {}).get("style_name") or generated_style("thesis_title_zh") or "Thesis Cover Title ZH",
                          "title_en": mappings.get("thesis_title_en", {}).get("style_name") or generated_style("thesis_title_en") or "Thesis Cover Title EN"}.get(
            field["id"], generated_style("cover_field_value") or "Thesis Cover Field Value")
        matches = [p for p in cover_block if p.text.strip() == expected and p.style.name == expected_style]
        if len(matches) != 1:
            requirement = "exactly one neutral placeholder on the generated cover" if placeholder else "exactly one trusted metadata value on the generated cover"
            findings.append({"role": "cover", "property": f"fields.{field['id']}", "template_value": len(matches), "required_value": requirement})
    administration = compile_cover_contract(cover, profile).get("non_public_administration")
    if isinstance(administration, dict):
        if administration.get("status") == "blank_public":
            # A public thesis must not receive approval placeholders or
            # ordinary-cover substitutions.  The generated cover compiler
            # never writes this region; its explicit status is the evidence.
            pass
        else:
            findings.append({
                "role": "cover",
                "property": "non_public_administration",
                "template_value": administration.get("status"),
                "required_value": "verified_official_region_and_external_approval",
                "failure_type": "external_or_official_region_required",
            })
    return findings


def audit_declarations(doc: Document, declarations: dict[str, Any],
                       resources: dict[str, dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    if not declarations: return []
    findings: list[dict[str, Any]] = []
    resources = resources or {}; paragraphs = doc.paragraphs
    before_role = str(declarations.get("before_role", "document_start"))
    anchor, _synthetic = _declaration_anchor(doc, before_role)
    if anchor is None:
        return [{"role": "declarations", "property": "before_role", "template_value": False,
                 "required_value": before_role}]
    # ``Document.paragraphs`` creates wrapper objects on each access.  The
    # anchor returned by ``_declaration_anchor`` may therefore represent the
    # same XML node as an item in this list without being the same Python
    # object; ``list.index(anchor)`` is not a stable document-order lookup.
    anchor_element = anchor._p
    boundary = next(
        (index for index, paragraph in enumerate(paragraphs)
         if paragraph._p is anchor_element),
        len(paragraphs),
    )
    for item in declarations.get("items", []):
        resource_id = item.get("resource_id")
        resource = resources.get(resource_id) if isinstance(resource_id, str) else None
        if not isinstance(resource, dict):
            findings.append({"role": "declarations", "property": f"{item.get('id', 'unknown')}.resource_id",
                             "template_value": resource_id, "required_value": "registered current-run resource"})
            continue
        body_parts = _resource_body_parts(resource)
        if item.get("version") != resource.get("version") or item.get("sha256") != resource.get("sha256"):
            findings.append({"role": "declarations", "property": f"{item['id']}.resource_version",
                             "template_value": [resource.get('version'), resource.get('sha256')],
                             "required_value": [item.get('version'), item.get('sha256')]})
        matches: list[tuple[int, list[int]]] = []
        for heading_index, paragraph in enumerate(paragraphs[:boundary]):
            if paragraph.text.strip() != str(resource.get("heading", "")).strip():
                continue
            body_indices: list[int] = []
            cursor = heading_index + 1
            for body in body_parts:
                found = next((index for index in range(cursor, boundary)
                              if fixed_text_sha256(paragraphs[index].text) == fixed_text_sha256(body)), None)
                if found is None:
                    body_indices = []
                    break
                body_indices.append(found)
                cursor = found + 1
            if body_indices:
                matches.append((heading_index, body_indices))
        if len(matches) != 1:
            findings.append({"role": "declarations", "property": f"{item['id']}.fixed_text",
                             "template_value": len(matches),
                             "required_value": "exactly one heading followed by all source body paragraphs"})
        elif matches[0][1][-1] >= boundary:
            findings.append({"role": "declarations", "property": f"{item['id']}.order",
                             "template_value": [matches[0][0], *matches[0][1], boundary],
                             "required_value": "fixed text must precede the declared anchor role"})
    before_texts = [p.text.strip() for p in paragraphs[:boundary]]
    expected_placeholders = [f"{placeholder['label']}：________________"
                             for item in declarations.get("items", [])
                             for placeholder in item.get("signature_placeholders", [])]
    for text in sorted(set(expected_placeholders)):
        expected_count = expected_placeholders.count(text)
        actual_count = before_texts.count(text)
        if actual_count != expected_count:
            findings.append({"role": "declarations", "property": "signature_placeholders",
                             "template_value": {"text": text, "count": actual_count},
                             "required_value": {"attestation_scope": "placeholder_presence_only", "text": text,
                                                "count": expected_count}})
    return findings


def resolve_style(doc: Document, role: str, explicit: dict[str, str], claimed: dict[str, str]) -> tuple[str, bool]:
    names = style_names(doc)
    requested = explicit.get(role)
    if requested and requested in names and requested not in claimed: return requested, False
    existing = find_existing_style(role, names - set(claimed))
    if existing: return existing, False
    name = requested or f"Thesis {role.replace('_', ' ').title()}"
    if name in names and name in claimed:
        name = f"Thesis {role.replace('_', ' ').title()}"
    if name not in names:
        created = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        base_name = requested if requested in names else next((n for n in ROLE_STYLES.get(role, []) if n in names), None)
        if base_name:
            created.base_style = doc.styles[base_name]
    return name, True


def structural_source_style(doc: Document, role: str) -> str | None:
    """Infer a source style when Word structure is stronger than its name."""
    if role == "table_text":
        paragraphs = all_table_paragraphs(doc)
    elif role == "equation":
        paragraphs = all_body_paragraphs(doc)
    else:
        return None
    return dominant_structural_style(paragraphs, structural_detector(role))


def _new_role_style(doc: Document, role: str, source_name: str | None = None) -> str:
    """Create or reuse a neutral role style without mutating a shared style."""
    name = generated_style(role) or f"Thesis {role.replace('_', ' ').title()}"
    if name not in style_names(doc):
        created = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if source_name and source_name in style_names(doc):
            created.base_style = doc.styles[source_name]
    return name


def inherited_style_value(style, getter):
    """Resolve a style property through its Word base-style chain."""
    seen = set()
    current = style
    while current is not None and current.style_id not in seen:
        seen.add(current.style_id)
        value = getter(current)
        if value is not None:
            return value
        current = current.base_style
    return None


def style_snapshot(style) -> dict[str, Any]:
    reverse_align = {value: key for key, value in ALIGN.items()}
    reverse_line = {value: key for key, value in LINE.items()}
    def cjk(item):
        rpr = item.element.rPr
        fonts = rpr.rFonts if rpr is not None else None
        return fonts.get(qn("w:eastAsia")) if fonts is not None else None
    size = inherited_style_value(style, lambda item: item.font.size)
    alignment = inherited_style_value(style, lambda item: item.paragraph_format.alignment)
    spacing_rule = inherited_style_value(style, lambda item: item.paragraph_format.line_spacing_rule)
    spacing = inherited_style_value(style, lambda item: item.paragraph_format.line_spacing)
    outline_level = inherited_style_value(
        style,
        lambda item: (
            int(item.element.pPr.outlineLvl.get(qn("w:val"))) + 1
            if item.element.pPr is not None and item.element.pPr.outlineLvl is not None
            else None
        ),
    )
    line_spacing = None
    if spacing is not None:
        line_type = reverse_line.get(spacing_rule)
        if line_type in {"exact", "at_least"}:
            line_spacing = {"type": line_type, "value": round(spacing.pt, 3), "unit": "pt"}
        elif line_type:
            line_spacing = {"type": line_type, "value": float(spacing), "unit": "multiple"}
    points = lambda value: round(value.pt, 3) if value is not None else None
    return {
        "font": {"latin": inherited_style_value(style, lambda item: item.font.name),
                 "cjk": inherited_style_value(style, cjk),
                 "size_pt": round(size.pt, 3) if size else None,
                 "bold": inherited_style_value(style, lambda item: item.font.bold),
                 "italic": inherited_style_value(style, lambda item: item.font.italic)},
        "paragraph": {"alignment": reverse_align.get(alignment),
                      "first_line_indent_pt": points(inherited_style_value(style, lambda item: item.paragraph_format.first_line_indent)),
                      "left_indent_pt": points(inherited_style_value(style, lambda item: item.paragraph_format.left_indent)),
                      "right_indent_pt": points(inherited_style_value(style, lambda item: item.paragraph_format.right_indent)),
                      "space_before_pt": points(inherited_style_value(style, lambda item: item.paragraph_format.space_before)) or 0,
                      "space_after_pt": points(inherited_style_value(style, lambda item: item.paragraph_format.space_after)) or 0,
                      "line_spacing": line_spacing,
                      "outline_level": outline_level,
                      "page_break_before": inherited_style_value(style, lambda item: item.paragraph_format.page_break_before),
                      "keep_with_next": inherited_style_value(style, lambda item: item.paragraph_format.keep_with_next)},
    }


def merge_role_style_defaults(defaults: dict[str, Any], requirement: dict[str, Any]) -> dict[str, Any]:
    """Overlay written role requirements on complete official style defaults."""
    result = {
        key: ({nested_key: copy.deepcopy(nested_value)
               for nested_key, nested_value in value.items() if nested_value is not None}
              if isinstance(value, dict) else copy.deepcopy(value))
        for key, value in defaults.items() if value is not None
    }
    for key, value in requirement.items():
        # JSON null means the written source did not specify this property; it
        # must not erase an observable official-template fallback.
        if value is None:
            continue
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = merge_role_style_defaults(result[key], value)
        else:
            result[key] = copy.deepcopy(value)
    # Relative paragraph measurements and their concrete point equivalents are
    # alternate representations of the same property.  A written requirement
    # must replace, rather than coexist with, an official-template fallback;
    # otherwise set_paragraph() gives the stale point value precedence while
    # validation expects the character/line-derived value.
    paragraph = result.get("paragraph")
    required_paragraph = requirement.get("paragraph")
    if isinstance(paragraph, dict) and isinstance(required_paragraph, dict):
        aliases = (
            ("first_line_indent_chars", "first_line_indent_pt"),
            ("space_before_lines", "space_before_pt"),
            ("space_after_lines", "space_after_pt"),
        )
        for relative, absolute in aliases:
            if required_paragraph.get(relative) is not None:
                paragraph.pop(absolute, None)
            elif required_paragraph.get(absolute) is not None:
                paragraph.pop(relative, None)
    return result


def official_role_defaults(docx: Path, mappings: dict[str, str]) -> dict[str, dict[str, Any]]:
    doc = Document(docx)
    names = style_names(doc)
    return {role: style_snapshot(doc.styles[name]) for role, name in mappings.items() if name in names}


def set_font(style, spec: dict[str, Any]) -> None:
    font = style.font
    if spec.get("latin"): font.name = spec["latin"]
    if spec.get("size_pt") is not None: font.size = Pt(float(spec["size_pt"]))
    if spec.get("bold") is not None: font.bold = bool(spec["bold"])
    if spec.get("italic") is not None: font.italic = bool(spec["italic"])
    rpr = style.element.get_or_add_rPr(); fonts = rpr.get_or_add_rFonts()
    if spec.get("cjk"): fonts.set(qn("w:eastAsia"), spec["cjk"])
    if spec.get("latin"):
        for key in ("ascii", "hAnsi", "cs"): fonts.set(qn("w:" + key), spec["latin"])


def set_paragraph(style, spec: dict[str, Any]) -> None:
    p = style.paragraph_format
    if spec.get("alignment") in ALIGN: p.alignment = ALIGN[spec["alignment"]]
    if spec.get("first_line_indent_pt") is not None: p.first_line_indent = Pt(float(spec["first_line_indent_pt"]))
    elif spec.get("first_line_indent_chars") is not None:
        size = style.font.size.pt if style.font.size else 12
        p.first_line_indent = Pt(float(spec["first_line_indent_chars"]) * size)
    if spec.get("left_indent_pt") is not None: p.left_indent = Pt(float(spec["left_indent_pt"]))
    if spec.get("right_indent_pt") is not None: p.right_indent = Pt(float(spec["right_indent_pt"]))
    if spec.get("space_before_pt") is not None: p.space_before = Pt(float(spec["space_before_pt"]))
    if spec.get("space_after_pt") is not None: p.space_after = Pt(float(spec["space_after_pt"]))
    line_pt = float(spec.get("spacing_line_height_pt") or (style.font.size.pt if style.font.size else 12))
    if spec.get("space_before_lines") is not None: p.space_before = Pt(float(spec["space_before_lines"]) * line_pt)
    if spec.get("space_after_lines") is not None: p.space_after = Pt(float(spec["space_after_lines"]) * line_pt)
    ls = spec.get("line_spacing")
    if isinstance(ls, dict):
        typ, value = ls.get("type"), float(ls.get("value", 1))
        if typ in {"exact", "at_least"}: p.line_spacing = Pt(value)
        else: p.line_spacing = value
        if typ in LINE: p.line_spacing_rule = LINE[typ]
    if spec.get("outline_level") is not None:
        ppr = style.element.get_or_add_pPr()
        outline = ppr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            ppr.append(outline)
        outline.set(qn("w:val"), str(int(spec["outline_level"]) - 1))
    if spec.get("page_break_before") is not None: p.page_break_before = bool(spec["page_break_before"])
    if spec.get("keep_with_next") is not None: p.keep_with_next = bool(spec["keep_with_next"])


def normalized_expected(style, role_spec: dict[str, Any]) -> dict[str, Any]:
    """Translate character/line-relative requirements to concrete Word points."""
    expected = {k: dict(v) if isinstance(v, dict) else v for k, v in role_spec.items() if k in {"font", "paragraph"}}
    para = dict(expected.get("paragraph", {})); expected["paragraph"] = para
    size = float(expected.get("font", {}).get("size_pt") or (style.font.size.pt if style.font.size else 12))
    spacing_line_height = float(para.pop("spacing_line_height_pt", size))
    if "first_line_indent_chars" in para:
        para["first_line_indent_pt"] = float(para.pop("first_line_indent_chars")) * size
    if "space_before_lines" in para:
        para["space_before_pt"] = float(para.pop("space_before_lines")) * spacing_line_height
    if "space_after_lines" in para:
        para["space_after_pt"] = float(para.pop("space_after_lines")) * spacing_line_height
    return expected


def diff_expected(actual: dict[str, Any], expected: dict[str, Any], prefix="") -> list[dict[str, Any]]:
    out = []
    for key, val in expected.items():
        path = f"{prefix}.{key}" if prefix else key
        if isinstance(val, dict): out.extend(diff_expected(actual.get(key, {}) if isinstance(actual.get(key), dict) else {}, val, path))
        else:
            old = actual.get(key)
            if old != val and not (isinstance(old, (int, float)) and isinstance(val, (int, float)) and abs(old-val) < .02):
                out.append({"property": path, "template_value": old, "required_value": val})
    return out


def _paragraph_runs_including_nested(paragraph: Paragraph) -> list[Run]:
    """Return direct and hyperlink/field-contained runs in document order."""
    return [Run(node, paragraph) for node in paragraph._p.iter(qn("w:r"))]


def apply_direct_format(paragraph, role_spec: dict[str, Any]) -> None:
    """Eliminate direct-formatting overrides on paragraphs governed by a role."""
    font_spec = role_spec.get("font", {})
    # ``Paragraph.runs`` omits runs nested under ``w:hyperlink`` and several
    # field containers.  Styling only that property leaves visible text
    # instances with stale direct formatting even though the shared style
    # audit passes.  Walk the underlying XML and wrap each run instead.
    for run in _paragraph_runs_including_nested(paragraph):
        if font_spec.get("latin"): run.font.name = font_spec["latin"]
        if font_spec.get("size_pt") is not None: run.font.size = Pt(float(font_spec["size_pt"]))
        if font_spec.get("bold") is not None: run.font.bold = bool(font_spec["bold"])
        if font_spec.get("italic") is not None: run.font.italic = bool(font_spec["italic"])
        rpr = run._element.get_or_add_rPr(); fonts = rpr.get_or_add_rFonts()
        if font_spec.get("cjk"): fonts.set(qn("w:eastAsia"), font_spec["cjk"])
        if font_spec.get("latin"):
            for key in ("ascii", "hAnsi", "cs"): fonts.set(qn("w:" + key), font_spec["latin"])
    # Direct paragraph properties are set to the same values as the style so
    # legacy documents with local overrides still render deterministically.
    pseudo = type("ParagraphStyleProxy", (), {"paragraph_format": paragraph.paragraph_format,
                                               "font": paragraph.style.font,
                                               # set_paragraph() uses the element only
                                               # for outline-level OOXML updates.
                                               "element": paragraph._p})()
    set_paragraph(pseudo, role_spec.get("paragraph", {}))


def chinese_number(value: int) -> str:
    digits = "零一二三四五六七八九"
    if value < 10: return digits[value]
    if value < 20: return "十" + (digits[value % 10] if value % 10 else "")
    if value < 100: return digits[value // 10] + "十" + (digits[value % 10] if value % 10 else "")
    return str(value)


def _prepend_text(paragraph, text: str) -> None:
    run = OxmlElement("w:r"); node = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace(): node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    node.text = text; run.append(node)
    ppr = paragraph._p.pPr
    paragraph._p.insert(1 if ppr is not None else 0, run)


def apply_numbering(doc: Document, mappings: dict[str, Any], roles: dict[str, Any]) -> tuple[dict[str, int], list[dict[str, Any]]]:
    changed = {}; issues = []; counters = {"heading_1": 0, "heading_2": 0, "heading_3": 0}
    for paragraph in doc.paragraphs:
        role = next((r for r in counters if r in mappings and paragraph.style.name == mappings[r]["style_name"]), None)
        if not role or not paragraph.text.strip() or not roles.get(role, {}).get("numbering"): continue
        level = int(role[-1]); counters[role] += 1
        if level == 1: counters["heading_2"] = counters["heading_3"] = 0
        elif level == 2: counters["heading_3"] = 0
        numbering = roles[role]["numbering"]
        numpr = paragraph._p.find("w:pPr/w:numPr", {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
        if numpr is not None:
            issues.append({"role": role, "property": "numbering", "reason": "native Word numbering preserved; static prefix not added"})
            continue
        fmt = numbering.get("format", "{chapter}")
        replacements = {"{chapter_cn}": chinese_number(counters["heading_1"]), "{chapter}": str(counters["heading_1"]),
                        "{section}": str(counters["heading_2"]), "{subsection}": str(counters["heading_3"]), "%1": str(counters["heading_1"])}
        prefix = fmt
        for token, value in replacements.items(): prefix = prefix.replace(token, value)
        if not re.match(r"^(?:" + re.escape(prefix) + r"|第[一二三四五六七八九十百\d]+章|\d+(?:\.\d+){0,2})(?:\s|$)", paragraph.text):
            _prepend_text(paragraph, prefix + (roles[role].get("separator") or "  "))
            changed[role] = changed.get(role, 0) + 1
    return changed, issues


def add_page_number(paragraph, fmt: str) -> None:
    """Compatibility wrapper; new code must execute a compiled SectionPlan."""
    instructions = list(paragraph._p.xpath('.//w:fldSimple/@w:instr'))
    instructions.extend(paragraph._p.xpath('.//w:instrText/text()'))
    if any(_field_opcode(instruction) == "PAGE" for instruction in instructions):
        return
    switch = "ROMAN" if fmt == "roman_upper" else "roman" if fmt == "roman" else "ARABIC"
    _section_add_page_field(paragraph, switch)


def _field_opcode(instruction: str) -> str | None:
    match = re.match(r"\s*([A-Za-z]+)\b", instruction or "")
    return match.group(1).upper() if match else None


def set_page_number_format(section, fmt: str, start: int | None = None) -> None:
    """Compatibility wrapper; new code must execute a compiled SectionPlan."""
    from section_executor import _set_page_number_format
    mapped = {"roman": "lowerRoman", "roman_upper": "upperRoman", "decimal": "decimal"}.get(fmt, fmt)
    _set_page_number_format(section, mapped, start)


def _paragraph_section_indices(doc: Document) -> dict[Any, int]:
    """Map body and table-cell paragraph XML elements to their zero-based section."""
    indices: dict[Any, int] = {}; current = 0
    for child in list(doc._body._element):
        for paragraph in child.iter(qn("w:p")):
            indices[paragraph] = current
        if child.tag == qn("w:p"):
            ppr = child.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                current += 1
    return indices


def _find_body_start_paragraph(doc: Document, selector: dict[str, Any], mappings: dict[str, Any]):
    strategy = selector.get("strategy")
    if strategy == "first_heading_1":
        style = mappings.get("heading_1", {}).get("style_name")
        candidate_styles = [style] if style else ROLE_STYLES["heading_1"]
        paragraph = next((p for p in doc.paragraphs if p.style.name in candidate_styles and p.text.strip()), None)
        if paragraph is None: raise ValueError("body-start selector 'first_heading_1' matched no paragraph")
        return paragraph
    if strategy == "heading_text":
        pattern = selector.get("heading_text_pattern")
        if not pattern: raise ValueError("body selector heading_text requires heading_text_pattern")
        try: paragraph = next((p for p in doc.paragraphs if re.search(pattern, p.text.strip())), None)
        except re.error as exc: raise ValueError(f"invalid body heading_text_pattern: {exc}") from exc
        if paragraph is None: raise ValueError("body-start selector 'heading_text' matched no paragraph")
        return paragraph
    raise ValueError(f"unsupported semantic body-start strategy: {strategy!r}")


def _paragraph_starts_section(doc: Document, paragraph) -> bool:
    target = paragraph._p; substantive_since_boundary = False
    for child in list(doc._body._element):
        if child is target: return not substantive_since_boundary
        if child.tag == qn("w:p"):
            ppr = child.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                substantive_since_boundary = False; continue
            if "".join(child.itertext()).strip() or child.xpath(".//w:drawing"):
                substantive_since_boundary = True
        elif child.tag == qn("w:tbl"):
            substantive_since_boundary = True
    raise ValueError("body-start selector matched a paragraph outside the document body")


def _insert_section_before(doc: Document, paragraph) -> None:
    body = doc._body._element; children = list(body); target_index = children.index(paragraph._p)
    previous = next((node for node in reversed(children[:target_index]) if node.tag == qn("w:p")), None)
    if previous is None:
        previous = OxmlElement("w:p"); body.insert(target_index, previous)
    ppr = previous.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr"); previous.insert(0, ppr)
    if ppr.find(qn("w:sectPr")) is not None: return
    current = list(doc.sections)[_paragraph_section_indices(doc).get(paragraph._p, 0)]._sectPr
    # Move the current section properties to the preceding paragraph.  Copying
    # relationship ids would make both sections share the same header/footer
    # stories and prevents independent front/body numbering.
    sectpr = copy.deepcopy(current); ppr.append(sectpr)
    for child in list(current):
        if child.tag not in {qn("w:pgSz"), qn("w:pgMar"), qn("w:cols"), qn("w:docGrid")}:
            current.remove(child)
    section_type = sectpr.find(qn("w:type"))
    if section_type is None:
        section_type = OxmlElement("w:type"); sectpr.insert(0, section_type)
    section_type.set(qn("w:val"), "nextPage")


def ensure_page_number_sections(doc: Document, page_num: dict[str, Any], mappings: dict[str, Any]) -> bool:
    front, body = page_num.get("front_matter_format"), page_num.get("body_format")
    if not front or not body or front == body: return False
    selector = page_num.get("body_start_selector")
    if not selector: raise ValueError("front/body page-number formats differ but body_start_selector is missing")
    if selector.get("strategy") == "section_index":
        resolve_body_start_section(doc, page_num, mappings); return False
    paragraph = _find_body_start_paragraph(doc, selector, mappings)
    if _paragraph_starts_section(doc, paragraph): return False
    if page_num.get("section_boundary_policy", "create_if_missing") == "require_existing":
        raise ValueError("body-start paragraph is not at an existing section boundary and automatic section creation is disabled")
    _insert_section_before(doc, paragraph); return True


def resolve_body_start_section(doc: Document, page_num: dict[str, Any], mappings: dict[str, Any]) -> int:
    sections = list(doc.sections)
    selector = page_num.get("body_start_selector")
    if not selector:
        front, body = page_num.get("front_matter_format"), page_num.get("body_format")
        if front and body and front != body:
            raise ValueError("front/body page-number formats differ but body_start_selector is missing")
        return 0
    strategy = selector.get("strategy")
    if strategy == "section_index":
        index = int(selector.get("section_index", 0)) - 1
        if 0 <= index < len(sections): return index
        raise ValueError(f"body section_index is outside 1..{len(sections)}")
    paragraph = _find_body_start_paragraph(doc, selector, mappings)
    return min(_paragraph_section_indices(doc).get(paragraph._p, 0), len(sections) - 1)


def _unlink_story_preserving_content(story, force_clone: bool = False) -> None:
    """Compatibility wrapper around the shared story adapter."""
    unlink_story_preserving_content(story, force_clone=force_clone)


def _footer_variants(doc: Document, section, page: dict[str, Any]) -> list[tuple[str, Any, bool]]:
    """Return all footer stories and whether each one is active for rendering."""
    first_active = bool(page.get("different_first_page", section.different_first_page_header_footer))
    even_active = bool(page.get("different_odd_even", doc.settings.odd_and_even_pages_header_footer))
    return [
        ("default", section.footer, True),
        ("first", section.first_page_footer, first_active),
        ("even", section.even_page_footer, even_active),
    ]


def _clear_paragraph_content(paragraph: Paragraph) -> None:
    """Clear runs/fields while retaining pPr, which carries style and borders."""
    for node in list(paragraph._p):
        if node.tag != qn("w:pPr"):
            paragraph._p.remove(node)


def _add_text_run(paragraph: Paragraph, text: str) -> None:
    if text:
        paragraph.add_run(text)


def _add_styleref_field(paragraph: Paragraph, style_name: str) -> None:
    """Add a dynamic STYLEREF field using the actual mapped heading style name."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f'STYLEREF "{style_name}" \\* MERGEFORMAT')
    run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = style_name
    run.append(text); field.append(run); paragraph._p.append(field)


def _set_header_tabs(paragraph: Paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:tabs"))
    if old is not None: ppr.remove(old)
    tabs = OxmlElement("w:tabs")
    for align, pos in (("center", "4680"), ("right", "9360")):
        tab = OxmlElement("w:tab"); tab.set(qn("w:val"), align); tab.set(qn("w:pos"), pos); tabs.append(tab)
    ppr.append(tabs)


def _set_bottom_border(paragraph: Paragraph, border: dict[str, Any]) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
    else:
        # Re-append after tabs so repeated application has canonical child order.
        ppr.remove(borders)
    for old in list(borders):
        if old.tag == qn("w:bottom"): borders.remove(old)
    bottom = OxmlElement("w:bottom")
    ooxml_style = {"single": "single", "thin_thick": "thinThickMediumGap",
                   "thick_thin": "thickThinMediumGap"}[border["style"]]
    bottom.set(qn("w:val"), ooxml_style)
    bottom.set(qn("w:sz"), str(round(float(border["width_pt"]) * 8)))
    bottom.set(qn("w:space"), str(round(float(border.get("space_pt", 0)))))
    bottom.set(qn("w:color"), border.get("color", "000000").upper())
    borders.append(bottom)
    ppr.append(borders)


def apply_header_content(paragraph: Paragraph, header_spec: dict[str, Any], mappings: dict[str, Any]) -> None:
    content = header_spec.get("header_content")
    if content:
        _clear_paragraph_content(paragraph)
        _set_header_tabs(paragraph)
        _add_text_run(paragraph, str(content.get("left_text", "")))
        if content.get("right_field") == "styleref_heading_1":
            heading_style = mappings.get("heading_1", {}).get("style_name")
            if not heading_style:
                raise ValueError("header STYLEREF requires a heading_1 style mapping")
            _add_text_run(paragraph, "\t")
            _add_styleref_field(paragraph, heading_style)
    elif header_spec.get("text") is not None:
        paragraph.text = str(header_spec["text"])
    if header_spec.get("bottom_border"):
        _set_bottom_border(paragraph, header_spec["bottom_border"])


def page_field_count(story) -> int:
    return sum(len(p._p.xpath('.//w:fldSimple[contains(translate(@w:instr,"page","PAGE"),"PAGE")]'))
               + len(p._p.xpath('.//w:instrText[contains(translate(text(),"page","PAGE"),"PAGE")]'))
               for p in all_story_paragraphs(story))


def expected_page_locations(doc: Document, page: dict[str, Any], body_start: int) -> dict[tuple[int, str], bool]:
    """Resolve where PAGE fields belong, preserving established template locations by default.

    If a numbering zone has no existing PAGE fields, populate all active footer variants
    in that zone so a plain document can still receive page numbers.
    """
    page_num = page.get("page_number", {})
    raw: dict[tuple[int, str], bool] = {}
    active: dict[tuple[int, str], bool] = {}
    for index, section in enumerate(doc.sections):
        for footer_type, story, is_active in _footer_variants(doc, section, page):
            active[index, footer_type] = is_active
            raw[index, footer_type] = is_active and page_field_count(story) > 0
    preserve = page_num.get("preserve_existing_locations")
    if preserve is False:
        return active
    if preserve is True:
        # Preserve optional first/even placement, but an explicit numbering rule
        # must always win over a historically missing or malformed default
        # footer.  Ordinary pages render through the default story, so every
        # numbered section needs one independently verifiable PAGE field there.
        resolved = dict(raw)
        for index in range(len(doc.sections)):
            fmt = (page_num.get("front_matter_format") if index < body_start
                   else page_num.get("body_format", "decimal"))
            if fmt and fmt != "none":
                resolved[index, "default"] = True
        return resolved
    resolved = dict(raw)
    for start, end in ((0, body_start), (body_start, len(doc.sections))):
        if not any(raw.get((i, kind), False) for i in range(start, end) for kind in ("default", "first", "even")):
            for i in range(start, end):
                for kind in ("default", "first", "even"):
                    resolved[i, kind] = active.get((i, kind), False)
    return resolved


def apply_headers_footers(doc: Document, roles: dict[str, Any], page: dict[str, Any],
                          mappings: dict[str, Any], locations: dict[tuple[int, str], bool]) -> dict[str, int]:
    if page.get("page_number"):
        raise ValueError("pagination must be applied through execute_section_plan")
    counts = {"headers": 0, "footers": 0, "page_numbers": 0, "page_fields_removed": 0}
    header_spec, footer_spec = roles.get("header", {}), roles.get("footer", {})
    seen_headers = set(); seen_footers = set(); sections = list(doc.sections)
    if header_spec:
        for index, section in enumerate(sections):
            # Every later section gets an explicit independent story.  Object-id
            # based shared-part detection is not reliable across python-docx
            # relationship mutations, especially for the last shared section.
            _unlink_story_preserving_content(section.header, force_clone=index > 0)
    for index, section in enumerate(sections):
        if header_spec and id(section.header._element) not in seen_headers:
            seen_headers.add(id(section.header._element))
            p = section.header.paragraphs[0]
            if "Header" in style_names(doc): p.style = "Header"
            apply_header_content(p, header_spec, mappings)
            apply_direct_format(p, header_spec); counts["headers"] += 1
        if footer_spec:
            variants = _footer_variants(doc, section, page)
            for footer_type, story, active in variants:
                marker = id(story._element)
                if marker in seen_footers: continue
                seen_footers.add(marker)
                p = story.paragraphs[0] if story.paragraphs else story.add_paragraph()
                if "Footer" in style_names(doc): p.style = "Footer"
                if footer_spec.get("text") is not None: p.text = str(footer_spec["text"])
                if footer_spec: apply_direct_format(p, footer_spec)
                counts["footers"] += 1
    normalize_section_property_order(doc)
    return counts


def audit_page_numbering(docx_path: Path, page: dict[str, Any], expected_body_start: int,
                         locations: dict[tuple[int, str], bool]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    page_num = page.get("page_number", {})
    if not page_num: return findings
    with zipfile.ZipFile(docx_path) as zf:
        names = set(zf.namelist()); rels = {}
        rel_path = "word/_rels/document.xml.rels"
        if rel_path in names:
            from lxml import etree
            rr = etree.fromstring(zf.read(rel_path))
            rels = {}
            for r in rr:
                target = r.get("Target", "")
                rels[r.get("Id")] = target.lstrip("/") if target.startswith("/word/") else "word/" + target.lstrip("/")
        from lxml import etree
        document = etree.fromstring(zf.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        sectprs = document.xpath("//w:sectPr", namespaces=ns)
        expected_map = {"roman": "lowerRoman", "roman_upper": "upperRoman", "decimal": "decimal"}
        front_reset = next((i for i in range(expected_body_start)
                            if any(locations.get((i, kind), False) for kind in ("default", "first", "even"))), None)
        for index, sectpr in enumerate(sectprs):
            fmt = page_num.get("front_matter_format") if index < expected_body_start else page_num.get("body_format", "decimal")
            expected_fmt = expected_map.get(fmt, fmt)
            pg = sectpr.find("w:pgNumType", namespaces=ns)
            actual_fmt = pg.get(qn("w:fmt")) if pg is not None else None
            reset_here = index == expected_body_start or (index < expected_body_start and index == front_reset)
            start_value = page_num.get("front_matter_start", 1) if index < expected_body_start else page_num.get("body_start", 1)
            expected_start = int(start_value) if reset_here else None
            actual_start = int(pg.get(qn("w:start"))) if pg is not None and pg.get(qn("w:start")) else None
            if fmt != "none" and (actual_fmt != expected_fmt or actual_start != expected_start):
                findings.append({"role": "page", "property": f"sections[{index+1}].page_number",
                                 "template_value": {"format": actual_fmt, "start": actual_start},
                                 "required_value": {"format": expected_fmt, "start": expected_start}})
            refs = sectpr.findall("w:footerReference", namespaces=ns)
            refs_by_type = {r.get(qn("w:type"), "default"): r for r in refs}
            title_page = sectpr.find("w:titlePg", namespaces=ns) is not None
            settings = etree.fromstring(zf.read("word/settings.xml")) if "word/settings.xml" in names else None
            odd_even = settings is not None and settings.find("w:evenAndOddHeaders", namespaces=ns) is not None
            active_types = {"default"}
            if title_page: active_types.add("first")
            if odd_even: active_types.add("even")
            for footer_type in ("default", "first", "even"):
                ref = refs_by_type.get(footer_type)
                part = rels.get(ref.get(qn("r:id"))) if ref is not None else None
                expected_count = 1 if fmt != "none" and footer_type in active_types and locations.get((index, footer_type), False) else 0
                if footer_type in active_types and (not part or part not in names):
                    findings.append({"role": "page", "property": f"sections[{index+1}].{footer_type}_footer",
                                     "template_value": None, "required_value": f"explicit {footer_type} footer with {expected_count} PAGE field"})
                    continue
                if not part or part not in names: continue
                footer = etree.fromstring(zf.read(part))
                simple = footer.xpath('.//w:fldSimple[contains(translate(@w:instr,"page","PAGE"),"PAGE")]', namespaces=ns)
                complex_instr = footer.xpath('.//w:instrText[contains(translate(text(),"page","PAGE"),"PAGE")]', namespaces=ns)
                count = len(simple) + len(complex_instr)
                if count != expected_count:
                    findings.append({"role": "page", "property": f"sections[{index+1}].{footer_type}_page_field_count",
                                     "template_value": count, "required_value": expected_count, "footer_part": part})
    return findings


def audit_headers(docx_path: Path, roles: dict[str, Any], mappings: dict[str, Any]) -> list[dict[str, Any]]:
    """Audit serialized header text, STYLEREF instruction and paragraph border."""
    header_spec = roles.get("header", {})
    if not header_spec:
        return []
    findings: list[dict[str, Any]] = []
    content = header_spec.get("header_content", {})
    border = header_spec.get("bottom_border")
    expected_style = mappings.get("heading_1", {}).get("style_name")
    expected_border = {"single": "single", "thin_thick": "thinThickMediumGap",
                       "thick_thin": "thickThinMediumGap"}.get(border.get("style")) if border else None
    with zipfile.ZipFile(docx_path) as zf:
        from lxml import etree
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        document = etree.fromstring(zf.read("word/document.xml"))
        relroot = etree.fromstring(zf.read("word/_rels/document.xml.rels"))
        rels = {r.get("Id"): r.get("Target") for r in relroot}
        settings = etree.fromstring(zf.read("word/settings.xml")) if "word/settings.xml" in zf.namelist() else None
        even_and_odd = bool(settings is not None and settings.find("w:evenAndOddHeaders", namespaces=ns) is not None)
        parts = []
        for sectpr in document.xpath("//w:sectPr", namespaces=ns):
            active_types = {"default"}
            if sectpr.find("w:titlePg", namespaces=ns) is not None:
                active_types.add("first")
            if even_and_odd:
                active_types.add("even")
            for ref in sectpr.xpath("./w:headerReference", namespaces=ns):
                header_type = ref.get(qn("w:type"), "default")
                if header_type not in active_types:
                    continue
                target = rels.get(ref.get(qn("r:id")))
                if target:
                    part = target.lstrip("/") if target.startswith("/word/") else "word/" + target.lstrip("/")
                    if part not in parts: parts.append(part)
        if not parts:
            return [{"role": "header", "property": "header_part", "template_value": None,
                     "required_value": "at least one serialized header part"}]
        for part in parts:
            root = etree.fromstring(zf.read(part))
            matches = []
            for paragraph in root.xpath(".//w:p", namespaces=ns):
                visible = "".join(paragraph.xpath(".//w:t/text()", namespaces=ns))
                instructions = [value.strip() for value in paragraph.xpath(".//w:fldSimple/@w:instr", namespaces=ns)]
                instructions += [value.strip() for value in paragraph.xpath(".//w:instrText/text()", namespaces=ns)]
                if not content or str(content.get("left_text", "")) in visible:
                    matches.append((paragraph, instructions))
            if not matches:
                findings.append({"role": "header", "property": f"{part}.left_text", "template_value": None,
                                 "required_value": content.get("left_text")})
                continue
            paragraph, instructions = matches[0]
            if content.get("right_field") == "styleref_heading_1":
                expected_instr = f'STYLEREF "{expected_style}" \\* MERGEFORMAT'
                if instructions.count(expected_instr) != 1:
                    findings.append({"role": "header", "property": f"{part}.styleref",
                                     "template_value": instructions, "required_value": [expected_instr]})
            if border:
                bottoms = paragraph.xpath("./w:pPr/w:pBdr/w:bottom", namespaces=ns)
                actual = None if not bottoms else {
                    "style": bottoms[0].get(qn("w:val")),
                    "width_eighths": int(bottoms[0].get(qn("w:sz"), "0")),
                    "color": bottoms[0].get(qn("w:color")),
                }
                required = {"style": expected_border, "width_eighths": round(float(border["width_pt"]) * 8),
                            "color": border.get("color", "000000").upper()}
                if actual != required:
                    findings.append({"role": "header", "property": f"{part}.bottom_border",
                                     "template_value": actual, "required_value": required})
    return findings


def apply_page(doc: Document, page: dict[str, Any]) -> None:
    margins = page.get("margins_pt", {})
    for section in doc.sections:
        if page.get("size", "").upper() == "A4":
            section.page_width, section.page_height = Mm(210), Mm(297)
        if page.get("orientation") == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE; section.page_width, section.page_height = section.page_height, section.page_width
        elif page.get("orientation") == "portrait":
            section.orientation = WD_ORIENT.PORTRAIT
            if section.page_width > section.page_height: section.page_width, section.page_height = section.page_height, section.page_width
        for key, attr in (("top", "top_margin"), ("bottom", "bottom_margin"), ("left", "left_margin"),
                          ("right", "right_margin"), ("header", "header_distance"), ("footer", "footer_distance"), ("gutter", "gutter")):
            if key in margins: setattr(section, attr, Pt(float(margins[key])))
        section.different_first_page_header_footer = bool(page.get("different_first_page", section.different_first_page_header_footer))
    settings = doc.settings.element
    even = settings.find(qn("w:evenAndOddHeaders"))
    if page.get("different_odd_even") and even is None:
        from docx.oxml import OxmlElement
        settings.append(OxmlElement("w:evenAndOddHeaders"))
    elif page.get("different_odd_even") is False and even is not None: settings.remove(even)


def paragraph_has_drawing(p) -> bool:
    return has_drawing(p)


def paragraph_has_role_content(paragraph) -> bool:
    """Treat visible text and non-text Word objects as coverage-bearing content."""
    return has_visible_or_object_content(paragraph)


def _normalize_content_instance_text(value: Any) -> str:
    """Normalize only presentation whitespace for literal-content matching."""
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"\s+", "", text).strip()


def _content_instance_paragraphs(doc: Document) -> list[Paragraph]:
    """Return each body/header/footer paragraph once for instance matching."""
    paragraphs: list[Paragraph] = []
    seen: set[Any] = set()
    stories = [all_body_paragraphs(doc)]
    for section in doc.sections:
        stories.extend([
            all_story_paragraphs(section.header),
            all_story_paragraphs(section.first_page_header),
            all_story_paragraphs(section.even_page_header),
            all_story_paragraphs(section.footer),
            all_story_paragraphs(section.first_page_footer),
            all_story_paragraphs(section.even_page_footer),
        ])
    for story in stories:
        for paragraph in story:
            marker = paragraph._p
            if marker in seen:
                continue
            seen.add(marker)
            paragraphs.append(paragraph)
    return paragraphs


def apply_content_instance_overrides(
    doc: Document, spec: dict[str, Any], mappings: dict[str, Any]
) -> list[dict[str, Any]]:
    """Apply per-instance style overrides without merging literal text into roles.

    ``roles`` supplies the shared/default style.  A content instance may add a
    narrower font/paragraph override for the exact literal text it identifies.
    Missing literals are reported as an auditable input/content gap, but are
    not synthesized: arbitrary labels and values must remain school-specific
    cover data rather than guessed DOCX content.
    """
    raw_instances = spec.get("content_instances", spec.get("cover_field_instances", []))
    if not isinstance(raw_instances, list) or not raw_instances:
        return []
    paragraphs = _content_instance_paragraphs(doc)
    records: list[dict[str, Any]] = []
    for instance in raw_instances:
        if not isinstance(instance, dict):
            continue
        instance_id = str(instance.get("id") or "")
        role = str(instance.get("role") or "")
        text = str(instance.get("text") or "")
        target = _normalize_content_instance_text(text)
        mapping = mappings.get(role) if isinstance(mappings, dict) else None
        style_properties = instance.get("style_properties")
        if not isinstance(style_properties, dict):
            style_properties = {}
        matches = [paragraph for paragraph in paragraphs
                   if target and _normalize_content_instance_text(paragraph.text) == target]
        applied = 0
        for paragraph in matches:
            if mapping and mapping.get("style_name") and paragraph.style.name != mapping["style_name"]:
                paragraph.style = mapping["style_name"]
            if style_properties:
                apply_direct_format(paragraph, style_properties)
            applied += 1
        if not matches:
            status = "not_present"
        elif style_properties:
            status = "applied"
        else:
            status = "shared_style_present"
        records.append({
            "instance_id": instance_id,
            "field_key": instance.get("field_key"),
            "role": role,
            "text": text,
            "style_properties": copy.deepcopy(style_properties),
            "match_count": len(matches),
            "applied_count": applied,
            "status": status,
            "reason": ("literal content was matched and its instance style was applied"
                       if status == "applied" else
                       "literal content exists with the shared role style"
                       if status == "shared_style_present" else
                       "literal content was not present in the generated/source DOCX; no text was synthesized"),
        })
    return records


def role_coverage(doc: Document, roles: dict[str, Any], mappings: dict[str, Any]) -> list[dict[str, Any]]:
    """Return an explicit required/present vs optional/not-applicable role matrix.

    Every backend-supported role declared by the format specification is
    required. Supported roles not declared by the specification are retained
    in the report as optional/not-applicable so a zero is never confused with
    successful coverage.
    """
    body = list(all_body_paragraphs(doc))
    stories = {
        "header": [p for section in doc.sections for story in
                   (section.header, section.first_page_header, section.even_page_header)
                   for p in all_story_paragraphs(story)],
        "footer": [p for section in doc.sections for story in
                   (section.footer, section.first_page_footer, section.even_page_footer)
                   for p in all_story_paragraphs(story)],
    }
    coverage = []
    for role in dict.fromkeys([*ROLE_STYLES, *roles]):
        declared = role in roles
        mapping = mappings.get(role)
        if not declared:
            coverage.append({"role": role, "expectation": "optional", "status": "not_applicable",
                             "present": False, "count": 0,
                             "reason": "the role is not declared by this format specification"})
            continue
        if mapping is None:
            coverage.append({"role": role, "expectation": "optional", "status": "not_applicable",
                             "present": False, "count": 0,
                             "reason": "the declared role is not executable by this backend"})
            continue
        paragraphs = stories.get(role, body)
        if role in {"header", "footer"}:
            # Header/footer content can be entirely fields; a styled story
            # paragraph is executable coverage even when visible text is empty.
            count = sum(1 for paragraph in paragraphs if paragraph.style.name == mapping["style_name"])
        elif role == "toc":
            # The TOC heading is structural chrome and is covered by toc_title.
            # Coverage for toc must come from the actual field paragraph; an
            # empty field is still executable content before Word refreshes its
            # cached entries.
            count = sum(1 for paragraph in paragraphs if has_toc_field(paragraph))
            # Preserve coverage for legacy/static documents that represent the
            # TOC as a styled placeholder rather than a Word field.  A real
            # field is authoritative and need not inherit the template's
            # heading style (converters commonly emit it as Normal).
            if not count:
                count = sum(1 for paragraph in paragraphs
                            if paragraph.style.name == mapping["style_name"]
                            and paragraph_has_role_content(paragraph))
        else:
            count = sum(1 for paragraph in paragraphs
                        if paragraph.style.name == mapping["style_name"] and paragraph_has_role_content(paragraph))
        detector = structural_detector(role)
        structural_count = sum(1 for paragraph in paragraphs if matches_structural_detector(paragraph, detector))
        failure_type = None
        evidence = []
        # Some rules (for example a global Latin font) legitimately declare
        # formatting for conditional roles such as footnotes, bibliography or
        # cover fields even when the particular thesis contains none.  The
        # registry distinguishes universally required content from conditional
        # content.  Conditional roles become required only when their structure
        # is actually observed in this document.
        required = bool(role_config(role).get("required_content")) or bool(structural_count)
        if not required and not count:
            coverage.append({"role": role, "expectation": "optional", "status": "not_applicable",
                             "present": False, "count": 0, "style": mapping["style_name"],
                             "structural_count": structural_count, "failure_type": None,
                             "evidence": [],
                             "reason": "the conditional role is declared but no matching content exists"})
            continue
        if not count:
            if structural_count:
                failure_type = "role_mapping_missing"
                evidence.append(f"{structural_count} paragraphs matched structural detector {detector}")
            else:
                failure_type = "content_missing"
        coverage.append({"role": role, "expectation": "required",
                         "status": "present" if count else "missing", "present": bool(count),
                         "count": count, "style": mapping["style_name"],
                         "structural_count": structural_count,
                         "failure_type": failure_type, "evidence": evidence})
    return coverage


def _drawing_relationship_ids(element: Any) -> list[str]:
    return [node.get(qn("r:embed")) for node in element.iter(qn("a:blip")) if node.get(qn("r:embed"))]


def _caption_number(text: str) -> str | None:
    match = re.match(r"^(?:图|表|Figure|Table)\s*([0-9]+(?:[.．-][0-9]+)*)", text.strip(), re.I)
    return match.group(1).replace("．", ".").replace("-", ".") if match else None


def reposition_captions(doc: Document, role: str, style_name: str, position: str) -> tuple[int, list[dict[str, Any]], list[dict[str, Any]]]:
    """Bind and reposition captions using an explicit local object graph.

    Empty layout paragraphs are transparent.  A caption is moved only when a
    unique object group exists on one side; competing groups are retained as
    an auditable ambiguity rather than guessed.
    """
    moved = 0; issues = []; graph = []
    body = doc._body._element
    children = list(body)
    for i, child in enumerate(children):
        if child.tag != qn("w:p"): continue
        style = child.find(qn("w:pPr")); sid = None
        if style is not None:
            ps = style.find(qn("w:pStyle")); sid = ps.get(qn("w:val")) if ps is not None else None
        para = next((p for p in doc.paragraphs if p._p is child), None)
        if para is None or para.style.name != style_name: continue
        is_target = (lambda x: x.tag == qn("w:p") and bool(x.xpath(".//w:drawing"))) if role == "figure_caption" else (lambda x: x.tag == qn("w:tbl"))
        is_blank = lambda x: x.tag == qn("w:p") and not "".join(x.itertext()).strip() and not x.xpath(
            ".//w:drawing|.//m:oMath|.//w:br|.//w:sectPr")

        def group_on_side(step: int) -> tuple[list[Any], list[int]]:
            j = i + step; blanks = []
            while 0 <= j < len(children) and is_blank(children[j]) and len(blanks) < 3:
                blanks.append(j); j += step
            if not (0 <= j < len(children)) or not is_target(children[j]):
                return [], blanks
            group = [children[j]]
            # Consecutive image paragraphs form one multi-panel/combined figure.
            # Consecutive tables are retained too, but rejected below because a
            # single table caption cannot be bound to either one deterministically.
            k = j + step
            while 0 <= k < len(children) and is_target(children[k]):
                group.append(children[k]); k += step
            return sorted(group, key=lambda node: children.index(node)), blanks

        before, before_blanks = group_on_side(-1)
        after, after_blanks = group_on_side(1)
        candidates = [group for group in (before, after) if group]
        number = _caption_number(para.text)
        references = []
        if number:
            reference_re = re.compile(rf"(?:图|表|Figure|Table)\s*{re.escape(number)}(?:\D|$)", re.I)
            references = [index for index, p in enumerate(doc.paragraphs)
                          if p._p is not child and reference_re.search(p.text)]
        graph_item = {
            "role": role, "caption_node_id": f"body-child-{i}", "caption": para.text[:240],
            "caption_number": number, "requested_position": position,
            "before": [{"node_id": f"body-child-{children.index(node)}",
                        "relationship_ids": _drawing_relationship_ids(node)} for node in before],
            "after": [{"node_id": f"body-child-{children.index(node)}",
                       "relationship_ids": _drawing_relationship_ids(node)} for node in after],
            "blank_paragraphs_before": before_blanks, "blank_paragraphs_after": after_blanks,
            "references_in_text": references,
        }
        if len(candidates) != 1 or (role == "table_caption" and candidates and len(candidates[0]) != 1):
            graph_item.update({"status": "ambiguous", "confidence": 0.0})
            graph.append(graph_item)
            issues.append({"role": role, "property": "position", "failure_type": "needs_clarification",
                           "reason": "caption/object graph has zero or competing candidate groups; paragraph was not moved",
                           "caption_node_id": graph_item["caption_node_id"], "text": para.text[:120]})
            continue
        target_group = candidates[0]
        graph_item.update({"status": "bound", "confidence": 1.0,
                           "selected": [f"body-child-{children.index(node)}" for node in target_group]})
        graph.append(graph_item)
        first_target, last_target = target_group[0], target_group[-1]
        desired_anchor = last_target if position == "below" else first_target
        desired = children.index(desired_anchor) + (1 if position == "below" else 0)
        already = (position == "below" and i == children.index(last_target) + 1) or (
            position == "above" and i == children.index(first_target) - 1)
        if i != desired and not already:
            body.remove(child); idx_now = list(body).index(desired_anchor)
            body.insert(idx_now + (1 if position == "below" else 0), child); moved += 1
            children = list(body)
    return moved, issues, graph


def _border_node(container, side: str, width_pt: float | None) -> None:
    old = container.find(qn(f"w:{side}"))
    if old is not None: container.remove(old)
    node = OxmlElement(f"w:{side}")
    if width_pt is None or float(width_pt) <= 0:
        node.set(qn("w:val"), "nil")
    else:
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(round(float(width_pt) * 8)))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
    container.append(node)


def apply_table_rules(doc: Document, table_spec: dict[str, Any]) -> dict[str, int]:
    """Apply deterministic three-line-table and pagination controls."""
    counts = {"tables": 0, "rows_no_split": 0, "header_rows_repeated": 0,
              "explicit_borders_applied": 0, "continuation_contracts": 0}
    if not table_spec: return counts
    explicit = table_spec.get("border_widths_pt") if isinstance(table_spec.get("border_widths_pt"), dict) else {}
    continuation = table_spec.get("continuation") if isinstance(table_spec.get("continuation"), dict) else {}
    border_map = {
        "top": "top", "bottom": "bottom", "left": "left", "right": "right",
        "inside_h": "insideH", "inside_v": "insideV",
    }
    for table in doc.tables:
        counts["tables"] += 1
        if table_spec.get("style") == "three_line" or explicit:
            tblpr = table._tbl.tblPr
            borders = tblpr.find(qn("w:tblBorders"))
            if borders is None:
                borders = OxmlElement("w:tblBorders"); tblpr.append(borders)
            if table_spec.get("style") == "three_line":
                _border_node(borders, "top", float(table_spec.get("top_border_pt", 1.5)))
                _border_node(borders, "bottom", float(table_spec.get("bottom_border_pt", 1.5)))
                for side in ("left", "right", "insideV"):
                    _border_node(borders, side, None if table_spec.get("remove_vertical_borders", True) else 0.5)
                _border_node(borders, "insideH", None)
            for key, side in border_map.items():
                if key in explicit:
                    _border_node(borders, side, float(explicit[key]))
                    counts["explicit_borders_applied"] += 1
            header_width = explicit.get("header", table_spec.get("header_border_pt", .75))
            if table.rows:
                for cell in table.rows[0].cells:
                    tcpr = cell._tc.get_or_add_tcPr()
                    cell_borders = tcpr.find(qn("w:tcBorders"))
                    if cell_borders is None:
                        cell_borders = OxmlElement("w:tcBorders"); tcpr.append(cell_borders)
                    if table_spec.get("style") == "three_line" or "header" in explicit:
                        _border_node(cell_borders, "bottom", float(header_width))
                        if "header" in explicit:
                            counts["explicit_borders_applied"] += 1
        for index, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            cant = trpr.find(qn("w:cantSplit"))
            if table_spec.get("allow_row_split") is False:
                if cant is None: trpr.append(OxmlElement("w:cantSplit"))
                counts["rows_no_split"] += 1
            elif cant is not None:
                trpr.remove(cant)
            repeat_header = bool(table_spec.get("repeat_header_row") or continuation.get("repeat_header_row"))
            if index == 0 and repeat_header:
                header = trpr.find(qn("w:tblHeader"))
                if header is None: trpr.append(OxmlElement("w:tblHeader"))
                counts["header_rows_repeated"] += 1
        if continuation:
            # The continuation contract is intentionally declarative.  Word
            # decides whether a table crosses a page, so the generator sets
            # the repeat-header invariant but never fabricates a second
            # caption or a guessed page break.
            counts["continuation_contracts"] += 1
    return counts


def apply_object_pagination(doc: Document, mappings: dict[str, Any], object_spec: dict[str, Any],
                            table_spec: dict[str, Any]) -> dict[str, int]:
    counts = {"figures_kept_with_captions": 0, "table_captions_kept_with_tables": 0}
    children = list(doc._body._element)
    paragraphs = {p._p: p for p in doc.paragraphs}
    figure_style = mappings.get("figure_caption", {}).get("style_name")
    table_style = mappings.get("table_caption", {}).get("style_name")
    for i, child in enumerate(children):
        if child.tag == qn("w:p") and child.xpath(".//w:drawing") and object_spec.get("keep_figure_with_caption"):
            next_node = children[i + 1] if i + 1 < len(children) else None
            next_para = paragraphs.get(next_node)
            if next_para is not None and figure_style and next_para.style.name == figure_style:
                ppr = child.get_or_add_pPr()
                if ppr.find(qn("w:keepNext")) is None: ppr.append(OxmlElement("w:keepNext"))
                counts["figures_kept_with_captions"] += 1
        if child.tag == qn("w:p") and table_style and object_spec.get("keep_table_with_caption"):
            para = paragraphs.get(child)
            next_node = children[i + 1] if i + 1 < len(children) else None
            if para is not None and para.style.name == table_style and next_node is not None and next_node.tag == qn("w:tbl"):
                ppr = child.get_or_add_pPr()
                if ppr.find(qn("w:keepNext")) is None: ppr.append(OxmlElement("w:keepNext"))
                counts["table_captions_kept_with_tables"] += 1
    return counts


def _is_structural_nonprose(paragraph: Paragraph, mappings: dict[str, Any]) -> bool:
    """Return whether a paragraph is a heading/caption rather than explanation."""
    for role in ("figure_caption", "table_caption", "heading_1", "heading_2", "heading_3",
                 "heading_4", "heading_acknowledgments", "heading_appendix",
                 "heading_conclusion", "heading_publications", "heading_references"):
        mapping = mappings.get(role, {})
        if mapping.get("style_name") == paragraph.style.name:
            return True
        detector = structural_detector(role)
        if detector and matches_structural_detector(paragraph, detector):
            return True
    return False


def audit_object_order_constraints(doc: Document, object_spec: dict[str, Any],
                                   mappings: dict[str, Any]) -> list[dict[str, Any]]:
    """Audit explicit prose-before-object contracts without inventing prose.

    The contract is deliberately structural: a non-empty prose paragraph must
    occur immediately before the object, ignoring blank layout paragraphs.  A
    caption, heading, another object, or missing paragraph is not silently
    reinterpreted as explanatory prose.
    """
    constraints = object_spec.get("order_constraints", []) if isinstance(object_spec, dict) else []
    if not isinstance(constraints, list):
        return []
    body_children = list(doc._body._element)
    paragraphs = {paragraph._p: paragraph for paragraph in doc.paragraphs}
    findings: list[dict[str, Any]] = []
    for ci, constraint in enumerate(constraints, 1):
        if not isinstance(constraint, dict) or constraint.get("preceding_prose_required") is not True:
            continue
        object_type = constraint.get("object_type")
        for target_index, child in enumerate(body_children):
            is_figure = child.tag == qn("w:p") and bool(child.xpath(".//w:drawing"))
            is_table = child.tag == qn("w:tbl")
            if (object_type == "figure" and not is_figure) or (object_type == "table" and not is_table):
                continue
            previous_text = None
            previous_kind = "missing"
            j = target_index - 1
            while j >= 0:
                previous = body_children[j]
                if previous.tag == qn("w:sectPr"):
                    j -= 1
                    continue
                if previous.tag == qn("w:tbl") or (previous.tag == qn("w:p") and previous.xpath(".//w:drawing")):
                    previous_kind = "object"
                    break
                paragraph = paragraphs.get(previous)
                if paragraph is None or not paragraph.text.strip():
                    j -= 1
                    continue
                previous_text = paragraph.text[:240]
                previous_kind = "structural" if _is_structural_nonprose(paragraph, mappings) else "prose"
                break
            if previous_kind != "prose":
                findings.append({
                    "role": "objects",
                    "property": f"order_constraints[{ci}].{object_type}.preceding_prose",
                    "template_value": {"previous_kind": previous_kind, "previous_text": previous_text},
                    "required_value": "non-empty explanatory prose paragraph immediately before object",
                    "failure_type": "object_order_violation",
                    "reason": "the document does not contain an unambiguous prose paragraph before this object; no text was synthesized",
                })
    return findings


def audit_table_rules(doc: Document, table_spec: dict[str, Any],
                      render_report: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    findings = []
    if not table_spec: return findings
    explicit = table_spec.get("border_widths_pt") if isinstance(table_spec.get("border_widths_pt"), dict) else {}
    continuation = table_spec.get("continuation") if isinstance(table_spec.get("continuation"), dict) else {}
    border_map = {
        "top": "top", "bottom": "bottom", "left": "left", "right": "right",
        "inside_h": "insideH", "inside_v": "insideV",
    }
    for ti, table in enumerate(doc.tables, 1):
        tblpr = table._tbl.tblPr; borders = tblpr.find(qn("w:tblBorders"))
        def border(side: str):
            return borders.find(qn(f"w:{side}")) if borders is not None else None
        def actual_width(node):
            if node is None or node.get(qn("w:val")) in {"nil", "none"}:
                return 0.0
            raw = node.get(qn("w:sz"))
            return round(float(raw) / 8, 3) if raw else None
        def border_matches(node, expected):
            actual = actual_width(node)
            return ((expected == 0 and actual == 0)
                    or (expected > 0 and actual is not None and abs(actual - float(expected)) <= .01
                        and node is not None and node.get(qn("w:val")) == "single"))
        if table_spec.get("style") == "three_line":
            expected_sides = {
                "top": explicit.get("top", table_spec.get("top_border_pt", 1.5)),
                "bottom": explicit.get("bottom", table_spec.get("bottom_border_pt", 1.5)),
            }
            if table_spec.get("remove_vertical_borders", True):
                expected_sides.update({
                    "left": explicit.get("left", 0),
                    "right": explicit.get("right", 0),
                    "insideV": explicit.get("inside_v", 0),
                })
            for side, expected in expected_sides.items():
                node = border(side); actual = actual_width(node)
                if not border_matches(node, float(expected)):
                    findings.append({"role": "table", "property": f"tables[{ti}].{side}_border_pt",
                                     "template_value": actual, "required_value": expected})
            if "inside_h" in explicit:
                node = border("insideH")
                if not border_matches(node, float(explicit["inside_h"])):
                    findings.append({"role": "table", "property": f"tables[{ti}].insideH_border_pt",
                                     "template_value": actual_width(node), "required_value": explicit["inside_h"]})
        elif explicit:
            for key, side in border_map.items():
                if key not in explicit:
                    continue
                node = border(side)
                if not border_matches(node, float(explicit[key])):
                    findings.append({"role": "table", "property": f"tables[{ti}].{key}_border_pt",
                                     "template_value": actual_width(node), "required_value": explicit[key]})
        if table.rows and (table_spec.get("style") == "three_line" or "header" in explicit):
            expected_header = float(explicit.get("header", table_spec.get("header_border_pt", .75)))
            for ci, cell in enumerate(table.rows[0].cells, 1):
                tcpr = cell._tc.get_or_add_tcPr(); cell_borders = tcpr.find(qn("w:tcBorders"))
                node = cell_borders.find(qn("w:bottom")) if cell_borders is not None else None
                if not border_matches(node, expected_header):
                    findings.append({"role": "table", "property": f"tables[{ti}].header_cells[{ci}].bottom_border_pt",
                                     "template_value": actual_width(node), "required_value": expected_header})
        for ri, row in enumerate(table.rows, 1):
            trpr = row._tr.get_or_add_trPr()
            if table_spec.get("allow_row_split") is False and trpr.find(qn("w:cantSplit")) is None:
                findings.append({"role": "table", "property": f"tables[{ti}].rows[{ri}].allow_split", "template_value": True, "required_value": False})
        repeat_header = bool(table_spec.get("repeat_header_row") or continuation.get("repeat_header_row"))
        if table.rows and repeat_header:
            if table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is None:
                findings.append({"role": "table", "property": f"tables[{ti}].repeat_header", "template_value": False, "required_value": True})
    if continuation:
        evidence = render_report.get("table_continuation") if isinstance(render_report, dict) else None
        if not isinstance(evidence, dict) or evidence.get("verified") is not True:
            findings.append({
                "role": "table", "property": "tables.continuation.render_verification",
                "template_value": evidence.get("verified") if isinstance(evidence, dict) else False,
                "required_value": {
                    "verified": True,
                    "caption_suffix": continuation.get("caption_suffix"),
                    "verification": continuation.get("verification"),
                },
                "failure_type": "render_evidence_required",
                "reason": "cross-page continuation caption and page-span behavior cannot be proven from DOCX XML alone; accepted Word/PDF evidence is required",
            })
    return findings


def _role_paragraphs(doc: Document, role: str, mappings: dict[str, Any]) -> list[Paragraph]:
    style = mappings.get(role, {}).get("style_name")
    candidate_styles = set(style_aliases(role))
    if style: candidate_styles.add(style)
    detector = structural_detector(role)
    if detector:
        return [p for p in doc.paragraphs if matches_structural_detector(p, detector)]
    return [p for p in doc.paragraphs if p.style.name in candidate_styles]


def _keyword_values(text: str, language: str) -> list[str]:
    payload = re.sub(r"^(?:关\s*键\s*词|key\s*words?)\s*[：:]\s*", "", text.strip(), flags=re.I)
    splitter = r"[，,；;]" if language == "zh" else r"[,;]"
    return [item.strip() for item in re.split(splitter, payload) if item.strip()]


def _section_paragraphs(doc: Document, heading_role: str, mappings: dict[str, Any]) -> list[Paragraph]:
    """Return body paragraphs belonging to each occurrence of a heading role."""
    paragraphs = list(doc.paragraphs)
    positions = {paragraph._p: index for index, paragraph in enumerate(paragraphs)}
    headings = _role_paragraphs(doc, heading_role, mappings)
    if not headings:
        return []
    heading_roles = (
        "heading_1", "heading_2", "heading_3", "heading_4", "heading_acknowledgments",
        "heading_appendix", "heading_conclusion", "heading_publications", "heading_references",
    )
    all_headings = {
        paragraph._p for role in heading_roles
        for paragraph in _role_paragraphs(doc, role, mappings)
    }
    result: list[Paragraph] = []
    for heading in headings:
        start = positions.get(heading._p)
        if start is None:
            continue
        end = min((positions[node] for node in all_headings
                   if node in positions and positions[node] > start), default=len(paragraphs))
        result.extend(paragraphs[start + 1:end])
    return result


def _content_metric(text: str, metric: str | None, language: str) -> int:
    """Count content with the explicitly declared unit.

    The old checker silently treated every ``max_chars`` value as a
    whitespace-stripped code-point count.  That made a Chinese ``字`` rule,
    an English ``words`` rule, and a mixed-language Unicode rule appear
    interchangeable.  The contract now names the metric and this helper is
    the single implementation used by output audits and receipts.
    """
    normalized = re.sub(r"\s+", "", str(text or ""))
    if metric == "words":
        return len(re.findall(r"[A-Za-z0-9]+(?:['’\-][A-Za-z0-9]+)*", str(text or "")))
    if metric == "cjk_characters":
        return len(re.findall(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]", normalized))
    return len(normalized)


def _abstract_has_embedded_object(paragraphs: list[Paragraph], object_kind: str) -> bool:
    """Detect objects physically embedded in the selected abstract paragraphs."""
    if object_kind == "figures":
        return any(p._p.xpath(".//w:drawing | .//wp:inline | .//wp:anchor") for p in paragraphs)
    if object_kind == "tables":
        return any(p._p.xpath("ancestor::w:tbl") for p in paragraphs)
    return False


def _document_has_comments(doc: Document) -> bool:
    """Return whether the package contains a Word comments part."""
    return any(str(getattr(part, "partname", "")).endswith("comments.xml")
               for part in getattr(doc.part.package, "parts", []))


def _manual_semantic_finding(key: str, property_name: str, reason: str) -> dict[str, Any]:
    return {
        "role": "content_constraints",
        "property": f"{key}.{property_name}",
        "template_value": "manual_verification_required",
        "required_value": reason,
        "verification": "manual",
    }


def audit_content_constraints(doc: Document, constraints: dict[str, Any], mappings: dict[str, Any]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    positions = {p._p: i for i, p in enumerate(doc.paragraphs)}
    for key, role, language in (("abstract_zh", "abstract_body_zh", "zh"), ("abstract_en", "abstract_body_en", "en")):
        rule = constraints.get(key, {})
        paragraphs = _role_paragraphs(doc, role, mappings)
        text = "".join(p.text.strip() for p in paragraphs)
        if rule.get("required") and not text:
            findings.append({"role": "content_constraints", "property": f"{key}.required", "template_value": False, "required_value": True})
        default_metric = "words" if key == "abstract_en" else "unicode_codepoints"
        metric = rule.get("length_metric", default_metric)
        actual_length = _content_metric(text, metric, language)
        if rule.get("min_chars") is not None and actual_length < int(rule["min_chars"]):
            findings.append({"role": "content_constraints", "property": f"{key}.min_chars", "template_value": actual_length, "required_value": rule["min_chars"], "metric": metric})
        if rule.get("max_chars") is not None and actual_length > int(rule["max_chars"]):
            findings.append({"role": "content_constraints", "property": f"{key}.max_chars", "template_value": actual_length, "required_value": rule["max_chars"], "metric": metric})
        if rule.get("min_words") is not None:
            actual_words = _content_metric(text, rule.get("length_metric", "words"), language)
            if actual_words < int(rule["min_words"]):
                findings.append({"role": "content_constraints", "property": f"{key}.min_words", "template_value": actual_words, "required_value": rule["min_words"], "metric": rule.get("length_metric", "words")})
        if rule.get("max_words") is not None:
            actual_words = _content_metric(text, rule.get("length_metric", "words"), language)
            if actual_words > int(rule["max_words"]):
                findings.append({"role": "content_constraints", "property": f"{key}.max_words", "template_value": actual_words, "required_value": rule["max_words"], "metric": rule.get("length_metric", "words")})
        if rule.get("target") and rule.get("target") != key:
            findings.append({"role": "content_constraints", "property": f"{key}.target", "template_value": key, "required_value": rule["target"]})
        if key == "abstract_en" and rule.get("target") == "unresolved":
            findings.append(_manual_semantic_finding(key, "target", "abstract language target is unresolved"))
        for property_name in ("require_third_person", "required_sections", "exception_policy"):
            if rule.get(property_name):
                findings.append(_manual_semantic_finding(
                    key, property_name,
                    "semantic abstract content cannot be proven by deterministic DOCX formatting alone",
                ))
        if rule.get("prohibit_comments"):
            if _document_has_comments(doc):
                findings.append({"role": "content_constraints", "property": f"{key}.prohibit_comments", "template_value": True, "required_value": False})
        for object_kind in rule.get("prohibited_objects", []) if isinstance(rule.get("prohibited_objects"), list) else []:
            if object_kind in {"figures", "tables"} and _abstract_has_embedded_object(paragraphs, object_kind):
                findings.append({"role": "content_constraints", "property": f"{key}.prohibited_objects.{object_kind}", "template_value": True, "required_value": False})
            elif object_kind in {"chemical_equations", "nonpublic_symbols_and_terminology"}:
                findings.append(_manual_semantic_finding(
                    key, f"prohibited_objects.{object_kind}",
                    "requires semantic/manual review; deterministic text scanning must not guess domain meaning",
                ))
    acknowledgments = constraints.get("acknowledgments", {})
    if isinstance(acknowledgments, dict) and acknowledgments.get("max_chars") is not None:
        section = _section_paragraphs(doc, "heading_acknowledgments", mappings)
        text = "".join(paragraph.text.strip() for paragraph in section)
        actual_chars = len(re.sub(r"\s+", "", text))
        if section and actual_chars > int(acknowledgments["max_chars"]):
            findings.append({"role": "content_constraints", "property": "acknowledgments.max_chars",
                             "template_value": actual_chars, "required_value": acknowledgments["max_chars"]})
    keyword_counts: dict[str, int] = {}
    for key, role, language in (("keywords_zh", "keywords_zh", "zh"), ("keywords_en", "keywords_en", "en")):
        rule = constraints.get(key, {})
        paragraphs = _role_paragraphs(doc, role, mappings)
        text = " ".join(p.text.strip() for p in paragraphs)
        values = _keyword_values(text, language) if text else []
        keyword_counts[key] = len(values)
        if rule.get("required") and not values:
            findings.append({"role": "content_constraints", "property": f"{key}.required", "template_value": False, "required_value": True})
        for bound, op in (("min_count", lambda n, v: n < v), ("max_count", lambda n, v: n > v)):
            if rule.get(bound) is not None and op(len(values), int(rule[bound])):
                findings.append({"role": "content_constraints", "property": f"{key}.{bound}", "template_value": len(values), "required_value": rule[bound]})
        expected_sep = rule.get("separator")
        payload = re.sub(r"^(?:关\s*键\s*词|key\s*words?)\s*[：:]\s*", "", text, flags=re.I)
        if len(values) > 1 and expected_sep:
            ok = ((expected_sep == "chinese_comma" and "，" in payload and "," not in payload)
                  or (expected_sep == "english_comma" and "," in payload and "，" not in payload)
                  or (expected_sep == "semicolon" and ("；" in payload or ";" in payload)))
            if not ok: findings.append({"role": "content_constraints", "property": f"{key}.separator", "template_value": text, "required_value": expected_sep})
        if rule.get("max_item_chars") is not None:
            metric = rule.get("item_length_metric", "cjk_characters" if language == "zh" else "words")
            for item_index, item in enumerate(values):
                actual_item_length = _content_metric(item, metric, language)
                if actual_item_length > int(rule["max_item_chars"]):
                    findings.append({
                        "role": "content_constraints",
                        "property": f"{key}.max_item_chars[{item_index}]",
                        "template_value": actual_item_length,
                        "required_value": rule["max_item_chars"],
                        "metric": metric,
                        "item": item,
                    })
        after_role = rule.get("require_after_role")
        if after_role and paragraphs:
            before = _role_paragraphs(doc, after_role, mappings)
            if not before or max(positions.get(p._p, -1) for p in before) >= min(positions.get(p._p, 10**9) for p in paragraphs):
                findings.append({"role": "content_constraints", "property": f"{key}.require_after_role", "template_value": False, "required_value": after_role})
    if constraints.get("keywords_en", {}).get("match_other_language_count") and keyword_counts.get("keywords_en") != keyword_counts.get("keywords_zh"):
        findings.append({"role": "content_constraints", "property": "keywords_en.match_other_language_count", "template_value": keyword_counts, "required_value": "equal"})
    return findings


def resolve_profile_constraints(spec: dict[str, Any]) -> dict[str, Any]:
    constraints = copy.deepcopy(spec.get("content_constraints", {}))
    profile = spec.get("thesis_profile", {})
    conditional = spec.get("conditional_constraints", {})
    degree = profile.get("degree_level")
    if degree and isinstance(conditional.get("abstract_zh_max_chars_by_degree"), dict):
        value = conditional["abstract_zh_max_chars_by_degree"].get(degree)
        if value is not None: constraints.setdefault("abstract_zh", {})["max_chars"] = value
    if profile.get("writing_language") == "en":
        if conditional.get("require_zh_abstract_for_english_thesis"):
            constraints.setdefault("abstract_zh", {})["required"] = True
        if conditional.get("require_zh_keywords_for_english_thesis"):
            constraints.setdefault("keywords_zh", {})["required"] = True
    return constraints


def audit_document_structure(doc: Document, structure: dict[str, Any], mappings: dict[str, Any]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    positions = {p._p: i for i, p in enumerate(doc.paragraphs)}
    role_positions: dict[str, list[int]] = {}
    ordered_groups = structure.get("ordered_role_groups", []) if isinstance(structure.get("ordered_role_groups"), list) else []
    sequences = []
    if isinstance(structure.get("ordered_roles"), list) and structure.get("ordered_roles"):
        sequences.append(structure["ordered_roles"])
    sequences.extend(group for group in ordered_groups if isinstance(group, list) and group)
    for role in set(structure.get("required_roles", [])) | {
        role for sequence in sequences for role in sequence
    }:
        role_positions[role] = [positions[p._p] for p in _role_paragraphs(doc, role, mappings)]
    for role in structure.get("required_roles", []):
        if not role_positions.get(role):
            findings.append({"role": "document_structure", "property": f"required_roles.{role}", "template_value": False, "required_value": True})
    for sequence_index, sequence in enumerate(sequences):
        previous = -1
        for role in sequence:
            if not role_positions.get(role):
                continue
            current = min(role_positions[role])
            if current < previous:
                findings.append({"role": "document_structure", "property": "ordered_roles" if sequence_index == 0 else f"ordered_role_groups[{sequence_index - 1}]", "template_value": role, "required_value": sequence})
            previous = max(previous, current)
    max_depth = structure.get("max_heading_depth")
    if max_depth is not None:
        for depth in range(int(max_depth) + 1, 10):
            if any(p.style.name in {f"Heading {depth}", f"Heading{depth}", f"标题 {depth}"} for p in doc.paragraphs):
                findings.append({"role": "document_structure", "property": "max_heading_depth", "template_value": depth, "required_value": max_depth})
    toc_depth = structure.get("toc_depth")
    if toc_depth is not None:
        instructions = " ".join(n.text or "" for p in doc.paragraphs for n in p._p.xpath('.//w:instrText[contains(text(), "TOC")]'))
        if instructions:
            match = re.search(r'\\o\s+"1-(\d+)"', instructions)
            actual = int(match.group(1)) if match else None
            if actual != int(toc_depth): findings.append({"role": "document_structure", "property": "toc_depth", "template_value": actual, "required_value": toc_depth})
        elif _role_paragraphs(doc, "toc", mappings):
            findings.append({"role": "document_structure", "property": "toc_depth", "template_value": "no TOC field", "required_value": toc_depth})
    return findings


def set_toc_depth(doc: Document, depth: int | None) -> int:
    if depth is None: return 0
    changed = 0
    for paragraph in doc.paragraphs:
        for node in paragraph._p.xpath('.//w:instrText[contains(text(), "TOC")]'):
            text = node.text or ""
            updated, count = re.subn(r'\\o\s+"1-\d+"', rf'\\o "1-{int(depth)}"', text)
            if not count: updated = text.rstrip() + f' \\o "1-{int(depth)}" '
            if updated != text:
                node.text = updated; changed += 1
    return changed


_APPENDIX_HEADING_RE = re.compile(
    r"^(?:(?:第[一二三四五六七八九十百千万\d]+章)|(?:\d+(?:\.\d+)*))?\s*"
    r"附\s*录(?:\s*([A-ZＡ-Ｚ]))?\s*(?:[:：-]\s*)?(.*)$",
    re.IGNORECASE,
)


def _appendix_heading_match(text: str) -> re.Match[str] | None:
    return _APPENDIX_HEADING_RE.match((text or "").strip())


def _appendix_headings(doc: Document) -> list[Paragraph]:
    return [p for p in doc.paragraphs if _appendix_heading_match(p.text)]


def apply_appendix_rules(doc: Document, appendix_spec: dict[str, Any]) -> dict[str, int]:
    """Normalize explicit appendix headings without manufacturing appendix content."""
    counts = {"headings": 0, "labels_normalized": 0, "page_breaks_set": 0,
              "titles_missing": 0}
    if not appendix_spec:
        return counts
    prefix = str(appendix_spec.get("label_prefix", "附录"))
    for index, paragraph in enumerate(_appendix_headings(doc)):
        counts["headings"] += 1
        if appendix_spec.get("label_style") == "alpha_upper":
            label = chr(ord("A") + index)
            match = _appendix_heading_match(paragraph.text)
            title = match.group(2).strip() if match else ""
            normalized = f"{prefix} {label}" + (f"  {title}" if title else "")
            if paragraph.text.strip() != normalized:
                paragraph.text = normalized
                counts["labels_normalized"] += 1
        if appendix_spec.get("page_break_each") and paragraph.paragraph_format.page_break_before is not True:
            paragraph.paragraph_format.page_break_before = True
            counts["page_breaks_set"] += 1
        if appendix_spec.get("per_appendix_title_required"):
            title = re.sub(r"^附\s*录(?:\s*[A-ZＡ-Ｚ])?\s*[:：-]?\s*", "", paragraph.text.strip()).strip()
            if not title:
                counts["titles_missing"] += 1
    return counts


def audit_appendix_rules(doc: Document, appendix_spec: dict[str, Any], profile: dict[str, Any]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    if not appendix_spec:
        return findings
    headings = _appendix_headings(doc)
    if (appendix_spec.get("required_when_profile_has_appendices")
            and profile.get("has_appendices") and not headings):
        findings.append({"role": "appendices", "property": "required", "template_value": False, "required_value": True})
    for index, paragraph in enumerate(headings):
        expected = chr(ord("A") + index)
        match = _appendix_heading_match(paragraph.text)
        label = match.group(1).upper() if match and match.group(1) else None
        if appendix_spec.get("label_style") == "alpha_upper" and label != expected:
            findings.append({"role": "appendices", "property": f"headings[{index + 1}].label",
                             "template_value": label if label else paragraph.text.strip(),
                             "required_value": expected})
        if appendix_spec.get("page_break_each") and paragraph.paragraph_format.page_break_before is not True:
            findings.append({"role": "appendices", "property": f"headings[{index + 1}].page_break_before",
                             "template_value": paragraph.paragraph_format.page_break_before, "required_value": True})
        if appendix_spec.get("per_appendix_title_required"):
            title = match.group(2).strip() if match else ""
            if not title:
                findings.append({"role": "appendices", "property": f"headings[{index + 1}].title",
                                 "template_value": "", "required_value": "non-empty appendix title",
                                 "failure_type": "appendix_title_missing",
                                 "reason": "appendix label is present but no title is attached; no title was synthesized"})
    return findings


def _receipt_put(target: dict[str, Any], path: str, value: Any) -> None:
    """Set a dotted property path in a nested receipt-evidence object."""
    parts = path.split(".")
    cursor = target
    for part in parts[:-1]:
        child = cursor.get(part)
        if not isinstance(child, dict):
            child = {}
            cursor[part] = child
        cursor = child
    cursor[parts[-1]] = value


def _receipt_requested_paths(requirements: list[dict[str, Any]]) -> dict[tuple[str, str], Any]:
    requested: dict[tuple[str, str], Any] = {}
    for requirement in requirements:
        role = str(requirement.get("role") or "")
        properties = requirement.get("properties")
        if not role or not isinstance(properties, dict):
            continue
        for property_path, expected in flatten(properties).items():
            requested.setdefault((role, property_path), expected)
    return requested


def _receipt_keyword_separator(text: str, language: str, values: list[str]) -> str | None:
    if len(values) <= 1:
        return None
    payload = re.sub(r"^(?:关\s*键\s*词|key\s*words?)\s*[：: ]*", "", text.strip(), flags=re.I)
    if language == "zh" and "，" in payload and "," not in payload:
        return "chinese_comma"
    if language == "en" and "," in payload and "，" not in payload:
        return "english_comma"
    if "；" in payload or ";" in payload:
        return "semicolon"
    return None


def _receipt_semantic_actuals(
    doc: Document,
    spec: dict[str, Any],
    mappings: dict[str, Any],
    requirements: list[dict[str, Any]],
    applied_section_plan: dict[str, Any],
    section_findings: list[dict[str, Any]],
    cover_contract: dict[str, Any],
) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    """Build post-serialization evidence for non-style requirement properties.

    The old receipt builder only had shared-style snapshots.  That was sound
    for font/paragraph properties but left page, cover, content-order, table,
    appendix, and section-numbering contracts as ``unverified`` even when the
    final OOXML already contained direct evidence.  This adapter deliberately
    derives only values that can be inspected from the serialized document or
    from a section plan whose serialized audit passed.  Missing evidence is
    left absent so the receipt remains ``unverified`` rather than being
    converted into a guessed success.
    """
    requested = _receipt_requested_paths(requirements)
    actual: dict[str, dict[str, Any]] = {}
    methods: dict[str, str] = {}

    def put(role: str, path: str, value: Any, method: str = "serialized_docx_semantic") -> None:
        if (role, path) not in requested:
            return
        _receipt_put(actual.setdefault(role, {}), path, value)
        methods[role] = method

    # Page size, margins, and header/footer distances are direct section XML
    # properties.  A4 is recognized by its serialized portrait dimensions.
    if doc.sections:
        section = doc.sections[0]
        width_pt = float(section.page_width.pt)
        height_pt = float(section.page_height.pt)
        short, long = sorted((width_pt, height_pt))
        if abs(short - 595.276) <= 0.2 and abs(long - 841.89) <= 0.2:
            put("page", "size", "A4")
        for key, attr in (
            ("top", "top_margin"), ("bottom", "bottom_margin"),
            ("left", "left_margin"), ("right", "right_margin"),
            ("header", "header_distance"), ("footer", "footer_distance"),
        ):
            if hasattr(section, attr):
                put("page", f"margins_pt.{key}", round(float(getattr(section, attr).pt), 3))

    # SectionPlan is compiled from the selected heading evidence and then
    # audited against the serialized DOCX.  Its values are therefore accepted
    # as section-level OOXML evidence, including the selector provenance.
    if applied_section_plan and not section_findings:
        planned = {item.get("zone"): item for item in applied_section_plan.get("sections", [])}
        front = planned.get("front")
        body = planned.get("body")
        if isinstance(front, dict):
            put("page", "page_number.front_matter_format", front.get("format"), "serialized_docx_section_plan")
            put("page", "page_number.front_matter_start", front.get("start"), "serialized_docx_section_plan")
        if isinstance(body, dict):
            put("page", "page_number.body_format", body.get("format"), "serialized_docx_section_plan")
            put("page", "page_number.body_start", body.get("start"), "serialized_docx_section_plan")
        page_number = spec.get("page", {}).get("page_number", {})
        for key, selector_key in (
            ("front_matter_start_selector", "front_matter_start_selector"),
            ("body_start_selector", "body_start_selector"),
        ):
            selector = page_number.get(selector_key)
            if isinstance(selector, dict):
                for child_key, child_value in flatten({key: selector}).items():
                    put("page", f"page_number.{child_key}", child_value,
                        "serialized_docx_section_plan")

    # Content constraints are checked from the final paragraph order and text,
    # using the same role detectors as the corresponding audits.
    constraints = resolve_profile_constraints(spec)
    positions = {paragraph._p: index for index, paragraph in enumerate(doc.paragraphs)}
    for key, role, language in (
        ("abstract_zh", "abstract_body_zh", "zh"),
        ("abstract_en", "abstract_body_en", "en"),
    ):
        paragraphs = _role_paragraphs(doc, role, mappings)
        text = "".join(paragraph.text.strip() for paragraph in paragraphs)
        put("content_constraints", f"{key}.required", bool(text))
        rule = constraints.get(key, {}) if isinstance(constraints.get(key), dict) else {}
        metric = rule.get("length_metric", "words" if key == "abstract_en" else "unicode_codepoints")
        measured = _content_metric(text, metric, language)
        put("content_constraints", f"{key}.length", measured, metric)
    for key, role, language in (("keywords_zh", "keywords_zh", "zh"), ("keywords_en", "keywords_en", "en")):
        paragraphs = _role_paragraphs(doc, role, mappings)
        text = " ".join(paragraph.text.strip() for paragraph in paragraphs)
        values = _keyword_values(text, language) if text else []
        put("content_constraints", f"{key}.required", bool(values))
        put("content_constraints", f"{key}.min_count", len(values))
        put("content_constraints", f"{key}.max_count", len(values))
        separator = _receipt_keyword_separator(text, language, values)
        if separator is not None:
            put("content_constraints", f"{key}.separator", separator)
        rule = constraints.get(key, {}) if isinstance(constraints.get(key), dict) else {}
        if rule.get("max_item_chars") is not None:
            metric = rule.get("item_length_metric", "cjk_characters" if language == "zh" else "words")
            lengths = [_content_metric(item, metric, language) for item in values]
            put("content_constraints", f"{key}.max_item_chars", max(lengths, default=0), metric)
        after_role = constraints.get(key, {}).get("require_after_role")
        if after_role:
            before = _role_paragraphs(doc, after_role, mappings)
            valid_order = bool(before and paragraphs and
                               max(positions.get(item._p, -1) for item in before) <
                               min(positions.get(item._p, 10**9) for item in paragraphs))
            put("content_constraints", f"{key}.require_after_role",
                after_role if valid_order else False)

    # TOC depth is read from the actual TOC field instruction.
    toc_depth = None
    for paragraph in doc.paragraphs:
        for node in paragraph._p.xpath('.//w:instrText[contains(text(), "TOC")]'):
            match = re.search(r'\\o\s+"1-(\d+)"', node.text or "")
            if match:
                toc_depth = int(match.group(1))
                break
        if toc_depth is not None:
            break
    if toc_depth is not None:
        put("document_structure", "toc_depth", toc_depth)
    required_roles = spec.get("document_structure", {}).get("required_roles", [])
    if required_roles:
        present = [role for role in required_roles
                   if _role_paragraphs(doc, role, mappings)]
        if len(present) == len(required_roles):
            put("document_structure", "required_roles", list(required_roles))

    # Cover evidence is bounded to the generated cover block; source titles
    # elsewhere in the thesis must not be mistaken for cover duplicates.
    cover = spec.get("cover", {})
    if cover:
        block = _generated_cover_block(doc, cover.get("institution", ""))
        nonempty = [(i, paragraph) for i, paragraph in enumerate(doc.paragraphs)
                    if paragraph.text.strip()]
        institution = [paragraph for paragraph in block
                       if paragraph.text.strip() == str(cover.get("institution", "")).strip()]
        first = nonempty[0][1] if nonempty else None
        if len(institution) == 1 and first is not None:
            same_path = (institution[0]._p.getroottree().getpath(institution[0]._p) ==
                         first._p.getroottree().getpath(first._p))
            if same_path:
                put("cover", "before_role", "document_start")
                put("cover", "institution", cover.get("institution"))
        if block and cover_contract:
            contract_by_id = {item.get("id"): item for item in cover_contract.get("fields", [])}
            for field in cover.get("fields", []):
                field_id = field.get("id")
                contract = contract_by_id.get(field_id)
                if not contract or not contract.get("value"):
                    continue
                expected_text = (contract["value"] if field_id in {"title_zh", "title_en"}
                                 else f"{contract.get('label')}：{contract['value']}")
                if any(paragraph.text.strip() == expected_text for paragraph in block):
                    continue
                break
            else:
                # The requirement may intentionally list a subset of the
                # executable cover fields (for example, title_en is separately
                # tracked as a structural role).  Validate that subset only.
                field_expected = requested.get(("cover", "fields"))
                if field_expected is not None:
                    put("cover", "fields", field_expected)
            placeholder = cover.get("missing_value_placeholder")
            if placeholder == cover_contract.get("missing_value_placeholder", placeholder):
                put("cover", "missing_value_placeholder", placeholder)
            policy = cover.get("missing_value_policy")
            if policy == cover_contract.get("missing_value_policy", policy):
                put("cover", "missing_value_policy", policy)

    # Table contracts are read from table/row properties after serialization.
    tables = list(doc.tables)
    if tables:
        def border_width(table, side: str) -> float:
            borders = table._tbl.tblPr.find(qn("w:tblBorders"))
            node = borders.find(qn(f"w:{side}")) if borders is not None else None
            if node is None or node.get(qn("w:val")) in {"nil", "none"}:
                return 0.0
            raw = node.get(qn("w:sz"))
            return round(float(raw) / 8, 3) if raw else 0.0
        no_vertical = all(
            border_width(table, side) == 0.0
            for table in tables for side in ("left", "right", "insideV")
        )
        three_line = all(
            border_width(table, "top") > 0 and border_width(table, "bottom") > 0 and
            border_width(table, side) == 0.0
            for table in tables for side in ("left", "right", "insideV")
        )
        repeated = all(
            bool(table.rows and table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None)
            for table in tables
        )
        put("table", "remove_vertical_borders", no_vertical)
        put("table", "style", "three_line" if three_line else None)
        put("table", "repeat_header_row", repeated)

    # Appendix labels are verified against the actual serialized headings.
    appendix_spec = spec.get("appendices", {})
    headings = _appendix_headings(doc)
    if headings:
        labels = []
        for paragraph in headings:
            match = _appendix_heading_match(paragraph.text)
            labels.append(match.group(1).upper() if match and match.group(1) else None)
        if all(label == chr(ord("A") + index) for index, label in enumerate(labels)):
            put("appendices", "label_prefix", appendix_spec.get("label_prefix", "附录"))
            put("appendices", "label_style", "alpha_upper")
        profile = spec.get("thesis_profile", {})
        if profile.get("has_appendices") is True:
            put("appendices", "required_when_profile_has_appendices", True)

    return actual, methods


def _set_equation_tabs(paragraph: Paragraph, center: int, right: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    tabs = ppr.find(qn("w:tabs"))
    if tabs is not None:
        ppr.remove(tabs)
    tabs = OxmlElement("w:tabs")
    for align, pos in (("center", center), ("right", right)):
        tab = OxmlElement("w:tab"); tab.set(qn("w:val"), align); tab.set(qn("w:pos"), str(pos)); tabs.append(tab)
    ppr.append(tabs)


def apply_equation_layout(doc: Document, equation_spec: dict[str, Any]) -> dict[str, int]:
    counts = {"equations": 0, "tabs_set": 0, "borders_removed": 0}
    if not equation_spec:
        return counts
    center = int(equation_spec.get("center_tab_twips", 4680))
    right = int(equation_spec.get("right_tab_twips", 9360))
    for paragraph in doc.paragraphs:
        if not is_display_equation(paragraph):
            continue
        counts["equations"] += 1
        if equation_spec.get("alignment") == "center" or equation_spec.get("number_alignment") == "right":
            _set_equation_tabs(paragraph, center, right); counts["tabs_set"] += 1
        if equation_spec.get("no_lines"):
            ppr = paragraph._p.get_or_add_pPr(); borders = ppr.find(qn("w:pBdr"))
            if borders is not None:
                ppr.remove(borders); counts["borders_removed"] += 1
    return counts


def audit_equation_layout(doc: Document, equation_spec: dict[str, Any]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    if not equation_spec:
        return findings
    center = str(int(equation_spec.get("center_tab_twips", 4680)))
    right = str(int(equation_spec.get("right_tab_twips", 9360)))
    for index, paragraph in enumerate((p for p in doc.paragraphs if is_display_equation(p)), 1):
        tab_nodes = paragraph._p.xpath("./w:pPr/w:tabs/w:tab")
        tabs = {(n.get(qn("w:val")), n.get(qn("w:pos"))) for n in tab_nodes}
        if equation_spec.get("alignment") == "center" and ("center", center) not in tabs:
            findings.append({"role": "equations", "property": f"equations[{index}].center_tab", "template_value": sorted(tabs), "required_value": center})
        if equation_spec.get("number_alignment") == "right" and ("right", right) not in tabs:
            findings.append({"role": "equations", "property": f"equations[{index}].right_tab", "template_value": sorted(tabs), "required_value": right})
        raw_text = paragraph.text
        text = raw_text.strip()
        number = re.search(r"[（(]\s*\d+(?:[.-]\d+)*\s*[）)]\s*$", raw_text)
        if equation_spec.get("number_parentheses") and not number:
            findings.append({"role": "equations", "property": f"equations[{index}].number_parentheses", "template_value": text, "required_value": "(number)"})
        if equation_spec.get("same_line") and (not number or "\n" in raw_text):
            findings.append({"role": "equations", "property": f"equations[{index}].same_line", "template_value": False, "required_value": True})
        # A right tab stop alone does not right-align anything.  Require the
        # equation paragraph to actually traverse a center tab and a right tab
        # before its trailing number, as produced by the conversion filter.
        if number and (equation_spec.get("alignment") == "center" or equation_spec.get("number_alignment") == "right"):
            before_number = raw_text[:number.start()]
            required_tabs = 2 if equation_spec.get("alignment") == "center" and equation_spec.get("number_alignment") == "right" else 1
            if before_number.count("\t") < required_tabs:
                findings.append({"role": "equations", "property": f"equations[{index}].tab_usage",
                                 "template_value": before_number.count("\t"), "required_value": f">={required_tabs}"})
        if equation_spec.get("no_lines") and paragraph._p.xpath("./w:pPr/w:pBdr/*[not(@w:val='nil') and not(@w:val='none')]"):
            findings.append({"role": "equations", "property": f"equations[{index}].no_lines", "template_value": False, "required_value": True})
        if equation_spec.get("no_lines"):
            leaders = [n.get(qn("w:leader")) for n in tab_nodes if n.get(qn("w:leader")) not in {None, "none"}]
            if leaders:
                findings.append({"role": "equations", "property": f"equations[{index}].tab_leader",
                                 "template_value": leaders, "required_value": "none"})
    return findings


def main(argv: list[str]) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("input", type=Path); p.add_argument("format_spec", type=Path); p.add_argument("output", type=Path)
    p.add_argument("--out-dir", type=Path, required=True); p.add_argument("--style-map", type=Path)
    p.add_argument("--official-template", type=Path,
                   help="official DOCX whose mapped role styles supply defaults when written rules are silent")
    p.add_argument("--neutral-reference", action="store_true",
                   help="apply a school-neutral reference style map using structural/source-style fallback")
    p.add_argument("--require-coverage", action="store_true",
                   help="treat zero matched paragraphs for critical roles as a validation error")
    p.add_argument("--compliance-mode", choices=["full", "supported_subset"], default=None,
                   help="override the format spec compliance mode")
    p.add_argument("--render-report", type=Path,
                   help="optional accepted Word/PDF render evidence JSON for submission readiness")
    p.add_argument("--require-submission-ready", action="store_true",
                   help="fail unless serialized and rendered submission evidence both pass")
    p.add_argument("--capability-report", type=Path,
                   help="capability-preflight report used to distinguish supplied inputs from missing inputs")
    p.add_argument("--preview-placeholders", action="store_true",
                   help="opt-in supported-subset preview: continue past unresolved requirement semantics while keeping submission_ready false")
    args = p.parse_args(argv)
    spec = load_json(args.format_spec)
    validation_errors = load_and_validate(spec, Path(__file__).resolve().parents[1] / "schema" / "format-spec.schema.json")
    if validation_errors: raise SystemExit("invalid format spec:\n" + "\n".join(validation_errors))
    explicit_raw = load_json(args.style_map) if args.style_map else {}
    source_mappings = explicit_raw.get("mappings", explicit_raw)
    explicit = {role: value.get("style_name") if isinstance(value, dict) else value
                for role, value in source_mappings.items()}
    defaults = official_role_defaults(args.official_template, explicit) if args.official_template else {}
    effective_roles = {
        role: merge_role_style_defaults(defaults.get(role, {}), role_spec)
        for role, role_spec in spec.get("roles", {}).items()
    }
    compliance_mode = args.compliance_mode or spec.get("compliance_mode", "supported_subset")
    blockers = format_spec_blockers(spec, compliance_mode)
    preview_bypassed_blockers = []
    if args.preview_placeholders:
        if compliance_mode != "supported_subset":
            raise SystemExit("--preview-placeholders requires --compliance-mode supported_subset")
        preview_bypassed_blockers = sorted(set(blockers) & {
            "status_needs_clarification", "unresolved_clauses", "open_questions",
        })
        blockers = [item for item in blockers if item not in {
            "status_needs_clarification", "unresolved_clauses", "open_questions",
        }]
    if blockers:
        raise SystemExit(f"refusing to apply a blocked format spec: {', '.join(blockers)}")
    doc = Document(args.input); args.out_dir.mkdir(parents=True, exist_ok=True)
    mappings = {}; conflicts = []; created = []; claimed = {}
    page_num = spec.get("page", {}).get("page_number", {})
    selector_mappings = ({"heading_1": {"style_name": explicit["heading_1"]}}
                         if explicit.get("heading_1") else {})
    try:
        section_created = ensure_page_number_sections(doc, page_num, selector_mappings) if page_num else False
        preliminary_body_start = resolve_body_start_section(doc, page_num, selector_mappings) if page_num else 0
    except ValueError as exc:
        raise SystemExit(f"page-number application blocked: {exc}") from exc
    section_indices = _paragraph_section_indices(doc)
    for role, role_spec in effective_roles.items():
        if role not in ROLE_STYLES and role not in explicit: continue
        structural_fallback = role not in explicit
        name, was_created = resolve_style(doc, role, explicit, claimed)
        source_name = name
        # Normal is shared by body prose, empty layout paragraphs, and often
        # table cells.  A semantic table rule must never shrink or otherwise
        # mutate Normal globally; route table content through its own style.
        if role == "table_text" and name == "Normal":
            name = _new_role_style(doc, role, source_name)
            was_created = name != source_name
        # A template may reuse Normal on covers, declarations, TOC and body. Do not
        # mutate that shared style globally when the requirement is body-scoped.
        if role == "body_text" and preliminary_body_start > 0:
            used_before = any(p.style.name == source_name and section_indices.get(p._p, 0) < preliminary_body_start
                              for p in all_body_paragraphs(doc))
            if used_before:
                scoped_name = "Thesis Body Text"
                if scoped_name in style_names(doc):
                    name = scoped_name; was_created = False
                else:
                    scoped = doc.styles.add_style(scoped_name, WD_STYLE_TYPE.PARAGRAPH)
                    scoped.base_style = doc.styles[source_name]
                    name = scoped_name; was_created = True
        claimed[name] = role; style = doc.styles[name]
        before = style_snapshot(doc.styles[source_name]); expected = normalized_expected(style, role_spec)
        for item in diff_expected(before, expected): conflicts.append({"type": "template_requirement_conflict", "role": role, "style": source_name, **item})
        if role_spec.get("font"): set_font(style, role_spec["font"])
        if role_spec.get("paragraph"): set_paragraph(style, role_spec["paragraph"])
        mappings[role] = {"style_name": name, "source_style_name": source_name, "created": was_created,
                          "structural_fallback": structural_fallback,
                          "scope_start_section": preliminary_body_start if role == "body_text" else 0,
                          "before": before, "after": style_snapshot(style)}
        if was_created: created.append(name)
    applied_paragraphs = {}
    execution_receipts: list[dict[str, Any]] = []
    nodes_by_element = {node.paragraph._p: node for node in iter_document_nodes(doc)}
    # More specific roles must claim paragraphs before broad body prose.  A
    # body style such as Normal is frequently shared by converter-emitted
    # captions; applying body_text first would relabel them and make both the
    # receipt and caption-object contracts internally contradictory.
    role_priority = {
        "figure_caption": 10, "table_caption": 10,
        # Semantic headings and keyword lines must claim a converter-emitted
        # generic Heading 1/Normal paragraph before broad roles do.
        "heading_acknowledgments": 15, "heading_appendix": 15,
        "heading_conclusion": 15, "heading_publications": 15,
        "heading_references": 15, "bibliography_heading": 15,
        "keywords_zh": 15, "keywords_en": 15,
        "equation": 20, "toc": 20, "table_text": 30, "body_text": 100,
    }
    claimed_elements: set[Any] = set()
    for role, mapping in sorted(mappings.items(), key=lambda item: (role_priority.get(item[0], 50), item[0])):
        count = 0
        role_spec = effective_roles[role]
        paragraphs = all_table_paragraphs(doc) if role == "table_text" else all_body_paragraphs(doc)
        for paragraph in paragraphs:
            if paragraph._p in claimed_elements:
                continue
            section_index = section_indices.get(paragraph._p, 0)
            if role == "body_text" and section_index < mapping.get("scope_start_section", 0):
                continue
            eligible_styles = {mapping.get("source_style_name", mapping["style_name"]), mapping["style_name"]}
            # Unstyled table cells normally inherit Normal and are legitimate table
            # body candidates.  Do not, however, seize cells explicitly assigned a
            # different semantic style (for example Body Text), because that would
            # manufacture coverage for a missing declared table_text role.
            detector = structural_detector(role)
            # Converter-emitted semantic content can legitimately carry a
            # source style that differs from an official template's explicit
            # target style.  Only narrow detectors are allowed to bridge that
            # gap.  Table cells need an extra guard: Normal/untyped cells are
            # eligible, while cells already assigned another semantic style
            # must not be seized merely because they live inside a table.
            narrow_structural_roles = {
                "thesis_title_zh", "thesis_title_en",
                "figure_caption", "table_caption", "equation", "keywords_zh", "keywords_en",
                "heading_acknowledgments", "heading_appendix", "heading_conclusion",
                "heading_publications", "heading_references", "bibliography_heading",
            }
            detector_match = bool(detector and matches_structural_detector(paragraph, detector))
            if role == "toc":
                # Do not map the visible “目录” heading to the executable TOC
                # role.  Map only the real field paragraph; its cached result
                # may legitimately be empty until Word/WPS updates fields.
                structural_candidate = has_toc_field(paragraph)
            elif role == "table_text":
                # Location is the authoritative discriminator for ordinary
                # table-cell prose.  Exclude explicit nested semantics, but do
                # not require the cell's inherited style to match the style
                # selected from the official template: many converters emit a
                # named Body Text style in every cell.
                source_style = paragraph.style.name
                # Known converter defaults are safe to migrate.  Unknown named
                # styles remain protected because they may encode a nested
                # semantic role not represented by the current specification.
                generic_table_styles = {"Normal", "Compact"}
                nested_semantic = (is_display_equation(paragraph)
                                   or matches_structural_detector(paragraph, "figure_caption")
                                   or matches_structural_detector(paragraph, "table_caption"))
                structural_candidate = detector_match and not nested_semantic and (
                    source_style in eligible_styles | generic_table_styles
                )
            else:
                structural_candidate = detector_match and (
                    mapping.get("structural_fallback") or role in narrow_structural_roles or args.neutral_reference
                )
            if args.neutral_reference and role == "body_text" and paragraph.style.name in {"Normal", "Body Text"}:
                structural_candidate = True
            eligible_candidate = paragraph.style.name in eligible_styles
            if role == "toc":
                eligible_candidate = has_toc_field(paragraph)
            if structural_candidate or (role not in {"table_text", "equation"} and eligible_candidate):
                old_style = paragraph.style.name
                if mapping["style_name"] != paragraph.style.name:
                    paragraph.style = mapping["style_name"]
                apply_direct_format(paragraph, role_spec); count += 1
                node = nodes_by_element.get(paragraph._p)
                execution_receipts.append({
                    "operation": "apply_role_format", "role": role,
                    "node_id": node.node_id if node else None,
                    "story": node.story if node else "document",
                    "container": node.container if node else "body",
                    "xml_order": node.xml_order if node else None,
                    "source_style_name": old_style,
                    "style_name": mapping["style_name"],
                    "value_kind": "object" if has_drawing(paragraph) else "text",
                    "value": paragraph.text[:240],
                    "_element": paragraph._p,
                })
                claimed_elements.add(paragraph._p)
        applied_paragraphs[role] = count
    numbering_changed, numbering_issues = apply_numbering(doc, mappings, effective_roles)
    apply_page(doc, spec.get("page", {}))
    try:
        # Compile only after semantic boundary creation, so the plan describes
        # the exact section topology that will be serialized.  The executor is
        # the sole PAGE-field and pgNumType mutator; header content remains a
        # separate story-content operation until its plan schema is introduced.
        with tempfile.NamedTemporaryFile(suffix=".docx") as handle:
            doc.save(handle.name)
            applied_section_plan = compile_section_plan(spec, handle.name)
        if not applied_section_plan.get("valid"):
            raise ValueError("compiled SectionPlan is invalid: " + json.dumps(
                applied_section_plan.get("findings", []), ensure_ascii=False))
        section_execution = execute_section_plan(
            doc, applied_section_plan,
            alignment=spec.get("page", {}).get("page_number", {}).get("alignment"),
        )
        expected_body_start = int(applied_section_plan["body_start_section"] or 1) - 1
        # Header/footer text formatting is intentionally invoked without a
        # page-number rule; pagination has already been applied from the plan.
        # PAGE-field placement/alignment belongs exclusively to the compiled
        # SectionPlan.  Do not let the generic footer role overwrite the
        # executor's centered footer paragraph with body-style alignment.
        story_roles = dict(effective_roles)
        story_roles.pop("footer", None)
        header_footer_changed = apply_headers_footers(
            doc, story_roles, {}, mappings, {},
        )
        # Stable report compatibility: these counters now come from the
        # SectionPlan execution evidence rather than a second pagination path.
        pagination_stories = section_execution.get("stories", [])
        header_footer_changed["page_fields_removed"] = sum(
            int(item.get("removed_page_fields", 0)) for item in pagination_stories
        )
        header_footer_changed["page_numbers"] = sum(
            1 for section in applied_section_plan.get("sections", []) if section.get("numbered")
        )
    except ValueError as exc:
        raise SystemExit(f"page-number application blocked: {exc}") from exc
    moved = {}; caption_graph: list[dict[str, Any]] = []
    for role in ("figure_caption", "table_caption"):
        pos = spec.get("roles", {}).get(role, {}).get("position")
        if pos and role in mappings:
            moved[role], position_issues, role_graph = reposition_captions(
                doc, role, mappings[role]["style_name"], pos)
            numbering_issues.extend(position_issues)
            caption_graph.extend(role_graph)
    table_changes = apply_table_rules(doc, spec.get("tables", {}))
    pagination_changes = apply_object_pagination(doc, mappings, spec.get("objects", {}), spec.get("tables", {}))
    toc_fields_updated = set_toc_depth(doc, spec.get("document_structure", {}).get("toc_depth"))
    appendix_changes = apply_appendix_rules(doc, spec.get("appendices", {}))
    equation_changes = apply_equation_layout(doc, spec.get("equations", {}))
    cover_contract = compile_cover_contract(spec.get("cover", {}), spec.get("thesis_profile", {})) if spec.get("cover") else {}
    cover_changes = apply_cover(doc, spec.get("cover", {}), spec.get("thesis_profile", {}), cover_contract, mappings)
    declaration_changes = apply_declarations(
        doc, spec.get("declarations", {}), resource_items(spec)
    )
    content_instance_audit = apply_content_instance_overrides(doc, spec, mappings)
    pending_content = insert_missing_content_placeholders(doc, spec, mappings)
    # Structural generators can insert content before already formatted
    # paragraphs.  Resolve stable node ids only after all such mutations.
    # lxml may hand out distinct Python proxy objects for the same underlying
    # paragraph, so a dictionary keyed by the proxy can miss after body
    # insertions.  Resolve by XML-element equality after generators finish;
    # this keeps positional node_ids synchronized with the serialized body.
    final_nodes = list(iter_document_nodes(doc))
    consumed_node_ids: set[str] = set()
    for receipt in execution_receipts:
        element = receipt.pop("_element")
        node = next((candidate for candidate in final_nodes
                     if candidate.node_id not in consumed_node_ids
                     and candidate.paragraph._p is element), None)
        if node is None:
            # Element proxy identity is not stable across every lxml traversal.
            # Fall back to the receipt's serialized semantic signature and
            # consume matches in document order, preserving duplicate text.
            node = next((candidate for candidate in final_nodes
                         if candidate.node_id not in consumed_node_ids
                         and candidate.paragraph.style.name == receipt["style_name"]
                         and candidate.paragraph.text[:240] == receipt["value"]
                         and ("object" if has_drawing(candidate.paragraph) else "text") == receipt["value_kind"]), None)
        if node:
            consumed_node_ids.add(node.node_id)
            receipt.update({"node_id": node.node_id, "story": node.story,
                            "container": node.container, "xml_order": node.xml_order})
        elif cover_contract and receipt["role"] in {"thesis_title_zh", "thesis_title_en"}:
            # The cover compiler intentionally replaces source title nodes with
            # contract-generated title nodes.  Preserve that causal history as
            # a superseded receipt rather than reporting a false round-trip
            # serialization failure for a paragraph that was deliberately removed.
            receipt.update({"node_id": None, "xml_order": None, "superseded": True,
                            "superseded_by": "cover_compiler"})
    staged_output = sibling_temp(args.output)
    try:
        doc.save(staged_output)
        canonicalize_docx_zip(staged_output)
        commit_files([(staged_output, args.output)])
    finally:
        staged_output.unlink(missing_ok=True)
    write = lambda name, data: atomic_write_text(
        args.out_dir / name,
        json.dumps(data, ensure_ascii=False, indent=2) + "\n",
    )
    write("style-map.json", mappings); write("conflicts.json", conflicts)
    write("section-plan-applied.json", applied_section_plan)
    write("section-execution.json", section_execution)
    write("caption-object-graph.json", {"schema_version": "1.0", "bindings": caption_graph})
    write("content-instance-audit.json", {
        "schema_version": "1.0",
        "policy": "match_existing_literal_text_only",
        "instances": content_instance_audit,
        "counts": {
            "total": len(content_instance_audit),
            "applied": sum(item.get("status") == "applied" for item in content_instance_audit),
            "shared_style_present": sum(item.get("status") == "shared_style_present" for item in content_instance_audit),
            "not_present": sum(item.get("status") == "not_present" for item in content_instance_audit),
        },
    })
    write("cover-contract.json", cover_contract)
    write("execution-receipts.json", {"schema_version": "1.0", "receipts": execution_receipts})
    # Re-open the serialized file before validation to catch OOXML round-trip errors.
    check = Document(args.output); findings = []; coverage_warnings = []
    render_report = load_json(args.render_report) if args.render_report else None
    serialized_nodes = {node.node_id: node for node in iter_document_nodes(check)}
    receipt_audit = []
    for receipt in execution_receipts:
        node = serialized_nodes.get(receipt.get("node_id"))
        ok = bool(receipt.get("superseded") or
                  (node and node.paragraph.style.name == receipt["style_name"]))
        receipt_audit.append({**receipt, "serialized": ok,
                              "serialized_style_name": node.paragraph.style.name if node else None})
        if not ok:
            findings.append({
                "role": receipt["role"], "property": "execution_receipt",
                "node_id": receipt.get("node_id"), "failure_type": "serialization_mismatch",
                "template_value": node.paragraph.style.name if node else None,
                "required_value": receipt["style_name"],
                "reason": "the applied semantic operation was not found after DOCX serialization",
            })
    write("execution-receipt-audit.json", {"schema_version": "1.0",
          "valid": all(item["serialized"] for item in receipt_audit), "receipts": receipt_audit})
    for role, mapping in mappings.items():
        style = check.styles[mapping["style_name"]]
        actual = style_snapshot(style); expected = normalized_expected(style, effective_roles[role])
        for item in diff_expected(actual, expected):
            findings.append({"role": role, "style": mapping["style_name"], **item})
    coverage = role_coverage(check, spec.get("roles", {}), mappings)
    for item in coverage:
        if item["expectation"] == "required" and item["status"] == "missing":
            issue = {"role": item["role"], "style": item.get("style"), "property": "coverage",
                     "template_value": item["count"], "required_value": ">=1 content-bearing paragraph",
                     "expectation": "required", "status": "missing",
                     "failure_type": item.get("failure_type", "content_missing"),
                     "structural_count": item.get("structural_count", 0),
                     "evidence": item.get("evidence", []),
                     "reason": "the document contains no content mapped to this declared executable role"}
            (findings if args.require_coverage else coverage_warnings).append(issue)
    serialized_body = list(all_body_paragraphs(check))
    for item in pending_content:
        serialized_count = sum(
            1 for paragraph in serialized_body
            if paragraph.style.name == item["style"]
            and paragraph.text.strip() == NEUTRAL_CONTENT_PLACEHOLDER
        )
        item["serialized_count"] = serialized_count
        if serialized_count < int(item.get("count", 1)):
            findings.append({
                "role": item["role"], "property": "content_placeholder",
                "template_value": serialized_count,
                "required_value": item.get("count", 1),
                "failure_type": "placeholder_serialization_mismatch",
                "reason": "the neutral content placeholder was not preserved after DOCX serialization",
            })
    for item in pending_content:
        coverage_warnings.append({
            "role": item["role"], "style": item["style"],
            "property": "content_placeholder", "template_value": item["placeholder"],
            "required_value": "user-supplied thesis content",
            "expectation": "pending", "status": "pending_user_content",
            "failure_type": "content_placeholder_pending",
            "reason": "formatting continued with a neutral placeholder; submission remains blocked",
        })
    write("pending-content.json", {
        "schema_version": "1.0",
        "placeholder_policy": "neutral_placeholder",
        "placeholder": NEUTRAL_CONTENT_PLACEHOLDER,
        "items": pending_content,
    })
    findings.extend(numbering_issues)
    section_findings = audit_plan_against_docx(applied_section_plan, args.output)
    write("section-plan-audit.json", {
        "schema_version": "1.0", "valid": not section_findings, "findings": section_findings,
    })
    for item in section_findings:
        findings.append({
            "role": "page", "property": item["code"],
            "template_value": next((e.get("value") for e in item.get("evidence", []) if e.get("kind") == "actual"), None),
            "required_value": next((e.get("value") for e in item.get("evidence", []) if e.get("kind") == "expected"), None),
            "section_plan_finding": item,
        })
    findings.extend(audit_headers(args.output, spec.get("roles", {}), mappings))
    findings.extend(audit_object_order_constraints(check, spec.get("objects", {}), mappings))
    findings.extend(audit_table_rules(check, spec.get("tables", {}), render_report))
    effective_content_constraints = resolve_profile_constraints(spec)
    findings.extend(audit_content_constraints(check, effective_content_constraints, mappings))
    findings.extend(audit_document_structure(check, spec.get("document_structure", {}), mappings))
    findings.extend(audit_appendix_rules(check, spec.get("appendices", {}), spec.get("thesis_profile", {})))
    findings.extend(audit_equation_layout(check, spec.get("equations", {})))
    findings.extend(audit_cover(check, spec.get("cover", {}), spec.get("thesis_profile", {}), mappings))
    findings.extend(audit_declarations(
        check, spec.get("declarations", {}), resource_items(spec)
    ))
    expected_page = spec.get("page", {})
    if expected_page:
        section = check.sections[0]
        for key, attr in (("top", "top_margin"), ("bottom", "bottom_margin"), ("left", "left_margin"), ("right", "right_margin")):
            if key in expected_page.get("margins_pt", {}):
                actual = round(getattr(section, attr).pt, 3); expected = expected_page["margins_pt"][key]
                if abs(actual-expected) > .05: findings.append({"role": "page", "property": f"margins_pt.{key}", "template_value": actual, "required_value": expected})
    missing_required = [item for item in coverage
                        if item["expectation"] == "required" and item["status"] == "missing"]
    serialized_docx_sha256 = hashlib.sha256(args.output.read_bytes()).hexdigest()
    receipt_requirements = []
    actual_by_role = {}
    for role, mapping in mappings.items():
        try:
            actual_by_role[role] = style_snapshot(check.styles[mapping["style_name"]])
        except KeyError:
            actual_by_role[role] = {}
    for requirement in spec.get("requirements", []):
        if not isinstance(requirement, dict):
            continue
        item = copy.deepcopy(requirement)
        role = item.get("role")
        if role in mappings and isinstance(item.get("properties"), dict):
            try:
                item["properties"] = normalized_expected(
                    check.styles[mappings[role]["style_name"]], item["properties"]
                )
            except KeyError:
                pass
        receipt_requirements.append(item)
    semantic_actuals, verification_methods = _receipt_semantic_actuals(
        check, spec, mappings, receipt_requirements, applied_section_plan,
        section_findings, cover_contract,
    )
    for role, evidence in semantic_actuals.items():
        actual_by_role.setdefault(role, {}).update(evidence)
    # Role-level structural receipts are derived from the same post-serialization
    # findings used by clause finalization.  Compute the role result before
    # building receipts; otherwise ``__requirement__`` receipts default to
    # false even when the role is present and no finding targets it.
    finding_roles = {item.get("role") for item in findings if item.get("role")}
    role_results = {
        item["role"]: item.get("status") == "present" and item["role"] not in finding_roles
        for item in coverage if item.get("expectation") == "required"
    }
    if spec.get("page"):
        role_results["page"] = "page" not in finding_roles
    if spec.get("tables"):
        role_results["table"] = "table" not in finding_roles and bool(check.tables)
    if spec.get("objects"):
        role_results["objects"] = "objects" not in finding_roles
    if spec.get("content_constraints") or spec.get("conditional_constraints"):
        role_results["content_constraints"] = "content_constraints" not in finding_roles
        role_results["conditional_constraints"] = "content_constraints" not in finding_roles
    if spec.get("document_structure"):
        role_results["document_structure"] = "document_structure" not in finding_roles
    if spec.get("appendices"):
        role_results["appendices"] = "appendices" not in finding_roles
    if spec.get("equations"):
        role_results["equations"] = "equations" not in finding_roles
    if spec.get("cover"):
        role_results["cover"] = "cover" not in finding_roles
    if spec.get("declarations"):
        role_results["declarations"] = "declarations" not in finding_roles
    property_receipts = build_property_receipts(
        receipt_requirements, mappings, actual_by_role,
        serialized_docx_sha256=serialized_docx_sha256,
        role_results=role_results,
        verification_methods=verification_methods,
        applicable_roles={
            item["role"] for item in coverage
            if item.get("status") == "present" or item.get("expectation") == "required"
        } | set(role_results),
    )
    property_receipt_audit = audit_property_receipts(
        property_receipts,
        expected_receipt_ids=expected_receipt_ids(
            receipt_requirements,
            applicable_roles={
                item["role"] for item in coverage
                if item.get("status") == "present" or item.get("expectation") == "required"
            } | set(role_results),
        ),
    )
    write("property-receipts.json", {
        "schema_version": "1.0", "receipts": property_receipts,
    })
    write("property-receipt-audit.json", property_receipt_audit)
    if compliance_mode == "full":
        for receipt in property_receipts:
            if receipt.get("status") != "verified":
                findings.append({
                    "role": receipt.get("role"),
                    "property": receipt.get("property_path"),
                    "target_locator": receipt.get("target_locator"),
                    "failure_type": "property_receipt_" + str(receipt.get("status")),
                    "template_value": receipt.get("actual"),
                    "required_value": receipt.get("expected"),
                    "reason": "every executable property requires a serialized DOCX receipt",
                })
    source_clause_records = spec.get("clause_compliance", [])
    if args.capability_report:
        capability_report = load_json(args.capability_report)
        source_clause_records = annotate_satisfied_inputs(
            source_clause_records, capability_report.get("clauses", []))
    verified_clause_records = finalize_records(
        source_clause_records, spec.get("requirements", []), role_results, finding_roles
    )
    compliance = compliance_report(verified_clause_records, compliance_mode, "validation")
    submission_audit = audit_submission_docx(args.output, spec, render_report)
    cover_pending_fields = cover_changes.get("metadata_pending_fields", [])
    metadata_pending = bool(cover_pending_fields)
    content_pending = bool(pending_content)
    if content_pending:
        compliance["docx_fully_compliant"] = False
        if compliance.get("overall_status") in {"passed", "supported_subset_passed"}:
            compliance["overall_status"] = "content_pending"
    effective_submission_ready = bool(
        submission_audit["submission_ready"] and not metadata_pending and not content_pending
    )
    legacy_role_coverage = not missing_required and not content_pending
    if compliance_mode == "supported_subset" and not source_clause_records:
        compliance["docx_fully_compliant"] = False
        compliance["overall_status"] = (
            "content_pending" if content_pending
            else ("supported_subset_passed" if not findings and legacy_role_coverage else "failed")
        )
    report = {"valid": not findings, "fully_covered": (compliance["docx_fully_compliant"] if source_clause_records else legacy_role_coverage),
              "pipeline_valid": not findings,
              "format_ready": (not findings and bool(compliance.get("format_ready"))
                               and bool(property_receipt_audit.get("valid"))),
              "supported_subset_valid": not findings and compliance["overall_status"] in {"passed", "supported_subset_passed", "input_pending"},
              "serialized_docx_valid": submission_audit["serialized_docx_valid"],
              "render_validation": submission_audit["render_validation"],
              "submission_ready": effective_submission_ready,
              "submission_status": ("content_pending" if content_pending
                                    else ("cover_metadata_pending" if metadata_pending else submission_audit["status"])),
              "overall_status": compliance["overall_status"],
              "docx_fully_compliant": compliance["docx_fully_compliant"],
              "compliance_mode": compliance_mode,
              "compliance_summary": {k: v for k, v in compliance.items() if k not in {"records", "external_checklist"}},
              "property_receipt_audit": property_receipt_audit,
              "findings": findings, "coverage_warnings": coverage_warnings, "styles_created": created,
              "role_coverage": coverage,
              "content_placeholder_policy": {
                  "placeholder": NEUTRAL_CONTENT_PLACEHOLDER,
                  "pending": content_pending,
                  "items": pending_content,
              },
              "pending_content": pending_content,
              "paragraphs_directly_formatted": applied_paragraphs,
              "numbered_paragraphs": numbering_changed, "header_footer_changes": header_footer_changed,
              "automatic_section_created": section_created,
              "caption_paragraphs_moved": moved,
              "table_changes": table_changes,
              "object_pagination_changes": pagination_changes,
              "toc_fields_updated": toc_fields_updated,
              "appendix_changes": appendix_changes,
              "equation_layout_changes": equation_changes,
              "cover_changes": cover_changes,
              "content_instance_audit": content_instance_audit,
              "cover_metadata": {"status": cover_changes.get("metadata_status", "not_applicable"),
                                   "pending_fields": cover_pending_fields},
              "declaration_changes": declaration_changes,
              "submission_audit": submission_audit,
              "unsupported_items": spec.get("completeness", {}).get("unsupported_items", []),
              "preview_placeholders": bool(args.preview_placeholders),
              "preview_bypassed_blockers": preview_bypassed_blockers,
              "output_docx": str(args.output)}
    write("clause-compliance-report.json", compliance)
    write("external-compliance-checklist.json", compliance["external_checklist"])
    write("submission-audit.json", submission_audit)
    write("validation-report.json", report); print(json.dumps(report, ensure_ascii=False))
    failed = (bool(findings) or (compliance_mode == "full" and not report["format_ready"])
              or (args.require_submission_ready and not effective_submission_ready))
    return 1 if failed else 0

if __name__ == "__main__": raise SystemExit(main(sys.argv[1:]))
