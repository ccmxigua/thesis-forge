#!/usr/bin/env python3
"""Apply only deterministic DOCX repairs and emit an evidence report.

This command never edits the input, never renders Word, and never claims that
the output is submission-ready.  A separate submission audit remains required.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from lxml import etree

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from apply_format_spec import (
    _unlink_story_preserving_content,
    all_story_paragraphs,
)
from format_spec_validation import load_and_validate
from section_executor import execute_section_plan
from section_model import audit_plan_against_docx, compile_section_plan


DEGREE_TEXT = {"doctor": "博士", "master": "硕士"}
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
COMMENT_ELEMENT_NAMES = {
    "commentRangeStart", "commentRangeEnd", "commentReference",
    "comment", "commentExt", "commentId", "commentEx",
}
COMMENT_PART_BASENAMES = {
    "comments.xml", "commentsExtended.xml", "commentsIds.xml", "people.xml",
}
COMMENT_RELATIONSHIP_TOKENS = ("/comments", "/people")
COMMENT_CONTENT_TYPE_TOKENS = ("comments", "person")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def remove_docx_comments(path: Path) -> dict[str, Any]:
    """Remove Word comments and every package-level reference to them.

    This intentionally runs after python-docx serialization because comments,
    extended-comment metadata, relationships, and content-type declarations
    live outside the main document object model.  The transformation is
    package-generic and does not inspect comment text or template identity.
    """
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}

    def relationship_owner(rels_name: str) -> str:
        if rels_name == "_rels/.rels":
            return ""
        match = re.fullmatch(r"(.+)/_rels/([^/]+)\.rels", rels_name)
        return f"{match.group(1)}/{match.group(2)}" if match else ""

    def resolve_target(rels_name: str, target: str) -> str:
        owner = relationship_owner(rels_name)
        base = Path(owner).parent if owner else Path()
        return str((base / target).as_posix()).lstrip("/")

    removed_part_set = {
        name for name in members
        if Path(name).name in COMMENT_PART_BASENAMES
        or re.fullmatch(r"word/comments(?:Extended|Ids)?\.xml", name, re.I)
    }

    # Discover nonstandard comment-part names from package declarations rather
    # than relying only on Word's conventional basenames.
    content_types = members.get("[Content_Types].xml")
    if content_types is not None:
        root = etree.fromstring(content_types)
        for node in root:
            part_name = (node.get("PartName") or "").lstrip("/")
            content_type = (node.get("ContentType") or "").casefold()
            if part_name and any(token in content_type for token in COMMENT_CONTENT_TYPE_TOKENS):
                removed_part_set.add(part_name)
    for name, payload in members.items():
        if not name.endswith(".rels"):
            continue
        try:
            root = etree.fromstring(payload)
        except etree.XMLSyntaxError:
            continue
        for rel in root:
            if (rel.get("TargetMode") or "").casefold() == "external":
                continue
            rel_type = (rel.get("Type") or "").casefold()
            if any(token in rel_type for token in COMMENT_RELATIONSHIP_TOKENS):
                removed_part_set.add(resolve_target(name, rel.get("Target") or ""))

    removed_parts = sorted(name for name in removed_part_set if name in members)
    removed_part_set = set(removed_parts)
    removed_sidecar_relationship_parts = sorted(
        name for part in removed_parts
        for name in [str(Path(part).parent / "_rels" / f"{Path(part).name}.rels")]
        if name in members
    )
    removed_markers = 0
    rewritten_xml_parts: list[str] = []

    for name, payload in list(members.items()):
        if name in removed_part_set or not name.endswith(".xml"):
            continue
        try:
            root = etree.fromstring(payload)
        except etree.XMLSyntaxError:
            continue
        changed = False
        for node in list(root.iter()):
            if etree.QName(node).localname not in COMMENT_ELEMENT_NAMES:
                continue
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
                removed_markers += 1
                changed = True
        if changed:
            members[name] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            rewritten_xml_parts.append(name)

    removed_relationships = 0
    for name, payload in list(members.items()):
        if not name.endswith(".rels"):
            continue
        try:
            root = etree.fromstring(payload)
        except etree.XMLSyntaxError:
            continue
        changed = False
        for rel in list(root):
            target = rel.get("Target") or ""
            rel_type = (rel.get("Type") or "").casefold()
            resolved_target = resolve_target(name, target)
            if resolved_target in removed_part_set or any(
                    token in rel_type for token in COMMENT_RELATIONSHIP_TOKENS):
                root.remove(rel)
                removed_relationships += 1
                changed = True
        if changed:
            members[name] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )

    removed_content_types = 0
    content_types = members.get("[Content_Types].xml")
    if content_types is not None:
        root = etree.fromstring(content_types)
        for node in list(root):
            part_name = (node.get("PartName") or "").lstrip("/")
            content_type = (node.get("ContentType") or "").casefold()
            if part_name in removed_part_set or any(
                    token in content_type for token in COMMENT_CONTENT_TYPE_TOKENS):
                root.remove(node)
                removed_content_types += 1
        members["[Content_Types].xml"] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )

    for name in removed_parts + removed_sidecar_relationship_parts:
        members.pop(name, None)

    # Package-level postcondition: no comment part, marker, relationship, or
    # content-type declaration may survive.  Fail before replacing the source
    # file so a partial cleanup can never be reported as successful.
    residuals: list[str] = []
    for name, payload in members.items():
        if Path(name).name in COMMENT_PART_BASENAMES:
            residuals.append(f"part:{name}")
        if not (name.endswith(".xml") or name.endswith(".rels")):
            continue
        try:
            root = etree.fromstring(payload)
        except etree.XMLSyntaxError:
            continue
        if name.endswith(".rels"):
            for rel in root:
                rel_type = (rel.get("Type") or "").casefold()
                resolved = resolve_target(name, rel.get("Target") or "")
                if (any(token in rel_type for token in COMMENT_RELATIONSHIP_TOKENS)
                        or resolved in removed_part_set):
                    residuals.append(f"relationship:{name}:{rel.get('Id') or ''}")
        elif name == "[Content_Types].xml":
            for node in root:
                content_type = (node.get("ContentType") or "").casefold()
                part_name = (node.get("PartName") or "").lstrip("/")
                if (any(token in content_type for token in COMMENT_CONTENT_TYPE_TOKENS)
                        or part_name in removed_part_set):
                    residuals.append(f"content-type:{part_name or content_type}")
        else:
            for node in root.iter():
                if etree.QName(node).localname in COMMENT_ELEMENT_NAMES:
                    residuals.append(f"marker:{name}:{etree.QName(node).localname}")
    if residuals:
        raise ValueError(f"comment cleanup postcondition failed: {sorted(set(residuals))}")

    replacement = path.with_name(f".{path.name}.comments-clean")
    with zipfile.ZipFile(replacement, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    replacement.replace(path)
    return {
        "removed_parts": removed_parts,
        "removed_part_count": len(removed_parts),
        "removed_sidecar_relationship_parts": removed_sidecar_relationship_parts,
        "removed_markers": removed_markers,
        "removed_relationships": removed_relationships,
        "removed_content_types": removed_content_types,
        "rewritten_xml_parts": sorted(rewritten_xml_parts),
    }


def _normalized_text(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def cleanup_declared_ranges(doc: Document, ranges: list[dict[str, Any]]) -> dict[str, Any]:
    """Remove uniquely declared body ranges without interpreting their prose."""
    body = doc._element.body
    removed: list[dict[str, Any]] = []
    for rule in ranges:
        children = list(body)
        paragraphs = [node for node in children if node.tag == qn("w:p")]
        start_text = _normalized_text(str(rule.get("start_text", "")))
        end_text = _normalized_text(str(rule.get("end_text", "")))
        starts = [node for node in paragraphs
                  if _normalized_text("".join(node.xpath(".//w:t/text()"))) == start_text]
        ends = [node for node in paragraphs
                if _normalized_text("".join(node.xpath(".//w:t/text()"))) == end_text]
        if len(starts) != 1 or len(ends) != 1:
            raise ValueError(
                f"cleanup range boundaries must be unique: {rule.get('start_text')!r} "
                f"({len(starts)}) -> {rule.get('end_text')!r} ({len(ends)})"
            )
        start_index = children.index(starts[0]); end_index = children.index(ends[0])
        if start_index > end_index:
            raise ValueError("cleanup range boundaries are inverted")
        first = start_index if rule.get("include_start") else start_index + 1
        last = end_index + 1 if rule.get("include_end") else end_index
        targets = children[first:last]
        removed_count = preserved_boundaries = 0
        for node in targets:
            if node.xpath(".//w:sectPr"):
                for child in list(node):
                    if child.tag != qn("w:pPr"):
                        node.remove(child)
                preserved_boundaries += 1
            else:
                body.remove(node)
                removed_count += 1
        removed.append({
            "start_text": rule.get("start_text"), "end_text": rule.get("end_text"),
            "removed_body_children": removed_count,
            "preserved_section_boundaries": preserved_boundaries,
        })
    return {"ranges": removed, "range_count": len(removed),
            "removed_body_children": sum(item["removed_body_children"] for item in removed)}


def normalize_heading1_boundaries(doc: Document, texts: list[str]) -> dict[str, Any]:
    """Apply Heading 1 only to unique non-TOC paragraphs named by configuration."""
    changed: list[dict[str, Any]] = []
    for configured in texts:
        target = _normalized_text(configured)
        matches = []
        for index, paragraph in enumerate(doc.paragraphs, 1):
            style_id = paragraph.style.style_id or ""
            style_name = paragraph.style.name or ""
            if re.match(r"^TOC\d*$", style_id, re.I) or re.match(r"^toc\s*\d*$", style_name, re.I):
                continue
            if _normalized_text(paragraph.text) == target:
                matches.append((index, paragraph))
        if len(matches) != 1:
            raise ValueError(f"Heading 1 boundary must be unique: {configured!r} matched {len(matches)}")
        index, paragraph = matches[0]
        old_style = paragraph.style.style_id
        paragraph.style = doc.styles["Heading 1"]
        changed.append({"paragraph": index, "text": paragraph.text,
                        "old_style_id": old_style, "new_style_id": paragraph.style.style_id})
    return {"count": len(changed), "changes": changed}


def cleanup_declared_header_texts(doc: Document, texts: list[str]) -> dict[str, Any]:
    """Clear uniquely declared header paragraphs without keyword heuristics."""
    changed: list[dict[str, Any]] = []
    for configured in texts:
        target = _normalized_text(configured)
        matches: list[tuple[int, str, Paragraph]] = []
        seen_parts: set[int] = set()
        for section_index, section in enumerate(doc.sections, 1):
            for variant, story in (("default", section.header),
                                   ("first", section.first_page_header),
                                   ("even", section.even_page_header)):
                marker = id(story._element)
                if marker in seen_parts:
                    continue
                seen_parts.add(marker)
                for paragraph in all_story_paragraphs(story):
                    if _normalized_text(paragraph.text) == target:
                        matches.append((section_index, variant, paragraph))
        if len(matches) != 1:
            raise ValueError(
                f"declared header text must be unique: {configured!r} matched {len(matches)}"
            )
        section_index, variant, paragraph = matches[0]
        for node in list(paragraph._p):
            if node.tag != qn("w:pPr"):
                paragraph._p.remove(node)
        changed.append({"text": configured, "section": section_index, "variant": variant})
    return {"count": len(changed), "changes": changed}


def _clear_story(story) -> Paragraph:
    paragraphs = all_story_paragraphs(story)
    if not paragraphs:
        return story.add_paragraph()
    target = paragraphs[0]
    for node in list(story._element):
        if node is not target._p:
            story._element.remove(node)
    for node in list(target._p):
        if node.tag != qn("w:pPr"):
            target._p.remove(node)
    return target


def configure_body_headers(doc: Document, spec: dict[str, Any]) -> dict[str, Any]:
    roles = spec.get("roles", {}) if isinstance(spec.get("roles"), dict) else {}
    header = roles.get("header", {}) if isinstance(roles.get("header"), dict) else {}
    content = header.get("header_content", {}) if isinstance(header.get("header_content"), dict) else {}
    if content.get("right_field") != "styleref_heading_1":
        return {"changed_sections": 0, "skipped": True}
    page_rule = spec.get("page", {}).get("page_number", {})
    if not isinstance(page_rule, dict):
        raise ValueError("page.page_number is required for body-only header configuration")
    with tempfile.NamedTemporaryFile(suffix=".docx") as handle:
        doc.save(handle.name)
        plan = compile_section_plan(spec, handle.name)
    if not plan.get("valid"):
        raise ValueError("cannot configure headers from an invalid SectionPlan")
    operand, style_evidence = _heading_style_reference(doc)
    even_text = content.get("left_text")
    doc.settings.odd_and_even_pages_header_footer = True
    changed = []
    for index, (section, section_plan) in enumerate(zip(doc.sections, plan["sections"]), 1):
        if section_plan.get("zone") != "body":
            continue
        _unlink_story_preserving_content(section.header, force_clone=index > 1)
        odd = _clear_story(section.header)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), f"STYLEREF {operand} \\* MERGEFORMAT")
        run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = ""
        run.append(text); field.append(run); odd._p.append(field)
        _unlink_story_preserving_content(section.even_page_header, force_clone=index > 1)
        even = _clear_story(section.even_page_header)
        if isinstance(even_text, str) and even_text:
            even.add_run(even_text)
        changed.append(index)
    return {"changed_sections": len(changed), "section_indices": changed,
            "styleref_operand": operand, "style_evidence": style_evidence,
            "even_header_text": even_text}


def _degree_header_text(spec: dict[str, Any]) -> tuple[str | None, str | None]:
    roles = spec.get("roles", {})
    header = roles.get("header", {}) if isinstance(roles, dict) else {}
    content = header.get("header_content", {}) if isinstance(header, dict) else {}
    raw = content.get("left_text") if isinstance(content, dict) else None
    degree = DEGREE_TEXT.get(spec.get("thesis_profile", {}).get("degree_level"))
    if not isinstance(raw, str) or not degree:
        return raw if isinstance(raw, str) else None, None
    resolved = re.sub(r"博士\s*/\s*硕士|硕士\s*/\s*博士", degree, raw)
    return raw, resolved if resolved != raw else None


def _replace_text_across_nodes(text_nodes: list[Any], old: str, new: str) -> int:
    """Replace an exact phrase even when Word splits it across ``w:t`` nodes.

    Only the differing middle of the old/new strings is rewritten.  This keeps
    stable prefixes and suffixes in their original runs and therefore preserves
    substantially more run formatting than flattening the whole paragraph.
    The caller must first establish that the full old phrase is present.
    """
    visible = "".join(node.text or "" for node in text_nodes)
    starts = [match.start() for match in re.finditer(re.escape(old), visible)]
    if not starts:
        return 0

    prefix_length = 0
    while prefix_length < min(len(old), len(new)) and old[prefix_length] == new[prefix_length]:
        prefix_length += 1
    suffix_length = 0
    while (suffix_length < len(old) - prefix_length
           and suffix_length < len(new) - prefix_length
           and old[-suffix_length - 1] == new[-suffix_length - 1]):
        suffix_length += 1
    old_middle = old[prefix_length:len(old) - suffix_length if suffix_length else len(old)]
    new_middle = new[prefix_length:len(new) - suffix_length if suffix_length else len(new)]
    if not old_middle:
        raise ValueError("cross-node replacement has no non-empty source delta")

    spans = []
    cursor = 0
    for node in text_nodes:
        text = node.text or ""
        spans.append((cursor, cursor + len(text), node))
        cursor += len(text)

    # Reverse order keeps earlier character offsets stable if a paragraph ever
    # contains more than one exact occurrence.
    for phrase_start in reversed(starts):
        start = phrase_start + prefix_length
        end = start + len(old_middle)
        affected = [(a, b, node) for a, b, node in spans if b > start and a < end]
        if not affected:
            raise ValueError("cross-node replacement delta could not be mapped to w:t nodes")
        first_a, _, first = affected[0]
        _, last_b, last = affected[-1]
        first_text = first.text or ""
        last_text = last.text or ""
        before = first_text[:start - first_a]
        after = last_text[len(last_text) - (last_b - end):] if last_b > end else ""
        if first is last:
            first.text = before + new_middle + after
        else:
            first.text = before + new_middle
            for _, _, node in affected[1:-1]:
                node.text = ""
            last.text = after
    return len(starts)


def _heading_style_reference(doc: Document) -> tuple[str, str]:
    """Return a locale-stable STYLEREF operand and the evidence used.

    Word accepts heading-level numbers for built-in Heading styles.  This is
    more portable than an English display name in a localized Word process.
    Custom styles continue to use their actual style name.
    """
    chapter_styles = []
    chapter = re.compile(r"^(?:第\s*[一二三四五六七八九十百零〇两\d]+\s*章|附录\s*[A-Z一二三四五六七八九十\d]+|后\s*记)")
    for paragraph in doc.paragraphs:
        style_id = paragraph.style.style_id or ""
        style_name = paragraph.style.name or ""
        if chapter.match(paragraph.text.strip()) and not (
                re.match(r"^TOC\d*$", style_id, re.I)
                or re.match(r"^toc\s*\d*$", style_name, re.I)
                or re.match(r"^toc\s*\d*$", style_id, re.I)
                or re.match(r"^TOC\d*$", style_name, re.I)
        ):
            chapter_styles.append(paragraph.style)
    if not chapter_styles:
        chapter_styles = [
            paragraph.style for paragraph in doc.paragraphs
            if (paragraph.style.style_id or "") == "Heading1" and paragraph.text.strip()
        ]
    if not chapter_styles:
        raise ValueError("no chapter heading paragraphs were found")
    style_ids = {style.style_id for style in chapter_styles}
    style_names = {style.name for style in chapter_styles}
    normalized_names = {re.sub(r"[\s_-]", "", name or "").casefold() for name in style_names}
    if len(style_ids) != 1 and normalized_names != {"heading1"}:
        raise ValueError(f"chapter headings use multiple styles: {sorted(style_ids)}")
    style = next((item for item in chapter_styles if (item.style_id or "") == "Heading1"), chapter_styles[0])
    match = re.fullmatch(r"Heading([1-9])", style.style_id or "")
    if match:
        return match.group(1), f"built-in style id {style.style_id}"
    if len(style_names) != 1 or not style.name:
        raise ValueError("chapter heading style has no unique usable name")
    escaped = style.name.replace('"', '\\"')
    return f'"{escaped}"', f"custom style name {style.name}"


def repair_headers(doc: Document, spec: dict[str, Any]) -> dict[str, Any]:
    raw_left, resolved_left = _degree_header_text(spec)
    roles = spec.get("roles", {})
    header = roles.get("header", {}) if isinstance(roles, dict) else {}
    content = header.get("header_content", {}) if isinstance(header, dict) else {}
    styleref_required = isinstance(content, dict) and content.get("right_field") == "styleref_heading_1"
    if not resolved_left and not styleref_required:
        return {"changed_parts": 0, "degree_placeholder_replacements": 0, "styleref_fields_rebuilt": 0}
    body_header_mode = bool(spec.get("cleanup", {}).get("body_headers_only"))
    operand, style_evidence = (
        (None, None) if body_header_mode else _heading_style_reference(doc)
    ) if styleref_required else (None, None)
    changed_markers: set[int] = set()
    degree_changes = field_changes = 0

    # Resolve the degree placeholder in every actually serialized header part.
    # Inactive first/even stories can be materialized by Word during a later
    # save, so ignoring them allows a hidden generic placeholder to reappear.
    seen_degree: set[int] = set()
    if resolved_left and raw_left:
        for section in doc.sections:
            for story in (section.header, section.first_page_header, section.even_page_header):
                marker = id(story._element)
                if marker in seen_degree:
                    continue
                seen_degree.add(marker)
                paragraphs = all_story_paragraphs(story)
                before = "\n".join("".join(p._p.xpath(".//w:t/text()")) for p in paragraphs)
                for paragraph in paragraphs:
                    text_nodes = paragraph._p.xpath(".//w:t")
                    visible = "".join(node.text or "" for node in text_nodes)
                    replacements_here = 0
                    for text in text_nodes:
                        if text.text:
                            replaced = text.text.replace(raw_left, resolved_left)
                            if replaced != text.text:
                                text.text = replaced
                                replacements_here += 1
                    if raw_left in visible and replacements_here == 0:
                        replacements_here = _replace_text_across_nodes(text_nodes, raw_left, resolved_left)
                    degree_changes += replacements_here
                after = "\n".join("".join(p._p.xpath(".//w:t/text()")) for p in paragraphs)
                if before != after:
                    changed_markers.add(marker)

    # Preserve the established STYLEREF repair contract: each section gets an
    # independent default header and exactly one proven simple-field location.
    if styleref_required and not body_header_mode:
        seen_fields: set[int] = set()
        for index, section in enumerate(doc.sections):
            _unlink_story_preserving_content(section.header, force_clone=index > 0)
            story = section.header
            marker = id(story._element)
            if marker in seen_fields:
                continue
            seen_fields.add(marker)
            paragraphs = all_story_paragraphs(story)
            simple_targets = [
                paragraph for paragraph in paragraphs
                if paragraph._p.xpath(
                    './/w:fldSimple[contains(translate(@w:instr,"styleref","STYLEREF"),"STYLEREF")]'
                )
            ]
            complex_targets = [
                paragraph for paragraph in paragraphs
                if paragraph._p.xpath(
                    './/w:instrText[contains(translate(text(),"styleref","STYLEREF"),"STYLEREF")]'
                )
            ]
            if complex_targets:
                raise ValueError(
                    f"section {index + 1} header uses a complex STYLEREF field; safe in-place "
                    "reconstruction is not supported"
                )
            if len(simple_targets) != 1:
                raise ValueError(
                    f"section {index + 1} header has {len(simple_targets)} simple STYLEREF locations; "
                    "a unique existing location is required"
                )
            paragraph = simple_targets[0]
            for field in list(paragraph._p.xpath(
                    './/w:fldSimple[contains(translate(@w:instr,"styleref","STYLEREF"),"STYLEREF")]')):
                field.getparent().remove(field)
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), f"STYLEREF {operand} \\* MERGEFORMAT")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = ""
            run.append(text)
            field.append(run)
            paragraph._p.append(field)
            field_changes += 1
            changed_markers.add(marker)
    return {
        "changed_parts": len(changed_markers),
        "degree_placeholder_replacements": degree_changes,
        "resolved_left_text": resolved_left,
        "styleref_fields_rebuilt": field_changes,
        "styleref_operand": operand,
        "style_evidence": style_evidence,
        "body_headers_only_deferred": body_header_mode,
    }


def repair_page_numbering(doc: Document, spec: dict[str, Any]) -> dict[str, Any]:
    rule = spec.get("page", {}).get("page_number", {})
    if not isinstance(rule, dict) or not rule:
        return {"changed_sections": 0}
    with tempfile.NamedTemporaryFile(suffix=".docx") as handle:
        doc.save(handle.name)
        plan = compile_section_plan(spec, handle.name)
    if not plan.get("valid"):
        raise ValueError("compiled SectionPlan is invalid: " + json.dumps(
            plan.get("findings", []), ensure_ascii=False))
    execution = execute_section_plan(doc, plan, alignment=rule.get("alignment"))
    execution["section_plan"] = plan
    execution["changed_sections"] = len(doc.sections)
    return execution


def repair_duplicate_equation_prefixes(doc: Document) -> dict[str, Any]:
    changes = []
    for paragraph_index, paragraph in enumerate(doc.paragraphs):
        children = list(paragraph._p)
        for index, child in enumerate(children):
            instruction = child.get(qn("w:instr")) if child.tag == qn("w:fldSimple") else None
            cached = "".join((node.text or "") for node in child.findall(".//" + qn("w:t")))
            anchor = child.get(qn("w:anchor")) if child.tag == qn("w:hyperlink") else None
            field_reference = bool(instruction and re.search(r"\b(?:REF|SEQ)\b", instruction, re.I))
            equation_hyperlink = bool(anchor and re.search(r"(?:^|_)eq(?:_|$)", anchor, re.I))
            if not field_reference and not equation_hyperlink:
                continue
            if not re.match(r"\s*式\s*[（(]", cached):
                continue
            previous = next((node for node in reversed(children[:index])
                             if node.findall(".//" + qn("w:t"))), None)
            if previous is None:
                continue
            texts = previous.findall(".//" + qn("w:t"))
            target = texts[-1]
            old = target.text or ""
            new = re.sub(r"式[\s\u00a0]*$", "", old, count=1)
            if new != old:
                target.text = new
                changes.append({"paragraph": paragraph_index + 1,
                                "reference": instruction.strip() if instruction else f"bookmark:{anchor}",
                                "removed_prefix": old[len(new):]})
    return {"changes": changes, "count": len(changes)}


def repair_bound_source_markers(doc: Document) -> dict[str, Any]:
    """Replace source-marker caches only when their existing bookmark proves the target."""
    root = doc._element
    bookmarks: dict[str, list[dict[str, str]]] = {}
    for start in root.xpath(".//w:bookmarkStart"):
        name = start.get(qn("w:name"))
        if not name:
            continue
        paragraph = start
        while paragraph is not None and paragraph.tag != qn("w:p"):
            paragraph = paragraph.getparent()
        if paragraph is None:
            continue
        text = "".join(node.text or "" for node in paragraph.findall(".//" + qn("w:t"))).strip()
        match = re.match(r"^(图|表)\s*(\d+(?:[-–—.]\d+)+)", text)
        if match:
            bookmarks.setdefault(name, []).append({"type": match.group(1), "number": match.group(2),
                                                   "caption": text})

    marker_pattern = re.compile(r"^\[(fig|tab|table):([^\]\n）)]*)[\]）)]?$", re.I)
    expected_type = {"fig": "图", "tab": "表", "table": "表"}
    changes = []
    deferred = []
    for paragraph_index, paragraph in enumerate(doc.paragraphs, 1):
        children = list(paragraph._p)
        for index, child in enumerate(children):
            if child.tag != qn("w:hyperlink"):
                continue
            texts = child.findall(".//" + qn("w:t"))
            visible = "".join(node.text or "" for node in texts).strip()
            match = marker_pattern.fullmatch(visible)
            if not match:
                continue
            anchor = child.get(qn("w:anchor"))
            candidates = bookmarks.get(anchor or "", [])
            kind = expected_type[match.group(1).lower()]
            previous = next((node for node in reversed(children[:index])
                             if node.findall(".//" + qn("w:t"))), None)
            previous_text = "" if previous is None else "".join(
                node.text or "" for node in previous.findall(".//" + qn("w:t"))
            )
            reason = None
            if len(candidates) != 1:
                reason = f"bookmark target count is {len(candidates)}"
            elif candidates[0]["type"] != kind:
                reason = f"marker type {kind} conflicts with caption type {candidates[0]['type']}"
            elif not re.search(rf"{re.escape(kind)}[\s\u00a0]*$", previous_text):
                reason = f"preceding context does not supply {kind}"
            if reason:
                deferred.append({"paragraph": paragraph_index, "marker": visible,
                                 "anchor": anchor, "reason": reason})
                continue
            replacement = candidates[0]["number"]
            texts[0].text = replacement
            for extra in texts[1:]:
                extra.text = ""
            changes.append({"paragraph": paragraph_index, "marker": visible, "anchor": anchor,
                            "caption": candidates[0]["caption"], "replacement": replacement})
    return {"changes": changes, "count": len(changes), "deferred": deferred}


def repair(input_path: Path, spec: dict[str, Any], output_path: Path) -> dict[str, Any]:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("output must differ from input")
    doc = Document(input_path)
    report = {
        "schema_version": "1.0",
        "evidence_type": "deterministic_docx_repair",
        "input": {"path": str(input_path.resolve()), "sha256": digest(input_path)},
        "repairs": {},
        "deferred": [
            {"code": "chapter_structure_requires_semantic_review",
             "reason": "Chapter declarations are not rewritten from text heuristics."},
        ],
        "submission_ready_claimed": False,
    }
    cleanup = spec.get("cleanup", {}) if isinstance(spec.get("cleanup"), dict) else {}
    report["repairs"]["declared_ranges"] = cleanup_declared_ranges(
        doc, cleanup.get("remove_ranges", [])) if cleanup.get("remove_ranges") else {
            "range_count": 0, "removed_body_children": 0
        }
    report["repairs"]["heading_1_boundaries"] = normalize_heading1_boundaries(
        doc, cleanup.get("heading_1_texts", [])) if cleanup.get("heading_1_texts") else {
            "count": 0, "changes": []
        }
    report["repairs"]["declared_header_texts"] = cleanup_declared_header_texts(
        doc, cleanup.get("remove_header_texts", [])) if cleanup.get("remove_header_texts") else {
            "count": 0, "changes": []
        }
    report["repairs"]["headers"] = repair_headers(doc, spec)
    report["repairs"]["page_numbering"] = repair_page_numbering(doc, spec)
    report["repairs"]["body_headers"] = configure_body_headers(doc, spec) if cleanup.get(
        "body_headers_only") else {"changed_sections": 0, "skipped": True}
    report["repairs"]["equation_reference_prefixes"] = repair_duplicate_equation_prefixes(doc)
    marker_repairs = repair_bound_source_markers(doc)
    report["repairs"]["bound_source_markers"] = marker_repairs
    if marker_repairs["deferred"]:
        report["deferred"].append({
            "code": "source_marker_requires_target_resolution",
            "reason": "One or more markers lacked a unique type-consistent bound caption target.",
            "evidence": marker_repairs["deferred"],
        })
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    report["repairs"]["comments"] = (
        remove_docx_comments(output_path) if cleanup.get("remove_comments") else {
            "skipped": True, "reason": "cleanup.remove_comments is not enabled"
        }
    )
    page_repair = report["repairs"].get("page_numbering", {})
    plan = page_repair.get("section_plan") if isinstance(page_repair, dict) else None
    if isinstance(plan, dict):
        findings = audit_plan_against_docx(plan, output_path)
        page_repair["audit"] = {"valid": not findings, "findings": findings}
    report["output"] = {"path": str(output_path.resolve()), "sha256": digest(output_path),
                        "size_bytes": output_path.stat().st_size}
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("format_spec", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    resolved_paths = {
        "input": args.input.expanduser().resolve(),
        "output": args.output.expanduser().resolve(),
        "report": args.report.expanduser().resolve(),
    }
    if len(set(resolved_paths.values())) != len(resolved_paths):
        parser.error("input DOCX, output DOCX, and repair report must be three distinct paths")
    spec = json.loads(args.format_spec.read_text(encoding="utf-8"))
    schema_path = Path(__file__).resolve().parents[1] / "schema" / "format-spec.schema.json"
    schema_errors = load_and_validate(spec, schema_path)
    if schema_errors:
        parser.error("invalid format spec:\n" + "\n".join(f"- {error}" for error in schema_errors))
    report = repair(resolved_paths["input"], spec, resolved_paths["output"])
    resolved_paths["report"].parent.mkdir(parents=True, exist_ok=True)
    resolved_paths["report"].write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
