#!/usr/bin/env python3
"""Shared OOXML-first semantic detectors for DOCX pipeline stages."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from collections.abc import Iterable, Iterator
from typing import Any

from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
try:
    from .role_registry import style_aliases
except ImportError:  # direct script execution
    from role_registry import style_aliases


@dataclass(frozen=True)
class DocumentNode:
    """Stable, shared description of a paragraph in a DOCX story/container."""

    node_id: str
    story: str
    container: str
    paragraph: Any
    xml_order: int


def _walk_table(table: Any, prefix: str, story: str, order: list[int], seen: set[int]) -> Iterator[DocumentNode]:
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            marker = id(cell._tc)
            if marker in seen:
                continue
            seen.add(marker)
            for child_index, child in enumerate(cell._tc.iterchildren()):
                if child.tag.endswith("}p"):
                    paragraph = Paragraph(child, cell)
                    yield DocumentNode(
                        f"{prefix}-row-{row_index}-cell-{cell_index}-paragraph-{child_index}",
                        story, "table_cell", paragraph, order[0],
                    )
                    order[0] += 1
                elif child.tag.endswith("}tbl"):
                    nested = Table(child, cell)
                    yield from _walk_table(
                        nested,
                        f"{prefix}-row-{row_index}-cell-{cell_index}-table-{child_index}",
                        story, order, seen,
                    )


def iter_document_nodes(doc: Any, *, include_stories: bool = False) -> Iterator[DocumentNode]:
    """Walk body paragraphs and tables in XML order, including nested tables.

    This is the canonical traversal used by both mutation and validation.  It
    deliberately avoids ``document.paragraphs`` because that API omits table
    cells and loses body XML order.
    """
    order = [0]
    seen_cells: set[int] = set()
    body = doc._body
    table_index = 0
    paragraph_index = 0
    for child in body._element.iterchildren():
        if child.tag.endswith("}p"):
            yield DocumentNode(
                f"body-paragraph-{paragraph_index}", "document", "body",
                Paragraph(child, body), order[0],
            )
            paragraph_index += 1
            order[0] += 1
        elif child.tag.endswith("}tbl"):
            yield from _walk_table(Table(child, body), f"body-table-{table_index}", "document", order, seen_cells)
            table_index += 1
    if include_stories:
        seen_parts: set[str] = set()
        for section_index, section in enumerate(doc.sections):
            for kind, story in (("header", section.header), ("header_first", section.first_page_header),
                                ("header_even", section.even_page_header), ("footer", section.footer),
                                ("footer_first", section.first_page_footer), ("footer_even", section.even_page_footer)):
                partname = str(story.part.partname)
                if partname in seen_parts:
                    continue
                seen_parts.add(partname)
                for index, paragraph in enumerate(story.paragraphs):
                    yield DocumentNode(f"section-{section_index}-{kind}-paragraph-{index}", kind, kind,
                                       paragraph, order[0])
                    order[0] += 1


def all_body_paragraphs(doc: Any) -> Iterator[Any]:
    yield from (node.paragraph for node in iter_document_nodes(doc))


def all_table_paragraphs(doc: Any) -> Iterator[Any]:
    yield from (node.paragraph for node in iter_document_nodes(doc) if node.container == "table_cell")


def all_story_paragraphs(story: Any) -> Iterator[Any]:
    yield from story.paragraphs
    for table in story.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def has_visible_or_object_content(paragraph: Any) -> bool:
    return bool(paragraph.text.strip() or paragraph._p.xpath(".//w:drawing|.//m:oMath|.//m:oMathPara"))


def has_drawing(paragraph: Any) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def is_display_equation(paragraph: Any) -> bool:
    return bool(paragraph._p.xpath(".//m:oMathPara"))


def is_table_cell_paragraph(paragraph: Any) -> bool:
    return bool(paragraph._p.xpath("ancestor::w:tc"))


def is_figure_caption(paragraph: Any) -> bool:
    text = paragraph.text.strip()
    # A prose reference such as “图 2.2 展示了……” is not a caption. Require
    # either a registered caption style or the generated caption separator.
    style = getattr(getattr(paragraph, "style", None), "name", "")
    return bool(
        re.match(r"^(图|Figure)\s*\d+[.．-]?\d*\s{2,}\S", text, re.I)
        or (re.match(r"^(图|Figure)\s*\d+", text, re.I)
            and re.search(r"caption|题注|图题", style, re.I))
    )


def is_table_caption(paragraph: Any) -> bool:
    text = paragraph.text.strip()
    style = getattr(getattr(paragraph, "style", None), "name", "")
    return bool(
        re.match(r"^(表|Table)\s*\d+[.．-]?\d*\s{2,}\S", text, re.I)
        or (re.match(r"^(表|Table)\s*\d+", text, re.I)
            and re.search(r"caption|题注|表题", style, re.I))
    )


def is_abstract_title_zh(paragraph: Any) -> bool:
    return bool(re.match(r"^摘\s*要$", paragraph.text.strip()))


def is_abstract_title_en(paragraph: Any) -> bool:
    return paragraph.text.strip().lower() == "abstract"


def is_keywords_zh(paragraph: Any) -> bool:
    return bool(re.match(r"^关\s*键\s*词\s*[：:]", paragraph.text.strip()))


def is_keywords_en(paragraph: Any) -> bool:
    return bool(re.match(r"^key\s*words?\s*:", paragraph.text.strip(), re.I))


def is_abstract_body_zh(paragraph: Any) -> bool:
    return paragraph.style.name in {"AbstractBodyCN", "Abstract Body CN", "中文摘要正文"}


def is_abstract_body_en(paragraph: Any) -> bool:
    if is_keywords_en(paragraph):
        return False
    return paragraph.style.name in {"AbstractBodyEN", "Abstract Body EN", "英文摘要正文"}


def is_toc_heading(paragraph: Any) -> bool:
    return bool(re.match(r"^目\s*录$", paragraph.text.strip()))


def has_toc_field(paragraph: Any) -> bool:
    """Return whether a paragraph contains a real Word TOC field.

    A TOC heading is only a label.  The executable/content role belongs to the
    paragraph containing the TOC field, even while Word has not materialized
    its cached entries yet.  Support both complex fields (``w:instrText``)
    and simple fields (``w:fldSimple``), as both occur in DOCX producers.
    """
    instructions = list(paragraph._p.xpath(".//w:instrText/text()"))
    instructions.extend(paragraph._p.xpath(".//w:fldSimple/@w:instr"))
    return any(re.search(r"""(?:^|[\s"'])TOC(?:$|[\s"'])""", str(value), re.I)
               for value in instructions)


def is_special_heading(paragraph: Any, role: str) -> bool:
    text = normalized_fixed_text(paragraph.text).strip("：:")
    patterns = {
        "heading_acknowledgments": r"^(致谢|后记)$",
        "heading_appendix": r"^(?:(?:第[一二三四五六七八九十百千万\d]+章)|(?:\d+(?:\.\d+)*))?附录(?:[A-ZＡ-Ｚ一二三四五六七八九十\d])?(?:[:：].*)?.*$",
        "heading_conclusion": r"^(结论|结语)$",
        "heading_publications": r"^(?:在学期间发表的)?(?:学术论文与研究成果|发表论文和科研情况|科研成果|研究成果|攻读学位期间取得的成果)$",
        "heading_references": r"^(?:(?:第[一二三四五六七八九十百千万\d]+章)|(?:\d+(?:\.\d+)*))?参考文献$",
    }
    return bool(re.match(patterns.get(role, r"a^"), text, re.I))


def normalized_fixed_text(text: str) -> str:
    return re.sub(r"\s+", "", text)


def fixed_text_sha256(text: str) -> str:
    return hashlib.sha256(normalized_fixed_text(text).encode("utf-8")).hexdigest()


def is_declaration_heading(paragraph: Any) -> bool:
    text = normalized_fixed_text(paragraph.text)
    return bool(
        re.search(r"(?:独创性|原创性)声明", text)
        or re.search(r"(?:学位论文|论文).*(?:授权|版权)", text)
    )


def is_declaration_body(paragraph: Any) -> bool:
    return paragraph.style.name in set(style_aliases("declaration_body"))


def is_signature_line_author(paragraph: Any) -> bool:
    return bool(re.search(r"(?:研究生|作者)签名\s*[：:]", paragraph.text))


def is_signature_line_supervisor(paragraph: Any) -> bool:
    return bool(re.search(r"(?:指导教师|导师)签名\s*[：:]", paragraph.text))


def is_signature_date_line(paragraph: Any) -> bool:
    return bool(re.search(r"日期\s*[：:]", paragraph.text))


def is_cover_field_label(paragraph: Any) -> bool:
    return paragraph.style.name in set(style_aliases("cover_field_label"))


def is_cover_field_value(paragraph: Any) -> bool:
    return paragraph.style.name in set(style_aliases("cover_field_value"))


def matches_structural_detector(paragraph: Any, detector: str | None) -> bool:
    if detector == "table_cell_paragraph":
        return is_table_cell_paragraph(paragraph) and has_visible_or_object_content(paragraph)
    if detector == "display_math":
        return is_display_equation(paragraph)
    if detector == "figure_caption":
        return is_figure_caption(paragraph)
    if detector == "table_caption":
        return is_table_caption(paragraph)
    if detector == "abstract_title_zh":
        return is_abstract_title_zh(paragraph)
    if detector == "abstract_title_en":
        return is_abstract_title_en(paragraph)
    if detector == "abstract_body_zh":
        return is_abstract_body_zh(paragraph)
    if detector == "abstract_body_en":
        return is_abstract_body_en(paragraph)
    if detector == "frontmatter_title_zh":
        return bool(paragraph.text.strip()) and paragraph.style.name in set(style_aliases("thesis_title_zh"))
    if detector == "frontmatter_title_en":
        return bool(paragraph.text.strip()) and paragraph.style.name in set(style_aliases("thesis_title_en"))
    if detector == "keywords_zh":
        return is_keywords_zh(paragraph)
    if detector == "keywords_en":
        return is_keywords_en(paragraph)
    if detector == "toc_heading":
        return is_toc_heading(paragraph)
    if detector in {
        "heading_acknowledgments", "heading_appendix", "heading_conclusion",
        "heading_publications", "heading_references",
    }:
        return is_special_heading(paragraph, detector)
    if detector == "declaration_heading":
        return is_declaration_heading(paragraph)
    if detector == "declaration_body":
        return is_declaration_body(paragraph)
    if detector == "signature_line_author":
        return is_signature_line_author(paragraph)
    if detector == "signature_line_supervisor":
        return is_signature_line_supervisor(paragraph)
    if detector == "signature_date_line":
        return is_signature_date_line(paragraph)
    if detector == "cover_field_label":
        return is_cover_field_label(paragraph)
    if detector == "cover_field_value":
        return is_cover_field_value(paragraph)
    return False


def count_styles(paragraphs: Iterable[Any], detector: str | None) -> dict[str, int]:
    counts: dict[str, int] = {}
    for paragraph in paragraphs:
        if matches_structural_detector(paragraph, detector):
            counts[paragraph.style.name] = counts.get(paragraph.style.name, 0) + 1
    return counts


def dominant_structural_style(paragraphs: Iterable[Any], detector: str | None) -> str | None:
    counts = count_styles(paragraphs, detector)
    return min(counts, key=lambda name: (-counts[name], name)) if counts else None
