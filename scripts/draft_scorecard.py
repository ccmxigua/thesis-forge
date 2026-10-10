"""Evidence-based draft grading; never a compliance or release certificate."""
from __future__ import annotations

import copy
import hashlib
import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from manual_review_display import (ensure_manual_review_styles, mark_manual_review_paragraph,
                                   MANUAL_REVIEW_STYLE, _effective_run_attribute)
from docx_semantics import all_body_paragraphs
from semantic_contract import sha256_json, strict_json_read
from pipeline_finding import integrity_findings
from responsibility_ledger import WEIGHTS, FORCES, APPLICABILITIES, ROUTES
from source_obligation_assessment import assessment_projection_valid, build_obligation_assessment

POLICY = "source_obligation_draft_scorecard_v3"
PREFIX = "【SC-"
STATUSES = ("failed", "unverified", "pending", "verified")


def _entry_score(status: str) -> int | None:
    """Only observed outcomes receive a numeric diagnostic score."""
    if status == "verified":
        return 100
    if status == "failed":
        return 0
    if status in {"unverified", "pending"}:
        return None
    raise ValueError("unknown scorecard status")


def _scorecard_semantics_valid(card: dict[str, Any]) -> bool:
    entries = card.get("entries")
    if not isinstance(entries, list) or any(not isinstance(item, dict) for item in entries):
        return False
    expected_counts = {status: 0 for status in STATUSES}
    for item in entries:
        status = item.get("status")
        if (not isinstance(status, str) or status not in expected_counts
                or item.get("score") != _entry_score(status)):
            return False
        expected_counts[status] += 1
    if card.get("status_counts") != expected_counts:
        return False
    if (card.get("item_count") != len(entries)
            or card.get("diagnostic_item_count") != len(entries)
            or card.get("verified_count") != expected_counts["verified"]
            or card.get("failed_count") != expected_counts["failed"]
            or card.get("unverified_count") != expected_counts["unverified"]
            or card.get("pending_count") != expected_counts["pending"]):
        return False
    history = card.get("semantic_history")
    if history is not None:
        if not isinstance(history, dict) or not isinstance(history.get("status_counts"), dict):
            return False
        historical = {status: 0 for status in STATUSES}
        for item in entries:
            old = item.get("historical_status")
            if old is None:
                continue  # Current rendered-output findings have no semantic predecessor.
            if old not in historical:
                return False
            historical[old] += 1
        if historical != history["status_counts"]:
            return False
        current = card.get("current_output_assessment")
        expected_current_status = (
            "issues_found" if expected_counts["failed"] else
            "unverified" if expected_counts["unverified"] or expected_counts["pending"] else
            "passed"
        )
        if (not isinstance(current, dict)
                or current.get("protocol") != "current_docx_rendered_instance_assessment_v1"
                or current.get("status_counts") != expected_counts
                or current.get("status") != expected_current_status
                or current.get("submission_ready") is not False
                or not isinstance(current.get("docx_sha256"), str)
                or not isinstance(current.get("rendered_report_sha256"), str)):
            return False
        attachment = card.get("rendered_output_audit")
        if (isinstance(attachment, dict)
                and current.get("rendered_report_sha256") != attachment.get("report_sha256")):
            return False
    if "obligation_assessment" in card:
        scope_assessment_valid = assessment_projection_valid(card.get("obligation_assessment"))
    else:
        # The original v3 scorecard contract predates object-scope assessment.
        # Keep its display audit readable while explicitly leaving scope
        # assessment absent; this compatibility path never grants release.
        scope_assessment_valid = (
            card.get("schema_version") in {"1.0", "3.0"}
            and card.get("policy") == POLICY
            and card.get("unique_source_obligation_count") is None
        )
    return card.get("submission_ready") is False and scope_assessment_valid


def build_scorecard(binding: dict[str, Any], receipt_audit: dict[str, Any],
                    manual_items: list[dict[str, Any]], findings: list[dict[str, Any]],
                    capability_findings: list[dict[str, Any]] | None = None,
                    *,
                    evaluation_units_by_requirement: dict[str, list[dict[str, Any]]] | None = None,
                    obligation_scope_inventory: dict[str, Any] | None = None,
                    expected_docx_sha256: str | None = None,
                    ) -> dict[str, Any]:
    """Known automatic source units count once; unknown semantics stay unweighted.

    Legacy check percent remains a separately named diagnostic, never silently
    upgraded to a normative score. A shared source unit requires explicit,
    code-generated identity; wording, role, confidence and array position do
    not establish equivalence. Every raw check remains independently visible.
    """
    if not isinstance(binding, dict) or not binding.get("run_id") or not binding.get("case_id"):
        raise ValueError("scorecard requires current run/case binding")
    for key in ("source_sha256", "format_spec_sha256"):
        if not isinstance(binding.get(key), str) or not re.fullmatch(r"[0-9a-f]{64}", binding[key]):
            raise ValueError("scorecard input hash missing: " + key)
    if any(receipt_audit.get(key) != 0 for key in ("missing_count", "unexpected_count", "duplicate_count")):
        raise ValueError("scorecard cannot bypass receipt inventory integrity")
    receipts = receipt_audit.get("receipts")
    expected = receipt_audit.get("expected_receipt_ids")
    if not isinstance(receipts, list) or not isinstance(expected, list):
        raise ValueError("scorecard requires complete expected receipt inventory")
    if any(not isinstance(value, str) or not value for value in expected):
        raise ValueError("scorecard expected receipt IDs are invalid")
    for rows in (manual_items, findings, capability_findings or []):
        if not isinstance(rows, list) or any(not isinstance(item, dict) for item in rows):
            raise ValueError("scorecard inventory is not a record list")
    if integrity_findings([*findings, *(capability_findings or [])]):
        raise ValueError("scorecard cannot bypass an integrity gate")
    if evaluation_units_by_requirement is not None and (
        not isinstance(evaluation_units_by_requirement, dict)
        or any(not isinstance(key, str) or not isinstance(value, list)
               for key, value in evaluation_units_by_requirement.items())
    ):
        raise ValueError("authoritative evaluation-unit inventory is invalid")
    ids = [item.get("receipt_id") for item in receipts if isinstance(item, dict)]
    if (len(ids) != len(receipts) or any(not isinstance(x, str) or not x for x in ids)
            or len(set(ids)) != len(ids) or len(set(expected)) != len(expected) or set(ids) != set(expected)):
        raise ValueError("scorecard receipt identities do not match")
    entries: list[dict[str, Any]] = []
    units: dict[str, dict[str, Any]] = {}
    diagnostic_keys: set[str] = set()

    def attach_unit(entry: dict, unit: dict) -> None:
        uid = unit["evaluation_unit_id"]
        if uid in units and units[uid]["identity"] != unit:
            raise ValueError("conflicting score unit metadata")
        if uid not in units:
            units[uid] = {"evaluation_unit_id": uid, "identity": copy.deepcopy(unit), "item_ids": [], "statuses": []}
        units[uid]["item_ids"].append(entry["item_id"])
        units[uid]["statuses"].append(entry["status"])
        entry.setdefault("evaluation_unit_ids", []).append(uid)

    def add(kind: str, identity: Any, label: str, status: str, detail: Any) -> dict | None:
        key = sha256_json({"kind": kind, "identity": identity})
        if key in diagnostic_keys:
            return None  # Exact duplicate diagnostics do not inflate any denominator.
        diagnostic_keys.add(key)
        entry = {
            "item_id": "SC-" + sha256_json({"kind": kind, "identity": identity, "binding": binding})[:24],
            "kind": kind, "label": label, "status": status,
            "score": _entry_score(status), "max_score": 100,
            "human_check_required": status != "verified", "detail": copy.deepcopy(detail),
        }
        entries.append(entry)
        return entry

    if evaluation_units_by_requirement is not None:
        for requirement_id, expected_units in evaluation_units_by_requirement.items():
            for unit in expected_units:
                if not isinstance(unit, dict) or not isinstance(unit.get("evaluation_unit_id"), str):
                    raise ValueError("authoritative evaluation unit is invalid")
                uid = unit["evaluation_unit_id"]
                if uid in units and units[uid]["identity"] != unit:
                    raise ValueError("conflicting authoritative evaluation unit metadata")
                units.setdefault(uid, {
                    "evaluation_unit_id": uid, "identity": copy.deepcopy(unit),
                    "item_ids": [], "statuses": [],
                })

    for item in receipts:
        if item.get("status") not in {"verified", "failed", "unverified"}:
            raise ValueError("unknown property receipt status")
        record = {key: copy.deepcopy(value) for key, value in item.items() if key != "serialized_docx_sha256"}
        entry = add("property", item["receipt_id"], str(item.get("role", "")) + "." + str(item.get("property_path", "")), item["status"], record)
        declared = item.get("evaluation_units")
        if declared is not None and (not isinstance(declared, list) or not declared):
            raise ValueError("invalid evaluation unit inventory")
        requirement_id = item.get("requirement_id")
        if declared is not None:
            if evaluation_units_by_requirement is None:
                raise ValueError("typed score units require a verified current-run source ledger")
            expected_units = evaluation_units_by_requirement.get(requirement_id)
            if not expected_units or declared != expected_units:
                raise ValueError("property receipt units do not match the verified source ledger")
        elif (evaluation_units_by_requirement is not None
              and requirement_id in evaluation_units_by_requirement):
            raise ValueError("property receipt omitted verified source units")
        for unit in declared or []:
            if (not isinstance(unit, dict) or unit.get("force") not in FORCES
                    or unit.get("route") not in ROUTES or unit.get("applicability") not in APPLICABILITIES
                    or unit.get("source_sha256") != binding["source_sha256"]
                    or not isinstance(unit.get("canonical_obligation_key"), str)
                    or not re.fullmatch(r"[0-9a-f]{64}", unit["canonical_obligation_key"])
                    or unit.get("evaluation_unit_id") != "EU-" + unit["canonical_obligation_key"][:32]
                    or not re.fullmatch(r"[0-9a-f]{64}", str(unit.get("source_span_sha256", "")))
                    or unit.get("semantic_basis") not in {"typed_source_atom", "legacy_unknown_dimensions"}
                    or (unit["semantic_basis"] != "typed_source_atom" and unit["force"] != "unknown")):
                raise ValueError("score unit is not bound to the current source atom")
            attach_unit(entry, unit)
        if not declared:
            # Requirement IDs are only legacy grouping, never semantic identity.
            # Keep them unweighted; splitting properties cannot create points.
            key = sha256_json({"legacy_requirement": item.get("requirement_id", item["receipt_id"]),
                               "source_sha256": binding["source_sha256"]})
            attach_unit(entry, {"evaluation_unit_id": "EU-" + key[:32], "canonical_obligation_key": key,
                               "force": "unknown", "applicability": "unknown", "route": "unknown",
                               "semantic_basis": "legacy_unknown_dimensions"})
    if evaluation_units_by_requirement is not None:
        for uid, row in units.items():
            if row["identity"].get("route") == "human":
                # A property receipt can verify a document property, but it
                # cannot complete a source/provenance decision. Keep this unit
                # pending and visible even if an associated property receipt
                # says "verified"; only a separate human-verification receipt
                # may eventually resolve that obligation.
                diagnostic = add(
                    "human_review", uid,
                    "来源或适用性需要人工核验", "pending",
                    {"evaluation_unit_id": uid, **copy.deepcopy(row["identity"])},
                )
                if diagnostic is not None:
                    row["item_ids"].append(diagnostic["item_id"])
                row["statuses"].append("pending")
                continue
            if row["statuses"]:
                continue
            diagnostic = add(
                "source_obligation", uid,
                "来源义务缺少可核验的文档属性回执", "unverified",
                {"evaluation_unit_id": uid, **copy.deepcopy(row["identity"])},
            )
            if diagnostic is not None:
                row["item_ids"].append(diagnostic["item_id"])
                row["statuses"].append("unverified")
        authoritative_ids = {
            unit.get("evaluation_unit_id")
            for rows in evaluation_units_by_requirement.values()
            for unit in rows if isinstance(unit, dict)
        }
        if any(uid not in authoritative_ids for uid, row in units.items()
               if row["identity"].get("semantic_basis") == "typed_source_atom"):
            raise ValueError("scorecard contains a typed unit outside the verified source inventory")
    for item in manual_items:
        add("human_review", item, str(item.get("source_code") or item.get("source_text") or "人工核验"), "pending", item)
    for item in findings:
        # This is a known observability gap, not evidence that the document's
        # drawing/math text uses the wrong font. Keep it visible and blocking
        # in the underlying audit, while withholding a compliance score.
        if item.get("failure_type") == "document_font_not_observable":
            add("capability", item,
                str(item.get("role", "")) + "." + str(item.get("property") or item.get("failure_type")),
                "unverified", item)
        else:
            add("finding", item,
                str(item.get("role", "")) + "." + str(item.get("property") or item.get("failure_type") or "未满足检查"),
                "failed", item)
    for item in capability_findings or []:
        add("capability", item, str(item.get("code") or "能力预检待处理"), "unverified", item)
    if len({item["item_id"] for item in entries}) != len(entries):
        raise ValueError("duplicate scorecard item")
    # Proven mismatches precede unknown and human work. Stable order preserves
    # source/receipt order within each state; unknown work never sorts by a
    # fabricated zero.
    status_order = {status: index for index, status in enumerate(STATUSES)}
    entries.sort(key=lambda item: status_order[item["status"]])
    status_counts = {status: sum(item["status"] == status for item in entries)
                     for status in STATUSES}
    earned = status_counts["verified"]
    possible = 100 * len(entries)
    source_units = []
    for uid in sorted(units):
        row = units[uid]
        identity, statuses = row["identity"], row.pop("statuses")
        status = ("pending" if identity["route"] == "human" else
                  "failed" if "failed" in statuses else "unverified" if "unverified" in statuses else "verified")
        bucket = ("H" if identity["route"] == "human" else "I" if identity["route"] == "input" else
                  "F" if identity["route"] in {"automatic", "check_only"} and identity["force"] != "unknown"
                  and identity["applicability"] == "applicable" else "U")
        weight = WEIGHTS[identity["force"]] if bucket == "F" else None
        source_units.append({**row, "ledger": bucket, "status": status,
                             "weight": weight, "earned_weight": weight if weight and status == "verified" else 0})
    total_weight = sum(u["weight"] or 0 for u in source_units)
    earned_weight = sum(u["earned_weight"] for u in source_units)
    ledgers = {letter: [u["evaluation_unit_id"] for u in source_units if u["ledger"] == letter] for letter in "FHIU"}
    ledgers["H"].extend(e["item_id"] for e in entries if e["kind"] == "human_review")
    # Findings without a proved unit reference are diagnostics, not new duties.
    ledgers["U"].extend(e["item_id"] for e in entries if e["kind"] in {"finding", "capability"})
    obligation_assessment = build_obligation_assessment(
        binding, evaluation_units_by_requirement, receipt_audit, obligation_scope_inventory,
        expected_docx_sha256,
    )
    return {
        "schema_version": "3.0", "policy": POLICY, "output_policy": "review_draft",
        "binding": copy.deepcopy(binding), "submission_ready": False,
        "human_check_required": True, "coverage_complete": False,
        "score_meaning": "verified_weight_of_known_applicable_automatic_source_units_not_release",
        "weighting": "product_policy_required_prohibited_5_recommended_2_optional_1_unknown_unweighted",
        "weight_policy": copy.deepcopy(WEIGHTS), "weight_policy_is_school_rule": False,
        "score": round(earned_weight * 100 / total_weight, 2) if total_weight else None,
        "earned_weight": earned_weight, "total_weight": total_weight,
        "observed_checks_verified_percent": round(earned * 100 / len(entries), 2) if entries else None,
        "evaluation_units": source_units, "ledgers": ledgers,
        "minimum_draft_score": None,
        "status_counts": status_counts,
        "verified_count": status_counts["verified"],
        "failed_count": status_counts["failed"],
        "unverified_count": status_counts["unverified"],
        "pending_count": status_counts["pending"],
        "item_count": len(entries), "diagnostic_item_count": len(entries),
        "evaluation_unit_count": len(source_units),
        "typed_source_unit_count": sum(
            unit["identity"].get("semantic_basis") == "typed_source_atom" for unit in source_units
        ),
        "legacy_unknown_unit_count": sum(
            unit["identity"].get("semantic_basis") == "legacy_unknown_dimensions" for unit in source_units
        ),
        "unique_source_obligation_count": (
            sum(unit["identity"].get("semantic_basis") == "typed_source_atom" for unit in source_units)
            if any(unit["identity"].get("semantic_basis") == "typed_source_atom" for unit in source_units)
            else None
        ),
        "diagnostic_count_meaning": (
            "Deduplicated diagnostic entries, not a count of independent source obligations. "
            "Legacy unknown-dimension evaluation units do not establish unique obligation identity."
        ),
        "obligation_assessment": obligation_assessment,
        "entries": entries,
    }


def scorecard_lines(card: dict[str, Any]) -> list[str]:
    scope_line = (_obligation_assessment_line(card.get("obligation_assessment")) + "\n"
                  if "obligation_assessment" in card else "")
    current = card.get("current_output_assessment")
    if isinstance(current, dict):
        history = card.get("semantic_history", {})
        historic_counts = history.get("status_counts", {}) if isinstance(history, dict) else {}
        head = (
            "当前生成稿实例核验（审查草稿，不可提交；不是来源语义审查或提交许可）\n"
            f"当前 DOCX 实例：已验证{card['verified_count']}项，已证实失败{card['failed_count']}项，"
            f"当前未核验{card['unverified_count']}项，待人工核验{card['pending_count']}项；"
            f"诊断对象{card['diagnostic_item_count']}项。\n"
            f"继承历史实例状态（仅供追溯）：已验证{historic_counts.get('verified', 0)}项，"
            f"失败{historic_counts.get('failed', 0)}项，未核验{historic_counts.get('unverified', 0)}项，"
            f"待人工{historic_counts.get('pending', 0)}项。\n"
            "历史 verified 不代表当前 DOCX 已验证；当前实例须有匹配 DOCX 哈希和对象/义务绑定的证据。\n"
            + ("历史" + scope_line if scope_line else "")
            + "当前实例核验不计算来源义务加权分；未知或缺少精确绑定的项目保持未核验，submission_ready=false。"
        )
    else:
        head = ("自动核验评分（审查草稿，不可提交）\n"
                f"得分：{card['score'] if card['score'] is not None else '未评分'}/100；"
                f"诊断条目{card['diagnostic_item_count']}项（不是独立义务数）；"
                f"已验证{card['verified_count']}项，已证实未满足{card['failed_count']}项，"
                f"未核验{card['unverified_count']}项，待人工核验{card['pending_count']}项。\n"
                f"来源评价单元{card['evaluation_unit_count']}个，其中可确认独立来源义务"
                f"{card['unique_source_obligation_count'] if card['unique_source_obligation_count'] is not None else '未知'}个。\n"
                + scope_line
                + "按已知适用的自动义务单元计分：强制/禁止5、建议2、已选择可选项1；这是产品策略，不是学校权重。"
                + "未知强制性/适用性不赋权，人工、待输入、未知问题单列；重复检查不重复加权。"
                + "所有未满足/待核验项须人工检查；高分不能抵消安全错误或授权提交。")
    statuses = {"verified": "已验证满足", "failed": "未满足", "unverified": "未核验", "pending": "待人工核验"}
    def details(item: dict[str, Any]) -> str:
        record = item["detail"]
        # Full producer records stay in JSON, not an unbounded paragraph in
        # the document. Render the actionable evidence, never guessed roles.
        fields = {key: record[key] for key in (
            "requirement_id", "clause_ids", "evidence_ids", "source_text", "source_code",
            "property_path", "property", "expected", "actual", "required_value",
            "evaluation_unit_id", "canonical_obligation_key",
            "template_value", "reason", "status_reason", "action", "message", "status",
            "current_output_instance",
        ) if key in record}
        instance = fields.get("current_output_instance")
        if isinstance(instance, dict):
            # Keep visible reasons and exact finding identities stable while
            # the DOCX/PDF hashes change during render/re-audit convergence.
            fields["current_output_instance"] = {
                key: value for key, value in instance.items()
                if key not in {"docx_sha256", "pdf_sha256", "rendered_report_sha256",
                               "serialized_docx_sha256"}
            }
        if item.get("historical_status") is not None:
            fields["historical_status"] = item["historical_status"]
        return json.dumps(fields, ensure_ascii=False, sort_keys=True)
    def score_text(item: dict[str, Any]) -> str:
        if item["status"] == "verified":
            return "100/100 已验证满足"
        if item["status"] == "failed":
            return "0/100 已证实未满足"
        if item["status"] == "unverified":
            return "未评分（未核验）"
        return "未评分（待人工核验）"
    return [head] + [
        f"【{item['item_id']}｜{'当前实例核验' if isinstance(current, dict) else '自动评分'}】{score_text(item)}：{item['label']}\n"
        + "具体检查/要求/实际值及原因：" + details(item)
        + ("\n处理：请对照原文和文档人工核验、修正后重新检查。" if item["human_check_required"] else "")
        for item in card["entries"]
    ]


def _obligation_assessment_line(assessment: dict[str, Any] | None) -> str:
    if not isinstance(assessment, dict):
        return "对象级来源义务：未评估。"
    if assessment.get("status") == "scope_inventory_missing":
        return (
            "对象级来源义务：范围清单缺失，不计算覆盖率或满足率；"
            f"旧式未知来源单元{assessment.get('legacy_unknown_unit_count', 0)}个，"
            f"结构化来源单元{assessment.get('typed_source_unit_count', 0)}个。"
        )
    if assessment.get("assessment_coverage") is None:
        return (
            "对象级来源义务：范围或适用性未完整确认，不计算覆盖率或满足率；"
            f"清单完整={assessment.get('inventory_complete') is True}，"
            f"旧式未知来源单元{assessment.get('legacy_unknown_unit_count', 0)}个。"
        )
    coverage = f"{assessment['assessment_coverage'] * 100:.1f}%"
    satisfaction = assessment.get("satisfaction_ratio")
    if satisfaction is None:
        return f"对象级核验覆盖率{coverage}；仍有未核验/待人工/未知项，不宣称部分满足或完整满足。"
    return (
        f"对象级核验覆盖率{coverage}；完整核验范围内满足率"
        f"{satisfaction * 100:.1f}%（不代表学校评分或发布授权）。"
    )


def append_scorecard(doc: Any, card: dict[str, Any]) -> None:
    ensure_manual_review_styles(doc)
    # A separate document-front section; never a guessed thesis role/location.
    anchor = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    for text in scorecard_lines(card):
        element = OxmlElement("w:p")
        anchor._p.addprevious(element)
        paragraph = Paragraph(element, anchor._parent)
        paragraph.style = doc.styles[MANUAL_REVIEW_STYLE]
        paragraph.add_run(text)
        mark_manual_review_paragraph(paragraph)


def audit_scorecard(path: Path, card: dict[str, Any]) -> dict[str, Any]:
    doc = Document(path)
    paragraphs = list(all_body_paragraphs(doc))
    semantic_valid = _scorecard_semantics_valid(card)
    docx_sha = hashlib.sha256(path.read_bytes()).hexdigest()
    serialized_audit = card.get("serialized_format_audit")
    rendered_audit = card.get("rendered_output_audit")

    def attached_audit_valid(value: Any, *, expected_protocol: str) -> bool:
        if value is None:
            return True
        if not isinstance(value, dict) or value.get("protocol") != expected_protocol:
            return False
        digest = value.get("audit_sha256")
        payload = {key: copy.deepcopy(child) for key, child in value.items() if key != "audit_sha256"}
        return (value.get("docx_sha256") == docx_sha
                and isinstance(digest, str) and sha256_json(payload) == digest)

    if isinstance(serialized_audit, dict):
        audit_digest = serialized_audit.get("audit_sha256")
        audit_payload = {key: copy.deepcopy(value) for key, value in serialized_audit.items()
                         if key != "audit_sha256"}
        semantic_valid = semantic_valid and (
            serialized_audit.get("protocol") == "serialized_format_repairs_audit_v1"
            and serialized_audit.get("docx_sha256") == docx_sha
            and isinstance(audit_digest, str)
            and sha256_json(audit_payload) == audit_digest
        )
    if rendered_audit is not None:
        semantic_valid = semantic_valid and attached_audit_valid(
            rendered_audit, expected_protocol="rendered_output_audit_attachment_v1",
        )
        if isinstance(rendered_audit, dict):
            semantic_valid = semantic_valid and (
                rendered_audit.get("submission_ready") is False
                and isinstance(rendered_audit.get("pdf_sha256"), str)
                and isinstance(rendered_audit.get("report"), dict)
                and sha256_json(rendered_audit["report"]) == rendered_audit.get("report_sha256")
                and card.get("current_output_assessment", {}).get("docx_sha256") == docx_sha
                and card.get("current_output_assessment", {}).get("pdf_sha256") == rendered_audit.get("pdf_sha256")
            )
    current_assessment = card.get("current_output_assessment")
    if isinstance(current_assessment, dict):
        semantic_valid = semantic_valid and current_assessment.get("docx_sha256") == docx_sha
    lines = scorecard_lines(card) if semantic_valid else []
    texts = [p.text for p in paragraphs]
    expected_ids = {item["item_id"] for item in card["entries"]}
    actual_ids = [p.text.split("｜", 1)[0][1:] for p in paragraphs if p.text.startswith(PREFIX)]
    valid = (semantic_valid and all(texts.count(line) == 1 for line in lines)
             and len(actual_ids) == len(expected_ids) and set(actual_ids) == expected_ids
             and all(p.runs and all(
                 _effective_run_attribute(run, p, doc, "vanish") in (None, "0", "false", "off")
                 and _effective_run_attribute(run, p, doc, "color") == "C00000"
                 and bool(_effective_run_attribute(run, p, doc, "rFonts", "eastAsia"))
                 for run in p.runs if run.text)
                     for p in paragraphs if p.text in lines))
    return {"valid": valid, "policy": POLICY, "submission_ready": False,
            "scorecard_sha256": sha256_json(card),
            "docx_sha256": docx_sha,
            "item_count": len(actual_ids), "visual_verification": "required",
            "obligation_assessment_status": (
                "recorded" if isinstance(card.get("obligation_assessment"), dict)
                else "legacy_not_recorded"
            )}


def attach_rendered_output_audit(scorecard_path: Path, report: dict[str, Any],
                                 docx_path: Path,
                                 property_receipt_audit: dict[str, Any] | None = None,
                                 *, historical_docx_sha256: str | None = None,
                                 rendered_requirement_bindings: dict[str, list[str]] | None = None
                                 ) -> dict[str, Any]:
    """Reconcile statuses and attach only after the visible scorecard is current.

    If reconciliation changes visible scorecard text, this updates the DOCX,
    saves the reconciled JSON, and raises so the caller must rerender and audit
    those new bytes before a final attachment is accepted.
    """
    card = strict_json_read(scorecard_path)
    docx_sha = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    report_payload = {key: value for key, value in report.items() if key != "audit_sha256"}
    report_digest = report.get("audit_sha256")
    report_final_docx = report.get("final_docx")
    if (report.get("protocol") != "rendered_format_audit_v1"
            or report.get("docx_sha256") != docx_sha
            or not isinstance(report_final_docx, dict)
            or report_final_docx.get("sha256") != docx_sha
            or not isinstance(report.get("pdf_sha256"), str)
            or re.fullmatch(r"[0-9a-f]{64}", report.get("pdf_sha256", "")) is None
            or report.get("submission_ready") is not False
            or report.get("field_refresh_claimed") is not False
            or not isinstance(report.get("findings"), list)
            or not isinstance(report_digest, str)
            or sha256_json(report_payload) != report_digest):
        raise ValueError("rendered-format report DOCX hash does not match scorecard document")
    card = reconcile_current_output_scorecard(
        card, report, docx_path,
        property_receipt_audit=property_receipt_audit,
        historical_docx_sha256=historical_docx_sha256,
        rendered_requirement_bindings=rendered_requirement_bindings,
    )
    history = card["semantic_history"]
    current_report_sha = sha256_json(report)
    for audit_key in ("serialized_format_audit", "rendered_output_audit"):
        old_audit = card.get(audit_key)
        stale = isinstance(old_audit, dict) and (
            old_audit.get("docx_sha256") != docx_sha
            or (audit_key == "rendered_output_audit"
                and old_audit.get("report_sha256") != current_report_sha)
        )
        if stale:
            history.setdefault(audit_key + "_history", []).append(copy.deepcopy(old_audit))
            card.pop(audit_key, None)
    display_audit = audit_scorecard(docx_path, card)
    if not display_audit["valid"]:
        sync_result = synchronize_scorecard_display(docx_path, card)
        scorecard_path.write_text(json.dumps(card, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if sync_result.get("changed"):
            raise ValueError("current scorecard display changed the DOCX; rerender and re-audit before attachment")
        raise ValueError("current scorecard display or bindings do not validate against the DOCX")
    attachment: dict[str, Any] = {
        "protocol": "rendered_output_audit_attachment_v1",
        "docx_sha256": docx_sha,
        "pdf_sha256": report.get("pdf_sha256"),
        "parent_scorecard_sha256": card["semantic_history"]["source_scorecard_sha256"],
        "report_sha256": sha256_json(report),
        "report": copy.deepcopy(report),
        "submission_ready": False,
    }
    attachment["audit_sha256"] = sha256_json(attachment)
    card["rendered_output_audit"] = attachment
    scorecard_path.write_text(json.dumps(card, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if not audit_scorecard(docx_path, card)["valid"]:
        raise ValueError("final rendered audit attachment does not validate against current scorecard and DOCX")
    return card


def _current_receipt_status(entry: dict[str, Any], docx_sha: str,
                             receipt_audit: dict[str, Any] | None) -> tuple[str, dict[str, Any]]:
    detail = entry.get("detail") if isinstance(entry.get("detail"), dict) else {}
    receipt_id = detail.get("receipt_id")
    if not isinstance(receipt_audit, dict) or not isinstance(receipt_audit.get("receipts"), list):
        return "unverified", {"reason": "current DOCX-bound property receipt inventory is absent"}
    matches = [item for item in receipt_audit["receipts"]
               if isinstance(item, dict) and item.get("receipt_id") == receipt_id]
    if len(matches) != 1:
        return "unverified", {"reason": "current property receipt is missing or duplicated"}
    receipt = matches[0]
    keys = ("requirement_id", "role", "property_path", "target_locator", "expected")
    if (receipt.get("serialized_docx_sha256") != docx_sha
            or not isinstance(receipt.get("target_locator"), str)
            or not receipt.get("target_locator")
            or any(receipt.get(key) != detail.get(key) for key in keys)):
        return "unverified", {"reason": "receipt identity or DOCX hash does not match this scorecard object"}
    status = receipt.get("status")
    if status not in STATUSES:
        return "unverified", {"reason": "current receipt status is unknown"}
    return status, {"receipt_id": receipt_id, "docx_sha256": docx_sha,
                    "role": receipt.get("role"), "property_path": receipt.get("property_path"),
                    "target_locator": receipt.get("target_locator"),
                    "receipt_status": status}


def _rendered_finding_identity(finding: dict[str, Any]) -> dict[str, Any]:
    code = finding.get("code")
    if code == "rendered_pdf_font_mismatch":
        expected = finding.get("expected") if isinstance(finding.get("expected"), dict) else {}
        return {"kind": "font", "role": expected.get("role"),
                "style_name": expected.get("style_name"),
                "script": expected.get("script"), "font": expected.get("name"),
                "requirement_ids": sorted(expected.get("requirement_ids", []))}
    if code and str(code).startswith("drawing_"):
        return {"kind": "drawing", "code": code,
                "drawing_index": finding.get("drawing_index"),
                "pixel_sha256": finding.get("pixel_sha256"),
                "media_sha256": finding.get("media_sha256"),
                "page": finding.get("page")}
    return {"kind": "rendered_finding", "code": code,
            "role": finding.get("role"), "property": finding.get("property"),
            "target_locator": finding.get("target_locator"),
            "object_sha256": finding.get("pixel_sha256") or finding.get("media_sha256")}


def reconcile_current_output_scorecard(
    card: dict[str, Any], report: dict[str, Any], docx_path: Path, *,
    property_receipt_audit: dict[str, Any] | None = None,
    historical_docx_sha256: str | None = None,
    rendered_requirement_bindings: dict[str, list[str]] | None = None,
) -> dict[str, Any]:
    """Separate inherited semantic/history results from exact current DOCX checks."""
    docx_sha = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    if report.get("docx_sha256") != docx_sha:
        raise ValueError("current output reconciliation requires the exact final DOCX hash")
    has_semantic_history = isinstance(card.get("semantic_history"), dict)
    if "semantic_history" not in card:
        history_counts = {status: card.get("status_counts", {}).get(status, 0) for status in STATUSES}
        card["semantic_history"] = {
            "protocol": "inherited_scorecard_history_v1",
            "source_scorecard_sha256": sha256_json(card),
            "historical_docx_sha256": historical_docx_sha256,
            "status_counts": history_counts,
            "score": card.get("score"),
            "earned_weight": card.get("earned_weight"),
            "total_weight": card.get("total_weight"),
            "coverage_complete": card.get("coverage_complete"),
            "unique_source_obligation_count": card.get("unique_source_obligation_count"),
            "evaluation_unit_statuses": {
                str(item.get("evaluation_unit_id")): item.get("status")
                for item in card.get("evaluation_units", []) if isinstance(item, dict)
            },
        }
    history = card["semantic_history"]
    # Output-only entries describe findings in the previously rendered PDF.
    # Recompute them from this exact report rather than carrying stale findings
    # forward. The prior audit itself is retained by attach_rendered_output_audit.
    card["entries"] = [entry for entry in card.get("entries", [])
                       if entry.get("kind") != "rendered_output_instance"]
    for entry in card.get("entries", []):
        if not has_semantic_history and entry.get("historical_status") is None:
            entry["historical_status"] = entry.get("status")
        old_detail_status = (entry.get("detail") or {}).get("status")
        if old_detail_status is not None:
            entry.setdefault("historical_detail_status", old_detail_status)
        if entry.get("kind") == "human_review" or entry.get("historical_status") == "pending":
            status, evidence = "pending", {"reason": "human review remains pending"}
        elif entry.get("kind") == "property":
            status, evidence = _current_receipt_status(entry, docx_sha, property_receipt_audit)
        else:
            status, evidence = "unverified", {"reason": "no exact current DOCX property receipt binds this entry"}
        entry["status"] = status
        entry["score"] = _entry_score(status)
        entry["human_check_required"] = status != "verified"
        if isinstance(entry.get("detail"), dict):
            entry["detail"]["status"] = status
            entry["detail"]["current_output_instance"] = evidence

    entries_by_id = {entry.get("item_id"): entry for entry in card.get("entries", [])}
    for finding in report.get("findings", []):
        if not isinstance(finding, dict):
            continue
        code = finding.get("code")
        identity = _rendered_finding_identity(finding)
        status = "unverified"
        matches: list[dict[str, Any]] = []
        if code == "rendered_pdf_font_mismatch":
            expected = finding.get("expected") if isinstance(finding.get("expected"), dict) else {}
            script = expected.get("script")
            path = {"cjk": "font.cjk", "latin": "font.latin"}.get(script)
            requirement_ids = expected.get("requirement_ids")
            bound_ids = (rendered_requirement_bindings or {}).get(sha256_json(finding))
            if (path and isinstance(expected.get("style_name"), str)
                    and isinstance(requirement_ids, list)
                    and all(isinstance(value, str) for value in requirement_ids)
                    and isinstance(bound_ids, list)
                    and all(isinstance(value, str) for value in bound_ids)
                    and sorted(set(requirement_ids)) == sorted(set(bound_ids))):
                locator = "style:" + expected["style_name"]
                matches = [entry for entry in card.get("entries", [])
                           if entry.get("kind") == "property"
                           and (entry.get("detail") or {}).get("requirement_id") in requirement_ids
                           and (entry.get("detail") or {}).get("role") == expected.get("role")
                           and (entry.get("detail") or {}).get("property_path") == path
                           and (entry.get("detail") or {}).get("target_locator") == locator
                           and (entry.get("detail") or {}).get("expected") == expected.get("name")]
            status = "failed"  # Explicit required-font mismatch in the bound rendered PDF.
        elif code in {"rendered_drawing_clipped_by_page", "rendered_drawing_extent_mismatch",
                      "drawing_count_changed", "drawing_media_changed",
                      "drawing_line_box_not_object_safe", "renderer_report_pdf_hash_mismatch"}:
            status = "failed"
        elif code in {"drawing_not_found_in_rendered_pdf", "drawing_match_ambiguous"}:
            status = "unverified"
        if matches:
            for entry in matches:
                entry["status"] = status
                entry["score"] = _entry_score(status)
                entry["human_check_required"] = True
                entry["detail"]["status"] = status
                entry["detail"]["current_output_instance"] = {
                    "docx_sha256": docx_sha, "pdf_sha256": report.get("pdf_sha256"),
                    "finding_code": code, "object_identity": identity,
                }
            continue
        stable_id = "SC-OUT-" + sha256_json(identity)[:24]
        if stable_id in entries_by_id:
            entry = entries_by_id[stable_id]
            entry["status"] = status
            entry["score"] = _entry_score(status)
            entry["human_check_required"] = True
            entry["detail"]["current_output_instance"] = {
                "docx_sha256": docx_sha, "pdf_sha256": report.get("pdf_sha256"),
                "finding_code": code, "object_identity": identity,
            }
            continue
        entry = {
            "item_id": stable_id, "kind": "rendered_output_instance",
            "label": str(code or "rendered output instance"), "status": status,
            "score": _entry_score(status), "max_score": 100,
            "human_check_required": True,
            "detail": {**copy.deepcopy(finding), "current_output_instance": {
                "docx_sha256": docx_sha, "pdf_sha256": report.get("pdf_sha256"),
                "finding_code": code, "object_identity": identity,
            }},
        }
        card.setdefault("entries", []).append(entry)
        entries_by_id[stable_id] = entry

    status_counts = {status: sum(item.get("status") == status for item in card["entries"])
                     for status in STATUSES}
    card["status_counts"] = status_counts
    for status in STATUSES:
        card[f"{status}_count"] = status_counts[status]
    card["item_count"] = len(card["entries"])
    card["diagnostic_item_count"] = len(card["entries"])
    card["score"] = None
    card["score_meaning"] = "current_output_instance_validation_only_no_semantic_or_release_score"
    card["earned_weight"] = 0
    card["total_weight"] = 0
    card["coverage_complete"] = False
    card["coverage_assessment_status"] = "current_output_instance_only_no_semantic_reassessment"
    card["unique_source_obligation_count"] = None
    card["observed_checks_verified_percent"] = (
        round(status_counts["verified"] * 100 / len(card["entries"]), 2) if card["entries"] else None
    )
    card["human_check_required"] = True
    card["submission_ready"] = False
    for unit in card.get("evaluation_units", []):
        if not isinstance(unit, dict):
            continue
        unit.setdefault("historical_status", unit.get("status"))
        linked = [entries_by_id.get(item_id) for item_id in unit.get("item_ids", [])]
        linked = [item for item in linked if isinstance(item, dict)]
        states = [item.get("status") for item in linked]
        unit["status"] = ("pending" if unit.get("identity", {}).get("route") == "human"
                           else "failed" if "failed" in states
                           else "unverified" if "unverified" in states
                           else "verified" if states else "unverified")
        unit["earned_weight"] = 0
    card["current_output_assessment"] = {
        "protocol": "current_docx_rendered_instance_assessment_v1",
        "docx_sha256": docx_sha, "pdf_sha256": report.get("pdf_sha256"),
        "rendered_report_sha256": sha256_json(report),
        "status": ("issues_found" if status_counts["failed"] else
                   "unverified" if status_counts["unverified"] or status_counts["pending"] else "passed"),
        "status_counts": status_counts, "submission_ready": False,
        "semantic_history_status_counts": copy.deepcopy(history["status_counts"]),
        "property_receipt_inventory_bound": bool(property_receipt_audit),
        "field_refresh_claimed": False,
    }
    card["semantic_history_status_counts"] = copy.deepcopy(history["status_counts"])
    return card


def synchronize_scorecard_display(docx_path: Path, card: dict[str, Any]) -> dict[str, Any]:
    """Rewrite the visible scorecard block to match current card statuses.

    This changes DOCX bytes when any status or count changes; callers must
    rerender, refresh the TOC cache, and attach a new hash-bound PDF audit.
    """
    doc = Document(docx_path)
    lines = scorecard_lines(card)
    paragraphs = list(all_body_paragraphs(doc))
    head = [p for p in paragraphs if p.text.startswith("自动核验评分（")
            or p.text.startswith("当前生成稿实例核验（")]
    if len(head) != 1:
        raise ValueError("scorecard display must have exactly one summary paragraph")
    actual_entries: dict[str, list[Any]] = {}
    for paragraph in paragraphs:
        if paragraph.text.startswith(PREFIX):
            item_id = paragraph.text.split("｜", 1)[0][1:]
            actual_entries.setdefault(item_id, []).append(paragraph)
    expected_ids = [item["item_id"] for item in card.get("entries", [])]
    if len(set(expected_ids)) != len(expected_ids):
        raise ValueError("scorecard contains duplicate item IDs")
    stale_output_ids = {item_id for item_id in set(actual_entries) - set(expected_ids)
                        if item_id.startswith("SC-OUT-")}
    extra_ids = set(actual_entries) - set(expected_ids) - stale_output_ids
    if extra_ids:
        raise ValueError("DOCX contains scorecard items absent from current scorecard")
    if any(len(actual_entries.get(item_id, [])) > 1 for item_id in expected_ids):
        raise ValueError("DOCX contains duplicate visible scorecard items")

    changed = False
    for item_id in stale_output_ids:
        for paragraph in actual_entries[item_id]:
            paragraph._p.getparent().remove(paragraph._p)
            changed = True

    def set_text(paragraph: Any, text: str) -> None:
        nonlocal changed
        if paragraph.text == text:
            return
        paragraph.clear()
        paragraph.style = doc.styles[MANUAL_REVIEW_STYLE]
        paragraph.add_run(text)
        mark_manual_review_paragraph(paragraph)
        changed = True

    set_text(head[0], lines[0])
    previous = head[0]
    for item, line in zip(card["entries"], lines[1:]):
        item_id = item["item_id"]
        existing = actual_entries.get(item_id, [])
        if existing:
            paragraph = existing[0]
            set_text(paragraph, line)
            # Keep the display order identical to the scorecard JSON order.
            if paragraph._p.getprevious() is not previous._p:
                paragraph._p.getparent().remove(paragraph._p)
                previous._p.addnext(paragraph._p)
                changed = True
        else:
            element = OxmlElement("w:p")
            previous._p.addnext(element)
            paragraph = Paragraph(element, previous._parent)
            paragraph.style = doc.styles[MANUAL_REVIEW_STYLE]
            paragraph.add_run(line)
            mark_manual_review_paragraph(paragraph)
            changed = True
        previous = paragraph
    if not changed:
        return {"changed": False, "docx_sha256": hashlib.sha256(docx_path.read_bytes()).hexdigest()}

    before_sha = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    descriptor, temporary = tempfile.mkstemp(prefix=docx_path.stem + ".scorecard-", suffix=".docx",
                                               dir=docx_path.parent)
    os.close(descriptor)
    temp_path = Path(temporary)
    try:
        doc.save(temp_path)
        os.replace(temp_path, docx_path)
    finally:
        temp_path.unlink(missing_ok=True)
    after_sha = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    if after_sha == before_sha:
        raise RuntimeError("scorecard display changed in memory but DOCX hash did not change")
    return {"changed": True, "before_docx_sha256": before_sha, "docx_sha256": after_sha,
            "rerender_required": True}


def validate_bound_scorecard(card: dict[str, Any], *, binding: dict[str, Any],
                             receipt_audit: dict[str, Any], manual_items: list[dict[str, Any]],
                             findings: list[dict[str, Any]], output: Path,
                             capability_findings: list[dict[str, Any]] | None = None,
                             evaluation_units_by_requirement: dict[str, list[dict[str, Any]]] | None = None,
                             obligation_scope_inventory: dict[str, Any] | None = None,
                             expected_docx_sha256: str | None = None) -> bool:
    try:
        expected = build_scorecard(
            binding, receipt_audit, manual_items, findings, capability_findings,
            evaluation_units_by_requirement=evaluation_units_by_requirement,
            obligation_scope_inventory=obligation_scope_inventory,
            expected_docx_sha256=expected_docx_sha256,
        )
    except (ValueError, TypeError):
        return False
    return sha256_json(card) == sha256_json(expected) and audit_scorecard(output, card)["valid"]
