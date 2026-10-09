"""Compile and independently audit source-scoped DOCX header rules.

Header text is not a document-wide singleton when the accepted source clauses
limit it to a document region.  This module preserves those clause scopes and
binds them to concrete section anchors in the supplied thesis DOCX.  Header
execution and OOXML relationship auditing live in ``apply_format_spec`` and
``header_scope_audit`` respectively; this file contains the shared source
anchor interpretation only.
"""
from __future__ import annotations

import copy
import hashlib
import json
import re
import unicodedata
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SCHEMA_PATH = Path(__file__).resolve().parents[1] / "schema" / "format-spec.schema.json"

# These are registered applicability facts in the existing semantic contract,
# not clause IDs or thesis-specific section numbers.
_APPLICABILITY_SCOPES = {
    ("source_inventory.table_of_contents", "present", None): ("toc", "toc_heading"),
    ("thesis_profile.has_appendices", "equals", True): ("appendix", "appendix_heading_1"),
}


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _normalized_text(value: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", value)).strip()


def _compact_text(value: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value)).casefold()


def _heading1(paragraph: Any) -> bool:
    style = _normalized_text(getattr(getattr(paragraph, "style", None), "name", "")).casefold()
    return style in {"heading 1", "heading1", "标题 1", "标题1"}


def _toc_heading(paragraph: Any) -> bool:
    style = _normalized_text(getattr(getattr(paragraph, "style", None), "name", "")).casefold()
    text = _compact_text(getattr(paragraph, "text", ""))
    return "toc" in style or "目录标题" in style or (
        text in {"目录", "目錄"} and style in {"heading 1", "heading1", "标题 1", "标题1"}
    )


def _appendix_heading(text: str) -> bool:
    return _compact_text(text).startswith("附录")


_CHAPTER_PREFIX = re.compile(
    r"^(?:第\s*[0-9一二三四五六七八九十百千〇零]+\s*章|"
    r"chapter\s+[0-9ivxlcdm]+|[0-9]+(?:[.．][0-9]+)*\s+\S)",
    re.IGNORECASE,
)


def _chapter_heading(text: str) -> bool:
    return bool(_CHAPTER_PREFIX.search(_normalized_text(text)))


def _condition_scope(condition: Mapping[str, Any]) -> tuple[str, str] | None:
    key = (
        condition.get("fact"), condition.get("operator"), condition.get("value"),
    )
    return _APPLICABILITY_SCOPES.get(key)


def _rule_signature(rule: Mapping[str, Any]) -> tuple[Any, ...]:
    return (
        rule.get("scope"), rule.get("selector"), rule.get("content_mode"),
        rule.get("header_content"), tuple(rule.get("requirement_ids", [])),
        tuple(rule.get("clause_ids", [])), tuple(rule.get("evidence_ids", [])),
    )


def compile_header_scope_rules(spec: Mapping[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Compile accepted clause applicability into typed header-scope rules.

    A legacy singleton ``roles.header.header_content`` is intentionally not
    treated as a scoped rule.  Conditional header requirements must identify a
    registered spatial scope; unknown or compound conditions fail closed.
    """
    findings: list[dict[str, Any]] = []
    requirements = spec.get("requirements", [])
    if not isinstance(requirements, list):
        return [], [{"code": "header_scope_requirements_invalid", "severity": "error",
                     "message": "format-spec requirements must be an array before header scopes can be compiled."}]

    by_scope: dict[str, dict[str, Any]] = {}
    unscoped_header_requirements: list[str] = []
    for requirement in requirements:
        if not isinstance(requirement, Mapping) or requirement.get("role") != "header":
            continue
        props = requirement.get("properties")
        content = props.get("header_content") if isinstance(props, Mapping) else None
        if not isinstance(content, Mapping) or not content:
            continue
        applicability = requirement.get("applicability")
        if not isinstance(applicability, Mapping):
            unscoped_header_requirements.append(str(requirement.get("id", "<missing-id>")))
            continue
        conditions = applicability.get("conditions")
        if applicability.get("status") != "conditional" or not isinstance(conditions, list) or not conditions:
            findings.append({
                "code": "header_scope_applicability_unresolved", "severity": "error",
                "requirement_id": requirement.get("id"),
                "message": "A header requirement has applicability but no supported conditional source scope.",
                "evidence": {"applicability": copy.deepcopy(applicability)},
            })
            continue
        resolved = [_condition_scope(item) if isinstance(item, Mapping) else None for item in conditions]
        if any(item is None for item in resolved) or len(set(resolved)) != 1:
            findings.append({
                "code": "header_scope_applicability_unmapped", "severity": "error",
                "requirement_id": requirement.get("id"),
                "message": "Header applicability conditions do not identify exactly one supported section scope.",
                "evidence": {"conditions": copy.deepcopy(conditions)},
            })
            continue
        scope, selector_kind = resolved[0]  # type: ignore[misc]
        clause_ids = requirement.get("clause_ids") if isinstance(requirement.get("clause_ids"), list) else []
        evidence_ids = requirement.get("evidence_ids") if isinstance(requirement.get("evidence_ids"), list) else []
        requirement_id = requirement.get("id")
        if not isinstance(requirement_id, str) or not clause_ids or not evidence_ids:
            findings.append({
                "code": "header_scope_provenance_missing", "severity": "error",
                "requirement_id": requirement_id,
                "message": "A scoped header rule must retain its requirement, clause, and evidence references.",
            })
            continue
        rule = {
            "scope": scope,
            "selector": {"kind": selector_kind},
            "content_mode": "fixed",
            "header_content": copy.deepcopy(dict(content)),
            "requirement_ids": [requirement_id],
            "clause_ids": sorted({str(item) for item in clause_ids}),
            "evidence_ids": sorted({str(item) for item in evidence_ids}),
        }
        prior = by_scope.get(scope)
        if prior is None:
            by_scope[scope] = rule
        elif prior["header_content"] == rule["header_content"] and prior["selector"] == rule["selector"]:
            for key in ("requirement_ids", "clause_ids", "evidence_ids"):
                prior[key] = sorted(set(prior[key]) | set(rule[key]))
        else:
            findings.append({
                "code": "header_scope_conflict", "severity": "error",
                "scope": scope,
                "message": "Multiple accepted header requirements provide conflicting content for the same source scope.",
                "evidence": {"requirements": [prior, rule]},
            })

    if by_scope and unscoped_header_requirements:
        findings.append({
            "code": "header_scope_unscoped_requirement", "severity": "error",
            "message": "Scoped and unscoped header content coexist; the unscoped requirement cannot be safely applied by section.",
            "evidence": {"requirement_ids": sorted(unscoped_header_requirements)},
        })

    derived = [by_scope[key] for key in sorted(by_scope)]
    explicit = spec.get("header_scope_rules")
    if explicit is not None:
        if not isinstance(explicit, list):
            findings.append({"code": "header_scope_rules_invalid", "severity": "error",
                             "message": "header_scope_rules must be an array."})
        else:
            explicit_by_scope = {str(item.get("scope")): item for item in explicit if isinstance(item, Mapping)}
            derived_by_scope = {item["scope"]: item for item in derived}
            if set(explicit_by_scope) != set(derived_by_scope) or any(
                _rule_signature(explicit_by_scope[key]) != _rule_signature(derived_by_scope[key])
                for key in set(explicit_by_scope) & set(derived_by_scope)
            ):
                findings.append({
                    "code": "header_scope_rule_source_mismatch", "severity": "error",
                    "message": "Persisted header_scope_rules differ from deterministic compilation of current accepted requirements.",
                    "evidence": {"persisted": explicit, "compiled": derived},
                })
            else:
                derived = [copy.deepcopy(explicit_by_scope[key]) for key in sorted(explicit_by_scope)]
    if any(item.get("severity") == "error" for item in findings):
        # A valid subset is not safe to execute when another accepted header
        # requirement is unmapped. Return no executable rules in that case.
        return [], findings
    return derived, findings


def _section_paragraphs(doc: Any) -> list[list[Any]]:
    sections: list[list[Any]] = [[] for _ in doc.sections]
    index = 0
    # Walk the body tree, not Document.paragraphs: the latter silently omits
    # table-cell paragraphs, and a heading in a table still belongs to the
    # section whose header will render it. A section break applies after its
    # paragraph, so record the paragraph before advancing the section index.
    from docx.text.paragraph import Paragraph

    body = doc._body._element
    for child in list(body):
        for element in child.iter(qn("w:p")):
            if index >= len(sections):
                return []
            paragraph = Paragraph(element, doc._body)
            sections[index].append(paragraph)
            ppr = element.find(qn("w:pPr"))
            sectpr = ppr.find(qn("w:sectPr")) if ppr is not None else None
            if sectpr is not None:
                # OOXML section boundaries in table-cell paragraphs are not a
                # supported source topology. Do not guess which table row is
                # the boundary for the body-level section properties.
                if child.tag != qn("w:p") or element.getparent() is not body:
                    return []
                index += 1
    if index != len(sections) - 1:
        return []
    return sections


def prepare_header_scope_boundaries(doc: Any, spec: Mapping[str, Any]) -> dict[str, Any]:
    """Isolate accepted fixed header scopes from preceding unrelated content.

    A Word header applies to the whole section. When an accepted TOC or
    appendix anchor starts inside a mixed source section, applying that header
    would leak it onto earlier pages. Add a next-page section boundary at the
    preceding paragraph, without changing source text. The caller keeps the
    original source bytes and records this deterministic layout transform.
    """
    rules, _ = compile_header_scope_rules(spec)
    rules_by_scope = {str(item["scope"]): item for item in rules}
    original_section_count = len(doc.sections)
    source_sections = _section_paragraphs(doc)
    findings: list[dict[str, Any]] = []
    if len(source_sections) != original_section_count:
        findings.append({
            "code": "header_scope_section_parse_failed", "severity": "error",
            "message": "Could not map source paragraphs before isolating header scopes.",
        })
        source_sections = [[] for _ in range(original_section_count)]

    candidates: list[tuple[int, int, Any, str, dict[str, Any]]] = []
    toc_rule = rules_by_scope.get("toc")
    appendix_rule = rules_by_scope.get("appendix")
    toc_count = 0
    for source_section_index, paragraphs in enumerate(source_sections, 1):
        nonblank = [(i, p) for i, p in enumerate(paragraphs) if str(p.text).strip()]
        toc_anchors = [(i, p) for i, p in nonblank if _toc_heading(p)]
        appendix_anchors = [
            (i, p) for i, p in nonblank
            if _heading1(p) and _appendix_heading(str(p.text))
        ]
        if toc_rule:
            toc_count += len(toc_anchors)
            if len(toc_anchors) > 1:
                findings.append({
                    "code": "header_scope_toc_anchor_ambiguous", "severity": "error",
                    "section_index": source_section_index,
                    "message": "The source section contains more than one table-of-contents title anchor.",
                    "evidence": {"paragraphs": [p.text for _, p in toc_anchors]},
                })
            elif toc_anchors:
                anchor_index, anchor = toc_anchors[0]
                if any(str(p.text).strip() for paragraph_index, p in nonblank if paragraph_index < anchor_index):
                    candidates.append((source_section_index, anchor_index, anchor, "toc", toc_rule))
        if appendix_rule and appendix_anchors:
            anchor_index, anchor = appendix_anchors[0]
            if any(str(p.text).strip() for paragraph_index, p in nonblank if paragraph_index < anchor_index):
                candidates.append((source_section_index, anchor_index, anchor, "appendix", appendix_rule))
    if toc_rule and toc_count > 1:
        findings.append({
            "code": "header_scope_toc_anchor_ambiguous", "severity": "error",
            "message": "More than one source table-of-contents title matches the accepted TOC header rule.",
            "evidence": {"matched_anchor_count": toc_count},
        })

    # Insert from the end so earlier paragraph and section indices remain stable.
    candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
    body = doc._body._element
    inserted_by_source_section: dict[int, int] = {}
    transforms: list[dict[str, Any]] = []
    for source_section_index, anchor_index, anchor, scope, rule in candidates:
        preceding = [
            str(paragraph.text).strip()
            for paragraph_index, paragraph in enumerate(source_sections[source_section_index - 1])
            if paragraph_index < anchor_index and str(paragraph.text).strip()
        ]
        if not preceding:
            continue
        element = anchor._p
        if element.getparent() is not body:
            findings.append({
                "code": "header_scope_boundary_unsupported_container", "severity": "error",
                "section_index": source_section_index, "scope": scope,
                "message": "A mixed-scope header anchor is inside a table and cannot be isolated with a safe body-level section boundary.",
                "evidence": {"anchor": str(anchor.text), "paragraph_index": anchor_index},
            })
            continue
        children = list(body)
        try:
            target_position = children.index(element)
        except ValueError:
            findings.append({
                "code": "header_scope_boundary_anchor_unresolved", "severity": "error",
                "section_index": source_section_index, "scope": scope,
                "message": "The source scope anchor could not be resolved to a body-level location.",
                "evidence": {"anchor": str(anchor.text), "paragraph_index": anchor_index},
            })
            continue
        if target_position == 0 or children[target_position - 1].tag != qn("w:p"):
            findings.append({
                "code": "header_scope_boundary_predecessor_unsupported", "severity": "error",
                "section_index": source_section_index, "scope": scope,
                "message": "The source scope anchor is not preceded by a direct body paragraph that can safely carry a section boundary.",
                "evidence": {"anchor": str(anchor.text), "paragraph_index": anchor_index},
            })
            continue
        predecessor = children[target_position - 1]
        ppr = predecessor.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            predecessor.insert(0, ppr)
        if ppr.find(qn("w:sectPr")) is not None:
            findings.append({
                "code": "header_scope_boundary_already_present", "severity": "error",
                "section_index": source_section_index, "scope": scope,
                "message": "A pre-existing section boundary conflicts with the expected source scope split.",
                "evidence": {"anchor": str(anchor.text), "paragraph_index": anchor_index},
            })
            continue
        source_section = doc.sections[source_section_index - 1]
        sectpr = copy.deepcopy(source_section._sectPr)
        section_type = sectpr.find(qn("w:type"))
        if section_type is None:
            section_type = OxmlElement("w:type")
            page_size = sectpr.find(qn("w:pgSz"))
            if page_size is None:
                sectpr.append(section_type)
            else:
                sectpr.insert(list(sectpr).index(page_size), section_type)
        section_type.set(qn("w:val"), "nextPage")
        ppr.append(sectpr)
        inserted_by_source_section[source_section_index] = inserted_by_source_section.get(source_section_index, 0) + 1
        previous_index = max(0, anchor_index - 1)
        previous_text = str(source_sections[source_section_index - 1][previous_index].text) if anchor_index else ""
        transforms.append({
            "kind": "insert_next_page_section_boundary_before_anchor",
            "scope": scope,
            "source_section_index": source_section_index,
            "source_anchor": _source_anchor(anchor, anchor_index),
            "predecessor_sha256": _sha256_text(previous_text),
            "requirement_ids": list(rule.get("requirement_ids", [])),
            "clause_ids": list(rule.get("clause_ids", [])),
            "evidence_ids": list(rule.get("evidence_ids", [])),
        })

    source_section_indices: list[int] = []
    for source_index in range(1, original_section_count + 1):
        source_section_indices.extend([source_index] * (1 + inserted_by_source_section.get(source_index, 0)))
    transforms.sort(key=lambda item: (
        int(item["source_section_index"]),
        int(item["source_anchor"]["paragraph_index"]),
    ))
    return {
        "source_section_indices": source_section_indices,
        "scope_transforms": transforms,
        "findings": findings,
    }


def _story_details(story: Any) -> dict[str, Any]:
    text = "\n".join(paragraph.text for paragraph in story.paragraphs)
    instructions: list[str] = []
    for paragraph in story.paragraphs:
        instructions.extend(
            str(value).strip() for value in paragraph._p.xpath(".//w:instrText/text()")
        )
        instructions.extend(
            str(value).strip() for value in paragraph._p.xpath(".//w:fldSimple/@w:instr")
        )
    return {"text": text, "instructions": instructions, "linked_to_previous": bool(story.is_linked_to_previous)}


def _active_variants(doc: Any, section: Any, spec: Mapping[str, Any]) -> dict[str, bool]:
    settings = doc.settings.element
    even_odd = settings.find(qn("w:evenAndOddHeaders")) is not None
    page = spec.get("page") if isinstance(spec.get("page"), Mapping) else {}
    desired_first = page.get("different_first_page") if isinstance(page, Mapping) else None
    desired_even = page.get("different_odd_even") if isinstance(page, Mapping) else None
    return {
        "default": True,
        "first": bool(section.different_first_page_header_footer if desired_first is None else desired_first),
        "even": bool(even_odd if desired_even is None else desired_even),
    }


def _source_anchor(paragraph: Any, paragraph_index: int) -> dict[str, Any]:
    text = str(paragraph.text)
    style = getattr(getattr(paragraph, "style", None), "name", "")
    return {
        "kind": "toc_heading" if _toc_heading(paragraph) else "heading_1",
        "paragraph_index": paragraph_index,
        "style": str(style),
        "text": text,
        "text_sha256": _sha256_text(text),
    }


def compile_header_scope_plan(source_docx: str | Path, spec: Mapping[str, Any],
                             format_spec_path: str | Path | None = None,
                             *, document: Any | None = None,
                             prepared_layout: Mapping[str, Any] | None = None) -> dict[str, Any]:
    """Resolve compiled header scope rules against the actual input DOCX."""
    path = Path(source_docx)
    source_bytes = path.read_bytes()
    source_sha = hashlib.sha256(source_bytes).hexdigest()
    spec_bytes = (
        Path(format_spec_path).read_bytes() if format_spec_path is not None
        else json.dumps(dict(spec), ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    )
    rules, findings = compile_header_scope_rules(spec)
    doc = document if document is not None else Document(path)
    layout = dict(prepared_layout) if prepared_layout is not None else prepare_header_scope_boundaries(doc, spec)
    findings.extend(copy.deepcopy(layout.get("findings", [])))
    if not rules:
        return {
            "schema_version": "1.1", "valid": not any(x.get("severity") == "error" for x in findings),
            "scoped": False, "source_sha256": source_sha,
            "format_spec_sha256": hashlib.sha256(spec_bytes).hexdigest(),
            "scope_rules": [], "sections": [], "scope_transforms": [], "findings": findings,
        }

    sections = _section_paragraphs(doc)
    if len(sections) != len(doc.sections):
        findings.append({"code": "header_scope_section_parse_failed", "severity": "error",
                         "message": "Could not map source paragraphs to every DOCX section."})
        sections = [[] for _ in doc.sections]
    source_section_indices = layout.get("source_section_indices", [])
    if len(source_section_indices) != len(sections):
        findings.append({"code": "header_scope_source_section_map_mismatch", "severity": "error",
                         "message": "The deterministic source layout transform does not map every prepared section back to the immutable source."})
        source_section_indices = list(range(1, len(sections) + 1))
    rules_by_scope = {str(item["scope"]): item for item in rules}
    toc_locations: list[tuple[int, Any, int]] = []
    appendix_locations: list[tuple[int, Any, int]] = []
    plan_sections: list[dict[str, Any]] = []
    profile = spec.get("thesis_profile") if isinstance(spec.get("thesis_profile"), Mapping) else {}

    for section_index, (section, paragraphs) in enumerate(zip(doc.sections, sections), 1):
        nonblank = [(i, p) for i, p in enumerate(paragraphs) if str(p.text).strip()]
        toc = [(i, p) for i, p in nonblank if _toc_heading(p)]
        headings = [(i, p) for i, p in nonblank if _heading1(p)]
        default_story = _story_details(section.header)
        active = _active_variants(doc, section, spec)
        source_anchor = None
        scope = "unscoped"
        content_mode = "preserve_source"
        fixed_content = None
        requirement_ids: list[str] = []
        clause_ids: list[str] = []
        evidence_ids: list[str] = []
        error = False

        if toc:
            if len(toc) != 1:
                findings.append({"code": "header_scope_toc_anchor_ambiguous", "severity": "error",
                                 "section_index": section_index,
                                 "message": "The source section contains more than one table-of-contents title anchor.",
                                 "evidence": {"paragraphs": [p.text for _, p in toc]}})
                error = True
            if headings:
                findings.append({"code": "header_scope_mixed_toc_and_heading", "severity": "error",
                                 "section_index": section_index,
                                 "message": "A single source section combines a table-of-contents anchor with Heading 1 scope; a section-level header cannot safely serve both.",
                                 "evidence": {"toc": [p.text for _, p in toc], "heading_1": [p.text for _, p in headings]}})
                error = True
            rule = rules_by_scope.get("toc")
            if rule is None:
                findings.append({"code": "header_scope_toc_rule_missing", "severity": "error",
                                 "section_index": section_index,
                                 "message": "A source table-of-contents anchor has no accepted TOC-scoped header rule."})
                error = True
            elif len(toc) == 1:
                scope = "toc"
                anchor_index, anchor_paragraph = toc[0]
                source_anchor = _source_anchor(anchor_paragraph, anchor_index)
                content_mode = str(rule["content_mode"])
                fixed_content = copy.deepcopy(rule["header_content"])
                requirement_ids = list(rule["requirement_ids"])
                clause_ids = list(rule["clause_ids"])
                evidence_ids = list(rule["evidence_ids"])
                toc_locations.append((section_index, anchor_paragraph, anchor_index))
        else:
            appendix_heads = [(i, p) for i, p in headings if _appendix_heading(p.text)]
            chapter_heads = [(i, p) for i, p in headings if _chapter_heading(p.text)]
            if appendix_heads:
                appendix_locations.extend((section_index, p, i) for i, p in appendix_heads)
                other_headings = [(i, p) for i, p in headings if not _appendix_heading(p.text)]
                if other_headings:
                    findings.append({"code": "header_scope_mixed_appendix_and_body", "severity": "error",
                                     "section_index": section_index,
                                     "message": "A source section combines appendix and non-appendix Heading 1 anchors, so their header scopes are not isolated.",
                                     "evidence": {"appendix": [p.text for _, p in appendix_heads],
                                                  "other_heading_1": [p.text for _, p in other_headings]}})
                    error = True
                rule = rules_by_scope.get("appendix")
                if rule is None:
                    findings.append({"code": "header_scope_appendix_rule_missing", "severity": "error",
                                     "section_index": section_index,
                                     "message": "An appendix source anchor has no accepted appendix-scoped header rule.",
                                     "evidence": {"anchors": [p.text for _, p in appendix_heads]}})
                    error = True
                else:
                    scope = "appendix"
                    anchor_index, anchor_paragraph = appendix_heads[0]
                    source_anchor = _source_anchor(anchor_paragraph, anchor_index)
                    content_mode = str(rule["content_mode"])
                    fixed_content = copy.deepcopy(rule["header_content"])
                    requirement_ids = list(rule["requirement_ids"])
                    clause_ids = list(rule["clause_ids"])
                    evidence_ids = list(rule["evidence_ids"])
            elif chapter_heads:
                if len(chapter_heads) != 1:
                    findings.append({"code": "header_scope_chapter_anchor_ambiguous", "severity": "error",
                                     "section_index": section_index,
                                     "message": "A static source section contains multiple chapter Heading 1 anchors; a single section header cannot be mapped to one chapter.",
                                     "evidence": {"anchors": [p.text for _, p in chapter_heads]}})
                    error = True
                anchor_index, anchor_paragraph = chapter_heads[0]
                source_anchor = _source_anchor(anchor_paragraph, anchor_index)
                scope = "chapter"
                rule = rules_by_scope.get("chapter")
                if rule is not None:
                    content_mode = str(rule["content_mode"])
                    fixed_content = copy.deepcopy(rule["header_content"])
                    requirement_ids = list(rule["requirement_ids"])
                    clause_ids = list(rule["clause_ids"])
                    evidence_ids = list(rule["evidence_ids"])
                else:
                    header_text = _normalized_text(default_story["text"])
                    heading_text = _normalized_text(anchor_paragraph.text)
                    if not header_text:
                        findings.append({"code": "header_scope_chapter_source_missing", "severity": "error",
                                         "section_index": section_index,
                                         "message": "The chapter anchor has no source-authored header and no accepted chapter header rule.",
                                         "evidence": {"anchor": anchor_paragraph.text}})
                        error = True
                    elif header_text != heading_text:
                        findings.append({"code": "header_scope_chapter_source_mismatch", "severity": "error",
                                         "section_index": section_index,
                                         "message": "The source chapter header does not match its Heading 1 anchor; no chapter title will be guessed.",
                                         "evidence": {"anchor": anchor_paragraph.text, "source_header": default_story["text"]}})
                        error = True
            elif headings:
                anchor_index, anchor_paragraph = headings[0]
                source_anchor = _source_anchor(anchor_paragraph, anchor_index)
                header_text = _normalized_text(default_story["text"])
                if header_text:
                    if header_text != _normalized_text(anchor_paragraph.text):
                        findings.append({"code": "header_scope_source_heading_mismatch", "severity": "error",
                                         "section_index": section_index,
                                         "message": "A source header is not anchored to the section's Heading 1; the mapping is unresolved.",
                                         "evidence": {"anchor": anchor_paragraph.text, "source_header": default_story["text"]}})
                        error = True
                    else:
                        scope = "source_heading"
                else:
                    scope = "unscoped"
            elif default_story["text"].strip():
                findings.append({"code": "header_scope_source_anchor_missing", "severity": "error",
                                 "section_index": section_index,
                                 "message": "The source contains header text but no TOC or Heading 1 anchor can bind it to a section."})
                error = True

        if error:
            scope = "unresolved"
        if fixed_content is not None and not isinstance(fixed_content.get("left_text"), str):
            findings.append({"code": "header_scope_fixed_text_missing", "severity": "error",
                             "section_index": section_index,
                             "scope": scope,
                             "message": "The scoped fixed header rule has no left_text value."})
            error = True
            scope = "unresolved"

        if source_anchor is not None and scope in {"toc", "chapter", "appendix", "source_heading"}:
            anchor_index = int(source_anchor.get("paragraph_index", 0))
            if any(i < anchor_index and str(paragraph.text).strip() for i, paragraph in nonblank):
                findings.append({
                    "code": "header_scope_anchor_not_section_start", "severity": "error",
                    "section_index": section_index, "scope": scope,
                    "message": "A scoped header anchor is preceded by unrelated content in the same section, so that header would appear outside its source scope.",
                    "evidence": {"source_anchor": source_anchor,
                                 "preceding_paragraphs": [str(p.text) for i, p in nonblank if i < anchor_index]},
                })
                error = True
                scope = "unresolved"

        variant_contracts: dict[str, dict[str, Any]] = {}
        for variant, is_active in active.items():
            source_story = _story_details(getattr(section, f"{variant}_page_header") if variant == "first" else
                                          getattr(section, "even_page_header") if variant == "even" else section.header)
            if content_mode == "fixed" and fixed_content is not None:
                expected = copy.deepcopy(fixed_content)
                expected_text = str(expected.get("left_text", ""))
                mode = "fixed"
            else:
                expected_text = source_story["text"]
                expected = {"left_text": expected_text}
                mode = "preserve_source"
            variant_contracts[variant] = {
                "active": bool(is_active), "content_mode": mode,
                "header_content": expected, "expected_text": expected_text,
                "source_instructions": list(source_story["instructions"]),
                "source_linked_to_previous": bool(source_story["linked_to_previous"]),
            }

        plan_sections.append({
            "section_index": section_index,
            "source_section_index": source_section_indices[section_index - 1],
            "scope": scope,
            "source_anchor": source_anchor,
            "content_mode": "fixed" if fixed_content is not None else "preserve_source",
            "header_content": copy.deepcopy(fixed_content) if fixed_content is not None else None,
            "requirement_ids": requirement_ids,
            "clause_ids": clause_ids,
            "evidence_ids": evidence_ids,
            "source_default_header_sha256": _sha256_text(default_story["text"]),
            "variants": variant_contracts,
        })

    toc_rule = rules_by_scope.get("toc")
    if toc_rule and not toc_locations:
        findings.append({"code": "header_scope_toc_anchor_missing", "severity": "error",
                         "message": "The accepted TOC-scoped header rule matched no table-of-contents anchor in the source DOCX.",
                         "evidence": {"requirement_ids": toc_rule["requirement_ids"],
                                      "clause_ids": toc_rule["clause_ids"]}})
    appendix_rule = rules_by_scope.get("appendix")
    has_appendices = profile.get("has_appendices")
    if appendix_rule and not appendix_locations:
        findings.append({"code": "header_scope_appendix_anchor_missing", "severity": "error",
                         "message": "The accepted appendix-scoped header rule matched no appendix Heading 1 anchor in the source DOCX.",
                         "evidence": {"requirement_ids": appendix_rule["requirement_ids"],
                                      "clause_ids": appendix_rule["clause_ids"]}})
    if appendix_locations and has_appendices is False:
        findings.append({"code": "header_scope_appendix_profile_conflict", "severity": "error",
                         "message": "The source contains appendix anchors while the bound thesis profile says has_appendices=false.",
                         "evidence": {"anchors": [p.text for _, p, _ in appendix_locations]}})
    if has_appendices is True and appendix_locations and not appendix_rule:
        findings.append({"code": "header_scope_appendix_rule_missing", "severity": "error",
                         "message": "The profile and source show appendices, but no accepted appendix header rule exists."})

    valid = not any(item.get("severity") == "error" for item in findings)
    return {
        "schema_version": "1.1", "valid": valid, "scoped": True,
        "source_sha256": source_sha,
        "format_spec_sha256": hashlib.sha256(spec_bytes).hexdigest(),
        "scope_rules": rules,
        "sections": plan_sections,
        "scope_transforms": layout.get("scope_transforms", []),
        "findings": findings,
    }


def strip_flat_header_content(spec: Mapping[str, Any]) -> dict[str, Any]:
    """Return style-only header roles when scoped rules are active."""
    roles = spec.get("roles") if isinstance(spec.get("roles"), Mapping) else {}
    result = copy.deepcopy(dict(roles))
    header = result.get("header")
    if isinstance(header, dict):
        header.pop("header_content", None)
        if not header:
            result.pop("header", None)
    return result
