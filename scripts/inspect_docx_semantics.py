#!/usr/bin/env python3
"""Inspect semantic content that survived into a DOCX using OOXML evidence."""
from __future__ import annotations

import argparse
import json
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from lxml import etree

try:
    from .docx_semantics import (
        all_body_paragraphs,
        all_table_paragraphs,
        has_visible_or_object_content,
        is_figure_caption,
        is_table_caption,
    )
    from .role_registry import normalize_style_name, style_aliases
except ImportError:  # direct script execution
    from docx_semantics import (
        all_body_paragraphs,
        all_table_paragraphs,
        has_visible_or_object_content,
        is_figure_caption,
        is_table_caption,
    )
    from role_registry import normalize_style_name, style_aliases

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
}


def _style_is(paragraph: Any, role: str, *, exclude: set[str] | None = None) -> bool:
    excluded = {normalize_style_name(value) for value in (exclude or set())}
    current = normalize_style_name(paragraph.style.name)
    aliases = {normalize_style_name(value) for value in style_aliases(role)} - excluded
    return current in aliases


def _table_text(doc: Document) -> tuple[int, str | None]:
    paragraphs = [p for p in all_table_paragraphs(doc) if has_visible_or_object_content(p)]
    if not paragraphs:
        return 0, None
    direct = [p for p in paragraphs if _style_is(p, "table_text")]
    if direct:
        return len(direct), direct[0].style.name
    # Frontmatter tables commonly use Normal while data cells consistently use
    # another style (for example Compact).  The dominant non-Normal style is a
    # stable structural signal for the actual table body.
    counts = Counter(p.style.name for p in paragraphs if normalize_style_name(p.style.name) != "normal")
    if not counts:
        return 0, None
    style = min(counts, key=lambda name: (-counts[name], name))
    return counts[style], style


def inspect(path: Path) -> dict[str, Any]:
    doc = Document(path)
    body = list(all_body_paragraphs(doc))
    table_text_count, table_style = _table_text(doc)
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    display_math = len(root.xpath(".//m:oMathPara", namespaces=NS))
    math_objects = len(root.xpath(".//m:oMath", namespaces=NS))
    drawings = len(root.xpath(".//w:drawing", namespaces=NS))
    title_zh = [p for p in doc.paragraphs if p.text.strip() and _style_is(p, "thesis_title_zh")]
    title_en = [p for p in doc.paragraphs if p.text.strip() and _style_is(p, "thesis_title_en", exclude={"Title"})]
    figure_captions = [p for p in doc.paragraphs if is_figure_caption(p)]
    table_captions = [p for p in doc.paragraphs if is_table_caption(p)]
    return {
        "schema_version": "1.0",
        "stage": "docx",
        "source_document": str(path.resolve()),
        "roles": {
            "thesis_title_zh": {"count": len(title_zh), "evidence": "registered-style"},
            "thesis_title_en": {"count": len(title_en), "evidence": "registered-style"},
            "figure": {"count": drawings, "evidence": "w:drawing"},
            "figure_caption": {"count": len(figure_captions), "evidence": "caption-text-pattern"},
            "table": {"count": len(doc.tables), "evidence": "w:tbl"},
            "table_caption": {"count": len(table_captions), "evidence": "caption-text-pattern"},
            "table_text": {"count": table_text_count, "evidence": "w:tc+dominant-style", "style": table_style},
            "equation": {"count": display_math, "evidence": "m:oMathPara"},
            "math_object": {"count": math_objects, "evidence": "m:oMath"},
        },
        "paragraphs": len(doc.paragraphs),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args(argv)
    result = inspect(args.input)
    rendered = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(rendered, encoding="utf-8")
    else:
        print(rendered, end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
