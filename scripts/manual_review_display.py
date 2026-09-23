"""Visible, non-semantic placement and serialized checks for review drafts.

The ledger is never edited here.  A display anchor is not an interpretation
of an unresolved clause, and a visible marker is not evidence of compliance.
"""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from docx_semantics import all_body_paragraphs, is_abstract_body_en, is_abstract_body_zh
from manual_review import HUMAN_MARKER_CATEGORIES

MANUAL_REVIEW_STYLE = "Thesis Manual Review"
MANUAL_REVIEW_PLACEHOLDER_STYLE = "Thesis Manual Review Placeholder"
MANUAL_REVIEW_RED = RGBColor(0xC0, 0x00, 0x00)
MANUAL_REVIEW_CJK_FONT = "Noto Sans SC"
MARKER_START = re.compile(r"^【(MR-\d{4})｜人工待审】")
DISPLAY_STYLES = {MANUAL_REVIEW_STYLE, MANUAL_REVIEW_PLACEHOLDER_STYLE}


def _marker_font(owner: Any) -> None:
    owner.font.name = MANUAL_REVIEW_CJK_FONT
    owner.font.color.rgb = MANUAL_REVIEW_RED
    owner.font.bold = True
    owner.font.hidden = False
    rpr = owner.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), MANUAL_REVIEW_CJK_FONT)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def ensure_manual_review_styles(doc: Document) -> None:
    for name in DISPLAY_STYLES:
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        _marker_font(style)
        style.font.size = Pt(10.5)
        para = style.paragraph_format
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.first_line_indent = para.left_indent = para.right_indent = Pt(0)
        para.space_before, para.space_after = Pt(6), Pt(6)
        para.line_spacing = 1.15
        para.keep_with_next = False
        para.keep_together = True
        para.page_break_before = False
        ppr = style.element.get_or_add_pPr()
        outline = ppr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            ppr.append(outline)
        outline.set(qn("w:val"), "9")  # never becomes a TOC heading


def mark_manual_review_paragraph(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        _marker_font(run)
        run.font.size = Pt(10.5)
        shading = run._r.get_or_add_rPr().find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            run._r.get_or_add_rPr().append(shading)
        shading.set(qn("w:fill"), "FFF2CC")


def _anchor_role(item: dict[str, Any], mappings: dict[str, Any]) -> str | None:
    # Free-text questions, missing official templates and ambiguous dimensions
    # have no trusted document location.  Never classify them by keywords.
    if item.get("source_type") not in {
        "format_constraint", "post_application_validation", "property_receipt",
    }:
        return None
    path = str(item.get("source_text") or "").split("（", 1)[0]
    path = path.removeprefix("content_constraints.")
    if path in {"abstract_zh.target", "abstract_en.target"}:
        return None
    # These are versioned code-owned property names, not natural-language
    # guesses about whether an English paragraph governs a Chinese abstract.
    aliases = {"abstract_zh": "abstract_body_zh", "abstract_en": "abstract_body_en"}
    root = path.split(".", 1)[0]
    role = aliases.get(root, root)
    return role if role in mappings or role in aliases.values() else None


def _anchor_paragraphs(doc: Document, role: str, mappings: dict[str, Any]) -> list[Paragraph]:
    mapping = mappings.get(role)
    names = {
        str(mapping[key]) for key in ("style_name", "source_style_name")
        if isinstance(mapping, dict) and mapping.get(key)
    }
    paragraphs = [p for p in all_body_paragraphs(doc) if p.style.name not in DISPLAY_STYLES]
    candidates = [p for p in paragraphs if p.style.name in names]
    if not candidates and role == "abstract_body_zh":
        candidates = [p for p in paragraphs if is_abstract_body_zh(p)]
    if not candidates and role == "abstract_body_en":
        candidates = [p for p in paragraphs if is_abstract_body_en(p)]
    return candidates


def _block_anchor(paragraph: Paragraph, doc: Document) -> Any:
    """Place table-related notes outside the top-level table, never in a cell."""
    node = paragraph._p
    while node.getparent() is not doc.element.body:
        parent = node.getparent()
        if parent is None:
            return None
        node = parent
    return node


def _text(item: dict[str, Any]) -> str:
    def compact(key: str, limit: int) -> str:
        value = " ".join(str(item.get(key) or "").split())
        return value if len(value) <= limit else value[:limit] + "…（完整内容见审查记录）"

    placeholder = str(item.get("placeholder_text") or "【待人工处理：请核对本项】")
    clauses = ", ".join(str(v) for v in item.get("clause_ids", []))
    questions = ", ".join(str(v) for v in item.get("question_ids", []))
    evidence = ", ".join(str(v) for v in item.get("evidence_ids", []))
    return (
        f"【{item['marker_id']}｜人工待审】{placeholder}"
        + (f"\n条款：{clauses}" if clauses else "")
        + (f"\n问题编号：{questions}" if questions else "")
        + (f"\n证据编号：{evidence}" if evidence else "")
        + f"\n原文/问题：{compact('source_text', 180)}"
        + f"\n待确认：{compact('reason', 180)}"
        + f"\n请在此人工处理：{compact('action', 140) or '核对后填写处理结果。'}"
    )


def _canonical_sha256(value: Any) -> str:
    payload = json.dumps(
        value, ensure_ascii=False, sort_keys=True,
        separators=(",", ":"), allow_nan=False,
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _items(ledger: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not isinstance(ledger, dict):
        return []
    items = ledger.get("items", [])
    ids = [item.get("marker_id") for item in items if isinstance(item, dict)]
    if len(ids) != len(items) or len(set(ids)) != len(ids) or any(
        not isinstance(value, str) or not re.fullmatch(r"MR-\d{4}", value) for value in ids
    ):
        raise ValueError("manual review markers require unique, valid ledger IDs")
    invalid_categories = sorted({
        str(item.get("category") or "")
        for item in items
        if item.get("category") not in HUMAN_MARKER_CATEGORIES
    })
    if invalid_categories:
        raise ValueError(
            "manual review display accepts only unresolved human decisions/inputs; "
            "technical diagnostics must stay in reports: " + ", ".join(invalid_categories)
        )
    return items


def insert_inline_manual_review_markers(
    doc: Document, ledger: dict[str, Any] | None,
    mappings: dict[str, Any] | None = None,
) -> dict[str, dict[str, Any]]:
    items = _items(ledger)
    mappings = mappings or {}
    ensure_manual_review_styles(doc)
    # Resolve all anchors before mutation.  One new marker can never become
    # the anchor of another item or influence role detection.
    anchors: dict[str, tuple[str, Paragraph, Any]] = {}
    for item in items:
        role = _anchor_role(item, mappings)
        candidates = _anchor_paragraphs(doc, role, mappings) if role else []
        if candidates:
            paragraph = candidates[-1]
            block = _block_anchor(paragraph, doc)
            if block is not None:
                anchors[item["marker_id"]] = (role, paragraph, block)
    locations, tails = {}, {}
    for item in items:
        marker_id = item["marker_id"]
        if marker_id not in anchors:
            continue
        role, original, block = anchors[marker_id]
        node = OxmlElement("w:p")
        tails.get(block, block).addnext(node)
        paragraph = Paragraph(node, doc._body)
        paragraph.style = MANUAL_REVIEW_PLACEHOLDER_STYLE
        paragraph.add_run(_text(item))
        mark_manual_review_paragraph(paragraph)
        tails[block] = node
        locations[marker_id] = {
            "location": "inline_after_role", "anchor_role": role,
            "anchor_text": original.text[:240],
            "anchor_kind": "table" if block.tag == qn("w:tbl") else "paragraph",
            "paragraph_text": paragraph.text,
        }
    return locations


def append_manual_review_markers(
    doc: Document, ledger: dict[str, Any] | None,
    inline_locations: dict[str, dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Compatibility entrypoint: unlocated markers now go at document FRONT.

    There is exactly one primary marker per ledger item, not a partial set of
    inline notes plus an end-only checklist masquerading as placeholders.
    """
    items = _items(ledger)
    if not items:
        return []
    locations = dict(inline_locations or {})
    ensure_manual_review_styles(doc)
    # An explicit front section makes the non-submission status visible even
    # when every issue has a body anchor.  A final break keeps the source's
    # first paragraph/table and section properties intact.
    front_nodes = []

    def front(text: str) -> Paragraph:
        node = OxmlElement("w:p")
        paragraph = Paragraph(node, doc._body)
        paragraph.style = MANUAL_REVIEW_STYLE
        paragraph.add_run(text)
        mark_manual_review_paragraph(paragraph)
        front_nodes.append(node)
        return paragraph

    front("待定位人工处理（审查草稿，不可提交）")
    front(
        f"本次共有 {len(items)} 项人工审查记录。明确位置的标记置于对应正文旁；"
        "下列项目尚无可靠位置，请按编号核对并自行处理。红字不替代原文，也不代表问题已经解决。"
    )
    for item in items:
        if item["marker_id"] in locations:
            continue
        paragraph = front(_text(item))
        locations[item["marker_id"]] = {
            "location": "document_front_unlocated", "anchor_role": None,
            "paragraph_text": paragraph.text,
        }
    if len(locations) == len(items) and all(v["location"] == "inline_after_role" for v in locations.values()):
        front("本次没有待定位项目；请在正文中按 MR 编号逐项处理。")
    break_node = OxmlElement("w:p")
    run = OxmlElement("w:r")
    page_break = OxmlElement("w:br")
    page_break.set(qn("w:type"), "page")
    run.append(page_break)
    break_node.append(run)
    front_nodes.append(break_node)
    for index, node in enumerate(front_nodes):
        doc.element.body.insert(index, node)
    binding = ledger.get("binding") if isinstance(ledger, dict) else {}
    binding = binding if isinstance(binding, dict) else {}
    binding_sha256 = _canonical_sha256(binding)
    return [{
        "marker_id": item["marker_id"], "category": item.get("category"),
        "clause_ids": item.get("clause_ids", []), "requirement_ids": item.get("requirement_ids", []),
        "question_ids": item.get("question_ids", []), "evidence_ids": item.get("evidence_ids", []),
        "ledger_binding_sha256": binding_sha256,
        "ledger_item_sha256": _canonical_sha256(item),
        "placeholder_text": item.get("placeholder_text"),
        **locations[item["marker_id"]],
        **({"inline": locations[item["marker_id"]]} if item["marker_id"] in (inline_locations or {}) else {}),
    } for item in items]


def _effective_run_attribute(run: Any, paragraph: Paragraph, doc: Document,
                             element: str, attribute: str = "val") -> str | None:
    """Resolve Word's inherited attributes after it removes redundant rPr."""
    properties = [run._r.rPr]
    for initial in (run.style, paragraph.style):
        style, seen = initial, set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            properties.append(style.element.rPr)
            style = style.base_style
    properties.extend(doc.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr"))
    for rpr in properties:
        if rpr is None:
            continue
        node = rpr.find(qn("w:" + element))
        if node is not None:
            value = node.get(qn("w:" + attribute))
            if value is not None:
                return value
            if element == "vanish":
                return "true"
    return None


def audit_manual_review_markers(path: Path, ledger: dict[str, Any]) -> dict[str, Any]:
    """Check the serialized DOCX, not only insertion-loop receipts.

    Package checks cannot prove font glyph rendering; renderer/visual QA is a
    separate mandatory step.  This audit deliberately does not claim it.
    """
    expected = {item["marker_id"] for item in _items(ledger)}
    expected_items = {item["marker_id"]: item for item in _items(ledger)}
    counts: Counter[str] = Counter()
    style_errors = []
    source_binding_errors = []
    marker_text: dict[str, str] = {}
    doc = Document(path)
    for paragraph in all_body_paragraphs(doc):
        match = MARKER_START.match(paragraph.text)
        if not match:
            continue
        marker_id = match.group(1)
        counts[marker_id] += 1
        marker_text[marker_id] = paragraph.text
        for run in paragraph.runs:
            if not run.text:
                continue
            color = _effective_run_attribute(run, paragraph, doc, "color")
            hidden = _effective_run_attribute(run, paragraph, doc, "vanish")
            cjk_font = _effective_run_attribute(run, paragraph, doc, "rFonts", "eastAsia")
            shading = _effective_run_attribute(run, paragraph, doc, "shd", "fill")
            if (color != "C00000" or hidden not in (None, "0", "false", "off")
                    or not cjk_font or shading != "FFF2CC"):
                style_errors.append(marker_id)
    missing, extra = sorted(expected - counts.keys()), sorted(counts.keys() - expected)
    duplicate = sorted(key for key, count in counts.items() if count != 1)
    marker_bindings = []
    ledger_binding = ledger.get("binding") if isinstance(ledger, dict) else {}
    ledger_binding = ledger_binding if isinstance(ledger_binding, dict) else {}
    ledger_binding_sha256 = _canonical_sha256(ledger_binding)
    for marker_id in sorted(expected):
        item = expected_items[marker_id]
        paragraph_text = marker_text.get(marker_id, "")
        question_ids = sorted(str(value) for value in item.get("question_ids", []) if value)
        evidence_ids = sorted(str(value) for value in item.get("evidence_ids", []) if value)
        expected_refs = (
            (f"问题编号：{', '.join(question_ids)}" if question_ids else None),
            (f"证据编号：{', '.join(evidence_ids)}" if evidence_ids else None),
        )
        if any(value is not None and value not in paragraph_text for value in expected_refs):
            source_binding_errors.append(marker_id)
        marker_bindings.append({
            "marker_id": marker_id,
            "question_ids": question_ids,
            "evidence_ids": evidence_ids,
            "clause_ids": sorted(str(value) for value in item.get("clause_ids", []) if value),
            "requirement_ids": sorted(str(value) for value in item.get("requirement_ids", []) if value),
            "ledger_binding_sha256": ledger_binding_sha256,
            "ledger_item_sha256": _canonical_sha256(item),
        })
    return {
        "valid": not (missing or extra or duplicate or style_errors or source_binding_errors),
        "expected_count": len(expected), "visible_marker_count": sum(counts.values()),
        "missing_ids": missing, "unexpected_ids": extra, "duplicate_ids": duplicate,
        "style_errors": sorted(set(style_errors)),
        "source_binding_errors": sorted(set(source_binding_errors)),
        "marker_bindings": marker_bindings,
        "docx_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "visual_verification": "required", "submission_ready": False,
    }
