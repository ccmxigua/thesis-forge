#!/usr/bin/env python3
"""Rebuild Tianjin University thesis sections, headers, and page numbering.

Policy is taken from the extracted TJU rules in this repository:
- cover material before the TOC is unnumbered and has no header/footer;
- front matter starts at the TOC and uses upper-case Roman numerals from I;
- body starts at the first numbered chapter and restarts at Arabic 1;
- body odd-page header is the current Heading 1 (STYLEREF 1);
- body even-page header is ``天津大学硕士学位论文``;
- numbered chapters and the bibliography start in odd-page sections.

The command never overwrites its input and emits an OOXML evidence report.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"w": W, "r": R}
BODY_HEADING = re.compile(r"^第\s*1\s*章")
NUMBERED_CHAPTER = re.compile(r"^第\s*[一二三四五六七八九十百零〇0-9]+\s*章")


def paragraph_style_name(paragraph) -> str:
    return paragraph.style.name if paragraph.style is not None else ""


def paragraph_is_heading_1(paragraph) -> bool:
    return paragraph_style_name(paragraph).replace(" ", "").lower() in {"heading1", "标题1"}


def find_unique(doc: Document, predicate, description: str):
    matches = [p for p in doc.paragraphs if predicate(p)]
    if len(matches) != 1:
        raise ValueError(f"expected one {description}, found {len(matches)}")
    return matches[0]


def paragraph_section_index(doc: Document, target) -> int:
    index = 0
    for child in doc._body._element:
        if child is target._p:
            return index
        if child.tag == qn("w:p"):
            ppr = child.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                index += 1
    raise ValueError("paragraph is not in the document body")


def strip_relationship_references(sectpr) -> None:
    for child in list(sectpr):
        if child.tag in {qn("w:headerReference"), qn("w:footerReference")}:
            sectpr.remove(child)


def insert_boundary_before(doc: Document, paragraph, section_type: WD_SECTION) -> bool:
    if paragraph_section_index(doc, paragraph) > 0:
        # This is sufficient for the controlled one-pass rebuild: every target
        # is initially in the terminal section, and targets are processed in
        # document order.
        previous = paragraph._p.getprevious()
        if previous is not None:
            ppr = previous.find(qn("w:pPr")) if previous.tag == qn("w:p") else None
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                return False
    body = doc._body._element
    children = list(body)
    target_index = children.index(paragraph._p)
    previous = next((node for node in reversed(children[:target_index]) if node.tag == qn("w:p")), None)
    if previous is None:
        previous = OxmlElement("w:p")
        body.insert(target_index, previous)
    ppr = previous.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        previous.insert(0, ppr)
    if ppr.find(qn("w:sectPr")) is not None:
        return False
    current = doc.sections[paragraph_section_index(doc, paragraph)]._sectPr
    boundary = copy.deepcopy(current)
    strip_relationship_references(boundary)
    ppr.append(boundary)
    typ = boundary.find(qn("w:type"))
    if typ is None:
        typ = OxmlElement("w:type")
        boundary.insert(0, typ)
    typ.set(qn("w:val"), "oddPage" if section_type == WD_SECTION.ODD_PAGE else "nextPage")
    return True


def clear_story(story) -> None:
    for node in list(story._element):
        story._element.remove(node)
    story._element.append(OxmlElement("w:p"))


def clear_paragraph(paragraph) -> None:
    for node in list(paragraph._p):
        if node.tag != qn("w:pPr"):
            paragraph._p.remove(node)


def set_run_font(run, size_pt: float) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), "宋体")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")


def add_simple_field(paragraph, instruction: str, fallback: str, size_pt: float) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "宋体")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(round(size_pt * 2)))
    rpr.extend([fonts, size])
    text = OxmlElement("w:t")
    text.text = fallback
    run.extend([rpr, text])
    field.append(run)
    paragraph._p.append(field)


def set_page_number(section, fmt: str | None, start: int | None = None) -> None:
    pg = section._sectPr.find(qn("w:pgNumType"))
    if fmt is None:
        if pg is not None:
            section._sectPr.remove(pg)
        return
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        section._sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    if start is None:
        pg.attrib.pop(qn("w:start"), None)
    else:
        pg.set(qn("w:start"), str(start))


def set_section_type(section, value: str) -> None:
    typ = section._sectPr.find(qn("w:type"))
    if typ is None:
        typ = OxmlElement("w:type")
        section._sectPr.insert(0, typ)
    typ.set(qn("w:val"), value)


def unlink_and_clear(story) -> None:
    story.is_linked_to_previous = False
    clear_story(story)


def configure_footer(section, instruction: str | None, fallback: str = "1") -> None:
    stories = (section.footer, section.first_page_footer, section.even_page_footer)
    for story in stories:
        unlink_and_clear(story)
    if instruction:
        # odd/even headers are enabled document-wide, so both default (odd)
        # and even footer stories must carry PAGE. Populate first as well to
        # make the section robust if different-first-page is later enabled.
        for story in stories:
            p = story.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            clear_paragraph(p)
            add_simple_field(p, instruction, fallback, 9)


def configure_header(section, *, body: bool) -> None:
    for story in (section.header, section.first_page_header, section.even_page_header):
        unlink_and_clear(story)
    if not body:
        return
    default = section.header.paragraphs[0]
    default.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_paragraph(default)
    add_simple_field(default, "STYLEREF 1 \\* MERGEFORMAT", "章标题", 10.5)
    even = section.even_page_header.paragraphs[0]
    even.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_paragraph(even)
    set_run_font(even.add_run("天津大学硕士学位论文"), 10.5)


def normalize_refs(doc: Document) -> None:
    tags = {qn("w:headerReference"), qn("w:footerReference")}
    for section in doc.sections:
        sectpr = section._sectPr
        refs = [node for node in list(sectpr) if node.tag in tags]
        for node in refs:
            sectpr.remove(node)
        for i, node in enumerate(refs):
            sectpr.insert(i, node)


def inspect_docx(path: Path) -> dict:
    with ZipFile(path) as z:
        root = __import__("lxml.etree", fromlist=["etree"]).fromstring(z.read("word/document.xml"))
        rels = __import__("lxml.etree", fromlist=["etree"]).fromstring(z.read("word/_rels/document.xml.rels"))
        relmap = {r.get("Id"): r.get("Target") for r in rels}
        rows = []
        page_fields = styleref_fields = 0
        for index, sect in enumerate(root.xpath("//w:sectPr", namespaces=NS), 1):
            typ = sect.find("w:type", NS)
            pg = sect.find("w:pgNumType", NS)
            stories = []
            for kind in ("header", "footer"):
                for ref in sect.findall(f"w:{kind}Reference", NS):
                    rid = ref.get(f"{{{R}}}id")
                    target = relmap.get(rid)
                    part = target if target and target.startswith("word/") else f"word/{target}" if target else None
                    if part and part in z.namelist():
                        story = __import__("lxml.etree", fromlist=["etree"]).fromstring(z.read(part))
                        instructions = story.xpath("//w:instrText/text()|//w:fldSimple/@w:instr", namespaces=NS)
                        page_fields += sum(1 for value in instructions if re.match(r"^\s*PAGE\b", value, re.I))
                        styleref_fields += sum(1 for value in instructions if re.match(r"^\s*STYLEREF\b", value, re.I))
                        text = "".join(story.xpath("//w:t/text()", namespaces=NS))
                    else:
                        instructions, text = [], ""
                    stories.append({"kind": kind, "type": ref.get(f"{{{W}}}type"), "part": part,
                                    "text": text, "instructions": instructions})
            rows.append({
                "section": index,
                "type": typ.get(f"{{{W}}}val") if typ is not None else "nextPage(default)",
                "page_number": dict(pg.attrib) if pg is not None else None,
                "stories": stories,
            })
        return {
            "artifact": str(path),
            "section_count": len(rows),
            "header_part_count": len([n for n in z.namelist() if re.match(r"word/header\d+\.xml$", n)]),
            "footer_part_count": len([n for n in z.namelist() if re.match(r"word/footer\d+\.xml$", n)]),
            "page_field_count": page_fields,
            "styleref_field_count": styleref_fields,
            "sections": rows,
        }


def rebuild(input_path: Path, output_path: Path, report_path: Path) -> None:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("input and output must differ")
    doc = Document(input_path)
    toc = find_unique(doc, lambda p: paragraph_style_name(p).lower() == "toc heading" and p.text.strip() == "目 录", "TOC heading")
    body = find_unique(doc, lambda p: paragraph_is_heading_1(p) and BODY_HEADING.match(p.text.strip()), "first chapter heading")
    later = [p for p in doc.paragraphs if paragraph_is_heading_1(p)
             and (NUMBERED_CHAPTER.match(p.text.strip()) or p.text.strip() == "参考文献")
             and p is not body]

    # Insert in document order; body nodes remain stable while boundaries are appended.
    insert_boundary_before(doc, toc, WD_SECTION.NEW_PAGE)
    insert_boundary_before(doc, body, WD_SECTION.ODD_PAGE)
    for heading in later:
        insert_boundary_before(doc, heading, WD_SECTION.ODD_PAGE)

    body_section = paragraph_section_index(doc, body)
    toc_section = paragraph_section_index(doc, toc)
    if toc_section != 1 or body_section != 2:
        raise ValueError(f"unexpected section zoning: toc={toc_section + 1}, body={body_section + 1}")

    doc.settings.odd_and_even_pages_header_footer = True
    for index, section in enumerate(doc.sections):
        section.different_first_page_header_footer = False
        if index == 0:
            set_section_type(section, "nextPage")
            set_page_number(section, None)
            configure_header(section, body=False)
            configure_footer(section, None)
        elif index == 1:
            set_section_type(section, "nextPage")
            set_page_number(section, "upperRoman", 1)
            configure_header(section, body=False)
            configure_footer(section, "PAGE \\* ROMAN", "I")
        else:
            set_section_type(section, "oddPage")
            set_page_number(section, "decimal", 1 if index == 2 else None)
            configure_header(section, body=True)
            configure_footer(section, "PAGE \\* ARABIC", "1")
    normalize_refs(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    report = inspect_docx(output_path)
    report["policy"] = {
        "cover": "unnumbered; no header/footer content",
        "front_matter": "TOC through before chapter 1; upper Roman from I; centered footer",
        "body": "chapter 1 through end; Arabic from 1; odd-page chapter starts",
        "odd_header": "STYLEREF 1",
        "even_header": "天津大学硕士学位论文",
    }
    report["body_start_section"] = body_section + 1
    report["front_matter_start_section"] = toc_section + 1
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    rebuild(args.input, args.output, args.report)
