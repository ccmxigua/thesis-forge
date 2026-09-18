#!/usr/bin/env python3
"""Deterministic official-template evidence and pre-generation reconciliation.

The written requirements remain authoritative.  A supplied official DOCX may
corroborate an explicit value, fill a silent scalar formatting value, or supply
an exact fixed-text resource that the written source explicitly delegates to
the template (for example, ``见例文``).  Missing sample content is never treated
as evidence that a feature is prohibited.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import unicodedata
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET

from docx import Document
from docx.oxml.ns import qn

from analyze_template_styles import analyze as analyze_styles


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
Q = lambda name: f"{{{W}}}{name}"
SCHEMA_VERSION = "1.0"
RESULT_STATES = {"agree", "conflict", "insufficient", "not_applicable"}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_record(path: Path, kind: str) -> dict[str, Any]:
    resolved = path.resolve()
    return {
        "kind": kind,
        "path": str(resolved),
        "bytes": resolved.stat().st_size,
        "sha256": _sha256(resolved),
    }


def _normalized_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"\s+", "", text).strip(" ：:。；;，,、.!！？?")


def _normalized_value(value: Any) -> Any:
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, (int, float)):
        return round(float(value), 3)
    if isinstance(value, str):
        normalized = _normalized_text(value).casefold()
        return {
            "宋体": "simsun", "黑体": "simhei", "楷体": "kaiti",
            "仿宋": "fangsong", "微软雅黑": "microsoftyahei",
            "timesnewroman": "timesnewroman",
        }.get(normalized, normalized)
    if isinstance(value, list):
        return [_normalized_value(item) for item in value]
    if isinstance(value, dict):
        return {key: _normalized_value(value[key]) for key in sorted(value)}
    return value


def _values_equal(first: Any, second: Any) -> bool:
    if (isinstance(first, (int, float)) and not isinstance(first, bool)
            and isinstance(second, (int, float)) and not isinstance(second, bool)):
        return abs(float(first) - float(second)) <= 0.05
    return _normalized_value(first) == _normalized_value(second)


def _provenance(source: dict[str, Any]) -> dict[str, Any]:
    return {
        "source_kind": source["kind"],
        "source_path": source["path"],
        "source_sha256": source["sha256"],
        "extractor": "template_reconciliation.extract_template_evidence",
        "extractor_version": SCHEMA_VERSION,
    }


def _paragraph_text(node: ET.Element) -> str:
    return "".join(item.text or "" for item in node.findall(".//w:t", NS)).strip()


def _paragraph_style_id(node: ET.Element) -> str | None:
    style = node.find("w:pPr/w:pStyle", NS)
    return style.get(Q("val")) if style is not None else None


def _paragraph_alignment(node: ET.Element) -> str | None:
    alignment = node.find("w:pPr/w:jc", NS)
    value = alignment.get(Q("val")) if alignment is not None else None
    aliases = {"both": "justify", "distribute": "distributed"}
    return aliases.get(value, value)


def _toggle_value(node: ET.Element | None) -> bool | None:
    if node is None:
        return None
    value = str(node.get(Q("val"), "1")).lower()
    return value not in {"0", "false", "off", "none"}


def _direct_paragraph_properties(node: ET.Element) -> dict[str, Any]:
    """Return only uniform, explicit paragraph/run properties.

    Missing run properties are not interpreted as false/default.  A run value
    is promoted only when every visible run explicitly serializes the same
    value, which keeps mixed direct formatting in the insufficient state.
    """
    result: dict[str, Any] = {}
    paragraph: dict[str, Any] = {}
    alignment = _paragraph_alignment(node)
    if alignment:
        paragraph["alignment"] = alignment
    ppr = node.find("w:pPr", NS)
    if ppr is not None:
        ind = ppr.find("w:ind", NS)
        if ind is not None:
            for attribute, key in (("firstLine", "first_line_indent_pt"),
                                   ("left", "left_indent_pt"), ("right", "right_indent_pt")):
                raw = ind.get(Q(attribute))
                if raw is not None:
                    paragraph[key] = round(int(raw) / 20, 3)
        spacing = ppr.find("w:spacing", NS)
        if spacing is not None:
            for attribute, key in (("before", "space_before_pt"), ("after", "space_after_pt")):
                raw = spacing.get(Q(attribute))
                if raw is not None:
                    paragraph[key] = round(int(raw) / 20, 3)
            line = spacing.get(Q("line"))
            rule = spacing.get(Q("lineRule"))
            if line is not None:
                if rule in {"exact", "atLeast"}:
                    paragraph["line_spacing"] = {
                        "type": "exact" if rule == "exact" else "at_least",
                        "value": round(int(line) / 20, 3), "unit": "pt",
                    }
                elif rule in {None, "auto"}:
                    multiple = round(int(line) / 240, 3)
                    paragraph["line_spacing"] = {
                        "type": "multiple", "value": multiple, "unit": "multiple",
                    }
    if paragraph:
        result["paragraph"] = paragraph

    run_values: list[dict[str, Any]] = []
    for run in node.findall("w:r", NS):
        if not _paragraph_text(run):
            continue
        rpr = run.find("w:rPr", NS)
        values: dict[str, Any] = {}
        if rpr is not None:
            fonts = rpr.find("w:rFonts", NS)
            if fonts is not None:
                cjk = fonts.get(Q("eastAsia"))
                latin = fonts.get(Q("ascii")) or fonts.get(Q("hAnsi"))
                if cjk:
                    values["cjk"] = cjk
                if latin:
                    values["latin"] = latin
            size = rpr.find("w:sz", NS)
            if size is not None and size.get(Q("val")):
                values["size_pt"] = round(int(size.get(Q("val"))) / 2, 3)
            for tag, key in (("b", "bold"), ("i", "italic")):
                toggled = _toggle_value(rpr.find(f"w:{tag}", NS))
                if toggled is not None:
                    values[key] = toggled
        run_values.append(values)
    if run_values:
        font: dict[str, Any] = {}
        for key in ("cjk", "latin", "size_pt", "bold", "italic"):
            values = [item.get(key) for item in run_values]
            if all(value is not None for value in values) and all(
                    _values_equal(values[0], value) for value in values[1:]):
                font[key] = values[0]
        if font:
            result["font"] = font
    return result


def _style_names(zf: zipfile.ZipFile) -> dict[str, str]:
    if "word/styles.xml" not in zf.namelist():
        return {}
    root = ET.fromstring(zf.read("word/styles.xml"))
    result: dict[str, str] = {}
    for style in root.findall("w:style", NS):
        style_id = style.get(Q("styleId"))
        name = style.find("w:name", NS)
        if style_id and name is not None and name.get(Q("val")):
            result[style_id] = str(name.get(Q("val")))
    return result


def _paragraph_records(zf: zipfile.ZipFile) -> list[dict[str, Any]]:
    styles = _style_names(zf)
    parts = ["word/document.xml", *sorted(
        name for name in zf.namelist()
        if re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
    )]
    records: list[dict[str, Any]] = []
    for part in parts:
        root = ET.fromstring(zf.read(part))
        for index, paragraph in enumerate(root.findall(".//w:p", NS)):
            text = _paragraph_text(paragraph)
            if not text:
                continue
            style_id = _paragraph_style_id(paragraph)
            records.append({
                "text": text,
                "normalized_text": _normalized_text(text),
                "style_id": style_id,
                "style_name": styles.get(style_id),
                "alignment": _paragraph_alignment(paragraph),
                "direct_properties": _direct_paragraph_properties(paragraph),
                "source_location": {
                    "part": part,
                    "paragraph_index": index,
                    "style_id": style_id,
                },
            })
    return records


def _clean_mapping_properties(value: Any) -> Any:
    if isinstance(value, dict):
        cleaned = {key: _clean_mapping_properties(item) for key, item in value.items()}
        return {key: item for key, item in cleaned.items() if item not in (None, {}, [])}
    if isinstance(value, list):
        return [_clean_mapping_properties(item) for item in value]
    if isinstance(value, str) and re.match(r"^[A-Z_]+ \(\d+\)$", value):
        aliases = {"BOTH": "justify", "DISTRIBUTE": "distributed"}
        name = value.split(" ", 1)[0]
        return aliases.get(name, name.lower())
    return value


def _overlay(target: dict[str, Any], update: dict[str, Any]) -> None:
    for key, value in update.items():
        if isinstance(value, dict) and isinstance(target.get(key), dict):
            _overlay(target[key], value)
        else:
            target[key] = copy.deepcopy(value)


def _inherited(style: Any, getter: Callable[[Any], Any]) -> Any:
    seen: set[str] = set()
    current = style
    while current is not None and current.style_id not in seen:
        seen.add(current.style_id)
        value = getter(current)
        if value is not None:
            return value
        current = current.base_style
    return None


def _style_properties(style: Any) -> dict[str, Any]:
    def cjk(current: Any) -> str | None:
        rpr = current.element.rPr
        fonts = rpr.rFonts if rpr is not None else None
        return fonts.get(qn("w:eastAsia")) if fonts is not None else None

    font: dict[str, Any] = {}
    latin = _inherited(style, lambda current: current.font.name)
    east_asia = _inherited(style, cjk)
    size = _inherited(style, lambda current: current.font.size)
    bold = _inherited(style, lambda current: current.font.bold)
    italic = _inherited(style, lambda current: current.font.italic)
    if latin is not None:
        font["latin"] = latin
    if east_asia is not None:
        font["cjk"] = east_asia
    if size is not None:
        font["size_pt"] = round(size.pt, 3)
    if bold is not None:
        font["bold"] = bool(bold)
    if italic is not None:
        font["italic"] = bool(italic)

    paragraph: dict[str, Any] = {}
    alignment = _inherited(style, lambda current: current.paragraph_format.alignment)
    first_indent = _inherited(style, lambda current: current.paragraph_format.first_line_indent)
    left_indent = _inherited(style, lambda current: current.paragraph_format.left_indent)
    right_indent = _inherited(style, lambda current: current.paragraph_format.right_indent)
    before = _inherited(style, lambda current: current.paragraph_format.space_before)
    after = _inherited(style, lambda current: current.paragraph_format.space_after)
    if alignment is not None:
        paragraph["alignment"] = _clean_mapping_properties(str(alignment))
    for key, length in (
        ("first_line_indent_pt", first_indent), ("left_indent_pt", left_indent),
        ("right_indent_pt", right_indent), ("space_before_pt", before),
        ("space_after_pt", after),
    ):
        if length is not None:
            paragraph[key] = round(length.pt, 3)
    result: dict[str, Any] = {}
    if font:
        result["font"] = font
    if paragraph:
        result["paragraph"] = paragraph
    return result


def _add_item(items: list[dict[str, Any]], source: dict[str, Any], **payload: Any) -> dict[str, Any]:
    item = {
        "id": f"TE{len(items) + 1:05d}",
        **payload,
        "provenance": _provenance(source),
    }
    if "source_location" not in item:
        raise ValueError("template evidence item is missing source_location")
    items.append(item)
    return item


def _signature_placeholders(texts: list[str]) -> list[dict[str, str]]:
    joined = "\n".join(texts)
    result: list[dict[str, str]] = []
    for role, pattern, label in (
        ("author", r"(?:作者|本人)(?:签名|签字)", "作者签名"),
        ("supervisor", r"(?:导师|指导教师)(?:签名|签字)", "导师签名"),
        ("date", r"(?:日期|年\s*月\s*日)", "日期"),
    ):
        if re.search(pattern, joined) and role not in {item["role"] for item in result}:
            result.append({"role": role, "label": label,
                           "attestation_scope": "placeholder_presence_only"})
    return result


def _declaration_blocks(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    body = [item for item in records if item["source_location"]["part"] == "word/document.xml"]
    heading_pattern = re.compile(r"(?:原创性|独创性|诚信|使用授权|版权授权|公开授权).{0,12}(?:声明|说明|书)$|^(?:声明|授权书)$")
    boundary_pattern = re.compile(r"^(?:摘\s*要|ABSTRACT|目\s*录|参考文献|第[一二三四五六七八九十\d]+章)$", re.I)
    headings = [index for index, item in enumerate(body)
                if len(item["normalized_text"]) <= 50 and heading_pattern.search(item["normalized_text"])]
    blocks: list[dict[str, Any]] = []
    for heading_pos, start in enumerate(headings):
        stop = headings[heading_pos + 1] if heading_pos + 1 < len(headings) else min(len(body), start + 14)
        candidates: list[dict[str, Any]] = []
        observed_stop = stop
        for offset, item in enumerate(body[start + 1:stop], start + 1):
            if boundary_pattern.match(item["normalized_text"]):
                observed_stop = offset
                break
            candidates.append(item)
        fixed_body = [item for item in candidates if not (
            len(item["normalized_text"]) <= 80 and re.search(
                r"(?:作者|本人|导师|指导教师)(?:签名|签字)|日期.{0,12}年?.{0,6}月?.{0,6}日?",
                item["text"],
            )
        )]
        if not fixed_body:
            continue
        later_texts = [item["normalized_text"] for item in body[observed_stop:]]
        if any(re.fullmatch(r"摘\s*要", text) for text in later_texts):
            before_role: str | None = "abstract_title_zh"
        elif start == 0:
            before_role = "document_start"
        else:
            before_role = None
        locations = [body[start]["source_location"], *[item["source_location"] for item in candidates]]
        blocks.append({
            "heading": body[start]["text"],
            "body_parts": [item["text"] for item in fixed_body],
            "signature_placeholders": _signature_placeholders([item["text"] for item in candidates]),
            "before_role": before_role,
            "source_location": {"part": "word/document.xml", "paragraphs": locations},
        })
    return blocks


def extract_template_evidence(path: Path) -> dict[str, Any]:
    """Extract only directly observable facts from an official DOCX."""
    source = file_record(path.resolve(), "official_template_docx")
    document = Document(path)
    with zipfile.ZipFile(path) as zf:
        records = _paragraph_records(zf)
        field_parts = ["word/document.xml", *sorted(
            name for name in zf.namelist()
            if re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
        )]
        field_roots = [(part, ET.fromstring(zf.read(part))) for part in field_parts]
        document_root = field_roots[0][1]

    items: list[dict[str, Any]] = []
    for section_index, section in enumerate(document.sections, 1):
        width_pt = round(section.page_width.pt, 3)
        height_pt = round(section.page_height.pt, 3)
        page_size = "A4" if (
            (abs(section.page_width.mm - 210) < 1 and abs(section.page_height.mm - 297) < 1)
            or (abs(section.page_width.mm - 297) < 1 and abs(section.page_height.mm - 210) < 1)
        ) else None
        location = {"part": "word/document.xml", "section_index": section_index,
                    "element": "w:sectPr"}
        _add_item(items, source, category="page_observation", semantic_role="page",
                  property="dimensions_pt", value={"width": width_pt, "height": height_pt},
                  source_location=location)
        if page_size:
            _add_item(items, source, category="page_setting", semantic_role="page",
                      property="size", value=page_size, source_location=location)
        _add_item(items, source, category="page_setting", semantic_role="page",
                  property="orientation",
                  value="landscape" if width_pt > height_pt else "portrait",
                  source_location=location)
        for key, length in (
            ("top", section.top_margin), ("bottom", section.bottom_margin),
            ("left", section.left_margin), ("right", section.right_margin),
            ("header", section.header_distance), ("footer", section.footer_distance),
            ("gutter", section.gutter),
        ):
            if length is not None:
                _add_item(items, source, category="page_setting", semantic_role="page",
                          property=f"margins_pt.{key}", value=round(length.pt, 3),
                          source_location=location)
        if section.different_first_page_header_footer:
            _add_item(items, source, category="page_setting", semantic_role="page",
                      property="different_first_page", value=True, source_location=location)
    if getattr(document.settings, "odd_and_even_pages_header_footer", False):
        _add_item(items, source, category="page_setting", semantic_role="page",
                  property="different_odd_even", value=True,
                  source_location={"part": "word/settings.xml", "element": "w:evenAndOddHeaders"})

    for section_index, sect_pr in enumerate(document_root.findall(".//w:sectPr", NS), 1):
        page_number_type = sect_pr.find("w:pgNumType", NS)
        if page_number_type is None:
            continue
        raw = {key: page_number_type.get(Q(key)) for key in ("fmt", "start", "chapStyle", "chapSep")
               if page_number_type.get(Q(key)) is not None}
        _add_item(items, source, category="page_observation", semantic_role="page",
                  property="page_number.serialized_section_format", value=raw,
                  source_location={"part": "word/document.xml", "section_index": section_index,
                                   "element": "w:sectPr/w:pgNumType"})

    page_field_locations: list[dict[str, Any]] = []
    for part, root in field_roots:
        instructions = [node.text or "" for node in root.findall(".//w:instrText", NS)]
        instructions += [node.get(Q("instr"), "") for node in root.findall(".//w:fldSimple", NS)]
        for field_index, instruction in enumerate(instructions):
            if re.search(r"\bPAGE\b", instruction, re.I):
                location = {"part": part, "field_index": field_index, "element": "w:instrText|w:fldSimple"}
                page_field_locations.append(location)
                _add_item(items, source, category="page_field", semantic_role="page",
                          property="page_number.preserve_existing_locations", value=True,
                          field_instruction=instruction.strip(), source_location=location)

    style_analysis = analyze_styles(path)
    style_by_name = {style.name: style for style in document.styles if style.type == 1}
    direct_role_styles: dict[str, str] = {}
    for role, mapping in sorted(style_analysis.get("mappings", {}).items()):
        style_name = mapping.get("style_name")
        matching = [item for item in records if item.get("style_name") == style_name]
        reasons = list(mapping.get("reasons") or [])
        direct = bool(matching) and (
            any(reason.startswith("sample-text:") for reason in reasons)
            or role in {"header", "footer", "table_text", "equation"}
            or any(reason.startswith("style-name:") for reason in reasons)
        )
        style = style_by_name.get(style_name)
        base_properties = _style_properties(style) if style is not None else {}
        if not direct:
            continue
        direct_role_styles[role] = style_name
        for observed in matching[:8]:
            properties = copy.deepcopy(base_properties)
            _overlay(properties, observed.get("direct_properties") or {})
            if not properties:
                continue
            _add_item(
                items, source, category="semantic_style", semantic_role=role,
                property="style_properties", value=properties, direct_semantic=True,
                style_name=style_name, mapping_score=mapping.get("score"), mapping_reasons=reasons,
                observed_texts=[observed["text"]],
                source_location={
                    "part": observed["source_location"]["part"],
                    "paragraph": observed["source_location"],
                    "style_location": {"part": "word/styles.xml", "style_name": style_name},
                },
            )

    ordered_roles: list[str] = []
    for record in records:
        role = next((candidate for candidate, style_name in direct_role_styles.items()
                     if style_name == record.get("style_name")), None)
        if role and role not in ordered_roles and record["source_location"]["part"] == "word/document.xml":
            ordered_roles.append(role)
            _add_item(items, source, category="structure", semantic_role=role,
                      property="observed_order", value=len(ordered_roles),
                      exact_text=record["text"], source_location=record["source_location"])
    if ordered_roles:
        _add_item(items, source, category="structure", semantic_role="document_structure",
                  property="ordered_roles", value=ordered_roles,
                  source_location={"part": "word/document.xml",
                                   "observations": [item["source_location"] for item in records
                                                    if item["source_location"]["part"] == "word/document.xml"]})

    for record in records:
        _add_item(items, source, category="fixed_text", semantic_role="fixed_text",
                  property="text", value=record["text"], normalized_text=record["normalized_text"],
                  style_name=record.get("style_name"), source_location=record["source_location"])
    declaration_ids: list[str] = []
    for block in _declaration_blocks(records):
        item = _add_item(items, source, category="fixed_body", semantic_role="declarations",
                         property="items", value={key: block[key] for key in
                         ("heading", "body_parts", "signature_placeholders", "before_role")},
                         source_location=block["source_location"])
        declaration_ids.append(item["id"])

    counts: dict[str, int] = {}
    for item in items:
        counts[item["category"]] = counts.get(item["category"], 0) + 1
    return {
        "schema_version": SCHEMA_VERSION,
        "status": "extracted",
        "source": source,
        "policy": {
            "observable_facts_only": True,
            "sample_absence_is_prohibition": False,
            "fixed_text_requires_exact_extraction": True,
        },
        "items": items,
        "declaration_evidence_ids": declaration_ids,
        "summary": {"item_count": len(items), "counts": counts,
                    "page_field_locations": len(page_field_locations)},
    }


def _flatten(value: Any, prefix: str = "") -> dict[str, Any]:
    if isinstance(value, dict):
        result: dict[str, Any] = {}
        for key, item in value.items():
            path = f"{prefix}.{key}" if prefix else key
            result.update(_flatten(item, path))
        return result
    return {prefix: value}


def _get_nested(value: dict[str, Any], path: str) -> tuple[bool, Any]:
    current: Any = value
    for key in path.split("."):
        if not isinstance(current, dict) or key not in current:
            return False, None
        current = current[key]
    return True, current


def _set_nested(value: dict[str, Any], path: str, item: Any) -> None:
    current = value
    parts = path.split(".")
    for key in parts[:-1]:
        current = current.setdefault(key, {})
    current[parts[-1]] = copy.deepcopy(item)


def _next_requirement_id(spec: dict[str, Any]) -> str:
    used = {str(item.get("id")) for item in spec.get("requirements", []) if isinstance(item, dict)}
    number = 1
    while f"R{number:05d}" in used:
        number += 1
    return f"R{number:05d}"


def _evidence_index(evidence: dict[str, Any]) -> dict[tuple[str, str], list[dict[str, Any]]]:
    result: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for item in evidence.get("items", []):
        role = item.get("semantic_role")
        category = item.get("category")
        if category == "semantic_style" and item.get("direct_semantic"):
            for path, value in _flatten(item.get("value", {})).items():
                result.setdefault((str(role), path), []).append({**item, "leaf_value": value})
        elif category in {"page_setting", "page_field", "structure"}:
            result.setdefault((str(role), str(item.get("property"))), []).append(
                {**item, "leaf_value": item.get("value")}
            )
    return result


def _comparison(role: str, path: str, written_value: Any,
                index: dict[tuple[str, str], list[dict[str, Any]]]) -> tuple[str, list[dict[str, Any]]]:
    if role == "document_structure" and path in {"ordered_roles", "required_roles"}:
        candidates = index.get((role, "ordered_roles"), [])
        if not candidates or not isinstance(written_value, list):
            return "insufficient", candidates
        states: list[str] = []
        for item in candidates:
            observed = item.get("leaf_value")
            if not isinstance(observed, list) or not all(value in observed for value in written_value):
                # A sample missing a role cannot prove that role is prohibited.
                states.append("insufficient")
                continue
            if path == "required_roles":
                states.append("agree")
                continue
            positions = [observed.index(value) for value in written_value]
            states.append("agree" if positions == sorted(positions) else "conflict")
        if states and all(state == "agree" for state in states):
            return "agree", candidates
        if states and all(state == "conflict" for state in states):
            return "conflict", candidates
        return "insufficient", candidates
    candidates = index.get((role, path), [])
    if not candidates:
        return "insufficient", []
    equal = [item for item in candidates if _values_equal(written_value, item.get("leaf_value"))]
    distinct = {_normalized_text(json.dumps(_normalized_value(item.get("leaf_value")), ensure_ascii=False, sort_keys=True))
                for item in candidates}
    if equal and len(equal) == len(candidates):
        return "agree", candidates
    if not equal and len(distinct) == 1:
        return "conflict", candidates
    return "insufficient", candidates


def _manual_markers(text: str) -> list[str]:
    known = [
        marker for marker in ("论文作者", "作者签名", "指导教师", "导师签名", "提交日期",
                              "原创性声明", "独创性声明", "使用授权声明")
        if marker in text
    ]
    if known:
        return known
    return [token for token in re.findall(r"[\u4e00-\u9fff]{2,12}", text)
            if token not in {"论文", "要求", "模板", "例文", "格式"}][:3]


def _manual_corroboration(text: str, evidence: dict[str, Any]) -> list[dict[str, Any]]:
    markers = _manual_markers(text)
    if not markers:
        return []
    fixed = [item for item in evidence.get("items", []) if item.get("category") == "fixed_text"]
    matched = [item for item in fixed if any(
        _normalized_text(marker) in str(item.get("normalized_text") or "") for marker in markers
    )]
    corpus = "".join(str(item.get("normalized_text") or "") for item in fixed)
    return matched if all(_normalized_text(marker) in corpus for marker in markers) else []


def _remove_unresolved(spec: dict[str, Any], clause_id: str, *, covered: bool) -> None:
    completeness = spec.get("completeness")
    if not isinstance(completeness, dict):
        return
    completeness["unresolved_clause_ids"] = [
        value for value in completeness.get("unresolved_clause_ids", []) if str(value) != clause_id
    ]
    if covered:
        values = {str(value) for value in completeness.get("covered_clause_ids", [])}
        values.add(clause_id)
        completeness["covered_clause_ids"] = sorted(values)


def _upsert_clause_record(spec: dict[str, Any], clause: dict[str, Any], status: str,
                          reason: str, requirement_ids: list[str]) -> None:
    records = spec.get("clause_compliance")
    if not isinstance(records, list):
        return
    record = next((item for item in records if str(item.get("clause_id")) == str(clause.get("id"))), None)
    if record is None:
        record = {"clause_id": str(clause["id"]), "evidence_ids": list(clause.get("evidence_ids", []))}
        records.append(record)
    record.update({"scope": "docx", "status": status, "requirement_ids": requirement_ids,
                   "reason": reason})
    if status == "pending_execution":
        record["enforcement"] = "generate_or_repair"
    else:
        record.pop("enforcement", None)


def _declaration_resolution(spec: dict[str, Any], clauses: list[dict[str, Any]],
                            questions: list[dict[str, Any]], evidence: dict[str, Any],
                            results: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    blocks = [item for item in evidence.get("items", [])
              if item.get("category") == "fixed_body" and item.get("semantic_role") == "declarations"]
    if not blocks or isinstance(spec.get("declarations"), dict):
        return questions, []
    clause = next((item for item in clauses if re.search(
        r"(?:声明|授权).{0,20}(?:见|参照|按照).{0,8}(?:例文|模板|样例)|(?:见|参照).{0,8}(?:例文|模板).{0,20}(?:声明|授权)",
        str(item.get("text") or ""))), None)
    if clause is None:
        return questions, []
    before_roles = {str((item.get("value") or {}).get("before_role")) for item in blocks
                    if (item.get("value") or {}).get("before_role")}
    if len(before_roles) != 1:
        results.append({
            "id": f"TR{len(results) + 1:05d}", "result": "insufficient",
            "action": "manual_review_retained", "clause_id": str(clause["id"]),
            "semantic_role": "declarations", "property": "before_role",
            "written": {"value": str(clause.get("text") or ""),
                        "evidence_ids": list(clause.get("evidence_ids", []))},
            "official_template": {"values": sorted(before_roles),
                                  "evidence_ids": [item["id"] for item in blocks]},
            "reason": "Exact declaration text was observable, but one safe insertion anchor was not; no resource was executed.",
        })
        return questions, []
    declaration_items: list[dict[str, Any]] = []
    template_ids: list[str] = []
    for index, block in enumerate(blocks, 1):
        value = block.get("value") or {}
        declaration_items.append({
            "id": f"template_statement_{index}",
            "heading": value["heading"],
            "body_parts": list(value["body_parts"]),
            "source_evidence_ids": [block["id"]],
            "signature_placeholders": copy.deepcopy(value.get("signature_placeholders") or []),
        })
        template_ids.append(block["id"])
    properties = {"before_role": next(iter(before_roles)), "items": declaration_items}
    requirement_id = _next_requirement_id(spec)
    source_ids = sorted({str(value) for value in clause.get("evidence_ids", [])} | set(template_ids))
    spec["declarations"] = copy.deepcopy(properties)
    spec.setdefault("requirements", []).append({
        "id": requirement_id, "role": "declarations", "properties": copy.deepcopy(properties),
        "evidence_ids": source_ids, "clause_ids": [str(clause["id"])],
        "resolved_by": "template", "confidence": 1.0,
        "source_text": str(clause.get("text") or ""),
        "reason": "The written requirement explicitly refers to the official example; exact declaration text was extracted from that DOCX.",
    })
    _remove_unresolved(spec, str(clause["id"]), covered=True)
    _upsert_clause_record(
        spec, clause, "pending_execution",
        "Exact declaration text was materialized from the official template explicitly referenced by the written clause.",
        [requirement_id],
    )
    remaining = [item for item in questions if str(item.get("clause_id")) != str(clause["id"])]
    results.append({
        "id": f"TR{len(results) + 1:05d}", "result": "agree", "action": "resource_from_template",
        "clause_id": str(clause["id"]), "requirement_id": requirement_id,
        "semantic_role": "declarations", "property": "items",
        "written": {"value": str(clause.get("text") or ""),
                    "evidence_ids": list(clause.get("evidence_ids", []))},
        "official_template": {"value": [item["value"] for item in blocks],
                              "evidence_ids": template_ids},
        "reason": "Written cross-reference and exact official-template declaration bodies agree.",
    })
    return remaining, template_ids


def reconcile_template(
    spec: dict[str, Any], clauses: list[dict[str, Any]], questions: list[dict[str, Any]],
    evidence: dict[str, Any], *, sources: dict[str, Any],
    property_parser: Callable[[str], dict[str, Any]] | None = None,
    page_property_parser: Callable[[str], dict[str, Any]] | None = None,
    role_candidates: Callable[[dict[str, Any]], list[str]] | None = None,
    fill_silent_values: bool = True,
) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    """Reconcile a merged semantic spec against direct official evidence."""
    working = copy.deepcopy(spec)
    remaining_questions = copy.deepcopy(questions)
    index = _evidence_index(evidence)
    results: list[dict[str, Any]] = []
    genuine_conflicts: list[dict[str, Any]] = []
    clause_map = {str(item.get("id")): item for item in clauses}

    for requirement in list(working.get("requirements", [])):
        role = str(requirement.get("role") or "")
        clause_ids = [str(value) for value in requirement.get("clause_ids", [])]
        if (requirement.get("applicability") or {}).get("status") == "excluded":
            results.append({
                "id": f"TR{len(results) + 1:05d}", "result": "not_applicable", "action": "none",
                "requirement_id": requirement.get("id"), "clause_ids": clause_ids,
                "semantic_role": role, "property": None,
                "reason": "The written applicability contract explicitly excludes this requirement.",
            })
            continue
        properties = requirement.get("properties")
        if not isinstance(properties, dict) or not properties:
            # Empty executable requirements are never retained.  A content
            # instance can be repaired only from its exact registered text.
            instances = {str(item.get("id")): item for item in working.get("content_instances", [])
                         if isinstance(item, dict)}
            texts = {str(instances[item_id].get("text") or "").strip()
                     for item_id in requirement.get("field_instance_ids", [])
                     if item_id in instances and str(instances[item_id].get("text") or "").strip()}
            if len(texts) == 1:
                requirement["properties"] = {"text": texts.pop()}
                properties = requirement["properties"]
            else:
                working["requirements"].remove(requirement)
                for clause_id in clause_ids:
                    clause = clause_map.get(clause_id, {"id": clause_id, "evidence_ids": []})
                    remaining_ids = [
                        str(item.get("id")) for item in working.get("requirements", [])
                        if isinstance(item, dict) and clause_id in {
                            str(value) for value in item.get("clause_ids", [])
                        }
                    ]
                    if remaining_ids:
                        _remove_unresolved(working, clause_id, covered=True)
                        _upsert_clause_record(
                            working, clause, "pending_execution",
                            "The empty requirement was removed; other non-empty requirements still execute this clause.",
                            remaining_ids,
                        )
                    else:
                        _remove_unresolved(working, clause_id, covered=False)
                        _upsert_clause_record(
                            working, clause, "unverifiable",
                            "No executable property was established; the clause remains an explicit manual check.", [],
                        )
                results.append({
                    "id": f"TR{len(results) + 1:05d}", "result": "insufficient",
                    "action": "manual_review_retained", "requirement_id": requirement.get("id"),
                    "clause_ids": clause_ids, "semantic_role": role, "property": None,
                    "reason": "An executable requirement with empty properties was removed; template absence or labels alone cannot create a universal rule.",
                })
                continue
        if role == "declarations":
            blocks = [item for item in evidence.get("items", [])
                      if item.get("category") == "fixed_body"
                      and item.get("semantic_role") == "declarations"]
            for written_item in properties.get("items", []) if isinstance(properties.get("items"), list) else []:
                if not isinstance(written_item, dict):
                    continue
                heading = str(written_item.get("heading") or "")
                body_parts = written_item.get("body_parts")
                if not isinstance(body_parts, list):
                    body = written_item.get("body")
                    body_parts = [body] if isinstance(body, str) and body.strip() else []
                # A declaration requirement can contain several items and
                # several clause occurrences.  Attach a clarification to
                # the conflicting item's source evidence, not the first
                # clause of the whole requirement.
                item_evidence_ids = {
                    str(value) for value in (written_item.get("source_evidence_ids") or [])
                    if value is not None
                }
                item_clause_ids = [
                    str(clause_id) for clause_id in clause_ids
                    if item_evidence_ids
                    and item_evidence_ids.intersection({
                        str(value) for value in (clause_map.get(str(clause_id), {}).get("evidence_ids") or [])
                    })
                ]
                item_question_clause_id = (
                    item_clause_ids[0] if item_clause_ids else (clause_ids[0] if clause_ids else None)
                )
                heading_matches = [item for item in blocks
                                   if _values_equal((item.get("value") or {}).get("heading"), heading)]
                exact_matches = [item for item in heading_matches
                                 if _values_equal((item.get("value") or {}).get("body_parts"), body_parts)]
                if exact_matches:
                    state = "agree"
                    reason = "Exact written declaration heading/body paragraphs match the official template block."
                    action = "corroborated"
                    observations = exact_matches
                elif len(heading_matches) == 1 and body_parts:
                    state = "conflict"
                    reason = "The declaration heading matches, but the exact written and official-template body paragraphs differ."
                    action = "clarification_required"
                    observations = heading_matches
                else:
                    state = "insufficient"
                    reason = "No uniquely corresponding official declaration body was directly observable; absence is not a prohibition."
                    action = "manual_review_retained"
                    observations = heading_matches
                result = {
                    "id": f"TR{len(results) + 1:05d}", "result": state, "action": action,
                    "requirement_id": requirement.get("id"), "clause_ids": clause_ids,
                    "semantic_role": role, "property": f"items.{written_item.get('id', 'unknown')}.body_parts",
                    "written": {"value": {"heading": heading, "body_parts": body_parts},
                                "evidence_ids": list(requirement.get("evidence_ids", []))},
                    "official_template": {"values": [item.get("value") for item in observations],
                                          "evidence_ids": [item["id"] for item in observations]},
                    "reason": reason,
                }
                results.append(result)
                if state == "conflict":
                    conflict = {
                        "type": "official_template_conflict", "clause_ids": clause_ids,
                        "requirement_id": requirement.get("id"), "role": role,
                        "property": result["property"], "written_value": result["written"]["value"],
                        "official_template_values": result["official_template"]["values"],
                        "official_template_evidence_ids": result["official_template"]["evidence_ids"],
                        "reason": reason,
                    }
                    genuine_conflicts.append(conflict)
                    remaining_questions.append({
                        "id": f"TQ{len(genuine_conflicts):04d}",
                        "clause_id": item_question_clause_id,
                        "requires_clarification": True,
                        "question": "书面声明正文与官方模板的精确正文冲突，请确认采用哪一版。",
                        "source_text": requirement.get("source_text", ""),
                        "candidate_roles": [role], "property": result["property"],
                        "written_value": result["written"]["value"],
                        "official_template_values": result["official_template"]["values"],
                        "evidence_ids": result["official_template"]["evidence_ids"],
                    })
            if "before_role" in properties:
                before_observations = [item for item in blocks
                                       if (item.get("value") or {}).get("before_role")]
                before_values = [(item.get("value") or {}).get("before_role")
                                 for item in before_observations]
                if before_values and all(_values_equal(before_values[0], value)
                                         for value in before_values[1:]):
                    before_state = ("agree" if _values_equal(properties.get("before_role"), before_values[0])
                                    else "conflict")
                else:
                    before_state = "insufficient"
                before_reason = (
                    "The written insertion anchor matches the directly observed declaration placement."
                    if before_state == "agree" else
                    "The written insertion anchor conflicts with the directly observed declaration placement."
                    if before_state == "conflict" else
                    "One unambiguous declaration insertion anchor was not directly observable."
                )
                before_result = {
                    "id": f"TR{len(results) + 1:05d}", "result": before_state,
                    "action": ("corroborated" if before_state == "agree" else
                               "clarification_required" if before_state == "conflict" else
                               "manual_review_retained"),
                    "requirement_id": requirement.get("id"),
                    "clause_ids": clause_ids, "semantic_role": role, "property": "before_role",
                    "written": {"value": properties.get("before_role"),
                                "evidence_ids": list(requirement.get("evidence_ids", []))},
                    "official_template": {"values": before_values,
                                          "evidence_ids": [item["id"] for item in before_observations]},
                    "reason": before_reason,
                }
                results.append(before_result)
                if before_state == "conflict":
                    conflict = {
                        "type": "official_template_conflict", "clause_ids": clause_ids,
                        "requirement_id": requirement.get("id"), "role": role,
                        "property": "before_role", "written_value": properties.get("before_role"),
                        "official_template_values": before_values,
                        "official_template_evidence_ids": before_result["official_template"]["evidence_ids"],
                        "reason": before_reason,
                    }
                    genuine_conflicts.append(conflict)
                    remaining_questions.append({
                        "id": f"TQ{len(genuine_conflicts):04d}",
                        "clause_id": clause_ids[0] if clause_ids else None,
                        "requires_clarification": True,
                        "question": "书面声明插入位置与官方模板冲突，请确认采用哪一项。",
                        "source_text": requirement.get("source_text", ""),
                        "candidate_roles": [role], "property": "before_role",
                        "written_value": properties.get("before_role"),
                        "official_template_values": before_values,
                        "evidence_ids": before_result["official_template"]["evidence_ids"],
                    })
            continue
        for path, written_value in _flatten(properties).items():
            state, candidates = _comparison(role, path, written_value, index)
            result = {
                "id": f"TR{len(results) + 1:05d}", "result": state,
                "action": "corroborated" if state == "agree" else "clarification_required" if state == "conflict" else "manual_review_retained",
                "requirement_id": requirement.get("id"), "clause_ids": clause_ids,
                "semantic_role": role, "property": path,
                "written": {"value": written_value, "evidence_ids": list(requirement.get("evidence_ids", []))},
                "official_template": {"values": [item.get("leaf_value") for item in candidates],
                                      "evidence_ids": [item["id"] for item in candidates]},
                "reason": ("Direct official-template evidence matches the written value." if state == "agree" else
                           "Direct official-template evidence contradicts the written value." if state == "conflict" else
                           "The official sample does not provide one unambiguous directly observable value; absence is not prohibition."),
            }
            results.append(result)
            if state == "conflict":
                conflict = {
                    "type": "official_template_conflict", "clause_ids": clause_ids,
                    "requirement_id": requirement.get("id"), "role": role, "property": path,
                    "written_value": written_value,
                    "official_template_values": [item.get("leaf_value") for item in candidates],
                    "official_template_evidence_ids": [item["id"] for item in candidates],
                    "reason": result["reason"],
                }
                genuine_conflicts.append(conflict)
                remaining_questions.append({
                    "id": f"TQ{len(genuine_conflicts):04d}",
                    "clause_id": clause_ids[0] if clause_ids else None,
                    "requires_clarification": True,
                    "question": "书面要求与官方 Word 模板的直接证据冲突，请确认应采用哪一项。",
                    "source_text": requirement.get("source_text", ""),
                    "candidate_roles": [role], "property": path,
                    "written_value": written_value,
                    "official_template_values": conflict["official_template_values"],
                    "evidence_ids": conflict["official_template_evidence_ids"],
                })

    # Manual-only semantic reviews have no executable requirement by design.
    # Corroborate exact observable labels when possible, while retaining the
    # manual/unverifiable clause state instead of fabricating properties.
    question_clause_ids = {str(item.get("clause_id")) for item in remaining_questions}
    for record in working.get("clause_compliance", []):
        if not isinstance(record, dict) or record.get("status") != "unverifiable":
            continue
        clause_id = str(record.get("clause_id") or "")
        if not clause_id or clause_id in question_clause_ids:
            continue
        clause = clause_map.get(clause_id)
        if clause is None:
            continue
        matched = _manual_corroboration(str(clause.get("text") or ""), evidence)
        results.append({
            "id": f"TR{len(results) + 1:05d}",
            "result": "agree" if matched else "insufficient",
            "action": "corroborated_manual" if matched else "manual_review_retained",
            "clause_id": clause_id, "semantic_role": "manual", "property": None,
            "written": {"value": clause.get("text"),
                        "evidence_ids": list(clause.get("evidence_ids", []))},
            "official_template": {"values": [item.get("value") for item in matched],
                                  "evidence_ids": [item["id"] for item in matched]},
            "reason": (
                "Direct fixed-text evidence corroborates the labels, while the clause correctly remains a non-executable manual check."
                if matched else
                "No complete direct corroboration was observable; the manual state is retained and sample absence is not a prohibition."
            ),
        })

    represented_not_applicable = {
        clause_id for item in results if item.get("result") == "not_applicable"
        for clause_id in ([str(item.get("clause_id"))] if item.get("clause_id") else
                          [str(value) for value in item.get("clause_ids", [])])
    }
    for record in working.get("clause_compliance", []):
        if not isinstance(record, dict) or record.get("status") != "not_applicable":
            continue
        clause_id = str(record.get("clause_id") or "")
        if not clause_id or clause_id in represented_not_applicable:
            continue
        results.append({
            "id": f"TR{len(results) + 1:05d}", "result": "not_applicable", "action": "none",
            "clause_id": clause_id, "semantic_role": None, "property": None,
            "written": {"evidence_ids": list(record.get("evidence_ids", []))},
            "official_template": {"evidence_ids": []},
            "reason": str(record.get("reason") or "The written applicability contract excludes this clause."),
        })

    # Resolve outstanding role questions only when one candidate role has
    # direct evidence agreeing on every explicit written property.
    for question in list(remaining_questions):
        # A genuine direct conflict must remain a user-visible clarification;
        # do not let the generic role resolver silently turn its clause into
        # an unverifiable manual check merely because one candidate is
        # observable in the official template.
        if question.get("requires_clarification"):
            continue
        clause_id = str(question.get("clause_id") or "")
        clause = clause_map.get(clause_id)
        if clause is None:
            continue
        text = str(question.get("source_text") or clause.get("text") or "")
        page_props = page_property_parser(text) if page_property_parser else {}
        props = page_props or (property_parser(text) if property_parser else {})
        candidates = [str(value) for value in question.get("candidate_roles", [])
                      if value not in {"unknown", "llm_primary_review"}]
        if page_props:
            candidates = ["page"]
        elif not candidates and role_candidates:
            candidates = role_candidates(clause)
        if not props:
            matched = _manual_corroboration(text, evidence)
            if matched:
                _remove_unresolved(working, clause_id, covered=False)
                _upsert_clause_record(
                    working, clause, "unverifiable",
                    "The official template corroborates the referenced labels/text, but no executable property was stated; manual verification is retained.", [],
                )
                results.append({
                    "id": f"TR{len(results) + 1:05d}", "result": "agree",
                    "action": "corroborated_manual", "clause_id": clause_id,
                    "semantic_role": candidates[0] if len(candidates) == 1 else "manual",
                    "property": None, "written": {"value": text, "evidence_ids": clause.get("evidence_ids", [])},
                    "official_template": {"values": [item.get("value") for item in matched],
                                          "evidence_ids": [item["id"] for item in matched]},
                    "reason": "Direct fixed-text evidence corroborates the written labels, but no universal executable formatting rule is inferred.",
                })
                remaining_questions.remove(question)
            else:
                results.append({
                    "id": f"TR{len(results) + 1:05d}", "result": "insufficient",
                    "action": "manual_review_retained", "clause_id": clause_id,
                    "semantic_role": candidates[0] if len(candidates) == 1 else "unknown",
                    "property": None,
                    "reason": "No explicit executable property and no complete direct fixed-text corroboration were observable.",
                })
            continue
        matching_roles: list[tuple[str, list[dict[str, Any]]]] = []
        fully_conflicting: list[tuple[str, list[dict[str, Any]]]] = []
        for role in candidates:
            observations: list[dict[str, Any]] = []
            states: list[str] = []
            for path, value in _flatten(props).items():
                state, evidence_items = _comparison(role, path, value, index)
                states.append(state); observations.extend(evidence_items)
            if states and all(state == "agree" for state in states):
                matching_roles.append((role, observations))
            elif states and all(state == "conflict" for state in states):
                fully_conflicting.append((role, observations))
        if len(matching_roles) == 1:
            role, observations = matching_roles[0]
            target = working.setdefault("page", {}) if role == "page" else working.setdefault("roles", {}).setdefault(role, {})
            for path, value in _flatten(props).items():
                _set_nested(target, path, value)
            requirement_id = _next_requirement_id(working)
            template_ids = sorted({item["id"] for item in observations})
            working.setdefault("requirements", []).append({
                "id": requirement_id, "role": role, "properties": copy.deepcopy(props),
                "evidence_ids": sorted(set(clause.get("evidence_ids", [])) | set(template_ids)),
                "clause_ids": [clause_id], "resolved_by": "template", "confidence": 1.0,
                "source_text": text,
                "reason": "The written property and one directly observed official-template semantic role agree.",
            })
            _remove_unresolved(working, clause_id, covered=True)
            _upsert_clause_record(working, clause, "pending_execution",
                                  "Written and official-template evidence agree.", [requirement_id])
            results.append({
                "id": f"TR{len(results) + 1:05d}", "result": "agree", "action": "resolved_by_template",
                "clause_id": clause_id, "requirement_id": requirement_id,
                "semantic_role": role, "property": sorted(_flatten(props)),
                "written": {"value": props, "evidence_ids": clause.get("evidence_ids", [])},
                "official_template": {"evidence_ids": template_ids},
                "reason": "Exactly one candidate semantic role matched every explicit written property.",
            })
            remaining_questions.remove(question)
        elif fully_conflicting and len(fully_conflicting) == len(candidates) == 1:
            role, observations = fully_conflicting[0]
            conflict = {
                "type": "official_template_conflict", "clause_ids": [clause_id],
                "role": role, "property": sorted(_flatten(props)), "written_value": props,
                "official_template_evidence_ids": sorted({item["id"] for item in observations}),
                "reason": "The only candidate role directly contradicts every explicit written property.",
            }
            genuine_conflicts.append(conflict)
            results.append({
                "id": f"TR{len(results) + 1:05d}", "result": "conflict",
                "action": "clarification_required", "clause_id": clause_id,
                "semantic_role": role, "property": sorted(_flatten(props)),
                "reason": conflict["reason"],
            })
        else:
            results.append({
                "id": f"TR{len(results) + 1:05d}", "result": "insufficient",
                "action": "manual_review_retained", "clause_id": clause_id,
                "semantic_role": candidates, "property": sorted(_flatten(props)),
                "reason": "Template evidence did not uniquely identify one agreeing semantic role; absence or ambiguity is not a prohibition.",
            })

    remaining_questions, declaration_template_ids = _declaration_resolution(
        working, clauses, remaining_questions, evidence, results
    )

    # Explicit silent-value policy: only scalar page/style values with direct
    # semantic observations may be filled.  Structure, resources, conditions,
    # and feature absence are never inferred from a sample.
    if fill_silent_values:
        for (role, path), observations in sorted(index.items()):
            if role == "document_structure" or path in {"observed_order", "ordered_roles"}:
                continue
            values = [item.get("leaf_value") for item in observations]
            if not values or not all(_values_equal(values[0], value) for value in values[1:]):
                continue
            target = working.setdefault("page", {}) if role == "page" else working.setdefault("roles", {}).setdefault(role, {})
            present, _ = _get_nested(target, path)
            if present:
                continue
            _set_nested(target, path, values[0])
            results.append({
                "id": f"TR{len(results) + 1:05d}", "result": "agree",
                "action": "filled_silent_value", "semantic_role": role, "property": path,
                "written": {"status": "silent"},
                "official_template": {"value": values[0],
                                      "evidence_ids": [item["id"] for item in observations]},
                "reason": "The explicit policy permits a silent scalar formatting value to be filled from unambiguous direct official-template evidence.",
            })

    if genuine_conflicts:
        working["status"] = "needs_clarification"
        working.setdefault("blocking_errors", []).extend(copy.deepcopy(genuine_conflicts))
    elif (not remaining_questions
          and not (working.get("completeness") or {}).get("unresolved_clause_ids")
          and not working.get("blocking_errors")):
        if working.get("status") == "needs_clarification":
            working["status"] = "semantic_resolved" if working.get("analysis_mode") == "llm_primary" else "rule_resolved"

    counts = {state: sum(1 for item in results if item.get("result") == state)
              for state in sorted(RESULT_STATES)}
    report = {
        "schema_version": SCHEMA_VERSION,
        "status": "conflict" if genuine_conflicts else "reconciled",
        "policy": {
            "written_explicit_requirements_authoritative": True,
            "official_template_may_corroborate": True,
            "fill_silent_scalar_formatting_values": bool(fill_silent_values),
            "fill_silent_structure_or_feature_absence": False,
            "sample_absence_is_prohibition": False,
            "unresolved_without_direct_agreement": "insufficient_manual",
        },
        "sources": copy.deepcopy(sources),
        "summary": {**counts, "results": len(results),
                    "genuine_conflicts": len(genuine_conflicts),
                    "resolved_questions": len(questions) - len(remaining_questions),
                    "remaining_questions": len(remaining_questions),
                    "declaration_template_evidence": len(declaration_template_ids)},
        "results": results,
        "genuine_conflicts": genuine_conflicts,
    }
    return working, remaining_questions, genuine_conflicts, report


def not_supplied_report(sources: dict[str, Any]) -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION, "status": "not_applicable",
        "policy": {"sample_absence_is_prohibition": False},
        "sources": copy.deepcopy(sources),
        "summary": {"agree": 0, "conflict": 0, "insufficient": 0,
                    "not_applicable": 0, "results": 0, "genuine_conflicts": 0,
                    "resolved_questions": 0, "remaining_questions": None,
                    "declaration_template_evidence": 0},
        "results": [], "genuine_conflicts": [],
    }


def reconciliation_markdown(report: dict[str, Any]) -> str:
    lines = [
        "# Official Template Reconciliation", "",
        f"Status: `{report.get('status')}`", "",
        "Written explicit requirements remain authoritative. Missing sample content is not treated as a prohibition.", "",
        "| Result | Action | Clause | Role | Property | Reason |",
        "|---|---|---|---|---|---|",
    ]
    for item in report.get("results", []):
        reason = str(item.get("reason") or "").replace("|", "\\|").replace("\n", " ")
        lines.append("| {result} | {action} | {clause} | {role} | {prop} | {reason} |".format(
            result=item.get("result", ""), action=item.get("action", ""),
            clause=item.get("clause_id") or ", ".join(item.get("clause_ids", [])),
            role=item.get("semantic_role", ""), prop=item.get("property", ""), reason=reason,
        ))
    if not report.get("results"):
        lines.append("| not_applicable | none |  |  |  | No official template evidence was supplied. |")
    return "\n".join(lines) + "\n"


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("official_template", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    evidence = extract_template_evidence(args.official_template.resolve())
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": evidence["status"], "items": evidence["summary"]["item_count"],
                      "output": str(args.out)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(__import__("sys").argv[1:]))
