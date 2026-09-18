"""School-neutral, high-coverage DOCX fixture for format-spec application tests."""
from __future__ import annotations

import base64
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

ROLE_STYLE_NAMES = {
    "thesis_title_zh": "ThesisTitle",
    "thesis_title_en": "EnglishTitle",
    "abstract_title_zh": "AbstractTitleCN",
    "abstract_body_zh": "AbstractBodyCN",
    "abstract_title_en": "AbstractTitleEN",
    "abstract_body_en": "AbstractBodyEN",
    "keywords_zh": "KeywordsLineCN",
    "keywords_en": "KeywordsLineEN",
    "heading_1": "Heading 1",
    "heading_2": "Heading 2",
    "heading_3": "Heading 3",
    "body_text": "Body Text",
    "figure_caption": "ImageCaption",
    "table_caption": "TableCaption",
    "table_text": "TableText",
    "equation": "EquationBlock",
    "bibliography_heading": "ReferencesHeading",
    "bibliography_entry": "Bibliography",
    "footnote": "Footnote Text",
    "header": "Header",
    "footer": "Footer",
    "toc": "TOC Heading",
}


def _ensure_styles(doc: Document) -> None:
    existing = {style.name for style in doc.styles}
    for name in ROLE_STYLE_NAMES.values():
        if name not in existing:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            existing.add(name)


def _add_page_field(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE \\* ARABIC")
    paragraph._p.append(field)


def build_neutral_thesis(path: Path) -> None:
    """Build a generic thesis with every paragraph role supported by the backend."""
    image = path.with_suffix(".png")
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="))

    doc = Document()
    _ensure_styles(doc)

    # Cover/title page: deliberately generic and institution-free.
    doc.add_paragraph("面向可验证文档处理的中性论文", ROLE_STYLE_NAMES["thesis_title_zh"])
    doc.add_paragraph("A Neutral Thesis on Verifiable Document Processing", ROLE_STYLE_NAMES["thesis_title_en"])
    doc.add_paragraph("Author: Example Researcher", ROLE_STYLE_NAMES["body_text"])

    # Front matter occupies its own section.
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("摘 要", ROLE_STYLE_NAMES["abstract_title_zh"])
    doc.add_paragraph("本文提供与学校无关的高覆盖测试内容。", ROLE_STYLE_NAMES["abstract_body_zh"])
    doc.add_paragraph("关键词：文档处理；验证；测试", ROLE_STYLE_NAMES["keywords_zh"])
    doc.add_paragraph("ABSTRACT", ROLE_STYLE_NAMES["abstract_title_en"])
    doc.add_paragraph("This fixture exercises all executable semantic roles.", ROLE_STYLE_NAMES["abstract_body_en"])
    doc.add_paragraph("Keywords: documents; validation; testing", ROLE_STYLE_NAMES["keywords_en"])
    doc.add_paragraph("目 录", ROLE_STYLE_NAMES["toc"])
    doc.add_paragraph("1 Introduction ................................ 1", ROLE_STYLE_NAMES["body_text"])

    # Main matter has multiple sections and heading depths.
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Introduction", ROLE_STYLE_NAMES["heading_1"])
    doc.add_paragraph("Motivation", ROLE_STYLE_NAMES["heading_2"])
    doc.add_paragraph("Validation Boundary", ROLE_STYLE_NAMES["heading_3"])
    doc.add_paragraph("This is the main body paragraph with neutral prose.", ROLE_STYLE_NAMES["body_text"])
    picture = doc.add_paragraph()
    picture.add_run().add_picture(str(image), width=Cm(1))
    doc.add_paragraph("Figure 1 Neutral processing pipeline", ROLE_STYLE_NAMES["figure_caption"])
    doc.add_paragraph("Table 1 Coverage matrix", ROLE_STYLE_NAMES["table_caption"])
    table = doc.add_table(rows=2, cols=2)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            paragraph = cell.paragraphs[0]
            paragraph.style = ROLE_STYLE_NAMES["table_text"]
            paragraph.add_run(f"cell {row_index + 1}-{col_index + 1}")
    doc.add_paragraph("E = mc²", ROLE_STYLE_NAMES["equation"])
    doc.add_paragraph("1 Fixture note used as a footnote-format test point.", ROLE_STYLE_NAMES["footnote"])

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("References", ROLE_STYLE_NAMES["bibliography_heading"])
    doc.add_paragraph("[1] Example Author. Neutral Document Testing. 2026.", ROLE_STYLE_NAMES["bibliography_entry"])
    doc.add_paragraph("Appendix A Reproducibility Material", ROLE_STYLE_NAMES["heading_1"])
    doc.add_paragraph("Appendix body content.", ROLE_STYLE_NAMES["body_text"])

    # Exercise header/footer stories independently of body paragraphs.
    for index, section in enumerate(doc.sections, start=1):
        header = section.header.paragraphs[0]
        header.style = ROLE_STYLE_NAMES["header"]
        header.text = f"Neutral Thesis Header {index}"
        footer = section.footer.paragraphs[0]
        footer.style = ROLE_STYLE_NAMES["footer"]
        footer.text = "Page "
        _add_page_field(footer)

    doc.save(path)
    image.unlink(missing_ok=True)


def make_full_role_spec() -> dict:
    """Return a valid spec declaring every role exercised by the fixture."""
    roles = {}
    for role in ROLE_STYLE_NAMES:
        roles[role] = {
            "font": {"latin": "Arial", "cjk": "SimSun", "size_pt": 10},
            "paragraph": {"alignment": "left"},
        }
    roles["figure_caption"]["position"] = "below"
    roles["table_caption"]["position"] = "above"
    return {
        "schema_version": "1.0",
        "source_document": "neutral-thesis-fixture",
        "analysis_mode": "llm_primary",
        "status": "semantic_resolved",
        "roles": roles,
        "requirements": [],
        "completeness": {
            "reviewed_by": "user",
            "covered_clause_ids": [],
            "ignored_clause_ids": [],
            "unresolved_clause_ids": [],
            "unsupported_items": [],
        },
    }


def make_neutral_style_map() -> dict:
    """Pin semantic roles to the fixture's generic styles without school templates."""
    return {"mappings": {role: {"style_name": style} for role, style in ROLE_STYLE_NAMES.items()}}
