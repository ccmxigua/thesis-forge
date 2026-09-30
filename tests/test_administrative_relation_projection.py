from __future__ import annotations

import copy
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import host_agent_bridge as bridge
from administrative_relation_projection import project_administrative_copies
from host_review_contract import _security_marking_qualifier_binding_errors
from requirements_engine import extract_document_evidence, split_clauses, build_llm_request
from semantic_contract import attach_request_provenance, sha256_file, sha256_json


def fixture(variant: int = 0, *, intervening_external: bool = False) -> tuple[dict, dict, dict]:
    """Create and extract two different DOCX sources, not relabeled BSU JSON."""
    labels = [("内部", 18, "月"), ("敏感", 5, "年")] if variant else [("限阅", 6, "月"), ("保密", 3, "年")]
    region = "非公开论文管理栏" if variant else "论文保密信息"
    policy = ("未获批准的一律视为公开学位论文（公开学位论文此项留空）" if variant
              else "未经批准的均为公开学位论文（公开的学位论文本项为空白）")
    consent = "非公开论文须由作者申请并经导师同意和管理部门批准"
    seal = "管理办公室盖章(有效)"
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "requirements.docx"
        doc = Document()
        for _ in range(variant + 1):
            doc.add_paragraph("排版说明")
        doc.add_paragraph(region)
        if intervening_external:
            doc.add_paragraph("另一管理区域的申请须获批准")
        doc.add_paragraph(consent + "。" + policy)
        table = doc.add_table(rows=3, cols=2)
        rows = [
            ("论文名称", "填写论文名称"),
            ("申请级别", " ".join(f"□{label}(≤{value}{unit})" for label, value, unit in labels)),
            ("申请编号", "填写编号"),
        ]
        for row, pair in zip(table.rows, rows):
            for cell, text in zip(row.cells, pair):
                cell.text = text
        doc.add_paragraph(seal)
        doc.add_paragraph(" ".join(f"{label}★{value}{unit}(可少于{value}{unit})" for label, value, unit in labels))
        doc.save(path)
        evidence = extract_document_evidence(path)
        clauses = split_clauses(evidence)
        request = build_llm_request([], clauses, evidence, {}, "full", contract_version="3.0")
        request = attach_request_provenance(request, source_sha256=sha256_file(path),
                                           evidence_doc=evidence, clauses=clauses, run_id=f"different-docx-{variant}")
    ids = {"region": next(c["id"] for c in clauses if c["text"] == region),
           "policy": next(c["id"] for c in clauses if c["text"] == policy),
           "consent": next(c["id"] for c in clauses if c["text"] == consent),
           "seal": next(c["id"] for c in clauses if c["text"] == seal)}
    if intervening_external:
        ids["other_external"] = next(c["id"] for c in clauses if c["text"] == "另一管理区域的申请须获批准")
    field_defs = [("title_zh", "论文名称", "title_zh"), ("security_marking", "申请级别", "security_marking"),
                  ("approval_number", "申请编号", "approval_number")]
    fields = [{"id": ident, "label": label, "value_from": f"thesis_profile.cover_metadata.{key}",
               "display_policy": "required", "order": i + 1} for i, (ident, label, key) in enumerate(field_defs)]
    properties = {"institution": "示例单位", "fields": [fields[0]],
                  "non_public_administration": {
                      "source_region": region, "fields": [{**f, "order": i + 1} for i, f in enumerate(fields[1:])],
                      "public_policy": "blank", "publication_default_policy": "unapproved_is_public",
                      "applicability": {"status": "conditional", "conditions": [
                          {"fact": "thesis_profile.security_level", "operator": "in", "value": ["restricted", "classified"]}]},
                      "security_marking_options": [{"label": label, "maximum_duration": {"value": value, "unit": unit},
                                                    "shorter_duration_allowed": True} for label, value, unit in labels]}}
    local_ids = [c["id"] for c in clauses if c["id"] not in ids.values()
                 and c["text"] != "排版说明"] + [ids["region"]]
    by_id = {c["id"]: c for c in clauses}

    def requirement(cids: list[str], props: dict, mode: str) -> dict:
        return {"role": "cover", "properties": copy.deepcopy(props), "clause_ids": cids,
                "evidence_ids": sorted({eid for cid in cids for eid in by_id[cid]["evidence_ids"]}),
                "confidence": 0.9, "reason": "Current source-backed administrative rule",
                "verification": {"mode": mode, "checks": ["Inspect the bound source duty."]}}

    reviews = []
    for clause in clauses:
        cid = clause["id"]
        external = cid in (ids["consent"], ids["seal"], ids.get("other_external"))
        info = clause["text"] == "排版说明"
        reviews.append({"clause_id": cid, "classification": "external_compliance" if external else "informational" if info else "executable",
                        "reason": "Source duty remains pending" if external else "Source structure",
                        "normative_basis": "external_duty" if external else "sample_content" if info else "template_structure",
                        "obligations": [] if info else [{"id": f"source-duty-{cid}", "status": "unverifiable" if external else "covered",
                                                       "reason": "Bound to the current source"}]})
    subset = copy.deepcopy(properties)
    subset["fields"] = []
    subset["non_public_administration"]["fields"] = subset["non_public_administration"]["fields"][:1]
    subset["non_public_administration"].pop("security_marking_options")
    response = {"contract_version": "3.0", "unsupported_items": [], "reported_conflicts": [],
                "requirements": [requirement(local_ids, properties, "static_docx"),
                  requirement([ids["consent"]], subset, "external"), requirement([ids["seal"]], subset, "external")],
                "clause_reviews": reviews}
    return response, request, ids


class AdministrativeRelationProjectionTests(unittest.TestCase):
    def test_two_extracted_docx_preserve_fields_and_dynamic_terms(self) -> None:
        hashes = []
        for variant in (0, 1):
            with self.subTest(variant=variant):
                raw, chunk, ids = fixture(variant)
                before = copy.deepcopy(raw)
                candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
                self.assertEqual(raw, before)
                self.assertEqual(len(candidate["requirements"]), 1)
                self.assertEqual(candidate["requirements"][0]["properties"], raw["requirements"][0]["properties"])
                self.assertIn(ids["policy"], candidate["requirements"][0]["clause_ids"])
                self.assertEqual(candidate["clause_reviews"], raw["clause_reviews"])
                self.assertTrue(audit["mechanical_repairs"][0]["external_actions_remain_pending"])
                hashes.append(chunk["provenance"]["source_sha256"])
                again, _ = bridge.prepare_native_response_candidate(candidate, chunk)
                self.assertEqual(candidate, again)
        self.assertNotEqual(*hashes)

    def test_duplicate_policy_copy_has_no_unique_operation(self) -> None:
        raw, chunk, ids = fixture()
        duplicate = copy.deepcopy(raw["requirements"][0])
        duplicate["clause_ids"] = [ids["policy"]]
        duplicate["evidence_ids"] = next(c["evidence_ids"] for c in chunk["clauses"] if c["id"] == ids["policy"])
        raw["requirements"].append(duplicate)
        candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(len(candidate["requirements"]), 1)
        self.assertEqual(audit["mechanical_repairs"][0]["removed_indexes"], [1, 2, 3])

    def test_unique_properties_or_meaning_never_disappear(self) -> None:
        for change in ("different_value", "unique_field", "condition", "existing_id", "prerequisite", "unknown_top_key", "unknown_nested_null", "reordered_fields"):
            with self.subTest(change=change):
                raw, chunk, _ = fixture()
                req = raw["requirements"][1]
                if change == "different_value": req["properties"]["institution"] = "另一单位"
                elif change == "unique_field": req["properties"]["font"] = {"size_pt": 16}
                elif change == "condition": req["applicability"] = {"status": "conditional", "conditions": []}
                elif change == "existing_id": req["existing_requirement_id"] = "existing-other"
                elif change == "prerequisite": req["input_prerequisites"] = [{"required_fields": ["other"]}]
                elif change == "unknown_top_key": req["unrecognized"] = True
                elif change == "unknown_nested_null": req["properties"]["non_public_administration"]["unrecognized"] = None
                else: req["properties"]["non_public_administration"]["fields"] = list(reversed(raw["requirements"][0]["properties"]["non_public_administration"]["fields"]))
                before = copy.deepcopy(raw)
                candidate, _ = project_administrative_copies(raw, chunk, validate=bridge.validate_host_agent_response)
                self.assertIsNone(candidate)
                self.assertEqual(raw, before)

    def test_stale_ambiguous_or_foreign_source_fails_closed(self) -> None:
        for change in ("hash", "offset", "evidence", "location", "other_region", "empty_inventory", "mixed", "conflict"):
            with self.subTest(change=change):
                raw, chunk, ids = fixture()
                clause = next(c for c in chunk["clauses"] if c["id"] == ids["policy"])
                if change == "hash": clause["source_span"]["source_sha256"] = "a" * 64
                elif change == "offset": clause["source_span"]["start_offset"] += 1
                elif change == "evidence": raw["requirements"][1]["evidence_ids"] = ["foreign-evidence"]
                elif change == "location": clause["source_span"]["location"]["child_index"] += 10
                elif change == "other_region": raw["requirements"].append(copy.deepcopy(raw["requirements"][0]))
                elif change == "empty_inventory": next(r for r in raw["clause_reviews"] if r["clause_id"] == ids["consent"])["obligations"] = []
                elif change == "mixed": next(r for r in raw["clause_reviews"] if r["clause_id"] == ids["consent"])["classification"] = "executable_with_external_check"
                else: raw["conflicts"] = [{"reason": "Contradicting rule"}]
                projected, _ = project_administrative_copies(raw, chunk, validate=bridge.validate_host_agent_response)
                self.assertIsNone(projected)

    def test_feedback_must_bind_complete_current_response(self) -> None:
        raw, chunk, _ = fixture()
        records = bridge.contract_error_records(bridge.validate_host_agent_response(raw, chunk), response=raw, chunk=chunk)
        for invalid in (records[:-1], [{**r, "response_sha256": "a" * 64} for r in records]):
            projected, _ = bridge._apply_safe_mechanical_repairs(raw, invalid, chunk=chunk)
            self.assertIsNone(projected)

    def test_an_external_heading_between_regions_is_not_proximity_authority(self) -> None:
        raw, chunk, _ = fixture(intervening_external=True)
        projected, _ = project_administrative_copies(raw, chunk, validate=bridge.validate_host_agent_response)
        self.assertIsNone(projected)

    def test_shorter_duration_is_bound_per_option_not_per_array(self) -> None:
        options = [{"label": "内部", "maximum_duration": {"value": 18, "unit": "月"}, "shorter_duration_allowed": True},
                   {"label": "敏感", "maximum_duration": {"value": 5, "unit": "年"}, "shorter_duration_allowed": True}]
        req = {"role": "cover", "clause_ids": ["dynamic-source"], "properties": {"non_public_administration": {"security_marking_options": options}}}
        clauses = [{"id": "dynamic-source", "text": "内部★18月(可少于18月)"}]
        errors = _security_marking_qualifier_binding_errors({"requirements": [req]}, clauses)
        self.assertEqual(len(errors), 1)
        self.assertIn("security_marking_options[1]", errors[0])


if __name__ == "__main__":
    unittest.main()
