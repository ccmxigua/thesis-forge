from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from assembly_executor import _section_invariants, execute_assembly_plan  # noqa: E402
from region_graph import compile_region_graph  # noqa: E402
from source_role_extractor import extract_source_roles, target_headings_from_profile  # noqa: E402
from template_profile import load_profile, resource  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
COMMENT_MARKERS = {"commentRangeStart", "commentRangeEnd", "commentReference"}


def _source_docx(path: Path) -> None:
    document = Document()
    for heading, content in (
        ("第1章 引言", "generated body"),
        ("致谢", "generated acknowledgments"),
        ("参考文献", "generated references"),
        ("附录A 补充材料", "generated appendix"),
        ("攻读硕士学位期间的学术成果", "generated academic outputs"),
    ):
        document.add_heading(heading, level=1)
        document.add_paragraph(content)
    document.save(path)


def _section_index_by_text(root: etree._Element, wanted: str) -> int:
    body = root.find("w:body", NS)
    assert body is not None
    section_index = 0
    for child in body:
        text = "".join(child.xpath(".//w:t/text() | .//w:instrText/text()", namespaces=NS))
        if (text.startswith(wanted) if wanted == "附录A" else text == wanted):
            return section_index
        if child.xpath("./w:pPr/w:sectPr | self::w:sectPr", namespaces=NS):
            section_index += 1
    raise AssertionError(f"assembled document has no anchor containing {wanted!r}")


class AssemblySanitizationTests(unittest.TestCase):
    def _assemble_official(self, base: Path) -> tuple[Path, Path, dict, dict]:
        profile_path = ROOT / "template_profiles/neau-academic-master-2026/profile.json"
        profile = load_profile(profile_path)
        template = resource(profile, "official_docx")
        assert template is not None
        source = base / "source.docx"
        output = base / "assembled.docx"
        _source_docx(source)
        source_roles = extract_source_roles(
            source,
            required_roles={"body", "references", "acknowledgments", "academic_outputs"},
            target_headings=target_headings_from_profile(profile_path),
        )
        self.assertEqual(source_roles["status"], "extracted", source_roles["findings"])
        plan_result = compile_region_graph(profile, template_path=template)
        self.assertEqual(plan_result["status"], "compiled", plan_result["findings"])
        result = execute_assembly_plan(
            plan_result,
            template,
            source,
            output,
            metadata={
                "cover_metadata": {
                    "title_zh": "测试中文题目",
                    "title_en": "Test English Thesis Title",
                    "author_name": "测试作者",
                    "student_id": "TEST0001",
                    "supervisor_name": "测试导师 教授",
                    "college_name": "测试学院",
                    "degree_discipline": "测试学科",
                    "program_name": "测试专业",
                    "completion_date": "2026-06",
                }
            },
            source_section_policy="discard",
            source_role_map=source_roles,
        )
        return template, output, plan_result, result

    def test_official_17_sections_do_not_collapse_to_9_after_assembly(self):
        with tempfile.TemporaryDirectory() as raw:
            template, output, plan_result, result = self._assemble_official(Path(raw))
            with zipfile.ZipFile(output) as archive:
                self.assertIsNone(archive.testzip())
                root = etree.fromstring(archive.read("word/document.xml"))

            self.assertEqual(result["status"], "assembled")
            self.assertEqual(result["effective_section_count"], 17)
            self.assertEqual(result["effective_section_invariants_verified"], 1)
            self.assertGreater(result["role_anchors_verified"], 0)
            self.assertEqual(len(root.xpath("//w:sectPr", namespaces=NS)), 17)
            self.assertEqual(
                _section_invariants(root),
                next(
                    item["sections"]
                    for item in plan_result["assembly_plan"]["structural_postconditions"]
                    if item.get("type") == "effective_section_invariants"
                ),
            )

            # These are the original template section-index anchors.  The
            # assembled source content must not shift them into a new
            # 9-section document or make their selectors disappear.
            self.assertEqual(_section_index_by_text(root, "1  引言"), 8)
            self.assertEqual(_section_index_by_text(root, "致  谢"), 13)
            self.assertEqual(_section_index_by_text(root, "参考文献"), 14)
            self.assertEqual(_section_index_by_text(root, "附录A"), 15)
            self.assertEqual(_section_index_by_text(root, "攻读硕士学位期间的学术成果"), 16)

            # Ensure the generated role payload survived and was not replaced
            # by a fabricated placeholder or by the official template body.
            text = "".join(root.xpath("//w:t/text()", namespaces=NS))
            self.assertIn("generated body", text)
            self.assertIn("generated academic outputs", text)
            self.assertTrue(template.is_file())

    def test_review_comments_are_cleaned_from_output_only(self):
        profile_path = ROOT / "template_profiles/neau-academic-master-2026/profile.json"
        profile = load_profile(profile_path)
        template = resource(profile, "official_docx")
        assert template is not None
        template_bytes_before = template.read_bytes()
        with zipfile.ZipFile(template) as archive:
            original_members = {name: archive.read(name) for name in archive.namelist()}
        original_comment_parts = sorted(
            name for name in original_members if name.lower().startswith("word/comments")
        )
        original_marker_count = sum(
            payload.count(marker.encode("ascii"))
            for payload in original_members.values()
            for marker in COMMENT_MARKERS
        )
        self.assertEqual(original_comment_parts, ["word/comments.xml", "word/commentsExtended.xml"])
        self.assertGreater(original_marker_count, 0)

        with tempfile.TemporaryDirectory() as raw:
            _, output, _, result = self._assemble_official(Path(raw))
            self.assertEqual(template.read_bytes(), template_bytes_before)
            with zipfile.ZipFile(template) as archive:
                self.assertEqual(archive.read("word/comments.xml"), original_members["word/comments.xml"])
            with zipfile.ZipFile(output) as archive:
                self.assertIsNone(archive.testzip())
                output_members = {name: archive.read(name) for name in archive.namelist()}

            self.assertEqual(result["comments_cleanup"]["status"], "cleaned")
            self.assertEqual(result["comments_cleanup"]["comment_parts_removed"], original_comment_parts)
            self.assertGreater(result["comments_cleanup"]["comment_markers_removed"], 0)
            self.assertFalse(any(name.lower().startswith("word/comments") for name in output_members))

            for name, payload in output_members.items():
                if not (name.endswith(".xml") or name.endswith(".rels")):
                    continue
                document = etree.fromstring(payload)
                self.assertFalse(
                    any(etree.QName(node).localname in COMMENT_MARKERS for node in document.iter()),
                    name,
                )
                if name.endswith(".rels"):
                    self.assertFalse(
                        any("comment" in str(node.get("Type", "")).lower() for node in document),
                        name,
                    )

            output_root = etree.fromstring(output_members["word/document.xml"])
            output_text = "".join(output_root.xpath("//w:t/text()", namespaces=NS))
            self.assertIn("generated references", output_text)


if __name__ == "__main__":
    unittest.main()
