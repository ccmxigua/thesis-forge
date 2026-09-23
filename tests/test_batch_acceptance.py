from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import batch_rerun_ten_schools as batch  # noqa: E402
from manual_review_display import append_manual_review_markers  # noqa: E402


class BatchAcceptanceTests(unittest.TestCase):
    def test_native_semantic_review_requires_case_run_and_exact_check_coverage(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            case_root = Path(td) / "case"
            artifacts = case_root / "work" / "application" / "native-semantic-review"
            artifacts.mkdir(parents=True)
            checks = [{
                "check_id": "abstract_zh.require_third_person",
                "document_text": "本文提出一种模型。",
            }]
            document_text_sha256 = __import__("hashlib").sha256(json.dumps(
                [(item["check_id"], item["document_text"]) for item in checks],
                ensure_ascii=False, separators=(",", ":"),
            ).encode("utf-8")).hexdigest()
            request = {
                "schema_version": "1.0",
                "protocol": "native_semantic_content_review_v1",
                "case_id": "case-a", "run_id": "run-a",
                "source_sha256": "a" * 64,
                "format_spec_sha256": "b" * 64,
                "document_text_sha256": document_text_sha256,
                "checks": checks,
            }
            response = {"results": [{
                "check_id": "abstract_zh.require_third_person",
                "verdict": "satisfied",
                "rationale": "The passage uses third-person wording.",
                "evidence_quotes": ["本文提出"],
            }]}
            request_path = artifacts / "request.json"
            response_path = artifacts / "response.json"
            request_path.write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
            response_path.write_text(json.dumps(response, ensure_ascii=False), encoding="utf-8")
            review = {
                "schema_version": "1.0",
                "protocol": "native_semantic_content_review_v1",
                "status": "completed",
                "case_id": "case-a", "run_id": "run-a",
                "checks": checks, "results": response["results"],
                "request_path": str(request_path), "response_path": str(response_path),
                "request_sha256": batch.sha256_json(request),
                "response_sha256": __import__("hashlib").sha256(response_path.read_bytes()).hexdigest(),
                "source_sha256": request["source_sha256"],
                "format_spec_sha256": request["format_spec_sha256"],
                "document_text_sha256": document_text_sha256,
            }
            kwargs = {"case_id": "case-a", "run_id": "run-a", "case_root": case_root}
            self.assertTrue(batch.native_semantic_review_integrity(review, **kwargs))
            self.assertFalse(batch.native_semantic_review_integrity(
                review, **{**kwargs, "case_id": "case-b"},
            ))
            self.assertFalse(batch.native_semantic_review_integrity(
                review, **{**kwargs, "run_id": "run-b"},
            ))
            review["results"] = []
            self.assertFalse(batch.native_semantic_review_integrity(review, **kwargs))

    def _result(self, root: Path, *, returncode: int = 0, render: bool = True,
                compliance_mode: str = "full") -> dict:
        work = root / "case" / "work"
        requirements = work / "requirements"
        apply_dir = work / "application"
        requirements.mkdir(parents=True)
        apply_dir.mkdir(parents=True)
        output = root / "case" / "generated.docx"
        Document().save(output)
        import hashlib
        output_sha256 = hashlib.sha256(output.read_bytes()).hexdigest()
        format_spec = requirements / "format-spec.json"
        format_spec.write_text("{}\n", encoding="utf-8")
        (requirements / "schema-validation.json").write_text(
            json.dumps({"valid": True, "errors": []}), encoding="utf-8"
        )
        capability = work / "capability-preflight.json"
        capability.write_text(json.dumps({"status": "passed", "findings": []}), encoding="utf-8")
        validation = apply_dir / "validation-report.json"
        validation.write_text(json.dumps({
            "valid": True,
            "format_ready": True,
            "serialized_docx_valid": True,
            "render_validation": {
                "rendered_verified": render,
                "evidence": {"source_docx": {"sha256": output_sha256}},
            },
            "submission_ready": render,
            "output_docx": str(output),
            "property_receipt_audit": {
                "valid": True,
                "receipt_count": 1,
                "verified_count": 1,
                "failed_count": 0,
                "unverified_count": 0,
                "missing_count": 0,
                "unexpected_count": 0,
                "duplicate_count": 0,
                "expected_receipt_ids": ["PR-R00001-0001"],
                "receipts": [{"receipt_id": "PR-R00001-0001", "status": "verified",
                              "serialized_docx_sha256": output_sha256}],
            },
        }), encoding="utf-8")
        comparison = apply_dir / "format-comparison.json"
        comparison.write_text(json.dumps({"status": "passed"}), encoding="utf-8")
        manifest = work / "pipeline-manifest.json"
        manifest.write_text(json.dumps({
            "status": "completed",
            "compliance_mode": compliance_mode,
            "output": str(output),
            "format_spec": str(format_spec),
            "capability_preflight": str(capability),
            "capability_preflight_status": "passed",
            "validation_report": str(validation),
            "format_comparison": str(comparison),
            "submission_ready": render,
        }), encoding="utf-8")
        return {
            "returncode": returncode,
            "fresh_run": {
                "run_id": "fresh-run",
                "cache_reused": False,
                "pipeline_manifest": str(manifest),
            },
        }

    def test_returncode_three_and_blocked_manifest_are_not_acceptance(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            result = self._result(root, returncode=3)
            manifest = Path(result["fresh_run"]["pipeline_manifest"])
            payload = json.loads(manifest.read_text(encoding="utf-8"))
            payload["status"] = "blocked"
            payload["blocking_reasons"] = ["open_questions"]
            manifest.write_text(json.dumps(payload), encoding="utf-8")
            accepted = batch.case_acceptance(result)
            self.assertFalse(accepted["accepted"])
            self.assertIn("pipeline_returncode:3", accepted["blockers"])
            self.assertIn("pipeline_status:blocked", accepted["blockers"])

    def test_missing_trusted_render_is_not_acceptance(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            accepted = batch.case_acceptance(self._result(Path(td), render=False))
            self.assertFalse(accepted["accepted"])
            self.assertIn("post_render_acceptance_required_missing", accepted["blockers"])

    def test_current_artifacts_can_pass_only_as_one_bound_set(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            accepted = batch.case_acceptance(self._result(Path(td)))
            self.assertFalse(accepted["accepted"])
            self.assertIn("post_render_acceptance_required_missing", accepted["blockers"])

    def test_full_compliance_requires_post_render_chain(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            accepted = batch.case_acceptance(self._result(Path(td)))
            self.assertFalse(accepted["accepted"])
            self.assertIn("post_render_acceptance_required_missing", accepted["blockers"])
            self.assertIn("post_word_render_manifest_missing", accepted["blockers"])

    def test_host_binding_failure_is_global_batch_stop(self) -> None:
        result = {
            "returncode": 2,
            "stages": {
                "host_agent": {
                    "stderr_tail": "host runtime mismatch: expected codex, observed openclaw",
                },
            },
        }
        self.assertEqual(batch.global_fatal_reason(result), "host runtime mismatch")

    def test_manifest_boundary_and_canonical_selection(self) -> None:
        if not batch.DEFAULT_MANIFEST.is_file():
            self.skipTest("external ten-school manifest is not included in this source checkout")
        _source, cases, _manifest = batch.load_manifest(batch.DEFAULT_MANIFEST)
        by_id = {case["id"]: case for case in cases}
        self.assertIsNone(by_id["ustb"]["style_template"])
        self.assertEqual(by_id["ustb"]["template_boundary"], "requirements_only")
        selected = batch.select_cases(cases, ["ustb", "bsu"])
        self.assertEqual([case["id"] for case in selected], ["bsu", "ustb"])
        command = batch.pipeline_command(
            by_id["ustb"], Path("source.tex"), Path("work"), Path("output.docx"),
            compliance_mode="full", prepare_host_review=True,
        )
        self.assertNotIn("--style-template", command)

    def test_pipeline_command_can_bind_review_and_execution_stages(self) -> None:
        command = batch.pipeline_command(
            {"id": "bsu", "requirements": Path("requirements.docx"),
             "analysis_mode": "llm_primary"},
            Path("source.tex"), Path("work"), Path("output.docx"),
            compliance_mode="full", prepare_host_review=False,
            requirements_dir=Path("work/execution/requirements"),
            llm_response=Path("work/review/host-agent-response.json"),
            run_id="run-1",
            host_agent_audit=Path("work/review/requirements/host-agent-run.json"),
            merge_receipt=Path("work/review/requirements/merge-receipt.json"),
        )
        self.assertIn("--requirements-dir", command)
        self.assertIn("--allow-existing-work", command)
        self.assertIn("--host-agent-audit", command)
        self.assertIn("--merge-receipt", command)
        self.assertIn("--case-id", command)
        self.assertIn("bsu", command)

    def test_pipeline_command_binds_post_format_review_to_explicit_host_model(self) -> None:
        command = batch.pipeline_command(
            {"id": "bsu", "requirements": Path("requirements.docx"),
             "analysis_mode": "llm_primary"},
            Path("source.tex"), Path("work"), Path("output.docx"),
            compliance_mode="full", prepare_host_review=False,
            semantic_review_runtime="codex", semantic_review_model="gpt-5.6-luna",
        )
        self.assertIn("--semantic-review-runtime", command)
        self.assertEqual(command[command.index("--semantic-review-runtime") + 1], "codex")
        self.assertEqual(command[command.index("--semantic-review-model") + 1], "gpt-5.6-luna")

    def test_pipeline_command_passes_explicit_source_bound_thesis_profile(self) -> None:
        command = batch.pipeline_command(
            {"id": "bsu", "requirements": Path("requirements.docx"),
             "analysis_mode": "llm_primary", "thesis_profile": Path("profile.json")},
            Path("source.tex"), Path("work"), Path("output.docx"),
            compliance_mode="full", prepare_host_review=True,
        )
        self.assertIn("--thesis-profile", command)
        self.assertEqual(command[command.index("--thesis-profile") + 1], "profile.json")

    def test_review_draft_with_manual_items_is_accepted_without_release_render(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            case_root = root / "case"
            work = case_root / "work"
            requirements = work / "requirements"
            apply_dir = work / "application"
            requirements.mkdir(parents=True)
            apply_dir.mkdir(parents=True)
            output = case_root / "generated.docx"
            Document().save(output)
            output_sha256 = __import__("hashlib").sha256(output.read_bytes()).hexdigest()
            format_spec = requirements / "format-spec.json"
            format_spec.write_text("{}\n", encoding="utf-8")
            (requirements / "schema-validation.json").write_text(
                json.dumps({"valid": True, "errors": []}), encoding="utf-8"
            )
            capability = work / "capability-preflight.json"
            capability.write_text(json.dumps({
                "status": "blocked",
                "findings": [{
                    "code": "capability.clause_gap",
                    "blocking": True,
                    "message": "C00076 requires manual review",
                    "evidence": [{"kind": "category", "value": "runtime_manual_unverifiable"}],
                }],
            }), encoding="utf-8")
            ledger = work / "manual-review-items.json"
            ledger.write_text(json.dumps({
                "schema_version": "1.0",
                "policy": "review_draft_only",
                "binding": {"run_id": "fresh-run"},
                "submission_ready": False,
                "items": [{"marker_id": "MR-0001", "category": "runtime_manual_unverifiable",
                           "source_codes": ["capability.clause_gap"]}],
                "summary": {"total": 1},
            }), encoding="utf-8")
            markers = apply_dir / "manual-review-markers.json"
            doc = Document(output)
            append_manual_review_markers(doc, json.loads(ledger.read_text()))
            doc.save(output)
            output_sha256 = __import__("hashlib").sha256(output.read_bytes()).hexdigest()
            markers.write_text(json.dumps({
                "schema_version": "1.0", "policy": "review_draft",
                "markers": [{"marker_id": "MR-0001"}],
            }), encoding="utf-8")
            validation = apply_dir / "validation-report.json"
            validation.write_text(json.dumps({
                "valid": True,
                "format_ready": True,
                "serialized_docx_valid": True,
                "diagnostic_draft_generated": True,
                "review_draft_ready": True,
                "review_draft_package_valid": True,
                "submission_ready": False,
                "findings": [],
                "native_semantic_content_review": {
                    "schema_version": "1.0",
                    "protocol": "native_semantic_content_review_v1",
                    "status": "not_required",
                    "case_id": "case",
                    "run_id": "fresh-run",
                    "checks": [],
                    "results": [],
                },
                "property_receipt_audit": {
                    "valid": True,
                    "receipt_count": 1,
                    "verified_count": 1,
                    "failed_count": 0,
                    "unverified_count": 0,
                    "missing_count": 0,
                    "unexpected_count": 0,
                    "duplicate_count": 0,
                    "expected_receipt_ids": ["PR-R00001-0001"],
                    "receipts": [{"receipt_id": "PR-R00001-0001", "status": "verified",
                                  "serialized_docx_sha256": output_sha256}],
                },
            }), encoding="utf-8")
            comparison = apply_dir / "format-comparison.json"
            comparison.write_text(json.dumps({"status": "review_draft_pending"}), encoding="utf-8")
            manifest = work / "pipeline-manifest.json"
            manifest.write_text(json.dumps({
                "status": "draft_manual_review",
                "output_policy": "review_draft",
                "execution_compliance_mode": "supported_subset",
                "output": str(output),
                "format_spec": str(format_spec),
                "capability_preflight": str(capability),
                "manual_review_items": str(ledger),
                "validation_report": str(validation),
                "format_comparison": str(comparison),
                "diagnostic_draft_generated": True,
                "review_draft_ready": True,
                "submission_ready": False,
                "case_id": "case",
                "requirements_extraction": {"run_id": "fresh-run"},
                "code_fingerprint": batch.runtime_code_fingerprint(),
            }), encoding="utf-8")
            result = {
                "returncode": 0,
                "analysis_mode": "rule_only",
                "fresh_run": {
                    "run_id": "fresh-run",
                    "cache_reused": False,
                    "pipeline_manifest": str(manifest),
                },
            }
            accepted = batch.case_acceptance(result, root=root)
            self.assertTrue(accepted["accepted"], accepted)
            self.assertEqual(accepted["status"], "accepted_review_draft")
            validation_payload = json.loads(validation.read_text(encoding="utf-8"))
            validation_payload["diagnostic_draft_generated"] = False
            validation.write_text(json.dumps(validation_payload), encoding="utf-8")
            missing_diagnostic = batch.case_acceptance(result, root=root)
            self.assertFalse(missing_diagnostic["accepted"])
            self.assertIn("review_draft_not_generated_in_validation", missing_diagnostic["blockers"])
            validation_payload["diagnostic_draft_generated"] = True
            review_payload = validation_payload["native_semantic_content_review"]
            review_payload["case_id"] = "another-case"
            validation.write_text(json.dumps(validation_payload), encoding="utf-8")
            wrong_case = batch.case_acceptance(result, root=root)
            self.assertFalse(wrong_case["accepted"])
            self.assertIn("review_draft_semantic_review_invalid_or_unbound", wrong_case["blockers"])
            review_payload["case_id"] = "case"
            review_payload["status"] = "completed"
            review_payload["checks"] = [{"check_id": "abstract_zh.require_third_person",
                                          "document_text": "本文提出一种模型。"}]
            review_payload["results"] = []
            validation.write_text(json.dumps(validation_payload), encoding="utf-8")
            incomplete_semantic = batch.case_acceptance(result, root=root)
            self.assertFalse(incomplete_semantic["accepted"])
            self.assertIn("review_draft_semantic_review_invalid_or_unbound", incomplete_semantic["blockers"])
            review_payload.update({"status": "not_required", "checks": [], "results": []})
            validation.write_text(json.dumps(validation_payload), encoding="utf-8")
            manifest_payload = json.loads(manifest.read_text(encoding="utf-8"))
            manifest_payload["diagnostic_draft_generated"] = False
            manifest.write_text(json.dumps(manifest_payload), encoding="utf-8")
            missing_manifest_diagnostic = batch.case_acceptance(result, root=root)
            self.assertFalse(missing_manifest_diagnostic["accepted"])
            self.assertIn("review_draft_not_generated_in_manifest", missing_manifest_diagnostic["blockers"])
            manifest_payload["diagnostic_draft_generated"] = True
            manifest.write_text(json.dumps(manifest_payload), encoding="utf-8")
            audit = validation_payload["property_receipt_audit"]
            expected_ids = audit.pop("expected_receipt_ids")
            validation.write_text(json.dumps(validation_payload), encoding="utf-8")
            missing_expected = batch.case_acceptance(result, root=root)
            self.assertFalse(missing_expected["accepted"])
            self.assertIn("review_draft_property_receipts_not_verified", missing_expected["blockers"])
            audit["expected_receipt_ids"] = expected_ids
            audit["receipt_count"] = 0
            validation.write_text(json.dumps(validation_payload), encoding="utf-8")
            inconsistent_count = batch.case_acceptance(result, root=root)
            self.assertFalse(inconsistent_count["accepted"])
            self.assertIn("review_draft_property_receipts_not_verified", inconsistent_count["blockers"])
            audit["receipt_count"] = 1
            validation.write_text(json.dumps(validation_payload), encoding="utf-8")
            capability_payload = json.loads(capability.read_text(encoding="utf-8"))
            capability_payload["findings"].append({
                "code": "format.backend_gap", "blocking": True,
                "evidence": [{"kind": "category", "value": "backend_capability_gap"}],
            })
            capability.write_text(json.dumps(capability_payload), encoding="utf-8")
            technical_capability = batch.case_acceptance(result, root=root)
            self.assertFalse(technical_capability["accepted"])
            self.assertIn("capability_technical_blocking_findings", technical_capability["blockers"])
            capability_payload["findings"].pop()
            capability.write_text(json.dumps(capability_payload), encoding="utf-8")
            # Forged JSON receipts cannot compensate for absent DOCX markers.
            Document().save(output)
            rejected = batch.case_acceptance(result, root=root)
            self.assertFalse(rejected["accepted"])
            self.assertIn("manual_review_serialized_markers_invalid", rejected["blockers"])

    def test_review_draft_technical_findings_and_failed_receipts_block_acceptance(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            result = self._result(root, render=False, compliance_mode="supported_subset")
            manifest_path = Path(result["fresh_run"]["pipeline_manifest"])
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            case_root = root / "case"
            work = case_root / "work"
            apply_dir = work / "application"
            manual = work / "manual-review-items.json"
            manual.write_text(json.dumps({
                "schema_version": "1.0",
                "policy": "review_draft_only",
                "binding": {"run_id": "fresh-run"},
                "submission_ready": False,
                "items": [{
                    "marker_id": "MR-0001",
                    "category": "input_prerequisite",
                    "source_codes": ["official_template_missing"],
                }],
                "summary": {"total": 1},
            }), encoding="utf-8")
            (apply_dir / "manual-review-markers.json").write_text(json.dumps({
                "schema_version": "1.0", "policy": "review_draft",
                "markers": [{"marker_id": "MR-0001"}],
            }), encoding="utf-8")
            output = case_root / "generated.docx"
            doc = Document(output)
            append_manual_review_markers(doc, json.loads(manual.read_text()))
            doc.save(output)
            validation = json.loads((apply_dir / "validation-report.json").read_text(encoding="utf-8"))
            validation.update({
                "valid": False,
                "format_ready": False,
                "review_draft_ready": True,
                "review_draft_package_valid": True,
                "submission_ready": False,
                "findings": [{
                    "role": "keywords_zh", "property": "separator",
                    "template_value": "，", "required_value": "semicolon",
                    "failure_type": "deterministic_format_validation",
                }],
                "property_receipt_audit": {
                    "valid": False,
                    "review_draft_diagnostic_only": True,
                    "receipt_count": 1,
                    "verified_count": 0,
                    "failed_count": 1,
                    "unverified_count": 0,
                    "missing_count": 0,
                    "unexpected_count": 0,
                    "duplicate_count": 0,
                    "expected_receipt_ids": ["PR-R00001-0001"],
                    "receipts": [{
                        "receipt_id": "PR-R00001-0001", "status": "failed",
                        "serialized_docx_sha256": __import__("hashlib").sha256(output.read_bytes()).hexdigest(),
                    }],
                },
            })
            (apply_dir / "validation-report.json").write_text(
                json.dumps(validation), encoding="utf-8"
            )
            comparison = apply_dir / "format-comparison.json"
            comparison.write_text(json.dumps({"status": "review_draft_pending"}), encoding="utf-8")
            manifest.update({
                "status": "draft_manual_review",
                "output_policy": "review_draft",
                "execution_compliance_mode": "supported_subset",
                "submission_ready": False,
                "case_id": "case",
                "requirements_extraction": {"run_id": "fresh-run"},
                "manual_review_items": str(manual),
                "validation_report": str(apply_dir / "validation-report.json"),
                "format_comparison": str(comparison),
                "review_draft_ready": True,
                "code_fingerprint": batch.runtime_code_fingerprint(),
            })
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            accepted = batch.case_acceptance(result, root=root)
            self.assertFalse(accepted["accepted"], accepted)
            self.assertIn("review_draft_technical_validation_not_passed", accepted["blockers"])
            self.assertIn("review_draft_property_receipts_not_verified", accepted["blockers"])
            self.assertFalse(accepted["checks"]["property_receipts"]["all_expected_receipts_verified"])
            self.assertFalse(validation["format_ready"])


if __name__ == "__main__":
    unittest.main()
