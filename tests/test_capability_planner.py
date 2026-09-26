#!/usr/bin/env python3
from __future__ import annotations

import importlib.util
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.format_spec_validation import load_and_validate
from scripts.requirements_engine import build_llm_request
from scripts.semantic_contract import attach_request_provenance

ROOT = Path(__file__).resolve().parents[1]
PY = sys.executable


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.path.insert(0, str(ROOT / "scripts"))
    try:
        spec.loader.exec_module(module)
    finally:
        sys.path.pop(0)
    return module


class CapabilityPlannerTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.planner = load_module("capability_planner", ROOT / "scripts" / "capability_planner.py")
        cls.pipeline = load_module("thesis_format_pipeline_capability_test", ROOT / "scripts" / "thesis_format_pipeline.py")
        cls.registry = json.loads((ROOT / "resources" / "backend-capabilities.default.json").read_text())

    def test_bundled_registry_and_findings_validate(self) -> None:
        errors = load_and_validate(self.registry, ROOT / "schema" / "backend-capability-registry.schema.json")
        self.assertEqual(errors, [])
        spec = {"requirements": [{"id": "R1", "role": "body_text",
                                  "properties": {"font": {"size_pt": 12}}}], "clause_compliance": []}
        report = self.planner.plan_capabilities(spec, self.registry)
        self.assertEqual(report["requirements"][0]["disposition"], "supported")
        self.assertEqual(report["status"], "ready")

    def test_extracted_clauses_without_compliance_fail_closed_in_full_mode(self) -> None:
        extracted = [{"id": "C1"}, {"id": "C2"}]
        spec = {"requirements": [], "clause_compliance": [
            {"clause_id": "C1", "scope": "informational", "status": "informational",
             "requirement_ids": []}
        ]}
        full = self.planner.plan_capabilities(
            spec, self.registry, "full", extracted_clauses=extracted)
        self.assertEqual(full["status"], "blocked")
        self.assertFalse(full["execution_ready"])
        self.assertEqual(full["findings"][0]["code"], "capability.missing_clause_compliance")
        self.assertTrue(full["findings"][0]["blocking"])
        self.assertEqual(full["summary"]["extracted_clauses"], 2)
        self.assertEqual(full["summary"]["reviewed_extracted_clauses"], 1)
        self.assertEqual(full["summary"]["missing_clause_records"], 1)
        evidence_by_kind = {item["kind"]: item["value"] for item in full["findings"][0]["evidence"]}
        self.assertEqual(evidence_by_kind["missing_clause_ids_sample"], ["C2"])
        self.assertEqual(load_and_validate(full["findings"][0],
                                           ROOT / "schema" / "pipeline-finding.schema.json"), [])

        subset = self.planner.plan_capabilities(
            spec, self.registry, "supported_subset", extracted_clauses=extracted)
        self.assertEqual(subset["status"], "gaps_present")
        self.assertTrue(subset["execution_ready"])
        self.assertFalse(subset["findings"][0]["blocking"])

    def test_unknown_property_is_a_stable_blocking_finding_in_full_mode(self) -> None:
        spec = {"requirements": [{"id": "R7", "role": "body_text",
                                  "properties": {"future_layout": {"axis": 1}},
                                  "clause_ids": ["C7"]}],
                "clause_compliance": [{"clause_id": "C7", "scope": "docx",
                                       "status": "pending_execution", "requirement_ids": ["R7"]}]}
        report = self.planner.plan_capabilities(spec, self.registry, "full")
        item = report["findings"][0]
        self.assertEqual(report["requirements"][0]["disposition"], "unknown")
        self.assertEqual(set(item), {"code", "stage", "severity", "blocking", "evidence", "message"})
        self.assertEqual(item["code"], "capability.requirement_unknown")
        self.assertTrue(item["blocking"])
        self.assertEqual(report["requirements"][0]["findings"], [item])
        self.assertFalse(report["execution_ready"])
        self.assertEqual(load_and_validate(item, ROOT / "schema" / "pipeline-finding.schema.json"), [])

    def test_registered_text_role_is_executable(self) -> None:
        spec = {"requirements": [{"id": "R1", "role": "thesis_author",
                                  "properties": {"font": {"size_pt": 15}},
                                  "clause_ids": ["C1"]}],
                "clause_compliance": [{"clause_id": "C1", "scope": "docx",
                                       "status": "pending_execution", "requirement_ids": ["R1"]}]}
        report = self.planner.plan_capabilities(spec, self.registry, "full")
        requirement = report["requirements"][0]
        self.assertEqual(requirement["disposition"], "supported")
        self.assertEqual(report["clauses"][0]["disposition"], "supported")
        self.assertEqual(report["status"], "ready")
        self.assertTrue(report["execution_ready"])

    def test_registered_appendix_layout_properties_are_executable(self) -> None:
        spec = {
            "requirements": [{
                "id": "R1",
                "role": "appendices",
                "properties": {
                    "label_style": "alpha_upper",
                    "page_break_each": True,
                    "per_appendix_title_required": True,
                },
                "clause_ids": ["C1"],
            }],
            "clause_compliance": [{
                "clause_id": "C1",
                "scope": "docx",
                "status": "pending_execution",
                "requirement_ids": ["R1"],
            }],
        }
        report = self.planner.plan_capabilities(spec, self.registry, "full")
        self.assertEqual(report["requirements"][0]["disposition"], "supported")
        self.assertEqual(report["clauses"][0]["disposition"], "supported")
        self.assertEqual(report["status"], "ready")
        self.assertTrue(report["execution_ready"])

    def test_footnote_style_role_is_executable(self) -> None:
        spec = {"requirements": [{"id": "R1", "role": "footnote",
                                  "properties": {"font": {"size_pt": 10.5}},
                                  "clause_ids": ["C1"]}],
                "clause_compliance": [{"clause_id": "C1", "scope": "docx",
                                       "status": "pending_execution", "requirement_ids": ["R1"]}]}
        report = self.planner.plan_capabilities(spec, self.registry, "full")
        requirement = report["requirements"][0]
        self.assertEqual(requirement["disposition"], "supported")
        self.assertEqual(report["clauses"][0]["disposition"], "supported")
        self.assertEqual(report["status"], "ready")
        self.assertTrue(report["execution_ready"])

    def test_explicit_unsupported_and_subset_mode_preserve_gap_without_blocking(self) -> None:
        registry = {"schema_version": "1.0", "backend": "test", "capabilities": [{
            "id": "no-body-font", "role_pattern": "^body_text$", "property_patterns": ["font.*"],
            "disposition": "unsupported", "message": "not implemented"}]}
        spec = {"requirements": [{"id": "R2", "role": "body_text",
                                  "properties": {"font": {"cjk": "SimSun"}}}], "clause_compliance": []}
        full = self.planner.plan_capabilities(spec, registry, "full")
        subset = self.planner.plan_capabilities(spec, registry, "supported_subset")
        self.assertEqual(full["status"], "blocked")
        self.assertEqual(subset["status"], "gaps_present")
        self.assertTrue(full["findings"][0]["blocking"])
        self.assertFalse(subset["findings"][0]["blocking"])
        self.assertTrue(subset["execution_ready"])

    def test_optional_inputs_can_resolve_conditional_capability(self) -> None:
        registry = {"schema_version": "1.0", "backend": "conditional", "capabilities": [{
            "id": "figures", "role_pattern": "^objects$", "property_patterns": ["keep_figure_with_caption"],
            "disposition": "supported", "requires_source_inventory": ["inventory.figures"]}]}
        spec = {"requirements": [{"id": "R3", "role": "objects",
                                  "properties": {"keep_figure_with_caption": True}}], "clause_compliance": []}
        missing = self.planner.plan_capabilities(spec, registry)
        present = self.planner.plan_capabilities(spec, registry, source_inventory={"inventory": {"figures": 2}})
        self.assertEqual(missing["requirements"][0]["disposition"], "unknown")
        self.assertEqual(missing["requirements"][0]["properties"][0]["missing_inputs"], ["inventory.figures"])
        self.assertEqual(present["requirements"][0]["disposition"], "supported")

    def test_false_applicability_is_not_applicable_and_not_executable(self) -> None:
        spec = {
            "requirements": [{
                "id": "R-conditional",
                "role": "body_text",
                "properties": {"font": {"size_pt": 12}},
                "clause_ids": ["C-conditional"],
                "applicability": {
                    "status": "conditional",
                    "conditions": [{
                        "fact": "runtime.host", "operator": "equals", "value": "codex",
                    }],
                },
            }],
            "clause_compliance": [{
                "clause_id": "C-conditional", "scope": "docx",
                "status": "pending_execution", "requirement_ids": ["R-conditional"],
            }],
        }
        report = self.planner.plan_capabilities(
            spec, self.registry, "full", source_inventory={"runtime": {"host": "openclaw"}}
        )
        requirement = report["requirements"][0]
        clause = report["clauses"][0]
        self.assertEqual(requirement["applicability_evaluation"]["result"], "false")
        self.assertEqual(requirement["category"], "external_not_applicable")
        self.assertEqual(requirement["findings"], [])
        self.assertEqual(clause["category"], "external_not_applicable")
        self.assertEqual(clause["disposition"], "not_applicable")
        self.assertEqual(report["applicability"]["excluded_requirement_ids"], ["R-conditional"])

    def test_unknown_applicability_blocks_full_execution(self) -> None:
        spec = {
            "requirements": [{
                "id": "R-unknown",
                "role": "body_text",
                "properties": {"font": {"size_pt": 12}},
                "applicability": {
                    "status": "conditional",
                    "conditions": [{
                        "fact": "runtime.host", "operator": "equals", "value": "codex",
                    }],
                },
            }],
            "clause_compliance": [],
        }
        report = self.planner.plan_capabilities(spec, self.registry, "full")
        self.assertEqual(report["applicability"]["unknown_requirement_ids"], ["R-unknown"])
        self.assertEqual(report["status"], "blocked")
        self.assertFalse(report["execution_ready"])

    def test_declared_requirement_prerequisite_is_classified_as_input_gap(self) -> None:
        spec = {"requirements": [{
            "id": "R-input", "role": "body_text",
            "properties": {"font": {"size_pt": 12}}, "clause_ids": ["C-input"],
            "input_prerequisites": [{
                "kind": "metadata", "key": "thesis_profile.degree_level",
                "required": True, "reason": "degree-specific requirement"
            }],
            "verification": {"mode": "static_docx", "checks": ["font_size"]}
        }], "clause_compliance": [{
            "clause_id": "C-input", "scope": "docx", "status": "pending_execution",
            "requirement_ids": ["R-input"], "evidence_ids": ["E-input"], "reason": "pending"
        }]}
        missing = self.planner.plan_capabilities(spec, self.registry, "full")
        item = missing["requirements"][0]
        self.assertEqual(item["disposition"], "unknown")
        self.assertEqual(item["category"], "input_prerequisite")
        self.assertEqual(item["missing_declared_inputs"], ["thesis_profile.degree_level"])
        self.assertEqual(missing["summary"]["requirement_input_prerequisites"], 1)
        present = self.planner.plan_capabilities(
            spec, self.registry, "full",
            source_inventory={"thesis_profile": {"degree_level": "doctor"}})
        self.assertEqual(present["requirements"][0]["disposition"], "supported")
        self.assertEqual(present["requirements"][0]["category"], "supported")

    def test_metadata_only_profile_satisfies_declared_prerequisite(self) -> None:
        spec = {"requirements": [{
            "id": "R-profile", "role": "body_text",
            "properties": {"font": {"size_pt": 12}}, "clause_ids": ["C-profile"],
            "input_prerequisites": [{
                "kind": "metadata", "key": "thesis_profile.degree_level",
                "required": True, "reason": "degree-specific rule"
            }],
        }], "clause_compliance": [{
            "clause_id": "C-profile", "scope": "docx", "status": "pending_execution",
            "requirement_ids": ["R-profile"], "evidence_ids": ["E-profile"],
        }]}
        report = self.planner.plan_capabilities(
            spec, self.registry, "full", metadata={"degree_level": "doctor"})
        item = report["requirements"][0]
        self.assertEqual(item["missing_declared_inputs"], [])
        self.assertEqual(item["category"], "supported")
        self.assertTrue(report["execution_ready"])
        self.assertTrue(report["inputs"]["metadata_provided"])

    def test_cover_metadata_object_satisfies_declared_prerequisite(self) -> None:
        spec = {"requirements": [{
            "id": "R-cover-profile", "role": "cover",
            "properties": {"institution": "——", "fields": []},
            "clause_ids": ["C-cover-profile"],
            "input_prerequisites": [{
                "kind": "metadata", "key": "thesis_profile.cover_metadata",
                "required": True, "reason": "cover metadata object required",
            }],
        }], "clause_compliance": [{
            "clause_id": "C-cover-profile", "scope": "docx", "status": "pending_execution",
            "requirement_ids": ["R-cover-profile"], "evidence_ids": ["E-cover-profile"],
        }]}
        report = self.planner.plan_capabilities(
            spec, self.registry, "full",
            metadata={"cover_metadata": {"title_zh": "测试题目"}})
        self.assertEqual(report["requirements"][0]["missing_declared_inputs"], [])
        self.assertEqual(report["requirements"][0]["category"], "supported")

    def test_runtime_anchor_inventory_satisfies_runtime_prerequisite(self) -> None:
        spec = {"requirements": [{
            "id": "R-anchor", "role": "declarations",
            "properties": {"before_role": "abstract_title_zh", "items": [{
                "id": "authorization", "signature_placeholders": []
            }]}, "clause_ids": ["C-anchor"],
            "input_prerequisites": [{
                "kind": "runtime", "key": "runtime.declaration_anchor",
                "required": True, "reason": "anchor must be observed"
            }],
        }], "clause_compliance": [{
            "clause_id": "C-anchor", "scope": "docx", "status": "pending_execution",
            "requirement_ids": ["R-anchor"], "evidence_ids": ["E-anchor"],
        }]}
        report = self.planner.plan_capabilities(
            spec, self.registry, "full",
            runtime_inventory={"declaration_anchor": "abstract_title_zh"})
        self.assertEqual(report["requirements"][0]["missing_declared_inputs"], [])
        self.assertEqual(report["requirements"][0]["category"], "supported")

    def test_nested_runtime_anchor_selection_satisfies_runtime_prerequisite(self) -> None:
        spec = {"requirements": [{
            "id": "R-anchor-selected", "role": "declarations",
            "properties": {"before_role": "abstract_title_zh", "items": [{
                "id": "authorization", "signature_placeholders": []
            }]}, "clause_ids": ["C-anchor-selected"],
            "input_prerequisites": [{
                "kind": "runtime", "key": "runtime.anchor_inventory.selected",
                "required": True, "reason": "selected anchor must be verified"
            }],
        }], "clause_compliance": [{
            "clause_id": "C-anchor-selected", "scope": "docx", "status": "pending_execution",
            "requirement_ids": ["R-anchor-selected"], "evidence_ids": ["E-anchor-selected"],
        }]}
        report = self.planner.plan_capabilities(
            spec, self.registry, "full",
            runtime_inventory={"anchor_inventory": {"selected": {
                "name": "abstract_title_zh", "binding_status": "verified",
            }}})
        self.assertEqual(report["requirements"][0]["missing_declared_inputs"], [])
        self.assertEqual(report["requirements"][0]["category"], "supported")

    def test_conflicting_profile_namespaces_block_instead_of_picking_one(self) -> None:
        spec = {"requirements": [{
            "id": "R-conflict", "role": "body_text",
            "properties": {"font": {"size_pt": 12}}, "clause_ids": ["C-conflict"],
            "input_prerequisites": [{
                "kind": "metadata", "key": "thesis_profile.cover_metadata.title_zh",
                "required": True, "reason": "title required"
            }],
        }], "clause_compliance": [{
            "clause_id": "C-conflict", "scope": "docx", "status": "pending_execution",
            "requirement_ids": ["R-conflict"], "evidence_ids": ["E-conflict"],
        }]}
        report = self.planner.plan_capabilities(
            spec, self.registry, "full",
            metadata={"cover_metadata": {"title_zh": "甲"}},
            source_inventory={"thesis_profile": {"cover_metadata": {"title_zh": "乙"}}})
        item = report["requirements"][0]
        self.assertEqual(item["category"], "input_prerequisite")
        self.assertEqual(len(item["input_conflicts"]), 1)
        self.assertFalse(report["execution_ready"])

    def test_zero_is_a_supplied_inventory_value(self) -> None:
        registry = {"schema_version": "1.0", "backend": "test", "capabilities": [{
            "id": "figures", "role_pattern": "^objects$",
            "property_patterns": ["keep_figure_with_caption"], "disposition": "supported",
            "requires_source_inventory": ["inventory.figures"],
        }]}
        spec = {"requirements": [{"id": "R-zero", "role": "objects",
                                  "properties": {"keep_figure_with_caption": True}}],
                "clause_compliance": []}
        report = self.planner.plan_capabilities(
            spec, registry, source_inventory={"inventory": {"figures": 0}})
        self.assertEqual(report["requirements"][0]["disposition"], "supported")

    def test_content_instance_missing_inventory_is_input_gap_not_backend_gap(self) -> None:
        spec = {"requirements": [{
            "id": "R00677", "role": "heading_publications",
            "field_instance_ids": ["publication-heading"],
            "input_prerequisites": [{
                "kind": "source_inventory", "key": "source_inventory.publications",
                "required": True, "reason": "publication heading needs source content"
            }],
            "clause_ids": ["C00395"],
        }], "clause_compliance": [{
            "clause_id": "C00395", "scope": "docx", "status": "pending_execution",
            "requirement_ids": ["R00677"], "evidence_ids": ["E00395"],
            "reason": "publication heading is representable subject to source content"
        }]}
        report = self.planner.plan_capabilities(spec, self.registry, "full")
        requirement = report["requirements"][0]
        clause = report["clauses"][0]
        self.assertEqual(requirement["category"], "input_prerequisite")
        self.assertEqual(requirement["disposition"], "unknown")
        self.assertEqual(requirement["findings"][0]["code"],
                         "capability.requirement_input_prerequisite")
        self.assertEqual(clause["category"], "input_prerequisite")
        self.assertEqual(clause["category_source"], "requirement")
        self.assertEqual(clause["findings"], [])
        self.assertEqual(report["summary"]["backend_capability_gaps"], 0)
        self.assertEqual(report["summary"]["input_prerequisites"], 1)
        self.assertEqual(report["summary"]["finding_backend_capability_gaps"], 0)
        self.assertEqual(report["summary"]["gaps"], 1)
        self.assertEqual(report["summary"]["finding_gaps"], 1)

    def test_clause_dispositions_include_unlinked_analysis_gaps(self) -> None:
        spec = {"requirements": [], "clause_compliance": [
            {"clause_id": "C1", "scope": "docx", "status": "unsupported_backend",
             "requirement_ids": [], "evidence_ids": ["E1"], "reason": "backend lacks adapter"},
            {"clause_id": "C2", "scope": "informational", "status": "informational",
             "requirement_ids": [], "evidence_ids": ["E2"], "reason": "advisory only"},
        ]}
        report = self.planner.plan_capabilities(spec, self.registry)
        self.assertEqual([item["disposition"] for item in report["clauses"]], ["unsupported", "not_applicable"])
        self.assertEqual([item["category"] for item in report["clauses"]],
                         ["backend_capability_gap", "external_not_applicable"])
        self.assertEqual(report["findings"][0]["code"], "capability.clause_backend_unsupported")
        self.assertEqual(report["clauses"][0]["findings"], [report["findings"][0]])
        self.assertEqual(report["clauses"][1]["findings"], [])
        self.assertEqual(report["summary"]["backend_capability_gaps"], 1)
        self.assertEqual(report["summary"]["external_not_applicable"], 1)

    def test_clause_gap_categories_have_stable_codes_and_traceable_evidence(self) -> None:
        records = [
            {"clause_id": "C-meta", "scope": "docx", "status": "requires_metadata",
             "requirement_ids": [], "evidence_ids": ["E-meta"], "reason": "author is absent"},
            {"clause_id": "C-source", "scope": "docx", "status": "requires_source_content",
             "requirement_ids": [], "evidence_ids": ["E-source"], "reason": "abstract is absent"},
            {"clause_id": "C-manual", "scope": "docx", "status": "unverifiable",
             "requirement_ids": [], "evidence_ids": ["E-manual"], "reason": "manual review required"},
        ]
        report = self.planner.plan_capabilities({"requirements": [], "clause_compliance": records},
                                                self.registry, "full")
        self.assertEqual([item["category"] for item in report["clauses"]], [
            "input_prerequisite", "input_prerequisite", "runtime_manual_unverifiable"])
        self.assertEqual([item["code"] for item in report["findings"]], [
            "capability.clause_input_requires_metadata",
            "capability.clause_input_requires_source_content",
            "capability.clause_runtime_unverifiable",
        ])
        for finding_item, record in zip(report["findings"], records):
            evidence_by_kind = {item["kind"]: item["value"] for item in finding_item["evidence"]}
            self.assertEqual(set(evidence_by_kind), {
                "clause_id", "source_status", "scope", "requirement_ids", "evidence_ids", "reason"})
            self.assertEqual(evidence_by_kind["clause_id"], record["clause_id"])
            self.assertEqual(evidence_by_kind["evidence_ids"], record["evidence_ids"])
            self.assertTrue(finding_item["blocking"])
            self.assertEqual(load_and_validate(finding_item,
                                               ROOT / "schema" / "pipeline-finding.schema.json"), [])
        self.assertEqual(report["summary"]["gaps"], 3)
        self.assertEqual(report["summary"]["backend_capability_gaps"], 0)
        self.assertEqual(report["summary"]["input_prerequisites"], 2)
        self.assertEqual(report["summary"]["runtime_manual_unverifiable"], 1)
        self.assertFalse(report["execution_ready"])

    def test_explicit_metadata_label_is_satisfied_by_normalized_metadata(self) -> None:
        extracted = [{"id": "C-meta", "text": "分类号", "source_text_full": "分类号："}]
        spec = {"requirements": [], "clause_compliance": [{
            "clause_id": "C-meta", "scope": "docx", "status": "requires_metadata",
            "requirement_ids": [], "evidence_ids": ["E-meta"],
            "reason": "metadata is needed"
        }]}
        report = self.planner.plan_capabilities(
            spec, self.registry, "full", extracted_clauses=extracted,
            metadata={"classification_number": "U491.1"})
        clause = report["clauses"][0]
        self.assertEqual(clause["disposition"], "supported")
        self.assertEqual(clause["category"], "supported")
        self.assertEqual(clause["metadata_fields"], ["classification_number"])
        self.assertEqual(clause["input_status"], "provided")
        self.assertEqual(clause["binding_status"], "bound")
        self.assertEqual(clause["output_status"], "unverified")
        self.assertEqual(report["summary"]["clause_input_prerequisites"], 0)
        self.assertTrue(report["execution_ready"])

    def test_semantic_metadata_binding_is_input_only(self) -> None:
        extracted = [{"id": "C-abstract", "text": "中文摘要应为300～1000字"}]
        spec = {"requirements": [], "clause_compliance": [{
            "clause_id": "C-abstract", "scope": "docx",
            "status": "requires_source_content", "requirement_ids": [],
            "evidence_ids": ["E-abstract"]
        }]}
        report = self.planner.plan_capabilities(
            spec, self.registry, "full", extracted_clauses=extracted,
            metadata={"abstract_zh": "一段摘要"})
        clause = report["clauses"][0]
        self.assertEqual(clause["category"], "supported")
        self.assertEqual(clause["input_status"], "provided")
        self.assertEqual(clause["output_status"], "unverified")
        self.assertEqual(clause["metadata_fields"], ["abstract_zh"])

    def test_unknown_metadata_label_remains_fail_closed(self) -> None:
        extracted = [{"id": "C-meta", "text": "学校代码", "source_text_full": "学校代码："}]
        spec = {"requirements": [], "clause_compliance": [{
            "clause_id": "C-meta", "scope": "docx", "status": "requires_metadata",
            "requirement_ids": [], "evidence_ids": ["E-meta"]
        }]}
        report = self.planner.plan_capabilities(
            spec, self.registry, "full", extracted_clauses=extracted,
            metadata={"classification_number": "U491.1"})
        self.assertEqual(report["clauses"][0]["category"], "input_prerequisite")
        self.assertFalse(report["execution_ready"])

    def test_one_requirement_bound_to_one_clause_has_separate_nonduplicated_counts(self) -> None:
        spec = {"requirements": [{"id": "R1", "role": "body_text",
                                  "properties": {"future_layout": {"axis": 1}},
                                  "clause_ids": ["C1"]}],
                "clause_compliance": [{"clause_id": "C1", "scope": "docx",
                                       "status": "pending_execution", "requirement_ids": ["R1"]}]}
        report = self.planner.plan_capabilities(spec, self.registry, "supported_subset")
        summary = report["summary"]
        self.assertEqual(summary["gaps"], 1)
        self.assertEqual(summary["finding_gaps"], 1)
        self.assertEqual(summary["requirement_gaps"], 1)
        self.assertEqual(summary["requirement_backend_capability_gaps"], 1)
        self.assertEqual(summary["clause_gaps"], 0)
        self.assertEqual(summary["clause_backend_capability_gaps"], 0)

    def test_multiple_clauses_bound_to_one_requirement_count_each_clause_once(self) -> None:
        spec = {"requirements": [{"id": "R1", "role": "body_text",
                                  "properties": {"future_layout": {"axis": 1}},
                                  "clause_ids": ["C1", "C2"]}],
                "clause_compliance": [
                    {"clause_id": "C1", "scope": "docx", "status": "pending_execution",
                     "requirement_ids": ["R1"]},
                    {"clause_id": "C2", "scope": "docx", "status": "pending_execution",
                     "requirement_ids": ["R1"]},
                ]}
        summary = self.planner.plan_capabilities(
            spec, self.registry, "supported_subset")["summary"]
        self.assertEqual(summary["gaps"], 1)
        self.assertEqual(summary["finding_gaps"], 1)
        self.assertEqual(summary["requirement_gaps"], 1)
        self.assertEqual(summary["clause_gaps"], 0)
        self.assertEqual(summary["clause_backend_capability_gaps"], 0)

    def test_unlinked_clause_and_external_clause_use_clause_only_gap_semantics(self) -> None:
        spec = {"requirements": [], "clause_compliance": [
            {"clause_id": "C-gap", "scope": "docx", "status": "requires_metadata",
             "requirement_ids": []},
            {"clause_id": "C-external", "scope": "external_submission",
             "status": "external_compliance", "requirement_ids": []},
        ]}
        summary = self.planner.plan_capabilities(
            spec, self.registry, "supported_subset")["summary"]
        self.assertEqual(summary["requirement_gaps"], 0)
        self.assertEqual(summary["clause_gaps"], 1)
        self.assertEqual(summary["clause_input_prerequisites"], 1)
        self.assertEqual(summary["clause_external_not_applicable"], 1)
        self.assertEqual(summary["external_not_applicable"], 1)
        self.assertEqual(summary["gaps"], 1)

    def test_duplicate_named_clause_records_do_not_inflate_clause_counts(self) -> None:
        records = [
            {"clause_id": "C1", "scope": "docx", "status": "requires_metadata",
             "requirement_ids": []},
            {"clause_id": "C1", "scope": "docx", "status": "requires_metadata",
             "requirement_ids": []},
        ]
        summary = self.planner.plan_capabilities(
            {"requirements": [], "clause_compliance": records}, self.registry,
            "supported_subset")["summary"]
        self.assertEqual(summary["gaps"], 1)
        self.assertEqual(summary["finding_gaps"], 2)
        self.assertEqual(summary["clause_gaps"], 1)
        self.assertEqual(summary["clause_input_prerequisites"], 1)

    def test_manifest_uses_clause_categories_but_keeps_legacy_capability_gaps(self) -> None:
        manifest = {}
        self.pipeline.record_capability_summary(manifest, {
            "gaps": 3,
            "backend_capability_gaps": 3,
            "clause_gaps": 2,
            "clause_backend_capability_gaps": 2,
            "clause_input_prerequisites": 0,
            "clause_runtime_manual_unverifiable": 0,
            "clause_external_not_applicable": 4,
            "requirement_gaps": 1,
        })
        self.assertEqual(manifest["capability_gaps"], 3)
        self.assertEqual(manifest["capability_clause_gaps"], 2)
        self.assertEqual(manifest["capability_requirement_gaps"], 1)
        self.assertEqual(manifest["capability_backend_gaps"], 2)
        self.assertEqual(manifest["capability_external_not_applicable"], 4)

    def test_pipeline_gate_blocks_full_but_not_supported_subset(self) -> None:
        report = {"findings": [{"code": "capability.requirement_unknown", "blocking": True}]}
        self.assertTrue(self.pipeline.capability_gate_blocked(report, "full"))
        self.assertFalse(self.pipeline.capability_gate_blocked(report, "supported_subset"))

    def test_semantic_contract_gate_blocks_invalid_response_before_capability(self) -> None:
        clauses = [{"id": "C1", "text": "正文使用宋体", "evidence_ids": ["E1"]}]
        evidence = {"evidence": [{"id": "E1", "text": "正文使用宋体", "kind": "paragraph"}]}
        self._bind_test_source_spans(clauses, evidence)
        request = build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
        )
        invalid = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "body_text", "properties": {"text": "正文使用宋体"},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "confidence": 0.9, "reason": "错误地为 informational 条款生成 requirement",
            }],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "informational",
                "reason": "这是示例内容，不是格式义务",
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            request_path = td / "llm-request.json"
            response_path = td / "llm-response.json"
            request_path.write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
            response_path.write_text(json.dumps(invalid, ensure_ascii=False), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "before capability planning"):
                self.pipeline.validate_semantic_contract_gate(
                    response_path=response_path,
                    request_path=request_path,
                    require_provenance=False,
                )

    def test_semantic_contract_gate_requires_current_provenance_for_release(self) -> None:
        clauses = [{"id": "C1", "text": "正文使用宋体", "evidence_ids": ["E1"]}]
        evidence = {"evidence": [{"id": "E1", "text": "正文使用宋体", "kind": "paragraph"}]}
        self._bind_test_source_spans(clauses, evidence)
        request = build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
        )
        request = attach_request_provenance(
            request, source_sha256="a" * 64, evidence_doc=evidence,
            clauses=clauses, run_id="run-current",
        )
        response = {
            "contract_version": "3.0", "requirements": [],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "informational",
                "reason": "这是示例内容，不是格式义务",
            }],
            "unsupported_items": [], "reported_conflicts": [],
            "provenance": {**request["provenance"], "run_id": "old-run"},
        }
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            request_path = td / "llm-request.json"
            response_path = td / "llm-response.json"
            request_path.write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
            response_path.write_text(json.dumps(response, ensure_ascii=False), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "before capability planning"):
                self.pipeline.validate_semantic_contract_gate(
                    response_path=response_path,
                    request_path=request_path,
                    require_provenance=True,
                )

    @staticmethod
    def _bind_test_source_spans(clauses: list[dict], evidence_doc: dict) -> None:
        evidence_by_id = {
            str(item.get("id")): item for item in evidence_doc.get("evidence", [])
            if isinstance(item, dict) and item.get("id")
        }
        cursors: dict[str, int] = {}
        for clause in clauses:
            evidence_id = str(clause["evidence_ids"][0])
            source = evidence_by_id[evidence_id]["text"]
            start = source.find(clause["text"], cursors.get(evidence_id, 0))
            if start < 0:
                raise AssertionError(f"test source span not found for {clause['id']}")
            end = start + len(clause["text"])
            clause["source_span"] = {
                "evidence_id": evidence_id, "start_offset": start, "end_offset": end,
                "text": source[start:end],
                "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
            }
            cursors[evidence_id] = end

    def test_pipeline_runs_preflight_before_style_and_records_subset_gaps(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); requirements = td / "requirements.docx"; target = td / "target.docx"
            rules = Document()
            for text in (
                "论文中文题目使用二号黑体，居中排列。",
                "一级标题使用小三号黑体，编号形式为“第一章”。",
                "正文中文使用小四号宋体，英文和数字使用 Times New Roman，行距固定值20磅。",
                "纸张采用A4，上页边距2.5厘米，下页边距2.0厘米，左页边距3.0厘米，右页边距2.0厘米。",
            ):
                rules.add_paragraph(text)
            rules.save(requirements)
            thesis = Document()
            thesis.add_paragraph("测试论文").style = "Title"
            thesis.add_paragraph("绪论").style = "Heading 1"
            thesis.add_paragraph("正文内容 with English 123.").style = "Normal"
            thesis.save(target)
            registry = json.loads(json.dumps(self.registry))
            registry["capabilities"].append({
                "id": "test-body-font-gap", "role_pattern": "^body_text$",
                "property_patterns": ["font.*"], "disposition": "unsupported",
                "message": "integration-test gap",
            })
            registry_path = td / "registry.json"
            registry_path.write_text(json.dumps(registry), encoding="utf-8")
            blocked = subprocess.run([
                PY, "scripts/thesis_format_pipeline.py", str(requirements), str(target), str(td / "blocked.docx"),
                "--work-dir", str(td / "blocked-work"), "--analysis-mode", "rule_only",
                "--compliance-mode", "full", "--capability-registry", str(registry_path),
            ], cwd=ROOT, text=True, capture_output=True)
            # Full compliance may no longer enter the pipeline through a
            # deterministic baseline mode; reject the invalid state before any
            # conversion or capability output is created.
            self.assertEqual(blocked.returncode, 2, blocked.stderr + blocked.stdout)
            self.assertIn("requires --analysis-mode llm_primary", blocked.stderr)
            self.assertFalse((td / "blocked-work").exists())

            result = subprocess.run([
                PY, "scripts/thesis_format_pipeline.py", str(requirements), str(target), str(td / "output.docx"),
                "--work-dir", str(td / "work"), "--analysis-mode", "rule_only",
                "--compliance-mode", "supported_subset", "--capability-registry", str(registry_path),
            ], cwd=ROOT, text=True, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((td / "work" / "capability-preflight.json").read_text())
            manifest = json.loads((td / "work" / "pipeline-manifest.json").read_text())
            self.assertEqual(report["status"], "gaps_present")
            self.assertGreater(report["summary"]["gaps"], 0)
            self.assertEqual(manifest["capability_preflight_status"], "gaps_present")
            self.assertGreater(manifest["capability_gaps"], 0)
            self.assertEqual(manifest["capability_backend_gaps"],
                             report["summary"]["clause_backend_capability_gaps"])
            self.assertEqual(manifest["capability_input_prerequisites"],
                             report["summary"]["clause_input_prerequisites"])
            names = [step["name"] for step in manifest["steps"]]
            self.assertLess(names.index("requirements"), names.index("capability_preflight"))
            self.assertLess(names.index("capability_preflight"), names.index("style_analysis"))
            self.assertTrue(Path(manifest["section_plan"]).is_file())
            self.assertTrue(manifest["section_plan_valid"])
            self.assertFalse((td / "work" / "assembly-plan.json").exists())

    def test_pipeline_region_graph_blocks_before_style_analysis(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            requirements = base / "requirements.docx"
            source = base / "source.docx"
            template = base / "template.docx"
            profile = base / "profile.json"
            output = base / "out.docx"
            work = base / "work"
            rules = Document()
            rules.add_paragraph("论文正文使用小四号宋体，行距固定值20磅。")
            rules.save(requirements)
            thesis = Document()
            thesis.add_paragraph("测试论文").style = "Title"
            thesis.add_paragraph("绪论").style = "Heading 1"
            thesis.save(source)
            official = Document()
            official.add_paragraph("唯一模板段落")
            official.save(template)
            profile.write_text(json.dumps({
                "schema_version": "1.0",
                "profile_id": "test-region-profile",
                "resources": [{"kind": "official_docx", "path": "template.docx"}],
                "document": {},
                "regions": {
                    "graph_id": "invalid-selector",
                    "nodes": [{
                        "id": "body", "kind": "dynamic_content",
                        "selector": {"text_equals": "不存在的段落"},
                    }],
                    "edges": [],
                },
            }, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([
                PY, "scripts/thesis_format_pipeline.py", str(requirements), str(source), str(output),
                "--work-dir", str(work), "--analysis-mode", "rule_only",
                "--template-profile", str(profile),
                "--compliance-mode", "supported_subset",
            ], cwd=ROOT, text=True, capture_output=True)
            self.assertEqual(result.returncode, 10, result.stderr + result.stdout)
            manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["reason"], "region graph preflight blocked execution")
            self.assertEqual(manifest["assembly_plan_status"], "invalid")
            self.assertTrue((work / "assembly-plan.json").is_file())
            self.assertNotIn("style_analysis", [step["name"] for step in manifest["steps"]])

    def test_required_template_structure_without_assembly_contract_blocks_preflight(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            requirements = base / "requirements.docx"
            source = base / "source.docx"
            template = base / "template.docx"
            profile = base / "profile.json"
            work = base / "work"
            rules = Document()
            rules.add_paragraph("论文正文使用小四号宋体，行距固定值20磅。")
            rules.save(requirements)
            thesis = Document()
            thesis.add_paragraph("测试论文").style = "Title"
            thesis.add_paragraph("正文内容").style = "Normal"
            thesis.save(source)
            official = Document()
            official.add_paragraph("固定封面")
            official.save(template)
            profile.write_text(json.dumps({
                "schema_version": "1.0",
                "profile_id": "required-structure-without-regions",
                "resources": [{"kind": "official_docx", "path": "template.docx"}],
                "structure": {"ordered_roles": [{
                    "role": "cover", "required": True,
                    "selector": {"kind": "paragraph", "text": "固定封面", "match": "exact"},
                }]},
            }, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([
                PY, "scripts/thesis_format_pipeline.py", str(requirements), str(source), str(base / "out.docx"),
                "--work-dir", str(work), "--analysis-mode", "rule_only",
                "--template-profile", str(profile),
                "--compliance-mode", "supported_subset",
            ], cwd=ROOT, text=True, capture_output=True)
            self.assertEqual(result.returncode, 10, result.stderr + result.stdout)
            manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
            plan = json.loads((work / "assembly-plan.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["reason"], "region graph preflight blocked execution")
            self.assertEqual(manifest["assembly_plan_status"], "invalid")
            self.assertEqual(plan["findings"][0]["code"], "assembly.contract_missing")
            self.assertNotIn("style_analysis", [step["name"] for step in manifest["steps"]])

    def test_compiled_region_graph_with_missing_execution_adapter_blocks_before_style_analysis(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            requirements = base / "requirements.docx"
            source = base / "source.docx"
            template = base / "template.docx"
            profile = base / "profile.json"
            work = base / "work"
            rules = Document(); rules.add_paragraph("论文正文使用小四号宋体，行距固定值20磅。"); rules.save(requirements)
            thesis = Document(); thesis.add_paragraph("测试论文").style = "Title"; thesis.save(source)
            official = Document(); official.add_paragraph("固定声明").style = "Title"; official.save(template)
            profile.write_text(json.dumps({
                "schema_version": "1.0", "profile_id": "compiled-but-generator-missing",
                "resources": [{"kind": "official_docx", "path": "template.docx"}],
                "regions": {
                    "graph_id": "generator-required",
                    "nodes": [
                        {"id": "declaration", "kind": "fixed_protected",
                         "selector": {"kind": "paragraph", "text": "固定声明", "match": "exact"}},
                        {"id": "toc", "kind": "generated", "content_role": "toc"},
                    ],
                    "edges": [{"kind": "order", "from": "declaration", "to": "toc"}],
                },
            }, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([
                PY, "scripts/thesis_format_pipeline.py", str(requirements), str(source), str(base / "out.docx"),
                "--work-dir", str(work), "--analysis-mode", "rule_only",
                "--template-profile", str(profile), "--compliance-mode", "supported_subset",
            ], cwd=ROOT, text=True, capture_output=True)
            self.assertEqual(result.returncode, 11, result.stderr + result.stdout)
            manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
            report = json.loads((work / "assembly-execution.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["reason"], "assembly plan execution blocked")
            self.assertEqual(manifest["assembly_plan_status"], "compiled")
            self.assertEqual(manifest["assembly_execution_status"], "blocked")
            self.assertEqual(report["findings"][0]["code"], "assembly.execution_blocked")
            names = [step["name"] for step in manifest["steps"]]
            self.assertIn("assembly_execution", names)
            self.assertNotIn("style_analysis", names)

    def test_executable_region_graph_feeds_assembled_source_to_style_application(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            requirements = base / "requirements.docx"
            source = base / "source.docx"
            template = base / "template.docx"
            profile = base / "profile.json"
            output = base / "out.docx"
            work = base / "work"
            rules = Document(); rules.add_paragraph("论文正文使用小四号宋体，行距固定值20磅。"); rules.save(requirements)
            thesis = Document(); thesis.add_paragraph("测试论文").style = "Title"; thesis.add_paragraph("装配正文内容。"); thesis.save(source)
            official = Document(); official.add_paragraph("固定声明").style = "Title"; official.add_paragraph("正文占位符"); official.save(template)
            template_sha = hashlib.sha256(template.read_bytes()).hexdigest()
            profile.write_text(json.dumps({
                "schema_version": "1.0", "profile_id": "assembly-success",
                "authority": {"organization": "test", "source_url": "https://example.test",
                              "effective_version": "1"},
                "resources": [{"id": "official", "kind": "official_docx", "path": "template.docx",
                               "sha256": template_sha}],
                "structure": {"ordered_roles": [{"role": "declaration", "required": True,
                    "selector": {"kind": "paragraph", "text": "固定声明", "match": "exact"}}]},
                "render_rules": {},
                "regions": {
                    "graph_id": "one-dynamic-body",
                    "nodes": [
                        {"id": "declaration", "kind": "fixed_protected",
                         "selector": {"kind": "paragraph", "text": "固定声明", "match": "exact"}},
                        {"id": "body", "kind": "dynamic_content", "content_role": "body",
                         "selector": {"kind": "paragraph", "text": "正文占位符", "match": "exact"}},
                    ],
                    "edges": [{"kind": "order", "from": "declaration", "to": "body"}],
                },
            }, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([
                PY, "scripts/thesis_format_pipeline.py", str(requirements), str(source), str(output),
                "--work-dir", str(work), "--analysis-mode", "rule_only",
                "--template-profile", str(profile), "--compliance-mode", "supported_subset",
            ], cwd=ROOT, text=True, capture_output=True)
            # The synthetic fixture deliberately has no Word render evidence,
            # so the independent final submission audit remains fail-closed.
            # Assembly and format application must nevertheless be reported as
            # completed; a submission-gate failure is not a structure failure.
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
            assembled = work / "assembled-source.docx"
            self.assertEqual(manifest["status"], "completed")
            self.assertTrue(manifest["official_template_structure_valid"])
            self.assertFalse(manifest["official_template_submission_ready"])
            self.assertFalse(manifest["submission_ready"])
            self.assertEqual(manifest["assembly_execution_status"], "assembled")
            self.assertTrue(assembled.is_file())
            assembled_text = "\n".join(paragraph.text for paragraph in Document(assembled).paragraphs)
            self.assertIn("固定声明", assembled_text)
            self.assertIn("装配正文内容。", assembled_text)
            self.assertNotIn("正文占位符", assembled_text)
            apply_step = next(step for step in manifest["steps"] if step["name"] == "apply_and_validate")
            self.assertEqual(Path(apply_step["command"][2]).resolve(), assembled.resolve())
            self.assertIn("official_template_audit", [step["name"] for step in manifest["steps"]])

    def test_multi_role_pipeline_blocks_before_assembly_when_academic_outputs_is_missing(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            requirements = base / "requirements.docx"
            source = base / "source.docx"
            template = base / "template.docx"
            profile = base / "profile.json"
            output = base / "out.docx"
            work = base / "work"
            rules = Document()
            rules.add_paragraph("论文正文使用小四号宋体，行距固定值20磅。")
            rules.save(requirements)
            thesis = Document()
            thesis.add_heading("1 引言", level=1)
            thesis.add_paragraph("正文内容。")
            thesis.add_heading("致谢", level=1)
            thesis.add_paragraph("致谢内容。")
            thesis.add_heading("参考文献", level=1)
            thesis.add_paragraph("参考文献内容。")
            thesis.save(source)
            official = Document()
            for heading, placeholder in [
                ("BODY", "old body"),
                ("ACKNOWLEDGMENTS", "old acknowledgments"),
                ("REFERENCES", "old references"),
                ("APPENDIX", "old appendix"),
                ("OUTPUTS", "old outputs"),
            ]:
                official.add_paragraph(heading)
                official.add_paragraph(placeholder)
            official.add_paragraph("END")
            official.save(template)
            template_sha = hashlib.sha256(template.read_bytes()).hexdigest()
            roles = ["body", "acknowledgments", "references", "appendices", "academic_outputs"]
            boundaries = ["BODY", "ACKNOWLEDGMENTS", "REFERENCES", "APPENDIX", "OUTPUTS", "END"]
            nodes = []
            for role, start, end in zip(roles, boundaries, boundaries[1:]):
                node = {
                    "id": role,
                    "kind": "optional" if role == "appendices" else "dynamic_content",
                    "content_role": role,
                    "start_selector": {"kind": "paragraph", "text": start, "match": "exact"},
                    "end_selector": {"kind": "paragraph", "text": end, "match": "exact"},
                    "range_policy": "replace_between",
                }
                if role == "appendices":
                    node["condition"] = "source_role.appendices.present"
                nodes.append(node)
            profile.write_text(json.dumps({
                "schema_version": "1.0", "profile_id": "missing-academic-outputs",
                "authority": {"organization": "test", "source_url": "https://example.test",
                              "effective_version": "1"},
                "resources": [{"id": "official", "kind": "official_docx", "path": "template.docx",
                               "sha256": template_sha}],
                "structure": {"ordered_roles": []},
                "render_rules": {},
                "regions": {
                    "graph_id": "missing-academic-outputs-regions",
                    "source_section_policy": "discard",
                    "nodes": nodes,
                    "edges": [{"kind": "order", "from": left, "to": right}
                              for left, right in zip(roles, roles[1:])],
                },
            }, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([
                PY, "scripts/thesis_format_pipeline.py", str(requirements), str(source), str(output),
                "--work-dir", str(work), "--analysis-mode", "rule_only",
                "--template-profile", str(profile), "--compliance-mode", "supported_subset",
            ], cwd=ROOT, text=True, capture_output=True)
            self.assertEqual(result.returncode, 11, result.stderr + result.stdout)
            manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
            role_map = json.loads((work / "source-role-map.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["reason"], "source role extraction blocked assembly")
            self.assertEqual(manifest["source_role_map_status"], "blocked")
            self.assertEqual(
                role_map["required_roles"],
                ["academic_outputs", "acknowledgments", "body", "references"],
            )
            self.assertIn(
                {"code": "source_role.required_missing", "role": "academic_outputs",
                 "message": "required source role is missing: academic_outputs"},
                role_map["findings"],
            )
            self.assertFalse((work / "assembled-source.docx").exists())
            self.assertFalse(output.exists())
            names = [step["name"] for step in manifest["steps"]]
            self.assertIn("source_role_extraction", names)
            self.assertNotIn("assembly_execution", names)
            self.assertNotIn("style_analysis", names)
            extraction_step = next(
                step for step in manifest["steps"] if step["name"] == "source_role_extraction"
            )
            required_arguments = [
                extraction_step["command"][index + 1]
                for index, value in enumerate(extraction_step["command"][:-1])
                if value == "--required-role"
            ]
            self.assertNotIn("appendices", required_arguments)

    def test_cli_writes_report_and_uses_blocking_exit_code(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); spec = td / "spec.json"; out = td / "report.json"
            spec.write_text(json.dumps({"requirements": [{"id": "R9", "role": "future_role",
                                                          "properties": {"x": True}}]}), encoding="utf-8")
            result = subprocess.run([PY, "scripts/capability_planner.py", str(spec), "--out", str(out)],
                                    cwd=ROOT, text=True, capture_output=True)
            self.assertEqual(result.returncode, 3, result.stderr + result.stdout)
            self.assertEqual(json.loads(out.read_text())["status"], "blocked")


if __name__ == "__main__":
    unittest.main()
