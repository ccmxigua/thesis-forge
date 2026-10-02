"""Source-bound primary condition proposals and serialized label-only covers."""
import copy
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))
import host_agent_bridge as bridge
import native_semantic_review as native
import requirements_engine as engine
import apply_format_spec as applier
from docx import Document
from semantic_contract import attach_request_provenance, sha256_json
from source_condition_reassessment import (
    condition_feedback, condition_reassessment, condition_proposal_budget_receipt,
    CODE, RULE_ID, RETRY_BUDGET_POLICY,
)


def current_schema_condition_fixture():
    """Synthetic current-schema case; the historical JSON stays untouched.

    The captured incident predates distinct discipline fields. Correct those
    unrelated bindings in this test copy, not in production or the raw file,
    then recompute payload fingerprints. Condition retry authority is still
    tested against the complete current payload, never ordinal-only refs.
    """
    data = json.loads((ROOT / "tests/fixtures/cover-condition-reassessment-incident.json").read_text())
    fingerprint_changes = {}
    for requirement in data["primary_candidate"]["requirements"]:
        payload = {"source_requirement_id": requirement.get("existing_requirement_id") or requirement.get("id"),
                   "role": requirement.get("role"), "properties": copy.deepcopy(requirement.get("properties")),
                   "evidence_ids": copy.deepcopy(requirement.get("evidence_ids") or []),
                   "verification": copy.deepcopy(requirement.get("verification"))}
        original = sha256_json(payload)
        for field in payload["properties"].get("fields", []):
            label = "".join(str(field.get("label", "")).split()).rstrip("：:")
            key = {"一级学科": "first_discipline", "二级学科": "second_discipline"}.get(label)
            if key:
                field.update(id=key, value_from=f"thesis_profile.cover_metadata.{key}")
        fingerprint_changes[original] = sha256_json(payload)
        requirement["properties"] = payload["properties"]
    for requirement in data["primary_raw"]["requirements"]:
        for field in requirement["properties"].get("fields", []):
            label = "".join(str(field.get("label", "")).split()).rstrip("：:")
            key = {"一级学科": "first_discipline", "二级学科": "second_discipline"}.get(label)
            if key:
                field.update(id=key, value_from=f"thesis_profile.cover_metadata.{key}")
    for refs in data["linked_requirement_fingerprints"].values():
        for reference, fingerprint in refs.items():
            assert fingerprint in fingerprint_changes
            refs[reference] = fingerprint_changes[fingerprint]
    return data


def incident():
    data = current_schema_condition_fixture()
    source = data["source"]
    evidence = {"evidence": list(source["evidence_context"].values())}
    chunk = engine.build_llm_request([], source["clauses"], evidence, {}, "full", contract_version="3.0")
    chunk.update(source)
    chunk.update(case_id="condition-offline", batch={"index": 1},
                 runtime_context={"code_fingerprint_sha256": sha256_json("offline-condition-code")})
    chunk = attach_request_provenance(chunk, source_sha256=sha256_json(evidence), evidence_doc=evidence,
                                     clauses=source["clauses"], run_id="condition-offline")
    raw = bridge.normalize_native_response(data["primary_raw"], chunk["response_schema"])
    candidate = bridge.prepare_native_response_candidate(raw, chunk)[0]
    candidate["provenance"] = copy.deepcopy(chunk["provenance"])
    request = native.build_obligation_coverage_request(candidate, chunk, run_id="condition-offline", chunk_index=1)
    request.update(attempt=1, provider_attempt=2)
    # Offline schema regeneration changes selectors. Rebind only when the
    # entire referenced payload (not its ordinal or just its role) matches.
    by_check = {check["check_id"]: check for check in request["checks"]}
    for result in data["independent"]["results"]:
        links = by_check[result["check_id"]]["review_context"]["linked_requirements"]
        fingerprints = data["linked_requirement_fingerprints"][result["check_id"]]
        for atom in result["identified_obligations"]:
            rebound = []
            for reference in atom["requirement_refs"]:
                matching = [link["requirement_ref"] for link in links
                    if sha256_json({k:v for k,v in link.items() if k != "requirement_ref"}) == fingerprints[reference]]
                assert len(matching) == 1
                rebound.append(matching[0])
            atom["requirement_refs"] = rebound
    record = condition_feedback(candidate, chunk, request, data["independent"])
    assert record is not None
    record["response_sha256"] = sha256_json(raw)
    return data, raw, candidate, chunk, record


def proposed(raw, record):
    result = copy.deepcopy(raw)
    for entry in record["condition_atoms"]:
        for review in result["clause_reviews"]:
            if review["clause_id"] == entry["clause_id"]:
                for atom in review["obligations"]:
                    if atom["id"] == entry["obligation_id"]:
                        atom["condition"] = None
        for requirement in result["requirements"]:
            if entry["clause_id"] not in requirement.get("clause_ids", []):
                continue
            props = requirement["properties"]
            for field in props.get("fields", []) + props.get("non_public_administration", {}).get("fields", []):
                if field["id"] == "security_marking":
                    field["label_display_policy"] = "always"
    return result


class ConditionReassessmentTests(unittest.TestCase):
    def proof(self, old, new, chunk, records):
        return condition_reassessment(old, new, records, bridge._retry_change_paths(old, new), chunk,
            prepare=bridge.prepare_native_response_candidate, validate=bridge.validate_host_agent_response)

    def test_captured_raw_and_projected_transitions_are_source_bound(self):
        _, raw, candidate, chunk, record = incident()
        frozen = copy.deepcopy((raw, candidate, chunk, record))
        new = proposed(raw, record)
        proofs = self.proof(raw, new, chunk, [record])
        self.assertEqual(len(proofs), 4)
        authorizations = []
        error, changes = bridge._retry_semantic_change_error(raw, new, [record],
            contract_version="3.0", chunk=chunk, authorization_out=authorizations)
        self.assertIsNone(error)
        self.assertEqual(len(changes), 4)
        self.assertTrue(all(a["rule_id"] == RULE_ID and a["independent_review_required"]
                            and not a["mechanical_equivalence_claimed"] for a in authorizations))
        compiled = bridge.prepare_native_response_candidate(new, chunk)[0]
        compiled["provenance"] = chunk["provenance"]
        self.assertEqual(len(self.proof(candidate, compiled, chunk, [record])), 4)
        self.assertEqual(bridge.validate_host_agent_response(compiled, chunk), [])
        self.assertEqual((raw, candidate, chunk, record), frozen)

    def test_separate_budget_requires_complete_current_feedback_and_raw_parent(self):
        _, raw, candidate, chunk, record = incident()
        frozen = copy.deepcopy((raw, candidate, chunk, record))
        receipt = condition_proposal_budget_receipt(candidate, chunk, [record], raw)
        self.assertEqual(receipt["policy_version"], RETRY_BUDGET_POLICY)
        self.assertEqual(receipt["proposal_limit"], 1)
        self.assertFalse(receipt["submission_ready"])
        self.assertEqual((raw, candidate, chunk, record), frozen)
        for defect in ("code_only", "run", "source", "candidate", "parent", "request", "review", "extra"):
            damaged = copy.deepcopy(record); records = [damaged]
            if defect == "code_only": records = [{"code": CODE}]
            elif defect == "run": damaged["run_id"] = "old-run"
            elif defect == "source": damaged["source_chunk_sha256"] = "0" * 64
            elif defect == "candidate": damaged["candidate_response_sha256"] = "0" * 64
            elif defect == "parent": damaged["response_sha256"] = "0" * 64
            elif defect == "request": damaged["review_request"]["run_id"] = "old-run"
            elif defect == "review": damaged["rejected_review"]["results"].pop()
            else: records.append({"code": "unrelated_error"})
            with self.subTest(defect=defect):
                self.assertIsNone(condition_proposal_budget_receipt(candidate, chunk, records, raw))

    def _late_condition_budget_case(self, *, outcome="accepted", ordinary_limit=2, early_rejection=True):
        """Offline bridge orchestration, not a real provider/Word acceptance."""
        data, raw, _, _, _ = incident()
        original_data = current_schema_condition_fixture()
        source = data["source"]
        evidence = {"evidence": list(source["evidence_context"].values())}
        with tempfile.TemporaryDirectory() as td:
            directory = Path(td) / "packet"; directory.mkdir()
            request = engine.build_llm_request([], source["clauses"], evidence, {}, "full", contract_version="3.0")
            request.update(case_id="budget-offline", batch={"index": 1},
                           runtime_context={"code_fingerprint_sha256": sha256_json("budget-offline-code")})
            request = attach_request_provenance(request, source_sha256=sha256_json(evidence), evidence_doc=evidence,
                clauses=source["clauses"], run_id="budget-offline")
            engine.prepare_host_agent_review_packets(request, source["clauses"], evidence, sha256_json(evidence),
                directory, chunk_size=100)
            chunk = json.loads((directory / "llm-request-chunks.json").read_text())[0]
            captured = {}; reviews = []; primary_calls = []

            def independent(candidate, current_chunk, **kwargs):
                reviews.append(kwargs["attempt"])
                captured["controller"] = kwargs["controller"]
                if "record" not in captured:
                    review_request = native.build_obligation_coverage_request(candidate, current_chunk,
                        run_id=kwargs["run_id"], chunk_index=kwargs["chunk_index"])
                    review_request.update(attempt=kwargs["attempt"], provider_attempt=2)
                    review = copy.deepcopy(original_data["independent"])
                    checks = {c["check_id"]: c for c in review_request["checks"]}
                    for result in review["results"]:
                        links = checks[result["check_id"]]["review_context"]["linked_requirements"]
                        for atom in result["identified_obligations"]:
                            rebound = []
                            for ref in atom["requirement_refs"]:
                                fingerprint = original_data["linked_requirement_fingerprints"][result["check_id"]][ref]
                                matches = [link["requirement_ref"] for link in links
                                    if sha256_json({k:v for k,v in link.items() if k != "requirement_ref"}) == fingerprint]
                                self.assertEqual(len(matches), 1)
                                rebound.append(matches[0])
                            atom["requirement_refs"] = rebound
                    record = condition_feedback(candidate, current_chunk, review_request, review)
                    self.assertIsNotNone(record)
                    captured["record"] = copy.deepcopy(record)
                    if outcome == "stale_feedback": record["source_chunk_sha256"] = "0" * 64
                    error = bridge.IndependentObligationReviewError("source-bound condition proposal needed")
                    error.retryable = True; error.error_records = [record]
                    raise error
                self.assertEqual(bridge.validate_host_agent_response(candidate, current_chunk), [])
                if outcome == "repeat_request":
                    error = bridge.IndependentObligationReviewError("another condition request must not reset the budget")
                    error.retryable = True; error.error_records = [{"code": CODE}]
                    raise error
                if outcome == "fresh_rejection":
                    error = bridge.IndependentObligationReviewError("fresh reviewer still rejects the candidate")
                    error.retryable = False
                    raise error
                from test_host_agent_bridge import HostAgentBridgeTests
                return HostAgentBridgeTests._fake_independent_review(candidate, current_chunk, **kwargs)

            def primary(_command, **kwargs):
                primary_calls.append(len(primary_calls) + 1)
                body = copy.deepcopy(raw) if "record" not in captured else proposed(raw, captured["record"])
                body["provenance"] = chunk["provenance"]
                if outcome == "unrelated_edit" and "record" in captured:
                    body["clause_reviews"][0]["reason"] += " unauthorized edit"
                text = ("not valid JSON" if early_rejection and len(primary_calls) == 1 else json.dumps(body))
                return subprocess.CompletedProcess(["mock-host"], 0, json.dumps({"runId": "offline-budget",
                    "status": "ok", "provider": "openai", "model": "gpt-5.6-luna",
                    "result": {"payloads": [{"text": text}]}}), "")

            actual_budget = bridge.condition_proposal_budget_receipt
            def reserve(*args):
                receipt = actual_budget(*args)
                if outcome == "cancel_after_reservation" and receipt is not None:
                    captured["controller"].request_stop("cancelled after reservation")
                return receipt

            output = Path(td) / "merged.json"
            with patch.dict(os.environ, {"THESIS_FORGE_HOST_RUNTIME": "openclaw"}), \
                 patch.object(bridge, "_run_command", side_effect=primary), \
                 patch.object(bridge, "condition_proposal_budget_receipt", side_effect=reserve), \
                 patch.object(bridge, "_run_independent_obligation_coverage_review", side_effect=independent):
                if outcome == "accepted":
                    bridge.run_bridge(directory, response_out=output, agent_id="main", timeout=1,
                        max_attempts=ordinary_limit, openclaw_bin="mock-host", model="openai/gpt-5.6-luna")
                else:
                    with self.assertRaises((ValueError, bridge.HostAgentCancelled,
                                            bridge.IndependentObligationReviewError)):
                        bridge.run_bridge(directory, response_out=output, agent_id="main", timeout=1,
                            max_attempts=ordinary_limit, openclaw_bin="mock-host", model="openai/gpt-5.6-luna")
            audit = json.loads((directory / "host-agent-run.json").read_text())
            self.assertEqual(output.exists(), outcome == "accepted")
            return audit, primary_calls, reviews

    def test_first_condition_proposal_survives_exhausted_ordinary_budget(self):
        audit, calls, reviews = self._late_condition_budget_case()
        self.assertEqual(calls, [1, 2, 3])
        self.assertEqual(reviews, [2, 3])
        self.assertEqual(audit["status"], "merged")
        budget = audit["chunk_runs"][0]["retry_budget"]
        self.assertEqual(budget["ordinary_attempts_started"], 2)
        self.assertEqual(budget["condition_proposals_started"], 1)
        attempts = audit["chunk_lifecycle"][0]["attempts"]
        self.assertIsNotNone(attempts[1]["condition_proposal_authorization"])
        self.assertEqual(attempts[2]["retry_budget"]["current_attempt_kind"], "condition_proposal")
        self.assertIsNotNone(attempts[2]["condition_proposal_reservation"])

    def test_bad_feedback_and_cancel_cannot_dispatch_extra_primary(self):
        for outcome in ("stale_feedback", "cancel_after_reservation"):
            with self.subTest(outcome=outcome):
                audit, calls, reviews = self._late_condition_budget_case(outcome=outcome)
                self.assertEqual(calls, [1, 2])
                self.assertEqual(reviews, [2])
                self.assertFalse(audit["merged_response_written"])

    def test_proposal_never_bypasses_fresh_review_or_retries_forever(self):
        for outcome in ("repeat_request", "fresh_rejection", "unrelated_edit"):
            with self.subTest(outcome=outcome):
                audit, calls, reviews = self._late_condition_budget_case(outcome=outcome)
                self.assertEqual(calls, [1, 2, 3])
                self.assertEqual(reviews, [2] if outcome == "unrelated_edit" else [2, 3])
                self.assertFalse(audit["merged_response_written"])
        # An early condition proposal is still one-shot, not followed by
        # unused ordinary slots or a second condition proposal.
        _, calls, reviews = self._late_condition_budget_case(
            outcome="repeat_request", ordinary_limit=3, early_rejection=False)
        self.assertEqual(calls, [1, 2])
        self.assertEqual(reviews, [1, 2])

    def test_not_a_reviewer_copy_or_a_pass(self):
        _, raw, candidate, chunk, record = incident()
        new = proposed(raw, record)
        # The primary may disagree again; authorization is only a proposal.
        review = next(r for r in new["clause_reviews"] if r["clause_id"] == record["condition_atoms"][0]["clause_id"])
        review["obligations"][0]["condition"] = "a newly proposed condition"
        self.assertIsNotNone(self.proof(raw, new, chunk, [record]))
        with self.assertRaises(native.TypedSourceAtomAlignmentError):
            native.validate_obligation_coverage_response(record["rejected_review"], record["review_request"]["checks"])

    def test_all_unrelated_semantics_relations_and_values_are_frozen(self):
        for change in ("actor", "action", "target", "force", "route", "status", "quote", "reason",
                       "remove_atom", "remove_requirement", "foreign_condition", "value_policy", "field_value", "field_order"):
            _, raw, _, chunk, record = incident()
            new = proposed(raw, record)
            review = next(r for r in new["clause_reviews"] if r["clause_id"] == record["condition_atoms"][0]["clause_id"])
            atom = review["obligations"][0]
            if change in {"actor", "action", "target", "force", "route", "status"}: atom[change] = "changed"
            elif change == "quote": atom["source_quote"] = "invented"
            elif change == "reason": review["reason"] += " repaired"
            elif change == "remove_atom": review["obligations"] = []
            elif change == "remove_requirement": new["requirements"].pop()
            elif change == "foreign_condition": new["clause_reviews"][0]["obligations"][0]["condition"] = "unrelated"
            else:
                field = new["requirements"][0]["properties"]["fields"][0]
                field[{"value_policy": "display_policy", "field_value": "value_from", "field_order": "order"}[change]] = "changed"
            with self.subTest(change=change): self.assertIsNone(self.proof(raw, new, chunk, [record]))

    def test_stale_or_incomplete_feedback_and_source_are_rejected(self):
        for change in ("run", "source", "candidate", "parent", "request", "review", "unknown_atom", "extra_record"):
            _, raw, _, chunk, record = incident()
            new = proposed(raw, record); records = [record]
            if change == "run": record["run_id"] = "old-run"
            elif change == "source": chunk["clauses"][0]["source_span"]["source_sha256"] = "0" * 64
            elif change == "candidate": record["candidate_response_sha256"] = "0" * 64
            elif change == "parent": record["response_sha256"] = "0" * 64
            elif change == "request": record["review_request"]["run_id"] = "old-run"
            elif change == "review": record["rejected_review"]["results"].pop()
            elif change == "unknown_atom": record["condition_atoms"][0]["obligation_id"] = "foreign"
            else: records.append({"code": "another_error"})
            with self.subTest(change=change): self.assertIsNone(self.proof(raw, new, chunk, records))

    def test_condition_disagreement_cannot_hide_other_review_errors(self):
        for defect in ("foreign_ref", "empty_ref", "foreign_quote", "wrong_disposition", "missing_machine_fact"):
            data, _, candidate, chunk, record = incident()
            rejected = copy.deepcopy(data["independent"])
            target = next(r for r in rejected["results"] if r["check_id"] == record["condition_atoms"][0]["clause_id"])
            atom = target["identified_obligations"][0]
            if defect == "foreign_ref": atom["requirement_refs"] = ["RR-foreign"]
            elif defect == "empty_ref": atom["requirement_refs"] = []
            elif defect == "foreign_quote": atom["source_quote"] = "not from this source"
            elif defect == "wrong_disposition": atom["disposition"] = "external_action_pending"
            else: target["machine_obligation_ids"] = ["foreign-fact"]
            with self.subTest(defect=defect):
                self.assertIsNone(condition_feedback(candidate, chunk, record["review_request"], rejected))

    def test_real_independent_exhaustion_routes_only_condition_disagreements(self):
        data, _, candidate, chunk, _ = incident()
        calls = []
        def reviewer(request, **kwargs):
            calls.append(request)
            # Rejected canonical payload only. This double cannot publish a
            # successful invocation/ledger or claim real provider freshness.
            result = copy.deepcopy(data["independent"])
            kwargs["output_dir"].mkdir(parents=True)
            bridge._write_json(kwargs["output_dir"] / "compiled-response.json", result)
            native.validate_obligation_coverage_response(result, request["checks"])
            raise AssertionError("the captured candidate must remain rejected")
        with tempfile.TemporaryDirectory() as td, patch.object(bridge, "run_native_semantic_review", side_effect=reviewer), patch.object(bridge.time, "sleep"):
            with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                bridge._run_independent_obligation_coverage_review(candidate, chunk, review_dir=Path(td),
                    run_id="condition-offline", chunk_index=1, attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                    timeout=5, agent_id="main", runner="exec", binary="codex", config_path=None,
                    controller=bridge.RunController())
            self.assertEqual(len(calls), 2)
            self.assertTrue(caught.exception.retryable)
            self.assertEqual(caught.exception.error_records[0]["code"], CODE)
            self.assertEqual(list(Path(td).rglob("obligation-analysis-ledger.json")), [])

    def test_dynamic_identifiers_do_not_select_a_school_or_clause_exception(self):
        _, _, candidate, chunk, original = incident()
        replacements = {c["id"]: f"other-school-clause-{i}" for i,c in enumerate(chunk["clauses"])}
        replacements.update({eid: f"other-evidence-{i}" for i,eid in enumerate(chunk["evidence_context"])})
        def rename(value):
            if isinstance(value, dict): return {replacements.get(k,k):rename(v) for k,v in value.items()}
            if isinstance(value, list): return [rename(v) for v in value]
            return replacements.get(value,value) if isinstance(value,str) else value
        candidate, chunk = rename(candidate), rename(chunk)
        # The synthetic input is a new extraction: its code-owned row hashes
        # must include the renamed evidence identities, not the prior run's.
        for item in chunk["evidence_context"].values():
            row = item.get("table_row_context")
            if isinstance(row, dict):
                row["row_sha256"] = sha256_json({k:v for k,v in row.items() if k != "row_sha256"})
        evidence = {"evidence": list(chunk["evidence_context"].values())}
        chunk = attach_request_provenance(chunk, source_sha256=sha256_json(evidence), evidence_doc=evidence,
            clauses=chunk["clauses"], run_id="other-school-run")
        candidate["provenance"] = chunk["provenance"]
        request = native.build_obligation_coverage_request(candidate, chunk, run_id="other-school-run", chunk_index=1)
        request.update(attempt=1, provider_attempt=2)
        disputed = {(replacements[item["clause_id"]], item["obligation_id"]) for item in original["condition_atoms"]}
        response = {"results": []}
        for check in request["checks"]:
            context = check["review_context"]
            # Synthetic alignment test, not provider or natural-language proof.
            atoms = []
            mismatch = False
            for primary in context["primary_obligations"]:
                atom = {k:copy.deepcopy(primary[k]) for k in
                    ("actor", "action", "target", "source_quote", "force", "applicability", "condition") if k in primary}
                if (check["check_id"], primary["id"]) in disputed:
                    atom.pop("condition", None); mismatch = True
                atoms.append({**atom, "primary_obligation_id": primary["id"], "obligation_summary": "synthetic fixture",
                    "requirement_refs": [r["requirement_ref"] for r in context["linked_requirements"]],
                    "disposition": "represented"})
            response["results"].append({"check_id": check["check_id"],
                "verdict": "incomplete" if mismatch else "consistent", "rationale": "synthetic fixture",
                "identified_obligations": atoms, "evidence_quotes": [check["document_text"]],
                "machine_obligation_ids": context["machine_obligation_ids"]})
        record = condition_feedback(candidate, chunk, request, response)
        self.assertIsNotNone(record)
        new = proposed(candidate, record)
        self.assertEqual(len(self.proof(candidate, new, chunk, [record])), 4)

    def test_empty_optional_value_label_is_visible_after_serialization_and_audited(self):
        cover = {"institution": "示例学校", "fields": [{"id": "program_name", "label": "专业：",
            "value_from": "thesis_profile.cover_metadata.program_name", "display_policy": "if_present",
            "label_display_policy": "always", "order": 1}]}
        doc = Document(); doc.add_paragraph("正文")
        contract = applier.compile_cover_contract(cover, {})
        report = applier.apply_cover(doc, cover, {}, contract)
        self.assertEqual(report["label_only_fields_written"], 1)
        self.assertEqual(report["trusted_fields_written"], 0)
        self.assertEqual(report["metadata_pending_fields"], [])
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "cover.docx"; doc.save(path); saved = Document(path)
            self.assertIn("专业：", [p.text for p in saved.paragraphs])
            self.assertEqual(applier.audit_cover(saved, cover, {}), [])
            requirements = [{"id": "cover-rule", "role": "cover", "properties": {"fields": cover["fields"]}}]
            spec = {"cover": cover, "requirements": requirements}
            actual, _ = applier._receipt_semantic_actuals(saved, spec, {}, requirements, {}, [], contract)
            self.assertEqual(actual["cover"]["fields"], cover["fields"])
            next(p for p in saved.paragraphs if p.text == "专业：").text = ""
            self.assertTrue(applier.audit_cover(saved, cover, {}))
            actual, _ = applier._receipt_semantic_actuals(saved, spec, {}, requirements, {}, [], contract)
            self.assertNotIn("fields", actual.get("cover", {}))
        # No declaration means exact legacy behavior, not a global rule.
        del cover["fields"][0]["label_display_policy"]
        legacy = Document(); legacy.add_paragraph("正文")
        applier.apply_cover(legacy, cover, {})
        self.assertNotIn("专业：", [p.text for p in legacy.paragraphs])

    def test_bridge_primary_retry_still_needs_a_new_independent_review(self):
        data, raw, _, _, _ = incident()
        source = data["source"]
        evidence = {"evidence": list(source["evidence_context"].values())}
        original_data = current_schema_condition_fixture()
        with tempfile.TemporaryDirectory() as td:
            directory = Path(td) / "packet"; directory.mkdir()
            request = engine.build_llm_request([], source["clauses"], evidence, {}, "full", contract_version="3.0")
            request.update(case_id="condition-orchestration", batch={"index": 1},
                           runtime_context={"code_fingerprint_sha256": sha256_json("offline-code")})
            request = attach_request_provenance(request, source_sha256=sha256_json(evidence), evidence_doc=evidence,
                clauses=source["clauses"], run_id="condition-orchestration")
            engine.prepare_host_agent_review_packets(request, source["clauses"], evidence, sha256_json(evidence),
                directory, chunk_size=100)
            chunks = json.loads((directory / "llm-request-chunks.json").read_text())
            self.assertEqual(len(chunks), 1)
            calls = []
            captured = {}
            def independent(candidate, chunk, **kwargs):
                calls.append(kwargs["attempt"])
                if kwargs["attempt"] == 1:
                    review_request = native.build_obligation_coverage_request(candidate, chunk,
                        run_id=kwargs["run_id"], chunk_index=kwargs["chunk_index"])
                    review_request.update(attempt=1, provider_attempt=2)
                    review = copy.deepcopy(original_data["independent"])
                    checks = {check["check_id"]: check for check in review_request["checks"]}
                    for result in review["results"]:
                        links = checks[result["check_id"]]["review_context"]["linked_requirements"]
                        for atom in result["identified_obligations"]:
                            rebound = []
                            for ref in atom["requirement_refs"]:
                                fingerprint = original_data["linked_requirement_fingerprints"][result["check_id"]][ref]
                                matching = [link["requirement_ref"] for link in links
                                    if sha256_json({k:v for k,v in link.items() if k != "requirement_ref"}) == fingerprint]
                                self.assertEqual(len(matching), 1)
                                rebound.append(matching[0])
                            atom["requirement_refs"] = rebound
                    record = condition_feedback(candidate, chunk, review_request, review)
                    self.assertIsNotNone(record)
                    captured["record"] = record
                    error = bridge.IndependentObligationReviewError("fresh primary condition proposal required")
                    error.retryable = True; error.error_records = [record]
                    raise error
                self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
                self.assertTrue(all(r["properties"]["non_public_administration"]["fields"][0]["label_display_policy"] == "always"
                    for r in candidate["requirements"] if r.get("field_key") in {"cover_outer", "cover_inner"}))
                error = bridge.IndependentObligationReviewError("fresh review still rejects unsupported meaning")
                error.retryable = False
                raise error
            def primary(_command, **kwargs):
                body = copy.deepcopy(raw) if not captured else proposed(raw, captured["record"])
                body["provenance"] = chunks[0]["provenance"]
                return subprocess.CompletedProcess(["mock-native-host"], 0, json.dumps({"runId": "offline-primary",
                    "status": "ok", "provider": "openai", "model": "gpt-5.6-luna",
                    "result": {"payloads": [{"text": json.dumps(body)}]}}), "")
            output = Path(td) / "merged.json"
            with patch.dict(os.environ, {"THESIS_FORGE_HOST_RUNTIME": "openclaw"}), \
                 patch.object(bridge, "_run_command", side_effect=primary) as producer, \
                 patch.object(bridge, "_run_independent_obligation_coverage_review", side_effect=independent):
                with self.assertRaises(bridge.IndependentObligationReviewError):
                    bridge.run_bridge(directory, response_out=output, agent_id="main", timeout=1,
                        max_attempts=2, openclaw_bin="mock-native-host", model="openai/gpt-5.6-luna")
            self.assertEqual(producer.call_count, 2)
            self.assertEqual(calls, [1, 2])
            self.assertFalse(output.exists())
            self.assertEqual(json.loads((directory / "host-agent-run.json").read_text())["status"], "failed")

    def test_administrative_label_never_claims_approval_or_trusted_value(self):
        cover = {"institution": "示例学校", "fields": [], "non_public_administration": {
            "applicability": {"status": "conditional", "conditions": []}, "public_policy": "blank",
            "source_region": "cover", "fields": [{"id": "security_marking", "label": "密  级：",
            "value_from": "thesis_profile.cover_metadata.security_marking", "display_policy": "if_present",
            "label_display_policy": "always", "order": 1}]}}
        for profile in ({"security_level": "public"}, {"security_level": "classified"}):
            contract = applier.compile_cover_contract(cover, profile)
            doc = Document(); doc.add_paragraph("正文")
            report = applier.apply_cover(doc, cover, profile, contract)
            self.assertIn("密  级：", [p.text for p in doc.paragraphs])
            self.assertEqual(report["trusted_fields_written"], 0)
            self.assertFalse(contract["non_public_administration"]["approval_status_verified"])
            with tempfile.TemporaryDirectory() as td:
                path = Path(td) / "admin.docx"; doc.save(path); saved = Document(path)
                self.assertFalse(any(f.get("property") == "fields.security_marking"
                    for f in applier.audit_cover(saved, cover, profile)))
                next(p for p in saved.paragraphs if p.text == "密  级：").text = ""
                self.assertTrue(any(f.get("property") == "fields.security_marking"
                    for f in applier.audit_cover(saved, cover, profile)))
            if profile["security_level"] == "classified":
                self.assertTrue(applier.audit_cover(doc, cover, profile))


if __name__ == "__main__": unittest.main()
