from __future__ import annotations

import sys
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from apply_format_spec import (  # noqa: E402
    audit_content_constraints,
    audit_nonblocking_guidance,
    manual_review_receipt_items,
    manual_review_validation_items,
    normalize_keyword_separators,
    semantic_content_review_findings,
    semantic_content_review_items,
)
from format_spec_validation import load_and_validate  # noqa: E402


class ContentConstraintMetricTests(unittest.TestCase):
    def _doc(self, abstract: str, keywords: str) -> Document:
        document = Document()
        abstract_style = document.styles.add_style("AbstractBodyCN", 1)
        keyword_style = document.styles.add_style("KeywordsFixture", 1)
        document.add_paragraph(abstract, abstract_style)
        document.add_paragraph(keywords, keyword_style)
        return document

    def test_chinese_bounds_use_declared_codepoint_metric(self) -> None:
        document = self._doc("甲" * 299, "关键词：甲")
        findings = audit_content_constraints(
            document,
            {
                "abstract_zh": {"required": True, "min_chars": 300, "max_chars": 1000,
                                "length_metric": "unicode_codepoints"},
                "keywords_zh": {"required": True},
            },
            {"abstract_body_zh": {"style_name": "AbstractBodyCN"}, "keywords_zh": {"style_name": "KeywordsFixture"}},
        )
        self.assertTrue(any(item["property"] == "abstract_zh.min_chars" for item in findings), findings)

    def test_general_length_and_keyword_guidance_is_reported_but_never_blocks(self) -> None:
        document = self._doc("短文", "关键词：甲；乙")
        constraints = {
            "abstract_zh": {"length_guidance": {
                "min_chars": 3, "max_chars": 8,
                "length_metric": "cjk_characters",
                "strength": "general_guidance",
                "exception_text": "必要时可略多",
            }},
            "keywords_zh": {"count_guidance": {
                "min_count": 3, "max_count": 8,
                "strength": "general_guidance",
            }},
        }
        mappings = {
            "abstract_body_zh": {"style_name": "AbstractBodyCN"},
            "keywords_zh": {"style_name": "KeywordsFixture"},
        }

        advisories = audit_nonblocking_guidance(document, constraints, mappings)
        findings = audit_content_constraints(document, constraints, mappings)

        self.assertEqual(len(advisories), 2)
        self.assertEqual([item["actual"] for item in advisories], [2, 2])
        self.assertTrue(all(item["policy"] == "non_blocking_general_guidance" for item in advisories))
        self.assertTrue(all(item["within_guidance"] is False for item in advisories))
        self.assertEqual(findings, [])

    def test_keyword_item_limit_reports_item_without_truncation(self) -> None:
        document = self._doc("摘要", "关键词：甲乙丙丁戊己庚辛")
        findings = audit_content_constraints(
            document,
            {
                "abstract_zh": {"required": True},
                "keywords_zh": {"required": True, "max_item_chars": 7,
                                 "item_length_metric": "cjk_characters"},
            },
            {"abstract_body_zh": {"style_name": "AbstractBodyCN"}, "keywords_zh": {"style_name": "KeywordsFixture"}},
        )
        violation = next(item for item in findings if "max_item_chars" in item["property"])
        self.assertEqual(violation["template_value"], 8)
        self.assertEqual(violation["item"], "甲乙丙丁戊己庚辛")

    def test_semicolon_separator_is_normalized_without_changing_keyword_text(self) -> None:
        document = Document()
        keyword_style = document.styles.add_style("KeywordsRepairFixture", 1)
        paragraph = document.add_paragraph("关键词：甲词，乙词; 丙词", keyword_style)
        before_values = ["甲词", "乙词", "丙词"]

        repairs = normalize_keyword_separators(
            document,
            {"keywords_zh": {"separator": "semicolon"}},
            {"keywords_zh": {"style_name": "KeywordsRepairFixture"}},
        )

        self.assertEqual(paragraph.text, "关键词：甲词；乙词； 丙词")
        self.assertEqual(repairs[0]["before"], "关键词：甲词，乙词; 丙词")
        self.assertEqual(repairs[0]["keyword_values_preserved"], True)
        audit = audit_content_constraints(
            document,
            {"keywords_zh": {"separator": "semicolon"}},
            {"keywords_zh": {"style_name": "KeywordsRepairFixture"}},
        )
        self.assertFalse(any(item["property"] == "keywords_zh.separator" for item in audit), audit)
        self.assertEqual(
            [item.strip() for item in paragraph.text.removeprefix("关键词：").split("；")],
            before_values,
        )

    def test_separator_normalization_preserves_run_formatting(self) -> None:
        document = Document()
        keyword_style = document.styles.add_style("KeywordsRunFixture", 1)
        paragraph = document.add_paragraph(style=keyword_style)
        first = paragraph.add_run("关键词：甲词，")
        first.bold = True
        second = paragraph.add_run("乙词; 丙词")
        second.italic = True

        repairs = normalize_keyword_separators(
            document,
            {"keywords_zh": {"separator": "semicolon"}},
            {"keywords_zh": {"style_name": "KeywordsRunFixture"}},
        )

        self.assertEqual(paragraph.text, "关键词：甲词；乙词； 丙词")
        self.assertTrue(first.bold)
        self.assertTrue(second.italic)
        self.assertEqual(len(repairs), 1)

    def test_separator_normalization_leaves_special_run_content_unchanged(self) -> None:
        document = Document()
        keyword_style = document.styles.add_style("KeywordsBreakFixture", 1)
        paragraph = document.add_paragraph(style=keyword_style)
        run = paragraph.add_run("关键词：甲词，乙词")
        run.add_break()
        run.add_text("丙词")

        repairs = normalize_keyword_separators(
            document,
            {"keywords_zh": {"separator": "semicolon"}},
            {"keywords_zh": {"style_name": "KeywordsBreakFixture"}},
        )

        self.assertEqual(repairs, [])
        self.assertIn("甲词，乙词", paragraph.text)
        self.assertIn("\n", paragraph.text)

    def test_english_keyword_cjk_character_metric_is_computed_deterministically(self) -> None:
        document = self._doc("摘要", "Keywords: 交通 flow; deep learning")
        findings = audit_content_constraints(
            document,
            {
                "abstract_zh": {"required": True},
                "keywords_en": {
                    "required": True,
                    "max_item_chars": 1,
                    "item_length_metric": "cjk_characters",
                },
            },
            {
                "abstract_body_zh": {"style_name": "AbstractBodyCN"},
                "keywords_en": {"style_name": "KeywordsFixture"},
            },
        )
        violation = next(item for item in findings
                         if item.get("property") == "keywords_en.max_item_chars[0]")
        self.assertEqual(violation["template_value"], 2)
        self.assertEqual(violation["metric"], "cjk_characters")
        self.assertFalse(any(item.get("verification") == "manual" for item in findings))

    def test_chinese_keyword_count_matching_is_checked_when_declared_on_chinese_rule(self) -> None:
        document = self._doc("摘要", "关键词：甲；乙")
        document.paragraphs[1].style = "KeywordsFixture"
        document.add_paragraph("Keywords: alpha", "KeywordsFixture")
        findings = audit_content_constraints(
            document,
            {"keywords_zh": {"match_other_language_count": True}},
            {"keywords_zh": {"style_name": "KeywordsFixture"},
             "keywords_en": {"style_name": "KeywordsFixture"}},
        )
        match = next(item for item in findings
                     if item.get("property") == "keywords.match_other_language_count")
        self.assertEqual(match["declared_by"], ["keywords_zh"])
        self.assertEqual(match["template_value"], {"keywords_zh": 2, "keywords_en": 1})
        markers = manual_review_validation_items(findings)
        self.assertEqual(len(markers), 1)
        self.assertEqual(markers[0]["category"], "semantic_content_review")
        self.assertIn("不能确定应补充、删除", markers[0]["reason"])

    def test_overlong_keyword_is_detected_but_term_rewrite_is_a_human_decision(self) -> None:
        document = self._doc("摘要", "关键词：交通流预测模型")
        findings = audit_content_constraints(
            document,
            {"keywords_zh": {"max_item_chars": 5, "item_length_metric": "cjk_characters"}},
            {"keywords_zh": {"style_name": "KeywordsFixture"}},
        )
        violation = next(item for item in findings if ".max_item_chars[" in item["property"])
        self.assertEqual(violation["template_value"], 7)
        markers = manual_review_validation_items(findings)
        self.assertEqual(len(markers), 1)
        self.assertIn("不能擅自删改", markers[0]["reason"])

    def test_unimplemented_exception_policy_is_rejected_by_format_spec_schema(self) -> None:
        spec = {
            "schema_version": "1.0", "source_document": "template.docx",
            "roles": {}, "requirements": [], "status": "semantic_resolved",
            "content_constraints": {"abstract_zh": {"exception_policy": "none"}},
        }
        errors = load_and_validate(spec, ROOT / "schema" / "format-spec.schema.json")
        self.assertTrue(any("exception_policy" in error for error in errors), errors)

    def test_semantic_uncertainty_is_red_but_confident_violation_is_a_finding(self) -> None:
        checks = [{
            "check_id": "abstract_zh.require_third_person",
            "source_requirements": [{
                "requirement_id": "R00001", "clause_ids": ["C00001"],
                "evidence_ids": ["E00001"],
            }],
        }, {
            "check_id": "abstract_en.required_sections",
            "source_requirements": [{
                "requirement_id": "R00002", "clause_ids": ["C00002"],
                "evidence_ids": ["E00002"],
            }],
        }]
        review = {"results": [
            {"check_id": "abstract_zh.require_third_person", "verdict": "uncertain",
             "rationale": "target unclear", "evidence_quotes": ["原文"]},
            {"check_id": "abstract_en.required_sections", "verdict": "noncompliant",
             "rationale": "conclusion absent", "evidence_quotes": ["abstract"]},
        ]}
        markers = semantic_content_review_items(review, checks)
        findings = semantic_content_review_findings(review, checks)
        self.assertEqual(len(markers), 1)
        self.assertEqual(markers[0]["clause_ids"], ["C00001"])
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["failure_type"], "native_semantic_noncompliance")
        self.assertEqual(findings[0]["clause_ids"], ["C00002"])

    def test_only_explicit_manual_findings_become_red_items(self) -> None:
        findings = manual_review_validation_items([
            {"role": "keywords_zh", "property": "separator", "template_value": "，",
             "required_value": "semicolon"},
            {"role": "keywords_en", "property": "metric", "verification": "human_decision",
             "template_value": "cjk_characters", "required_value": "human scope decision"},
            {"role": "keywords_en", "property": "metric", "verification": "manual",
             "template_value": "cjk_characters", "required_value": "legacy generic manual"},
        ])
        self.assertEqual(len(findings), 1)
        self.assertIn("keywords_en.metric", findings[0]["source_text"])

    def test_unverified_property_receipts_stay_out_of_red_ledger(self) -> None:
        self.assertEqual(manual_review_receipt_items([{
            "receipt_id": "PR-R00001-0001", "status": "unverified",
        }]), [])


if __name__ == "__main__":
    unittest.main()
