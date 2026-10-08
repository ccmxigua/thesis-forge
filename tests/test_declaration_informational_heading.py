"""Physical declaration identity does not require a semantic heading edge."""
import copy
import io
import json
from pathlib import Path
import sys
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import host_agent_bridge as bridge
from fixed_declaration_source import derive_fixed_declaration_candidates, matches_declaration_render_selection
from requirements_engine import build_llm_request, merge_llm_primary
from source_literal_binding import materialize_source_fragment_literals
from resource_registry import materialize_declaration_resources, resource_items
from apply_format_spec import apply_declarations, audit_declarations
from docx import Document


def incident(*, fragments=False):
    data = json.loads((ROOT / "tests/fixtures/declaration-informational-heading-incident.json").read_text())
    source, raw = data["source"], data["raw"]
    chunk = build_llm_request([], source["clauses"],
        {"evidence": list(source["evidence_context"].values())}, {}, "full", contract_version="3.0")
    chunk.update(source)
    group = derive_fixed_declaration_candidates(chunk["clauses"], chunk["evidence_context"],
        anchor=chunk["declaration_anchor_preference"])[0]
    if fragments:
        raw["requirements"][0]["source_fragment_clause_ids"] = group["clause_ids"]
    return raw, chunk, group


class DeclarationInformationalHeadingTests(unittest.TestCase):
    def test_captured_render_sources_do_not_expand_execution_edges(self):
        for fragments in (False, True):
            with self.subTest(fragments=fragments):
                raw, chunk, group = incident(fragments=fragments)
                frozen = copy.deepcopy(raw)
                normalized = bridge.normalize_native_response(raw, chunk["response_schema"])
                candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
                self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
                self.assertEqual(candidate["clause_reviews"], normalized["clause_reviews"])
                req = candidate["requirements"][0]
                self.assertEqual(req["clause_ids"], raw["requirements"][0]["clause_ids"])
                self.assertEqual(req["evidence_ids"], raw["requirements"][0]["evidence_ids"])
                self.assertNotIn(group["heading_clause_id"], req["clause_ids"])
                projection = audit["declaration_source_text_projections"][0]
                self.assertIn(group["heading_clause_id"], projection["rendered_only_clause_ids"])
                self.assertEqual(projection["executable_clause_ids"], req["clause_ids"])
                self.assertEqual(raw, frozen)
                if fragments:
                    _, audits, errors = materialize_source_fragment_literals(candidate,
                        chunk["clauses"], chunk["evidence_context"])
                    self.assertEqual(errors, [])
                    self.assertFalse(audits[0]["render_only_selection"]["external_actions_verified"])

    def test_serialization_prints_each_source_paragraph_once_not_an_attestation(self):
        raw, chunk, group = incident(fragments=True)
        candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        declaration = candidate["requirements"][0]["properties"]["items"][0]
        self.assertEqual(
            declaration["heading"],
            chunk["evidence_context"][group["heading_evidence_ids"][0]]["text"],
        )
        self.assertEqual(
            declaration["body_parts"],
            [chunk["evidence_context"][eid]["text"] for eid in group["body_evidence_ids"]],
        )
        spec = {"schema_version": "1.0", "source_document": "current.docx", "status": "semantic_resolved",
                "roles": {}, "requirements": [], "declarations": candidate["requirements"][0]["properties"]}
        materialize_declaration_resources(spec, "informational-heading-offline",
            evidence={"evidence": list(chunk["evidence_context"].values())})
        doc = Document(); doc.add_paragraph("摘要", "Heading 1")
        apply_declarations(doc, spec["declarations"], resource_items(spec))
        memory = io.BytesIO(); doc.save(memory); saved = Document(memory)
        self.assertEqual(audit_declarations(saved, spec["declarations"], resource_items(spec)), [])
        for eid in group["evidence_ids"]:
            self.assertEqual([p.text for p in saved.paragraphs].count(chunk["evidence_context"][eid]["text"]), 1)
        for review in candidate["clause_reviews"]:
            if review["classification"] == "external_compliance":
                self.assertTrue(review["obligations"])
                self.assertTrue(all(a["status"] == "unverifiable" and a["route"] == "human"
                                    for a in review["obligations"]))

    def test_authorization_with_mixed_external_action_materializes_from_exact_source(self):
        raw, chunk, group = incident()
        raw = copy.deepcopy(raw)
        clause_id = "C00058"
        clause = next(item for item in chunk["clauses"] if item["id"] == clause_id)
        span = clause["source_span"]
        review = next(item for item in raw["clause_reviews"] if item["clause_id"] == clause_id)
        review["classification"] = "executable_with_external_check"
        review["reason"] = "The source text is printed; actual author consent remains a separate human action."
        review["obligations"] = [
            {"id": "authorization-text", "status": "covered",
             "reason": "The fixed text is printed exactly."},
            {"id": "author-consent", "status": "unverifiable",
             "actor": "author", "action": "confirm consent",
             "target": "authorization", "reason": "Printing does not establish consent."},
        ]
        semantic_obligations = copy.deepcopy(review["obligations"])

        raw_errors = bridge.validate_host_agent_response(raw, chunk)
        self.assertTrue(any("mixed_declaration_requires_explicit_route" in error for error in raw_errors), raw_errors)
        self.assertTrue(any("mixed_declaration_requires_current_source_quote" in error for error in raw_errors), raw_errors)
        candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
        self.assertEqual(len(audit["declaration_source_text_projections"]), 1)
        projection = audit["declaration_source_text_projections"][0]
        self.assertEqual(projection["clause_ids"], group["clause_ids"])
        declaration = candidate["requirements"][0]["properties"]["items"][0]
        self.assertEqual(
            declaration["body_parts"],
            [chunk["evidence_context"][eid]["text"] for eid in group["body_evidence_ids"]],
        )
        source_quote = chunk["evidence_context"][span["evidence_id"]]["text"][
            span["start_offset"]:span["end_offset"]
        ]
        candidate_review = next(
            item for item in candidate["clause_reviews"] if item["clause_id"] == clause_id
        )
        self.assertEqual(
            {atom["route"] for atom in candidate_review["obligations"]}, {"automatic", "human"},
        )
        self.assertEqual(
            {atom["source_quote"] for atom in candidate_review["obligations"]}, {source_quote},
        )
        self.assertEqual(
            {(atom["status"], atom["route"]) for atom in candidate_review["obligations"]},
            {("covered", "automatic"), ("unverifiable", "human")},
        )
        for original, projected in zip(semantic_obligations, candidate_review["obligations"]):
            for field in ("id", "status", "reason", "actor", "action", "target"):
                self.assertEqual(projected.get(field), original.get(field), field)
        projections = audit["declaration_source_text_projections"][0]["obligation_field_projections"]
        self.assertEqual({entry["field"] for entry in projections}, {"route", "source_quote"})

    def test_declaration_mixed_action_with_conflicting_route_or_unbound_quote_is_rejected(self):
        raw, chunk, _ = incident()
        raw = copy.deepcopy(raw)
        review = next(item for item in raw["clause_reviews"] if item["clause_id"] == "C00058")
        review["classification"] = "executable_with_external_check"
        review["reason"] = "The printed authorization text and a separate human action are both present."
        review["obligations"] = [
            {"id": "text", "status": "covered", "route": "human",
             "source_quote": "an unrelated source", "reason": "Printed."},
            {"id": "consent", "status": "unverifiable", "route": "automatic",
             "source_quote": "another source", "reason": "Requires author action."},
        ]
        errors = bridge.validate_host_agent_response(raw, chunk)
        self.assertTrue(any("responsibility_route_conflict" in error for error in errors), errors)
        self.assertTrue(any("must_equal_current_source_subspan" in error for error in errors), errors)
        projected, audits = bridge._materialize_fixed_declaration_source_text(raw, chunk)
        self.assertEqual(projected, raw)
        self.assertEqual(audits, [])

    def test_nonadjacent_source_fragment_is_not_grouped_as_one_declaration(self):
        raw, chunk, _ = incident()
        chunk = copy.deepcopy(chunk)
        clause = next(item for item in chunk["clauses"] if item["id"] == "C00055")
        span = clause["source_span"]
        evidence = chunk["evidence_context"][span["evidence_id"]]
        # Keep the source text/hash exact while creating a physical-node gap.
        # A source grouping must stop instead of joining over unreviewed text.
        for location in (clause.get("location"), span.get("location"), evidence.get("location")):
            location["child_index"] += 1
        self.assertEqual(
            derive_fixed_declaration_candidates(
                chunk["clauses"], chunk["evidence_context"],
                anchor=chunk["declaration_anchor_preference"],
            ),
            [],
        )
        projected, audits = bridge._materialize_fixed_declaration_source_text(raw, chunk)
        self.assertEqual(projected, raw)
        self.assertEqual(audits, [])

    def test_unresolved_requirement_edge_cannot_be_materialized(self):
        raw, chunk, _ = incident()
        raw = copy.deepcopy(raw)
        review = next(item for item in raw["clause_reviews"] if item["clause_id"] == "C00058")
        review["classification"] = "unresolved"
        review["obligations"] = []
        projected, audits = bridge._materialize_fixed_declaration_source_text(raw, chunk)
        self.assertEqual(projected, raw)
        self.assertEqual(audits, [])

    def test_structure_predicate_rejects_empty_foreign_duplicate_and_reordered_edges(self):
        raw, _, group = incident()
        edges = raw["requirements"][0]["clause_ids"]
        self.assertTrue(matches_declaration_render_selection(group, edges, group["evidence_ids"]))
        for bad in ([], edges + ["foreign"], edges + [edges[0]], list(reversed(edges))):
            self.assertFalse(matches_declaration_render_selection(group, bad, group["evidence_ids"]))
        for bad in ([], group["evidence_ids"][:-1], list(reversed(group["evidence_ids"])),
                    group["evidence_ids"] + [group["evidence_ids"][0]]):
            self.assertFalse(matches_declaration_render_selection(group, edges, bad))

    def test_merge_keeps_render_only_title_and_external_actions_out_of_execution(self):
        raw, chunk, group = incident(fragments=True)
        candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        # Merge consumes full extraction clauses, not the compact wire packet.
        clauses = copy.deepcopy(chunk["clauses"])
        for clause in clauses:
            clause["source_evidence_text"] = chunk["evidence_context"][clause["source_span"]["evidence_id"]]["text"]
        spec, conflicts, _ = merge_llm_primary(Path("current.docx"),
            {"schema_version": "1.0", "roles": {}, "requirements": [], "content_instances": []},
            clauses, candidate, set(chunk["evidence_context"]))
        self.assertEqual(conflicts, [])
        self.assertEqual(spec["requirements"][0]["clause_ids"], candidate["requirements"][0]["clause_ids"])
        self.assertNotIn(group["heading_clause_id"], spec["requirements"][0]["clause_ids"])
        external = {r["clause_id"] for r in candidate["clause_reviews"] if r["classification"] == "external_compliance"}
        self.assertFalse(external & set(spec["requirements"][0]["clause_ids"]))

    def test_informational_heading_cannot_become_an_execution_edge(self):
        raw, chunk, group = incident()
        raw["requirements"][0]["clause_ids"].insert(0, group["heading_clause_id"])
        raw["requirements"][0]["evidence_ids"].insert(0, group["heading_evidence_ids"][0])
        # Structural matching alone is not permission: classification validation rejects it.
        with self.assertRaises(ValueError):
            bridge.prepare_native_response_candidate(raw, chunk)

    def test_pending_route_cannot_expand_render_selector_and_forged_quote_rejects(self):
        raw, chunk, _ = incident(fragments=True)
        projected = bridge._materialize_fixed_declaration_source_text(raw, chunk)[0]
        pending = next(r for r in projected["clause_reviews"] if r["classification"] == "external_compliance")
        pending["obligations"][0]["route"] = "automatic"
        self.assertTrue(materialize_source_fragment_literals(projected, chunk["clauses"], chunk["evidence_context"])[2])
        raw, chunk, _ = incident(fragments=True)
        next(r for r in raw["clause_reviews"] if r["classification"] == "external_compliance")["obligations"][0]["source_quote"] = "foreign approval"
        with self.assertRaises(ValueError):
            bridge.prepare_native_response_candidate(raw, chunk)

    def test_current_source_and_human_inventory_faults_still_reject(self):
        for fault in ("stale_hash", "stale_location", "unknown_evidence", "duplicate_evidence",
                      "reordered_evidence", "missing_paragraph", "duplicate_clause", "invented_literal",
                      "no_pending", "human_as_covered", "informational_with_atom", "unresolved_heading"):
            with self.subTest(fault=fault):
                raw, chunk, group = incident(fragments=True)
                declaration = raw["requirements"][0]["properties"]["items"][0]
                heading = chunk["clauses"][0]
                pending = next(r for r in raw["clause_reviews"] if r["classification"] == "external_compliance")
                info = next(r for r in raw["clause_reviews"] if r["clause_id"] == group["heading_clause_id"])
                if fault == "stale_hash": heading["source_span"]["source_sha256"] = "0" * 64
                elif fault == "stale_location": heading["source_span"]["location"]["child_index"] += 100
                elif fault == "unknown_evidence": declaration["source_evidence_ids"][0] = "foreign"
                elif fault == "duplicate_evidence": declaration["source_evidence_ids"].append(declaration["source_evidence_ids"][0])
                elif fault == "reordered_evidence": declaration["source_evidence_ids"].reverse()
                elif fault == "missing_paragraph": declaration["source_evidence_ids"].pop()
                elif fault == "duplicate_clause": chunk["clauses"].append(copy.deepcopy(heading))
                elif fault == "invented_literal": declaration["heading"] = "forged"
                elif fault == "no_pending": pending["obligations"] = []
                elif fault == "human_as_covered": pending["obligations"][0]["status"] = "covered"
                elif fault == "informational_with_atom": info["obligations"] = copy.deepcopy(pending["obligations"])
                else: info["classification"] = "unresolved"
                with self.assertRaises(ValueError):
                    bridge.prepare_native_response_candidate(raw, chunk)

    def test_materializer_does_not_reseal_bad_render_only_heading_source(self):
        for position in (0, 1, -1):
            with self.subTest(position=position):
                raw, chunk, _ = incident()
                chunk["clauses"][position]["source_span"]["source_sha256"] = "0" * 64
                projected, audits = bridge._materialize_fixed_declaration_source_text(raw, chunk)
                self.assertEqual(projected, raw)
                self.assertEqual(audits, [])
                with self.assertRaises(ValueError):
                    bridge.prepare_native_response_candidate(raw, chunk)

    def test_generic_current_identifiers_not_a_school_or_clause_exception(self):
        raw, chunk, _ = incident(fragments=True)
        mapping = {c["id"]: "current-" + c["id"] for c in chunk["clauses"]}
        mapping.update({eid: "current-" + eid for eid in chunk["evidence_context"]})
        def rename(value):
            if isinstance(value, dict): return {mapping.get(k, k): rename(v) for k, v in value.items()}
            if isinstance(value, list): return [rename(v) for v in value]
            return mapping.get(value, value) if isinstance(value, str) else value
        raw, chunk = rename(raw), rename(chunk)
        candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])


if __name__ == "__main__": unittest.main()
