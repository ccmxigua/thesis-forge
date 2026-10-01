"""Render-source edges are not claims of completed legal/human obligations."""
import copy
import io
import json
from pathlib import Path
import sys
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import host_agent_bridge as bridge
from requirements_engine import build_llm_request
from source_literal_binding import materialize_source_fragment_literals
from resource_registry import materialize_declaration_resources, resource_items
from apply_format_spec import apply_declarations, audit_declarations
from docx import Document


def incident():
    data = json.loads((ROOT / "tests/fixtures/declaration-render-only-incident.json").read_text())
    source = data["source"]
    evidence = {"evidence": list(source["evidence_context"].values())}
    chunk = build_llm_request([], source["clauses"], evidence, {}, "full", contract_version="3.0")
    chunk.update(source)
    return data["raw"], chunk


class DeclarationRenderSelectionTests(unittest.TestCase):
    def test_captured_complete_render_selector_preserves_external_pending_edges(self):
        raw, chunk = incident(); frozen = copy.deepcopy(raw)
        candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
        normalized = bridge.normalize_native_response(raw, chunk["response_schema"])
        native_text = bridge._materialize_fixed_declaration_source_text(normalized, chunk)[0]
        source_only, _, source_errors = materialize_source_fragment_literals(
            native_text, chunk["clauses"], chunk["evidence_context"])
        self.assertEqual(source_errors, [])
        self.assertEqual(source_only["clause_reviews"], normalized["clause_reviews"])
        self.assertEqual(candidate["requirements"][0]["clause_ids"], raw["requirements"][0]["clause_ids"])
        _, audits, errors = materialize_source_fragment_literals(candidate, chunk["clauses"], chunk["evidence_context"])
        self.assertEqual(errors, [])
        self.assertEqual(len(audits), 1)
        audit = audits[0]["render_only_selection"]
        self.assertEqual(len(audit["render_only_clause_ids"]), 1)
        self.assertFalse(audit["external_actions_verified"])
        self.assertEqual(raw, frozen)

    def test_serialized_docx_prints_entire_source_once_without_signing(self):
        raw, chunk = incident(); candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        spec = {"schema_version": "1.0", "source_document": "current.docx", "status": "semantic_resolved",
                "roles": {}, "requirements": [], "declarations": candidate["requirements"][0]["properties"]}
        materialize_declaration_resources(spec, "render-only-offline", evidence={"evidence": list(chunk["evidence_context"].values())})
        doc = Document(); doc.add_paragraph("摘要", "Heading 1")
        apply_declarations(doc, spec["declarations"], resource_items(spec))
        memory = io.BytesIO(); doc.save(memory); saved = Document(memory)
        self.assertEqual(audit_declarations(saved, spec["declarations"], resource_items(spec)), [])
        for body in next(iter(resource_items(spec).values()))["body_parts"]:
            self.assertEqual([p.text for p in saved.paragraphs].count(body), 1)
        self.assertTrue(all(atom["status"] == "unverifiable" and atom["route"] == "human"
            for review in candidate["clause_reviews"] if review["classification"] == "external_compliance"
            for atom in review["obligations"]))

    def test_incomplete_prose_selector_still_fails(self):
        raw, chunk = incident()
        raw["requirements"][0]["source_fragment_clause_ids"] = raw["requirements"][0]["clause_ids"]
        with self.assertRaisesRegex(ValueError, "omits_evidence_text"):
            bridge.prepare_native_response_candidate(raw, chunk)

    def test_foreign_ambiguous_or_unproved_human_inventory_cannot_expand_scope(self):
        for change in ("unknown_clause", "signature_clause", "no_pending", "covered", "automatic",
                       "wrong_evidence", "duplicate_clause", "stale_source", "not_declaration"):
            with self.subTest(change=change):
                raw, chunk = incident()
                req = raw["requirements"][0]
                extra = next(cid for cid in req["source_fragment_clause_ids"] if cid not in req["clause_ids"])
                review = next(r for r in raw["clause_reviews"] if r["clause_id"] == extra)
                if change == "unknown_clause": req["source_fragment_clause_ids"].append("foreign")
                elif change == "signature_clause": req["source_fragment_clause_ids"].append(chunk["clauses"][-1]["id"])
                elif change == "no_pending": review["obligations"] = []
                elif change == "covered": review["obligations"][0]["status"] = "covered"
                elif change == "automatic": review["obligations"][0]["route"] = "automatic"
                elif change == "wrong_evidence": req["properties"]["items"][0]["source_evidence_ids"] = ["foreign"]
                elif change == "duplicate_clause": chunk["clauses"].append(copy.deepcopy(chunk["clauses"][0]))
                elif change == "stale_source": chunk["clauses"][0]["source_span"]["source_sha256"] = "0" * 64
                else: req["role"] = "body_text"
                projected = bridge._materialize_fixed_declaration_source_text(raw, chunk)[0]
                _, _, errors = materialize_source_fragment_literals(projected, chunk["clauses"], chunk["evidence_context"])
                self.assertTrue(errors)

    def test_role_native_text_cannot_be_shortened_or_rewritten(self):
        raw, chunk = incident()
        candidate = bridge._materialize_fixed_declaration_source_text(raw, chunk)[0]
        candidate["requirements"][0]["properties"]["items"][0]["body_parts"] = ["shortened declaration"]
        _, _, errors = materialize_source_fragment_literals(candidate, chunk["clauses"], chunk["evidence_context"])
        self.assertTrue(errors)

    def test_current_ids_not_school_specific(self):
        raw, chunk = incident()
        replacements = {c["id"]: "different-" + c["id"] for c in chunk["clauses"]}
        replacements.update({eid: "different-" + eid for eid in chunk["evidence_context"]})
        def rename(value):
            if isinstance(value, dict): return {replacements.get(k,k): rename(v) for k,v in value.items()}
            if isinstance(value, list): return [rename(v) for v in value]
            return replacements.get(value, value) if isinstance(value, str) else value
        raw, chunk = rename(raw), rename(chunk)
        candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])


if __name__ == "__main__": unittest.main()
