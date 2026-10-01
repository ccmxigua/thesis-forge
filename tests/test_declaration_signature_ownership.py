"""Captured signature hash typo cannot remain a model copying responsibility."""
import copy
import hashlib
import io
import json
from pathlib import Path
import sys
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import host_agent_bridge as bridge
from requirements_engine import build_llm_request
from host_review_schema import native_output_schema, native_schema_support_errors
from format_spec_validation import validate_instance
from resource_registry import materialize_declaration_resources, resource_items
from apply_format_spec import apply_declarations, audit_declarations
from docx import Document


def incident():
    data = json.loads((ROOT / "tests/fixtures/declaration-signature-hash-incident.json").read_text())
    source = data["source"]
    chunk = build_llm_request([], source["clauses"],
        {"evidence": list(source["evidence_context"].values())}, {}, "full", contract_version="3.0")
    chunk.update(source)
    return data, chunk


def model_proposal(data):
    proposal = copy.deepcopy(data["parent"])
    proposal["requirements"][0]["properties"]["items"][0]["source_signature_lines"] = None
    return proposal


class DeclarationSignatureOwnershipTests(unittest.TestCase):
    def test_captured_parent_typo_stays_rejected_not_resealed(self):
        data, chunk = incident()
        parent = copy.deepcopy(data["parent"])
        self.assertEqual(bridge._retry_change_paths(parent, data["retry"]),
            ["$.requirements[0].properties.items[0].source_signature_lines[0].source_sha256"])
        with self.assertRaisesRegex(ValueError, "must_equal_unique_adjacent_blank_source_lines"):
            bridge.prepare_native_response_candidate(parent, chunk)
        self.assertEqual(parent, data["parent"])

    def test_native_wire_only_accepts_null_local_and_resource_contract_unchanged(self):
        data, chunk = incident()
        local = chunk["response_schema"]; frozen = copy.deepcopy(local)
        wire = native_output_schema(local)
        self.assertEqual(native_schema_support_errors(wire), [])
        prop = wire["$defs"]["declarationItem"]["properties"]["source_signature_lines"]
        self.assertEqual(prop["type"], "null")
        self.assertIn("Code-owned", prop["description"])
        check = {"$ref": "#/$defs/declarationItem", "$defs": wire["$defs"]}
        for supplied in (data["parent"], data["retry"]):
            # validate_instance returns ERROR STRINGS, not a validity boolean.
            errors = validate_instance(supplied["requirements"][0]["properties"]["items"][0], check)
            self.assertIn("$.source_signature_lines: expected null, got array", errors)
        self.assertEqual(validate_instance(model_proposal(data)["requirements"][0]["properties"]["items"][0], check), [])
        self.assertEqual(local, frozen)
        for definition in ("declarationItem", "resourceEntry"):
            self.assertEqual(local["$defs"][definition]["properties"]["source_signature_lines"]["type"], "array")

    def test_null_proposal_gets_exact_current_source_not_model_hash(self):
        data, chunk = incident(); raw = model_proposal(data); frozen = copy.deepcopy(raw)
        candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
        lines = candidate["requirements"][0]["properties"]["items"][0]["source_signature_lines"]
        self.assertEqual(len(lines), 1)
        evidence = chunk["evidence_context"][lines[0]["source_evidence_id"]]
        self.assertEqual(lines[0]["text"], evidence["text"])
        self.assertEqual(lines[0]["source_sha256"], hashlib.sha256(evidence["text"].encode()).hexdigest())
        self.assertEqual(lines[0]["attestation_scope"], "placeholder_presence_only")
        projection = audit["declaration_source_text_projections"][0]
        self.assertEqual(projection["source_signature_lines"], lines)
        self.assertNotEqual(projection["response_before_sha256"], projection["response_after_sha256"])
        self.assertEqual(raw, frozen)
        self.assertEqual(candidate["clause_reviews"], bridge.normalize_native_response(raw, chunk["response_schema"])["clause_reviews"])

    def test_signature_source_identity_does_not_require_a_heading_execution_edge(self):
        data, chunk = incident(); raw = model_proposal(data)
        group = bridge._fixed_declaration_candidates(chunk["clauses"], chunk["evidence_context"],
            anchor=chunk["declaration_anchor_preference"])[0]
        heading = group["heading_clause_id"]
        req = raw["requirements"][0]
        req["clause_ids"].remove(heading)
        req["evidence_ids"] = [eid for eid in req["evidence_ids"] if eid not in group["heading_evidence_ids"]]
        review = next(r for r in raw["clause_reviews"] if r["clause_id"] == heading)
        review.update(classification="informational", obligations=[], normative_basis="insufficient")
        candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
        self.assertTrue(candidate["requirements"][0]["properties"]["items"][0]["source_signature_lines"])
        self.assertNotIn(heading, candidate["requirements"][0]["clause_ids"])
        self.assertEqual(candidate["clause_reviews"], bridge.normalize_native_response(raw, chunk["response_schema"])["clause_reviews"])

    def test_source_identity_faults_still_fail_closed(self):
        for mutation in ("stale_hash", "foreign_evidence", "forged_text"):
            data, chunk = incident(); raw = copy.deepcopy(data["retry"])
            line = raw["requirements"][0]["properties"]["items"][0]["source_signature_lines"][0]
            if mutation == "stale_hash": line["source_sha256"] = "0" * 64
            elif mutation == "foreign_evidence": line["source_evidence_id"] = "foreign"
            else: line["text"] += " signed"
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                bridge.prepare_native_response_candidate(raw, chunk)

    def test_nonadjacent_filled_or_completed_lines_not_materialized(self):
        for mutation in ("nonadjacent", "filled", "completed"):
            data, chunk = incident(); raw = model_proposal(data)
            eid = data["retry"]["requirements"][0]["properties"]["items"][0]["source_signature_lines"][0]["source_evidence_id"]
            if mutation == "nonadjacent": chunk["evidence_context"][eid]["location"]["child_index"] += 10
            elif mutation == "filled": chunk["evidence_context"][eid]["text"] += " 张三"
            else:
                signature = next(c for c in chunk["clauses"] if eid in c["evidence_ids"])
                review = next(r for r in raw["clause_reviews"] if r["clause_id"] == signature["id"])
                review["classification"] = "covered"
            with self.subTest(mutation=mutation):
                projected, _ = bridge._materialize_fixed_declaration_source_text(raw, chunk)
                self.assertFalse(projected["requirements"][0]["properties"]["items"][0].get("source_signature_lines"))

    def test_wire_rule_is_not_a_global_property_name_filter(self):
        local = {"type": "object", "properties": {"source_signature_lines": {"type": "string"}},
                 "required": ["source_signature_lines"], "additionalProperties": False}
        self.assertEqual(native_output_schema(local)["properties"]["source_signature_lines"], {"type": "string"})

    def test_portable_prompt_states_code_ownership(self):
        _, chunk = incident()
        self.assertTrue(any("source_signature_lines is code-owned" in instruction for instruction in chunk["instructions"]))
        self.assertIn("emit null", bridge._BASE_CONTRACT_REPAIR_RULES[0])

    def test_different_current_ids_do_not_need_a_school_specific_rule(self):
        data, chunk = incident(); raw = model_proposal(data)
        replacements = {c["id"]: "new-" + c["id"] for c in chunk["clauses"]}
        replacements.update({eid: "new-" + eid for eid in chunk["evidence_context"]})
        def rename(value):
            if isinstance(value, dict): return {replacements.get(k, k): rename(v) for k, v in value.items()}
            if isinstance(value, list): return [rename(v) for v in value]
            return replacements.get(value, value) if isinstance(value, str) else value
        source = rename(data["source"]); raw = rename(raw)
        rebuilt = build_llm_request([], source["clauses"],
            {"evidence": list(source["evidence_context"].values())}, {}, "full", contract_version="3.0")
        rebuilt.update(source)
        candidate, _ = bridge.prepare_native_response_candidate(raw, rebuilt)
        self.assertEqual(bridge.validate_host_agent_response(candidate, rebuilt), [])
        self.assertTrue(candidate["requirements"][0]["properties"]["items"][0]["source_signature_lines"][0]["source_evidence_id"].startswith("new-"))

    def test_serialized_docx_has_one_blank_line_not_a_completed_attestation(self):
        data, chunk = incident(); candidate, _ = bridge.prepare_native_response_candidate(model_proposal(data), chunk)
        spec = {"schema_version": "1.0", "source_document": "current.docx", "status": "semantic_resolved",
                "roles": {}, "requirements": [], "declarations": candidate["requirements"][0]["properties"]}
        materialize_declaration_resources(spec, "signature-offline", evidence={"evidence": list(chunk["evidence_context"].values())})
        doc = Document(); doc.add_paragraph("摘要", "Heading 1")
        apply_declarations(doc, spec["declarations"], resource_items(spec))
        memory = io.BytesIO(); doc.save(memory); saved = Document(memory)
        self.assertEqual(audit_declarations(saved, spec["declarations"], resource_items(spec)), [])
        line = candidate["requirements"][0]["properties"]["items"][0]["source_signature_lines"][0]
        self.assertEqual([p.text for p in saved.paragraphs].count(line["text"]), 1)
        self.assertTrue(all(atom["status"] == "unverifiable" and atom["route"] == "human"
            for review in candidate["clause_reviews"] if review["classification"] == "external_compliance"
            for atom in review["obligations"]))


if __name__ == "__main__": unittest.main()
