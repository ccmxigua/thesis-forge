from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt
from PIL import Image
import io

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from derive_review_draft_repairs import derive  # noqa: E402
from draft_scorecard import append_scorecard, build_scorecard  # noqa: E402
from manual_review import build_manual_review_ledger  # noqa: E402


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def image_data() -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (32, 32), "navy").save(stream, format="PNG")
    return stream.getvalue()


class DerivedReviewDraftRepairTests(unittest.TestCase):
    def test_fresh_child_repairs_caption_and_drawing_without_model_or_parent_mutation(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_thesis = root / "source-thesis.docx"
            requirements = root / "requirements.docx"
            parent = root / "parent-review.docx"
            format_spec = root / "format-spec.json"
            clauses = root / "clauses.json"
            extraction = root / "extraction-manifest.json"
            style_map = root / "style-map.json"
            scorecard_path = root / "scorecard.json"
            ledger_path = root / "manual-review-ledger.json"
            output = root / "derived" / "review-baseline.docx"
            report_path = root / "derived" / "repair-report.json"

            thesis = Document()
            thesis.add_paragraph("Sample source thesis")
            thesis.save(source_thesis)
            req = Document()
            req.add_paragraph("Synthetic requirement source")
            req.save(requirements)

            spec = {
                "schema_version": "1.0", "source_document": str(requirements.resolve()),
                "run_id": "test-run", "status": "rule_resolved", "roles": {
                    "figure_caption": {"separator": " "},
                    "table_caption": {"position": "above", "separator": " "},
                }, "requirements": [],
                "semantic_review_provenance": {
                    "version": "1.0", "origin": "fresh_host_agent",
                    "run_id": "test-run", "source_sha256": sha(requirements),
                    "evidence_sha256": "a" * 64, "clause_sha256": "b" * 64,
                    "request_sha256": "c" * 64,
                },
            }
            format_spec.write_text(json.dumps(spec), encoding="utf-8")
            clauses.write_text("[]\n", encoding="utf-8")
            extraction.write_text(json.dumps({
                "status": "completed", "run_id": "test-run",
                "normalized_source_document": str(requirements.resolve()),
                "normalized_source_sha256": sha(requirements),
                "requirement_clauses_sha256": sha(clauses),
            }), encoding="utf-8")
            style_map.write_text(json.dumps({
                "figure_caption": {"style_name": "Figure Caption"},
                "table_caption": {"style_name": "Table Caption"},
            }), encoding="utf-8")

            document = Document()
            normal = document.styles["Normal"]
            normal.paragraph_format.line_spacing = Pt(20)
            normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            figure_style = document.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
            table_style = document.styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
            drawing = document.add_paragraph()
            drawing.add_run().add_picture(io.BytesIO(image_data()), width=Inches(1))
            caption = document.add_paragraph("图1    示例图")
            caption.style = figure_style
            table_caption = document.add_paragraph("表1    示例表")
            table_caption.style = table_style
            document.add_table(rows=1, cols=1).cell(0, 0).text = "sample"
            parent_binding = {
                "run_id": "test-run", "case_id": "case-test",
                "source_sha256": sha(requirements),
                "clause_sha256": "a" * 64,
                "evidence_sha256": "b" * 64,
                "request_sha256": None,
                "format_spec_sha256": sha(format_spec),
                "requirements_sha256": sha(requirements),
                "input_source_sha256": sha(source_thesis),
                "official_template_sha256": None,
                "official_template_source": "not_supplied",
            }
            card = build_scorecard(parent_binding, {
                "missing_count": 0, "unexpected_count": 0, "duplicate_count": 0,
                "expected_receipt_ids": [], "receipts": [],
            }, [], [])
            append_scorecard(document, card)
            document.save(parent)
            scorecard_path.write_text(json.dumps(card), encoding="utf-8")
            ledger = build_manual_review_ledger({}, [], binding=parent_binding)
            ledger_path.write_text(json.dumps(ledger), encoding="utf-8")
            parent_before = sha(parent)
            input_hashes = {p: sha(p) for p in (source_thesis, requirements, format_spec,
                                                clauses, extraction, style_map, scorecard_path,
                                                ledger_path)}

            report = derive(source_thesis, parent, format_spec, clauses, extraction,
                            style_map, scorecard_path, ledger_path, output, report_path)
            self.assertEqual(report["status"], "passed", report["findings"])
            self.assertFalse(report["model_request_made"])
            self.assertFalse(report["semantic_review_is_new"])
            self.assertEqual(report["semantic_review_lineage"]["parent_docx_sha256"], parent_before)
            self.assertTrue(report["input_hashes_unchanged"])
            self.assertEqual({p: sha(p) for p in input_hashes}, input_hashes)
            self.assertEqual(sha(parent), parent_before)
            final_doc = Document(output)
            caption_text = next(p.text for p in final_doc.paragraphs if p.text.startswith("图1"))
            self.assertEqual(caption_text, "图1 示例图")
            self.assertEqual(report["audits"]["drawing_line_boxes"]["findings"], [])
            self.assertTrue((report_path.parent / "inherited-draft-scorecard.json").is_file())


if __name__ == "__main__":
    unittest.main()
