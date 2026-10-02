"""Distinct discipline inputs stay distinct from schema through serialized cover."""
from __future__ import annotations

import copy
import io
import json
from pathlib import Path
import sys

from docx import Document
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import apply_format_spec
import format_contract_guards
from format_spec_validation import validate_instance
from extract_semantic_metadata import extract, normalize_semantic_metadata
from input_resolver import input_value_type_valid, resolve_metadata


def cover():
    return {"institution": "测试大学", "fields": [
        {"id": key, "label": label, "value_from": f"thesis_profile.cover_metadata.{key}",
         "display_policy": "required", "order": index}
        for index, (key, label) in enumerate([
            ("first_discipline", "一级学科"), ("second_discipline", "二级学科")], 1)
    ]}


def metadata():
    return {"trust": {"source": "source_document", "confirmed": True},
            "author_name": "测试作者", "title_zh": "测试题目", "title_en": "Test title",
            "student_id": "TEST", "supervisor_name": "测试导师", "completion_date": "2026-06",
            "first_discipline": "交通运输工程", "second_discipline": "交通信息工程及控制"}


def test_schema_catalog_and_field_binding():
    profile_schema = json.loads((ROOT / "schema/thesis-profile.schema.json").read_text())
    format_schema = json.loads((ROOT / "schema/format-spec.schema.json").read_text())
    profile = {"schema_version": "1.0", "degree_level": "master", "writing_language": "zh",
               "cover_metadata": metadata()}
    assert validate_instance(profile, profile_schema) == []
    assert validate_instance(profile["cover_metadata"], format_schema["$defs"]["coverMetadata"], format_schema) == []
    assert validate_instance(cover(), format_schema["$defs"]["coverSpec"], format_schema) == []
    assert format_contract_guards.cover_binding_errors({"cover": cover()}) == []
    for key in ("first_discipline", "second_discipline"):
        assert key in format_contract_guards.REGISTERED_COVER_FIELDS
        invalid = copy.deepcopy(profile)
        invalid["cover_metadata"][key] = ""
        assert validate_instance(invalid, profile_schema)
    invalid = cover()
    invalid["fields"][0]["id"] = "invented_discipline"
    assert validate_instance(invalid, format_schema["$defs"]["coverSpec"], format_schema)


def test_explicit_source_disciplines_reach_profile_and_registered_inputs():
    raw = extract(ROOT / "tests/sample-thesis.tex")
    profile = normalize_semantic_metadata(raw)
    catalog = format_contract_guards.registered_input_catalog()
    for key in ("first_discipline", "second_discipline"):
        path = f"thesis_profile.cover_metadata.{key}"
        assert profile["cover_metadata"][key] == raw[key]
        assert profile["provenance"]["field_sources"][f"cover_metadata.{key}"] == key
        assert path in catalog["metadata"]
        assert resolve_metadata(profile, key) == raw[key]
        assert input_value_type_valid(path, raw[key])
        for wrong in (False, 0, [], {}, None):
            assert not input_value_type_valid(path, wrong)
    raw.pop("first_discipline")
    raw.pop("second_discipline")
    raw.update(major="不可推断", degree_discipline="不可推断", field_name="不可推断")
    absent = normalize_semantic_metadata(raw)["cover_metadata"]
    assert "first_discipline" not in absent
    assert "second_discipline" not in absent


@pytest.mark.parametrize("id_key,source_key", [
    ("degree_discipline", "degree_discipline"), ("field_name", "field_name"),
    ("second_discipline", "second_discipline"), ("first_discipline", "second_discipline")])
def test_discipline_label_cannot_bind_to_another_value(id_key, source_key):
    invalid = cover()
    invalid["fields"][0].update(id=id_key, value_from=f"thesis_profile.cover_metadata.{source_key}")
    assert format_contract_guards.cover_binding_errors({"cover": invalid})


def test_historical_discipline_aliases_still_reject_without_rewriting_fixture():
    fixture = json.loads((ROOT / "tests/fixtures/cover-condition-reassessment-incident.json").read_text())
    old = fixture["primary_candidate"]["requirements"][0]["properties"]
    frozen = copy.deepcopy(old)
    errors = format_contract_guards.cover_binding_errors({"cover": old})
    assert any("first_discipline" in error for error in errors)
    assert any("second_discipline" in error for error in errors)
    assert old == frozen


def test_discipline_serialization_and_receipt_detects_missing_or_swapped_values():
    definition = cover()
    profile = {"cover_metadata": metadata()}
    contract = apply_format_spec.compile_cover_contract(definition, profile)
    assert [item["value"] for item in contract["fields"]] == [
        "交通运输工程", "交通信息工程及控制"]
    doc = Document()
    doc.add_paragraph("正文保留")
    counts = apply_format_spec.apply_cover(doc, definition, profile, contract)
    assert counts["trusted_fields_written"] == 2
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    serialized = Document(buffer)
    assert [p.text for p in serialized.paragraphs][:3] == [
        "测试大学", "一级学科：交通运输工程", "二级学科：交通信息工程及控制"]
    spec = {"cover": definition}
    reqs = [{"role": "cover", "properties": {"fields": definition["fields"]}}]
    actual, _ = apply_format_spec._receipt_semantic_actuals(serialized, spec, {}, reqs, {}, [], contract)
    assert actual["cover"]["fields"] == definition["fields"]
    serialized.paragraphs[1].text = "一级学科：交通信息工程及控制"
    actual, _ = apply_format_spec._receipt_semantic_actuals(serialized, spec, {}, reqs, {}, [], contract)
    assert "fields" not in actual.get("cover", {})


@pytest.mark.parametrize("missing", ["first_discipline", "second_discipline", "both", "trust"])
def test_missing_discipline_is_placeholder_never_a_major_alias(missing):
    values = metadata()
    values.update(degree_discipline="不能替代一级学科", field_name="不能替代二级学科")
    if missing == "both":
        values.pop("first_discipline")
        values.pop("second_discipline")
    else:
        values.pop(missing)
    profile = {"cover_metadata": values}
    contract = apply_format_spec.compile_cover_contract(cover(), profile)
    pending = {item["id"] for item in contract["fields"] if item["value_kind"] == "placeholder"}
    expected = {"first_discipline", "second_discipline"} if missing in {"both", "trust"} else {missing}
    assert pending == expected
    doc = Document()
    doc.add_paragraph("正文")
    counts = apply_format_spec.apply_cover(doc, cover(), profile, contract)
    assert set(counts["metadata_pending_fields"]) == expected
    assert all(item["value"] == "——" for item in contract["fields"] if item["id"] in expected)


@pytest.mark.parametrize("key", ["first_discipline", "second_discipline"])
@pytest.mark.parametrize("wrong", [False, 0, ["甲"], {"name": "甲"}])
def test_direct_compiler_and_executor_reject_structured_discipline_values(key, wrong):
    values = metadata()
    values[key] = wrong
    profile = {"cover_metadata": values}
    with pytest.raises(ValueError, match=f"{key} must be a string"):
        apply_format_spec.compile_cover_contract(cover(), profile)
    doc = Document()
    doc.add_paragraph("原文保留")
    with pytest.raises(ValueError, match=f"{key} must be a string"):
        apply_format_spec.apply_cover(doc, cover(), profile)
    assert [p.text for p in doc.paragraphs] == ["原文保留"]
