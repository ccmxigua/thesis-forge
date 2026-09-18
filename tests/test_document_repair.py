from __future__ import annotations

import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import document_repair
from document_repair import repair
from submission_audit import audit_docx


def field(paragraph, instruction: str, cached: str) -> None:
    node = OxmlElement("w:fldSimple"); node.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = cached
    run.append(text); node.append(run); paragraph._p.append(node)


class DocumentRepairTest(unittest.TestCase):
    def test_deterministic_header_page_and_equation_repairs(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "repaired.docx"
            doc = Document()
            doc.add_heading("摘 要", 1)
            doc.add_paragraph("前置内容")
            doc.add_section(WD_SECTION.NEW_PAGE)
            doc.add_heading("第1章 引言", 1)
            paragraph = doc.add_paragraph("参见式\u00a0")
            field(paragraph, "REF eq_one \\h", "式（1.1）")
            hyperlink_paragraph = doc.add_paragraph("另见式\u00a0")
            hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("w:anchor"), "sample_eq_two")
            run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "式（1.2）"
            run.append(text); hyperlink.append(run); hyperlink_paragraph._p.append(hyperlink)
            marker_paragraph = doc.add_paragraph("参见图\u00a0")
            marker_link = OxmlElement("w:hyperlink"); marker_link.set(qn("w:anchor"), "sample_fig_one")
            run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "[fig:one)"
            run.append(text); marker_link.append(run); marker_paragraph._p.append(marker_link)
            caption = doc.add_paragraph("图1-1    合成示意图")
            bookmark = OxmlElement("w:bookmarkStart"); bookmark.set(qn("w:id"), "21")
            bookmark.set(qn("w:name"), "sample_fig_one")
            bookmark_end = OxmlElement("w:bookmarkEnd"); bookmark_end.set(qn("w:id"), "21")
            caption._p.insert(1, bookmark); caption._p.insert(2, bookmark_end)
            for section in doc.sections:
                section.header.paragraphs[0].add_run("某大学博士/硕士学位论文")
                field(section.header.paragraphs[0], 'STYLEREF "Heading 1" \\* MERGEFORMAT', "Heading 1")
                field(section.footer.paragraphs[0], "PAGE \\* ROMAN", "I")
            doc.save(source)
            spec = {
                "thesis_profile": {"degree_level": "doctor"},
                "roles": {"header": {"header_content": {
                    "left_text": "某大学博士/硕士学位论文", "right_field": "styleref_heading_1"}}},
                "page": {"page_number": {"front_matter_format": "roman_upper", "body_format": "decimal",
                    "front_matter_start": 1, "body_start": 1, "alignment": "center",
                    "body_start_selector": {"strategy": "section_index", "section_index": 2}}},
            }
            report = repair(source, spec, output)
            self.assertFalse(report["submission_ready_claimed"])
            self.assertEqual(report["repairs"]["headers"]["styleref_operand"], "1")
            self.assertEqual(report["repairs"]["equation_reference_prefixes"]["count"], 2)
            self.assertEqual(report["repairs"]["bound_source_markers"]["count"], 1)
            audit = audit_docx(output, spec)
            codes = {issue["code"] for issue in audit["issues"]}
            self.assertNotIn("unresolved_degree_header_placeholder", codes)
            self.assertNotIn("page_field_section_coverage_or_format_mismatch", codes)
            self.assertNotIn("page_fields_not_centered", codes)
            self.assertNotIn("duplicate_equation_reference_prefix", codes)
            self.assertFalse(audit["submission_ready"])
            self.assertEqual(audit["render_validation"]["status"], "not_run")
            # The independent PDF policy resolves the declared degree too; it
            # must not continue expecting the unresolved slash placeholder.
            self.assertEqual(spec["thesis_profile"]["degree_level"], "doctor")

    def test_cleanup_removes_comments_package_wide(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "commented.docx"; output = root / "clean.docx"
            doc = Document(); paragraph = doc.add_paragraph("正文")
            start = OxmlElement("w:commentRangeStart"); start.set(qn("w:id"), "0")
            end = OxmlElement("w:commentRangeEnd"); end.set(qn("w:id"), "0")
            reference = OxmlElement("w:commentReference"); reference.set(qn("w:id"), "0")
            paragraph._p.insert(1, start); paragraph._p.append(end); paragraph._p.append(reference)
            doc.save(source)
            with zipfile.ZipFile(source) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            custom_comment_part = "word/review/reviewerNotes.xml"
            members[custom_comment_part] = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:comment w:id="0" w:author="Reviewer"><w:p><w:r><w:t>remove me</w:t></w:r></w:p></w:comment>'
                '</w:comments>'
            ).encode()
            members["word/review/_rels/reviewerNotes.xml.rels"] = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rIdPeople" '
                'Type="http://schemas.microsoft.com/office/2011/relationships/people" '
                'Target="../people.xml"/></Relationships>'
            ).encode()
            members["word/people.xml"] = b'<w:people xmlns:w="http://schemas.microsoft.com/office/word/2012/wordml"/>'
            rels = members["word/_rels/document.xml.rels"].decode()
            rels = rels.replace(
                "</Relationships>",
                '<Relationship Id="rIdComment" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" '
                'Target="review/reviewerNotes.xml"/></Relationships>',
            )
            members["word/_rels/document.xml.rels"] = rels.encode()
            types = members["[Content_Types].xml"].decode()
            types = types.replace(
                "</Types>",
                f'<Override PartName="/{custom_comment_part}" '
                'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
                '<Override PartName="/word/people.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.person+xml"/>'
                '</Types>',
            )
            members["[Content_Types].xml"] = types.encode()
            with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
                for name, payload in members.items():
                    archive.writestr(name, payload)

            report = repair(source, {"cleanup": {"remove_comments": True}}, output)
            cleanup = report["repairs"]["comments"]
            self.assertEqual(cleanup["removed_part_count"], 2)
            self.assertEqual(cleanup["removed_markers"], 3)
            self.assertEqual(
                cleanup["removed_sidecar_relationship_parts"],
                ["word/review/_rels/reviewerNotes.xml.rels"],
            )
            with zipfile.ZipFile(output) as archive:
                names = set(archive.namelist())
                self.assertNotIn(custom_comment_part, names)
                self.assertNotIn("word/review/_rels/reviewerNotes.xml.rels", names)
                self.assertNotIn("word/people.xml", names)
                self.assertNotIn("comments", archive.read("word/document.xml").decode().casefold())
                self.assertNotIn("relationships/comments", archive.read("word/_rels/document.xml.rels").decode())
                self.assertNotIn("comments+xml", archive.read("[Content_Types].xml").decode())
            audit = audit_docx(output, {})
            self.assertNotIn("unresolved_document_comments", {x["code"] for x in audit["issues"]})

    def test_declared_range_heading_normalization_and_body_headers_are_deterministic(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "output.docx"
            doc = Document()
            doc.add_paragraph("封面")
            doc.add_section(WD_SECTION.NEW_PAGE); doc.add_paragraph("目 录")
            doc.add_paragraph("CONTENTS"); doc.add_paragraph("Title 1")
            doc.add_section(WD_SECTION.NEW_PAGE)
            doc.add_paragraph("1 引言"); doc.add_paragraph("正文")
            doc.add_paragraph("致 谢")
            for section in doc.sections:
                section.header.paragraphs[0].text = "旧奇数页眉"
                section.even_page_header.paragraphs[0].text = "旧偶数页眉"
            doc.save(source)
            spec = {
                "roles": {"header": {"header_content": {
                    "left_text": "测试大学硕士学位论文",
                    "right_field": "styleref_heading_1",
                }}},
                "cleanup": {
                    "body_headers_only": True,
                    "remove_ranges": [{
                        "start_text": "CONTENTS", "end_text": "1 引言",
                        "include_start": True, "include_end": False,
                    }],
                    "heading_1_texts": ["1 引言", "致 谢"],
                },
                "page": {"different_odd_even": True, "page_number": {
                    "front_matter_format": "roman_upper", "body_format": "decimal",
                    "front_matter_start_selector": {"strategy": "section_index", "section_index": 2},
                    "body_start_selector": {"strategy": "section_index", "section_index": 3},
                    "front_matter_start": 1, "body_start": 1, "alignment": "center",
                }},
            }
            report = repair(source, spec, output)
            self.assertEqual(report["repairs"]["declared_ranges"]["range_count"], 1)
            self.assertEqual(report["repairs"]["heading_1_boundaries"]["count"], 2)
            self.assertEqual(report["repairs"]["body_headers"]["section_indices"], [3])
            repaired = Document(output)
            self.assertNotIn("CONTENTS", [p.text for p in repaired.paragraphs])
            self.assertEqual([p.style.style_id for p in repaired.paragraphs if p.text in {"1 引言", "致 谢"}],
                             ["Heading1", "Heading1"])
            self.assertNotIn("STYLEREF", repaired.sections[0].header._element.xml)
            self.assertIn("STYLEREF 1", repaired.sections[2].header._element.xml)
            self.assertIn("测试大学硕士学位论文", repaired.sections[2].even_page_header._element.xml)
            plan = report["repairs"]["page_numbering"]["section_plan"]
            self.assertEqual([s["zone"] for s in plan["sections"]], ["cover", "front", "body"])
            self.assertEqual(report["repairs"]["page_numbering"]["audit"]["findings"], [])

    def test_refuses_in_place_repair(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "source.docx"; Document().save(path)
            with self.assertRaisesRegex(ValueError, "output must differ"):
                repair(path, {}, path)

    def test_degree_header_repair_handles_phrase_split_across_runs(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "repaired.docx"
            doc = Document()
            header = doc.sections[0].header.paragraphs[0]
            for value in ["某大学", "博士", "/", "硕士", "学位论文"]:
                header.add_run(value)
            doc.save(source)
            spec = {"thesis_profile": {"degree_level": "master"}, "roles": {
                "header": {"header_content": {"left_text": "某大学博士/硕士学位论文"}}}}
            report = repair(source, spec, output)
            self.assertEqual(report["repairs"]["headers"]["degree_placeholder_replacements"], 1)
            repaired = Document(output)
            self.assertEqual(repaired.sections[0].header.paragraphs[0].text, "某大学硕士学位论文")
            self.assertNotIn("博士/硕士", repaired.sections[0].header._element.xml)

    def test_degree_header_repair_covers_default_first_and_even_variants(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "repaired.docx"
            doc = Document()
            section = doc.sections[0]
            section.header.paragraphs[0].text = "某大学博士/硕士学位论文"
            section.first_page_header.paragraphs[0].text = "某大学博士/硕士学位论文"
            section.even_page_header.paragraphs[0].text = "某大学博士/硕士学位论文"
            doc.save(source)
            spec = {"thesis_profile": {"degree_level": "master"}, "roles": {
                "header": {"header_content": {"left_text": "某大学博士/硕士学位论文"}}}}
            report = repair(source, spec, output)
            self.assertEqual(report["repairs"]["headers"]["degree_placeholder_replacements"], 3)
            repaired = Document(output)
            for story in (repaired.sections[0].header, repaired.sections[0].first_page_header,
                          repaired.sections[0].even_page_header):
                self.assertIn("某大学硕士学位论文", story._element.xml)
                self.assertNotIn("博士/硕士", story._element.xml)

    def test_source_marker_without_unique_bound_caption_is_deferred(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "repaired.docx"
            doc = Document()
            doc.add_heading("第1章 引言", 1)
            paragraph = doc.add_paragraph("参见图\u00a0")
            hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("w:anchor"), "missing_fig")
            run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "[fig:missing)"
            run.append(text); hyperlink.append(run); paragraph._p.append(hyperlink)
            doc.save(source)
            report = repair(source, {}, output)
            marker = report["repairs"]["bound_source_markers"]
            self.assertEqual(marker["count"], 0)
            self.assertEqual(len(marker["deferred"]), 1)
            self.assertIn("target count is 0", marker["deferred"][0]["reason"])
            repaired = Document(output)
            self.assertIn("[fig:missing)", repaired.paragraphs[1].text)
            self.assertTrue(any(x["code"] == "source_marker_requires_target_resolution"
                                for x in report["deferred"]))

    def test_cli_rejects_invalid_format_spec_before_writing_output(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "output.docx"
            spec = root / "invalid.json"; report = root / "report.json"
            Document().save(source)
            spec.write_text(json.dumps({"page": {"page_number": {"body_format": "binary"}}}))
            argv = ["document_repair.py", str(source), str(spec), str(output), "--report", str(report)]
            with patch.object(sys, "argv", argv):
                with self.assertRaises(SystemExit) as raised:
                    document_repair.main()
            self.assertEqual(raised.exception.code, 2)
            self.assertFalse(output.exists())
            self.assertFalse(report.exists())

    def test_cli_refuses_report_path_aliasing_input(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "output.docx"
            spec = root / "spec.json"; Document().save(source); original = source.read_bytes()
            spec.write_text("{}")
            argv = ["document_repair.py", str(source), str(spec), str(output), "--report", str(source)]
            with patch.object(sys, "argv", argv):
                with self.assertRaises(SystemExit) as raised:
                    document_repair.main()
            self.assertEqual(raised.exception.code, 2)
            self.assertEqual(source.read_bytes(), original)
            self.assertFalse(output.exists())

    def test_page_repair_handles_active_first_footer_and_rejects_ambiguous_text(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "output.docx"
            doc = Document(); doc.add_heading("第1章 引言", 1)
            doc.sections[0].different_first_page_header_footer = True
            doc.save(source)
            spec = {"page": {"different_first_page": True, "page_number": {
                "body_format": "decimal", "alignment": "center"}}}
            report = repair(source, spec, output)
            active = [x for x in report["repairs"]["page_numbering"]["stories"] if x["active"]]
            self.assertEqual({x["variant"] for x in active}, {"default", "first"})
            repaired = Document(output)
            self.assertIn("PAGE", repaired.sections[0].footer._element.xml)
            self.assertIn("PAGE", repaired.sections[0].first_page_footer._element.xml)

            ambiguous = root / "ambiguous.docx"; rejected = root / "rejected.docx"
            doc = Document(); doc.add_heading("第1章 引言", 1)
            doc.sections[0].footer.paragraphs[0].text = "Confidential"
            doc.save(ambiguous)
            with self.assertRaisesRegex(ValueError, "automatic placement would be ambiguous"):
                repair(ambiguous, spec, rejected)
            self.assertFalse(rejected.exists())

    def test_styleref_style_detection_ignores_cached_toc_entries(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "output.docx"
            doc = Document()
            toc_style = doc.styles.add_style("TOC 1", 1)
            doc.add_paragraph("第1章 引言 ........ 1", style=toc_style)
            doc.add_heading("第1章 引言", 1)
            field(doc.sections[0].header.paragraphs[0], 'STYLEREF "Heading 1"', "Heading 1")
            doc.save(source)
            spec = {"roles": {"header": {"header_content": {"right_field": "styleref_heading_1"}}}}
            report = repair(source, spec, output)
            self.assertEqual(report["repairs"]["headers"]["styleref_operand"], "1")

    def test_page_repair_preserves_non_page_fields_and_removes_none_format(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; output = root / "output.docx"
            doc = Document(); doc.add_heading("第1章 引言", 1)
            footer = doc.sections[0].footer.paragraphs[0]
            field(footer, "PAGE \\* ARABIC", "1")
            field(footer, "NUMPAGES \\* ARABIC", "10")
            field(footer, "PAGEREF target \\h", "5")
            doc.save(source)
            spec = {"page": {"page_number": {"body_format": "none"}}}
            report = repair(source, spec, output)
            self.assertEqual(report["repairs"]["page_numbering"]["stories"][0]["removed_page_fields"], 1)
            xml = Document(output).sections[0].footer._element.xml
            self.assertNotIn("PAGE \\* ARABIC", xml)
            self.assertIn("NUMPAGES", xml)
            self.assertIn("PAGEREF", xml)


if __name__ == "__main__":
    unittest.main()
