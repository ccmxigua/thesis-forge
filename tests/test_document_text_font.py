from __future__ import annotations

import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import document_text_font as fonts
import host_agent_bridge as bridge
import requirements_engine as engine
import test_host_agent_bridge as bridge_fixtures
from semantic_contract import attach_request_provenance
import post_generation_format_audit as post_audit


class DocumentTextFontTests(unittest.TestCase):
    @staticmethod
    def fixture():
        source = "论文中出现英文时需要使用Times New Roman字体"
        cid, eid = "dynamic-global-clause", "dynamic-evidence"
        clause = {"id": cid, "text": source, "evidence_ids": [eid],
                  "source_span": {"evidence_id": eid, "start_offset": 0, "end_offset": len(source),
                                  "text": source, "source_sha256": hashlib.sha256(source.encode()).hexdigest()}}
        catalog = [{"id": f"dynamic-rule-{index}", "role": role, "source_text": source,
                    "properties": {"font": {"latin": "Times New Roman"}},
                    "applicability": fonts.compile_document_font_applicability(source),
                    "clause_ids": [cid], "evidence_ids": [eid]}
                   for index, role in enumerate(fonts.TEXT_FONT_ROLES)]
        chunk = {"clauses": [clause], "evidence_context": {eid: {"id": eid, "text": source}},
                 "rule_spec": {"requirements": catalog}, "provenance": {"run_id": "current-run"}}
        body = next(rule for rule in catalog if rule["role"] == "body_text")
        response = {"contract_version": "3.0", "requirements": [{
            "existing_requirement_id": body["id"], "role": "body_text", "properties": body["properties"],
            "applicability": copy.deepcopy(body["applicability"]),
            "clause_ids": [cid], "evidence_ids": [eid], "confidence": 0.98, "reason": "Body-only coverage."}],
            "clause_reviews": [{"clause_id": cid, "classification": "covered", "reason": "Font rule."}]}
        return response, chunk

    def test_source_scope_completion_is_dynamic_audited_and_idempotent(self):
        response, chunk = self.fixture()
        original = copy.deepcopy(response)
        projected, audit = fonts.materialize_document_font_references(response, chunk)
        self.assertEqual(response, original)
        self.assertEqual({r["role"] for r in projected["requirements"]}, set(fonts.TEXT_FONT_ROLES))
        self.assertEqual(len(audit), 1)
        self.assertEqual(audit[0]["provenance"], chunk["provenance"])
        self.assertEqual(audit[0]["source_span"], chunk["clauses"][0]["source_span"])
        self.assertFalse(audit[0]["submission_ready"])
        self.assertEqual(fonts.materialize_document_font_references(projected, chunk), (projected, []))
        self.assertTrue(fonts.document_font_scope_errors(response, chunk))
        self.assertEqual(fonts.document_font_scope_errors(projected, chunk), [])
        expected = fonts.compile_document_font_applicability(chunk["clauses"][0]["text"])
        self.assertTrue(all(r["applicability"] == expected for r in projected["requirements"]))
        self.assertEqual(audit[0]["source_applicability"], expected)
        self.assertEqual(audit[0]["source_applicability_sha256"], fonts._digest(expected))

    def test_changed_missing_or_extra_conditions_are_not_completed_or_executed(self):
        expected = fonts.compile_document_font_applicability("论文中出现英文时需要使用Times New Roman字体")
        extra = copy.deepcopy(expected)
        extra["conditions"].append({"fact": "thesis_profile.degree", "operator": "equals", "value": "doctor"})
        for declaration in (None, {}, {"status": "always"}, {"status": "excluded"},
                            {"status": "conditional", "conditions": []}, extra,
                            {**expected, "exceptions": ["正文除外"]},
                            {**expected, "unknown": True},
                            {"status": "conditional", "conditions": [{"fact": "runtime.english_text",
                                                                       "operator": "present", "value": None}]},
                            {"status": "conditional", "conditions": [{"fact": "source_inventory.english_text",
                                                                       "operator": "equals", "value": True}]}):
            with self.subTest(declaration=declaration):
                response, chunk = self.fixture()
                response["requirements"][0]["applicability"] = declaration
                self.assertEqual(fonts.materialize_document_font_references(response, chunk), (response, []))
                self.assertTrue(fonts.document_font_scope_errors(response, chunk))
                spec = self.spec()
                spec["requirements"][0]["applicability"] = declaration
                doc = Document(); doc.add_paragraph("English")
                before = etree.tostring(doc.element)
                with self.assertRaisesRegex(ValueError, "incomplete"):
                    fonts.apply_document_font(doc, spec)
                self.assertEqual(etree.tostring(doc.element), before)

    def test_empty_native_exceptions_are_preserved_but_prerequisites_are_not_guessed(self):
        for empty in (None, []):
            response, chunk = self.fixture()
            response["requirements"][0]["applicability"]["exceptions"] = empty
            candidate, audit = fonts.materialize_document_font_references(response, chunk)
            self.assertEqual(len(audit), 1)
            self.assertIn("exceptions", candidate["requirements"][0]["applicability"])
            self.assertEqual(fonts.document_font_scope_errors(candidate, chunk), [])
        for target in ("response", "catalog"):
            response, chunk = self.fixture()
            item = response["requirements"][0] if target == "response" else chunk["rule_spec"]["requirements"][0]
            item["input_prerequisites"] = [{"kind": "runtime", "key": "runtime.unknown", "required": True}]
            self.assertEqual(fonts.materialize_document_font_references(response, chunk), (response, []))

    def test_catalog_cannot_drop_or_change_source_condition(self):
        for declaration in (None, {"status": "always"}, {"status": "conditional", "conditions": []}):
            response, chunk = self.fixture()
            chunk["rule_spec"]["requirements"][0]["applicability"] = declaration
            self.assertEqual(fonts.materialize_document_font_references(response, chunk), (response, []))

    def test_unconditional_sources_do_not_gain_an_invented_condition(self):
        for source in ("所有英文采用Arial字体。", "全文西文使用Calibri字体", "论文中的英文采用Helvetica字体"):
            with self.subTest(source=source):
                response, chunk = self.fixture()
                font = fonts.compile_document_latin_font(source)
                clause = chunk["clauses"][0]
                clause["text"] = source
                clause["source_span"].update(text=source, end_offset=len(source),
                                             source_sha256=hashlib.sha256(source.encode()).hexdigest())
                chunk["evidence_context"][clause["evidence_ids"][0]]["text"] = source
                for rule in chunk["rule_spec"]["requirements"]:
                    rule.update(source_text=source, properties={"font": {"latin": font}})
                    rule.pop("applicability")
                item = response["requirements"][0]
                item["properties"] = {"font": {"latin": font}}
                # A condition imported from another source is not accepted.
                self.assertEqual(fonts.materialize_document_font_references(response, chunk), (response, []))
                item.pop("applicability")
                candidate, audit = fonts.materialize_document_font_references(response, chunk)
                self.assertEqual(len(audit), 1)
                self.assertEqual(fonts.document_font_scope_errors(candidate, chunk), [])
                self.assertFalse(any("applicability" in r for r in candidate["requirements"]))
                self.assertEqual(fonts.document_font_policy(chunk["rule_spec"]), font)

    def test_rule_catalog_compiles_conditions_without_clause_id_or_font_special_cases(self):
        for font in ("Times New Roman", "Arial", "Calibri", "Cambria", "Helvetica"):
            response, chunk = self.fixture()
            clause = chunk["clauses"][0]
            clause.update(text=f"论文中出现英文时必须采用{font}字体", context_before="", context_after="")
            spec, questions, conflicts = engine.build_rule_result(Path("other-school.docx"), [clause])
            self.assertEqual(questions, [])
            self.assertEqual(conflicts, [])
            self.assertEqual({r["role"] for r in spec["requirements"]}, set(fonts.TEXT_FONT_ROLES))
            self.assertTrue(all(r["applicability"] == fonts.compile_document_font_applicability(clause["text"])
                                for r in spec["requirements"]))
            self.assertEqual(fonts.document_font_policy(spec), font)

    def test_missing_inventory_fact_stays_unknown_and_zero_runs_is_not_global_approval(self):
        from applicability import evaluate_applicability
        declaration = self.spec()["requirements"][0]["applicability"]
        self.assertEqual(evaluate_applicability(declaration, source_inventory={})["result"], "unknown")
        self.assertEqual(evaluate_applicability(declaration, source_inventory={"english_text": True})["result"], "true")
        doc = Document(); doc.add_paragraph("只有中文和123")
        before = etree.tostring(doc.element)
        receipt = fonts.apply_document_font(doc, self.spec())
        self.assertEqual(receipt["formatted_run_count"], 0)
        self.assertFalse(receipt["applicability_evidence"]["global_source_inventory_inferred"])
        self.assertEqual(etree.tostring(doc.element), before)

    def test_capability_preflight_does_not_default_unknown_inventory_to_true(self):
        import capability_planner
        root = Path(__file__).resolve().parents[1]
        registry = json.loads((root / "resources/backend-capabilities.default.json").read_text())
        spec = self.spec()
        missing = capability_planner.plan_capabilities(spec, registry, "full", source_inventory={})
        observed = capability_planner.plan_capabilities(spec, registry, "full",
                                                       source_inventory={"english_text": True})
        self.assertEqual(missing["status"], "blocked")
        self.assertFalse(missing["execution_ready"])
        self.assertTrue(all(r["applicability_evaluation"]["result"] == "unknown"
                            for r in missing["requirements"]))
        self.assertTrue(all(r["applicability_evaluation"]["result"] == "true"
                            for r in observed["requirements"]))
        self.assertTrue(all(r["category"] == "runtime_manual_unverifiable"
                            for r in missing["requirements"]))

    def test_stale_missing_conflicting_and_conditional_inputs_cannot_be_completed(self):
        base, chunk = self.fixture()
        cases = []
        stale = copy.deepcopy(chunk); stale["clauses"][0]["source_span"]["source_sha256"] = "0" * 64
        cases.append((base, stale))
        partial = copy.deepcopy(chunk); partial["rule_spec"]["requirements"].pop()
        cases.append((base, partial))
        conflict = copy.deepcopy(base); conflict["requirements"][0]["properties"]["font"]["latin"] = "Arial"
        cases.append((conflict, chunk))
        duplicate = copy.deepcopy(base); duplicate["requirements"].append(copy.deepcopy(duplicate["requirements"][0]))
        cases.append((duplicate, chunk))
        conditional = copy.deepcopy(base); conditional["requirements"][0]["applicability"] = {"conditions": []}
        cases.append((conditional, chunk))
        manual = copy.deepcopy(base); manual["clause_reviews"][0]["classification"] = "unresolved"
        cases.append((manual, chunk))
        for response, current in cases:
            with self.subTest(case=cases.index((response, current))):
                self.assertEqual(fonts.materialize_document_font_references(response, current), (response, []))

    def test_parser_does_not_turn_examples_local_rules_or_alternatives_into_global_rules(self):
        for text in ("正文英文使用Times New Roman字体", "如果需要，论文中出现英文时需要使用Times New Roman字体",
                     "示例：论文中出现英文时需要使用Times New Roman字体", "所有英文使用Times New Roman或Arial字体",
                     "所有英文使用Times New Roman or Arial", "全文英文使用Unknown Font字体",
                     "论文中出现英文时使用 Arial，并由导师签字确认"):
            with self.subTest(text=text):
                self.assertIsNone(fonts.compile_document_latin_font(text))
        self.assertEqual(fonts.compile_document_latin_font("所有英文采用Arial字体。"), "Arial")

    def test_malformed_candidate_payload_is_reported_not_crashed_or_completed(self):
        for properties in (None, [], {"font": None}):
            response, chunk = self.fixture()
            response["requirements"][0]["properties"] = properties
            self.assertTrue(fonts.document_font_scope_errors(response, chunk))
        response, chunk = self.fixture()
        response["requirements"][0]["evidence_ids"] = ["foreign-evidence"]
        self.assertEqual(fonts.materialize_document_font_references(response, chunk), (response, []))

    def test_generic_footnotes_and_mixed_cjk_runs_preserve_east_asian_font(self):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
        doc = Document()
        run = doc.add_paragraph("中文 English").runs[0]
        run.font.name = "Arial"
        rf = run._r.get_or_add_rPr().find(qn("w:rFonts"))
        rf.set(qn("w:eastAsia"), "SimSun")
        rf.set(qn("w:asciiTheme"), "minorHAnsi")
        note = etree.Element(qn("w:footnotes"), nsmap={"w": fonts.W})
        item = etree.SubElement(note, qn("w:footnote")); item.set(qn("w:id"), "1")
        p = etree.SubElement(item, qn("w:p")); r = etree.SubElement(p, qn("w:r"))
        etree.SubElement(r, qn("w:t")).text = "Footnote English"
        part = Part(PackURI("/word/footnotes.xml"), CONTENT_TYPE.WML_FOOTNOTES, etree.tostring(note), doc.part.package)
        doc.part.relate_to(part, RELATIONSHIP_TYPE.FOOTNOTES)
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "with-notes.docx"
            fonts.apply_document_font(doc, self.spec()); doc.save(path)
            self.assertEqual(fonts.audit_document_font(path, self.spec()), [])
            self.assertEqual(rf.get(qn("w:eastAsia")), "SimSun")
            self.assertIsNone(rf.get(qn("w:asciiTheme")))
            self.assertEqual(run.text, "中文 English")

    def test_final_format_comparison_detects_serialized_run_override(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); doc = Document(); doc.add_paragraph("English").runs[0].font.name = "Arial"
            path = td / "final.docx"; doc.save(path)
            spec = td / "spec.json"; spec.write_text(json.dumps(self.spec()))
            mappings = td / "map.json"; mappings.write_text("{}")
            report = lambda: post_audit.build_report(path, path, spec, mappings, mappings, requirements_only=True)
            self.assertEqual(report()["status"], "failed")
            fonts.apply_document_font(doc, self.spec()); doc.save(path)
            self.assertEqual(report()["status"], "passed")
            doc.paragraphs[0].runs[0].font.name = "Arial"; doc.save(path)
            self.assertEqual(report()["status"], "failed")

    @staticmethod
    def spec():
        _, chunk = DocumentTextFontTests.fixture()
        return {"requirements": copy.deepcopy(chunk["rule_spec"]["requirements"])}

    def test_role_style_is_not_run_evidence_and_all_editable_stories_are_formatted(self):
        spec = self.spec(); doc = Document()
        chinese = doc.add_paragraph("中文不改").runs[0]
        chinese.font.name = "SimSun"
        original_chinese = etree.tostring(chinese._r)
        doc.add_paragraph("Body English").runs[0].font.name = "Arial"
        doc.add_heading("English title", 1)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table English"
        doc.sections[0].header.paragraphs[0].text = "Header English"
        doc.sections[0].footer.paragraphs[0].text = "Footer English"
        # Text boxes and hyperlinked runs are not visible to Paragraph.runs.
        holder = doc.add_paragraph()._p
        box = OxmlElement("w:txbxContent"); p = OxmlElement("w:p"); run = OxmlElement("w:r")
        text = OxmlElement("w:t"); text.text = "Textbox English"; run.append(text); p.append(run); box.append(p); holder.append(box)
        link = OxmlElement("w:hyperlink"); link_run = OxmlElement("w:r"); t = OxmlElement("w:t")
        t.text = "Hyperlink English"; link_run.append(t); link.append(link_run); holder.append(link)
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "draft.docx"; doc.save(path)
            self.assertTrue(fonts.audit_document_font(path, spec))
            receipt = fonts.apply_document_font(doc, spec); doc.save(path)
            self.assertGreaterEqual(receipt["formatted_run_count"], 7)
            self.assertEqual(fonts.audit_document_font(path, spec), [])
            self.assertEqual(etree.tostring(chinese._r), original_chinese)
            check = Document(path)
            check.paragraphs[1].runs[0].font.name = "Arial"; check.save(path)
            self.assertTrue(any(f["failure_type"] == "document_font_run_mismatch"
                                for f in fonts.audit_document_font(path, spec)))

    def test_incomplete_catalog_and_local_exceptions_fail_closed(self):
        spec = self.spec(); spec["requirements"].pop()
        with self.assertRaisesRegex(ValueError, "incomplete"):
            fonts.document_font_policy(spec)
        spec = self.spec(); spec["requirements"].append({"role": "abstract_body_en", "properties": {"font": {"latin": "Arial"}}})
        with self.assertRaisesRegex(ValueError, "conflict"):
            fonts.apply_document_font(Document(), spec)

    def test_math_and_drawing_text_are_not_falsely_verified_as_word_runs(self):
        doc = Document(); p = doc.add_paragraph()._p
        for ns in (fonts.M, fonts.A):
            t = etree.Element(f"{{{ns}}}t"); t.text = "English"; p.append(t)
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "draft.docx"
            fonts.apply_document_font(doc, self.spec()); doc.save(path)
            findings = fonts.audit_document_font(path, self.spec())
            self.assertEqual(len(findings), 2)
            self.assertTrue(all(f["failure_type"] == "document_font_not_observable" for f in findings))

    def test_real_candidate_and_merge_preserve_complete_source_scope(self):
        helper = bridge_fixtures.HostAgentBridgeTests()
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = helper._packet(Path(td), source="论文中出现英文时需要使用Times New Roman字体", contract_version="3.0")
            catalog, _, _ = engine.build_rule_result(Path("dynamic.docx"), chunk["clauses"])
            evidence = {"evidence": list(chunk["evidence_context"].values())}
            request = engine.build_llm_request([], chunk["clauses"], evidence, catalog, "full", contract_version="3.0",
                                             runtime_context={"code_fingerprint_sha256": "f" * 64})
            request["case_id"] = "standalone"
            request = attach_request_provenance(request, source_sha256="a" * 64, evidence_doc=evidence,
                                                clauses=chunk["clauses"], run_id="run-font-test")
            review_dir = Path(td) / "complete-catalog"
            engine.prepare_host_agent_review_packets(request, chunk["clauses"], evidence, "a" * 64, review_dir, chunk_size=1)
            chunk = json.loads((review_dir / "llm-request-chunks.json").read_text())[0]
            rule = next(r for r in chunk["rule_spec"]["requirements"] if r["role"] == "body_text")
            raw = {"contract_version": "3.0", "provenance": chunk["provenance"], "requirements": [{
                "existing_requirement_id": rule["id"], "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "role": rule["role"], "properties": rule["properties"],
                "applicability": copy.deepcopy(rule["applicability"]),
                "confidence": .98, "reason": "Reuse one role."}], "clause_reviews": [{
                    "clause_id": "C1", "classification": "covered", "reason": "Font rule.",
                    "normative_basis": "explicit_normative_text",
                    "obligations": [{"id": "font", "status": "covered", "reason": "English font."}]}],
                   "unsupported_items": [], "reported_conflicts": []}
            candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
            self.assertEqual(len(candidate["requirements"]), len(fonts.TEXT_FONT_ROLES))
            self.assertEqual(len(audit["document_font_projections"]), 1)
            self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
            spec, _, _ = engine.merge_llm_primary(Path("dynamic.docx"), chunk["rule_spec"], chunk["clauses"], candidate, {"E1"})
            self.assertEqual(fonts.document_font_policy(spec), "Times New Roman")
            self.assertTrue(all(r["applicability"] == rule["applicability"] for r in spec["requirements"]))
            doc = Document(); doc.add_paragraph("中文 English"); doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table English"
            receipt = fonts.apply_document_font(doc, spec)
            path = Path(td) / "merged.docx"; doc.save(path)
            self.assertEqual(receipt["formatted_run_count"], 2)
            self.assertEqual(fonts.audit_document_font(path, spec), [])
