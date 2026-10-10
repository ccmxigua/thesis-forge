from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from draft_scorecard import (append_scorecard, attach_rendered_output_audit,
                             audit_scorecard, build_scorecard, scorecard_lines,
                             validate_bound_scorecard)
from native_semantic_review import (NativeSemanticReviewError,
    validate_draft_dispute_envelope, validate_obligation_coverage_response)
from thesis_format_pipeline import enforce_obligation_review_output_policy


class DraftScorecardTests(unittest.TestCase):
    def setUp(self):
        self.binding = {"run_id": "current-run", "case_id": "case-test",
                        "source_sha256": "a" * 64, "format_spec_sha256": "b" * 64}
        self.audit = {"missing_count": 0, "unexpected_count": 0, "duplicate_count": 0,
                      "expected_receipt_ids": ["r1", "r2", "r3"], "receipts": [
            {"receipt_id": "r1", "status": "verified", "role": "body",
             "property_path": "font.size_pt", "actual": 12, "expected": 12},
            {"receipt_id": "r2", "status": "failed", "role": "keywords",
             "property_path": "separator", "actual": "chinese_comma", "expected": "semicolon"},
            {"receipt_id": "r3", "status": "unverified", "role": "abstract",
             "property_path": "third_person", "reason": "尚未人工确认"},
        ]}
        self.manual = [{"source_code": "source-dispute", "clause_ids": ["source-clause"],
                        "source_text": "来源存在歧义", "reason": "待核验", "action": "人工检查"}]
        self.findings = [{"role": "keywords", "property": "separator",
                          "template_value": "，", "required_value": "；"}]

    def card(self):
        return build_scorecard(self.binding, self.audit, self.manual, self.findings)

    def test_equal_observed_check_scores_not_confidence_or_release(self):
        card = self.card()
        self.assertIsNone(card["score"])
        self.assertEqual(card["observed_checks_verified_percent"], 20)
        self.assertEqual(card["verified_count"], 1)
        self.assertEqual(card["status_counts"], {
            "failed": 2, "unverified": 1, "pending": 1, "verified": 1,
        })
        self.assertEqual(card["failed_count"], 2)
        self.assertEqual(card["unverified_count"], 1)
        self.assertEqual(card["pending_count"], 1)
        self.assertEqual(card["diagnostic_item_count"], 5)
        self.assertEqual(card["evaluation_unit_count"], 3)
        self.assertIsNone(card["unique_source_obligation_count"])
        self.assertFalse(card["submission_ready"])
        self.assertFalse(card["coverage_complete"])
        self.assertTrue(card["human_check_required"])
        self.assertIn("semicolon", "\n".join(scorecard_lines(card)))
        self.assertIn("chinese_comma", "\n".join(scorecard_lines(card)))
        self.assertIn("来源存在歧义", "\n".join(scorecard_lines(card)))
        lines = "\n".join(scorecard_lines(card))
        self.assertIn("诊断条目5项（不是独立义务数）", lines)
        self.assertIn("已证实未满足2项", lines)
        self.assertIn("未核验1项", lines)
        self.assertIn("待人工核验1项", lines)
        by_id = {item["detail"].get("receipt_id"): item for item in card["entries"]
                 if item["kind"] == "property"}
        self.assertEqual(by_id["r1"]["score"], 100)
        self.assertEqual(by_id["r2"]["score"], 0)
        self.assertIsNone(by_id["r3"]["score"])
        self.assertIn("未评分（未核验）", lines)
        self.assertIn("未评分（待人工核验）", lines)
        self.assertNotIn("0/100 未核验", lines)
        self.assertNotIn("0/100 待人工核验", lines)

    def test_unverified_receipt_reason_is_visible_in_review_text(self):
        self.audit["receipts"][2]["status_reason"] = "来源封面实例尚未绑定"
        lines = "\n".join(scorecard_lines(self.card()))
        self.assertIn("来源封面实例尚未绑定", lines)

    def test_unobservable_drawing_font_is_capability_unknown_not_failure(self):
        finding = {"role": "all_text", "property": "font.latin.non_word_text",
                   "failure_type": "document_font_not_observable",
                   "reason": "Math/drawing Latin text needs a separate font executor and verifier."}
        card = build_scorecard(self.binding, self.audit, [], [finding])
        entry = next(item for item in card["entries"] if item["detail"].get("failure_type"))
        self.assertEqual(entry["kind"], "capability")
        self.assertEqual(entry["status"], "unverified")
        self.assertIsNone(entry["score"])
        self.assertEqual(card["failed_count"], 1)  # only the separate proven keyword mismatch
        self.assertIn("Math/drawing", "\n".join(scorecard_lines(card)))

    def test_audit_rejects_numeric_score_on_unknown_or_pending_items(self):
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "draft.docx"
            card = self.card()
            doc = Document()
            append_scorecard(doc, card)
            doc.save(output)
            self.assertTrue(audit_scorecard(output, card)["valid"])
            for status in ("unverified", "pending"):
                tampered = copy.deepcopy(card)
                row = next(item for item in tampered["entries"] if item["status"] == status)
                row["score"] = 0
                with self.subTest(status=status):
                    self.assertFalse(audit_scorecard(output, tampered)["valid"])

    def test_legacy_v3_scorecard_without_scope_assessment_remains_nonrelease_auditable(self):
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "legacy-draft.docx"
            card = self.card()
            card.pop("obligation_assessment", None)
            doc = Document()
            append_scorecard(doc, card)
            doc.save(output)
            audit = audit_scorecard(output, card)
            self.assertTrue(audit["valid"])
            self.assertEqual(audit["obligation_assessment_status"], "legacy_not_recorded")
            self.assertFalse(audit["submission_ready"])
            self.assertNotIn("对象级来源义务：未评估。", "\n".join(scorecard_lines(card)))

    def test_rendered_audit_reconciles_visible_scorecard_then_requires_final_rerender(self):
        import hashlib
        import json
        from semantic_contract import sha256_json

        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            output = root / "draft.docx"
            scorecard_path = root / "scorecard.json"
            card = self.card()
            property_entry = next(item for item in card["entries"]
                                  if item["kind"] == "property" and item["detail"].get("receipt_id") == "r1")
            property_entry["detail"].update({
                "requirement_id": "R1", "role": "body", "property_path": "font.cjk",
                "target_locator": "style:Body", "expected": "SimSun",
            })
            doc = Document()
            append_scorecard(doc, card)
            doc.save(output)
            scorecard_path.write_text(json.dumps(card), encoding="utf-8")
            before = hashlib.sha256(output.read_bytes()).hexdigest()
            property_receipts = {"expected_receipt_ids": ["r1"], "receipts": [{
                "receipt_id": "r1", "requirement_id": "R1", "role": "body",
                "property_path": "font.cjk", "target_locator": "style:Body",
                "expected": "SimSun", "actual": "SimSun", "status": "verified",
                "serialized_docx_sha256": before,
            }]}

            def report_for_current_docx():
                digest = hashlib.sha256(output.read_bytes()).hexdigest()
                report = {
                    "protocol": "rendered_format_audit_v1", "docx_sha256": digest,
                    "final_docx": {"sha256": digest}, "pdf_sha256": "a" * 64,
                    "findings": [
                        {"code": "rendered_pdf_font_mismatch", "expected": {
                            "role": "body", "style_name": "Body", "script": "cjk",
                            "name": "SimSun", "requirement_ids": ["R1"],
                        }},
                        {"code": "drawing_not_found_in_rendered_pdf", "drawing_index": 1,
                         "candidate_count": 0, "media_sha256": "b" * 64},
                    ],
                    "submission_ready": False, "field_refresh_claimed": False,
                }
                report["audit_sha256"] = sha256_json(report)
                return report

            def report_bindings(report):
                finding = report["findings"][0]
                return {sha256_json(finding): ["R1"]}

            first_report = report_for_current_docx()
            with self.assertRaisesRegex(ValueError, "rerender and re-audit"):
                attach_rendered_output_audit(
                    scorecard_path, first_report, output, property_receipts,
                    rendered_requirement_bindings=report_bindings(first_report),
                )
            self.assertNotEqual(hashlib.sha256(output.read_bytes()).hexdigest(), before)
            intermediate = json.loads(scorecard_path.read_text())
            self.assertNotIn("rendered_output_audit", intermediate)
            self.assertTrue(any(p.text.startswith("当前生成稿实例核验（")
                                for p in Document(output).paragraphs))

            # The visible status block changes the serialized DOCX; a new
            # property pass must bind receipts to those exact new bytes.
            property_receipts["receipts"][0]["serialized_docx_sha256"] = hashlib.sha256(
                output.read_bytes()).hexdigest()
            final_report = report_for_current_docx()
            attached = attach_rendered_output_audit(
                scorecard_path, final_report, output, property_receipts,
                rendered_requirement_bindings=report_bindings(final_report),
            )
            current_property = next(item for item in attached["entries"]
                                    if item["kind"] == "property" and item["detail"].get("receipt_id") == "r1")
            self.assertEqual(current_property["historical_status"], "verified")
            self.assertEqual(current_property["status"], "failed")
            output_only = [item for item in attached["entries"]
                           if item["kind"] == "rendered_output_instance"]
            self.assertEqual(len(output_only), 1)
            self.assertEqual(output_only[0]["status"], "unverified")
            self.assertTrue(audit_scorecard(output, attached)["valid"])

            tampered_report = copy.deepcopy(attached)
            tampered_report["rendered_output_audit"]["report"]["findings"].append(
                {"code": "tampered_render_finding"})
            self.assertFalse(audit_scorecard(output, tampered_report)["valid"])

    def test_rendered_audit_does_not_bind_wrong_requirement_or_stale_docx(self):
        import hashlib
        import json
        from semantic_contract import sha256_json

        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            output = root / "draft.docx"
            scorecard_path = root / "scorecard.json"
            card = self.card()
            entry = next(item for item in card["entries"]
                         if item["kind"] == "property" and item["detail"].get("receipt_id") == "r1")
            entry["detail"].update({"requirement_id": "R1", "role": "body",
                                    "property_path": "font.cjk", "target_locator": "style:Body",
                                    "expected": "SimSun"})
            doc = Document()
            append_scorecard(doc, card)
            doc.save(output)
            scorecard_path.write_text(json.dumps(card), encoding="utf-8")
            digest = hashlib.sha256(output.read_bytes()).hexdigest()
            report = {
                "protocol": "rendered_format_audit_v1", "docx_sha256": digest,
                "final_docx": {"sha256": digest}, "pdf_sha256": "a" * 64,
                "findings": [{"code": "rendered_pdf_font_mismatch", "expected": {
                    "role": "body", "style_name": "Body", "script": "cjk",
                    "name": "SimSun", "requirement_ids": ["R-other"],
                }}],
                "submission_ready": False, "field_refresh_claimed": False,
            }
            report["audit_sha256"] = sha256_json(report)
            receipt_audit = {"expected_receipt_ids": ["r1"], "receipts": [{
                "receipt_id": "r1", "requirement_id": "R1", "role": "body",
                "property_path": "font.cjk", "target_locator": "style:Body",
                "expected": "SimSun", "actual": "SimSun", "status": "verified",
                "serialized_docx_sha256": digest,
            }]}
            with self.assertRaisesRegex(ValueError, "rerender and re-audit"):
                attach_rendered_output_audit(
                    scorecard_path, report, output, receipt_audit,
                    rendered_requirement_bindings={sha256_json(report["findings"][0]): ["R1"]},
                )
            updated = json.loads(scorecard_path.read_text())
            row = next(item for item in updated["entries"]
                       if item["kind"] == "property" and item["detail"].get("receipt_id") == "r1")
            self.assertEqual(row["status"], "verified")
            self.assertTrue(any(item["kind"] == "rendered_output_instance"
                                and item["status"] == "failed" for item in updated["entries"]))
            stale = copy.deepcopy(report)
            stale["final_docx"]["sha256"] = "c" * 64
            stale["audit_sha256"] = sha256_json({k: v for k, v in stale.items() if k != "audit_sha256"})
            with self.assertRaisesRegex(ValueError, "DOCX hash"):
                attach_rendered_output_audit(
                    scorecard_path, stale, output, receipt_audit,
                    rendered_requirement_bindings={sha256_json(stale["findings"][0]): ["R1"]},
                )

    def test_even_perfect_score_requires_final_human_review(self):
        for receipt in self.audit["receipts"]:
            receipt["status"] = "verified"
        card = build_scorecard(self.binding, self.audit, [], [])
        self.assertIsNone(card["score"])
        self.assertEqual(card["observed_checks_verified_percent"], 100)
        self.assertTrue(card["human_check_required"])
        self.assertFalse(card["submission_ready"])

    def test_empty_inventory_is_unscored_not_perfect(self):
        self.audit.update(receipts=[], expected_receipt_ids=[])
        self.assertIsNone(build_scorecard(self.binding, self.audit, [], [])["score"])

    def test_missing_duplicate_unknown_receipts_fail_closed(self):
        mutations = [
            lambda x: x.update(missing_count=1),
            lambda x: x.pop("missing_count"),
            lambda x: x["receipts"].pop(),
            lambda x: x["receipts"].append(copy.deepcopy(x["receipts"][0])),
            lambda x: x["receipts"][0].update(status="model_says_passed"),
            lambda x: x["expected_receipt_ids"].append({"bad": "id"}),
        ]
        for mutate in mutations:
            with self.subTest(mutation=mutate):
                audit = copy.deepcopy(self.audit)
                mutate(audit)
                with self.assertRaises(ValueError):
                    build_scorecard(self.binding, audit, self.manual, self.findings)

    def test_capability_warning_and_blocker_both_preserved(self):
        findings = [{"code": "cap.warning", "blocking": False},
                    {"code": "cap.backend", "blocking": True}]
        card = build_scorecard(self.binding, self.audit, [], [], findings)
        self.assertEqual(len([x for x in card["entries"] if x["kind"] == "capability"]), 2)
        self.assertIsNone(card["score"])
        self.assertEqual(card["observed_checks_verified_percent"], 20)
        self.assertEqual(findings, [x["detail"] for x in card["entries"] if x["kind"] == "capability"])

    def test_docx_serialization_bound_exact_visible_text(self):
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "draft.docx"
            doc = Document()
            doc.add_paragraph("保留原始正文")
            card = self.card()
            append_scorecard(doc, card)
            doc.save(output)
            self.assertTrue(validate_bound_scorecard(card, binding=self.binding,
                receipt_audit=self.audit, manual_items=self.manual, findings=self.findings, output=output))
            self.assertIn("保留原始正文", [p.text for p in Document(output).paragraphs])
            forged = copy.deepcopy(card)
            forged["score"] = 100
            self.assertFalse(validate_bound_scorecard(forged, binding=self.binding,
                receipt_audit=self.audit, manual_items=self.manual, findings=self.findings, output=output))
            for field in ("run_id", "case_id", "source_sha256", "format_spec_sha256"):
                binding = copy.deepcopy(self.binding)
                binding[field] = "f" * 64
                with self.subTest(field=field):
                    self.assertFalse(validate_bound_scorecard(card, binding=binding,
                        receipt_audit=self.audit, manual_items=self.manual, findings=self.findings, output=output))
            self.assertFalse(validate_bound_scorecard(card, binding=self.binding,
                receipt_audit=self.audit, manual_items=[], findings=self.findings, output=output))

    def test_hidden_missing_duplicate_extra_or_black_score_markers_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "draft.docx"
            card = self.card()
            for tamper in ("hidden", "missing", "duplicate", "extra", "black"):
                with self.subTest(tamper=tamper):
                    doc = Document()
                    append_scorecard(doc, card)
                    paragraph = doc.paragraphs[1]
                    if tamper == "hidden": paragraph.runs[0].font.hidden = True
                    if tamper == "black": paragraph.runs[0].font.color.rgb = RGBColor(0, 0, 0)
                    if tamper == "missing": paragraph._p.getparent().remove(paragraph._p)
                    if tamper == "duplicate": doc.add_paragraph(paragraph.text)
                    if tamper == "extra": doc.add_paragraph("【SC-forged｜自动评分】100/100")
                    doc.save(output)
                    self.assertFalse(audit_scorecard(output, card)["valid"])

    def test_inherited_word_styles_are_resolved_and_table_duplicates_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "draft.docx"
            card = self.card()
            doc = Document()
            append_scorecard(doc, card)
            for paragraph in doc.paragraphs:
                for run in paragraph.runs:
                    if run._r.rPr is not None:
                        run._r.remove(run._r.rPr)
            doc.save(output)
            self.assertTrue(audit_scorecard(output, card)["valid"])
            doc.styles[doc.paragraphs[0].style.name].font.hidden = True
            doc.save(output)
            self.assertFalse(audit_scorecard(output, card)["valid"])
            doc.styles[doc.paragraphs[0].style.name].font.hidden = False
            doc.add_table(1, 1).cell(0, 0).text = doc.paragraphs[1].text
            doc.save(output)
            self.assertFalse(audit_scorecard(output, card)["valid"])


class DraftCoverageDisputeTests(unittest.TestCase):
    def setUp(self):
        self.source = "学校可以检索论文，可以复制、保存和汇编论文。"
        self.check = {"check_id": "external-random", "document_text": self.source,
            "review_context": {"classification": "external_compliance", "requires_requirement": False,
                "linked_requirements": [], "primary_obligations": [
                    {"id": "broad-action", "status": "unverifiable", "reason": "人工核验"}]}}
        self.response = {"results": [{"check_id": "external-random", "verdict": "incomplete",
            "machine_obligation_ids": [],
            "rationale": "义务拆分不一致，需要核验", "evidence_quotes": [self.source],
            "identified_obligations": [
                {"source_quote": quote, "disposition": "unrepresented", "requirement_refs": [],
                 "primary_obligation_id": "broad-action"}
                for quote in ("检索论文", "复制", "保存和汇编论文")]}]}

    def test_draft_keeps_dispute_without_inventing_requirement_or_reclassifying(self):
        before = copy.deepcopy(self.response)
        results = validate_obligation_coverage_response(self.response, [self.check], allow_draft_disputes=True)
        self.assertEqual(self.response, before)
        self.assertEqual(results[0]["verdict"], "incomplete")
        self.assertEqual(enforce_obligation_review_output_policy(results, output_policy="review_draft"), ["external-random"])
        with self.assertRaises(ValueError):
            enforce_obligation_review_output_policy(results, output_policy="submission")
        with self.assertRaises(NativeSemanticReviewError):
            validate_obligation_coverage_response(self.response, [self.check])

    def test_draft_cannot_bypass_quote_reference_schema_or_known_fact_checks(self):
        for kind in ("quote", "reference", "check_id", "duplicate", "known_facts", "status"):
            with self.subTest(kind=kind):
                response, check = copy.deepcopy(self.response), copy.deepcopy(self.check)
                row = response["results"][0]
                if kind == "quote": row["identified_obligations"][0]["source_quote"] = "源中没有此文字"
                if kind == "reference": row["identified_obligations"][0]["requirement_refs"] = ["foreign"]
                if kind == "check_id": row["check_id"] = "foreign-clause"
                if kind == "duplicate": response["results"].append(copy.deepcopy(row))
                if kind == "known_facts": check["review_context"]["machine_obligation_ids"] = ["local-fact"]
                if kind == "status": check["review_context"]["primary_obligations"][0]["status"] = "covered"
                with self.assertRaises(NativeSemanticReviewError):
                    validate_obligation_coverage_response(response, [check], allow_draft_disputes=True)

    def test_completed_dispute_requires_policy_incomplete_flag_and_nonrelease(self):
        envelope = {"status": "completed_with_disputes", "coverage_complete": False,
                    "submission_ready": False, "results": self.response["results"]}
        request = {"output_policy": "review_draft"}
        validate_draft_dispute_envelope(envelope, request, output_policy="review_draft")
        for key, value in (("status", "completed"), ("coverage_complete", True),
                           ("submission_ready", True), ("results", [])):
            with self.subTest(key=key):
                forged = copy.deepcopy(envelope); forged[key] = value
                with self.assertRaises(ValueError):
                    validate_draft_dispute_envelope(forged, request, output_policy="review_draft")
        with self.assertRaises(ValueError):
            validate_draft_dispute_envelope(envelope, {}, output_policy="review_draft")
        with self.assertRaises(ValueError):
            validate_draft_dispute_envelope(envelope, request, output_policy="submission")


if __name__ == "__main__":
    unittest.main()
