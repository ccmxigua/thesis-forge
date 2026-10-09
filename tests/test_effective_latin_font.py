from __future__ import annotations

import sys
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from apply_format_spec import (  # noqa: E402
    _receipt_semantic_actuals,
    _serialized_run_latin_font,
    effective_style_latin_font,
    set_font,
    style_snapshot,
)


class EffectiveLatinFontTests(unittest.TestCase):
    def test_inherited_theme_font_resolves_and_explicit_style_override_wins(self):
        doc = Document()
        child = doc.styles.add_style("Inherited Thesis Role", WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = doc.styles["Normal"]

        self.assertEqual(effective_style_latin_font(child, doc), "Cambria")
        self.assertIsNone(style_snapshot(child, doc)["font"]["latin"])

        set_font(child, {"latin": "Times New Roman"})
        self.assertEqual(effective_style_latin_font(child, doc), "Times New Roman")
        self.assertEqual(style_snapshot(child, doc)["font"]["latin"], "Times New Roman")

    def test_inconsistent_script_slots_fail_closed(self):
        doc = Document()
        style = doc.styles.add_style("Mixed Thesis Role", WD_STYLE_TYPE.PARAGRAPH)
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), "Times New Roman")
        rfonts.set(qn("w:hAnsi"), "Arial")

        self.assertIsNone(effective_style_latin_font(style, doc))
        # style_snapshot preserves the legacy primary-slot view; executable
        # effective-font evidence uses the fail-closed two-slot resolver.
        self.assertEqual(style_snapshot(style, doc)["font"]["latin"], "Times New Roman")

    def test_serialized_latin_runs_override_inherited_theme_font(self):
        doc = Document()
        style = doc.styles.add_style("Latin Thesis Role", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        paragraph = doc.add_paragraph("References")
        paragraph.style = style
        run = paragraph.runs[0]
        run.font.name = "Times New Roman"

        self.assertEqual(_serialized_run_latin_font(run, paragraph, doc), "Times New Roman")

    def test_role_font_audit_marks_empty_latin_scope_unknown_and_observed_run_concrete(self):
        requirements = [{
            "id": "R_BIB_FONT", "role": "bibliography_heading",
            "properties": {"font": {"latin": "Times New Roman"}},
        }]
        mappings = {"bibliography_heading": {"style_name": "Bibliography Heading"}}
        spec = {"thesis_profile": {}}

        doc = Document()
        style = doc.styles.add_style("Bibliography Heading", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        paragraph = doc.add_paragraph("参考文献")
        paragraph.style = style
        actual, _, observability, font_actuals = _receipt_semantic_actuals(
            doc, spec, mappings, requirements, {}, [], {},
        )
        self.assertNotIn("font", actual.get("bibliography_heading", {}))
        self.assertEqual(font_actuals["bibliography_heading"], "Cambria")
        self.assertFalse(observability["bibliography_heading"]["font.latin"]["observed"])

        body_doc = Document()
        body_style = body_doc.styles.add_style("Thesis Body Text", WD_STYLE_TYPE.PARAGRAPH)
        body_style.base_style = body_doc.styles["Normal"]
        body_doc.add_paragraph("References", style=body_style)
        body_requirements = [{
            "id": "R_BODY_FONT", "role": "body_text",
            "properties": {"font": {"latin": "Times New Roman"}},
        }]
        actual, _, observability, font_actuals = _receipt_semantic_actuals(
            body_doc, spec, {"body_text": {"style_name": "Thesis Body Text"}},
            body_requirements, {}, [], {},
        )
        self.assertTrue(observability["body_text"]["font.latin"]["observed"])
        self.assertEqual(font_actuals["body_text"], "Cambria")


if __name__ == "__main__":
    unittest.main()
