"""Native editable math and actual PDF defects require independent checks."""
from __future__ import annotations

import copy
import contextlib
import io
import json
from pathlib import Path
import shutil
import sys
import tempfile
import subprocess
import unittest
from unittest import mock

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import fitz

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/"scripts"))
import apply_format_spec as afs
import equation_compat as compat
import equation_reflow as reflow
import audit_equation_export as export


def node(tag, value=None):
    n = OxmlElement(tag)
    if value is not None:
        n.set(qn("m:val"), value)
    return n


def sample(*, array=False, grow=None, shape=None, rows=2):
    doc = Document(); p = doc.add_paragraph(); group = node("m:oMathPara")
    math = node("m:oMath"); delimiter = node("m:d"); pr = node("m:dPr")
    pr.extend([node("m:begChr", "{" if array else "["), node("m:endChr", "" if array else "]")])
    if grow is not None:
        pr.append(node("m:grow", grow))
    if shape is not None:
        pr.append(node("m:shp", shape))
    delimiter.append(pr); operand = node("m:e"); container = node("m:eqArr" if array else "m:m")
    for row in range(rows):
        item = node("m:e" if array else "m:mr")
        for col in range(1 if array else 2):
            cell = item if array else node("m:e"); run = node("m:r")
            text = node("m:t"); text.text = f"y{row+1}=x{row+1}+1" if array else str(row*2+col+1)
            run.append(text); cell.append(run)
            if not array:
                item.append(cell)
        container.append(item)
    operand.append(container); delimiter.append(operand); math.append(delimiter); group.append(math); p._p.append(group)
    return doc, p, group


class CompatibilityTests(unittest.TestCase):
    def test_defaults_are_explicit_and_inverse_is_exact_for_matrix_and_piecewise(self):
        for array in [False, True]:
            doc, p, original = sample(array=array); before = compat.canonical(original)
            updated = compat.materialize_growth(original); proof = compat.proof(original, updated, 0)
            self.assertEqual(before, compat.canonical(original))
            self.assertEqual(updated.find(".//"+qn("m:grow")).get(qn("m:val")), "1")
            self.assertEqual(compat.canonical(compat.restore(proof, updated)), before)
            self.assertEqual(compat.canonical(compat.materialize_growth(updated)), compat.canonical(updated))
            source_hashes = afs._equation_math_signature(doc); p._p.replace(original, updated)
            self.assertTrue(reflow.integrity_matches(source_hashes, afs._equation_math_nodes(doc), [proof]))
            for corruption in ["text", "cell", "font", "shape"]:
                altered = copy.deepcopy(updated)
                if corruption == "text":
                    altered.find(".//"+qn("m:t")).text = "WRONG"
                elif corruption == "cell":
                    parent = altered.find(".//"+qn("m:eqArr" if array else "m:mr")); parent.insert(0, parent[-1])
                elif corruption == "font":
                    altered.find(".//"+qn("m:r")).insert(0, node("m:sty", "p"))
                else:
                    altered.find(".//"+qn("m:dPr")).append(node("m:shp", "match"))
                self.assertFalse(compat.validates(proof, altered), corruption)
                with self.assertRaises(ValueError):
                    compat.restore(proof, altered)

    def test_author_settings_and_single_rows_are_preserved(self):
        for grow in ["0", "false", "1"]:
            for shape in [None, "match", "centered"]:
                _, _, source = sample(grow=grow, shape=shape)
                self.assertEqual(compat.canonical(source), compat.canonical(compat.materialize_growth(source)))
        _, _, source = sample(rows=1)
        self.assertEqual(compat.canonical(source), compat.canonical(compat.materialize_growth(source)))
        _, _, source = sample(shape="match")
        updated = compat.materialize_growth(source)
        self.assertEqual([n.tag for n in updated.find(".//"+qn("m:dPr"))][-2:], [qn("m:grow"), qn("m:shp")])

    def test_failed_render_restores_original_and_does_not_publish_growth_proof(self):
        doc, p, _ = sample(); before = compat.canonical(p._p)
        with mock.patch.object(afs, '_equation_render_observation', return_value={
                'findings':[{'code':'equation_geometry_unavailable','requires_manual_review':True}]}):
            result = afs.apply_equation_layout(doc,{})
        self.assertEqual(result['layout_repairs'],[])
        self.assertEqual(result['equation_tab_layout_unresolved'],1)
        self.assertEqual(compat.canonical(p._p),before)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice required")
    def test_formatter_protects_inline_bookmark_and_has_stable_second_pass(self):
        doc, p, source = sample(); bookmark = OxmlElement("w:bookmarkStart")
        bookmark.set(qn("w:id"), "4"); bookmark.set(qn("w:name"), "native_math")
        p._p.insert(0, bookmark); inline = doc.add_paragraph("Inline ")
        m = node("m:oMath"); r = node("m:r"); t = node("m:t"); t.text = "a+b"; r.append(t); m.append(r); inline._p.append(m)
        inline_before = compat.canonical(inline._p); original = afs._equation_math_signature(doc)
        first = afs.apply_equation_layout(doc, {})
        self.assertEqual(first["equation_tab_layout_unresolved"], 0, first)
        self.assertEqual(len(first["layout_repairs"]), 1)
        self.assertTrue(reflow.integrity_matches(original, afs._equation_math_nodes(doc), first["layout_repairs"]))
        self.assertEqual(inline_before, compat.canonical(inline._p))
        self.assertEqual(p._p.find(qn("w:bookmarkStart")).get(qn("w:name")), "native_math")
        before = compat.canonical(doc._element); second = afs.apply_equation_layout(doc, {})
        self.assertEqual(second["layout_repairs"], []); self.assertEqual(before, compat.canonical(doc._element))

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice required")
    def test_five_cli_passes_do_not_drift_after_growth_normalization(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source = root/"input.docx"; sample()[0].save(source)
            original_bytes = source.read_bytes(); spec = root/"spec.json"
            spec.write_text(json.dumps({"schema_version":"1.0", "source_document":"Synthetic compatibility regression",
                "status":"semantic_resolved", "requirements":[], "roles":{"equation":{}}, "equations":{}}))
            outputs = []
            for index in range(5):
                out = root/f"pass-{index}.docx"
                with contextlib.redirect_stdout(io.StringIO()):
                    code = afs.main([str(source),str(spec),str(out),"--out-dir",str(root/f"audit-{index}")])
                self.assertEqual(code, 0)
                outputs.append(out.read_bytes()); source = out
            self.assertTrue(all(data == outputs[0] for data in outputs))
            self.assertEqual((root/"input.docx").read_bytes(), original_bytes)


class ActualPdfTests(unittest.TestCase):
    @unittest.skipUnless(shutil.which('soffice'), 'LibreOffice required')
    def test_stretched_font_delimiter_uses_actual_ink_not_logical_font_box(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);doc,_,_=sample(rows=4);source=root/'source.docx';doc.save(source)
            result=subprocess.run(['soffice','--headless','--norestore','--nodefault',
                '-env:UserInstallation='+(root/'profile').as_uri(),'--convert-to','pdf','--outdir',str(root),str(source)],capture_output=True,timeout=60)
            self.assertEqual(result.returncode,0,result.stderr);pdf=root/'source.pdf'
            with fitz.open(pdf) as rendered:
                self.assertEqual(len(rendered),1);bounds=list(rendered[0].rect)
            regions={'source_docx_sha256':export.sha(source),'pdf_sha256':export.sha(pdf),
                'regions':[{'math_index':0,'delimiter_index':0,'page':1,'bbox':bounds}]}
            report=export.audit(source,pdf,regions)
            self.assertTrue(report['export_checks_passed'],report)

    def make_pdf(self, directory, height):
        path = directory/"export.pdf"; pdf = fitz.open(); page = pdf.new_page(width=240, height=240)
        for row, values in enumerate(["12", "34"]):
            for col, value in enumerate(values):
                page.insert_text((90+col*25, 100+row*24), value, fontsize=12)
        middle = 106; y0 = middle-height/2; y1 = middle+height/2
        for x, sign in [(84, 1), (130, -1)]:
            shape = page.new_shape(); shape.draw_polyline([(x+sign*3,y0),(x,y0),(x,y1),(x+sign*3,y1)])
            shape.finish(color=(0,0,0)); shape.commit()
        pdf.save(path); pdf.close(); return path

    def test_short_vector_delimiters_fail_and_growing_ones_pass_on_real_pdf(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); doc, _, _ = sample(); source = root/"source.docx"; doc.save(source)
            for height, expected in [(10, False), (55, True)]:
                pdf = self.make_pdf(root, height)
                regions = {"source_docx_sha256": export.sha(source), "pdf_sha256": export.sha(pdf),
                           "regions": [{"math_index":0,"delimiter_index":0,"page":1,"bbox":[79,75,137,140]}]}
                report = export.audit(source,pdf,regions)
                self.assertEqual(report["export_checks_passed"], expected, report)
                if not expected:
                    self.assertIn("delimiter_does_not_cover_operand", {f["code"] for f in report["findings"]})
                self.assertFalse(report["whole_document_visual_review"])
                if pdf.exists():
                    pdf.unlink()

    def test_unmapped_or_unbound_regions_and_changed_cell_order_are_blocked(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); doc, _, source_math = sample(); source = root/"source.docx"; doc.save(source)
            pdf = self.make_pdf(root,55)
            self.assertFalse(export.audit(source,pdf)["export_checks_passed"])
            self.assertFalse(export.audit(source,pdf,{"source_docx_sha256":"wrong"})["export_checks_passed"])
            source_math.find(".//"+qn("m:t")).text="9";doc.save(source)
            regions={"source_docx_sha256":export.sha(source),"pdf_sha256":export.sha(pdf),"regions":[{"math_index":0,"delimiter_index":0,"page":1,"bbox":[79,75,137,140]}]}
            report=export.audit(source,pdf,regions)
            self.assertIn("delimiter_operand_text_unverified",{f["code"] for f in report["findings"]})

    def test_actual_broken_tounicode_is_blocked_without_replacing_visible_letters(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);doc=Document();doc.add_paragraph('y');source=root/'source.docx';doc.save(source)
            pdf=fitz.open();page=pdf.new_page();page.insert_font(fontname='fixture',fontbuffer=fitz.Font('cjk').buffer)
            page.insert_text((40,40),'y',fontname='fixture',fontsize=12)
            font=page.get_fonts()[0];xref=int(pdf.xref_get_key(font[0],'ToUnicode')[1].split()[0])
            glyph=page.get_texttrace()[0]['chars'][0][1]
            cmap=f'/CIDInit /ProcSet findresource begin 12 dict begin begincmap /CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def /CMapName /Test def /CMapType 2 def 1 begincodespacerange <0000> <FFFF> endcodespacerange 1 beginbfchar <{glyph:04X}> <FFFD> endbfchar endcmap CMapName currentdict /CMap defineresource pop end end'
            pdf.update_stream(xref,cmap.encode());path=root/'broken.pdf';pdf.save(path);pdf.close()
            report=export.audit(source,path)
            self.assertIn('pdf_unicode_mapping_unresolved',{f['code'] for f in report['findings']})


if __name__ == "__main__":
    unittest.main()
