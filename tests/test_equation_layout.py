from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Twips
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from apply_format_spec import (
    _equation_tab_position_twips, apply_equation_layout, audit_equation_layout,
)
from docx_semantics import is_display_equation


def add_math(paragraph, label="", *, direct=False):
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x+y=1"
    run.append(text)
    math.append(run)
    if direct:
        paragraph._p.append(math)
    else:
        group = OxmlElement("m:oMathPara")
        group.append(math)
        paragraph._p.append(group)
    if label:
        paragraph.add_run(label)
    return paragraph


def canonical(node):
    return etree.tostring(node, method="c14n")


def tab_positions(paragraph):
    return {tab.get(qn("w:val")): int(tab.get(qn("w:pos")))
            for tab in paragraph._p.xpath("./w:pPr/w:tabs/w:tab")}


class EquationLayoutTests(unittest.TestCase):
    def test_default_center_preserves_inline_and_mixed_math(self):
        doc = Document()
        inline = doc.add_paragraph("before ")
        run = inline.add_run()
        run._r.append(OxmlElement("m:oMath"))
        inline.add_run(" after")
        mixed = add_math(doc.add_paragraph("prose "))
        direct_inline = add_math(doc.add_paragraph("inline "), direct=True)
        display = add_math(doc.add_paragraph())
        direct_display = add_math(doc.add_paragraph(), direct=True)
        protected = [canonical(p._p) for p in (inline, mixed, direct_inline)]
        counts = apply_equation_layout(doc, {})
        self.assertEqual(counts["equations"], 2)
        self.assertEqual(protected, [canonical(p._p) for p in (inline, mixed, direct_inline)])
        self.assertEqual(display.text, "")
        self.assertEqual(display.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(direct_display.text, "\t")
        self.assertEqual(tab_positions(direct_display)["center"], 4320)
        self.assertTrue(all(check["status"] == "passed" for check in counts["geometry_checks"]))
        self.assertEqual(audit_equation_layout(doc, {}), [])

    def test_multiline_math_and_review_anchors_survive_serialization(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        group = parse_xml('''<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
          <m:oMathParaPr><m:jc m:val="left"/></m:oMathParaPr><m:oMath><m:eqArr>
          <m:e><m:r><m:rPr><m:aln/></m:rPr><m:t>a=</m:t></m:r>
          <m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:e>
          <m:e><m:r><m:rPr><m:aln/></m:rPr><m:t>=b</m:t></m:r>
          <m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>n</m:t></m:r></m:sub></m:sSub></m:e>
          </m:eqArr></m:oMath></m:oMathPara>''')
        for tag in ("bookmarkStart", "commentRangeStart"):
            node = OxmlElement(f"w:{tag}")
            node.set(qn("w:id"), "7")
            if tag == "bookmarkStart":
                node.set(qn("w:name"), "eq-multiline")
            paragraph._p.append(node)
        paragraph._p.append(group)
        for tag in ("commentRangeEnd", "bookmarkEnd"):
            node = OxmlElement(f"w:{tag}")
            node.set(qn("w:id"), "7")
            paragraph._p.append(node)
        label = paragraph.add_run("   (12.10)")
        label.bold = True
        before_math = canonical(group)
        before_anchors = [canonical(n) for n in paragraph._p if "Range" in n.tag or "bookmark" in n.tag]
        apply_equation_layout(doc, {})
        buffer = io.BytesIO()
        doc.save(buffer)
        result = Document(buffer).paragraphs[0]
        self.assertEqual(canonical(result._p.xpath("./m:oMathPara")[0]), before_math)
        self.assertEqual(before_anchors, [canonical(n) for n in result._p if "Range" in n.tag or "bookmark" in n.tag])
        self.assertEqual(result.text, "\t\t(12.10)")
        self.assertTrue(next(r for r in result.runs if "12.10" in r.text).bold)
        self.assertEqual(audit_equation_layout(Document(buffer), {}), [])

    def test_padding_variants_and_second_application_are_idempotent(self):
        for parts in (("(2.4)",), ("  (2.4)",), ("  ", "(2.4)"),
                      ("\t", "\t(2.4)"), ("(2.", "4)", "  ")):
            with self.subTest(parts=parts):
                doc = Document()
                p = add_math(doc.add_paragraph())
                for text in parts:
                    p.add_run(text)
                apply_equation_layout(doc, {})
                self.assertEqual(p.text.strip(), "(2.4)")
                self.assertEqual(p.text.count("\t"), 2)
                self.assertEqual(audit_equation_layout(doc, {}), [])
                before = canonical(p._p)
                counts = apply_equation_layout(doc, {})
                self.assertEqual(canonical(p._p), before)
                self.assertEqual(counts["tabs_set"], 0)
                self.assertEqual(counts["center_tabs_materialized"], 0)
                self.assertEqual(counts["number_tabs_materialized"], 0)

    def test_fixed_table_uses_grid_instead_of_stale_cell_width(self):
        # Reproduces v007's grid/tcW discrepancy, at two different widths.
        for grid, stale in ((3600, 4680), (2880, 6000)):
            with self.subTest(grid=grid):
                doc = Document()
                table = doc.add_table(rows=1, cols=2)
                table.autofit = False
                table.columns[0].width = Twips(grid)
                cell = table.cell(0, 0)
                cell._tc.tcPr.tcW.w = stale
                p = add_math(cell.paragraphs[0], "   （3.2）")
                apply_equation_layout(doc, {})
                self.assertEqual(tab_positions(p), {"center": round((grid - 216) / 2), "right": grid - 216})
                self.assertEqual(cell._tc.tcPr.tcW.w, stale)
                self.assertEqual(audit_equation_layout(doc, {}), [])

    def test_grid_before_and_spans_select_the_actual_tracks(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        table.autofit = False
        for column, width in zip(table.columns, (1000, 1500, 2000, 2500)):
            column.width = Twips(width)
        row = table.rows[0]._tr
        cells = row.findall(qn("w:tc"))
        target = cells[1]
        for extra in cells[2:]:
            row.remove(extra)
        before = OxmlElement("w:gridBefore")
        before.set(qn("w:val"), "1")
        row.get_or_add_trPr().append(before)
        span = OxmlElement("w:gridSpan")
        span.set(qn("w:val"), "2")
        target.get_or_add_tcPr().append(span)
        from docx.text.paragraph import Paragraph
        paragraph = Paragraph(target.find(qn("w:p")), table)
        self.assertEqual(_equation_tab_position_twips(doc, paragraph, {}), 2000 + 2500 - 216)
        before.set(qn("w:val"), "-1")
        self.assertIsNone(_equation_tab_position_twips(doc, paragraph, {}))

    def test_merged_fixed_cell_width(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        table.autofit = False
        for column in table.columns:
            column.width = Inches(1.5)
        cell = table.cell(0, 0).merge(table.cell(0, 1))
        p = add_math(cell.paragraphs[0], "(1)")
        self.assertEqual(apply_equation_layout(doc, {})["equations"], 1)
        self.assertEqual(tab_positions(p)["right"], 4104)

    def test_nested_table_and_one_edge_margin_override(self):
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        outer.autofit = False
        table = outer.cell(0, 0).add_table(rows=1, cols=1)
        table.autofit = False
        table.columns[0].width = Twips(3000)
        cell = table.cell(0, 0)
        margins = OxmlElement("w:tcMar")
        left = OxmlElement("w:left")
        left.set(qn("w:type"), "dxa")
        left.set(qn("w:w"), "36")
        margins.append(left)
        cell._tc.get_or_add_tcPr().append(margins)
        p = add_math(cell.paragraphs[0], "(1)")
        apply_equation_layout(doc, {})
        self.assertEqual(tab_positions(p)["right"], 3000 - 36 - 108)
        self.assertEqual(audit_equation_layout(doc, {}), [])

    def test_ambiguous_or_invalid_width_does_not_create_number_tabs(self):
        for case in ("auto_conflict", "missing_grid", "percent_margin", "negative_margin", "oversize_stop"):
            with self.subTest(case=case):
                doc = Document()
                table = doc.add_table(rows=1, cols=2)
                table.autofit = case == "auto_conflict"
                table.columns[0].width = Twips(2400)
                p = add_math(table.cell(0, 0).paragraphs[0], "(1)")
                spec = {}
                if case == "missing_grid":
                    table._tbl.remove(table._tbl.tblGrid)
                    # Direct width resolution is also safe without a grid.
                    self.assertIsNone(_equation_tab_position_twips(doc, p, spec))
                    continue
                if case in {"percent_margin", "negative_margin"}:
                    mar = OxmlElement("w:tcMar")
                    edge = OxmlElement("w:left")
                    edge.set(qn("w:type"), "pct" if case == "percent_margin" else "dxa")
                    edge.set(qn("w:w"), "100" if case == "percent_margin" else "-100")
                    mar.append(edge)
                    p._p.getparent().get_or_add_tcPr().append(mar)
                if case == "oversize_stop":
                    spec = {"right_tab_twips": 9000}
                self.assertEqual(apply_equation_layout(doc, spec)["equation_tab_layout_unresolved"], 1)
                self.assertEqual(tab_positions(p), {})
                self.assertTrue(audit_equation_layout(doc, spec))

    def test_current_section_width_and_incidental_indents(self):
        doc = Document()
        first = add_math(doc.add_paragraph(), "(1)")
        first.paragraph_format.left_indent = Twips(900)
        first.paragraph_format.first_line_indent = Twips(240)
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.page_width = Inches(10)
        section.left_margin = section.right_margin = Inches(1)
        section.gutter = Twips(120)
        second = add_math(doc.add_paragraph(), "(2)")
        apply_equation_layout(doc, {})
        self.assertEqual(tab_positions(first)["right"], 8640)
        self.assertEqual(tab_positions(second)["right"], 11400)
        self.assertEqual(first.paragraph_format.left_indent.twips, 0)
        self.assertEqual(first.paragraph_format.first_line_indent.twips, 0)
        self.assertEqual(audit_equation_layout(doc, {}), [])

    def test_declared_role_alignment_and_indents_are_respected(self):
        doc = Document()
        p = add_math(doc.add_paragraph(), direct=True)
        p.paragraph_format.left_indent = Twips(360)
        props = {"alignment": "left", "left_indent_pt": 18}
        apply_equation_layout(doc, {}, paragraph_spec=props)
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(p.paragraph_format.left_indent.twips, 360)
        self.assertEqual(tab_positions(p), {})
        self.assertEqual(audit_equation_layout(doc, {}, paragraph_spec=props), [])

    def test_unsupported_label_or_field_stays_intact_and_fails_audit(self):
        for label, field in (("(T.1)", False), ("(1)", True)):
            doc = Document()
            p = add_math(doc.add_paragraph(), label)
            if field:
                node = OxmlElement("w:fldChar")
                node.set(qn("w:fldCharType"), "begin")
                p.add_run()._r.append(node)
            before = canonical(p._p)
            self.assertEqual(apply_equation_layout(doc, {})["equation_tab_layout_unresolved"], 1)
            self.assertEqual(canonical(p._p), before)
            self.assertTrue(audit_equation_layout(doc, {}))

    def test_audit_detects_removed_and_preempted_tabs_including_inheritance(self):
        doc = Document()
        p = add_math(doc.add_paragraph(), "(1)")
        apply_equation_layout(doc, {})
        right_tab = p._p.xpath("./w:r[w:tab]")[-1]
        p._p.remove(right_tab)
        self.assertTrue(any(f["property"].endswith(".tab_usage") for f in audit_equation_layout(doc, {})))
        apply_equation_layout(doc, {})
        p.style.paragraph_format.tab_stops.add_tab_stop(Twips(1000), WD_TAB_ALIGNMENT.LEFT)
        self.assertTrue(any(f["property"].endswith(".center_tab_preempted") for f in audit_equation_layout(doc, {})))

    def test_multiple_math_objects_are_not_recentered(self):
        doc = Document()
        p = add_math(doc.add_paragraph(), direct=True)
        add_math(p, direct=True)
        before = canonical(p._p)
        self.assertFalse(is_display_equation(p))
        self.assertEqual(apply_equation_layout(doc, {})["equations"], 0)
        self.assertEqual(canonical(p._p), before)


if __name__ == "__main__":
    unittest.main()
