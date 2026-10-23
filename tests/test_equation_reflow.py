"""Reflow can change layout, never the protected expression or source file."""
from __future__ import annotations

import contextlib
import copy
import hashlib
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import apply_format_spec as afs
import equation_reflow as reflow


def node(tag,**attrs):
    value=OxmlElement(tag)
    for name,item in attrs.items():value.set(qn(name.replace('_',':')),str(item))
    return value


def math_run(text):
    run=node('m:r');props=node('w:rPr');props.append(node('w:sz',w_val='24'));run.append(props)
    value=node('m:t');value.text=text;run.append(value);return run


def make_doc(text,label='(1)',width=None):
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.left_margin=section.right_margin=Inches(1)
    section.top_margin=section.bottom_margin=Inches(.75)
    doc.styles['Normal'].font.size=Pt(11)
    inline=doc.add_paragraph('INLINE CONTROL ');im=node('m:oMath');im.append(math_run('u+v'));inline._p.append(im);inline.add_run(' stays inline.')
    if width:
        table=doc.add_table(rows=1,cols=2);table.autofit=False
        table.columns[0].width=Pt(width);table.columns[1].width=Pt(468-width)
        table.cell(0,1).text='NEIGHBOR';p=table.cell(0,0).paragraphs[0]
    else:p=doc.add_paragraph()
    group=node('m:oMathPara');pr=node('m:oMathParaPr');pr.append(node('m:jc',m_val='left'));group.append(pr)
    math=node('m:oMath');math.append(math_run(text));group.append(math)
    p._p.append(node('w:bookmarkStart',w_id='47',w_name='equation_source'))
    p._p.append(group);p._p.append(node('w:bookmarkEnd',w_id='47'))
    if label:p.add_run('   '+label)
    return doc,p,group


def source_hash(node):return hashlib.sha256(reflow.canonical(node)).hexdigest()
def ordered_text(node):return ''.join(x.text or '' for x in node.iter(qn('m:t')))
def config():return {'schema_version':'1.0','source_document':'Synthetic reflow regression','status':'semantic_resolved','requirements':[],'roles':{'equation':{}},'equations':{'same_line':True,'number_alignment':'right'}}


class ReflowIntegrityTests(unittest.TestCase):
    def test_only_original_binary_addition_boundaries_can_be_used(self):
        _,_,source=make_doc('Q=a1+a2+a3+a4')
        before=reflow.canonical(source);cuts=[ordered_text(source).index('+')]
        result=reflow.reflow(source,cuts);proof=reflow.proof(source,result,cuts,1)
        self.assertEqual(ordered_text(source),ordered_text(result))
        rows=result.findall('.//'+qn('m:eqArr')+'/'+qn('m:e'))
        self.assertEqual(len(rows),2);self.assertTrue(ordered_text(rows[1]).startswith('+'))
        self.assertEqual(reflow.canonical(source),before)
        self.assertTrue(reflow.validates(proof,result))
        self.assertTrue(all(r.find(qn('w:rPr')).find(qn('w:sz')).get(qn('w:val'))=='24' for r in result.iter(qn('m:r'))))
        with self.assertRaises(ValueError):reflow.reflow(source,[2])
        with self.assertRaises(ValueError):reflow.reflow(source,cuts+cuts)

    def test_ambiguous_text_structures_and_authored_alignment_are_not_reflowed(self):
        for text in ['Q=a+(b+c)','Q=a+b*c','Q=a-b-c','Q=a++b','a+b=c+d','a=b=c','Q=sin(a)+b','Q=a/b+c','Q=1e+3+b']:
            with self.subTest(text=text):
                _,_,math=make_doc(text);before=reflow.canonical(math)
                self.assertEqual(list(reflow.candidate_breaks(math,1000,200)),[])
                self.assertEqual(reflow.canonical(math),before)
        for tag in ['m:aln','m:brk','m:lit','m:nor']:
            _,_,math=make_doc('Q=a+b+c');props=node('m:rPr');props.append(node(tag));math.find(qn('m:oMath')).find(qn('m:r')).insert(0,props)
            self.assertEqual(list(reflow.candidate_breaks(math,1000,200)),[])
        _,_,math=make_doc('Q=a+b+c');math.find(qn('m:oMath')).append(node('m:f'))
        self.assertEqual(list(reflow.candidate_breaks(math,1000,200)),[])
        _,_,math=make_doc('Q=a+b+c');math.find(qn('m:oMath')).find(qn('m:r')).tail='unparsed math'
        self.assertEqual(list(reflow.candidate_breaks(math,1000,200)),[])

    def test_corrupt_expression_font_rows_or_proof_is_rejected(self):
        _,_,source=make_doc('Q=a1+a2+a3+a4');cuts=[4,7];result=reflow.reflow(source,cuts);proof=reflow.proof(source,result,cuts,1)
        for kind in ['operator','missing_operand','reordered_rows','smaller_font']:
            with self.subTest(kind=kind):
                changed=copy.deepcopy(result)
                if kind=='operator':changed.find('.//'+qn('m:t')).text='Q-a1'
                elif kind=='missing_operand':changed.find('.//'+qn('m:t')).text='Q='
                elif kind=='smaller_font':changed.find('.//'+qn('w:sz')).set(qn('w:val'),'12')
                else:
                    array=changed.find('.//'+qn('m:eqArr'));array.insert(0,array[-1])
                self.assertFalse(reflow.validates(proof,changed))
        wrong=copy.deepcopy(proof);wrong['source_math_xml']=wrong['source_math_xml'].replace('a4','b4')
        self.assertFalse(reflow.validates(wrong,result))
        wrong=copy.deepcopy(proof);wrong['break_offsets']=[2]
        self.assertFalse(reflow.validates(wrong,result))

    def test_global_integrity_still_protects_inline_math_and_count_and_order(self):
        doc,p,source=make_doc('Q=a1+a2+a3+a4');original=afs._equation_math_signature(doc)
        # Capture the original namespace context before unlinking, exactly as
        # the formatter does for its staged source-integrity proof.
        snapshot=etree.fromstring(reflow.canonical(source))
        result=reflow.reflow(snapshot,[4]);record=reflow.proof(snapshot,result,[4],1);p._p.replace(source,result)
        self.assertFalse(reflow.integrity_matches(original,afs._equation_math_nodes(doc),[]))
        self.assertTrue(reflow.integrity_matches(original,afs._equation_math_nodes(doc),[record]))
        self.assertFalse(reflow.integrity_matches(original,afs._equation_math_nodes(doc),[record,record]))
        self.assertFalse(reflow.integrity_matches(original,afs._equation_math_nodes(doc)[1:],[record]))
        afs._equation_math_nodes(doc)[0].find('.//'+qn('m:t')).text='v+u'
        self.assertFalse(reflow.integrity_matches(original,afs._equation_math_nodes(doc),[record]))

    def test_invalid_measures_produce_no_proposals(self):
        _,_,source=make_doc('Q=a1+a2+a3+a4')
        for natural,target in [(1000,0),(1000,-1),(float('nan'),20),(1000,float('inf'))]:
            self.assertEqual(list(reflow.candidate_breaks(source,natural,target)),[])


@unittest.skipUnless(shutil.which('soffice'),'LibreOffice is required for rendered reflow tests')
class RenderedReflowTests(unittest.TestCase):
    def test_long_sums_and_number_clearance_are_repaired_at_unchanged_font_size(self):
        for terms,label,width in [(42,'(6)',None),(14,'(1234567890.1234567890)',None),(12,'(13)',144)]:
            with self.subTest(terms=terms,width=width):
                text='S='+'+'.join(f'x{i}' for i in range(1,terms+1));doc,p,source=make_doc(text,label,width)
                original=afs._equation_math_signature(doc);inline=reflow.canonical(doc.paragraphs[0]._p)
                anchors=[reflow.canonical(n) for n in p._p if n.tag in {qn('w:bookmarkStart'),qn('w:bookmarkEnd')}]
                receipt=afs.apply_equation_layout(doc,{'same_line':True})
                self.assertEqual(receipt['equation_tab_layout_unresolved'],0,receipt)
                self.assertEqual(len(receipt['layout_repairs']),1)
                self.assertTrue(reflow.integrity_matches(original,afs._equation_math_nodes(doc),receipt['layout_repairs']))
                current=p._p.find(qn('m:oMathPara'));self.assertEqual(ordered_text(current),text)
                self.assertEqual(reflow.canonical(doc.paragraphs[0]._p),inline)
                self.assertEqual([reflow.canonical(n) for n in p._p if n.tag in {qn('w:bookmarkStart'),qn('w:bookmarkEnd')}],anchors)
                geometry=receipt['geometry_checks'][0];self.assertEqual(geometry['status'],'passed')
                math_box=geometry['context_math_bbox'];label_box=geometry['context_label_bbox']
                container=width-10.8 if width else 468
                self.assertAlmostEqual((math_box[0]+math_box[2])/2,36+container/2,delta=2)
                self.assertAlmostEqual(label_box[2],36+container,delta=2)
                self.assertFalse(math_box[2]>label_box[0] and math_box[0]<label_box[2])
                self.assertEqual(afs.audit_equation_layout(doc,{'same_line':True}),[])

    def test_opt_out_and_unbreakable_or_number_wider_than_cell_still_preserve_source(self):
        for text,label,width,spec in [('S='+'+'.join('x'+str(i) for i in range(42)),'(6)',None,{'wrap_simple_sums':False}),('S='+'X'*180,'(1)',None,{}),('S=a+b+c','(1234567890.1234567890)',72,{})]:
            with self.subTest(width=width,spec=spec):
                doc,p,_=make_doc(text,label,width);before=reflow.canonical(p._p)
                result=afs.apply_equation_layout(doc,spec)
                self.assertEqual(result['equation_tab_layout_unresolved'],1)
                self.assertEqual(result['layout_repairs'],[]);self.assertEqual(reflow.canonical(p._p),before)

    def test_five_full_cli_passes_are_byte_stable_after_a_real_reflow(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'input.docx';text='S='+'+'.join('x'+str(i) for i in range(1,43))
            doc,_,_=make_doc(text);doc.save(source);source_bytes=source.read_bytes();conf=root/'spec.json';conf.write_text(json.dumps(config()))
            outputs=[]
            for i in range(1,6):
                output=root/f'pass-{i}.docx'
                with contextlib.redirect_stdout(io.StringIO()):code=afs.main([str(source),str(conf),str(output),'--out-dir',str(root/f'audit-{i}')])
                self.assertEqual(code,0);current=Document(output)
                self.assertEqual(ordered_text(afs._equation_math_nodes(current)[1]),text)
                self.assertEqual(ordered_text(afs._equation_math_nodes(current)[0]),'u+v')
                report=json.loads((root/f'audit-{i}/validation-report.json').read_text())
                self.assertTrue(report['valid']);outputs.append(output.read_bytes());source=output
            self.assertTrue(all(value==outputs[0] for value in outputs[1:]))
            self.assertEqual((root/'input.docx').read_bytes(),source_bytes)

    def test_staged_math_corruption_is_blocked_and_existing_output_is_preserved(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'input.docx';output=root/'existing.docx'
            doc,_,_=make_doc('S='+'+'.join('x'+str(i) for i in range(1,43)));doc.save(source)
            before=source.read_bytes();output.write_bytes(b'existing output');conf=root/'spec.json';conf.write_text(json.dumps(config()))
            original=afs.apply_document_font
            def corrupt(document,spec):
                result=original(document,spec)
                afs._equation_math_nodes(document)[1].find('.//'+qn('m:t')).text='CORRUPTED'
                return result
            with mock.patch.object(afs,'apply_document_font',side_effect=corrupt),contextlib.redirect_stdout(io.StringIO()):
                code=afs.main([str(source),str(conf),str(output),'--out-dir',str(root/'audit')])
            self.assertEqual(code,1);self.assertEqual(source.read_bytes(),before);self.assertEqual(output.read_bytes(),b'existing output')
            report=json.loads((root/'audit/validation-report.json').read_text())
            self.assertFalse(report['valid']);self.assertFalse(report['output_updated'])
            self.assertTrue(any(x.get('failure_type')=='equation_math_content_changed' for x in report['findings']))


if __name__=='__main__':unittest.main()
