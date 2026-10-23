"""Rendered regression cases for equation geometry and atomic output safety."""
from __future__ import annotations

import contextlib
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import apply_format_spec as afs
import equation_geometry as geometry


def node(tag, **attributes):
    result = OxmlElement(tag)
    for name, value in attributes.items():
        result.set(qn(name.replace("_", ":")), str(value))
    return result


def math_run(text):
    run = node("m:r")
    item = node("m:t"); item.text = text; run.append(item)
    return run


def add_equation(paragraph, text="x+y=1", label="(1)", rows=0, native_center=False):
    group = node("m:oMathPara")
    if not native_center:
        props = node("m:oMathParaPr"); props.append(node("m:jc", m_val="left")); group.append(props)
    math = node("m:oMath")
    if rows:
        array = node("m:eqArr")
        for index in range(rows):
            element = node("m:e"); element.append(math_run(f"y{index+1}=x{index+1}+{index+1}")); array.append(element)
        math.append(array)
    else:
        math.append(math_run(text))
    group.append(math); paragraph._p.append(group)
    paragraph._p.append(node("w:bookmarkStart", w_id="42", w_name="equation_control"))
    if label:
        paragraph.add_run("   " + label)
    paragraph._p.append(node("w:bookmarkEnd", w_id="42"))
    return paragraph


def document():
    doc = Document(); section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(.75)
    doc.styles["Normal"].font.size = Pt(11)
    return doc


def canonical(element):
    return etree.tostring(element, method="c14n")


def failure_codes(receipt):
    return {finding["code"] for item in receipt["manual_review_items"] for finding in item["findings"]}


def spec(paragraph=None, font=None):
    return {"schema_version": "1.0", "source_document": "Synthetic equation regression",
            "status": "semantic_resolved", "requirements": [],
            "roles": {"equation": {"paragraph": paragraph or {}, "font": font or {}}},
            "equations": {"alignment": "center", "number_alignment": "right", "same_line": True}}


class EquationSafetyWithoutRendererTests(unittest.TestCase):
    def test_actual_ink_keeps_a_math_rule_crossing_the_number(self):
        import fitz
        pdf = fitz.open(); page = pdf.new_page()
        sentinel = 0x13A5C7; color = tuple(c/255 for c in fitz.sRGB_to_rgb(sentinel))
        page.insert_text((80, 70), "x+y=1", fontsize=12)
        page.insert_text((180, 70), "(1)", fontsize=12, color=color)
        page.draw_line((100, 65), (195, 65), color=(0, 0, 0))
        page.draw_line((180, 74), (195, 74), color=color)
        math_box, label_box = geometry._context_ink(page, sentinel)
        self.assertGreaterEqual(math_box.x1, 195)
        self.assertLess(math_box.y1, 74)
        self.assertGreaterEqual(label_box.y1, 74)
        self.assertTrue(math_box.intersects(label_box))
        pdf.close()

    def test_unavailable_renderer_restores_source_and_fails_closed(self):
        doc = document(); paragraph = add_equation(doc.add_paragraph())
        before = canonical(paragraph._p)
        with mock.patch.object(geometry.shutil, "which", return_value=None):
            receipt = afs.apply_equation_layout(doc, {})
        self.assertEqual(receipt["equation_tab_layout_unresolved"], 1)
        self.assertIn("equation_geometry_unavailable", failure_codes(receipt))
        self.assertEqual(canonical(paragraph._p), before)

    def test_unequal_columns_and_unknown_outer_table_require_review(self):
        for kind in ("unequal_columns", "outer_autofit", "conditional_table_style"):
            with self.subTest(kind=kind):
                doc = document()
                if kind == "unequal_columns":
                    cols = doc.sections[0]._sectPr.find(qn("w:cols"))
                    cols.set(qn("w:num"), "2"); cols.set(qn("w:equalWidth"), "0")
                    cols.extend([node("w:col", w_w=3000, w_space=360), node("w:col", w_w=6000)])
                    paragraph = add_equation(doc.add_paragraph())
                else:
                    outer = doc.add_table(rows=1, cols=1)
                    if kind == "outer_autofit":
                        table = outer.cell(0, 0).add_table(rows=1, cols=1)
                    else:
                        outer.autofit = False
                        table = outer
                        from docx.enum.style import WD_STYLE_TYPE
                        style = doc.styles.add_style("Conditional Equation Table", WD_STYLE_TYPE.TABLE)
                        conditional = node("w:tblStylePr", w_type="firstRow")
                        rpr = node("w:rPr"); rpr.append(node("w:sz", w_val="60")); conditional.append(rpr)
                        style.element.append(conditional); table.style = style
                    table.autofit = False
                    table.columns[0].width = Inches(3)
                    paragraph = add_equation(table.cell(0, 0).paragraphs[0])
                before = canonical(paragraph._p)
                receipt = afs.apply_equation_layout(doc, {})
                self.assertEqual(receipt["equation_tab_layout_unresolved"], 1)
                self.assertEqual(canonical(paragraph._p), before)

    def test_empty_fixture_does_not_claim_rendered_geometry(self):
        doc = document(); p = doc.add_paragraph(); p._p.append(node("m:oMath"))
        with mock.patch.object(geometry, "_render", side_effect=AssertionError("empty fixture rendered")):
            receipt = afs.apply_equation_layout(doc, {})
        self.assertEqual(receipt["geometry_checks"][0]["status"], "empty_placeholder")
        self.assertFalse(receipt["geometry_checks"][0]["probe_geometry_verified"])

    def test_cli_failure_does_not_replace_an_existing_output(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source = root/"input.docx"; output = root/"existing.docx"
            doc = document(); add_equation(doc.add_paragraph()); doc.save(source)
            source_bytes = source.read_bytes(); output.write_bytes(b"existing output must survive")
            config = root/"spec.json"; config.write_text(json.dumps(spec()))
            with mock.patch.object(geometry.shutil, "which", return_value=None), contextlib.redirect_stdout(io.StringIO()):
                code = afs.main([str(source), str(config), str(output), "--out-dir", str(root/"audit")])
            self.assertEqual(code, 1)
            self.assertEqual(source.read_bytes(), source_bytes)
            self.assertEqual(output.read_bytes(), b"existing output must survive")
            report = json.loads((root/"audit/validation-report.json").read_text())
            self.assertFalse(report["valid"]); self.assertFalse(report["output_updated"])
            self.assertTrue(report["existing_output_preserved"])
            self.assertIsNone(report["output_docx"])


@unittest.skipUnless(shutil.which("soffice"), "LibreOffice is required for actual equation geometry")
class RenderedEquationSafetyTests(unittest.TestCase):
    def test_incidental_fixed_line_height_cannot_hide_sum_limits(self):
        for explicit in (False, True):
            with self.subTest(explicit=explicit):
                doc=document();p=add_equation(doc.add_paragraph(),label="(1)")
                math=p._p.xpath("./m:oMathPara/m:oMath")[0]
                for child in list(math): math.remove(child)
                nary=node("m:nary");pr=node("m:naryPr")
                pr.extend([node("m:chr",m_val="∑"),node("m:limLoc",m_val="undOvr"),node("m:subHide",m_val="off"),node("m:supHide",m_val="off")])
                nary.append(pr)
                for tag,text in (("m:sub","i=1"),("m:sup","N"),("m:e","x+i")):
                    item=node(tag);item.append(math_run(text));nary.append(item)
                math.append(nary)
                p.paragraph_format.line_spacing=Pt(8)
                from docx.enum.text import WD_LINE_SPACING
                p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
                before=canonical(p._p);math_before=afs._equation_math_signature(doc)
                props={"line_spacing":{"type":"exact","value":8}} if explicit else {}
                receipt=afs.apply_equation_layout(doc,{},paragraph_spec=props)
                self.assertEqual(receipt["equation_tab_layout_unresolved"],1 if explicit else 0,receipt)
                self.assertEqual(afs._equation_math_signature(doc),math_before)
                if explicit:
                    self.assertEqual(canonical(p._p),before)
                else:
                    self.assertTrue(receipt["geometry_checks"][0]["probe_geometry_verified"])

    def test_derived_zero_character_indent_overrides_base_but_active_unknown_still_blocks(self):
        from docx.enum.style import WD_STYLE_TYPE
        for override in (True, False):
            with self.subTest(override=override):
                doc=document()
                doc.styles["Normal"].element.get_or_add_pPr().append(node("w:ind", w_firstLineChars="200"))
                style=doc.styles.add_style("Equation with inherited indent", WD_STYLE_TYPE.PARAGRAPH)
                style.base_style=doc.styles["Normal"]
                if override:
                    style.element.get_or_add_pPr().append(node("w:ind", w_firstLineChars="0"))
                p=add_equation(doc.add_paragraph(style=style));before=canonical(p._p)
                result=afs.apply_equation_layout(doc,{})
                self.assertEqual(result["equation_tab_layout_unresolved"],0 if override else 1)
                if override:
                    self.assertTrue(result["geometry_checks"][0]["probe_geometry_verified"])
                else:
                    self.assertEqual(canonical(p._p),before)
                    self.assertIn("equation_character_or_directional_indent_requires_review",failure_codes(result))

    def test_indents_and_two_column_measures_are_rendered_at_their_actual_centers(self):
        for left, right, columns in ((36,36,1), (48,12,1), (0,0,2)):
            with self.subTest(left=left, right=right, columns=columns):
                doc = document(); p = add_equation(doc.add_paragraph())
                p.paragraph_format.left_indent = Pt(left); p.paragraph_format.right_indent = Pt(right)
                if columns == 2:
                    cols = doc.sections[0]._sectPr.find(qn("w:cols"))
                    cols.set(qn("w:num"), "2"); cols.set(qn("w:space"), "360")
                props = {"left_indent_pt":left, "right_indent_pt":right, "first_line_indent_pt":0}
                receipt = afs.apply_equation_layout(doc, {}, paragraph_spec=props)
                self.assertEqual(receipt["equation_tab_layout_unresolved"], 0, receipt)
                result = receipt["geometry_checks"][0]
                self.assertTrue(result["probe_geometry_verified"])
                width = (9360-360)/2 if columns==2 else 9360
                expected_center = 36 + (left*20 + width-right*20)/40
                actual = result["context_math_bbox"]
                self.assertAlmostEqual((actual[0]+actual[2])/2, expected_center, delta=2)
                self.assertEqual(afs.audit_equation_layout(doc, {}, paragraph_spec=props), [])
                self.assertFalse(result["whole_document_visual_review"])
                self.assertFalse(result["word_wps_verified"])

    def test_wide_number_long_math_and_tall_math_are_not_published_as_safe(self):
        for kind in ("wide_number", "long", "tall", "exact_row"):
            with self.subTest(kind=kind):
                doc = document()
                if kind in {"wide_number", "exact_row"}:
                    table = doc.add_table(rows=1, cols=2); table.autofit=False
                    table.columns[0].width = Inches(1 if kind=="wide_number" else 3)
                    if kind == "exact_row":
                        table.rows[0]._tr.get_or_add_trPr().append(node("w:trHeight", w_val="160", w_hRule="exact"))
                    p = table.cell(0,0).paragraphs[0]
                else:
                    p = doc.add_paragraph()
                # An indivisible identifier has no safe additive break. Flat
                # sums have their own positive rendered regression coverage.
                text = "S="+"X"*180 if kind=="long" else "x+y=1"
                add_equation(p, text, "(1234567890.1234567890)" if kind=="wide_number" else "(1)", rows=58 if kind=="tall" else 0)
                before = canonical(p._p); before_math = afs._equation_math_signature(doc)
                receipt = afs.apply_equation_layout(doc, {})
                self.assertEqual(receipt["equation_tab_layout_unresolved"], 1, receipt)
                self.assertEqual(canonical(p._p), before)
                self.assertEqual(afs._equation_math_signature(doc), before_math)
                expected = {"wide_number":"equation_number_has_no_clearance", "long":"equation_too_wide"}
                if kind in expected: self.assertIn(expected[kind], failure_codes(receipt))
                self.assertTrue(failure_codes(receipt))

    def test_same_line_audit_catches_automatic_wrap_without_newline(self):
        doc = document(); p = add_equation(doc.add_paragraph(), "S="+"+".join(f"x{i}" for i in range(1,15)), "(1234567890.1234567890)")
        p.paragraph_format.left_indent = p.paragraph_format.right_indent = Twips(0)
        p.paragraph_format.first_line_indent = Twips(0)
        p.alignment = afs.ALIGN["left"]
        p._p.insert(1, node("w:r")); p._p[1].append(node("w:tab"))
        label = next(run for run in p.runs if "123456" in run.text)
        label.text = "\t(1234567890.1234567890)"
        afs._set_equation_tab_stop(p,"center",4680); afs._set_equation_tab_stop(p,"right",9360)
        self.assertNotIn("\n",p.text)
        findings = afs.audit_equation_layout(doc, {"same_line":True})
        codes = {item.get("failure_type") for item in findings}
        self.assertIn("equation_rendered_number_not_same_line",codes)

    def test_unnumbered_native_center_group_uses_one_centering_mechanism(self):
        doc = document(); p = add_equation(doc.add_paragraph(), label="", native_center=True)
        before = afs._equation_math_signature(doc)
        receipt = afs.apply_equation_layout(doc, {})
        self.assertEqual(receipt["equation_tab_layout_unresolved"],0)
        self.assertNotIn("\t",p.text)
        self.assertTrue(receipt["geometry_checks"][0]["probe_geometry_verified"])
        self.assertEqual(afs._equation_math_signature(doc),before)

    def test_declared_left_does_not_accept_native_center_group_as_left_aligned(self):
        doc = document(); p = add_equation(doc.add_paragraph(),label="",native_center=True)
        before=canonical(p._p)
        result=afs.apply_equation_layout(doc,{},paragraph_spec={"alignment":"left"})
        self.assertEqual(result["equation_tab_layout_unresolved"],1)
        self.assertIn("equation_rendered_not_left_aligned",failure_codes(result))
        self.assertEqual(canonical(p._p),before)

    def test_cache_is_bound_to_required_center_and_saves_its_probe(self):
        doc = document(); p = add_equation(doc.add_paragraph())
        receipt=afs.apply_equation_layout(doc,{})
        self.assertEqual(receipt["equation_tab_layout_unresolved"],0)
        with tempfile.TemporaryDirectory() as temp:
            result=geometry.verify_equation(doc,p,width=9360,height=13680,left=0,right=9360,center=4000,label="(1)",evidence_dir=temp)
            self.assertIn("equation_rendered_not_centered",{f["code"] for f in result["findings"]})
            evidence=Path(temp)/result["validation_key"]
            self.assertTrue((evidence/"probe.docx").is_file()); self.assertTrue((evidence/"probe.pdf").is_file())

    def test_cli_five_passes_are_byte_identical_and_preserve_math_bookmarks_inline(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp); source=root/"source.docx"
            doc=document(); doc.add_paragraph("Adjacent text remains editable.")
            inline=doc.add_paragraph("Inline control A "); m=node("m:oMath");m.append(math_run("u+v"));inline._p.append(m);inline.add_run(" B.")
            p=add_equation(doc.add_paragraph());p.paragraph_format.left_indent=p.paragraph_format.right_indent=Pt(36)
            doc.save(source); original_math=afs._equation_math_signature(doc)
            original_inline=canonical(inline._p)
            anchors=[canonical(n) for n in p._p.xpath("./w:bookmarkStart|./w:bookmarkEnd")]
            config=root/"spec.json";config.write_text(json.dumps(spec({"left_indent_pt":36,"right_indent_pt":36,"first_line_indent_pt":0},{"latin":"Times New Roman","size_pt":12})))
            outputs=[]
            for index in range(1,6):
                output=root/f"pass-{index}.docx"
                with contextlib.redirect_stdout(io.StringIO()):
                    code=afs.main([str(source),str(config),str(output),"--out-dir",str(root/f"audit-{index}")])
                self.assertEqual(code,0)
                current=Document(output);self.assertEqual(afs._equation_math_signature(current),original_math)
                self.assertEqual(canonical(current.paragraphs[1]._p),original_inline)
                self.assertEqual([canonical(n) for n in current.paragraphs[2]._p.xpath("./w:bookmarkStart|./w:bookmarkEnd")],anchors)
                outputs.append(output.read_bytes());source=output
            self.assertTrue(all(data==outputs[0] for data in outputs[1:]))


if __name__=="__main__":
    unittest.main()
