from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from apply_format_spec import apply_headers_footers
from format_spec_validation import load_and_validate
from header_scope import compile_header_scope_plan, compile_header_scope_rules
from header_scope_audit import audit_scoped_headers, build_header_scope_review_ledger


def header_spec(*, include_toc: bool = True, include_appendix: bool = True) -> dict:
    requirements = []
    if include_toc:
        requirements.append({
            "id": "R-TOC-ARBITRARY", "role": "header",
            "resolved_by": "llm", "confidence": 0.97,
            "properties": {"header_content": {"left_text": "目  录"}},
            "evidence_ids": ["E-TOC"], "clause_ids": ["C-TOC"],
            "applicability": {"status": "conditional", "conditions": [{
                "fact": "source_inventory.table_of_contents", "operator": "present", "value": None,
            }], "exceptions": []},
        })
    if include_appendix:
        requirements.append({
            "id": "R-APPENDIX-ARBITRARY", "role": "header",
            "resolved_by": "llm", "confidence": 0.97,
            "properties": {"header_content": {"left_text": "附  录"}},
            "evidence_ids": ["E-APPENDIX"], "clause_ids": ["C-APPENDIX"],
            "applicability": {"status": "conditional", "conditions": [{
                "fact": "thesis_profile.has_appendices", "operator": "equals", "value": True,
            }], "exceptions": []},
        })
    return {
        "schema_version": "1.0", "source_document": "official-requirements.docx",
        "status": "semantic_resolved", "roles": {"header": {
            # This deliberately represents the old lossy projection. The scoped
            # executor must use the requirement-level rules below it.
            "header_content": {"left_text": "目  录"},
        }},
        "requirements": requirements,
        "thesis_profile": {
            "schema_version": "1.0", "degree_level": "master",
            "writing_language": "zh", "has_appendices": include_appendix,
        },
    }


def make_two_section_source(path: Path, *, active_variants: bool = False,
                            link_duplicate_chapter: bool = False) -> None:
    document = Document()
    toc = document.add_paragraph("目 录", style="TOC Heading")
    toc.paragraph_format.page_break_before = True
    section = document.add_section(WD_SECTION.NEW_PAGE)
    heading_text = "第1章 引言"
    document.add_paragraph(heading_text, style="Heading 1")
    if link_duplicate_chapter:
        # The inherited static header is valid only because the following
        # source anchor is identical; the executor must split the linked part.
        document.add_paragraph(heading_text, style="Heading 1")
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = heading_text
    if active_variants:
        document.settings.odd_and_even_pages_header_footer = True
        for item in document.sections:
            item.different_first_page_header_footer = True
        document.sections[0].first_page_header.paragraphs[0].text = "TOC first source value"
        document.sections[0].even_page_header.paragraphs[0].text = "TOC even source value"
        section.first_page_header.is_linked_to_previous = False
        section.first_page_header.paragraphs[0].text = heading_text + " 首"
        section.even_page_header.is_linked_to_previous = False
        section.even_page_header.paragraphs[0].text = heading_text + " 偶"
    document.save(path)


def generate_scoped(source: Path, output: Path, spec: dict) -> dict:
    plan = compile_header_scope_plan(source, spec)
    if not plan["valid"]:
        raise AssertionError(json.dumps(plan["findings"], ensure_ascii=False))
    document = Document(source)
    result = apply_headers_footers(
        document, spec.get("roles", {}), {}, {}, {}, header_scope_plan=plan,
    )
    document.save(output)
    return {"plan": plan, "execution": result}


def rewrite_package(source: Path, destination: Path, mutate) -> None:
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                from lxml import etree
                root = etree.fromstring(data)
                mutate(root)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)


class HeaderScopeTest(unittest.TestCase):
    def test_requirements_compile_distinct_scopes_without_clause_id_constants(self):
        spec = header_spec()
        rules, findings = compile_header_scope_rules(spec)
        self.assertFalse(findings)
        self.assertEqual({rule["scope"] for rule in rules}, {"toc", "appendix"})
        self.assertEqual(
            {clause for rule in rules for clause in rule["clause_ids"]},
            {"C-TOC", "C-APPENDIX"},
        )
        compiled = {**spec, "header_scope_rules": rules}
        self.assertEqual(
            load_and_validate(compiled, ROOT / "schema" / "format-spec.schema.json"), [],
        )

    def test_two_section_counterexample_generates_distinct_toc_and_chapter_headers(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, output = root / "source.docx", root / "review.docx"
            make_two_section_source(source)
            spec = header_spec(include_appendix=False)
            generated = generate_scoped(source, output, spec)
            self.assertEqual(generated["plan"]["sections"][0]["scope"], "toc")
            self.assertEqual(generated["plan"]["sections"][1]["scope"], "chapter")
            self.assertEqual(
                load_and_validate(
                    generated["plan"], ROOT / "schema" / "header-scope-plan.schema.json",
                ),
                [],
            )
            result = audit_scoped_headers(source, output, spec)
            self.assertTrue(result["valid"], json.dumps(result["findings"], ensure_ascii=False))
            self.assertEqual(len(result["sections"]), 2)
            chapter_header = result["sections"][1]["variants"]["default"]
            self.assertEqual(chapter_header["expected_text"], "第1章 引言")
            self.assertEqual(chapter_header["actual_text"], "第1章 引言")
            self.assertTrue(chapter_header["relationship_id"])
            self.assertEqual(chapter_header["relationship_type"], "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header")
            doc = Document(output)
            self.assertEqual(doc.sections[0].header.paragraphs[0].text, "目  录")
            self.assertEqual(doc.sections[1].header.paragraphs[0].text, "第1章 引言")

    def test_source_anchor_inside_table_is_mapped_and_independently_audited(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, output = root / "source.docx", root / "review.docx"
            document = Document()
            cell = document.add_table(rows=1, cols=1).cell(0, 0)
            cell.paragraphs[0].text = "目 录"
            cell.paragraphs[0].style = "TOC Heading"
            document.save(source)
            spec = header_spec(include_appendix=False)
            plan = generate_scoped(source, output, spec)["plan"]
            self.assertEqual(plan["sections"][0]["scope"], "toc")
            result = audit_scoped_headers(source, output, spec)
            self.assertTrue(result["valid"], json.dumps(result["findings"], ensure_ascii=False))

    def test_appendix_rule_overrides_source_header_only_on_appendix_anchor(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "source.docx"
            document = Document()
            document.add_paragraph("第1章 引言", style="Heading 1")
            document.sections[0].header.paragraphs[0].text = "第1章 引言"
            section = document.add_section(WD_SECTION.NEW_PAGE)
            document.add_paragraph("附录A 实验环境", style="Heading 1")
            section.header.is_linked_to_previous = False
            section.header.paragraphs[0].text = "附录A 实验环境"
            document.save(source)
            spec = header_spec(include_toc=False)
            output = root / "review.docx"
            plan = generate_scoped(source, output, spec)["plan"]
            self.assertEqual([x["scope"] for x in plan["sections"]], ["chapter", "appendix"])
            result = audit_scoped_headers(source, output, spec)
            self.assertTrue(result["valid"], json.dumps(result["findings"], ensure_ascii=False))
            doc = Document(output)
            self.assertEqual(doc.sections[0].header.paragraphs[0].text, "第1章 引言")
            self.assertEqual(doc.sections[1].header.paragraphs[0].text, "附  录")

    def test_first_even_and_linked_to_previous_are_materialized_and_audited(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, output = root / "source.docx", root / "review.docx"
            make_two_section_source(source, active_variants=True)
            spec = header_spec(include_appendix=False)
            generate_scoped(source, output, spec)
            result = audit_scoped_headers(source, output, spec)
            self.assertTrue(result["valid"], json.dumps(result["findings"], ensure_ascii=False))
            doc = Document(output)
            for variant in ("header", "first_page_header", "even_page_header"):
                self.assertFalse(getattr(doc.sections[1], variant).is_linked_to_previous)
            self.assertEqual(doc.sections[0].first_page_header.paragraphs[0].text, "目  录")
            self.assertEqual(doc.sections[0].even_page_header.paragraphs[0].text, "目  录")
            self.assertEqual(doc.sections[1].first_page_header.paragraphs[0].text, "第1章 引言 首")
            self.assertEqual(doc.sections[1].even_page_header.paragraphs[0].text, "第1章 引言 偶")

    def test_source_linked_header_is_split_when_scopes_are_distinct(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, output = root / "source.docx", root / "review.docx"
            document = Document()
            document.add_paragraph("目 录", style="TOC Heading")
            document.sections[0].header.paragraphs[0].text = "TOC source header"
            second = document.add_section(WD_SECTION.NEW_PAGE)
            document.add_paragraph("附录A 实验环境", style="Heading 1")
            second.header.is_linked_to_previous = True
            document.save(source)
            spec = header_spec()
            plan = compile_header_scope_plan(source, spec)
            self.assertTrue(plan["valid"], json.dumps(plan["findings"], ensure_ascii=False))
            generated = generate_scoped(source, output, spec)
            self.assertTrue(audit_scoped_headers(source, output, spec)["valid"])
            result = Document(output)
            self.assertFalse(result.sections[1].header.is_linked_to_previous)
            self.assertEqual(result.sections[1].header.paragraphs[0].text, "附  录")

    def test_tampered_header_text_and_relationship_are_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, good = root / "source.docx", root / "good.docx"
            make_two_section_source(source)
            spec = header_spec(include_appendix=False)
            generate_scoped(source, good, spec)
            doc = Document(good)
            doc.sections[1].header.paragraphs[0].text = "目  录"
            bad_text = root / "bad-text.docx"
            doc.save(bad_text)
            text_audit = audit_scoped_headers(source, bad_text, spec)
            self.assertFalse(text_audit["valid"])
            self.assertIn("header_scope_text_mismatch", {x["code"] for x in text_audit["findings"]})
            text_ledger = build_header_scope_review_ledger(text_audit)
            self.assertEqual(text_ledger["status"], "blocked")
            mismatch = next(item for item in text_ledger["items"] if item["code"] == "header_scope_text_mismatch")
            self.assertEqual(mismatch["section_index"], 2)
            self.assertEqual(mismatch["scope"], "chapter")
            self.assertEqual(mismatch["source_anchor"]["text"], "第1章 引言")
            self.assertEqual(mismatch["expected_text"], "第1章 引言")
            self.assertEqual(mismatch["actual_text"], "目  录")
            self.assertEqual(mismatch["clause_ids"], [])
            self.assertEqual(
                load_and_validate(text_ledger, ROOT / "schema" / "header-scope-review-ledger.schema.json"),
                [],
            )

            bad_rel = root / "bad-rel.docx"
            def corrupt_relationship(xml):
                refs = xml.xpath("//w:sectPr", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})[1].xpath(
                    "./w:headerReference[@w:type='default']",
                    namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
                )
                refs[0].set(qn("r:id"), "rIdMissing")
            rewrite_package(good, bad_rel, corrupt_relationship)
            rel_audit = audit_scoped_headers(source, bad_rel, spec)
            self.assertFalse(rel_audit["valid"])
            self.assertIn("header_scope_relationship_missing", {x["code"] for x in rel_audit["findings"]})

    def test_first_even_missing_links_cannot_hide_cross_section_scope_bleed(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, good, bad = root / "source.docx", root / "good.docx", root / "missing-first-even.docx"
            make_two_section_source(source, active_variants=True)
            spec = header_spec(include_appendix=False)
            generate_scoped(source, good, spec)

            def remove_first_even_refs(xml):
                sectprs = xml.xpath("//w:sectPr", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
                for variant in ("first", "even"):
                    refs = sectprs[1].xpath(
                        f"./w:headerReference[@w:type='{variant}']",
                        namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
                    )
                    self.assertEqual(len(refs), 1)
                    refs[0].getparent().remove(refs[0])

            rewrite_package(good, bad, remove_first_even_refs)
            result = audit_scoped_headers(source, bad, spec)
            mismatches = [x for x in result["findings"] if x.get("code") == "header_scope_text_mismatch"]
            self.assertEqual({x.get("variant") for x in mismatches}, {"first", "even"})
            ledger = build_header_scope_review_ledger(result)
            self.assertEqual(ledger["status"], "blocked")
            ledger_mismatches = [x for x in ledger["items"] if x["code"] == "header_scope_text_mismatch"]
            self.assertEqual({x.get("variant") for x in ledger_mismatches}, {"first", "even"})
            self.assertEqual(
                load_and_validate(ledger, ROOT / "schema" / "header-scope-review-ledger.schema.json"),
                [],
            )

    def test_shared_header_part_across_different_scopes_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, good, bad = root / "source.docx", root / "good.docx", root / "shared.docx"
            make_two_section_source(source)
            spec = header_spec(include_appendix=False)
            generate_scoped(source, good, spec)
            def share(xml):
                refs = xml.xpath("//w:sectPr/w:headerReference[@w:type='default']", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
                refs[1].set(qn("r:id"), refs[0].get(qn("r:id")))
            rewrite_package(good, bad, share)
            result = audit_scoped_headers(source, bad, spec)
            codes = {item["code"] for item in result["findings"]}
            self.assertFalse(result["valid"])
            self.assertTrue({"header_scope_text_mismatch", "header_scope_shared_part_collision"} & codes)

    def test_missing_toc_anchor_and_mismatched_chapter_source_fail_closed(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            no_toc = root / "no-toc.docx"
            document = Document()
            document.add_paragraph("第1章 引言", style="Heading 1")
            document.sections[0].header.paragraphs[0].text = "第1章 引言"
            document.save(no_toc)
            plan = compile_header_scope_plan(no_toc, header_spec(include_appendix=False))
            self.assertFalse(plan["valid"])
            self.assertIn("header_scope_toc_anchor_missing", {x["code"] for x in plan["findings"]})

            mismatch = root / "mismatch.docx"
            document = Document()
            document.add_paragraph("第1章 引言", style="Heading 1")
            document.sections[0].header.paragraphs[0].text = "第2章 错误标题"
            document.save(mismatch)
            plan = compile_header_scope_plan(mismatch, header_spec())
            self.assertFalse(plan["valid"])
            self.assertIn("header_scope_chapter_source_mismatch", {x["code"] for x in plan["findings"]})
            self.assertIn("header_scope_chapter_source_mismatch", {x["code"] for x in plan["findings"]})

    def test_unknown_applicability_and_conflicting_scopes_are_rejected(self):
        spec = header_spec()
        spec["requirements"][0]["applicability"]["conditions"][0]["fact"] = "unknown.page_region"
        rules, findings = compile_header_scope_rules(spec)
        self.assertFalse(rules)
        self.assertIn("header_scope_applicability_unmapped", {x["code"] for x in findings})

        spec = header_spec(include_appendix=False)
        duplicate = copy.deepcopy(spec["requirements"][0])
        duplicate["id"] = "R-TOC-CONFLICT"
        duplicate["properties"]["header_content"]["left_text"] = "章节文字"
        duplicate["clause_ids"] = ["C-OTHER"]
        duplicate["evidence_ids"] = ["E-OTHER"]
        spec["requirements"].append(duplicate)
        rules, findings = compile_header_scope_rules(spec)
        self.assertFalse(rules)
        self.assertIn("header_scope_conflict", {x["code"] for x in findings})


if __name__ == "__main__":
    unittest.main()
