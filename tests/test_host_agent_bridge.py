from __future__ import annotations

import copy
from concurrent.futures import ALL_COMPLETED
import hashlib
import json
import os
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import host_agent_bridge as bridge  # noqa: E402
import requirements_engine as engine  # noqa: E402
import thesis_format_pipeline as pipeline  # noqa: E402
from manual_review import build_manual_review_ledger  # noqa: E402
from manual_review_display import (  # noqa: E402
    append_manual_review_markers,
    audit_manual_review_markers,
)
from semantic_contract import (  # noqa: E402
    attach_request_provenance,
    request_body_sha256,
    request_envelope_sha256,
    sha256_file,
    sha256_json,
)
from semantic_source_references import (  # noqa: E402
    REFERENCE_PROTOCOL,
    build_source_reference_packet,
    compile_source_reference_response,
)


def bind_mock_review_to_source_spans(review_result: dict, request: dict, output_dir: Path) -> dict:
    """Give mocked native-review results the same request-bound compilation receipt as production."""
    packet = build_source_reference_packet(request)
    checks = {item["check_id"]: item for item in packet["checks"]}
    results = copy.deepcopy(review_result.get("results") or [])
    selections = []
    for result in results:
        check_id = result["check_id"]
        source_spans = checks[check_id]["source_spans"]

        def resolve_span(quote: str) -> dict:
            matches = [span for span in source_spans if quote in span["text"]]
            if not matches:
                raise AssertionError(f"fixture quote is not in source spans: {quote!r}")
            return min(matches, key=lambda span: len(span["text"]))

        selected_spans = []
        for obligation_index, obligation in enumerate(result.get("identified_obligations", [])):
            span = resolve_span(obligation["source_quote"])
            obligation["source_quote"] = span["text"]
            selected_spans.append({
                "obligation_index": obligation_index,
                "source_ref": span["ref_id"],
                "span": copy.deepcopy(span),
            })
        evidence_spans = []
        for index, quote in enumerate(result.get("evidence_quotes", [])):
            span = resolve_span(quote)
            result["evidence_quotes"][index] = span["text"]
            evidence_spans.append(span)
        selections.append({
            "check_id": check_id,
            "spans": list({span["ref_id"]: copy.deepcopy(span)
                           for span in [*evidence_spans, *[item["span"] for item in selected_spans]]}.values()),
            "obligations": selected_spans,
        })
    output_dir.mkdir(parents=True, exist_ok=True)
    request_sha = bridge.sha256_json(request)
    compilation = {
        "protocol": "semantic_source_references_v2",
        "run_id": request.get("run_id"),
        "request_sha256": request_sha,
        "packet_sha256": bridge.sha256_json(packet),
        "raw_response_sha256": "0" * 64,
        "compiled_response_sha256": bridge.sha256_json({"results": results}),
        "selections": selections,
        "semantic_verdicts_unchanged": True,
    }
    compilation_path = output_dir / "source-reference-compilation.json"
    bridge._write_json(compilation_path, compilation)
    return {
        **review_result,
        "request_sha256": request_sha,
        "results": results,
        "canonical_response_sha256": bridge.sha256_json({"results": results}),
        "source_reference_compilation_path": str(compilation_path),
        "source_reference_compilation_sha256": bridge.sha256_file(compilation_path),
    }


def exact_source_clause(clause_id: str, source: str, evidence_id: str = "E1") -> dict:
    """Build a clause fixture whose exact span is bound to its evidence text."""
    return {
        "id": clause_id,
        "text": source,
        "evidence_ids": [evidence_id],
        "source_span": {
            "evidence_id": evidence_id,
            "start_offset": 0,
            "end_offset": len(source),
            "text": source,
            "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
        },
    }


def bind_declaration_fixture_spans(chunk: dict) -> None:
    """Supply exact current source occurrences for old compact declaration fixtures."""
    for clause in chunk["clauses"]:
        if "source_span" in clause:
            continue
        eid = clause["evidence_ids"][0]
        evidence = chunk["evidence_context"][eid]
        source, text = evidence["text"], clause["text"]
        start = source.index(text)
        clause["source_span"] = {
            "evidence_id": eid, "start_offset": start, "end_offset": start + len(text),
            "text": text, "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
            **({"location": copy.deepcopy(evidence["location"])} if "location" in evidence else {}),
        }


class HostAgentBridgeTests(unittest.TestCase):
    def setUp(self) -> None:
        self._host_runtime_env = patch.dict(
            os.environ,
            {"THESIS_FORGE_HOST_RUNTIME": "openclaw"},
        )
        self._host_runtime_env.start()
        self._independent_review_patch = patch.object(
            bridge,
            "_run_independent_obligation_coverage_review",
            side_effect=self._fake_independent_review,
        )
        self._independent_review_patch.start()
        self.addCleanup(self._independent_review_patch.stop)

    def tearDown(self) -> None:
        self._host_runtime_env.stop()

    def test_primary_failure_selection_is_audited_not_called_chronological(self) -> None:
        selection = bridge._completed_batch_failure_selection(
            [8, 2, 5], [8, 2],
        )
        self.assertEqual(selection["policy"], bridge.FAILURE_SELECTION_POLICY)
        self.assertEqual(selection["completed_batch_chunk_indexes"], [2, 5, 8])
        self.assertEqual(selection["failed_chunk_indexes_in_batch"], [2, 8])
        self.assertEqual(selection["primary_chunk_index"], 2)
        self.assertIs(selection["chronological_first_failure_claimed"], False)
        with self.assertRaisesRegex(ValueError, "failures in the completed batch"):
            bridge._completed_batch_failure_selection([2, 5], [8])

    @unittest.skipUnless(os.name == "posix", "POSIX process groups required")
    def test_cancellation_signals_owned_group_after_leader_exits(self) -> None:
        process = Mock(pid=54321)
        process.poll.return_value = 0
        with patch.object(bridge.os, "killpg") as killpg:
            bridge.RunController._terminate(process)
        killpg.assert_called_once_with(process.pid, bridge.signal.SIGTERM)

    @staticmethod
    def _fake_independent_review(
        response: dict, chunk: dict, *, review_dir: Path, run_id: str,
        chunk_index: int, attempt: int, result_builder=None, **_kwargs,
    ) -> dict:
        """Offline, source-bound reviewer double for bridge orchestration tests."""
        response_sha = bridge._response_sha256(response)
        provenance = response.get("provenance")
        out_dir = review_dir / f"independent-review-chunk-{chunk_index:04d}-attempt-{attempt:02d}"
        out_dir.mkdir(parents=True, exist_ok=True)
        audit_path = out_dir / "coverage-audit.json"
        request = bridge.build_obligation_coverage_request(
            response, chunk, run_id=run_id, chunk_index=chunk_index,
        )
        request["attempt"] = attempt
        request["provider_attempt"] = 1
        source_packet = build_source_reference_packet(request)
        raw_response = {"results": []}
        for check in source_packet["checks"]:
            source_ref = check["source_spans"][0]["ref_id"]
            default_result = {
                "check_id": check["check_id"],
                "verdict": "consistent",
                "rationale": "offline bridge integration fixture",
                "evidence_refs": [source_ref],
                "identified_obligations": [],
            }
            raw_response["results"].append(
                result_builder(check, source_ref) if result_builder else default_result
            )
        compiled_response, compilation = bridge.compile_source_reference_response(
            raw_response, request, bridge.OBLIGATION_COVERAGE_SCHEMA, coverage=True,
        )
        if result_builder:
            bridge.validate_obligation_coverage_response(
                copy.deepcopy(compiled_response), request["checks"],
            )
        request_path = out_dir / "request.json"
        source_packet_path = out_dir / "source-reference-packet.json"
        raw_response_path = out_dir / "raw-response.json"
        compiled_response_path = out_dir / "compiled-response.json"
        response_path = out_dir / "response.json"
        compilation_path = out_dir / "source-reference-compilation.json"
        for path, payload in (
            (request_path, request), (source_packet_path, source_packet),
            (raw_response_path, raw_response), (compiled_response_path, compiled_response),
            (response_path, compiled_response), (compilation_path, compilation),
        ):
            bridge._write_json(path, payload)
        request_sha = bridge.sha256_json(request)
        response_file_sha = bridge.sha256_file(response_path)
        review_audit = {
            "status": "completed", "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
            "adapter_id": "openclaw", "host_runtime": "openclaw",
            "request_path": str(request_path.resolve()), "request_sha256": request_sha,
            "response_path": str(response_path.resolve()), "response_sha256": response_file_sha,
            "raw_response_path": str(raw_response_path.resolve()),
            "raw_response_file_sha256": bridge.sha256_file(raw_response_path),
            "compiled_response_path": str(compiled_response_path.resolve()),
            "compiled_response_sha256": bridge.sha256_file(compiled_response_path),
            "canonical_response_sha256": bridge.sha256_json(compiled_response),
            "source_reference_protocol": "semantic_source_references_v2",
            "source_reference_packet_path": str(source_packet_path.resolve()),
            "source_reference_packet_sha256": bridge.sha256_file(source_packet_path),
            "source_reference_compilation_path": str(compilation_path.resolve()),
            "source_reference_compilation_sha256": bridge.sha256_file(compilation_path),
            "results": copy.deepcopy(compiled_response["results"]),
        }
        ledger_pointer = bridge._write_obligation_analysis_ledger(
            review_audit, response, chunk,
            coverage_request=request, output_dir=out_dir, review_dir=review_dir,
            run_id=run_id, chunk_index=chunk_index, attempt=attempt,
        )
        envelope = {
            "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
            "status": "completed", "run_id": provenance.get("run_id") if isinstance(provenance, dict) else None,
            "chunk_index": chunk_index, "attempt": attempt, "provider_attempt": 1,
            "retry_feedback": None,
            "candidate_response_sha256": response_sha,
            "provenance": copy.deepcopy(provenance),
            "review_request_sha256": request_sha,
            "review_response_sha256": response_file_sha,
            "review_audit": review_audit,
            "results": copy.deepcopy(compiled_response["results"]),
            "obligation_analysis_ledger": ledger_pointer,
        }
        bridge._write_json(audit_path, envelope)
        return {
            "status": "completed", "protocol": envelope["protocol"],
            "audit_path": audit_path.relative_to(review_dir).as_posix(),
            "audit_sha256": bridge.sha256_file(audit_path), "run_id": run_id,
            "chunk_index": chunk_index, "candidate_response_sha256": response_sha,
            "review_request_sha256": request_sha,
            "review_response_sha256": response_file_sha,
            "retry_feedback": None,
            "obligation_analysis_ledger_path": ledger_pointer["path"],
            "obligation_analysis_ledger_sha256": ledger_pointer["sha256"],
            "obligation_analysis_ledger_status": ledger_pointer["status"],
            "summary": {"consistent": len(compiled_response["results"]), "incomplete": 0, "uncertain": 0},
        }

    def test_independent_coverage_gate_persists_run_and_candidate_binding(self) -> None:
        self._independent_review_patch.stop()
        source = "表格应居中"
        provenance = {
            "run_id": "run-coverage-test", "source_sha256": "a" * 64,
            "clause_sha256": "b" * 64, "evidence_sha256": "c" * 64,
            "request_sha256": "d" * 64,
        }
        chunk = {
            "provenance": provenance,
            "clauses": [{
                "id": "C1", "text": source, "evidence_ids": ["E1"],
                "source_span": {
                    "evidence_id": "E1", "start_offset": 0, "end_offset": len(source),
                    "text": source, "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
                },
            }],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [{
                "clause_id": "C1", "classification": "covered", "reason": "centered",
                "obligations": [{"id": "O1", "status": "covered", "reason": "center"}],
            }],
            "requirements": [{
                "id": "R1", "role": "table", "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "properties": {"paragraph": {"alignment": "center"}},
                "verification": {"checker_ids": ["docx.property_receipts"]},
            }],
        }
        reviewer_result = {
            "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
            "status": "completed",
            "response_sha256": "f" * 64,
            "results": [{
                "check_id": "C1", "verdict": "consistent", "rationale": "source represented",
                "evidence_quotes": [source], "identified_obligations": [{
                    "source_quote": source, "disposition": "represented", "requirement_refs": ["RR-test"],
                }],
                "machine_obligation_ids": ["table_caption.alignment_center"],
            }],
            "summary": {"consistent": 1, "incomplete": 0, "uncertain": 0},
        }
        reviewer_result["canonical_response_sha256"] = sha256_json({
            "results": reviewer_result["results"],
        })
        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td)
            coverage_request = bridge.build_obligation_coverage_request(
                response, chunk, run_id="run-coverage-test", chunk_index=1,
            )
            coverage_request["attempt"] = 1
            coverage_request["provider_attempt"] = 1
            source_packet = build_source_reference_packet(coverage_request)
            source_span = source_packet["checks"][0]["source_spans"][0]
            reviewer_result["request_sha256"] = bridge.sha256_json(coverage_request)
            compilation_path = (
                review_dir / "independent-review-chunk-0001-attempt-01"
                / "source-reference-compilation.json"
            )
            compilation_path.parent.mkdir(parents=True)
            compilation = {
                "protocol": "semantic_source_references_v2",
                "run_id": "run-coverage-test",
                "request_sha256": reviewer_result["request_sha256"],
                "packet_sha256": bridge.sha256_json(source_packet),
                "compiled_response_sha256": reviewer_result["canonical_response_sha256"],
                "selections": [{
                    "check_id": "C1",
                    "spans": [source_span],
                    "obligations": [{
                        "obligation_index": 0,
                        "source_ref": source_span["ref_id"],
                        "span": source_span,
                    }],
                }],
            }
            bridge._write_json(compilation_path, compilation)
            reviewer_result["source_reference_compilation_path"] = str(compilation_path)
            reviewer_result["source_reference_compilation_sha256"] = bridge.sha256_file(
                compilation_path
            )
            with patch.object(bridge, "run_native_semantic_review", return_value=reviewer_result):
                pointer = bridge._run_independent_obligation_coverage_review(
                    response, chunk, review_dir=review_dir, run_id="run-coverage-test",
                    chunk_index=1, attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                    timeout=10, agent_id="main", runner="exec", binary="codex",
                    config_path=None, controller=bridge.RunController(),
                )
            audit_path = review_dir / pointer["audit_path"]
            envelope = json.loads(audit_path.read_text(encoding="utf-8"))
            self.assertEqual(pointer["status"], "completed")
            self.assertEqual(pointer["audit_sha256"], bridge.sha256_file(audit_path))
            self.assertEqual(envelope["candidate_response_sha256"], bridge._response_sha256(response))
            self.assertEqual(envelope["provenance"], provenance)
            self.assertEqual(envelope["run_id"], "run-coverage-test")
            ledger = json.loads(
                (review_dir / pointer["obligation_analysis_ledger_path"]).read_text(encoding="utf-8")
            )
            self.assertFalse(ledger["submission_ready"])
            self.assertFalse(ledger["obligations"][0]["execution_authorized"])
            self.assertEqual(ledger["obligations"][0]["source_quote"], source)

    def test_c00076_scope_unresolved_is_persisted_only_as_non_executable_analysis(self) -> None:
        self._independent_review_patch.stop()
        source = (
            "The Chinese abstract is usually written in third person, "
            "300 to 1,000 words, without comment and explanation."
        )
        provenance = {
            "run_id": "run-c76-ledger", "source_sha256": "a" * 64,
            "clause_sha256": "b" * 64, "evidence_sha256": "c" * 64,
            "request_sha256": "d" * 64,
        }
        chunk = {
            "provenance": provenance,
            "clauses": [exact_source_clause("C00076", source, "E00076")],
            "evidence_context": {"E00076": {"id": "E00076", "text": source}},
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [{
                "clause_id": "C00076", "classification": "unresolved",
                "reason": "The target abstract and metric remain ambiguous.",
            }],
            "requirements": [],
        }
        reviewer_result = {
            "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
            "status": "completed", "response_sha256": "f" * 64,
            "results": [{
                "check_id": "C00076", "verdict": "manual_review_required",
                "rationale": "The scope depends on the registered target and metric ambiguity.",
                "evidence_quotes": ["300 to 1,000 words", "without comment and explanation"],
                "machine_obligation_ids": [],
                "identified_obligations": [
                    {
                        "source_quote": "300 to 1,000 words", "disposition": "scope_unresolved",
                        "obligation_summary": "The word-count guidance has an unresolved target.",
                        "scope_dependency_codes": ["abstract_target_metric_ambiguity"],
                        "scope_dependency_dimensions": ["target", "metric"],
                        "requirement_refs": [],
                    },
                    {
                        "source_quote": "without comment and explanation",
                        "disposition": "unrepresented", "requirement_refs": [],
                    },
                ],
            }],
            "summary": {"consistent": 0, "incomplete": 0, "uncertain": 0},
        }
        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td)
            with patch.object(
                bridge, "run_native_semantic_review",
                side_effect=lambda request, **kwargs: bind_mock_review_to_source_spans(
                    reviewer_result, request, kwargs["output_dir"],
                ),
            ):
                pointer = bridge._run_independent_obligation_coverage_review(
                    response, chunk, review_dir=review_dir, run_id="run-c76-ledger",
                    chunk_index=1, attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                    timeout=10, agent_id="main", runner="exec", binary="codex",
                    config_path=None, controller=bridge.RunController(),
                )
            ledger = json.loads(
                (review_dir / pointer["obligation_analysis_ledger_path"]).read_text(encoding="utf-8")
            )
            self.assertEqual(ledger["run_id"], "run-c76-ledger")
            self.assertEqual(ledger["candidate_response_sha256"], bridge._response_sha256(response))
            self.assertFalse(ledger["submission_ready"])
            by_disposition = {item["disposition"]: item for item in ledger["obligations"]}
            self.assertEqual(set(by_disposition), {"scope_unresolved", "unrepresented"})
            for obligation in by_disposition.values():
                self.assertEqual(obligation["requirement_refs"], [])
                self.assertFalse(obligation["execution_authorized"])

    def test_independent_coverage_retries_capacity_once_with_same_model_and_fresh_attempt_dir(self) -> None:
        self._independent_review_patch.stop()
        source = "表格应居中"
        provenance = {
            "run_id": "run-capacity-retry",
            "source_sha256": "a" * 64,
            "clause_sha256": "b" * 64,
            "evidence_sha256": "c" * 64,
            "request_sha256": "d" * 64,
        }
        chunk = {
            "provenance": provenance,
            "clauses": [exact_source_clause("C1", source)],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [{
                "clause_id": "C1", "classification": "covered", "reason": "centered",
                "obligations": [{"id": "O1", "status": "covered", "reason": "center"}],
            }],
            "requirements": [{
                "id": "R1", "role": "table", "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "properties": {"paragraph": {"alignment": "center"}},
                "verification": {"checker_ids": ["docx.property_receipts"]},
            }],
        }
        review_result = {
            "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
            "status": "completed", "request_sha256": "e" * 64,
            "response_sha256": "f" * 64,
            "results": [{
                "check_id": "C1", "verdict": "consistent", "rationale": "source represented",
                "evidence_quotes": [source], "identified_obligations": [{
                    "source_quote": source, "disposition": "represented", "requirement_refs": ["RR-test"],
                }], "machine_obligation_ids": ["table_caption.alignment_center"],
            }],
            "summary": {"consistent": 1, "incomplete": 0, "uncertain": 0},
        }
        calls: list[tuple[dict, dict]] = []

        def fail_once_for_capacity(request: dict, **kwargs: dict) -> dict:
            calls.append((copy.deepcopy(request), kwargs))
            if len(calls) == 1:
                raise bridge.RetryableNativeSemanticReviewError(
                    "Codex reported model capacity", retry_code="model_capacity",
                )
            return bind_mock_review_to_source_spans(
                review_result, request, kwargs["output_dir"],
            )

        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td)
            with patch.object(bridge, "run_native_semantic_review", side_effect=fail_once_for_capacity), \
                    patch.object(bridge.time, "sleep") as sleep:
                pointer = bridge._run_independent_obligation_coverage_review(
                    response, chunk, review_dir=review_dir, run_id="run-capacity-retry",
                    chunk_index=3, attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                    timeout=10, agent_id="main", runner="exec", binary="codex",
                    config_path=None, controller=bridge.RunController(), codex_reasoning_effort="max",
                )

            self.assertEqual([kwargs["reasoning_effort"] for _, kwargs in calls], ["max"] * 2)
            self.assertEqual(len(calls), 2)
            self.assertEqual([kwargs["model"] for _, kwargs in calls], ["gpt-5.6-luna"] * 2)
            self.assertNotEqual(calls[0][1]["output_dir"], calls[1][1]["output_dir"])
            self.assertEqual([request["provider_attempt"] for request, _ in calls], [1, 2])
            self.assertEqual([request["attempt"] for request, _ in calls], [1, 1])
            self.assertEqual([request["provenance"] for request, _ in calls], [provenance, provenance])
            sleep.assert_called_once_with(bridge.INDEPENDENT_REVIEW_RETRY_BACKOFF_SECONDS)

            first_audit = review_dir / "independent-review-chunk-0003-attempt-01" / "coverage-audit.json"
            failed = json.loads(first_audit.read_text(encoding="utf-8"))
            self.assertEqual(failed["retry_code"], "model_capacity")
            self.assertTrue(failed["retryable"])

            accepted_audit = json.loads((review_dir / pointer["audit_path"]).read_text(encoding="utf-8"))
            self.assertEqual(pointer["provider_attempt"], 2)
            self.assertEqual(accepted_audit["provider_attempt_history"][0]["retry_code"], "model_capacity")
            self.assertEqual(accepted_audit["candidate_response_sha256"], bridge._response_sha256(response))

    @staticmethod
    def _missing_inventory_review_case() -> tuple[dict, dict, dict]:
        source = "提交的学位论文电子版与纸质本论文的内容一致，如因不同造成不良后果由本人自负"
        provenance = {
            "run_id": "run-empty-inventory-correction",
            "source_sha256": "a" * 64,
            "clause_sha256": "b" * 64,
            "evidence_sha256": "c" * 64,
            "request_sha256": "d" * 64,
        }
        chunk = {
            "provenance": provenance,
            "clauses": [{
                "id": "C00061", "text": source, "evidence_ids": ["E1"],
                "source_span": {
                    "evidence_id": "E1", "start_offset": 0, "end_offset": len(source),
                    "text": source, "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
                },
            }],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [{
                "clause_id": "C00061", "classification": "executable",
                "reason": "This declaration must be preserved.",
                "obligations": [{"id": "O1", "status": "covered", "reason": "fixed text"}],
            }],
            "requirements": [{
                "id": "R1", "role": "declarations", "clause_ids": ["C00061"],
                "evidence_ids": ["E1"], "properties": {"declarations": {"body": source}},
                "verification": {"checker_ids": ["docx.declaration_text"]},
            }],
        }
        request = bridge.build_obligation_coverage_request(
            response, chunk, run_id="run-empty-inventory-correction", chunk_index=4,
        )
        requirement_ref = request["checks"][0]["review_context"]["linked_requirements"][0]["requirement_ref"]
        valid_review = {
            "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
            "status": "completed", "request_sha256": "e" * 64,
            "response_sha256": "f" * 64,
            "results": [{
                "check_id": "C00061", "verdict": "consistent",
                "rationale": "The fixed declaration is represented by its linked requirement.",
                "evidence_quotes": [source],
                "identified_obligations": [{
                    "source_quote": source, "disposition": "represented",
                    "requirement_refs": [requirement_ref],
                }],
                "machine_obligation_ids": [],
            }],
            "summary": {"consistent": 1, "incomplete": 0, "uncertain": 0},
        }
        return chunk, response, valid_review

    def test_independent_coverage_corrects_only_empty_executable_inventory_once(self) -> None:
        self._independent_review_patch.stop()
        chunk, response, valid_review = self._missing_inventory_review_case()
        initial_candidate_sha = bridge._response_sha256(response)
        calls: list[tuple[dict, dict]] = []

        def reject_then_review(request: dict, **kwargs: dict) -> dict:
            calls.append((copy.deepcopy(request), kwargs))
            if len(calls) == 1:
                raise bridge.MissingExecutableObligationInventoryError(["C00061"])
            return bind_mock_review_to_source_spans(
                valid_review, request, kwargs["output_dir"],
            )

        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td)
            with patch.object(bridge, "run_native_semantic_review", side_effect=reject_then_review), \
                    patch.object(bridge.time, "sleep") as sleep:
                pointer = bridge._run_independent_obligation_coverage_review(
                    response, chunk, review_dir=review_dir,
                    run_id="run-empty-inventory-correction", chunk_index=4, attempt=1,
                    host_runtime="codex", model="gpt-5.6-luna", timeout=10,
                    agent_id="main", runner="exec", binary="codex", config_path=None,
                    controller=bridge.RunController(),
                )

            self.assertEqual(len(calls), 2)
            self.assertEqual(calls[0][0]["provenance"], calls[1][0]["provenance"])
            self.assertEqual([item[0]["provider_attempt"] for item in calls], [1, 2])
            self.assertNotEqual(calls[0][1]["output_dir"], calls[1][1]["output_dir"])
            self.assertNotIn("retry_feedback", calls[0][0])
            self.assertEqual(calls[1][0]["retry_feedback"], {
                "code": bridge.MissingExecutableObligationInventoryError.code,
                "clause_ids": ["C00061"],
            })
            sleep.assert_called_once_with(bridge.INDEPENDENT_REVIEW_RETRY_BACKOFF_SECONDS)
            self.assertEqual(bridge._response_sha256(response), initial_candidate_sha)
            self.assertEqual(pointer["candidate_response_sha256"], initial_candidate_sha)

            first_audit_path = review_dir / (
                "independent-review-chunk-0004-attempt-01/coverage-audit.json"
            )
            first_audit = json.loads(first_audit_path.read_text(encoding="utf-8"))
            self.assertEqual(first_audit["status"], "rejected")
            self.assertTrue(first_audit["retryable"])
            self.assertEqual(first_audit["missing_clause_ids"], ["C00061"])

            accepted_audit = json.loads((review_dir / pointer["audit_path"]).read_text(encoding="utf-8"))
            self.assertEqual(accepted_audit["status"], "completed")
            self.assertEqual(accepted_audit["candidate_response_sha256"], initial_candidate_sha)
            self.assertEqual(accepted_audit["provider_attempt_history"][0]["status"], "semantic_contract_rejected")
            self.assertEqual(accepted_audit["retry_feedback"]["clause_ids"], ["C00061"])

    def test_independent_coverage_empty_inventory_correction_exhaustion_fails_closed(self) -> None:
        self._independent_review_patch.stop()
        chunk, response, _ = self._missing_inventory_review_case()
        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td)
            with patch.object(
                bridge, "run_native_semantic_review",
                side_effect=bridge.MissingExecutableObligationInventoryError(["C00061"]),
            ) as review_call, patch.object(bridge.time, "sleep") as sleep:
                with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                    bridge._run_independent_obligation_coverage_review(
                        response, chunk, review_dir=review_dir,
                        run_id="run-empty-inventory-correction", chunk_index=4, attempt=1,
                        host_runtime="codex", model="gpt-5.6-luna", timeout=10,
                        agent_id="main", runner="exec", binary="codex", config_path=None,
                        controller=bridge.RunController(),
                    )

            self.assertEqual(review_call.call_count, 2)
            sleep.assert_called_once_with(bridge.INDEPENDENT_REVIEW_RETRY_BACKOFF_SECONDS)
            self.assertEqual(
                caught.exception.error_records[0]["code"],
                "independent_obligation_review_correction_exhausted",
            )
            self.assertEqual(len(caught.exception.error_records[0]["provider_attempt_history"]), 2)
            for suffix in ("", "-provider-attempt-02"):
                audit_path = review_dir / (
                    "independent-review-chunk-0004-attempt-01" + suffix + "/coverage-audit.json"
                )
                self.assertEqual(json.loads(audit_path.read_text(encoding="utf-8"))["status"], "rejected")

    def test_inconsistent_incomplete_disposition_rereviews_unchanged_candidate_once(self) -> None:
        self._independent_review_patch.stop()
        chunk, response, valid_review = self._missing_inventory_review_case()
        calls: list[dict] = []

        def reject_then_review(request: dict, **kwargs: dict) -> dict:
            calls.append(copy.deepcopy(request))
            if len(calls) == 1:
                raise bridge.InconsistentObligationVerdictError(["C00061"])
            return bind_mock_review_to_source_spans(
                valid_review, request, kwargs["output_dir"],
            )

        with tempfile.TemporaryDirectory() as td:
            with patch.object(bridge, "run_native_semantic_review", side_effect=reject_then_review), \
                    patch.object(bridge.time, "sleep"):
                pointer = bridge._run_independent_obligation_coverage_review(
                    response, chunk, review_dir=Path(td),
                    run_id="run-empty-inventory-correction", chunk_index=4, attempt=1,
                    host_runtime="codex", model="gpt-5.6-luna", timeout=10,
                    agent_id="main", runner="exec", binary="codex", config_path=None,
                    controller=bridge.RunController(),
                )
            self.assertEqual(len(calls), 2)
            self.assertEqual(calls[0]["provenance"], calls[1]["provenance"])
            self.assertEqual(calls[1]["retry_feedback"], {
                "code": bridge.InconsistentObligationVerdictError.code,
                "clause_ids": ["C00061"],
            })
            self.assertEqual(pointer["candidate_response_sha256"], bridge._response_sha256(response))

    def test_unlinked_represented_review_rereads_same_label_once_and_exhausts(self) -> None:
        self._independent_review_patch.stop()
        for exhausted in (False, True):
            with self.subTest(exhausted=exhausted), tempfile.TemporaryDirectory() as td:
                review_dir, chunk = self._packet(Path(td) / "requirements", source="论文题目",
                                                  contract_version="3.0")
                response = {"contract_version": "3.0", "provenance": chunk["provenance"],
                    "requirements": [], "clause_reviews": [{"clause_id": "C1",
                        "classification": "informational", "reason": "Source row label."}],
                    "unsupported_items": [], "reported_conflicts": []}
                before = copy.deepcopy(response)
                calls = []

                def review(request: dict, *, output_dir: Path, **kwargs: object) -> dict:
                    calls.append(copy.deepcopy(request))
                    if len(calls) == 1 or exhausted:
                        raise bridge.UnlinkedRepresentedObligationError(["C1"])
                    return bind_mock_review_to_source_spans({"status": "completed",
                        "results": [{"check_id": "C1", "verdict": "consistent",
                            "rationale": "The source is a label, not a normative duty.",
                            "evidence_quotes": ["论文题目"], "machine_obligation_ids": [],
                            "identified_obligations": []}], "summary": {"consistent": 1}},
                        request, output_dir)

                with patch.object(bridge, "run_native_semantic_review", side_effect=review), \
                        patch.object(bridge.time, "sleep"):
                    args = dict(review_dir=review_dir, run_id=chunk["provenance"]["run_id"],
                        chunk_index=1, attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                        timeout=10, agent_id="main", runner="exec", binary="codex", config_path=None,
                        controller=bridge.RunController(), output_policy="review_draft")
                    if exhausted:
                        with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                            bridge._run_independent_obligation_coverage_review(response, chunk, **args)
                        self.assertFalse(caught.exception.retryable)
                    else:
                        pointer = bridge._run_independent_obligation_coverage_review(response, chunk, **args)
                        self.assertEqual(pointer["status"], "completed")
                self.assertEqual(len(calls), 2)
                self.assertEqual(calls[0]["checks"], calls[1]["checks"])
                self.assertEqual(calls[0]["provenance"], calls[1]["provenance"])
                self.assertEqual(calls[1]["retry_feedback"], {
                    "code": bridge.UnlinkedRepresentedObligationError.code, "clause_ids": ["C1"]})
                self.assertEqual(response, before)
                rejected = json.loads((review_dir / "independent-review-chunk-0001-attempt-01/coverage-audit.json").read_text())
                self.assertEqual(rejected["status"], "rejected")
                self.assertEqual(rejected["candidate_response_sha256"], bridge._response_sha256(before))

    def test_mislabelled_source_verification_rereviews_same_candidate_once(self) -> None:
        self._independent_review_patch.stop()
        source = (
            "关键词是为了便于做文献索引和检索工作而从论文中选取出来用以表示全文主题内容信息的"
            "单词或术语，在论文中有明确出处"
        )
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(
                Path(td) / "requirements", source=source, contract_version="3.0",
            )
            response = {
                "contract_version": "3.0", "provenance": chunk["provenance"],
                "requirements": [], "clause_reviews": [{
                    "clause_id": "C1", "classification": "requires_source_verification",
                    "reason": "The existing keywords require human traceability verification.",
                    "obligations": [],
                }], "unsupported_items": [], "reported_conflicts": [],
            }
            accepted_review = {
                "status": "completed", "results": [{
                    "check_id": "C1", "verdict": "source_content_verification_pending",
                    "rationale": "Human verification of keyword origin is outstanding.",
                    "evidence_quotes": [source], "machine_obligation_ids": [],
                    "identified_obligations": [{
                        "source_quote": source,
                        "disposition": "source_content_verification_pending",
                        "requirement_refs": [],
                    }],
                }], "summary": {"source_content_verification_pending": 1},
            }
            calls: list[dict] = []

            def first_mislabels_then_corrects(request: dict, *, output_dir: Path, **_kwargs: object) -> dict:
                calls.append(copy.deepcopy(request))
                if len(calls) == 1:
                    raise bridge.SourceVerificationMislabelledAsAuthoringError(["C1"])
                return bind_mock_review_to_source_spans(accepted_review, request, output_dir)

            with patch.object(bridge, "run_native_semantic_review", side_effect=first_mislabels_then_corrects), \
                    patch.object(bridge.time, "sleep"):
                pointer = bridge._run_independent_obligation_coverage_review(
                    response, chunk, review_dir=review_dir,
                    run_id=chunk["provenance"]["run_id"], chunk_index=1,
                    attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                    timeout=10, agent_id="main", runner="exec", binary="codex",
                    config_path=None, controller=bridge.RunController(),
                )
            self.assertEqual(len(calls), 2)
            self.assertEqual(calls[0]["checks"], calls[1]["checks"])
            self.assertEqual(calls[0]["provenance"], calls[1]["provenance"])
            self.assertEqual(calls[1]["retry_feedback"], {
                "code": bridge.SourceVerificationMislabelledAsAuthoringError.code,
                "clause_ids": ["C1"],
            })
            self.assertEqual(pointer["status"], "completed")
            self.assertEqual(pointer["candidate_response_sha256"], bridge._response_sha256(response))
            rejected = json.loads((review_dir / "independent-review-chunk-0001-attempt-01/coverage-audit.json").read_text(encoding="utf-8"))
            self.assertEqual(rejected["status"], "rejected")
            self.assertTrue(rejected["source_bound_retry_authorized"])

    def test_mislabelled_source_verification_without_registered_code_does_not_retry(self) -> None:
        self._independent_review_patch.stop()
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(
                Path(td) / "requirements", source="表格应居中", contract_version="3.0",
            )
            response = {
                "contract_version": "3.0", "provenance": chunk["provenance"],
                "requirements": [], "clause_reviews": [{
                    "clause_id": "C1", "classification": "requires_source_content",
                    "reason": "pending", "obligations": [],
                }], "unsupported_items": [], "reported_conflicts": [],
            }
            with patch.object(
                bridge, "run_native_semantic_review",
                side_effect=bridge.SourceVerificationMislabelledAsAuthoringError(["C1"]),
            ) as call:
                with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                    bridge._run_independent_obligation_coverage_review(
                        response, chunk, review_dir=review_dir,
                        run_id=chunk["provenance"]["run_id"], chunk_index=1,
                        attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                        timeout=10, agent_id="main", runner="exec", binary="codex",
                        config_path=None, controller=bridge.RunController(),
                    )
            self.assertEqual(call.call_count, 1)
            self.assertFalse(caught.exception.retryable)
            rejected = json.loads((review_dir / "independent-review-chunk-0001-attempt-01/coverage-audit.json").read_text(encoding="utf-8"))
            self.assertFalse(rejected["source_bound_retry_authorized"])

    def test_inconsistent_incomplete_disposition_exhaustion_fails_closed(self) -> None:
        self._independent_review_patch.stop()
        chunk, response, _ = self._missing_inventory_review_case()
        with tempfile.TemporaryDirectory() as td:
            with patch.object(
                bridge, "run_native_semantic_review",
                side_effect=bridge.InconsistentObligationVerdictError(["C00061"]),
            ) as review_call, patch.object(bridge.time, "sleep"):
                with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                    bridge._run_independent_obligation_coverage_review(
                        response, chunk, review_dir=Path(td),
                        run_id="run-empty-inventory-correction", chunk_index=4, attempt=1,
                        host_runtime="codex", model="gpt-5.6-luna", timeout=10,
                        agent_id="main", runner="exec", binary="codex", config_path=None,
                        controller=bridge.RunController(),
                    )
            self.assertEqual(review_call.call_count, 2)
            self.assertEqual(
                caught.exception.error_records[0]["retry_code"],
                bridge.InconsistentObligationVerdictError.code,
            )

    def test_independent_coverage_does_not_retry_unclassified_provider_or_contract_errors(self) -> None:
        self._independent_review_patch.stop()
        source = "表格应居中"
        provenance = {"run_id": "run-no-retry"}
        chunk = {
            "provenance": provenance,
            "clauses": [exact_source_clause("C1", source)],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [{"clause_id": "C1", "classification": "covered", "reason": "centered"}],
            "requirements": [{
                "id": "R1", "role": "table", "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "properties": {"paragraph": {"alignment": "center"}},
            }],
        }
        with tempfile.TemporaryDirectory() as td:
            with patch.object(
                bridge, "run_native_semantic_review", side_effect=RuntimeError("invalid response contract"),
            ) as review_call:
                with self.assertRaises(bridge.IndependentObligationReviewError):
                    bridge._run_independent_obligation_coverage_review(
                        response, chunk, review_dir=Path(td), run_id="run-no-retry",
                        chunk_index=1, attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                        timeout=10, agent_id="main", runner="exec", binary="codex",
                        config_path=None, controller=bridge.RunController(),
                    )
            self.assertEqual(review_call.call_count, 1)

    def test_independent_coverage_capacity_exhaustion_fails_closed_after_two_calls(self) -> None:
        self._independent_review_patch.stop()
        source = "表格应居中"
        provenance = {"run_id": "run-capacity-exhausted"}
        chunk = {
            "provenance": provenance,
            "clauses": [exact_source_clause("C1", source)],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [{"clause_id": "C1", "classification": "covered", "reason": "centered"}],
            "requirements": [{
                "id": "R1", "role": "table", "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "properties": {"paragraph": {"alignment": "center"}},
            }],
        }
        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td)
            with patch.object(
                bridge,
                "run_native_semantic_review",
                side_effect=bridge.RetryableNativeSemanticReviewError(
                    "Codex reported model capacity", retry_code="model_capacity",
                ),
            ) as review_call, patch.object(bridge.time, "sleep") as sleep:
                with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                    bridge._run_independent_obligation_coverage_review(
                        response, chunk, review_dir=review_dir, run_id="run-capacity-exhausted",
                        chunk_index=2, attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                        timeout=10, agent_id="main", runner="exec", binary="codex",
                        config_path=None, controller=bridge.RunController(),
                    )

            self.assertEqual(review_call.call_count, bridge.INDEPENDENT_REVIEW_PROVIDER_MAX_ATTEMPTS)
            sleep.assert_called_once_with(bridge.INDEPENDENT_REVIEW_RETRY_BACKOFF_SECONDS)
            record = caught.exception.error_records[0]
            self.assertEqual(record["code"], "independent_obligation_review_retry_exhausted")
            self.assertEqual(record["provider_attempts"], 2)
            self.assertEqual(len(record["provider_attempt_history"]), 2)
            for suffix in ("", "-provider-attempt-02"):
                attempt_dir = review_dir / ("independent-review-chunk-0002-attempt-01" + suffix)
                audit_path = attempt_dir / "coverage-audit.json"
                self.assertTrue(audit_path.is_file())
                self.assertEqual(json.loads(audit_path.read_text())["status"], "failed")

    def test_independent_coverage_gate_rejects_unrepresented_source_obligation(self) -> None:
        self._independent_review_patch.stop()
        source = "表格应居中"
        provenance = {"run_id": "run-coverage-test"}
        chunk = {
            "provenance": provenance,
            "clauses": [exact_source_clause("C1", source)],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [{"clause_id": "C1", "classification": "covered", "reason": "covered"}],
            "requirements": [{
                "id": "R1", "role": "table", "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "properties": {},
            }],
        }
        reviewer_result = {
            "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
            "status": "completed", "request_sha256": "e" * 64,
            "response_sha256": "f" * 64,
            "results": [{
                "check_id": "C1", "verdict": "incomplete", "rationale": "alignment missing",
                "evidence_quotes": [source], "identified_obligations": [{
                    "source_quote": source, "disposition": "unrepresented", "requirement_refs": [],
                }], "machine_obligation_ids": ["table_caption.alignment_center"],
            }],
            "summary": {"consistent": 0, "incomplete": 1, "uncertain": 0},
        }
        with tempfile.TemporaryDirectory() as td:
            with patch.object(
                bridge, "run_native_semantic_review",
                side_effect=lambda request, **kwargs: bind_mock_review_to_source_spans(
                    reviewer_result, request, kwargs["output_dir"],
                ),
            ):
                with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                    bridge._run_independent_obligation_coverage_review(
                        response, chunk, review_dir=Path(td), run_id="run-coverage-test",
                        chunk_index=1, attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                        timeout=10, agent_id="main", runner="exec", binary="codex",
                        config_path=None, controller=bridge.RunController(),
                    )
            self.assertEqual(
                caught.exception.error_records[0]["code"],
                "independent_obligation_review_incomplete",
            )
            self.assertTrue(caught.exception.retryable)
            pointer = caught.exception.independent_review_audit
            envelope = json.loads((Path(td) / pointer["audit_path"]).read_text(encoding="utf-8"))
            self.assertEqual(pointer["status"], "rejected")
            self.assertEqual(envelope["status"], "rejected")

    def test_executable_gap_routes_to_primary_without_rewriting_other_inventories(self) -> None:
        self._independent_review_patch.stop()
        source_missing = "表格应居中"
        source_represented = "图题应居中"
        provenance = {"run_id": "run-primary-gap"}
        chunk = {
            "provenance": provenance,
            "clauses": [
                exact_source_clause("C1", source_missing, "E1"),
                exact_source_clause("C2", source_represented, "E2"),
            ],
            "evidence_context": {
                "E1": {"id": "E1", "text": source_missing},
                "E2": {"id": "E2", "text": source_represented},
            },
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [
                {"clause_id": "C1", "classification": "covered", "reason": "covered"},
                {"clause_id": "C2", "classification": "covered", "reason": "covered"},
            ],
            "requirements": [
                {"id": "R1", "role": "table", "clause_ids": ["C1"],
                 "evidence_ids": ["E1"], "properties": {}},
                {"id": "R2", "role": "figure_caption", "clause_ids": ["C2"],
                 "evidence_ids": ["E2"], "properties": {}},
            ],
        }
        initial_response_sha = bridge._response_sha256(response)

        def independent_review(request: dict, **kwargs: dict) -> dict:
            checks = {item["check_id"]: item for item in request["checks"]}
            represented_ref = checks["C2"]["review_context"]["linked_requirements"][0][
                "requirement_ref"
            ]
            review_result = {
                "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
                "status": "completed", "response_sha256": "f" * 64,
                "results": [
                    {
                        "check_id": "C1", "verdict": "incomplete",
                        "rationale": "the candidate does not preserve the alignment obligation",
                        "evidence_quotes": [source_missing],
                        "identified_obligations": [{
                            "source_quote": source_missing,
                            "disposition": "unrepresented", "requirement_refs": [],
                        }],
                        "machine_obligation_ids": checks["C1"]["review_context"][
                            "machine_obligation_ids"
                        ],
                    },
                    {
                        "check_id": "C2", "verdict": "consistent",
                        "rationale": "the linked requirement represents this source obligation",
                        "evidence_quotes": [source_represented],
                        "identified_obligations": [{
                            "source_quote": source_represented,
                            "disposition": "represented",
                            "requirement_refs": [represented_ref],
                        }],
                        "machine_obligation_ids": checks["C2"]["review_context"][
                            "machine_obligation_ids"
                        ],
                    },
                ],
                "summary": {"consistent": 1, "incomplete": 1, "uncertain": 0},
            }
            bound = bind_mock_review_to_source_spans(
                review_result, request, kwargs["output_dir"],
            )
            bridge.validate_obligation_coverage_response(
                {"results": copy.deepcopy(bound["results"])}, request["checks"],
            )
            return bound

        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td)
            with patch.object(bridge, "run_native_semantic_review", side_effect=independent_review) as review_call, \
                    patch.object(bridge.time, "sleep") as sleep:
                with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                    bridge._run_independent_obligation_coverage_review(
                        response, chunk, review_dir=review_dir,
                        run_id="run-primary-gap", chunk_index=1, attempt=1,
                        host_runtime="codex", model="gpt-5.6-luna", timeout=10,
                        agent_id="main", runner="exec", binary="codex",
                        config_path=None, controller=bridge.RunController(),
                    )
            review_call.assert_called_once()
            sleep.assert_not_called()
            self.assertTrue(getattr(caught.exception, "retryable", False), str(caught.exception))
            self.assertEqual(bridge._response_sha256(response), initial_response_sha)
            record = caught.exception.error_records[0]
            self.assertEqual(record["clause_id"], "C1")
            self.assertEqual(record["primary_retry_authorization"],
                             "executable_requirement_completion")
            pointer = caught.exception.independent_review_audit
            self.assertEqual(pointer["provider_attempt"], 1)
            ledger = json.loads((review_dir / pointer["obligation_analysis_ledger_path"]).read_text())
            self.assertEqual({item["check_id"] for item in ledger["obligations"]}, {"C1", "C2"})
            self.assertFalse((review_dir / "independent-review-chunk-0001-attempt-01-provider-attempt-02").exists())

    @staticmethod
    def _external_compliance_review_case() -> tuple[dict, dict, dict, str]:
        source = "学位论文作者签名： 年 月 日"
        provenance = {
            "run_id": "run-external-correction",
            "source_sha256": "a" * 64,
            "clause_sha256": "b" * 64,
            "evidence_sha256": "c" * 64,
            "request_sha256": "d" * 64,
        }
        chunk = {
            "provenance": provenance,
            "clauses": [{
                "id": "C00037", "text": source, "evidence_ids": ["E1"],
                "source_span": {
                    "evidence_id": "E1", "start_offset": 0, "end_offset": len(source),
                    "text": source, "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
                },
            }],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [{
                "clause_id": "C00037", "classification": "external_compliance",
                "reason": "The author must sign and date the declaration.",
                "obligations": [],
            }],
            "requirements": [],
        }
        request = bridge.build_obligation_coverage_request(
            response, chunk, run_id="run-external-correction", chunk_index=2,
        )
        machine_ids = request["checks"][0]["review_context"]["machine_obligation_ids"]
        accepted_review = {
            "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
            "status": "completed", "request_sha256": "e" * 64,
            "response_sha256": "f" * 64,
            "results": [{
                "check_id": "C00037", "verdict": "external_compliance_pending",
                "rationale": "The author signature and date are external actions, not DOCX work.",
                "evidence_quotes": [source], "machine_obligation_ids": machine_ids,
                "identified_obligations": [{
                    "source_quote": source, "disposition": "external_action_pending",
                    "requirement_refs": [],
                }],
            }],
            "summary": {"consistent": 0, "incomplete": 0, "uncertain": 0},
        }
        return chunk, response, accepted_review, source

    def test_external_compliance_incomplete_review_gets_one_same_candidate_correction(self) -> None:
        self._independent_review_patch.stop()
        chunk, response, accepted_review, source = self._external_compliance_review_case()
        initial_response_sha = bridge._response_sha256(response)
        correction = bridge.ExternalComplianceCorrectionRequiredError([{
            "check_id": "C00037", "source_quotes": [source],
        }])
        calls: list[tuple[dict, dict]] = []

        def correct_external_review(request: dict, **kwargs: dict) -> dict:
            calls.append((copy.deepcopy(request), kwargs))
            if len(calls) == 1:
                raise correction
            return bind_mock_review_to_source_spans(
                accepted_review, request, kwargs["output_dir"],
            )

        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td)
            with patch.object(bridge, "run_native_semantic_review", side_effect=correct_external_review), \
                    patch.object(bridge.time, "sleep") as sleep:
                pointer = bridge._run_independent_obligation_coverage_review(
                    response, chunk, review_dir=review_dir, run_id="run-external-correction",
                    chunk_index=2, attempt=1, host_runtime="codex", model="gpt-5.6-luna",
                    timeout=10, agent_id="main", runner="exec", binary="codex",
                    config_path=None, controller=bridge.RunController(),
                )

            self.assertEqual(len(calls), 2)
            self.assertEqual([call[0]["provider_attempt"] for call in calls], [1, 2])
            self.assertEqual([call[1]["model"] for call in calls], ["gpt-5.6-luna"] * 2)
            self.assertNotEqual(calls[0][1]["output_dir"], calls[1][1]["output_dir"])
            self.assertEqual(calls[0][0]["provenance"], calls[1][0]["provenance"])
            self.assertEqual(calls[0][0]["checks"], calls[1][0]["checks"])
            self.assertNotIn("retry_feedback", calls[0][0])
            self.assertEqual(calls[1][0]["retry_feedback"], {
                "code": bridge.ExternalComplianceCorrectionRequiredError.code,
                "checks": [{"check_id": "C00037", "source_quotes": [source]}],
            })
            sleep.assert_called_once_with(bridge.INDEPENDENT_REVIEW_RETRY_BACKOFF_SECONDS)
            self.assertEqual(bridge._response_sha256(response), initial_response_sha)
            self.assertEqual(response["requirements"], [])
            self.assertEqual(response["clause_reviews"][0]["obligations"], [])
            self.assertEqual(pointer["status"], "completed")
            self.assertEqual(pointer["provider_attempt"], 2)
            self.assertEqual(pointer["candidate_response_sha256"], initial_response_sha)

            first_audit_path = review_dir / (
                "independent-review-chunk-0002-attempt-01/coverage-audit.json"
            )
            first_audit = json.loads(first_audit_path.read_text(encoding="utf-8"))
            self.assertEqual(first_audit["status"], "rejected")
            self.assertTrue(first_audit["retryable"])
            accepted_audit = json.loads(
                (review_dir / pointer["audit_path"]).read_text(encoding="utf-8")
            )
            self.assertEqual(accepted_audit["status"], "completed")
            self.assertEqual(
                accepted_audit["provider_attempt_history"][0]["retry_code"],
                bridge.ExternalComplianceCorrectionRequiredError.code,
            )
            self.assertEqual(accepted_audit["candidate_response_sha256"], initial_response_sha)

    def test_typed_external_disposition_correction_keeps_candidate_and_two_call_limit(self):
        self._independent_review_patch.stop()
        from test_native_semantic_review import NativeSemanticReviewTests
        import native_semantic_review as native
        for outcome in ("accepted", "repeat", "foreign_quote"):
            chunk, response, accepted_review, _ = self._external_compliance_review_case()
            check, rejected, pending = NativeSemanticReviewTests._typed_external_disposition_case()
            source = check["document_text"]
            chunk["clauses"][0].update(text=source, source_span={
                "evidence_id": "E1", "start_offset": 0, "end_offset": len(source), "text": source,
                "source_sha256": hashlib.sha256(source.encode()).hexdigest()})
            chunk["evidence_context"]["E1"]["text"] = source
            response["clause_reviews"][0]["obligations"] = copy.deepcopy(check["review_context"]["primary_obligations"])
            for payload in (rejected, pending): payload["results"][0]["check_id"] = "C00037"
            accepted_review["results"] = copy.deepcopy(pending["results"])
            frozen = copy.deepcopy((chunk, response))
            calls = []
            def review(request, **kwargs):
                calls.append(copy.deepcopy(request))
                body = copy.deepcopy(rejected if len(calls) == 1 or outcome == "repeat" else pending)
                if len(calls) == 2 and outcome == "foreign_quote":
                    body["results"][0]["identified_obligations"][0]["source_quote"] = "foreign source"
                native.validate_obligation_coverage_response(body, request["checks"])
                native._validate_external_compliance_retry_result(body, request)
                return bind_mock_review_to_source_spans(accepted_review, request, kwargs["output_dir"])
            with self.subTest(outcome=outcome), tempfile.TemporaryDirectory() as td:
                review_dir = Path(td)
                with patch.object(bridge, "run_native_semantic_review", side_effect=review), patch.object(bridge.time, "sleep"):
                    kwargs = dict(review_dir=review_dir, run_id="run-external-correction", chunk_index=2,
                        attempt=1, host_runtime="codex", model="gpt-5.6-luna", timeout=10,
                        agent_id="main", runner="exec", binary="codex", config_path=None,
                        controller=bridge.RunController())
                    if outcome == "accepted":
                        pointer = bridge._run_independent_obligation_coverage_review(response, chunk, **kwargs)
                        self.assertEqual(pointer["provider_attempt"], 2)
                        self.assertEqual(pointer["status"], "completed")
                    else:
                        with self.assertRaises(bridge.IndependentObligationReviewError):
                            bridge._run_independent_obligation_coverage_review(response, chunk, **kwargs)
                        self.assertFalse(list(review_dir.glob("**/obligation-analysis-ledger.json")))
                self.assertEqual(len(calls), 2)
                self.assertEqual(calls[0]["checks"], calls[1]["checks"])
                self.assertEqual(calls[0]["provenance"], calls[1]["provenance"])
                self.assertEqual((chunk, response), frozen)
                self.assertEqual(calls[1]["retry_feedback"]["checks"][0]["rejected_result"], rejected["results"][0])

    def test_external_compliance_correction_exhaustion_fails_closed(self) -> None:
        self._independent_review_patch.stop()
        chunk, response, _, source = self._external_compliance_review_case()
        correction = bridge.ExternalComplianceCorrectionRequiredError([{
            "check_id": "C00037", "source_quotes": [source],
        }])
        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td)
            with patch.object(
                bridge, "run_native_semantic_review", side_effect=[correction, correction],
            ) as review_call, patch.object(bridge.time, "sleep") as sleep:
                with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                    bridge._run_independent_obligation_coverage_review(
                        response, chunk, review_dir=review_dir,
                        run_id="run-external-correction", chunk_index=2, attempt=1,
                        host_runtime="codex", model="gpt-5.6-luna", timeout=10,
                        agent_id="main", runner="exec", binary="codex", config_path=None,
                        controller=bridge.RunController(),
                    )

            self.assertEqual(review_call.call_count, bridge.INDEPENDENT_REVIEW_PROVIDER_MAX_ATTEMPTS)
            sleep.assert_called_once_with(bridge.INDEPENDENT_REVIEW_RETRY_BACKOFF_SECONDS)
            record = caught.exception.error_records[0]
            self.assertEqual(record["code"], "independent_obligation_review_correction_exhausted")
            self.assertEqual(record["retry_code"], bridge.ExternalComplianceCorrectionRequiredError.code)
            self.assertEqual(record["check_ids"], ["C00037"])
            self.assertFalse(getattr(caught.exception, "retryable", False))
            for suffix in ("", "-provider-attempt-02"):
                audit_path = review_dir / (
                    "independent-review-chunk-0002-attempt-01" + suffix + "/coverage-audit.json"
                )
                audit = json.loads(audit_path.read_text(encoding="utf-8"))
                self.assertEqual(audit["status"], "rejected")
                if suffix:
                    self.assertFalse(audit["retryable"])
                else:
                    self.assertTrue(audit["retryable"])

    def _packet(
        self, directory: Path, *, contract_version: str | None = None,
        source: str = "正文使用宋体",
    ) -> tuple[Path, dict]:
        directory.mkdir(parents=True, exist_ok=True)
        clauses = [{
            "id": "C1", "text": source, "evidence_ids": ["E1"],
            "source_kind": "paragraph", "location": {}, "part_index": 0,
        }]
        evidence = {"evidence": [{"id": "E1", "text": source, "kind": "paragraph", "location": {}}]}
        clauses = self._bind_test_source_spans(clauses, evidence)
        if contract_version is None:
            request = engine.build_llm_request(
                [], clauses, evidence, {}, "full",
                runtime_context={"code_fingerprint_sha256": "f" * 64},
            )
        else:
            request = engine.build_llm_request(
                [], clauses, evidence, {}, "full", contract_version=contract_version,
                runtime_context={"code_fingerprint_sha256": "f" * 64},
            )
        request["case_id"] = "standalone"
        request = attach_request_provenance(
            request, source_sha256="a" * 64, evidence_doc=evidence, clauses=clauses,
            run_id="run-bridge-test",
        )
        engine.prepare_host_agent_review_packets(
            request, clauses, evidence, "a" * 64, directory, chunk_size=1,
        )
        chunk = json.loads((directory / "llm-request-chunks.json").read_text(encoding="utf-8"))[0]
        return directory, chunk

    def _validated_chunk_projection_sha256(self, review_dir: Path, chunk: dict) -> str:
        full_request = json.loads((review_dir / "llm-request.json").read_text(encoding="utf-8"))
        chunks = json.loads(
            (review_dir / "llm-request-chunks.json").read_text(encoding="utf-8")
        )
        manifest = json.loads(
            (review_dir / "host-agent-review-manifest.json").read_text(encoding="utf-8")
        )
        engine.validate_host_review_chunk_source_projection(full_request, chunks, manifest)
        index = chunk.get("batch", {}).get("index", 0) - 1
        self.assertGreaterEqual(index, 0)
        self.assertEqual(chunks[index], chunk)
        return bridge._response_sha256(chunk)

    @staticmethod
    def _bind_test_source_spans(clauses: list[dict], evidence_doc: dict) -> list[dict]:
        evidence_by_id = {
            str(item.get("id")): item for item in evidence_doc.get("evidence", [])
            if isinstance(item, dict) and item.get("id")
        }
        for clause in clauses:
            evidence_ids = clause.get("evidence_ids") or []
            if len(evidence_ids) != 1:
                raise AssertionError("test source span requires one evidence ID")
            evidence = evidence_by_id[str(evidence_ids[0])]
            source = evidence["text"]
            exact = str(clause["text"])
            start = source.find(exact)
            if start < 0 or source.find(exact, start + 1) >= 0:
                raise AssertionError(f"test clause is not a unique evidence substring: {exact!r}")
            clause["source_span"] = {
                "evidence_id": str(evidence_ids[0]), "start_offset": start,
                "end_offset": start + len(exact), "text": exact,
                "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
            }
        return clauses

    def _response(self, chunk: dict) -> dict:
        return {
            "contract_version": "2.1",
            "provenance": chunk["provenance"],
            "requirements": [],
            "clause_reviews": [{
                "clause_id": "C1",
                "classification": "informational",
                "requirement_indexes": [],
                "reason": "The current evidence was reviewed.",
            }],
            "unsupported_items": [],
            "reported_conflicts": [],
        }

    def _executable_response(self, chunk: dict, *, invalid_verification: bool = False) -> dict:
        response = self._response(chunk)
        response["requirements"] = [{
            "role": "body_text",
            "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
            "clause_ids": ["C1"],
            "evidence_ids": ["E1"],
            "confidence": 0.9,
            "reason": "The clause specifies the body-text font.",
            "verification": "word_render" if invalid_verification else {
                "mode": "word_render",
                "checks": ["Check the body-text font."],
            },
        }]
        response["clause_reviews"] = [{
            "clause_id": "C1",
            "classification": "executable",
            "requirement_indexes": [0],
            "reason": "The clause is executable in DOCX.",
        }]
        return response

    def test_compact_model_packet_preserves_the_machine_readable_contract(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(Path(td) / "requirements")
            chunk["source_continuity_context"] = {
                "policy": "orientation_only_not_for_review",
                "preceding_clauses": [{"id": "C0", "text": "前一段"}],
                "following_clauses": [{"id": "C2", "text": "后一段"}],
            }
            packet = bridge.compact_model_packet(chunk)
            contract = packet["requirement_contract"]
            self.assertEqual(contract["allowed_roles"], chunk["requirement_contract"]["allowed_roles"])
            self.assertEqual(
                contract["role_properties_schema"],
                chunk["requirement_contract"]["role_properties_schema"],
            )
            self.assertIn("roleSpec", contract["$defs"])
            self.assertIn("verificationSpec", packet["response_schema"]["$defs"])
            self.assertIn("inputPrerequisiteSpec", packet["response_schema"]["$defs"])
            self.assertNotIn("provenance", packet)
            self.assertEqual(
                packet["source_continuity_context"]["following_clauses"][0]["id"],
                "C2",
            )

    def test_chunk_packets_expose_bounded_continuity_without_expanding_scope(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            directory = Path(td) / "requirements"
            clauses = [
                {"id": f"C{i}", "text": f"段落{i}", "evidence_ids": [f"E{i}"],
                 "source_kind": "paragraph", "location": {"order": i}, "part_index": 0}
                for i in range(1, 5)
            ]
            evidence = {"evidence": [
                {"id": f"E{i}", "text": f"段落{i}", "kind": "paragraph"}
                for i in range(1, 5)
            ]}
            request = engine.build_llm_request([], clauses, evidence, {}, "full")
            request = attach_request_provenance(
                request, source_sha256="a" * 64, evidence_doc=evidence,
                clauses=clauses, run_id="run-continuity-test",
            )
            engine.prepare_host_agent_review_packets(
                request, clauses, evidence, "a" * 64, directory, chunk_size=2,
            )
            chunks = json.loads((directory / "llm-request-chunks.json").read_text())
            self.assertEqual(
                [item["id"] for item in chunks[0]["source_continuity_context"]["following_clauses"]],
                ["C3", "C4"],
            )
            self.assertEqual(
                [item["id"] for item in chunks[1]["source_continuity_context"]["preceding_clauses"]],
                ["C1", "C2"],
            )
            self.assertEqual(chunks[0]["batch"]["clause_ids"], ["C1", "C2"])
            self.assertEqual(chunks[1]["batch"]["clause_ids"], ["C3", "C4"])

    def test_chunk_builder_keeps_one_physical_source_occurrence_atomic(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td) / "requirements"
            clauses = [
                {"id": "C1", "text": "声明开头", "evidence_ids": ["E1"],
                 "source_kind": "paragraph", "location": {"order": 1}, "part_index": 0},
                {"id": "C2", "text": "声明续句", "evidence_ids": ["E1"],
                 "source_kind": "paragraph", "location": {"order": 1}, "part_index": 1},
                {"id": "C3", "text": "另一条格式", "evidence_ids": ["E2"],
                 "source_kind": "paragraph", "location": {"order": 2}, "part_index": 0},
            ]
            evidence = {"evidence": [
                {"id": "E1", "text": "声明开头声明续句", "kind": "paragraph",
                 "location": {"order": 1}},
                {"id": "E2", "text": "另一条格式", "kind": "paragraph",
                 "location": {"order": 2}},
            ]}
            request = engine.build_llm_request([], clauses, evidence, {}, "full")
            request = attach_request_provenance(
                request, source_sha256="a" * 64, evidence_doc=evidence,
                clauses=clauses, run_id="run-source-atomic-test",
            )
            manifest = engine.prepare_host_agent_review_packets(
                request, clauses, evidence, "a" * 64, review_dir, chunk_size=1,
            )
            chunks = json.loads((review_dir / "llm-request-chunks.json").read_text())
            self.assertEqual(manifest["chunk_count"], 2)
            self.assertEqual(chunks[0]["batch"]["clause_ids"], ["C1", "C2"])
            self.assertEqual(chunks[1]["batch"]["clause_ids"], ["C3"])
            engine.validate_host_review_chunk_source_projection(request, chunks, manifest)

    def test_fixed_declaration_heading_and_body_remain_in_one_source_bound_chunk(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td) / "requirements"
            texts = [
                "封面格式", "学位论文使用授权书", "本人同意提交论文电子版。",
                "本人承诺：论文已完成", "电子版与纸质版一致。",
                "学位论文作者暨授权人签字", "年    月    日", "摘  要",
            ]
            evidence_ids = ["E0", "E1", "E2", "E3", "E3", "E4", "E5", "E6"]
            clauses = [
                {"id": f"C{index}", "text": text,
                 "evidence_ids": [evidence_ids[index]], "source_kind": "paragraph",
                 "location": {"part": "document", "order": index if index < 4 else index - 1}}
                for index, text in enumerate(texts)
            ]
            evidence = {"evidence": [
                {"id": f"E{index}", "text": source, "kind": "paragraph"}
                for index, source in enumerate([
                    "封面格式", "学位论文使用授权书", "本人同意提交论文电子版。",
                    "本人承诺：论文已完成；电子版与纸质版一致。",
                    "学位论文作者暨授权人签字", "年    月    日", "摘  要",
                ])
            ]}
            request = engine.build_llm_request([], clauses, evidence, {}, "full")
            request = attach_request_provenance(
                request, source_sha256="a" * 64, evidence_doc=evidence,
                clauses=clauses, run_id="run-cross-chunk-declaration",
            )
            manifest = engine.prepare_host_agent_review_packets(
                request, clauses, evidence, "a" * 64, review_dir, chunk_size=3,
            )
            chunks = json.loads((review_dir / "llm-request-chunks.json").read_text())
            self.assertEqual([chunk["batch"]["clause_ids"] for chunk in chunks], [
                ["C0"], ["C1", "C2", "C3", "C4", "C5", "C6"], ["C7"],
            ])
            self.assertEqual(manifest["chunk_count"], 3)
            candidate = bridge.compact_model_packet(chunks[1])["fixed_declaration_candidates"]
            self.assertEqual(len(candidate), 1)
            self.assertEqual(candidate[0]["clause_ids"], ["C1", "C2", "C3", "C4"])
            self.assertEqual(candidate[0]["evidence_ids"], ["E1", "E2", "E3"])
            self.assertEqual(candidate[0]["body_evidence_ids"], ["E2", "E3"])
            engine.validate_host_review_chunk_source_projection(request, chunks, manifest)

    def test_adjacent_fixed_declarations_do_not_absorb_each_other(self) -> None:
        clauses = [
            {"id": "C1", "text": "原创性声明", "evidence_ids": ["E1"]},
            {"id": "C2", "text": "本人独立完成。", "evidence_ids": ["E2"]},
            {"id": "C3", "text": "学位论文使用授权书", "evidence_ids": ["E3"]},
            {"id": "C4", "text": "本人同意提交电子版。", "evidence_ids": ["E4"]},
            {"id": "C5", "text": "摘要", "evidence_ids": ["E5"]},
        ]
        chunks, _ = engine._source_atomic_chunks(clauses, 1)
        self.assertEqual([[item["id"] for item in chunk] for chunk in chunks], [
            ["C1", "C2"], ["C3", "C4"], ["C5"],
        ])
        evidence = {
            item["evidence_ids"][0]: {
                "id": item["evidence_ids"][0], "text": item["text"],
            }
            for item in clauses
        }
        first = bridge._fixed_declaration_candidates(clauses, evidence, anchor="abstract_title_zh")
        self.assertEqual([item["clause_ids"] for item in first], [
            ["C1", "C2"], ["C3", "C4"],
        ])

    def test_fixed_declaration_stops_before_other_document_sections(self) -> None:
        for section in ("致谢", "附录A", "后记", "在学期间发表的学术论文与研究成果"):
            with self.subTest(section=section):
                clauses = [
                    {"id": "C1", "text": "学位论文使用授权书", "evidence_ids": ["E1"]},
                    {"id": "C2", "text": "本人授权学校保存论文。", "evidence_ids": ["E2"]},
                    {"id": "C3", "text": section, "evidence_ids": ["E3"]},
                    {"id": "C4", "text": "本节正文与授权无关。", "evidence_ids": ["E4"]},
                ]
                evidence = {
                    item["evidence_ids"][0]: {"text": item["text"]} for item in clauses
                }
                candidates = bridge._fixed_declaration_candidates(
                    clauses, evidence, anchor="abstract_title_zh",
                )
                self.assertEqual([item["clause_ids"] for item in candidates], [["C1", "C2"]])
                chunks, _ = engine._source_atomic_chunks(clauses, 1)
                self.assertEqual([item["id"] for item in chunks[0]], ["C1", "C2"])

    def test_declaration_anchor_preference_is_bound_to_current_structure(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            directory = Path(td) / "requirements"
            clauses = [{
                "id": "C1", "text": "固定声明正文", "evidence_ids": ["E1"],
                "source_kind": "paragraph", "location": {}, "part_index": 0,
            }]
            evidence = {
                "evidence": [{"id": "E1", "text": "固定声明正文", "kind": "paragraph"}],
                "structure_evidence": {"sections": [{
                    "first_paragraphs": [{"text": "摘要", "style_name": "Abstract Title CN"}],
                    "last_paragraphs": [],
                }]},
            }
            request = engine.build_llm_request([], clauses, evidence, {}, "full")
            self.assertEqual(request["declaration_anchor_candidates"], [
                "document_start", "abstract_title_zh",
            ])
            self.assertEqual(request["declaration_anchor_preference"], "abstract_title_zh")
            request = attach_request_provenance(
                request, source_sha256="a" * 64, evidence_doc=evidence,
                clauses=clauses, run_id="run-anchor-test",
            )
            engine.prepare_host_agent_review_packets(
                request, clauses, evidence, "a" * 64, directory, chunk_size=1,
            )
            chunk = json.loads((directory / "llm-request-chunks.json").read_text())[0]
            self.assertEqual(chunk["declaration_anchor_preference"], "abstract_title_zh")

    def test_preflight_rejects_invalid_declaration_anchor_and_signature_only_block(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(Path(td) / "requirements")
            chunk["clauses"][0]["text"] = "作者姓名"
            chunk["evidence_context"]["E1"]["text"] = "作者姓名"
            chunk["clauses"] = self._bind_test_source_spans(
                chunk["clauses"], {"evidence": [chunk["evidence_context"]["E1"]]},
            )
            response = self._response(chunk)
            response["requirements"] = [{
                "role": "declarations",
                "properties": {
                    "before_role": "declarations",
                    "items": [{
                        "id": "author_signature",
                        "heading": "作者姓名",
                        "body_parts": ["年 月 日于北体大"],
                        "source_evidence_ids": ["E1"],
                        "signature_placeholders": [],
                    }],
                },
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "confidence": 0.9, "reason": "固定文本",
            }]
            response["clause_reviews"] = [{
                "clause_id": "C1", "classification": "executable",
                "requirement_indexes": [0], "reason": "声明结构",
            }]
            errors = bridge.validate_host_agent_response(response, chunk)
            self.assertTrue(any("before_role" in error for error in errors), errors)
            self.assertTrue(any("generic author/date/signature" in error for error in errors), errors)

    def test_preflight_rejects_response_schema_drift(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._executable_response(chunk, invalid_verification=True)
            errors = bridge.validate_host_agent_response(response, chunk)
            self.assertTrue(any("verification" in error for error in errors), errors)

    def test_preflight_rejects_sample_bibliography_as_executable(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(Path(td) / "requirements")
            chunk["clauses"][0].update({
                "text": "[1] 示例作者. 示例文献[J]. 示例期刊, 2020.",
                "source_text_full": "[1] 示例作者. 示例文献[J]. 示例期刊, 2020.",
                "context_before": ["参考文献"],
                "context_after": [],
            })
            chunk["evidence_context"]["E1"]["text"] = chunk["clauses"][0]["text"]
            chunk["clauses"] = self._bind_test_source_spans(
                chunk["clauses"], {"evidence": [chunk["evidence_context"]["E1"]]},
            )
            response = self._executable_response(chunk)
            errors = bridge.validate_host_agent_response(response, chunk)
            self.assertTrue(
                any("sample_content_cannot_be_executable" in error for error in errors),
                errors,
            )

    def test_preflight_rejects_informational_as_normative_basis(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            response["clause_reviews"][0]["normative_basis"] = "informational"
            errors = bridge.validate_host_agent_response(response, chunk)
            self.assertTrue(any("normative_basis" in error for error in errors), errors)

    def test_keyword_origin_is_projected_to_human_verification_not_authoring(self) -> None:
        source = (
            "关键词是为了便于做文献索引和检索工作而从论文中选取出来用以表示全文主题内容信息的"
            "单词或术语，在论文中有明确出处"
        )
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(
                Path(td) / "requirements", source=source, contract_version="3.0",
            )
            clause_id = chunk["clauses"][0]["id"]
            response = {
                "contract_version": chunk["contract_version"],
                "provenance": chunk["provenance"],
                "requirements": [],
                "clause_reviews": [{
                    "clause_id": clause_id,
                    "classification": "requires_source_content",
                    "normative_basis": "explicit_normative_text",
                    "reason": "The keyword's thesis origin must be verified against the manuscript.",
                    "obligations": [{
                        "id": f"{clause_id}-obligation-1",
                        "status": "requires_source_content",
                        "reason": "The keyword source is not verifiable from this formatting packet.",
                    }],
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }

            accepted, audit = bridge.prepare_native_response_candidate(response, chunk)
            self.assertEqual(bridge.validate_host_agent_response(accepted, chunk), [])
            self.assertEqual(
                audit["source_verification_classification_policy_version"],
                "source-verification-classification-v6",
            )
            self.assertEqual(
                accepted["clause_reviews"][0]["classification"],
                "requires_source_verification",
            )
            self.assertEqual(accepted["clause_reviews"][0]["obligations"], [])
            self.assertEqual(accepted["requirements"], [])
            self.assertEqual(
                audit["source_verification_classification_projections"][0]["authorization"],
                "registered_source_verification_without_authoring_instruction_v1",
            )

            request = bridge.build_obligation_coverage_request(
                accepted, chunk, run_id=chunk["provenance"]["run_id"], chunk_index=4,
            )
            check = request["checks"][0]
            self.assertEqual(check["review_context"]["classification"], "requires_source_verification")
            self.assertEqual(check["review_context"]["primary_obligations"], [])
            self.assertEqual(
                check["review_context"]["source_content_verification_codes"],
                ["keyword_source_traceability_verification"],
            )
            verification_pending = {"results": [{
                "check_id": clause_id,
                "verdict": "source_content_verification_pending",
                "rationale": "A human must verify the selected keywords against the thesis manuscript.",
                "evidence_quotes": [source],
                "machine_obligation_ids": [],
                "identified_obligations": [{
                    "source_quote": source,
                    "disposition": "source_content_verification_pending",
                    "requirement_refs": [],
                }],
            }]}
            validated = bridge.validate_obligation_coverage_response(
                verification_pending, request["checks"],
            )
            self.assertEqual(validated[0]["verdict"], "source_content_verification_pending")

            falsely_authoring = copy.deepcopy(verification_pending)
            falsely_authoring["results"][0]["verdict"] = "source_content_pending"
            falsely_authoring["results"][0]["identified_obligations"][0][
                "disposition"
            ] = "authoring_content_pending"
            with self.assertRaisesRegex(
                bridge.SourceVerificationMislabelledAsAuthoringError,
                "source-bound existing-content check as authoring",
            ):
                bridge.validate_obligation_coverage_response(
                    falsely_authoring, request["checks"],
                )

    def test_preflight_rejects_requirement_index_on_nonexecutable_review(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._executable_response(chunk)
            response["clause_reviews"][0]["classification"] = "informational"
            response["clause_reviews"][0]["requirement_indexes"] = [0]
            errors = bridge.validate_host_agent_response(response, chunk)
            self.assertTrue(
                any("nonexecutable_review_must_not_reference_requirement" in error for error in errors),
                errors,
            )

    def test_preflight_reports_exact_clause_and_matching_requirement_indexes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(Path(td) / "requirements")
            chunk["clauses"].append({
                "id": "C2", "text": "标题居中", "evidence_ids": ["E2"],
                "source_kind": "paragraph", "location": {}, "part_index": 0,
            })
            chunk["evidence_context"]["E2"] = {
                "id": "E2", "text": "标题居中", "kind": "paragraph",
            }
            response = self._executable_response(chunk)
            response["requirements"].append({
                "role": "body_text",
                "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                "clause_ids": ["C2"],
                "evidence_ids": ["E2"],
                "confidence": 0.9,
                "reason": "The clause specifies the heading alignment.",
                "verification": {
                    "mode": "word_render",
                    "checks": ["Check the heading alignment."],
                },
            })
            response["clause_reviews"][0]["requirement_indexes"] = [0, 1]
            response["clause_reviews"].append({
                "clause_id": "C2",
                "classification": "executable",
                "requirement_indexes": [1],
                "reason": "The clause is executable in DOCX.",
            })

            errors = bridge.validate_host_agent_response(response, chunk)
            backed_errors = [
                error for error in errors
                if "requirement_index_not_backed_by_clause" in error
            ]
            self.assertEqual(len(backed_errors), 1, errors)
            self.assertIn("clause_id=C1", backed_errors[0])
            self.assertIn("requirement_index=1", backed_errors[0])
            self.assertIn("matching_indexes=[0]", backed_errors[0])

            response["clause_reviews"][0]["requirement_indexes"] = [0]
            self.assertFalse(
                any(
                    "requirement_index_not_backed_by_clause" in error
                    for error in bridge.validate_host_agent_response(response, chunk)
                )
            )

    def test_host_prompt_exposes_mechanical_contract_rules(self) -> None:
        prompt = bridge._host_prompt(
            request_path=Path("request.json"),
            chunk_path=Path("chunk.json"),
            response_path=Path("response.json"),
            run_id="run-1",
            chunk_index=1,
            chunk_count=1,
        )
        self.assertIn("classification and normative_basis are different fields", prompt)
        self.assertIn("Only covered, executable, verify_existing, and executable_with_external_check", prompt)
        self.assertIn("require_after_role", prompt)
        self.assertIn("A single clause may support multiple requirements", prompt)
        self.assertIn("Role boundary for equations", prompt)
        self.assertIn("The bridge materializes heading/body_parts", prompt)
        self.assertIn("do not copy normalized clause fragments or repeat an evidence paragraph", prompt)
        self.assertIn("partial_clause_coverage error never authorizes changing classification", prompt)
        self.assertIn("Never emit a spare", prompt)
        retry = bridge._host_prompt(
            request_path=Path("request.json"),
            chunk_path=Path("chunk.json"),
            response_path=Path("response.json"),
            run_id="run-1",
            chunk_index=1,
            chunk_count=1,
            retry_hint="informational normative_basis nonexecutable_review_must_not_reference_requirement",
        )
        self.assertIn("targeted contract repair rules", retry)
        self.assertIn("Remove normative_basis", retry)
        self.assertIn("never replace it with another guessed value", retry)

        with tempfile.TemporaryDirectory() as td:
            parent_path = Path(td) / "parent-response.raw.json"
            parent_path.write_text('{"contract_version":"3.0","requirements":[],"clause_reviews":[]}', encoding="utf-8")
            retry_with_baseline = bridge._host_prompt(
                request_path=Path("request.json"),
                chunk_path=Path("chunk.json"),
                response_path=Path("response.json"),
                run_id="run-1",
                chunk_index=1,
                chunk_count=1,
                attempt=2,
                retry_hint="cover_binding_violation",
                retry_parent_response_sha256=bridge.sha256_file(parent_path),
                retry_parent_response_path=parent_path,
                retry_error_records=[{
                    "code": "cover_binding_violation",
                    "json_pointer": "$.requirements[0].properties.fields[2]",
                }],
            )
        self.assertIn(str(parent_path), retry_with_baseline)
        self.assertIn("Preserve every non-error semantic field", retry_with_baseline)
        self.assertIn("Do not rewrite explanatory reasons", retry_with_baseline)
        self.assertIn("Do not split, merge, add", retry_with_baseline)
        self.assertIn("FINAL RETRY INVARIANT", retry_with_baseline)
        self.assertIn("do not turn an unresolved or informational review into executable", retry_with_baseline)

    def test_retry_prompt_embeds_parent_semantics_without_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            parent_path = Path(td) / "parent-response.raw.json"
            parent_path.write_text(json.dumps({
                "contract_version": "3.0",
                "provenance": {"run_id": "must-not-be-copied"},
                "requirements": [{
                    "role": "body_text", "properties": {"font": {"cjk": "SimSun"}},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                    "reason": "preserve this semantic payload",
                }],
                "clause_reviews": [{
                    "clause_id": "C1", "classification": "executable",
                    "reason": "preserve this review",
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }, ensure_ascii=False), encoding="utf-8")
            prompt = bridge._host_prompt(
                request_path=Path(td) / "request.json",
                chunk_path=Path(td) / "chunk.json",
                response_path=Path(td) / "response.json",
                run_id="run-1", chunk_index=1, chunk_count=1, attempt=2,
                retry_hint="local contract validation failed",
                retry_parent_response_sha256=bridge.sha256_file(parent_path),
                retry_parent_response_path=parent_path,
                retry_error_records=[{
                    "code": "missing_derived_requirement",
                    "json_pointer": "$.clause_reviews[0]",
                    "raw_error": "executable_review_requires_derived_requirement",
                }],
            )
        self.assertIn("<immutable_parent_response_without_provenance>", prompt)
        self.assertIn("preserve this semantic payload", prompt)
        self.assertIn("preserve this review", prompt)
        self.assertNotIn("must-not-be-copied", prompt)
        self.assertIn("Do not regenerate the response from the chunk", prompt)

    def test_retry_guidance_targets_exact_invalid_property_without_contradiction(self) -> None:
        retry = bridge._contract_repair_guidance(
            "local response contract validation failed: "
            "$.clause_reviews[12].normative_basis: 'informational' is not in "
            "['explicit_normative_text']; "
            "$.requirements[3].properties: unknown property 'style_hint'",
            include_base=False,
        )
        self.assertIn("clause_reviews[12]", retry)
        self.assertIn("remove the entire normative_basis property", retry)
        self.assertIn("Delete only the unknown property 'style_hint'", retry)
        self.assertNotIn("Do not repair by deleting evidence, clearing indexes", retry)

    def test_retry_guidance_routes_keyword_payload_gaps_to_nested_content_constraints(self) -> None:
        raw_error = (
            "$.clause_reviews[9]: partial_clause_coverage:"
            "keywords_zh.count_guidance,keywords_zh.max_item_chars"
        )
        retry = bridge._contract_repair_guidance(
            f"local response contract validation failed: {raw_error}",
            include_base=False,
        )
        self.assertIn(
            "content_constraints requirement at properties.keywords_zh",
            retry,
        )
        self.assertIn("count_guidance", retry)
        self.assertIn("general_guidance", retry)
        self.assertIn("max_item_chars", retry)
        self.assertIn("do not treat a top-level keywords role as covering", retry)
        self.assertIn("fail closed", retry)

        structured = bridge._structured_contract_repair_guidance([{
            "code": "partial_clause_coverage",
            "json_pointer": "$.clause_reviews[9]",
            "raw_error": raw_error,
        }], contract_version="3.0")
        self.assertIn("content_constraints requirement", structured)
        self.assertIn("only explicit mandatory ranges", structured)

    def test_current_source_compiles_missing_keyword_and_retains_ambiguous_heading(self) -> None:
        heading = "摘  要"
        keyword = "关键词在摘要内容后另起一行，一般3～8个，之间用分号分开"
        evidence_doc = {"evidence": [
            {"id": "E_ANCHOR", "kind": "paragraph", "text": "摘要",
             "style_name": "Abstract Title CN",
             "location": {"part": "document", "child_index": 1}},
            {"id": "E_HEADING", "kind": "paragraph", "text": heading,
             "style_name": "Heading 1",
             "location": {"part": "document", "child_index": 2}},
            {"id": "E_KEYWORD", "kind": "paragraph", "text": keyword,
             "location": {"part": "document", "child_index": 3}},
        ]}
        clauses = [
            {"id": "C_HEADING", "text": "摘 要", "source_text_full": "摘 要",
             "evidence_ids": ["E_HEADING"], "source_kind": "paragraph",
             "location": evidence_doc["evidence"][1]["location"],
             "source_span": {"evidence_id": "E_HEADING", "start_offset": 0,
                             "end_offset": len(heading), "text": heading,
                             "source_sha256": hashlib.sha256(heading.encode()).hexdigest()}},
            {"id": "C_KEYWORD", "text": keyword, "source_text_full": keyword,
             "evidence_ids": ["E_KEYWORD"], "source_kind": "paragraph",
             "location": evidence_doc["evidence"][2]["location"],
             "source_span": {"evidence_id": "E_KEYWORD", "start_offset": 0,
                             "end_offset": len(keyword), "text": keyword,
                             "source_sha256": hashlib.sha256(keyword.encode()).hexdigest()}},
        ]
        anchor_inventory = {
            "status": "verified", "source": {"sha256": "a" * 64},
            "anchors": {"abstract_title_zh": {
                "anchor_type": "semantic_role", "binding_status": "verified",
                "match_count": 1,
                "matches": [{"evidence_id": "E_ANCHOR", "text": "摘要"}],
            }},
        }
        chunk = engine.build_llm_request(
            [], clauses, evidence_doc, {}, "full", contract_version="3.0",
            runtime_context={
                "code_fingerprint_sha256": "9" * 64,
                "runtime_inventory": {"anchor_inventory": anchor_inventory},
            },
        )
        chunk = attach_request_provenance(
            chunk, source_sha256="a" * 64, evidence_doc=evidence_doc,
            clauses=chunk["clauses"], run_id="dynamic-source-projection-test",
        )
        raw = {
            "contract_version": "3.0", "requirements": [],
            "clause_reviews": [{
                "clause_id": clause_id, "classification": "executable",
                "normative_basis": "explicit_normative_text", "reason": "Model claim.",
                "obligations": [{"id": duty, "status": "covered", "reason": "Model claim."}],
            } for clause_id, duty in (
                ("C_HEADING", "heading"), ("C_KEYWORD", "keyword"),
            )],
            "unsupported_items": [], "reported_conflicts": [],
        }
        candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(len(candidate["requirements"]), 1)
        requirement = candidate["requirements"][0]
        self.assertEqual(requirement["clause_ids"], ["C_KEYWORD"])
        self.assertEqual(requirement["properties"]["keywords_zh"]["separator"], "semicolon")
        self.assertEqual(candidate["clause_reviews"][0]["classification"], "unresolved")
        self.assertEqual(audit["source_heading_binding_projections"][0]["action"],
                         "retain_unresolved_heading_binding")
        self.assertEqual(audit["source_keyword_constraint_projections"][0]["authorization"],
                         "verified_complete_current_source_without_parent_v1")
        self.assertEqual(raw["requirements"], [])

        # A later model attempt may omit a prior requirement. Recompilation
        # must be audited, and raw retry authorization must still see the
        # model's deletion rather than silently treating it as unchanged.
        parent_raw = copy.deepcopy(raw)
        parent_raw["requirements"] = copy.deepcopy(candidate["requirements"])
        error, changed = bridge._retry_semantic_change_error(
            parent_raw, raw,
            [{"code": "missing_derived_requirement",
              "json_pointer": "$.clause_reviews[0]",
              "raw_error": "$.clause_reviews[0]: executable_review_requires_derived_requirement"}],
            contract_version=bridge.HOST_REVIEW_CONTRACT_V3, chunk=chunk,
        )
        self.assertIsNotNone(error)
        self.assertIn("$.requirements", changed)

        # A distinct invalid field may reject the candidate after the
        # source-driven projections. Their audit must survive the exception.
        invalid_raw = copy.deepcopy(raw)
        invalid_raw["clause_reviews"][1]["normative_basis"] = "informational"
        with self.assertRaises(ValueError) as rejected:
            bridge.prepare_native_response_candidate(invalid_raw, chunk)
        failure = rejected.exception
        self.assertEqual(
            failure.source_keyword_constraint_projections[0]["authorization"],
            "verified_complete_current_source_without_parent_v1",
        )
        self.assertEqual(
            failure.source_heading_binding_projections[0]["action"],
            "retain_unresolved_heading_binding",
        )
        self.assertEqual(
            failure.source_keyword_constraint_projection_policy_version,
            bridge.SOURCE_KEYWORD_CONSTRAINT_PROJECTION_POLICY_VERSION,
        )

    def test_keyword_source_constraints_pass_the_production_candidate_boundary(self) -> None:
        evidence_doc = {"evidence": [
            {
                "id": "E1", "kind": "paragraph",
                "text": "关键词在摘要内容后另起一行，一般3～8个，之间用分号分开",
            },
            {
                "id": "E2", "kind": "paragraph",
                "text": "关键词：术语；最多7个汉字；最少3组，最多8组",
            },
            {
                "id": "E3", "kind": "paragraph",
                "text": (
                    "Keywords in the abstract content after another line, generally 3~8, "
                    "separated by semicolons."
                ),
            },
        ]}
        clauses = self._bind_test_source_spans([
            {"id": "C00069", "text": evidence_doc["evidence"][0]["text"], "evidence_ids": ["E1"]},
            {
                "id": "C00071", "text": "最多7个汉字",
                "source_text_full": evidence_doc["evidence"][1]["text"],
                "evidence_ids": ["E2"],
            },
            {
                "id": "C00072", "text": "最少3组，最多8组",
                "source_text_full": evidence_doc["evidence"][1]["text"],
                "evidence_ids": ["E2"],
            },
            {"id": "C00077", "text": evidence_doc["evidence"][2]["text"], "evidence_ids": ["E3"]},
        ], evidence_doc)
        chunk = engine.build_llm_request(
            [], clauses, evidence_doc, {}, "full", contract_version="3.0",
            runtime_context={"code_fingerprint_sha256": "9" * 64},
        )
        chunk = attach_request_provenance(
            chunk, source_sha256="a" * 64, evidence_doc=evidence_doc,
            clauses=chunk["clauses"], run_id="keyword-source-projection-test",
        )

        def native_keyword_requirement(role: str, clause_ids: list[str], evidence_ids: list[str]) -> dict:
            role_properties = {
                "font": None, "paragraph": None, "numbering": None,
                "position": None, "prefix": None, "separator": "semicolon",
                "style_hint": None, "text": None, "header_content": None,
                "bottom_border": None,
            }
            return {
                "existing_requirement_id": None, "field_key": None,
                "clause_ids": clause_ids, "source_fragment_clause_ids": None,
                "evidence_ids": evidence_ids, "confidence": 0.97,
                "reason": "The cited source explicitly specifies keyword formatting.",
                "applicability": None, "input_prerequisites": None,
                "verification": {
                    "mode": "static_docx",
                    "checks": ["Verify the keyword presentation against the cited source."],
                    "checker_ids": None,
                },
                "role": role, "properties": role_properties,
            }

        reviews = []
        for clause_id, duty in (
            ("C00069", "placement, general count guidance, and separator"),
            ("C00071", "Chinese-character item length"),
            ("C00072", "mandatory keyword group count"),
            ("C00077", "English keyword placement, count guidance, and separator"),
        ):
            reviews.append({
                "clause_id": clause_id, "classification": "executable",
                "normative_basis": "explicit_normative_text",
                "reason": f"The linked source requirement represents {duty}.",
                "obligations": [{
                    "id": f"{clause_id}-duty", "status": "covered",
                    "reason": f"The linked requirement represents {duty}.",
                }],
            })
        raw_response = {
            "contract_version": "3.0",
            "requirements": [
                native_keyword_requirement(
                    "keywords_zh", ["C00069", "C00071", "C00072"], ["E1", "E2"],
                ),
                native_keyword_requirement("keywords_en", ["C00077"], ["E3"]),
            ],
            "clause_reviews": reviews,
            "unsupported_items": [], "reported_conflicts": [],
        }

        accepted, audit = bridge.prepare_native_response_candidate(raw_response, chunk)
        self.assertEqual(bridge.validate_host_agent_response(accepted, chunk), [])
        self.assertEqual(len(audit["source_keyword_constraint_projections"]), 2)
        projected = [
            item for item in accepted["requirements"]
            if item.get("role") == "content_constraints"
        ]
        self.assertEqual(len(projected), 2)
        zh = next(item["properties"]["keywords_zh"] for item in projected
                  if "keywords_zh" in item["properties"])
        en = next(item["properties"]["keywords_en"] for item in projected
                  if "keywords_en" in item["properties"])
        self.assertEqual(zh["count_guidance"], {
            "min_count": 3, "max_count": 8, "strength": "general_guidance",
        })
        self.assertEqual((zh["min_count"], zh["max_count"], zh["max_item_chars"]), (3, 8, 7))
        self.assertEqual((en["count_guidance"]["strength"], en["separator"]),
                         ("general_guidance", "semicolon"))
        self.assertNotIn("min_count", en)
        self.assertNotIn("max_count", en)

        # Reproduce the observed BSU topology: the top-level keyword style
        # cites only C00069, while one existing content constraint cites the
        # split C00071/C00072 source clauses and initially omits hard bounds.
        split_response = copy.deepcopy(accepted)
        zh_style = next(item for item in split_response["requirements"]
                        if item["role"] == "keywords_zh")
        zh_style["clause_ids"] = ["C00069"]
        zh_style["evidence_ids"] = ["E1"]
        zh_constraint = next(item for item in split_response["requirements"]
                             if item["role"] == "content_constraints"
                             and isinstance(item["properties"].get("keywords_zh"), dict))
        zh_constraint["properties"]["keywords_zh"]["min_count"] = None
        zh_constraint["properties"]["keywords_zh"]["max_count"] = None

        restored, split_audit = bridge.prepare_native_response_candidate(split_response, chunk)
        restored_zh = next(item["properties"]["keywords_zh"] for item in restored["requirements"]
                           if item["role"] == "content_constraints"
                           and isinstance(item["properties"].get("keywords_zh"), dict))
        self.assertEqual((restored_zh["min_count"], restored_zh["max_count"]), (3, 8))
        self.assertEqual(bridge.validate_host_agent_response(restored, chunk), [])
        self.assertEqual(len(split_audit["source_keyword_constraint_projections"]), 1)

    def test_native_guidance_check_survives_candidate_boundary_without_empty_check_bypass(self) -> None:
        source = "关键词一般4～9个"
        evidence = {"evidence": [{"id": "E_CURRENT", "kind": "paragraph", "text": source}]}
        clauses = self._bind_test_source_spans([
            {"id": "C_CURRENT", "text": source, "evidence_ids": ["E_CURRENT"]},
        ], evidence)
        chunk = engine.build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
            runtime_context={"code_fingerprint_sha256": "9" * 64},
        )
        chunk = attach_request_provenance(
            chunk, source_sha256="a" * 64, evidence_doc=evidence,
            clauses=chunk["clauses"], run_id="current-guidance-check-test",
        )
        check = "Verify the keyword count against the source-qualified 4-to-9 guidance and retain independent checks."
        raw = {
            "contract_version": "3.0", "requirements": [{
                "role": "content_constraints", "clause_ids": ["C_CURRENT"], "evidence_ids": ["E_CURRENT"],
                "reason": "The source states a qualified recommendation, not a hard bound.", "confidence": 0.9,
                "properties": {"keywords_zh": {"count_guidance": {
                    "min_count": 4, "max_count": 9, "strength": "general_guidance"}}},
                "verification": {"mode": "manual", "checks": [check]},
            }], "clause_reviews": [{
                "clause_id": "C_CURRENT", "classification": "executable", "normative_basis": "explicit_normative_text",
                "reason": "The linked property preserves the source recommendation.",
                "obligations": [{"id": "count-guidance", "status": "covered", "reason": "Bound as general guidance."}],
            }], "unsupported_items": [], "reported_conflicts": [],
        }
        original = copy.deepcopy(raw)
        accepted, _audit = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(raw, original)
        self.assertEqual(bridge.validate_host_agent_response(accepted, chunk), [])
        self.assertEqual(accepted["requirements"][0]["verification"]["checks"], [check])
        self.assertNotIn("min_count", accepted["requirements"][0]["properties"]["keywords_zh"])
        self.assertEqual(accepted["clause_reviews"], original["clause_reviews"])
        again, _again_audit = bridge.prepare_native_response_candidate(accepted, chunk)
        self.assertEqual(again, accepted)
        invalid = copy.deepcopy(raw)
        invalid["requirements"][0]["verification"]["checks"] = []
        with self.assertRaisesRegex(ValueError, "verification.checks"):
            bridge.prepare_native_response_candidate(invalid, chunk)

    def test_ambiguous_conflict_target_is_omitted_without_choosing_a_language(self) -> None:
        sources = {
            "E74": "The following English is not correct.",
            "E76": "The Chinese abstract is usually 300 to 1,000 words.",
        }
        evidence = {"evidence": [
            {"id": evidence_id, "kind": "paragraph", "text": source}
            for evidence_id, source in sources.items()
        ]}
        clauses = self._bind_test_source_spans([
            {"id": "C74", "text": sources["E74"], "evidence_ids": ["E74"]},
            {"id": "C76", "text": sources["E76"], "evidence_ids": ["E76"]},
        ], evidence)
        chunk = engine.build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
            runtime_context={"code_fingerprint_sha256": "9" * 64},
        )
        chunk = attach_request_provenance(
            chunk, source_sha256="a" * 64, evidence_doc=evidence,
            clauses=chunk["clauses"], run_id="unresolved-target-test",
        )
        raw = {
            "contract_version": "3.0",
            "requirements": [],
            "clause_reviews": [
                {"clause_id": clause["id"], "classification": "unresolved",
                 "normative_basis": "insufficient", "reason": "The target is ambiguous."}
                for clause in clauses
            ],
            "unsupported_items": [],
            "reported_conflicts": [{
                "type": "semantic_conflict",
                "reason": "The Chinese or English abstract target is unresolved.",
                "clause_ids": ["C74", "C76"], "evidence_ids": ["E74", "E76"],
                "target": {"role": "content_constraints",
                           "property": "abstract_zh_or_abstract_en"},
                "status": "unresolved",
            }],
        }
        self.assertIn("unknown_registered_role_property", str(
            bridge.validate_host_agent_response(raw, chunk)
        ))

        accepted, audit = bridge.prepare_native_response_candidate(raw, chunk)
        conflict = accepted["reported_conflicts"][0]
        self.assertNotIn("target", conflict)
        self.assertIn("content_constraints.abstract_zh | content_constraints.abstract_en",
                      conflict["reason"])
        self.assertEqual(conflict["status"], "unresolved")
        self.assertEqual(accepted["requirements"], [])
        self.assertTrue(all(item["classification"] == "unresolved"
                            for item in accepted["clause_reviews"]))
        self.assertEqual(bridge.validate_host_agent_response(accepted, chunk), [])
        repairs = audit["mechanical_repairs"]
        self.assertEqual(len(repairs), 1)
        self.assertEqual(repairs[0]["original_target"], raw["reported_conflicts"][0]["target"])
        self.assertEqual(repairs[0]["run_id"], "unresolved-target-test")
        self.assertEqual(repairs[0]["source_sha256"], "a" * 64)
        merged, _, _ = engine.merge_llm_primary(
            Path("synthetic-source"),
            {"schema_version": "1.0", "source_document": "synthetic-source",
             "roles": {}, "page": {}, "requirements": [], "content_instances": []},
            clauses, accepted, set(sources),
        )
        self.assertEqual(merged["semantic_conflicts"], accepted["reported_conflicts"])
        self.assertEqual(merged["status"], "needs_clarification")
        self.assertTrue(any(item.get("type") == "llm_reported_conflict"
                            for item in merged["blocking_errors"]))

        narrowed = copy.deepcopy(raw)
        narrowed["reported_conflicts"][0]["target"]["property"] = "abstract_en"
        records = bridge.contract_error_records(
            bridge.validate_host_agent_response(raw, chunk), response=raw, chunk=chunk,
        )
        drift_error, changed = bridge._retry_semantic_change_error(
            raw, narrowed, records, contract_version="3.0", chunk=chunk,
        )
        self.assertIsNotNone(drift_error)
        self.assertIn("$.top_level.reported_conflicts[0].target.property", changed)

        for change in (
            "unknown_alternative", "candidate_value", "covered_review",
            "unbound_evidence", "stale_record",
        ):
            with self.subTest(change=change):
                invalid = copy.deepcopy(raw)
                if change == "unknown_alternative":
                    invalid["reported_conflicts"][0]["target"]["property"] = (
                        "abstract_zh_or_unknown"
                    )
                elif change == "candidate_value":
                    invalid["reported_conflicts"][0]["candidates"] = [
                        {"evidence_id": "E74", "value": "possibly Chinese"},
                    ]
                elif change == "covered_review":
                    invalid["clause_reviews"][0]["classification"] = "executable"
                elif change == "unbound_evidence":
                    invalid["reported_conflicts"][0]["evidence_ids"] = ["E74"]
                invalid_record = {
                    "code": "contract_validation_error",
                    "json_pointer": "$.reported_conflicts[0].target",
                    "raw_error": (
                        "$.reported_conflicts[0].target: unknown_registered_role_property"
                    ),
                    "response_sha256": bridge._response_sha256(invalid),
                }
                if change == "stale_record":
                    invalid_record["response_sha256"] = "0" * 64
                rejected, _ = bridge._apply_safe_mechanical_repairs(
                    invalid, [invalid_record], chunk=chunk,
                )
                self.assertIsNone(rejected)

    def test_retry_guidance_targets_exact_requirement_clause_mapping(self) -> None:
        retry = bridge._contract_repair_guidance(
            "local response contract validation failed: "
            "$.clause_reviews[2]: requirement_index_not_backed_by_clause:"
            "clause_id=C00243:requirement_index=2:matching_indexes=[1]",
            include_base=False,
        )
        self.assertIn("clause_reviews[2]", retry)
        self.assertIn("C00243", retry)
        self.assertIn("requirements[2].clause_ids", retry)
        self.assertIn("matching requirement indexes are [1]", retry)
        self.assertIn("Do not copy a neighboring clause's index", retry)

    def test_retry_guidance_requires_payload_instead_of_field_key_only(self) -> None:
        retry = bridge._contract_repair_guidance(
            "local response contract validation failed: "
            "$.requirements[0].properties: must_include_semantic_payload",
            include_base=False,
        )
        self.assertIn("non-empty role-specific properties object", retry)
        self.assertIn("field_key alone", retry)
        self.assertIn("properties.text", retry)

    def test_retry_guidance_requires_real_semantic_obligation_inventory(self) -> None:
        retry = bridge._contract_repair_guidance(
            "$.clause_reviews[0].obligations: executable_review_requires_non_empty_inventory",
            include_base=False,
        )
        self.assertIn("non-empty array of distinct semantic duties", retry)
        self.assertIn("Do not", retry)
        self.assertIn("generic placeholder", retry)

    def test_retry_guidance_binds_prerequisites_and_admin_condition(self) -> None:
        retry = bridge._contract_repair_guidance(
            "local response contract validation failed: "
            "$.requirements[0].input_prerequisites[0].key: does not match "
            "'^(thesis_profile|source_inventory|template_profile|runtime)\\.'; "
            "$.requirements[1].properties.non_public_administration.applicability: "
            "must be conditional on thesis_profile.security_level",
            include_base=False,
        )
        self.assertIn("runtime_context.*", retry)
        self.assertIn("operator equals or in", retry)
        self.assertIn("exact whole-profile key thesis_profile", retry)
        self.assertIn("Do not replace thesis_profile with thesis_profile.cover_metadata", retry)

    def test_invalid_retry_cannot_change_classification_to_escape_a_contract_error(self) -> None:
        previous = self._executable_response({"provenance": {}})
        current = self._executable_response({"provenance": {}})
        current["clause_reviews"][0]["classification"] = "informational"
        current["clause_reviews"][0].pop("requirement_indexes", None)
        change_error, changed = bridge._retry_semantic_change_error(
            previous,
            current,
            [{"code": "input_prerequisite_namespace"}],
            contract_version="2.1",
        )
        self.assertTrue(changed)
        self.assertIsNotNone(change_error)
        self.assertIn("semantic re-review", str(change_error))

    def test_v3_missing_obligation_inventory_retry_is_narrow_and_source_bound(self) -> None:
        source = "学位论文作者签名： 年 月 日"
        provenance = {
            "run_id": "run-inventory-retry", "source_sha256": "a" * 64,
            "clause_sha256": "b" * 64, "evidence_sha256": "c" * 64,
            "request_sha256": "d" * 64,
        }
        chunk = {
            "provenance": provenance,
            "batch": {"index": 2},
            "case_id": "case-inventory-retry",
            "response_schema": {"type": "object", "title": "retry test"},
            "runtime_context": {"code_fingerprint_sha256": "9" * 64},
            "clauses": [exact_source_clause("C00037", source)],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        previous = {
            "contract_version": "3.0", "provenance": provenance,
            "requirements": [], "unsupported_items": [], "reported_conflicts": [],
            "clause_reviews": [{
                "clause_id": "C00037", "classification": "external_compliance",
                "normative_basis": "external_duty",
                "reason": "The author must sign and date the declaration.",
                "obligations": None,
            }],
        }
        current = copy.deepcopy(previous)
        current["clause_reviews"][0]["obligations"] = [{
            "id": "actual_signature_and_date", "status": "unverifiable",
            "reason": "The source requires an actual author signature and date.",
        }]
        records = [{
            "code": "executable_review_obligations_missing",
            "json_pointer": "$.clause_reviews[0].obligations",
            "clause_id": "C00037",
            "response_sha256": bridge._response_sha256(previous),
            "raw_error": "review_requires_non_empty_source_inventory",
        }]

        authorization: list[dict] = []
        error, changed = bridge._retry_semantic_change_error(
            previous, current, records, contract_version="3.0", chunk=chunk,
            authorization_out=authorization,
        )
        self.assertIsNone(error)
        self.assertEqual(changed, ["$.clause_reviews[0].obligations"])
        self.assertEqual(len(authorization), 1)
        self.assertEqual(authorization[0]["rule_id"], "v3_source_inventory_completion")
        self.assertTrue(authorization[0]["source_binding_complete"])
        self.assertEqual(authorization[0]["source_binding"]["run_id"], "run-inventory-retry")
        self.assertEqual(
            authorization[0]["source_evidence_bindings"][0]["clause_id"], "C00037",
        )

        # The allowance is not a blanket semantic retry grant.
        unsafe_candidates = []
        changed_classification = copy.deepcopy(current)
        changed_classification["clause_reviews"][0]["classification"] = "informational"
        unsafe_candidates.append(changed_classification)
        changed_reason = copy.deepcopy(current)
        changed_reason["clause_reviews"][0]["reason"] = "Unrelated rewrite"
        unsafe_candidates.append(changed_reason)
        covered_external_action = copy.deepcopy(current)
        covered_external_action["clause_reviews"][0]["obligations"][0]["status"] = "covered"
        unsafe_candidates.append(covered_external_action)
        unresolved_external_action = copy.deepcopy(current)
        unresolved_external_action["clause_reviews"][0]["obligations"][0]["status"] = "unresolved"
        unsafe_candidates.append(unresolved_external_action)
        empty_inventory = copy.deepcopy(current)
        empty_inventory["clause_reviews"][0]["obligations"] = []
        unsafe_candidates.append(empty_inventory)
        duplicate_ids = copy.deepcopy(current)
        duplicate_ids["clause_reviews"][0]["obligations"].append(
            copy.deepcopy(duplicate_ids["clause_reviews"][0]["obligations"][0])
        )
        unsafe_candidates.append(duplicate_ids)
        changed_requirement_graph = copy.deepcopy(current)
        changed_requirement_graph["requirements"].append({
            "role": "paragraph", "properties": {"text": source},
            "clause_ids": ["C00037"], "evidence_ids": ["E1"],
            "reason": "Unauthorized new requirement edge",
        })
        unsafe_candidates.append(changed_requirement_graph)

        for candidate in unsafe_candidates:
            with self.subTest(candidate=candidate):
                candidate_error, _ = bridge._retry_semantic_change_error(
                    previous, candidate, records, contract_version="3.0", chunk=chunk,
                )
                self.assertIsNotNone(candidate_error)

        stale_record = copy.deepcopy(records)
        stale_record[0]["response_sha256"] = "0" * 64
        stale_error, _ = bridge._retry_semantic_change_error(
            previous, current, stale_record, contract_version="3.0", chunk=chunk,
        )
        self.assertIsNotNone(stale_error)

        incomplete_fingerprints = copy.deepcopy(chunk)
        incomplete_fingerprints["runtime_context"].pop("code_fingerprint_sha256")
        fingerprint_error, _ = bridge._retry_semantic_change_error(
            previous, current, records, contract_version="3.0", chunk=incomplete_fingerprints,
        )
        self.assertIsNotNone(fingerprint_error)

        mismatched_evidence = copy.deepcopy(chunk)
        mismatched_evidence["evidence_context"]["E1"]["id"] = "E2"
        evidence_error, _ = bridge._retry_semantic_change_error(
            previous, current, records, contract_version="3.0", chunk=mismatched_evidence,
        )
        self.assertIsNotNone(evidence_error)

        duplicate_evidence_ids = copy.deepcopy(chunk)
        duplicate_evidence_ids["clauses"][0]["evidence_ids"] = ["E1", "E1"]
        duplicate_evidence_error, _ = bridge._retry_semantic_change_error(
            previous, current, records, contract_version="3.0", chunk=duplicate_evidence_ids,
        )
        self.assertIsNotNone(duplicate_evidence_error)

        mismatched_span = copy.deepcopy(chunk)
        mismatched_span["clauses"][0]["source_span"]["text"] = "不匹配的来源切片"
        invalid_spans = [mismatched_span]
        bool_offset = copy.deepcopy(chunk)
        bool_offset["clauses"][0]["source_span"]["start_offset"] = True
        invalid_spans.append(bool_offset)
        out_of_bounds = copy.deepcopy(chunk)
        out_of_bounds["clauses"][0]["source_span"]["end_offset"] = len(source) + 1
        invalid_spans.append(out_of_bounds)
        wrong_source_hash = copy.deepcopy(chunk)
        wrong_source_hash["clauses"][0]["source_span"]["source_sha256"] = "0" * 64
        invalid_spans.append(wrong_source_hash)
        unlinked_span_evidence = copy.deepcopy(chunk)
        unlinked_span_evidence["clauses"][0]["source_span"]["evidence_id"] = "E2"
        invalid_spans.append(unlinked_span_evidence)
        for invalid_chunk in invalid_spans:
            with self.subTest(source_span=invalid_chunk["clauses"][0]["source_span"]):
                span_error, _ = bridge._retry_semantic_change_error(
                    previous, current, records, contract_version="3.0", chunk=invalid_chunk,
                )
                self.assertIsNotNone(span_error)

    def test_external_action_retry_adds_only_inventory_then_projects_source_echo(self) -> None:
        source = "北京体育大学学位评定委员会办公室盖章(有效)"
        evidence_doc = {"evidence": [{"id": "E1", "text": source, "kind": "paragraph"}]}
        chunk = engine.build_llm_request(
            [], [exact_source_clause("C00049", source)], evidence_doc, {}, "full",
            contract_version="3.0",
            runtime_context={"code_fingerprint_sha256": "9" * 64},
        )
        chunk["batch"] = {"index": 3, "count": 22}
        chunk["case_id"] = "case-external-obligation-retry"
        chunk = attach_request_provenance(
            chunk, source_sha256="a" * 64, evidence_doc=evidence_doc,
            clauses=chunk["clauses"], run_id="run-external-obligation-retry",
        )
        parent = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "existing_requirement_id": None,
                "role": "body_text", "properties": {
                    "font": None, "paragraph": None, "numbering": None,
                    "position": None, "prefix": None, "separator": None,
                    "style_hint": None, "text": source,
                    "header_content": None, "bottom_border": None,
                },
                "clause_ids": ["C00049"], "evidence_ids": ["E1"],
                "confidence": 0.95, "reason": "The source names a real-world stamp.",
                "applicability": {
                    "status": "always", "conditions": None, "exceptions": None,
                },
                "verification": None,
            }],
            "clause_reviews": [{
                "clause_id": "C00049", "classification": "external_compliance",
                "normative_basis": "external_duty", "reason": "An office must apply the stamp.",
                "obligations": None,
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        parent = bridge.normalize_native_response(parent, chunk["response_schema"])
        errors = bridge.validate_host_agent_response(parent, chunk)
        self.assertTrue(errors)
        validator_records = bridge.contract_error_records(
            errors, response=parent, chunk=chunk,
        )
        inventory_records = bridge._external_action_obligation_retry_records(
            parent, validator_records, chunk,
        )
        self.assertEqual(inventory_records, [])
        self.assertEqual(
            [(item["clause_id"], item["json_pointer"]) for item in validator_records
             if item["code"] == "external_action_obligations_missing"],
            [("C00049", "$.clause_reviews[0].obligations")],
        )
        with self.assertRaises(ValueError) as rejected:
            bridge.prepare_native_response_candidate(parent, chunk)
        self.assertIn(
            "external_action_obligations_missing",
            {item["code"] for item in rejected.exception.error_records},
        )

        model_retry = copy.deepcopy(parent)
        model_retry["clause_reviews"][0]["obligations"] = [{
            "id": "committee_stamp", "status": "unverifiable",
            "reason": "The physical stamp must be applied by the named office.",
        }]
        # Even if the model follows the old deletion instruction, the retry
        # projection must ignore that extra change and copy only the inventory.
        model_retry["requirements"] = []
        model_retry = bridge.normalize_native_response(model_retry, chunk["response_schema"])
        all_retry_records = [*validator_records, *inventory_records]
        unprojected_error, _ = bridge._retry_semantic_change_error(
            parent, model_retry, all_retry_records,
            contract_version="3.0", chunk=chunk,
        )
        self.assertIsNotNone(unprojected_error)
        projected, projection_audit = bridge._project_validator_targeted_obligation_fields(
            parent, model_retry, all_retry_records,
        )
        self.assertIsNotNone(projected)
        self.assertEqual(projected["requirements"], parent["requirements"])
        self.assertEqual(
            projection_audit["discarded_unrequested_paths"], ["$.requirements"],
        )
        targeted_records = [
            item for item in all_retry_records
            if item["code"] == "external_action_obligations_missing"
        ]
        falsely_covered = copy.deepcopy(model_retry)
        falsely_covered["clause_reviews"][0]["obligations"][0]["status"] = "covered"
        falsely_covered = bridge.normalize_native_response(
            falsely_covered, chunk["response_schema"],
        )
        rejected_covered_projection, _ = bridge._project_validator_targeted_obligation_fields(
            parent, falsely_covered, all_retry_records,
        )
        self.assertIsNone(rejected_covered_projection)
        rejected_covered_error, _ = bridge._retry_semantic_change_error(
            parent,
            {**parent, "clause_reviews": falsely_covered["clause_reviews"]},
            targeted_records, contract_version="3.0", chunk=chunk,
        )
        self.assertIsNotNone(rejected_covered_error)
        authorizations: list[dict] = []
        drift_error, changed_paths = bridge._retry_semantic_change_error(
            parent, projected, targeted_records,
            contract_version="3.0", chunk=chunk,
            authorization_out=authorizations,
        )
        self.assertIsNone(drift_error)
        self.assertEqual(changed_paths, ["$.clause_reviews[0].obligations"])
        self.assertEqual(authorizations[0]["rule_id"], "v3_source_inventory_completion")
        self.assertTrue(authorizations[0]["source_binding_complete"])
        self.assertEqual(
            authorizations[0]["source_binding"]["run_id"],
            "run-external-obligation-retry",
        )

        accepted, candidate_audit = bridge.prepare_native_response_candidate(projected, chunk)
        self.assertEqual(accepted["requirements"], [])
        self.assertEqual(
            accepted["clause_reviews"][0]["classification"], "external_compliance",
        )
        self.assertEqual(
            accepted["clause_reviews"][0]["obligations"][0]["status"], "unverifiable",
        )
        self.assertEqual(bridge.validate_host_agent_response(accepted, chunk), [])
        self.assertIn(
            "external_action_relation_projection_v3",
            {item["rule_id"] for item in candidate_audit["mechanical_repairs"]},
        )
        self.assertEqual(
            candidate_audit["mechanical_repairs"][0]["removed_requirements"],
            parent["requirements"],
        )
        guidance = bridge._structured_contract_repair_guidance(
            all_retry_records, contract_version="3.0",
        )
        self.assertIn("explicitly supported by this exact clause", guidance)
        self.assertIn("unverifiable", guidance)
        self.assertIn("do not remove the source-only requirement in this retry", guidance)
        with tempfile.TemporaryDirectory() as td:
            directory = Path(td)
            chunk_path = directory / "chunk.json"
            parent_path = directory / "parent.json"
            bridge._write_json(chunk_path, chunk)
            bridge._write_json(parent_path, parent)
            prompt = bridge._host_prompt(
                request_path=directory / "request.json",
                chunk_path=chunk_path,
                response_path=directory / "response.json",
                run_id="run-external-obligation-retry",
                chunk_index=3, chunk_count=22, attempt=2,
                retry_hint="non_requirement_classification_relation",
                retry_parent_response_sha256=bridge.sha256_file(parent_path),
                retry_parent_response_path=parent_path,
                retry_error_records=all_retry_records,
            )
            self.assertIn("do not remove the source-only requirement in this retry", prompt)
            self.assertIn("preserve every requirement", prompt)
        self.assertNotIn("remove only this requirement object", prompt)

    def test_external_action_projection_rejects_non_null_or_conditional_payload(self) -> None:
        source = "北京体育大学学位评定委员会办公室盖章(有效)"
        evidence = {"evidence": [{"id": "E1", "text": source, "kind": "paragraph"}]}
        clauses = self._bind_test_source_spans(
            [exact_source_clause("C1", source)], evidence,
        )
        chunk = engine.build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
        )
        chunk = attach_request_provenance(
            chunk, source_sha256="f" * 64, evidence_doc=evidence,
            clauses=clauses, run_id="external-action-null-payload-rejection",
        )
        baseline = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "existing_requirement_id": None, "field_key": None,
                "role": "body_text", "properties": {"text": source},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "confidence": 0.95, "reason": "The source names an external action.",
                "applicability": {
                    "status": "always", "conditions": None, "exceptions": None,
                },
                "input_prerequisites": None, "verification": {
                    "mode": "external", "checks": ["Confirm the physical action."],
                    "checker_ids": None,
                },
            }],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "external_compliance",
                "normative_basis": "external_duty", "reason": "A physical action is required.",
                "obligations": [{
                    "id": "office_stamp", "status": "unverifiable",
                    "reason": "The real-world stamp must be applied by the office.",
                }],
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        for label, mutate in (
            ("non_null_role_property", lambda item: item["properties"].update(prefix="must remain")),
            ("conditional_applicability", lambda item: item.update(applicability={
                "status": "not_applicable", "conditions": None, "exceptions": None,
            })),
            ("local_verification", lambda item: item.update(verification={
                "mode": "local", "checks": ["Verify in the DOCX."], "checker_ids": None,
            })),
            ("input_prerequisite", lambda item: item.update(input_prerequisites=[{
                "kind": "source_content", "key": "source_inventory.thesis_title_zh",
                "reason": "A required source value.",
            }])),
        ):
            candidate = copy.deepcopy(baseline)
            mutate(candidate["requirements"][0])
            candidate = bridge.normalize_native_response(candidate, chunk["response_schema"])
            with self.subTest(case=label):
                self.assertFalse(
                    bridge._is_source_only_external_requirement(candidate["requirements"][0], chunk)
                )

        null_only = copy.deepcopy(baseline["requirements"][0])
        for property_name in (
            "font", "paragraph", "numbering", "position", "prefix", "separator",
            "style_hint", "header_content", "bottom_border",
        ):
            null_only["properties"][property_name] = None
        self.assertTrue(bridge._is_source_only_external_requirement(null_only, chunk))

    def test_external_action_projection_rejects_multi_clause_partial_source_echo(self) -> None:
        source_a = "北京体育大学学位评定委员会办公室盖章(有效)"
        source_b = "纸质材料由研究生院办公室核验后接收"
        evidence_doc = {"evidence": [
            {"id": "E-A", "text": source_a, "kind": "paragraph"},
            {"id": "E-B", "text": source_b, "kind": "paragraph"},
        ]}
        clauses = [
            exact_source_clause("C-A", source_a, "E-A"),
            exact_source_clause("C-B", source_b, "E-B"),
        ]
        chunk = engine.build_llm_request(
            [], clauses, evidence_doc, {}, "full", contract_version="3.0",
            runtime_context={"code_fingerprint_sha256": "a" * 64},
        )
        chunk = attach_request_provenance(
            chunk, source_sha256="b" * 64, evidence_doc=evidence_doc,
            clauses=clauses, run_id="run-multi-clause-external-projection",
        )
        response = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "existing_requirement_id": None,
                "role": "body_text", "properties": {"text": source_a},
                # This echoes only C-A while linking both C-A and C-B.
                "clause_ids": ["C-A", "C-B"], "evidence_ids": ["E-A", "E-B"],
                "confidence": 0.95,
                "reason": "The source text is linked to two external actions.",
                "verification": None,
            }],
            "clause_reviews": [
                {"clause_id": clause_id,
                 "classification": "external_compliance",
                 "normative_basis": "external_duty",
                 "reason": "The action is completed outside DOCX generation.",
                 "obligations": None}
                for clause_id in ("C-A", "C-B")
            ],
            "unsupported_items": [], "reported_conflicts": [],
        }
        response = bridge.normalize_native_response(response, chunk["response_schema"])
        errors = bridge.validate_host_agent_response(response, chunk)
        validator_records = bridge.contract_error_records(
            errors, response=response, chunk=chunk,
        )
        self.assertIn(
            "non_requirement_classification_relation",
            {record["code"] for record in validator_records},
        )
        self.assertEqual(
            bridge._external_action_obligation_retry_records(
                response, validator_records, chunk,
            ),
            [],
            "multi-clause edges must not receive source-inventory-only retry authorization",
        )

        # Even with complete-looking external inventories, exact equality to
        # one linked source cannot authorize deleting the combined requirement.
        projected_candidate = copy.deepcopy(response)
        projected_candidate["clause_reviews"][0]["obligations"] = [{
            "id": "office_stamp", "status": "unverifiable",
            "reason": "The named office must apply the stamp.",
        }]
        projected_candidate["clause_reviews"][1]["obligations"] = [{
            "id": "paper_receipt", "status": "unverifiable",
            "reason": "The office must receive the physical materials.",
        }]
        projected_candidate = bridge.normalize_native_response(
            projected_candidate, chunk["response_schema"],
        )
        candidate_records = bridge.contract_error_records(
            bridge.validate_host_agent_response(projected_candidate, chunk),
            response=projected_candidate, chunk=chunk,
        )
        self.assertIsNone(bridge._project_external_action_requirements(
            projected_candidate, candidate_records, chunk,
        )[0])

    def test_retry_projection_keeps_only_exact_validator_targeted_inventory_fields(self) -> None:
        parent = {
            "contract_version": "3.0", "requirements": [], "unsupported_items": [],
            "clause_reviews": [
                {"clause_id": "C00037", "classification": "executable",
                 "reason": "parent reason", "obligations": None},
                {"clause_id": "C00040", "classification": "external_compliance",
                 "normative_basis": "external_duty", "reason": "preserve this reason",
                 "obligations": [{"id": "affirm", "status": "unverifiable", "reason": "source duty"}]},
            ],
        }
        model_retry = copy.deepcopy(parent)
        model_retry["clause_reviews"][0]["obligations"] = [
            {"id": "signature", "status": "covered", "reason": "source duty represented"},
        ]
        model_retry["clause_reviews"][1]["reason"] = "unrequested semantic rewrite"
        records = [{
            "code": "executable_review_obligations_missing",
            "json_pointer": "$.clause_reviews[0].obligations",
            "clause_id": "C00037",
            "response_sha256": bridge._response_sha256(parent),
            "raw_error": "review_requires_non_empty_source_inventory",
        }]

        projected, audit = bridge._project_validator_targeted_obligation_fields(
            parent, model_retry, records,
        )
        self.assertIsNotNone(projected)
        self.assertEqual(projected["clause_reviews"][0]["obligations"],
                         model_retry["clause_reviews"][0]["obligations"])
        self.assertEqual(projected["clause_reviews"][1]["reason"], "preserve this reason")
        self.assertEqual(audit["status"], "projected")
        self.assertEqual(audit["applied_paths"], ["$.clause_reviews[0].obligations"])
        self.assertEqual(audit["discarded_unrequested_paths"], ["$.clause_reviews[1].reason"])

        stale_record = copy.deepcopy(records)
        stale_record[0]["response_sha256"] = "0" * 64
        broad_record = copy.deepcopy(records)
        broad_record[0]["json_pointer"] = "$.clause_reviews"
        wrong_clause_record = copy.deepcopy(records)
        wrong_clause_record[0]["clause_id"] = "C00040"
        reordered = copy.deepcopy(model_retry)
        reordered["clause_reviews"].reverse()
        not_a_completion = copy.deepcopy(model_retry)
        not_a_completion["clause_reviews"][0]["obligations"] = None
        for candidate, candidate_records in (
            (model_retry, stale_record), (model_retry, broad_record),
            (model_retry, wrong_clause_record), (reordered, records),
            (not_a_completion, records),
        ):
            with self.subTest(records=candidate_records, candidate=candidate):
                rejected, rejected_audit = bridge._project_validator_targeted_obligation_fields(
                    parent, candidate, candidate_records,
                )
                self.assertIsNone(rejected)
                self.assertEqual(rejected_audit["status"], "blocked")

    def test_external_null_inventory_is_valid_but_mixed_requirement_edge_stays_blocked(self) -> None:
        sources = {
            "C00046": "审批表编号需由外部审批流程核验。",
            "C00049": "办公室盖章状态需由外部审批流程核验。",
        }
        clauses = [{
            "id": clause_id, "text": source, "evidence_ids": [f"E{clause_id[-2:] or '1'}"],
            "source_kind": "paragraph", "location": {}, "part_index": 0,
        } for clause_id, source in sources.items()]
        evidence = {"evidence": [{
            "id": clause["evidence_ids"][0], "text": clause["text"], "kind": "paragraph",
        } for clause in clauses]}
        clauses = self._bind_test_source_spans(clauses, evidence)
        chunk = engine.build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
            runtime_context={"code_fingerprint_sha256": "9" * 64},
        )
        chunk["batch"] = {"index": 1}
        chunk["case_id"] = "case-mixed-external-inventory-retry"
        chunk = attach_request_provenance(
            chunk, source_sha256="e" * 64, evidence_doc=evidence,
            clauses=clauses, run_id="mixed-external-inventory-retry",
        )
        parent = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "role": "body_text",
                "properties": {"text": sources["C00049"]},
                "clause_ids": list(sources),
                "evidence_ids": [clause["evidence_ids"][0] for clause in clauses],
                "confidence": 0.98,
                "reason": "两条证据都要求外部审批流程核验。",
                "verification": {
                    "mode": "external", "checks": ["由外部审批流程核验"],
                },
            }],
            "clause_reviews": [{
                "clause_id": clause_id, "classification": "external_compliance",
                "normative_basis": "external_duty", "reason": "需外部流程核验。",
                "obligations": None,
            } for clause_id in sources],
            "unsupported_items": [], "reported_conflicts": [],
        }
        errors = bridge.validate_host_agent_response(parent, chunk)
        records = bridge.contract_error_records(errors, response=parent, chunk=chunk)
        record_codes = {record["code"] for record in records}
        self.assertIn("non_requirement_classification_relation", record_codes)
        self.assertNotIn("executable_review_obligations_missing", record_codes)
        self.assertEqual(
            [review["obligations"] for review in parent["clause_reviews"]],
            [None, None],
        )

        model_retry = copy.deepcopy(parent)
        for review in model_retry["clause_reviews"]:
            review["obligations"] = [{
                "id": f"{review['clause_id']}-1", "status": "unverifiable",
                "reason": "该外部审批状态不能由 DOCX 本身证明。",
            }]
        retry_error, changed_paths = bridge._retry_semantic_change_error(
            parent, model_retry, records, contract_version="3.0", chunk=chunk,
        )
        self.assertIsNotNone(retry_error)
        self.assertTrue(all("obligations" in path for path in changed_paths))

    def test_v3_executable_inventory_completion_is_revalidated_and_independently_reviewed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td) / "requirements"
            review_dir.mkdir(parents=True)
            source = "正文使用宋体。"
            unrelated_source = "论文作者须在声明页亲笔签名并填写日期。"
            clauses = [{
                "id": "C00037", "text": source, "evidence_ids": ["E1"],
                "source_kind": "paragraph", "location": {}, "part_index": 0,
            }, {
                "id": "C00040", "text": unrelated_source, "evidence_ids": ["E2"],
                "source_kind": "paragraph", "location": {}, "part_index": 0,
            }]
            evidence = {"evidence": [
                {"id": "E1", "text": source, "kind": "paragraph"},
                {"id": "E2", "text": unrelated_source, "kind": "paragraph"},
            ]}
            clauses = self._bind_test_source_spans(clauses, evidence)
            request = engine.build_llm_request(
                [], clauses, evidence, {}, "full", contract_version="3.0",
                runtime_context={"code_fingerprint_sha256": "f" * 64},
            )
            request = attach_request_provenance(
                request, source_sha256="a" * 64, evidence_doc=evidence,
                clauses=clauses, run_id="run-inventory-bridge-test",
            )
            engine.prepare_host_agent_review_packets(
                request, clauses, evidence, "a" * 64, review_dir, chunk_size=2,
            )
            first = {
                "contract_version": "3.0",
                "requirements": [{
                    "role": "body_text",
                    "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                    "clause_ids": ["C00037"], "evidence_ids": ["E1"],
                    "confidence": 0.98,
                    "reason": "The source explicitly requires SimSun body text.",
                    "verification": {
                        "mode": "word_render", "checks": ["Check body-text font."],
                    },
                }],
                "clause_reviews": [
                    {
                        "clause_id": "C00037", "classification": "executable",
                        "normative_basis": "explicit_normative_text",
                        "reason": "The body-text font can be set in DOCX.",
                        "obligations": None,
                    },
                    {
                        "clause_id": "C00040", "classification": "external_compliance",
                        "normative_basis": "external_duty",
                        "reason": "The author affirms the accuracy of the information.",
                        "obligations": [{
                            "id": "affirm_accuracy", "status": "unverifiable",
                            "reason": "The author must personally affirm the statement.",
                        }],
                    },
                ],
                "unsupported_items": [], "reported_conflicts": [],
            }
            second = copy.deepcopy(first)
            second["clause_reviews"][0]["obligations"] = [{
                "id": "body_font", "status": "covered",
                "reason": "The SimSun font requirement is represented by the DOCX requirement.",
            }]
            second["clause_reviews"][1]["reason"] = (
                "Unrelated rewrite that was not named by the validator."
            )
            envelopes = [
                {"runId": "inventory-retry-1", "status": "ok", "provider": "openai",
                 "model": "gpt-5.6-luna", "result": {"payloads": [{"text": json.dumps(first)}]}},
                {"runId": "inventory-retry-2", "status": "ok", "provider": "openai",
                 "model": "gpt-5.6-luna", "result": {"payloads": [{"text": json.dumps(second)}]}},
            ]
            host_results = [
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(item), "")
                for item in envelopes
            ]
            independent_candidates: list[dict] = []

            def review_candidate(candidate: dict, chunk: dict, **kwargs: dict) -> dict:
                independent_candidates.append(copy.deepcopy(candidate))
                return self._fake_independent_review(candidate, chunk, **kwargs)

            with patch.object(bridge, "_run_independent_obligation_coverage_review",
                              side_effect=review_candidate), \
                    patch.object(bridge, "_run_command", side_effect=host_results) as host_call:
                audit = bridge.run_bridge(
                    review_dir, response_out=Path(td) / "host-agent-response.json",
                    agent_id="main", timeout=1, max_attempts=2,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                )

            self.assertEqual(host_call.call_count, 2)
            self.assertEqual(audit["status"], "merged")
            self.assertEqual(len(independent_candidates), 1)
            accepted_reviews = independent_candidates[0]["clause_reviews"]
            accepted_review = accepted_reviews[0]
            self.assertEqual(accepted_review["classification"], "executable")
            self.assertEqual(accepted_review["obligations"], second["clause_reviews"][0]["obligations"])
            self.assertEqual(
                accepted_reviews[1]["reason"], first["clause_reviews"][1]["reason"],
                "unrequested model drift must not enter the validated candidate",
            )
            retry_comparison = audit["chunk_runs"][0]["retry_stage_comparison"]
            targeted_projection = retry_comparison["validator_targeted_field_projection"]
            self.assertEqual(targeted_projection["status"], "projected")
            self.assertEqual(
                targeted_projection["applied_paths"],
                ["$.clause_reviews[0].obligations"],
            )
            self.assertIn(
                "$.clause_reviews[1].reason",
                targeted_projection["discarded_unrequested_paths"],
            )
            self.assertEqual(
                retry_comparison["raw_semantic_changed_paths"],
                [
                    "$.clause_reviews[0].obligations", "$.clause_reviews[1].reason",
                ],
            )
            self.assertEqual(
                retry_comparison["authorized_projected_changed_paths"],
                ["$.clause_reviews[0].obligations"],
            )
            raw_retry = json.loads(
                Path(audit["chunk_runs"][0]["raw_response_path"]).read_text(encoding="utf-8")
            )
            self.assertEqual(
                raw_retry["clause_reviews"][1]["reason"],
                second["clause_reviews"][1]["reason"],
                "the raw provider response must remain available as unmodified evidence",
            )
            self.assertEqual(raw_retry["requirements"], second["requirements"])
            self.assertEqual(
                independent_candidates[0]["requirements"], second["requirements"],
            )
            self.assertEqual(
                audit["chunk_runs"][0]["semantic_retry_authorizations"]["paths"][0]["rule_id"],
                "v3_source_inventory_completion",
            )

    def test_retry_stage_pair_compares_raw_to_raw_and_cannot_hide_raw_drift(self) -> None:
        raw = {
            "contract_version": bridge.HOST_REVIEW_CONTRACT_V3,
            "requirements": [],
            "clause_reviews": [{
                "clause_id": "C00066", "classification": "executable",
                "reason": "model-authored source summary",
                "obligations": [{"id": "quality", "status": "covered"}],
            }],
            "unsupported_items": [],
        }
        projected = copy.deepcopy(raw)
        projected["clause_reviews"][0]["reason"] = "code-owned complete abstract projection"
        projected["clause_reviews"][0]["obligations"] = [
            {"id": f"abstract_zh.quality_guidance:{index}", "status": "covered"}
            for index in range(5)
        ]
        with tempfile.TemporaryDirectory() as td:
            parent_raw_path = Path(td) / "attempt-01.raw.json"
            current_raw_path = Path(td) / "attempt-02.raw.json"
            parent_raw_path.write_text(json.dumps(raw), encoding="utf-8")
            current_raw_path.write_text(json.dumps(raw), encoding="utf-8")
            parent_loaded, current_loaded = bridge._load_normalized_retry_raw_pair(
                parent_raw_path, current_raw_path, {},
            )
        error, changed = bridge._retry_semantic_change_error(
            parent_loaded, current_loaded, [],
            contract_version=bridge.HOST_REVIEW_CONTRACT_V3,
        )
        self.assertIsNone(error)
        self.assertEqual(changed, [])
        # The old cross-stage pairing would report the code projection as
        # model drift; a real raw change must still fail even if projection
        # later overwrites that field to the same candidate value.
        self.assertTrue(bridge._retry_change_paths(raw, projected))
        changed_raw = copy.deepcopy(raw)
        changed_raw["clause_reviews"][0]["reason"] = "different model assertion"
        error, changed = bridge._retry_semantic_change_error(
            raw, changed_raw, [],
            contract_version=bridge.HOST_REVIEW_CONTRACT_V3,
        )
        self.assertIsNotNone(error)
        self.assertIn("$.clause_reviews[0].reason", changed)

    def test_retry_parent_prompt_rejects_a_hash_for_a_different_file(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            parent_path = Path(td) / "parent.raw.json"
            parent_path.write_text('{"requirements":[]}', encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "path/hash binding mismatch"):
                bridge._host_prompt(
                    request_path=Path(td) / "request.json",
                    chunk_path=Path(td) / "chunk.json",
                    response_path=Path(td) / "candidate.json",
                    run_id="run-hash", chunk_index=1, chunk_count=1, attempt=2,
                    retry_hint="test",
                    retry_parent_response_sha256="0" * 64,
                    retry_parent_response_path=parent_path,
                )

    def test_no_progress_requires_same_raw_plan_candidate_and_complete_fingerprints(self) -> None:
        raw = {"clause_reviews": [{"clause_id": "C1", "reason": "same"}]}
        errors = [{"code": "contract_validation_error", "json_pointer": "$.x"}]
        fingerprints = {
            "run_id": "run-retry-test",
            "case_id": "case-retry-test",
            "source_sha256": "a" * 64,
            "clause_sha256": "b" * 64,
            "evidence_sha256": "c" * 64,
            "request_sha256": "d" * 64,
            "chunk_index": 1,
            "chunk_sha256": "9" * 64,
            "schema_sha256": "e" * 64,
            "code_fingerprint_sha256": "f" * 64,
        }
        self.assertTrue(bridge._retry_has_no_progress(
            raw, copy.deepcopy(raw), errors, copy.deepcopy(errors),
            "candidate-sha", "candidate-sha", fingerprints, copy.deepcopy(fingerprints),
        ))
        incomplete_fingerprints = dict(fingerprints)
        incomplete_fingerprints.pop("code_fingerprint_sha256")
        self.assertFalse(bridge._retry_has_no_progress(
            raw, copy.deepcopy(raw), errors, copy.deepcopy(errors),
            "candidate-sha", "candidate-sha", incomplete_fingerprints,
            copy.deepcopy(incomplete_fingerprints),
        ))
        self.assertFalse(bridge._retry_has_no_progress(
            raw, copy.deepcopy(raw), errors, copy.deepcopy(errors),
            "parent-candidate", "new-candidate", fingerprints, copy.deepcopy(fingerprints),
        ))
        for changed_key, changed_value in (
            ("run_id", "another-run"),
            ("case_id", "another-case"),
            ("chunk_index", 2),
            ("chunk_sha256", "8" * 64),
            ("schema_sha256", "7" * 64),
            ("code_fingerprint_sha256", "6" * 64),
        ):
            with self.subTest(changed_key=changed_key):
                changed_fingerprints = {**fingerprints, changed_key: changed_value}
                self.assertFalse(bridge._retry_has_no_progress(
                    raw, copy.deepcopy(raw), errors, copy.deepcopy(errors),
                    "candidate-sha", "candidate-sha", fingerprints, changed_fingerprints,
                ))

    def test_authoring_primary_retry_eligibility_requires_exact_cited_source(self) -> None:
        source = "以下示例内容请作者根据需要自行撰写真实研究内容。"
        evidence = {"E1": {"id": "E1", "text": source}}
        clause = {
            "id": "C1", "text": "规范化后的语义片段", "evidence_ids": ["E1"],
            "source_span": {
                "evidence_id": "E1", "start_offset": 0, "end_offset": len(source),
                "text": source, "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
            },
        }
        review_context = {
            "classification": "informational", "requires_requirement": False,
            "linked_requirements": [], "cited_evidence": evidence,
        }
        missing = [{
            "disposition": "unrepresented", "source_quote": source,
        }]
        self.assertTrue(bridge._v3_authoring_content_retry_is_source_bound(
            review_context, clause, evidence, missing,
        ))
        pending = [{**missing[0], "disposition": "authoring_content_pending"}]
        self.assertTrue(bridge._v3_authoring_content_retry_is_source_bound(
            review_context, clause, evidence, pending,
        ))

        clause_with_context = copy.deepcopy(clause)
        clause_with_context["evidence_ids"] = ["E1", "E2"]
        evidence_with_context = {
            **evidence,
            "E2": {"id": "E2", "text": "相邻说明，不包含作者撰写指令。"},
        }
        context_review = {
            **review_context,
            "cited_evidence": evidence_with_context,
        }
        self.assertTrue(bridge._v3_authoring_content_retry_is_source_bound(
            context_review, clause_with_context, evidence_with_context, missing,
        ))

        linked = {**review_context, "linked_requirements": [{"requirement_ref": "RR1"}]}
        self.assertFalse(bridge._v3_authoring_content_retry_is_source_bound(
            linked, clause, evidence, missing,
        ))
        bad_quote = [{**missing[0], "source_quote": "图题应居中"}]
        self.assertFalse(bridge._v3_authoring_content_retry_is_source_bound(
            review_context, clause, evidence, bad_quote,
        ))
        stale_span_clause = copy.deepcopy(clause)
        stale_span_clause["source_span"]["text"] = "以下示例内容"
        self.assertFalse(bridge._v3_authoring_content_retry_is_source_bound(
            review_context, stale_span_clause, evidence, missing,
        ))

    def test_informational_authoring_pending_makes_bounded_primary_retryable(self) -> None:
        self._independent_review_patch.stop()
        source = "这些文字都是编的，根据需要自己撰写"
        provenance = {
            "run_id": "run-authoring-primary-retry", "source_sha256": "a" * 64,
            "evidence_sha256": "b" * 64, "clause_sha256": "c" * 64,
            "request_sha256": "d" * 64,
        }
        chunk = {
            "provenance": provenance,
            "clauses": [{
                "id": "C1", "text": "规范化后的语义片段", "evidence_ids": ["E1"],
                "source_span": {
                    "evidence_id": "E1", "start_offset": 0, "end_offset": len(source),
                    "text": source,
                    "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
                },
            }],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        response = {
            "provenance": provenance,
            "clause_reviews": [{
                "clause_id": "C1", "classification": "informational",
                "reason": "The sample was mistaken for descriptive text.",
            }],
            "requirements": [],
        }
        incomplete = {
            "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
            "status": "completed", "response_sha256": "f" * 64,
            "results": [{
                "check_id": "C1", "verdict": "incomplete",
                "rationale": "The author-input obligation is unrepresented.",
                "evidence_quotes": [source], "machine_obligation_ids": [],
                "identified_obligations": [{
                    "source_quote": source, "disposition": "authoring_content_pending",
                    "obligation_summary": "The author must supply genuine thesis content.",
                    "requirement_refs": [],
                }],
            }],
            "summary": {"consistent": 0, "incomplete": 1, "uncertain": 0},
        }
        calls = []

        def reject_both_attempts(request: dict, **kwargs: dict) -> dict:
            calls.append(copy.deepcopy(request))
            return bind_mock_review_to_source_spans(
                incomplete, request, kwargs["output_dir"],
            )

        with tempfile.TemporaryDirectory() as td:
            with patch.object(bridge, "run_native_semantic_review", side_effect=reject_both_attempts), \
                    patch.object(bridge.time, "sleep"):
                with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                    bridge._run_independent_obligation_coverage_review(
                        response, chunk, review_dir=Path(td),
                        run_id="run-authoring-primary-retry", chunk_index=1, attempt=1,
                        host_runtime="codex", model="gpt-5.6-luna", timeout=5,
                        agent_id="main", runner="exec", binary="codex", config_path=None,
                        controller=bridge.RunController(),
                    )
        self.assertEqual([item["provider_attempt"] for item in calls], [1])
        self.assertTrue(caught.exception.retryable)
        error_record = caught.exception.error_records[0]
        self.assertEqual(
            error_record["primary_retry_authorization"],
            "source_bound_authoring_content_reclassification_v1",
        )

    def test_authoring_primary_retry_rechecks_complete_candidate_and_keeps_quality_gap(self) -> None:
        self._independent_review_patch.stop()
        source = "本部分主要撰写国内的研究现状，不能是文献资料的简单摘录，需要分类、总结、归纳"
        review_calls = []
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements", source=source, contract_version="3.0")
            parent = {
                "contract_version": "3.0", "provenance": chunk["provenance"],
                "requirements": [], "clause_reviews": [{
                    "clause_id": "C1", "classification": "informational", "reason": "Primary interpretation.",
                }], "unsupported_items": [], "reported_conflicts": [],
            }
            corrected = copy.deepcopy(parent)
            corrected["clause_reviews"][0]["classification"] = "requires_source_content"

            def primary_envelope(payload):
                return subprocess.CompletedProcess(["offline-adapter-double"], 0, json.dumps({
                    "runId": "offline-authoring-retry", "status": "ok", "provider": "openai",
                    "model": "gpt-5.6-luna", "result": {"payloads": [{"text": json.dumps(payload)}]},
                }), "")

            def source_reviewer(request, **kwargs):
                review_calls.append(copy.deepcopy(request))
                obligations = [
                    {"source_quote": source, "disposition": "authoring_content_pending",
                     "obligation_summary": "Write the research-status section.", "requirement_refs": []},
                    {"source_quote": source, "disposition": "authoring_content_pending"
                     if request["attempt"] == 1 else "unrepresented",
                     "obligation_summary": "Synthesize rather than copy literature.", "requirement_refs": []},
                ]
                result = {"protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL, "status": "completed",
                          "response_sha256": bridge.sha256_json(obligations), "summary": {}, "results": [{
                              "check_id": "C1", "verdict": "incomplete", "rationale": "Source duties remain.",
                              "identified_obligations": obligations, "evidence_quotes": [source],
                              "machine_obligation_ids": [],
                          }]}
                bound = bind_mock_review_to_source_spans(result, request, kwargs["output_dir"])
                bridge.validate_obligation_coverage_response({"results": bound["results"]}, request["checks"])
                return bound

            with patch.object(bridge, "_run_command", side_effect=[primary_envelope(parent), primary_envelope(corrected)]) as primary, \
                    patch.object(bridge, "run_native_semantic_review", side_effect=source_reviewer), \
                    patch.object(bridge.time, "sleep"):
                with self.assertRaises(bridge.IndependentObligationReviewError):
                    bridge.run_bridge(review_dir, response_out=Path(td) / "accepted.json", agent_id="main",
                                      timeout=1, max_attempts=2, openclaw_bin="openclaw",
                                      model="openai/gpt-5.6-luna")
            self.assertEqual(primary.call_count, 2)
            self.assertEqual([call["attempt"] for call in review_calls], [1, 2, 2])
            self.assertEqual([call["checks"][0]["review_context"]["classification"] for call in review_calls],
                             ["informational", "requires_source_content", "requires_source_content"])
            self.assertFalse((Path(td) / "accepted.json").exists())
            failure = json.loads((review_dir / "host-agent-run.json").read_text())
            self.assertEqual(failure["status"], "failed")
            self.assertEqual(failure["chunk_lifecycle"][0]["attempts"][1]["status"], "failed")
            ledger = json.loads(next(review_dir.glob("*attempt-02/obligation-analysis-ledger.json")).read_text())
            self.assertFalse(ledger["submission_ready"])
            self.assertEqual({entry["disposition"] for entry in ledger["obligations"]},
                             {"authoring_content_pending", "unrepresented"})

    def test_author_content_reclassification_retry_is_exactly_source_bound(self) -> None:
        source = "以下示例内容是编写的，请作者根据需要自行撰写真实研究内容。"
        provenance = {
            "run_id": "run-authoring-retry", "source_sha256": "a" * 64,
            "evidence_sha256": "b" * 64, "clause_sha256": "c" * 64,
            "request_sha256": "d" * 64,
        }
        parent = {
            "contract_version": "3.0",
            "requirements": [],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "informational",
                "reason": "Sample wording was treated as descriptive.",
            }],
        }
        candidate = copy.deepcopy(parent)
        candidate["clause_reviews"][0]["classification"] = "requires_source_content"
        chunk = {
            "provenance": provenance, "batch": {"index": 1},
            "case_id": "case-authoring-retry",
            "response_schema": {"type": "object"},
            "runtime_context": {"code_fingerprint_sha256": "9" * 64},
            "clauses": [exact_source_clause("C1", source)],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        record = {
            "code": "independent_obligation_review_incomplete", "clause_id": "C1",
            "json_pointer": "$.clause_reviews[0].classification",
            "baseline_classification": "informational", "evidence_ids": ["E1"],
            "missing_source_quotes": ["请作者根据需要自行撰写真实研究内容"],
            "candidate_response_sha256": bridge._response_sha256(
                bridge._bind_current_invocation_provenance(parent, provenance)
            ),
            "candidate_semantic_sha256": bridge._response_sha256(bridge._semantic_retry_view(parent)),
            "review_request_sha256": "e" * 64, "review_response_sha256": "f" * 64,
        }
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, repair_audit = bridge._v3_authoring_content_reclassification_response(
                parent, candidate, [record], chunk=chunk,
            )
            changed_paths = bridge._retry_change_paths(parent, candidate)
            change_error, authorized_paths = bridge._retry_semantic_change_error(
                parent, candidate, [record], contract_version="3.0", chunk=chunk,
            )
            authorization_ledger = bridge._retry_authorization_ledger(
                parent, candidate, [record], changed_paths,
                contract_version="3.0", chunk=chunk,
            )

        self.assertEqual(repaired, candidate)
        self.assertEqual(repair_audit["affected_clause_ids"], ["C1"])
        self.assertEqual(changed_paths, ["$.clause_reviews[0].classification"])
        self.assertIsNone(change_error)
        self.assertEqual(authorized_paths, changed_paths)
        self.assertTrue(authorization_ledger[0]["source_binding_complete"])

        multi_evidence_chunk = copy.deepcopy(chunk)
        multi_evidence_chunk["clauses"][0]["evidence_ids"] = ["E1", "E2"]
        multi_evidence_chunk["evidence_context"]["E2"] = {
            "id": "E2", "text": "相邻上下文不重复作者撰写指令。",
        }
        multi_evidence_record = {**record, "evidence_ids": ["E1", "E2"]}
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            multi_repaired, _ = bridge._v3_authoring_content_reclassification_response(
                parent, candidate, [multi_evidence_record], chunk=multi_evidence_chunk,
            )
        self.assertEqual(multi_repaired, candidate)

        changed_reason = copy.deepcopy(candidate)
        changed_reason["clause_reviews"][0]["reason"] = "Unrelated rewrite"
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertIsNone(bridge._v3_authoring_content_reclassification_response(
                parent, changed_reason, [record], chunk=chunk,
            )[0])

        forged_record = {**record, "candidate_response_sha256": "0" * 64}
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertIsNone(bridge._v3_authoring_content_reclassification_response(
                parent, candidate, [forged_record], chunk=chunk,
            )[0])

        unrelated_source = "图题应居中，编号按章节递增。"
        unrelated_chunk = copy.deepcopy(chunk)
        unrelated_chunk["clauses"][0]["text"] = unrelated_source
        unrelated_chunk["evidence_context"]["E1"]["text"] = unrelated_source
        unrelated_parent = copy.deepcopy(parent)
        unrelated_candidate = copy.deepcopy(candidate)
        unrelated_record = {
            **record,
            "missing_source_quotes": [unrelated_source],
            "candidate_response_sha256": bridge._response_sha256(
                bridge._bind_current_invocation_provenance(unrelated_parent, provenance)
            ),
            "candidate_semantic_sha256": bridge._response_sha256(
                bridge._semantic_retry_view(unrelated_parent)
            ),
        }
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertIsNone(bridge._v3_authoring_content_reclassification_response(
                unrelated_parent, unrelated_candidate, [unrelated_record], chunk=unrelated_chunk,
            )[0])

    def test_source_verification_reclassification_retry_is_narrow_and_source_bound(self) -> None:
        source = "关键词须源自论文，人工核验现有关键词的出处。"
        provenance = {
            "run_id": "run-source-verification-retry", "source_sha256": "a" * 64,
            "evidence_sha256": "b" * 64, "clause_sha256": "c" * 64,
            "request_sha256": "d" * 64,
        }
        parent = {
            "contract_version": "3.0",
            "requirements": [],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "informational",
                "reason": "This was mistaken for explanatory text.",
            }],
        }
        candidate = copy.deepcopy(parent)
        candidate["clause_reviews"][0]["classification"] = "requires_source_verification"
        chunk = {
            "provenance": provenance, "batch": {"index": 1},
            "case_id": "case-source-verification-retry",
            "response_schema": {"type": "object"},
            "runtime_context": {"code_fingerprint_sha256": "9" * 64},
            "clauses": [exact_source_clause("C1", source)],
            "evidence_context": {"E1": {"id": "E1", "text": source}},
        }
        authorization = "source_bound_existing_content_verification_reclassification_v1"
        record = {
            "code": "independent_obligation_review_incomplete", "clause_id": "C1",
            "json_pointer": "$.clause_reviews[0].classification",
            "baseline_classification": "informational", "evidence_ids": ["E1"],
            "missing_source_quotes": ["关键词须源自论文"],
            "candidate_response_sha256": bridge._response_sha256(
                bridge._bind_current_invocation_provenance(parent, provenance)
            ),
            "candidate_semantic_sha256": bridge._response_sha256(
                bridge._semantic_retry_view(parent)
            ),
            "review_request_sha256": "e" * 64,
            "review_response_sha256": "f" * 64,
            "source_reference_compilation_sha256": "a" * 64,
            "primary_retry_authorization": authorization,
        }

        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, audit = bridge._v3_source_verification_reclassification_response(
                parent, candidate, [record], chunk=chunk,
            )
            changed_paths = bridge._retry_change_paths(parent, candidate)
            change_error, authorized_paths = bridge._retry_semantic_change_error(
                parent, candidate, [record], contract_version="3.0", chunk=chunk,
            )

        self.assertEqual(repaired, candidate)
        self.assertEqual(audit["rule_id"], authorization)
        self.assertEqual(audit["affected_clause_ids"], ["C1"])
        self.assertEqual(audit["changed_field"], "clause_reviews[].classification")
        self.assertEqual(changed_paths, ["$.clause_reviews[0].classification"])
        self.assertIsNone(change_error)
        self.assertEqual(authorized_paths, changed_paths)

        source_content_parent = copy.deepcopy(parent)
        source_content_parent["clause_reviews"][0]["classification"] = "requires_source_content"
        source_content_candidate = copy.deepcopy(source_content_parent)
        source_content_candidate["clause_reviews"][0]["classification"] = "requires_source_verification"
        source_content_record = {
            **record,
            "baseline_classification": "requires_source_content",
            "candidate_response_sha256": bridge._response_sha256(
                bridge._bind_current_invocation_provenance(source_content_parent, provenance)
            ),
            "candidate_semantic_sha256": bridge._response_sha256(
                bridge._semantic_retry_view(source_content_parent)
            ),
        }
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            source_content_repair, source_content_audit = (
                bridge._v3_source_verification_reclassification_response(
                    source_content_parent, source_content_candidate,
                    [source_content_record], chunk=chunk,
                )
            )
        self.assertEqual(source_content_repair, source_content_candidate)
        self.assertEqual(source_content_audit["from"], "requires_source_content")

        authoring_parent = copy.deepcopy(source_content_parent)
        authoring_parent["clause_reviews"][0]["obligations"] = [{
            "id": "O1", "status": "unverifiable", "reason": "The author must provide content.",
        }]
        authoring_candidate = copy.deepcopy(authoring_parent)
        authoring_candidate["clause_reviews"][0]["classification"] = "requires_source_verification"
        authoring_record = {
            **source_content_record,
            "candidate_response_sha256": bridge._response_sha256(
                bridge._bind_current_invocation_provenance(authoring_parent, provenance)
            ),
            "candidate_semantic_sha256": bridge._response_sha256(
                bridge._semantic_retry_view(authoring_parent)
            ),
        }
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            unsafe_repair, _ = bridge._v3_source_verification_reclassification_response(
                authoring_parent, authoring_candidate, [authoring_record], chunk=chunk,
            )
        self.assertIsNone(unsafe_repair)

        unresolved_parent = copy.deepcopy(authoring_parent)
        unresolved_parent["clause_reviews"][0]["classification"] = "unresolved"
        unresolved_parent["clause_reviews"][0]["obligations"][0]["status"] = "unresolved"
        unresolved_candidate = copy.deepcopy(unresolved_parent)
        unresolved_candidate["clause_reviews"][0]["classification"] = "requires_source_verification"
        unresolved_record = {
            **record, "baseline_classification": "unresolved",
            "candidate_response_sha256": bridge._response_sha256(
                bridge._bind_current_invocation_provenance(unresolved_parent, provenance)
            ),
            "candidate_semantic_sha256": bridge._response_sha256(
                bridge._semantic_retry_view(unresolved_parent)
            ),
        }
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            rejected, _ = bridge._v3_source_verification_reclassification_response(
                unresolved_parent, unresolved_candidate, [unresolved_record], chunk=chunk,
            )
        self.assertIsNone(rejected)

        changed_reason = copy.deepcopy(candidate)
        changed_reason["clause_reviews"][0]["reason"] = "A second semantic edit."
        bad_evidence = {**record, "evidence_ids": ["E2"]}
        bad_quote = {**record, "missing_source_quotes": ["此句不在当前来源中"]}
        stale_parent = {**record, "candidate_semantic_sha256": "0" * 64}
        unauthorized = {**record, "primary_retry_authorization": "generic_retry"}
        changed_graph = copy.deepcopy(candidate)
        changed_graph["requirements"] = [{"role": "body_text", "properties": {}}]
        rejected_cases = [
            (parent, changed_reason, [record], chunk),
            (parent, candidate, [bad_evidence], chunk),
            (parent, candidate, [bad_quote], chunk),
            (parent, candidate, [stale_parent], chunk),
            (parent, candidate, [unauthorized], chunk),
            (parent, changed_graph, [record], chunk),
            (parent, candidate, [record, copy.deepcopy(record)], chunk),
        ]
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            for case_parent, case_candidate, case_records, case_chunk in rejected_cases:
                with self.subTest(record_count=len(case_records)):
                    result, _ = bridge._v3_source_verification_reclassification_response(
                        case_parent, case_candidate, case_records, chunk=case_chunk,
                    )
                    self.assertIsNone(result)

    def test_source_verification_retry_uses_source_materialized_candidate_identity(self) -> None:
        keyword_source = "关键词须在论文中有明确出处"
        provenance = {
            "run_id": "run-source-projection", "case_id": "case-source-projection",
            "source_sha256": "a" * 64, "evidence_sha256": "b" * 64,
            "clause_sha256": "c" * 64, "request_sha256": "d" * 64,
            "chunk_sha256": "e" * 64,
        }
        chunk = {
            "provenance": provenance, "batch": {"index": 1},
            "case_id": "case-source-projection",
            "response_schema": {"type": "object"},
            "runtime_context": {"code_fingerprint_sha256": "9" * 64},
            "clauses": [
                {"id": "C1", "text": "学位论文使用授权书", "evidence_ids": ["E1"]},
                {"id": "C2", "text": "固定正文", "evidence_ids": ["E2"]},
                {"id": "C3", "text": "摘要", "evidence_ids": ["E3"]},
                exact_source_clause("C4", keyword_source, "E4"),
            ],
            "evidence_context": {
                "E1": {"id": "E1", "text": "学位论文使用授权书"},
                "E2": {"id": "E2", "text": "固定正文。"},
                "E3": {"id": "E3", "text": "摘要"},
                "E4": {"id": "E4", "text": keyword_source},
            },
            "declaration_anchor_preference": "abstract_title_zh",
        }
        parent_raw = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "declarations", "existing_requirement_id": None,
                "clause_ids": ["C1", "C2"], "evidence_ids": ["E1", "E2"],
                "properties": {"before_role": "abstract_title_zh", "items": [{
                    "id": "authorization", "heading": "学位论文使用授权书",
                    "body": None, "body_parts": ["固定正文"],
                    "source_evidence_ids": ["E1", "E2"],
                    "signature_placeholders": [],
                }]},
            }],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
                {"clause_id": "C3", "classification": "informational"},
                {"clause_id": "C4", "classification": "informational",
                 "reason": "Existing keyword origin needs human verification."},
            ],
            "unsupported_items": [], "reported_conflicts": [],
        }
        retry_raw = copy.deepcopy(parent_raw)
        retry_raw["clause_reviews"][3]["classification"] = "requires_source_verification"
        bind_declaration_fixture_spans(chunk)
        parent_candidate, parent_projection = bridge._materialize_fixed_declaration_source_text(
            parent_raw, chunk,
        )
        retry_candidate, retry_projection = bridge._materialize_fixed_declaration_source_text(
            retry_raw, chunk,
        )
        self.assertTrue(parent_projection)
        self.assertTrue(retry_projection)
        self.assertNotEqual(parent_raw, parent_candidate)
        record = {
            "code": "independent_obligation_review_incomplete", "clause_id": "C4",
            "json_pointer": "$.clause_reviews[3].classification",
            "baseline_classification": "informational", "evidence_ids": ["E4"],
            "missing_source_quotes": [keyword_source],
            "candidate_response_sha256": bridge._response_sha256(
                bridge._bind_current_invocation_provenance(parent_candidate, provenance)
            ),
            "candidate_semantic_sha256": bridge._response_sha256(
                bridge._semantic_retry_view(parent_candidate)
            ),
            "review_request_sha256": "1" * 64,
            "review_response_sha256": "2" * 64,
            "source_reference_compilation_sha256": "3" * 64,
            "primary_retry_authorization": (
                "source_bound_existing_content_verification_reclassification_v1"
            ),
        }
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            direct_candidate, direct_audit = bridge._v3_source_verification_reclassification_response(
                parent_candidate, retry_candidate, [record], chunk=chunk,
            )
            self.assertIsNotNone(direct_candidate, direct_audit)
            rejected_without_projection, _ = bridge._retry_semantic_change_error(
                parent_raw, retry_raw, [record], contract_version="3.0", chunk=chunk,
            )
            accepted, changed = bridge._retry_semantic_change_error(
                parent_raw, retry_raw, [record], contract_version="3.0", chunk=chunk,
                comparison_previous_response=parent_candidate,
                comparison_current_response=retry_candidate,
            )
            tampered = copy.deepcopy(parent_candidate)
            tampered["requirements"][0]["properties"]["items"][0]["body_parts"] = ["伪造正文"]
            rejected_tampered, _ = bridge._retry_semantic_change_error(
                parent_raw, retry_raw, [record], contract_version="3.0", chunk=chunk,
                comparison_previous_response=tampered,
                comparison_current_response=retry_candidate,
            )
            stale_record = {**record, "candidate_semantic_sha256": "0" * 64}
            rejected_stale, _ = bridge._retry_semantic_change_error(
                parent_raw, retry_raw, [stale_record], contract_version="3.0", chunk=chunk,
                comparison_previous_response=parent_candidate,
                comparison_current_response=retry_candidate,
            )
        self.assertIsNotNone(rejected_without_projection)
        self.assertIsNone(accepted)
        self.assertEqual(changed, ["$.clause_reviews[3].classification"])
        self.assertIsNotNone(rejected_tampered)
        self.assertIsNotNone(rejected_stale)

    def test_source_verification_retry_fails_closed_when_prevalidation_compilation_is_tampered(self) -> None:
        self._independent_review_patch.stop()
        source = "关键词须源自论文，人工核验现有关键词是否可追溯至正文。"
        with tempfile.TemporaryDirectory() as td:
            root = Path(td).resolve()
            review_dir, chunk = self._packet(
                root / "requirements", contract_version="3.0", source=source,
            )
            candidate = {
                "contract_version": "3.0", "provenance": chunk["provenance"],
                "requirements": [],
                "clause_reviews": [{
                    "clause_id": "C1", "classification": "informational",
                    "reason": "The rule only describes keywords.",
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }

            def tampered_review(request: dict, *, output_dir: Path, **_kwargs: object) -> dict:
                source_packet = build_source_reference_packet(request)
                source_span = next(
                    span for span in source_packet["checks"][0]["source_spans"]
                    if span["text"] == source
                )
                raw_response = {"results": [{
                    "check_id": "C1",
                    "verdict": "source_content_verification_pending",
                    "rationale": "A human must verify the existing keyword traceability.",
                    "evidence_refs": [source_span["ref_id"]],
                    "identified_obligations": [{
                        "source_ref": source_span["ref_id"],
                        "disposition": "source_content_verification_pending",
                        "obligation_summary": "Verify that each existing keyword is traceable to the thesis.",
                        "requirement_refs": [],
                    }],
                }]}
                compiled_response, compilation = compile_source_reference_response(
                    raw_response, request, bridge.OBLIGATION_COVERAGE_SCHEMA, coverage=True,
                )
                paths = {
                    "request.json": request,
                    "raw-response.json": raw_response,
                    "source-reference-packet.json": source_packet,
                    "compiled-response.json": compiled_response,
                }
                for name, payload in paths.items():
                    bridge._write_json(output_dir / name, payload)
                compilation["compiled_response_sha256"] = "0" * 64
                bridge._write_json(output_dir / "source-reference-compilation.json", compilation)
                bridge.validate_obligation_coverage_response(
                    copy.deepcopy(compiled_response), request["checks"],
                )
                raise AssertionError("expected source-verification classification correction")

            with patch.object(bridge, "run_native_semantic_review", side_effect=tampered_review):
                with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                    bridge._run_independent_obligation_coverage_review(
                        candidate, chunk, review_dir=review_dir,
                        run_id=chunk["provenance"]["run_id"], chunk_index=1,
                        attempt=1, host_runtime="openclaw", model="openai/gpt-5.6-luna",
                        timeout=5, agent_id="main", runner="exec", binary=None,
                        config_path=None, controller=SimpleNamespace(check=lambda: None),
                    )

            self.assertFalse(caught.exception.retryable)
            pointer = caught.exception.independent_review_audit
            failure_audit = json.loads((review_dir / pointer["audit_path"]).read_text(encoding="utf-8"))
            self.assertFalse(failure_audit["retryable"])
            self.assertFalse(failure_audit["error_records"][0]["primary_repairable"])
            self.assertIsNone(failure_audit["error_records"][0]["source_reference_compilation_sha256"])

    def test_source_verification_correction_reaches_bound_red_release_marker(self) -> None:
        self._source_verification_red_marker_roundtrip(project_unresolved=False)

    def test_unresolved_source_check_reaches_bound_red_marker_without_primary_retry(self) -> None:
        self._source_verification_red_marker_roundtrip(project_unresolved=True)

    def test_unresolved_projection_remains_audited_after_independent_failure(self) -> None:
        source = "关键词须源自论文并有明确出处。"
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            review_dir, chunk = self._packet(root / "requirements", contract_version="3.0", source=source)
            raw = {
                "contract_version": "3.0", "provenance": chunk["provenance"],
                "requirements": [], "unsupported_items": [], "reported_conflicts": [],
                "clause_reviews": [{
                    "clause_id": "C1", "classification": "unresolved",
                    "normative_basis": "explicit_normative_text", "reason": "Cannot verify.",
                    "obligations": [{"id": "O1", "status": "unresolved", "reason": "Verify origin."}],
                }],
            }
            envelope = {"runId": "offline-failure", "status": "ok", "provider": "openai",
                        "model": "gpt-5.6-luna", "result": {"payloads": [{"text": json.dumps(raw)}]}}
            result = subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelope), "")
            with patch.object(bridge, "_run_command", return_value=result), \
                    patch.object(bridge, "_run_independent_obligation_coverage_review",
                                 side_effect=RuntimeError("independent source review unavailable")):
                with self.assertRaisesRegex(ValueError, "independent source review unavailable"):
                    bridge.run_bridge(review_dir, response_out=root / "response.json",
                                      agent_id="main", timeout=1, max_attempts=1,
                                      openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                                      host_runtime="openclaw")
            audit = json.loads((review_dir / "host-agent-run.json").read_text())
            attempt = audit["chunk_lifecycle"][0]["attempts"][0]
            projection = attempt["source_verification_classification_projections"][0]
            self.assertEqual(projection["original_primary_review"], raw["clause_reviews"][0])
            self.assertFalse(projection["submission_ready"])
            self.assertFalse((root / "response.json").exists())

    def _source_verification_red_marker_roundtrip(self, *, project_unresolved: bool) -> None:
        """Exercise conflict, authorized retry, accepted receipt, and visible non-release marker offline."""
        source = "论文中的关键词须源自论文，并可追溯至对应原文。"
        with tempfile.TemporaryDirectory() as td:
            root = Path(td).resolve()
            review_dir, chunk = self._packet(
                root / "requirements", contract_version="3.0", source=source,
            )
            misclassified_as_author_input = {
                "contract_version": "3.0",
                "provenance": chunk["provenance"],
                "requirements": [],
                "clause_reviews": [{
                    "clause_id": "C1", "classification": "requires_source_content",
                    "reason": "现有论文正文未包含在该审查请求中，关键词出处需人工核验。",
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }
            corrected = copy.deepcopy(misclassified_as_author_input)
            if project_unresolved:
                review = misclassified_as_author_input["clause_reviews"][0]
                review.update({
                    "classification": "unresolved", "normative_basis": "explicit_normative_text",
                    "obligations": [{"id": "origin-check", "status": "unresolved",
                                     "reason": "The backend cannot verify thesis origin."}],
                })
            corrected["clause_reviews"][0]["classification"] = "requires_source_verification"
            envelopes = [
                {"runId": "offline-source-verification-1", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(misclassified_as_author_input, ensure_ascii=False)}]}},
                {"runId": "offline-source-verification-2", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(corrected, ensure_ascii=False)}]}},
            ]
            primary_results = [
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(item), "")
                for item in envelopes
            ]
            native_calls = 0

            def offline_native_review(request: dict, *, output_dir: Path, **_kwargs: object) -> dict:
                nonlocal native_calls
                native_calls += 1
                output_dir.mkdir(parents=True, exist_ok=True)
                source_packet = build_source_reference_packet(request)
                request_path = output_dir / "request.json"
                compiled_path = output_dir / "compiled-response.json"
                compilation_path = output_dir / "source-reference-compilation.json"
                if native_calls == 1 and not project_unresolved:
                    # Reproduce the production failure lifecycle: persist the
                    # source-bound pre-validation compilation, then let the
                    # validator raise the narrow classification correction.
                    check = source_packet["checks"][0]
                    exact_span = next(span for span in check["source_spans"] if span["text"] == source)
                    raw_response = {"results": [{
                        "check_id": "C1",
                        "verdict": "source_content_verification_pending",
                        "rationale": "关键词出处须由人工在论文正文中核验。",
                        "evidence_refs": [exact_span["ref_id"]],
                        "identified_obligations": [{
                            "source_ref": exact_span["ref_id"],
                            "disposition": "source_content_verification_pending",
                            "obligation_summary": "核验论文现有关键词是否能追溯到对应原文。",
                            "requirement_refs": [],
                        }],
                    }]}
                    compiled_response, compilation = compile_source_reference_response(
                        raw_response, request, bridge.OBLIGATION_COVERAGE_SCHEMA, coverage=True,
                    )
                    packet_path = output_dir / "source-reference-packet.json"
                    raw_path = output_dir / "raw-response.json"
                    for path, payload in (
                        (request_path, request), (packet_path, source_packet),
                        (raw_path, raw_response), (compiled_path, compiled_response),
                        (compilation_path, compilation),
                    ):
                        bridge._write_json(path, payload)
                    bridge.validate_obligation_coverage_response(
                        copy.deepcopy(compiled_response), request["checks"],
                    )
                    raise AssertionError("expected source-verification classification correction")

                check = source_packet["checks"][0]
                exact_span = next(span for span in check["source_spans"] if span["text"] == source)
                raw_response = {"results": [{
                    "check_id": "C1",
                    "verdict": "source_content_verification_pending",
                    "rationale": "关键词出处须由人工在论文正文中核验。",
                    "evidence_refs": [exact_span["ref_id"]],
                    "identified_obligations": [{
                        "source_ref": exact_span["ref_id"],
                        "disposition": "source_content_verification_pending",
                        "obligation_summary": "核验论文现有关键词是否能追溯到对应原文。",
                        "requirement_refs": [],
                    }],
                }]}
                compiled_response, compilation = compile_source_reference_response(
                    raw_response, request, bridge.OBLIGATION_COVERAGE_SCHEMA, coverage=True,
                )
                bridge.validate_obligation_coverage_response(
                    copy.deepcopy(compiled_response), request["checks"],
                )
                packet_path = output_dir / "source-reference-packet.json"
                raw_path = output_dir / "raw-response.json"
                response_path = output_dir / "response.json"
                for path, payload in (
                    (request_path, request),
                    (packet_path, source_packet),
                    (raw_path, raw_response),
                    (compiled_path, compiled_response),
                    (response_path, compiled_response),
                    (compilation_path, compilation),
                ):
                    bridge._write_json(path, payload)
                return {
                    "status": "completed",
                    "protocol": bridge.OBLIGATION_COVERAGE_PROTOCOL,
                    "adapter_id": "openclaw", "host_runtime": "openclaw",
                    "case_id": request.get("case_id"), "run_id": request["run_id"],
                    "request_path": str(request_path.resolve()),
                    "request_sha256": sha256_json(request),
                    "response_path": str(response_path.resolve()),
                    "response_sha256": sha256_file(response_path),
                    "canonical_response_sha256": sha256_json(compiled_response),
                    "compiled_response_path": str(compiled_path.resolve()),
                    "compiled_response_sha256": sha256_file(compiled_path),
                    "raw_response_path": str(raw_path.resolve()),
                    "raw_response_file_sha256": sha256_file(raw_path),
                    "source_reference_protocol": REFERENCE_PROTOCOL,
                    "source_reference_packet_path": str(packet_path.resolve()),
                    "source_reference_packet_sha256": sha256_file(packet_path),
                    "source_reference_compilation_path": str(compilation_path.resolve()),
                    "source_reference_compilation_sha256": sha256_file(compilation_path),
                    "results": compiled_response["results"],
                    "summary": {"consistent": 0, "incomplete": 0, "uncertain": 0},
                }

            self._independent_review_patch.stop()
            response_path = root / "merged-response.json"
            with patch.object(bridge, "_run_command", side_effect=primary_results) as primary_call, \
                    patch.object(bridge, "run_native_semantic_review", side_effect=offline_native_review):
                bridge_audit = bridge.run_bridge(
                    review_dir, response_out=response_path,
                    agent_id="main", timeout=1, max_attempts=2,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                    host_runtime="openclaw",
                )

            self.assertEqual(primary_call.call_count, 1 if project_unresolved else 2)
            self.assertEqual(native_calls, 1 if project_unresolved else 2)
            self.assertEqual(bridge_audit["status"], "merged")
            chunk_audit = bridge_audit["chunk_runs"][0]
            self.assertEqual(chunk_audit["candidate_status"], "accepted_after_independent_review")
            if project_unresolved:
                projection = chunk_audit["source_verification_classification_projections"][0]
                self.assertEqual(projection["before_classification"], "unresolved")
                self.assertTrue(projection["current_evidence_binding_verified"])
                self.assertEqual(projection["original_primary_review"],
                                 misclassified_as_author_input["clause_reviews"][0])
                self.assertFalse(projection["submission_ready"])
            else:
                retry_authorization = chunk_audit["semantic_retry_authorizations"]
                self.assertEqual(
                    retry_authorization["status"],
                    "accepted_after_provenance_and_contract_validation",
                )
                self.assertEqual(
                    retry_authorization["paths"][0]["rule_id"],
                    "v3_source_bound_existing_content_verification_reclassification",
                )
                self.assertEqual(
                    retry_authorization["paths"][0]["authorized_changed_paths"],
                    ["$.clause_reviews[0].classification"],
                )
            merged = json.loads(response_path.read_text(encoding="utf-8"))
            self.assertEqual(merged["clause_reviews"][0]["classification"], "requires_source_verification")
            self.assertEqual(merged["requirements"], misclassified_as_author_input["requirements"])

            request_path = review_dir / "llm-request.json"
            full_request = json.loads(request_path.read_text(encoding="utf-8"))
            extraction_manifest = {
                "run_id": full_request["provenance"]["run_id"],
                "llm_request_body_sha256": request_body_sha256(full_request),
                "llm_request_envelope_sha256": request_envelope_sha256(full_request),
                "llm_request_file_sha256": sha256_file(request_path),
                "runtime_context": full_request.get("runtime_context"),
            }
            receipts = pipeline.validate_host_review_receipts(
                response_path=response_path,
                audit_path=review_dir / "host-agent-run.json",
                receipt_path=review_dir / "merge-receipt.json",
                extraction_manifest=extraction_manifest,
                work=root,
                output_policy="review_draft",
            )
            independent_reviews = receipts["independent_obligation_reviews"]
            pending = independent_reviews[0]["source_content_verification_items"]
            self.assertEqual(len(pending), 1)
            self.assertEqual(pending[0]["analysis_obligation_identity"]["run_id"], extraction_manifest["run_id"])
            self.assertEqual(pending[0]["requirement_refs"], [])
            self.assertFalse(pending[0]["execution_authorized"])

            evidence_doc = {"evidence": list(chunk["evidence_context"].values())}
            gates = pipeline._source_content_verification_release_gates(
                independent_reviews,
                clauses=chunk["clauses"],
                evidence_doc=evidence_doc,
                expected_run_id=extraction_manifest["run_id"],
            )
            self.assertEqual(len(gates), 1)
            self.assertEqual(gates[0]["category"], "semantic_content_review")
            self.assertFalse(gates[0]["execution_authorized"])
            self.assertIn("仍未通过", gates[0]["action"])

            provenance = chunk["provenance"]
            manual_ledger = build_manual_review_ledger(
                {}, [], release_gates=gates,
                binding={
                    "case_id": pending[0]["analysis_obligation_identity"]["case_id"],
                    "run_id": extraction_manifest["run_id"],
                    "source_sha256": provenance["source_sha256"],
                    "clause_sha256": provenance["clause_sha256"],
                    "evidence_sha256": provenance["evidence_sha256"],
                    "request_sha256": extraction_manifest["llm_request_body_sha256"],
                    "requirements_sha256": sha256_json([]),
                    "input_source_sha256": provenance["source_sha256"],
                    "format_spec_sha256": "0" * 64,
                    "official_template_sha256": None,
                    "official_template_source": "not_supplied",
                },
            )
            self.assertFalse(manual_ledger["submission_ready"])
            docx_path = root / "source-verification-draft.docx"
            document = Document()
            marker_receipts = append_manual_review_markers(document, manual_ledger)
            document.save(docx_path)
            marker_audit = audit_manual_review_markers(docx_path, manual_ledger)
            self.assertEqual(len(marker_receipts), 1)
            self.assertTrue(marker_audit["valid"])
            self.assertEqual(marker_audit["visible_marker_count"], 1)
            self.assertEqual(marker_audit["style_errors"], [])
            self.assertFalse(marker_audit["submission_ready"])
            self.assertEqual(
                manual_ledger["items"][0]["marker_id"],
                marker_audit["marker_bindings"][0]["marker_id"],
            )

    def test_retry_allows_only_completion_of_an_empty_requirement_payload(self) -> None:
        previous = self._executable_response({"provenance": {}})
        previous["requirements"][0]["properties"] = {}
        current = self._executable_response({"provenance": {}})
        records = [{
            "code": "empty_requirement_properties",
            "json_pointer": "$.requirements[0].properties",
        }]
        self.assertTrue(
            bridge._retry_changes_allowed(
                records,
                ["$.requirements[0].properties.font"],
                contract_version="2.1",
                previous_response=previous,
                current_response=current,
            )
        )
        current["requirements"][0]["properties"]["font"]["size_pt"] = 11
        self.assertFalse(
            bridge._retry_changes_allowed(
                records,
                ["$.requirements[0].properties.font.size_pt"],
                contract_version="2.1",
                previous_response=current,
                current_response=current,
            )
        )

    def test_empty_payload_fill_may_rewrite_only_the_requirement_reason(self) -> None:
        previous = self._executable_response({"provenance": {}})
        previous["requirements"][0]["properties"] = {}
        previous["requirements"][0]["reason"] = "首轮解释"
        current = self._executable_response({"provenance": {}})
        current["requirements"][0]["reason"] = "重试后的解释"
        records = [{
            "code": "empty_requirement_properties",
            "json_pointer": "$.requirements[0].properties",
        }]
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(
            changed,
            ["$.requirements[0].properties.font", "$.requirements[0].reason"],
        )
        self.assertTrue(bridge._retry_changes_allowed(
            records,
            changed,
            contract_version="2.1",
            previous_response=previous,
            current_response=current,
        ))
        current["requirements"][0]["clause_ids"] = ["C999"]
        changed = bridge._retry_change_paths(previous, current)
        self.assertFalse(bridge._retry_changes_allowed(
            records,
            changed,
            contract_version="2.1",
            previous_response=previous,
            current_response=current,
        ))

    def test_retry_allows_empty_payload_fill_and_validator_named_property_removal(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "equations", "properties": {},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                },
                {
                    "role": "equation", "properties": {"style": "three_line"},
                    "clause_ids": ["C2"], "evidence_ids": ["E2"],
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
            ],
            "unsupported_items": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"][0]["properties"] = {"same_line": False}
        current["requirements"][1]["properties"] = {}
        records = [
            {"code": "empty_requirement_properties", "json_pointer": "$.requirements[0].properties"},
            {
                "code": "unknown_property",
                "json_pointer": "$.requirements[1].properties",
                "raw_error": "unknown property 'style'",
            },
        ]
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(
            changed,
            ["$.requirements[0].properties.same_line", "$.requirements[1].properties.style"],
        )
        self.assertTrue(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current,
        ))

    def test_retry_allows_only_exact_fixed_text_repair(self) -> None:
        previous = {
            "requirements": [
                {
                    "role": "abstract_title_zh", "properties": {"heading": "wrong"},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                },
                {
                    "role": "abstract_title_zh", "properties": {"heading": "wrong other"},
                    "clause_ids": ["C2"], "evidence_ids": ["E2"],
                },
            ],
            "clause_reviews": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"][0]["properties"]["heading"] = "Exact cited text"
        records = [{
            "code": "fixed_text_evidence_mismatch",
            "json_pointer": "$.requirements[0].properties.heading",
        }]
        chunk = {"evidence_context": {"E1": {"text": "Exact cited text"}}}
        self.assertTrue(
            bridge._retry_changes_allowed(
                records,
                ["$.requirements[0].properties.heading"],
                contract_version="2.1",
                previous_response=previous,
                current_response=current,
                chunk=chunk,
            )
        )
        current["requirements"][1]["properties"]["heading"] = "unrelated change"
        self.assertFalse(
            bridge._retry_changes_allowed(
                records,
                ["$.requirements[0].properties.heading", "$.requirements[1].properties.heading"],
                contract_version="2.1",
                previous_response=previous,
                current_response=current,
                chunk=chunk,
            )
        )
        self.assertFalse(
            bridge._retry_changes_allowed(
                records,
                ["$.clause_reviews[0].classification"],
                contract_version="2.1",
                previous_response=previous,
                current_response=current,
                chunk=chunk,
            )
        )

    def test_cover_binding_retry_requires_schema_directed_proof(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "cover",
                "properties": {
                    "institution": "——",
                    "fields": [{"id": "security_marking", "label": "密级"}],
                },
                "clause_ids": ["C1"],
                "evidence_ids": ["E1"],
            }],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "executable",
                "normative_basis": "template_structure",
            }],
            "unsupported_items": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"][0]["properties"]["fields"] = []
        current["requirements"][0]["properties"]["non_public_administration"] = {
            "fields": [{"id": "security_marking", "label": "密级"}],
        }
        current["clause_reviews"][0]["reason"] = "已将密级字段迁移到行政区域。"
        records = [{
            "code": "cover_binding_violation",
            "json_pointer": "$.requirements[0].properties.fields[0]",
        }]
        changed = bridge._retry_change_paths(previous, current)
        self.assertTrue(changed)
        self.assertFalse(
            bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current,
            )
        )

        current["clause_reviews"][0]["classification"] = "not_applicable"
        changed = bridge._retry_change_paths(previous, current)
        self.assertFalse(
            bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current,
            )
        )

    def test_v3_retry_allows_only_new_requirement_for_missing_executable_clause(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "cover_field_label", "properties": {"text": "标题"},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
            }],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
            ],
            "unsupported_items": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"].append({
            "role": "declarations", "properties": {
                "before_role": "abstract_title_zh", "items": [{
                    "id": "authorization", "heading": "授权书",
                    "body_parts": ["固定正文"], "source_evidence_ids": ["E2"],
                    "signature_placeholders": [],
                }],
            }, "clause_ids": ["C2"], "evidence_ids": ["E2"],
        })
        records = [{
            "code": "requirement_relation_mismatch",
            "json_pointer": "$.clause_reviews[1]",
        }]
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(changed, ["$.requirements"])
        self.assertTrue(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current,
        ))

        current["requirements"][0]["properties"]["text"] = "改过的旧要求"
        changed = bridge._retry_change_paths(previous, current)
        self.assertFalse(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current,
        ))

    def test_v3_retry_allows_only_first_occurrence_evidence_deduplication(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "declarations", "properties": {"text": "授权书"},
                "clause_ids": ["C1"], "evidence_ids": ["E1", "E2", "E2"],
            }],
            "clause_reviews": [{"clause_id": "C1", "classification": "executable"}],
            "unsupported_items": [], "reported_conflicts": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"][0]["evidence_ids"] = ["E1", "E2"]
        records = [{
            "code": "duplicate_evidence_ids",
            "json_pointer": "$.requirements[0].evidence_ids",
        }]
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(changed, ["$.requirements"])
        self.assertTrue(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current,
        ))

        current["requirements"][0]["evidence_ids"] = ["E2", "E1"]
        changed = bridge._retry_change_paths(previous, current)
        self.assertFalse(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current,
        ))
        current = json.loads(json.dumps(previous))
        current["requirements"][0]["properties"]["text"] = "改动后的授权书"
        current["requirements"][0]["evidence_ids"] = ["E1", "E2"]
        changed = bridge._retry_change_paths(previous, current)
        self.assertFalse(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current,
        ))

    def test_v3_relation_completion_restores_requirements_after_reason_rewrite(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "heading_2",
                    "properties": {"text": "5.1 研究结论"},
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                    "confidence": 0.98,
                    "reason": "原始解释",
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"][0]["reason"] = "重试后的同义解释"
        current["requirements"].append({
            "role": "heading_3",
            "properties": {"text": "5.1.1 研究结论"},
            "clause_ids": ["C2"],
            "evidence_ids": ["E2"],
            "confidence": 0.97,
            "reason": "新增 requirement 的解释",
        })
        records = [{
            "code": "missing_derived_requirement",
            "json_pointer": "$.clause_reviews[1]",
            "raw_error": "$.clause_reviews[1]: executable_review_requires_derived_requirement",
        }]
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, audit = bridge._v3_relation_completion_response(
                previous, current, records, chunk={},
            )
        self.assertIsNotNone(repaired)
        self.assertEqual(repaired["requirements"][0]["reason"], "原始解释")
        self.assertEqual(
            [item["clause_ids"] for item in repaired["requirements"]],
            [["C1"], ["C2"]],
        )
        self.assertEqual(audit["preserved_requirement_count"], 1)
        self.assertEqual(audit["added_requirement_count"], 1)

        current["requirements"][0]["properties"]["text"] = "改过的旧 requirement"
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, audit = bridge._v3_relation_completion_response(
                previous, current, records, chunk={},
            )
        self.assertIsNone(repaired)
        self.assertIsNone(audit)

    def test_v3_relation_completion_cannot_authorize_top_level_drift(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "heading_2", "properties": {"text": "5.1 结论"},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "confidence": 0.9, "reason": "已有要求",
            }],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"].append({
            "role": "heading_3", "properties": {"text": "5.1.1 结论"},
            "clause_ids": ["C2"], "evidence_ids": ["E2"],
            "confidence": 0.9, "reason": "新增要求",
        })
        records = [{
            "code": "missing_derived_requirement",
            "json_pointer": "$.clause_reviews[1]",
            "raw_error": "$.clause_reviews[1]: executable_review_requires_derived_requirement",
        }]
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, _audit = bridge._v3_relation_completion_response(
                previous, current, records, chunk={},
            )
            self.assertIsNotNone(repaired)

            wrong_contract = json.loads(json.dumps(current))
            wrong_contract["contract_version"] = "2.1"
            repaired, audit = bridge._v3_relation_completion_response(
                previous, wrong_contract, records, chunk={},
            )
            self.assertIsNone(repaired)
            self.assertIsNone(audit)
            error, changed = bridge._retry_semantic_change_error(
                previous, wrong_contract, records,
                contract_version="3.0", chunk={},
            )
            self.assertIsNotNone(error)
            self.assertIn("$.contract_version", changed)

            extra_field = json.loads(json.dumps(current))
            extra_field["unrecognized_semantic_field"] = {"status": "changed"}
            repaired, audit = bridge._v3_relation_completion_response(
                previous, extra_field, records, chunk={},
            )
            self.assertIsNone(repaired)
            self.assertIsNone(audit)
            error, changed = bridge._retry_semantic_change_error(
                previous, extra_field, records,
                contract_version="3.0", chunk={},
            )
            self.assertIsNotNone(error)
            self.assertIn("$.top_level.unrecognized_semantic_field", changed)

            authorization: list[dict] = []
            bound_chunk = {
                "provenance": {
                    "run_id": "run-retry-test", "source_sha256": "a" * 64,
                    "evidence_sha256": "b" * 64, "clause_sha256": "c" * 64,
                    "request_sha256": "d" * 64,
                },
                "case_id": "case-retry-test",
                "batch": {"index": 1},
                "response_schema": {"type": "object"},
                "runtime_context": {"code_fingerprint_sha256": "9" * 64},
                "clauses": [
                    {"id": "C1", "text": "原条款一", "evidence_ids": ["E1"]},
                    {"id": "C2", "text": "原条款二", "evidence_ids": ["E2"]},
                ],
                "evidence_context": {
                    "E1": {"id": "E1", "text": "原条款一"},
                    "E2": {"id": "E2", "text": "原条款二"},
                },
            }
            error, changed = bridge._retry_semantic_change_error(
                previous, current, records,
                contract_version="3.0", chunk=bound_chunk,
                authorization_out=authorization,
            )
            self.assertIsNone(error)
            self.assertEqual(changed, ["$.requirements"])
            self.assertEqual(len(authorization), 1)
            entry = authorization[0]
            self.assertEqual(entry["authorization_type"], "named_response_rule")
            self.assertEqual(entry["old_value_sha256"], bridge._response_sha256(previous["requirements"]))
            self.assertEqual(entry["new_value_sha256"], bridge._response_sha256(current["requirements"]))
            self.assertFalse(entry["validator_records"])
            self.assertTrue(entry["source_binding_complete"])
            self.assertEqual(entry["source_binding"]["run_id"], "run-retry-test")
            self.assertEqual(
                {item["clause_id"] for item in entry["source_evidence_bindings"]},
                {"C1", "C2"},
            )
            self.assertEqual(entry["transform"]["kind"], "constrained_model_retry")
            self.assertTrue(entry["error_bundle_sha256"])
            self.assertEqual(
                entry["response_rule_evidence"][0]["code"],
                "missing_derived_requirement",
            )

    def test_retry_reorder_cannot_transfer_a_field_repair_grant(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {"role": "heading_1", "properties": {"text": "一"},
                 "clause_ids": ["C1"], "evidence_ids": ["E1"], "reason": "原文一"},
                {"role": "heading_2", "properties": {"text": "二"},
                 "clause_ids": ["C2"], "evidence_ids": ["E2"], "reason": "原文二"},
            ],
            "clause_reviews": [], "unsupported_items": [], "reported_conflicts": [],
        }
        reordered = json.loads(json.dumps(previous))
        reordered["requirements"].reverse()
        error, changed = bridge._retry_semantic_change_error(
            previous, reordered, [{"code": "test_only"}], contract_version="3.0",
        )
        self.assertIsNone(error)
        self.assertEqual(changed, [])

        reordered["requirements"][0]["reason"] = "夹带的语义变化"
        error, changed = bridge._retry_semantic_change_error(
            previous, reordered, [{"code": "contract_validation_error"}],
            contract_version="3.0",
        )
        self.assertIsNotNone(error)
        self.assertTrue(changed)

    def test_invalid_normative_basis_does_not_authorize_neighboring_review_change(self) -> None:
        previous = {
            "contract_version": "3.0", "requirements": [], "unsupported_items": [],
            "reported_conflicts": [],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable",
                 "normative_basis": "informational", "reason": "条款一"},
                {"clause_id": "C2", "classification": "executable",
                 "normative_basis": "explicit_normative_text", "reason": "条款二"},
            ],
        }
        current = copy.deepcopy(previous)
        current["clause_reviews"][0]["normative_basis"] = "explicit_normative_text"
        # This second change is legal according to the enum, but no validator
        # finding authorizes changing C2.
        current["clause_reviews"][1]["normative_basis"] = "template_structure"
        records = [{
            "code": "normative_basis_invalid",
            "json_pointer": "$.clause_reviews[0].normative_basis",
            "clause_id": "C1",
        }]
        error, changed = bridge._retry_semantic_change_error(
            previous, current, records, contract_version="3.0",
        )
        self.assertIsNotNone(error)
        self.assertIn("$.clause_reviews[1].normative_basis", changed)

    def test_clause_review_reorder_or_duplicate_identity_cannot_transfer_grant(self) -> None:
        previous = {
            "contract_version": "3.0", "requirements": [], "unsupported_items": [],
            "reported_conflicts": [],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable", "reason": "一"},
                {"clause_id": "C2", "classification": "executable", "reason": "二"},
            ],
        }
        reordered = copy.deepcopy(previous)
        reordered["clause_reviews"].reverse()
        reordered["clause_reviews"][0]["reason"] = "夹带变化"
        error, changed = bridge._retry_semantic_change_error(
            previous, reordered, [{"code": "missing_clause_review"}],
            contract_version="3.0",
        )
        self.assertIsNotNone(error)
        self.assertTrue(changed)

        duplicated = copy.deepcopy(previous)
        duplicated["clause_reviews"][1]["clause_id"] = "C1"
        error, changed = bridge._retry_semantic_change_error(
            previous, duplicated, [{"code": "missing_clause_review"}],
            contract_version="3.0",
        )
        self.assertIsNotNone(error)
        self.assertTrue(changed)

    def test_initial_dispatch_cancellation_writes_failure_audit(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, _chunk = self._packet(Path(td) / "requirements")

            class CancelBeforeDispatch(bridge.RunController):
                def check(self) -> None:
                    raise bridge.HostAgentCancelled("cancelled before initial dispatch")

            response_out = Path(td) / "host-agent-response.json"
            with patch.object(bridge, "_resolve_openclaw", return_value="openclaw"), \
                 patch.object(bridge, "RunController", return_value=CancelBeforeDispatch()):
                with self.assertRaisesRegex(bridge.HostAgentCancelled, "before initial dispatch"):
                    bridge.run_bridge(
                        review_dir, response_out=response_out,
                        agent_id="main", timeout=1, max_attempts=1,
                        openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                    )

            failure = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertEqual(failure["status"], "cancelled")
            self.assertEqual(failure["terminal_status"], "cancelled")
            self.assertFalse(failure["merged_response_written"])
            self.assertFalse(response_out.exists())
            lifecycle = failure["chunk_lifecycle"][0]
            self.assertEqual(lifecycle["status"], "not_started")
            self.assertEqual(lifecycle["dispatch_state"], "not_dispatched")

    def test_interruption_finalizer_closes_nested_running_and_retrying_attempts(self) -> None:
        record = {"attempts": [
            {"attempt": 1, "status": "failed"},
            {"attempt": 2, "status": "retrying"},
            {"attempt": 3, "status": "running"},
        ]}
        bridge._finalize_interrupted_attempts(record, "operator interrupted")
        self.assertEqual(
            [item["status"] for item in record["attempts"]],
            ["failed", "terminated", "terminated"],
        )
        for attempt in record["attempts"][1:]:
            self.assertEqual(attempt["remote_operation_state"], "unknown")
            self.assertEqual(attempt["termination_reason"], "operator interrupted")

    def test_interruption_preserves_completed_failed_attempt_before_retry(self) -> None:
        completed = {
            "attempt": 1, "status": "retrying", "finished_at": "2026-09-29T12:00:00+00:00",
            "error": "local contract failed", "error_records": [{"code": "contract_validation_error"}],
        }
        record = {"attempts": [completed, {"attempt": 2, "status": "running"}]}
        bridge._finalize_interrupted_attempts(record, "sibling failed")
        self.assertEqual(completed["status"], "failed")
        self.assertEqual(completed["finished_at"], "2026-09-29T12:00:00+00:00")
        self.assertEqual(completed["error"], "local contract failed")
        self.assertEqual(completed["subsequent_retry_cancelled_reason"], "sibling failed")
        self.assertNotIn("termination_reason", completed)
        self.assertEqual(record["attempts"][1]["status"], "terminated")

    def test_completed_sibling_is_harvested_before_concurrent_failure_audit(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir = Path(td) / "requirements"
            clauses = [
                {"id": "C1", "text": "正文使用宋体", "evidence_ids": ["E1"],
                 "source_kind": "paragraph", "location": {}, "part_index": 0},
                {"id": "C2", "text": "正文使用黑体", "evidence_ids": ["E2"],
                 "source_kind": "paragraph", "location": {}, "part_index": 0},
            ]
            evidence = {"evidence": [
                {"id": "E1", "text": clauses[0]["text"], "kind": "paragraph"},
                {"id": "E2", "text": clauses[1]["text"], "kind": "paragraph"},
            ]}
            clauses = self._bind_test_source_spans(clauses, evidence)
            request = engine.build_llm_request([], clauses, evidence, {}, "full")
            request = attach_request_provenance(
                request, source_sha256="a" * 64, evidence_doc=evidence, clauses=clauses,
                run_id="concurrent-failure-audit-test",
            )
            engine.prepare_host_agent_review_packets(
                request, clauses, evidence, "a" * 64, review_dir, chunk_size=1,
            )

            def fake_chunk(**kwargs):
                index = kwargs["chunk_index"]
                if index == 1:
                    raise ValueError("synthetic first chunk failure")
                chunk = kwargs["chunk"]
                clause = chunk["clauses"][0]
                response = {
                    "contract_version": chunk["contract_version"],
                    "provenance": chunk["provenance"],
                    "requirements": [],
                    "clause_reviews": [{
                        "clause_id": clause["id"], "classification": "informational",
                        "requirement_indexes": [], "reason": "Completed sibling review.",
                    }],
                    "unsupported_items": [], "reported_conflicts": [],
                }
                kwargs["response_path"].write_text(
                    json.dumps(response, ensure_ascii=False), encoding="utf-8",
                )
                return {"chunk_index": index, "attempt": 1, "status": "ok"}

            response_out = Path(td) / "host-agent-response.json"
            real_wait = bridge.wait

            def wait_until_all_done(futures, return_when):
                return real_wait(futures, return_when=ALL_COMPLETED)

            with patch.object(bridge, "_resolve_openclaw", return_value="openclaw"), \
                 patch.object(bridge, "run_host_agent_chunk", side_effect=fake_chunk), \
                 patch.object(bridge, "wait", side_effect=wait_until_all_done):
                with self.assertRaisesRegex(ValueError, "synthetic first chunk failure"):
                    bridge.run_bridge(
                        review_dir, response_out=response_out,
                        agent_id="main", timeout=2, max_attempts=1,
                        max_concurrency=2, openclaw_bin="openclaw",
                        model="openai/gpt-5.6-luna",
                    )

            failure = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertFalse(failure["merged_response_written"])
            self.assertFalse(response_out.exists())
            self.assertEqual(failure["completed_chunk_indexes"], [2])
            self.assertEqual(failure["in_flight_chunk_indexes"], [])
            self.assertEqual([item["chunk_index"] for item in failure["chunk_runs"]], [2])
            lifecycle = {item["chunk_index"]: item for item in failure["chunk_lifecycle"]}
            self.assertEqual(lifecycle[1]["status"], "failed")
            self.assertEqual(lifecycle[2]["status"], "completed")

    def test_v3_relation_completion_handles_multiple_missing_clauses_and_empty_placeholder(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "body_text", "properties": {"font": {"cjk": "SimSun"}},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                    "confidence": 0.9, "reason": "保留的已有要求",
                },
                {
                    "role": "abstract_title_zh", "properties": {},
                    "clause_ids": [], "evidence_ids": [], "confidence": 0,
                    "reason": "",
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
                {"clause_id": "C3", "classification": "executable"},
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = {
            "contract_version": "3.0",
            "requirements": [
                json.loads(json.dumps(previous["requirements"][0])),
                json.loads(json.dumps(previous["requirements"][1])),
                {
                    "role": "content_constraints", "properties": {"keywords_zh": {"max_chars": 10}},
                    "clause_ids": ["C2"], "evidence_ids": ["E2"],
                    "confidence": 0.9, "reason": "C2 的证据支持关键词限制",
                },
                {
                    "role": "content_constraints", "properties": {"keywords_en": {"max_chars": 10}},
                    "clause_ids": ["C3"], "evidence_ids": ["E3"],
                    "confidence": 0.9, "reason": "C3 的证据支持英文关键词限制",
                },
            ],
            "clause_reviews": previous["clause_reviews"],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        records = [
            {
                "code": "empty_requirement_properties",
                "json_pointer": "$.requirements[1].properties",
                "raw_error": "$.requirements[1].properties: must_include_semantic_payload",
            },
            {
                "code": "missing_derived_requirement",
                "json_pointer": "$.clause_reviews[1]",
                "raw_error": "$.clause_reviews[1]: executable_review_requires_derived_requirement",
            },
            {
                "code": "missing_derived_requirement",
                "json_pointer": "$.clause_reviews[2]",
                "raw_error": "$.clause_reviews[2]: executable_review_requires_derived_requirement",
            },
            {
                "code": "requirement_relation_mismatch",
                "json_pointer": "$.requirements[1]",
                "raw_error": "requirements_not_referenced_by_clause_review:1",
            },
        ]
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, audit = bridge._v3_relation_completion_response(
                previous, current, records, chunk={},
            )
        self.assertIsNotNone(repaired)
        self.assertEqual(
            [item["clause_ids"] for item in repaired["requirements"]],
            [["C1"], ["C2"], ["C3"]],
        )
        self.assertEqual(audit["missing_clause_ids"], ["C2", "C3"])
        self.assertEqual(audit["removed_unbound_placeholder_count"], 1)
        self.assertEqual(audit["added_requirement_count"], 2)

        current["requirements"][2]["clause_ids"] = ["C9"]
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, audit = bridge._v3_relation_completion_response(
                previous, current, records, chunk={},
            )
        self.assertIsNone(repaired)
        self.assertIsNone(audit)

    def test_v3_retry_allows_only_proven_informational_projection(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {"role": "body_text", "properties": {"text": "保留"}, "clause_ids": ["C1"]},
                {"role": "body_text", "properties": {"text": "删除一"}, "clause_ids": ["C2"]},
                {"role": "body_text", "properties": {"text": "删除二"}, "clause_ids": ["C3"]},
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "informational"},
                {"clause_id": "C3", "classification": "informational"},
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"] = [json.loads(json.dumps(previous["requirements"][0]))]
        records = [
            {
                "code": "informational_requirement_forbidden",
                "json_pointer": "$.requirements[1]",
                "requirement_index": 1,
                "relation_category": "informational_only",
                "mechanically_removable": True,
                "response_sha256": bridge._response_sha256(previous),
            },
            {
                "code": "informational_requirement_forbidden",
                "json_pointer": "$.requirements[2]",
                "requirement_index": 2,
                "relation_category": "informational_only",
                "mechanically_removable": True,
                "response_sha256": bridge._response_sha256(previous),
            },
        ]
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(changed, ["$.requirements"])
        self.assertTrue(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current,
            chunk={"clauses": [{"id": f"C{i}"} for i in range(1, 4)]},
        ))
        mismatched_pointer = copy.deepcopy(records)
        mismatched_pointer[0]["json_pointer"] = "$.requirements[999]"
        self.assertFalse(bridge._retry_changes_allowed(
            mismatched_pointer, changed, contract_version="3.0",
            previous_response=previous, current_response=current,
            chunk={"clauses": [{"id": f"C{i}"} for i in range(1, 4)]},
        ))
        self.assertFalse(bridge._retry_changes_allowed(
            records + [copy.deepcopy(records[0])], changed, contract_version="3.0",
            previous_response=previous, current_response=current,
            chunk={"clauses": [{"id": f"C{i}"} for i in range(1, 4)]},
        ))
        current["requirements"][0]["properties"]["text"] = "改动了保留项"
        changed = bridge._retry_change_paths(previous, current)
        self.assertFalse(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current,
            chunk={"clauses": [{"id": f"C{i}"} for i in range(1, 4)]},
        ))

    def test_v3_retry_rejects_generic_unresolved_requirement_removal(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "body_text", "properties": {"text": "保留"},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                },
                {
                    "role": "content_constraints", "properties": {"text": "未解决"},
                    "clause_ids": ["C2"], "evidence_ids": ["E2"],
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "unresolved"},
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = json.loads(json.dumps(previous, ensure_ascii=False))
        current["requirements"].pop(1)
        records = [{
            "code": "non_requirement_classification_relation",
            "json_pointer": "$.requirements[1]",
            "requirement_index": 1,
            "relation_category": "non_requirement_classification",
            "clause_ids": ["C2"],
        }]
        chunk = {"clauses": [{"id": "C1"}, {"id": "C2"}]}
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(changed, ["$.requirements"])
        self.assertFalse(bridge._v3_non_requirement_projection_allowed(
            previous, current, records, changed, chunk=chunk,
        ))
        self.assertFalse(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current, chunk=chunk,
        ))

        unsafe = json.loads(json.dumps(previous, ensure_ascii=False))
        unsafe["requirements"].pop(0)
        self.assertFalse(bridge._retry_changes_allowed(
            records, bridge._retry_change_paths(previous, unsafe),
            contract_version="3.0", previous_response=previous,
            current_response=unsafe, chunk=chunk,
        ))

    def test_v3_uncovered_obligation_retry_drops_stale_empty_requirement_shell(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "figure_caption",
                "properties": {"position": "below"},
                "clause_ids": ["C1"],
                "evidence_ids": ["E1"],
                "existing_requirement_id": "R1",
                "reason": "现有图题要求。",
            }],
            "clause_reviews": [{
                "clause_id": "C1",
                "classification": "verify_existing",
                "obligations": [{
                    "id": "C1.separator",
                    "status": "unverifiable",
                    "reason": "分隔要求需要人工确认。",
                }],
                "reason": "当前候选只能覆盖部分要求。",
                "normative_basis": "explicit_normative_text",
            }],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = copy.deepcopy(previous)
        current["clause_reviews"][0]["classification"] = "unsupported_backend"
        # This is the exact invalid intermediate shape emitted by the native
        # retry: it removes the last clause edge but leaves the old reference
        # and evidence behind. The bridge must drop this shell, not preserve it.
        current["requirements"][0]["clause_ids"] = []
        records = [{
            "code": "executable_review_obligations_uncovered",
            "json_pointer": "$.clause_reviews[0].obligations",
            "clause_id": "C1",
            "raw_error": "$.clause_reviews[0].obligations: executable_review_requires_all_obligations_covered",
        }]
        chunk = {"clauses": [{"id": "C1"}]}
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, audit = bridge._v3_uncovered_obligation_reclassification_response(
                previous, current, records, chunk=chunk,
            )
        self.assertIsNotNone(repaired)
        assert repaired is not None
        self.assertEqual(repaired["requirements"], [])
        self.assertEqual(
            repaired["clause_reviews"][0]["classification"], "unsupported_backend",
        )
        self.assertEqual(audit["dropped_stale_empty_requirement_count"], 1)

        # A payload edit next to the stale shell is still semantic drift and
        # must remain fail-closed.
        current["requirements"][0]["reason"] = "被篡改的说明"
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            rejected, _ = bridge._v3_uncovered_obligation_reclassification_response(
                previous, current, records, chunk=chunk,
            )
        self.assertIsNone(rejected)

    def test_v3_retry_allows_informational_projection_with_named_payload_cleanup(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "body_text", "properties": {"style": "bad"},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                },
                {
                    "role": "body_text", "properties": {"text": "未解决"},
                    "clause_ids": ["C2"], "evidence_ids": ["E2"],
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "informational"},
            ],
            "unsupported_items": [],
        }
        current = json.loads(json.dumps(previous, ensure_ascii=False))
        current["requirements"] = [{
            "role": "body_text", "properties": {},
            "clause_ids": ["C1"], "evidence_ids": ["E1"],
        }]
        records = [
            {
                "code": "unknown_property",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "unknown property 'style'",
            },
            {
                "code": "informational_requirement_forbidden",
                "json_pointer": "$.requirements[1]",
                "requirement_index": 1,
                "relation_category": "informational_only",
                "mechanically_removable": True,
                "clause_ids": ["C2"],
                "response_sha256": bridge._response_sha256(previous),
            },
        ]
        records[0]["response_sha256"] = bridge._response_sha256(previous)
        chunk = {
            "clauses": [
                {"id": "C1", "evidence_ids": ["E1"]},
                {"id": "C2", "evidence_ids": ["E2"]},
            ],
            "evidence_context": {
                "E1": {"text": "固定正文"},
                "E2": {"text": "未解决"},
            },
            "requirement_contract": {
                "role_properties_schema": {"body_text": {"$ref": "#/$defs/roleSpec"}},
            },
        }
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(changed, ["$.requirements"])
        self.assertTrue(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current, chunk=chunk,
        ))

        current["requirements"][0]["properties"] = {"text": "猜测正文"}
        self.assertFalse(bridge._retry_changes_allowed(
            records, bridge._retry_change_paths(previous, current),
            contract_version="3.0", previous_response=previous,
            current_response=current, chunk=chunk,
        ))

    def test_v3_retry_allows_only_exact_fixed_declaration_completion(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "body_text",
                    "properties": {"text": "保留 requirement"},
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                },
                {
                    "role": "table_text",
                    "properties": {},
                    "clause_ids": [],
                    "evidence_ids": [],
                    "confidence": 0,
                    "reason": "",
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable", "reason": "保留"},
                {"clause_id": "C2", "classification": "executable", "reason": "原始标题"},
                {"clause_id": "C3", "classification": "executable", "reason": "原始正文"},
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = json.loads(json.dumps(previous, ensure_ascii=False))
        current["requirements"] = [
            json.loads(json.dumps(previous["requirements"][0], ensure_ascii=False)),
            {
                "role": "declarations",
                "properties": {
                    "before_role": "abstract_title_zh",
                    "items": [{
                        "heading": "学位论文使用授权书",
                        "body_parts": ["本人同意提交电子版"],
                        "source_evidence_ids": ["E2", "E3"],
                    }],
                },
                "clause_ids": ["C2", "C3"],
                "evidence_ids": ["E2", "E3"],
                "confidence": 0.99,
                "reason": "由当前证据固定生成",
            },
        ]
        current["clause_reviews"][2]["reason"] = "当前证据明确该固定声明正文"
        records = [
            {
                "code": "empty_requirement_properties",
                "json_pointer": "$.requirements[1].properties",
                "raw_error": "must_include_semantic_payload",
            },
            {
                "code": "requirement_relation_mismatch",
                "json_pointer": "$.requirements[1]",
                "raw_error": "requirements_not_referenced_by_clause_review:1",
            },
            {
                "code": "missing_derived_requirement",
                "json_pointer": "$.clause_reviews[1]",
                "raw_error": "executable_review_requires_derived_requirement",
            },
            {
                "code": "missing_derived_requirement",
                "json_pointer": "$.clause_reviews[2]",
                "raw_error": "executable_review_requires_derived_requirement",
            },
        ]
        chunk = {
            "declaration_anchor_preference": "abstract_title_zh",
            "clauses": [
                {"id": "C1", "text": "保留 requirement", "evidence_ids": ["E1"]},
                {"id": "C2", "text": "学位论文使用授权书", "evidence_ids": ["E2"]},
                {"id": "C3", "text": "本人同意提交电子版", "evidence_ids": ["E3"]},
                {"id": "C4", "text": "摘要", "evidence_ids": ["E4"]},
            ],
            "evidence_context": {
                "E1": {"text": "保留 requirement"},
                "E2": {"text": "学位论文使用授权书"},
                "E3": {"text": "本人同意提交电子版"},
                "E4": {"text": "摘要"},
            },
        }
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(
            changed,
            ["$.clause_reviews[2].reason", "$.requirements"],
        )
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertTrue(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))

        current["requirements"][1]["properties"]["items"][0]["body_parts"] = ["猜测正文"]
        changed = bridge._retry_change_paths(previous, current)
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertFalse(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))

    def test_v3_retry_restores_omitted_baseline_requirements_before_accepting_additions(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "cover_field_label", "properties": {"text": "标题"},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                },
                {
                    "role": "body_text", "properties": {"font": {"latin": "Times New Roman"}},
                    "clause_ids": ["C3"], "evidence_ids": ["E3"],
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
                {"clause_id": "C3", "classification": "informational"},
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = {
            "contract_version": "3.0",
            "requirements": [
                json.loads(json.dumps(previous["requirements"][0])),
                {
                    "role": "heading_1", "properties": {"text": "第一章"},
                    "clause_ids": ["C2"], "evidence_ids": ["E2"],
                },
            ],
            "clause_reviews": previous["clause_reviews"],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        records = [{
            "code": "requirement_relation_mismatch",
            "json_pointer": "$.clause_reviews[1]",
            "raw_error": "$.clause_reviews[1]: executable_review_requires_derived_requirement",
        }]
        chunk = {}
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, audit = bridge._v3_relation_completion_response(
                previous, current, records, chunk=chunk,
            )
        self.assertIsNotNone(repaired)
        self.assertEqual(
            [item["clause_ids"] for item in repaired["requirements"]],
            [["C1"], ["C3"], ["C2"]],
        )
        self.assertEqual(audit["preserved_requirement_count"], 2)
        self.assertEqual(audit["added_requirement_count"], 1)

        current["requirements"][0]["properties"]["text"] = "改过的旧要求"
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, audit = bridge._v3_relation_completion_response(
                previous, current, records, chunk=chunk,
            )
        self.assertIsNone(repaired)
        self.assertIsNone(audit)

    def test_v3_relation_completion_accepts_exact_text_after_unknown_property_removal(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "cover_field_label",
                    "properties": {"style": "three_line"},
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                },
                {
                    "role": "body_text",
                    "properties": {"text": "正文"},
                    "clause_ids": ["C2"],
                    "evidence_ids": ["E2"],
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
                {"clause_id": "C3", "classification": "executable"},
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = json.loads(json.dumps(previous, ensure_ascii=False))
        current["requirements"] = [
            {
                "role": "cover_field_label",
                "properties": {"text": "论文题目"},
                "clause_ids": ["C1"],
                "evidence_ids": ["E1"],
            },
            current["requirements"][1],
            {
                "role": "body_text",
                "properties": {"text": "新增正文"},
                "clause_ids": ["C3"],
                "evidence_ids": ["E3"],
            },
        ]
        records = [
            {
                "code": "unknown_property",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "unknown property 'style'",
            },
            {
                "code": "missing_derived_requirement",
                "json_pointer": "$.clause_reviews[2]",
                "clause_id": "C3",
                "raw_error": "$.clause_reviews[2]: executable_review_requires_derived_requirement",
            },
        ]
        chunk = {
            "clauses": [
                {"id": "C1", "evidence_ids": ["E1"]},
                {"id": "C2", "evidence_ids": ["E2"]},
                {"id": "C3", "evidence_ids": ["E3"]},
            ],
            "evidence_context": {
                "E1": {"text": "论文题目"},
                "E2": {"text": "正文"},
                "E3": {"text": "新增正文"},
            },
            "requirement_contract": {
                "role_properties_schema": {
                    "cover_field_label": {"$ref": "#/$defs/roleSpec"},
                    "body_text": {"$ref": "#/$defs/roleSpec"},
                },
            },
        }
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            repaired, audit = bridge._v3_relation_completion_response(
                previous, current, records, chunk=chunk,
            )
        self.assertIsNotNone(repaired)
        self.assertIsNotNone(audit)
        assert repaired is not None
        self.assertEqual(repaired["requirements"][0]["properties"], {"text": "论文题目"})
        self.assertEqual(repaired["requirements"][-1]["clause_ids"], ["C3"])

    def test_v3_retry_allows_bounded_completion_of_truncated_response(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "cover_field_label", "properties": {},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
            }],
            "clause_reviews": [], "unsupported_items": [],
        }
        current = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "cover_field_label", "properties": {"text": "标题"},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
            }, {
                "role": "body_text", "properties": {"text": "正文"},
                "clause_ids": ["C2"], "evidence_ids": ["E2"],
            }],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "unsupported_backend"},
            ],
            "unsupported_items": ["C2：后端无法验证该要求"],
        }
        records = [
            {"code": "empty_requirement_properties"},
            {
                "code": "contract_validation_error",
                "raw_error": "clause_reviews_must_cover_each_chunk_clause_exactly_once",
            },
            {
                "code": "requirement_relation_mismatch",
                "raw_error": "requirements_not_referenced_by_clause_review:0",
            },
        ]
        chunk = {
            "evidence_context": {
                "E1": {"text": "标题"}, "E2": {"text": "正文"},
            },
            "requirement_contract": {
                "role_properties_schema": {
                    "cover_field_label": {"$ref": "#/$defs/roleSpec"},
                    "body_text": {"$ref": "#/$defs/roleSpec"},
                },
            },
        }
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(
            changed,
            ["$.clause_reviews", "$.requirements", "$.unsupported_items"],
        )
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertTrue(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))

        current["unsupported_items"] = ["C9：不属于当前 chunk"]
        changed = bridge._retry_change_paths(previous, current)
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertFalse(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))

    def test_v3_retry_allows_truncated_completion_with_fixed_text_and_placeholder(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "declarations",
                    "properties": {
                        "before_role": "abstract_title_zh",
                        "items": [{
                            "id": "authorization",
                            "body_parts": ["截短的固定声明正文"],
                        }],
                    },
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                    "confidence": 0.9,
                    "reason": "保留声明关系",
                },
                {
                    "role": "heading_1",
                    "properties": {"style": "three_line"},
                    "clause_ids": [],
                    "evidence_ids": [],
                    "confidence": 0,
                    "reason": "",
                },
            ],
            "clause_reviews": [],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        current = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "declarations",
                    "properties": {
                        "before_role": "abstract_title_zh",
                        "items": [{
                            "id": "authorization",
                            "body_parts": ["完整固定声明正文"],
                        }],
                    },
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                    "confidence": 0.9,
                    "reason": "保留声明关系",
                },
                {
                    "role": "body_text",
                    "properties": {"text": "正文"},
                    "clause_ids": ["C2"],
                    "evidence_ids": ["E2"],
                    "confidence": 0.8,
                    "reason": "补齐当前条款",
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        records = [
            {
                "code": "contract_validation_error",
                "json_pointer": "$.requirements[1].reason",
                "raw_error": "$.requirements[1].reason: is shorter than 1 characters",
            },
            {
                "code": "fixed_text_evidence_mismatch",
                "json_pointer": "$.requirements[0].properties.items[0].body_parts[0]",
                "raw_error": "must equal a complete cited source-evidence text",
            },
            {
                "code": "unknown_property",
                "json_pointer": "$.requirements[1].properties",
                "raw_error": "unknown property 'style'",
            },
            {
                "code": "schema_contract_violation",
                "json_pointer": "$.requirements[1].clause_ids",
                "raw_error": "must_be_non_empty",
            },
            {
                "code": "schema_contract_violation",
                "json_pointer": "$.requirements[1].evidence_ids",
                "raw_error": "must_be_non_empty",
            },
            {
                "code": "missing_clause_review",
                "json_pointer": "$.requirements[0]",
                "raw_error": "$.requirements[0]:missing_clause_review:clause_ids=C1",
            },
            {
                "code": "contract_validation_error",
                "raw_error": "clause_reviews_must_cover_each_chunk_clause_exactly_once",
            },
            {
                "code": "requirement_relation_mismatch",
                "json_pointer": "$.requirements[1]",
                "raw_error": "requirements_not_referenced_by_clause_review:1",
                "relation_category": "missing_clause_relation",
            },
        ]
        chunk = {
            "clauses": [{"id": "C1"}, {"id": "C2"}],
            "evidence_context": {
                "E1": {"text": "完整固定声明正文"},
                "E2": {"text": "正文"},
            },
        }
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(changed, ["$.clause_reviews", "$.requirements"])
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertTrue(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))

        current["requirements"][0]["properties"]["items"][0]["body_parts"] = ["模型猜测正文"]
        changed = bridge._retry_change_paths(previous, current)
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertFalse(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))

    def test_v3_retry_allows_only_validator_named_missing_clause_review(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "table", "properties": {"style": "three_line"},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
            }, {
                "role": "table", "properties": {"border_widths_pt": {"top": 1.5}},
                "clause_ids": ["C2"], "evidence_ids": ["E2"],
            }],
            "clause_reviews": [{"clause_id": "C1", "classification": "executable"}],
            "unsupported_items": [], "reported_conflicts": [],
        }
        current = json.loads(json.dumps(previous))
        current["clause_reviews"].append({
            "clause_id": "C2", "classification": "informational",
        })
        records = [{
            "code": "missing_clause_review",
            "json_pointer": "$.requirements[0]",
            "raw_error": "$.requirements[0]:missing_clause_review:clause_ids=C2",
        }, {
            "code": "contract_validation_error",
            "raw_error": "clause_reviews_must_cover_each_chunk_clause_exactly_once",
        }]
        chunk = {"clauses": [{"id": "C1"}, {"id": "C2"}]}
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(changed, ["$.clause_reviews"])
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertTrue(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))

        current["clause_reviews"][0]["classification"] = "informational"
        changed = bridge._retry_change_paths(previous, current)
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertFalse(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))

    def test_v3_retry_discards_only_unbound_truncated_placeholder(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "body_text", "properties": {"style": "three_line"},
                "clause_ids": [], "evidence_ids": [], "reason": "",
            }],
            "clause_reviews": [], "unsupported_items": [],
        }
        current = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "cover_field_label", "properties": {"text": "标题"},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
            }],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "executable",
            }],
            "unsupported_items": [],
        }
        records = [
            {"code": "contract_validation_error", "raw_error": "$.requirements[0].reason: is shorter than 1 characters"},
            {"code": "unknown_property", "raw_error": "$.requirements[0].properties: unknown property 'style'"},
            {"code": "schema_contract_violation", "raw_error": "$.requirements[0].clause_ids: must_be_non_empty"},
            {"code": "schema_contract_violation", "raw_error": "$.requirements[0].evidence_ids: must_be_non_empty"},
            {"code": "contract_validation_error", "raw_error": "clause_reviews_must_cover_each_chunk_clause_exactly_once"},
        ]
        chunk = {
            "evidence_context": {"E1": {"text": "标题"}},
            "requirement_contract": {
                "role_properties_schema": {
                    "cover_field_label": {"$ref": "#/$defs/roleSpec"},
                },
            },
        }
        changed = bridge._retry_change_paths(previous, current)
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertTrue(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))
        previous["requirements"][0]["clause_ids"] = ["C0"]
        self.assertFalse(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current, chunk=chunk,
        ))

    def test_v3_retry_allows_schema_directed_cover_completion_only(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "cover", "properties": {
                    "before_role": "document_start", "items": [],
                }, "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "verification": {
                    "mode": "static_docx", "checks": ["保留检查", ""],
                },
            }],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "executable",
            }],
            "unsupported_items": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"][0]["properties"] = {
            "institution": "——",
            "fields": [{
                "id": "title_zh", "label": "论文题目：",
                "value_from": "thesis_profile.cover_metadata.title_zh",
                "display_policy": "required", "order": 1,
            }],
            "before_role": "document_start",
            "missing_value_policy": "placeholder",
            "missing_value_placeholder": "——",
            "layout_id": "linear",
        }
        records = [
            {"code": "contract_validation_error", "json_pointer": "$.requirements[0].properties"},
            {"code": "schema_contract_violation", "json_pointer": "$.requirements[0].properties"},
            {"code": "unknown_property", "json_pointer": "$.requirements[0].properties"},
            {"code": "cover_binding_violation", "json_pointer": "$.requirements[0].properties.fields"},
            {
                "code": "contract_validation_error",
                "json_pointer": "$.requirements[0].verification.checks",
                "raw_error": "$.requirements[0].verification.checks[1]: is shorter than 1 characters",
            },
        ]
        current["requirements"][0]["verification"]["checks"] = ["保留检查"]
        changed = bridge._retry_change_paths(previous, current)
        chunk = {"response_schema": {"type": "object"}}
        with patch.object(bridge, "validate_host_agent_response", return_value=[]):
            self.assertTrue(bridge._retry_changes_allowed(
                records, changed, contract_version="3.0",
                previous_response=previous, current_response=current, chunk=chunk,
            ))
        current["clause_reviews"][0]["classification"] = "not_applicable"
        changed = bridge._retry_change_paths(previous, current)
        self.assertFalse(bridge._retry_changes_allowed(
            records, changed, contract_version="3.0",
            previous_response=previous, current_response=current, chunk=chunk,
        ))

    def test_mechanical_repair_removes_informational_only_requirement(self) -> None:
        response = {
            "requirements": [{
                "role": "body_text", "properties": {"text": "签名"},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
            }],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "informational",
            }],
        }
        record = {
            "code": "informational_requirement_forbidden",
            "json_pointer": "$.requirements[0]",
            "requirement_index": 0,
            "relation_category": "informational_only",
            "mechanically_removable": True,
            "response_sha256": bridge._response_sha256(response),
            "raw_error": "requirements_not_referenced_by_clause_review:0",
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(response, [record])
        self.assertEqual(repaired["requirements"], [])
        self.assertEqual(repairs[0]["removed_clause_ids"], ["C1"])
        self.assertEqual(repairs[0]["removed_requirement"]["clause_ids"], ["C1"])
        self.assertEqual(
            repairs[0]["removed_requirement_sha256"],
            bridge._response_sha256(response["requirements"][0]),
        )
        self.assertEqual(len(response["requirements"]), 1)
        for field, bad_value in (
            ("mechanically_removable", False),
            ("relation_category", "non_requirement_classification"),
            ("response_sha256", "0" * 64),
            ("json_pointer", "$.requirements[9]"),
        ):
            forged = copy.deepcopy(record)
            forged[field] = bad_value
            rejected, _ = bridge._apply_safe_mechanical_repairs(response, [forged])
            self.assertIsNone(rejected, field)
        self.assertIsNone(bridge._apply_safe_mechanical_repairs(
            response, [record, copy.deepcopy(record)],
        )[0])

    def test_external_pending_requirement_edge_is_projected_with_source_bound_audit(self) -> None:
        source = "北京体育大学学位评定委员会办公室盖章(有效)"
        clauses = [{
            "id": "C1", "text": source, "evidence_ids": ["E1"],
            "source_kind": "paragraph", "location": {}, "part_index": 0,
        }]
        evidence = {"evidence": [{"id": "E1", "text": source, "kind": "paragraph"}]}
        clauses = self._bind_test_source_spans(clauses, evidence)
        chunk = engine.build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
        )
        chunk = attach_request_provenance(
            chunk, source_sha256="e" * 64, evidence_doc=evidence,
            clauses=clauses, run_id="external-action-projection-test",
        )
        response = {
            "contract_version": "3.0",
            "provenance": chunk["provenance"],
            "requirements": [{
                "existing_requirement_id": None,
                "role": "body_text", "properties": {"text": source},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "confidence": 0.95, "reason": "The source requires a real-world stamp.",
                # Native structured output emits null for both optional
                # fields when no existing selector/local check applies.
                "verification": None,
            }],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "external_compliance",
                "normative_basis": "external_duty", "reason": "This is a real-world action.",
                "obligations": [{
                    "id": "stamp", "status": "unverifiable",
                    "reason": "The real-world stamp cannot be applied in DOCX.",
                }],
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        normalized_response = bridge.normalize_native_response(
            response, chunk["response_schema"],
        )
        errors = bridge.validate_host_agent_response(normalized_response, chunk)
        records = bridge.contract_error_records(
            errors, response=normalized_response, chunk=chunk,
        )
        self.assertEqual({item["code"] for item in records}, {
            "non_requirement_classification_relation",
        })

        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            normalized_response, records, chunk=chunk,
        )
        self.assertIsNotNone(repaired)
        self.assertEqual(repaired["requirements"], [])
        self.assertEqual(repaired["clause_reviews"], normalized_response["clause_reviews"])
        self.assertEqual(repairs[0]["rule_id"], "external_action_relation_projection_v3")
        self.assertEqual(repairs[0]["removed_requirements"], normalized_response["requirements"])
        self.assertTrue(repairs[0]["external_actions_remain_pending"])
        self.assertEqual(bridge.validate_host_agent_response(repaired, chunk), [])

        accepted, candidate_audit = bridge.prepare_native_response_candidate(response, chunk)
        self.assertEqual(accepted["requirements"], [])
        self.assertEqual(accepted["clause_reviews"], response["clause_reviews"])
        self.assertEqual(bridge.validate_host_agent_response(accepted, chunk), [])
        self.assertIn(
            "external_action_relation_projection_v3",
            {repair["rule_id"] for repair in candidate_audit["mechanical_repairs"]},
        )

        # Missing and explicit-null native fields are semantically identical;
        # neither grants identity or verification authority to the model.
        for identity_mode in ("missing", "null"):
            for verification_mode in ("missing", "null", "external"):
                native_variant = copy.deepcopy(response)
                requirement = native_variant["requirements"][0]
                if identity_mode == "missing":
                    requirement.pop("existing_requirement_id", None)
                if verification_mode == "missing":
                    requirement.pop("verification", None)
                elif verification_mode == "external":
                    requirement["verification"] = {
                        "mode": "external", "checks": ["Obtain and verify the actual stamp."],
                    }
                native_variant = bridge.normalize_native_response(
                    native_variant, chunk["response_schema"],
                )
                variant_errors = bridge.validate_host_agent_response(native_variant, chunk)
                variant_records = bridge.contract_error_records(
                    variant_errors, response=native_variant, chunk=chunk,
                )
                with self.subTest(identity=identity_mode, verification=verification_mode):
                    variant_repaired, variant_audit = bridge._apply_safe_mechanical_repairs(
                        native_variant, variant_records, chunk=chunk,
                    )
                    self.assertIsNotNone(variant_repaired)
                    self.assertEqual(variant_repaired["requirements"], [])
                    self.assertEqual(
                        variant_audit[0]["rule_id"],
                        "external_action_relation_projection_v3",
                    )

        identified = copy.deepcopy(response)
        identified["requirements"][0]["existing_requirement_id"] = "R-existing"
        identified = bridge.normalize_native_response(identified, chunk["response_schema"])
        identified_records = bridge.contract_error_records(
            bridge.validate_host_agent_response(identified, chunk),
            response=identified, chunk=chunk,
        )
        self.assertIsNone(bridge._project_external_action_requirements(
            identified, identified_records, chunk,
        )[0])

        local_verification = copy.deepcopy(response)
        local_verification["requirements"][0]["verification"] = {
            "mode": "static_docx", "checks": ["look for the seal in the document"],
        }
        local_verification = bridge.normalize_native_response(
            local_verification, chunk["response_schema"],
        )
        local_records = bridge.contract_error_records(
            bridge.validate_host_agent_response(local_verification, chunk),
            response=local_verification, chunk=chunk,
        )
        self.assertIsNone(bridge._project_external_action_requirements(
            local_verification, local_records, chunk,
        )[0])

        # Even a schema-valid body_text shell is not removable if its payload
        # is only a fragment rather than an exact complete source echo.
        local_payload = copy.deepcopy(normalized_response)
        local_payload["requirements"][0]["properties"]["text"] = source[:8]
        local_payload_errors = bridge.validate_host_agent_response(local_payload, chunk)
        local_payload_records = bridge.contract_error_records(
            local_payload_errors, response=local_payload, chunk=chunk,
        )
        self.assertEqual(
            {item["code"] for item in local_payload_records},
            {"non_requirement_classification_relation"},
            local_payload_errors,
        )
        self.assertIsNone(bridge._project_external_action_requirements(
            local_payload, local_payload_records, chunk,
        )[0])

        local_prerequisite = copy.deepcopy(normalized_response)
        local_prerequisite["requirements"][0]["input_prerequisites"] = [{
            "kind": "metadata", "key": "thesis_profile.cover_metadata.title_zh",
            "required": True, "reason": "A local cover title must be supplied.",
        }]
        prerequisite_errors = bridge.validate_host_agent_response(local_prerequisite, chunk)
        prerequisite_records = bridge.contract_error_records(
            prerequisite_errors, response=local_prerequisite, chunk=chunk,
        )
        self.assertEqual(
            {item["code"] for item in prerequisite_records},
            {"non_requirement_classification_relation"},
        )
        self.assertIsNone(bridge._project_external_action_requirements(
            local_prerequisite, prerequisite_records, chunk,
        )[0])

        orphan = copy.deepcopy(response)
        orphan["requirements"][0]["clause_ids"] = []
        orphan["requirements"][0]["evidence_ids"] = []
        orphan = bridge.normalize_native_response(orphan, chunk["response_schema"])
        orphan_errors = bridge.validate_host_agent_response(orphan, chunk)
        orphan_records = bridge.contract_error_records(
            orphan_errors, response=orphan, chunk=chunk,
        )
        self.assertTrue(any(
            item["relation_category"] == "missing_clause_relation"
            for item in orphan_records
            if item["code"] == "requirement_relation_mismatch"
        ))
        self.assertIsNone(bridge._project_external_action_requirements(
            orphan, orphan_records, chunk=chunk,
        )[0], "external-action projection must not guess a relation for an orphan")
        orphan_repaired, orphan_audit = bridge._apply_safe_mechanical_repairs(
            orphan, orphan_records, chunk=chunk,
        )
        self.assertIsNotNone(orphan_repaired)
        self.assertEqual(orphan_repaired["requirements"], [])
        self.assertEqual(
            orphan_repaired["clause_reviews"], orphan["clause_reviews"],
            "unbound-orphan cleanup must preserve every semantic clause review",
        )
        self.assertTrue(any(
            item.get("rule_id") == "remove_unbound_non_placeholder_requirement_v1"
            and item.get("reason")
            for item in orphan_audit
        ), "orphan cleanup must be explicit in the repair audit")

        stale = copy.deepcopy(records)
        stale[0]["response_sha256"] = "0" * 64
        self.assertIsNone(bridge._apply_safe_mechanical_repairs(
            normalized_response, stale, chunk=chunk,
        )[0])
        for field, bad_value in (
            ("relation_category", "informational_only"),
            ("mechanically_removable", True),
            ("json_pointer", "$.requirements[9]"),
        ):
            forged = copy.deepcopy(records)
            forged[0][field] = bad_value
            self.assertIsNone(bridge._project_external_action_requirements(
                normalized_response, forged, chunk,
            )[0], field)
        forged_diagnostic = copy.deepcopy(records)
        forged_diagnostic[0]["raw_error"] = "fabricated validator diagnostic"
        self.assertIsNone(bridge._project_external_action_requirements(
            normalized_response, forged_diagnostic, chunk,
        )[0])
        self.assertIsNone(bridge._project_external_action_requirements(
            normalized_response, records + [copy.deepcopy(records[0])], chunk,
        )[0])
        for evidence_item in (
            {},
            {"id": "E2", "text": source},
            {"id": "E1", "text": "不匹配的来源正文"},
        ):
            bad_chunk = copy.deepcopy(chunk)
            bad_chunk["evidence_context"]["E1"] = evidence_item
            self.assertIsNone(bridge._project_external_action_requirements(
                normalized_response, records, bad_chunk,
            )[0], evidence_item)
        bad_span_chunk = copy.deepcopy(chunk)
        bad_span_chunk["clauses"][0]["source_span"]["text"] = "被篡改的条款原文"
        self.assertIsNone(bridge._project_external_action_requirements(
            normalized_response, records, bad_span_chunk,
        )[0])
        missing_obligations = copy.deepcopy(normalized_response)
        missing_obligations["clause_reviews"][0]["obligations"] = []
        missing_errors = bridge.validate_host_agent_response(missing_obligations, chunk)
        missing_records = bridge.contract_error_records(
            missing_errors, response=missing_obligations, chunk=chunk,
        )
        relation_only = [
            item for item in missing_records
            if item["code"] == "non_requirement_classification_relation"
        ]
        self.assertFalse(any(
            item["code"] == "executable_review_obligations_missing"
            for item in missing_records
        ))
        self.assertTrue(relation_only)
        self.assertIsNone(bridge._project_external_action_requirements(
            missing_obligations, relation_only, chunk,
        )[0])

    def test_null_only_external_shells_are_removed_without_losing_pending_reviews(self) -> None:
        sources = [
            "非公开材料须经导师同意、作者申请和主管部门批准",
            "纸质审批页须由主管办公室盖章方有效",
        ]
        evidence_doc = {"evidence": [
            {"id": f"E{index}", "text": source, "kind": "paragraph"}
            for index, source in enumerate(sources, 1)
        ]}
        clauses = self._bind_test_source_spans([
            exact_source_clause(f"C{index}", source, f"E{index}")
            for index, source in enumerate(sources, 1)
        ], evidence_doc)
        chunk = engine.build_llm_request(
            [], clauses, evidence_doc, {}, "full", contract_version="3.0",
        )
        chunk = attach_request_provenance(
            chunk, source_sha256="c" * 64, evidence_doc=evidence_doc,
            clauses=clauses, run_id="null-external-shell-test",
        )
        response = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "existing_requirement_id": None, "field_key": None,
                "clause_ids": [f"C{index}"], "evidence_ids": [f"E{index}"],
                "role": "content_constraints", "properties": {
                    "abstract_zh": None, "abstract_en": None,
                    "keywords_zh": None, "keywords_en": None,
                    "acknowledgments": None,
                },
                "reason": "The source requires a real-world action.",
                "confidence": 1,
                "verification": {"mode": "external", "checks": [
                    "Verify the actual consent or stamp."
                ], "checker_ids": None},
            } for index in (1, 2)],
            "clause_reviews": [{
                "clause_id": f"C{index}", "classification": "external_compliance",
                "normative_basis": "external_duty",
                "reason": "The real-world action remains pending.",
                "obligations": [{
                    "id": f"external_action_{index}", "status": "unverifiable",
                    "reason": "The real-world action cannot be completed in DOCX.",
                }],
            } for index in (1, 2)],
            "unsupported_items": [], "reported_conflicts": [],
        }
        # Native normalization turns these nullable children into nested
        # empty dictionaries, not a top-level empty properties object.
        response["requirements"][0]["properties"] = {
            "abstract_zh": {"required": None, "min_chars": None},
            "abstract_en": {"required": None},
            "keywords_zh": {"required": None},
            "keywords_en": {"required": None},
            "acknowledgments": {"max_chars": None},
        }
        normalized = bridge.normalize_native_response(response, chunk["response_schema"])
        records = bridge.contract_error_records(
            bridge.validate_host_agent_response(normalized, chunk),
            response=normalized, chunk=chunk,
        )
        self.assertEqual(
            [record["code"] for record in records].count("empty_requirement_properties"), 2,
        )
        self.assertEqual(
            [record["code"] for record in records].count("non_requirement_classification_relation"), 2,
        )
        projected, repairs = bridge._apply_safe_mechanical_repairs(
            normalized, records, chunk=chunk,
        )
        self.assertIsNotNone(projected)
        self.assertEqual(projected["requirements"], [])
        self.assertEqual(projected["clause_reviews"], normalized["clause_reviews"])
        self.assertEqual(bridge.validate_host_agent_response(projected, chunk), [])
        self.assertEqual(repairs[0]["projection_kinds"], {
            "0": "null_external_shell", "1": "null_external_shell",
        })
        self.assertEqual(repairs[0]["removed_requirements"], normalized["requirements"])
        accepted, audit = bridge.prepare_native_response_candidate(response, chunk)
        self.assertEqual(accepted["requirements"], [])
        self.assertEqual(len(audit["mechanical_repairs"]), 1)

        # A conditional approval/stamp duty remains a manual external action.
        # Its applicability is retained in the audit, not treated as an
        # executable document property or silently dropped.
        conditional = copy.deepcopy(normalized)
        for requirement in conditional["requirements"]:
            requirement["applicability"] = {
                "status": "conditional",
                "conditions": [{
                    "fact": "thesis_profile.security_level",
                    "operator": "in",
                    "value": ["restricted", "classified"],
                }],
                "exceptions": ["公开论文不适用"],
            }
        conditional_records = bridge.contract_error_records(
            bridge.validate_host_agent_response(conditional, chunk),
            response=conditional, chunk=chunk,
        )
        conditional_projected, conditional_repairs = bridge._project_external_action_requirements(
            conditional, conditional_records, chunk,
        )
        self.assertIsNotNone(conditional_projected)
        self.assertEqual(conditional_projected["requirements"], [])
        self.assertEqual(conditional_projected["clause_reviews"], conditional["clause_reviews"])
        self.assertEqual(
            conditional_repairs[0]["pending_conditional_applicability"],
            {str(index): requirement["applicability"]
             for index, requirement in enumerate(conditional["requirements"])},
        )
        self.assertEqual(bridge.validate_host_agent_response(conditional_projected, chunk), [])

        for label, mutate in (
            ("existing_identity", lambda item: item.update(existing_requirement_id="R-old")),
            ("invalid_conditional", lambda item: item.update(applicability={
                "status": "conditional", "conditions": [{"fact": "security_level"}],
            })),
            ("local_verification", lambda item: item["verification"].update(mode="static_docx")),
            ("checker_binding", lambda item: item["verification"].update(checker_ids=["checker"])),
            ("non_null_payload", lambda item: item["properties"].update(abstract_zh={"required": True})),
            ("cross_evidence", lambda item: item.update(evidence_ids=["E2"])),
        ):
            changed = copy.deepcopy(normalized)
            mutate(changed["requirements"][0])
            changed_records = bridge.contract_error_records(
                bridge.validate_host_agent_response(changed, chunk),
                response=changed, chunk=chunk,
            )
            with self.subTest(label=label):
                self.assertIsNone(bridge._project_external_action_requirements(
                    changed, changed_records, chunk,
                )[0])
        stale_records = copy.deepcopy(records)
        stale_records[0]["response_sha256"] = "0" * 64
        self.assertIsNone(bridge._project_external_action_requirements(
            normalized, stale_records, chunk,
        )[0])
        missing_inventory = copy.deepcopy(normalized)
        missing_inventory["clause_reviews"][0]["obligations"] = []
        missing_records = bridge.contract_error_records(
            bridge.validate_host_agent_response(missing_inventory, chunk),
            response=missing_inventory, chunk=chunk,
        )
        self.assertIsNone(bridge._project_external_action_requirements(
            missing_inventory, missing_records, chunk,
        )[0])

    def test_external_shell_retry_guidance_does_not_authorize_model_deletion(self) -> None:
        parent = {
            "contract_version": "3.0",
            "requirements": [
                {"role": "cover", "clause_ids": ["C00048"],
                 "evidence_ids": ["E00037", "E00045"],
                 "properties": {"public_policy": "blank"}},
                {"role": "content_constraints", "clause_ids": ["C00039"],
                 "evidence_ids": ["E00037"],
                 "properties": {"abstract_zh": {}}},
            ],
            "clause_reviews": [{"clause_id": "C00039",
                                "classification": "external_compliance",
                                "obligations": [{"id": "instructor_consent",
                                                 "status": "unverifiable"}]}],
        }
        records = [{
            "code": "non_requirement_classification_relation",
            "json_pointer": "$.requirements[1]", "requirement_index": 1,
            "mechanically_removable": False,
        }, {
            "code": "empty_requirement_properties",
            "json_pointer": "$.requirements[1].properties",
        }]
        guidance = bridge._structured_contract_repair_guidance(
            records, contract_version="3.0",
        )
        self.assertIn("do not delete", guidance)
        self.assertIn("Do not invent a property", guidance)
        self.assertNotIn("remove only this requirement object", guidance)
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            packet = root / "chunk.json"
            packet.write_text('{"contract_version":"3.0"}', encoding="utf-8")
            parent_path = root / "parent.json"
            parent_path.write_text(json.dumps(parent), encoding="utf-8")
            prompt = bridge._host_prompt(
                request_path=root / "request.json", chunk_path=packet,
                response_path=root / "response.json", run_id="run",
                chunk_index=3, chunk_count=20, attempt=2,
                retry_hint="contract failed", retry_parent_response_path=parent_path,
                retry_error_records=records,
            )
        self.assertIn("The bridge alone can project a proved empty DOCX shell", prompt)
        self.assertNotIn("remove only the validator-identified requirement objects whose clause_ids", prompt)

        # A model retry that also drops the preserved cover evidence edge is
        # not equivalent to the code-owned removal of the external shell.
        unsafe_retry = copy.deepcopy(parent)
        unsafe_retry["requirements"] = [copy.deepcopy(parent["requirements"][0])]
        unsafe_retry["requirements"][0]["evidence_ids"].remove("E00045")
        self.assertNotEqual(unsafe_retry["requirements"], parent["requirements"][:1])
        error, paths = bridge._retry_semantic_change_error(
            parent, unsafe_retry, records, contract_version="3.0",
        )
        self.assertIsNotNone(error)
        self.assertEqual(paths, ["$.requirements"])
        self.assertEqual(error.error_records[0]["preserved_source_edge_changes"], [{
            "previous_requirement_index": 0,
            "current_requirement_index": 0,
            "role": "cover",
            "clause_ids": ["C00048"],
            "removed_evidence_ids": ["E00045"],
            "added_evidence_ids": [],
            "previous_evidence_ids": ["E00037", "E00045"],
            "current_evidence_ids": ["E00037"],
        }])

    def test_external_shell_and_independent_unknown_property_compose(self) -> None:
        sources = [
            "非公开材料须经导师同意、作者申请和主管部门批准",
            "正文使用宋体",
        ]
        evidence_doc = {"evidence": [
            {"id": f"E{index}", "text": source, "kind": "paragraph"}
            for index, source in enumerate(sources, 1)
        ]}
        clauses = self._bind_test_source_spans([
            exact_source_clause(f"C{index}", source, f"E{index}")
            for index, source in enumerate(sources, 1)
        ], evidence_doc)
        chunk = engine.build_llm_request(
            [], clauses, evidence_doc, {}, "full", contract_version="3.0",
        )
        chunk["case_id"] = "standalone"
        chunk = attach_request_provenance(
            chunk, source_sha256="c" * 64, evidence_doc=evidence_doc,
            clauses=clauses, run_id="composed-external-shell-test",
        )
        response = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "role": "content_constraints",
                "properties": {
                    "abstract_zh": {"required": None, "min_chars": None},
                    "abstract_en": {"required": None},
                    "keywords_zh": {"required": None},
                    "keywords_en": {"required": None},
                    "acknowledgments": {"max_chars": None},
                },
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "reason": "The source requires external consent and approval.",
                "confidence": 0.99,
                "verification": {"mode": "external", "checks": ["Check real approval."]},
            }, {
                "role": "body_text",
                "properties": {"font": {"cjk": "SimSun"}, "style": "three_line"},
                "clause_ids": ["C2"], "evidence_ids": ["E2"],
                "reason": "The source sets the body font.", "confidence": 0.9,
                "verification": {"mode": "word_render", "checks": ["Check font."]},
            }],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "external_compliance",
                "normative_basis": "external_duty", "reason": "External approvals remain pending.",
                "obligations": [
                    {"id": action, "status": "unverifiable", "reason": "Requires a real person."}
                    for action in ("instructor_consent", "author_application", "department_approval")
                ],
            }, {
                "clause_id": "C2", "classification": "executable",
                "reason": "The font can be formatted.",
                "obligations": [{
                    "id": "body_font", "status": "covered",
                    "reason": "The font property covers this source clause.",
                }],
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        candidate = bridge.normalize_native_response(response, chunk["response_schema"])
        records = bridge.contract_error_records(
            bridge.validate_host_agent_response(candidate, chunk),
            response=candidate, chunk=chunk,
        )
        self.assertIn("non_requirement_classification_relation", [item["code"] for item in records])
        self.assertIn("unknown_property", [item["code"] for item in records])
        mechanically_repaired, _repairs = bridge._apply_safe_mechanical_repairs(
            candidate, records, chunk=chunk,
        )
        self.assertIsNotNone(mechanically_repaired, records)
        accepted, audit = bridge.prepare_native_response_candidate(response, chunk)
        self.assertEqual(len(accepted["requirements"]), 1)
        self.assertEqual(accepted["requirements"][0]["clause_ids"], ["C2"])
        self.assertEqual(accepted["requirements"][0]["properties"], {"font": {"cjk": "SimSun"}})
        self.assertEqual(accepted["clause_reviews"], candidate["clause_reviews"])
        self.assertEqual(bridge.validate_host_agent_response(accepted, chunk), [])
        self.assertTrue(any(item.get("rule_id") == "external_action_relation_projection_v3"
                            for item in audit["mechanical_repairs"]))

        # Validator order is not part of repair authority. A stale record,
        # an invalid property on the deleted object, or an independent
        # unrepairable error must never be hidden by the external deletion.
        reversed_candidate, _ = bridge._apply_safe_mechanical_repairs(
            candidate, list(reversed(records)), chunk=chunk,
        )
        self.assertEqual(reversed_candidate, accepted)
        stale_records = copy.deepcopy(records)
        stale_records[0]["response_sha256"] = "0" * 64
        self.assertIsNone(bridge._apply_safe_mechanical_repairs(
            candidate, stale_records, chunk=chunk,
        )[0])
        overlapping = copy.deepcopy(response)
        overlapping["requirements"][0]["properties"]["style"] = "three_line"
        with self.assertRaises(ValueError):
            bridge.prepare_native_response_candidate(overlapping, chunk)
        unrepairable = copy.deepcopy(response)
        unrepairable["requirements"][1]["properties"]["font"]["cjk"] = 123
        with self.assertRaises(ValueError):
            bridge.prepare_native_response_candidate(unrepairable, chunk)

        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            review_dir = root / "requirements"
            engine.prepare_host_agent_review_packets(
                chunk, clauses, evidence_doc, "c" * 64,
                review_dir, chunk_size=2,
            )
            packet = json.loads(
                (review_dir / "llm-request-chunks.json").read_text(encoding="utf-8")
            )[0]
            packet_raw = copy.deepcopy(response)
            packet_raw["provenance"] = packet["provenance"]
            packet_accepted, _ = bridge.prepare_native_response_candidate(
                packet_raw, packet,
            )
            coverage_request = bridge.build_obligation_coverage_request(
                packet_accepted, packet,
                run_id=packet["provenance"]["run_id"], chunk_index=1,
            )
            checks = {item["check_id"]: item for item in coverage_request["checks"]}
            self.assertEqual(checks["C1"]["review_context"]["linked_requirements"], [])
            self.assertEqual(len(checks["C2"]["review_context"]["linked_requirements"]), 1)
            actions = ("导师同意", "作者申请", "主管部门批准")
            action_ids = ("instructor_consent", "author_application", "department_approval")
            coverage_response = {"results": [{
                "check_id": "C1", "verdict": "external_compliance_pending",
                "rationale": "Each real-world approval remains unverified.",
                "evidence_quotes": [sources[0]],
                "machine_obligation_ids": checks["C1"]["review_context"]["machine_obligation_ids"],
                "identified_obligations": [{
                    "source_quote": quote,
                    "disposition": "external_action_pending",
                    "primary_obligation_id": identifier,
                    "requirement_refs": [],
                } for quote, identifier in zip(actions, action_ids)],
            }, {
                "check_id": "C2", "verdict": "consistent",
                "rationale": "The retained font property represents this source.",
                "evidence_quotes": [sources[1]],
                "machine_obligation_ids": checks["C2"]["review_context"]["machine_obligation_ids"],
                "identified_obligations": [{
                    "source_quote": sources[1], "disposition": "represented",
                    "requirement_refs": [checks["C2"]["review_context"]["linked_requirements"][0]["requirement_ref"]],
                }],
            }]}
            validated_coverage = bridge.validate_obligation_coverage_response(
                coverage_response, coverage_request["checks"],
            )
            self.assertEqual(
                [item["primary_obligation_id"] for item in validated_coverage[0]["identified_obligations"]],
                list(action_ids),
            )
            collapsed_coverage = copy.deepcopy(coverage_response)
            collapsed_coverage["results"][0]["identified_obligations"] = [
                collapsed_coverage["results"][0]["identified_obligations"][0]
            ]
            with self.assertRaises(bridge.NativeSemanticReviewError):
                bridge.validate_obligation_coverage_response(
                    collapsed_coverage, coverage_request["checks"],
                )
            def source_bound_review_result(check: dict, source_ref: str) -> dict:
                context = check["review_context"]
                if check["check_id"] == "C1":
                    summaries = {
                        "instructor_consent": "核验导师同意记录",
                        "author_application": "核验作者申请记录",
                        "department_approval": "核验主管部门批准记录",
                    }
                    return {
                        "check_id": "C1", "verdict": "external_compliance_pending",
                        "rationale": "Three actual external approvals remain unverified.",
                        "evidence_refs": [source_ref],
                        "identified_obligations": [{
                            "source_ref": source_ref,
                            "disposition": "external_action_pending",
                            "primary_obligation_id": action,
                            "obligation_summary": summaries[action],
                            "requirement_refs": [],
                        } for action in action_ids],
                    }
                return {
                    "check_id": "C2", "verdict": "consistent",
                    "rationale": "The retained font property represents this clause.",
                    "evidence_refs": [source_ref],
                    "identified_obligations": [{
                        "source_ref": source_ref, "disposition": "represented",
                        "requirement_refs": [
                            context["linked_requirements"][0]["requirement_ref"]
                        ],
                    }],
                }

            primary_envelope = {
                "runId": "offline-composite-repair", "status": "ok",
                "provider": "openai", "model": "gpt-5.6-luna",
                "result": {"payloads": [{"text": json.dumps(packet_raw, ensure_ascii=False)}]},
            }
            primary_result = subprocess.CompletedProcess(
                ["openclaw"], 0, json.dumps(primary_envelope), "",
            )
            self._independent_review_patch.stop()
            response_path = review_dir / "llm-response.json"
            with patch.object(bridge, "_run_command", return_value=primary_result) as primary_call, \
                    patch.object(
                        bridge, "_run_independent_obligation_coverage_review",
                        side_effect=lambda *args, **kwargs: self._fake_independent_review(
                            *args, result_builder=source_bound_review_result, **kwargs,
                        ),
                    ):
                run_audit = bridge.run_bridge(
                    review_dir, response_out=response_path,
                    agent_id="main", timeout=1, max_attempts=1,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                    host_runtime="openclaw",
                )
            self.assertEqual(primary_call.call_count, 1)
            self.assertEqual(run_audit["status"], "merged")
            self.assertEqual(
                [item.get("rule_id") or item.get("code")
                 for item in run_audit["chunk_runs"][0]["mechanical_repairs"]],
                ["unknown_property", "external_action_relation_projection_v3"],
            )
            merged = json.loads(response_path.read_text(encoding="utf-8"))
            self.assertEqual([item["clause_ids"] for item in merged["requirements"]], [["C2"]])
            self.assertEqual(
                [item["clause_id"] for item in merged["clause_reviews"]],
                ["C1", "C2"],
            )
            self.assertEqual(
                [item["id"] for item in merged["clause_reviews"][0]["obligations"]],
                ["instructor_consent", "author_application", "department_approval"],
            )
            receipt = json.loads((review_dir / "merge-receipt.json").read_text(encoding="utf-8"))
            ledger = json.loads((review_dir / "semantic-review-ledger.json").read_text(encoding="utf-8"))
            self.assertEqual(receipt["aggregate_sha256"], engine.sha256_json(merged))
            self.assertEqual(receipt["semantic_review_ledger_sha256"], engine.sha256_json(ledger))

            request_path = review_dir / "llm-request.json"
            full_request = json.loads(request_path.read_text(encoding="utf-8"))
            extraction_manifest = {
                "run_id": full_request["provenance"]["run_id"],
                "llm_request_body_sha256": request_body_sha256(full_request),
                "llm_request_envelope_sha256": request_envelope_sha256(full_request),
                "llm_request_file_sha256": sha256_file(request_path),
                "runtime_context": full_request.get("runtime_context"),
            }
            receipts = pipeline.validate_host_review_receipts(
                response_path=response_path, audit_path=review_dir / "host-agent-run.json",
                receipt_path=review_dir / "merge-receipt.json",
                extraction_manifest=extraction_manifest, work=root,
                output_policy="review_draft",
            )
            independent = receipts["independent_obligation_reviews"]
            self.assertEqual(len(independent), 1)
            self.assertTrue(independent[0]["submission_blocked_by_source_content_verification"])
            self.assertEqual(independent[0]["mixed_external_items"], [])
            pending = independent[0]["source_content_verification_items"]
            self.assertEqual(len(pending), 3)
            self.assertEqual(
                {item["obligation_summary"] for item in pending},
                {"核验导师同意记录", "核验作者申请记录", "核验主管部门批准记录"},
            )
            self.assertEqual(len({item["analysis_obligation_id"] for item in pending}), 3)
            gates = pipeline._source_content_verification_release_gates(
                independent, clauses=clauses, evidence_doc=evidence_doc,
                expected_run_id=extraction_manifest["run_id"],
            )
            self.assertEqual(len(gates), 3)
            self.assertEqual(len({item["analysis_obligation_id"] for item in gates}), 3)
            self.assertTrue(all(not item["execution_authorized"] for item in gates))
            with self.assertRaises(ValueError):
                pipeline.validate_host_review_receipts(
                    response_path=response_path, audit_path=review_dir / "host-agent-run.json",
                    receipt_path=review_dir / "merge-receipt.json",
                    extraction_manifest=extraction_manifest, work=root,
                    output_policy="submission",
                )
            stale_independent = copy.deepcopy(independent)
            stale_independent[0]["source_content_verification_items"][0]["source_text_sha256"] = "0" * 64
            with self.assertRaises(ValueError):
                pipeline._source_content_verification_release_gates(
                    stale_independent, clauses=clauses, evidence_doc=evidence_doc,
                    expected_run_id=extraction_manifest["run_id"],
                )
            provenance = packet["provenance"]
            manual_ledger = build_manual_review_ledger(
                {}, [], release_gates=gates,
                binding={
                    "case_id": "standalone", "run_id": extraction_manifest["run_id"],
                    "source_sha256": provenance["source_sha256"],
                    "clause_sha256": provenance["clause_sha256"],
                    "evidence_sha256": provenance["evidence_sha256"],
                    "request_sha256": extraction_manifest["llm_request_body_sha256"],
                    "requirements_sha256": sha256_json(merged["requirements"]),
                    "input_source_sha256": provenance["source_sha256"],
                    "format_spec_sha256": "0" * 64,
                    "official_template_sha256": None,
                    "official_template_source": "not_supplied",
                },
            )
            self.assertFalse(manual_ledger["submission_ready"])
            draft_path = review_dir / "composite-external-draft.docx"
            draft = Document()
            marker_receipts = append_manual_review_markers(draft, manual_ledger)
            draft.save(draft_path)
            marker_audit = audit_manual_review_markers(draft_path, manual_ledger)
            self.assertEqual(len(marker_receipts), 3)
            self.assertTrue(marker_audit["valid"])
            self.assertEqual(marker_audit["visible_marker_count"], 3)
            self.assertEqual(marker_audit["style_errors"], [])
            draft_text = "\n".join(paragraph.text for paragraph in Document(draft_path).paragraphs)
            self.assertIn("核验导师同意记录", draft_text)
            self.assertIn("核验作者申请记录", draft_text)
            self.assertIn("核验主管部门批准记录", draft_text)
            self.assertFalse(marker_audit["submission_ready"])

    def test_external_projection_refuses_a_clause_with_a_local_field_action_cue(self) -> None:
        source = "封面须写明学号，并由导师签字盖章"
        clauses = [{
            "id": "C-MIXED", "text": source, "evidence_ids": ["E-MIXED"],
            "source_kind": "paragraph", "location": {}, "part_index": 0,
        }]
        evidence = {"evidence": [{"id": "E-MIXED", "text": source, "kind": "paragraph"}]}
        clauses = self._bind_test_source_spans(clauses, evidence)
        chunk = engine.build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
        )
        chunk = attach_request_provenance(
            chunk, source_sha256="f" * 64, evidence_doc=evidence,
            clauses=clauses, run_id="mixed-external-action-projection-test",
        )
        response = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "role": "body_text", "properties": {"text": source},
                "clause_ids": ["C-MIXED"], "evidence_ids": ["E-MIXED"],
                "confidence": 0.95, "reason": "The clause is classified as external.",
            }],
            "clause_reviews": [{
                "clause_id": "C-MIXED", "classification": "external_compliance",
                "normative_basis": "external_duty", "reason": "The seal is external.",
                # Deliberately incomplete inventory to reproduce the adversarial case.
                "obligations": [{
                    "id": "advisor_stamp", "status": "unverifiable",
                    "reason": "The signature/seal is outside DOCX generation.",
                }],
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        normalized = bridge.normalize_native_response(response, chunk["response_schema"])
        errors = bridge.validate_host_agent_response(normalized, chunk)
        records = bridge.contract_error_records(errors, response=normalized, chunk=chunk)
        self.assertTrue(any(
            item["code"] == "non_requirement_classification_relation" for item in records
        ))
        self.assertIsNone(
            bridge._project_external_action_requirements(normalized, records, chunk)[0],
        )

        # These semantically equivalent local actions use different wording
        # from the original cue and must be rejected by the actual projection
        # path, not only by the standalone lexical compiler test.
        for index, variant_source in enumerate((
            "签章后保留签名栏。",
            "签字后在首页保留落款。",
            "装订并保留签字页。",
        ), start=1):
            with self.subTest(source=variant_source):
                variant_clause_id = f"C-MIXED-{index}"
                variant_evidence_id = f"E-MIXED-{index}"
                variant_clauses = [{
                    "id": variant_clause_id, "text": variant_source,
                    "evidence_ids": [variant_evidence_id],
                    "source_kind": "paragraph", "location": {}, "part_index": 0,
                }]
                variant_evidence = {"evidence": [{
                    "id": variant_evidence_id, "text": variant_source, "kind": "paragraph",
                }]}
                variant_clauses = self._bind_test_source_spans(
                    variant_clauses, variant_evidence,
                )
                variant_chunk = engine.build_llm_request(
                    [], variant_clauses, variant_evidence, {}, "full", contract_version="3.0",
                )
                variant_chunk = attach_request_provenance(
                    variant_chunk, source_sha256="a" * 64,
                    evidence_doc=variant_evidence, clauses=variant_clauses,
                    run_id=f"mixed-action-variant-{index}",
                )
                variant_response = {
                    "contract_version": "3.0",
                    "provenance": variant_chunk["provenance"],
                    "requirements": [{
                        "role": "body_text", "properties": {"text": variant_source},
                        "clause_ids": [variant_clause_id],
                        "evidence_ids": [variant_evidence_id],
                        "confidence": 0.95, "reason": "The clause is classified as external.",
                    }],
                    "clause_reviews": [{
                        "clause_id": variant_clause_id,
                        "classification": "external_compliance",
                        "normative_basis": "external_duty",
                        "reason": "The physical signature or binding is external.",
                        "obligations": [{
                            "id": "physical_action", "status": "unverifiable",
                            "reason": "The physical action is outside DOCX generation.",
                        }],
                    }],
                    "unsupported_items": [], "reported_conflicts": [],
                }
                variant_response = bridge.normalize_native_response(
                    variant_response, variant_chunk["response_schema"],
                )
                variant_errors = bridge.validate_host_agent_response(
                    variant_response, variant_chunk,
                )
                variant_records = bridge.contract_error_records(
                    variant_errors, response=variant_response, chunk=variant_chunk,
                )
                self.assertIn(
                    "non_requirement_classification_relation",
                    {item["code"] for item in variant_records},
                )
                self.assertIsNone(bridge._project_external_action_requirements(
                    variant_response, variant_records, variant_chunk,
                )[0])

    def test_external_pending_projection_refuses_mixed_edges_and_existing_identity(self) -> None:
        response = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "body_text", "properties": {"text": "外部事项"},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "verification": {"mode": "external", "checks": ["人工核验"]},
            }],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "external_compliance"},
                {"clause_id": "C2", "classification": "executable"},
            ],
        }
        chunk = {"clauses": [
            {"id": "C1", "evidence_ids": ["E1"]},
            {"id": "C2", "evidence_ids": ["E2"]},
        ], "evidence_context": {
            "E1": {"id": "E1", "text": "外部事项"},
            "E2": {"id": "E2", "text": "正文事项"},
        }}
        mixed = copy.deepcopy(response)
        mixed["requirements"][0]["clause_ids"] = ["C1", "C2"]
        records = [{"code": "non_requirement_classification_relation",
                    "json_pointer": "$.requirements[0]", "requirement_index": 0,
                    "response_sha256": bridge._response_sha256(mixed)}]
        self.assertIsNone(bridge._project_external_action_requirements(mixed, records, chunk)[0])

        identified = copy.deepcopy(response)
        identified["requirements"][0]["existing_requirement_id"] = "R1"
        records[0]["response_sha256"] = bridge._response_sha256(identified)
        self.assertIsNone(bridge._project_external_action_requirements(identified, records, chunk)[0])

    def test_mechanical_repair_removes_batch_informational_requirements_by_mask(self) -> None:
        response = {
            "requirements": [
                {"role": "body_text", "properties": {"text": "一"}, "clause_ids": ["C1"]},
                {"role": "body_text", "properties": {"text": "二"}, "clause_ids": ["C2"]},
                {"role": "body_text", "properties": {"text": "三"}, "clause_ids": ["C3"]},
                {"role": "body_text", "properties": {"text": "四"}, "clause_ids": ["C4"]},
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "informational"},
                {"clause_id": "C2", "classification": "informational"},
                {"clause_id": "C3", "classification": "executable"},
                {"clause_id": "C4", "classification": "informational"},
            ],
        }
        parent_sha256 = bridge._response_sha256(response)
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [
                {
                    "code": "informational_requirement_forbidden",
                    "json_pointer": f"$.requirements[{index}]",
                    "requirement_index": index,
                    "relation_category": "informational_only",
                    "mechanically_removable": True,
                    "response_sha256": parent_sha256,
                }
                for index in (0, 1, 3)
            ],
        )
        self.assertEqual(
            [item["clause_ids"] for item in repaired["requirements"]], [["C3"]]
        )
        self.assertEqual(repairs[-1]["removed_requirement_indexes"], [0, 1, 3])
        self.assertEqual(repairs[-1]["projection"], "ordered_mask")
        self.assertEqual(
            repairs[-1]["original_index_to_repaired_index"],
            {"0": None, "1": None, "2": 0, "3": None},
        )
        self.assertEqual(
            repairs[-1]["removed_requirement_fingerprints"],
            [
                bridge._response_sha256(response["requirements"][0]),
                bridge._response_sha256(response["requirements"][1]),
                bridge._response_sha256(response["requirements"][3]),
            ],
        )
        self.assertEqual(
            repairs[-1]["source_response_sha256"],
            bridge._response_sha256(response),
        )
        self.assertEqual(
            repairs[-1]["repaired_response_sha256"],
            bridge._response_sha256(repaired),
        )
        self.assertEqual(len(response["requirements"]), 4)

    def test_mechanical_repair_removes_only_unbound_empty_requirement_placeholder(self) -> None:
        response = {
            "requirements": [
                {
                    "role": "declarations",
                    "properties": {"items": [{"heading": "固定声明"}]},
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                    "reason": "固定声明文本有来源依据。",
                    "confidence": 0.9,
                },
                {
                    "role": "body_text",
                    "properties": {
                        "style": None,
                        "top_border_pt": None,
                        "header_border_pt": None,
                        "bottom_border_pt": None,
                        "remove_vertical_borders": None,
                        "repeat_header_row": None,
                        "allow_row_split": None,
                        "keep_with_caption": None,
                        "border_widths_pt": None,
                        "continuation": None,
                    },
                    "clause_ids": [],
                    "evidence_ids": [],
                    "confidence": 0,
                    "reason": "",
                },
            ],
            "clause_reviews": [{"clause_id": "C1", "classification": "covered"}],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "contract_validation_error",
                "json_pointer": "$.requirements[1].reason",
                "raw_error": "$.requirements[1].reason: is shorter than 1 characters",
            }, {
                "code": "empty_requirement_properties",
                "json_pointer": "$.requirements[1].properties",
                "raw_error": "$.requirements[1].properties: must_include_semantic_payload",
            }, {
                "code": "schema_contract_violation",
                "json_pointer": "$.requirements[1].clause_ids",
                "raw_error": "$.requirements[1].clause_ids: must_be_non_empty",
            }, {
                "code": "schema_contract_violation",
                "json_pointer": "$.requirements[1].evidence_ids",
                "raw_error": "$.requirements[1].evidence_ids: must_be_non_empty",
            }, {
                "code": "requirement_relation_mismatch",
                "json_pointer": "$.requirements[1]",
                "raw_error": "requirements_not_referenced_by_clause_review:1",
                "relation_category": "missing_clause_relation",
            }],
        )
        self.assertIsNotNone(repaired)
        self.assertEqual(len(repaired["requirements"]), 1)
        self.assertEqual(repaired["requirements"][0]["role"], "declarations")
        self.assertEqual(
            repairs[-1]["rule_id"],
            "remove_unbound_empty_requirement_placeholder_v1",
        )
        self.assertEqual(repairs[-1]["removed_requirement_indexes"], [1])

        for field, value in (("clause_ids", ["C2"]), ("evidence_ids", ["E2"]), ("reason", "有语义")):
            response["requirements"][1][field] = value
            rejected, _ = bridge._apply_safe_mechanical_repairs(
                response,
                [{
                    "code": "empty_requirement_properties",
                    "json_pointer": "$.requirements[1].properties",
                    "raw_error": "must_include_semantic_payload",
                }],
            )
            self.assertIsNone(rejected)
            response["requirements"][1][field] = [] if field != "reason" else ""

    def test_mechanical_repair_removes_schema_invalid_unbound_placeholder_only(self) -> None:
        placeholder = {
            "existing_requirement_id": None, "field_key": None,
            "clause_ids": [], "source_fragment_clause_ids": None,
            "evidence_ids": [], "confidence": 0.0, "reason": "",
            "applicability": None, "input_prerequisites": None,
            "verification": None, "role": "body_text",
            "properties": {"font": None, "text": None},
        }
        response = {
            "requirements": [
                {"role": "body_text", "properties": {"text": "原文"},
                 "clause_ids": ["C1"], "evidence_ids": ["E1"]},
                placeholder,
            ],
            "clause_reviews": [{"clause_id": "C1", "classification": "covered"}],
        }

        def records(candidate: dict) -> list[dict]:
            sha = bridge._response_sha256(candidate)
            return [{
                "code": code, "json_pointer": pointer,
                "raw_error": raw, "response_sha256": sha,
                **extra,
            } for code, pointer, raw, extra in [
                ("contract_validation_error", "$.requirements[1]",
                 "$.requirements[1]: must match at least one schema in anyOf", {}),
                ("contract_validation_error", "$.requirements[1].reason",
                 "$.requirements[1].reason: is shorter than 1 characters", {}),
                ("empty_requirement_properties", "$.requirements[1].properties",
                 "$.requirements[1].properties: must_include_semantic_payload", {}),
                ("schema_contract_violation", "$.requirements[1].clause_ids",
                 "$.requirements[1].clause_ids: must_be_non_empty", {}),
                ("schema_contract_violation", "$.requirements[1].evidence_ids",
                 "$.requirements[1].evidence_ids: must_be_non_empty", {}),
                ("requirement_relation_mismatch", "$.requirements[1]",
                 "requirements_not_referenced_by_clause_review:1",
                 {"relation_category": "missing_clause_relation"}),
            ]]

        repaired, audit = bridge._apply_safe_mechanical_repairs(response, records(response))
        self.assertIsNotNone(repaired)
        self.assertEqual(repaired["requirements"], response["requirements"][:1])
        self.assertEqual(response["requirements"][1], placeholder)
        self.assertEqual(audit[-1]["removed_requirement_indexes"], [1])
        self.assertEqual(audit[-1]["source_response_sha256"], bridge._response_sha256(response))

        for field, value in (
            ("source_fragment_clause_ids", ["C1"]),
            ("reason", "有待确认的语义"),
            ("unexpected_semantic_field", "不可丢弃"),
        ):
            changed = copy.deepcopy(response)
            changed["requirements"][1][field] = value
            self.assertIsNone(bridge._apply_safe_mechanical_repairs(changed, records(changed))[0])

        wrong_hash = records(response)
        wrong_hash[0]["response_sha256"] = "0" * 64
        self.assertIsNone(bridge._apply_safe_mechanical_repairs(response, wrong_hash)[0])
        unrelated_error = records(response) + [{
            "code": "contract_validation_error",
            "json_pointer": "$.requirements[0]",
            "raw_error": "$.requirements[0]: must match at least one schema in anyOf",
        }]
        self.assertIsNone(bridge._apply_safe_mechanical_repairs_one_rule(
            response, unrelated_error,
        )[0])

    def test_unbound_cover_schema_shell_is_pruned_with_independent_errors(self) -> None:
        chunk = {"requirement_contract": {
            "role_properties_schema": {"cover": {"$ref": "#/$defs/coverSpec"}},
            "$defs": {"coverSpec": {
                "type": "object", "additionalProperties": False,
                "required": ["institution", "fields"],
                "properties": {
                    "institution": {"type": "string"}, "fields": {"type": "array"},
                    "non_public_administration": {"type": "object"},
                    "before_role": {"enum": ["document_start"]},
                    "missing_value_policy": {"const": "placeholder"},
                    "missing_value_placeholder": {"type": "string"},
                    "layout_id": {"enum": ["linear"]},
                },
            }},
        }}
        shell = {
            "role": "cover", "properties": {
                "institution": "", "fields": [], "non_public_administration": None,
                "before_role": None, "missing_value_policy": None,
                "missing_value_placeholder": None, "layout_id": None,
            },
            "clause_ids": [], "evidence_ids": [], "confidence": 0.0,
            "reason": "", "verification": None,
        }
        response = {
            "requirements": [
                {"role": "cover", "properties": {"institution": "", "fields": [{"order": 0}]},
                 "clause_ids": ["C1"], "evidence_ids": ["E1"]},
                shell,
            ],
            "clause_reviews": [{"clause_id": "C1", "classification": "covered"}],
        }
        sha = bridge._response_sha256(response)
        records = [{
            "code": "contract_validation_error", "json_pointer": "$.requirements[0]",
            "response_sha256": sha, "raw_error": "independent cover error",
        }, {
            "code": "cover_institution_placeholder",
            "json_pointer": "$.requirements[1].properties.institution",
            "response_sha256": sha, "raw_error": "empty institution",
        }, {
            "code": "requirement_relation_mismatch",
            "json_pointer": "$.requirements[1]", "requirement_index": 1,
            "raw_error": "requirements_not_referenced_by_clause_review:1",
            "response_sha256": sha, "relation_category": "missing_clause_relation",
            "mechanically_removable": True,
            "mechanical_removal_basis": "no_clause_or_evidence_binding",
        }]
        pruned, audit = bridge._prune_unbound_empty_schema_shells(response, records, chunk)
        self.assertEqual(pruned["requirements"], response["requirements"][:1])
        self.assertEqual(audit[0]["rule_id"], "remove_source_unbound_empty_schema_shell_v1")
        self.assertEqual(response["requirements"][1], shell)

        for change in (
            {"properties": {**shell["properties"], "approved": False}},
            {"properties": {**shell["properties"], "limit": 0}},
            {"properties": {"fields": []}},
            {"properties": {**shell["properties"], "before_role": ""}},
            {"clause_ids": ["C1"]}, {"reason": "human decision"},
        ):
            changed = copy.deepcopy(response)
            changed["requirements"][1].update(change)
            bound_records = copy.deepcopy(records)
            bound_sha = bridge._response_sha256(changed)
            for record in bound_records:
                record["response_sha256"] = bound_sha
            self.assertIsNone(bridge._prune_unbound_empty_schema_shells(
                changed, bound_records, chunk,
            )[0])
        tampered = copy.deepcopy(records)
        tampered[-1]["response_sha256"] = "0" * 64
        self.assertIsNone(bridge._prune_unbound_empty_schema_shells(response, tampered, chunk)[0])
        mixed_external = copy.deepcopy(records) + [{
            "code": "non_requirement_classification_relation",
            "json_pointer": "$.requirements[0]", "response_sha256": sha,
        }]
        self.assertIsNone(bridge._prune_unbound_empty_schema_shells(
            response, mixed_external, chunk,
        )[0])

    def test_exact_duplicate_requirements_only_collapse_identical_v3_items(self) -> None:
        item = {
            "role": "cover", "properties": {"fields": [{"id": "title_zh"}]},
            "clause_ids": ["C1"], "evidence_ids": ["E1"],
            "confidence": 0.9, "reason": "Same source-bound cover rule.",
        }
        response = {
            "contract_version": "3.0", "requirements": [copy.deepcopy(item), copy.deepcopy(item)],
            "clause_reviews": [{"clause_id": "C1", "classification": "covered"}],
        }
        original = copy.deepcopy(response)
        projected, audit = bridge._project_exact_duplicate_requirements(response)
        self.assertEqual(projected["requirements"], [item])
        self.assertEqual(audit["removed"][0]["removed_index"], 1)
        self.assertFalse(audit["accepted"])
        self.assertEqual(response, original)

        for change in (
            {"clause_ids": ["C2"]},
            {"evidence_ids": ["E2"]},
            {"properties": {"fields": [{"id": "security_marking"}]}},
        ):
            distinct = copy.deepcopy(response)
            distinct["requirements"][1].update(change)
            unchanged, no_audit = bridge._project_exact_duplicate_requirements(distinct)
            self.assertEqual(unchanged, distinct)
            self.assertIsNone(no_audit)
        legacy = copy.deepcopy(response)
        legacy["contract_version"] = "2.1"
        self.assertIsNone(bridge._project_exact_duplicate_requirements(legacy)[1])
        indexed = copy.deepcopy(response)
        indexed["clause_reviews"][0]["requirement_indexes"] = [0, 1]
        self.assertIsNone(bridge._project_exact_duplicate_requirements(indexed)[1])

    def test_unbound_semantic_orphan_is_audited_then_full_contract_is_revalidated(self) -> None:
        source = "北京体育大学学位评定委员会办公室盖章(有效)"
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(
                Path(td) / "review", contract_version="3.0", source=source,
            )
            raw = {
                "contract_version": "3.0",
                "requirements": [{
                    "existing_requirement_id": None,
                    "field_key": None,
                    "clause_ids": [],
                    "evidence_ids": [],
                    "confidence": 0.0,
                    "reason": "No additional executable requirement is needed for the external stamping duty.",
                    "applicability": None,
                    "input_prerequisites": None,
                    "verification": None,
                    "role": "body_text",
                    "properties": {
                        "font": None,
                        "paragraph": None,
                        "numbering": None,
                        "position": None,
                        "prefix": None,
                        "separator": None,
                        "style_hint": "external compliance only",
                        "text": source,
                        "header_content": None,
                        "bottom_border": None,
                    },
                }],
                "clause_reviews": [{
                    "clause_id": "C1",
                    "classification": "external_compliance",
                    "reason": "An actual office stamp is an external duty.",
                    "obligations": [{
                        "id": "office_stamp", "status": "unverifiable",
                        "reason": "A real office stamp remains pending outside DOCX generation.",
                    }],
                    "normative_basis": "external_duty",
                }],
                "unsupported_items": [],
                "reported_conflicts": [],
            }
            accepted, audit = bridge.prepare_native_response_candidate(raw, chunk)
            self.assertEqual(accepted["requirements"], [])
            self.assertEqual(accepted["clause_reviews"][0]["classification"], "external_compliance")
            self.assertEqual(bridge.validate_host_agent_response(accepted, chunk), [])
            repair = next(
                item for item in audit["mechanical_repairs"]
                if item.get("rule_id") == "remove_unbound_non_placeholder_requirement_v1"
                and "removed_requirement" in item
            )
            self.assertEqual(repair["removed_requirement"]["properties"]["text"], source)
            normalized_parent = bridge.normalize_native_response(raw, chunk["response_schema"])
            self.assertEqual(
                repair["removed_requirement_sha256"],
                bridge._response_sha256(normalized_parent["requirements"][0]),
            )

            bound_orphan = copy.deepcopy(raw)
            bound_orphan["requirements"][0]["evidence_ids"] = ["E1"]
            normalized = bridge.normalize_native_response(
                bound_orphan, chunk["response_schema"],
            )
            bound_errors = bridge.validate_host_agent_response(normalized, chunk)
            bound_records = bridge.contract_error_records(
                bound_errors, response=normalized, chunk=chunk,
            )
            bound_relation = next(
                item for item in bound_records
                if item.get("code") == "requirement_relation_mismatch"
            )
            self.assertFalse(bound_relation["mechanically_removable"])
            self.assertIsNone(bridge._apply_safe_mechanical_repairs(
                normalized, bound_records, chunk=chunk,
            )[0])

            executable = copy.deepcopy(raw)
            executable["clause_reviews"][0]["classification"] = "executable"
            executable["clause_reviews"][0]["normative_basis"] = "explicit_normative_text"
            executable["clause_reviews"][0]["obligations"] = [{
                "id": "office_stamp", "status": "covered",
                "reason": "The response claims to cover this source duty.",
            }]
            with self.assertRaises(ValueError) as blocked:
                bridge.prepare_native_response_candidate(executable, chunk)
            self.assertIn(
                "executable_review_requires_derived_requirement",
                str(getattr(blocked.exception, "error_records", [])),
            )

    def test_unbound_cover_shell_with_fields_error_is_removed_before_retry(self) -> None:
        source = "北京体育大学学位评定委员会办公室盖章(有效)"
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(
                Path(td) / "review", contract_version="3.0", source=source,
            )
            raw = {
                "contract_version": "3.0",
                "provenance": chunk["provenance"],
                "requirements": [{
                    "existing_requirement_id": None,
                    "field_key": None,
                    "clause_ids": [],
                    "evidence_ids": [],
                    "confidence": 0.99,
                    "reason": "占位",
                    "role": "cover",
                    "properties": {
                        "institution": "——",
                        "fields": [],
                        "non_public_administration": None,
                        "missing_value_policy": "placeholder",
                        "missing_value_placeholder": "——",
                        "layout_id": "linear",
                    },
                }],
                "clause_reviews": [{
                    "clause_id": "C1",
                    "classification": "external_compliance",
                    "reason": "The office stamp is an external duty.",
                    "normative_basis": "external_duty",
                    "obligations": [{
                        "id": "office_stamp", "status": "unverifiable",
                        "reason": "A real office stamp remains pending.",
                    }],
                }],
                "unsupported_items": [],
                "reported_conflicts": [],
            }
            normalized = bridge.normalize_native_response(raw, chunk["response_schema"])
            errors = bridge.validate_host_agent_response(normalized, chunk)
            records = bridge.contract_error_records(
                errors, response=normalized, chunk=chunk,
            )
            self.assertEqual(
                {item["code"] for item in records},
                {
                    "cover_binding_violation", "schema_contract_violation",
                    "requirement_relation_mismatch",
                },
                errors,
            )
            candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
            self.assertEqual(candidate["requirements"], [])
            self.assertEqual(candidate["clause_reviews"], normalized["clause_reviews"])
            self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
            self.assertTrue(any(
                item.get("rule_id") == "remove_unbound_non_placeholder_requirement_v1"
                and item.get("removed_requirement", {}).get("role") == "cover"
                for item in audit["mechanical_repairs"]
            ))

            wrong_error = copy.deepcopy(records)
            next(item for item in wrong_error if item["code"] == "cover_binding_violation")[
                "raw_error"
            ] = "fabricated cover error"
            self.assertIsNone(bridge._apply_safe_mechanical_repairs(
                normalized, wrong_error, chunk=chunk,
            )[0])
            stale = copy.deepcopy(records)
            next(item for item in stale if item["code"] == "requirement_relation_mismatch")[
                "response_sha256"
            ] = "0" * 64
            self.assertIsNone(bridge._apply_safe_mechanical_repairs(
                normalized, stale, chunk=chunk,
            )[0])

            bound = copy.deepcopy(raw)
            bound["requirements"][0]["clause_ids"] = ["C1"]
            bound["requirements"][0]["evidence_ids"] = ["E1"]
            with self.assertRaises(ValueError):
                bridge.prepare_native_response_candidate(bound, chunk)

    def test_public_cover_blank_does_not_need_empty_conditional_requirement(self) -> None:
        source = "未经批准的均为公开学位论文（公开的学位论文本项为空白）"
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(
                Path(td) / "review", contract_version="3.0", source=source,
            )
            raw = {
                "contract_version": "3.0",
                "requirements": [
                    {
                        "role": "cover",
                        "properties": {
                            "institution": "——", "fields": [],
                            "non_public_administration": {
                                "applicability": {
                                    "status": "conditional",
                                    "conditions": [{
                                        "fact": "thesis_profile.security_level",
                                        "operator": "in",
                                        "value": ["restricted", "classified"],
                                    }],
                                },
                                "fields": [{
                                    "id": "security_marking", "label": "申请密级",
                                    "value_from": "thesis_profile.cover_metadata.security_marking",
                                    "display_policy": "required", "order": 1,
                                }],
                                "public_policy": "blank",
                                "source_region": "非公开学位论文标注说明",
                            },
                            "before_role": "document_start",
                            "missing_value_policy": "placeholder",
                            "missing_value_placeholder": "——",
                            "layout_id": "linear",
                        },
                        "clause_ids": ["C1"], "evidence_ids": ["E1"],
                        "reason": "The cover leaves the administrative region blank for public theses.",
                        "confidence": 0.9,
                        "verification": {
                            "mode": "static_docx", "checks": ["Check the public cover blank policy."],
                        },
                    },
                    {
                        "role": "conditional_constraints", "properties": {
                            "abstract_zh_max_chars_by_degree": None,
                            "require_zh_abstract_for_english_thesis": None,
                            "require_zh_keywords_for_english_thesis": None,
                        },
                        "clause_ids": ["C1"], "evidence_ids": ["E1"],
                        "reason": "The public-cover blank condition is already in the cover.",
                        "confidence": 0.9,
                        "applicability": {"status": "always"},
                        "verification": {
                            "mode": "static_docx", "checks": ["Check the public blank policy."],
                        },
                    },
                ],
                "clause_reviews": [{
                    "clause_id": "C1", "classification": "executable",
                    "reason": "The public administrative item must be blank.",
                    "normative_basis": "explicit_normative_text",
                    "obligations": [{
                        "id": "public_blank", "status": "covered",
                        "reason": "The cover has a blank public policy.",
                    }],
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }
            normalized = bridge.normalize_native_response(raw, chunk["response_schema"])
            errors = bridge.validate_host_agent_response(normalized, chunk)
            self.assertEqual(
                errors,
                [
                    "$.requirements[1].properties: must_include_semantic_payload",
                    "$.clause_reviews[0]: partial_clause_coverage:"
                    "cover.publication_default.unapproved_is_public",
                ],
            )
            accepted, audit = bridge.prepare_native_response_candidate(raw, chunk)
            self.assertEqual(len(accepted["requirements"]), 1)
            expected_properties = copy.deepcopy(normalized["requirements"][0]["properties"])
            expected_properties["non_public_administration"]["publication_default_policy"] = (
                "unapproved_is_public"
            )
            self.assertEqual(
                accepted["requirements"][0]["properties"],
                expected_properties,
            )
            self.assertEqual(
                accepted["requirements"][0]["verification"]["checks"][0],
                "Check the public cover blank policy.",
            )
            self.assertEqual(
                accepted["requirements"][0]["verification"]["checks"][-1],
                "Check the public blank policy.",
            )
            self.assertIn(
                "cover_non_public_administration",
                accepted["requirements"][0]["verification"]["checker_ids"],
            )
            self.assertEqual(accepted["clause_reviews"], normalized["clause_reviews"])
            self.assertEqual(bridge.validate_host_agent_response(accepted, chunk), [])
            repair = next(item for item in audit["mechanical_repairs"] if item.get(
                "rule_id") == "remove_redundant_public_cover_condition_v1")
            self.assertEqual(repair["source_clause_ids"], ["C1"])
            self.assertEqual(repair["source_evidence_ids"], ["E1"])
            self.assertEqual(repair["run_id"], chunk["provenance"]["run_id"])
            self.assertEqual(repair["case_id"], chunk["case_id"])
            self.assertEqual(repair["source_sha256"], chunk["provenance"]["source_sha256"])
            self.assertEqual(repair["clause_sha256"], chunk["provenance"]["clause_sha256"])
            self.assertEqual(repair["evidence_sha256"], chunk["provenance"]["evidence_sha256"])
            self.assertEqual(repair["removed_requirement"]["role"], "conditional_constraints")
            self.assertEqual(repair["transferred_verification_checks"], ["Check the public blank policy."])
            review_request = bridge.build_obligation_coverage_request(
                accepted, chunk, run_id=chunk["provenance"]["run_id"], chunk_index=1,
            )
            check = next(item for item in review_request["checks"] if item["check_id"] == "C1")
            self.assertEqual(len(check["review_context"]["linked_requirements"]), 1)
            self.assertEqual(check["review_context"]["linked_requirements"][0]["role"], "cover")
            self.assertEqual(check["review_context"]["primary_obligations"][0]["status"], "covered")
            self.assertTrue(any(
                "conditional_constraints" in instruction and "public-blank" in instruction
                for instruction in chunk["instructions"]
            ))
            self.assertTrue(any(
                "conditional_constraints" in rule and "public-blank" in rule
                for rule in bridge._BASE_CONTRACT_REPAIR_RULES
            ))

            baseline_hash = bridge._response_sha256(normalized)
            record = bridge.contract_error_records(
                errors, response=normalized, chunk=chunk,
            )[0]
            def projected(candidate: dict, *, packet: dict | None = None,
                          original_record: dict | None = None) -> dict | None:
                candidate_hash = bridge._response_sha256(candidate)
                current_record = copy.deepcopy(original_record or record)
                if original_record is None:
                    current_record["response_sha256"] = candidate_hash
                return bridge._project_redundant_public_cover_condition(
                    candidate, [current_record], packet or chunk, candidate_hash,
                )[0]

            stale = copy.deepcopy(record)
            stale["response_sha256"] = "0" * 64
            self.assertIsNone(projected(normalized, original_record=stale))
            self.assertEqual(record["response_sha256"], baseline_hash)
            for label, mutate in (
                ("source-bound other obligation", lambda item: item["clause_reviews"][0][
                    "obligations"].append({"id": "other", "status": "covered", "reason": "Other"})),
                ("wrong cover policy", lambda item: item["requirements"][0][
                    "properties"]["non_public_administration"].update({"public_policy": "show"})),
                ("public included in cover applicability", lambda item: item[
                    "requirements"][0]["properties"]["non_public_administration"][
                    "applicability"]["conditions"][0]["value"].append("public")),
                ("cover missing source edge", lambda item: item["requirements"][0].update(
                    {"clause_ids": []})),
                ("nonempty independent property", lambda item: item["requirements"][1].update(
                    {"properties": {"require_zh_abstract_for_english_thesis": True}})),
                ("existing requirement id", lambda item: item["requirements"][1].update(
                    {"existing_requirement_id": "R1"})),
                ("independent verification check", lambda item: item[
                    "requirements"][1]["verification"]["checks"].append(
                    "Verify a separate declaration signature.")),
                ("different verification mode", lambda item: item[
                    "requirements"][1]["verification"].update({"mode": "word_render"})),
                ("different run provenance", lambda item: item.update({
                    "provenance": {"run_id": "stale-run"},
                })),
                ("two competing cover requirements", lambda item: item[
                    "requirements"].append(copy.deepcopy(item["requirements"][0]))),
                ("unresolved review", lambda item: item["clause_reviews"][0].update(
                    {"classification": "unresolved"})),
            ):
                altered = copy.deepcopy(normalized)
                mutate(altered)
                self.assertIsNone(projected(altered), label)

            other_source = "公开的学位论文本项为空白；摘要还必须控制字数"
            other_chunk = copy.deepcopy(chunk)
            other_chunk["evidence_context"]["E1"]["text"] = other_source
            other_chunk["clauses"][0]["text"] = other_source
            self._bind_test_source_spans(
                other_chunk["clauses"],
                {"evidence": [other_chunk["evidence_context"]["E1"]]},
            )
            self.assertIsNone(projected(normalized, packet=other_chunk))

            for key, value in (("run_id", ""), ("source_sha256", "stale")):
                invalid_identity_chunk = copy.deepcopy(chunk)
                invalid_identity_chunk["provenance"][key] = value
                self.assertIsNone(projected(normalized, packet=invalid_identity_chunk), key)

            external_cover = copy.deepcopy(normalized)
            external_cover["requirements"][0]["verification"]["mode"] = "external"
            external_projected = projected(external_cover)
            self.assertIsNotNone(external_projected)
            self.assertEqual(
                external_projected["requirements"][0]["verification"]["mode"],
                "external",
            )
            self.assertIn(
                "Check the public blank policy.",
                external_projected["requirements"][0]["verification"]["checks"],
            )

            reversed_raw = copy.deepcopy(raw)
            reversed_raw["requirements"].reverse()
            reversed_candidate, reversed_audit = bridge.prepare_native_response_candidate(
                reversed_raw, chunk,
            )
            self.assertEqual(len(reversed_candidate["requirements"]), 1)
            self.assertEqual(bridge.validate_host_agent_response(reversed_candidate, chunk), [])
            reversed_repair = next(item for item in reversed_audit["mechanical_repairs"] if item.get(
                "rule_id") == "remove_redundant_public_cover_condition_v1")
            self.assertEqual(reversed_repair["cover_requirement_index_after"], 0)

    def test_mechanical_repair_never_removes_mixed_or_noninformational_relations(self) -> None:
        cases = [
            (
                [{"role": "body_text", "properties": {"text": "混合"}, "clause_ids": ["C1", "C2"]}],
                [
                    {"clause_id": "C1", "classification": "informational"},
                    {"clause_id": "C2", "classification": "executable"},
                ],
            ),
            (
                [{"role": "body_text", "properties": {"text": "外部"}, "clause_ids": ["C1"]}],
                [{"clause_id": "C1", "classification": "external_compliance"}],
            ),
            (
                [{"role": "body_text", "properties": {"text": "缺审查"}, "clause_ids": ["C1"]}],
                [],
            ),
        ]
        for requirements, reviews in cases:
            response = {"requirements": requirements, "clause_reviews": reviews}
            repaired, repairs = bridge._apply_safe_mechanical_repairs(
                response,
                [{
                    "code": "informational_requirement_forbidden",
                    "requirement_index": 0,
                }],
            )
            self.assertIsNone(repaired)
            self.assertEqual(repairs, [])

    def test_mechanical_repair_normalizes_explicitly_optional_continuation_caption(self) -> None:
        response = {
            "requirements": [{
                "role": "table",
                "properties": {"continuation": {
                    "caption_suffix": "(续)",
                    "repeat_header_row": True,
                    "caption_required_on_continuation": True,
                    "verification": "word_render",
                }},
                "clause_ids": ["C00243"],
                "evidence_ids": ["E00174"],
            }],
            "clause_reviews": [{"clause_id": "C00243", "classification": "executable"}],
        }
        chunk = {"clauses": [{
            "id": "C00243",
            "text": "表序后跟表题(可省略)和“(续)”，居中置于表上方，续表均应重复表头",
        }]}
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "partial_clause_coverage",
                "json_pointer": "$.clause_reviews[0]",
                "raw_error": "$.clause_reviews[0]: partial_clause_coverage:table.continuation.caption_optional",
            }],
            chunk=chunk,
        )
        self.assertIsNotNone(repaired)
        self.assertFalse(
            repaired["requirements"][0]["properties"]["continuation"]["caption_required_on_continuation"]
        )
        self.assertEqual(repairs[0]["rule_id"], "compile_continuation_caption_requirement_v1")
        self.assertTrue(response["requirements"][0]["properties"]["continuation"]["caption_required_on_continuation"])

    def test_bsu_chunk13_mixed_errors_replay_through_production_candidate_boundary(self) -> None:
        """Replay the captured BSU failure shape without a provider call.

        The real packet's two existing caption selectors are retained, while
        their native nullable properties are empty and the continuation
        requirement incorrectly says captions are required.  This exercises
        the same candidate preparation function used by run_host_agent_chunk.
        """
        clauses = [
            {
                "id": "C00243",
                "text": "表序后跟表题(可省略)和“(续)”，居中置于表上方，续表均应重复表头",
                "evidence_ids": ["E00174"], "source_kind": "paragraph",
                "location": {}, "part_index": 0,
            },
            {
                "id": "C00246",
                "text": "表题间空1个字距，居中置于表的上方",
                "evidence_ids": ["E00175"], "source_kind": "paragraph",
                "location": {}, "part_index": 0,
            },
        ]
        evidence = {"evidence": [
            {"id": "E00174", "text": "如某表需要转页接排时，在随后的各页上应重复表序。表序后跟表题(可省略)和“(续)”，居中置于表上方，续表均应重复表头。", "kind": "paragraph"},
            {"id": "E00175", "text": "说明3：表序与表题，表序即表的编号，由“表”和从“1”开始的阿拉伯数字组成；表题即表的名称，应简明，置于表序之后，表序和表题间空1个字距，居中置于表的上方。", "kind": "paragraph"},
        ]}
        clauses = self._bind_test_source_spans(clauses, evidence)
        existing = [
            {
                "id": "R00530", "role": "table_caption",
                "properties": {"position": "above", "paragraph": {"alignment": "center"}},
                "clause_ids": ["C00243"], "evidence_ids": ["E00174"],
                "source_text": clauses[0]["text"],
                "verification": {"mode": "static_docx", "checks": ["核对表题位置和对齐。"], "checker_ids": ["docx.property_receipts"]},
            },
            {
                "id": "R00531", "role": "table_caption",
                "properties": {"position": "above", "paragraph": {"alignment": "center"}},
                "clause_ids": ["C00246"], "evidence_ids": ["E00175"],
                "source_text": clauses[1]["text"],
                "verification": {"mode": "static_docx", "checks": ["核对表题位置和对齐。"], "checker_ids": ["docx.property_receipts"]},
            },
        ]
        request = engine.build_llm_request(
            [], clauses, evidence, {"requirements": existing}, "full",
            contract_version="3.0",
        )
        request = attach_request_provenance(
            request, source_sha256="b" * 64, evidence_doc=evidence,
            clauses=clauses, run_id="bsu-chunk13-offline-regression",
        )
        chunk = engine._build_host_review_chunks(
            request, clauses, evidence, "b" * 64, chunk_size=20,
        )[0]

        null_role_properties = {
            "font": None, "paragraph": None, "numbering": None,
            "position": None, "prefix": None, "separator": None,
            "style_hint": None, "text": None, "header_content": None,
            "bottom_border": None,
        }
        raw_response = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "table",
                    "properties": {
                        "style": None, "top_border_pt": None,
                        "header_border_pt": None, "bottom_border_pt": None,
                        "remove_vertical_borders": None, "repeat_header_row": None,
                        "allow_row_split": None, "keep_with_caption": None,
                        "border_widths_pt": None,
                        "continuation": {
                            "caption_suffix": "(续)",
                            "repeat_header_row": True,
                            "caption_required_on_continuation": True,
                            "verification": "word_render",
                        },
                    },
                    "clause_ids": ["C00243"], "evidence_ids": ["E00174"],
                    "confidence": 0.98,
                    "reason": "The continuation table properties are supported by the cited source.",
                    "verification": {
                        "mode": "word_render", "checks": ["核对续表属性。"],
                        "checker_ids": ["docx.property_receipts", "docx.word_render"],
                    },
                },
                {
                    "existing_requirement_id": "R00530", "role": "table_caption",
                    "properties": copy.deepcopy(null_role_properties),
                    "clause_ids": ["C00243"], "evidence_ids": ["E00174"],
                    "confidence": 0.98,
                    "reason": "The selected existing caption requirement is supported by this occurrence.",
                    "verification": copy.deepcopy(existing[0]["verification"]),
                },
                {
                    "existing_requirement_id": "R00531", "role": "table_caption",
                    "properties": copy.deepcopy(null_role_properties),
                    "clause_ids": ["C00246"], "evidence_ids": ["E00175"],
                    "confidence": 0.98,
                    "reason": "The selected existing caption requirement is supported by this occurrence.",
                    "verification": copy.deepcopy(existing[1]["verification"]),
                },
            ],
            "clause_reviews": [
                {
                    "clause_id": "C00243", "classification": "executable",
                    "normative_basis": "explicit_normative_text",
                    "reason": "The explicit continuation-table rules are represented by linked requirements.",
                    "obligations": [
                        {"id": "suffix", "status": "covered", "reason": "The continuation suffix is represented."},
                        {"id": "header", "status": "covered", "reason": "The repeated header is represented."},
                        {"id": "caption", "status": "covered", "reason": "The caption position is represented."},
                    ],
                },
                {
                    "clause_id": "C00246", "classification": "executable",
                    "normative_basis": "explicit_normative_text",
                    "reason": "The explicit caption placement rule is represented by a linked requirement.",
                    "obligations": [
                        {"id": "caption-placement", "status": "covered", "reason": "The caption position is represented."},
                    ],
                },
            ],
            "unsupported_items": [], "reported_conflicts": [],
        }
        raw_before = copy.deepcopy(raw_response)
        normalized = bridge.normalize_native_response(raw_response, chunk["response_schema"])
        self.assertEqual(normalized["requirements"][1]["properties"], {})
        self.assertEqual(normalized["requirements"][2]["properties"], {})
        self.assertEqual(raw_response, raw_before, "normalization must not mutate captured raw response")

        candidate, audit = bridge.prepare_native_response_candidate(raw_response, chunk)
        self.assertEqual(
            [item["existing_requirement_id"] for item in audit["existing_requirement_payload_projections"]],
            ["R00530", "R00531"],
        )
        self.assertEqual(
            candidate["requirements"][1]["properties"], existing[0]["properties"],
        )
        self.assertEqual(
            candidate["requirements"][2]["properties"], existing[1]["properties"],
        )
        self.assertFalse(
            candidate["requirements"][0]["properties"]["continuation"][
                "caption_required_on_continuation"
            ]
        )
        self.assertEqual(audit["mechanical_repair_revalidation"]["status"], "passed")
        self.assertEqual(audit["mechanical_repairs"][0]["planner_id"], "composable_source_bound_patch_plan_v1")
        self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
        self.assertEqual(raw_response, raw_before, "candidate preparation must preserve the original response")

    def test_source_literal_whitespace_projection_uses_only_unique_current_source_text(self) -> None:
        cases = (
            ("学    号：", "学 号：", "cover_field_label"),
            ("硕  士 学 位 论 文", "硕士 学位 论 文", "thesis_type_zh"),
            ("年　月", "年月", "cover_field_label"),
        )
        for source, model_text, role in cases:
            with self.subTest(source=source, role=role), tempfile.TemporaryDirectory() as td:
                _, chunk = self._packet(
                    Path(td) / "requirements", contract_version="3.0", source=source,
                )
                review_dir = Path(td) / "requirements"
                validated_chunk_sha256 = self._validated_chunk_projection_sha256(
                    review_dir, chunk,
                )
                raw_response = {
                    "contract_version": "3.0",
                    "provenance": copy.deepcopy(chunk["provenance"]),
                    "requirements": [{
                        "role": role,
                        "properties": {"text": model_text},
                        "clause_ids": ["C1"],
                        "evidence_ids": ["E1"],
                        "confidence": 0.98,
                        "reason": "Preserve the exact source literal.",
                        "verification": {
                            "mode": "word_render",
                            "checks": ["Verify the literal after DOCX rendering."],
                            "checker_ids": ["docx.property_receipts", "docx.word_render"],
                        },
                    }],
                    "clause_reviews": [{
                        "clause_id": "C1",
                        "classification": "executable",
                        "normative_basis": "explicit_normative_text",
                        "reason": "The source label is represented exactly.",
                        "obligations": [{
                            "id": "literal",
                            "status": "covered",
                            "reason": "The exact source label is represented.",
                        }],
                    }],
                    "unsupported_items": [],
                    "reported_conflicts": [],
                }
                raw_before = copy.deepcopy(raw_response)
                candidate, audit = bridge.prepare_native_response_candidate(
                    raw_response, chunk,
                    source_projection_validation_sha256=validated_chunk_sha256,
                )

                self.assertEqual(candidate["requirements"][0]["properties"]["text"], source)
                self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
                self.assertEqual(raw_response, raw_before, "the raw model response must stay unchanged")
                projection = audit["source_literal_whitespace_projections"]
                self.assertEqual(len(projection), 1)
                self.assertEqual(projection[0]["rule_id"], "source_literal_whitespace_projection_v1")
                self.assertEqual(projection[0]["json_pointer"], "$.requirements[0].properties.text")
                self.assertEqual(projection[0]["run_id"], chunk["provenance"]["run_id"])
                self.assertEqual(projection[0]["clause_id"], "C1")
                self.assertEqual(projection[0]["evidence_id"], "E1")
                self.assertEqual(projection[0]["change_kind"], "unicode_whitespace_only")
                self.assertNotEqual(
                    projection[0]["response_before_sha256"],
                    projection[0]["response_after_sha256"],
                )

    def test_source_literal_projection_never_repairs_punctuation_or_content(self) -> None:
        rejected_texts = ("学 号;", "学X号：", "学 号：额外")
        for model_text in rejected_texts:
            with self.subTest(model_text=model_text), tempfile.TemporaryDirectory() as td:
                review_dir, chunk = self._packet(
                    Path(td) / "requirements", contract_version="3.0", source="学    号：",
                )
                validated_chunk_sha256 = self._validated_chunk_projection_sha256(
                    review_dir, chunk,
                )
                response = {
                    "contract_version": "3.0",
                    "provenance": copy.deepcopy(chunk["provenance"]),
                    "requirements": [{
                        "role": "cover_field_label", "properties": {"text": model_text},
                        "clause_ids": ["C1"], "evidence_ids": ["E1"],
                        "confidence": 0.98, "reason": "Preserve the exact source literal.",
                        "verification": {
                            "mode": "word_render", "checks": ["Verify the literal."],
                            "checker_ids": ["docx.property_receipts", "docx.word_render"],
                        },
                    }],
                    "clause_reviews": [{
                        "clause_id": "C1", "classification": "executable",
                        "normative_basis": "explicit_normative_text",
                        "reason": "The source label is represented.",
                        "obligations": [{
                            "id": "literal", "status": "covered",
                            "reason": "The source label is represented.",
                        }],
                    }],
                    "unsupported_items": [], "reported_conflicts": [],
                }
                with self.assertRaisesRegex(ValueError, "must_be_exact_substring_of_cited_source_span"):
                    bridge.prepare_native_response_candidate(
                        response, chunk,
                        source_projection_validation_sha256=validated_chunk_sha256,
                    )

    def test_source_literal_projection_requires_valid_cited_primary_span(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(
                Path(td) / "requirements", contract_version="3.0", source="学    号：",
            )
            validated_chunk_sha256 = self._validated_chunk_projection_sha256(
                review_dir, chunk,
            )
            response = {
                "contract_version": "3.0",
                "provenance": copy.deepcopy(chunk["provenance"]),
                "requirements": [{
                    "role": "cover_field_label", "properties": {"text": "学 号："},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                    "confidence": 0.98, "reason": "Preserve the exact source literal.",
                    "verification": {
                        "mode": "word_render", "checks": ["Verify the literal."],
                        "checker_ids": ["docx.property_receipts", "docx.word_render"],
                    },
                }],
                "clause_reviews": [{
                    "clause_id": "C1", "classification": "executable",
                    "normative_basis": "explicit_normative_text",
                    "reason": "The source label is represented.",
                    "obligations": [{
                        "id": "literal", "status": "covered",
                        "reason": "The source label is represented.",
                    }],
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }
            bad_evidence = copy.deepcopy(response)
            bad_evidence["requirements"][0]["evidence_ids"] = ["E-FOREIGN"]
            with self.assertRaises(ValueError):
                bridge.prepare_native_response_candidate(
                    bad_evidence, chunk,
                    source_projection_validation_sha256=validated_chunk_sha256,
                )

            bad_span_chunk = copy.deepcopy(chunk)
            bad_span_chunk["clauses"][0]["source_span"]["source_sha256"] = "f" * 64
            bad_span_chunk_sha256 = bridge._response_sha256(bad_span_chunk)
            with self.assertRaises(ValueError):
                bridge.prepare_native_response_candidate(
                    response, bad_span_chunk,
                    source_projection_validation_sha256=bad_span_chunk_sha256,
                )

            stale_run_chunk = copy.deepcopy(chunk)
            stale_run_chunk.pop("provenance")
            unchanged, projections = bridge._project_source_literal_whitespace_only(
                bridge.normalize_native_response(response, chunk["response_schema"]),
                stale_run_chunk,
                ["$.requirements[0].properties.text: must_be_exact_substring_of_cited_source_span"],
                source_projection_validation_sha256=bridge._response_sha256(stale_run_chunk),
            )
            self.assertEqual(unchanged["requirements"][0]["properties"]["text"], "学 号：")
            self.assertEqual(projections, [])

    def test_source_literal_projection_fails_closed_on_unverified_or_ambiguous_binding(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _, chunk = self._packet(
                Path(td) / "requirements", contract_version="3.0", source="字段：学    号：填写",
            )
            response = {
                "provenance": copy.deepcopy(chunk["provenance"]),
                "requirements": [{
                    "properties": {"text": "学 号"},
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                }],
            }
            exact_error = [
                "$.requirements[0].properties.text: "
                "must_be_exact_substring_of_cited_source_span"
            ]

            unchanged, projections = bridge._project_source_literal_whitespace_only(
                response,
                chunk,
                exact_error,
                source_projection_validation_sha256="0" * 64,
            )
            self.assertEqual(unchanged, response)
            self.assertEqual(projections, [])

            ambiguous_chunk = copy.deepcopy(chunk)
            source = ambiguous_chunk["evidence_context"]["E1"]["text"]
            start = source.index("学    号")
            end = start + len("学    号")
            span = ambiguous_chunk["clauses"][0]["source_span"]
            span.update({
                "start_offset": start,
                "end_offset": end,
                "text": source[start:end],
            })
            duplicate_clause = copy.deepcopy(ambiguous_chunk["clauses"][0])
            duplicate_clause["id"] = "C2"
            ambiguous_chunk["clauses"].append(duplicate_clause)
            ambiguous_response = copy.deepcopy(response)
            ambiguous_response["requirements"][0]["clause_ids"] = ["C1", "C2"]

            unchanged, projections = bridge._project_source_literal_whitespace_only(
                ambiguous_response,
                ambiguous_chunk,
                exact_error,
                source_projection_validation_sha256=bridge._response_sha256(ambiguous_chunk),
            )
            self.assertEqual(unchanged, ambiguous_response)
            self.assertEqual(projections, [])

    def test_mechanical_repair_compiles_explicit_empty_table_caption_properties(self) -> None:
        response = {
            "requirements": [
                {
                    "role": "table",
                    "properties": {"continuation": {"caption_suffix": "(续)"}},
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                },
                {
                    "role": "table_caption",
                    "properties": {},
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                },
                {
                    "role": "table_caption",
                    "properties": {},
                    "clause_ids": ["C2"],
                    "evidence_ids": ["E2"],
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
            ],
        }
        chunk = {
            "clauses": [
                {"id": "C1", "text": "表序后跟表题，居中置于表上方"},
                {"id": "C2", "text": "表题居中置于表的上方"},
            ],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [
                {
                    "code": "empty_requirement_properties",
                    "json_pointer": "$.requirements[1].properties",
                },
                {
                    "code": "empty_requirement_properties",
                    "json_pointer": "$.requirements[2].properties",
                },
                {
                    "code": "partial_clause_coverage",
                    "json_pointer": "$.clause_reviews[0]",
                    "raw_error": (
                        "$.clause_reviews[0]: partial_clause_coverage:"
                        "table_caption.position:above,table_caption.paragraph.alignment:center"
                    ),
                },
                {
                    "code": "partial_clause_coverage",
                    "json_pointer": "$.clause_reviews[1]",
                    "raw_error": (
                        "$.clause_reviews[1]: partial_clause_coverage:"
                        "table_caption.paragraph.alignment:center"
                    ),
                },
            ],
            chunk=chunk,
        )
        self.assertIsNotNone(repaired)
        assert repaired is not None
        self.assertEqual(
            repaired["requirements"][1]["properties"],
            {"position": "above", "paragraph": {"alignment": "center"}},
        )
        self.assertEqual(
            repaired["requirements"][2]["properties"],
            {"position": "above", "paragraph": {"alignment": "center"}},
        )
        self.assertEqual(
            {item["rule_id"] for item in repairs},
            {
                "compile_explicit_table_caption_position_v1",
                "compile_explicit_table_caption_alignment_v1",
            },
        )

        unrelated = json.loads(json.dumps(response, ensure_ascii=False))
        unrelated["requirements"][1]["clause_ids"] = ["C3"]
        rejected, _ = bridge._apply_safe_mechanical_repairs(
            unrelated,
            [
                {
                    "code": "empty_requirement_properties",
                    "json_pointer": "$.requirements[1].properties",
                },
                {
                    "code": "partial_clause_coverage",
                    "json_pointer": "$.clause_reviews[0]",
                    "raw_error": (
                        "$.clause_reviews[0]: partial_clause_coverage:"
                        "table_caption.position:above"
                    ),
                },
            ],
            chunk=chunk,
        )
        self.assertIsNone(rejected)

    def test_fixed_declaration_candidate_is_derived_from_exact_chunk_evidence(self) -> None:
        packet = bridge.compact_model_packet({
            "contract_version": "3.0",
            "clauses": [
                {"id": "C1", "text": "学位论文使用授权书", "evidence_ids": ["E1"]},
                {"id": "C2", "text": "本人同意提交电子版", "evidence_ids": ["E2"]},
                {"id": "C3", "text": "学位论文作者暨授权人签字", "evidence_ids": ["E3"]},
                {"id": "C4", "text": "摘要", "evidence_ids": ["E4"]},
            ],
            "evidence_context": {
                "E1": {"id": "E1", "kind": "paragraph", "text": "学位论文使用授权书"},
                "E2": {"id": "E2", "kind": "paragraph", "text": "本人同意提交电子版"},
                "E3": {"id": "E3", "kind": "paragraph", "text": "学位论文作者暨授权人签字"},
                "E4": {"id": "E4", "kind": "paragraph", "text": "摘要"},
            },
            "declaration_anchor_preference": "abstract_title_zh",
            "requirement_contract": {},
        })
        self.assertEqual(packet["fixed_declaration_candidates"][0]["clause_ids"], ["C1", "C2"])
        self.assertEqual(packet["fixed_declaration_candidates"][0]["before_role"], "abstract_title_zh")

    def test_fixed_declaration_materialization_restores_source_paragraphs_and_deduplicates_split_evidence(self) -> None:
        source = {
            "clauses": [
                {
                    "id": "C1", "text": "学位论文使用授权书",
                    "source_text_full": "学位论文使用授权书",
                    "evidence_ids": ["E1"],
                },
                {
                    "id": "C2", "text": "第一段正文",
                    "source_text_full": "第一段正文。",
                    "evidence_ids": ["E2"],
                },
                {
                    "id": "C3", "text": "第二段前半部分",
                    "source_text_full": "第二段前半部分；第二段后半部分；",
                    "evidence_ids": ["E3"],
                },
                {
                    "id": "C4", "text": "第二段后半部分",
                    "source_text_full": "第二段前半部分；第二段后半部分；",
                    "evidence_ids": ["E3"],
                },
                {
                    "id": "C5", "text": "摘要", "evidence_ids": ["E4"],
                },
            ],
            "evidence_context": {
                "E1": {"id": "E1", "text": "学位论文使用授权书"},
                "E2": {"id": "E2", "text": "第一段正文。"},
                "E3": {"id": "E3", "text": "第二段前半部分；第二段后半部分；"},
                "E4": {"id": "E4", "text": "摘要"},
            },
            "declaration_anchor_preference": "abstract_title_zh",
            "provenance": {
                "run_id": "run-declaration-materialization",
                "case_id": "case-declaration-materialization",
                "source_sha256": "a" * 64,
                "clause_sha256": "b" * 64,
                "evidence_sha256": "c" * 64,
                "request_sha256": "d" * 64,
                "chunk_sha256": "e" * 64,
            },
        }
        bind_declaration_fixture_spans(source)
        candidate = bridge._fixed_declaration_candidates(
            source["clauses"], source["evidence_context"],
            anchor=source["declaration_anchor_preference"],
        )[0]
        self.assertEqual(candidate["clause_ids"], ["C1", "C2", "C3", "C4"])
        self.assertEqual(candidate["body_evidence_ids"], ["E2", "E3"])

        response = {
            "requirements": [{
                "role": "declarations",
                "existing_requirement_id": None,
                "clause_ids": ["C1", "C2", "C3", "C4"],
                "evidence_ids": ["E1", "E2", "E3"],
                "properties": {
                    "before_role": "abstract_title_zh",
                    "items": [{
                        "id": "authorization",
                        "heading": "学位论文使用授权书",
                        "body": None,
                        "body_parts": ["第一段正文", "第二段前半部分", "第二段后半部分"],
                        "source_evidence_ids": ["E1", "E2", "E3"],
                        "signature_placeholders": [],
                    }],
                },
            }],
            "clause_reviews": [
                {"clause_id": clause_id, "classification": "executable"}
                for clause_id in ("C1", "C2", "C3", "C4")
            ],
        }
        projected, audits = bridge._materialize_fixed_declaration_source_text(response, source)
        item = projected["requirements"][0]["properties"]["items"][0]
        self.assertEqual(item["heading"], "学位论文使用授权书")
        self.assertEqual(item["body_parts"], ["第一段正文。", "第二段前半部分；第二段后半部分；"])
        self.assertEqual(len(audits), 1)
        self.assertEqual(audits[0]["rule_id"], "fixed_declaration_source_text_materialization_v1")
        self.assertEqual(audits[0]["run_id"], "run-declaration-materialization")
        self.assertEqual(audits[0]["body_evidence_ids"], ["E2", "E3"])
        self.assertEqual(
            response["requirements"][0]["properties"]["items"][0]["body_parts"],
            ["第一段正文", "第二段前半部分", "第二段后半部分"],
        )

        retried_raw = json.loads(json.dumps(response))
        retried_raw["requirements"][0]["properties"]["items"][0]["body_parts"] = [
            "第一段正文。", "第二段前半部分；第二段后半部分；",
            "第二段前半部分；第二段后半部分；",
        ]
        parent_comparison, _ = bridge._materialize_fixed_declaration_source_text(response, source)
        retry_comparison, _ = bridge._materialize_fixed_declaration_source_text(retried_raw, source)
        error, changed_paths = bridge._retry_semantic_change_error(
            response, retried_raw, [],
            contract_version=bridge.HOST_REVIEW_CONTRACT_V3,
            chunk=source,
            comparison_previous_response=parent_comparison,
            comparison_current_response=retry_comparison,
        )
        self.assertIsNone(error)
        self.assertEqual(changed_paths, [])

    def test_fixed_declaration_materialization_refuses_unbound_model_text_or_citations(self) -> None:
        source = {
            "clauses": [
                {"id": "C1", "text": "学位论文使用授权书", "evidence_ids": ["E1"]},
                {"id": "C2", "text": "固定正文", "evidence_ids": ["E2"]},
                {"id": "C3", "text": "摘要", "evidence_ids": ["E3"]},
            ],
            "evidence_context": {
                "E1": {"id": "E1", "text": "学位论文使用授权书"},
                "E2": {"id": "E2", "text": "固定正文。"},
                "E3": {"id": "E3", "text": "摘要"},
            },
            "declaration_anchor_preference": "abstract_title_zh",
        }
        response = {
            "requirements": [{
                "role": "declarations", "existing_requirement_id": None,
                "clause_ids": ["C1", "C2"], "evidence_ids": ["E1", "E2"],
                "properties": {"before_role": "abstract_title_zh", "items": [{
                    "id": "authorization", "heading": "学位论文使用授权书",
                    "body": None, "body_parts": ["模型编造正文"],
                    "source_evidence_ids": ["E1", "E2"],
                    "signature_placeholders": [],
                }]},
            }],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "executable"},
            ],
        }
        bind_declaration_fixture_spans(source)
        unchanged, audits = bridge._materialize_fixed_declaration_source_text(response, source)
        self.assertEqual(unchanged, response)
        self.assertEqual(audits, [])

        bad_citation = json.loads(json.dumps(response))
        bad_citation["requirements"][0]["evidence_ids"] = ["E1"]
        unchanged, audits = bridge._materialize_fixed_declaration_source_text(bad_citation, source)
        self.assertEqual(unchanged, bad_citation)
        self.assertEqual(audits, [])

    def test_non_public_administrative_heading_is_not_a_fixed_declaration(self) -> None:
        packet = bridge.compact_model_packet({
            "contract_version": "3.0",
            "clauses": [
                {"id": "C1", "text": "非公开学位论文标注说明", "evidence_ids": ["E1"]},
                {"id": "C2", "text": "根据北京体育大学有关规定，非公开学位论文须经批准方能标注。", "evidence_ids": ["E2"]},
                {"id": "C3", "text": "摘要", "evidence_ids": ["E3"]},
            ],
            "evidence_context": {
                "E1": {"id": "E1", "kind": "paragraph", "text": "非公开学位论文标注说明"},
                "E2": {"id": "E2", "kind": "paragraph", "text": "根据北京体育大学有关规定，非公开学位论文须经批准方能标注。"},
                "E3": {"id": "E3", "kind": "paragraph", "text": "摘要"},
            },
            "declaration_anchor_preference": "abstract_title_zh",
            "requirement_contract": {},
        })
        self.assertEqual(packet["fixed_declaration_candidates"], [])
        chunks, _ = engine._source_atomic_chunks(packet["clauses"], 1)
        self.assertEqual([[item["id"] for item in chunk] for chunk in chunks], [
            ["C1", "C2"], ["C3"],
        ])

    def test_fixed_declaration_candidate_stops_before_separate_table_and_seal(self) -> None:
        clauses = [
            {"id": "C1", "text": "学位论文使用授权书", "source_kind": "paragraph", "evidence_ids": ["E1"]},
            {"id": "C2", "text": "固定授权正文", "source_kind": "paragraph", "evidence_ids": ["E2"]},
            {"id": "C3", "text": "审批表编号", "source_kind": "table_cell", "evidence_ids": ["E3"]},
            {"id": "C4", "text": "办公室盖章", "source_kind": "paragraph", "evidence_ids": ["E4"]},
        ]
        evidence = {
            item["evidence_ids"][0]: {"text": item["text"]} for item in clauses
        }
        candidates = bridge._fixed_declaration_candidates(clauses, evidence, anchor="abstract_title_zh")
        self.assertEqual([item["clause_ids"] for item in candidates], [["C1", "C2"]])

    def test_bsu_admin_region_is_read_together_but_not_a_fixed_declaration(self) -> None:
        clauses = [
            {"id": "C00038", "text": "非公开学位论文标注说明", "evidence_ids": ["E00036"]},
            {"id": "C00039", "text": "非公开学位论文须经指导教师同意、作者本人申请和相关部门批准方能标注", "evidence_ids": ["E00037"]},
            {"id": "C00049", "text": "北京体育大学学位评定委员会办公室盖章(有效)", "evidence_ids": ["E00046"]},
            {"id": "C00053", "text": "学位论文使用授权书", "evidence_ids": ["E00050"]},
            {"id": "C00054", "text": "本论文的固定授权正文。", "evidence_ids": ["E00051"]},
            {"id": "C00062", "text": "摘要", "evidence_ids": ["E00060"]},
        ]
        evidence = {
            item["evidence_ids"][0]: {"id": item["evidence_ids"][0], "text": item["text"]}
            for item in clauses
        }
        chunks, _ = engine._source_atomic_chunks(clauses, 1)
        self.assertEqual([[item["id"] for item in chunk] for chunk in chunks], [
            ["C00038", "C00039", "C00049"], ["C00053", "C00054"], ["C00062"],
        ])
        candidates = bridge._fixed_declaration_candidates(
            clauses, evidence, anchor="abstract_title_zh",
        )
        self.assertEqual([candidate["clause_ids"] for candidate in candidates], [
            ["C00053", "C00054"],
        ])

    def test_declaration_text_binding_keeps_external_clause_out_of_requirement_edge(self) -> None:
        source = {
            "clauses": [
                {"id": "C1", "text": "学位论文使用授权书", "evidence_ids": ["E1"]},
                {"id": "C2", "text": "须由作者本人签署", "evidence_ids": ["E2"]},
                {"id": "C3", "text": "固定正文须完整展示", "evidence_ids": ["E3"]},
                {"id": "C4", "text": "摘要", "evidence_ids": ["E4"]},
            ],
            "evidence_context": {
                "E1": {"id": "E1", "text": "学位论文使用授权书"},
                "E2": {"id": "E2", "text": "须由作者本人签署。"},
                "E3": {"id": "E3", "text": "固定正文须完整展示。"},
                "E4": {"id": "E4", "text": "摘要"},
            },
            "declaration_anchor_preference": "abstract_title_zh",
        }
        response = {
            "requirements": [{
                "role": "declarations", "clause_ids": ["C1", "C3"],
                "evidence_ids": ["E1", "E3"], "existing_requirement_id": None,
                "properties": {"before_role": "abstract_title_zh", "items": [{
                    "id": "authorization", "source_evidence_ids": ["E1", "E2", "E3"],
                    "signature_placeholders": [], "heading": None, "body_parts": None,
                }]},
            }],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "external_compliance"},
                {"clause_id": "C3", "classification": "executable"},
            ],
        }
        bind_declaration_fixture_spans(source)
        projected, audits = bridge._materialize_fixed_declaration_source_text(response, source)
        item = projected["requirements"][0]["properties"]["items"][0]
        self.assertEqual(item["heading"], "学位论文使用授权书")
        self.assertEqual(item["body_parts"], ["须由作者本人签署。", "固定正文须完整展示。"])
        self.assertEqual(projected["requirements"][0]["clause_ids"], ["C1", "C3"])
        self.assertEqual(projected["clause_reviews"][1]["classification"], "external_compliance")
        self.assertEqual(audits[0]["rendered_only_clause_ids"], ["C2"])
        self.assertEqual(response["requirements"][0]["properties"]["items"][0]["body_parts"], None)

        bad = copy.deepcopy(response)
        bad["requirements"][0]["properties"]["items"][0]["source_evidence_ids"] = ["E1", "E4"]
        unchanged, audits = bridge._materialize_fixed_declaration_source_text(bad, source)
        self.assertEqual(unchanged, bad)
        self.assertEqual(audits, [])

    def test_mixed_relation_is_not_a_mechanical_retry_plan(self) -> None:
        records = [
            {"code": "mixed_execution_classification_relation", "json_pointer": "$.requirements[0]"},
            {"code": "missing_derived_requirement", "json_pointer": "$.clause_reviews[0]",
             "blocked_by_parent_relation": {"requirement_indexes": [0]}},
        ]
        self.assertTrue(bridge._requires_fresh_semantic_split(records))
        guidance = bridge._structured_contract_repair_guidance(records, contract_version="3.0")
        self.assertIn("do not add a second or title-only requirement", guidance)
        previous = {"contract_version": "3.0", "requirements": [], "clause_reviews": []}
        current = {**previous, "requirements": [{"role": "declarations"}]}
        self.assertFalse(bridge._retry_changes_allowed(
            records, ["$.requirements"], contract_version="3.0",
            previous_response=previous, current_response=current, chunk={},
        ))
        self.assertEqual(bridge._v3_relation_completion_response(
            previous, current, records, chunk={},
        ), (None, None))

    def test_external_pending_and_unbound_orphan_never_form_relation_addition_retry(self) -> None:
        """The two simultaneous BSU failures must not yield conflicting retry advice."""
        previous = {
            "contract_version": "3.0",
            "requirements": [
                {"role": "cover", "properties": {"non_public_administration": {"fields": [{"id": "approval_number"}]}},
                 "clause_ids": ["C1"], "evidence_ids": ["E1"], "reason": "cover"},
                {"role": "content_constraints", "properties": {"abstract_zh": {"required": None}},
                 "clause_ids": ["C2"], "evidence_ids": ["E1"], "reason": "external approval",
                 "applicability": {"status": "conditional"},
                 "verification": {"mode": "external"}},
                {"role": "cover", "properties": {"institution": "——", "fields": [],
                 "missing_value_policy": "placeholder"}, "clause_ids": [], "evidence_ids": [],
                 "reason": "placeholder"},
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable"},
                {"clause_id": "C2", "classification": "external_compliance",
                 "obligations": [{"id": "approval", "status": "unverifiable"}]},
            ],
            "unsupported_items": [], "reported_conflicts": [],
        }
        current = copy.deepcopy(previous)
        current["requirements"] = [copy.deepcopy(previous["requirements"][0])]
        parent_sha = bridge._response_sha256(previous)
        records = [
            {"code": "cover_binding_violation", "json_pointer": "$.requirements[2].properties.fields",
             "response_sha256": parent_sha},
            {"code": "schema_contract_violation", "json_pointer": "$.requirements[2].clause_ids",
             "response_sha256": parent_sha},
            {"code": "schema_contract_violation", "json_pointer": "$.requirements[2].evidence_ids",
             "response_sha256": parent_sha},
            {"code": "non_requirement_classification_relation",
             "json_pointer": "$.requirements[1]", "requirement_index": 1,
             "relation_category": "non_requirement_classification", "mechanically_removable": False,
             "clause_classifications": {"C2": ["external_compliance"]},
             "response_sha256": parent_sha},
            {"code": "requirement_relation_mismatch", "json_pointer": "$.requirements[2]",
             "requirement_index": 2, "relation_category": "missing_clause_relation",
             "mechanically_removable": True, "mechanical_removal_basis": "no_clause_or_evidence_binding",
             "clause_ids": [], "response_sha256": parent_sha},
        ]
        self.assertEqual(
            bridge._fresh_semantic_split_reason(records),
            "external_pending_requirement_plus_unbound_orphan",
        )
        guidance = bridge._structured_contract_repair_guidance(records, contract_version="3.0")
        self.assertIn("not a missing executable clause", guidance)
        self.assertIn("Never assign it a guessed clause/evidence", guidance)
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(changed, ["$.requirements"])
        error, _ = bridge._retry_semantic_change_error(
            previous, current, records, contract_version="3.0",
        )
        self.assertIsNotNone(error)
        self.assertEqual(error.error_records[0]["code"], "semantic_retry_change")
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            packet_path = root / "chunk.json"
            packet_path.write_text('{"contract_version":"3.0"}', encoding="utf-8")
            parent_path = root / "parent.json"
            parent_path.write_text(json.dumps(previous), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "incompatible semantic errors"):
                bridge._host_prompt(
                    request_path=root / "request.json", chunk_path=packet_path,
                    response_path=root / "response.json", run_id="run", chunk_index=1,
                    chunk_count=1, attempt=2, retry_hint="contract failed",
                    retry_parent_response_path=parent_path,
                    retry_parent_response_sha256=bridge.sha256_file(parent_path),
                    retry_error_records=records,
                )

        stale = copy.deepcopy(records)
        stale[-1]["response_sha256"] = "0" * 64
        self.assertIsNone(bridge._fresh_semantic_split_reason(stale))
        without_orphan = records[:-1]
        self.assertIsNone(bridge._fresh_semantic_split_reason(without_orphan))
        non_external = copy.deepcopy(records)
        non_external[-2]["clause_classifications"] = {"C2": ["unresolved"]}
        self.assertIsNone(bridge._fresh_semantic_split_reason(non_external))

    def test_unknown_property_is_removed_deterministically_without_semantic_retry(self) -> None:
        response = {
            "requirements": [{"properties": {"style": {"name": "bad"}, "text": "标题"}}],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "unknown_property",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "$.requirements[0].properties: unknown property 'style'",
            }],
        )
        self.assertEqual(repairs[0]["removed_property"], "style")
        self.assertIsNotNone(repaired)
        self.assertNotIn("style", repaired["requirements"][0]["properties"])
        self.assertEqual(repaired["requirements"][0]["properties"]["text"], "标题")
        self.assertIn("style", response["requirements"][0]["properties"])

    def test_retry_allows_exact_text_fill_after_unknown_property_removal(self) -> None:
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "abstract_title_zh",
                "properties": {"style": "three_line"},
                "clause_ids": ["C1"],
                "evidence_ids": ["E1"],
            }],
            "clause_reviews": [{"clause_id": "C1", "classification": "executable"}],
            "unsupported_items": [],
        }
        current = json.loads(json.dumps(previous))
        current["requirements"][0]["properties"] = {"text": "摘  要"}
        records = [{
            "code": "unknown_property",
            "json_pointer": "$.requirements[0].properties",
            "raw_error": "unknown property 'style'",
        }]
        chunk = {
            "evidence_context": {"E1": {"text": "摘  要"}},
            "requirement_contract": {
                "role_properties_schema": {
                    "abstract_title_zh": {"$ref": "#/$defs/roleSpec"},
                },
            },
        }
        changed = bridge._retry_change_paths(previous, current)
        self.assertEqual(
            changed,
            [
                "$.requirements[0].properties.style",
                "$.requirements[0].properties.text",
            ],
        )
        self.assertTrue(bridge._retry_changes_allowed(
            records,
            changed,
            contract_version="3.0",
            previous_response=previous,
            current_response=current,
            chunk=chunk,
        ))
        current["requirements"][0]["properties"]["text"] = "猜测标题"
        changed = bridge._retry_change_paths(previous, current)
        self.assertFalse(bridge._retry_changes_allowed(
            records,
            changed,
            contract_version="3.0",
            previous_response=previous,
            current_response=current,
            chunk=chunk,
        ))

    def test_mixed_contract_errors_yield_only_a_partial_candidate(self) -> None:
        response = {"requirements": [{"properties": {"style": {}}}]}
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [
                {
                    "code": "unknown_property",
                    "json_pointer": "$.requirements[0].properties",
                    "raw_error": "unknown property 'style'",
                },
                {
                    "code": "requirement_relation_mismatch",
                    "json_pointer": "$.requirements",
                    "raw_error": "requirements_not_referenced_by_clause_review:0",
                },
            ],
        )
        self.assertIsNotNone(repaired)
        self.assertEqual(repaired["requirements"][0]["properties"], {})
        self.assertEqual(response["requirements"][0]["properties"], {"style": {}})
        self.assertTrue(all(item["partial_candidate_only"] for item in repairs))
        self.assertTrue(all(item["planner_id"] == "composable_source_bound_patch_plan_v1" for item in repairs))

    def test_unbacked_evidence_id_is_removed_deterministically(self) -> None:
        response = {
            "requirements": [{
                "evidence_ids": ["E1", "E2"],
                "clause_ids": ["C1"],
            }],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "evidence_relation_mismatch",
                "json_pointer": "$.requirements[0].evidence_ids",
                "raw_error": "$.requirements[0].evidence_ids: not_backed_by_clause:E2",
            }],
        )
        self.assertEqual(repairs[0]["removed_evidence_id"], "E2")
        self.assertEqual(repaired["requirements"][0]["evidence_ids"], ["E1"])
        self.assertEqual(response["requirements"][0]["evidence_ids"], ["E1", "E2"])

    def test_empty_role_payload_is_filled_from_exact_evidence_text(self) -> None:
        response = {
            "requirements": [{
                "role": "abstract_title_zh", "evidence_ids": ["E1"],
                "properties": {},
            }],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "empty_requirement_properties",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "must_include_semantic_payload",
            }],
            chunk={
                "evidence_context": {
                    "E1": {"style_name": "heading 1", "text": "摘 要"},
                },
                "requirement_contract": {
                    "role_properties_schema": {
                        "abstract_title_zh": {"$ref": "#/$defs/roleSpec"},
                    },
                },
            },
        )
        self.assertEqual(repaired["requirements"][0]["properties"]["text"], "摘 要")
        self.assertEqual(repairs[0]["filled_property"], "text")
        self.assertNotIn("text", response["requirements"][0]["properties"])

    def test_exact_acknowledgments_limit_is_compiled_into_nested_payload(self) -> None:
        response = {
            "requirements": [{
                "role": "content_constraints",
                "properties": {},
                "clause_ids": ["C1"],
                "evidence_ids": ["E1"],
            }],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "empty_requirement_properties",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "must_include_semantic_payload",
            }],
            chunk={
                "clauses": [{
                    "id": "C1", "text": "字数一般不超过500字", "evidence_ids": ["E1"],
                }],
                "evidence_context": {
                    "E1": {"id": "E1", "text": "字数一般不超过500字"},
                },
                "requirement_contract": {},
            },
        )
        self.assertIsNotNone(repaired)
        self.assertEqual(
            repaired["requirements"][0]["properties"],
            {"acknowledgments": {"max_chars": 500}},
        )
        self.assertEqual(repairs[0]["rule_id"], "compile_exact_acknowledgments_max_chars_v1")

        response["requirements"][0]["clause_ids"] = ["C2"]
        rejected, _ = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "empty_requirement_properties",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "must_include_semantic_payload",
            }],
            chunk={
                "clauses": [{"id": "C2", "text": "致谢内容应简短", "evidence_ids": ["E1"]}],
                "evidence_context": {"E1": {"id": "E1", "text": "致谢内容应简短"}},
                "requirement_contract": {},
            },
        )
        self.assertIsNone(rejected)

    def test_exact_appendix_page_break_is_compiled_into_declared_payload(self) -> None:
        response = {
            "requirements": [{
                "role": "appendices",
                "properties": {
                    "required_when_profile_has_appendices": None,
                    "label_style": None,
                    "label_prefix": None,
                    "page_break_each": None,
                    "per_appendix_title_required": None,
                },
                "clause_ids": ["C1"],
                "evidence_ids": ["E1"],
            }],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "empty_requirement_properties",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "must_include_semantic_payload",
            }],
            chunk={
                "clauses": [{
                    "id": "C1", "text": "附录放在正文之后另起页", "evidence_ids": ["E1"],
                }],
                "evidence_context": {
                    "E1": {"id": "E1", "text": "附录放在正文之后另起页"},
                },
                "requirement_contract": {},
            },
        )
        self.assertIsNotNone(repaired)
        self.assertEqual(
            repaired["requirements"][0]["properties"],
            {"page_break_each": True},
        )
        self.assertEqual(repairs[0]["rule_id"], "compile_exact_appendix_page_break_v1")

        response["requirements"][0]["clause_ids"] = ["C2"]
        rejected, _ = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "empty_requirement_properties",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "must_include_semantic_payload",
            }],
            chunk={
                "clauses": [{"id": "C2", "text": "附录放在正文之后"}],
                "evidence_context": {"E1": {"id": "E1", "text": "附录放在正文之后"}},
                "requirement_contract": {},
            },
        )
        self.assertIsNone(rejected)

    def test_exact_appendix_label_title_is_compiled_into_declared_payload(self) -> None:
        response = {
            "requirements": [{
                "role": "appendices",
                "properties": {
                    "required_when_profile_has_appendices": None,
                    "label_style": None,
                    "label_prefix": None,
                    "page_break_each": None,
                    "per_appendix_title_required": None,
                },
                "clause_ids": ["C1"],
                "evidence_ids": ["E1"],
            }],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "empty_requirement_properties",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "must_include_semantic_payload",
            }],
            chunk={
                "clauses": [{
                    "id": "C1",
                    "text": "附录的序号用A，B，C，…系列，如附录A，附录B等，每个附录应有标题",
                    "evidence_ids": ["E1"],
                }],
                "evidence_context": {
                    "E1": {
                        "id": "E1",
                        "text": "附录的序号用A，B，C，…系列，如附录A，附录B等，每个附录应有标题",
                    },
                },
                "requirement_contract": {},
            },
        )
        self.assertIsNotNone(repaired)
        self.assertEqual(
            repaired["requirements"][0]["properties"],
            {"label_style": "alpha_upper", "per_appendix_title_required": True},
        )
        self.assertEqual(repairs[0]["rule_id"], "compile_exact_appendix_label_title_v1")

    def test_empty_administrative_cover_institution_uses_neutral_placeholder(self) -> None:
        response = {
            "requirements": [{
                "role": "cover",
                "properties": {
                    "institution": "",
                    "fields": [{"id": "title_zh"}],
                    "non_public_administration": None,
                    "missing_value_policy": "placeholder",
                    "missing_value_placeholder": "——",
                },
                "clause_ids": ["C1"],
                "evidence_ids": ["E1"],
            }],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "cover_institution_placeholder",
                "json_pointer": "$.requirements[0].properties",
            }, {
                "code": "cover_institution_placeholder",
                "json_pointer": "$.requirements[0].properties.institution",
            }],
        )
        self.assertIsNotNone(repaired)
        self.assertEqual(
            repaired["requirements"][0]["properties"]["institution"],
            "——",
        )
        self.assertEqual(
            repairs[0]["rule_id"],
            "compile_neutral_cover_institution_placeholder_v1",
        )

        response["requirements"][0]["properties"].pop("missing_value_placeholder")
        rejected, _ = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "cover_institution_placeholder",
                "json_pointer": "$.requirements[0].properties.institution",
            }],
        )
        self.assertIsNone(rejected)

    def test_zero_based_cover_orders_shift_only_complete_source_bound_groups(self) -> None:
        labels = ["论文题目", "申请密级", "保密期限", "审批表编号"]
        chunk = {
            "case_id": "case-test",
            "provenance": {
                "run_id": "run-test", "source_sha256": "a" * 64,
                "clause_sha256": "b" * 64, "evidence_sha256": "c" * 64,
                "request_sha256": "d" * 64,
            },
            "clauses": [
                {"id": f"C{i}", "text": label, "evidence_ids": [f"E{i}"]}
                for i, label in enumerate(labels, 1)
            ],
            "evidence_context": {
                f"E{i}": {"id": f"E{i}", "text": label}
                for i, label in enumerate(labels, 1)
            },
        }
        response = {
            "requirements": [{
                "role": "cover", "clause_ids": ["C1", "C2", "C3", "C4"],
                "evidence_ids": ["E1", "E2", "E3", "E4"],
                "properties": {
                    "institution": "", "missing_value_policy": "placeholder",
                    "missing_value_placeholder": "——",
                    "fields": [{"id": "title_zh", "label": "论文题目", "order": 0}],
                    "non_public_administration": {"fields": [
                        {"id": "security_marking", "label": "申请密级", "order": 0},
                        {"id": "embargo_start", "label": "保密期限", "order": 1},
                        {"id": "embargo_until", "label": "保密期限", "order": 2},
                        {"id": "approval_number", "label": "审批表编号", "order": 3},
                    ]},
                },
            }],
        }

        def records(candidate: dict) -> list[dict]:
            sha = bridge._response_sha256(candidate)
            return [{"code": code, "json_pointer": pointer,
                     "raw_error": raw, "response_sha256": sha}
                    for code, pointer, raw in (
                ("contract_validation_error", "$.requirements[0]",
                 "$.requirements[0]: must match at least one schema in anyOf"),
                ("cover_institution_placeholder", "$.requirements[0].properties.institution",
                 "$.requirements[0].properties.institution: is shorter than 1 characters"),
                ("contract_validation_error", "$.requirements[0].properties.fields[0].order",
                 "$.requirements[0].properties.fields[0].order: must be >= 1"),
                ("cover_binding_violation", "$.requirements[0].properties.non_public_administration.fields[0].order",
                 "$.requirements[0].properties.non_public_administration.fields[0].order: must be >= 1"),
            )]

        repaired, audit = bridge._apply_safe_mechanical_repairs(
            response, records(response), chunk=chunk,
        )
        self.assertIsNotNone(repaired)
        assert repaired is not None
        props = repaired["requirements"][0]["properties"]
        self.assertEqual(props["institution"], "——")
        self.assertEqual([field["order"] for field in props["fields"]], [1])
        self.assertEqual(
            [field["order"] for field in props["non_public_administration"]["fields"]],
            [1, 2, 3, 4],
        )
        self.assertEqual(response["requirements"][0]["properties"]["institution"], "")
        self.assertEqual(audit[0]["rule_id"], "normalize_source_bound_zero_based_cover_order_v1")
        self.assertEqual(audit[0]["source_response_sha256"], bridge._response_sha256(response))
        self.assertEqual(audit[0]["run_id"], "run-test")

        for mutation in ("duplicate_order", "order_gap", "wrong_source", "source_reversed",
                         "boolean_order", "missing_placeholder_policy", "uncited_evidence"):
            candidate = copy.deepcopy(response)
            fields = candidate["requirements"][0]["properties"]["non_public_administration"]["fields"]
            if mutation == "duplicate_order":
                fields[1]["order"] = 0
            elif mutation == "order_gap":
                fields[2]["order"] = 4
            elif mutation == "wrong_source":
                fields[0]["label"] = "未引用的标签"
            elif mutation == "source_reversed":
                fields[0], fields[-1] = fields[-1], fields[0]
                for index, field in enumerate(fields):
                    field["order"] = index
            elif mutation == "boolean_order":
                fields[0]["order"] = False
            elif mutation == "missing_placeholder_policy":
                candidate["requirements"][0]["properties"].pop("missing_value_policy")
            elif mutation == "uncited_evidence":
                candidate["requirements"][0]["evidence_ids"].remove("E3")
            self.assertIsNone(
                bridge._apply_safe_mechanical_repairs(candidate, records(candidate), chunk=chunk)[0],
                mutation,
            )

        stale = records(response)
        stale[0]["response_sha256"] = "0" * 64
        self.assertIsNone(bridge._apply_safe_mechanical_repairs(response, stale, chunk=chunk)[0])
        unrelated = records(response) + [{
            "code": "contract_validation_error", "json_pointer": "$.requirements[0].reason",
            "raw_error": "$.requirements[0].reason: is shorter than 1 characters",
            "response_sha256": bridge._response_sha256(response),
        }]
        self.assertIsNone(bridge._apply_safe_mechanical_repairs(response, unrelated, chunk=chunk)[0])

    def test_cover_security_marking_is_migrated_only_from_exact_linked_source(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(
                Path(td) / "requirements", contract_version="3.0",
            )
            clause_source = "封面密  级"
            evidence_source = "封面密  级："
            chunk["clauses"] = [{
                "id": "C1", "text": clause_source, "evidence_ids": ["E1"],
                "source_kind": "paragraph", "location": {}, "part_index": 0,
                "source_span": {
                    "evidence_id": "E1", "start_offset": 0,
                    "end_offset": len(clause_source), "text": clause_source,
                    "source_sha256": hashlib.sha256(evidence_source.encode("utf-8")).hexdigest(),
                },
            }]
            chunk["evidence_context"] = {
                "E1": {"id": "E1", "text": evidence_source},
            }
            contract_version = chunk["contract_version"]
            security_field = {
                "id": "security_marking", "label": "密  级：",
                "value_from": "thesis_profile.cover_metadata.security_marking",
                "display_policy": "if_present", "order": 3,
            }
            response = {
                "contract_version": contract_version,
                "provenance": chunk["provenance"],
                "requirements": [{
                    "role": "cover",
                    "properties": {
                        "institution": "",
                        "fields": [
                            {
                                "id": "classification_number", "label": "分类号：",
                                "value_from": "thesis_profile.cover_metadata.classification_number",
                                "display_policy": "required", "order": 1,
                            },
                            {
                                "id": "unit_code", "label": "学校代码：",
                                "value_from": "thesis_profile.cover_metadata.unit_code",
                                "display_policy": "if_present", "order": 2,
                            },
                            security_field,
                        ],
                        "non_public_administration": None,
                        "before_role": "document_start",
                        "missing_value_policy": "placeholder",
                        "missing_value_placeholder": "——",
                        "layout_id": "linear",
                    },
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                    "confidence": 0.98,
                    "reason": "来源明确标注封面密级字段。",
                    "applicability": {"status": "always"},
                    "input_prerequisites": [],
                    "verification": {"mode": "word_render", "checks": ["核验封面字段"]},
                }],
                "clause_reviews": [{
                    "clause_id": "C1", "classification": "executable",
                    "reason": "来源明确规定封面密级字段。",
                    "normative_basis": "template_structure",
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }
            if contract_version == "2.1":
                response["clause_reviews"][0]["requirement_indexes"] = [0]
            else:
                response["clause_reviews"][0]["obligations"] = [{
                    "id": "cover_security_label", "status": "covered",
                    "reason": "精确字段标签已由来源支持。",
                }]

            errors = bridge.validate_host_agent_response(response, chunk)
            self.assertTrue(errors)
            records = bridge.contract_error_records(errors, response=response, chunk=chunk)
            self.assertIn("cover_binding_violation", {item["code"] for item in records})
            repaired, audit = bridge._apply_safe_mechanical_repairs(
                response, records, chunk=chunk,
            )
            self.assertIsNotNone(repaired, repr(records))
            assert repaired is not None
            self.assertEqual(bridge.validate_host_agent_response(repaired, chunk), [])

            repaired_properties = repaired["requirements"][0]["properties"]
            self.assertEqual(repaired_properties["institution"], "——")
            self.assertEqual(
                repaired_properties["fields"],
                response["requirements"][0]["properties"]["fields"][:2],
            )
            admin = repaired_properties["non_public_administration"]
            self.assertEqual(admin["fields"], [security_field])
            self.assertEqual(admin["public_policy"], "blank")
            self.assertEqual(admin["source_region"], "cover")
            self.assertEqual(
                admin["applicability"]["conditions"],
                [{
                    "fact": "thesis_profile.security_level", "operator": "in",
                    "value": ["restricted", "classified"],
                }],
            )
            migration = next(
                item for item in audit
                if item.get("rule_id") == "source_bound_cover_security_marking_migration_v1"
            )
            self.assertEqual(migration["source_clause_ids"], ["C1"])
            self.assertEqual(migration["source_evidence_ids"], ["E1"])
            self.assertEqual(
                migration["authorized_changed_paths"],
                [
                    "$.requirements[0].properties.fields",
                    "$.requirements[0].properties.non_public_administration",
                ],
            )

            candidate, candidate_audit = bridge.prepare_native_response_candidate(
                response, chunk,
            )
            self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
            self.assertEqual(
                candidate_audit["mechanical_repairs"][0]["rule_id"],
                "source_bound_cover_security_marking_migration_v1",
            )

    def test_cover_security_marking_migration_fails_closed_without_exact_source_binding(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(Path(td) / "requirements")
            source = "封面设置密级字段"
            chunk["clauses"] = [{
                "id": "C1", "text": source, "evidence_ids": ["E1"],
                "source_kind": "paragraph", "location": {}, "part_index": 0,
            }]
            chunk["evidence_context"] = {"E1": {"id": "E1", "text": source}}
            response = {
                "contract_version": chunk["contract_version"],
                "requirements": [{
                    "role": "cover",
                    "properties": {
                        "institution": "——",
                        "fields": [{
                            "id": "security_marking", "label": "密级：",
                            "value_from": "thesis_profile.cover_metadata.security_marking",
                            "display_policy": "if_present", "order": 1,
                        }],
                        "non_public_administration": None,
                    },
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                    "confidence": 0.9, "reason": "封面字段",
                }],
                "clause_reviews": [],
            }
            record = {
                "code": "cover_binding_violation",
                "json_pointer": "$.requirements[0].properties.fields[0]",
                "raw_error": "$.requirements[0].properties.fields[0]: administrative label '密级' must be declared under cover.non_public_administration, not ordinary cover.fields",
                "response_sha256": bridge._response_sha256(response),
            }
            self.assertIsNone(
                bridge._project_source_bound_cover_security_marking(
                    copy.deepcopy(response), record, chunk=chunk,
                    baseline_response_sha256=bridge._response_sha256(response),
                )
            )
            self.assertIsNone(
                bridge._apply_safe_mechanical_repairs(response, [record], chunk=chunk)[0]
            )

    def test_missing_admin_fields_become_bound_manual_review_without_guessing(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            _review_dir, chunk = self._packet(Path(td) / "requirements")
            chunk["clauses"] = [
                {"id": "C1", "text": "正文使用宋体", "evidence_ids": ["E1"]},
                {"id": "C2", "text": "非公开论文须经审批，公开论文该项为空白", "evidence_ids": ["E2"]},
            ]
            chunk["evidence_context"] = {
                "E1": {"id": "E1", "text": "正文使用宋体"},
                "E2": {"id": "E2", "text": "非公开论文须经审批，公开论文该项为空白"},
            }
            # This fixture extends the packet after _packet() built its
            # source-scoped response schema; keep the schema's ID enum aligned.
            chunk["response_schema"] = engine.build_llm_request(
                [], chunk["clauses"],
                {"evidence": list(chunk["evidence_context"].values())},
                {}, "full", contract_version="2.1",
            )["response_schema"]
            response = {
                "contract_version": "2.1",
                "requirements": [{
                    "role": "cover",
                    "properties": {
                        "institution": "——",
                        "fields": [],
                        "non_public_administration": {
                            "applicability": {
                                "status": "conditional",
                                "conditions": [{
                                    "fact": "thesis_profile.security_level",
                                    "operator": "in",
                                    "value": ["restricted", "classified"],
                                }],
                            },
                            "fields": [],
                            "public_policy": "blank",
                            "source_region": "E2",
                        },
                    },
                    "clause_ids": ["C1", "C2"],
                    "evidence_ids": ["E1", "E2"],
                    "confidence": 0.9,
                    "reason": "行政说明",
                    "verification": {"mode": "static_docx", "checks": ["检查"]},
                }, {
                    "role": "body_text",
                    "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                    "confidence": 0.9,
                    "reason": "正文格式",
                    "verification": {"mode": "static_docx", "checks": ["检查"]},
                }],
                "clause_reviews": [{
                    "clause_id": "C1", "classification": "executable",
                    "requirement_indexes": [1], "reason": "正文格式",
                    "obligations": [{
                        "id": "fixed_statement", "status": "covered",
                        "reason": "固定文本",
                    }],
                }, {
                    "clause_id": "C2", "classification": "executable",
                    "requirement_indexes": [0], "reason": "行政规则",
                    "obligations": [{
                        "id": "public_blank_policy", "status": "covered",
                        "reason": "公开论文留白",
                    }],
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }
            records = [{
                "code": "contract_validation_error",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "must match at least one schema in anyOf",
            }, {
                "code": "non_public_administration_fields_missing",
                "json_pointer": "$.requirements[0].properties.non_public_administration.fields",
                "raw_error": "requires at least 1 items",
            }]
            repaired, repairs = bridge._apply_safe_mechanical_repairs(
                response, records, chunk=chunk,
            )
            self.assertIsNotNone(repaired)
            assert repaired is not None
            self.assertEqual(
                [item["clause_id"] for item in repaired["clause_reviews"]
                 if item["classification"] == "requires_source_content"],
                ["C2"],
            )
            self.assertEqual(repaired["clause_reviews"][1]["requirement_indexes"], [])
            self.assertEqual(repaired["requirements"][0]["clause_ids"], ["C1"])
            self.assertEqual(
                repairs[0]["rule_id"],
                "downgrade_unlabeled_admin_region_to_manual_review_v1",
            )
            self.assertEqual(bridge.validate_host_agent_response(repaired, chunk), [])

            chunk["clauses"][1]["text"] = "非公开论文需填写审批表编号和批准日期"
            chunk["evidence_context"]["E2"]["text"] = chunk["clauses"][1]["text"]
            rejected, _ = bridge._apply_safe_mechanical_repairs(
                response, records, chunk=chunk,
            )
            self.assertIsNone(rejected)

    def test_contract_error_records_name_missing_admin_fields_explicitly(self) -> None:
        response = {"requirements": [{
            "role": "cover",
            "properties": {"non_public_administration": {"fields": []}},
        }]}
        records = bridge.contract_error_records([
            "$.requirements[0].properties.non_public_administration.fields: requires at least 1 items",
        ], response=response, chunk={"clauses": []})
        self.assertEqual(records[0]["code"], "non_public_administration_fields_missing")
        self.assertTrue(records[0]["semantic_review_required"])

    def test_combined_mechanical_repairs_do_not_require_semantic_retry(self) -> None:
        response = {
            "contract_version": "3.0",
            "requirements": [
                {
                    "role": "cover_field_label",
                    "properties": {"style": "three_line", "continuation": None},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                },
                {
                    "role": "cover",
                    "properties": {
                        "institution": "",
                        "fields": [],
                        "missing_value_policy": "placeholder",
                        "missing_value_placeholder": "——",
                    },
                    "clause_ids": ["C2"], "evidence_ids": ["E2"],
                },
                {
                    "role": "declarations",
                    "properties": {"items": [{
                        "id": "authorization",
                        "source_evidence_ids": ["E3", "E3", "E4"],
                    }]},
                    "clause_ids": ["C3"], "evidence_ids": ["E3", "E4"],
                },
            ],
            "clause_reviews": [], "unsupported_items": [], "reported_conflicts": [],
        }
        records = [
            {
                "code": "unknown_property",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": "unknown property 'style'",
            },
            {
                "code": "cover_institution_placeholder",
                "json_pointer": "$.requirements[1].properties",
            },
            {
                "code": "contract_validation_error",
                "json_pointer": "$.requirements[2].properties",
                "raw_error": "must match at least one schema in anyOf",
            },
            {
                "code": "contract_validation_error",
                "json_pointer": "$.requirements[2].properties.items[0].source_evidence_ids",
                "raw_error": "items must be unique",
            },
        ]
        repaired, repairs = bridge._apply_safe_mechanical_repairs(response, records)
        self.assertIsNotNone(repaired)
        assert repaired is not None
        self.assertNotIn("style", repaired["requirements"][0]["properties"])
        self.assertEqual(
            repaired["requirements"][1]["properties"]["institution"], "——",
        )
        self.assertEqual(
            repaired["requirements"][2]["properties"]["items"][0]["source_evidence_ids"],
            ["E3", "E4"],
        )
        self.assertTrue(any(item.get("rule_id") == "deduplicate_first_occurrence_source_evidence_ids_v1"
                            for item in repairs))

    def test_unknown_layout_properties_are_removed_then_exact_text_is_filled(self) -> None:
        response = {
            "requirements": [{
                "role": "cover_field_label", "evidence_ids": ["E1"],
                "properties": {
                    "style": "three_line",
                    "remove_vertical_borders": False,
                    "repeat_header_row": False,
                    "allow_row_split": False,
                    "keep_with_caption": False,
                },
            }],
        }
        records = [
            {
                "code": "unknown_property",
                "json_pointer": "$.requirements[0].properties",
                "raw_error": f"unknown property '{name}'",
            }
            for name in (
                "style", "remove_vertical_borders", "repeat_header_row",
                "allow_row_split", "keep_with_caption",
            )
        ]
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            records,
            chunk={
                "evidence_context": {"E1": {"text": "论文题目（中文）"}},
                "requirement_contract": {
                    "role_properties_schema": {
                        "cover_field_label": {"$ref": "#/$defs/roleSpec"},
                    },
                },
            },
        )
        self.assertEqual(
            repaired["requirements"][0]["properties"],
            {"text": "论文题目（中文）"},
        )
        self.assertEqual(
            [item["code"] for item in repairs],
            ["unknown_property"] * 5 + ["empty_requirement_properties"],
        )

    def test_english_presence_fact_is_compiled_only_from_explicit_clause(self) -> None:
        response = {
            "requirements": [{
                "role": "body_text",
                "properties": {"font": {"latin": "Times New Roman"}},
                "clause_ids": ["C1"],
                "applicability": {
                    "status": "conditional",
                    "conditions": [{
                        "fact": "English text appears in the thesis",
                        "operator": "present",
                        "value": None,
                    }],
                    "exceptions": None,
                },
            }],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "applicability_fact_namespace",
                "json_pointer": "$.requirements[0].applicability.conditions[0].fact",
                "raw_error": (
                    "$.requirements[0].applicability.conditions[0].fact: "
                    "does not match the registered fact namespace"
                ),
            }],
            chunk={
                "clauses": [{
                    "id": "C1",
                    "text": "论文中出现英文时需要使用Times New Roman字体",
                }],
            },
        )
        self.assertEqual(
            repaired["requirements"][0]["applicability"]["conditions"][0]["fact"],
            "source_inventory.english_text",
        )
        self.assertEqual(repairs[0]["rule_id"], "explicit_english_presence_times_new_roman")
        self.assertEqual(
            response["requirements"][0]["applicability"]["conditions"][0]["fact"],
            "English text appears in the thesis",
        )

    def test_english_presence_fact_is_not_guessed_for_other_conditions(self) -> None:
        response = {
            "requirements": [{
                "role": "body_text",
                "properties": {"font": {"latin": "Times New Roman"}},
                "clause_ids": ["C1"],
                "applicability": {
                    "status": "conditional",
                    "conditions": [{
                        "fact": "some prose condition",
                        "operator": "present",
                        "value": None,
                    }],
                },
            }],
        }
        repaired, repairs = bridge._apply_safe_mechanical_repairs(
            response,
            [{
                "code": "applicability_fact_namespace",
                "json_pointer": "$.requirements[0].applicability.conditions[0].fact",
                "raw_error": "does not match the registered fact namespace",
            }],
            chunk={"clauses": [{"id": "C1", "text": "正文采用小四号字体"}]},
        )
        self.assertIsNone(repaired)
        self.assertEqual(repairs, [])

    def test_bridge_retries_locally_rejected_contract_in_a_new_session(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(
                Path(td) / "requirements", contract_version="3.0",
            )
            invalid = self._executable_response(chunk, invalid_verification=True)
            valid = self._executable_response(chunk)
            for response in (invalid, valid):
                response["contract_version"] = "3.0"
                response["clause_reviews"][0].pop("requirement_indexes", None)
                response["clause_reviews"][0].update({
                    "normative_basis": "explicit_normative_text",
                    "obligations": [{
                        "id": "body-font", "status": "covered",
                        "reason": "The current source requirement is represented.",
                    }],
                })
                response["requirements"][0]["role"] = "cover_field_label"
                response["requirements"][0]["properties"] = {
                    "text": "正文 使用宋体",
                }
            envelopes = [
                {"runId": "openclaw-run-contract-retry-1", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(invalid)}]}},
                {"runId": "openclaw-run-contract-retry-2", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(valid)}]}},
            ]
            response_out = Path(td) / "host-agent-response.json"
            fake_results = [
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelopes[0]), ""),
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelopes[1]), ""),
            ]
            with patch.object(bridge, "_run_command", side_effect=fake_results) as run:
                audit = bridge.run_bridge(
                    review_dir, response_out=response_out,
                    agent_id="main", timeout=1, max_attempts=2,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                )
            self.assertEqual(audit["status"], "merged")
            self.assertEqual(audit["chunk_runs"][0]["attempt"], 2)
            self.assertIn("local response contract validation failed",
                          audit["chunk_runs"][0]["attempt_failures"][0])
            retry_ledger = audit["chunk_runs"][0]["semantic_retry_authorizations"]
            self.assertEqual(
                retry_ledger["status"],
                "accepted_after_provenance_and_contract_validation",
            )
            self.assertEqual(
                retry_ledger["paths"][0]["path"],
                "$.requirements[0].verification",
            )
            self.assertTrue(retry_ledger["paths"][0]["validator_records"])
            literal_projections = audit["chunk_runs"][0][
                "source_literal_whitespace_projections"
            ]
            self.assertEqual(len(literal_projections), 1)
            self.assertEqual(
                literal_projections[0]["source_projection_validation"]["status"],
                "matched",
            )
            accepted = json.loads(response_out.read_text(encoding="utf-8"))
            self.assertEqual(
                accepted["requirements"][0]["properties"]["text"], "正文使用宋体",
            )
            raw_retry = json.loads(Path(audit["chunk_runs"][0]["raw_response_path"]).read_text(
                encoding="utf-8",
            ))
            self.assertEqual(
                raw_retry["requirements"][0]["properties"]["text"], "正文 使用宋体",
            )
            second_prompt = Path(
                run.call_args_list[1].args[0][run.call_args_list[1].args[0].index("--message-file") + 1]
            ).read_text(encoding="utf-8")
            self.assertIn("local contract validation failed", second_prompt)

    def test_successful_retry_fails_closed_if_raw_response_artifact_is_missing(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            invalid = self._executable_response(chunk, invalid_verification=True)
            valid = self._executable_response(chunk)
            envelopes = [
                {"runId": "openclaw-run-raw-integrity-1", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(invalid)}]}},
                {"runId": "openclaw-run-raw-integrity-2", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(valid)}]}},
            ]
            response_out = Path(td) / "host-agent-response.json"
            fake_results = [
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelopes[0]), ""),
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelopes[1]), ""),
            ]
            real_atomic_write_text = bridge.atomic_write_text

            def omit_second_attempt_raw(path: Path, text: str, **kwargs: object) -> None:
                if ".attempt-02.raw.json" in Path(path).name:
                    return
                real_atomic_write_text(path, text, **kwargs)

            with patch.object(bridge, "atomic_write_text", side_effect=omit_second_attempt_raw):
                with patch.object(bridge, "_run_command", side_effect=fake_results) as run:
                    with self.assertRaisesRegex(
                        bridge.RetryRawArtifactIntegrityError,
                        "current raw response artifact is missing",
                    ):
                        bridge.run_bridge(
                            review_dir, response_out=response_out,
                            agent_id="main", timeout=1, max_attempts=3,
                            openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                        )

            self.assertEqual(run.call_count, 2)
            self.assertFalse(response_out.exists())
            failure = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertEqual(failure["terminal_error"]["type"], "RetryRawArtifactIntegrityError")
            self.assertTrue(failure["primary_error"]["records"])
            self.assertTrue(any(
                record.get("code") == "retry_raw_artifact_missing"
                for record in failure["chunk_lifecycle"][0]["structured_error_records"]
            ))
            self.assertTrue(any(
                item.get("kind") == "retry_raw_artifact_integrity"
                for item in failure["secondary_errors"]
            ))
            self.assertEqual(
                failure["chunk_lifecycle"][0]["attempts"][-1]["retry_authorizing_error_records"],
                failure["primary_error"]["records"],
            )

    def test_retryable_independent_review_incompleteness_uses_bounded_primary_retry(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            envelope = {
                "runId": "openclaw-run-obligation-retry", "status": "ok",
                "provider": "openai", "model": "gpt-5.6-luna",
                "result": {"payloads": [{"text": json.dumps(response)}]},
            }
            response_out = Path(td) / "host-agent-response.json"
            fake_results = [
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelope), ""),
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelope), ""),
            ]
            calls = 0

            def fail_independent_once(candidate, current_chunk, **kwargs):
                nonlocal calls
                calls += 1
                if calls == 1:
                    error = bridge.IndependentObligationReviewError(
                        "source obligation was not represented",
                    )
                    error.error_records = [{
                        "code": "independent_obligation_review_incomplete",
                        "clause_id": "C1", "baseline_classification": "informational",
                        "json_pointer": "$.clause_reviews[0].classification",
                        "missing_source_quotes": ["正文使用宋体"],
                        "candidate_response_sha256": bridge._response_sha256(candidate),
                    }]
                    error.independent_review_audit = {"status": "rejected", "audit_path": "rejected.json"}
                    error.retryable = True
                    raise error
                return self._fake_independent_review(candidate, current_chunk, **kwargs)

            self._independent_review_patch.stop()
            with patch.object(bridge, "_run_independent_obligation_coverage_review", side_effect=fail_independent_once), \
                    patch.object(bridge, "_run_command", side_effect=fake_results) as run:
                audit = bridge.run_bridge(
                    review_dir, response_out=response_out,
                    agent_id="main", timeout=1, max_attempts=2,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                )

            self.assertEqual(audit["status"], "merged")
            self.assertEqual(run.call_count, 2)
            self.assertEqual(calls, 2)
            self.assertEqual(audit["chunk_lifecycle"][0]["attempts"][0]["status"], "retrying")
            self.assertEqual(
                audit["chunk_lifecycle"][0]["attempts"][0]["error_records"][0]["code"],
                "independent_obligation_review_incomplete",
            )
    def test_retry_cannot_hide_a_semantic_classification_change(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            invalid = self._executable_response(chunk, invalid_verification=True)
            semantic_change = self._response(chunk)
            envelopes = [
                {"runId": "openclaw-run-semantic-retry-1", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(invalid)}]}},
                {"runId": "openclaw-run-semantic-retry-2", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(semantic_change)}]}},
            ]
            response_out = Path(td) / "host-agent-response.json"
            fake_results = [
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelopes[0]), ""),
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelopes[1]), ""),
            ]
            with patch.object(bridge, "_run_command", side_effect=fake_results):
                with self.assertRaisesRegex(ValueError, "semantic re-review"):
                    bridge.run_bridge(
                        review_dir, response_out=response_out,
                        agent_id="main", timeout=1, max_attempts=2,
                        openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                    )
            failure = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertFalse(failure["merged_response_written"])
            self.assertTrue(failure["structured_error_records"])
            self.assertEqual(
                failure["structured_error_records"][-1]["code"],
                "semantic_retry_change",
            )
            # The terminal structured record retains the independently
            # detected raw-to-raw semantic drift; the original contract error
            # remains separately attached as the primary retry authorizer.
            self.assertEqual(failure["primary_error"]["code"], "contract_validation_error")
            self.assertEqual(failure["primary_error"]["chunk_index"], 1)
            self.assertTrue(failure["primary_error"]["message"])
            self.assertTrue(any(
                record.get("code") == "schema_contract_violation"
                for record in failure["primary_error"]["records"]
            ))
            self.assertIn("semantic re-review", failure["terminal_error"]["message"])
            self.assertEqual(failure["secondary_errors"][0]["kind"], "retry_semantic_drift")
            self.assertEqual(failure["secondary_errors"][0]["chunk_index"], 1)
            lifecycle = failure["chunk_lifecycle"][0]
            self.assertTrue(lifecycle["secondary_retry_drift_failures"])
            self.assertIn(
                "semantic re-review",
                lifecycle["secondary_retry_drift_failures"][-1]["error"],
            )

    def test_bridge_binds_missing_native_provenance_and_preserves_raw_response(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            del response["provenance"]
            envelope = {
                "runId": "openclaw-run-provenance-bind",
                "status": "ok",
                "provider": "openai", "model": "gpt-5.6-luna",
                "result": {"payloads": [{"text": json.dumps(response)}]},
            }
            response_out = Path(td) / "host-agent-response.json"
            fake = subprocess.CompletedProcess(
                ["openclaw"], 0, json.dumps(envelope), "",
            )
            with patch.object(bridge, "_run_command", return_value=fake):
                audit = bridge.run_bridge(
                    review_dir, response_out=response_out,
                    agent_id="main", timeout=1, max_attempts=1,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                )
            self.assertTrue(response_out.exists())
            raw_responses = list(review_dir.glob("*.raw.json"))
            self.assertEqual(len(raw_responses), 1)
            raw = json.loads(raw_responses[0].read_text(encoding="utf-8"))
            self.assertNotIn("provenance", raw)
            merged = json.loads(response_out.read_text(encoding="utf-8"))
            full_request = json.loads(
                (review_dir / "llm-request.json").read_text(encoding="utf-8")
            )
            self.assertEqual(merged["provenance"], full_request["provenance"])
            self.assertEqual(
                audit["chunk_runs"][0]["provenance_binding"],
                "bridge_generated",
            )
            self.assertEqual(
                audit["chunk_runs"][0]["provenance_mismatch_fields"],
                [],
            )

    def test_bridge_rejects_conflicting_native_provenance_without_rebinding(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            response["provenance"] = dict(chunk["provenance"])
            response["provenance"]["request_sha256"] = "0" * 64
            envelope = {
                "runId": "openclaw-run-provenance-conflict",
                "status": "ok",
                "provider": "openai", "model": "gpt-5.6-luna",
                "result": {"payloads": [{"text": json.dumps(response)}]},
            }
            response_out = Path(td) / "host-agent-response.json"
            fake = subprocess.CompletedProcess(
                ["openclaw"], 0, json.dumps(envelope), "",
            )
            with patch.object(bridge, "_run_command", return_value=fake):
                with self.assertRaisesRegex(bridge.HostAgentProvenanceMismatch, "provenance conflict"):
                    bridge.run_bridge(
                        review_dir, response_out=response_out,
                        agent_id="main", timeout=1, max_attempts=2,
                        openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                    )
            self.assertFalse(response_out.exists())
            audit = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertEqual(audit["error_type"], "HostAgentProvenanceMismatch")
            self.assertFalse(audit["merged_response_written"])

    def test_semantic_retry_projection_rebinds_current_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            directory = Path(td) / "requirements"
            clauses = [{
                "id": "C1", "text": "图题应简明并置于图序之后", "evidence_ids": ["E1"],
                "source_kind": "paragraph", "location": {}, "part_index": 0,
            }]
            evidence = {
                "evidence": [{"id": "E1", "text": clauses[0]["text"], "kind": "paragraph"}],
            }
            clauses = self._bind_test_source_spans(clauses, evidence)
            request = engine.build_llm_request(
                [], clauses, evidence, {}, "full", contract_version="3.0",
                runtime_context={"code_fingerprint_sha256": "f" * 64},
            )
            request = attach_request_provenance(
                request, source_sha256="a" * 64, evidence_doc=evidence,
                clauses=clauses, run_id="run-v3-provenance-retry",
            )
            engine.prepare_host_agent_review_packets(
                request, clauses, evidence, "a" * 64, directory, chunk_size=1,
            )
            chunk = json.loads(
                (directory / "llm-request-chunks.json").read_text(encoding="utf-8")
            )[0]

            def review(classification: str) -> dict:
                return {
                    "clause_id": "C1",
                    "classification": classification,
                    "obligations": [
                        {"id": "caption_after_number", "status": "covered", "reason": "位置已由 requirement 表示。"},
                        {"id": "concise_caption", "status": "unverifiable", "reason": "简明性需要人工判断。"},
                    ],
                    "reason": "当前证据已审查。",
                    "normative_basis": "explicit_normative_text",
                }

            first = {
                "contract_version": "3.0",
                "requirements": [{
                    "role": "figure_caption",
                    "properties": {"position": "below"},
                    "clause_ids": ["C1"],
                    "evidence_ids": ["E1"],
                    "confidence": 0.9,
                    "reason": "证据支持图题位置。",
                    "verification": {"mode": "static_docx", "checks": ["检查图题位置。"]},
                }],
                "clause_reviews": [review("executable")],
                "unsupported_items": [],
                "reported_conflicts": [],
            }
            second = {
                "contract_version": "3.0",
                "requirements": [],
                "clause_reviews": [review("unverifiable")],
                "unsupported_items": [],
                "reported_conflicts": [],
            }
            envelopes = [
                {"runId": "openclaw-v3-provenance-retry-1", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(first, ensure_ascii=False)}]}},
                {"runId": "openclaw-v3-provenance-retry-2", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(second, ensure_ascii=False)}]}},
            ]
            response_out = Path(td) / "host-agent-response.json"
            fake_results = [
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelopes[0]), ""),
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelopes[1]), ""),
            ]
            with patch.object(bridge, "_run_command", side_effect=fake_results):
                audit = bridge.run_bridge(
                    directory, response_out=response_out,
                    agent_id="main", timeout=1, max_attempts=2,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                )

            merged = json.loads(response_out.read_text(encoding="utf-8"))
            full_request = json.loads(
                (directory / "llm-request.json").read_text(encoding="utf-8")
            )
            self.assertEqual(merged["provenance"], full_request["provenance"])
            self.assertEqual(merged["clause_reviews"][0]["classification"], "unverifiable")
            self.assertEqual(merged["requirements"], [])
            self.assertEqual(audit["status"], "merged")
            self.assertEqual(audit["chunk_runs"][0]["provenance_binding"], "bridge_generated")

    def test_bridge_rejects_tampered_chunk_before_starting_native_agent(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, _chunk = self._packet(Path(td) / "requirements")
            chunks_path = review_dir / "llm-request-chunks.json"
            chunks = json.loads(chunks_path.read_text(encoding="utf-8"))
            chunks[0]["clauses"][0]["text"] = "被篡改的正文"
            chunks_path.write_text(json.dumps(chunks, ensure_ascii=False), encoding="utf-8")
            with patch.object(bridge, "_run_command") as run:
                with self.assertRaisesRegex(ValueError, "source projection"):
                    bridge.run_bridge(
                        review_dir, response_out=Path(td) / "host-agent-response.json",
                        agent_id="main", timeout=1, max_attempts=1,
                        openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                    )
            run.assert_not_called()

    def test_bridge_requires_declared_host_runtime(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, _chunk = self._packet(Path(td) / "requirements")
            with patch.dict(os.environ, {}, clear=True):
                with self.assertRaisesRegex(RuntimeError, "THESIS_FORGE_HOST_RUNTIME"):
                    bridge.run_bridge(review_dir, openclaw_bin="openclaw")

    def test_bridge_refuses_host_runtime_mismatch(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, _chunk = self._packet(Path(td) / "requirements")
            with self.assertRaisesRegex(RuntimeError, "host runtime mismatch"):
                bridge.run_bridge(
                    review_dir, host_runtime="codex", openclaw_bin="openclaw",
                )

    def test_bridge_refuses_recent_parent_session_guess(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, _chunk = self._packet(Path(td) / "requirements")
            with self.assertRaisesRegex(RuntimeError, "refusing to select a recent"):
                bridge.run_bridge(
                    review_dir, inherit_parent_model=True,
                    openclaw_bin="openclaw",
                )

    def test_bridge_refuses_unbound_gateway_default(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, _chunk = self._packet(Path(td) / "requirements")
            with self.assertRaisesRegex(RuntimeError, "explicit model route"):
                bridge.run_bridge(review_dir, openclaw_bin="openclaw")

    def test_failed_bridge_persists_failure_audit_without_merged_response(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            response["clause_reviews"][0]["requirement_indexes"] = [0]
            envelope = {
                "runId": "openclaw-run-final-contract-failure",
                "status": "ok",
                "provider": "openai", "model": "gpt-5.6-luna",
                "result": {"payloads": [{"text": json.dumps(response)}]},
            }
            response_out = Path(td) / "host-agent-response.json"
            fake = subprocess.CompletedProcess(
                ["openclaw"], 0, json.dumps(envelope), "",
            )
            with patch.object(bridge, "_run_command", return_value=fake):
                with self.assertRaises(ValueError):
                    bridge.run_bridge(
                        review_dir, response_out=response_out,
                        agent_id="main", timeout=1, max_attempts=1,
                        openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                    )
            failure = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertEqual(failure["status"], "failed")
            self.assertFalse(failure["merged_response_written"])
            self.assertFalse(response_out.exists())
            self.assertEqual(len(failure["chunk_lifecycle"]), 1)
            self.assertEqual(failure["chunk_lifecycle"][0]["status"], "failed")
            self.assertEqual(failure["chunk_lifecycle"][0]["attempts"][0]["status"], "failed")

    def test_bridge_calls_openclaw_without_delivery_and_merges_fresh_response(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            envelope = {
                "runId": "openclaw-run-test",
                "status": "ok",
                "provider": "openai", "model": "gpt-5.6-luna",
                "result": {"payloads": [{"text": json.dumps(response, ensure_ascii=False)}]},
            }
            response_out = Path(td) / "host-agent-response.json"
            fake = subprocess.CompletedProcess(
                ["openclaw"], 0, json.dumps(envelope), "",
            )
            with patch.object(bridge, "_run_command", return_value=fake) as run:
                audit = bridge.run_bridge(
                    review_dir, response_out=response_out,
                    agent_id="main", timeout=1, openclaw_bin="openclaw",
                    model="openai/gpt-5.6-luna", runner="gateway",
                )
            self.assertEqual(audit["status"], "merged")
            self.assertEqual(audit["chunk_count"], 1)
            self.assertEqual(json.loads(response_out.read_text(encoding="utf-8"))["contract_version"], "2.1")
            command = run.call_args.args[0]
            self.assertIn("--session-key", command)
            self.assertIn("--message-file", command)
            self.assertNotIn("--deliver", command)
            self.assertTrue((review_dir / "host-agent-run.json").is_file())

    def test_bridge_calls_native_codex_without_openclaw_route_or_session(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            events = "\n".join([
                json.dumps({"type": "thread.started", "thread_id": "codex-thread"}),
                json.dumps({
                    "type": "item.completed",
                    "item": {"type": "agent_message", "text": json.dumps(response)},
                }),
                json.dumps({"type": "turn.completed", "usage": {"output_tokens": 1}}),
            ])
            response_out = Path(td) / "host-agent-response.json"
            fake = subprocess.CompletedProcess(["codex", "exec"], 0, events, "")
            with patch.dict(os.environ, {"THESIS_FORGE_HOST_RUNTIME": "codex"}):
                with patch.object(bridge, "_run_command", return_value=fake) as run:
                    audit = bridge.run_bridge(
                        review_dir, response_out=response_out,
                        timeout=1, max_attempts=1,
                        codex_bin=sys.executable,
                        allow_prompt_only=True,
                    )
            command = run.call_args.args[0]
            self.assertEqual(command[1], "exec")
            self.assertIn("--json", command)
            self.assertIn("--sandbox", command)
            self.assertEqual(command[command.index("--model") + 1], "gpt-6-luna")
            self.assertEqual(command[-1], "-")
            prompt_files = list(review_dir.rglob("prompt-0001-attempt-01.txt"))
            self.assertEqual(len(prompt_files), 1)
            self.assertEqual(run.call_args.kwargs["input_text"], prompt_files[0].read_text(encoding="utf-8"))
            self.assertNotIn("openclaw", " ".join(command).lower())
            self.assertNotIn("--session-key", command)
            self.assertEqual(audit["adapter_id"], "codex")
            self.assertEqual(audit["model"], "gpt-6-luna")
            self.assertEqual(audit["model_source"], "project-default-native-codex-model")
            self.assertEqual(audit["route_visibility"], "unobservable")
            self.assertEqual(audit["observed_routes"], ["unobservable"])
            self.assertEqual(audit["chunk_runs"][0]["actual_route"], "unobservable")
            self.assertEqual(json.loads(response_out.read_text(encoding="utf-8"))["contract_version"], "2.1")

    def test_bridge_refuses_prompt_only_codex_by_default(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, _chunk = self._packet(Path(td) / "requirements")
            with patch.dict(os.environ, {"THESIS_FORGE_HOST_RUNTIME": "codex"}):
                with patch.object(bridge.codex_adapter, "probe_capabilities", return_value={
                    "output_schema_supported": False,
                    "structured_output_mode": "prompt_only",
                }):
                    with self.assertRaisesRegex(RuntimeError, "refusing prompt-only"):
                        bridge.run_bridge(
                            review_dir,
                            host_runtime="codex",
                            codex_bin=sys.executable,
                            max_attempts=1,
                        )

    def test_cli_does_not_inject_parent_inheritance_for_codex(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            with patch.dict(os.environ, {"THESIS_FORGE_HOST_RUNTIME": "codex"}):
                with patch.object(bridge, "run_bridge", return_value={"status": "mocked"}) as run:
                    code = bridge.main([
                        td,
                        "--host-runtime", "codex",
                    ])
            self.assertEqual(code, 0)
            self.assertFalse(run.call_args.kwargs["inherit_parent_model"])
            self.assertIsNone(run.call_args.kwargs["codex_model"])

    def test_resolve_parent_model_copies_provider_and_model_override(self) -> None:
        parent_key = "agent:main:telegram:direct:chat:thread:38479"
        sessions = {
            "sessions": [{
                "key": parent_key,
                "agentId": "main",
                "kind": "direct",
                "modelOverride": "gpt-5.6-luna",
                "providerOverride": "openai",
                "updatedAt": 123,
            }],
        }
        fake = subprocess.CompletedProcess(
            ["openclaw", "sessions"], 0, json.dumps(sessions), "",
        )
        with patch.object(bridge, "_run_command", return_value=fake) as run:
            resolved = bridge.resolve_parent_model(
                "openclaw", agent_id="main", parent_session_key=parent_key,
            )
        self.assertEqual(resolved["model"], "openai/gpt-5.6-luna")
        self.assertEqual(resolved["source"], "parent-session-override")
        self.assertEqual(resolved["parent_model_override"], "gpt-5.6-luna")
        self.assertEqual(resolved["parent_provider_override"], "openai")
        self.assertEqual(run.call_args.args[0][0:3], ["openclaw", "sessions", "--json"])
        command = run.call_args.args[0]
        self.assertEqual(command[command.index("--agent") + 1], "main")

    def test_bridge_passes_inherited_parent_route_to_each_chunk(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            parent_key = "agent:main:telegram:direct:chat:thread:38479"
            sessions = {
                "sessions": [{
                    "key": parent_key,
                    "agentId": "main",
                    "kind": "direct",
                    "modelOverride": "gpt-5.6-luna",
                    "providerOverride": "openai",
                    "updatedAt": 123,
                }],
            }
            envelope = {
                "runId": "openclaw-run-inherit-test",
                "status": "ok",
                "result": {
                    "payloads": [{"text": json.dumps(response)}],
                    "meta": {"agentMeta": {
                        "provider": "openai", "model": "gpt-5.6-luna",
                    }},
                },
            }
            session_result = subprocess.CompletedProcess(
                ["openclaw", "sessions"], 0, json.dumps(sessions), "",
            )
            agent_result = subprocess.CompletedProcess(
                ["openclaw", "agent"], 0, json.dumps(envelope), "",
            )
            response_out = Path(td) / "host-agent-response.json"

            def fake_run(command, **kwargs):
                if "sessions" in command:
                    return session_result
                return agent_result

            with patch.object(bridge, "_run_command", side_effect=fake_run) as run:
                audit = bridge.run_bridge(
                    review_dir, response_out=response_out,
                    agent_id="main", timeout=1, openclaw_bin="openclaw",
                    inherit_parent_model=True, parent_session_key=parent_key,
                    auth_env_only=True,
                )
            command = next(
                call.args[0] for call in run.call_args_list if "agent" in call.args[0]
            )
            self.assertEqual(command[command.index("--model") + 1], "openai/gpt-5.6-luna")
            self.assertIn("--auth-env-only", command)
            self.assertEqual(audit["model"], "openai/gpt-5.6-luna")
            self.assertEqual(audit["model_source"], "parent-session-override")
            self.assertEqual(audit["chunk_runs"][0]["actual_route"], "openai/gpt-5.6-luna")

    def test_bridge_can_use_gateway_runner_without_delivery(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            envelope = {
                "runId": "openclaw-run-gateway-test",
                "status": "ok",
                "result": {
                    "payloads": [{"text": json.dumps(response)}],
                    "meta": {"agentMeta": {
                        "provider": "openai", "model": "gpt-5.6-luna",
                    }},
                },
            }
            response_out = Path(td) / "host-agent-response.json"
            fake = subprocess.CompletedProcess(
                ["openclaw", "agent"], 0, json.dumps(envelope), "",
            )
            with patch.object(bridge, "_run_command", return_value=fake) as run:
                audit = bridge.run_bridge(
                    review_dir, response_out=response_out,
                    agent_id="main", timeout=1, openclaw_bin="openclaw",
                    model="openai/gpt-5.6-luna", runner="gateway",
                )
            command = run.call_args.args[0]
            self.assertIn("--agent", command)
            self.assertIn("--session-key", command)
            self.assertIn("--model", command)
            self.assertNotIn("--deliver", command)
            session_key = command[command.index("--session-key") + 1]
            self.assertTrue(session_key.startswith("agent:main:thesis-host-agent:"))
            self.assertEqual(audit["runner"], "gateway")
            self.assertEqual(audit["chunk_runs"][0]["runner"], "gateway-agent")

    def test_resolve_parent_model_uses_effective_route_without_override(self) -> None:
        parent_key = "agent:main:telegram:direct:chat:thread:38479"
        sessions = {
            "sessions": [{
                "key": parent_key,
                "agentId": "main",
                "kind": "direct",
                "modelProvider": "sub2api",
                "model": "gpt-5.6-sol",
                "updatedAt": 123,
            }],
        }
        fake = subprocess.CompletedProcess(
            ["openclaw", "sessions"], 0, json.dumps(sessions), "",
        )
        with patch.object(bridge, "_run_command", return_value=fake):
            resolved = bridge.resolve_parent_model(
                "openclaw", agent_id="main", parent_session_key=parent_key,
            )
        self.assertEqual(resolved["model"], "sub2api/gpt-5.6-sol")
        self.assertEqual(resolved["source"], "parent-session-effective")
        self.assertEqual(resolved["parent_effective_provider"], "sub2api")
        self.assertEqual(resolved["parent_effective_model"], "gpt-5.6-sol")

    def test_resolve_parent_model_refuses_unresolved_effective_route(self) -> None:
        parent_key = "agent:main:telegram:direct:chat:thread:38479"
        sessions = {
            "sessions": [{
                "key": parent_key,
                "agentId": "main",
                "kind": "direct",
                "updatedAt": 123,
            }],
        }
        fake = subprocess.CompletedProcess(
            ["openclaw", "sessions"], 0, json.dumps(sessions), "",
        )
        with patch.object(bridge, "_run_command", return_value=fake):
            with self.assertRaisesRegex(ValueError, "refusing to use the gateway default"):
                bridge.resolve_parent_model(
                    "openclaw", agent_id="main", parent_session_key=parent_key,
                )

    def test_bridge_fails_closed_on_child_route_mismatch_without_retry(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            parent_key = "agent:main:telegram:direct:chat:thread:38479"
            sessions = {
                "sessions": [{
                    "key": parent_key,
                    "agentId": "main",
                    "kind": "direct",
                    "modelProvider": "openai",
                    "model": "gpt-5.6-luna",
                    "updatedAt": 123,
                }],
            }
            envelope = {
                "runId": "openclaw-run-route-mismatch",
                "status": "ok",
                "provider": "sub2api",
                "model": "gpt-5.6-sol",
                "payloads": [{"text": json.dumps(response)}],
            }
            session_result = subprocess.CompletedProcess(
                ["openclaw", "sessions"], 0, json.dumps(sessions), "",
            )
            agent_result = subprocess.CompletedProcess(
                ["openclaw", "agent", "exec"], 0, json.dumps(envelope), "",
            )
            calls = []

            def fake_run(command, **kwargs):
                calls.append(command)
                return session_result if "sessions" in command else agent_result

            with patch.object(bridge, "_run_command", side_effect=fake_run) as run:
                with self.assertRaises(bridge.HostAgentRouteMismatch):
                    bridge.run_bridge(
                        review_dir,
                        response_out=Path(td) / "host-agent-response.json",
                        agent_id="main", timeout=1, max_attempts=2,
                        openclaw_bin="openclaw",
                        inherit_parent_model=True,
                        parent_session_key=parent_key,
                    )
            # One owned session-discovery call plus one child call; a route
            # mismatch must still not trigger a second child attempt.
            self.assertEqual(len(calls), 2)
            self.assertIn("sessions", calls[0])
            self.assertIn("agent", calls[1])
            self.assertFalse((Path(td) / "host-agent-response.json").exists())

    def test_bridge_refuses_existing_chunk_response_instead_of_reusing_it(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response_name = json.loads(
                (review_dir / "host-agent-review-manifest.json").read_text(encoding="utf-8")
            )["response_files"][0]
            (review_dir / response_name).write_text("{}\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "refusing to reuse an existing chunk response"):
                bridge.run_bridge(
                    review_dir, response_out=Path(td) / "response.json", timeout=1,
                    model="openai/gpt-5.6-luna",
                )

    def test_bridge_retries_rejected_json_in_a_new_session(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            envelope = {
                "runId": "openclaw-run-retry-test",
                "status": "ok",
                "provider": "openai", "model": "gpt-5.6-luna",
                "result": {"payloads": [{"text": json.dumps(response, ensure_ascii=False)}]},
            }
            response_out = Path(td) / "host-agent-response.json"
            fake_results = [
                subprocess.CompletedProcess(["openclaw"], 0, "{broken", ""),
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(envelope), ""),
            ]
            with patch.object(bridge, "_run_command", side_effect=fake_results) as run:
                audit = bridge.run_bridge(
                    review_dir, response_out=response_out,
                    agent_id="main", timeout=1, max_attempts=2,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                    runner="gateway",
                )
            self.assertEqual(audit["status"], "merged")
            self.assertEqual(audit["chunk_runs"][0]["attempt"], 2)
            self.assertEqual(len(audit["chunk_runs"][0]["attempt_failures"]), 1)
            self.assertEqual(
                audit["chunk_runs"][0]["retry_stage_comparison"]["status"],
                "no_prior_semantic_response",
            )
            self.assertRegex(
                audit["chunk_runs"][0]["retry_stage_comparison"]["parent_raw_envelope_sha256"],
                r"^[0-9a-f]{64}$",
            )
            self.assertEqual(run.call_count, 2)
            first_command = run.call_args_list[0].args[0]
            second_command = run.call_args_list[1].args[0]
            first_key = first_command[first_command.index("--session-key") + 1]
            second_key = second_command[second_command.index("--session-key") + 1]
            self.assertNotEqual(first_key, second_key)
            self.assertIn("attempt-01", first_key)
            self.assertIn("attempt-02", second_key)

    def test_retry_compares_latest_decoded_response_across_intervening_parse_failure(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            envelopes = [
                {
                    "runId": "openclaw-run-retry-chain-1", "status": "ok",
                    "provider": "openai", "model": "gpt-5.6-luna",
                    "result": {"payloads": [{"text": json.dumps(response)}]},
                },
                "{broken",
                {
                    "runId": "openclaw-run-retry-chain-3", "status": "ok",
                    "provider": "openai", "model": "gpt-5.6-luna",
                    "result": {"payloads": [{"text": json.dumps(response)}]},
                },
            ]
            fake_results = [
                subprocess.CompletedProcess(
                    ["openclaw"], 0,
                    json.dumps(item) if isinstance(item, dict) else item,
                    "",
                )
                for item in envelopes
            ]
            response_out = Path(td) / "host-agent-response.json"
            review_calls = 0

            def fail_independent_once(candidate, current_chunk, **kwargs):
                nonlocal review_calls
                review_calls += 1
                if review_calls == 1:
                    error = bridge.IndependentObligationReviewError(
                        "source obligation was not represented",
                    )
                    error.error_records = [{
                        "code": "independent_obligation_review_incomplete",
                        "clause_id": "C1", "baseline_classification": "informational",
                        "json_pointer": "$.clause_reviews[0].classification",
                        "missing_source_quotes": ["正文使用宋体"],
                        "candidate_response_sha256": bridge._response_sha256(candidate),
                    }]
                    error.independent_review_audit = {
                        "status": "rejected", "audit_path": "rejected.json",
                    }
                    error.retryable = True
                    raise error
                return self._fake_independent_review(candidate, current_chunk, **kwargs)

            self._independent_review_patch.stop()
            with patch.object(
                bridge, "_run_independent_obligation_coverage_review",
                side_effect=fail_independent_once,
            ), patch.object(bridge, "_run_command", side_effect=fake_results) as run:
                audit = bridge.run_bridge(
                    review_dir, response_out=response_out,
                    agent_id="main", timeout=1, max_attempts=3,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                )

            self.assertEqual(audit["status"], "merged")
            self.assertEqual(run.call_count, 3)
            self.assertEqual(review_calls, 2)
            successful_attempt = audit["chunk_runs"][0]
            self.assertEqual(successful_attempt["attempt"], 3)
            comparison = successful_attempt["retry_stage_comparison"]
            self.assertEqual(comparison["raw_parent_attempt"], 1)
            self.assertEqual(comparison["raw_candidate_attempt"], 3)
            self.assertEqual(comparison["authorization_parent_attempt"], 1)
            first_blocker = audit["chunk_lifecycle"][0]["attempts"][0]["error_records"]
            self.assertEqual(
                comparison["authorization_error_records_sha256"],
                bridge._response_sha256(first_blocker),
            )
            self.assertEqual(len(comparison["intervening_attempt_receipts"]), 1)
            self.assertEqual(comparison["intervening_attempt_receipts"][0]["attempt"], 2)
            self.assertEqual(
                comparison["intervening_attempt_receipts"][0]["kind"],
                "no_semantic_response",
            )

    def test_retry_will_not_continue_after_a_prior_decoded_raw_artifact_is_removed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            invalid = self._executable_response(chunk, invalid_verification=True)
            valid = self._response(chunk)
            envelopes = [
                {"runId": "openclaw-run-raw-chain-1", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(invalid)}]}},
                "{broken",
                {"runId": "openclaw-run-raw-chain-3", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(valid)}]}},
            ]
            response_name = json.loads(
                (review_dir / "host-agent-review-manifest.json").read_text(encoding="utf-8")
            )["response_files"][0]
            response_path = review_dir / response_name
            first_raw_path = response_path.with_name(
                f"{response_path.stem}.attempt-01.raw{response_path.suffix}"
            )
            fake_results = [
                subprocess.CompletedProcess(
                    ["openclaw"], 0,
                    json.dumps(item) if isinstance(item, dict) else item,
                    "",
                )
                for item in envelopes
            ]
            real_parse_result = bridge.openclaw_adapter.parse_result
            parse_calls = 0

            def remove_parent_after_retry_envelope_is_persisted(stdout: str):
                nonlocal parse_calls
                parse_calls += 1
                if parse_calls == 2:
                    first_raw_path.unlink()
                return real_parse_result(stdout)

            with patch.object(
                bridge.openclaw_adapter, "parse_result",
                side_effect=remove_parent_after_retry_envelope_is_persisted,
            ):
                with patch.object(bridge, "_run_command", side_effect=fake_results) as run:
                    with self.assertRaisesRegex(
                        bridge.RetryRawArtifactIntegrityError,
                        "decoded raw response artifact is missing",
                    ):
                        bridge.run_bridge(
                            review_dir,
                            response_out=Path(td) / "host-agent-response.json",
                            agent_id="main", timeout=1, max_attempts=3,
                            openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                        )

            self.assertEqual(run.call_count, 2)
            failure = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertEqual(failure["terminal_error"]["type"], "RetryRawArtifactIntegrityError")
            self.assertTrue(any(
                record.get("code") == "retry_raw_artifact_receipt_mismatch"
                for record in failure["chunk_lifecycle"][0]["structured_error_records"]
            ))
            self.assertTrue(any(
                record.get("code") == "schema_contract_violation"
                for record in failure["primary_error"]["records"]
            ), repr(failure["primary_error"]["records"]))

    def test_retry_will_not_dispatch_after_a_prior_decoded_raw_artifact_is_modified(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            invalid = self._executable_response(chunk, invalid_verification=True)
            valid = self._response(chunk)
            envelopes = [
                {"runId": "openclaw-run-raw-modified-1", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(invalid)}]}},
                {"runId": "openclaw-run-raw-modified-2", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(valid)}]}},
            ]
            response_name = json.loads(
                (review_dir / "host-agent-review-manifest.json").read_text(encoding="utf-8")
            )["response_files"][0]
            response_path = review_dir / response_name
            first_raw_path = response_path.with_name(
                f"{response_path.stem}.attempt-01.raw{response_path.suffix}"
            )
            fake_results = [
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(item), "")
                for item in envelopes
            ]
            real_validate_receipt = bridge._validate_retry_attempt_artifact
            validation_calls = 0

            def mutate_raw_before_retry_receipt_validation(
                path: Path, attempt_number: int, attempt_record: dict,
            ):
                nonlocal validation_calls
                validation_calls += 1
                if validation_calls == 1:
                    raw = json.loads(first_raw_path.read_text(encoding="utf-8"))
                    raw["clause_reviews"][0]["reason"] = "tampered after initial failure"
                    first_raw_path.write_text(json.dumps(raw), encoding="utf-8")
                return real_validate_receipt(path, attempt_number, attempt_record)

            review_calls = 0

            def count_review_calls(candidate, current_chunk, **kwargs):
                nonlocal review_calls
                review_calls += 1
                return self._fake_independent_review(candidate, current_chunk, **kwargs)

            self._independent_review_patch.stop()
            with patch.object(
                bridge, "_validate_retry_attempt_artifact",
                side_effect=mutate_raw_before_retry_receipt_validation,
            ), patch.object(
                bridge, "_run_independent_obligation_coverage_review",
                side_effect=count_review_calls,
            ):
                with patch.object(bridge, "_run_command", side_effect=fake_results) as run:
                    with self.assertRaisesRegex(
                        bridge.RetryRawArtifactIntegrityError,
                        "decoded raw response artifact hash differs from its receipt",
                    ):
                        bridge.run_bridge(
                            review_dir,
                            response_out=Path(td) / "host-agent-response.json",
                            agent_id="main", timeout=1, max_attempts=2,
                            openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                        )

            self.assertEqual(run.call_count, 1, "modified parent evidence must stop before retry dispatch")
            self.assertEqual(review_calls, 0, "tampered parent evidence must not reach independent review")
            failure = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertEqual(failure["terminal_error"]["type"], "RetryRawArtifactIntegrityError")
            self.assertTrue(any(
                record.get("code") == "retry_raw_artifact_receipt_mismatch"
                for record in failure["chunk_lifecycle"][0]["structured_error_records"]
            ))

    def test_retry_will_not_continue_after_a_prior_validated_candidate_is_tampered(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            valid = self._executable_response(chunk)
            envelopes = [
                {"runId": "openclaw-run-candidate-chain-1", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(valid)}]}},
                {"runId": "openclaw-run-candidate-chain-2", "status": "ok",
                 "provider": "openai", "model": "gpt-5.6-luna",
                 "result": {"payloads": [{"text": json.dumps(valid)}]}},
            ]
            response_name = json.loads(
                (review_dir / "host-agent-review-manifest.json").read_text(encoding="utf-8")
            )["response_files"][0]
            response_path = review_dir / response_name
            first_candidate_path = response_path.with_name(
                f"{response_path.stem}.attempt-01{response_path.suffix}"
            )
            fake_results = [
                subprocess.CompletedProcess(["openclaw"], 0, json.dumps(item), "")
                for item in envelopes
            ]
            real_validate_receipt = bridge._validate_retry_attempt_artifact
            validation_calls = 0

            def mutate_candidate_before_receipt_validation(
                path: Path, attempt_number: int, attempt_record: dict,
            ):
                nonlocal validation_calls
                validation_calls += 1
                if validation_calls == 1:
                    first_candidate_path.write_text("tampered candidate", encoding="utf-8")
                return real_validate_receipt(path, attempt_number, attempt_record)

            review_calls = 0

            def fail_first_independent_review(candidate, current_chunk, **kwargs):
                nonlocal review_calls
                review_calls += 1
                if review_calls == 1:
                    error = bridge.IndependentObligationReviewError(
                        "source obligation was not represented",
                    )
                    error.error_records = [{
                        "code": "independent_obligation_review_incomplete",
                        "clause_id": "C1",
                        "json_pointer": "$.clause_reviews[0]",
                    }]
                    error.retryable = True
                    raise error
                return self._fake_independent_review(candidate, current_chunk, **kwargs)

            self._independent_review_patch.stop()
            with patch.object(
                bridge, "_validate_retry_attempt_artifact",
                side_effect=mutate_candidate_before_receipt_validation,
            ), patch.object(
                bridge, "_run_independent_obligation_coverage_review",
                side_effect=fail_first_independent_review,
            ):
                with patch.object(bridge, "_run_command", side_effect=fake_results) as run:
                    with self.assertRaisesRegex(
                        bridge.RetryRawArtifactIntegrityError,
                        "validated candidate artifact hash differs from its receipt",
                    ):
                        bridge.run_bridge(
                            review_dir,
                            response_out=Path(td) / "host-agent-response.json",
                            agent_id="main", timeout=1, max_attempts=2,
                            openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                        )

            self.assertEqual(run.call_count, 1)
            self.assertEqual(review_calls, 1)
            failure = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertEqual(failure["terminal_error"]["type"], "RetryRawArtifactIntegrityError")
            self.assertTrue(any(
                record.get("code") == "retry_candidate_artifact_receipt_mismatch"
                for record in failure["chunk_lifecycle"][0]["structured_error_records"]
            ))

    def test_retry_refuses_a_mutated_raw_envelope_from_a_parse_failure(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            valid = self._response(chunk)
            fake_results = [
                subprocess.CompletedProcess(["openclaw"], 0, "{broken", ""),
                subprocess.CompletedProcess(
                    ["openclaw"], 0,
                    json.dumps({
                        "runId": "openclaw-run-envelope-tamper", "status": "ok",
                        "provider": "openai", "model": "gpt-5.6-luna",
                        "result": {"payloads": [{"text": json.dumps(valid)}]},
                    }),
                    "",
                ),
            ]
            response_out = Path(td) / "host-agent-response.json"
            real_validate_receipt = bridge._validate_retry_attempt_artifact
            validation_calls = 0

            def mutate_envelope_before_receipt_validation(
                response_path: Path, attempt_number: int, attempt_record: dict,
            ):
                nonlocal validation_calls
                validation_calls += 1
                if validation_calls == 1:
                    envelope_path = response_path.with_name(
                        f"{response_path.stem}.attempt-01.raw-envelope.txt"
                    )
                    envelope_path.write_text("tampered after parse receipt", encoding="utf-8")
                return real_validate_receipt(response_path, attempt_number, attempt_record)

            with patch.object(
                bridge, "_validate_retry_attempt_artifact",
                side_effect=mutate_envelope_before_receipt_validation,
            ):
                with patch.object(bridge, "_run_command", side_effect=fake_results) as run:
                    with self.assertRaisesRegex(
                        bridge.RetryRawArtifactIntegrityError,
                        "raw envelope artifact hash differs from its receipt",
                    ):
                        bridge.run_bridge(
                            review_dir, response_out=response_out,
                            agent_id="main", timeout=1, max_attempts=2,
                            openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                        )

            self.assertEqual(run.call_count, 1)
            failure = json.loads((review_dir / "host-agent-run.json").read_text(encoding="utf-8"))
            self.assertEqual(failure["terminal_error"]["type"], "RetryRawArtifactIntegrityError")
            self.assertTrue(any(
                record.get("code") == "retry_raw_envelope_receipt_mismatch"
                for record in failure["chunk_lifecycle"][0]["structured_error_records"]
            ))

    def test_no_semantic_response_retry_receipt_requires_invocation_binding(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response_path = review_dir / "llm-response-chunk-0001.json"
            envelope_path = response_path.with_name(
                f"{response_path.stem}.attempt-01.raw-envelope.txt"
            )
            envelope_path.write_text("provider response that did not decode", encoding="utf-8")
            binding = bridge._retry_input_fingerprints(chunk)
            error_record = {
                "code": "host_response_parse_error",
                "raw_envelope_path": str(envelope_path.resolve()),
                "raw_envelope_sha256": bridge.sha256_file(envelope_path),
            }
            attempt_record = {
                "retry_input_fingerprints": binding,
                "no_semantic_response_invocation_fingerprints": copy.deepcopy(binding),
                "error_records": [error_record],
            }
            receipt = bridge._validate_retry_attempt_artifact(
                response_path, 1, attempt_record,
            )
            self.assertEqual(receipt["kind"], "no_semantic_response")

            mismatched = copy.deepcopy(attempt_record)
            mismatched["no_semantic_response_invocation_fingerprints"]["run_id"] = "stale-run"
            with self.assertRaisesRegex(
                bridge.RetryRawArtifactIntegrityError,
                "raw envelope receipt does not match the attempt invocation fingerprints",
            ):
                bridge._validate_retry_attempt_artifact(response_path, 1, mismatched)

            missing_binding = copy.deepcopy(attempt_record)
            del missing_binding["no_semantic_response_invocation_fingerprints"]
            with self.assertRaisesRegex(
                bridge.RetryRawArtifactIntegrityError,
                "raw envelope receipt does not match the attempt invocation fingerprints",
            ):
                bridge._validate_retry_attempt_artifact(response_path, 1, missing_binding)

    def test_run_command_kills_and_reaps_the_process_group_on_timeout(self) -> None:
        controller = bridge.RunController()
        timeout_result = subprocess.CompletedProcess(
            ["openclaw"], 124, "partial stdout", "partial stderr\n[process-timeout] command exceeded 7s and was terminated",
        )
        with patch.object(bridge, "run_process", return_value=timeout_result) as run:
            with self.assertRaises(subprocess.TimeoutExpired) as raised:
                bridge._run_command(["openclaw"], timeout=7, controller=controller)
        run.assert_called_once_with(
            ["openclaw"], cwd=bridge.ROOT, timeout=7, controller=controller,
        )
        self.assertEqual(raised.exception.output, "partial stdout")
        self.assertIn("partial stderr", str(raised.exception.stderr))

    def test_generic_validator_path_does_not_authorize_retry_value_change(self) -> None:
        previous = {
            "contract_version": "3.0", "requirements": [], "unsupported_items": [],
            "reported_conflicts": [],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "executable",
                "normative_basis": "explicit_normative_text",
            }],
        }
        current = copy.deepcopy(previous)
        current["clause_reviews"][0]["normative_basis"] = "template_structure"
        self.assertFalse(bridge._retry_changes_allowed(
            [{
                "code": "contract_validation_error",
                "json_pointer": "$.clause_reviews[0].normative_basis",
            }],
            ["$.clause_reviews[0].normative_basis"],
            contract_version="3.0", previous_response=previous,
            current_response=current,
        ))

    def test_source_fragment_retry_authorizes_only_exact_current_source_composition(self) -> None:
        title = "硕 士 学 位 论 文"
        degree = "（学术学位）"
        left_location = {"part": "document", "child_index": 4, "order": 2}
        right_location = {"part": "document", "child_index": 5, "order": 3}
        clauses = [
            {
                "id": "C1", "text": title, "evidence_ids": ["E1"],
                "location": left_location, "source_kind": "paragraph",
                "source_evidence_text": title,
                "source_span": {
                    "evidence_id": "E1", "start_offset": 0, "end_offset": len(title),
                    "text": title, "source_sha256": hashlib.sha256(title.encode()).hexdigest(),
                    "location": left_location,
                },
            },
            {
                "id": "C2", "text": degree, "evidence_ids": ["E2"],
                "location": right_location, "source_kind": "paragraph",
                "source_evidence_text": degree,
                "source_span": {
                    "evidence_id": "E2", "start_offset": 0, "end_offset": len(degree),
                    "text": degree, "source_sha256": hashlib.sha256(degree.encode()).hexdigest(),
                    "location": right_location,
                },
            },
        ]
        evidence_context = {
            "E1": {"id": "E1", "kind": "paragraph", "text": title, "location": left_location},
            "E2": {"id": "E2", "kind": "paragraph", "text": degree, "location": right_location},
        }
        chunk = {
            "provenance": {
                "run_id": "run-fragment-retry", "case_id": "case-fragment-retry",
                "source_sha256": "a" * 64, "clause_sha256": "b" * 64,
                "evidence_sha256": "c" * 64, "request_sha256": "d" * 64,
            },
            "case_id": "case-fragment-retry", "batch": {"index": 2},
            "runtime_context": {"code_fingerprint_sha256": "9" * 64},
            "response_schema": {"type": "object", "title": "source fragment retry"},
            "requirement_contract": {
                "role_properties_schema": {"title": {"$ref": "#/$defs/roleSpec"}},
            },
            "clauses": clauses, "evidence_context": evidence_context,
        }
        previous = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "role": "title", "properties": {"text": title + degree},
                "clause_ids": ["C1", "C2"], "evidence_ids": ["E1", "E2"],
                "reason": "The two adjacent source paragraphs form the title.",
            }],
            "clause_reviews": [], "unsupported_items": [], "reported_conflicts": [],
        }
        current = copy.deepcopy(previous)
        current["requirements"][0]["source_fragment_clause_ids"] = ["C1", "C2"]
        current["requirements"][0]["properties"]["text"] = title + "\n" + degree
        records = [{
            "code": "source_fragment_binding_violation",
            "json_pointer": "$.requirements[0].properties.text",
            "response_sha256": bridge._response_sha256(previous),
            "raw_error": "source_fragment_binding_required_for_cross_source_literal",
        }]
        authorization: list[dict] = []
        error, changed = bridge._retry_semantic_change_error(
            previous, current, records, contract_version="3.0", chunk=chunk,
            authorization_out=authorization,
        )
        self.assertIsNone(error)
        self.assertEqual(set(changed), {
            "$.requirements[0].source_fragment_clause_ids",
            "$.requirements[0].properties.text",
        })
        self.assertEqual(len(authorization), 2)
        self.assertTrue(all(
            item["rule_id"] == "v3_source_fragment_binding_selection"
            and item["source_binding_complete"]
            for item in authorization
        ))

        unsafe = []
        bad_text = copy.deepcopy(current)
        bad_text["requirements"][0]["properties"]["text"] = title + degree + "猜测"
        unsafe.append(bad_text)
        wrong_order = copy.deepcopy(current)
        wrong_order["requirements"][0]["source_fragment_clause_ids"] = ["C2", "C1"]
        unsafe.append(wrong_order)
        extra_reason = copy.deepcopy(current)
        extra_reason["requirements"][0]["reason"] = "rewritten reason"
        unsafe.append(extra_reason)
        for candidate in unsafe:
            with self.subTest(candidate=candidate):
                rejected, _ = bridge._retry_semantic_change_error(
                    previous, candidate, records, contract_version="3.0", chunk=chunk,
                )
                self.assertIsNotNone(rejected)
        stale_records = copy.deepcopy(records)
        stale_records[0]["response_sha256"] = "0" * 64
        stale_error, _ = bridge._retry_semantic_change_error(
            previous, current, stale_records, contract_version="3.0", chunk=chunk,
        )
        self.assertIsNotNone(stale_error)

    def test_source_fragment_retry_repairs_targeted_literal_with_unchanged_selector(self) -> None:
        # An arbitrary label, not a school/clause-ID special case. The semantic
        # slice excludes punctuation, while the fixed source label retains it.
        label = "档案编号"
        source = label + "："
        location = {"part": "document", "table_child_index": 7,
                    "row_index": 2, "col_index": 0, "paragraph_index": 0, "order": 9}
        clause = {
            "id": "CL-label", "text": label, "evidence_ids": ["EV-label"],
            "source_kind": "table_cell", "location": location,
            "source_evidence_text": source,
            "source_span": {
                "evidence_id": "EV-label", "start_offset": 0,
                "end_offset": len(label), "text": label,
                "source_sha256": hashlib.sha256(source.encode()).hexdigest(),
                "location": location,
            },
        }
        chunk = {
            "provenance": {
                "run_id": "literal-retry", "case_id": "arbitrary-case",
                "source_sha256": "a" * 64, "clause_sha256": "b" * 64,
                "evidence_sha256": "c" * 64, "request_sha256": "d" * 64,
            },
            "case_id": "arbitrary-case", "batch": {"index": 1},
            "runtime_context": {"code_fingerprint_sha256": "9" * 64},
            "response_schema": {"type": "object"},
            "clauses": [clause],
            "evidence_context": {"EV-label": {
                "id": "EV-label", "kind": "table_cell", "text": source,
                "location": location,
            }},
            "requirement_contract": {"role_properties_schema": {
                "cover_field_label": {"$ref": "#/$defs/roleSpec"},
            }},
        }
        previous = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "role": "cover_field_label", "properties": {"text": label},
                "clause_ids": ["CL-label"], "evidence_ids": ["EV-label"],
                "source_fragment_clause_ids": ["CL-label"],
                "reason": "Preserve the fixed label from the current table cell.",
            }],
            "clause_reviews": [], "unsupported_items": [], "reported_conflicts": [],
        }
        records = [{
            "code": "source_fragment_binding_violation",
            "json_pointer": "$.requirements[0].properties.text",
            "response_sha256": bridge._response_sha256(previous),
            "raw_error": "source_fragment_literal_conflict",
        }]
        for literal in (None, source):
            with self.subTest(literal=literal):
                current = copy.deepcopy(previous)
                current["requirements"][0]["properties"]["text"] = literal
                authorizations = []
                error, changed = bridge._retry_semantic_change_error(
                    previous, current, records, contract_version="3.0", chunk=chunk,
                    authorization_out=authorizations,
                )
                self.assertIsNone(error)
                self.assertEqual(changed, ["$.requirements[0].properties.text"])
                self.assertEqual(len(authorizations), 1)
                self.assertTrue(authorizations[0]["source_binding_complete"])
                projected, _, errors = bridge.materialize_source_fragment_literals(
                    current, chunk["clauses"], chunk["evidence_context"],
                )
                self.assertEqual(errors, [])
                self.assertEqual(projected["requirements"][0]["properties"]["text"], source)
                self.assertEqual(current["requirements"][0]["properties"]["text"], literal)

        current = copy.deepcopy(previous)
        current["requirements"][0]["properties"]["text"] = None
        variants = []
        for key, value in (
            ("reason", "unrequested rewrite"),
            ("clause_ids", ["foreign-clause"]),
            ("evidence_ids", ["foreign-evidence"]),
            ("source_fragment_clause_ids", None),
            ("role", "body_text"),
        ):
            candidate = copy.deepcopy(current)
            candidate["requirements"][0][key] = value
            variants.append((key, candidate, records, chunk))
        candidate = copy.deepcopy(current)
        candidate["requirements"][0]["properties"]["text"] = source + "猜测"
        variants.append(("guessed_literal", candidate, records, chunk))
        stale = copy.deepcopy(records)
        stale[0]["response_sha256"] = "0" * 64
        variants.append(("stale_parent", current, stale, chunk))
        selector_error = copy.deepcopy(records)
        selector_error[0]["json_pointer"] = "$.requirements[0].source_fragment_clause_ids"
        variants.append(("selector_error_cannot_authorize_text_only", current, selector_error, chunk))
        stale_source = copy.deepcopy(chunk)
        stale_source["clauses"][0]["source_span"]["source_sha256"] = "0" * 64
        variants.append(("stale_source", current, records, stale_source))
        for name, candidate, error_records, source_chunk in variants:
            with self.subTest(rejection=name):
                error, _ = bridge._retry_semantic_change_error(
                    previous, candidate, error_records, contract_version="3.0", chunk=source_chunk,
                )
                self.assertIsNotNone(error)

        # Two independent source instances must have independent authorizations.
        multi_chunk = copy.deepcopy(chunk)
        second_clause = copy.deepcopy(clause)
        second_clause["id"] = "CL-second"
        second_clause["evidence_ids"] = ["EV-second"]
        second_clause["source_span"]["evidence_id"] = "EV-second"
        multi_chunk["clauses"].append(second_clause)
        multi_chunk["evidence_context"]["EV-second"] = {
            **copy.deepcopy(chunk["evidence_context"]["EV-label"]), "id": "EV-second",
        }
        multi_previous = copy.deepcopy(previous)
        second_requirement = copy.deepcopy(previous["requirements"][0])
        second_requirement["clause_ids"] = ["CL-second"]
        second_requirement["evidence_ids"] = ["EV-second"]
        second_requirement.pop("source_fragment_clause_ids")
        multi_previous["requirements"].append(second_requirement)
        multi_current = copy.deepcopy(multi_previous)
        multi_current["requirements"][0]["properties"]["text"] = None
        multi_current["requirements"][1]["properties"]["text"] = source
        multi_current["requirements"][1]["source_fragment_clause_ids"] = ["CL-second"]
        multi_records = [
            {**records[0], "response_sha256": bridge._response_sha256(multi_previous)},
            {**records[0], "response_sha256": bridge._response_sha256(multi_previous),
             "json_pointer": "$.requirements[1].source_fragment_clause_ids"},
        ]
        error, _ = bridge._retry_semantic_change_error(
            multi_previous, multi_current, multi_records, contract_version="3.0", chunk=multi_chunk,
        )
        self.assertIsNone(error)
        multi_previous["requirements"][1]["source_fragment_clause_ids"] = ["CL-second"]
        for record in multi_records:
            record["response_sha256"] = bridge._response_sha256(multi_previous)
        # Its selector is now unchanged: index 1 has no text-targeted error.
        error, _ = bridge._retry_semantic_change_error(
            multi_previous, multi_current, multi_records, contract_version="3.0", chunk=multi_chunk,
        )
        self.assertIsNotNone(error)
        multi_records.append({
            **multi_records[1], "json_pointer": "$.requirements[1].properties.text",
        })
        error, _ = bridge._retry_semantic_change_error(
            multi_previous, multi_current, multi_records, contract_version="3.0", chunk=multi_chunk,
        )
        self.assertIsNone(error)

    def test_source_fragment_retry_rejects_null_selector_and_unrequested_obligation_rewrite(self) -> None:
        label = "答辩委员会"
        source = label + "："
        location = {"part": "document", "child_index": 12, "order": 8}
        clause = {
            "id": "C1", "text": label, "evidence_ids": ["E1"],
            "source_kind": "paragraph", "location": location,
            "source_evidence_text": source,
            "source_span": {
                "evidence_id": "E1", "start_offset": 0, "end_offset": len(label),
                "text": label, "source_sha256": hashlib.sha256(source.encode()).hexdigest(),
                "location": location,
            },
        }
        chunk = {
            "clauses": [clause],
            "evidence_context": {"E1": {
                "id": "E1", "kind": "paragraph", "text": source,
                "location": location,
            }},
            "requirement_contract": {
                "role_properties_schema": {
                    "cover_field_label": {"$ref": "#/$defs/roleSpec"},
                },
            },
        }
        previous = {
            "contract_version": "3.0",
            "requirements": [{
                "role": "cover_field_label", "properties": {"text": source},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
                "source_fragment_clause_ids": ["C1"],
            }],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "executable",
                "normative_basis": "explicit_normative_text",
                "reason": "The source label is represented exactly.",
                "obligations": [{
                    "id": "label", "status": "covered",
                    "reason": "The defense committee section label is represented by the exact cover_field_label text.",
                }],
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        records = [{
            "code": "source_fragment_binding_violation",
            "json_pointer": "$.requirements[0].properties.text",
            "response_sha256": bridge._response_sha256(previous),
            "raw_error": "source_fragment_literal_conflict",
        }]

        rewritten_reason = copy.deepcopy(previous)
        rewritten_reason["clause_reviews"][0]["obligations"][0]["reason"] = (
            "The defense committee section label is represented by the cover_field_label text property."
        )
        changed = bridge._retry_change_paths(previous, rewritten_reason)
        self.assertFalse(bridge._source_fragment_binding_retry_allowed(
            previous, rewritten_reason, records, changed, chunk=chunk,
        ))

        dropped_selector = copy.deepcopy(rewritten_reason)
        dropped_selector["requirements"][0]["source_fragment_clause_ids"] = None
        changed = bridge._retry_change_paths(previous, dropped_selector)
        self.assertFalse(bridge._source_fragment_binding_retry_allowed(
            previous, dropped_selector, records, changed, chunk=chunk,
        ))

    def test_declaration_source_fragment_retry_corrects_selector_without_text_or_link_drift(self) -> None:
        title = "声明标题"
        body = "完整声明正文。"
        left_location = {"part": "document", "child_index": 4, "order": 2}
        right_location = {"part": "document", "child_index": 8, "order": 6}
        clauses = [
            {
                "id": "C1", "text": title, "evidence_ids": ["E1"],
                "location": left_location, "source_kind": "paragraph",
                "source_evidence_text": title,
                "source_span": {
                    "evidence_id": "E1", "start_offset": 0, "end_offset": len(title),
                    "text": title, "source_sha256": hashlib.sha256(title.encode()).hexdigest(),
                    "location": left_location,
                },
            },
            {
                "id": "C2", "text": body, "evidence_ids": ["E2"],
                "location": right_location, "source_kind": "paragraph",
                "source_evidence_text": body,
                "source_span": {
                    "evidence_id": "E2", "start_offset": 0, "end_offset": len(body),
                    "text": body, "source_sha256": hashlib.sha256(body.encode()).hexdigest(),
                    "location": right_location,
                },
            },
        ]
        evidence_context = {
            "E1": {"id": "E1", "kind": "paragraph", "text": title, "location": left_location},
            "E2": {"id": "E2", "kind": "paragraph", "text": body, "location": right_location},
        }
        chunk = {
            "provenance": {
                "run_id": "run-declaration-fragment-retry", "case_id": "case-declaration-fragment-retry",
                "source_sha256": "a" * 64, "clause_sha256": "b" * 64,
                "evidence_sha256": "c" * 64, "request_sha256": "d" * 64,
            },
            "case_id": "case-declaration-fragment-retry", "batch": {"index": 2},
            "runtime_context": {"code_fingerprint_sha256": "9" * 64},
            "response_schema": {"type": "object", "title": "declaration source fragment retry"},
            "requirement_contract": {
                "role_properties_schema": {"declarations": {"$ref": "#/$defs/declarationsSpec"}},
            },
            "clauses": clauses, "evidence_context": evidence_context,
        }
        previous = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [{
                "role": "declarations",
                "properties": {
                    "before_role": "document_start",
                    "items": [{
                        "id": "originality", "heading": title, "body_parts": [body],
                        "source_evidence_ids": ["E1", "E2"],
                    }],
                },
                "clause_ids": ["C1", "C2"], "evidence_ids": ["E1", "E2"],
                "source_fragment_clause_ids": ["C1", "C2"],
                "reason": "Both fixed declaration paragraphs are preserved exactly.",
            }],
            "clause_reviews": [], "unsupported_items": [], "reported_conflicts": [],
        }
        current = copy.deepcopy(previous)
        current["requirements"][0]["source_fragment_clause_ids"] = ["C1"]
        records = [{
            "code": "source_fragment_binding_violation",
            "json_pointer": "$.requirements[0].source_fragment_clause_ids",
            "response_sha256": bridge._response_sha256(previous),
            "raw_error": "cross_evidence_boundary_unproven",
        }]

        error, changed = bridge._retry_semantic_change_error(
            previous, current, records, contract_version="3.0", chunk=chunk,
        )

        self.assertIsNone(error)
        self.assertEqual(changed, ["$.requirements[0].source_fragment_clause_ids"])
        self.assertNotIn("text", current["requirements"][0]["properties"])

        dropped_clause_link = copy.deepcopy(current)
        dropped_clause_link["requirements"][0]["clause_ids"] = ["C1"]
        rejected, _ = bridge._retry_semantic_change_error(
            previous, dropped_clause_link, records, contract_version="3.0", chunk=chunk,
        )
        self.assertIsNotNone(rejected)

    def test_retry_artifact_receipts_include_unaccepted_repair_base_without_promoting_it(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response_path = review_dir / "llm-response-chunk-0001.json"
            raw_path = response_path.with_name(
                f"{response_path.stem}.attempt-01.raw{response_path.suffix}"
            )
            repair_base_path = response_path.with_name(
                f"{response_path.stem}.attempt-01.repair-base{response_path.suffix}"
            )
            raw = {"contract_version": "3.0", "requirements": [], "clause_reviews": []}
            repair_base = copy.deepcopy(raw)
            repair_base["requirements"] = [{
                "role": "title", "properties": {"text": "source-backed partial candidate"},
                "clause_ids": ["C1"], "evidence_ids": ["E1"],
            }]
            raw_path.write_text(json.dumps(raw), encoding="utf-8")
            repair_base_path.write_text(json.dumps(repair_base), encoding="utf-8")
            residual = [{
                "code": "schema_contract_violation",
                "response_sha256": bridge._response_sha256(repair_base),
            }]
            snapshots = [
                bridge._attempt_stage_snapshot("decoded_raw", raw, path=raw_path, chunk=chunk),
                bridge._attempt_stage_snapshot(
                    "repair_base", repair_base, path=repair_base_path, chunk=chunk,
                    projection_audit={
                        "accepted": False,
                        "initial_error_records_sha256": bridge._response_sha256([]),
                        "resolved_error_records_sha256": bridge._response_sha256([]),
                        "residual_error_records_sha256": bridge._response_sha256(residual),
                        "repair_authorization_error_records_sha256": bridge._response_sha256(residual),
                    }, accepted=False,
                ),
            ]
            attempt = {
                "retry_input_fingerprints": bridge._retry_input_fingerprints(chunk),
                "stage_snapshots": snapshots,
                "error_records": copy.deepcopy(residual),
                "initial_error_records": [], "resolved_error_records": [],
                "residual_error_records": copy.deepcopy(residual),
                "retry_authorizing_error_records": copy.deepcopy(residual),
            }
            receipt = bridge._validate_retry_attempt_artifact(response_path, 1, attempt)
            repair_receipt = receipt["unaccepted_repair_base_receipt"]
            self.assertFalse(repair_receipt["accepted"])
            self.assertEqual(repair_receipt["kind"], "unaccepted_repair_base")
            self.assertEqual(receipt["kind"], "decoded_raw")
            self.assertNotEqual(receipt["path"], repair_receipt["path"])
            selected = bridge._retry_semantic_parent_receipt([receipt])
            self.assertIsNotNone(selected)
            self.assertEqual(selected["path"], repair_receipt["path"])
            self.assertEqual(selected["sha256"], repair_receipt["sha256"])
            self.assertFalse(selected["accepted"])
            self.assertEqual(
                bridge._retry_semantic_parent_receipt([
                    receipt, {"kind": "no_semantic_response", "attempt": 2, "path": "unused"},
                ])["path"], repair_receipt["path"],
            )
            replay_error = ValueError("remaining local contract errors")
            replay_error.repair_base_candidate = copy.deepcopy(repair_base)
            replay_error.retry_authorizing_error_records = copy.deepcopy(residual)
            with patch.object(
                bridge, "prepare_native_response_candidate", side_effect=replay_error,
            ):
                proof = bridge._prove_retry_repair_base_replay(
                    selected, chunk, residual,
                    source_projection_validation_sha256=None,
                )
            self.assertFalse(proof["accepted"])
            self.assertEqual(proof["repair_base_sha256"], bridge._response_sha256(repair_base))
            replay_error.retry_authorizing_error_records[0]["code"] = "different_issue"
            with patch.object(
                bridge, "prepare_native_response_candidate", side_effect=replay_error,
            ), self.assertRaisesRegex(
                bridge.RetryRawArtifactIntegrityError, "cannot be replayed",
            ):
                bridge._prove_retry_repair_base_replay(
                    selected, chunk, residual,
                    source_projection_validation_sha256=None,
                )

            tampered = copy.deepcopy(attempt)
            tampered["stage_snapshots"][1]["accepted"] = True
            with self.assertRaisesRegex(
                bridge.RetryRawArtifactIntegrityError,
                "explicitly marked unaccepted",
            ):
                bridge._validate_retry_attempt_artifact(response_path, 1, tampered)

            missing_digest = copy.deepcopy(attempt)
            missing_digest["retry_authorizing_error_records"][0]["response_sha256"] = None
            with self.assertRaisesRegex(
                bridge.RetryRawArtifactIntegrityError, "authorization records",
            ):
                bridge._validate_retry_attempt_artifact(response_path, 1, missing_digest)

            changed_ledger = copy.deepcopy(attempt)
            changed_ledger["retry_authorizing_error_records"][0]["code"] = "different_issue"
            with self.assertRaisesRegex(
                bridge.RetryRawArtifactIntegrityError, "error ledger differs",
            ):
                bridge._validate_retry_attempt_artifact(response_path, 1, changed_ledger)

    def test_retry_prompt_and_drift_use_same_verified_repair_baseline(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            baseline = self._response(chunk)
            original = copy.deepcopy(baseline)
            original["requirements"] = [{
                "role": "body_text", "properties": {}, "clause_ids": [],
                "evidence_ids": [], "reason": "",
            }]
            records = [{
                "code": "contract_validation_error", "json_pointer": "$.requirements[0]",
                "raw_error": "empty provider shell",
                "response_sha256": bridge._response_sha256(baseline),
            }]
            actual_prepare = bridge.prepare_native_response_candidate
            prepare_calls = 0

            def partial_once(value, current_chunk, **kwargs):
                nonlocal prepare_calls
                prepare_calls += 1
                if prepare_calls <= 2:
                    error = ValueError("local response contract validation failed")
                    error.error_records = copy.deepcopy(records)
                    error.initial_error_records = copy.deepcopy(records)
                    error.resolved_error_records = []
                    error.residual_error_records = copy.deepcopy(records)
                    error.retry_authorizing_error_records = copy.deepcopy(records)
                    error.repair_base_candidate = copy.deepcopy(baseline)
                    raise error
                return actual_prepare(value, current_chunk, **kwargs)

            def envelope(payload):
                return subprocess.CompletedProcess(["openclaw"], 0, json.dumps({
                    "runId": "offline-retry", "status": "ok",
                    "provider": "openai", "model": "gpt-5.6-luna",
                    "result": {"payloads": [{"text": json.dumps(payload)}]},
                }), "")

            with patch.object(
                bridge, "_run_command", side_effect=[envelope(original), envelope(baseline)],
            ), patch.object(
                bridge, "prepare_native_response_candidate", side_effect=partial_once,
            ):
                audit = bridge.run_bridge(
                    review_dir, response_out=Path(td) / "accepted.json",
                    agent_id="main", timeout=1, max_attempts=2,
                    openclaw_bin="openclaw", model="openai/gpt-5.6-luna",
                )

            self.assertEqual(audit["status"], "merged")
            lifecycle = audit["chunk_lifecycle"][0]
            self.assertEqual(lifecycle["retry_parent_stage"], "unaccepted_repair_base")
            repair = lifecycle["attempts"][0]["repair_base_snapshot"]
            prompt = (review_dir / "host-agent-prompts" / "prompt-0001-attempt-02.txt").read_text()
            self.assertIn(repair["file_bytes_sha256"], prompt)
            self.assertIn('"requirements": []', prompt)
            self.assertNotIn('"reason": ""', prompt)
            comparison = audit["chunk_runs"][0]["retry_stage_comparison"]
            self.assertEqual(comparison["semantic_parent_stage"], "unaccepted_repair_base")
            self.assertEqual(comparison["raw_parent_file_sha256"], repair["file_bytes_sha256"])
            self.assertFalse(comparison["original_raw_observation"]["authorization_basis"])

    def test_completed_chunk_set_rejects_non_integer_and_duplicate_indexes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            response = {"requirements": [], "clause_reviews": []}
            response_path = root / "accepted.json"
            response_path.write_text(json.dumps(response), encoding="utf-8")
            digest = bridge._response_sha256(response)
            lifecycle = {
                1: {"status": "completed", "remote_operation_state": "completed"},
                2: {"status": "completed", "remote_operation_state": "completed"},
            }
            audits = [
                {"chunk_index": 1, "accepted_response_sha256": digest, "response_path": str(response_path)},
                {"chunk_index": 1, "accepted_response_sha256": digest, "response_path": str(response_path)},
            ]
            with self.assertRaisesRegex(ValueError, "each chunk exactly once"):
                bridge._validate_completed_chunk_set(root, ["accepted.json", "accepted.json"], lifecycle, audits)
            audits[1]["chunk_index"] = []
            with self.assertRaisesRegex(ValueError, "each chunk exactly once"):
                bridge._validate_completed_chunk_set(root, ["accepted.json", "accepted.json"], lifecycle, audits)

    def test_cancellation_after_raw_response_never_publishes_accepted_response(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            review_dir, chunk = self._packet(Path(td) / "requirements")
            response = self._response(chunk)
            envelope = {
                "runId": "openclaw-run-cancel-race",
                "status": "ok",
                "result": {"payloads": [{"text": json.dumps(response)}]},
            }

            class CancelAfterRaw(bridge.RunController):
                def __init__(self) -> None:
                    super().__init__()
                    self.check_count = 0

                def check(self) -> None:
                    self.check_count += 1
                    if self.check_count >= 3:
                        raise bridge.HostAgentCancelled("fatal sibling failure")

            controller = CancelAfterRaw()
            response_path = review_dir / "llm-response-chunk-0001.attempt-01.json"
            with patch.object(
                bridge, "_run_command",
                return_value=subprocess.CompletedProcess(
                    ["openclaw"], 0, json.dumps(envelope), ""
                ),
            ):
                with self.assertRaises(bridge.HostAgentCancelled):
                    bridge.run_host_agent_chunk(
                        request_path=review_dir / "llm-request.json",
                        chunk_path=review_dir / "llm-request-chunks.json",
                        chunk=chunk,
                        response_path=response_path,
                        run_id="run-bridge-test",
                        chunk_index=1,
                        chunk_count=1,
                        agent_id="main",
                        timeout=1,
                        openclaw_bin="openclaw",
                        prompt_path=review_dir / "prompt.txt",
                        controller=controller,
                    )
            self.assertFalse(response_path.exists())
            self.assertTrue(response_path.with_name(
                f"{response_path.stem}.raw{response_path.suffix}"
            ).exists())

    def test_parse_openclaw_result_accepts_one_json_fence(self) -> None:
        envelope = {
            "status": "ok",
            "result": {"payloads": [{"text": "```json\n{\"ok\": true}\n```"}]},
        }
        response, _ = bridge.parse_openclaw_result(json.dumps(envelope))
        self.assertEqual(response, {"ok": True})

    def test_parse_openclaw_exec_result_reads_root_payloads_and_route(self) -> None:
        envelope = {
            "status": "ok",
            "provider": "openai",
            "model": "gpt-5.6-luna",
            "payloads": [{"text": "{\"ok\": true}"}],
        }
        response, parsed = bridge.parse_openclaw_result(json.dumps(envelope))
        self.assertEqual(response, {"ok": True})
        route = bridge.verify_host_agent_route(parsed, "openai/gpt-5.6-luna")
        self.assertEqual(route["route"], "openai/gpt-5.6-luna")


if __name__ == "__main__":
    unittest.main()
