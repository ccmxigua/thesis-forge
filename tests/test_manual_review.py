from __future__ import annotations

import hashlib
import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from apply_format_spec import (  # noqa: E402
    append_manual_review_markers,
    insert_inline_manual_review_markers,
    audit_manual_review_markers,
)
from format_spec_validation import load_and_validate  # noqa: E402
from manual_review import (  # noqa: E402
    build_manual_review_ledger,
    filter_manual_marker_ledger,
)
import requirements_engine as engine  # noqa: E402
from submission_audit import PLACEHOLDER_PATTERNS  # noqa: E402


class ManualReviewTests(unittest.TestCase):
    def test_technical_findings_do_not_become_human_red_markers(self) -> None:
        ledger = build_manual_review_ledger({
            "findings": [
                {
                    "code": "backend_gap", "blocking": True,
                    "message": "format execution unavailable",
                    "evidence": [{"kind": "category", "value": "backend_capability_gap"}],
                },
                {
                    "code": "runtime_unknown", "blocking": True,
                    "message": "abstract target cannot be inferred",
                    "evidence": [{"kind": "category", "value": "runtime_manual_unverifiable"}],
                },
            ],
        }, [], release_gates=[{
            "source_code": "render_not_run", "category": "render_validation",
            "source_text": "word render not run",
        }])
        self.assertEqual([item["source_code"] for item in ledger["items"]], ["runtime_unknown"])
        self.assertEqual(ledger["summary"]["total"], 1)

    def test_legacy_technical_items_are_removed_but_human_decisions_remain(self) -> None:
        ledger = {
            "items": [
                {"marker_id": "MR-0001", "category": "property_receipt", "source_code": "receipt"},
                {"marker_id": "MR-0002", "category": "input_prerequisite", "source_code": "template"},
                {"marker_id": "MR-0003", "category": "semantic_content_review", "source_code": "abstract"},
            ],
        }
        filtered = filter_manual_marker_ledger(ledger)
        self.assertEqual([item["source_code"] for item in filtered["items"]], ["abstract", "template"])
        self.assertEqual([item["marker_id"] for item in filtered["items"]], ["MR-0001", "MR-0002"])
        self.assertEqual(filtered["summary"]["by_category"], {
            "semantic_content_review": 1, "input_prerequisite": 1,
        })

    def test_ledger_preserves_binding_and_merges_duplicate_question(self) -> None:
        report = {
            "findings": [{
                "code": "capability.clause_gap",
                "blocking": True,
                "message": "C00074 requires manual review",
                "evidence": [
                    {"kind": "clause_id", "value": "C00074"},
                    {"kind": "evidence_id", "value": "E00074"},
                    {"kind": "category", "value": "runtime_manual_unverifiable"},
                ],
            }],
        }
        questions = [{
            "question_id": "Q00074",
            "clause_id": "C00074",
            "evidence_id": "E00074",
            "question": "英文目标是否明确？",
            "source_text": "The Chinese abstract is 300 to 1,000 words",
        }]
        ledger = build_manual_review_ledger(
            report,
            questions,
            binding={
                "case_id": "bsu",
                "run_id": "run-1",
                "source_sha256": "0" * 64,
                "clause_sha256": "1" * 64,
                "evidence_sha256": "2" * 64,
            },
        )
        self.assertEqual(ledger["binding"]["run_id"], "run-1")
        self.assertFalse(ledger["submission_ready"])
        self.assertEqual(ledger["summary"]["total"], 2)
        self.assertEqual(
            {tuple(item["clause_ids"]) for item in ledger["items"]},
            {("C00074",)},
        )
        question_item = next(item for item in ledger["items"] if item["source_type"] == "open_question")
        self.assertEqual(question_item["source_text"], "The Chinese abstract is 300 to 1,000 words")
        self.assertEqual(
            load_and_validate(ledger, ROOT / "schema" / "manual-review-ledger.schema.json"),
            [],
        )

    def test_producer_question_shape_preserves_id_and_all_evidence_ids(self) -> None:
        ledger = build_manual_review_ledger({}, [{
            "id": "Q-PRODUCER-1",
            "clause_id": "C-PRODUCER-1",
            "question": "这条规则的适用对象是什么？",
            "source_text": "每个关键词不超过限定长度",
            "evidence_ids": ["E-PRODUCER-1", "E-PRODUCER-2"],
        }])
        item = next(item for item in ledger["items"] if item["source_type"] == "open_question")
        self.assertEqual(item["question_ids"], ["Q-PRODUCER-1"])
        self.assertEqual(item["evidence_ids"], ["E-PRODUCER-1", "E-PRODUCER-2"])

    def test_requirements_producer_to_ledger_to_serialized_marker_keeps_source_binding(self) -> None:
        clauses = [{
            "id": "C-PRODUCER-1",
            "text": "某一格式描述暂时无法映射到已注册语义角色",
            "evidence_ids": ["E-PRODUCER-1", "E-PRODUCER-2"],
            "context_before": "前置语境", "context_after": "后续语境",
        }]
        evidence_doc = {"evidence": [
            {"id": "E-PRODUCER-1", "text": clauses[0]["text"], "kind": "paragraph"},
            {"id": "E-PRODUCER-2", "text": clauses[0]["text"], "kind": "paragraph"},
        ]}
        # Keep semantic role discovery under test control while invoking the
        # actual rule-result producer and its real question-record shape.
        with patch.object(engine, "parse_properties", return_value={"style_hint": "unknown"}), \
                patch.object(engine, "identify_role", return_value=(None, ["unknown"])):
            _spec, questions, conflicts = engine.build_rule_result(
                Path("producer-fixture.docx"), clauses,
            )
        self.assertEqual(conflicts, [])
        self.assertEqual(questions[0]["question_id"], "Q0001")
        self.assertEqual(questions[0]["evidence_ids"], ["E-PRODUCER-1", "E-PRODUCER-2"])

        ledger = build_manual_review_ledger(
            {}, questions, clauses=clauses, evidence_doc=evidence_doc,
            binding={
                "case_id": "producer-fixture", "run_id": "question-marker-e2e",
                "source_sha256": "a" * 64, "clause_sha256": "b" * 64,
                "evidence_sha256": "c" * 64,
            },
        )
        self.assertEqual(len(ledger["items"]), 1)
        item = ledger["items"][0]
        self.assertEqual(item["question_ids"], ["Q0001"])
        self.assertEqual(item["clause_ids"], ["C-PRODUCER-1"])
        self.assertEqual(item["evidence_ids"], ["E-PRODUCER-1", "E-PRODUCER-2"])

        document = Document()
        receipts = append_manual_review_markers(document, ledger)
        self.assertEqual(receipts[0]["question_ids"], ["Q0001"])
        self.assertEqual(receipts[0]["evidence_ids"], ["E-PRODUCER-1", "E-PRODUCER-2"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "producer-question-marker.docx"
            document.save(path)
            audit = audit_manual_review_markers(path, ledger)
        self.assertTrue(audit["valid"], audit)
        self.assertEqual(audit["marker_bindings"][0]["question_ids"], ["Q0001"])
        self.assertEqual(
            audit["marker_bindings"][0]["evidence_ids"],
            ["E-PRODUCER-1", "E-PRODUCER-2"],
        )

    def test_pipeline_question_binding_checks_clause_and_evidence_before_marker(self) -> None:
        clauses = [{
            "id": "C1", "text": "关键词最多七个汉字", "evidence_ids": ["E1"],
        }]
        evidence_doc = {"evidence": [{"id": "E1", "text": "关键词最多七个汉字"}]}
        ledger = build_manual_review_ledger({}, [{
            "id": "Q1", "clause_id": "C1", "evidence_ids": ["E1"],
            "question": "这里的字符如何计数？",
        }], clauses=clauses, evidence_doc=evidence_doc)
        item = next(item for item in ledger["items"] if item["source_type"] == "open_question")
        self.assertEqual(item["question_ids"], ["Q1"])
        self.assertEqual(item["evidence_ids"], ["E1"])
        self.assertEqual(item["source_text"], "关键词最多七个汉字")

        with self.assertRaisesRegex(ValueError, "unknown source clause"):
            build_manual_review_ledger({}, [{
                "id": "Q1", "clause_id": "C404", "evidence_ids": ["E1"],
                "question": "无法绑定的问题",
            }], clauses=clauses, evidence_doc=evidence_doc)

    def test_run_global_diagnostic_is_not_emitted_as_author_red_marker(self) -> None:
        ledger = build_manual_review_ledger({}, [{
            "question_id": "Q-RUN", "scope": "global", "evidence_ids": [],
            "question": "需要 Host Agent 审查",
        }], clauses=[])
        self.assertEqual(ledger["items"], [])

    def test_conflicting_question_aliases_fail_closed(self) -> None:
        with self.assertRaisesRegex(ValueError, "question_id conflicts"):
            build_manual_review_ledger({}, [{
                "question_id": "Q1", "id": "Q2", "question": "确认问题",
            }])
        with self.assertRaisesRegex(ValueError, "evidence_id conflicts"):
            build_manual_review_ledger({}, [{
                "question_id": "Q1", "evidence_ids": ["E1"], "evidence_id": "E2",
                "question": "确认问题",
            }])

    def test_question_identity_and_evidence_are_required_and_unique(self) -> None:
        with self.assertRaisesRegex(ValueError, "question id"):
            build_manual_review_ledger({}, [{"question": "没有稳定 ID"}])
        with self.assertRaisesRegex(ValueError, "duplicate question_id"):
            build_manual_review_ledger({}, [
                {"question_id": "Q1", "question": "问题一"},
                {"id": "Q1", "question": "问题二"},
            ])
        with self.assertRaisesRegex(ValueError, "contains duplicates"):
            build_manual_review_ledger({}, [{
                "question_id": "Q1", "evidence_ids": ["E1", "E1"],
            }])

    def test_markers_are_visible_and_red(self) -> None:
        document = Document()
        ledger = {
            "schema_version": "1.0",
            "policy": "review_draft_only",
            "binding": {"run_id": "run-1"},
            "submission_ready": False,
            "items": [{
                "marker_id": "MR-0001",
                "status": "pending_manual_review",
                "marker_required": True,
                "category": "input_prerequisite",
                "clause_ids": ["C00076"],
                "requirement_ids": [],
                "source_text": "300–1000 words",
                "reason": "目标语言和单位需要人工确认",
                "action": "请确认后回填",
                "placeholder_text": "【待人工处理：C00076】",
                "original_blocking": True,
            }],
        }
        receipts = append_manual_review_markers(document, ledger)
        self.assertEqual(len(receipts), 1)
        marker = next(p for p in document.paragraphs if p.text.startswith("【MR-0001"))
        self.assertIn("【待人工处理：C00076】", marker.text)
        self.assertEqual(receipts[0]["placeholder_text"], "【待人工处理：C00076】")
        run = marker.runs[0]
        self.assertEqual(str(run.font.color.rgb), "C00000")
        shading = run._r.rPr.find(qn("w:shd"))
        self.assertIsNotNone(shading)
        self.assertEqual(shading.get(qn("w:fill")), "FFF2CC")

    def test_official_template_gate_is_a_run_bound_red_placeholder(self) -> None:
        ledger = build_manual_review_ledger(
            {}, [],
            release_gates=[{
                "source_code": "official_template_missing",
                "category": "input_prerequisite",
                "source_text": "官方版式模板未提供",
                "reason": "当前仅有中性参考文档。",
                "action": "提供官方模板后重新运行。",
                "placeholder_text": "【待提供：官方版式模板】",
            }],
            binding={
                "case_id": "bsu",
                "run_id": "run-1",
                "source_sha256": "0" * 64,
                "clause_sha256": "1" * 64,
                "evidence_sha256": "2" * 64,
            },
        )
        self.assertEqual(ledger["visual_policy"]["text_color"], "C00000")
        self.assertFalse(ledger["submission_ready"])
        item = ledger["items"][0]
        self.assertEqual(item["source_code"], "official_template_missing")
        self.assertEqual(item["placeholder_text"], "【待提供：官方版式模板】")
        self.assertEqual(
            load_and_validate(ledger, ROOT / "schema" / "manual-review-ledger.schema.json"),
            [],
        )

        document = Document()
        receipts = append_manual_review_markers(document, ledger)
        self.assertEqual(receipts[0]["placeholder_text"], "【待提供：官方版式模板】")
        marker = next(p for p in document.paragraphs if p.text.startswith("【MR-0001"))
        self.assertIn("【待提供：官方版式模板】", marker.text)
        self.assertEqual(str(marker.runs[0].font.color.rgb), "C00000")

    @staticmethod
    def _item(index=1, **values):
        return {
            "marker_id": f"MR-{index:04d}", "status": "pending_manual_review",
            "marker_required": True, "category": "runtime_manual_unverifiable",
            "source_type": "open_question", "clause_ids": [], "requirement_ids": [],
            "source_text": "适用对象需要确认", "reason": "不能推测具体解释",
            "action": "请人工核对后填写", "placeholder_text": "【待人工处理：保留原文】",
            "original_blocking": True, **values,
        }

    def test_inline_marker_is_inserted_after_safe_abstract_anchor(self) -> None:
        document = Document()
        document.styles.add_style("Abstract Body CN", 1)
        document.add_paragraph("中文摘要正文示例", "Abstract Body CN")
        ledger = {"items": [self._item(
            source_type="format_constraint", source_text="abstract_zh.require_third_person",
        )]}
        original = copy.deepcopy(ledger)
        locations = insert_inline_manual_review_markers(document, ledger, {})
        self.assertEqual(locations["MR-0001"]["location"], "inline_after_role")
        self.assertIn("【MR-0001｜人工待审】", document.paragraphs[1].text)
        self.assertIn("请在此人工处理", document.paragraphs[1].text)
        self.assertEqual(str(document.paragraphs[1].runs[0].font.color.rgb), "C00000")
        self.assertEqual(ledger, original)

    def test_english_abstract_fallback_has_imported_detector(self) -> None:
        document = Document()
        document.styles.add_style("Abstract Body EN", 1)
        document.add_paragraph("English abstract body", "Abstract Body EN")
        locations = insert_inline_manual_review_markers(document, {"items": [self._item(
            source_type="format_constraint", source_text="abstract_en.required_sections",
        )]}, {})
        self.assertEqual(locations["MR-0001"]["anchor_role"], "abstract_body_en")

    def test_ambiguous_measurement_and_abstract_are_front_unlocated(self) -> None:
        document = Document()
        document.add_paragraph("论文原始正文")
        ledger = {"items": [
            self._item(source_text="3cm左右", clause_ids=["C00421"]),
            self._item(2, source_text="The Chinese abstract is 300 to 1,000 words",
                       clause_ids=["C00076"]),
            self._item(3, source_type="release_gate", source_code="official_template_missing"),
            self._item(4, source_type="format_constraint", source_text="abstract_zh.target"),
        ]}
        before = copy.deepcopy(ledger)
        locations = insert_inline_manual_review_markers(document, ledger, {})
        self.assertEqual(locations, {})
        receipts = append_manual_review_markers(document, ledger, locations)
        self.assertEqual({r["location"] for r in receipts}, {"document_front_unlocated"})
        self.assertEqual(document.paragraphs[0].text, "待定位人工处理（审查草稿，不可提交）")
        self.assertEqual(document.paragraphs[-1].text, "论文原始正文")
        self.assertIn("3cm左右", document.paragraphs[2].text)
        self.assertEqual(ledger, before)

    def test_only_human_decision_categories_can_be_serialized_as_markers(self) -> None:
        document = Document()
        document.add_paragraph("源文档内容不应被改写")
        categories = ["input_prerequisite", "runtime_manual_unverifiable",
                      "confirmed_semantic_issue", "semantic_content_review"]
        ledger = {"items": [self._item(i, category=category)
                            for i, category in enumerate(categories, 1)]}
        append_manual_review_markers(document, ledger)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "markers.docx"
            document.save(path)
            result = audit_manual_review_markers(path, ledger)
        self.assertTrue(result["valid"], result)
        self.assertEqual(result["visible_marker_count"], len(categories))
        self.assertFalse(result["submission_ready"])
        self.assertEqual(result["visual_verification"], "required")

    def test_serialized_markers_keep_question_and_evidence_binding(self) -> None:
        ledger = {
            "binding": {
                "case_id": "bsu", "run_id": "run-marker-binding",
                "source_sha256": "0" * 64, "clause_sha256": "1" * 64,
                "evidence_sha256": "2" * 64,
            },
            "items": [self._item(
                source_type="open_question", question_ids=["Q00076"],
                evidence_ids=["E00065"], clause_ids=["C00076"],
            )],
        }
        document = Document()
        receipts = append_manual_review_markers(document, ledger)
        self.assertEqual(receipts[0]["question_ids"], ["Q00076"])
        self.assertEqual(receipts[0]["evidence_ids"], ["E00065"])
        marker = next(p for p in document.paragraphs if p.text.startswith("【MR-"))
        self.assertIn("问题编号：Q00076", marker.text)
        self.assertIn("证据编号：E00065", marker.text)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bound-markers.docx"
            document.save(path)
            audit = audit_manual_review_markers(path, ledger)
            self.assertTrue(audit["valid"], audit)
            self.assertEqual(audit["marker_bindings"][0]["question_ids"], ["Q00076"])
            self.assertTrue(audit["marker_bindings"][0]["ledger_item_sha256"])

            stale_ledger = copy.deepcopy(ledger)
            stale_ledger["items"][0]["question_ids"] = ["Q09999"]
            stale = audit_manual_review_markers(path, stale_ledger)
            self.assertFalse(stale["valid"])
            self.assertEqual(stale["source_binding_errors"], ["MR-0001"])

        for category in ("format_validation", "property_receipt",
                         "backend_capability_gap", "render_validation"):
            with self.subTest(category=category):
                with self.assertRaisesRegex(ValueError, "technical diagnostics"):
                    append_manual_review_markers(
                        Document(), {"items": [self._item(category=category)]},
                    )

    def test_marker_style_overrides_inherited_hidden_and_clipped_text(self) -> None:
        document = Document()
        document.styles["Normal"].font.hidden = True
        document.styles["Normal"].paragraph_format.line_spacing = Pt(1)
        ledger = {"items": [self._item()]}
        append_manual_review_markers(document, ledger)
        paragraph = next(p for p in document.paragraphs if p.text.startswith("【MR-"))
        run = paragraph.runs[0]
        self.assertIs(run.font.hidden, False)
        self.assertEqual(run._r.rPr.rFonts.get(qn("w:eastAsia")), "Noto Sans SC")
        self.assertEqual(paragraph.style.paragraph_format.line_spacing, 1.15)

    def test_markers_after_table_stay_outside_cells_and_anchors_are_stable(self) -> None:
        document = Document()
        document.styles.add_style("Test Table Body", 1)
        table = document.add_table(rows=1, cols=1)
        cell_paragraph = table.cell(0, 0).paragraphs[0]
        cell_paragraph.style = "Test Table Body"
        cell_paragraph.text = "表格原文"
        ledger = {"items": [self._item(i, source_type="property_receipt",
                                       source_text="table_body.font.cjk") for i in (1, 2)]}
        locations = insert_inline_manual_review_markers(
            document, ledger, {"table_body": {"style_name": "Test Table Body"}},
        )
        self.assertEqual(table.cell(0, 0).text, "表格原文")
        self.assertEqual([v["anchor_text"] for v in locations.values()], ["表格原文"] * 2)
        self.assertEqual([v["anchor_kind"] for v in locations.values()], ["table"] * 2)
        self.assertEqual(len(document.paragraphs), 2)
        append_manual_review_markers(document, ledger, locations)
        self.assertTrue(document.element.body.index(table._tbl) <
                        document.element.body.index(document.paragraphs[-2]._p))

    def test_serialized_audit_rejects_missing_duplicate_and_nonred_markers(self) -> None:
        for mutation in ("missing", "duplicate", "nonred"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as tmp:
                document = Document()
                ledger = {"items": [self._item()]}
                append_manual_review_markers(document, ledger)
                p = next(p for p in document.paragraphs if p.text.startswith("【MR-"))
                if mutation == "missing":
                    p._p.getparent().remove(p._p)
                elif mutation == "duplicate":
                    p._p.addnext(copy.deepcopy(p._p))
                else:
                    p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
                path = Path(tmp) / "bad.docx"
                document.save(path)
                self.assertFalse(audit_manual_review_markers(path, ledger)["valid"])

    def test_duplicate_ledger_ids_are_rejected_before_insertion(self) -> None:
        document = Document()
        with self.assertRaises(ValueError):
            insert_inline_manual_review_markers(
                document, {"items": [self._item(), self._item()]}, {},
            )

    def test_word_normalized_marker_uses_inherited_color_and_font(self) -> None:
        document = Document()
        ledger = {"items": [self._item()]}
        append_manual_review_markers(document, ledger)
        paragraph = next(p for p in document.paragraphs if p.text.startswith("【MR-"))
        rpr = paragraph.runs[0]._r.rPr
        for tag in ("rFonts", "color", "vanish"):
            node = rpr.find(qn("w:" + tag))
            if node is not None:
                rpr.remove(node)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "normalized.docx"
            document.save(path)
            self.assertTrue(audit_manual_review_markers(path, ledger)["valid"])

    def test_nonblocking_confirmed_semantic_issue_still_gets_marker(self) -> None:
        ledger = build_manual_review_ledger(
            {
                "findings": [{
                    "code": "capability.confirmed_semantic_issue",
                    "blocking": False,
                    "message": "Clause C00074 remains semantically unresolved.",
                    "evidence": [
                        {"kind": "category", "value": "confirmed_semantic_issue"},
                        {"kind": "clause_id", "value": "C00074"},
                    ],
                }],
            },
            [],
            binding={
                "case_id": "bsu",
                "run_id": "run-1",
                "source_sha256": "0" * 64,
                "clause_sha256": "1" * 64,
                "evidence_sha256": "2" * 64,
            },
        )
        self.assertEqual(ledger["summary"]["total"], 1)
        self.assertEqual(ledger["items"][0]["category"], "confirmed_semantic_issue")
        self.assertFalse(ledger["items"][0]["original_blocking"])

    def test_strict_audit_recognizes_manual_marker_as_placeholder(self) -> None:
        pattern = dict(PLACEHOLDER_PATTERNS)["manual_review_marker"]
        self.assertIsNotNone(pattern.search("【MR-0001｜人工待审】"))


if __name__ == "__main__":
    unittest.main()
