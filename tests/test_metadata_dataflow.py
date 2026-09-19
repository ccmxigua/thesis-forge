from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from scripts.extract_semantic_metadata import extract, normalize_semantic_metadata
from scripts.format_spec_validation import load_and_validate, validate_instance
import thesis_format_pipeline  # noqa: E402


def make_requirements(path: Path) -> None:
    document = Document()
    for text in (
        "论文中文题目使用二号黑体，居中排列。",
        "一级标题使用小三号黑体，编号形式为“第一章”。",
        "正文中文使用小四号宋体，英文和数字使用 Times New Roman，行距固定值20磅。",
        "图题使用五号宋体，居中排列，置于图下方。",
        "纸张采用A4，上页边距2.5厘米，下页边距2.0厘米，左页边距3.0厘米，右页边距2.0厘米。",
    ):
        document.add_paragraph(text)
    document.save(path)


class MetadataDataflowTest(unittest.TestCase):
    def test_pipeline_closes_running_manifest_on_unexpected_exception(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            requirements = base / "requirements.docx"
            work = base / "work"
            make_requirements(requirements)
            argv = [
                str(requirements), "tests/sample-thesis.tex", str(base / "out.docx"),
                "--work-dir", str(work), "--analysis-mode", "rule_only",
                "--compliance-mode", "supported_subset",
            ]
            with patch.object(thesis_format_pipeline, "run_step", side_effect=RuntimeError("synthetic failure")):
                with self.assertRaisesRegex(RuntimeError, "synthetic failure"):
                    thesis_format_pipeline.main(argv)
            manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["status"], "failed")
            self.assertEqual(manifest["error_type"], "RuntimeError")
            self.assertEqual(manifest["error"], "synthetic failure")
            self.assertIn("finished_at", manifest)

    def test_sample_semantics_normalize_to_schema_valid_profile_with_provenance(self) -> None:
        source = ROOT / "tests" / "sample-thesis.tex"
        profile = normalize_semantic_metadata(extract(source))
        errors = load_and_validate(profile, ROOT / "schema" / "thesis-profile.schema.json")
        self.assertEqual(errors, [])
        self.assertEqual(profile["degree_level"], "master")
        self.assertEqual(profile["degree_category"], "academic")
        self.assertEqual(profile["writing_language"], "zh")
        self.assertEqual(profile["co_supervisor_count"], 1)
        self.assertEqual(profile["completion_date"], "2026-06")
        self.assertTrue(profile["cover_metadata"])
        self.assertEqual(profile["cover_metadata"]["author_name"], "测试学生甲")
        self.assertEqual(profile["cover_metadata"]["trust"]["source"], "source_document")
        self.assertEqual(profile["provenance"]["trust"]["confirmed"], True)
        self.assertEqual(len(profile["provenance"]["source_sha256"]), 64)
        self.assertEqual(profile["metadata_status"], "complete")
        self.assertEqual(profile["pending_fields"], [])
        serialized = json.dumps(profile, ensure_ascii=False)
        self.assertNotIn("待确认", serialized)
        self.assertNotIn("——", serialized)

    def test_missing_user_fields_are_pending_without_personal_data_placeholders(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "incomplete.tex"
            source.write_text(
                r"""\documentclass{article}
\title{仅有题目}
\englishtitle{Title Only}
\begin{document}
正文。
\end{document}
""",
                encoding="utf-8",
            )
            profile = normalize_semantic_metadata(extract(source))
            errors = load_and_validate(profile, ROOT / "schema" / "thesis-profile.schema.json")
            self.assertEqual(errors, [])
            self.assertEqual(profile["metadata_status"], "pending")
            self.assertNotIn("cover_metadata", profile)
            for field in (
                "cover_metadata.author_name",
                "cover_metadata.student_id",
                "cover_metadata.supervisor_name",
                "cover_metadata.completion_date",
            ):
                self.assertIn(field, profile["pending_fields"])
            serialized = json.dumps(profile, ensure_ascii=False)
            for placeholder in ("待确认", "——", "张三", "000000"):
                self.assertNotIn(placeholder, serialized)

    def test_pending_profile_is_accepted_by_format_spec_nested_schema_but_partial_cover_is_not(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "incomplete.tex"
            source.write_text(r"\title{题目}\englishtitle{Title}\begin{document}\end{document}", encoding="utf-8")
            profile = normalize_semantic_metadata(extract(source))
            format_schema = json.loads((ROOT / "schema" / "format-spec.schema.json").read_text(encoding="utf-8"))
            nested_errors = validate_instance(profile, format_schema["$defs"]["thesisProfile"], format_schema)
            self.assertEqual(nested_errors, [])

            partial = dict(profile)
            partial["cover_metadata"] = {
                "trust": {"source": "source_document", "confirmed": True},
            }
            profile_errors = load_and_validate(partial, ROOT / "schema" / "thesis-profile.schema.json")
            self.assertTrue(any("missing required property" in error for error in profile_errors))

    def test_tex_pipeline_records_and_reuses_one_canonical_profile(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            requirements = base / "requirements.docx"
            output = base / "formatted.docx"
            work = base / "work"
            make_requirements(requirements)
            result = subprocess.run(
                [
                    sys.executable, "scripts/thesis_format_pipeline.py",
                    str(requirements), "tests/sample-thesis.tex", str(output),
                    "--work-dir", str(work), "--analysis-mode", "rule_only",
                    "--compliance-mode", "supported_subset",
                ],
                cwd=ROOT,
                text=True,
                capture_output=True,
            )
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            profile_path = work / "thesis-profile.json"
            spec_path = work / "requirements" / "format-spec.json"
            capability_path = work / "capability-preflight.json"
            manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
            profile = json.loads(profile_path.read_text(encoding="utf-8"))
            spec = json.loads(spec_path.read_text(encoding="utf-8"))
            capability_step = next(step for step in manifest["steps"] if step["name"] == "capability_preflight")
            semantic_step = next(step for step in manifest["steps"] if step["name"] == "semantic_metadata")
            self.assertEqual(spec["thesis_profile"], profile)
            self.assertEqual(Path(manifest["inputs"]["thesis_profile"]["path"]).resolve(), profile_path.resolve())
            self.assertEqual(Path(manifest["inputs"]["semantic_metadata"]["path"]), (work / "semantic-metadata.json").resolve())
            self.assertEqual(
                Path(semantic_step["command"][semantic_step["command"].index("--thesis-profile-out") + 1]).resolve(),
                profile_path.resolve(),
            )
            metadata_index = capability_step["command"].index("--metadata")
            self.assertEqual(Path(capability_step["command"][metadata_index + 1]).resolve(), profile_path.resolve())
            self.assertEqual(manifest["metadata_status"], profile["metadata_status"])
            self.assertTrue(capability_path.exists())

    def test_reused_pipeline_stage_rejects_code_runtime_drift(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            requirements = base / "requirements.docx"
            target = base / "target.docx"
            work = base / "work"
            make_requirements(requirements)
            first = subprocess.run(
                [
                    sys.executable, "scripts/thesis_format_pipeline.py",
                    str(requirements), "tests/sample-thesis.tex", str(target),
                    "--work-dir", str(work), "--analysis-mode", "rule_only",
                    "--compliance-mode", "supported_subset",
                ], cwd=ROOT, text=True, capture_output=True,
            )
            self.assertEqual(first.returncode, 0, first.stderr + first.stdout)
            manifest_path = work / "pipeline-manifest.json"
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            manifest["code_fingerprint"]["sha256"] = "0" * 64
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            second = subprocess.run(
                [
                    sys.executable, "scripts/thesis_format_pipeline.py",
                    str(requirements), "tests/sample-thesis.tex", str(base / "second.docx"),
                    "--work-dir", str(work), "--analysis-mode", "rule_only",
                    "--compliance-mode", "supported_subset", "--allow-existing-work",
                ], cwd=ROOT, text=True, capture_output=True,
            )
            self.assertNotEqual(second.returncode, 0)
            self.assertIn("code/runtime fingerprint", second.stderr)

    def test_tex_pipeline_honors_explicit_source_bound_profile(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            requirements = base / "requirements.docx"
            profile_path = base / "confirmed-profile.json"
            output = base / "formatted.docx"
            work = base / "work"
            make_requirements(requirements)
            profile = normalize_semantic_metadata(extract(ROOT / "tests" / "sample-thesis.tex"))
            profile["degree_category"] = "professional"
            profile["provenance"]["trust"]["source"] = "user_confirmed"
            profile["provenance"]["field_sources"]["degree_category"] = "user_confirmed"
            profile["metadata_status"] = "complete"
            profile["pending_fields"] = []
            profile["pending_metadata"] = []
            profile_path.write_text(json.dumps(profile, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run(
                [
                    sys.executable, "scripts/thesis_format_pipeline.py",
                    str(requirements), "tests/sample-thesis.tex", str(output),
                    "--work-dir", str(work), "--analysis-mode", "rule_only",
                    "--compliance-mode", "supported_subset",
                    "--case-id", "profile-flow",
                    "--thesis-profile", str(profile_path),
                ],
                cwd=ROOT, text=True, capture_output=True,
            )
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
            extraction = json.loads((work / "requirements" / "extraction-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["case_id"], "profile-flow")
            self.assertEqual(extraction["case_id"], "profile-flow")
            self.assertEqual(extraction["runtime_context"]["case_id"], "profile-flow")
            canonical = json.loads((work / "thesis-profile.json").read_text(encoding="utf-8"))
            self.assertEqual(canonical, profile)
            manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(
                manifest["inputs"]["thesis_profile_source"]["path"],
                str(profile_path.resolve()),
            )


if __name__ == "__main__":
    unittest.main()
