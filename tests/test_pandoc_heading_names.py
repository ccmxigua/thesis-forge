from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from postprocess_docx import _patch_styles, qn


class PandocHeadingNameTests(unittest.TestCase):
    def test_all_ooxml_custom_style_true_values_preserve_the_name(self):
        for custom_value in ('1', 'true', 'on'):
            with self.subTest(customStyle=custom_value), tempfile.TemporaryDirectory() as td:
                root = Path(td)
                (root / 'word').mkdir()
                styles_path = root / 'word/styles.xml'
                styles = ET.Element(qn('w', 'styles'))
                heading = ET.SubElement(styles, qn('w', 'style'), {
                    qn('w', 'type'): 'paragraph', qn('w', 'styleId'): 'Heading1',
                    qn('w', 'customStyle'): custom_value})
                ET.SubElement(heading, qn('w', 'name'), {qn('w', 'val'): 'Heading 1'})
                ET.ElementTree(styles).write(styles_path, encoding='utf-8', xml_declaration=True)
                _patch_styles(root)
                checked = ET.parse(styles_path).getroot().find(qn('w', 'style'))
                self.assertEqual(checked.find(qn('w', 'name')).get(qn('w', 'val')), 'Heading 1')
                self.assertEqual(checked.get(qn('w', 'customStyle')), custom_value)

    def test_canonical_heading_name_collision_rejects_without_writing(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root / 'word').mkdir()
            styles_path = root / 'word/styles.xml'
            styles = ET.Element(qn('w', 'styles'))
            for style_id, name in (('Heading1', 'Heading 1'), ('OtherHeading', 'heading 1')):
                style = ET.SubElement(styles, qn('w', 'style'), {
                    qn('w', 'type'): 'paragraph', qn('w', 'styleId'): style_id})
                ET.SubElement(style, qn('w', 'name'), {qn('w', 'val'): name})
            ET.ElementTree(styles).write(styles_path, encoding='utf-8', xml_declaration=True)
            original = styles_path.read_bytes()
            with self.assertRaisesRegex(ValueError, 'ambiguous built-in heading style'):
                _patch_styles(root)
            self.assertEqual(styles_path.read_bytes(), original)

    def test_converter_heading_name_round_trips_through_python_docx(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / 'source.docx'
            doc = Document()
            doc.add_paragraph('Chapter', 'Heading 1')
            doc.save(source)
            unpacked = root / 'unpacked'
            with zipfile.ZipFile(source) as archive:
                archive.extractall(unpacked)
            styles = unpacked / 'word/styles.xml'
            tree = ET.parse(styles)
            for style in tree.getroot().findall(qn('w', 'style')):
                if style.get(qn('w', 'styleId')) == 'Heading1':
                    style.find(qn('w', 'name')).set(qn('w', 'val'), 'Heading 1')
                if style.get(qn('w', 'styleId')) == 'Heading2':
                    style.find(qn('w', 'name')).set(qn('w', 'val'), 'Custom heading label')
            tree.write(styles, encoding='utf-8', xml_declaration=True)
            _patch_styles(unpacked)
            result = root / 'result.docx'
            with zipfile.ZipFile(result, 'w') as archive:
                for path in unpacked.rglob('*'):
                    if path.is_file():
                        archive.write(path, path.relative_to(unpacked))
            verified = Document(result)
            self.assertEqual(verified.styles['Heading 1'].style_id, 'Heading1')
            self.assertEqual(verified.paragraphs[0].style.style_id, 'Heading1')
            self.assertEqual(verified.paragraphs[0].text, 'Chapter')
            self.assertEqual(verified.styles['Custom heading label'].style_id, 'Heading2')
