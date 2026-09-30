"""Human work stays atomic, source-bound and non-executable through DOCX markers."""
from __future__ import annotations

import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))
import host_agent_bridge as bridge
import native_semantic_review as native
import requirements_engine as engine
import thesis_format_pipeline as pipeline
from pending_source_work import (
    ANTI_EXCERPT, REPETITION_SCOPE, SECTION_CHOICE,
    compile_pending_source_work, pending_work_inventory_is_bound,
)
from semantic_contract import attach_request_provenance, request_body_sha256, sha256_json
import test_authoring_correction_routing as authoring_tests
import test_table_structure_review as receipt_tests
import test_manual_review as manual_tests
from manual_review import build_manual_review_ledger
from apply_format_spec import append_manual_review_markers, audit_manual_review_markers

AUTHOR_SOURCE = "本部分主要撰写国内的研究现状，不能是文献资料的简单摘录，计算在重复率内，需要分类、总结、归纳"
CHOICE_SOURCE = "若论文研究确实无国外资料，本部分也可删除"


def fixture(source=AUTHOR_SOURCE, *, classification=None):
    facts = compile_pending_source_work(source)
    authoring = bool(facts) and facts[0]["code"] != SECTION_CHOICE
    chunk, candidate, response = authoring_tests.AuthoringCorrectionRoutingTests.fixture(
        classification or ("requires_source_content" if authoring else "requires_source_verification"),
        source=source,
    )
    # Dynamic IDs are deliberately unrelated to BSU or its clause numbering.
    result = response["results"][0]
    result["verdict"] = "source_content_pending" if authoring else "source_content_verification_pending"
    result["identified_obligations"] = ([{
        "source_quote": source, "disposition": "authoring_content_pending",
        "obligation_summary": "撰写研究现状，分类、总结和归纳。", "requirement_refs": [],
    }] if authoring else []) + [{
        "source_quote": source, "disposition": "source_content_verification_pending",
        "pending_work_code": fact["code"], "obligation_summary": fact["source_quote"],
        "requirement_refs": [],
    } for fact in facts]
    return chunk, candidate, response


class PendingSourceWorkTests(unittest.TestCase):
    def test_closed_grammars_preserve_current_offsets_hashes_and_unverified_condition(self):
        for area in ("国内", "国外"):
            source = AUTHOR_SOURCE.replace("国内", area)
            facts = compile_pending_source_work(source)
            self.assertEqual([f["code"] for f in facts], [ANTI_EXCERPT, REPETITION_SCOPE])
            for fact in facts:
                self.assertEqual(source[fact["start"]:fact["end"]], fact["source_quote"])
                self.assertEqual(fact["source_sha256"], hashlib.sha256(source.encode()).hexdigest())
                self.assertFalse(fact["execution_authorized"])
        fact = compile_pending_source_work(CHOICE_SOURCE)[0]
        self.assertFalse(fact["condition_verified"])
        self.assertEqual(fact["source_quote"], CHOICE_SOURCE)
        self.assertEqual(fact["decision"], "pending_human_decision")
        self.assertEqual(fact["code"], SECTION_CHOICE)

    def test_unknown_quoted_examples_negated_authoring_and_added_conditions_do_not_compile(self):
        for source in (
            "示例：" + AUTHOR_SOURCE, "“" + AUTHOR_SOURCE + "”", "如果需要，" + AUTHOR_SOURCE,
            AUTHOR_SOURCE.replace("主要撰写", "不得撰写"),
            "示例：" + CHOICE_SOURCE, "“" + CHOICE_SOURCE + "”",
            CHOICE_SOURCE + "，但本专业必须保留", "本部分也可删除", "若资料不全，本部分可删除",
        ):
            with self.subTest(source=source):
                self.assertEqual(compile_pending_source_work(source), [])
        self.assertEqual(compile_pending_source_work(
            "本部分主要撰写国内的研究现状，引用说明：“计算在重复率内，不能是文献资料的简单摘录”"
        ), [])

    def test_valid_mixed_inventory_is_pending_not_requirements_or_compliance(self):
        for source in (AUTHOR_SOURCE, CHOICE_SOURCE):
            chunk, candidate, response = fixture(source)
            before = copy.deepcopy((candidate, response))
            checks = native.build_obligation_coverage_request(candidate, chunk, run_id="new-run", chunk_index=1)["checks"]
            results = native.validate_obligation_coverage_response(response, checks)
            self.assertEqual(results, response["results"])
            self.assertEqual((candidate, response), before)
            self.assertEqual(candidate["requirements"], [])
            pipeline.enforce_obligation_review_output_policy(results, output_policy="review_draft")
            with self.assertRaisesRegex(ValueError, "submission output is blocked"):
                pipeline.enforce_obligation_review_output_policy(results, output_policy="submission")

    def test_wrong_missing_duplicate_code_range_or_refs_cannot_be_promoted(self):
        for mutation in ("missing", "duplicate", "unknown", "wrong_quote", "refs", "wrong_disposition", "consistent", "all_authoring", "no_code", "empty"):
            chunk, candidate, response = fixture()
            result = response["results"][0]
            items = result["identified_obligations"]
            if mutation == "missing":
                items.pop()
            elif mutation == "duplicate":
                items.append(copy.deepcopy(items[-1]))
            elif mutation == "unknown":
                items[-1]["pending_work_code"] = "unknown_work"
            elif mutation == "wrong_quote":
                items[-1]["source_quote"] = "需要分类、总结、归纳"
            elif mutation == "refs":
                items[-1]["requirement_refs"] = ["forged"]
            elif mutation == "wrong_disposition":
                items[-1]["disposition"] = "authoring_content_pending"
            elif mutation == "consistent":
                result["verdict"] = "consistent"
            elif mutation == "all_authoring":
                for item in items:
                    item.pop("pending_work_code", None)
                    item["disposition"] = "authoring_content_pending"
            elif mutation == "no_code":
                items[-1].pop("pending_work_code")
            elif mutation == "empty":
                result.update(verdict="consistent", identified_obligations=[])
                candidate["clause_reviews"][0]["classification"] = "informational"
            checks = native.build_obligation_coverage_request(candidate, chunk, run_id="new-run", chunk_index=1)["checks"]
            with self.subTest(mutation=mutation), self.assertRaises(native.NativeSemanticReviewError):
                native.validate_obligation_coverage_response(response, checks)

    def test_stale_context_cannot_authorize_human_work(self):
        chunk, candidate, response = fixture()
        checks = native.build_obligation_coverage_request(candidate, chunk, run_id="new-run", chunk_index=1)["checks"]
        for key, value in (("source_sha256", "0" * 64), ("execution_authorized", True), ("source_quote", "邻接条款")):
            bad = copy.deepcopy(checks)
            bad[0]["review_context"]["pending_source_work"][0][key] = value
            with self.subTest(key=key), self.assertRaisesRegex(native.NativeSemanticReviewError, "stale"):
                native.validate_obligation_coverage_response(copy.deepcopy(response), bad)

    def test_duplicate_source_atoms_fail_closed_instead_of_collapsing(self):
        source = AUTHOR_SOURCE + "，计算在重复率内"
        chunk, candidate, response = fixture(source)
        self.assertFalse(pending_work_inventory_is_bound(source,
            response["results"][0]["identified_obligations"], authoring=True))
        checks = native.build_obligation_coverage_request(candidate, chunk, run_id="new-run", chunk_index=1)["checks"]
        with self.assertRaises(native.NativeSemanticReviewError):
            native.validate_obligation_coverage_response(response, checks)

    def test_unrepresented_duty_is_not_automatically_converted_to_pending(self):
        chunk, candidate, response = authoring_tests.AuthoringCorrectionRoutingTests.fixture("requires_source_content", source=AUTHOR_SOURCE)
        original = copy.deepcopy(response)
        with tempfile.TemporaryDirectory() as td:
            with self.assertRaises(bridge.IndependentObligationReviewError):
                authoring_tests.AuthoringCorrectionRoutingTests.run_review(candidate, chunk, td, response, [])
        self.assertEqual(response, original)
        self.assertFalse(pending_work_inventory_is_bound(
            AUTHOR_SOURCE, response["results"][0]["identified_obligations"], authoring=True))

    def test_conditional_primary_classification_requires_bounded_correction_not_pass(self):
        chunk, candidate, response = fixture(CHOICE_SOURCE, classification="informational")
        checks = native.build_obligation_coverage_request(candidate, chunk, run_id="new-run", chunk_index=1)["checks"]
        before = copy.deepcopy(candidate)
        with self.assertRaises(native.SourceVerificationClassificationCorrectionRequiredError):
            native.validate_obligation_coverage_response(response, checks)
        self.assertEqual(candidate, before)
        self.assertEqual(candidate["requirements"], [])

    def test_incomplete_is_preserved_for_scored_draft_but_never_submission(self):
        chunk, candidate, response = fixture(CHOICE_SOURCE)
        result = response["results"][0]
        result["verdict"] = "incomplete"
        item = result["identified_obligations"][0]
        item.pop("pending_work_code")
        item["disposition"] = "unrepresented"
        checks = native.build_obligation_coverage_request(candidate, chunk, run_id="new-run", chunk_index=1)["checks"]
        results = native.validate_obligation_coverage_response(response, checks)
        self.assertEqual(results[0]["verdict"], "incomplete")
        self.assertEqual(results[0]["identified_obligations"][0]["disposition"], "unrepresented")
        with self.assertRaisesRegex(ValueError, "is incomplete"):
            pipeline.enforce_obligation_review_output_policy(results, output_policy="submission")
        self.assertEqual(pipeline.enforce_obligation_review_output_policy(
            results, output_policy="review_draft"), [results[0]["check_id"]])
        self.assertEqual(results[0]["verdict"], "incomplete")

    def test_valid_inventory_verdict_retry_is_bounded_and_changes_no_candidate_fields(self):
        chunk, candidate, first = fixture()
        first["results"][0]["verdict"] = "incomplete"
        second = copy.deepcopy(first)
        second["results"][0]["verdict"] = "source_content_pending"
        before = copy.deepcopy(candidate)
        calls = []
        with tempfile.TemporaryDirectory() as td:
            pointer = receipt_tests.TableStructureReviewTests.run_receipt_review(
                candidate, chunk, td, [first, second], calls)
            self.assertEqual(pointer["provider_attempt"], 2)
            self.assertEqual(len(calls), 2)
            self.assertEqual(calls[0]["checks"], calls[1]["checks"])
            self.assertEqual(candidate, before)
            ledger = json.loads((Path(td) / pointer["obligation_analysis_ledger_path"]).read_text())
            self.assertEqual(len(ledger["obligations"]), 3)
            self.assertFalse(ledger["submission_ready"])
        with tempfile.TemporaryDirectory() as td:
            with self.assertRaises(bridge.IndependentObligationReviewError):
                receipt_tests.TableStructureReviewTests.run_receipt_review(
                    candidate, chunk, td, [first, first], [])

    def test_current_source_conditional_retry_changes_only_classification(self):
        chunk, parent, response = fixture(CHOICE_SOURCE, classification="informational")
        parent.update(unsupported_items=[], reported_conflicts=[])
        # Real current response schema, not a mocked host validator.
        full = engine.build_llm_request([], chunk["clauses"],
            {"evidence": list(chunk["evidence_context"].values())}, {}, "full", contract_version="3.0")
        full = attach_request_provenance(full, source_sha256="a" * 64,
            evidence_doc={"evidence": list(chunk["evidence_context"].values())},
            clauses=chunk["clauses"], run_id="new-run")
        chunk["provenance"] = copy.deepcopy(full["provenance"])
        parent["provenance"] = copy.deepcopy(full["provenance"])
        chunk["requirement_contract"] = full["requirement_contract"]
        chunk["response_schema"] = full["response_schema"]
        current = copy.deepcopy(parent)
        current["clause_reviews"][0]["classification"] = "requires_source_verification"
        clause_id = chunk["clauses"][0]["id"]
        record = {
            "code": "independent_obligation_review_incomplete", "clause_id": clause_id,
            "json_pointer": "$.clause_reviews[0].classification",
            "baseline_classification": "informational", "evidence_ids": ["E-dynamic"],
            "missing_source_quotes": [CHOICE_SOURCE],
            "candidate_response_sha256": bridge._response_sha256(
                bridge._bind_current_invocation_provenance(parent, chunk["provenance"])),
            "candidate_semantic_sha256": bridge._response_sha256(bridge._semantic_retry_view(parent)),
            "review_request_sha256": "e" * 64, "review_response_sha256": "f" * 64,
            "source_reference_compilation_sha256": "a" * 64,
            "primary_retry_authorization": "source_bound_existing_content_verification_reclassification_v1",
        }
        repaired, audit = bridge._v3_source_verification_reclassification_response(parent, current, [record], chunk=chunk)
        self.assertEqual(repaired, current)
        self.assertFalse(audit["submission_ready"])
        for mutation in ("reason", "requirements", "parent_hash", "quote"):
            changed = copy.deepcopy(current)
            bad_record = copy.deepcopy(record)
            if mutation == "reason":
                changed["clause_reviews"][0]["reason"] = "Condition is met; delete the chapter."
            elif mutation == "requirements":
                changed["requirements"].append({"role": "body_text", "properties": {}})
            elif mutation == "parent_hash":
                bad_record["candidate_response_sha256"] = "0" * 64
            elif mutation == "quote":
                bad_record["missing_source_quotes"] = ["不存在的相邻来源"]
            with self.subTest(mutation=mutation):
                rejected, _ = bridge._v3_source_verification_reclassification_response(parent, changed, [bad_record], chunk=chunk)
                self.assertIsNone(rejected)

    def test_real_receipt_reconstruction_manual_ledger_and_serialized_red_markers(self):
        for source, expected_count in ((AUTHOR_SOURCE, 6), (AUTHOR_SOURCE.replace("国内", "国外"), 6), (CHOICE_SOURCE, 1)):
            chunk, candidate, response = fixture(source)
            if expected_count == 6:
                # Keep all four independently identified authoring duties,
                # not only the top-level writing instruction. This mirrors
                # the incident's atomic inventory, with dynamic IDs.
                authoring_item = response["results"][0]["identified_obligations"][0]
                response["results"][0]["identified_obligations"].extend(
                    {**copy.deepcopy(authoring_item), "obligation_summary": summary}
                    for summary in ("分类组织研究现状", "总结研究现状", "归纳研究现状")
                )
            evidence_doc = {"evidence": list(chunk["evidence_context"].values())}
            full = engine.build_llm_request([], chunk["clauses"], evidence_doc, {}, "full", contract_version="3.0")
            full["case_id"] = chunk["case_id"]
            full = attach_request_provenance(full, source_sha256="a" * 64,
                evidence_doc=evidence_doc, clauses=chunk["clauses"], run_id="new-run")
            with self.subTest(source=source), tempfile.TemporaryDirectory() as td:
                root = Path(td)
                engine.prepare_host_agent_review_packets(full, chunk["clauses"], evidence_doc,
                                                         "a" * 64, root, chunk_size=1)
                chunk = json.loads((root / "llm-request-chunks.json").read_text())[0]
                candidate["provenance"] = copy.deepcopy(chunk["provenance"])
                candidate.update(unsupported_items=[], reported_conflicts=[])
                self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
                pointer = receipt_tests.TableStructureReviewTests.run_receipt_review(candidate, chunk, td, [response], [])
                envelope = json.loads((root / pointer["audit_path"]).read_text())
                bridge._validate_completed_obligation_ledger_chain(root, envelope, pointer, candidate, chunk,
                                                                   chunk_index=1, attempt=1)
                candidate_path = root / "accepted-candidate.json"
                bridge._write_json(candidate_path, candidate)
                audit = {"chunk_count": 1, "adapter_id": "codex", "host_runtime": "codex",
                         "chunk_lifecycle": [{"chunk_index": 1, "status": "completed", "remote_operation_state": "completed"}],
                         "chunk_runs": [{"chunk_index": 1, "response_path": str(candidate_path),
                                         "accepted_response_sha256": sha256_json(candidate),
                                         "independent_obligation_review": pointer}]}
                reviews = pipeline._validate_independent_obligation_receipts(
                    audit=audit, review_root=root, expected_run_id="new-run",
                    expected_request_body_sha=request_body_sha256(full),
                    expected_request_envelope_sha=None, expected_request_file_sha=None,
                    output_policy="review_draft")
                gates = pipeline._source_content_pending_release_gates(reviews,
                    clauses=chunk["clauses"], evidence_doc=evidence_doc, expected_run_id="new-run")
                gates += pipeline._source_content_verification_release_gates(reviews,
                    clauses=chunk["clauses"], evidence_doc=evidence_doc, expected_run_id="new-run")
                self.assertEqual(len(gates), expected_count)
                self.assertEqual({g.get("pending_work_code") for g in gates if g.get("pending_work_code")},
                                 {f["code"] for f in compile_pending_source_work(source)})
                ledger = build_manual_review_ledger({}, [], release_gates=gates,
                    binding=manual_tests.ManualReviewTests._binding(run_id="new-run", case_id=chunk["case_id"]))
                self.assertEqual(len(ledger["items"]), expected_count)
                self.assertFalse(ledger["submission_ready"])
                self.assertEqual(len({i["manual_obligation_id"] for i in ledger["items"]}), expected_count)
                document = Document()
                document.add_paragraph("原有章节与论文内容不删除。")
                append_manual_review_markers(document, ledger)
                path = root / "pending-draft.docx"
                document.save(path)
                marker_audit = audit_manual_review_markers(path, ledger)
                self.assertTrue(marker_audit["valid"], marker_audit)
                self.assertIn("原有章节与论文内容不删除。", [p.text for p in Document(path).paragraphs])
                self.assertTrue(all(g["execution_authorized"] is False for g in gates))
                for mutation in ("code", "hash", "evidence", "duplicate"):
                    bad_reviews = copy.deepcopy(reviews)
                    pending = bad_reviews[0]["source_content_verification_items"][0]
                    if mutation == "code":
                        pending["pending_work_code"] = "unregistered_code"
                    elif mutation == "hash":
                        pending["source_text_sha256"] = "0" * 64
                    elif mutation == "evidence":
                        pending["evidence_ids"] = ["other-evidence"]
                    elif mutation == "duplicate":
                        bad_reviews[0]["source_content_verification_items"].append(copy.deepcopy(pending))
                    with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                        pipeline._source_content_verification_release_gates(bad_reviews,
                            clauses=chunk["clauses"], evidence_doc=evidence_doc, expected_run_id="new-run")
                for bad_run in ("old-run", "different-run"):
                    with self.assertRaises(ValueError):
                        pipeline._source_content_verification_release_gates(reviews,
                            clauses=chunk["clauses"], evidence_doc=evidence_doc, expected_run_id=bad_run)


if __name__ == "__main__":
    unittest.main()
