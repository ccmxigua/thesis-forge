from __future__ import annotations

import copy
import argparse
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from docx import Document
from types import SimpleNamespace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import requirements_engine as engine  # noqa: E402
import offline_review_receipt  # noqa: E402
import apply_format_spec as apply_spec  # noqa: E402
import host_agent_bridge as bridge  # noqa: E402
from semantic_contract import attach_request_provenance, sha256_file, sha256_json  # noqa: E402
from thesis_format_pipeline import (  # noqa: E402
    runtime_code_fingerprint,
    validate_offline_merge_receipt,
    validate_semantic_review_configuration,
)


def _write(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False) + "\n", encoding="utf-8")


class PortableReviewDraftTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.work = Path(self.temporary.name) / "run"
        review = self.work / "review"
        requirements = review / "requirements"
        self.response_path = review / "host-agent-response.json"
        self.receipt_path = requirements / "merge-receipt.json"
        self.ledger_path = requirements / "semantic-review-ledger.json"
        self.marker_path = requirements / "merge-commit.json"
        self.response = {"contract_version": "3.0", "requirements": [], "clause_reviews": []}
        _write(self.response_path, self.response)
        ledger = {"response_sha256": sha256_json(self.response)}
        _write(self.ledger_path, ledger)
        self.extraction = {
            "run_id": "fresh-run", "llm_request_body_sha256": "a" * 64,
            "llm_request_envelope_sha256": "b" * 64,
            "llm_request_file_sha256": "c" * 64,
            "runtime_context": None,
        }
        self.receipt = {
            "status": "merged", "protocol": "host_agent_semantic_review",
            "run_id": "fresh-run", "request_body_sha256": "a" * 64,
            "request_envelope_sha256": "b" * 64,
            "request_file_sha256": "c" * 64, "runtime_context": None,
            "aggregate_sha256": sha256_json(self.response),
            "merged_response_path": str(self.response_path),
            "semantic_review_ledger_path": str(self.ledger_path),
            "semantic_review_ledger_sha256": sha256_json(ledger),
            "merge_commit_path": str(self.marker_path),
        }
        _write(self.receipt_path, self.receipt)
        self._write_marker()

    def _write_marker(self) -> None:
        marker = {
            "status": "committed", "protocol": "host_agent_semantic_review_merge",
            "run_id": "fresh-run", "aggregate_sha256": sha256_json(self.response),
            "artifacts": [
                {"path": path.relative_to(self.work / "review").as_posix(),
                 "sha256": sha256_file(path)}
                for path in (self.response_path, self.receipt_path, self.ledger_path)
            ],
        }
        _write(self.marker_path, marker)

    def _check(self) -> dict:
        return validate_offline_merge_receipt(
            response_path=self.response_path,
            receipt_path=self.receipt_path,
            extraction_manifest=self.extraction,
            work=self.work,
        )

    def test_current_run_merge_is_analysis_only(self) -> None:
        result = self._check()
        self.assertEqual(result["status"], "offline_merged_without_independent_review")
        self.assertIs(result["submission_ready"], False)
        self.assertIs(result["independent_review_verified"], False)
        self.assertIs(result["provider_model_verified"], False)
        self.assertEqual(result["aggregate_sha256"], sha256_json(self.response))

    def test_evaluation_units_accept_only_a_revalidated_offline_merge_receipt(self) -> None:
        extraction_path = self.work / "review" / "requirements" / "extraction-manifest.json"
        _write(extraction_path, self.extraction)
        source_clauses_path = self.work / "review" / "requirements" / "requirement-clauses.json"
        source_clauses_path.write_text("[]\n", encoding="utf-8")
        verified_offline = self._check()
        pipeline_manifest = {
            "run_id": "fresh-run", "work_dir": str(self.work),
            "offline_review_receipt": verified_offline,
            "source_clause_record": {
                "path": str(source_clauses_path.resolve()),
                "bytes": source_clauses_path.stat().st_size,
                "sha256": sha256_file(source_clauses_path),
            },
        }
        spec = {"requirements": [{"id": "R1", "evaluation_units": []}]}
        with patch.object(apply_spec, "validate_requirement_evaluation_units", return_value=([], {})) as validate_units:
            errors, _ = apply_spec.validate_current_evaluation_units(
                spec, pipeline_manifest=pipeline_manifest,
                source_clauses_path=source_clauses_path,
                source_extraction_manifest_path=extraction_path,
            )
        self.assertEqual(errors, [])
        self.assertEqual(validate_units.call_args.kwargs["expected_response_sha256"],
                         sha256_json(self.response))

        pipeline_manifest["offline_review_receipt"] = copy.deepcopy(verified_offline)
        pipeline_manifest["offline_review_receipt"]["response"]["sha256"] = "f" * 64
        errors, _ = apply_spec.validate_current_evaluation_units(
            spec, pipeline_manifest=pipeline_manifest,
            source_clauses_path=source_clauses_path,
            source_extraction_manifest_path=extraction_path,
        )
        self.assertTrue(any("offline_review_receipt_invalid" in item for item in errors))

    def test_review_spec_without_evaluation_units_does_not_synthesize_projection(self) -> None:
        extraction_path = self.work / "review" / "requirements" / "extraction-manifest.json"
        _write(extraction_path, self.extraction)
        source_clauses_path = self.work / "review" / "requirements" / "requirement-clauses.json"
        source_clauses_path.write_text("[]\n", encoding="utf-8")
        pipeline_manifest = {
            "run_id": "fresh-run", "work_dir": str(self.work),
            "offline_review_receipt": self._check(),
            "source_clause_record": {
                "path": str(source_clauses_path.resolve()),
                "bytes": source_clauses_path.stat().st_size,
                "sha256": sha256_file(source_clauses_path),
            },
        }
        with patch.object(apply_spec, "validate_requirement_evaluation_units") as validate_units:
            errors, projected = apply_spec.validate_current_evaluation_units(
                {"requirements": [{"id": "R1"}]},
                pipeline_manifest=pipeline_manifest,
                source_clauses_path=source_clauses_path,
                source_extraction_manifest_path=extraction_path,
            )
        self.assertEqual(errors, [])
        self.assertEqual(projected, {})
        validate_units.assert_not_called()

    def test_portable_cli_works_without_codex_or_other_native_cli(self) -> None:
        extraction_path = self.work / "review" / "requirements" / "extraction-manifest.json"
        _write(extraction_path, self.extraction)
        result = subprocess.run(
            [sys.executable, str(ROOT / "scripts" / "offline_review_receipt.py"),
             "--work-dir", str(self.work), "--response", str(self.response_path),
             "--receipt", str(self.receipt_path),
             "--extraction-manifest", str(extraction_path)],
            env={"PATH": ""}, capture_output=True, text=True, check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertFalse(json.loads(result.stdout)["submission_ready"])

    def test_portable_cli_rejects_duplicate_keys_and_nonfinite_numbers(self) -> None:
        for raw in ('{"run_id":"fresh-run","run_id":"stale"}',
                    '{"run_id":NaN}'):
            with self.subTest(raw=raw):
                extraction_path = self.work / "review" / "requirements" / "extraction-manifest.json"
                extraction_path.write_text(raw, encoding="utf-8")
                with self.assertRaises(ValueError):
                    offline_review_receipt._object(extraction_path, label="offline extraction manifest")

    def test_non_object_receipt_and_response_fail_closed(self) -> None:
        for path in (self.receipt_path, self.response_path):
            with self.subTest(path=path):
                original = path.read_bytes()
                path.write_text("[]", encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "must be a JSON object"):
                    self._check()
                path.write_bytes(original)

    def test_old_run_and_byte_changes_are_rejected(self) -> None:
        self.extraction["run_id"] = "different-run"
        with self.assertRaisesRegex(ValueError, "not bound"):
            self._check()
        self.extraction["run_id"] = "fresh-run"
        self.response["requirements"].append({"role": "cover"})
        _write(self.response_path, self.response)
        with self.assertRaisesRegex(ValueError, "not bound"):
            self._check()

    def test_modified_ledger_or_receipt_cannot_be_resealed_without_marker(self) -> None:
        _write(self.ledger_path, {"response_sha256": "0" * 64})
        with self.assertRaisesRegex(ValueError, "ledger"):
            self._check()
        _write(self.ledger_path, {"response_sha256": sha256_json(self.response)})
        changed = copy.deepcopy(self.receipt)
        changed["semantic_review_ledger_path"] = str(self.work / "other.json")
        _write(self.receipt_path, changed)
        with self.assertRaises(ValueError):
            self._check()

    def test_marker_path_traversal_is_rejected(self) -> None:
        marker = json.loads(self.marker_path.read_text(encoding="utf-8"))
        marker["artifacts"][0]["path"] = "../../outside.json"
        _write(self.marker_path, marker)
        with self.assertRaises(ValueError):
            self._check()

    def test_real_packet_merge_and_receipt_need_no_native_cli(self) -> None:
        portable_work = Path(self.temporary.name) / "portable-run"
        review_dir = portable_work / "review" / "requirements"
        response_path = portable_work / "review" / "host-agent-response.json"
        receipt_path = review_dir / "merge-receipt.json"
        source = "本节为说明性标题。"
        digest = hashlib.sha256(source.encode("utf-8")).hexdigest()
        clauses = [{
            "id": "C1", "text": source, "evidence_ids": ["E1"],
            "source_kind": "paragraph", "location": {"part": "document", "order": 1},
            "source_span": {
                "evidence_id": "E1", "start_offset": 0, "end_offset": len(source),
                "text": source, "source_sha256": digest,
            },
        }]
        evidence = {"evidence": [{"id": "E1", "text": source, "kind": "paragraph"}]}
        request = engine.build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
        )
        request["runtime_context"] = {"code_fingerprint_sha256": "c" * 64}
        request = attach_request_provenance(
            request, source_sha256=digest, evidence_doc=evidence,
            clauses=clauses, run_id="portable-agent-run",
        )
        engine.prepare_host_agent_review_packets(
            request, clauses, evidence, digest, review_dir, chunk_size=1,
        )
        chunk = json.loads(
            (review_dir / "llm-request-chunks.json").read_text(encoding="utf-8")
        )[0]
        response = {
            "contract_version": "3.0", "provenance": chunk["provenance"],
            "requirements": [],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "informational",
                "reason": "The cited paragraph is a section label, not a format instruction.",
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        _write(review_dir / chunk["batch"]["response_filename"], response)
        engine.merge_host_agent_review_packets(
            review_dir, response_out=response_path,
        )
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        extraction = {
            "run_id": "portable-agent-run",
            "llm_request_body_sha256": receipt["request_body_sha256"],
            "llm_request_envelope_sha256": receipt["request_envelope_sha256"],
            "llm_request_file_sha256": receipt["request_file_sha256"],
            "runtime_context": receipt["runtime_context"],
        }
        _write(review_dir / "extraction-manifest.json", extraction)
        result = validate_offline_merge_receipt(
            response_path=response_path, receipt_path=receipt_path,
            extraction_manifest=extraction, work=portable_work,
        )
        self.assertEqual(result["run_id"], "portable-agent-run")
        self.assertFalse(result["submission_ready"])
        parent_identity = offline_review_receipt.validate_offline_parent_merge_receipt(
            parent_receipt_path=receipt_path,
        )
        self.assertEqual(parent_identity["parent_run_id"], "portable-agent-run")
        self.assertEqual(
            parent_identity["inherited_request_code_fingerprint_sha256"], "c" * 64,
        )

        child_work = Path(self.temporary.name) / "derived-run"
        child_artifacts = child_work / "review" / "requirements"
        child_response_path = child_artifacts / "host-agent-response.json"
        child_receipt_path = child_artifacts / "merge-receipt.json"
        engine.merge_host_agent_review_packets(
            review_dir, response_out=child_response_path,
            response_projector=bridge._materialize_fixed_declaration_source_text,
            artifact_dir=child_artifacts, parent_receipt_path=receipt_path,
            derivation_code_fingerprint_sha256=runtime_code_fingerprint()["sha256"],
        )
        with self.assertRaisesRegex(ValueError, "fixed declaration projection audit"):
            offline_review_receipt.validate_offline_parent_merge_receipt(
                parent_receipt_path=receipt_path,
                child_receipt_path=child_receipt_path,
                child_response_path=child_response_path,
                child_work=child_work,
                current_code_fingerprint_sha256=runtime_code_fingerprint()["sha256"],
            )

    def _make_deterministic_declaration_child(self) -> tuple[Path, Path, Path, str]:
        incident = json.loads(
            (ROOT / "tests" / "fixtures" / "declaration-informational-heading-incident.json")
            .read_text(encoding="utf-8")
        )
        source = copy.deepcopy(incident["source"])
        clauses = source["clauses"]
        evidence = {
            "evidence": list(source["evidence_context"].values()),
            "structure_evidence": {"sections": [{
                "first_paragraphs": [{"text": "摘要", "style_name": "Abstract Title CN"}],
            }]},
        }
        run_id = "deterministic-declaration-parent-run"
        request = engine.build_llm_request(
            [], clauses, evidence, {}, "full", contract_version="3.0",
        )
        request.update(source)
        request["runtime_context"] = {"code_fingerprint_sha256": "c" * 64}
        request = attach_request_provenance(
            request, source_sha256="d" * 64, evidence_doc=evidence,
            clauses=clauses, run_id=run_id,
        )
        root = Path(self.temporary.name) / "deterministic-projection"
        parent_work = root / "parent"
        parent_dir = parent_work / "review" / "requirements"
        engine.prepare_host_agent_review_packets(
            request, clauses, evidence, "d" * 64, parent_dir, chunk_size=100,
        )
        chunk = json.loads((parent_dir / "llm-request-chunks.json").read_text())[0]
        raw_response = bridge.normalize_native_response(
            copy.deepcopy(incident["raw"]), chunk["response_schema"],
        )
        raw_response["provenance"] = chunk["provenance"]
        _write(parent_dir / chunk["batch"]["response_filename"], raw_response)
        parent_response = parent_dir / "parent-aggregate.json"
        engine.merge_host_agent_review_packets(parent_dir, response_out=parent_response)
        parent_receipt = parent_dir / "merge-receipt.json"
        parent_value = json.loads(parent_receipt.read_text(encoding="utf-8"))
        _write(parent_dir / "extraction-manifest.json", {
            "run_id": run_id,
            "llm_request_body_sha256": parent_value["request_body_sha256"],
            "llm_request_envelope_sha256": parent_value["request_envelope_sha256"],
            "llm_request_file_sha256": parent_value["request_file_sha256"],
            "runtime_context": parent_value["runtime_context"],
        })

        current_fingerprint = runtime_code_fingerprint()["sha256"]
        child_work = root / "child"
        child_dir = child_work / "review" / "requirements"
        child_response = child_dir / "host-agent-response.json"
        merged = subprocess.run(
            [
                sys.executable, str(ROOT / "scripts" / "merge_host_agent_review.py"),
                str(parent_dir), "--response-out", str(child_response),
                "--artifact-dir", str(child_dir), "--parent-receipt", str(parent_receipt),
            ],
            capture_output=True, text=True, check=False,
        )
        if merged.returncode:
            raise AssertionError(merged.stderr[-2500:] + merged.stdout[-1500:])
        return parent_receipt, child_work, child_response, current_fingerprint

    def test_parent_projection_is_replayed_from_raw_bytes_and_child_is_byte_bound(self) -> None:
        parent_receipt, child_work, child_response, current_fingerprint = (
            self._make_deterministic_declaration_child()
        )
        child_receipt = child_work / "review" / "requirements" / "merge-receipt.json"
        verified = offline_review_receipt.validate_offline_parent_merge_receipt(
            parent_receipt_path=parent_receipt,
            child_receipt_path=child_receipt,
            child_response_path=child_response,
            child_work=child_work,
            current_code_fingerprint_sha256=current_fingerprint,
        )
        self.assertTrue(verified["child_verification"]["deterministic_projection_replay_verified"])
        self.assertFalse(verified["model_request_made"])
        self.assertFalse(verified["submission_ready"])
        child = json.loads(child_receipt.read_text(encoding="utf-8"))
        projection = child["declaration_source_text_projection"][0]
        self.assertEqual(projection["projected_response_serialization"], "canonical_json_utf8_v1")
        self.assertEqual(
            projection["projected_response_bytes_sha256"],
            projection["projected_response_sha256"],
        )

        original_response_bytes = child_response.read_bytes()
        original_receipt_bytes = child_receipt.read_bytes()

        def restore() -> tuple[dict, dict]:
            child_response.write_bytes(original_response_bytes)
            child_value = json.loads(original_receipt_bytes)
            child_receipt.write_bytes(original_receipt_bytes)
            return child_value, child

        # Even a tampered response with a freshly recomputed aggregate digest
        # must fail because it cannot be reproduced from the immutable packet.
        tampered_response = json.loads(original_response_bytes)
        tampered_response["requirements"][0]["properties"]["items"][0]["body_parts"][0] += "已篡改"
        _write(child_response, tampered_response)
        tampered_receipt = json.loads(original_receipt_bytes)
        tampered_receipt["aggregate_sha256"] = sha256_json(tampered_response)
        _write(child_receipt, tampered_receipt)
        with self.assertRaisesRegex(ValueError, "differ from deterministic parent replay"):
            offline_review_receipt.validate_offline_parent_merge_receipt(
                parent_receipt_path=parent_receipt, child_receipt_path=child_receipt,
                child_response_path=child_response, child_work=child_work,
                current_code_fingerprint_sha256=current_fingerprint,
            )
        restore()

        # Path, aggregate-summary, and code-identity substitution are rejected
        # before consuming the projection hash claim.
        mutations = [
            ("parent_merge_receipt_path", str(child_receipt)),
            ("parent_aggregate_sha256", "f" * 64),
            ("parent_request_code_fingerprint_sha256", "e" * 64),
            ("derivation_code_fingerprint_sha256", "c" * 64),
        ]
        for field, value in mutations:
            with self.subTest(field=field):
                changed, _ = restore()
                changed["derivation"][field] = value
                _write(child_receipt, changed)
                with self.assertRaisesRegex(ValueError, "not bound to the verified parent run"):
                    offline_review_receipt.validate_offline_parent_merge_receipt(
                        parent_receipt_path=parent_receipt, child_receipt_path=child_receipt,
                        child_response_path=child_response, child_work=child_work,
                        current_code_fingerprint_sha256=current_fingerprint,
                    )
        restore()

        missing_audit, _ = restore()
        missing_audit["declaration_source_text_projection"] = []
        _write(child_receipt, missing_audit)
        with self.assertRaisesRegex(ValueError, "no fixed declaration projection audit"):
            offline_review_receipt.validate_offline_parent_merge_receipt(
                parent_receipt_path=parent_receipt, child_receipt_path=child_receipt,
                child_response_path=child_response, child_work=child_work,
                current_code_fingerprint_sha256=current_fingerprint,
            )
        restore()

        altered_projection_hash, _ = restore()
        projection = altered_projection_hash["declaration_source_text_projection"][0]
        projection["projected_response_sha256"] = "0" * 64
        projection["projected_response_bytes_sha256"] = "0" * 64
        _write(child_receipt, altered_projection_hash)
        with self.assertRaisesRegex(ValueError, "differs from deterministic parent replay"):
            offline_review_receipt.validate_offline_parent_merge_receipt(
                parent_receipt_path=parent_receipt, child_receipt_path=child_receipt,
                child_response_path=child_response, child_work=child_work,
                current_code_fingerprint_sha256=current_fingerprint,
            )
        restore()

    def test_parent_lineage_rejects_non_review_draft_policy(self) -> None:
        args = SimpleNamespace(
            prepare_host_review=True, llm_response=None,
            host_agent_audit=None, merge_receipt=None,
            allow_offline_review=False, output_policy="submission",
            require_submission_ready=False, strict_release=False,
            offline_merge_receipt=None,
            offline_parent_merge_receipt=Path("parent/merge-receipt.json"),
            run_id="same-run", analysis_mode="llm_primary", compliance_mode="full",
            template_profile=None, render_report=None, allow_unresolved=False,
            preview_placeholders=False,
        )
        with self.assertRaises(SystemExit):
            validate_semantic_review_configuration(args, argparse.ArgumentParser())

    def test_real_wrapper_prepare_merge_and_continue_without_native_cli(self) -> None:
        self._real_wrapper_flow(defaults=False)

    def test_ordinary_conversation_defaults_work_without_native_cli(self) -> None:
        self._real_wrapper_flow(defaults=True)

    def _real_wrapper_flow(self, *, defaults: bool) -> None:
        root = Path(self.temporary.name) / "cli-smoke"
        root.mkdir()
        requirements = root / "requirements.docx"
        source = root / "source.docx"
        output = root / "review.docx"
        work = root / "work"
        for path, text in ((requirements, "本节为说明性标题。"),
                           (source, "测试论文正文。")):
            document = Document()
            document.add_paragraph(text)
            document.save(path)
        # Relative caller-owned paths outside the checkout and no installed
        # native CLI: the current agent supplies JSON, not a provider override.
        env = dict(os.environ, PATH="", THESIS_FORGE_HOST_RUNTIME="ordinary-agent")
        base = [sys.executable, str(ROOT / "scripts" / "thesis_format.py"),
                requirements.name, source.name, output.name,
                "--work-dir", work.name]
        preparation = base[:-2] if defaults else base + ["--prepare-agent-review"]
        prepared = subprocess.run(preparation,
                                  cwd=root, env=env, capture_output=True, text=True, check=False)
        self.assertEqual(prepared.returncode, 0, prepared.stderr[-1200:] + prepared.stdout[-1200:])
        if defaults:
            context = json.loads(prepared.stdout.splitlines()[0])
            self.assertEqual(context["model_selection"], "unchanged_by_skill")
            work = Path(context["work_dir"])
            self.assertEqual(work.parent, (root / "build").resolve())
            base[-1] = str(work)
        review_dir = work / "review" / "requirements"
        chunks = json.loads((review_dir / "llm-request-chunks.json").read_text(encoding="utf-8"))
        self.assertTrue(chunks)
        for chunk in chunks:
            response = {
                "contract_version": "3.0", "provenance": chunk["provenance"],
                "requirements": [],
                "clause_reviews": [
                    {"clause_id": clause["id"], "classification": "informational",
                     "reason": "The cited text is a section label, not a formatting instruction."}
                    for clause in chunk["clauses"]
                ],
                "unsupported_items": [], "reported_conflicts": [],
            }
            _write(review_dir / chunk["batch"]["response_filename"], response)
        merged_path = work / "review" / "host-agent-response.json"
        merged = subprocess.run(
            [sys.executable, str(ROOT / "scripts" / "merge_host_agent_review.py"),
             str(review_dir), "--response-out", str(merged_path)],
            cwd=root, env=env, capture_output=True, text=True, check=False,
        )
        self.assertEqual(merged.returncode, 0, merged.stderr[-1200:] + merged.stdout[-1200:])
        continued = subprocess.run(
            base + ["--llm-response", str(merged_path)]
            + ([] if defaults else ["--offline-review-draft"]),
            cwd=root, env=env, capture_output=True, text=True, check=False,
        )
        manifest = json.loads((work / "pipeline-manifest.json").read_text(encoding="utf-8"))
        self.assertNotIn("refusing to reuse a non-empty work directory", continued.stderr)
        self.assertEqual(manifest["offline_review_receipt"]["status"],
                         "offline_merged_without_independent_review")
        self.assertFalse(manifest.get("submission_ready", False))
        self.assertEqual(continued.returncode, 0,
                         continued.stderr[-1200:] + continued.stdout[-1200:])
        self.assertTrue(output.is_file())
        self.assertEqual(manifest["output_policy"], "review_draft")


if __name__ == "__main__":
    unittest.main()
