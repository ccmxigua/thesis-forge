#!/usr/bin/env python3
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
PY = sys.executable


def make_docx(path: Path, size: float, alignment=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    doc = Document()
    style = doc.styles["Heading 1"]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.alignment = alignment
    paragraph = doc.add_paragraph("第一章 绪论")
    paragraph.style = style
    doc.save(path)


def write_inputs(root: Path, *, official_size: float = 15, generated_size: float = 15,
                 required_size: float = 15) -> tuple[Path, Path, Path, Path, Path]:
    official = root / "official.docx"
    generated = root / "generated.docx"
    spec = root / "format-spec.json"
    official_map = root / "official-style-map.json"
    generated_map = root / "generated-style-map.json"
    make_docx(official, official_size)
    make_docx(generated, generated_size)
    spec.write_text(json.dumps({
        "schema_version": "1.0", "source_document": "rules.docx", "status": "rule_resolved",
        "roles": {"heading_1": {
            "font": {"size_pt": required_size, "bold": True},
            "paragraph": {"alignment": "center"},
        }},
        "requirements": [],
    }), encoding="utf-8")
    official_map.write_text(json.dumps({"mappings": {"heading_1": {"style_name": "Heading 1"}}}), encoding="utf-8")
    generated_map.write_text(json.dumps({"heading_1": {"style_name": "Heading 1"}}), encoding="utf-8")
    return generated, official, spec, official_map, generated_map


def run_audit(root: Path, inputs: tuple[Path, Path, Path, Path, Path], strict: bool = True):
    generated, official, spec, official_map, generated_map = inputs
    report = root / "format-comparison.json"
    markdown = root / "FORMAT-COMPARISON.md"
    command = [
        PY, "scripts/post_generation_format_audit.py", str(generated),
        "--official-template", str(official), "--format-spec", str(spec),
        "--official-style-map", str(official_map), "--generated-style-map", str(generated_map),
        "--out", str(report), "--markdown", str(markdown),
    ]
    if strict:
        command.append("--strict")
    result = subprocess.run(command, cwd=ROOT, text=True, capture_output=True)
    return result, json.loads(report.read_text(encoding="utf-8")), markdown


class PostGenerationFormatAuditTest(unittest.TestCase):
    def test_outline_level_is_compared_as_one_based_requirement(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            inputs = write_inputs(root)
            for path in inputs[:2]:
                doc = Document(path)
                ppr = doc.styles["Heading 1"].element.get_or_add_pPr()
                outline = ppr.find(qn("w:outlineLvl"))
                if outline is None:
                    outline = OxmlElement("w:outlineLvl")
                    ppr.append(outline)
                outline.set(qn("w:val"), "0")
                doc.save(path)
            spec = json.loads(inputs[2].read_text(encoding="utf-8"))
            spec["roles"]["heading_1"]["paragraph"]["outline_level"] = 1
            inputs[2].write_text(json.dumps(spec), encoding="utf-8")
            result, report, _ = run_audit(root, inputs)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            row = next(row for row in report["comparisons"]
                       if row["property"] == "paragraph.outline_level")
            self.assertEqual(row["generated_value"], 1)
            self.assertEqual(row["status"], "pass")

    def test_matching_requirement_template_and_output_pass(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            result, report, markdown = run_audit(root, write_inputs(root))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertEqual(report["status"], "passed")
            self.assertEqual(report["summary"]["blocking"], 0)
            self.assertGreaterEqual(report["summary"]["pass"], 3)
            self.assertTrue(markdown.exists())
            schema = json.loads((ROOT / "schema/format-comparison.schema.json").read_text(encoding="utf-8"))
            self.assertEqual(schema["properties"]["schema_version"]["const"], report["schema_version"])
            self.assertIn(report["status"], schema["properties"]["status"]["enum"])
            allowed = schema["properties"]["comparisons"]["items"]["properties"]["status"]["enum"]
            self.assertTrue(all(row["status"] in allowed for row in report["comparisons"]))

    def test_generated_value_mismatch_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            result, report, _ = run_audit(root, write_inputs(root, generated_size=14))
            self.assertEqual(result.returncode, 1)
            self.assertEqual(report["status"], "failed")
            failures = [row for row in report["comparisons"] if row["status"] == "fail"]
            self.assertTrue(any(row["property"] == "font.size_pt" for row in failures))

    def test_written_rule_and_official_template_conflict_is_explicit(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            result, report, _ = run_audit(root, write_inputs(root, official_size=16, generated_size=15, required_size=15))
            self.assertEqual(result.returncode, 0)
            self.assertEqual(report["status"], "passed")
            conflicts = [row for row in report["comparisons"] if row["status"] == "conflict"]
            self.assertTrue(any(row["property"] == "font.size_pt" for row in conflicts))
            row = next(row for row in conflicts if row["property"] == "font.size_pt")
            self.assertEqual(row["authoritative_source"], "written_requirement")
            self.assertEqual(row["generated_value"], 15.0)
            self.assertFalse(row["blocking"])


if __name__ == "__main__":
    unittest.main()
