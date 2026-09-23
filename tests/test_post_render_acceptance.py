from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
import fitz

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import batch_rerun_ten_schools as batch  # noqa: E402
import post_render_acceptance as post_render  # noqa: E402
from thesis_format_pipeline import runtime_code_fingerprint  # noqa: E402


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_valid_pdf(path: Path, text: str = "post-word") -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(str(path))
    document.close()


class PostRenderAcceptanceTests(unittest.TestCase):
    def test_evidence_loader_rejects_duplicate_json_keys(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "evidence.json"
            path.write_text('{"valid":true,"valid":false}', encoding="utf-8")
            self.assertIsNone(post_render.read_object(path))

    def _pre_render_fixture(self, root: Path) -> tuple[Path, Path, str]:
        source = root / "generated.docx"
        Document().save(source)
        digest = _sha256(source)
        validation = root / "validation-report.json"
        validation.write_text(json.dumps({
            "valid": True,
            "format_ready": True,
            "serialized_docx_valid": True,
            "output_docx": str(source),
            "property_receipt_audit": {
                "valid": True,
                "receipts": [{"serialized_docx_sha256": digest}],
            },
        }), encoding="utf-8")
        return source, validation, digest

    def _batch_fixture(self, root: Path) -> dict:
        case_dir = root / "case"
        work = case_dir / "work"
        requirements = work / "requirements"
        apply_dir = work / "application"
        requirements.mkdir(parents=True)
        apply_dir.mkdir(parents=True)
        output = case_dir / "generated.docx"
        Document().save(output)
        digest = _sha256(output)
        format_spec = requirements / "format-spec.json"
        format_spec.write_text("{}\n", encoding="utf-8")
        (requirements / "schema-validation.json").write_text(
            json.dumps({"valid": True, "errors": []}), encoding="utf-8"
        )
        capability = work / "capability-preflight.json"
        capability.write_text(json.dumps({"status": "passed", "findings": []}), encoding="utf-8")
        validation = apply_dir / "validation-report.json"
        validation.write_text(json.dumps({
            "valid": True,
            "format_ready": True,
            "serialized_docx_valid": True,
            "output_docx": str(output),
            "property_receipt_audit": {
                "valid": True,
                "receipts": [{"serialized_docx_sha256": digest}],
            },
        }), encoding="utf-8")
        comparison = apply_dir / "format-comparison.json"
        comparison.write_text(json.dumps({"status": "passed"}), encoding="utf-8")
        manifest = work / "pipeline-manifest.json"
        manifest.write_text(json.dumps({
            "status": "completed",
            "compliance_mode": "full",
            "case_id": "fixture-case",
            "inputs": {"baseline_authority": "fallback_input_not_official"},
            "requirements_extraction": {"run_id": "fresh-run"},
            "code_fingerprint": runtime_code_fingerprint(),
            "output": str(output),
            "format_spec": str(format_spec),
            "capability_preflight": str(capability),
            "capability_preflight_status": "passed",
            "validation_report": str(validation),
            "format_comparison": str(comparison),
        }), encoding="utf-8")
        return {
            "returncode": 0,
            "fresh_run": {
                "run_id": "fresh-run",
                "cache_reused": False,
                "pipeline_manifest": str(manifest),
            },
        }

    def test_pre_render_receipts_are_bound_to_the_pre_render_docx(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            source, validation, digest = self._pre_render_fixture(Path(td))
            checks, blockers = post_render._pre_render_checks(source, validation)
            self.assertEqual(blockers, [])
            self.assertEqual(checks["pre_render_artifact"]["sha256"], digest)

            validation.write_text(validation.read_text(encoding="utf-8").replace(digest, "0" * 64), encoding="utf-8")
            _checks, blockers = post_render._pre_render_checks(source, validation)
            self.assertIn("pre_render_property_receipt_artifact_hash_mismatch", blockers)

    def test_batch_acceptance_allows_distinct_post_word_hashes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            case_result = self._batch_fixture(root)
            manifest_path = Path(case_result["fresh_run"]["pipeline_manifest"])
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            pre_render = Path(manifest["output"])
            format_spec = Path(manifest["format_spec"])
            final = pre_render.with_name("final-word.docx")
            final_doc = Document()
            final_doc.add_paragraph("post-word")
            final_doc.save(final)
            final_digest = _sha256(final)
            acceptance = final.parent / "post-render-acceptance.json"
            comparison = final.parent / "final-format-comparison.json"
            pdf = final.parent / "final.pdf"
            _write_valid_pdf(pdf)
            render_report = final.parent / "word-render-report.json"
            render_report.write_text(json.dumps({
                "source_docx": {"path": str(final), "sha256": final_digest},
                "rendered_pdf": {"path": str(pdf), "sha256": _sha256(pdf)},
                "case_id": "fixture-case",
                "run_id": "fresh-run",
            }), encoding="utf-8")
            submission_audit = final.parent / "final-submission-audit.json"
            submission_audit.write_text(json.dumps({
                "submission_ready": True,
                "artifact": str(final),
                "render_validation": {
                    "rendered_verified": True,
                    "evidence": {
                        "source_docx_sha256": final_digest,
                        "rendered_pdf": {"sha256": _sha256(pdf)},
                    },
                },
            }), encoding="utf-8")
            markdown = final.parent / "FINAL-FORMAT-COMPARISON.md"
            markdown.write_text("passed\n", encoding="utf-8")
            visual_audit = final.parent / "pdf-visual-audit.json"
            visual_audit.write_text(json.dumps({
                "status": "passed",
                "pdf": {"path": str(pdf), "sha256": _sha256(pdf)},
            }), encoding="utf-8")
            comparison.write_text(json.dumps({
                "status": "passed",
                "inputs": {
                    "generated_docx": str(final),
                    "generated_docx_sha256": final_digest,
                    "format_spec": str(format_spec),
                    "official_template": None,
                    "official_template_sha256": None,
                },
            }), encoding="utf-8")
            acceptance.write_text(json.dumps({
                "status": "accepted",
                "blockers": [],
                "case_id": "fixture-case",
                "run_id": "fresh-run",
                "pre_render_docx_sha256": _sha256(pre_render),
                "post_render_docx": str(final),
                "post_render_docx_sha256": final_digest,
                "rendered_pdf": str(pdf),
                "rendered_pdf_sha256": _sha256(pdf),
                "render_report": str(render_report),
                "visual_audit": str(visual_audit),
                "submission_audit": str(submission_audit),
                "format_comparison": str(comparison),
            }), encoding="utf-8")
            manifest.update({
                "pre_render_output": str(pre_render),
                "post_render_status": "accepted",
                "post_render_acceptance": str(acceptance),
                "post_word_render": {
                    "pre_render_docx": str(pre_render),
                    "final_docx": str(final),
                    "pdf": str(pdf),
                    "render_report": str(render_report),
                    "visual_audit": str(visual_audit),
                    "submission_audit": str(submission_audit),
                    "format_comparison": str(comparison),
                    "format_comparison_markdown": str(markdown),
                    "requirements_only": True,
                },
                "output": str(final),
                "format_comparison": str(comparison),
            })
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            accepted = batch.case_acceptance(case_result, root=root)
            self.assertTrue(accepted["accepted"], accepted)
            self.assertNotEqual(_sha256(pre_render), final_digest)

    def test_batch_acceptance_rejects_final_hash_mismatch(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            case_result = self._batch_fixture(root)
            manifest_path = Path(case_result["fresh_run"]["pipeline_manifest"])
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            pre_render = Path(manifest["output"])
            final = pre_render.with_name("final-word.docx")
            final_doc = Document()
            final_doc.add_paragraph("post-word")
            final_doc.save(final)
            acceptance = final.parent / "post-render-acceptance.json"
            comparison = final.parent / "final-format-comparison.json"
            pdf = final.parent / "final.pdf"
            _write_valid_pdf(pdf)
            render_report = final.parent / "word-render-report.json"
            render_report.write_text(json.dumps({
                "source_docx": {"path": str(final), "sha256": _sha256(final)},
                "rendered_pdf": {"path": str(pdf), "sha256": _sha256(pdf)},
            }), encoding="utf-8")
            submission_audit = final.parent / "final-submission-audit.json"
            submission_audit.write_text(json.dumps({
                "submission_ready": True,
                "render_validation": {"rendered_verified": True},
            }), encoding="utf-8")
            markdown = final.parent / "FINAL-FORMAT-COMPARISON.md"
            markdown.write_text("passed\n", encoding="utf-8")
            visual_audit = final.parent / "pdf-visual-audit.json"
            visual_audit.write_text(json.dumps({
                "status": "passed",
                "pdf": {"path": str(pdf), "sha256": _sha256(pdf)},
            }), encoding="utf-8")
            comparison.write_text(json.dumps({
                "status": "passed",
                "inputs": {"generated_docx_sha256": _sha256(final)},
            }), encoding="utf-8")
            acceptance.write_text(json.dumps({
                "status": "accepted",
                "blockers": [],
                "pre_render_docx_sha256": _sha256(pre_render),
                "post_render_docx": str(final),
                "post_render_docx_sha256": "0" * 64,
                "rendered_pdf": str(pdf),
                "rendered_pdf_sha256": _sha256(pdf),
                "render_report": str(render_report),
                "visual_audit": str(visual_audit),
                "submission_audit": str(submission_audit),
                "format_comparison": str(comparison),
            }), encoding="utf-8")
            manifest.update({
                "pre_render_output": str(pre_render),
                "post_render_status": "accepted",
                "post_render_acceptance": str(acceptance),
                "post_word_render": {
                    "pre_render_docx": str(pre_render),
                    "final_docx": str(final),
                    "pdf": str(pdf),
                    "render_report": str(render_report),
                    "visual_audit": str(visual_audit),
                    "submission_audit": str(submission_audit),
                    "format_comparison": str(comparison),
                    "format_comparison_markdown": str(markdown),
                },
                "output": str(final),
                "format_comparison": str(comparison),
            })
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            accepted = batch.case_acceptance(case_result, root=root)
            self.assertFalse(accepted["accepted"])
            self.assertIn("post_render_acceptance_artifact_hash_mismatch", accepted["blockers"])


if __name__ == "__main__":
    unittest.main()
