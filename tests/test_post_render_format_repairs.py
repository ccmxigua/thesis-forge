from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from apply_format_spec import (  # noqa: E402
    apply_drawing_line_box_safety,
    apply_table_rules,
    audit_caption_separators,
    audit_drawing_line_boxes,
    audit_table_rules,
    normalize_caption_separators,
)


def png_bytes(size: tuple[int, int] = (128, 128)) -> bytes:
    image = Image.new("RGB", size, (25, 70, 110))
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def table_spec() -> dict:
    return {
        "style": "three_line", "scope": "captioned_tables",
        "caption_position": "above",
        "top_border_pt": 1.5, "header_border_pt": .5, "bottom_border_pt": 1.5,
        "remove_vertical_borders": True,
        "border_widths_pt": {
            "top": 1.5, "header": .5, "bottom": 1.5,
            "left": 0, "right": 0, "inside_h": 0, "inside_v": 0,
        },
    }


def solid_border(cell, side: str, eighth_points: int = 8) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    node = OxmlElement(f"w:{side}")
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(eighth_points))
    borders.append(node)


class DrawingLineBoxRepairTests(unittest.TestCase):
    def test_exact_shared_style_is_overridden_locally_and_nested_drawings_are_safe(self) -> None:
        doc = Document()
        normal = doc.styles["Normal"]
        normal.paragraph_format.line_spacing = Pt(20)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        body_para = doc.add_paragraph()
        body_para.add_run().add_picture(io.BytesIO(png_bytes()), width=Inches(1))
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].add_run().add_picture(io.BytesIO(png_bytes((80, 80))), width=Inches(.5))

        before = audit_drawing_line_boxes(doc)
        self.assertEqual(len(before), 2)
        self.assertTrue(all(row["failure_type"] == "drawing_line_box_clips_object" for row in before))

        changes = apply_drawing_line_box_safety(doc)
        self.assertEqual(len(changes), 2)
        self.assertTrue(all(row["changed"] for row in changes))
        self.assertEqual(normal.paragraph_format.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
        self.assertEqual(audit_drawing_line_boxes(doc), [])

    def test_single_spacing_with_an_inline_object_is_not_assumed_safe(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.add_run().add_picture(io.BytesIO(png_bytes()), width=Inches(1))
        self.assertTrue(audit_drawing_line_boxes(doc))
        apply_drawing_line_box_safety(doc)
        self.assertEqual(audit_drawing_line_boxes(doc), [])


class TableBorderRepairTests(unittest.TestCase):
    def test_captioned_scope_is_repaired_and_direct_cell_overrides_are_audited(self) -> None:
        doc = Document()
        caption_style = doc.styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption_style.base_style = doc.styles["Normal"]
        mappings = {"table_caption": {"style_name": "Table Caption"}}
        selected_tables = []
        for number in (1, 2):
            paragraph = doc.add_paragraph(f"表2-{number} caption {number}")
            paragraph.style = caption_style
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "header"
            table.cell(1, 0).text = "value"
            solid_border(table.cell(1, 1), "left", 4)
            selected_tables.append(table)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "uncaptioned cover table"

        changes = apply_table_rules(doc, table_spec(), mappings)
        self.assertEqual(changes["selected_table_count"], 2)
        self.assertEqual(changes["scope_findings"], [])
        self.assertEqual(changes["scope_exclusions"], [{
            "table_index": 3,
            "reason": "no mapped caption on the source-declared side",
        }])
        self.assertEqual(audit_table_rules(doc, table_spec(), mappings=mappings), [])

        solid_border(selected_tables[1].cell(1, 1), "bottom", 8)
        findings = audit_table_rules(doc, table_spec(), mappings=mappings)
        self.assertTrue(any(item.get("failure_type") == "three_line_table_cell_override"
                            for item in findings), findings)

    def test_empty_captioned_scope_fails_closed(self) -> None:
        doc = Document()
        doc.add_table(rows=1, cols=1)
        mappings = {"table_caption": {"style_name": "Table Caption"}}
        changes = apply_table_rules(doc, table_spec(), mappings)
        findings = audit_table_rules(doc, table_spec(), mappings=mappings)
        self.assertEqual(changes["selected_table_count"], 0)
        self.assertTrue(any(item.get("failure_type") == "table_scope_unresolved"
                            for item in findings), findings)


class CaptionSeparatorRepairTests(unittest.TestCase):
    def _doc_with_caption(self, *, split: bool) -> Document:
        doc = Document()
        style = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        paragraph = doc.add_paragraph()
        paragraph.style = style
        if split:
            paragraph.add_run("图2.1  ")
            paragraph.add_run("  示例图片")
        else:
            paragraph.add_run("图2.1    示例图片")
        return doc

    def test_single_run_separator_is_normalized_without_changing_caption_words(self) -> None:
        doc = self._doc_with_caption(split=False)
        original_words = "".join(doc.paragraphs[0].text.split())
        roles = {"figure_caption": {"separator": " "}}
        mappings = {"figure_caption": {"style_name": "Figure Caption"}}
        repairs = normalize_caption_separators(doc, roles, mappings)
        self.assertEqual(repairs[0]["status"], "normalized")
        self.assertEqual(doc.paragraphs[0].text, "图2.1 示例图片")
        self.assertEqual("".join(doc.paragraphs[0].text.split()), original_words)
        self.assertEqual(audit_caption_separators(doc, roles, mappings), [])

    def test_separator_crossing_runs_is_left_unchanged_and_reported(self) -> None:
        doc = self._doc_with_caption(split=True)
        before = doc.paragraphs[0].text
        roles = {"figure_caption": {"separator": " "}}
        mappings = {"figure_caption": {"style_name": "Figure Caption"}}
        repairs = normalize_caption_separators(doc, roles, mappings)
        self.assertEqual(repairs[0]["status"], "not_materialized")
        self.assertEqual(doc.paragraphs[0].text, before)
        findings = audit_caption_separators(doc, roles, mappings)
        self.assertEqual(findings[0]["failure_type"], "caption_separator_not_materialized")


if __name__ == "__main__":
    unittest.main()
