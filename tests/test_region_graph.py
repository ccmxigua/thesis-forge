from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from region_graph import CANONICAL_OOXML_HASH_POLICY, compile_region_graph  # noqa: E402
from assembly_executor import execute_assembly_plan  # noqa: E402
from format_spec_validation import load_and_validate  # noqa: E402
from source_role_extractor import extract_source_roles  # noqa: E402


def _template(path: Path, *, duplicate_anchor: bool = False) -> None:
    document = Document()
    document.add_paragraph("Protected declaration", style="Title")
    document.add_paragraph("Body placeholder")
    if duplicate_anchor:
        document.add_paragraph("Body placeholder")
    document.add_paragraph("Appendix placeholder")
    document.save(path)


def _range_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("Protected declaration", style="Title")
    document.add_paragraph("BODY START")
    document.add_paragraph("Old body paragraph one")
    document.add_paragraph("Old body paragraph two")
    document.add_paragraph("BODY END")
    document.add_paragraph("Protected closing")
    document.save(path)


def _multi_role_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("BODY")
    document.add_paragraph("old body")
    document.add_paragraph("REFERENCES")
    document.add_paragraph("old references")
    document.add_paragraph("ACKNOWLEDGMENTS")
    document.add_paragraph("old acknowledgments")
    document.add_paragraph("OUTPUTS")
    document.add_paragraph("old outputs")
    document.add_paragraph("END")
    document.save(path)


def _multi_role_profile() -> dict:
    nodes = []
    names = ["body", "references", "acknowledgments", "academic_outputs"]
    boundaries = ["BODY", "REFERENCES", "ACKNOWLEDGMENTS", "OUTPUTS", "END"]
    for role, start, end in zip(names, boundaries, boundaries[1:]):
        nodes.append({
            "id": role, "kind": "dynamic_content", "content_role": role,
            "start_selector": {"kind": "paragraph", "text": start, "match": "exact"},
            "end_selector": {"kind": "paragraph", "text": end, "match": "exact"},
            "range_policy": "replace_between",
        })
    return {
        "profile_id": "multi-role-fixture",
        "regions": {
            "graph_id": "multi-role-fixture-regions",
            "nodes": nodes,
            "edges": [
                {"kind": "order", "from": left, "to": right}
                for left, right in zip(names, names[1:])
            ],
        },
    }


def _optional_appendix_template(path: Path) -> None:
    document = Document()
    for heading, content in (
        ("BODY", "old body"),
        ("ACKNOWLEDGMENTS", "old acknowledgments"),
        ("REFERENCES", "old references"),
        ("APPENDIX", "old appendix placeholder"),
        ("OUTPUTS", "old outputs"),
        ("END", "protected end"),
    ):
        document.add_paragraph(heading)
        document.add_paragraph(content)
    document.save(path)


def _optional_appendix_profile() -> dict:
    names = ["body", "acknowledgments", "references", "appendices", "academic_outputs"]
    boundaries = ["BODY", "ACKNOWLEDGMENTS", "REFERENCES", "APPENDIX", "OUTPUTS", "END"]
    nodes = []
    for role, start, end in zip(names, boundaries, boundaries[1:]):
        node = {
            "id": role,
            "kind": "optional" if role == "appendices" else "dynamic_content",
            "content_role": role,
            "start_selector": {"kind": "paragraph", "text": start, "match": "exact"},
            "end_selector": {"kind": "paragraph", "text": end, "match": "exact"},
            "range_policy": "replace_between",
        }
        if role == "appendices":
            node["condition"] = "source_role.appendices.present"
        nodes.append(node)
    return {
        "profile_id": "optional-appendix-fixture",
        "regions": {
            "graph_id": "optional-appendix-fixture-regions",
            "source_section_policy": "discard",
            "nodes": nodes,
            "edges": [
                {"kind": "order", "from": left, "to": right}
                for left, right in zip(names, names[1:])
            ],
        },
    }


def _range_profile() -> dict:
    return {
        "profile_id": "range-fixture",
        "regions": {
            "graph_id": "range-fixture-regions",
            "nodes": [
                {"id": "declaration", "kind": "fixed_protected",
                 "selector": {"kind": "paragraph", "text": "Protected declaration", "match": "exact"}},
                {"id": "body", "kind": "dynamic_content", "content_role": "body",
                 "start_selector": {"kind": "paragraph", "text": "BODY START", "match": "exact"},
                 "end_selector": {"kind": "paragraph", "text": "BODY END", "match": "exact"},
                 "range_policy": "replace_between"},
                {"id": "closing", "kind": "fixed_protected",
                 "selector": {"kind": "paragraph", "text": "Protected closing", "match": "exact"}},
            ],
            "edges": [
                {"kind": "order", "from": "declaration", "to": "body"},
                {"kind": "order", "from": "body", "to": "closing"},
            ],
        },
    }


def _profile() -> dict:
    return {
        "profile_id": "neutral-fixture",
        "regions": {
            "graph_id": "neutral-thesis-regions",
            "nodes": [
                {"id": "declaration", "kind": "fixed_protected",
                 "selector": {"kind": "paragraph", "text": "Protected declaration", "match": "exact"}},
                {"id": "body", "kind": "dynamic_content", "content_role": "body",
                 "selector": {"kind": "paragraph", "text": "Body placeholder", "match": "exact"}},
                {"id": "references", "kind": "generated", "content_role": "references"},
                {"id": "appendix", "kind": "optional", "condition": "metadata.has_appendix",
                 "selector": {"kind": "paragraph", "text": "Appendix placeholder", "match": "exact"}},
            ],
            "edges": [
                {"kind": "order", "from": "declaration", "to": "body"},
                {"kind": "boundary", "from": "body", "to": "references",
                 "policy": {"page_break": True}},
                {"kind": "section_policy", "from": "references", "to": "appendix",
                 "policy": {"section_break": "next_page", "inherit_headers": False}},
            ],
        },
    }


class RegionGraphTests(unittest.TestCase):
    def test_compiles_dag_to_assembly_plan_and_fixed_hash_assertion(self):
        with tempfile.TemporaryDirectory() as raw:
            template = Path(raw) / "neutral.docx"
            _template(template)
            result = compile_region_graph(_profile(), template_path=template)

        self.assertEqual(result["status"], "compiled")
        self.assertEqual(result["findings"], [])
        plan = result["assembly_plan"]
        self.assertEqual([item["node_id"] for item in plan["operations"]],
                         ["declaration", "body", "references", "appendix"])
        self.assertEqual([item["action"] for item in plan["operations"]],
                         ["preserve", "replace_content", "generate_content", "conditionally_assemble"])
        self.assertEqual(plan["operations"][2]["prerequisites"][0]["relation"], "boundary")
        self.assertEqual(plan["operations"][3]["prerequisites"][0]["relation"], "section_policy")
        assertion = plan["protection_assertions"][0]
        self.assertEqual(assertion["node_id"], "declaration")
        self.assertEqual(assertion["hash_policy"], CANONICAL_OOXML_HASH_POLICY)
        self.assertRegex(assertion["expected_sha256"], r"^[0-9a-f]{64}$")
        self.assertEqual(assertion["locator"]["part"], "word/document.xml")
        self.assertTrue(any(item["type"] == "boundary" for item in plan["structural_postconditions"]))
        self.assertTrue(any(item["type"] == "section_policy" for item in plan["structural_postconditions"]))

    def test_compiled_plan_validates_against_machine_schema(self):
        with tempfile.TemporaryDirectory() as raw:
            template = Path(raw) / "neutral.docx"
            _template(template)
            result = compile_region_graph(_profile(), template_path=template)
        self.assertEqual(load_and_validate(result, ROOT / "schema" / "assembly-plan.schema.json"), [])

    def test_fixed_range_preserves_official_layout_without_redrawing(self):
        with tempfile.TemporaryDirectory() as raw:
            template = Path(raw) / "neutral.docx"
            _range_template(template)
            profile = {
                "profile_id": "neutral-range-protection",
                "regions": {
                    "graph_id": "neutral-range-protection",
                    "nodes": [{
                        "id": "official_cover",
                        "kind": "fixed_protected",
                        "start_selector": {"kind": "paragraph", "text": "BODY START", "match": "exact"},
                        "end_selector": {"kind": "paragraph", "text": "BODY END", "match": "exact"},
                        "range_policy": "preserve_between",
                    }],
                    "edges": [],
                },
            }
            compiled = compile_region_graph(profile, template_path=template)
            self.assertEqual(compiled["status"], "compiled")
            self.assertEqual(load_and_validate(compiled, ROOT / "schema" / "assembly-plan.schema.json"), [])
            operation = compiled["assembly_plan"]["operations"][0]
            self.assertEqual(operation["action"], "preserve")
            self.assertEqual(operation["range_locator"]["policy"], "preserve_between")
            assertion = compiled["assembly_plan"]["protection_assertions"][0]
            self.assertIsNone(assertion["locator"])
            self.assertEqual(assertion["range_locator"]["policy"], "preserve_between")

    def test_fixed_region_metadata_binding_replaces_text_without_redrawing_template(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "neutral.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _template(template)
            Document().save(source)
            profile = _profile()
            profile["regions"]["nodes"] = [profile["regions"]["nodes"][0]]
            profile["regions"]["edges"] = []
            profile["regions"]["nodes"][0]["metadata_bindings"] = [{
                "field": "cover_metadata.title_zh",
                "selector": {"kind": "paragraph", "text": "Protected declaration", "match": "exact"},
                "required": True,
                "missing_policy": "error",
            }]
            compiled = compile_region_graph(profile, template_path=template)
            self.assertEqual(load_and_validate(compiled, ROOT / "schema" / "assembly-plan.schema.json"), [])
            result = execute_assembly_plan(
                compiled, template, source, output,
                metadata={"cover_metadata": {"title_zh": "正式论文标题"}},
            )
            text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
            self.assertIn("正式论文标题", text)
            self.assertEqual(result["metadata_bindings"][0]["status"], "replaced")
            self.assertEqual(result["protected_regions_verified"], 1)

    def test_executor_replaces_dynamic_anchor_and_preserves_protected_hash(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "neutral.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _template(template)
            incoming = Document()
            incoming.add_heading("第1章 装配后的正文", level=1)
            incoming.add_paragraph("来自显式绑定 source DOCX 的内容。")
            incoming.save(source)
            profile = _profile()
            # Generated and optional adapters are tested separately; this plan
            # exercises the minimum fully executable preserve+replace contract.
            profile["regions"]["nodes"] = profile["regions"]["nodes"][:2]
            profile["regions"]["edges"] = profile["regions"]["edges"][:1]
            compiled = compile_region_graph(profile, template_path=template)
            result = execute_assembly_plan(compiled, template, source, output)
            text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
            self.assertEqual(result["status"], "assembled")
            self.assertEqual(result["protected_regions_verified"], 1)
            self.assertEqual(result["structural_postconditions_verified"], 1)
            self.assertIn("Protected declaration", text)
            self.assertIn("第1章 装配后的正文", text)
            self.assertNotIn("Body placeholder", text)

    def test_executor_imports_image_relationship_and_remaps_drawing_ids(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "range-template.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            image = temp / "pixel.png"
            _range_template(template)
            # One-pixel PNG fixture; no image library dependency is required.
            image.write_bytes(__import__("base64").b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZlVQAAAAASUVORK5CYII="
            ))
            incoming = Document()
            incoming.add_paragraph("Image follows")
            incoming.add_picture(str(image), width=Inches(0.2))
            incoming.save(source)

            compiled = compile_region_graph(_range_profile(), template_path=template)
            result = execute_assembly_plan(compiled, template, source, output)
            self.assertEqual(result["status"], "assembled")
            self.assertEqual(len(result["source_import"]["relationships_imported"]), 1)
            self.assertEqual(result["source_import"]["relationships_imported"][0]["kind"], "image")
            self.assertEqual(len(result["source_import"]["parts_imported"]), 1)
            with zipfile.ZipFile(output) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            root = etree.fromstring(members["word/document.xml"])
            relroot = etree.fromstring(members["word/_rels/document.xml.rels"])
            rels = {node.get("Id"): node.get("Target") for node in relroot}
            embeds = root.xpath("//a:blip/@r:embed", namespaces={
                "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            })
            self.assertEqual(len(embeds), 1)
            target = rels[embeds[0]]
            self.assertIn("word/" + target, members)
            drawing_ids = [int(value) for value in root.xpath(
                "//wp:docPr/@id", namespaces={
                    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                })]
            self.assertEqual(len(drawing_ids), len(set(drawing_ids)))

    def test_executor_imports_external_hyperlink_relationship(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "range-template.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _range_template(template)
            incoming = Document()
            paragraph = incoming.add_paragraph("Link: ")
            relationship_id = incoming.part.relate_to(
                "https://example.test/thesis",
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), relationship_id)
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "evidence"
            run.append(text); hyperlink.append(run); paragraph._p.append(hyperlink)
            incoming.save(source)

            compiled = compile_region_graph(_range_profile(), template_path=template)
            result = execute_assembly_plan(compiled, template, source, output)
            imported = result["source_import"]["relationships_imported"]
            self.assertEqual([item["kind"] for item in imported], ["hyperlink"])
            with zipfile.ZipFile(output) as archive:
                root = etree.fromstring(archive.read("word/document.xml"))
                relroot = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
            rid = root.xpath("//w:hyperlink/@r:id", namespaces={
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            })[0]
            relationship = next(node for node in relroot if node.get("Id") == rid)
            self.assertEqual(relationship.get("Target"), "https://example.test/thesis")
            self.assertEqual(relationship.get("TargetMode"), "External")

    def test_executor_imports_missing_style_dependency_closure(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "range-template.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _range_template(template)
            incoming = Document()
            base = incoming.styles.add_style("ImportedBase", WD_STYLE_TYPE.PARAGRAPH)
            child = incoming.styles.add_style("ImportedChild", WD_STYLE_TYPE.PARAGRAPH)
            child.base_style = base
            incoming.add_paragraph("Styled content", style="ImportedChild")
            incoming.save(source)

            compiled = compile_region_graph(_range_profile(), template_path=template)
            result = execute_assembly_plan(compiled, template, source, output)
            self.assertEqual(
                set(result["source_import"]["styles_imported"]),
                {"ImportedBase", "ImportedChild"},
            )
            with zipfile.ZipFile(output) as archive:
                document_root = etree.fromstring(archive.read("word/document.xml"))
                styles_root = etree.fromstring(archive.read("word/styles.xml"))
            used = set(document_root.xpath("//w:pStyle/@w:val", namespaces={
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            }))
            defined = set(styles_root.xpath("//w:style/@w:styleId", namespaces={
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            }))
            self.assertTrue({"ImportedBase", "ImportedChild"} <= defined)
            self.assertFalse(used - defined)

    def test_executor_rejects_numbering_in_imported_style_without_partial_output(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "range-template.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _range_template(template)
            incoming = Document()
            numbered = incoming.styles.add_style("NumberedImportedStyle", WD_STYLE_TYPE.PARAGRAPH)
            properties = numbered.element.get_or_add_pPr()
            num_properties = OxmlElement("w:numPr")
            num_id = OxmlElement("w:numId")
            num_id.set(qn("w:val"), "1")
            num_properties.append(num_id)
            properties.append(num_properties)
            incoming.add_paragraph("Indirect numbering", style="NumberedImportedStyle")
            incoming.save(source)

            compiled = compile_region_graph(_range_profile(), template_path=template)
            with self.assertRaisesRegex(ValueError, "numbering import adapter"):
                execute_assembly_plan(compiled, template, source, output)
            self.assertFalse(output.exists())

    def test_executor_normalizes_only_implicit_table_normal_dependency(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "range-template.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _range_template(template)
            incoming = Document()
            table_style = incoming.styles.add_style("ImportedTable", WD_STYLE_TYPE.TABLE)
            based_on = OxmlElement("w:basedOn")
            based_on.set(qn("w:val"), "TableNormal")
            table_style.element.insert(0, based_on)
            table = incoming.add_table(rows=1, cols=1)
            table.style = "ImportedTable"
            table.cell(0, 0).text = "Imported cell"
            incoming.save(source)
            # python-docx may materialize the built-in TableNormal definition
            # while saving. Remove it so the fixture matches the real source
            # package, whose ImportedTable-equivalent has an implicit dangling
            # basedOn reference to Word's built-in TableNormal.
            for package in (source, template):
                with zipfile.ZipFile(package) as archive:
                    members = {name: archive.read(name) for name in archive.namelist()}
                styles_root = etree.fromstring(members["word/styles.xml"])
                for node in styles_root.xpath(
                    "./w:style[@w:styleId='TableNormal']",
                    namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
                ):
                    styles_root.remove(node)
                members["word/styles.xml"] = etree.tostring(
                    styles_root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
                with zipfile.ZipFile(package, "w", zipfile.ZIP_DEFLATED) as archive:
                    for name, payload in members.items():
                        archive.writestr(name, payload)

            compiled = compile_region_graph(_range_profile(), template_path=template)
            result = execute_assembly_plan(compiled, template, source, output)
            self.assertEqual(result["source_import"]["style_dependencies_normalized"], [{
                "style_id": "ImportedTable",
                "dependency": "TableNormal",
                "action": "remove_implicit_builtin_reference",
            }])
            with zipfile.ZipFile(output) as archive:
                styles_root = etree.fromstring(archive.read("word/styles.xml"))
            self.assertEqual(styles_root.xpath(
                "./w:style[@w:styleId='ImportedTable']/w:basedOn/@w:val",
                namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
            ), [])

    def test_executor_requires_explicit_policy_to_discard_source_sections(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "range-template.docx"
            source = temp / "source.docx"
            rejected = temp / "rejected.docx"
            accepted = temp / "accepted.docx"
            _range_template(template)
            incoming = Document()
            incoming.add_paragraph("Front source section")
            incoming.add_section()
            incoming.add_paragraph("Body source section")
            incoming.save(source)
            compiled = compile_region_graph(_range_profile(), template_path=template)

            with self.assertRaisesRegex(ValueError, "may not contain section properties"):
                execute_assembly_plan(compiled, template, source, rejected)
            self.assertFalse(rejected.exists())
            result = execute_assembly_plan(
                compiled, template, source, accepted, source_section_policy="discard"
            )
            self.assertGreater(result["source_import"]["source_sections_discarded"], 0)
            with zipfile.ZipFile(accepted) as archive:
                root = etree.fromstring(archive.read("word/document.xml"))
            # Only the destination template's terminal section remains.
            self.assertEqual(len(root.xpath("//w:sectPr", namespaces={
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            })), 1)

    def test_executor_imports_direct_numbering_with_collision_safe_remap(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "range-template.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _range_template(template)
            incoming = Document()
            paragraph = incoming.add_paragraph("Numbered")
            properties = paragraph._p.get_or_add_pPr()
            num_properties = OxmlElement("w:numPr")
            num_id = OxmlElement("w:numId")
            num_id.set(qn("w:val"), "1")
            num_properties.append(num_id)
            properties.append(num_properties)
            incoming.save(source)
            compiled = compile_region_graph(_range_profile(), template_path=template)
            result = execute_assembly_plan(compiled, template, source, output)
            imported = result["source_import"]["numbering_imported"]
            self.assertEqual(len(imported), 1)
            self.assertEqual(imported[0]["source_num_id"], "1")
            with zipfile.ZipFile(output) as archive:
                document_root = etree.fromstring(archive.read("word/document.xml"))
                numbering_root = etree.fromstring(archive.read("word/numbering.xml"))
            namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            used = document_root.xpath("//w:numPr/w:numId/@w:val", namespaces=namespace)
            defined = numbering_root.xpath("./w:num/@w:numId", namespaces=namespace)
            self.assertIn(imported[0]["destination_num_id"], used)
            self.assertTrue(set(used) <= set(defined))

    def test_executor_replaces_exclusive_range_and_preserves_boundaries(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "range-template.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _range_template(template)
            incoming = Document()
            incoming.add_heading("第1章 新正文", level=1)
            incoming.add_paragraph("范围装配后的正文。")
            incoming.save(source)

            compiled = compile_region_graph(_range_profile(), template_path=template)
            operation = next(item for item in compiled["assembly_plan"]["operations"]
                             if item["node_id"] == "body")
            self.assertIsNone(operation["locator"])
            self.assertEqual(operation["range_locator"]["policy"], "replace_between")

            result = execute_assembly_plan(compiled, template, source, output)
            paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
            self.assertEqual(result["status"], "assembled")
            self.assertEqual(result["protected_regions_verified"], 2)
            self.assertIn("BODY START", paragraphs)
            self.assertIn("BODY END", paragraphs)
            self.assertIn("第1章 新正文", paragraphs)
            self.assertIn("范围装配后的正文。", paragraphs)
            self.assertNotIn("Old body paragraph one", paragraphs)
            self.assertNotIn("Old body paragraph two", paragraphs)
            self.assertLess(paragraphs.index("BODY START"), paragraphs.index("第1章 新正文"))
            self.assertLess(paragraphs.index("第1章 新正文"), paragraphs.index("BODY END"))

    def test_executor_replaces_to_body_end_without_removing_template_sectpr(self):
        with tempfile.TemporaryDirectory() as raw:
            base = Path(raw)
            template = base / "template.docx"
            source = base / "source.docx"
            output = base / "output.docx"
            document = Document()
            document.add_paragraph("OUTPUTS")
            document.add_paragraph("old output")
            document.sections[0].top_margin = Inches(1.23)
            document.save(template)
            incoming = Document()
            incoming.add_paragraph("new output")
            incoming.sections[0].top_margin = Inches(0.55)
            incoming.save(source)
            profile = {
                "profile_id": "body-end-fixture",
                "regions": {
                    "graph_id": "body-end-fixture-regions",
                    "source_section_policy": "discard",
                    "nodes": [{
                        "id": "academic_outputs", "kind": "dynamic_content",
                        "content_role": "academic_outputs",
                        "start_selector": {"kind": "paragraph", "text": "OUTPUTS", "match": "exact"},
                        "end_boundary": "body_end", "range_policy": "replace_between",
                    }],
                    "edges": [],
                },
            }
            compiled = compile_region_graph(profile, template_path=template)
            self.assertEqual(compiled["status"], "compiled", compiled["findings"])
            self.assertEqual(
                compiled["assembly_plan"]["operations"][0]["range_locator"]["end"],
                {"boundary": "body_end"},
            )
            report = execute_assembly_plan(
                compiled, template, source, output, source_section_policy="discard"
            )
            self.assertEqual(report["status"], "assembled")
            result = Document(output)
            self.assertEqual([paragraph.text for paragraph in result.paragraphs], ["OUTPUTS", "new output"])
            self.assertAlmostEqual(result.sections[0].top_margin.inches, 1.23, places=2)
            with zipfile.ZipFile(output) as archive:
                root = etree.fromstring(archive.read("word/document.xml"))
            self.assertEqual(len(root.xpath("/w:document/w:body/w:sectPr", namespaces={
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            })), 1)

    def test_executor_rejects_stale_range_boundary_without_partial_output(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "range-template.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _range_template(template)
            _template(source)
            compiled = compile_region_graph(_range_profile(), template_path=template)
            operation = next(item for item in compiled["assembly_plan"]["operations"]
                             if item["node_id"] == "body")
            operation["range_locator"]["end"]["normalized_text_sha256"] = "0" * 64
            with self.assertRaisesRegex(ValueError, "stale assembly locator identity"):
                execute_assembly_plan(compiled, template, source, output)
            self.assertFalse(output.exists())

    def test_compiler_rejects_reversed_range(self):
        with tempfile.TemporaryDirectory() as raw:
            template = Path(raw) / "range-template.docx"
            _range_template(template)
            profile = _range_profile()
            body = profile["regions"]["nodes"][1]
            body["start_selector"], body["end_selector"] = (
                body["end_selector"], body["start_selector"]
            )
            result = compile_region_graph(profile, template_path=template)
        self.assertEqual(result["status"], "invalid")
        self.assertIn("region_graph.range_invalid", {item["code"] for item in result["findings"]})
        self.assertIsNone(result["assembly_plan"])

    def test_compiler_allows_adjacent_ranges_to_share_boundary(self):
        with tempfile.TemporaryDirectory() as raw:
            template = Path(raw) / "multi-role.docx"
            _multi_role_template(template)
            result = compile_region_graph(_multi_role_profile(), template_path=template)
        self.assertEqual(result["status"], "compiled")
        self.assertEqual(result["findings"], [])
        self.assertEqual(
            [item["content_role"] for item in result["assembly_plan"]["operations"]],
            ["body", "references", "acknowledgments", "academic_outputs"],
        )

    def test_executor_imports_multiple_source_role_ranges_without_duplicate_headings(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "multi-role.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _multi_role_template(template)
            incoming = Document()
            incoming.add_heading("第1章 引言", level=1)
            incoming.add_paragraph("new body")
            incoming.add_heading("3.5 参考文献", level=2)
            incoming.add_paragraph("new reference")
            incoming.add_heading("后 记", level=1)
            incoming.add_paragraph("new thanks")
            incoming.add_heading("在学期间发表的学术论文与研究成果", level=1)
            incoming.add_paragraph("new output")
            incoming.save(source)

            compiled = compile_region_graph(_multi_role_profile(), template_path=template)
            source_roles = extract_source_roles(
                source,
                required_roles={"body", "references", "acknowledgments", "academic_outputs"},
            )
            result = execute_assembly_plan(
                compiled, template, source, output, source_role_map=source_roles
            )
            paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]

        self.assertEqual(result["status"], "assembled")
        self.assertEqual(set(result["source_import"]), {
            "body", "references", "acknowledgments", "academic_outputs"
        })
        self.assertEqual(paragraphs, [
            "BODY", "new body", "REFERENCES", "new reference",
            "ACKNOWLEDGMENTS", "new thanks", "OUTPUTS", "new output", "END",
        ])

    def test_optional_appendix_range_is_removed_when_source_has_no_appendix(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "template.docx"
            source = temp / "source.docx"
            output = temp / "output.docx"
            _optional_appendix_template(template)
            incoming = Document()
            incoming.add_heading("第1章 引言", level=1)
            incoming.add_paragraph("new body")
            incoming.add_heading("致谢", level=1)
            incoming.add_paragraph("new thanks")
            incoming.add_heading("参考文献", level=1)
            incoming.add_paragraph("new reference")
            incoming.add_heading("攻读硕士学位期间的学术成果", level=1)
            incoming.add_paragraph("new output")
            incoming.save(source)

            compiled = compile_region_graph(_optional_appendix_profile(), template_path=template)
            self.assertEqual(compiled["status"], "compiled", compiled["findings"])
            source_roles = extract_source_roles(
                source,
                required_roles={"body", "references", "acknowledgments", "academic_outputs"},
            )
            report = execute_assembly_plan(
                compiled, template, source, output,
                source_section_policy="discard", source_role_map=source_roles,
            )
            paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]

        self.assertEqual(report["status"], "assembled")
        self.assertNotIn("APPENDIX", paragraphs)
        self.assertNotIn("old appendix placeholder", paragraphs)
        self.assertIn("OUTPUTS", paragraphs)
        self.assertIn("new output", paragraphs)

    def test_optional_appendix_range_imports_content_when_source_has_appendix(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "template.docx"
            source = temp / "source.docx"
            output = temp / "output.docx"
            _optional_appendix_template(template)
            incoming = Document()
            incoming.add_heading("第1章 引言", level=1)
            incoming.add_paragraph("new body")
            incoming.add_heading("致谢", level=1)
            incoming.add_paragraph("new thanks")
            incoming.add_heading("参考文献", level=1)
            incoming.add_paragraph("new reference")
            incoming.add_heading("附录A 补充材料", level=1)
            incoming.add_paragraph("new appendix")
            incoming.add_heading("攻读硕士学位期间的学术成果", level=1)
            incoming.add_paragraph("new output")
            incoming.save(source)

            compiled = compile_region_graph(_optional_appendix_profile(), template_path=template)
            source_roles = extract_source_roles(
                source,
                required_roles={"body", "references", "acknowledgments", "academic_outputs"},
            )
            report = execute_assembly_plan(
                compiled, template, source, output,
                source_section_policy="discard", source_role_map=source_roles,
            )
            paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]

        self.assertEqual(report["status"], "assembled")
        self.assertIn("APPENDIX", paragraphs)
        self.assertIn("new appendix", paragraphs)
        self.assertNotIn("old appendix placeholder", paragraphs)
        self.assertLess(paragraphs.index("APPENDIX"), paragraphs.index("OUTPUTS"))

    def test_compiler_rejects_overlapping_dynamic_ranges(self):
        with tempfile.TemporaryDirectory() as raw:
            template = Path(raw) / "range-template.docx"
            _range_template(template)
            profile = _range_profile()
            profile["regions"]["nodes"].insert(2, {
                "id": "overlap",
                "kind": "dynamic_content",
                "content_role": "overlap",
                "start_selector": {"kind": "paragraph", "text": "Old body paragraph one", "match": "exact"},
                "end_selector": {"kind": "paragraph", "text": "Protected closing", "match": "exact"},
                "range_policy": "replace_between",
            })
            profile["regions"]["edges"] = []
            result = compile_region_graph(profile, template_path=template)
        self.assertEqual(result["status"], "invalid")
        self.assertIn("region_graph.range_overlap", {item["code"] for item in result["findings"]})
        self.assertIsNone(result["assembly_plan"])

    def test_compiler_rejects_range_containing_protected_region(self):
        with tempfile.TemporaryDirectory() as raw:
            template = Path(raw) / "range-template.docx"
            _range_template(template)
            profile = _range_profile()
            body = profile["regions"]["nodes"][1]
            body["end_selector"] = {
                "kind": "paragraph", "text": "Protected closing", "match": "exact"
            }
            result = compile_region_graph(profile, template_path=template)
        self.assertEqual(result["status"], "invalid")
        self.assertIn("region_graph.range_protected_overlap",
                      {item["code"] for item in result["findings"]})
        self.assertIsNone(result["assembly_plan"])

    def test_executor_blocks_unsupported_generator_without_partial_output(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "neutral.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _template(template); _template(source)
            profile = _profile()
            profile["regions"]["nodes"] = profile["regions"]["nodes"][:3]
            profile["regions"]["edges"] = [
                edge for edge in profile["regions"]["edges"][:2] if edge["kind"] == "order"
            ]
            compiled = compile_region_graph(profile, template_path=template)
            with self.assertRaisesRegex(ValueError, "no generator adapter"):
                execute_assembly_plan(compiled, template, source, output, metadata={"has_appendix": True})
            self.assertFalse(output.exists())

    def test_executor_blocks_declared_boundary_policy_until_adapter_exists(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "neutral.docx"; source = temp / "source.docx"; output = temp / "out.docx"
            _template(template); _template(source)
            profile = _profile()
            profile["regions"]["nodes"] = profile["regions"]["nodes"][:3]
            profile["regions"]["edges"] = profile["regions"]["edges"][:2]
            compiled = compile_region_graph(profile, template_path=template)
            with self.assertRaisesRegex(ValueError, "boundary/section-policy"):
                execute_assembly_plan(compiled, template, source, output)
            self.assertFalse(output.exists())

    def test_executor_rejects_multiple_dynamic_regions_until_role_extractors_exist(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            template = temp / "neutral.docx"
            source = temp / "source.docx"
            output = temp / "assembled.docx"
            _template(template); _template(source)
            profile = _profile()
            profile["regions"]["nodes"] = [
                profile["regions"]["nodes"][0],
                profile["regions"]["nodes"][1],
                {"id": "appendix", "kind": "dynamic_content", "content_role": "appendix",
                 "selector": {"kind": "paragraph", "text": "Appendix placeholder", "match": "exact"}},
            ]
            profile["regions"]["edges"] = [
                {"kind": "order", "from": "declaration", "to": "body"},
                {"kind": "order", "from": "body", "to": "appendix"},
            ]
            compiled = compile_region_graph(profile, template_path=template)
            with self.assertRaisesRegex(ValueError, "multiple replace_content"):
                execute_assembly_plan(compiled, template, source, output)
            self.assertFalse(output.exists())

    def test_selector_preflight_fails_on_zero_or_multiple_matches(self):
        for mode in ("missing", "ambiguous"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as raw:
                template = Path(raw) / "neutral.docx"
                _template(template, duplicate_anchor=(mode == "ambiguous"))
                profile = _profile()
                if mode == "missing":
                    profile["regions"]["nodes"][1]["selector"]["text"] = "Not present"
                result = compile_region_graph(profile, template_path=template)
                self.assertEqual(result["status"], "invalid")
                finding = next(item for item in result["findings"]
                               if item["code"] == "region_graph.selector_not_unique")
                self.assertIn({"kind": "node_id", "value": "body"}, finding["evidence"])
                self.assertIsNone(result["assembly_plan"])

    def test_selector_collision_is_rejected_even_when_each_selector_is_unique(self):
        with tempfile.TemporaryDirectory() as raw:
            template = Path(raw) / "neutral.docx"
            _template(template)
            profile = _profile()
            profile["regions"]["nodes"][3]["selector"] = copy.deepcopy(
                profile["regions"]["nodes"][1]["selector"]
            )
            result = compile_region_graph(profile, template_path=template)
        self.assertIn("region_graph.selector_collision", {item["code"] for item in result["findings"]})

    def test_cycle_duplicate_node_and_missing_endpoint_are_machine_findings(self):
        with tempfile.TemporaryDirectory() as raw:
            template = Path(raw) / "neutral.docx"
            _template(template)
            profile = _profile()
            profile["regions"]["nodes"].append(
                {"id": "body", "kind": "generated", "content_role": "duplicate"}
            )
            profile["regions"]["edges"].extend([
                {"kind": "order", "from": "declaration", "to": "body"},
                {"kind": "order", "from": "appendix", "to": "declaration"},
                {"kind": "order", "from": "missing", "to": "body"},
            ])
            result = compile_region_graph(profile, template_path=template)
        codes = {item["code"] for item in result["findings"]}
        self.assertTrue({"region_graph.node_duplicate", "region_graph.edge_duplicate",
                         "region_graph.edge_node_missing", "region_graph.cycle"} <= codes)
        self.assertIsNone(result["assembly_plan"])

    def test_generated_only_graph_does_not_require_template(self):
        profile = {
            "profile_id": "generated-only",
            "regions": {
                "graph_id": "generated-only",
                "nodes": [
                    {"id": "toc", "kind": "generated", "content_role": "toc"},
                    {"id": "body", "kind": "generated", "content_role": "body"},
                ],
                "edges": [{"kind": "order", "from": "toc", "to": "body"}],
            },
        }
        result = compile_region_graph(profile)
        self.assertEqual(result["status"], "compiled")
        self.assertEqual([item["node_id"] for item in result["assembly_plan"]["operations"]], ["toc", "body"])

    def test_schema_rejects_unknown_node_and_edge_kinds(self):
        profile = _profile()
        profile["regions"]["nodes"][0]["kind"] = "unknown_kind"
        profile["regions"]["edges"][0]["kind"] = "teleport"
        result = compile_region_graph(profile)
        self.assertEqual(result["status"], "invalid")
        self.assertEqual({item["code"] for item in result["findings"]}, {"region_graph.schema_invalid"})


if __name__ == "__main__":
    unittest.main()
