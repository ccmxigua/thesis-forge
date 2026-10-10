from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from rendered_format_audit import audit  # noqa: E402


def make_image() -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (40, 40), (11, 72, 150)).save(stream, format="PNG")
    return stream.getvalue()


def make_docx(path: Path, *, object_safe: bool) -> None:
    doc = Document()
    paragraph = doc.add_paragraph("Rendered format check")
    if object_safe:
        paragraph.paragraph_format.line_spacing = Pt(72)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    paragraph.add_run().add_picture(io.BytesIO(make_image()), width=Inches(1))
    doc.save(path)


def make_pdf(path: Path, *, clip_image: bool = False) -> None:
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_text((72, 72), "Rendered format check", fontname="helv")
    rect = fitz.Rect(-10, -10, 62, 62) if clip_image else fitz.Rect(100, 100, 172, 172)
    page.insert_image(rect, stream=make_image())
    pdf.save(path)
    pdf.close()


class RenderedFormatAuditTests(unittest.TestCase):
    def _run(self, root: Path, *, expected_font: str = "Helvetica",
             clip_image: bool = False, renderer_pdf_hash: str | None = None) -> dict:
        source = root / "source.docx"
        final = root / "final.docx"
        pdf = root / "final.pdf"
        spec = root / "format-spec.json"
        style_map = root / "style-map.json"
        make_docx(source, object_safe=False)
        make_docx(final, object_safe=True)
        make_pdf(pdf, clip_image=clip_image)
        spec.write_text(json.dumps({"roles": {"body_text": {
            "font": {"latin": expected_font},
        }}, "requirements": [
            {"id": "R1", "role": "body_text", "properties": {"font": {"latin": expected_font}}},
            {"id": "R2", "role": "body_text", "properties": {"font": {"latin": "Different Font"}}},
        ]}), encoding="utf-8")
        style_map.write_text(json.dumps({"body_text": {"style_name": "Normal"}}), encoding="utf-8")
        renderer_report = None
        if renderer_pdf_hash:
            renderer_report = root / "renderer-report.json"
            renderer_report.write_text(json.dumps({"final_pdf": {"pdf_sha256": renderer_pdf_hash}}),
                                       encoding="utf-8")
        return audit(source, final, pdf, spec, renderer_report, style_map)

    def test_audit_binds_pdf_and_docx_and_reports_observed_font_resources(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            report = self._run(Path(td))
            self.assertEqual(report["status"], "passed", report["findings"])
            self.assertGreater(report["pdf"]["page_count"], 0)
            self.assertTrue(report["actual_pdf_font_resources"])
            self.assertEqual(report["expected_pdf_fonts"][0]["name"], "Helvetica")
            self.assertEqual(report["expected_pdf_fonts"][0]["requirement_ids"], ["R1"])
            self.assertFalse(report["submission_ready"])
            self.assertFalse(report["field_refresh_claimed"])
            self.assertEqual(len(report["audit_sha256"]), 64)

    def test_font_substitution_and_clipped_image_are_specific_findings(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            report = self._run(Path(td), expected_font="Expected CJK Font", clip_image=True)
            codes = {item["code"] for item in report["findings"]}
            self.assertIn("rendered_pdf_font_mismatch", codes)
            self.assertIn("rendered_drawing_clipped_by_page", codes)
            self.assertEqual(report["status"], "issues_found")

    def test_stale_renderer_report_pdf_hash_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            report = self._run(Path(td), renderer_pdf_hash="0" * 64)
            self.assertIn("renderer_report_pdf_hash_mismatch",
                          {item["code"] for item in report["findings"]})


if __name__ == "__main__":
    unittest.main()
