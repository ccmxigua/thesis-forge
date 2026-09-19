#!/usr/bin/env python3
from __future__ import annotations

import base64
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

from tests.neutral_thesis_fixture import (
    ROLE_STYLE_NAMES, build_neutral_thesis, make_full_role_spec, make_neutral_style_map,
)

ROOT = Path(__file__).resolve().parents[1]
PY = sys.executable
PANDOC_BIN = os.environ.get("PANDOC") or shutil.which("pandoc") or "pandoc"
sys.path.insert(0, str(ROOT / "scripts"))
import apply_format_spec
import requirements_engine
from resource_registry import materialize_declaration_resources


def run(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run([PY, *args], cwd=ROOT, check=True, text=True, capture_output=True)


def run_raw(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run([PY, *args], cwd=ROOT, check=False, text=True, capture_output=True)


def make_requirements(path: Path, ambiguous: bool = False) -> None:
    d = Document()
    lines = [
        "标题使用二号黑体，居中排列，段前、段后各空一行。" if ambiguous else "论文中文题目使用二号黑体，居中排列，段前、段后各空一行。",
        "一级标题使用小三号黑体，编号形式为“第一章”。",
        "正文中文使用小四号宋体，英文和数字使用 Times New Roman，行距固定值20磅。",
        "图题使用五号宋体，居中排列，置于图下方。",
        "纸张采用A4，上页边距2.5厘米，下页边距2.0厘米，左页边距3.0厘米，右页边距2.0厘米。",
    ]
    for text in lines: d.add_paragraph(text)
    d.save(path)


def make_target(path: Path) -> None:
    image = path.with_suffix(".png")
    image.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="))
    d = Document()
    for name in ("ImageCaption", "TableCaption"):
        if name not in {s.name for s in d.styles}: d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    p = d.add_paragraph("基于规则和受约束模型的论文格式转换研究"); p.style = "Title"
    p = d.add_paragraph("绪论"); p.style = "Heading 1"
    p = d.add_paragraph("这是中文正文 with English and 123 numbers."); p.style = "Normal"
    p = d.add_paragraph(); p.add_run().add_picture(str(image), width=Cm(1))
    p = d.add_paragraph("图1-1 系统架构"); p.style = "ImageCaption"
    d.save(path)


def add_page_field(paragraph, switch: str = "ARABIC") -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f"PAGE \\* {switch}")
    paragraph._p.append(field)


def add_sdt_page_field(story, switch: str = "ARABIC") -> None:
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f"PAGE \\* {switch}")
    paragraph.append(field); content.append(paragraph); sdt.append(content)
    story._element.append(sdt)


def make_multi_section_target(path: Path) -> None:
    d = Document()
    d.sections[0].different_first_page_header_footer = True
    d.settings.odd_and_even_pages_header_footer = True
    d.add_paragraph("封面")
    first_footer = d.sections[0].footer.paragraphs[0]
    add_page_field(first_footer, "roman")
    add_page_field(first_footer, "ROMAN")
    add_sdt_page_field(d.sections[0].footer, "roman")
    add_page_field(d.sections[0].first_page_footer.paragraphs[0], "roman")
    add_page_field(d.sections[0].even_page_footer.paragraphs[0], "roman")
    d.add_section(WD_SECTION.NEW_PAGE)
    d.add_paragraph("摘要")
    second_footer = d.sections[1].footer.paragraphs[0]
    add_page_field(second_footer, "roman")
    # Real templates may contain an explicit, valid footer part with no w:p nodes.
    empty_story = d.sections[1].first_page_footer
    for paragraph in list(empty_story.paragraphs):
        empty_story._element.remove(paragraph._p)
    d.add_section(WD_SECTION.NEW_PAGE)
    heading = d.add_paragraph("第一章 绪论"); heading.style = "Heading 1"
    d.add_paragraph("正文内容").style = "Normal"
    third_footer = d.sections[2].footer.paragraphs[0]
    add_page_field(third_footer, "ARABIC")
    d.save(path)


class RequirementsPipelineTest(unittest.TestCase):
    def test_preprocess_infers_xrefs_without_aux(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            source.write_text(r"""\chapter{理论}
\begin{figure}\caption{示意图}\label{fig:a}\end{figure}
见图~\ref{fig:a}。
\chapter{方法}
\begin{equation}x=1\label{eq:a}\end{equation}
见式~\eqref{eq:a}。
""", encoding="utf-8")
            run("scripts/preprocess_tex.py", str(source), str(output))
            text = output.read_text(encoding="utf-8")
            self.assertIn("TJUFE_XREF:fig:a|1.1", text)
            self.assertIn("TJUFE_XREF:eq:a|(2.1)", text)
            self.assertNotIn("[fig:a)", text)

    def test_preprocess_merges_current_structural_labels_with_stale_aux(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            stale_aux = td / "stale.aux"
            source.write_text(r"""\chapter{理论}
\section{模型}\label{sec:model}
见\ref{sec:model}。
""", encoding="utf-8")
            stale_aux.write_text(r"\newlabel{other}{{9.9}{99}{旧标签}{}{}}", encoding="utf-8")
            run("scripts/preprocess_tex.py", str(source), str(output), "--aux", str(stale_aux))
            text = output.read_text(encoding="utf-8")
            self.assertIn("TJUFE_XREF:sec:model|1.1", text)
            self.assertNotIn("Reference [sec:model)", text)

    def test_preprocess_does_not_auto_select_unrelated_aux(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "current.tex"; output = td / "preprocessed.tex"
            source.write_text(r"""\chapter{结果}
\begin{figure}\caption{图}\label{fig:current}\end{figure}
参见\zcref{fig:current}。
""", encoding="utf-8")
            (td / "older-revision.aux").write_text(
                r"\newlabel{fig:current}{{3.99}{42}{旧图}{figure.99}{}}",
                encoding="utf-8",
            )
            run("scripts/preprocess_tex.py", str(source), str(output))
            text = output.read_text(encoding="utf-8")
            self.assertIn("TJUFE_XREF:fig:current|Figure 1.1", text)
            self.assertNotIn("Figure 3.99", text)

    def test_preprocess_numbers_current_theorem_labels_without_aux(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            source.write_text(r"""\chapter{理论}
\begin{proposition}\label{prop:one}结论一。\end{proposition}
\begin{corollary}\label{cor:two}结论二。\end{corollary}
见\zcref{prop:one}与\zcref{cor:two}。
""", encoding="utf-8")
            run("scripts/preprocess_tex.py", str(source), str(output))
            text = output.read_text(encoding="utf-8")
            self.assertIn("TJUFE_XREF:prop:one|Proposition 1.1", text)
            self.assertIn("TJUFE_XREF:cor:two|Corollary 1.2", text)
            self.assertNotIn("[prop:one)", text)
            self.assertNotIn("[cor:two)", text)

    def test_preprocess_rejects_duplicate_source_labels(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            source.write_text(
                r"\begin{equation}a=1\label{eq:dup}\end{equation}"
                "\n"
                r"\begin{equation}b=2\label{eq:dup}\end{equation}",
                encoding="utf-8",
            )
            result = run_raw("scripts/preprocess_tex.py", str(source), str(output))
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("duplicate LaTeX labels are not allowed", result.stderr)
            self.assertIn("eq:dup at lines 1,2", result.stderr)

    def test_preprocess_preserves_supported_newcommand_and_tracks_include_graph(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "main.tex"; child = td / "parts" / "body.tex"
            output = td / "preprocessed.tex"; dependencies = td / "dependencies.json"
            child.parent.mkdir()
            child.write_text(
                r"""\newcommand{\R}{\mathbb{R}}
子文件正文 $x\in\R$。
""", encoding="utf-8")
            source.write_text(
                r"""\documentclass{article}
\newcommand{\vect}[1]{\mathbf{#1}}
\begin{document}
\title{使用 $\vect{x}$ 的题目}
\input{parts/body}
\begin{verbatim}
\input{missing-file}
\ref{not-a-reference}
\end{verbatim}
\end{document}
""", encoding="utf-8")

            run("scripts/preprocess_tex.py", str(source), str(output),
                "--dependency-manifest", str(dependencies))
            text = output.read_text(encoding="utf-8")
            manifest = json.loads(dependencies.read_text(encoding="utf-8"))

            self.assertIn(r"\newcommand{\vect}[1]{\mathbf{#1}}", text)
            self.assertIn(r"\newcommand{\R}{\mathbb{R}}", text)
            self.assertIn("子文件正文", text)
            self.assertIn(r"\input{missing-file}", text)
            self.assertIn(str(source.resolve()), {item["path"] for item in manifest["files"]})
            self.assertIn(str(child.resolve()), {item["path"] for item in manifest["files"]})

    def test_preprocess_rejects_include_cycle_and_mixed_equation_rows(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); main = td / "main.tex"; child = td / "child.tex"
            output = td / "out.tex"
            main.write_text(r"\input{child}", encoding="utf-8")
            child.write_text(r"\input{main}", encoding="utf-8")
            cycle = run_raw("scripts/preprocess_tex.py", str(main), str(output))
            self.assertNotEqual(cycle.returncode, 0)
            self.assertIn("cyclic TeX include graph", cycle.stderr)

            main.write_text(
                r"""\begin{align}
a&=1\\
b&=2\notag
\end{align}
""", encoding="utf-8")
            mixed = run_raw("scripts/preprocess_tex.py", str(main), str(output))
            self.assertNotEqual(mixed.returncode, 0)
            self.assertIn("mixed numbered and unnumbered rows", mixed.stderr)

    def test_preprocess_does_not_number_starred_equations(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            source.write_text(
                r"""\chapter{理论}
\begin{equation*}x=1\label{eq:star}\end{equation*}
见式~\eqref{eq:star}。
\begin{equation}y=2\label{eq:numbered}\end{equation}
见式~\eqref{eq:numbered}。
""", encoding="utf-8")
            run("scripts/preprocess_tex.py", str(source), str(output))
            text = output.read_text(encoding="utf-8")

            self.assertIn("TJUFE_XREF:eq:star|[eq:star)", text)
            self.assertIn("TJUFE_XREF:eq:numbered|(1.1)", text)
            self.assertEqual(text.count("TJUFE_EQCONTROL:unnumbered"), 1)

    def test_preprocess_can_preserve_citations_for_citeproc(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            source.write_text(
                r"""正文 \citep{alpha}，以及 \citet{beta}。
\addbibresource{complete.bib}
\printbibliography[heading=bibintoc]
""",
                encoding="utf-8",
            )
            run("scripts/preprocess_tex.py", str(source), str(output), "--preserve-citations")
            text = output.read_text(encoding="utf-8")
            self.assertIn(r"\citep{alpha}", text)
            self.assertIn(r"\citet{beta}", text)
            self.assertNotIn(r"\printbibliography", text)
            self.assertNotIn(r"\addbibresource", text)

    def test_preprocess_expands_theorem_environments_before_docx_writer(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            source.write_text(
                r"""\begin{remark}
\begin{enumerate}\item 实质正文 \citep{alpha}\end{enumerate}
综合段落。
\end{remark}
\begin{lemma}[\citet{beta}]Lemma body.\end{lemma}
""",
                encoding="utf-8",
            )
            run("scripts/preprocess_tex.py", str(source), str(output), "--preserve-citations")
            text = output.read_text(encoding="utf-8")
            self.assertNotIn(r"\begin{remark}", text)
            self.assertNotIn(r"\begin{lemma}", text)
            self.assertIn("实质正文", text)
            self.assertIn(r"\citep{alpha}", text)
            self.assertIn(r"\citet{beta}", text)

            native = subprocess.run(
                [PANDOC_BIN, str(output), "-f", "latex+raw_tex", "-t", "plain"],
                cwd=ROOT, check=True, text=True, capture_output=True,
            ).stdout
            self.assertIn("实质正文", native)
            self.assertIn("Lemma body", native)

    def test_preprocess_preserves_text_macro_payload_used_in_prose(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            source.write_text(
                r'''其中分别为“\text{U\_BS}”和“\text{N\_Model}”。
数学中的 $\text{TC}=\kappa S|u|$ 应保持为数学命令。
\begin{equation}\text{Var}(X)=1\end{equation}
''',
                encoding="utf-8",
            )
            run("scripts/preprocess_tex.py", str(source), str(output))
            text = output.read_text(encoding="utf-8")
            self.assertIn(r"“U\_BS”和“N\_Model”", text)
            self.assertIn(r"$\text{TC}=\kappa S|u|$", text)
            self.assertIn(r"\begin{equation}\text{Var}(X)=1\end{equation}", text)

            plain = subprocess.run(
                [PANDOC_BIN, str(output), "-f", "latex+raw_tex", "-t", "plain"],
                cwd=ROOT, check=True, text=True, capture_output=True,
            ).stdout
            self.assertIn("U_BS", plain)
            self.assertIn("N_Model", plain)

    def test_preprocess_expands_datetime2_dates_before_pandoc(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            docx = td / "output.docx"
            source.write_text(
                r'''\documentclass{article}
\usepackage{datetime2}
\begin{document}
数据范围从 \DTMdate{2018-01-01} 到 \DTMdate{2020-06-30}。
记录日期为\DTMdate { 2024-11-01 }。
\end{document}
''',
                encoding="utf-8",
            )
            run("scripts/preprocess_tex.py", str(source), str(output))
            text = output.read_text(encoding="utf-8")
            self.assertIn("从 2018-01-01 到 2020-06-30", text)
            self.assertIn("记录日期为2024-11-01", text)
            self.assertNotIn(r"\DTMdate", text)

            subprocess.run(
                [PANDOC_BIN, str(output), "-f", "latex+raw_tex", "-t", "docx", "-o", str(docx)],
                cwd=ROOT, check=True, text=True, capture_output=True,
            )
            visible = "\n".join(p.text for p in Document(docx).paragraphs)
            self.assertIn("从 2018-01-01 到 2020-06-30", visible)
            self.assertIn("记录日期为2024-11-01", visible)

    def test_convert_enables_internal_links_for_citeproc_citations(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.tex"
            bibliography = td / "sources.bib"
            output = td / "output.docx"
            source.write_text(
                r"""\documentclass{article}
\begin{document}
正文引用\cite{alpha}。
\printbibliography
\end{document}
""",
                encoding="utf-8",
            )
            bibliography.write_text(
                "@article{alpha, author={Alpha, Alice}, title={Linked Citation}, "
                "journal={Test Journal}, year={2026}}\n",
                encoding="utf-8",
            )
            subprocess.run(
                [str(ROOT / "convert.sh"), str(source), str(output),
                 f"--bibliography={bibliography}"],
                cwd=ROOT,
                check=True,
                text=True,
                capture_output=True,
            )
            with zipfile.ZipFile(output) as archive:
                root = ET.fromstring(archive.read("word/document.xml"))
            citation_fields = [
                node for node in root.findall(".//" + qn("w:instrText"))
                if re.search(r'HYPERLINK\s+\\l\s+"REF[0-9]{4}"', node.text or "")
            ]
            self.assertEqual(len(citation_fields), 1)
            self.assertIn('HYPERLINK \\l "REF0001"', citation_fields[0].text)
            field_parent = next(
                paragraph for paragraph in root.findall(".//" + qn("w:p"))
                if citation_fields[0] in paragraph.findall(".//" + qn("w:instrText"))
            )
            self.assertEqual(
                "".join(text.text or "" for text in field_parent.findall(".//" + qn("w:t"))),
                "正文引用[1]。",
            )
            bookmark_names = {
                node.get(qn("w:name"))
                for node in root.findall(".//" + qn("w:bookmarkStart"))
            }
            self.assertIn("REF0001", bookmark_names)
            self.assertNotIn("ref-alpha", bookmark_names)

            bibliography_paragraphs = [
                paragraph for paragraph in root.findall(".//" + qn("w:p"))
                if paragraph.find("./" + qn("w:pPr") + "/" + qn("w:pStyle")) is not None
                and paragraph.find("./" + qn("w:pPr") + "/" + qn("w:pStyle")).get(qn("w:val"))
                == "Bibliography"
            ]
            self.assertEqual(len(bibliography_paragraphs), 1)
            paragraph = bibliography_paragraphs[0]
            rendered = "".join(text.text or "" for text in paragraph.findall(".//" + qn("w:t")))
            self.assertTrue(rendered.startswith("[1] "))
            self.assertFalse(rendered.startswith("[1]  "))
            self.assertEqual(paragraph.findall(".//" + qn("w:tab")), [])

    def test_caption_detector_does_not_count_prose_cross_reference(self) -> None:
        from scripts.docx_semantics import is_figure_caption, is_table_caption
        doc = Document()
        figure_caption = doc.add_paragraph("图2-1    模型结构", style="Caption")
        figure_reference = doc.add_paragraph("图 2.1 展示了模型结构。")
        table_caption = doc.add_paragraph("表2-1    数据统计", style="Caption")
        table_reference = doc.add_paragraph("表 2.1 总结了数据。")
        self.assertTrue(is_figure_caption(figure_caption))
        self.assertFalse(is_figure_caption(figure_reference))
        self.assertTrue(is_table_caption(table_caption))
        self.assertFalse(is_table_caption(table_reference))

    def test_sample_thesis_serializes_without_xref_markers(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); output = td / "sample.docx"; audit = td / "audit.json"
            subprocess.run([str(ROOT / "convert.sh"), "tests/sample-thesis.tex", str(output)],
                           cwd=ROOT, check=True, text=True, capture_output=True)
            result = run_raw("scripts/submission_audit.py", str(output), "--out", str(audit))
            self.assertIn(result.returncode, {0, 1, 2})
            report = json.loads(audit.read_text(encoding="utf-8"))
            codes = {issue["code"] for issue in report["issues"]}
            self.assertNotIn("unresolved_source_marker_bracket_reference", codes)
            self.assertNotIn("duplicate_equation_reference_prefix", codes)
            # The fixture's narrative declaration now matches its three
            # numbered body chapters; this regression remains scoped to clean
            # cross-reference serialization and must not reintroduce a chapter
            # count mismatch.
            self.assertNotIn("chapter_count_claim_mismatch", codes)

    def test_pipeline_merges_unspecified_official_page_baseline(self) -> None:
        spec_module = importlib.util.spec_from_file_location(
            "thesis_format_pipeline", ROOT / "scripts" / "thesis_format_pipeline.py")
        module = importlib.util.module_from_spec(spec_module)
        assert spec_module.loader is not None
        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            spec_module.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        merge_official_page_baseline = module.merge_official_page_baseline
        with tempfile.TemporaryDirectory() as td:
            target = Path(td) / "official.docx"
            doc = Document(); section = doc.sections[0]
            section.top_margin = Pt(70); section.bottom_margin = Pt(71)
            section.left_margin = Pt(72); section.right_margin = Pt(73)
            doc.save(target)
            spec = {"page": {"margins_pt": {"left": 90}}}
            added = merge_official_page_baseline(spec, target)
            self.assertEqual(spec["page"]["margins_pt"]["left"], 90)
            self.assertAlmostEqual(spec["page"]["margins_pt"]["top"], 70, places=2)
            self.assertIn("margins_pt.top", added)
            self.assertNotIn("margins_pt.left", added)

    def test_rule_parser_and_ambiguous_question(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; out = td / "out"
            make_requirements(req); run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "rule_only")
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "rule_resolved")
            self.assertEqual(spec["roles"]["thesis_title_zh"]["font"]["size_pt"], 22)
            self.assertEqual(spec["roles"]["body_text"]["font"]["latin"], "Times New Roman")
            self.assertEqual(spec["roles"]["body_text"]["paragraph"]["line_spacing"]["value"], 20)
            self.assertEqual(spec["roles"]["figure_caption"]["position"], "below")
            self.assertAlmostEqual(spec["page"]["margins_pt"]["left"], 3 * 72 / 2.54, places=2)
            make_requirements(req, ambiguous=True); run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "rule_only")
            self.assertEqual(json.loads((out / "format-spec.json").read_text())["status"], "needs_clarification")
            self.assertTrue(json.loads((out / "questions.json").read_text()))

    def test_llm_contract_resolves_role_only(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; out = td / "out"; response = td / "response.json"
            make_requirements(req, ambiguous=True); run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "rule_only")
            q = json.loads((out / "questions.json").read_text())[0]
            response.write_text(json.dumps({"resolutions": [{"question_id": q["id"], "unresolved": False,
                "role": "thesis_title_zh", "evidence_ids": q["evidence_ids"], "reason": "context"}]}, ensure_ascii=False))
            run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "rule_only", "--llm-response", str(response))
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "semantic_resolved")
            self.assertEqual(spec["roles"]["thesis_title_zh"]["font"]["size_pt"], 22)
            self.assertTrue(any(r["role"] == "thesis_title_zh" and r["resolved_by"] == "llm"
                                for r in spec["requirements"]))

    def test_style_evidence_and_docx_round_trip(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            analysis = td / "analysis"; applied = td / "applied"; style_map = td / "style-map.json"; output = td / "formatted.docx"
            make_requirements(req); make_target(target)
            run("scripts/requirements_engine.py", str(req), "--out", str(analysis), "--analysis-mode", "rule_only")
            run("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            mapped = json.loads(style_map.read_text())["mappings"]
            self.assertEqual(mapped["body_text"]["style_name"], "Normal")
            run("scripts/apply_format_spec.py", str(target), str(analysis / "format-spec.json"), str(output), "--out-dir", str(applied), "--style-map", str(style_map))
            report = json.loads((applied / "validation-report.json").read_text())
            self.assertTrue(report["valid"]); self.assertGreaterEqual(report["paragraphs_directly_formatted"]["body_text"], 1)
            d = Document(output)
            self.assertTrue(d.paragraphs[1].text.startswith("第一章"))
            self.assertEqual(d.styles["Normal"].font.size.pt, 12)
            self.assertAlmostEqual(d.sections[0].left_margin.cm, 3.0, places=1)

    def test_exact_bibliography_sample_allows_heading_style_reuse(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; style_map = td / "style-map.json"
            doc = Document()
            doc.add_paragraph("正文", style="Heading 2")
            doc.add_paragraph("参考文献", style="Heading 2")
            doc.save(target)
            run("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            result = json.loads(style_map.read_text())
            self.assertEqual(result["mappings"]["heading_2"]["style_name"], "Heading 2")
            self.assertEqual(result["mappings"]["bibliography_heading"]["style_name"], "Heading 2")
            self.assertFalse(any(q["role"] == "bibliography_heading" for q in result["questions"]))

    def test_preprocessor_preserves_bilingual_title_metadata_from_preamble(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; output = td / "preprocessed.tex"
            source.write_text(
                "\\documentclass{ctexbook}\n"
                "\\title{基于机器学习的城市交通流预测方法研究}\n"
                "\\englishtitle{Research on Urban Traffic Flow Prediction Methods}\n"
                "\\begin{document}\n正文\n\\end{document}\n",
                encoding="utf-8",
            )
            run("scripts/preprocess_tex.py", str(source), str(output))
            text = output.read_text(encoding="utf-8")
            self.assertIn("\\title{基于机器学习的城市交通流预测方法研究}", text)
            self.assertIn("\\begin{english-title}\nResearch on Urban Traffic Flow Prediction Methods\n"
                          "\\end{english-title}", text)
            self.assertNotIn("\\documentclass", text)

    def test_style_analyzer_prefers_unique_table_specific_style_over_normal_tie(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; style_map = td / "style-map.json"
            doc = Document(); doc.styles.add_style("Compact", WD_STYLE_TYPE.PARAGRAPH)
            table = doc.add_table(rows=2, cols=2)
            for index, cell in enumerate(table._cells):
                cell.paragraphs[0].text = str(index)
                cell.paragraphs[0].style = "Compact" if index % 2 == 0 else "Normal"
            doc.save(target)
            result = run_raw("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads(style_map.read_text())
            self.assertEqual(report["status"], "rule_resolved")
            self.assertEqual(report["mappings"]["table_text"]["style_name"], "Compact")
            self.assertFalse(any(q["role"] == "table_text" for q in report["questions"]))

    def test_style_analyzer_uses_structural_table_and_equation_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; style_map = td / "style-map.json"
            doc = Document()
            compact = doc.styles.add_style("Compact", WD_STYLE_TYPE.PARAGRAPH)
            equation_style = doc.styles.add_style("Display Formula", WD_STYLE_TYPE.PARAGRAPH)
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).paragraphs[0].text = "表格正文"
            table.cell(0, 0).paragraphs[0].style = compact
            equation = doc.add_paragraph(style=equation_style)
            math_para = OxmlElement("m:oMathPara")
            math = OxmlElement("m:oMath")
            math_para.append(math); equation._p.append(math_para)
            doc.save(target)
            run("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            mappings = json.loads(style_map.read_text(encoding="utf-8"))["mappings"]
            self.assertEqual(mappings["table_text"]["style_name"], "Compact")
            self.assertEqual(mappings["equation"]["style_name"], "Display Formula")

    def test_english_thesis_title_does_not_claim_abstract_title_en_style(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "titles.docx"; style_map = td / "style-map.json"
            doc = Document()
            english_title = doc.styles.add_style("English Title", WD_STYLE_TYPE.PARAGRAPH)
            abstract_title = doc.styles.add_style("Abstract Title EN", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("A Thesis Title", english_title)
            doc.add_paragraph("Abstract", abstract_title)
            doc.save(target)
            run("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            result = json.loads(style_map.read_text(encoding="utf-8"))
            self.assertEqual(result["mappings"]["thesis_title_en"]["style_name"], "English Title")
            self.assertEqual(result["mappings"]["abstract_title_en"]["style_name"], "Abstract Title EN")
            self.assertNotIn("thesis_title_en", {q["role"] for q in result["questions"]})
            self.assertNotIn("abstract_title_en", {q["role"] for q in result["questions"]})

    def test_style_analyzer_allows_text_disambiguated_heading_style_reuse(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "shared-headings.docx"; style_map = td / "style-map.json"
            doc = Document()
            shared = doc.styles["Heading 1"]
            for text in ("摘 要", "ABSTRACT", "1 引言", "参考文献"):
                doc.add_paragraph(text, shared)
            doc.save(target)
            run("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            result = json.loads(style_map.read_text(encoding="utf-8"))
            for role in ("heading_1", "abstract_title_zh", "abstract_title_en", "bibliography_heading"):
                self.assertEqual(result["mappings"][role]["style_name"], "Heading 1")
                self.assertNotIn(role, {q["role"] for q in result["questions"]})

    def test_style_analyzer_does_not_treat_toc_style_as_body_heading(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "toc.docx"; style_map = td / "style-map.json"
            doc = Document()
            toc = doc.styles.add_style("目录 11", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("1 引言", toc)
            doc.add_paragraph("1 正文章节", doc.styles["Heading 1"])
            doc.save(target)
            run("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            result = json.loads(style_map.read_text(encoding="utf-8"))
            self.assertEqual(result["mappings"]["heading_1"]["style_name"], "Heading 1")
            candidates = [q for q in result["questions"] if q["role"] == "heading_1"]
            self.assertFalse(candidates)

    def test_style_analyzer_does_not_treat_table_of_figures_as_caption(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "captions.docx"; style_map = td / "style-map.json"
            doc = Document()
            list_style = doc.styles.add_style("table of figures", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("图 2.1 图目录条目\t2", list_style)
            doc.add_paragraph("图 2.1 正文图题", doc.styles["Caption"])
            doc.save(target)
            run("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            result = json.loads(style_map.read_text(encoding="utf-8"))
            self.assertEqual(result["mappings"]["figure_caption"]["style_name"], "Caption")
            self.assertNotIn("figure_caption", {q["role"] for q in result["questions"]})

    def test_style_analyzer_recognizes_spaced_chinese_bibliography_heading(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "references.docx"; style_map = td / "style-map.json"
            doc = Document()
            heading = doc.styles.add_style("参考文献", WD_STYLE_TYPE.PARAGRAPH)
            doc.styles.add_style("参考文献标题", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("参 考 文 献", heading)
            doc.save(target)
            run("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            result = json.loads(style_map.read_text(encoding="utf-8"))
            self.assertEqual(result["mappings"]["bibliography_heading"]["style_name"], "参考文献")
            self.assertNotIn("bibliography_heading", {q["role"] for q in result["questions"]})

    def test_style_analyzer_allows_exact_abstract_text_on_normal_style(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "normal-abstract.docx"; style_map = td / "style-map.json"
            doc = Document()
            doc.add_paragraph("普通正文", doc.styles["Normal"])
            doc.add_paragraph("摘 要", doc.styles["Normal"])
            doc.save(target)
            run("scripts/analyze_template_styles.py", str(target), "--out", str(style_map))
            result = json.loads(style_map.read_text(encoding="utf-8"))
            self.assertEqual(result["mappings"]["abstract_title_zh"]["style_name"], "Normal")
            self.assertNotIn("abstract_title_zh", {q["role"] for q in result["questions"]})

    def test_lua_filter_emits_visible_english_thesis_title(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.tex"; preprocessed = td / "preprocessed.tex"
            output = td / "output.docx"
            source.write_text(
                "\\documentclass{ctexbook}\n"
                "\\title{中文论文题名}\n"
                "\\englishtitle{Visible English Thesis Title}\n"
                "\\begin{document}\n正文\n\\end{document}\n",
                encoding="utf-8",
            )
            run("scripts/preprocess_tex.py", str(source), str(preprocessed))
            subprocess.run([
                PANDOC_BIN, str(preprocessed), "--from=latex+raw_tex", "--to=docx", "--standalone",
                f"--reference-doc={ROOT / 'reference.docx'}",
                f"--lua-filter={ROOT / 'filters' / 'thesis-v2.lua'}", "-o", str(output),
            ], cwd=ROOT, check=True, capture_output=True, text=True)
            visible = [(p.text, p.style.name) for p in Document(output).paragraphs if p.text.strip()]
            self.assertIn(("Visible English Thesis Title", "English Title"), visible)

    def test_equation_and_compact_table_content_are_structurally_formatted(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; spec_path = td / "spec.json"
            output = td / "formatted.docx"; audit = td / "audit"
            doc = Document()
            compact = doc.styles.add_style("Compact", WD_STYLE_TYPE.PARAGRAPH)
            display = doc.styles.add_style("Display Formula", WD_STYLE_TYPE.PARAGRAPH)
            table = doc.add_table(rows=1, cols=2)
            for i, cell in enumerate(table.rows[0].cells):
                cell.paragraphs[0].text = f"单元格{i + 1}"
                cell.paragraphs[0].style = compact
            equation = doc.add_paragraph(style=display)
            math_para = OxmlElement("m:oMathPara")
            math_para.append(OxmlElement("m:oMath")); equation._p.append(math_para)
            doc.save(target)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "test", "status": "rule_resolved",
                "roles": {
                    "table_text": {"font": {"cjk": "SimSun", "size_pt": 10.5}},
                    "equation": {"font": {"latin": "Times New Roman", "size_pt": 10.5}},
                },
                "requirements": [],
            }, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(output),
                             "--out-dir", str(audit), "--require-coverage")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((audit / "validation-report.json").read_text(encoding="utf-8"))
            coverage = {item["role"]: item for item in report["role_coverage"]}
            self.assertEqual(coverage["table_text"]["count"], 2)
            self.assertEqual(coverage["equation"]["count"], 1)
            self.assertEqual(report["paragraphs_directly_formatted"]["table_text"], 2)
            self.assertEqual(report["paragraphs_directly_formatted"]["equation"], 1)

    def test_table_text_does_not_mutate_normal_and_receipts_survive_serialization(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document(); doc.add_paragraph("ordinary body", "Normal")
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).paragraphs[0].text = "table content"
            doc.save(source)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "test", "status": "rule_resolved",
                "roles": {"table_text": {"font": {"size_pt": 9}}}, "requirements": [],
            }), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--require-coverage")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            check = Document(output)
            self.assertEqual(check.paragraphs[0].style.name, "Normal")
            self.assertEqual(check.tables[0].cell(0, 0).paragraphs[0].style.name, "Thesis Table Text")
            receipt_audit = json.loads((audit / "execution-receipt-audit.json").read_text())
            self.assertTrue(receipt_audit["valid"])
            receipt = next(item for item in receipt_audit["receipts"] if item["role"] == "table_text")
            self.assertTrue(receipt["node_id"].startswith("body-table-0-row-0-cell-0-paragraph-"))

    def test_explicit_caption_mapping_restyles_structurally_detected_source_caption(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; spec_path = td / "spec.json"
            style_map = td / "style-map.json"; output = td / "formatted.docx"; audit = td / "audit"
            doc = Document(); doc.styles.add_style("Image Caption", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("图2-1 模型结构", "Image Caption"); doc.save(target)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "test", "status": "rule_resolved",
                "roles": {"figure_caption": {"font": {"size_pt": 10.5}}}, "requirements": [],
            }, ensure_ascii=False), encoding="utf-8")
            style_map.write_text(json.dumps({
                "mappings": {"figure_caption": {"style_name": "Caption"}}
            }), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(output),
                             "--out-dir", str(audit), "--style-map", str(style_map), "--require-coverage")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertEqual(Document(output).paragraphs[0].style.name, "Caption")
            report = json.loads((audit / "validation-report.json").read_text())
            self.assertEqual(report["paragraphs_directly_formatted"]["figure_caption"], 1)

    def test_caption_object_graph_binds_across_blank_paragraph_and_is_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document(); drawing = doc.add_paragraph(); drawing.add_run().add_picture(
                str(ROOT / "tests" / "assets" / "frontmatter-contact.jpg"), width=Inches(.2))
            doc.add_paragraph("")
            doc.add_paragraph("图1-1 流程", "Caption"); doc.save(source)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "test", "status": "rule_resolved",
                "roles": {"figure_caption": {"position": "below"}}, "requirements": [],
            }, ensure_ascii=False), encoding="utf-8")
            first = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                            "--out-dir", str(audit), "--require-coverage")
            self.assertEqual(first.returncode, 0, first.stderr + first.stdout)
            graph = json.loads((audit / "caption-object-graph.json").read_text())["bindings"]
            self.assertEqual(graph[0]["status"], "bound")
            self.assertTrue(graph[0]["before"][0]["relationship_ids"])
            second = run_raw("scripts/apply_format_spec.py", str(output), str(spec_path), str(td / "twice.docx"),
                             "--out-dir", str(td / "audit2"), "--require-coverage")
            self.assertEqual(second.returncode, 0, second.stderr + second.stdout)
            report = json.loads((td / "audit2" / "validation-report.json").read_text())
            self.assertEqual(report["caption_paragraphs_moved"]["figure_caption"], 0)

    def test_caption_object_graph_rejects_competing_tables(self) -> None:
        doc = Document(); doc.add_table(rows=1, cols=1); doc.add_table(rows=1, cols=1)
        caption = doc.add_paragraph("表1-1 数据", "Caption")
        moved, issues, graph = apply_format_spec.reposition_captions(
            doc, "table_caption", caption.style.name, "above")
        self.assertEqual(moved, 0)
        self.assertEqual(graph[0]["status"], "ambiguous")
        self.assertEqual(issues[0]["failure_type"], "needs_clarification")

    def test_cover_contract_distinguishes_trusted_placeholder_and_omitted_values(self) -> None:
        cover = {
            "institution": "Test University", "before_role": "abstract_title_zh",
            "missing_value_policy": "placeholder", "missing_value_placeholder": "——",
            "fields": [
                {"id": "title_zh", "label": "题目", "value_from": "thesis_profile.cover_metadata.title_zh", "display_policy": "required", "order": 1},
                {"id": "subtitle_zh", "label": "副题", "value_from": "thesis_profile.cover_metadata.subtitle_zh", "display_policy": "if_present", "order": 2},
            ],
        }
        pending = apply_format_spec.compile_cover_contract(cover, {})
        self.assertEqual([item["value_kind"] for item in pending["fields"]], ["placeholder", "omitted"])
        trusted = apply_format_spec.compile_cover_contract(cover, {"cover_metadata": {
            "trust": {"confirmed": True, "source": "user_confirmed"}, "title_zh": "真实题名",
        }})
        self.assertEqual(trusted["metadata_status"], "trusted")
        self.assertEqual(trusted["fields"][0]["value_kind"], "trusted")
        self.assertEqual(trusted["fields"][0]["value"], "真实题名")

    def test_explicit_keyword_and_table_mappings_migrate_only_narrow_structural_content(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; spec_path = td / "spec.json"
            style_map = td / "style-map.json"; output = td / "formatted.docx"; audit = td / "audit"
            doc = Document()
            doc.styles.add_style("Other Semantic", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("Keywords: documents, validation, testing", "Normal")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).paragraphs[0].text = "ordinary table content"
            table.cell(0, 1).paragraphs[0].text = "protected semantic content"
            table.cell(0, 1).paragraphs[0].style = "Other Semantic"
            doc.save(target)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "test", "status": "rule_resolved",
                "roles": {
                    "keywords_en": {"font": {"latin": "Times New Roman", "size_pt": 12}},
                    "table_text": {"font": {"latin": "Times New Roman", "size_pt": 10.5}},
                },
                "requirements": [],
            }, ensure_ascii=False), encoding="utf-8")
            style_map.write_text(json.dumps({"mappings": {
                "keywords_en": {"style_name": "Official Keywords EN"},
                "table_text": {"style_name": "Official Table Text"},
            }}), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(output),
                             "--out-dir", str(audit), "--style-map", str(style_map), "--require-coverage")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            check = Document(output)
            self.assertEqual(check.paragraphs[0].style.name, "Official Keywords EN")
            self.assertEqual(check.tables[0].cell(0, 0).paragraphs[0].style.name, "Official Table Text")
            self.assertEqual(check.tables[0].cell(0, 1).paragraphs[0].style.name, "Other Semantic")
            coverage = {item["role"]: item for item in
                        json.loads((audit / "validation-report.json").read_text())["role_coverage"]}
            self.assertEqual(coverage["keywords_en"]["count"], 1)
            self.assertEqual(coverage["table_text"]["count"], 1)

    def test_relative_indent_requirement_replaces_official_absolute_fallback(self) -> None:
        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            from apply_format_spec import merge_role_style_defaults
        finally:
            sys.path.pop(0)
        merged = merge_role_style_defaults(
            {"font": {"size_pt": 12}, "paragraph": {"first_line_indent_pt": 10}},
            {"paragraph": {"first_line_indent_chars": 2}},
        )
        self.assertNotIn("first_line_indent_pt", merged["paragraph"])
        self.assertEqual(merged["paragraph"]["first_line_indent_chars"], 2)

    def test_appendix_labels_page_breaks_and_equation_layout_are_applied_and_audited(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; spec_path = td / "spec.json"
            first = td / "first.docx"; second = td / "second.docx"
            doc = Document()
            doc.add_paragraph("第4章 附录A 实验配置", "Heading 1")
            doc.add_paragraph("第5章 附录B 补充结果", "Heading 1")
            equation = doc.add_paragraph()
            math_para = OxmlElement("m:oMathPara"); math_para.append(OxmlElement("m:oMath")); equation._p.append(math_para)
            equation.add_run("\t\t（3.1）")
            ppr = equation._p.get_or_add_pPr(); border = OxmlElement("w:pBdr"); border.append(OxmlElement("w:bottom")); ppr.append(border)
            doc.save(target)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                "roles": {}, "requirements": [],
                "thesis_profile": {"schema_version": "1.0", "degree_level": "doctor",
                                   "writing_language": "zh", "has_appendices": True},
                "appendices": {"required_when_profile_has_appendices": True, "label_style": "alpha_upper",
                               "label_prefix": "附录", "page_break_each": True,
                               "per_appendix_title_required": True},
                "equations": {"alignment": "center", "number_alignment": "right", "number_parentheses": True,
                              "same_line": True, "no_lines": True, "center_tab_twips": 4500,
                              "right_tab_twips": 9000}
            }, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(first),
                             "--out-dir", str(td / "audit-1"))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((td / "audit-1" / "validation-report.json").read_text())
            self.assertTrue(report["valid"], report["findings"])
            self.assertEqual(report["appendix_changes"]["headings"], 2)
            self.assertEqual(report["equation_layout_changes"]["equations"], 1)
            check = Document(first)
            self.assertEqual([p.text for p in check.paragraphs[:2]], ["附录 A  实验配置", "附录 B  补充结果"])
            self.assertTrue(all(p.paragraph_format.page_break_before for p in check.paragraphs[:2]))
            result = run_raw("scripts/apply_format_spec.py", str(first), str(spec_path), str(second),
                             "--out-dir", str(td / "audit-2"))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertEqual(first.read_bytes(), second.read_bytes())

    def test_cover_cleanup_does_not_delete_source_title_style(self) -> None:
        doc = Document()
        for name in ("Thesis Cover Institution", "Thesis Cover Field Value"):
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        institution = doc.add_paragraph("天津大学", "Thesis Cover Institution")
        doc.add_paragraph("题目", "Title")
        field = doc.add_paragraph("指导教师：——", "Thesis Cover Field Value")
        field.add_run().add_break(WD_BREAK.PAGE)
        source_title = doc.add_paragraph("源论文中文题目", "Title")

        removed = apply_format_spec._remove_generated_cover_block(
            doc,
            {"Thesis Cover Institution", "Thesis Cover Field Value", "Title"},
            "天津大学",
        )

        self.assertEqual(removed, 3)
        self.assertEqual([(p.style.name, p.text) for p in doc.paragraphs], [("Title", source_title.text)])

    def test_equation_layout_rejects_a_number_not_driven_by_right_tab(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; spec_path = td / "spec.json"
            doc = Document(); equation = doc.add_paragraph()
            math_para = OxmlElement("m:oMathPara"); math_para.append(OxmlElement("m:oMath")); equation._p.append(math_para)
            equation.add_run("（3.1）")
            doc.save(target)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                "roles": {}, "requirements": [],
                "equations": {"alignment": "center", "number_alignment": "right", "number_parentheses": True,
                              "same_line": True, "no_lines": True}
            }), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(td / "out.docx"),
                             "--out-dir", str(td / "audit"))
            self.assertNotEqual(result.returncode, 0)
            findings = json.loads((td / "audit" / "validation-report.json").read_text())["findings"]
            self.assertTrue(any(item.get("property") == "equations[1].tab_usage" for item in findings))

    def test_v2_overlays_merge(self) -> None:
        for overlay in sorted((ROOT / "schema").glob("*overlay*.yaml")):
            with self.subTest(overlay=overlay.name), tempfile.TemporaryDirectory() as td:
                report = Path(td) / "report.json"
                run("scripts/config_v2.py", "schema/config-schema-v2.yaml", str(overlay), str(Path(td) / "effective.yaml"), "--report", str(report))
                self.assertTrue(json.loads(report.read_text())["valid"])

    def test_negation_and_multi_role_segmentation(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; out = td / "out"
            d = Document()
            d.add_paragraph("一级标题使用黑体且不斜体，二级标题使用宋体且不加粗。")
            d.add_paragraph("正文无需分页，使用小四号宋体。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "llm_primary")
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["roles"]["heading_1"]["font"]["cjk"], "SimHei")
            self.assertFalse(spec["roles"]["heading_1"]["font"]["italic"])
            self.assertEqual(spec["roles"]["heading_2"]["font"]["cjk"], "SimSun")
            self.assertFalse(spec["roles"]["heading_2"]["font"]["bold"])
            self.assertFalse(spec["roles"]["body_text"]["paragraph"]["page_break_before"])

    def test_page_number_clauses_do_not_pollute_paragraph_roles(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; out = td / "out"
            d = Document()
            d.add_paragraph("前置部分页码使用小写罗马数字，正文页码使用阿拉伯数字，页码居中。")
            d.add_paragraph("正文中文使用小四号宋体，两端对齐。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "rule_only")
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "rule_resolved")
            self.assertEqual(spec["page"]["page_number"], {
                "alignment": "center", "front_matter_format": "roman", "body_format": "decimal"
            })
            self.assertEqual(spec["roles"]["body_text"]["paragraph"]["alignment"], "justify")
            self.assertNotIn("page_break_before", spec["roles"]["body_text"]["paragraph"])
            self.assertFalse(json.loads((out / "questions.json").read_text()))
            self.assertFalse(json.loads((out / "conflicts.json").read_text()))

            d = Document()
            d.add_paragraph("前置部分页码使用小写罗马数字。")
            d.add_paragraph("正文页码使用阿拉伯数字。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "rule_only")
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "rule_resolved")
            self.assertNotIn("body_text", spec["roles"])
            self.assertEqual(spec["page"]["page_number"]["front_matter_format"], "roman")
            self.assertEqual(spec["page"]["page_number"]["body_format"], "decimal")
            self.assertFalse(json.loads((out / "questions.json").read_text()))

            d = Document()
            d.add_paragraph("正文页码使用阿拉伯数字，正文使用小四号宋体并两端对齐。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "rule_only")
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "rule_resolved")
            self.assertEqual(spec["page"]["page_number"]["body_format"], "decimal")
            self.assertEqual(spec["roles"]["body_text"]["font"]["size_pt"], 12)
            self.assertEqual(spec["roles"]["body_text"]["paragraph"]["alignment"], "justify")

    def test_page_prefix_before_header_footer_roles_is_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; out = td / "out"
            d = Document()
            d.add_paragraph("页面采用A4纸，上边距30毫米，下边距25毫米，左边距30毫米，右边距25毫米，页眉距顶端23毫米，页脚距底端18毫米。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "llm_primary")
            request = json.loads((out / "llm-request.json").read_text())
            self.assertTrue(any("A4纸" in c["text"] and "上边距30毫米" in c["text"] for c in request["clauses"]))
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["page"]["size"], "A4")
            margins = spec["page"]["margins_pt"]
            self.assertAlmostEqual(margins["top"], 30 * 72 / 25.4, places=3)
            self.assertAlmostEqual(margins["bottom"], 25 * 72 / 25.4, places=3)
            self.assertAlmostEqual(margins["left"], 30 * 72 / 25.4, places=3)
            self.assertAlmostEqual(margins["right"], 25 * 72 / 25.4, places=3)

    def test_invalid_format_spec_is_rejected_before_docx_mutation(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; spec = td / "invalid.json"; output = td / "output.docx"
            make_target(target)
            spec.write_text(json.dumps({"schema_version": "1.0", "source_document": "x", "status": "rule_resolved",
                "roles": {"body_text": {"font": {"size_pt": -1}}}, "requirements": []}), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec), str(output), "--out-dir", str(td / "audit"))
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("invalid format spec", result.stderr)
            self.assertFalse(output.exists())

    def test_template_mode_reports_coverage_without_false_attribute_failure(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "empty-template.docx"
            make_requirements(req); Document().save(target)
            run("scripts/requirements_engine.py", str(req), "--out", str(td / "spec"), "--analysis-mode", "rule_only")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(td / "spec" / "format-spec.json"),
                             str(td / "formatted.docx"), "--out-dir", str(td / "audit"))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((td / "audit" / "validation-report.json").read_text())
            self.assertTrue(report["valid"])
            self.assertFalse(report["fully_covered"])
            self.assertTrue(report["coverage_warnings"])

            strict = run_raw("scripts/apply_format_spec.py", str(target), str(td / "spec" / "format-spec.json"),
                             str(td / "strict.docx"), "--out-dir", str(td / "strict-audit"), "--require-coverage")
            self.assertNotEqual(strict.returncode, 0)
            strict_report = json.loads((td / "strict-audit" / "validation-report.json").read_text())
            self.assertFalse(strict_report["valid"])
            self.assertFalse(strict_report["fully_covered"])

    def test_neutral_thesis_fixture_covers_every_declared_executable_role(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "neutral-thesis.docx"; spec_path = td / "spec.json"
            output = td / "formatted.docx"; audit = td / "audit"; style_map = td / "style-map.json"
            build_neutral_thesis(target)
            spec_path.write_text(json.dumps(make_full_role_spec(), ensure_ascii=False), encoding="utf-8")
            style_map.write_text(json.dumps(make_neutral_style_map()), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(output),
                             "--out-dir", str(audit), "--style-map", str(style_map), "--require-coverage")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((audit / "validation-report.json").read_text())
            self.assertTrue(report["valid"], report["findings"])
            self.assertTrue(report["fully_covered"])
            required = {item["role"]: item for item in report["role_coverage"]
                        if item["expectation"] == "required"}
            self.assertEqual(set(required), set(ROLE_STYLE_NAMES))
            self.assertTrue(all(item["status"] == "present" and item["count"] > 0
                                for item in required.values()))
            self.assertGreaterEqual(required["table_text"]["count"], 4)
            self.assertGreaterEqual(len(Document(output).sections), 4)

    def test_zero_count_content_roles_use_placeholders_but_keep_mapping_gaps_visible(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "incomplete-neutral-thesis.docx"; spec_path = td / "spec.json"
            style_map = td / "style-map.json"
            build_neutral_thesis(target)
            doc = Document(target)
            for paragraph in doc.paragraphs:
                if paragraph.style.name in {ROLE_STYLE_NAMES["thesis_title_en"], ROLE_STYLE_NAMES["heading_2"],
                                             ROLE_STYLE_NAMES["heading_3"]}:
                    paragraph.text = ""
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.style = ROLE_STYLE_NAMES["body_text"]
            doc.save(target)
            spec_path.write_text(json.dumps(make_full_role_spec(), ensure_ascii=False), encoding="utf-8")
            style_map.write_text(json.dumps(make_neutral_style_map()), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(td / "formatted.docx"),
                             "--out-dir", str(td / "audit"), "--style-map", str(style_map), "--require-coverage")
            self.assertNotEqual(result.returncode, 0)
            report = json.loads((td / "audit" / "validation-report.json").read_text())
            self.assertFalse(report["valid"])
            self.assertFalse(report["fully_covered"])
            coverage = {item["role"]: item for item in report["role_coverage"]}
            for role in ("table_text", "thesis_title_en", "heading_2", "heading_3"):
                self.assertEqual(coverage[role]["expectation"], "required")
            self.assertEqual(coverage["table_text"]["status"], "missing")
            self.assertEqual(coverage["thesis_title_en"]["status"], "missing")
            self.assertEqual(coverage["heading_2"]["status"], "present")
            self.assertEqual(coverage["heading_3"]["status"], "present")
            self.assertEqual(coverage["table_text"]["count"], 0)
            self.assertEqual(coverage["thesis_title_en"]["count"], 0)
            self.assertEqual(coverage["table_text"]["failure_type"], "role_mapping_missing")
            self.assertGreaterEqual(coverage["table_text"]["structural_count"], 1)
            self.assertEqual(coverage["thesis_title_en"]["failure_type"], "content_missing")
            missing_findings = {item["role"] for item in report["findings"]
                                if item.get("property") == "coverage"}
            self.assertTrue({"table_text", "thesis_title_en"} <= missing_findings)
            self.assertFalse({"heading_2", "heading_3"} & missing_findings)
            pending = {item["role"]: item for item in report["pending_content"]}
            self.assertEqual(set(pending), {"heading_2", "heading_3"})
            self.assertTrue(all(item["placeholder"] == "——" for item in pending.values()))
            self.assertFalse(report["submission_ready"])

    def test_missing_body_content_gets_neutral_placeholder_without_blocking_formatting(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "missing-body.docx"; spec_path = td / "spec.json"
            style_map = td / "style-map.json"; output = td / "formatted.docx"; audit = td / "audit"
            build_neutral_thesis(target)
            doc = Document(target)
            for paragraph in doc.paragraphs:
                if paragraph.style.name == ROLE_STYLE_NAMES["body_text"]:
                    paragraph.text = ""
            doc.save(target)
            spec_path.write_text(json.dumps(make_full_role_spec(), ensure_ascii=False), encoding="utf-8")
            style_map.write_text(json.dumps(make_neutral_style_map()), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(output),
                             "--out-dir", str(audit), "--style-map", str(style_map), "--require-coverage")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((audit / "validation-report.json").read_text())
            self.assertTrue(report["valid"], report["findings"])
            self.assertFalse(report["fully_covered"])
            self.assertFalse(report["submission_ready"])
            self.assertEqual(report["submission_status"], "content_pending")
            pending = {item["role"]: item for item in report["pending_content"]}
            self.assertEqual(set(pending), {"body_text"})
            self.assertEqual(pending["body_text"]["placeholder"], "——")
            self.assertTrue(any(paragraph.style.name == ROLE_STYLE_NAMES["body_text"]
                                and paragraph.text.strip() == "——"
                                for paragraph in Document(output).paragraphs))

    def test_undeclared_roles_are_optional_and_not_applicable(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "body-only.docx"; spec_path = td / "spec.json"
            doc = Document(); doc.add_paragraph("Only body content", "Normal"); doc.save(target)
            spec = make_full_role_spec(); spec["roles"] = {"body_text": spec["roles"]["body_text"]}
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(td / "formatted.docx"),
                             "--out-dir", str(td / "audit"), "--require-coverage")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((td / "audit" / "validation-report.json").read_text())
            coverage = {item["role"]: item for item in report["role_coverage"]}
            self.assertEqual(coverage["body_text"]["status"], "present")
            self.assertEqual(coverage["table_text"]["expectation"], "optional")
            self.assertEqual(coverage["table_text"]["status"], "not_applicable")
            self.assertTrue(report["fully_covered"])

    def test_invalid_overlay_does_not_write_effective_config(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); overlay = td / "bad.yaml"; effective = td / "effective.yaml"; report = td / "report.json"
            overlay.write_text("font_sizes:\n  body: not-a-number\n", encoding="utf-8")
            result = run_raw("scripts/config_v2.py", "schema/config-schema-v2.yaml", str(overlay), str(effective), "--report", str(report))
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(effective.exists())
            self.assertFalse(json.loads(report.read_text())["valid"])

            overlay.write_text("font_sizes:\n  unknown_role: 12\n", encoding="utf-8")
            result = run_raw("scripts/config_v2.py", "schema/config-schema-v2.yaml", str(overlay), str(effective), "--report", str(report), "--strict")
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(effective.exists())
            self.assertIn("font_sizes.unknown_role", json.loads(report.read_text())["unknown_fields"])

    def test_duplicate_requirement_ids_and_unsupported_inline_are_rejected(self) -> None:
        from scripts.format_spec_validation import load_and_validate
        req = {"id": "R00001", "role": "figure_caption", "properties": {"position": "inline"},
               "evidence_ids": ["E00001"], "resolved_by": "rule", "confidence": 0.98}
        spec = {"schema_version": "1.0", "source_document": "x", "status": "rule_resolved",
                "roles": {"figure_caption": {"position": "inline"}}, "requirements": [req, dict(req)]}
        errors = load_and_validate(spec, ROOT / "schema" / "format-spec.schema.json")
        self.assertTrue(any("inline placement" in e for e in errors))
        self.assertTrue(any("duplicate" in e for e in errors))

        body_spec = {"schema_version": "1.0", "source_document": "x", "status": "rule_resolved",
                     "roles": {"body_text": {"position": "inline"}}, "requirements": []}
        body_errors = load_and_validate(body_spec, ROOT / "schema" / "format-spec.schema.json")
        self.assertFalse(any("inline placement" in e for e in body_errors), body_errors)

    def test_direct_line_spacing_is_parsed_without_fixed_value_keyword(self) -> None:
        from scripts.requirements_engine import parse_properties
        properties = parse_properties("论文题目仿宋14磅，行距16磅，段前段后0磅")
        self.assertEqual(properties["paragraph"]["line_spacing"],
                         {"type": "exact", "value": 16.0, "unit": "pt"})

    def test_fresh_extraction_refuses_to_rebuild_over_host_review_artifacts(self) -> None:
        from scripts.requirements_engine import prepare_fresh_extraction
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "requirements.docx"
            source.write_bytes(b"source")
            out = td / "requirements"
            out.mkdir()
            (out / "merge-receipt.json").write_text("{}", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "immutable Host Agent review artifacts"):
                prepare_fresh_extraction(source, out)

    def test_alternate_textbox_is_not_duplicated_in_outer_paragraph_evidence(self) -> None:
        from scripts.requirements_engine import extract_document_evidence
        document_xml = """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="wps">
  <w:body>
    <w:p>
      <w:r><w:t>外层文字</w:t></w:r>
      <mc:AlternateContent>
        <mc:Choice Requires="wps"><w:drawing><w:txbxContent><w:p><w:r><w:t>论文题目仿宋14磅，行距16磅</w:t></w:r></w:p></w:txbxContent></w:drawing></mc:Choice>
        <mc:Fallback><w:pict><w:txbxContent><w:p><w:r><w:t>论文题目仿宋14磅，行距16磅</w:t></w:r></w:p></w:txbxContent></w:pict></mc:Fallback>
      </mc:AlternateContent>
    </w:p>
    <w:sectPr/>
  </w:body>
</w:document>"""
        with tempfile.TemporaryDirectory() as td:
            docx = Path(td) / "alternate.docx"
            with zipfile.ZipFile(docx, "w") as archive:
                archive.writestr("word/document.xml", document_xml)
            evidence = extract_document_evidence(docx)["evidence"]
        outer = [item for item in evidence if item["kind"] == "paragraph"]
        textboxes = [item for item in evidence if item["kind"] == "textbox"]
        self.assertEqual([item["text"] for item in outer], ["外层文字"])
        self.assertEqual(len(textboxes), 1)
        self.assertEqual(textboxes[0]["location"]["alternate_branch"], "Choice")

    def test_zero_line_spacing_does_not_require_line_height(self) -> None:
        from scripts.format_spec_validation import load_and_validate
        spec = {"schema_version": "1.0", "source_document": "x", "status": "rule_resolved",
                "roles": {"toc": {"paragraph": {"space_before_lines": 0.0,
                                                    "space_after_lines": 0.0}}},
                "requirements": []}
        errors = load_and_validate(spec, ROOT / "schema" / "format-spec.schema.json")
        self.assertFalse(any("spacing_line_height_pt" in error for error in errors), errors)

        spec["roles"]["toc"]["paragraph"]["space_after_lines"] = 1.0
        errors = load_and_validate(spec, ROOT / "schema" / "format-spec.schema.json")
        self.assertTrue(any("spacing_line_height_pt" in error for error in errors), errors)

    def test_chinese_size_name_requires_formatting_context(self) -> None:
        import sys
        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            from requirements_engine import parse_properties
            citation = "天宫二号遥感图像自然景物分类科学数据[DS/OL]"
            self.assertEqual(parse_properties(citation), {})
            self.assertEqual(parse_properties("论文中文题目使用二号黑体")["font"]["size_pt"], 22)
        finally:
            sys.path.pop(0)

    def test_common_template_wording_maps_to_executable_roles_without_false_page_conflicts(self) -> None:
        import sys
        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            from requirements_engine import identify_role, infer_role_from_context, parse_page_properties, parse_properties
            self.assertEqual(identify_role("摘要内容用小四号宋体字书写，两端对齐")[0],
                             "abstract_body_zh")
            self.assertEqual(identify_role("英文摘要部分的题头为“ABSTRACT”，用Times New Roman三号字，居中")[0],
                             "abstract_title_en")
            self.assertEqual(identify_role("论文中出现英文时需要使用Times New Roman字体")[0],
                             "all_text")
            self.assertEqual(identify_role("页码采用Times New Roman五号字体，数字两侧不加修饰线")[0],
                             "footer")
            self.assertNotIn("page_number", parse_page_properties(
                "目录中的标题与页码之间用省略号连接，页码右对齐顶格编排"))
            self.assertNotIn("size_pt", parse_properties("页眉文字之下划横线，线粗1磅").get("font", {}))
            self.assertNotIn("numbering", parse_properties("编号后跟（续），如表1（续）"))
            clause = {"context_before": ["英文摘要"], "context_after": ["论文第二页为英文摘要"]}
            self.assertEqual(infer_role_from_context(
                clause, "题目用三号Times New Roman字体，居中排"), "abstract_title_en")
            self.assertEqual(infer_role_from_context(
                {"context_before": ["论文段落的文字"], "context_after": []},
                "宋体小四号，两端对齐，首行缩进2字符，1.3倍行距"), "body_text")
            self.assertEqual(infer_role_from_context(
                {"source_text_full": "附录正文样式为宋体小四。1.3倍行距，段前0.1行，段后0.1行",
                 "context_before": ["书写格式说明"], "context_after": []},
                "1.3倍行距，段前0.1行，段后0.1行"), "body_text")
            self.assertEqual(infer_role_from_context(
                {"source_kind": "textbox", "source_text_full": "仿宋14磅，行距16磅",
                 "context_before": ["研究生姓名"], "context_after": []},
                "仿宋14磅，行距16磅"), "cover_field_value")
            self.assertEqual(parse_properties("引文内容可用楷体"), {})
            self.assertNotIn("italic", parse_properties("量的符号一律采用斜体" ).get("font", {}))
            self.assertEqual(identify_role("表内字体为宋体，小五号")[0], "table_text")
            self.assertEqual(identify_role("标题“附录A 附录内容名称”样式为黑体小三")[0], "heading_1")
        finally:
            sys.path.pop(0)

    def test_style_analyzer_prefers_builtin_header_on_exact_tie(self) -> None:
        import sys
        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            from analyze_template_styles import analyze
            with tempfile.TemporaryDirectory() as td:
                path = Path(td) / "headers.docx"
                doc = Document()
                doc.styles.add_style("页眉2", 1)
                doc.sections[0].header.paragraphs[0].style = "Header"
                doc.save(path)
                result = analyze(path)
                self.assertEqual(result["mappings"]["header"]["style_name"], "Header")
                self.assertFalse(any(q.get("role") == "header" for q in result["questions"]))
        finally:
            sys.path.pop(0)

    def test_style_analyzer_prefers_builtin_heading_on_exact_tie(self) -> None:
        import sys
        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            from analyze_template_styles import analyze
            with tempfile.TemporaryDirectory() as td:
                path = Path(td) / "headings.docx"
                doc = Document()
                doc.styles.add_style("heading1", 1)
                doc.styles.add_style("heading2", 1)
                doc.save(path)
                result = analyze(path)
                self.assertEqual(result["mappings"]["heading_1"]["style_name"], "Heading 1")
                self.assertEqual(result["mappings"]["heading_2"]["style_name"], "Heading 2")
                self.assertFalse(any(q.get("role") in {"heading_1", "heading_2"}
                                     for q in result["questions"]))
        finally:
            sys.path.pop(0)

    def test_style_analyzer_does_not_promote_table_metadata_to_heading(self) -> None:
        import sys
        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            from analyze_template_styles import analyze
            with tempfile.TemporaryDirectory() as td:
                path = Path(td) / "table-metadata.docx"
                doc = Document()
                table = doc.add_table(rows=1, cols=1)
                table.cell(0, 0).paragraphs[0].add_run("20   年   月   日")
                doc.save(path)
                result = analyze(path)
                self.assertEqual(result["mappings"]["heading_1"]["style_name"], "Heading 1")
                self.assertFalse(any(q.get("role") == "heading_1" for q in result["questions"]))
        finally:
            sys.path.pop(0)

    def test_unified_pipeline_completes_and_blocks_unresolved_requirements(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            make_requirements(req); make_target(target)
            result = run_raw("scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "formatted.docx"),
                             "--work-dir", str(td / "work"), "--analysis-mode", "rule_only",
                             "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            manifest = json.loads((td / "work" / "pipeline-manifest.json").read_text())
            self.assertEqual(manifest["status"], "completed")
            self.assertTrue((td / "formatted.docx").exists())
            section_artifacts = manifest["section_execution"]
            self.assertEqual(
                set(section_artifacts), {"plan", "execution", "audit"})
            self.assertTrue(all(Path(path).exists() for path in section_artifacts.values()))
            apply_step = next(step for step in manifest["steps"] if step["name"] == "apply_and_validate")
            self.assertEqual(apply_step["artifacts"], section_artifacts)

            make_requirements(req, ambiguous=True)
            result = run_raw("scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "blocked.docx"),
                             "--work-dir", str(td / "blocked-work"), "--analysis-mode", "rule_only",
                             "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 3)
            self.assertFalse((td / "blocked.docx").exists())
            self.assertEqual(json.loads((td / "blocked-work" / "pipeline-manifest.json").read_text())["status"], "blocked")

    def test_requirements_engine_reextracts_docx_and_invalidates_old_optional_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; out = td / "requirements-out"
            make_requirements(req)
            first = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                            "--analysis-mode", "rule_only")
            self.assertEqual(first.returncode, 0, first.stderr + first.stdout)
            first_manifest = json.loads((out / "extraction-manifest.json").read_text())
            (out / "llm-response.raw.json").write_text('{"stale": true}', encoding="utf-8")
            (out / "llm-merge-audit.json").write_text('[]', encoding="utf-8")
            doc = Document(req)
            doc.add_paragraph("附录标题使用黑体小三号。")
            doc.save(req)
            second = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                            "--analysis-mode", "rule_only")
            self.assertEqual(second.returncode, 0, second.stderr + second.stdout)
            second_manifest = json.loads((out / "extraction-manifest.json").read_text())
            self.assertNotEqual(first_manifest["run_id"], second_manifest["run_id"])
            self.assertNotEqual(first_manifest["source_sha256"], second_manifest["source_sha256"])
            self.assertGreater(second_manifest["clause_count"], first_manifest["clause_count"])
            self.assertFalse(second_manifest["cache_reused"])
            self.assertIn("llm-response.raw.json", second_manifest["invalidated_prior_artifacts"])
            self.assertIn("llm-merge-audit.json", second_manifest["invalidated_prior_artifacts"])
            self.assertFalse((out / "llm-response.raw.json").exists())
            self.assertFalse((out / "llm-merge-audit.json").exists())
            second_spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(second_spec["run_id"], second_manifest["run_id"])

    def test_unified_pipeline_reextracts_requirements_when_work_dir_is_reused(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            work = td / "work"
            make_requirements(req); make_target(target)
            first = run_raw("scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "first.docx"),
                            "--work-dir", str(work), "--analysis-mode", "rule_only",
                            "--compliance-mode", "supported_subset")
            self.assertEqual(first.returncode, 0, first.stderr + first.stdout)
            first_extraction = json.loads((work / "requirements" / "extraction-manifest.json").read_text())
            stale = work / "requirements" / "must-not-survive.txt"
            stale.write_text("old extraction", encoding="utf-8")
            doc = Document(req)
            doc.add_paragraph("附录标题使用黑体小三号。")
            doc.save(req)
            second = run_raw("scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "second.docx"),
                             "--work-dir", str(work), "--analysis-mode", "rule_only",
                             "--compliance-mode", "supported_subset",
                             "--allow-existing-work")
            self.assertEqual(second.returncode, 0, second.stderr + second.stdout)
            second_extraction = json.loads((work / "requirements" / "extraction-manifest.json").read_text())
            manifest = json.loads((work / "pipeline-manifest.json").read_text())
            # Fresh extraction invalidates only known generated artifacts.  An
            # unrelated file is preserved so rebuilding a stage cannot erase a
            # user's evidence or an immutable host-review artifact by accident.
            self.assertTrue(stale.exists(), "unrelated stage files must be preserved")
            self.assertNotEqual(first_extraction["run_id"], second_extraction["run_id"])
            self.assertNotEqual(first_extraction["source_sha256"], second_extraction["source_sha256"])
            self.assertEqual(manifest["requirements_extraction"]["policy"],
                             "fresh_required_docx_extraction")
            self.assertFalse(manifest["requirements_extraction"]["cache_reused"])
            self.assertEqual(manifest["requirements_extraction"]["run_id"], second_extraction["run_id"])
            self.assertEqual(manifest["requirements_extraction"]["source_sha256"],
                             second_extraction["source_sha256"])

    def test_latex_stage_reuses_source_docx_bound_to_same_existing_run(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            req = td / "requirements.docx"
            work = td / "work"
            make_requirements(req)
            first = run_raw(
                "scripts/thesis_format_pipeline.py", str(req), "tests/sample-thesis.tex",
                str(td / "first.docx"), "--work-dir", str(work),
                "--analysis-mode", "rule_only", "--compliance-mode", "supported_subset",
            )
            self.assertEqual(first.returncode, 0, first.stderr + first.stdout)
            first_manifest = json.loads((work / "pipeline-manifest.json").read_text())
            first_intermediate = first_manifest["intermediate_docx"]

            second = run_raw(
                "scripts/thesis_format_pipeline.py", str(req), "tests/sample-thesis.tex",
                str(td / "second.docx"), "--work-dir", str(work),
                "--analysis-mode", "rule_only", "--compliance-mode", "supported_subset",
                "--allow-existing-work",
            )
            self.assertEqual(second.returncode, 0, second.stderr + second.stdout)
            second_manifest = json.loads((work / "pipeline-manifest.json").read_text())
            self.assertEqual(second_manifest["steps"][0]["name"], "latex_to_docx_reuse")
            self.assertEqual(second_manifest["steps"][0]["reused"], True)
            self.assertEqual(second_manifest["intermediate_docx"], first_intermediate)

    def test_unified_pipeline_accepts_tex_and_records_end_to_end_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; output = td / "formatted.docx"
            make_requirements(req)
            result = run_raw(
                "scripts/thesis_format_pipeline.py", str(req), "tests/sample-thesis.tex", str(output),
                "--work-dir", str(td / "work"), "--analysis-mode", "rule_only",
                "--compliance-mode", "supported_subset",
            )
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertTrue(output.exists())
            manifest = json.loads((td / "work" / "pipeline-manifest.json").read_text())
            self.assertEqual(manifest["schema_version"], "1.1")
            self.assertEqual(manifest["pipeline_level"], "latex_end_to_end")
            self.assertEqual(manifest["inputs"]["source_kind"], "latex")
            self.assertEqual(Path(manifest["inputs"]["source"]["path"]),
                             (ROOT / "tests" / "sample-thesis.tex").resolve())
            self.assertEqual(len(manifest["inputs"]["source"]["sha256"]), 64)
            self.assertEqual(len(manifest["intermediate_docx"]["sha256"]), 64)
            self.assertEqual(len(manifest["output_artifact"]["sha256"]), 64)
            self.assertEqual(manifest["steps"][0]["name"], "latex_to_docx")
            self.assertTrue(Path(manifest["intermediate_docx"]["path"]).exists())
            generated = Document(output)
            text = "\n".join(p.text for p in generated.paragraphs)
            self.assertIn("基于机器学习的城市交通流预测方法研究", text)
            self.assertGreaterEqual(len(generated.tables), 2)

    def test_docx_input_is_explicitly_labeled_as_formatting_stage(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            make_requirements(req); make_target(target)
            result = run_raw(
                "scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "formatted.docx"),
                "--work-dir", str(td / "work"), "--analysis-mode", "rule_only",
                "--compliance-mode", "supported_subset",
            )
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            manifest = json.loads((td / "work" / "pipeline-manifest.json").read_text())
            self.assertEqual(manifest["pipeline_level"], "docx_formatting_stage")
            self.assertEqual(manifest["inputs"]["source_kind"], "docx")
            self.assertNotIn("intermediate_docx", manifest)

    def test_full_pipeline_rejects_missing_semantic_review_before_work(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            make_requirements(req); make_target(target)
            result = run_raw("scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "output.docx"),
                             "--work-dir", str(td / "work"))
            self.assertEqual(result.returncode, 2, result.stderr + result.stdout)
            self.assertIn("full compliance requires a fresh complete host-Agent clause review", result.stderr)
            self.assertFalse((td / "work").exists())
            self.assertFalse((td / "output.docx").exists())

    def test_host_agent_preparation_is_provider_free_and_stops_before_output(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            make_requirements(req); make_target(target)
            result = run_raw(
                "scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "output.docx"),
                "--work-dir", str(td / "work"), "--prepare-host-review", "--host-review-chunk-size", "2",
            )
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            manifest = json.loads((td / "work" / "pipeline-manifest.json").read_text())
            self.assertEqual(manifest["status"], "host_review_required")
            review_manifest = json.loads(
                (td / "work" / "requirements" / "host-agent-review-manifest.json").read_text()
            )
            self.assertEqual(review_manifest["protocol"], "host_agent_semantic_review")
            self.assertGreaterEqual(review_manifest["chunk_count"], 1)
            self.assertFalse((td / "output.docx").exists())

    def test_full_pipeline_rejects_rule_modes_before_work(self) -> None:
        for mode in ("rule_only", "known_template"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as td:
                td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
                make_requirements(req); make_target(target)
                result = run_raw(
                    "scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "output.docx"),
                    "--work-dir", str(td / "work"), "--analysis-mode", mode,
                    "--compliance-mode", "full",
                )
                self.assertEqual(result.returncode, 2, result.stderr + result.stdout)
                self.assertIn("requires --analysis-mode llm_primary", result.stderr)
                self.assertFalse((td / "work").exists())

    def test_dual_page_numbering_is_section_aware_and_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; spec_path = td / "spec.json"
            first = td / "first.docx"; second = td / "second.docx"
            make_multi_section_target(target)
            spec = {
                "schema_version": "1.0", "source_document": "test", "analysis_mode": "llm_primary",
                "status": "semantic_resolved", "roles": {}, "requirements": [],
                "completeness": {"reviewed_by": "llm", "covered_clause_ids": [],
                                 "ignored_clause_ids": [], "unresolved_clause_ids": [], "unsupported_items": []},
                "page": {"page_number": {"alignment": "center", "front_matter_format": "roman",
                    "body_format": "decimal", "front_matter_start": 1, "body_start": 1,
                    "body_start_selector": {"strategy": "section_index", "section_index": 3},
                    "replace_existing_fields": True, "preserve_existing_locations": True},
                    "different_first_page": True, "different_odd_even": True}
            }
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(first),
                             "--out-dir", str(td / "audit-1"))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((td / "audit-1" / "validation-report.json").read_text())
            self.assertTrue(report["valid"], report["findings"])
            self.assertGreaterEqual(report["header_footer_changes"]["page_fields_removed"], 5)

            result = run_raw("scripts/apply_format_spec.py", str(first), str(spec_path), str(second),
                             "--out-dir", str(td / "audit-2"))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((td / "audit-2" / "validation-report.json").read_text())
            self.assertTrue(report["valid"], report["findings"])
            self.assertEqual(report["header_footer_changes"]["page_numbers"], 3)
            self.assertEqual(first.read_bytes(), second.read_bytes())

    def test_body_text_scope_does_not_mutate_front_matter_and_is_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; spec_path = td / "spec.json"; style_map = td / "style-map.json"
            first = td / "first.docx"; second = td / "second.docx"
            d = Document(); d.add_paragraph("封面信息").style = "Normal"
            d.add_section(WD_SECTION.NEW_PAGE)
            d.add_paragraph("正文内容").style = "Normal"; d.save(target)
            spec = {
                "schema_version": "1.0", "source_document": "test", "analysis_mode": "llm_primary",
                "status": "semantic_resolved", "requirements": [],
                "roles": {"body_text": {"font": {"cjk": "SimSun", "latin": "Times New Roman", "size_pt": 12},
                    "paragraph": {"alignment": "justify", "line_spacing": {"type": "exact", "value": 20, "unit": "pt"}}}},
                "completeness": {"reviewed_by": "llm", "covered_clause_ids": [], "ignored_clause_ids": [],
                                 "unresolved_clause_ids": [], "unsupported_items": []},
                "page": {"page_number": {"front_matter_format": "none", "body_format": "decimal", "body_start": 1,
                    "body_start_selector": {"strategy": "section_index", "section_index": 2},
                    "preserve_existing_locations": False}}
            }
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            style_map.write_text(json.dumps({"mappings": {"body_text": {"style_name": "Normal"}}}), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(first),
                             "--out-dir", str(td / "audit-1"), "--style-map", str(style_map), "--require-coverage")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            result = run_raw("scripts/apply_format_spec.py", str(first), str(spec_path), str(second),
                             "--out-dir", str(td / "audit-2"), "--style-map", str(style_map), "--require-coverage")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            check = Document(second)
            cover = next(p for p in check.paragraphs if p.text == "封面信息")
            body = next(p for p in check.paragraphs if p.text == "正文内容")
            self.assertEqual(cover.style.name, "Normal")
            self.assertEqual(body.style.name, "Thesis Body Text")
            self.assertEqual(json.loads((td / "audit-2" / "validation-report.json").read_text())["styles_created"], [])
            self.assertEqual(first.read_bytes(), second.read_bytes())

    def test_llm_primary_21_requires_complete_reasoned_clause_reviews(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; preview = td / "preview"; response = td / "response.json"
            d = Document(); d.add_paragraph("正文中文使用小四号宋体，两端对齐。"); d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview), "--analysis-mode", "llm_primary")
            clauses = json.loads((preview / "requirement-clauses.json").read_text())
            self.assertEqual(len(clauses), 1)
            clause = clauses[0]
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [{"role": "body_text", "properties": {"font": {"cjk": "SimSun", "size_pt": 12},
                    "paragraph": {"alignment": "justify"}}, "clause_ids": [clause["id"]],
                    "evidence_ids": clause["evidence_ids"], "confidence": 0.98, "reason": "明确的正文格式条款"}],
                "clause_reviews": [{"clause_id": clause["id"], "classification": "covered",
                    "requirement_indexes": [0], "reason": "由需求0完整覆盖"}],
                "unsupported_items": [], "reported_conflicts": []
            }, ensure_ascii=False), encoding="utf-8")
            out = td / "accepted"
            result = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                             "--analysis-mode", "llm_primary", "--llm-response", str(response))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "semantic_resolved")
            self.assertEqual(spec["completeness"]["covered_clause_ids"], [clause["id"]])

            broken = json.loads(response.read_text())
            broken["clause_reviews"][0]["reason"] = ""
            broken["requirements"][0]["confidence"] = True
            response.write_text(json.dumps(broken, ensure_ascii=False), encoding="utf-8")
            blocked = td / "blocked"
            result = run_raw("scripts/requirements_engine.py", str(req), "--out", str(blocked),
                             "--analysis-mode", "llm_primary", "--llm-response", str(response))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertEqual(json.loads((blocked / "format-spec.json").read_text())["status"], "needs_clarification")
            reasons = json.dumps(json.loads((blocked / "conflicts.json").read_text()), ensure_ascii=False)
            self.assertIn("missing_reason", reasons)
            self.assertIn("invalid_confidence", reasons)

    def test_complete_21_review_is_applied_independently_of_baseline_mode(self) -> None:
        """rule_only/known_template select a baseline, not a review contract."""
        for mode in ("rule_only", "known_template"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as td:
                td = Path(td); req = td / "requirements.docx"; preview = td / "preview"
                response = td / "response.json"; out = td / "out"
                d = Document(); d.add_paragraph("本页为模板填写操作说明。"); d.save(req)
                run("scripts/requirements_engine.py", str(req), "--out", str(preview),
                    "--analysis-mode", mode)
                clauses = json.loads((preview / "requirement-clauses.json").read_text())
                self.assertEqual(len(clauses), 1)
                response.write_text(json.dumps({
                    "contract_version": "2.1", "requirements": [],
                    "clause_reviews": [{
                        "clause_id": clauses[0]["id"], "classification": "informational",
                        "requirement_indexes": [], "reason": "模板操作说明，不产生独立格式义务"
                    }],
                    "unsupported_items": [], "reported_conflicts": []
                }, ensure_ascii=False), encoding="utf-8")
                result = run_raw(
                    "scripts/requirements_engine.py", str(req), "--out", str(out),
                    "--analysis-mode", mode, "--llm-response", str(response),
                )
                self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
                spec = json.loads((out / "format-spec.json").read_text())
                self.assertEqual(spec["status"], "semantic_resolved")
                self.assertEqual(len(spec["clause_compliance"]), 1)
                self.assertEqual(spec["clause_compliance"][0]["status"], "informational")
                self.assertEqual(spec["completeness"]["reviewed_by"], "llm")

    def test_btbu_numbered_bibliography_sample_content_cannot_be_executable(self) -> None:
        """Filled BTBU reference entries are sample content, not obligations."""
        source = ROOT / "inputs" / "ten-school-templates" / "btbu-requirements.docx"
        if not source.is_file():
            self.skipTest("external BTBU template is not included in this source checkout")
        evidence = requirements_engine.extract_document_evidence(source)
        clauses = requirements_engine.split_clauses(evidence)
        reference_ids = [f"C{i:05d}" for i in range(250, 256)]
        references = [clause for clause in clauses if clause["id"] in reference_ids]
        self.assertEqual([clause["id"] for clause in references], reference_ids)

        requirements = []
        reference_indexes = {}
        for clause in references:
            reference_indexes[clause["id"]] = len(requirements)
            requirements.append({
                "role": "body_text",
                "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                "clause_ids": [clause["id"]],
                "evidence_ids": clause["evidence_ids"],
                "confidence": 0.99,
                "reason": "The exact numbered entry appears in the input template.",
            })
        response = {
            "contract_version": "2.1",
            "requirements": requirements,
            "clause_reviews": [
                *[
                    {
                        "clause_id": clause["id"],
                        "classification": "executable",
                        "requirement_indexes": [reference_indexes[clause["id"]]],
                        "reason": "The exact numbered entry appears in the input template.",
                    }
                    for clause in references
                ],
                *[
                    {
                        "clause_id": clause["id"],
                        "classification": "informational",
                        "requirement_indexes": [],
                        "reason": "No independent formatting obligation was identified.",
                    }
                    for clause in clauses if clause["id"] not in reference_ids
                ],
            ],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        spec, conflicts, audit = requirements_engine.merge_llm_primary(
            source,
            {"schema_version": "1.0", "roles": {}, "page": {},
             "requirements": [], "content_instances": []},
            clauses,
            response,
            {item["id"] for item in evidence["evidence"]},
        )

        self.assertFalse(
            any(item.get("type") in {"llm_contract", "llm_internal_conflict", "completeness"}
                for item in conflicts),
            conflicts,
        )
        self.assertEqual(spec["requirements"], [])
        records = {item["clause_id"]: item for item in spec["clause_compliance"]}
        self.assertTrue(all(records[cid]["status"] == "informational" for cid in reference_ids))
        guards = [item for item in audit if item.get("type") == "normative_scope_guard"]
        self.assertEqual(len(guards), 1)
        self.assertEqual(
            {item["clause_id"] for item in guards[0]["changes"]},
            set(reference_ids),
        )

    def test_llm_primary_persists_strict_execution_contract_fields(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; preview = td / "preview"
            response = td / "response.json"; out = td / "out"
            d = Document(); d.add_paragraph("仅博士论文正文使用小四号宋体，并须经 Word 刷新后检查分页。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview),
                "--analysis-mode", "llm_primary")
            clause = json.loads((preview / "requirement-clauses.json").read_text())[0]
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [{
                    "role": "body_text",
                    "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                    "clause_ids": [clause["id"]], "evidence_ids": clause["evidence_ids"],
                    "confidence": .98, "reason": "博士论文条件下的正文格式要求",
                    "applicability": {"status": "conditional", "conditions": [{
                        "fact": "thesis_profile.degree_level", "operator": "equals", "value": "doctor"
                    }], "exceptions": []},
                    "input_prerequisites": [{
                        "kind": "metadata", "key": "thesis_profile.degree_level",
                        "required": True, "reason": "需要学位层次判断条件是否成立"
                    }],
                    "verification": {"mode": "word_render", "checks": ["pagination_refreshed"]}
                }],
                "clause_reviews": [{"clause_id": clause["id"], "classification": "covered",
                    "requirement_indexes": [0], "reason": "由需求0覆盖"}],
                "unsupported_items": [], "reported_conflicts": []
            }, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                             "--analysis-mode", "llm_primary", "--llm-response", str(response))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            requirement = json.loads((out / "format-spec.json").read_text())["requirements"][0]
            self.assertEqual(requirement["applicability"]["status"], "conditional")
            self.assertEqual(requirement["input_prerequisites"][0]["key"],
                             "thesis_profile.degree_level")
            self.assertEqual(requirement["verification"]["mode"], "word_render")
            self.assertEqual(requirement["reason"], "博士论文条件下的正文格式要求")

    def test_llm_primary_rejects_unknown_role_and_invalid_contract_namespace(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; preview = td / "preview"
            response = td / "response.json"; out = td / "out"
            d = Document(); d.add_paragraph("正文中文使用小四号宋体。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview),
                "--analysis-mode", "llm_primary")
            clause = json.loads((preview / "requirement-clauses.json").read_text())[0]
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [{
                    "role": "ujs_special_body", "properties": {"font": {"cjk": "SimSun"}},
                    "clause_ids": [clause["id"]], "evidence_ids": clause["evidence_ids"],
                    "confidence": .9, "reason": "非法学校专用角色",
                    "input_prerequisites": [{"kind": "metadata", "key": "shell.command",
                        "required": True, "reason": "非法命名空间"}]
                }],
                "clause_reviews": [{"clause_id": clause["id"], "classification": "covered",
                    "requirement_indexes": [0], "reason": "由需求0覆盖"}],
                "unsupported_items": [], "reported_conflicts": []
            }, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                             "--analysis-mode", "llm_primary", "--llm-response", str(response))
            self.assertEqual(result.returncode, 1, result.stderr + result.stdout)
            self.assertEqual(json.loads((out / "format-spec.json").read_text())["status"],
                             "needs_clarification")
            conflicts = json.dumps(json.loads((out / "conflicts.json").read_text()), ensure_ascii=False)
            self.assertIn("unknown_requirement_role", conflicts)
            self.assertIn("$.input_prerequisites[0].key", conflicts)
            self.assertIn("shell.command", conflicts)

    def test_rejected_requirement_does_not_shift_later_review_indexes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; preview = td / "preview"
            response = td / "response.json"; out = td / "out"
            d = Document()
            d.add_paragraph("作者字段标签使用五号字。")
            d.add_paragraph("正文中文使用小四号宋体。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview),
                "--analysis-mode", "llm_primary")
            clauses = json.loads((preview / "requirement-clauses.json").read_text())
            self.assertEqual(len(clauses), 2)
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [
                    {"role": "not_a_registered_role", "properties": {"font": {"size_pt": 10.5}},
                     "clause_ids": [clauses[0]["id"]], "evidence_ids": clauses[0]["evidence_ids"],
                     "confidence": .9, "reason": "用于验证拒绝首项时索引不漂移"},
                    {"role": "body_text", "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                     "clause_ids": [clauses[1]["id"]], "evidence_ids": clauses[1]["evidence_ids"],
                     "confidence": .98, "reason": "明确的正文格式条款"},
                ],
                "clause_reviews": [
                    {"clause_id": clauses[0]["id"], "classification": "covered",
                     "requirement_indexes": [0], "reason": "由被拒绝的需求0引用"},
                    {"clause_id": clauses[1]["id"], "classification": "covered",
                     "requirement_indexes": [1], "reason": "由有效需求1引用"},
                ],
                "unsupported_items": [], "reported_conflicts": []
            }, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                             "--analysis-mode", "llm_primary", "--llm-response", str(response))
            self.assertEqual(result.returncode, 1, result.stderr + result.stdout)
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(len(spec["requirements"]), 1)
            self.assertEqual(spec["requirements"][0]["role"], "body_text")
            records = {item["clause_id"]: item for item in spec["clause_compliance"]}
            self.assertEqual(records[clauses[1]["id"]]["requirement_ids"], ["R00001"])
            self.assertEqual(records[clauses[0]["id"]]["requirement_ids"], [])
            self.assertEqual(records[clauses[0]["id"]]["status"], "unresolved")
            self.assertIn("rejected", records[clauses[0]["id"]]["reason"])

    def test_declaration_fragments_merge_by_item_id_and_materialize(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; preview = td / "preview"
            response = td / "response.json"; out = td / "out"
            d = Document()
            d.add_paragraph("独创性声明")
            d.add_paragraph("本人声明本论文为本人在导师指导下完成的研究成果。")
            d.add_paragraph("使用授权说明")
            d.add_paragraph("本人同意按照学校规定保存和使用本学位论文。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview),
                "--analysis-mode", "llm_primary")
            clauses = json.loads((preview / "requirement-clauses.json").read_text())
            self.assertEqual(len(clauses), 4)
            originality_heading, originality_body, authorization_heading, authorization_body = clauses
            evidence_items = {
                item["id"]: item["text"]
                for item in json.loads((preview / "document-evidence.json").read_text())["evidence"]
            }
            source_text = lambda clause: evidence_items[clause["evidence_ids"][0]]
            authorization_evidence = authorization_heading["evidence_ids"] + authorization_body["evidence_ids"]
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [
                    {
                        "role": "declarations",
                        "properties": {
                            "before_role": "abstract_title_zh",
                            "items": [{
                                "id": "originality", "heading": source_text(originality_heading),
                                "source_evidence_ids": originality_heading["evidence_ids"],
                                "signature_placeholders": [],
                            }],
                        },
                        "clause_ids": [originality_heading["id"]],
                        "evidence_ids": originality_heading["evidence_ids"],
                        "confidence": .98, "reason": "保留原文固定声明标题",
                    },
                    {
                        "role": "declarations",
                        "properties": {
                            "before_role": "abstract_title_zh",
                            "items": [{
                                "id": "originality",
                                "body_parts": [source_text(originality_body)],
                                "source_evidence_ids": originality_body["evidence_ids"],
                                "signature_placeholders": [],
                            }, {
                                "id": "authorization", "heading": source_text(authorization_heading),
                                "body_parts": [source_text(authorization_body)],
                                "source_evidence_ids": authorization_evidence,
                                "signature_placeholders": [],
                            }],
                        },
                        "clause_ids": [originality_body["id"], authorization_heading["id"], authorization_body["id"]],
                        "evidence_ids": originality_body["evidence_ids"] + authorization_evidence,
                        "confidence": .98, "reason": "保留原文授权声明",
                    },
                ],
                "clause_reviews": [
                    {"clause_id": originality_heading["id"], "classification": "covered",
                     "requirement_indexes": [0], "reason": "由声明需求0覆盖"},
                    {"clause_id": originality_body["id"], "classification": "covered",
                     "requirement_indexes": [1], "reason": "由声明需求1覆盖"},
                    {"clause_id": authorization_heading["id"], "classification": "covered",
                     "requirement_indexes": [1], "reason": "由声明需求1覆盖"},
                    {"clause_id": authorization_body["id"], "classification": "covered",
                     "requirement_indexes": [1], "reason": "由声明需求1覆盖"},
                ],
                "unsupported_items": [], "reported_conflicts": [],
            }, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                             "--analysis-mode", "llm_primary", "--llm-response", str(response))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            spec = json.loads((out / "format-spec.json").read_text())
            items = {item["id"]: item for item in spec["declarations"]["items"]}
            self.assertEqual(set(items), {"originality", "authorization"})
            self.assertEqual(items["originality"]["version"], f"run-{spec['run_id']}")
            registry_item = spec["resource_registry"]["items"][items["originality"]["resource_id"]]
            self.assertEqual(registry_item["heading"], "独创性声明")
            self.assertEqual(registry_item["body_parts"], ["本人声明本论文为本人在导师指导下完成的研究成果。"])
            validation = json.loads((out / "schema-validation.json").read_text())
            self.assertTrue(validation["valid"], validation["errors"])

    def test_registered_pending_semantic_role_is_accepted_for_capability_preflight(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; preview = td / "preview"
            response = td / "response.json"; out = td / "out"
            d = Document(); d.add_paragraph("作者字段标签使用五号字。"); d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview),
                "--analysis-mode", "llm_primary")
            clause = json.loads((preview / "requirement-clauses.json").read_text())[0]
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [{
                    "role": "cover_field_label", "properties": {"font": {"size_pt": 10.5}},
                    "clause_ids": [clause["id"]], "evidence_ids": clause["evidence_ids"],
                    "confidence": .95, "reason": "已注册的封面字段标签语义要求"
                }],
                "clause_reviews": [{
                    "clause_id": clause["id"], "classification": "covered",
                    "requirement_indexes": [0], "reason": "由需求0覆盖"
                }],
                "unsupported_items": [], "reported_conflicts": []
            }, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                             "--analysis-mode", "llm_primary", "--llm-response", str(response))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "semantic_resolved")
            self.assertEqual(spec["requirements"][0]["role"], "cover_field_label")

    def test_reviewed_nonexecutable_outcomes_do_not_become_open_questions(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; preview = td / "preview"
            response = td / "response.json"; out = td / "out"
            d = Document()
            d.add_paragraph("学号按学生提供的学籍编号填写。")
            d.add_paragraph("论文题目应准确概括研究内容。")
            d.add_paragraph("论文须使用学校规定颜色的实体封皮装订。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview),
                "--analysis-mode", "llm_primary")
            clauses = json.loads((preview / "requirement-clauses.json").read_text())
            classifications = ["requires_metadata", "unverifiable", "external_compliance"]
            response.write_text(json.dumps({
                "contract_version": "2.1", "requirements": [],
                "clause_reviews": [{
                    "clause_id": clause["id"], "classification": classification,
                    "requirement_indexes": [], "reason": reason,
                } for clause, classification, reason in zip(clauses, classifications, [
                    "需要可信学籍元数据，不能从文档猜测。",
                    "准确概括研究内容属于学术判断。",
                    "实体封皮颜色和装订需要外部实物核验。",
                ])],
                "unsupported_items": [], "reported_conflicts": [],
            }, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                             "--analysis-mode", "llm_primary", "--llm-response", str(response))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "semantic_resolved")
            self.assertEqual(spec["completeness"]["unresolved_clause_ids"], [])
            self.assertEqual(json.loads((out / "questions.json").read_text()), [])
            states = {record["clause_id"]: record["status"] for record in spec["clause_compliance"]}
            self.assertEqual(list(states.values()), classifications)

    def test_unsupported_clause_is_preserved_without_blocking_supported_requirements(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            preview = td / "preview"; response = td / "response.json"; work = td / "work"
            d = Document()
            d.add_paragraph("正文中文使用小四号宋体，英文和数字使用 Times New Roman，行距固定值20磅。")
            d.add_paragraph("论文须使用学校专用防伪纸张印刷。")
            d.save(req); make_target(target)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview), "--analysis-mode", "llm_primary")
            clauses = json.loads((preview / "requirement-clauses.json").read_text())
            supported, unsupported = clauses
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [{"role": "body_text", "properties": {
                    "font": {"cjk": "SimSun", "latin": "Times New Roman", "size_pt": 12},
                    "paragraph": {"line_spacing": {"type": "exact", "value": 20, "unit": "pt"}}},
                    "clause_ids": [supported["id"]], "evidence_ids": supported["evidence_ids"],
                    "confidence": 0.99, "reason": "可执行正文格式"}],
                "clause_reviews": [
                    {"clause_id": supported["id"], "classification": "covered",
                     "requirement_indexes": [0], "reason": "由需求0覆盖"},
                    {"clause_id": unsupported["id"], "classification": "unsupported",
                     "requirement_indexes": [], "reason": "当前 DOCX 格式后端无法选择实体防伪纸张"}],
                "unsupported_items": ["需要学校专用防伪纸张"], "reported_conflicts": []
            }, ensure_ascii=False), encoding="utf-8")

            result = run_raw("scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "formatted.docx"),
                             "--work-dir", str(work), "--analysis-mode", "llm_primary",
                             "--llm-response", str(response), "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            spec = json.loads((work / "requirements" / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "semantic_resolved")
            self.assertEqual(spec["completeness"]["unsupported_items"],
                             ["需要学校专用防伪纸张", unsupported["id"]])
            manifest = json.loads((work / "pipeline-manifest.json").read_text())
            self.assertEqual(manifest["status"], "completed")
            self.assertEqual(manifest["unsupported_items"], spec["completeness"]["unsupported_items"])
            report = json.loads((work / "application" / "validation-report.json").read_text())
            self.assertEqual(report["unsupported_items"], spec["completeness"]["unsupported_items"])
            self.assertTrue((td / "formatted.docx").exists())

    def test_unsupported_items_block_full_format_spec_but_not_supported_subset(self) -> None:
        spec = {
            "status": "semantic_resolved",
            "completeness": {"unsupported_items": ["实体防伪纸张"]},
            "clause_compliance": [],
        }
        self.assertIn("unsupported_items", apply_format_spec.format_spec_blockers(spec, "full"))
        self.assertNotIn(
            "unsupported_items",
            apply_format_spec.format_spec_blockers(spec, "supported_subset"),
        )

    def test_llm_cannot_claim_physical_cover_or_actual_signature_as_docx_execution(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; preview = td / "preview"; response = td / "response.json"
            d = Document()
            d.add_paragraph("论文封面底色必须为学校规定颜色。")
            d.add_paragraph("研究生须在独创性声明页亲笔签名。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview), "--analysis-mode", "llm_primary")
            clauses = json.loads((preview / "requirement-clauses.json").read_text())
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [
                    {"role": "cover", "properties": {"institution": "中国农业大学", "fields": [{
                        "id": "title_zh", "label": "中文题目",
                        "value_from": "thesis_profile.cover_metadata.title_zh", "display_policy": "required", "order": 1}]},
                     "clause_ids": [clauses[0]["id"]], "evidence_ids": clauses[0]["evidence_ids"], "confidence": .9},
                    {"role": "declarations", "properties": {"before_role": "abstract_title_zh", "items": []},
                     "clause_ids": [clauses[1]["id"]], "evidence_ids": clauses[1]["evidence_ids"], "confidence": .9},
                ],
                "clause_reviews": [
                    {"clause_id": clauses[0]["id"], "classification": "executable", "requirement_indexes": [0], "reason": "设置页面底色"},
                    {"clause_id": clauses[1]["id"], "classification": "verify_existing", "requirement_indexes": [1], "reason": "发现签名栏"},
                ],
                "unsupported_items": [], "reported_conflicts": []
            }, ensure_ascii=False), encoding="utf-8")
            out = td / "out"
            result = run_raw("scripts/requirements_engine.py", str(req), "--out", str(out),
                             "--analysis-mode", "llm_primary", "--llm-response", str(response))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "needs_clarification")
            conflicts = json.dumps(json.loads((out / "conflicts.json").read_text()), ensure_ascii=False)
            self.assertIn("external_artifact_cannot_be_docx_executable_or_verified", conflicts)

    def test_margin_purpose_phrase_is_not_external_artifact_guard(self) -> None:
        clauses = [
            {"id": "C1", "text": "纸型为A4，页边距上、左为2.5厘米。", "evidence_ids": ["E1"]},
            {"id": "C2", "text": "下、右为2厘米，以便装订。", "evidence_ids": ["E1"]},
        ]
        response = {
            "contract_version": "2.1",
            "requirements": [{
                "role": "page",
                "properties": {
                    "size": "A4",
                    "margins_pt": {"top": 70.87, "left": 70.87, "bottom": 56.69, "right": 56.69},
                },
                "clause_ids": ["C1", "C2"], "evidence_ids": ["E1"],
                "confidence": 0.99, "reason": "页型和页边距可由 DOCX 静态属性执行并核验。",
                "verification": {"mode": "static_docx", "checks": ["核验纸型和四边页边距。"]},
            }],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable", "requirement_indexes": [0], "reason": "页型和上左边距。"},
                {"clause_id": "C2", "classification": "executable", "requirement_indexes": [0], "reason": "下右边距及装订目的。"},
            ],
            "unsupported_items": [], "reported_conflicts": [],
        }
        self.assertFalse(requirements_engine._requires_external_artifact_verification(clauses[1]["text"]))
        spec, conflicts, audit = requirements_engine.merge_llm_primary(
            Path("synthetic-source"),
            {"schema_version": "1.0", "roles": {}, "page": {}, "requirements": [], "content_instances": []},
            clauses, response, {"E1"},
        )
        hard_types = {"llm_contract", "llm_internal_conflict", "completeness"}
        self.assertFalse(any(item.get("type") in hard_types for item in conflicts), conflicts)
        self.assertEqual(spec["status"], "semantic_resolved")
        self.assertEqual(
            [item.get("accepted") for item in audit if "accepted" in item],
            [True],
        )
        self.assertEqual(spec["requirements"][0]["clause_ids"], ["C1", "C2"])
        records = {item["clause_id"]: item for item in spec["clause_compliance"]}
        self.assertEqual(records["C1"]["status"], "pending_execution")
        self.assertEqual(records["C2"]["status"], "pending_execution")

    def test_executable_requirement_may_cite_external_companion_clause(self) -> None:
        clauses = [
            {"id": "C1", "text": "独创性声明页应包含固定正文。", "evidence_ids": ["E1"]},
            {"id": "C2", "text": "研究生须在独创性声明页亲笔签名。", "evidence_ids": ["E2"]},
        ]
        response = {
            "contract_version": "2.1",
            "requirements": [{
                "role": "body_text",
                "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                "clause_ids": ["C1", "C2"], "evidence_ids": ["E1", "E2"],
                "confidence": 0.99, "reason": "固定正文在DOCX中保留，签名作为空白占位。",
            }],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable",
                 "requirement_indexes": [0], "reason": "固定正文属于DOCX结构。"},
                {"clause_id": "C2", "classification": "external_compliance",
                 "requirement_indexes": [], "reason": "实际签名不能由DOCX证明。"},
            ],
            "unsupported_items": [], "reported_conflicts": [],
        }
        spec, conflicts, audit = requirements_engine.merge_llm_primary(
            Path("synthetic-source"),
            {"schema_version": "1.0", "roles": {}, "page": {},
             "requirements": [], "content_instances": []},
            clauses, response, {"E1", "E2"},
        )
        hard_types = {"llm_contract", "llm_internal_conflict", "completeness"}
        self.assertFalse(any(item.get("type") in hard_types for item in conflicts), conflicts)
        self.assertEqual(spec["status"], "semantic_resolved")
        records = {item["clause_id"]: item for item in spec["clause_compliance"]}
        self.assertEqual(records["C1"]["status"], "pending_execution")
        self.assertEqual(records["C2"]["status"], "external_compliance")
        self.assertEqual(requirements_engine.validate_spec(spec, {"E1", "E2"}), [])

    def test_llm_primary_preserves_clause_scoped_top_level_requirements(self) -> None:
        clauses = [
            {"id": "C1", "text": "博士论文中文关键词应为5至8个。", "evidence_ids": ["E1"]},
            {"id": "C2", "text": "目录应按封面、正文顺序列出。", "evidence_ids": ["E2"]},
        ]
        response = {
            "contract_version": "2.1",
            "requirements": [
                {
                    "role": "content_constraints",
                    "properties": {"keywords_zh": {"required": True, "min_count": 5, "max_count": 8}},
                    "clause_ids": ["C1"], "evidence_ids": ["E1"], "confidence": 0.98,
                    "reason": "博士论文中文关键词数量取源文规定的5至8个。",
                    "applicability": {"status": "conditional", "conditions": [{
                        "fact": "thesis_profile.degree_level", "operator": "equals", "value": "doctor",
                    }]},
                    "input_prerequisites": [{
                        "kind": "metadata", "key": "thesis_profile.degree_level", "required": True,
                        "reason": "关键词数量依赖学位层次。",
                    }],
                    "verification": {"mode": "manual", "checks": ["按学位层次核验中文关键词数量。"]},
                },
                {
                    "role": "document_structure",
                    "properties": {
                        "required_roles": ["cover", "body_text"],
                        "ordered_roles": ["cover", "body_text"],
                    },
                    "clause_ids": ["C2"], "evidence_ids": ["E2"], "confidence": 0.97,
                    "reason": "目录顺序由已注册的封面和正文角色表达。",
                    "verification": {"mode": "static_docx", "checks": ["核验目录中封面和正文的顺序。"]},
                },
            ],
            "clause_reviews": [
                {"clause_id": "C1", "classification": "executable", "requirement_indexes": [0], "reason": "有明确关键词数量约束。"},
                {"clause_id": "C2", "classification": "executable", "requirement_indexes": [1], "reason": "有明确目录顺序约束。"},
            ],
            "unsupported_items": [], "reported_conflicts": [],
        }
        rule_spec = {
            "schema_version": "1.0", "roles": {}, "page": {}, "requirements": [], "content_instances": [],
            "content_constraints": {"keywords_zh": {"required": True, "min_count": 3, "max_count": 5}},
            "document_structure": {"required_roles": ["baseline"], "ordered_roles": ["baseline"]},
        }
        spec, conflicts, audit = requirements_engine.merge_llm_primary(
            Path("synthetic-source"), rule_spec, clauses, response, {"E1", "E2"},
        )
        hard_types = {"llm_contract", "llm_internal_conflict", "completeness"}
        self.assertFalse(any(item.get("type") in hard_types for item in conflicts), conflicts)
        self.assertEqual(sum(1 for item in audit if item.get("accepted")), 2)
        self.assertEqual(spec["status"], "semantic_resolved")
        self.assertEqual(spec["content_constraints"]["keywords_zh"],
                         {"required": True, "min_count": 5, "max_count": 8})
        self.assertEqual(spec["document_structure"]["required_roles"], ["cover", "body_text"])
        self.assertTrue(any(item.get("type") == "rule_llm_conflict" for item in conflicts))
        by_clause = {
            clause_id: next(req for req in spec["requirements"] if clause_id in req["clause_ids"])
            for clause_id in ("C1", "C2")
        }
        self.assertEqual(by_clause["C1"]["properties"]["keywords_zh"]["min_count"], 5)
        self.assertEqual(by_clause["C1"]["applicability"]["status"], "conditional")
        self.assertEqual(by_clause["C1"]["input_prerequisites"][0]["key"], "thesis_profile.degree_level")
        self.assertEqual(by_clause["C2"]["properties"]["ordered_roles"], ["cover", "body_text"])

    def test_llm_primary_unions_scoped_cover_object_and_order_fragments(self) -> None:
        clauses = [
            {"id": "C1", "text": "封面应包含中文题目。", "evidence_ids": ["E1"]},
            {"id": "C2", "text": "封面应包含英文题目。", "evidence_ids": ["E2"]},
            {"id": "C3", "text": "图形前应有正文说明。", "evidence_ids": ["E3"]},
            {"id": "C4", "text": "表格前应有正文说明。", "evidence_ids": ["E4"]},
            {"id": "C5", "text": "正文后接附录。", "evidence_ids": ["E5"]},
            {"id": "C6", "text": "参考文献和致谢单独成组。", "evidence_ids": ["E6"]},
        ]
        field = lambda field_id, label, order: {
            "id": field_id,
            "label": label,
            "value_from": f"thesis_profile.cover_metadata.{field_id}",
            "display_policy": "required",
            "order": order,
        }
        requirements = [
            {"role": "cover", "properties": {"institution": "——", "fields": [field("title_zh", "论文题目：", 1)]},
             "clause_ids": ["C1"], "evidence_ids": ["E1"], "confidence": .99, "reason": "中文题目是封面字段。"},
            {"role": "cover", "properties": {"institution": "测试大学", "fields": [field("title_en", "English title", 2)]},
             "clause_ids": ["C2"], "evidence_ids": ["E2"], "confidence": .99, "reason": "英文题目是封面字段。"},
            {"role": "objects", "properties": {"order_constraints": [{"object_type": "figure", "preceding_prose_required": True}]},
             "clause_ids": ["C3"], "evidence_ids": ["E3"], "confidence": .99, "reason": "图形需要前置正文。"},
            {"role": "objects", "properties": {"order_constraints": [{"object_type": "table", "preceding_prose_required": True}]},
             "clause_ids": ["C4"], "evidence_ids": ["E4"], "confidence": .99, "reason": "表格需要前置正文。"},
            {"role": "document_structure", "properties": {"ordered_roles": ["body_text", "appendices"]},
             "clause_ids": ["C5"], "evidence_ids": ["E5"], "confidence": .99, "reason": "正文和附录保持顺序。"},
            {"role": "document_structure", "properties": {"ordered_roles": ["heading_references", "heading_acknowledgments"]},
             "clause_ids": ["C6"], "evidence_ids": ["E6"], "confidence": .99, "reason": "后置部分保留独立顺序。"},
        ]
        response = {
            "contract_version": "2.1", "requirements": requirements,
            "clause_reviews": [
                {"clause_id": clause["id"], "classification": "executable",
                 "requirement_indexes": [index], "reason": "有明确的结构性约束。"}
                for index, clause in enumerate(clauses)
            ],
            "unsupported_items": [], "reported_conflicts": [],
        }
        spec, conflicts, audit = requirements_engine.merge_llm_primary(
            Path("synthetic-source"),
            {"schema_version": "1.0", "roles": {}, "page": {}, "requirements": [], "content_instances": []},
            clauses, response, {f"E{index}" for index in range(1, 7)},
        )
        self.assertEqual(sum(1 for item in audit if item.get("accepted")), 6)
        self.assertFalse(any(item.get("type") in {"llm_contract", "llm_internal_conflict", "completeness"}
                             for item in conflicts), conflicts)
        self.assertEqual({item["id"] for item in spec["cover"]["fields"]}, {"title_zh", "title_en"})
        self.assertEqual(spec["cover"]["institution"], "测试大学")
        self.assertEqual({item["object_type"] for item in spec["objects"]["order_constraints"]}, {"figure", "table"})
        self.assertEqual(spec["document_structure"]["ordered_roles"], ["body_text", "appendices"])
        self.assertEqual(spec["document_structure"]["ordered_role_groups"],
                         [["heading_references", "heading_acknowledgments"]])

    def test_full_compliance_blocks_applicable_backend_gap(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            preview = td / "preview"; response = td / "response.json"; work = td / "work"
            d = Document()
            d.add_paragraph("正文中文使用小四号宋体。")
            d.add_paragraph("图形不能跨页显示。")
            d.save(req); make_target(target)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview), "--analysis-mode", "llm_primary")
            supported, unsupported = json.loads((preview / "requirement-clauses.json").read_text())
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [{"role": "body_text", "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                    "clause_ids": [supported["id"]], "evidence_ids": supported["evidence_ids"],
                    "confidence": .99, "reason": "可执行正文格式"}],
                "clause_reviews": [
                    {"clause_id": supported["id"], "classification": "executable", "requirement_indexes": [0], "reason": "自动应用并验证"},
                    {"clause_id": unsupported["id"], "classification": "unsupported_backend", "requirement_indexes": [], "reason": "当前后端没有分页几何验证器"}],
                "unsupported_items": [], "reported_conflicts": []
            }, ensure_ascii=False), encoding="utf-8")
            preflight = run_raw(
                "scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "preflight.docx"),
                "--work-dir", str(td / "preflight-work"), "--llm-response", str(response),
                "--analysis-mode", "llm_primary", "--compliance-mode", "supported_subset",
            )
            self.assertEqual(preflight.returncode, 0, preflight.stderr + preflight.stdout)
            bound = json.loads(response.read_text(encoding="utf-8"))
            bound["provenance"] = json.loads(
                (td / "preflight-work" / "requirements" / "llm-request.json").read_text(encoding="utf-8")
            )["provenance"]
            response.write_text(json.dumps(bound, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "output.docx"),
                             "--work-dir", str(work), "--llm-response", str(response),
                             "--run-id", bound["provenance"]["run_id"],
                             "--analysis-mode", "llm_primary", "--compliance-mode", "full",
                             "--allow-offline-review")
            self.assertEqual(result.returncode, 3, result.stderr + result.stdout)
            manifest = json.loads((work / "pipeline-manifest.json").read_text())
            self.assertIn("full_compliance_analysis_failed", manifest["blocking_reasons"])
            self.assertFalse((td / "output.docx").exists())

    def test_full_compliance_emits_clause_and_external_reports(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            preview = td / "preview"; response = td / "response.json"; work = td / "work"
            d = Document()
            d.add_paragraph("正文中文使用小四号宋体。")
            d.add_paragraph("论文统一双面打印装订成册。")
            d.save(req); make_target(target)
            run("scripts/requirements_engine.py", str(req), "--out", str(preview), "--analysis-mode", "llm_primary")
            supported, external = json.loads((preview / "requirement-clauses.json").read_text())
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "requirements": [{"role": "body_text", "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                    "clause_ids": [supported["id"]], "evidence_ids": supported["evidence_ids"],
                    "confidence": .99, "reason": "可执行正文格式"}],
                "clause_reviews": [
                    {"clause_id": supported["id"], "classification": "executable", "requirement_indexes": [0], "reason": "自动应用并验证"},
                    {"clause_id": external["id"], "classification": "external_compliance", "requirement_indexes": [], "reason": "真实打印与装订不属于 DOCX 文件属性"}],
                "unsupported_items": [], "reported_conflicts": []
            }, ensure_ascii=False), encoding="utf-8")
            # Formal full mode now requires the response to be bound to the
            # exact request produced with this target DOCX.  Generate that
            # request through a supported-subset preflight, then carry only
            # its provenance block into the offline fixture.
            preflight = run_raw(
                "scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "preflight.docx"),
                "--work-dir", str(td / "preflight-work"), "--llm-response", str(response),
                "--analysis-mode", "llm_primary", "--compliance-mode", "supported_subset",
            )
            self.assertEqual(preflight.returncode, 0, preflight.stderr + preflight.stdout)
            bound = json.loads(response.read_text(encoding="utf-8"))
            bound["provenance"] = json.loads(
                (td / "preflight-work" / "requirements" / "llm-request.json").read_text(encoding="utf-8")
            )["provenance"]
            response.write_text(json.dumps(bound, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/thesis_format_pipeline.py", str(req), str(target), str(td / "output.docx"),
                             "--work-dir", str(work), "--llm-response", str(response),
                             "--run-id", bound["provenance"]["run_id"],
                             "--analysis-mode", "llm_primary", "--compliance-mode", "full",
                             "--allow-offline-review")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((work / "application" / "validation-report.json").read_text())
            self.assertTrue(report["docx_fully_compliant"])
            self.assertEqual(report["overall_status"], "passed")
            clauses = json.loads((work / "application" / "clause-compliance-report.json").read_text())
            self.assertEqual(clauses["docx_compliance"]["counts"]["generated_and_verified"], 1)
            checklist = json.loads((work / "application" / "external-compliance-checklist.json").read_text())
            self.assertEqual([item["clause_id"] for item in checklist], [external["id"]])

    def test_three_line_table_pagination_rules_are_applied_and_audited(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document(); doc.add_paragraph("表1 测试表")
            table = doc.add_table(rows=3, cols=2)
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells): cell.text = f"{ri}-{ci}"
            doc.save(source)
            spec = {"schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                    "analysis_mode": "known_template", "roles": {}, "requirements": [],
                    "tables": {"style": "three_line", "top_border_pt": 1.5,
                               "header_border_pt": .75, "bottom_border_pt": 1.5,
                               "remove_vertical_borders": True, "repeat_header_row": True,
                               "allow_row_split": False}}
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((audit / "validation-report.json").read_text())
            self.assertTrue(report["valid"], report["findings"])
            self.assertEqual(report["table_changes"]["rows_no_split"], 3)
            with zipfile.ZipFile(output) as zf:
                xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn('w:sz="12"', xml)
            self.assertIn('w:sz="6"', xml)
            self.assertEqual(xml.count("<w:cantSplit"), 3)
            self.assertEqual(xml.count("<w:tblHeader"), 1)
            self.assertIn('<w:insideV w:val="nil"', xml)

    def test_explicit_table_border_map_is_applied_without_guessing(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document(); table = doc.add_table(rows=2, cols=2)
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells): cell.text = f"{ri}-{ci}"
            doc.save(source)
            spec = {"schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                    "analysis_mode": "known_template", "roles": {}, "requirements": [],
                    "tables": {"border_widths_pt": {
                        "top": 1.0, "header": .5, "bottom": 2.0,
                        "left": 0, "right": 0, "inside_h": .25, "inside_v": 0,
                    }}}
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((audit / "validation-report.json").read_text())
            self.assertTrue(report["valid"], report["findings"])
            self.assertEqual(report["table_changes"]["explicit_borders_applied"], 8)
            with zipfile.ZipFile(output) as zf:
                xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn('<w:top w:val="single" w:sz="8"', xml)
            self.assertIn('<w:bottom w:val="single" w:sz="16"', xml)
            self.assertIn('<w:insideH w:val="single" w:sz="2"', xml)
            self.assertIn('<w:insideV w:val="nil"', xml)

    def test_object_order_contract_requires_unambiguous_preceding_prose(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document()
            doc.add_paragraph().add_run().add_picture(str(ROOT / "tests" / "assets" / "frontmatter-contact.jpg"), width=Inches(.2))
            doc.add_paragraph("图1-1 流程图")
            doc.save(source)
            spec = {"schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                    "analysis_mode": "known_template", "roles": {}, "requirements": [],
                    "objects": {"order_constraints": [{"object_type": "figure",
                                                          "preceding_prose_required": True}]}}
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0)
            findings = json.loads((audit / "validation-report.json").read_text())["findings"]
            self.assertTrue(any(item["property"].endswith("preceding_prose") for item in findings))

            good = Document(); good.add_paragraph("图示展示了研究流程。")
            good.add_paragraph().add_run().add_picture(str(ROOT / "tests" / "assets" / "frontmatter-contact.jpg"), width=Inches(.2))
            good.save(source)
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(td / "good.docx"),
                             "--out-dir", str(td / "good-audit"), "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)

    def test_continued_table_requires_render_evidence_and_repeats_header(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document(); table = doc.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "表头"; table.cell(1, 0).text = "数据"
            doc.save(source)
            spec = {"schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                    "analysis_mode": "known_template", "roles": {}, "requirements": [],
                    "tables": {"continuation": {"caption_suffix": "（续）",
                                                   "repeat_header_row": True,
                                                   "caption_required_on_continuation": True,
                                                   "verification": "pdf_render"}}}
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0)
            report = json.loads((audit / "validation-report.json").read_text())
            self.assertTrue(any(item["property"] == "tables.continuation.render_verification"
                                for item in report["findings"]))
            self.assertEqual(report["table_changes"]["header_rows_repeated"], 1)

    def test_acknowledgments_section_length_is_audited(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document(); doc.add_paragraph("致谢", "Heading 1"); doc.add_paragraph("感谢导师和同学")
            doc.add_paragraph("参考文献", "Heading 1"); doc.save(source)
            spec = {"schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                    "analysis_mode": "known_template", "roles": {}, "requirements": [],
                    "content_constraints": {"acknowledgments": {"max_chars": 5}}}
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0)
            findings = json.loads((audit / "validation-report.json").read_text())["findings"]
            self.assertTrue(any(item["property"] == "acknowledgments.max_chars" for item in findings))

    def test_abstract_keyword_and_document_structure_constraints(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document()
            for name in ("AbstractBodyCN", "KeywordsLineCN", "AbstractBodyEN", "KeywordsLineEN", "TOCHeading"):
                doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("摘 要", doc.styles.add_style("AbstractTitleCN", WD_STYLE_TYPE.PARAGRAPH))
            doc.add_paragraph("测试摘要内容。", "AbstractBodyCN")
            doc.add_paragraph("关键词：测试，验证，文档", "KeywordsLineCN")
            doc.add_paragraph("Abstract", doc.styles.add_style("AbstractTitleEN", WD_STYLE_TYPE.PARAGRAPH))
            doc.add_paragraph("Test abstract content.", "AbstractBodyEN")
            doc.add_paragraph("Keywords: test, validation, document", "KeywordsLineEN")
            doc.add_paragraph("目 录", "TOCHeading")
            doc.add_paragraph("第一章", "Heading 1")
            doc.add_paragraph("第一节", "Heading 2")
            doc.add_paragraph("第一小节", "Heading 3")
            doc.save(source)
            spec = {"schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                    "analysis_mode": "known_template", "roles": {}, "requirements": [],
                    "content_constraints": {
                        "abstract_zh": {"required": True, "max_chars": 800},
                        "abstract_en": {"required": True},
                        "keywords_zh": {"required": True, "min_count": 3, "max_count": 5,
                                        "separator": "chinese_comma", "require_after_role": "abstract_body_zh"},
                        "keywords_en": {"required": True, "min_count": 3, "max_count": 5,
                                        "separator": "english_comma", "require_after_role": "abstract_body_en",
                                        "match_other_language_count": True}},
                    "document_structure": {"required_roles": ["abstract_title_zh", "abstract_body_zh", "keywords_zh",
                        "abstract_title_en", "abstract_body_en", "keywords_en", "toc"],
                        "ordered_roles": ["abstract_title_zh", "abstract_body_zh", "keywords_zh",
                            "abstract_title_en", "abstract_body_en", "keywords_en", "toc"],
                        "max_heading_depth": 3}}
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertTrue(json.loads((audit / "validation-report.json").read_text())["valid"])

            bad = Document(source)
            next(p for p in bad.paragraphs if p.text.startswith("关键词")).text = "关键词：测试,验证"
            bad_path = td / "bad.docx"; bad.save(bad_path)
            result = run_raw("scripts/apply_format_spec.py", str(bad_path), str(spec_path), str(td / "bad-out.docx"),
                             "--out-dir", str(td / "bad-audit"), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0)
            findings = json.loads((td / "bad-audit" / "validation-report.json").read_text())["findings"]
            self.assertTrue(any(x["property"] == "keywords_zh.min_count" for x in findings))
            self.assertTrue(any(x["property"] == "keywords_zh.separator" for x in findings))

    def test_thesis_profile_resolves_conditional_abstract_limit(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document(); doc.styles.add_style("AbstractBodyCN", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("测试摘要内容", "AbstractBodyCN"); doc.save(source)
            spec = {"schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                    "analysis_mode": "known_template", "roles": {}, "requirements": [],
                    "thesis_profile": {"schema_version": "1.0", "degree_level": "master", "writing_language": "zh"},
                    "conditional_constraints": {"abstract_zh_max_chars_by_degree": {"master": 4, "doctor": 1500}}}
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0)
            findings = json.loads((audit / "validation-report.json").read_text())["findings"]
            self.assertTrue(any(x["property"] == "abstract_zh.max_chars" and x["required_value"] == 4 for x in findings))

    def test_toc_depth_is_rewritten_and_verified(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document(); p = doc.add_paragraph("目 录")
            run = p.add_run(); begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); run._r.append(begin)
            instr = OxmlElement("w:instrText"); instr.text = ' TOC \\o "1-3" \\h \\z \\u '; run._r.append(instr)
            end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); run._r.append(end); doc.save(source)
            spec = {"schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                    "analysis_mode": "known_template", "roles": {}, "requirements": [],
                    "document_structure": {"toc_depth": 2}}
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((audit / "validation-report.json").read_text())
            self.assertEqual(report["toc_fields_updated"], 1)
            with zipfile.ZipFile(output) as zf: xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn('TOC \\o "1-2"', xml)

    def test_toc_field_is_content_role_even_without_cached_entries(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document()
            doc.add_paragraph("目 录", "TOC Heading")
            field = doc.add_paragraph()
            begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
            field.add_run()._r.append(begin)
            instr = OxmlElement("w:instrText"); instr.text = ' TOC \\o "1-3" \\h \\z \\u '
            field.add_run()._r.append(instr)
            end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
            field.add_run()._r.append(end)
            doc.save(source)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "test", "status": "rule_resolved",
                "roles": {"toc": {"font": {"size_pt": 12}}}, "requirements": [],
            }), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--require-coverage", "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((audit / "validation-report.json").read_text(encoding="utf-8"))
            coverage = {item["role"]: item for item in report["role_coverage"]}
            self.assertEqual(coverage["toc"]["status"], "present")
            self.assertEqual(coverage["toc"]["count"], 1)
            self.assertEqual(Document(output).paragraphs[1].style.name, "TOC Heading")

    def test_content_instance_style_override_is_applied_without_synthesizing_text(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; audit = td / "audit"
            doc = Document(); doc.add_paragraph("分类号"); doc.save(source)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "test",
                "status": "semantic_resolved", "analysis_mode": "llm_primary",
                "roles": {"cover_field_label": {"font": {"size_pt": 10.5}}},
                "content_instances": [{
                    "id": "CFI-test", "field_key": "cover_field_label:分类号",
                    "role": "cover_field_label", "text": "分类号", "order": 1,
                    "clause_ids": ["C1"], "evidence_ids": ["E1"],
                    "style_properties": {"font": {"size_pt": 16}},
                    "reason": "instance override",
                }],
                "requirements": [{
                    "id": "R00001", "role": "cover_field_label",
                    "properties": {"text": "分类号"},
                    "field_instance_ids": ["CFI-test"], "evidence_ids": ["E1"],
                    "clause_ids": ["C1"], "resolved_by": "llm", "confidence": .98,
                }],
            }, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            rendered = Document(output)
            self.assertEqual(len(rendered.paragraphs), 1)
            self.assertAlmostEqual(rendered.paragraphs[0].runs[0].font.size.pt, 16, places=2)
            instance_audit = json.loads((audit / "content-instance-audit.json").read_text())
            self.assertEqual(instance_audit["counts"]["applied"], 1)
            self.assertEqual(instance_audit["instances"][0]["status"], "applied")


    def test_legacy_unsupported_status_applies_when_no_concrete_blocker_exists(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; output = td / "output.docx"
            spec_path = td / "spec.json"; make_target(target)
            spec_path.write_text(json.dumps({
                "schema_version": "1.0", "source_document": "legacy", "status": "unsupported",
                "roles": {"body_text": {"font": {"cjk": "SimSun", "size_pt": 12}}},
                "requirements": [], "completeness": {"reviewed_by": "llm", "covered_clause_ids": [],
                    "ignored_clause_ids": [], "unresolved_clause_ids": [],
                    "unsupported_items": ["legacy non-executable clause"]}
            }), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(target), str(spec_path), str(output),
                             "--out-dir", str(td / "audit"))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertTrue(output.exists())
    def test_combined_figure_table_title_merge_is_atomic(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; out = td / "out"
            d = Document()
            d.add_paragraph("表题使用五号宋体。")
            d.add_paragraph("图表题使用小四号黑体。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(out))
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "needs_clarification")
            self.assertEqual(spec["roles"]["table_caption"]["font"], {"cjk": "SimSun", "size_pt": 10.5})
            self.assertNotIn("figure_caption", spec["roles"], "a failed group merge must not leave a half-merge")
            conflicts = json.loads((out / "conflicts.json").read_text())
            self.assertTrue(any(item["role"] == "table_caption" for item in conflicts))

    def test_heading_spacing_lines_are_strict_and_converted_to_points(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; target = td / "target.docx"
            spec_dir = td / "spec"; audit = td / "audit"; output = td / "output.docx"
            d = Document(); d.add_paragraph("一级标题使用小三号黑体，段前一行，段后半行。"); d.save(req)
            t = Document(); p = t.add_paragraph("绪论"); p.style = "Heading 1"; t.save(target)
            run("scripts/requirements_engine.py", str(req), "--out", str(spec_dir), "--analysis-mode", "rule_only")
            spec = json.loads((spec_dir / "format-spec.json").read_text())
            paragraph = spec["roles"]["heading_1"]["paragraph"]
            self.assertEqual(paragraph["space_before_lines"], 1)
            self.assertEqual(paragraph["space_after_lines"], .5)
            self.assertEqual(paragraph["spacing_line_height_pt"], 15)
            run("scripts/apply_format_spec.py", str(target), str(spec_dir / "format-spec.json"), str(output),
                "--out-dir", str(audit), "--require-coverage")
            check = Document(output)
            self.assertAlmostEqual(check.styles["Heading 1"].paragraph_format.space_before.pt, 15, places=2)
            self.assertAlmostEqual(check.styles["Heading 1"].paragraph_format.space_after.pt, 7.5, places=2)

    def test_line_spacing_height_is_completed_after_cross_clause_role_merge(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); req = td / "requirements.docx"; out = td / "out"
            d = Document()
            d.add_paragraph("目录段前、段后各空0行。")
            d.add_paragraph("目录使用小四号宋体。")
            d.save(req)
            run("scripts/requirements_engine.py", str(req), "--out", str(out), "--analysis-mode", "rule_only")
            validation = json.loads((out / "schema-validation.json").read_text())
            self.assertTrue(validation["valid"], validation["errors"])
            paragraph = json.loads((out / "format-spec.json").read_text())["roles"]["toc"]["paragraph"]
            self.assertEqual(paragraph["space_before_lines"], 0)
            self.assertEqual(paragraph["space_after_lines"], 0)
            self.assertEqual(paragraph["spacing_line_height_pt"], 12)

    def _fixed_text_cover_declaration_spec(self) -> dict:
        resources = {
            "originality": {
                "heading": "独创性声明",
                "body": "本人声明本论文为本人在导师指导下完成的研究成果。",
            },
            "authorization": {
                "heading": "关于学位论文使用授权的说明",
                "body": "本人同意按照学校规定保存和使用本学位论文。",
            },
        }
        fields = [
            ("unit_code", "单位代码", "required"), ("title_zh", "中文题目", "required"),
            ("title_en", "英文题目", "required"), ("author_name", "作者", "required"),
            ("student_id", "学号", "required"), ("program_name", "专业", "if_present"),
            ("supervisor_name", "指导教师", "required"), ("co_supervisors", "合作指导教师", "if_present"),
            ("completion_date", "完成日期", "required"), ("classification_number", "分类号", "if_present"),
        ]
        items = []
        placeholders = {
            "originality": [("author", "研究生签名"), ("date", "日期")],
            "authorization": [("author", "研究生签名"), ("supervisor", "指导教师签名"), ("date", "日期")],
        }
        for item_id in ("originality", "authorization"):
            resource = resources[item_id]
            items.append({"id": item_id, "heading": resource["heading"], "body_parts": [resource["body"]],
                          "source_evidence_ids": [f"fixture-{item_id}"], "signature_placeholders": [
                              {"role": role, "label": label, "attestation_scope": "placeholder_presence_only"}
                              for role, label in placeholders[item_id]]})
        spec = {
            "schema_version": "1.0", "run_id": "test-fixed-text-run", "source_document": "generic-template", "status": "semantic_resolved",
            "analysis_mode": "known_template", "compliance_mode": "supported_subset", "roles": {},
            "thesis_profile": {"schema_version": "1.0", "degree_level": "master", "degree_category": "academic",
                "writing_language": "zh", "security_level": "public", "student_id": "S20260001",
                "completion_date": "2026-06", "co_supervisor_count": 1, "cover_metadata": {
                    "trust": {"source": "user_confirmed", "confirmed": True}, "unit_code": "10019",
                    "title_zh": "示例智能系统研究", "title_en": "Research on Example Intelligent Systems",
                    "author_name": "张三", "student_id": "S20260001", "program_name": "农业工程",
                    "supervisor_name": "李四", "co_supervisors": [{"name": "王五", "kind": "academic"}],
                    "completion_date": "2026-06"}},
            "cover": {"institution": "示例研究院", "before_role": "abstract_title_zh", "fields": [
                {"id": field_id, "label": label, "value_from": f"thesis_profile.cover_metadata.{field_id}",
                 "display_policy": policy, "order": i + 1} for i, (field_id, label, policy) in enumerate(fields)]},
            "declarations": {"before_role": "abstract_title_zh", "items": items}, "requirements": []}
        return materialize_declaration_resources(spec, "test-fixed-text-run")

    def test_fixed_text_cover_declarations_round_trip_idempotency_and_signature_boundary(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; first = td / "first.docx"; second = td / "second.docx"
            doc = Document(); doc.styles.add_style("AbstractTitleCN", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("摘 要", "AbstractTitleCN"); doc.add_paragraph("摘要正文"); doc.save(source)
            spec = self._fixed_text_cover_declaration_spec(); spec_path = td / "spec.json"
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            for input_path, output_path, audit in ((source, first, td / "audit-1"), (first, second, td / "audit-2")):
                result = run_raw("scripts/apply_format_spec.py", str(input_path), str(spec_path), str(output_path),
                                 "--out-dir", str(audit), "--compliance-mode", "supported_subset")
                self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
                report = json.loads((audit / "validation-report.json").read_text())
                self.assertTrue(report["valid"], report["findings"])
                self.assertEqual(report["declaration_changes"]["items_written"], 2)
            self.assertEqual(first.read_bytes(), second.read_bytes())
            reopened = Document(second); texts = [p.text for p in reopened.paragraphs]
            self.assertLess(texts.index("独创性声明"), texts.index("摘 要"))
            self.assertEqual(texts.count("示例研究院"), 1)
            self.assertEqual(texts.count("示例智能系统研究"), 1)
            self.assertTrue(any("研究生签名：________________" == text for text in texts))
            self.assertTrue(any("指导教师签名：________________" == text for text in texts))
            self.assertFalse(any(text.strip() in {"张三", "李四"} for text in texts if "签名" in text),
                             "placeholder lines must not prefill names or claim a real signature")

    def test_fixed_text_cover_without_instance_metadata_preserves_structure_with_neutral_placeholders(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "out.docx"; audit = td / "audit"
            doc = Document(); doc.styles.add_style("AbstractTitleCN", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("摘 要", "AbstractTitleCN"); doc.save(source)
            spec = self._fixed_text_cover_declaration_spec()
            spec["cover"]["missing_value_policy"] = "placeholder"
            spec["cover"]["missing_value_placeholder"] = "——"
            del spec["thesis_profile"]["cover_metadata"]
            spec["thesis_profile"].pop("student_id", None)
            spec["thesis_profile"].pop("completion_date", None)
            spec["thesis_profile"].pop("co_supervisor_count", None)
            spec_path = td / "placeholder.json"
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((audit / "validation-report.json").read_text())
            self.assertTrue(report["valid"], report["findings"])
            self.assertFalse(report["submission_ready"])
            self.assertEqual(report["submission_status"], "cover_metadata_pending")
            self.assertEqual(report["cover_metadata"]["status"], "pending")
            self.assertEqual(report["cover_changes"]["trusted_fields_written"], 0)
            self.assertEqual(report["cover_changes"]["placeholder_fields_written"], 7)
            texts = [paragraph.text for paragraph in Document(output).paragraphs]
            self.assertEqual(texts.count("示例研究院"), 1)
            self.assertEqual(texts.count("——"), 2)
            self.assertIn("作者：——", texts)
            self.assertIn("学号：——", texts)
            self.assertIn("指导教师：——", texts)
            self.assertIn("完成日期：——", texts)
            self.assertNotIn("专业：——", texts, "if_present fields stay absent rather than pretending to be supplied")

    def test_cover_starts_document_for_body_fragment_without_front_matter_or_abstract(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "fragment.docx"; output = td / "out.docx"; audit = td / "audit"
            doc = Document()
            doc.add_heading("第一章 绪论", level=1)
            doc.add_paragraph("这是直接从第一章开始的正文片段。")
            doc.add_heading("第三章 实验", level=1)
            doc.add_paragraph("abstract")
            doc.save(source)
            spec = self._fixed_text_cover_declaration_spec()
            spec.pop("declarations")
            spec["cover"]["missing_value_policy"] = "placeholder"
            spec["cover"]["missing_value_placeholder"] = "——"
            del spec["thesis_profile"]["cover_metadata"]
            spec["thesis_profile"].pop("student_id", None)
            spec["thesis_profile"].pop("completion_date", None)
            spec["thesis_profile"].pop("co_supervisor_count", None)
            spec_path = td / "fragment-spec.json"
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            report = json.loads((audit / "validation-report.json").read_text())
            self.assertTrue(report["valid"], report["findings"])
            self.assertEqual(report["cover_changes"]["placement"], "document_start")
            self.assertEqual(report["cover_changes"]["page_break_written"], 1)
            contract = json.loads((audit / "cover-contract.json").read_text())
            self.assertEqual(contract["before_role"], "document_start")
            reopened = Document(output)
            nonempty = [p for p in reopened.paragraphs if p.text.strip()]
            self.assertEqual(nonempty[0].text, "示例研究院")
            chapter = next(p for p in reopened.paragraphs if p.text == "第一章 绪论")
            cover_last = chapter._p.getprevious()
            while cover_last is not None and cover_last.tag != qn("w:p"):
                cover_last = cover_last.getprevious()
            self.assertIsNotNone(cover_last)
            self.assertTrue(bool(cover_last.xpath('.//w:br[@w:type="page"]')))
            self.assertLess([p.text for p in reopened.paragraphs].index("示例研究院"),
                            [p.text for p in reopened.paragraphs].index("第一章 绪论"))

    def test_official_cover_layout_refuses_synthetic_redrawing(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "fragment.docx"; output = td / "out.docx"; audit = td / "audit"
            doc = Document(); doc.add_heading("第一章 绪论", level=1); doc.add_paragraph("正文"); doc.save(source)
            spec = self._fixed_text_cover_declaration_spec(); spec.pop("declarations")
            spec["cover"]["layout_id"] = "official_template_regions"
            spec_path = td / "spec.json"; spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("synthetic cover drawing is prohibited", result.stderr)
            self.assertFalse(output.exists())

    def test_pipeline_does_not_infer_cover_contract_from_requirement_keywords(self) -> None:
        module_spec = importlib.util.spec_from_file_location(
            "thesis_format_pipeline_cover_fallback", ROOT / "scripts" / "thesis_format_pipeline.py")
        pipeline = importlib.util.module_from_spec(module_spec)
        assert module_spec.loader is not None
        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            module_spec.loader.exec_module(pipeline)
        finally:
            sys.path.pop(0)
        spec = {"schema_version": "1.0", "roles": {}, "requirements": []}
        clauses = [
            {"text": "（一）封面"},
            {"text": "天津财经大学博士毕业（学位）论文"},
            {"text": "封面上方为论文题目，下方依次为专业名称、作者学号、论文作者、指导教师，最末为提交论文日期"},
            {"text": "扉页上方为论文题目（中英文对照）"},
        ]
        self.assertFalse(pipeline.ensure_declared_cover(spec, clauses))
        self.assertNotIn("cover", spec)
        spec["cover"] = {"before_role": "abstract_title_zh"}
        self.assertTrue(pipeline.ensure_declared_cover(spec, clauses))
        self.assertEqual(spec["cover"]["before_role"], "document_start")
        self.assertFalse(pipeline.ensure_declared_cover(spec, clauses), "an existing cover must not be overwritten")

    def test_fixed_text_schema_rejects_untrusted_missing_or_conflicting_cover_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "out.docx"; audit = td / "audit"
            doc = Document(); doc.styles.add_style("AbstractTitleCN", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("摘 要", "AbstractTitleCN"); doc.save(source)
            spec = self._fixed_text_cover_declaration_spec()
            spec["thesis_profile"]["cover_metadata"]["mystery"] = "unsafe"
            spec_path = td / "bad.json"; spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0); self.assertIn("unknown property 'mystery'", result.stderr)
            spec = self._fixed_text_cover_declaration_spec(); del spec["thesis_profile"]["cover_metadata"]["supervisor_name"]
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0); self.assertIn("supervisor_name", result.stderr)

            spec = self._fixed_text_cover_declaration_spec(); del spec["thesis_profile"]["cover_metadata"]["unit_code"]
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0); self.assertIn("unit_code", result.stderr)

            spec = self._fixed_text_cover_declaration_spec()
            spec["thesis_profile"]["cover_metadata"]["author_name"] = "   "
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0); self.assertIn("author_name", result.stderr)

    def test_fixed_text_schema_rejects_misbound_fields_and_noncanonical_declaration_placeholders(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "out.docx"; audit = td / "audit"
            doc = Document(); doc.styles.add_style("AbstractTitleCN", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("摘 要", "AbstractTitleCN"); doc.save(source)
            spec = self._fixed_text_cover_declaration_spec()
            spec["cover"]["fields"][0]["value_from"] = "thesis_profile.cover_metadata.title_zh"
            spec["declarations"]["items"][0]["resource_id"] = spec["declarations"]["items"][1]["resource_id"]
            spec["declarations"]["items"][1]["signature_placeholders"][0]["label"] = "研究生已签名"
            spec_path = td / "bad.json"; spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("must bind to its own field id", result.stderr)
            self.assertIn("does not match bound resource", result.stderr)
            self.assertIn("canonical blank placeholder set", result.stderr)

    def test_fixed_text_declaration_failure_report_and_full_subset_modes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); source = td / "source.docx"; output = td / "out.docx"; audit = td / "audit"
            doc = Document(); doc.add_paragraph("正文而无中文摘要标题"); doc.save(source)
            spec = self._fixed_text_cover_declaration_spec(); spec_path = td / "spec.json"
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                             "--out-dir", str(audit), "--compliance-mode", "supported_subset")
            self.assertNotEqual(result.returncode, 0)
            findings = json.loads((audit / "validation-report.json").read_text())["findings"]
            self.assertTrue(any(item["role"] == "declarations" and item["property"] == "before_role" for item in findings))

            spec = self._fixed_text_cover_declaration_spec(); spec["compliance_mode"] = "full"
            spec["clause_compliance"] = [{"clause_id": "C00041", "evidence_ids": ["E1"],
                "scope": "docx", "status": "unsupported_backend", "requirement_ids": [],
                "reason": "实体封皮颜色只能外部核验"}]
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            full = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(td / "full.docx"),
                           "--out-dir", str(td / "full-audit"), "--compliance-mode", "full")
            subset = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(td / "subset.docx"),
                             "--out-dir", str(td / "subset-audit"), "--compliance-mode", "supported_subset")
            self.assertNotEqual(full.returncode, 0); self.assertIn("full_compliance_analysis_failed", full.stderr)
            self.assertNotIn("full_compliance_analysis_failed", subset.stderr)

    def test_real_docx_table_text_dynamic_header_ooxml_and_idempotency(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); target = td / "target.docx"; spec_path = td / "spec.json"
            first = td / "first.docx"; second = td / "second.docx"
            d = Document(); heading = d.add_paragraph("第一章 绪论"); heading.style = "Heading 1"
            table = d.add_table(rows=2, cols=2)
            for i, cell in enumerate(cell for row in table.rows for cell in row.cells):
                cell.text = f"单元格{i + 1}"
            protected = OxmlElement("w:fldSimple"); protected.set(qn("w:instr"), "DATE \\@ yyyy-MM-dd")
            table.cell(0, 0).paragraphs[0]._p.append(protected)
            d.save(target)
            spec = {
                "schema_version": "1.0", "source_document": "test", "status": "semantic_resolved",
                "roles": {
                    "heading_1": {"font": {"cjk": "SimHei", "size_pt": 15}},
                    "table_text": {"font": {"cjk": "SimSun", "latin": "Times New Roman", "size_pt": 10.5},
                                   "paragraph": {"alignment": "center"}},
                    "header": {"font": {"cjk": "SimSun", "size_pt": 10.5},
                               "header_content": {"left_text": "中国农业大学", "right_field": "styleref_heading_1"},
                               "bottom_border": {"style": "thin_thick", "width_pt": 3,
                                                 "space_pt": 0, "color": "000000"}}
                },
                "requirements": []
            }
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            for source, output, audit in ((target, first, td / "audit-1"), (first, second, td / "audit-2")):
                result = run_raw("scripts/apply_format_spec.py", str(source), str(spec_path), str(output),
                                 "--out-dir", str(audit))
                self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
                report = json.loads((audit / "validation-report.json").read_text())
                self.assertTrue(report["valid"], report["findings"])
                self.assertEqual(report["paragraphs_directly_formatted"]["table_text"], 4)
            self.assertEqual(first.read_bytes(), second.read_bytes())
            check = Document(second)
            for cell in (cell for row in check.tables[0].rows for cell in row.cells):
                paragraph = cell.paragraphs[0]
                self.assertEqual(paragraph.style.name, "Thesis Table Text")
                self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            with zipfile.ZipFile(second) as zf:
                header_name = next(name for name in zf.namelist() if name.startswith("word/header") and name.endswith(".xml"))
                header_xml = zf.read(header_name).decode("utf-8")
                document_xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn("中国农业大学", header_xml)
            self.assertIn('STYLEREF &quot;Heading 1&quot; \\* MERGEFORMAT', header_xml)
            self.assertIn('w:val="thinThickMediumGap"', header_xml)
            self.assertIn('w:sz="24"', header_xml)
            self.assertEqual(header_xml.count("STYLEREF"), 1)
            self.assertIn("DATE \\@ yyyy-MM-dd", document_xml, "unrelated fields in table cells must be preserved")


if __name__ == "__main__": unittest.main()
