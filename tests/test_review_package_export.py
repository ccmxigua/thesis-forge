from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from docx import Document  # noqa: E402
from draft_scorecard import (append_scorecard, build_scorecard, reconcile_external_scorecard,
                             scorecard_lines)  # noqa: E402
from manual_review_display import _text as full_marker_text  # noqa: E402
from review_package_export import (export_review_package, strip_scorecard_display,
                                   _manual_marker_projections,
                                   _load_source_review_bundle, _build_source_crosswalk,
                                   _render_human_review_section, _actionable_manual_text,
                                   _is_reference_sample_numbering_issue, _issue_summary)  # noqa: E402
from semantic_contract import sha256_json  # noqa: E402


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


class ReviewPackageExportTests(unittest.TestCase):
    def make_card(self, source_sha: str, manual_items: list[dict] | None = None) -> dict:
        binding = {
            "run_id": "parent-run", "case_id": "bsu-test",
            "source_sha256": source_sha, "input_source_sha256": source_sha,
            "format_spec_sha256": "b" * 64,
        }
        receipt_audit = {
            "missing_count": 0, "unexpected_count": 0, "duplicate_count": 0,
            "expected_receipt_ids": ["r1"], "receipts": [{
                "receipt_id": "r1", "status": "unverified", "role": "body",
                "property_path": "font.cjk", "expected": "SimSun",
            }],
        }
        manual = manual_items or [{
            "marker_id": "MR-0001", "source_code": "missing_cover_metadata:unit_code",
            "source_text": "unit_code 尚未确认", "reason": "用户尚未确认",
            "action": "请确认 unit_code", "clause_ids": [],
            "manual_obligation_id": "MO-test-0001", "placeholder_text": "【待提供：unit_code】",
        }]
        return build_scorecard(binding, receipt_audit, manual, [])

    def make_source_review_bundle(self, root: Path) -> tuple[Path, Path]:
        source_docx = root / "requirements.docx"
        full_text = ("自行确定是否使用四级标题。论文中出现英文时需要使用Times New Roman字体。"
                     "这个结构可以根据自己论文需要进行调整。这些文字都是编的，根据需要自己撰写。")
        doc = Document()
        doc.add_paragraph("3.3.3 研究创新点")
        doc.add_paragraph(full_text)
        doc.save(source_docx)
        requirements_dir = root / "requirements"
        requirements_dir.mkdir()
        normalization_id = "norm-test-1"
        source_sha = digest(source_docx)
        provenance = {
            "version": "1.0", "origin": "fresh_host_agent", "source_sha256": source_sha,
            "evidence_sha256": "a" * 64, "clause_sha256": "b" * 64,
            "request_sha256": "c" * 64, "run_id": "run-test-1",
        }
        spec = {
            "schema_version": "1.0", "source_document": str(source_docx.resolve()),
            "run_id": "run-test-1", "semantic_review_provenance": provenance,
            "semantic_review_provenance_valid": False,
            "requirements": [{
                "id": "R00489", "role": "thesis_affiliation_date",
                "properties": {"font": {"latin": "Times New Roman"}},
                "evidence_ids": ["E00119"], "clause_ids": ["C00176"],
                "source_text": "论文中出现英文时需要使用Times New Roman字体",
                "applicability": {"status": "conditional", "conditions": [{
                    "fact": "source_inventory.english_text", "operator": "present", "value": None,
                }], "exceptions": []},
            }],
        }
        clause = [{
            "id": "C00176", "text": "论文中出现英文时需要使用Times New Roman字体",
            "source_text_full": full_text, "source_evidence_text": full_text,
            "source_span": {
                "evidence_id": "E00119", "start_offset": 13, "end_offset": 42,
                "text": "论文中出现英文时需要使用Times New Roman字体",
                "source_sha256": hashlib.sha256(full_text.encode()).hexdigest(),
                "location": {"part": "document", "child_index": 1, "order": 1},
            },
            "evidence_ids": ["E00119"], "context_before": ["3.3.3 研究创新点"],
            "context_after": ["4 研究过程与分析"],
            "location": {"part": "document", "child_index": 1, "order": 1},
        }]
        spec_path = requirements_dir / "format-spec.json"
        clauses_path = requirements_dir / "requirement-clauses.json"
        input_path = requirements_dir / "requirements-input-manifest.json"
        extraction_path = requirements_dir / "extraction-manifest.json"
        write_json(spec_path, spec)
        write_json(clauses_path, clause)
        write_json(input_path, {
            "normalization_id": normalization_id,
            "original": {"path": str(source_docx.resolve()), "sha256": source_sha},
            "normalized": {"path": str(source_docx.resolve()), "sha256": source_sha},
        })
        write_json(extraction_path, {
            "run_id": "run-test-1", "source_document": str(source_docx.resolve()),
            "source_sha256": source_sha, "format_spec_sha256": digest(spec_path),
            "requirement_clauses_sha256": digest(clauses_path),
            "semantic_review_provenance": provenance,
            "requirements_input_normalization": {
                "normalization_id": normalization_id,
                "original": {"sha256": source_sha},
            },
        })
        return requirements_dir, source_docx

    def make_review_entry(self, *, requirement_ids=None, clause_ids=None,
                          evidence_ids=None) -> dict:
        return {
            "item_id": "SC-test", "kind": "human_review", "status": "pending",
            "detail": {
                "marker_id": "MR-0001", "source_type": "capability_finding",
                "source_code": "capability.requirement_unknown",
                "category": "runtime_manual_unverifiable", "requirement_ids": requirement_ids or [],
                "clause_ids": clause_ids or [], "evidence_ids": evidence_ids or [],
                "question_ids": [], "source_text": "runtime_manual_unverifiable",
                "reason": "The runtime check is not implemented.",
                "action": "核对论文与学校来源。", "original_blocking": False,
            },
        }

    def make_review_card(self, entry: dict) -> dict:
        return {"entries": [entry]}

    def test_strip_removes_only_exact_scorecard_and_retains_manual_marker(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "source.docx"
            baseline = root / "baseline.docx"
            output = root / "paper.docx"
            scorecard_path = root / "scorecard.json"
            report_path = root / "split-report.json"
            Document().save(source)
            card = self.make_card(digest(source))
            doc = Document()
            doc.add_paragraph("原始论文正文")
            append_scorecard(doc, card)
            detail = next(e["detail"] for e in card["entries"] if e["kind"] == "human_review")
            doc.add_paragraph(full_marker_text(detail))
            doc.add_paragraph("来源绑定的论文正文仍保留")
            doc.save(baseline)
            write_json(scorecard_path, card)

            result = strip_scorecard_display(source, baseline, scorecard_path, output, report_path)
            self.assertEqual(result["removed_scorecard_item_ids"],
                             [entry["item_id"] for entry in card["entries"]])
            self.assertEqual(result["retained_manual_review_marker_count"], 1)
            self.assertEqual(result["changed_ooxml_parts"], ["word/document.xml"])
            text = "\n".join(p.text for p in Document(output).paragraphs)
            self.assertNotIn("【SC-", text)
            self.assertNotIn("自动核验评分（", text)
            self.assertIn("【MR-0001｜人工待审】", text)
            self.assertIn("详见同编号独立审查台账。", text)
            self.assertNotIn("请在此人工处理：", text)
            self.assertNotIn("原文/问题：", text)
            self.assertIn("来源绑定的论文正文仍保留", text)
            self.assertEqual(result["compacted_manual_review_marker_count"], 1)

    def test_strip_compacts_multiple_markers_and_preserves_order_and_ids(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, baseline, output = (root / "source.docx", root / "baseline.docx",
                                        root / "paper.docx")
            card_path, report = root / "scorecard.json", root / "report.json"
            Document().save(source)
            manual = [
                {"marker_id": "MR-0001", "source_code": "one", "source_text": "原文甲",
                 "reason": "理由甲", "action": "处理甲", "clause_ids": [],
                 "manual_obligation_id": "MO-test-0001"},
                {"marker_id": "MR-0002", "source_code": "two", "source_text": "原文乙",
                 "reason": "理由乙", "action": "处理乙", "clause_ids": [],
                 "manual_obligation_id": "MO-test-0002"},
            ]
            card = self.make_card(digest(source), manual)
            doc = Document()
            append_scorecard(doc, card)
            details = [e["detail"] for e in card["entries"] if e["kind"] == "human_review"]
            for detail in details:
                doc.add_paragraph(full_marker_text(detail))
            doc.add_paragraph("正文")
            doc.save(baseline)
            write_json(card_path, card)

            result = strip_scorecard_display(source, baseline, card_path, output, report)
            projections = _manual_marker_projections(card)
            texts = [p.text for p in Document(output).paragraphs]
            marker_texts = [text for text in texts if text.startswith("【MR-")]
            self.assertEqual(marker_texts,
                             [projections[detail["marker_id"]]["paper_text"]
                              for detail in details])
            self.assertEqual(result["compacted_manual_review_marker_ids"],
                             [detail["marker_id"] for detail in details])

    def test_strip_rejects_scorecard_or_manual_marker_mismatch(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "source.docx"
            baseline = root / "baseline.docx"
            output = root / "paper.docx"
            card_path = root / "scorecard.json"
            report = root / "report.json"
            Document().save(source)
            card = self.make_card(digest(source))
            doc = Document()
            doc.add_paragraph("body")
            append_scorecard(doc, card)
            doc.add_paragraph("【MR-0002｜人工待审】错误标记")
            doc.save(baseline)
            write_json(card_path, card)
            with self.assertRaisesRegex(ValueError, "manual-review markers"):
                strip_scorecard_display(source, baseline, card_path, output, report)

            doc = Document()
            doc.add_paragraph("body")
            append_scorecard(doc, card)
            doc.paragraphs[0].text = "tampered summary"
            detail = next(e["detail"] for e in card["entries"] if e["kind"] == "human_review")
            doc.add_paragraph(full_marker_text(detail))
            doc.save(baseline)
            with self.assertRaisesRegex(ValueError, "front block"):
                strip_scorecard_display(source, baseline, card_path, output, report)

            doc = Document()
            doc.add_paragraph("body")
            append_scorecard(doc, card)
            doc.add_paragraph(full_marker_text(detail).replace("unit_code", "tampered"))
            doc.save(baseline)
            with self.assertRaisesRegex(ValueError, "exact projection of the bound ledger"):
                strip_scorecard_display(source, baseline, card_path, output, report)

    def test_external_ledger_exports_all_entries_and_current_binding(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "source.docx"
            paper = root / "paper.docx"
            scorecard_path = root / "scorecard.json"
            reconciliation_path = root / "reconciliation.json"
            ledger_docx = root / "review-ledger.docx"
            ledger_json = root / "review-ledger.json"
            summary = root / "summary.md"
            Document().save(source)
            card = self.make_card(digest(source))
            doc = Document()
            detail = next(e["detail"] for e in card["entries"] if e["kind"] == "human_review")
            doc.add_paragraph(_manual_marker_projections(card)[detail["marker_id"]]["paper_text"])
            doc.add_paragraph("论文正文")
            doc.save(paper)
            report = {
                "protocol": "rendered_format_audit_v1",
                "docx_sha256": digest(paper),
                "final_docx": {"sha256": digest(paper)},
                "pdf_sha256": "c" * 64,
                "findings": [{"code": "drawing_not_found_in_rendered_pdf", "drawing_index": 1}],
                "submission_ready": False, "field_refresh_claimed": False,
            }
            report["audit_sha256"] = sha256_json(report)
            current_card = reconcile_external_scorecard(card, report, paper)
            write_json(scorecard_path, current_card)
            reconciliation = {
                "status": "complete", "submission_ready": False,
                "model_request_made": False, "semantic_review_is_new": False,
                "current_output": {"docx_sha256": digest(paper), "pdf_sha256": "c" * 64},
                "scorecard_reconciliation": {"status": "external_ledger_only"},
            }
            write_json(reconciliation_path, reconciliation)

            result = export_review_package(source, paper, scorecard_path, reconciliation_path,
                                           ledger_docx, ledger_json, summary)
            self.assertEqual(result["item_count"], current_card["item_count"])
            self.assertFalse(result["submission_ready"])
            wrapper = json.loads(ledger_json.read_text())
            self.assertEqual(len(wrapper["scorecard"]["entries"]), current_card["item_count"])
            self.assertEqual(wrapper["paper_docx"]["sha256"], digest(paper))
            self.assertIn("培养单位代码", summary.read_text())
            self.assertNotIn("unit_code", summary.read_text())
            self.assertIn("封面培养单位代码待确认", summary.read_text())
            self.assertIn("字体相关项仍未全部核验", summary.read_text())
            self.assertIn("未核验/待人工项不计为通过", summary.read_text())
            editable = Document(ledger_docx)
            self.assertEqual(len(editable.tables[-1].rows) - 1, current_card["item_count"])

    def test_reader_guide_resolves_bound_source_and_flags_scope_without_machine_ids(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            requirements_dir, requirements_docx = self.make_source_review_bundle(root)
            bundle = _load_source_review_bundle(requirements_dir, requirements_docx)
            card = self.make_review_card(self.make_review_entry(requirement_ids=["R00489"]))
            crosswalk = _build_source_crosswalk(card, bundle)
            self.assertEqual(crosswalk[0]["crosswalk_status"], "resolved")
            record = crosswalk[0]["resolved_source_records"][0]
            self.assertEqual(record["clause_id"], "C00176")
            self.assertEqual(record["evidence_ids"], ["E00119"])
            self.assertEqual(record["role"], "thesis_affiliation_date")
            self.assertFalse(bundle["semantic_review_provenance_valid"])

            guide = "\n".join(_render_human_review_section(
                card, bundle, crosswalk, Path("technical.json"),
            ))
            self.assertIn("论文中出现英文时需要使用Times New Roman字体", guide)
            self.assertIn("论文单位与日期栏", guide)
            self.assertIn("论文中出现英文时", guide)
            self.assertIn("逐一检查上述区域中实际出现的英文", guide)
            self.assertNotIn("适用范围需先确认", guide)
            self.assertNotIn("研究创新点", guide)
            self.assertNotIn("这些文字都是编的", guide)
            self.assertIn("语义来源校验标记无效", guide)
            self.assertIn("该标记本身不表示来源原文不清楚", guide)
            self.assertIn("不表示论文内容错误，也不表示已经通过", guide)
            self.assertNotIn("R00489", guide)
            self.assertNotIn("C00176", guide)
            self.assertNotIn("E00119", guide)
            self.assertFalse(card.get("submission_ready", False))

    def test_reader_guide_gives_searchable_abstract_review_actions(self):
        entries = [
            {
                "item_id": "SC-abstract-1", "kind": "human_review", "status": "pending",
                "detail": {
                    "marker_id": "MR-1", "source_type": "current_session_semantic_review",
                    "source_code": "semantic_review:abstract_zh.quality_guidance",
                    "source_text": "中文摘要是论文内容的简要陈述，300～1000字；其内容包括目的、方法、成果和结论。",
                    "reason": "当前会话 Agent判定 uncertain：结构说明可能影响摘要凝练度。；原文证据：本文的创新点主要体现在三个方面 / 全文包含三章编号正文，系统地阐述了研究过程",
                    "action": "当前会话 Agent无法可靠判断此项；请人工对照权威条款核实后重新运行。",
                },
            },
            {
                "item_id": "SC-abstract-2", "kind": "human_review", "status": "pending",
                "detail": {
                    "marker_id": "MR-2", "source_type": "current_session_semantic_review",
                    "source_code": "semantic_review:abstract_zh.required_sections",
                    "source_text": "中文摘要内容应包括目的、方法、成果和结论。",
                    "reason": "当前会话 Agent判定 uncertain：末段没有清晰独立的结论句。；原文证据：实验结果表明模型精度优于基线 / 本文的创新点主要体现在三个方面",
                    "action": "当前会话 Agent无法可靠判断此项；请人工对照权威条款核实后重新运行。",
                },
            },
        ]
        card = {"entries": entries}
        crosswalk = [entry["detail"] for entry in entries]
        guide = "\n".join(_render_human_review_section(card, None, crosswalk, None))
        self.assertIn("可在论文中搜索这些摘要原句", guide)
        self.assertIn("全文包含三章编号正文，系统地阐述了研究过程", guide)
        self.assertIn("研究目的、方法和主要结果", guide)
        self.assertIn("独立结论句", guide)
        self.assertIn("与正文结论段对照", guide)
        self.assertNotIn("Agent", guide)
        self.assertNotIn("中文摘要是论文内容的简要陈述", guide)

    def test_reader_guide_shortens_repeated_abstract_source_citation(self):
        long_quote = ("The Chinese abstract is a brief statement of the content of the paper, "
                      "which is usually written in third person tone, 300 to 1,000 words, "
                      "without comment and explanation, is a short essay with independence and integrity.")
        detail = {
            "marker_id": "MR-3", "source_type": "open_question", "source_code": "open_question",
            "source_text": long_quote, "reason": "该条款证据不足，需要人工确认",
            "action": "请人工确认该条款的适用对象、数值单位或权威解释。",
        }
        projected = {**detail, "resolved_source_records": [{"clause_text": long_quote}]}
        guide = "\n".join(_render_human_review_section(
            {"entries": [{"item_id": "SC-3", "kind": "human_review", "detail": detail}]},
            None, [projected], None,
        ))
        self.assertIn("摘要语言与长度计量待核", guide)
        self.assertIn("Chinese abstract，却使用 words 计数", guide)
        self.assertIn("英文摘要", guide)
        self.assertNotIn(long_quote, guide)

    def test_spine_textbox_rewrite_is_specific_and_layout_actionable(self):
        item = {
            "source_type": "release_gate",
            "source_text": "C00454 [E00341]: 论文题目；C00456 [E00343]: 研究生姓名；C00458 [E00345]: 3cm左右",
            "reason": "rule_spec 中 R00602/R00603 将其映射为 cover_field_value；这些文字在独立文本框中，作用域可能是书脊。",
            "action": "对照来源证据人工核对冲突并记录结论。",
        }
        source, reason, action = _actionable_manual_text(item)
        self.assertEqual(source, "模板文本框的封面或书脊归属待核")
        self.assertIn("论文题目、研究生姓名和约3cm标注", reason)
        self.assertIn("位置和方向", action)
        self.assertNotIn("cover_field_value", " ".join((source, reason, action)))
        self.assertNotIn("R00602", " ".join((source, reason, action)))

        printability = {
            "source_text": "书脊是否能印取决于纸张、装订和实际页数。",
            "reason": "实际打印条件尚未确认。", "action": "人工核对。",
        }
        self.assertEqual(_actionable_manual_text(printability), (
            printability["source_text"], printability["reason"], printability["action"],
        ))

    def test_reader_guide_uses_plain_labels_for_metadata_and_keyword_rules(self):
        unit = {
            "source_text": "封面字段 unit_code 缺少已确认的 metadata",
            "reason": "当前运行没有提供 unit_code；草稿使用中性占位符。",
            "action": "补充 thesis_profile.cover_metadata.unit_code。",
        }
        source, reason, action = _actionable_manual_text(unit)
        self.assertEqual(source, "封面培养单位代码待确认")
        self.assertIn("中性占位符", reason)
        self.assertIn("提供或确认", action)
        self.assertNotIn("unit_code", " ".join((source, reason, action)))
        self.assertNotIn("metadata", " ".join((source, reason, action)))

        keyword = {
            "source_text": "关键词是为了便于文献检索工作而从论文中选取出来表示全文主题内容信息的单词或术语，在论文中有明确出处",
            "reason": "来源要求关键词和论文主题相符。", "action": "逐项核实。",
        }
        source, reason, action = _actionable_manual_text(keyword)
        self.assertEqual(source, "关键词须能代表论文主题并有文内出处")
        self.assertIn("逐个搜索关键词", action)
        self.assertNotIn("关键词是为了便于", source)
        full_quote = keyword["source_text"]
        detail = {
            **keyword, "marker_id": "MR-keyword", "source_type": "current_session_semantic_review",
            "source_code": "semantic_review:keywords.provenance",
        }
        projected = {**detail, "resolved_source_records": [{"clause_text": full_quote}]}
        guide = "\n".join(_render_human_review_section(
            {"entries": [{"item_id": "SC-keyword", "kind": "human_review", "detail": detail}]},
            None, [projected], None,
        ))
        self.assertIn("关键词须能代表论文主题并有文内出处", guide)
        self.assertNotIn(full_quote, guide)

    def test_reference_category_examples_are_not_reported_as_confirmed_conflicts(self):
        items = []
        crosswalk = []
        for index in range(4):
            detail = {
                "marker_id": f"MR-{index}", "source_type": "release_gate",
                "source_code": f"reported_conflict:{index}", "category": "semantic_content_review",
                "source_text": "[1] 分类示例条目；[2] 另一示例。",
                "reason": "参考文献示例按类别编号，切换类别时示例重新编号；需确认最终规则。",
                "action": "对照来源证据人工核对冲突并记录结论。",
            }
            items.append({"item_id": f"SC-{index}", "kind": "human_review", "status": "pending",
                          "detail": detail})
            crosswalk.append(detail)
        self.assertTrue(all(_is_reference_sample_numbering_issue(row) for row in crosswalk))
        guide = "\n".join(_render_human_review_section({"entries": items}, None, crosswalk, None))
        self.assertIn("有 4 条既有待办", guide)
        self.assertIn("示例编号本身不能证明", guide)
        self.assertIn("不把它们列为已证实冲突", guide)
        self.assertNotIn("切换类别时示例重新编号", guide)

    def test_issue_summary_keeps_unlocated_font_and_drawing_findings_qualified(self):
        entries = []
        for role in ("figure_caption", "table_caption", "table_text"):
            entries.append({
                "item_id": f"SC-{role}", "kind": "property", "status": "failed",
                "label": role + ".font.cjk",
                "detail": {"role": role, "property_path": "font.cjk", "expected": "SimSun",
                           "actual": "SimSun", "current_output_instance": {
                               "finding_code": "rendered_pdf_font_mismatch"}},
            })
        entries.append({
            "item_id": "SC-cover-font", "kind": "rendered_output_instance", "status": "failed",
            "label": "cover field font", "detail": {
                "code": "rendered_pdf_font_mismatch",
                "expected": {"role": "cover_field_value", "script": "cjk", "name": "FangSong"},
            },
        })
        for index in range(1, 4):
            entries.append({
                "item_id": f"SC-drawing-{index}", "kind": "rendered_output_instance",
                "status": "unverified", "label": "drawing",
                "detail": {"code": "drawing_not_found_in_rendered_pdf", "candidate_count": 0,
                           "current_output_instance": {"object_identity": {"page": None}}},
            })
        card = {
            "entries": entries,
            "status_counts": {"verified": 0, "failed": 4, "unverified": 3, "pending": 0},
            "rendered_output_audit": {"report": {"actual_pdf_font_resources": [
                {"base_font": "ABAAAA+NotoSerifCJKsc-Regular"},
                {"base_font": "FAAAAA+NotoSansCJKhk-Regular"},
            ]}},
        }
        guide = _issue_summary(
            card, {"current_output": {}}, {"output_docx": {"path": "draft.docx"}},
            Path("ledger.json"), Path("ledger.docx"),
        )
        self.assertIn("自动检查标记的失败项（待定位）", guide)
        self.assertIn("DOCX 保存的样式值为 SimSun", guide)
        self.assertIn("Noto Serif CJK SC", guide)
        self.assertIn("Noto Sans CJK HK", guide)
        self.assertIn("实际 PDF 字体和位置仍未知", guide)
        self.assertIn("PDF 自动比对未匹配到 3 个源文档图形", guide)
        self.assertIn("这不能证明图形缺失", guide)
        self.assertNotIn("源文档图形缺失", guide)

    def test_reader_guide_keeps_missing_or_conflicting_source_links_unresolved(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            requirements_dir, requirements_docx = self.make_source_review_bundle(root)
            bundle = _load_source_review_bundle(requirements_dir, requirements_docx)
            missing = self.make_review_card(self.make_review_entry(requirement_ids=["R99999"]))
            missing_crosswalk = _build_source_crosswalk(missing, bundle)
            self.assertEqual(missing_crosswalk[0]["crosswalk_status"], "conflict")
            self.assertIn("requirement_id_not_found", missing_crosswalk[0]["crosswalk_problems"])
            missing_guide = "\n".join(_render_human_review_section(
                missing, bundle, missing_crosswalk, None,
            ))
            self.assertIn("来源条款尚未关联", missing_guide)
            self.assertEqual(missing_guide.count("来源条款尚未关联"), 1)
            self.assertNotIn("Times New Roman", missing_guide)
            self.assertNotIn("已经通过", missing_guide.replace("不表示已经通过", ""))
            self.assertNotIn("已判定符合", missing_guide)

            conflict = self.make_review_card(self.make_review_entry(
                requirement_ids=["R00489"], evidence_ids=["E-FOREIGN"],
            ))
            conflict_crosswalk = _build_source_crosswalk(conflict, bundle)
            self.assertEqual(conflict_crosswalk[0]["crosswalk_status"], "conflict")
            self.assertIn("scorecard_evidence_refs_conflict_with_requirement",
                          conflict_crosswalk[0]["crosswalk_problems"])

    def test_reader_guide_rejects_clause_quote_not_present_in_bound_docx(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            requirements_dir, requirements_docx = self.make_source_review_bundle(root)
            clause_path = requirements_dir / "requirement-clauses.json"
            clauses = json.loads(clause_path.read_text())
            clauses[0]["source_text_full"] = "伪造的学校要求"
            clauses[0]["source_evidence_text"] = "伪造的学校要求"
            clauses[0]["source_span"]["source_sha256"] = hashlib.sha256(
                "伪造的学校要求".encode()
            ).hexdigest()
            write_json(clause_path, clauses)
            extraction_path = requirements_dir / "extraction-manifest.json"
            extraction = json.loads(extraction_path.read_text())
            extraction["requirement_clauses_sha256"] = digest(clause_path)
            write_json(extraction_path, extraction)
            bundle = _load_source_review_bundle(requirements_dir, requirements_docx)
            card = self.make_review_card(self.make_review_entry(requirement_ids=["R00489"]))
            crosswalk = _build_source_crosswalk(card, bundle)
            self.assertEqual(crosswalk[0]["crosswalk_status"], "conflict")
            self.assertIn("clause_text_not_bound_to_source_docx", crosswalk[0]["crosswalk_problems"])
            self.assertEqual(crosswalk[0]["resolved_source_records"], [])

    def test_export_writes_reader_report_and_separate_bound_technical_appendix(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            paper_source = root / "paper-source.docx"
            paper = root / "paper.docx"
            requirements_dir, requirements_docx = self.make_source_review_bundle(root)
            Document().save(paper_source)
            manual = [{
                "marker_id": "MR-0001", "source_code": "capability.requirement_unknown",
                "source_type": "capability_finding", "category": "runtime_manual_unverifiable",
                "source_text": "Requirement R00489 is not execution-ready.",
                "reason": "runtime_manual_unverifiable", "action": "请核对原始条款。",
                "clause_ids": [], "requirement_ids": ["R00489"],
                "question_ids": [], "evidence_ids": [],
                "manual_obligation_id": "MO-test-0001",
                "placeholder_text": "【待人工处理】",
            }]
            card = self.make_card(digest(paper_source), manual)
            marker_projection = _manual_marker_projections(card)["MR-0001"]["paper_text"]
            doc = Document()
            doc.add_paragraph(marker_projection)
            doc.add_paragraph("paper body")
            doc.save(paper)
            render_report = {
                "protocol": "rendered_format_audit_v1", "docx_sha256": digest(paper),
                "final_docx": {"sha256": digest(paper)}, "pdf_sha256": "c" * 64,
                "findings": [], "submission_ready": False, "field_refresh_claimed": False,
            }
            render_report["audit_sha256"] = sha256_json(render_report)
            card = reconcile_external_scorecard(card, render_report, paper)
            scorecard = root / "scorecard.json"
            write_json(scorecard, card)
            reconciliation = root / "reconciliation.json"
            write_json(reconciliation, {
                "status": "complete", "submission_ready": False,
                "model_request_made": False, "semantic_review_is_new": False,
                "current_output": {"docx_sha256": digest(paper), "pdf_sha256": "c" * 64},
            })
            ledger_docx, ledger_json = root / "ledger.docx", root / "ledger.json"
            guide, technical = root / "reader-guide.md", root / "reader-guide-technical.json"

            result = export_review_package(
                paper_source, paper, scorecard, reconciliation, ledger_docx, ledger_json, guide,
                source_requirements_dir=requirements_dir,
                requirements_docx_path=requirements_docx,
                technical_json_path=technical,
            )
            report_text = guide.read_text(encoding="utf8")
            technical_payload = json.loads(technical.read_text(encoding="utf8"))
            self.assertFalse(result["submission_ready"])
            self.assertIn("来源原文", report_text)
            self.assertIn("这条文字明确规定论文中出现英文时使用 Times New Roman", report_text)
            self.assertIn("逐一检查上述区域中实际出现的英文", report_text)
            self.assertNotIn("适用范围需先确认", report_text)
            self.assertNotIn("SC-test", report_text)
            self.assertNotIn("R00489", report_text)
            self.assertNotIn("C00176", report_text)
            self.assertNotIn("E00119", report_text)
            crosswalk = technical_payload["human_review_crosswalk"][0]
            source_record = crosswalk["resolved_source_records"][0]
            self.assertEqual(source_record["requirement_id"], "R00489")
            self.assertEqual(source_record["clause_id"], "C00176")
            self.assertEqual(source_record["evidence_ids"], ["E00119"])
            self.assertFalse(technical_payload["source_requirement_bundle"]["format_spec"][
                "semantic_review_provenance_valid"
            ])
            self.assertEqual(technical_payload["reader_report"]["sha256"], digest(guide))

    def test_external_package_rejects_stale_output_hash_and_release_claim(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "source.docx"
            paper = root / "paper.docx"
            Document().save(source)
            doc = Document()
            doc.add_paragraph("paper")
            doc.save(paper)
            card = self.make_card(digest(source))
            report = {"protocol": "rendered_format_audit_v1", "docx_sha256": digest(paper),
                      "final_docx": {"sha256": digest(paper)}, "pdf_sha256": "c" * 64,
                      "findings": [], "submission_ready": False, "field_refresh_claimed": False}
            report["audit_sha256"] = sha256_json(report)
            write_json(root / "scorecard.json", reconcile_external_scorecard(card, report, paper))
            recon = {"status": "complete", "submission_ready": False,
                     "model_request_made": False, "semantic_review_is_new": False,
                     "current_output": {"docx_sha256": "d" * 64, "pdf_sha256": "c" * 64}}
            write_json(root / "recon.json", recon)
            with self.assertRaisesRegex(ValueError, "not bound"):
                export_review_package(source, paper, root / "scorecard.json", root / "recon.json",
                                      root / "ledger.docx", root / "ledger.json", root / "summary.md")
            recon["current_output"]["docx_sha256"] = digest(paper)
            recon["submission_ready"] = True
            write_json(root / "recon.json", recon)
            with self.assertRaisesRegex(ValueError, "inherited, non-model, and non-submission"):
                export_review_package(source, paper, root / "scorecard.json", root / "recon.json",
                                      root / "ledger.docx", root / "ledger.json", root / "summary.md")


if __name__ == "__main__":
    unittest.main()
