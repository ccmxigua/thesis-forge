from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from page_field_ops import inspect_docx, parse_field_instruction
from section_executor import execute_section_plan
from section_model import audit_plan_against_docx, compile_section_plan


def set_page_format(section, fmt: str | None, start: int | None = None) -> None:
    old = section._sectPr.find(qn("w:pgNumType"))
    if old is not None:
        section._sectPr.remove(old)
    if fmt is None:
        return
    node = OxmlElement("w:pgNumType")
    node.set(qn("w:fmt"), {"roman": "lowerRoman", "roman_upper": "upperRoman", "decimal": "decimal"}[fmt])
    if start is not None:
        node.set(qn("w:start"), str(start))
    section._sectPr.append(node)


def add_page(story, switch: str, *, complex_field: bool = False) -> None:
    paragraph = story.paragraphs[0]
    if not complex_field:
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), f" PAGE \\* {switch} ")
        paragraph._p.append(field)
        return
    for kind, text in (("begin", None), (None, " PAGE \\* "), (None, switch), ("separate", None), ("end", None)):
        run = OxmlElement("w:r")
        if kind:
            char = OxmlElement("w:fldChar")
            char.set(qn("w:fldCharType"), kind)
            run.append(char)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run.append(instr)
        paragraph._p.append(run)


def make_sections(count: int) -> Document:
    doc = Document()
    doc.add_paragraph("封面")
    for index in range(2, count + 1):
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_paragraph(f"Section {index}")
    return doc


def codes(findings: list[dict]) -> set[str]:
    return {item["code"] for item in findings}


class SectionModelTests(unittest.TestCase):
    def save(self, doc: Document, directory: str, name: str = "test.docx") -> Path:
        path = Path(directory) / name
        doc.save(path)
        return path

    def test_cover_none_then_decimal_body_selected_by_section_index(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(2)
            set_page_format(doc.sections[0], None)
            set_page_format(doc.sections[1], "decimal", 1)
            doc.sections[1].footer.is_linked_to_previous = False
            add_page(doc.sections[1].footer, "ARABIC")
            path = self.save(doc, td)
            spec = {"page": {"page_number": {
                "front_matter_format": "none", "body_format": "decimal", "body_start": 1,
                "body_start_selector": {"strategy": "section_index", "section_index": 2},
            }}}
            plan = compile_section_plan(spec, path)
            self.assertTrue(plan["valid"], plan["findings"])
            self.assertEqual([s["format"] for s in plan["sections"]], ["none", "decimal"])
            self.assertEqual([s["restart"] for s in plan["sections"]], [False, True])
            self.assertEqual(audit_plan_against_docx(plan, path), [])

    def test_unnumbered_cover_roman_front_and_decimal_body_have_explicit_restarts(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(4)
            path = self.save(doc, td)
            spec = {"page": {"page_number": {
                "front_matter_format": "roman_upper", "body_format": "decimal",
                "front_matter_start": 1, "body_start": 1,
                "front_matter_start_selector": {"strategy": "section_index", "section_index": 3},
                "body_start_selector": {"strategy": "section_index", "section_index": 4},
            }}}
            plan = compile_section_plan(spec, path)
            self.assertTrue(plan["valid"], plan["findings"])
            self.assertEqual(plan["front_matter_start_section"], 3)
            self.assertEqual([s["zone"] for s in plan["sections"]],
                             ["cover", "cover", "front", "body"])
            self.assertEqual([s["format"] for s in plan["sections"]],
                             ["none", "none", "roman_upper", "decimal"])
            self.assertEqual([(s["restart"], s["start"]) for s in plan["sections"]],
                             [(False, None), (False, None), (True, 1), (True, 1)])
            execute_section_plan(doc, plan, alignment="center")
            output = self.save(doc, td, "output.docx")
            self.assertEqual(audit_plan_against_docx(plan, output), [])

    def test_lower_and_upper_roman_and_decimal_matrix_with_two_restarts(self) -> None:
        cases = (("roman", "roman"), ("roman_upper", "ROMAN"))
        for front_format, switch in cases:
            with self.subTest(front_format=front_format), tempfile.TemporaryDirectory() as td:
                doc = make_sections(3)
                for index, section in enumerate(doc.sections):
                    section.footer.is_linked_to_previous = False
                    if index < 2:
                        set_page_format(section, front_format, 4 if index == 0 else None)
                        add_page(section.footer, switch, complex_field=index == 1)
                    else:
                        set_page_format(section, "decimal", 1)
                        add_page(section.footer, "ARABIC")
                path = self.save(doc, td)
                spec = {"page": {"page_number": {
                    "front_matter_format": front_format, "body_format": "decimal",
                    "front_matter_start": 4, "body_start": 1,
                    "body_start_selector": {"strategy": "section_index", "section_index": 3},
                }}}
                plan = compile_section_plan(spec, path)
                self.assertEqual([(s["restart"], s["start"]) for s in plan["sections"]],
                                 [(True, 4), (False, None), (True, 1)])
                self.assertEqual(audit_plan_against_docx(plan, path), [])

    def test_first_heading_and_heading_text_selectors(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(3)
            body_heading = doc.paragraphs[-1]
            body_heading.text = "第1章 绪论"
            body_heading.style = "Heading 1"
            path = self.save(doc, td)
            base = {"front_matter_format": "roman", "body_format": "decimal"}
            first = compile_section_plan({"page": {"page_number": {
                **base, "body_start_selector": {"strategy": "first_heading_1"},
            }}}, path)
            text = compile_section_plan({"page": {"page_number": {
                **base, "body_start_selector": {"strategy": "heading_text", "heading_text_pattern": "绪论$"},
            }}}, path)
            self.assertEqual(first["body_start_section"], 3)
            self.assertEqual(text["body_start_section"], 3)

    def test_heading_selectors_fail_closed_when_multiple_matches_exist(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(3)
            for paragraph, text in zip(
                (doc.paragraphs[-2], doc.paragraphs[-1]),
                ("第1章 绪论", "第2章 方法"),
            ):
                paragraph.text = text
                paragraph.style = "Heading 1"
            path = self.save(doc, td)
            spec = {"page": {"page_number": {
                "front_matter_format": "roman", "body_format": "decimal",
                "body_start_selector": {"strategy": "first_heading_1"},
            }}}
            plan = compile_section_plan(spec, path)
            self.assertFalse(plan["valid"])
            self.assertIn("body_selector_ambiguous", codes(plan["findings"]))

    def test_default_first_even_headers_footers_and_settings_are_parsed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(1)
            doc.sections[0].different_first_page_header_footer = True
            doc.settings.odd_and_even_pages_header_footer = True
            set_page_format(doc.sections[0], "roman_upper", 1)
            for accessor in ("footer", "first_page_footer", "even_page_footer"):
                story = getattr(doc.sections[0], accessor)
                story.is_linked_to_previous = False
                add_page(story, "ROMAN")
            for accessor in ("header", "first_page_header", "even_page_header"):
                story = getattr(doc.sections[0], accessor)
                story.is_linked_to_previous = False
                story.paragraphs[0].add_run(accessor)
            path = self.save(doc, td)
            evidence = inspect_docx(path)
            section = evidence["sections"][0]
            self.assertTrue(section["titlePg"])
            self.assertTrue(section["evenAndOdd"])
            for kind in ("header", "footer"):
                self.assertTrue(all(section["stories"][kind][variant]["explicit"] for variant in ("default", "first", "even")))
            spec = {"page": {"different_first_page": True, "different_odd_even": True,
                             "page_number": {"body_format": "roman_upper", "body_start": 1}}}
            plan = compile_section_plan(spec, path)
            self.assertEqual(audit_plan_against_docx(plan, path), [])

    def test_linked_footer_inheritance_is_explicit_in_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(2)
            set_page_format(doc.sections[0], "decimal", 1)
            set_page_format(doc.sections[1], "decimal")
            add_page(doc.sections[0].footer, "ARABIC")
            self.assertTrue(doc.sections[1].footer.is_linked_to_previous)
            path = self.save(doc, td)
            evidence = inspect_docx(path)
            inherited = evidence["sections"][1]["stories"]["footer"]["default"]
            self.assertTrue(inherited["linked_to_previous"])
            self.assertEqual(len(inherited["page_fields"]), 1)
            plan = compile_section_plan({"page": {"page_number": {"body_format": "decimal"}}}, evidence)
            self.assertEqual(audit_plan_against_docx(plan, evidence), [])

    def test_preserve_existing_page_contract_does_not_rewrite_fields_or_numbering(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(2)
            for index, section in enumerate(doc.sections):
                section.footer.is_linked_to_previous = False
                set_page_format(section, "roman_upper" if index == 0 else "decimal", 1)
                add_page(section.footer, "ROMAN" if index == 0 else "ARABIC")
            source = self.save(doc, td, "source.docx")
            before = inspect_docx(source)
            plan = compile_section_plan(
                {"page": {"page_number": {"preserve_existing_locations": True}}},
                source,
            )
            self.assertTrue(plan["valid"], plan["findings"])
            self.assertEqual(plan["page_properties_policy"], "preserve_existing")
            result = execute_section_plan(doc, plan, alignment="center")
            output = self.save(doc, td, "output.docx")
            after = inspect_docx(output)
            self.assertEqual(result["status"], "applied")
            self.assertEqual(audit_plan_against_docx(plan, output), [])
            self.assertEqual(
                [section["page_number"] for section in before["sections"]],
                [section["page_number"] for section in after["sections"]],
            )
            self.assertEqual(
                [section["stories"]["footer"]["default"]["page_fields"] for section in before["sections"]],
                [section["stories"]["footer"]["default"]["page_fields"] for section in after["sections"]],
            )

    def test_missing_page_field_and_wrong_switch_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(2)
            for section in doc.sections:
                section.footer.is_linked_to_previous = False
                set_page_format(section, "decimal", 1 if section is doc.sections[0] else None)
            add_page(doc.sections[1].footer, "roman")
            path = self.save(doc, td)
            plan = compile_section_plan({"page": {"page_number": {"body_format": "decimal"}}}, path)
            findings = audit_plan_against_docx(plan, path)
            self.assertIn("page_field_count_mismatch", codes(findings))
            self.assertIn("page_field_switch_mismatch", codes(findings))
            self.assertTrue(all(item["severity"] == "error" for item in findings))

    def test_invalid_selector_and_unsupported_format_are_findings_not_guesses(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = self.save(make_sections(2), td)
            plan = compile_section_plan({"page": {"page_number": {
                "front_matter_format": "roman", "body_format": "binary",
                "body_start_selector": {"strategy": "section_index", "section_index": 99},
            }}}, path)
            self.assertFalse(plan["valid"])
            self.assertIsNone(plan["body_start_section"])
            self.assertIn("body_selector_out_of_range", codes(plan["findings"]))
            self.assertIn("page_number_format_unsupported", codes(plan["findings"]))

    def test_page_opcode_and_switch_parser_excludes_related_fields(self) -> None:
        self.assertTrue(parse_field_instruction(" PAGE \\* roman ")["is_page"])
        self.assertEqual(parse_field_instruction("PAGE \\* ROMAN")["format_switch"], "ROMAN")
        self.assertFalse(parse_field_instruction("PAGEREF target \\h")["is_page"])
        self.assertFalse(parse_field_instruction("NUMPAGES \\* ARABIC")["is_page"])

    def test_plan_is_json_serializable_and_schema_shape_is_stable(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = self.save(make_sections(1), td)
            plan = compile_section_plan({"page": {"page_number": {"body_format": "none"}}}, path)
            json.dumps(plan)
            schema = json.loads((ROOT / "schema" / "section-plan.schema.json").read_text())
            self.assertEqual(schema["title"], "SectionPlan")
            self.assertEqual(plan["schema_version"], "1.0")

    def test_alignment_only_page_rule_compiles_to_explicit_unnumbered_plan(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = self.save(make_sections(2), td)
            plan = compile_section_plan({"page": {"page_number": {"alignment": "center"}}}, path)
            self.assertTrue(plan["valid"], plan["findings"])
            self.assertEqual([section["format"] for section in plan["sections"]], ["none", "none"])
            self.assertEqual([section["numbered"] for section in plan["sections"]], [False, False])
            self.assertNotIn("page_number_format_missing", codes(plan["findings"]))

    def test_executor_rejects_invalid_or_topology_mismatched_plan(self) -> None:
        doc = make_sections(1)
        with self.assertRaisesRegex(ValueError, "invalid SectionPlan"):
            execute_section_plan(doc, {"valid": False, "sections": []})
        with tempfile.TemporaryDirectory() as td:
            path = self.save(doc, td)
            plan = compile_section_plan(
                {"page": {"page_number": {"body_format": "decimal"}}}, path)
            plan["sections"].append(dict(plan["sections"][0], section_index=2))
            with self.assertRaisesRegex(ValueError, "section count"):
                execute_section_plan(doc, plan)

    def test_executor_applies_all_footer_variants_and_audits_serialized_result(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(2)
            doc.settings.odd_and_even_pages_header_footer = True
            for section in doc.sections:
                section.different_first_page_header_footer = True
                for accessor in ("footer", "first_page_footer", "even_page_footer"):
                    story = getattr(section, accessor)
                    story.is_linked_to_previous = False
                    add_page(story, "roman", complex_field=accessor == "even_page_footer")
            source = self.save(doc, td, "source.docx")
            spec = {"page": {
                "different_first_page": True, "different_odd_even": True,
                "page_number": {
                    "front_matter_format": "none", "body_format": "decimal", "body_start": 1,
                    "body_start_selector": {"strategy": "section_index", "section_index": 2},
                },
            }}
            plan = compile_section_plan(spec, source)
            self.assertTrue(plan["valid"], plan["findings"])
            result = execute_section_plan(doc, plan, alignment="center")
            output = self.save(doc, td, "output.docx")
            self.assertEqual(result["status"], "applied")
            self.assertEqual(sum(item["removed_page_fields"] for item in result["stories"]), 6)
            self.assertEqual(audit_plan_against_docx(plan, output), [])
            evidence = inspect_docx(output)
            for variant in ("default", "first", "even"):
                self.assertEqual(
                    evidence["sections"][0]["stories"]["footer"][variant]["page_fields"], [])
                self.assertEqual(
                    len(evidence["sections"][1]["stories"]["footer"][variant]["page_fields"]), 1)

    def test_executor_replaces_page_only_textbox_with_unclipped_footer_paragraph(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            doc = make_sections(1)
            footer = doc.sections[0].footer
            paragraph = footer.paragraphs[0]
            drawing = OxmlElement("w:drawing")
            textbox = OxmlElement("w:txbxContent")
            nested = OxmlElement("w:p")
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE \\* ARABIC")
            run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "1"
            run.append(text); field.append(run); nested.append(field); textbox.append(nested)
            drawing.append(textbox); paragraph._p.append(drawing)
            source = self.save(doc, td, "source.docx")
            plan = compile_section_plan(
                {"page": {"page_number": {"body_format": "decimal", "body_start": 1}}},
                source,
            )
            result = execute_section_plan(doc, plan, alignment="center")
            output = self.save(doc, td, "output.docx")
            self.assertEqual(result["status"], "applied")
            self.assertFalse(doc.sections[0].footer._element.xpath(".//w:txbxContent"))
            self.assertFalse(doc.sections[0].footer._element.xpath(".//w:drawing"))
            self.assertEqual(len(doc.sections[0].footer._element.xpath("./w:p")), 1)
            self.assertEqual(audit_plan_against_docx(plan, output), [])


if __name__ == "__main__":
    unittest.main()
