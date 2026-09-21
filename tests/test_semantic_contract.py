#!/usr/bin/env python3
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement

from scripts.extract_semantic_metadata import extract
from scripts.inspect_docx_semantics import inspect
from scripts.role_registry import normalize_style_name, style_matches

ROOT = Path(__file__).resolve().parents[1]
PY = sys.executable


class SemanticContractTest(unittest.TestCase):
    def test_role_registry_normalizes_style_id_and_display_name(self) -> None:
        self.assertEqual(normalize_style_name("Equation Block"), normalize_style_name("equation_block"))
        self.assertTrue(style_matches("equation", style_id="EquationBlock", style_name="Equation Block"))
        self.assertTrue(style_matches("thesis_title_en", style_name="EnglishTitle"))

    def test_sample_source_semantic_metadata_has_golden_inventory(self) -> None:
        result = extract(ROOT / "tests" / "sample-thesis.tex")
        self.assertTrue(result["title_zh"])
        self.assertTrue(result["title_en"])
        self.assertEqual(result["author"], "测试学生甲")
        self.assertEqual(result["student_id"], "TEST20260001")
        self.assertEqual(result["advisor"], "测试导师乙 教授")
        self.assertEqual(result["school"], "智能交通与数据科学学院")
        self.assertEqual(result["major"], "交通信息工程及控制")
        self.assertEqual(result["first_discipline"], "交通运输工程")
        self.assertEqual(result["second_discipline"], "交通信息工程及控制")
        self.assertEqual(result["classification_number"], "U491.1")
        self.assertEqual(result["confidentiality_level"], "公开")
        self.assertEqual(result["submit_date"], "2026年6月")
        self.assertEqual(result["defense_date"], "2026年5月20日")
        self.assertEqual(result["degree_conferral_date"], "2026年6月30日")
        self.assertEqual(result["degree_category"], "academic")
        self.assertEqual(result["degree_display"], "学术学位工学硕士学位论文")
        self.assertTrue(result["abstract_zh"])
        self.assertTrue(result["abstract_en"])
        self.assertTrue(result["english_text"])
        self.assertEqual(result["keywords_zh"], ["交通流预测", "时空图卷积网络", "注意力机制", "深度学习", "智能交通系统"])
        self.assertEqual(result["keywords_en"], ["traffic flow prediction", "spatio-temporal graph convolutional network", "attention mechanism", "deep learning", "intelligent transportation system"])
        self.assertEqual(result["college"], result["school"])
        self.assertEqual(result["discipline"], result["major"])
        self.assertEqual(result["class_no"], result["classification_number"])
        self.assertEqual(result["submit_date_cn"], result["submit_date"])
        self.assertEqual(result["inventory"], {"figures": 3, "tables": 2, "display_equations": 6})
        self.assertEqual(len(result["bibliography_entries"]), 10)
        self.assertEqual(result["bibliography_entries"][0]["key"], "arellano1991")
        self.assertTrue(result["bibliography_entries"][0]["text"])

    def test_explicit_first_and_second_discipline_macros_are_independent(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "disciplines.tex"
            source.write_text(
                r"""\documentclass{article}
\firstdiscipline{交通运输工程}
\seconddiscipline{交通信息工程及控制}
""",
                encoding="utf-8",
            )
            result = extract(source)
            self.assertEqual(result["first_discipline"], "交通运输工程")
            self.assertEqual(result["second_discipline"], "交通信息工程及控制")
            self.assertEqual(result["discipline"], result["major"])

    def test_optional_v10_metadata_is_extracted_only_from_explicit_source_macros(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "metadata-v10.tex"
            source.write_text(
                r"""\documentclass{article}
\title{题目}
\englishtitle{Title}
\confidentialityperiodstart{2026-01}
\confidentialityperiodend{2028-01}
\authorpostgraduationdestination{继续深造}
\employmentunit{某研究院}
\contactphone{010-00000000}
\contactaddress{某市某路}
\postalcode{100000}
\disciplinecategory{工学}
\enterpriseadvisor{企业导师甲}
\unitcode{10008}
\unitaddress{权威配置地址}
\completiondate{2026-06}
\defensecommitteemember{委员甲}{教授}{某大学}{主席}
""",
                encoding="utf-8",
            )
            result = extract(source)
            self.assertEqual(result["confidentiality_period"], {"start": "2026-01", "end": "2028-01"})
            self.assertEqual(result["author_post_graduation_destination"], "继续深造")
            self.assertEqual(result["discipline_category"], "工学")
            self.assertEqual(result["co_supervisors"], [{"name": "企业导师甲", "kind": "enterprise"}])
            self.assertEqual(result["unit_code"], "10008")
            self.assertEqual(result["unit_address"], "权威配置地址")
            self.assertEqual(result["completion_date"], "2026-06")
            self.assertEqual(result["defense_committee"][0]["professional_title"], "教授")

    def test_docx_inspector_uses_ooxml_structure(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "semantic.docx"
            doc = Document()
            doc.styles.add_style("English Title", WD_STYLE_TYPE.PARAGRAPH)
            doc.styles.add_style("Compact", WD_STYLE_TYPE.PARAGRAPH)
            doc.styles.add_style("Equation Block", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("中文题名", "Title")
            doc.add_paragraph("English Thesis Title", "English Title")
            table = doc.add_table(rows=1, cols=2)
            for index, cell in enumerate(table.rows[0].cells):
                cell.paragraphs[0].text = f"cell-{index}"
                cell.paragraphs[0].style = "Compact"
            equation = doc.add_paragraph(style="Equation Block")
            math_para = OxmlElement("m:oMathPara")
            math_para.append(OxmlElement("m:oMath")); equation._p.append(math_para)
            doc.save(path)
            result = inspect(path)["roles"]
            self.assertEqual(result["thesis_title_zh"]["count"], 1)
            self.assertEqual(result["thesis_title_en"]["count"], 1)
            self.assertEqual(result["table_text"]["count"], 2)
            self.assertEqual(result["equation"]["count"], 1)

    def test_golden_sample_conversion_contract(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            result = subprocess.run([
                PY, "scripts/run_golden_e2e.py",
                "--source", "tests/sample-thesis.tex",
                "--expected", "tests/golden/sample-thesis.expected.json",
                "--out-dir", td,
            ], cwd=ROOT, text=True, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            manifest = json.loads((Path(td) / "golden-e2e-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["status"], "completed")
            self.assertEqual(manifest["failures"], [])
            document = Document(Path(td) / "generic.docx")
            output_text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            # These synthetic values prove that metadata extracted from the
            # TeX preamble reaches the generated DOCX, rather than merely
            # surviving in semantic-metadata.json.
            for expected_value in (
                "测试学生甲",
                "TEST20260001",
                "测试导师乙 教授",
                "测试合作导师丙 副教授",
                "智能交通与数据科学学院",
                "交通信息工程及控制",
                "交通大数据与智能计算",
                "U491.1",
                "656.13",
                "公开",
                "2026年5月20日",
                "2026年6月30日",
            ):
                self.assertIn(expected_value, output_text)


if __name__ == "__main__":
    unittest.main()
