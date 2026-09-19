from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from semantic_contract import (  # noqa: E402
    attach_request_provenance,
    request_body_sha256,
    request_envelope_sha256,
)


def run_engine(source: Path, out: Path, *extra: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "scripts/requirements_engine.py", str(source), "--out", str(out), *extra],
        cwd=ROOT, text=True, capture_output=True, check=False,
    )


class SemanticProvenanceTests(unittest.TestCase):
    def test_request_hash_domains_are_explicit_and_stable(self) -> None:
        clauses = [{"id": "C1"}]
        request = {"contract_version": "2.1", "clauses": clauses}
        bound = attach_request_provenance(
            request,
            source_sha256="a" * 64,
            evidence_doc={"evidence": []},
            clauses=clauses,
            run_id="run-1",
        )
        self.assertEqual(
            bound["provenance"]["request_sha256"], request_body_sha256(bound)
        )
        self.assertNotEqual(request_body_sha256(bound), request_envelope_sha256(bound))
        self.assertNotEqual(
            request_envelope_sha256(bound),
            request_envelope_sha256({**bound, "execution_policy": "fresh_run_no_cache"}),
        )

    def test_fresh_bound_response_is_accepted_and_evidence_context_is_rich(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "requirements.docx"
            preview = root / "preview"
            response = root / "response.json"
            accepted = root / "accepted"
            document = Document()
            paragraph = document.add_paragraph()
            run = paragraph.add_run("正文中文使用小四号宋体。")
            run.bold = True
            document.save(source)

            first = run_engine(source, preview, "--analysis-mode", "llm_primary")
            self.assertEqual(first.returncode, 0, first.stderr + first.stdout)
            clauses = json.loads((preview / "requirement-clauses.json").read_text())
            request = json.loads((preview / "llm-request.json").read_text())
            evidence_context = json.loads((preview / "evidence-context.json").read_text())
            self.assertIn("style_id", evidence_context["evidence"][0])
            self.assertEqual(evidence_context["evidence"][0]["runs"][0]["format"]["bold"], True)
            self.assertNotIn("evidence_context", request["clauses"][0])
            self.assertIn(clauses[0]["evidence_ids"][0], request["evidence_context"])
            response.write_text(json.dumps({
                "contract_version": "2.1",
                "provenance": request["provenance"],
                "requirements": [{
                    "role": "body_text",
                    "properties": {"font": {"cjk": "SimSun", "size_pt": 12}},
                    "clause_ids": [clauses[0]["id"]],
                    "evidence_ids": clauses[0]["evidence_ids"],
                    "confidence": 0.99,
                    "reason": "fresh bound response",
                }],
                "clause_reviews": [{
                    "clause_id": clauses[0]["id"],
                    "classification": "executable",
                    "requirement_indexes": [0],
                    "reason": "fresh bound response",
                }],
                "unsupported_items": [],
                "reported_conflicts": [],
            }, ensure_ascii=False), encoding="utf-8")
            accepted_result = run_engine(
                source, accepted, "--analysis-mode", "llm_primary",
                "--strict-provenance", "--run-id", request["provenance"]["run_id"],
                "--llm-response", str(response),
            )
            self.assertEqual(accepted_result.returncode, 0, accepted_result.stderr + accepted_result.stdout)
            accepted_spec = json.loads((accepted / "format-spec.json").read_text())
            self.assertEqual(accepted_spec["status"], "semantic_resolved")
            self.assertTrue(accepted_spec["semantic_review_provenance_valid"])

    def test_stale_bound_response_is_rejected_in_strict_mode(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "requirements.docx"
            preview = root / "preview"
            response = root / "response.json"
            rejected = root / "rejected"
            document = Document()
            document.add_paragraph("正文中文使用小四号宋体。")
            document.save(source)
            first = run_engine(source, preview, "--analysis-mode", "llm_primary")
            self.assertEqual(first.returncode, 0, first.stderr + first.stdout)
            clauses = json.loads((preview / "requirement-clauses.json").read_text())
            request = json.loads((preview / "llm-request.json").read_text())
            stale = request["provenance"].copy()
            stale["clause_sha256"] = "0" * 64
            response.write_text(json.dumps({
                "contract_version": "2.1", "provenance": stale,
                "requirements": [],
                "clause_reviews": [{
                    "clause_id": clauses[0]["id"], "classification": "informational",
                    "requirement_indexes": [], "reason": "stale response",
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }, ensure_ascii=False), encoding="utf-8")
            result = run_engine(
                source, rejected, "--analysis-mode", "llm_primary",
                "--strict-provenance", "--llm-response", str(response),
            )
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            spec = json.loads((rejected / "format-spec.json").read_text())
            self.assertEqual(spec["status"], "needs_clarification")
            self.assertFalse(spec["semantic_review_provenance_valid"])
            conflicts = json.loads((rejected / "conflicts.json").read_text())
            self.assertIn("provenance_clause_sha256_mismatch", {
                item["reason"] for item in conflicts if item.get("type") == "llm_provenance"
            })

    def test_same_input_response_is_rejected_without_explicit_current_run_binding(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "requirements.docx"
            packet = root / "packet"
            response = root / "response.json"
            rejected = root / "rejected"
            document = Document()
            document.add_paragraph("正文中文使用小四号宋体。")
            document.save(source)

            first = run_engine(source, packet, "--analysis-mode", "llm_primary")
            self.assertEqual(first.returncode, 0, first.stderr + first.stdout)
            clauses = json.loads((packet / "requirement-clauses.json").read_text())
            request = json.loads((packet / "llm-request.json").read_text())
            response.write_text(json.dumps({
                "contract_version": "2.1", "provenance": request["provenance"],
                "requirements": [],
                "clause_reviews": [{
                    "clause_id": clauses[0]["id"], "classification": "informational",
                    "requirement_indexes": [], "reason": "response belongs to the first run",
                }],
                "unsupported_items": [], "reported_conflicts": [],
            }, ensure_ascii=False), encoding="utf-8")

            result = run_engine(
                source, rejected, "--analysis-mode", "llm_primary",
                "--strict-provenance", "--llm-response", str(response),
            )
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            conflicts = json.loads((rejected / "conflicts.json").read_text())
            reasons = {item["reason"] for item in conflicts if item.get("type") == "llm_provenance"}
            self.assertIn("provenance_run_id_mismatch", reasons)


if __name__ == "__main__":
    unittest.main()
