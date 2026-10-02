"""Captured placeholder-only duplicates preserve current print and human work."""
import copy
import hashlib
import io
import json
from pathlib import Path
import sys
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import host_agent_bridge as bridge
from declaration_signature_projection import project_signature_only_declarations
from requirements_engine import build_llm_request, merge_llm_primary
from resource_registry import materialize_declaration_resources, resource_items
from apply_format_spec import apply_declarations, audit_declarations
from docx import Document


def incident():
    case = json.loads((ROOT / "tests/fixtures/signature-placeholder-incident.json").read_text())
    source = case["source"]
    chunk = build_llm_request([], source["clauses"],
        {"evidence": list(source["evidence_context"].values())}, {}, "full", contract_version="3.0")
    chunk.update(source)
    return case["attempts"], chunk


def compiled_parent():
    attempts, chunk = incident()
    parent = bridge.normalize_native_response(attempts[1], chunk["response_schema"])
    parent, _ = bridge._materialize_fixed_declaration_source_text(parent, chunk)
    return parent, chunk


class SignaturePlaceholderProjectionTests(unittest.TestCase):
    def test_captured_both_attempts_pass_production_preparation_without_discharging_humans(self):
        attempts, chunk = incident()
        for raw in attempts:
            with self.subTest(attempt=attempts.index(raw)):
                frozen = copy.deepcopy(raw)
                candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
                repairs = [r for r in audit["mechanical_repairs"]
                           if r["rule_id"] == "source_bound_signature_block_projection_v2"]
                self.assertEqual(len(repairs), 1)
                proof = repairs[0]["proofs"][0]
                self.assertEqual(proof["removed_print_representation"], "source_bound_placeholders")
                self.assertEqual(proof["verification_mode"], "static_docx")
                self.assertTrue(proof["all_removed_print_operations_already_retained"])
                self.assertTrue(repairs[0]["independent_review_required"])
                self.assertFalse(repairs[0]["submission_ready"])
                self.assertEqual(raw, frozen)
                self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
                self.assertNotIn(raw["requirements"][1]["properties"]["items"][0]["id"],
                    [item["id"] for r in candidate["requirements"] for item in r["properties"]["items"]])
                original = raw["clause_reviews"][-1]
                retained = next(r for r in candidate["clause_reviews"] if r["clause_id"] == original["clause_id"])
                self.assertEqual(retained["classification"], "external_compliance")
                self.assertEqual(len(retained["obligations"]), 2)
                for old, new in zip(original["obligations"], retained["obligations"]):
                    for field in ("id", "actor", "action", "target", "condition", "force", "applicability", "status", "route"):
                        self.assertEqual(old.get(field), new.get(field))
                    self.assertEqual(new["status"], "unverifiable")
                    self.assertEqual(new["route"], "human")

    def test_only_redundant_print_objects_removed_reviews_and_owner_payload_frozen(self):
        parent, chunk = compiled_parent(); frozen = copy.deepcopy(parent)
        candidate, audit = project_signature_only_declarations(parent, chunk, validate=bridge.validate_host_agent_response)
        self.assertIsNotNone(candidate)
        self.assertEqual(candidate["clause_reviews"], parent["clause_reviews"])
        self.assertEqual(len(candidate["requirements"]), 1)
        expected = copy.deepcopy(parent["requirements"][0])
        expected["clause_ids"] = [c["id"] for c in chunk["clauses"]
                                  if c["id"] in set(parent["requirements"][0]["clause_ids"] + parent["requirements"][2]["clause_ids"])]
        expected["evidence_ids"] = list(dict.fromkeys(eid for c in chunk["clauses"]
            if c["id"] in expected["clause_ids"] for eid in c["evidence_ids"]))
        self.assertEqual(candidate["requirements"], [expected])
        self.assertEqual(parent, frozen)
        self.assertEqual(audit[0]["source_response_sha256"], bridge._response_sha256(parent))
        self.assertEqual(audit[0]["source_chunk_sha256"], bridge._response_sha256(chunk))

    def test_nonmaterialized_heading_is_rejected_or_bound_to_unique_existing_print_entity(self):
        for change in ("valid", "no_owner", "foreign_source", "extra_placeholder", "different_text",
                       "registered_checker", "source_hash", "condition", "existing_id", "pending_heading"):
            parent, chunk = compiled_parent()
            parent["requirements"].pop(1)  # No signature duplication: same heading-only path.
            heading = parent["requirements"][1]
            item = heading["properties"]["items"][0]
            if change == "no_owner": parent["requirements"].pop(0)
            elif change == "foreign_source": item["source_evidence_ids"] = ["foreign"]
            elif change == "extra_placeholder": item["signature_placeholders"] = [{"role":"supervisor", "label":"导师签名", "attestation_scope":"placeholder_presence_only"}]
            elif change == "different_text": item["heading"] = "Different declaration"
            elif change == "registered_checker": heading["verification"]["checker_ids"] = ["independent_check"]
            elif change == "source_hash": chunk["clauses"][0]["source_span"]["source_sha256"] = "0" * 64
            elif change == "condition": heading["applicability"] = {"status":"not_applicable", "conditions":[], "exceptions":[]}
            elif change == "existing_id": heading["existing_requirement_id"] = "R-existing"
            elif change == "pending_heading": parent["clause_reviews"][0]["obligations"][0]["status"] = "unverifiable"
            with self.subTest(change=change):
                errors = bridge.validate_host_agent_response(parent, chunk)
                self.assertTrue(errors)
                if change == "valid":
                    self.assertTrue(any("declaration_source_text_not_materialized" in e for e in errors))
                candidate, audit = project_signature_only_declarations(parent, chunk, validate=bridge.validate_host_agent_response)
                if change == "valid":
                    self.assertIsNotNone(candidate)
                    self.assertEqual(candidate["clause_reviews"], parent["clause_reviews"])
                    self.assertEqual(audit[0]["rule_id"], "source_bound_declaration_heading_coalescence_v1")
                    self.assertTrue(audit[0]["proofs"][0]["all_edges_previously_supplied"])
                else: self.assertIsNone(candidate)

    def test_captured_whole_candidates_merge_without_declaration_conflicts(self):
        attempts, chunk = incident()
        for raw in attempts:
            with self.subTest(attempt=attempts.index(raw)):
                candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
                spec, conflicts, _ = merge_llm_primary(Path("current.docx"), {}, chunk["clauses"], candidate,
                    set(chunk["evidence_context"]))
                self.assertEqual(conflicts, [])
                self.assertEqual(len(spec["declarations"]["items"]), 1)
                self.assertEqual(spec["declarations"]["items"][0]["source_signature_lines"],
                    candidate["requirements"][0]["properties"]["items"][0]["source_signature_lines"])
                materialize_declaration_resources(spec, "whole-current-candidate", evidence={"evidence": list(chunk["evidence_context"].values())})
                doc = Document(); doc.add_paragraph("摘要", "Heading 1")
                apply_declarations(doc, spec["declarations"], resource_items(spec))
                buffer = io.BytesIO(); doc.save(buffer); saved = Document(buffer)
                self.assertEqual(audit_declarations(saved, spec["declarations"], resource_items(spec)), [])

    def test_both_body_and_placeholder_forms_work_with_external_or_static_presence_checks(self):
        for mode in ("external", "static_docx"):
            for body in (False, True):
                parent, chunk = compiled_parent(); target = parent["requirements"][1]
                target["verification"]["mode"] = mode
                if body:
                    target["properties"]["items"][0]["body_parts"] = [chunk["evidence_context"][target["evidence_ids"][0]]["text"]]
                with self.subTest(mode=mode, body=body):
                    candidate, _ = project_signature_only_declarations(parent, chunk, validate=bridge.validate_host_agent_response)
                    self.assertIsNotNone(candidate)

    def test_unproved_print_or_pending_work_cannot_be_deleted(self):
        cases = ("no_owner", "two_owners", "foreign_hash", "foreign_evidence", "location",
            "filled_line", "foreign_label", "empty_placeholders", "extra_body", "existing_id",
            "field_key", "condition", "prerequisite", "registered_checker", "different_anchor",
            "covered_atom", "empty_atoms", "changed_quote", "conflict", "unrelated_error")
        for change in cases:
            parent, chunk = compiled_parent(); target = parent["requirements"][1]
            item = target["properties"]["items"][0]
            review = parent["clause_reviews"][-1]
            line = parent["requirements"][0]["properties"]["items"][0]["source_signature_lines"][0]
            if change == "no_owner": parent["requirements"].pop(0)
            elif change == "two_owners":
                twin = copy.deepcopy(parent["requirements"][0]); twin["properties"]["items"][0]["id"] += "-twin"
                parent["requirements"].append(twin)
            elif change == "foreign_hash": line["source_sha256"] = "0" * 64
            elif change == "foreign_evidence": item["source_evidence_ids"] = ["foreign"]
            elif change == "location": chunk["evidence_context"][line["source_evidence_id"]]["location"]["child_index"] += 10
            elif change == "filled_line": chunk["evidence_context"][line["source_evidence_id"]]["text"] += " 张三"
            elif change == "foreign_label": item["signature_placeholders"][0]["label"] = "导师已批准"
            elif change == "empty_placeholders": item["signature_placeholders"] = []
            elif change == "extra_body": item["body_parts"] = ["独立的不可删除正文"]
            elif change == "existing_id": target["existing_requirement_id"] = "R-current-existing"
            elif change == "field_key": target["field_key"] = "independent-instance"
            elif change == "condition": target["applicability"] = {"status": "not_applicable", "conditions": [], "exceptions": []}
            elif change == "prerequisite": target["input_prerequisites"] = [{"kind": "metadata", "key": "author_name"}]
            elif change == "registered_checker": target["verification"]["checker_ids"] = ["verify_external_approval"]
            elif change == "different_anchor": target["properties"]["before_role"] = "document_start"
            elif change == "covered_atom": review["obligations"][0]["status"] = "covered"
            elif change == "empty_atoms": review["obligations"] = []
            elif change == "changed_quote": review["obligations"][0]["source_quote"] = "不存在的签署语句"
            elif change == "conflict": parent["reported_conflicts"] = [{"reason": "conflicting owner"}]
            else: parent["requirements"][0]["properties"]["unknown_operation"] = True
            frozen = copy.deepcopy(parent)
            with self.subTest(change=change):
                candidate, audit = project_signature_only_declarations(parent, chunk, validate=bridge.validate_host_agent_response)
                self.assertIsNone(candidate); self.assertEqual(audit, [])
                self.assertEqual(parent, frozen)

    def test_stale_or_incomplete_feedback_never_authorizes_production_repair(self):
        parent, chunk = compiled_parent()
        records = bridge.contract_error_records(bridge.validate_host_agent_response(parent, chunk), response=parent, chunk=chunk)
        for changed in ([], [{**records[0], "response_sha256": "0" * 64}]):
            with self.subTest(records=changed):
                self.assertIsNone(bridge._apply_safe_mechanical_repairs(parent, changed, chunk=chunk)[0])

    def test_identical_text_at_another_physical_source_is_not_the_same_owner(self):
        parent, chunk = compiled_parent()
        mapping = {c["id"]: "second-" + c["id"] for c in chunk["clauses"]}
        mapping.update({eid: "second-" + eid for eid in chunk["evidence_context"]})
        def rename(value):
            if isinstance(value, dict): return {mapping.get(k, k): rename(v) for k, v in value.items()}
            if isinstance(value, list): return [rename(v) for v in value]
            return mapping.get(value, value) if isinstance(value, str) else value
        clones = rename(copy.deepcopy(chunk["clauses"]))
        evidence = rename(copy.deepcopy(chunk["evidence_context"]))
        reviews = rename(copy.deepcopy(parent["clause_reviews"]))
        for c in clones:
            for location in (c["location"], c["source_span"]["location"]):
                location["child_index"] += 100
                location["order"] += 100
        for record in evidence.values():
            record["location"]["child_index"] += 100
            record["location"]["order"] += 100
        for review in reviews:
            for atom in review.get("obligations", []): atom["id"] = "second-" + atom["id"]
        chunk["clauses"].extend(clones); chunk["evidence_context"].update(evidence)
        parent["clause_reviews"].extend(reviews)
        other = rename(copy.deepcopy(parent["requirements"][0]))
        other["properties"]["items"][0]["id"] = "second-print-instance"
        other["clause_ids"].insert(0, clones[0]["id"])
        other["evidence_ids"].insert(0, clones[0]["evidence_ids"][0])
        parent["requirements"].append(other)
        # Remove the original complete owner. The second one has byte-identical
        # heading/body/blank line, but different current evidence and locations.
        parent["requirements"].pop(0)
        errors = bridge.validate_host_agent_response(parent, chunk)
        self.assertTrue(any("declaration_source_text_not_materialized" in e for e in errors))
        self.assertFalse(any("unknown" in e or "duplicate" in e for e in errors), errors)
        candidate, audit = project_signature_only_declarations(parent, chunk, validate=bridge.validate_host_agent_response)
        self.assertIsNone(candidate); self.assertEqual(audit, [])

    def test_renamed_source_and_altered_current_blank_spacing_not_id_special_cases(self):
        attempts, chunk = incident(); raw = attempts[1]
        remap = {c["id"]: f"current-{i}" for i, c in enumerate(chunk["clauses"])}
        remap.update({eid: f"evidence-{i}" for i, eid in enumerate(chunk["evidence_context"])})
        def rename(value):
            if isinstance(value, dict): return {remap.get(k, k): rename(v) for k, v in value.items()}
            if isinstance(value, list): return [rename(v) for v in value]
            return remap.get(value, value) if isinstance(value, str) else value
        source = rename({k:chunk[k] for k in ("clauses", "evidence_context", "declaration_anchor_preference")})
        raw = rename(raw); clause = source["clauses"][-1]; eid = clause["evidence_ids"][0]
        text = source["evidence_context"][eid]["text"].replace("                       ", "        ")
        source["evidence_context"][eid]["text"] = text
        clause["source_span"].update(text=text, end_offset=len(text), source_sha256=hashlib.sha256(text.encode()).hexdigest())
        rebuilt = build_llm_request([], source["clauses"], {"evidence": list(source["evidence_context"].values())}, {}, "full", contract_version="3.0")
        rebuilt.update(source)
        candidate, audit = bridge.prepare_native_response_candidate(raw, rebuilt)
        proof = next(r for r in audit["mechanical_repairs"] if r["rule_id"] == "source_bound_signature_block_projection_v2")["proofs"][0]
        self.assertEqual(proof["source_signature_lines"][0]["text"], text)
        self.assertNotIn(clause["id"], [cid for r in candidate["requirements"] for cid in r["clause_ids"]])

    def test_retained_resource_prints_blank_line_once_after_serialization(self):
        attempts, chunk = incident(); candidate, _ = bridge.prepare_native_response_candidate(attempts[1], chunk)
        # Isolate the retained complete resource, not the whole BSU pipeline.
        retained = next(r for r in candidate["requirements"] if r["properties"]["items"][0].get("source_signature_lines"))
        spec = {"schema_version": "1.0", "source_document": "current.docx", "status": "semantic_resolved",
                "roles": {}, "requirements": [], "declarations": copy.deepcopy(retained["properties"])}
        materialize_declaration_resources(spec, "current-signature-projection", evidence={"evidence": list(chunk["evidence_context"].values())})
        doc = Document(); doc.add_paragraph("摘要", "Heading 1")
        apply_declarations(doc, spec["declarations"], resource_items(spec))
        buffer = io.BytesIO(); doc.save(buffer); saved = Document(buffer)
        line = retained["properties"]["items"][0]["source_signature_lines"][0]["text"]
        self.assertEqual([p.text for p in saved.paragraphs].count(line), 1)
        self.assertEqual(audit_declarations(saved, spec["declarations"], resource_items(spec)), [])


if __name__ == "__main__": unittest.main()
