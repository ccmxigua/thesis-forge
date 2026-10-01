"""Raw native incident -> source preservation -> serialized DOCX audit."""
import copy
import hashlib
import io
import json
from pathlib import Path
import sys
import unittest
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from docx import Document
import host_agent_bridge as bridge
from requirements_engine import build_llm_request
from resource_registry import materialize_declaration_resources, resource_items
from format_spec_validation import load_and_validate
from apply_format_spec import apply_declarations, audit_declarations


def incident():
    case = json.loads((ROOT / 'tests/fixtures/signature-source-incident.json').read_text())
    source = case['chunk']
    chunk = build_llm_request([], source['clauses'], {'evidence': list(source['evidence_context'].values())}, {}, 'full', contract_version='3.0')
    chunk.update(source)
    return case['candidate'], chunk


class SignatureSourceLinesTests(unittest.TestCase):
    def test_captured_response_preserves_line_and_pending_actions(self):
        raw, chunk = incident(); frozen = copy.deepcopy(raw)
        candidate, audit = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(raw, frozen)
        normalized = bridge.normalize_native_response(raw, chunk['response_schema'])
        self.assertEqual(candidate['clause_reviews'], normalized['clause_reviews'])
        self.assertEqual(len(candidate['requirements']), 1)
        kept = candidate['requirements'][0]
        self.assertEqual(kept['clause_ids'], raw['requirements'][0]['clause_ids'])
        lines = kept['properties']['items'][0]['source_signature_lines']
        self.assertEqual([l['text'] for l in lines], raw['requirements'][1]['properties']['items'][0]['body_parts'])
        self.assertEqual(lines[0]['attestation_scope'], 'placeholder_presence_only')
        self.assertEqual(bridge.validate_host_agent_response(candidate, chunk), [])
        self.assertIn('source_bound_signature_block_projection_v1', json.dumps(audit))

    def test_current_resource_and_serialized_docx_have_exact_spaces_once(self):
        raw, chunk = incident()
        candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        spec = {'schema_version':'1.0','source_document':'current.docx','status':'semantic_resolved',
            'roles':{},'requirements':[], 'declarations':candidate['requirements'][0]['properties']}
        evidence = {'evidence': list(chunk['evidence_context'].values())}
        spec = materialize_declaration_resources(spec, 'signature-current-run', evidence=evidence)
        self.assertEqual(load_and_validate(spec, ROOT / 'schema/format-spec.schema.json'), [])
        doc = Document(); doc.add_paragraph('摘要', 'Heading 1')
        counts = apply_declarations(doc, spec['declarations'], resource_items(spec))
        memory = io.BytesIO(); doc.save(memory); reopened = Document(memory)
        self.assertEqual(audit_declarations(reopened, spec['declarations'], resource_items(spec)), [])
        line = next(iter(resource_items(spec).values()))['source_signature_lines'][0]['text']
        self.assertEqual([p.text for p in reopened.paragraphs].count(line), 1)
        self.assertEqual(counts['placeholders_written'], 1)
        signature = next(p for p in reopened.paragraphs if p.text == line)
        signature.text = line.replace('                       ', ' ')
        self.assertTrue(audit_declarations(reopened, spec['declarations'], resource_items(spec)))
        changed = copy.deepcopy(spec)
        next(iter(resource_items(changed).values()))['source_signature_lines'][0]['text'] += '张三'
        with self.assertRaises(ValueError):
            materialize_declaration_resources(changed, 'signature-current-run', evidence=evidence)
        changed = copy.deepcopy(spec)
        changed['declarations']['items'][0]['source_signature_lines'] = []
        self.assertTrue(load_and_validate(changed, ROOT / 'schema/format-spec.schema.json'))
        with self.assertRaises(ValueError):
            materialize_declaration_resources(changed, 'signature-current-run', evidence=evidence)
        with self.assertRaises(ValueError):
            apply_declarations(reopened, changed['declarations'], resource_items(changed))

    def test_extra_blank_placeholder_and_changed_physical_source_are_preserved_or_rejected(self):
        raw, chunk = incident()
        raw['requirements'][0]['properties']['items'][0]['signature_placeholders'] = [
            {'role':'supervisor', 'label':'导师签名', 'attestation_scope':'placeholder_presence_only'}]
        candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        spec = {'schema_version':'1.0','source_document':'current.docx','status':'semantic_resolved',
            'roles':{},'requirements':[], 'declarations':candidate['requirements'][0]['properties']}
        evidence = {'evidence':list(chunk['evidence_context'].values())}
        materialize_declaration_resources(spec, 'extra-label-current-run', evidence=evidence)
        doc = Document(); doc.add_paragraph('摘要', 'Heading 1')
        apply_declarations(doc, spec['declarations'], resource_items(spec))
        self.assertIn('导师签名：________________', [p.text for p in doc.paragraphs])
        self.assertEqual(audit_declarations(doc, spec['declarations'], resource_items(spec)), [])
        changed = copy.deepcopy(evidence)
        changed['evidence'][-1]['location']['child_index'] += 10
        with self.assertRaisesRegex(ValueError, 'not adjacent'):
            materialize_declaration_resources(copy.deepcopy(spec), 'extra-label-current-run', evidence=changed)

    def test_projection_rejects_unproved_ownership_payload_or_duties(self):
        for change in ('foreign_hash','different_table','no_heading','two_owners','completed_signature',
                       'extra_operation','scope','empty_pending','covered_pending','changed_quote','stale_feedback'):
            with self.subTest(change=change):
                raw, chunk = incident()
                target = raw['requirements'][1]
                signature = chunk['clauses'][-1]
                if change == 'foreign_hash': signature['source_span']['source_sha256'] = '0' * 64
                elif change == 'different_table': chunk['evidence_context'][signature['evidence_ids'][0]]['location']['child_index'] += 1
                elif change == 'no_heading': raw['requirements'].pop(0)
                elif change == 'two_owners':
                    other = copy.deepcopy(raw['requirements'][0])
                    other['properties']['items'][0]['id'] = 'different-complete-owner'
                    raw['requirements'].insert(0, other)
                elif change == 'completed_signature': target['properties']['items'][0]['body_parts'][0] += ' 张三'
                elif change == 'extra_operation': target['field_key'] = 'unique-user-instance'
                elif change == 'scope': target['applicability'] = {'status':'not_applicable','conditions':[],'exceptions':[]}
                elif change == 'empty_pending': raw['clause_reviews'][-1]['obligations'] = []
                elif change == 'covered_pending': raw['clause_reviews'][-1]['obligations'][0]['status'] = 'covered'
                elif change == 'changed_quote': raw['clause_reviews'][-1]['obligations'][0]['source_quote'] = '无关内容'
                else:
                    candidate, _ = bridge._materialize_fixed_declaration_source_text(raw, chunk)
                    errors = bridge.validate_host_agent_response(candidate, chunk)
                    records = bridge.contract_error_records(errors, response=candidate, chunk=chunk)
                    records[0]['response_sha256'] = '0' * 64
                    self.assertIsNone(bridge._apply_safe_mechanical_repairs(candidate, records, chunk=chunk)[0])
                    continue
                with self.assertRaises(ValueError): bridge.prepare_native_response_candidate(raw, chunk)

    def test_different_ids_and_current_line_are_not_a_bsu_special_case(self):
        raw, chunk = incident()
        # Re-extract a separate current source with a different author label.
        old = chunk['evidence_context'][chunk['clauses'][-1]['evidence_ids'][0]]['text']
        new = '研究生签字：            年   月   日'
        encoded = json.dumps({'raw':raw,'source':{'clauses':chunk['clauses'],'evidence_context':chunk['evidence_context']}},ensure_ascii=False)
        encoded = encoded.replace(old, new).replace('学位论文作者签名','研究生签字')
        for n in range(33, 39): encoded = encoded.replace(f'C{n:05}', f'current-clause-{n}')
        for n in range(33, 36): encoded = encoded.replace(f'E{n:05}', f'current-evidence-{n}')
        data = json.loads(encoded); source = data['source']
        c = source['clauses'][-1]; c['source_span']['end_offset'] = len(new)
        c['source_span']['source_sha256'] = hashlib.sha256(new.encode('utf-8')).hexdigest()
        c['source_span']['text'] = new
        raw = data['raw']; target = raw['requirements'][-1]['properties']['items'][0]
        target['signature_placeholders'] = [{'role':'author','label':'研究生签字','attestation_scope':'placeholder_presence_only'},
            {'role':'date','label':'年 月 日','attestation_scope':'placeholder_presence_only'}]
        chunk = build_llm_request([], source['clauses'], {'evidence':list(source['evidence_context'].values())}, {}, 'full', contract_version='3.0')
        chunk.update(source); chunk['declaration_anchor_preference'] = 'abstract_title_zh'
        candidate, _ = bridge.prepare_native_response_candidate(raw, chunk)
        self.assertEqual(candidate['requirements'][0]['properties']['items'][0]['source_signature_lines'][0]['text'], new)

if __name__ == '__main__': unittest.main()
