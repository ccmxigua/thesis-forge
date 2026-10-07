"""Offline observed disagreement, production replay and non-release consumers."""
from __future__ import annotations

import copy
from contextlib import ExitStack
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import native_semantic_review as native
import host_agent_bridge as bridge
import thesis_format_pipeline as pipeline
from source_inventory_dispute import inventory_existence_dispute, build_inventory_existence_disputes
from source_classification_dispute import (
    informational_uncertainty_dispute,
    template_example_alignment_dispute,
    build_informational_uncertainty_disputes,
)
from semantic_contract import sha256_json, request_body_sha256
from tests import test_empty_inventory_verdict as transport
from tests import test_independent_retry_scope as wire
from tests import test_host_agent_bridge as packets
from tests import test_manual_review as manual_fixtures
from manual_review import build_manual_review_ledger
from manual_review_display import append_manual_review_markers, audit_manual_review_markers
from draft_scorecard import build_scorecard, append_scorecard, audit_scorecard
from format_spec_validation import load_and_validate
from docx import Document
import apply_format_spec


class InventoryDisputeTests(unittest.TestCase):
    def setUp(self):
        self.fixture = json.loads((ROOT / "tests/fixtures/source-inventory-existence-dispute.json").read_text())
        self.check = self.fixture["checks"][0]
        self.result = self.fixture["response"]["results"][0]

    def validate(self, check=None, result=None, *, draft=True):
        return native.validate_obligation_coverage_response(
            {"results": [copy.deepcopy(result or self.result)]}, [copy.deepcopy(check or self.check)],
            allow_draft_disputes=draft,
        )

    def test_observed_no_duty_assessment_retained_without_rewriting_either_party(self):
        before = copy.deepcopy((self.check, self.result))
        self.assertEqual(self.validate(), [self.result])
        dispute = inventory_existence_dispute(self.check, self.result)
        self.assertIsNotNone(dispute)
        self.assertEqual(dispute["primary_review_context"], self.check["review_context"])
        self.assertEqual(dispute["independent_result"], self.result)
        self.assertFalse(dispute["coverage_complete"])
        self.assertFalse(dispute["execution_authorized"])
        self.assertFalse(dispute["submission_ready"])
        self.assertEqual(before, (self.check, self.result))
        with self.assertRaises(native.NativeSemanticReviewError): self.validate(draft=False)

    def test_bad_bindings_inventories_and_structural_errors_are_not_disputes(self):
        def binding(c): return c["review_context"]["primary_obligation_quote_bindings"][0]
        mutations = {
            "old hash": lambda c,r: binding(c)["source_binding"]["clause_binding"]["source_fragments"][0].update(source_sha256="0"*64),
            "wrong evidence": lambda c,r: binding(c)["source_binding"]["clause_binding"]["source_fragments"][0].update(evidence_id="foreign"),
            "wrong occurrence": lambda c,r: binding(c)["source_binding"].update(quote_start_offset=1),
            "wrong literal": lambda c,r: c["review_context"]["primary_obligations"][0].update(source_quote="日期"),
            "duplicate primary": lambda c,r: c["review_context"]["primary_obligations"].append(copy.deepcopy(c["review_context"]["primary_obligations"][0])),
            "missing primary": lambda c,r: c["review_context"].update(primary_obligations=[]),
            "covered": lambda c,r: c["review_context"]["primary_obligations"][0].update(status="covered"),
            "automatic": lambda c,r: c["review_context"]["primary_obligations"][0].update(route="automatic"),
            "unknown force": lambda c,r: c["review_context"]["primary_obligations"][0].update(force="unknown"),
            "mixed requirement": lambda c,r: c["review_context"].update(linked_requirements=[{"requirement_ref":"rr"}]),
            "machine fact": lambda c,r: c["review_context"].update(machine_obligation_ids=["known"]),
            "empty rationale": lambda c,r: r.update(rationale=" "),
            "foreign quote": lambda c,r: r.update(evidence_quotes=["unrelated"]),
            "foreign clause": lambda c,r: r.update(check_id="foreign"),
            "uncertain empty": lambda c,r: r.update(verdict="uncertain"),
            "fake field": lambda c,r: r.update(submission_ready=True),
        }
        for name, mutate in mutations.items():
            with self.subTest(name=name):
                check, result = copy.deepcopy(self.check), copy.deepcopy(self.result)
                mutate(check, result)
                with self.assertRaises(native.NativeSemanticReviewError): self.validate(check, result)
        with self.assertRaises(native.NativeSemanticReviewError):
            native.validate_obligation_coverage_response({"results":[self.result,self.result]}, [self.check], allow_draft_disputes=True)

    def test_invalid_sibling_not_hidden_by_dispute(self):
        check, result = copy.deepcopy(self.check), copy.deepcopy(self.result)
        check["check_id"] = result["check_id"] = "sibling"
        result["evidence_quotes"] = ["bad"]
        with self.assertRaises(native.NativeSemanticReviewError):
            native.validate_obligation_coverage_response({"results":[self.result,result]}, [self.check,check], allow_draft_disputes=True)

    def test_envelope_reconstructs_dispute_and_blocks_forged_complete_or_policy(self):
        disputes = build_inventory_existence_disputes([self.check], [self.result])
        request = {"checks":[self.check],"output_policy":"review_draft"}
        envelope = {"results":[self.result],"source_inventory_disputes":disputes,
                    "status":"completed_with_disputes","coverage_complete":False,"submission_ready":False}
        native.validate_draft_dispute_envelope(envelope,request,output_policy="review_draft")
        for key,value in (("status","completed"),("coverage_complete",True),("submission_ready",True),
                          ("source_inventory_disputes",[]),("results",[])):
            forged=copy.deepcopy(envelope);forged[key]=value
            with self.subTest(key=key),self.assertRaises(ValueError):
                native.validate_draft_dispute_envelope(forged,request,output_policy="review_draft")
        for policy in ("submission","supported_subset"):
            with self.subTest(policy=policy),self.assertRaises(ValueError):
                native.validate_draft_dispute_envelope(envelope,request,output_policy=policy)
        with self.assertRaises(ValueError):
            pipeline.enforce_obligation_review_output_policy([self.result],output_policy="submission",source_inventory_disputes=disputes)

    def make_packet(self,directory,source):
        _,chunk=packets.HostAgentBridgeTests()._packet(directory,source=source,contract_version="3.0")
        context=copy.deepcopy(self.check["review_context"])
        atom=context["primary_obligations"][0]
        atom.update(id="arbitrary-human-duty",source_quote=source,action="Complete this field",target="Form field")
        candidate={"contract_version":"3.0","provenance":chunk["provenance"],"requirements":[],
            "clause_reviews":[{"clause_id":"C1","classification":"external_compliance","normative_basis":"external_duty",
                               "reason":"Primary proposes a human duty.","obligations":[atom]}],
            "unsupported_items":[],"reported_conflicts":[]}
        request=native.build_obligation_coverage_request(candidate,chunk,run_id=chunk["provenance"]["run_id"],chunk_index=1)
        request.update(attempt=1,provider_attempt=1,output_policy="review_draft")
        result={**copy.deepcopy(self.result),"check_id":"C1","evidence_quotes":[source],
                "rationale":"The source is a label without an explicit action."}
        return chunk,candidate,request,result

    def test_general_nondate_sources_and_current_known_facts(self):
        for source,allowed in (
                ("联系电话：",True),("Reviewer code:",True),
                ("未经批准的均为公开学位论文（公开的学位论文本项为空白）",False),
                ("关键词须源自论文",False),
                ("The following English is not correct.",False)):
            with self.subTest(source=source),tempfile.TemporaryDirectory() as tmp:
                chunk,candidate,request,result=self.make_packet(Path(tmp),source)
                if allowed:
                    self.assertIsNotNone(inventory_existence_dispute(request["checks"][0],result))
                    native.validate_obligation_coverage_response({"results":[result]},request["checks"],allow_draft_disputes=True)
                else:
                    self.assertIsNone(inventory_existence_dispute(request["checks"][0],result))
                    with self.assertRaises(native.NativeSemanticReviewError):
                        native.validate_obligation_coverage_response({"results":[result]},request["checks"],allow_draft_disputes=True)

    def test_production_bridge_both_consumers_manual_docx_and_resealed_tamper(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory=Path(tmp)/"packet"
            chunk,candidate,request,result=self.make_packet(directory,"年    月    日")
            before=copy.deepcopy(candidate)
            raw=transport.enveloped(wire.RetryScopeTests().wire({"results":[result]},request))
            with ExitStack() as stack:
                transport.EmptyInventoryVerdictTests().mock_transport(stack,[raw])
                pointer=bridge._run_independent_obligation_coverage_review(candidate,chunk,review_dir=directory,
                    run_id=request["run_id"],chunk_index=1,attempt=1,host_runtime="codex",model="gpt-6-luna",
                    timeout=5,agent_id="main",runner="exec",binary="codex",config_path=None,
                    controller=bridge.RunController(),output_policy="review_draft")
            self.assertEqual(candidate,before)
            self.assertEqual(pointer["provider_attempt"],1)
            self.assertEqual(pointer["status"],"completed_with_disputes")
            envelope=json.loads((directory/pointer["audit_path"]).read_text())
            def bridge_consume():
                bridge._validate_completed_obligation_ledger_chain(directory,envelope,pointer,candidate,chunk,
                    chunk_index=1,attempt=1,output_policy="review_draft")
            bridge_consume()
            full=json.loads((directory/"llm-request.json").read_text())
            candidate_path=directory/"accepted-candidate.json";bridge._write_json(candidate_path,candidate)
            audit={"chunk_count":1,"adapter_id":"codex","host_runtime":"codex",
                "chunk_lifecycle":[{"chunk_index":1,"status":"completed","remote_operation_state":"completed"}],
                "chunk_runs":[{"chunk_index":1,"response_path":str(candidate_path),"accepted_response_sha256":sha256_json(candidate),
                               "independent_obligation_review":pointer}]}
            def consume(policy="review_draft"):
                return pipeline._validate_independent_obligation_receipts(audit=audit,review_root=directory,
                    expected_run_id=request["run_id"],expected_request_body_sha=request_body_sha256(full),
                    expected_request_envelope_sha=None,expected_request_file_sha=None,output_policy=policy)
            receipts=consume()
            self.assertEqual(receipts[0]["manual_review_required_clause_ids"],["C1"])
            self.assertTrue(receipts[0]["submission_blocked_by_manual_review"])
            with self.assertRaises(ValueError):consume("submission")
            self.verify_manual_draft(receipts,chunk,Path(tmp)/"draft.docx")
            ledger_path=directory/pointer["obligation_analysis_ledger_path"]
            ledger=json.loads(ledger_path.read_text())
            self.assertEqual(ledger["obligations"],[])  # no invented AO to satisfy a schema
            saved=copy.deepcopy(ledger)
            for mutation in ("deleted","changed_primary","stale_run"):
                ledger=copy.deepcopy(saved)
                if mutation=="deleted":ledger["source_inventory_disputes"]=[]
                if mutation=="changed_primary":ledger["source_inventory_disputes"][0]["primary_review_context"]["primary_obligations"]=[]
                if mutation=="stale_run":ledger["run_id"]="old-run"
                bridge._write_json(ledger_path,ledger)
                pointer["obligation_analysis_ledger_sha256"]=bridge.sha256_file(ledger_path)
                envelope["obligation_analysis_ledger"]["sha256"]=bridge.sha256_file(ledger_path)
                bridge._write_json(directory/pointer["audit_path"],envelope)
                pointer["audit_sha256"]=bridge.sha256_file(directory/pointer["audit_path"])
                with self.subTest(mutation=mutation):
                    with self.assertRaises(ValueError):bridge_consume()
                    with self.assertRaises(ValueError):consume()

    def verify_manual_draft(self,receipts,chunk,path):
        gates=pipeline._inventory_dispute_release_gates(receipts,clauses=chunk["clauses"])
        self.assertEqual(len(gates),1)
        binding=manual_fixtures.ManualReviewTests._binding(run_id=receipts[0]["run_id"])
        ledger=build_manual_review_ledger({},[],binding=binding,release_gates=gates)
        self.assertEqual(load_and_validate(ledger,ROOT/"schema/manual-review-ledger.schema.json"),[])
        self.assertEqual(len(ledger["items"]),1)
        card=build_scorecard(binding,{"missing_count":0,"unexpected_count":0,"duplicate_count":0,
                                     "receipts":[],"expected_receipt_ids":[]},ledger["items"],[])
        self.assertFalse(card["submission_ready"])
        doc=Document();doc.add_paragraph("保留原文")
        append_manual_review_markers(doc,ledger);append_scorecard(doc,card);doc.save(path)
        self.assertTrue(audit_manual_review_markers(path,ledger)["valid"])
        self.assertTrue(audit_scorecard(path,card)["valid"])
        text="\n".join(p.text for p in Document(path).paragraphs)
        self.assertIn("义务是否存在尚未确认",text)
        self.assertIn("不要按主审推测",text)
        self.assertIn("保留原文",text)
        # Both opinions participate in identity, so changing one invalidates
        # existing DOCX markers rather than silently changing the human task.
        changed_receipts=copy.deepcopy(receipts)
        changed_receipts[0]["source_inventory_disputes"][0]["independent_result"]["rationale"]="changed"
        changed=pipeline._inventory_dispute_release_gates(changed_receipts,clauses=chunk["clauses"])
        stale=build_manual_review_ledger({},[],binding=binding,release_gates=changed)
        self.assertNotEqual(stale["items"][0]["manual_obligation_id"],ledger["items"][0]["manual_obligation_id"])
        self.assertFalse(audit_manual_review_markers(path,stale)["valid"])


class ClassificationDisputeTests(unittest.TestCase):
    def packet(self, directory):
        _, chunk = packets.HostAgentBridgeTests()._packet(
            directory, source="10043", contract_version="3.0",
        )
        candidate = {
            "contract_version": "3.0",
            "provenance": copy.deepcopy(chunk["provenance"]),
            "requirements": [],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "informational",
                "normative_basis": "source_content",
                "reason": "The isolated source value does not state an instruction.",
                "obligations": [],
            }],
            "unsupported_items": [],
            "reported_conflicts": [],
        }
        request = native.build_obligation_coverage_request(
            candidate, chunk, run_id=chunk["provenance"]["run_id"], chunk_index=1,
        )
        request.update(attempt=1, provider_attempt=1, output_policy="review_draft")
        check = request["checks"][0]
        result = {
            "check_id": check["check_id"],
            "verdict": "uncertain",
            "rationale": "The isolated value may have a field meaning, but the source does not settle it.",
            "identified_obligations": [{
                "force": "unknown", "applicability": "unknown", "target": "10043",
                "obligation_summary": "The field meaning and any preservation duty are not established.",
                "requirement_refs": [], "disposition": "ambiguous", "source_quote": "10043",
            }],
            "evidence_quotes": ["10043"],
            "machine_obligation_ids": [],
        }
        return chunk, candidate, request, result

    def template_example_packet(self, directory):
        _, chunk = packets.HostAgentBridgeTests()._packet(
            directory, source="培养单位", contract_version="3.0",
        )
        candidate = {
            "contract_version": "3.0",
            "provenance": copy.deepcopy(chunk["provenance"]),
            "requirements": [],
            "clause_reviews": [{
                "clause_id": "C1", "classification": "informational",
                "normative_basis": "template_structure",
                "reason": "This is an optional template cover example, not a separate requirement.",
                "obligations": [{
                    "id": "C1-O1", "status": "covered",
                    "reason": "The field appears as an optional template example.",
                    "actor": "template", "action": "display",
                    "target": "cover field label 培养单位", "source_quote": "培养单位",
                    "force": "optional", "applicability": "applicable", "route": "example",
                }],
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        request = native.build_obligation_coverage_request(
            candidate, chunk, run_id=chunk["provenance"]["run_id"], chunk_index=1,
        )
        request.update(attempt=1, provider_attempt=1, output_policy="review_draft")
        check = request["checks"][0]
        result = {
            "check_id": check["check_id"], "verdict": "consistent",
            "rationale": "The source is an informational cover label and names no independent duty.",
            "identified_obligations": [], "evidence_quotes": [check["document_text"]],
            "machine_obligation_ids": [],
        }
        return chunk, candidate, request, result

    def test_narrow_uncertainty_conflict_is_preserved_only_in_review_draft(self):
        with tempfile.TemporaryDirectory() as tmp:
            _, _, request, result = self.packet(Path(tmp) / "packet")
            check = request["checks"][0]
            before = copy.deepcopy((check, result))
            validated = native.validate_obligation_coverage_response(
                {"results": [result]}, [check], allow_draft_disputes=True,
            )
            self.assertEqual(validated, [result])
            dispute = informational_uncertainty_dispute(check, result)
            self.assertIsNotNone(dispute)
            self.assertEqual(dispute["primary_review_context"], check["review_context"])
            self.assertEqual(dispute["independent_result"], result)
            self.assertFalse(dispute["coverage_complete"])
            self.assertFalse(dispute["execution_authorized"])
            self.assertFalse(dispute["submission_ready"])
            self.assertEqual(before, (check, result))

            envelope = {
                "results": [result],
                "source_classification_disputes": [dispute],
                "status": "completed_with_disputes",
                "coverage_complete": False,
                "submission_ready": False,
            }
            native.validate_draft_dispute_envelope(
                envelope, request, output_policy="review_draft",
            )
            with self.assertRaises(native.NativeSemanticReviewError):
                native.validate_obligation_coverage_response(
                    {"results": [result]}, [check], allow_draft_disputes=False,
                )
            for mutation in (
                lambda x: x.update(source_classification_disputes=[]),
                lambda x: x.update(status="completed"),
                lambda x: x.update(coverage_complete=True),
                lambda x: x.update(submission_ready=True),
            ):
                forged = copy.deepcopy(envelope)
                mutation(forged)
                with self.assertRaises(ValueError):
                    native.validate_draft_dispute_envelope(
                        forged, request, output_policy="review_draft",
                    )

            altered = copy.deepcopy(result)
            altered["identified_obligations"][0]["force"] = "required"
            self.assertIsNone(informational_uncertainty_dispute(check, altered))

    def test_source_bound_reported_conflict_gets_draft_marker_and_only_that_blocker_is_relaxed(self):
        source_text = "正文使用宋体或仿宋；两处规定冲突。"
        source_sha256 = hashlib.sha256(source_text.encode("utf-8")).hexdigest()
        clauses = [{
            "id": "C-conflict", "text": source_text, "evidence_ids": ["E-conflict"],
            "source_span": {
                "evidence_id": "E-conflict", "start_offset": 0,
                "end_offset": len(source_text), "text": source_text,
                "source_sha256": source_sha256,
            },
        }]
        evidence_doc = {"evidence": [{"id": "E-conflict", "text": source_text}]}
        conflict = {
            "type": "source_conflict", "reason": "两个来源位置给出了不同字体。",
            "clause_ids": ["C-conflict"], "evidence_ids": ["E-conflict"],
            "status": "requires_human_review",
        }
        spec = {
            "run_id": "conflict-draft-run", "status": "needs_clarification",
            "completeness": {"unresolved_clause_ids": ["C-conflict"]},
            "semantic_conflicts": [conflict],
            "blocking_errors": [{
                "type": "llm_reported_conflict", "conflict_index": 0,
                "conflict_type": conflict["type"], "status": conflict["status"],
                "reason": conflict["reason"], "clause_ids": ["C-conflict"],
                "evidence_ids": ["E-conflict"],
            }],
        }
        gates = pipeline._reported_conflict_release_gates(
            [conflict], clauses=clauses, evidence_doc=evidence_doc,
        )
        ledger = build_manual_review_ledger(
            {}, [], binding=manual_fixtures.ManualReviewTests._binding(
                run_id=spec["run_id"],
            ), release_gates=gates,
        )
        self.assertEqual(len(ledger["items"]), 1)
        marker = ledger["items"][0]
        self.assertEqual(marker["category"], "semantic_content_review")
        self.assertEqual(marker["clause_ids"], ["C-conflict"])
        self.assertEqual(marker["evidence_ids"], ["E-conflict"])
        self.assertFalse(ledger["submission_ready"])

        draft_blockers = pipeline.requirement_blockers(
            spec, [], "supported_subset", review_draft_manual_ledger=ledger,
            allow_source_bound_review_draft_conflicts=True,
        )
        self.assertNotIn("blocking_errors", draft_blockers)
        preview_bypasses = pipeline._source_bound_review_draft_preview_bypasses(
            spec, [], draft_blockers, ledger,
        )
        self.assertIn("unresolved_clauses", preview_bypasses)
        submission_blockers = pipeline.requirement_blockers(
            spec, [], "supported_subset", review_draft_manual_ledger=ledger,
        )
        self.assertIn("blocking_errors", submission_blockers)
        full_mode_blockers = pipeline.requirement_blockers(
            spec, [], "full", review_draft_manual_ledger=ledger,
            allow_source_bound_review_draft_conflicts=True,
        )
        self.assertIn("blocking_errors", full_mode_blockers)
        application_blockers = apply_format_spec.format_spec_blockers(
            spec, "supported_subset", review_draft_manual_ledger=ledger,
            allow_source_bound_review_draft_conflicts=True,
        )
        self.assertNotIn("blocking_errors", application_blockers)
        self.assertIn("blocking_errors", apply_format_spec.format_spec_blockers(
            spec, "supported_subset", allow_source_bound_review_draft_conflicts=True,
        ))

        unbound = copy.deepcopy(ledger)
        unbound["binding"]["run_id"] = "stale-run"
        self.assertIn("blocking_errors", pipeline.requirement_blockers(
            spec, [], "supported_subset", review_draft_manual_ledger=unbound,
            allow_source_bound_review_draft_conflicts=True,
        ))
        self.assertNotIn("unresolved_clauses", pipeline._source_bound_review_draft_preview_bypasses(
            spec, [], ["unresolved_clauses"], unbound,
        ))
        mixed_errors = copy.deepcopy(spec)
        mixed_errors["blocking_errors"].append({"type": "llm_contract", "reason": "bad contract"})
        self.assertIn("blocking_errors", pipeline.requirement_blockers(
            mixed_errors, [], "supported_subset", review_draft_manual_ledger=ledger,
            allow_source_bound_review_draft_conflicts=True,
        ))
        with self.assertRaisesRegex(ValueError, "evidence outside its clauses"):
            pipeline._reported_conflict_release_gates(
                [{**conflict, "evidence_ids": ["E-unrelated"]}],
                clauses=clauses, evidence_doc=evidence_doc,
            )
        with self.assertRaisesRegex(ValueError, "invalid source span"):
            pipeline._reported_conflict_release_gates(
                [conflict], clauses=[{
                    **clauses[0], "source_span": {
                        **clauses[0]["source_span"], "source_sha256": "0" * 64,
                    },
                }], evidence_doc=evidence_doc,
            )
        global_question = [{"question_id": "Q-global", "scope": "global", "question": "workflow issue"}]
        self.assertNotIn("open_questions", pipeline._source_bound_review_draft_preview_bypasses(
            spec, global_question, ["open_questions"], ledger,
        ))
        bound_question = {
            "question_id": "Q-source", "scope": "clause", "clause_id": "C-conflict",
            "evidence_ids": ["E-conflict"], "question": "请人工判断来源冲突的适用规则。",
        }
        question_spec = {"run_id": "question-run", "status": "semantic_resolved"}
        question_ledger = build_manual_review_ledger(
            {}, [bound_question], binding=manual_fixtures.ManualReviewTests._binding(
                run_id=question_spec["run_id"],
            ),
        )
        question_blockers = pipeline.requirement_blockers(
            question_spec, [bound_question], "supported_subset",
        )
        self.assertIn("open_questions", question_blockers)
        self.assertIn("open_questions", pipeline._source_bound_review_draft_preview_bypasses(
            question_spec, [bound_question], question_blockers, question_ledger,
        ))

    def test_informational_primary_atom_with_empty_independent_inventory_stays_strict(self):
        with tempfile.TemporaryDirectory() as tmp:
            _, _, request, _ = self.packet(Path(tmp) / "packet")
            check = copy.deepcopy(request["checks"][0])
            check["review_context"]["primary_obligations"] = [{
                "id": "C1-O1", "status": "covered", "actor": "template",
                "action": "display", "target": "cover field label 10043",
                "source_quote": "10043", "force": "optional",
                "applicability": "applicable", "route": "example",
            }]
            result = {
                "check_id": "C1", "verdict": "consistent",
                "rationale": "The source is only an informational cover label.",
                "identified_obligations": [], "evidence_quotes": ["10043"],
                "machine_obligation_ids": [],
            }
            self.assertIsNone(informational_uncertainty_dispute(check, result))
            with self.assertRaises(native.TypedSourceAtomAlignmentError):
                native.validate_obligation_coverage_response(
                    {"results": [result]}, [check], allow_draft_disputes=True,
                )

    def test_optional_template_example_is_retained_only_after_bound_final_retry(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp) / "packet"
            chunk, candidate, first_request, result = self.template_example_packet(directory)
            check = first_request["checks"][0]
            with self.assertRaises(native.TypedSourceAtomAlignmentError):
                native.validate_obligation_coverage_response(
                    {"results": [result]}, first_request["checks"],
                    allow_draft_disputes=True, review_request=first_request,
                )

            raw = transport.enveloped(wire.RetryScopeTests().wire({"results": [result]}, first_request))
            with ExitStack() as stack:
                transport.EmptyInventoryVerdictTests().mock_transport(stack, [raw])
                def current_request_process(command, **kwargs):
                    last = Path(command[command.index("--output-last-message") + 1])
                    child = last.parent
                    owner = child.parent if child.name.startswith("native-batch-") else child
                    current_request = json.loads((owner / "request.json").read_text())
                    current_raw = transport.enveloped(
                        wire.RetryScopeTests().wire({"results": [result]}, current_request)
                    )
                    text = json.dumps(current_raw, ensure_ascii=False)
                    last.write_text(text, encoding="utf-8")
                    events = [
                        {"type": "item.completed", "item": {"type": "agent_message", "text": text}},
                        {"type": "turn.completed"},
                    ]
                    return transport.CompletedProcess(command, 0,
                        "\n".join(json.dumps(event) for event in events), "")
                stack.enter_context(patch.object(native, "run_process", side_effect=current_request_process))
                stack.enter_context(patch.object(bridge, "INDEPENDENT_REVIEW_RETRY_BACKOFF_SECONDS", 0))
                pointer = bridge._run_independent_obligation_coverage_review(
                    candidate, chunk, review_dir=directory,
                    run_id=first_request["run_id"], chunk_index=1, attempt=1,
                    host_runtime="codex", model="gpt-6-luna", timeout=5,
                    agent_id="main", runner="exec", binary="codex", config_path=None,
                    controller=bridge.RunController(), output_policy="review_draft",
                )

            self.assertEqual(pointer["provider_attempt"], 2)
            self.assertEqual(pointer["status"], "completed_with_disputes")
            envelope = json.loads((directory / pointer["audit_path"]).read_text())
            self.assertFalse(envelope["coverage_complete"])
            self.assertFalse(envelope["submission_ready"])
            self.assertEqual(len(envelope["source_classification_disputes"]), 1)
            dispute = envelope["source_classification_disputes"][0]
            self.assertEqual(dispute["policy"], "source_bound_template_example_alignment_dispute_v1")
            self.assertFalse(dispute["execution_authorized"])
            self.assertFalse(dispute["submission_ready"])
            retry_dir = directory / "independent-review-chunk-0001-attempt-01-provider-attempt-02"
            retry_request = json.loads((retry_dir / "request.json").read_text())
            self.assertTrue(native.typed_alignment_retry_feedback_is_bound(retry_request))
            self.assertEqual(
                native.validate_obligation_coverage_response(
                    {"results": [result]}, retry_request["checks"],
                    allow_draft_disputes=True, review_request=retry_request,
                ),
                [result],
            )

            bridge._validate_completed_obligation_ledger_chain(
                directory, envelope, pointer, candidate, chunk,
                chunk_index=1, attempt=1, output_policy="review_draft",
            )
            candidate_path = directory / "accepted-candidate.json"
            bridge._write_json(candidate_path, candidate)
            full_request = json.loads((directory / "llm-request.json").read_text())
            audit = {
                "chunk_count": 1, "adapter_id": "codex", "host_runtime": "codex",
                "chunk_lifecycle": [{
                    "chunk_index": 1, "status": "completed", "remote_operation_state": "completed",
                }],
                "chunk_runs": [{
                    "chunk_index": 1, "response_path": str(candidate_path),
                    "accepted_response_sha256": sha256_json(candidate),
                    "independent_obligation_review": pointer,
                }],
            }
            receipts = pipeline._validate_independent_obligation_receipts(
                audit=audit, review_root=directory, expected_run_id=first_request["run_id"],
                expected_request_body_sha=request_body_sha256(full_request),
                expected_request_envelope_sha=None, expected_request_file_sha=None,
                output_policy="review_draft",
            )
            self.assertEqual(receipts[0]["source_classification_disputes"], [dispute])
            self.assertEqual(receipts[0]["manual_review_required_clause_ids"], ["C1"])
            with self.assertRaises(ValueError):
                pipeline.enforce_obligation_review_output_policy(
                    [result], output_policy="submission", source_classification_disputes=[dispute],
                )

            gates = pipeline._classification_dispute_release_gates(receipts, clauses=chunk["clauses"])
            self.assertEqual(len(gates), 1)
            self.assertIn("可选的模板示例字段", gates[0]["reason"])
            binding = manual_fixtures.ManualReviewTests._binding(run_id=receipts[0]["run_id"])
            ledger = build_manual_review_ledger({}, [], binding=binding, release_gates=gates)
            self.assertEqual(load_and_validate(ledger, ROOT / "schema/manual-review-ledger.schema.json"), [])
            self.assertFalse(ledger["submission_ready"])
            docx_path = Path(tmp) / "template-example-draft.docx"
            doc = Document()
            doc.add_paragraph("保留模板来源")
            append_manual_review_markers(doc, ledger)
            doc.save(docx_path)
            self.assertTrue(audit_manual_review_markers(docx_path, ledger)["valid"])
            marker_text = "\n".join(item.text for item in Document(docx_path).paragraphs)
            self.assertIn("模板示例字段是否应保留", marker_text)
            self.assertIn("确认前不可提交", marker_text)

            unbound = copy.deepcopy(retry_request)
            unbound["retry_feedback"]["checks_sha256"] = "0" * 64
            self.assertIsNone(template_example_alignment_dispute(check, result, request=unbound))
            with self.assertRaises(native.TypedSourceAtomAlignmentError):
                native.validate_obligation_coverage_response(
                    {"results": [result]}, unbound["checks"],
                    allow_draft_disputes=True, review_request=unbound,
                )

    def test_production_receipts_and_visible_nonrelease_gate_reconstruct_the_dispute(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp) / "packet"
            chunk, candidate, request, result = self.packet(directory)
            raw = transport.enveloped(wire.RetryScopeTests().wire({"results": [result]}, request))
            with ExitStack() as stack:
                transport.EmptyInventoryVerdictTests().mock_transport(stack, [raw])
                pointer = bridge._run_independent_obligation_coverage_review(
                    candidate, chunk, review_dir=directory,
                    run_id=request["run_id"], chunk_index=1, attempt=1,
                    host_runtime="codex", model="gpt-6-luna", timeout=5,
                    agent_id="main", runner="exec", binary="codex", config_path=None,
                    controller=bridge.RunController(), output_policy="review_draft",
                )
            self.assertEqual(pointer["status"], "completed_with_disputes")
            envelope = json.loads((directory / pointer["audit_path"]).read_text())
            self.assertFalse(envelope["coverage_complete"])
            self.assertFalse(envelope["submission_ready"])
            self.assertEqual(len(envelope["source_classification_disputes"]), 1)
            self.assertEqual(envelope["source_classification_disputes"][0]["independent_result"], result)

            def bridge_consume():
                bridge._validate_completed_obligation_ledger_chain(
                    directory, envelope, pointer, candidate, chunk,
                    chunk_index=1, attempt=1, output_policy="review_draft",
                )

            bridge_consume()
            candidate_path = directory / "accepted-candidate.json"
            bridge._write_json(candidate_path, candidate)
            full_request = json.loads((directory / "llm-request.json").read_text())
            audit = {
                "chunk_count": 1, "adapter_id": "codex", "host_runtime": "codex",
                "chunk_lifecycle": [{
                    "chunk_index": 1, "status": "completed", "remote_operation_state": "completed",
                }],
                "chunk_runs": [{
                    "chunk_index": 1, "response_path": str(candidate_path),
                    "accepted_response_sha256": sha256_json(candidate),
                    "independent_obligation_review": pointer,
                }],
            }
            receipts = pipeline._validate_independent_obligation_receipts(
                audit=audit, review_root=directory, expected_run_id=request["run_id"],
                expected_request_body_sha=request_body_sha256(full_request),
                expected_request_envelope_sha=None, expected_request_file_sha=None,
                output_policy="review_draft",
            )
            self.assertEqual(receipts[0]["source_classification_disputes"], envelope["source_classification_disputes"])
            self.assertEqual(receipts[0]["manual_review_required_clause_ids"], ["C1"])
            with self.assertRaises(ValueError):
                pipeline.enforce_obligation_review_output_policy(
                    [result], output_policy="submission",
                    source_classification_disputes=envelope["source_classification_disputes"],
                )

            gates = pipeline._classification_dispute_release_gates(
                receipts, clauses=chunk["clauses"],
            )
            self.assertEqual(len(gates), 1)
            binding = manual_fixtures.ManualReviewTests._binding(run_id=receipts[0]["run_id"])
            ledger = build_manual_review_ledger({}, [], binding=binding, release_gates=gates)
            self.assertEqual(load_and_validate(ledger, ROOT / "schema/manual-review-ledger.schema.json"), [])
            self.assertFalse(ledger["submission_ready"])
            docx_path = Path(tmp) / "draft.docx"
            doc = Document()
            doc.add_paragraph("保留来源原文")
            append_manual_review_markers(doc, ledger)
            doc.save(docx_path)
            self.assertTrue(audit_manual_review_markers(docx_path, ledger)["valid"])
            rendered_text = "\n".join(item.text for item in Document(docx_path).paragraphs)
            self.assertIn("待人工判断", rendered_text)
            self.assertIn("确认前不可提交", rendered_text)

            forged = copy.deepcopy(envelope)
            forged["source_classification_disputes"] = []
            with self.assertRaises(ValueError):
                native.validate_draft_dispute_envelope(
                    forged, request, output_policy="review_draft",
                )


if __name__=="__main__":unittest.main()
