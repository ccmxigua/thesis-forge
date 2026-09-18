from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from source_role_extractor import extract_source_roles  # noqa: E402


class SourceRoleExtractorTests(unittest.TestCase):
    def _document(self, path: Path, *, duplicate_references: bool = False,
                  academic_outputs: bool = True) -> None:
        document = Document()
        document.add_heading("第1章 引言", level=1)
        document.add_paragraph("正文")
        document.add_heading("3.5 参考文献", level=2)
        document.add_paragraph("文献一")
        if duplicate_references:
            document.add_heading("参考文献", level=1)
        document.add_heading("附录A 实验", level=1)
        document.add_paragraph("附录正文")
        document.add_heading("后 记", level=1)
        document.add_paragraph("致谢正文")
        if academic_outputs:
            document.add_heading("在学期间发表的学术论文与研究成果", level=1)
            document.add_paragraph("成果一")
        document.save(path)

    def test_extracts_ordered_half_open_ranges_and_normalization_actions(self):
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "source.docx"
            self._document(path)
            result = extract_source_roles(path, target_headings={
                "body": "第1章 引言",
                "references": "参考文献",
                "appendices": "附录A",
                "acknowledgments": "致  谢",
                "academic_outputs": "攻读硕士学位期间的学术成果",
            })

        self.assertEqual(result["status"], "extracted")
        self.assertEqual(result["findings"], [])
        self.assertEqual(list(result["roles"]), [
            "body", "references", "appendices", "acknowledgments", "academic_outputs"
        ])
        ranges = [
            (item["start_body_child_index"], item["end_body_child_index"])
            for item in result["roles"].values()
        ]
        self.assertTrue(all(start < end for start, end in ranges))
        self.assertTrue(all(left[1] == right[0] for left, right in zip(ranges, ranges[1:])))
        self.assertTrue(all(
            item["content_start_body_child_index"] == item["start_body_child_index"] + 1
            and item["heading_policy"] == "preserve_template_boundary"
            for item in result["roles"].values()
        ))
        self.assertEqual(
            result["roles"]["references"]["normalization_actions"][0],
            {"action": "replace_heading_text", "from": "3.5 参考文献", "to": "参考文献"},
        )
        self.assertEqual(
            result["roles"]["acknowledgments"]["normalization_actions"][0]["to"],
            "致  谢",
        )
        self.assertEqual(
            result["roles"]["academic_outputs"]["normalization_actions"][0]["to"],
            "攻读硕士学位期间的学术成果",
        )

    def test_empty_target_boundary_keeps_source_body_heading(self):
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "source.docx"
            self._document(path)
            result = extract_source_roles(path, target_headings={
                "references": "参考文献",
                "appendices": "附录A",
                "acknowledgments": "致  谢",
                "academic_outputs": "攻读硕士学位期间的学术成果",
            })

        self.assertEqual(result["status"], "extracted")
        body = result["roles"]["body"]
        self.assertEqual(body["content_start_body_child_index"], body["start_body_child_index"])
        self.assertEqual(body["heading_policy"], "include_source_boundary")

    def test_whole_document_role_spans_from_abstract_to_document_end(self):
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "source.docx"
            document = Document()
            document.add_paragraph("源论文封面")
            document.add_paragraph("摘 要", style="Title")
            document.add_paragraph("摘要正文")
            document.add_heading("第1章 引言", level=1)
            document.add_paragraph("正文")
            document.save(path)
            result = extract_source_roles(
                path,
                required_roles={"body"},
                target_headings={},
                whole_document_roles={"body"},
            )

        self.assertEqual(result["status"], "extracted")
        body = result["roles"]["body"]
        self.assertEqual(body["start_body_child_index"], 1)
        self.assertEqual(body["content_start_body_child_index"], 1)
        self.assertEqual(body["end_body_child_index"], 5)
        self.assertEqual(body["source_document_content_start"], 1)

    def test_whole_document_role_falls_back_to_first_paragraph_without_heading(self):
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "minimal-source.docx"
            document = Document()
            document.add_paragraph("测试论文", style="Title")
            document.add_paragraph("装配正文内容。")
            document.save(path)
            result = extract_source_roles(
                path,
                required_roles={"body"},
                target_headings={},
                whole_document_roles={"body"},
            )

        self.assertEqual(result["status"], "extracted")
        self.assertEqual(result["findings"], [])
        self.assertEqual(result["roles"]["body"]["start_body_child_index"], 0)
        self.assertEqual(result["roles"]["body"]["end_body_child_index"], 2)
        self.assertEqual(result["roles"]["body"]["normalization_actions"], [])

    def test_missing_required_academic_outputs_fails_closed(self):
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "source.docx"
            self._document(path, academic_outputs=False)
            result = extract_source_roles(path)

        self.assertEqual(result["status"], "blocked")
        self.assertIn("body", result["roles"])
        self.assertEqual(
            [(item["code"], item["role"]) for item in result["findings"]],
            [("source_role.required_missing", "academic_outputs")],
        )

    def test_duplicate_role_heading_fails_closed_with_candidates(self):
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "source.docx"
            self._document(path, duplicate_references=True)
            result = extract_source_roles(path)

        self.assertEqual(result["status"], "blocked")
        finding = next(item for item in result["findings"] if item["role"] == "references")
        self.assertEqual(finding["code"], "source_role.heading_not_unique")
        self.assertEqual(len(finding["candidates"]), 2)

    def test_real_neau_source_identifies_roles_and_exposes_missing_required_role(self):
        path = ROOT / "build/architecture-regression-20260719/neau-structural-preflight-fix/work/source-conversion/source.docx"
        if not path.exists():
            self.skipTest("NEAU integration source is not present")
        result = extract_source_roles(path)

        self.assertEqual(result["status"], "blocked")
        self.assertEqual(set(result["roles"]), {"body", "references", "appendices", "acknowledgments"})
        self.assertEqual(
            [(item["code"], item["role"]) for item in result["findings"]],
            [("source_role.required_missing", "academic_outputs")],
        )
        self.assertEqual(" ".join(result["roles"]["references"]["heading"]["text"].split()), "3.5 参考文献")
        self.assertEqual(result["roles"]["acknowledgments"]["heading"]["text"], "后 记")


if __name__ == "__main__":
    unittest.main()
