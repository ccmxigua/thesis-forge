from __future__ import annotations

import hashlib
import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from source_scope_binding import (  # noqa: E402
    diagnose_source_scope_reference,
    validate_source_scope_binding,
)
import requirements_engine  # noqa: E402
from property_receipts import audit_property_receipts, build_property_receipts  # noqa: E402
import apply_format_spec  # noqa: E402


def clause(clause_id: str, text: str, source: str, *, evidence_id: str,
           location: dict, start: int = 0) -> dict:
    return {
        "id": clause_id,
        "text": text,
        "source_evidence_text": source,
        "source_kind": "table_cell" if "table_child_index" in location else "paragraph",
        "location": dict(location),
        "evidence_ids": [evidence_id],
        "source_span": {
            "evidence_id": evidence_id,
            "start_offset": start,
            "end_offset": start + len(text),
            "text": text,
            "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest(),
            "location": dict(location),
        },
    }


class SourceScopeBindingTests(unittest.TestCase):
    def test_tjufe_cover_four_field_labels_are_typed_without_capturing_date(self) -> None:
        for literal in ("专业名称：", "作者学号：", "论文作者：", "指导教师："):
            with self.subTest(literal=literal):
                self.assertEqual(
                    requirements_engine.identify_role(literal)[0], "cover_field_label",
                )
        self.assertEqual(
            requirements_engine.identify_role("专业名称至指导教师标题中文字体为宋体")[0],
            "cover_field_label",
        )
        self.assertEqual(
            requirements_engine.identify_role("论文提交日期为宋体，四号，居中对齐")[0],
            None,
        )
        self.assertNotEqual(
            requirements_engine.identify_role("所属学院至论文作者标题加粗")[0],
            "cover_field_label",
        )

    def test_tjufe_cover_rules_report_resolved_scope_but_unverified_when_target_absent(self) -> None:
        cover_fields_text = "自上而下为：专业名称、作者学号、论文作者、指导教师"
        cover_fields = clause(
            "C-FIELDS", cover_fields_text, cover_fields_text,
            evidence_id="E-FIELDS",
            location={"part": "document", "child_index": 1, "paragraph": 0},
        )
        cover_fields["context_before"] = ["（一）封面"]
        label_text = "专业名称至指导教师标题中文字体为宋体，四号，分散对齐，无缩进，单倍行距"
        label_rule = clause(
            "C-LABEL-RULE", label_text, label_text,
            evidence_id="E-LABEL-RULE",
            location={"part": "document", "child_index": 2, "paragraph": 0},
        )
        label_rule["context_before"] = ["第三章 学位论文的书写、排版及印刷要求", "（一）封面"]
        label_rule["context_after"] = []
        value_text = "对应填写的内容为宋体，四号，居中对齐，无缩进，单倍行距"
        value_rule = clause(
            "C-VALUE-RULE", value_text, value_text,
            evidence_id="E-VALUE-RULE",
            location={"part": "document", "child_index": 3, "paragraph": 0},
        )
        value_rule["context_before"] = [label_text, "（一）封面"]
        value_rule["context_after"] = []
        date_text = "论文提交日期为宋体，四号，居中对齐，无缩进，单倍行距"
        date_rule = clause(
            "C-DATE-RULE", date_text, date_text,
            evidence_id="E-DATE-RULE",
            location={"part": "document", "child_index": 4, "paragraph": 0},
        )
        date_rule["context_before"] = [value_text, "（一）封面"]
        date_rule["context_after"] = []
        title_value_text = "对应填写的内容为宋体，四号，加粗，居中对齐，无缩进，单倍行距"
        title_value_rule = clause(
            "C-TITLE-VALUE-RULE", title_value_text, title_value_text,
            evidence_id="E-TITLE-VALUE-RULE",
            location={"part": "document", "child_index": 5, "paragraph": 0},
        )
        title_value_rule["context_before"] = ["（二）扉页", "所属学院至论文作者标题"]
        title_value_rule["context_after"] = ["（三）摘要"]

        clauses = [cover_fields, label_rule, value_rule, date_rule, title_value_rule]
        spec, questions, conflicts = requirements_engine.build_rule_result(
            Path("synthetic-tjufe-source.docx"), clauses,
        )
        self.assertFalse(conflicts)
        self.assertNotIn("cover_field_label", spec["roles"])
        self.assertNotIn("cover_field_value", spec["roles"])
        self.assertFalse(any(item["role"] in {"cover_field_label", "cover_field_value"}
                             for item in spec["requirements"]))
        question_by_clause = {item["clause_id"]: item for item in questions}
        for clause_id in ("C-LABEL-RULE", "C-VALUE-RULE"):
            resolution = question_by_clause[clause_id]["source_scope_resolution"]
            self.assertEqual(resolution["rule_interpretation_status"], "resolved")
            self.assertEqual(resolution["target_status"], "target_absent")
            self.assertEqual(resolution["verification_status"], "unverified")
            self.assertEqual(
                resolution["field_labels"],
                ["专业名称", "作者学号", "论文作者", "指导教师"],
            )
        self.assertEqual(
            question_by_clause["C-VALUE-RULE"]["candidate_roles"], ["cover_field_value"],
        )
        self.assertEqual(
            question_by_clause["C-TITLE-VALUE-RULE"]["candidate_roles"],
            ["cover_field_value"],
        )
        self.assertEqual(
            question_by_clause["C-TITLE-VALUE-RULE"]["source_scope_diagnosis"]["status"],
            "pending_input",
        )
        self.assertNotIn(
            "C-DATE-RULE",
            {clause_id for item in spec["requirements"]
             if item["role"] == "cover_field_label" for clause_id in item["clause_ids"]},
        )

    def test_tjufe_cover_absence_suppresses_stale_scope_diagnosis_for_labels_and_values(self) -> None:
        fields_text = "自上而下为：专业名称、作者学号、论文作者、指导教师"
        fields = clause(
            "C-FIELDS", fields_text, fields_text, evidence_id="E-FIELDS",
            location={"part": "document", "child_index": 1, "paragraph": 0},
        )
        fields["context_before"] = ["（一）封面"]
        fields["context_after"] = []
        label_text = "专业名称至指导教师标题中文字体为宋体，四号，分散对齐，无缩进，单倍行距"
        label = clause(
            "C-LABEL", label_text, label_text, evidence_id="E-LABEL",
            location={"part": "document", "child_index": 2, "paragraph": 0},
        )
        label["context_before"] = ["（一）封面"]
        label["context_after"] = []
        value_text = "对应填写的内容为宋体，四号，居中对齐，无缩进，单倍行距"
        value = clause(
            "C-VALUE", value_text, value_text, evidence_id="E-VALUE",
            location={"part": "document", "child_index": 3, "paragraph": 0},
        )
        value["context_before"] = [label_text, "（一）封面"]
        value["context_after"] = []

        _, questions, conflicts = requirements_engine.build_rule_result(
            Path("synthetic-tjufe-source.docx"), [fields, label, value],
            target_clauses=[],
        )
        self.assertEqual(conflicts, [])
        by_clause = {item["clause_id"]: item for item in questions}
        for clause_id in ("C-LABEL", "C-VALUE"):
            with self.subTest(clause_id=clause_id):
                self.assertEqual(
                    by_clause[clause_id]["source_scope_resolution"]["target_status"],
                    "target_absent",
                )
                self.assertEqual(
                    by_clause[clause_id]["source_scope_resolution"]["verification_status"],
                    "unverified",
                )
                self.assertNotIn("source_scope_diagnosis", by_clause[clause_id])

    def test_title_page_range_requires_hash_bound_ordered_template_endpoints(self) -> None:
        text = "分类号至密级标题及内容：中文字体为宋体，西文字体为Times New Roman，四号"
        source_clause = clause(
            "C-CLASSIFICATION", text, text, evidence_id="E-CLASSIFICATION",
            location={"part": "document", "child_index": 10, "paragraph": 0},
        )
        source_clause["context_before"] = ["（二）扉页"]
        source_clause["context_after"] = []
        template_sha = "a" * 64
        official_template = {
            "source": {"kind": "official_template_docx", "sha256": template_sha},
            "items": [
                {
                    "id": "TE-A", "category": "fixed_text", "normalized_text": "分类号:XXXX",
                    "provenance": {"source_kind": "official_template_docx", "source_sha256": template_sha},
                    "source_location": {"part": "word/document.xml", "paragraph_index": 42},
                },
                {
                    "id": "TE-B", "category": "fixed_text", "normalized_text": "密级:XXXX",
                    "provenance": {"source_kind": "official_template_docx", "source_sha256": template_sha},
                    "source_location": {"part": "word/document.xml", "paragraph_index": 43},
                },
            ],
        }
        spec, questions, conflicts = requirements_engine.build_rule_result(
            Path("synthetic-tjufe-source.docx"), [source_clause],
            target_clauses=[], official_template_evidence=official_template,
        )
        self.assertEqual(conflicts, [])
        item = next(q for q in questions if q["clause_id"] == "C-CLASSIFICATION")
        self.assertEqual(item["source_scope_resolution"]["field_labels"], ["分类号", "密级"])
        self.assertEqual(item["source_scope_resolution"]["template_evidence_ids"], ["TE-A", "TE-B"])
        self.assertEqual(item["source_scope_resolution"]["target_status"], "target_absent")
        self.assertEqual(item["source_scope_resolution"]["verification_status"], "unverified")
        self.assertNotIn("thesis_classified_index", spec["roles"])
        self.assertNotIn("source_scope_diagnosis", item)

    def test_title_page_range_fails_closed_without_unique_official_template_binding(self) -> None:
        text = "分类号至密级标题及内容：中文字体为宋体，四号"
        source_clause = clause(
            "C-CLASSIFICATION", text, text, evidence_id="E-CLASSIFICATION",
            location={"part": "document", "child_index": 10, "paragraph": 0},
        )
        source_clause["context_before"] = ["（二）扉页"]
        source_clause["context_after"] = []
        duplicated = {
            "source": {"kind": "official_template_docx", "sha256": "a" * 64},
            "items": [
                {
                    "id": "TE-A1", "category": "fixed_text", "normalized_text": "分类号:XXXX",
                    "provenance": {"source_kind": "official_template_docx", "source_sha256": "a" * 64},
                    "source_location": {"part": "word/document.xml", "paragraph_index": 42},
                },
                {
                    "id": "TE-A2", "category": "fixed_text", "normalized_text": "分类号:YYYY",
                    "provenance": {"source_kind": "official_template_docx", "source_sha256": "a" * 64},
                    "source_location": {"part": "word/document.xml", "paragraph_index": 44},
                },
                {
                    "id": "TE-B", "category": "fixed_text", "normalized_text": "密级:XXXX",
                    "provenance": {"source_kind": "official_template_docx", "source_sha256": "a" * 64},
                    "source_location": {"part": "word/document.xml", "paragraph_index": 43},
                },
            ],
        }
        _, questions, _ = requirements_engine.build_rule_result(
            Path("synthetic-tjufe-source.docx"), [source_clause],
            target_clauses=[], official_template_evidence=duplicated,
        )
        item = next(q for q in questions if q["clause_id"] == "C-CLASSIFICATION")
        self.assertNotIn("target_status", item["source_scope_resolution"])
        self.assertEqual(item["source_scope_resolution"]["status"], "semantic_ambiguity")

    def test_conditional_ranges_and_anaphoric_successor_remain_ambiguous(self) -> None:
        first_text = "所属学院至论文作者标题（博士扉页的论文作者至论文评阅人标题）中文字体为宋体"
        second_text = "对应填写的内容为宋体，四号，加粗"
        source = first_text + "。" + second_text
        location = {"part": "document", "child_index": 56, "paragraph": 0}
        first = clause(
            "C-HEADER-RANGE", first_text, source, evidence_id="E-TITLE",
            location=location, start=0,
        )
        first["context_before"] = ["（二）扉页"]
        second = clause(
            "C-CORRESPONDING-CONTENT", second_text, source, evidence_id="E-TITLE",
            location=location, start=len(first_text) + 1,
        )
        second["context_before"] = ["（二）扉页"]
        records = [first, second]

        range_diagnosis = diagnose_source_scope_reference(first, records)
        self.assertEqual(range_diagnosis["classification"], "semantic_ambiguity")
        self.assertEqual(range_diagnosis["source_ranges"], [
            {"left": "所属学院", "right": "论文作者"},
            {"left": "论文作者", "right": "论文评阅人", "condition": "博士扉页"},
        ])
        anaphora_diagnosis = diagnose_source_scope_reference(second, records)
        self.assertEqual(anaphora_diagnosis["classification"], "semantic_ambiguity")
        self.assertEqual(anaphora_diagnosis["candidate_clause_ids"], ["C-HEADER-RANGE"])
        self.assertEqual(
            anaphora_diagnosis["candidates"][0]["antecedent_source_ranges"],
            range_diagnosis["source_ranges"],
        )

    def test_explicit_title_page_range_does_not_select_cover_or_unanchored_candidates(self) -> None:
        current_text = "论文作者至论文评阅人标题使用宋体"
        current = {"id": "C-CURRENT", "text": current_text,
                   "context_before": ["（二）扉页"], "evidence_ids": ["E-CURRENT"]}
        cover = {"id": "C-COVER", "text": "论文作者：导师姓名",
                 "context_before": ["（一）封面"], "evidence_ids": ["E-COVER"]}
        unanchored = {"id": "C-UNANCHORED", "text": "论文评阅人：专家姓名",
                      "evidence_ids": ["E-UNKNOWN"]}

        result = diagnose_source_scope_reference(current, [cover, unanchored, current])

        self.assertEqual(result["classification"], "parser_capability_gap")
        self.assertEqual(result["candidate_clause_ids"], [])

    def test_primary_merge_cannot_erase_source_scope_ambiguity_or_target_absence(self) -> None:
        deterministic = [
            {
                "question_id": "Q-RULE-1", "clause_id": "C-TARGET",
                "question": "目标字段缺失，状态未核验。",
                "source_scope_resolution": {
                    "rule_interpretation_status": "resolved",
                    "target_status": "target_absent",
                    "verification_status": "unverified",
                    "field_labels": ["甲", "乙"],
                },
                "source_scope_diagnosis": {
                    "status": "pending_input", "classification": "human_confirmation_required",
                    "reason": "stale_scope_diagnosis",
                    "resolution_policy": "human_selects_exact_clause_ids_before_execution",
                },
            },
            {
                "question_id": "Q-RULE-2", "clause_id": "C-RANGE",
                "question": "两种条件范围待人工消歧。",
                "source_scope_diagnosis": {
                    "status": "pending_input", "classification": "semantic_ambiguity",
                    "reason": "source_clause_contains_conditional_range_alternatives",
                    "resolution_policy": "human_selects_exact_clause_ids_before_execution",
                },
            },
        ]
        primary = [{
            "question_id": "Q-PRIMARY-1", "clause_id": "C-TARGET",
            "question": "模型认为格式规则已理解。",
        }]

        result = requirements_engine.preserve_deterministic_source_scope_questions(
            primary, deterministic,
        )

        by_clause = {item["clause_id"]: item for item in result}
        self.assertEqual(set(by_clause), {"C-TARGET", "C-RANGE"})
        self.assertEqual(by_clause["C-TARGET"]["source_scope_resolution"]["target_status"],
                         "target_absent")
        self.assertNotIn("source_scope_diagnosis", by_clause["C-TARGET"])
        self.assertEqual(by_clause["C-TARGET"]["source_scope_question"],
                         "目标字段缺失，状态未核验。")
        self.assertEqual(by_clause["C-RANGE"]["source_scope_diagnosis"]["classification"],
                         "semantic_ambiguity")
        self.assertEqual(len({item["question_id"] for item in result}), 2)

    def test_tjufe_cover_ordered_range_compiles_only_matching_cover_cells(self) -> None:
        fields_text = "专业名称、作者学号、论文作者、指导教师"
        clauses: list[dict] = []

        def add(clause_id: str, text: str, location: dict, context: list[str] | None = None) -> None:
            item = clause(
                clause_id, text, text, evidence_id=f"E-{clause_id}", location=location,
            )
            item["context_before"] = list(context or [])
            item["context_after"] = []
            clauses.append(item)

        add(
            "C-FIELDS",
            f"自上而下为：{fields_text}，最末一行为提交论文日期",
            {"part": "document", "child_index": 1, "paragraph": 0},
            ["（一）封面"],
        )
        add(
            "C-LABEL-RULE",
            "专业名称至指导教师标题中文字体为宋体，西文字体为Times New Roman，四号，"
            "分散对齐，无缩进，单倍行距",
            {"part": "document", "child_index": 2, "paragraph": 0},
            ["（一）封面"],
        )
        add(
            "C-VALUE-RULE",
            "对应填写的内容为宋体，四号，居中对齐，无缩进，单倍行距",
            {"part": "document", "child_index": 3, "paragraph": 0},
            ["（一）封面"],
        )
        labels = ["专业名称：", "作者学号：", "论文作者：", "指导教师："]
        values = ["应用数学", "12345", "张三", "李四"]
        cover_ids: set[str] = set()
        for row, (label, value) in enumerate(zip(labels, values)):
            label_id, value_id = f"C-LABEL-{row}", f"C-VALUE-{row}"
            cover_ids.update({label_id, value_id})
            add(label_id, label, {
                "part": "document", "table_child_index": 7, "row": row,
                "column": 0, "paragraph": 0,
            }, ["（一）封面"])
            add(value_id, value, {
                "part": "document", "table_child_index": 7, "row": row,
                "column": 1, "paragraph": 0,
            }, ["（一）封面"])
        add("C-DATE-LABEL", "提交日期：", {
            "part": "document", "table_child_index": 7, "row": 4,
            "column": 0, "paragraph": 0,
        }, ["（一）封面"])
        add("C-DATE-VALUE", "二〇二五年五月", {
            "part": "document", "table_child_index": 7, "row": 4,
            "column": 1, "paragraph": 0,
        }, ["（一）封面"])

        # Same field labels on the title page must not make the cover rule
        # ambiguous or receive the cover's direct formatting.
        for row, (label, value) in enumerate(zip(labels, values)):
            add(f"T-LABEL-{row}", label, {
                "part": "document", "table_child_index": 12, "row": row,
                "column": 0, "paragraph": 0,
            }, ["（二）扉页"])
            add(f"T-VALUE-{row}", value, {
                "part": "document", "table_child_index": 12, "row": row,
                "column": 1, "paragraph": 0,
            }, ["（二）扉页"])
        title_rule = "对应填写的内容为宋体，四号，加粗，居中对齐，无缩进，单倍行距"
        add("C-TITLE-VALUE-RULE", title_rule, {
            "part": "document", "child_index": 20, "paragraph": 0,
        }, ["（二）扉页"])

        spec, questions, conflicts = requirements_engine.build_rule_result(
            Path("synthetic-tjufe-composite.docx"), clauses,
        )
        self.assertEqual(conflicts, [])
        requirements = {
            item["role"]: item for item in spec["requirements"]
            if item["role"] in {"cover_field_label", "cover_field_value"}
        }
        self.assertEqual(set(requirements), {"cover_field_label", "cover_field_value"})
        self.assertEqual(
            set(requirements["cover_field_label"]["source_scope"]["target_clause_ids"]),
            {f"C-LABEL-{row}" for row in range(4)},
        )
        self.assertEqual(
            set(requirements["cover_field_value"]["source_scope"]["target_clause_ids"]),
            {f"C-VALUE-{row}" for row in range(4)},
        )
        self.assertTrue(cover_ids <= set(
            requirements["cover_field_label"]["source_scope"]["target_clause_ids"]
            + requirements["cover_field_value"]["source_scope"]["target_clause_ids"]
        ))
        self.assertEqual(
            requirements["cover_field_label"]["properties"]["paragraph"]["alignment"],
            "distributed",
        )
        self.assertEqual(
            requirements["cover_field_value"]["properties"]["paragraph"]["alignment"],
            "center",
        )
        for role in ("cover_field_label", "cover_field_value"):
            paragraph = requirements[role]["properties"]["paragraph"]
            self.assertEqual(paragraph["first_line_indent_pt"], 0.0)
            self.assertEqual(paragraph["left_indent_pt"], 0.0)
            self.assertEqual(paragraph["right_indent_pt"], 0.0)
        self.assertEqual(
            requirements["cover_field_label"]["properties"]["font"]["cjk"], "SimSun",
        )
        self.assertNotIn("cover_field_label", spec["roles"])
        self.assertNotIn("cover_field_value", spec["roles"])
        self.assertFalse(any(
            target_id.startswith(("C-DATE", "T-"))
            for item in requirements.values()
            for target_id in item["source_scope"]["target_clause_ids"]
        ))
        question_by_clause = {item["clause_id"]: item for item in questions}
        self.assertIn("C-TITLE-VALUE-RULE", question_by_clause)
        self.assertEqual(
            question_by_clause["C-TITLE-VALUE-RULE"]["source_scope_resolution"]["target_status"],
            "parser_capability_gap",
        )
        self.assertFalse(any(
            instance["clause_ids"][0].startswith("T-") for instance in spec["content_instances"]
        ))
        evidence_ids = {value for item in clauses for value in item["evidence_ids"]}
        self.assertEqual(requirements_engine.validate_spec(spec, evidence_ids), [])

    def test_source_scoped_receipts_reject_wrong_font_and_alignment(self) -> None:
        target_clause_ids = [f"C-VALUE-{index}" for index in range(1, 5)]
        instance_ids = [f"I-VALUE-{index}" for index in range(1, 5)]
        requirement = {
            "id": "R-VALUES",
            "role": "cover_field_value",
            "properties": {
                "font": {"cjk": "SimSun"},
                "paragraph": {"alignment": "center"},
            },
            "clause_ids": target_clause_ids,
            "field_instance_ids": instance_ids,
            "source_scope": {
                "kind": "explicit_clause_set",
                "anchor_clause_ids": [],
                "target_clause_ids": target_clause_ids,
            },
        }
        source_sha256 = "b" * 64
        scope_sha256 = "a" * 64

        def observations(*, tampered_instance_id: str | None = None,
                         font: str = "SimSun", alignment: str = "center") -> list[dict]:
            rows = []
            for clause_id, instance_id in zip(target_clause_ids, instance_ids):
                actual_font = font if instance_id == tampered_instance_id else "SimSun"
                actual_alignment = alignment if instance_id == tampered_instance_id else "center"
                rows.append({
                    "field_instance_id": instance_id,
                    "actual": {
                        "font": {"cjk": actual_font},
                        "paragraph": {"alignment": actual_alignment},
                    },
                    "source_bindings": [{
                        "clause_id": clause_id,
                        "source_scope_requirement_id": "R-VALUES",
                        "source_scope_binding_sha256": scope_sha256,
                        "source_scope_source_document_sha256": source_sha256,
                        "source_scope_kind": "explicit_clause_set",
                        "source_scope_target_clause_ids": target_clause_ids,
                    }],
                })
            return rows

        def receipt_status(*, tampered_instance_id: str | None = None,
                           font: str = "SimSun", alignment: str = "center") -> dict:
            receipts = build_property_receipts(
                [requirement], {}, {},
                serialized_docx_sha256="c" * 64,
                actual_by_requirement={"R-VALUES": observations(
                    tampered_instance_id=tampered_instance_id,
                    font=font, alignment=alignment,
                )},
                applicable_roles=apply_format_spec._source_scoped_receipt_roles([requirement]),
                verification_method="serialized_docx_exact_content_instance",
            )
            return audit_property_receipts(receipts)

        self.assertTrue(receipt_status()["valid"])
        self.assertEqual(
            apply_format_spec._source_scoped_receipt_roles([requirement]),
            {"cover_field_value"},
        )
        for instance_id in instance_ids:
            for kwargs, failed_property in (
                ({"font": "Arial"}, "font.cjk"),
                ({"alignment": "left"}, "paragraph.alignment"),
            ):
                with self.subTest(instance_id=instance_id, kwargs=kwargs):
                    report = receipt_status(tampered_instance_id=instance_id, **kwargs)
                    self.assertFalse(report["valid"])
                    self.assertTrue(any(
                        item.get("property_path") == failed_property
                        and item.get("status") == "failed"
                        for item in report["failures"]
                    ))

    def test_separate_target_source_projection_binds_and_fails_closed_on_substitution(self) -> None:
        with tempfile.TemporaryDirectory(prefix="source-scope-separate-inputs-") as directory:
            root = Path(directory)
            requirements_path = root / "requirements.docx"
            requirements_doc = Document()
            requirements_doc.add_paragraph("（一）封面")
            requirements_doc.add_paragraph(
                "签名栏自上而下为：专业名称、作者学号、论文作者、指导教师，最末一行为提交论文日期"
            )
            requirements_doc.add_paragraph(
                "专业名称至指导教师标题中文字体为宋体，西文字体为Times New Roman，四号，"
                "分散对齐，无缩进，单倍行距"
            )
            requirements_doc.add_paragraph(
                "对应填写的内容为宋体，四号，居中对齐，无缩进，单倍行距"
            )
            requirements_doc.save(requirements_path)

            target_path = root / "target.docx"
            target_doc = Document()
            target_doc.add_paragraph("（一）封面")
            cover = target_doc.add_table(rows=5, cols=2)
            for index, (label, value) in enumerate(zip(
                ("专业名称：", "作者学号：", "论文作者：", "指导教师：", "提交论文日期："),
                ("管理学", "2026", "作者甲", "导师乙", "二○二六年五月"),
            )):
                cover.cell(index, 0).text = label
                cover.cell(index, 1).text = value
            target_doc.add_page_break()
            target_doc.add_paragraph("（二）扉页")
            title_page = target_doc.add_table(rows=4, cols=2)
            for index, (label, value) in enumerate(zip(
                ("专业名称：", "作者学号：", "论文作者：", "指导教师："),
                ("管理学", "2026", "作者甲", "导师乙"),
            )):
                title_page.cell(index, 0).text = label
                title_page.cell(index, 1).text = value
            target_doc.save(target_path)

            source_evidence = requirements_engine.extract_document_evidence(requirements_path)
            source_evidence["source_document"] = str(requirements_path.resolve())
            clauses = requirements_engine.split_clauses(source_evidence)
            structure_evidence = requirements_engine.extract_document_evidence(target_path)
            source_record = {
                "kind": "target_thesis_structure_docx",
                "path": str(target_path.resolve()),
                "bytes": target_path.stat().st_size,
                "sha256": hashlib.sha256(target_path.read_bytes()).hexdigest(),
                "same_artifact_as_requirements_source": False,
            }
            projection = requirements_engine.build_structure_target_evidence(
                structure_evidence, source_record,
            )
            source_record["projection_sha256"] = requirements_engine.sha256_json({
                "evidence": projection["evidence"], "clauses": projection["clauses"],
            })
            projection["source"] = copy.deepcopy(source_record)
            source_evidence["structure_source_document_evidence"] = projection
            spec, questions, conflicts = requirements_engine.build_rule_result(
                requirements_path, clauses, target_clauses=projection["clauses"],
            )
            spec["run_id"] = "separate-source-test-run"
            spec["analysis_mode"] = "rule_only"
            focused_requirements = [
                item for item in spec["requirements"]
                if item.get("role") in {"cover_field_label", "cover_field_value"}
                and isinstance(item.get("source_scope"), dict)
            ]
            self.assertEqual(len(focused_requirements), 2)
            target_instances = {
                str(value) for item in focused_requirements
                for value in item.get("field_instance_ids", [])
            }
            self.assertEqual(len(target_instances), 8)
            self.assertTrue(all(
                instance.get("clause_ids", [""])[0].startswith("STRUCT-")
                for instance in spec["content_instances"]
                if instance.get("id") in target_instances
            ))

            evidence_context = requirements_engine.evidence_payload(source_evidence)
            clauses_path = root / "requirement-clauses.json"
            evidence_path = root / "document-evidence.json"
            context_path = root / "evidence-context.json"
            clauses_path.write_text(json.dumps(clauses, ensure_ascii=False), encoding="utf-8")
            evidence_path.write_text(json.dumps(source_evidence, ensure_ascii=False), encoding="utf-8")
            context_path.write_text(json.dumps(evidence_context, ensure_ascii=False), encoding="utf-8")
            file_sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
            manifest = {
                "status": "completed", "cache_reused": False,
                "run_id": spec["run_id"],
                "source_document": str(requirements_path.resolve()),
                "source_sha256": file_sha(requirements_path),
                "requirement_clauses_sha256": file_sha(clauses_path),
                "document_evidence_sha256": file_sha(evidence_path),
                "evidence_sha256": file_sha(context_path),
                "sources": {"structure_source": source_record},
            }

            def validate(current_evidence=source_evidence, current_manifest=manifest):
                evidence_bytes = json.dumps(current_evidence, ensure_ascii=False).encode("utf-8")
                return apply_format_spec.validate_content_instance_source_bindings(
                    spec, clauses, current_evidence, current_manifest, evidence_context,
                    clause_artifact_sha256=file_sha(clauses_path),
                    document_evidence_artifact_sha256=hashlib.sha256(evidence_bytes).hexdigest(),
                    evidence_context_artifact_sha256=file_sha(context_path),
                )

            errors, bindings = validate()
            self.assertEqual(errors, [])
            self.assertEqual(len(bindings), 8)
            self.assertTrue(all(
                binding["source_document"] == str(target_path.resolve())
                and binding["source_document_sha256"] == source_record["sha256"]
                for values in bindings.values() for binding in values
            ))

            missing_projection = copy.deepcopy(source_evidence)
            missing_projection.pop("structure_source_document_evidence")
            missing_errors, _ = validate(current_evidence=missing_projection)
            self.assertIn("structure_source_projection_missing", missing_errors)

            substituted_path_manifest = copy.deepcopy(manifest)
            substituted_path_manifest["sources"]["structure_source"]["path"] = str(
                requirements_path.resolve()
            )
            path_errors, _ = validate(current_manifest=substituted_path_manifest)
            self.assertIn("structure_source_projection_manifest_mismatch", path_errors)

            substituted_hash_manifest = copy.deepcopy(manifest)
            substituted_hash_manifest["sources"]["structure_source"]["sha256"] = "0" * 64
            hash_errors, _ = validate(current_manifest=substituted_hash_manifest)
            self.assertIn("structure_source_projection_manifest_mismatch", hash_errors)
            self.assertIn("structure_source_file_hash_mismatch", hash_errors)

            target_path.write_bytes(target_path.read_bytes() + b"tampered")
            changed_file_errors, _ = validate()
            self.assertIn("structure_source_file_hash_mismatch", changed_file_errors)

    def test_serialized_scope_audit_uses_unique_table_context_for_duplicate_label(self) -> None:
        doc = Document()
        doc.add_paragraph("（一）封面")
        cover = doc.add_table(rows=2, cols=2)
        cover.cell(0, 0).text = "论文作者："
        cover.cell(0, 1).text = "张三"
        cover.cell(1, 0).text = "指导教师："
        cover.cell(1, 1).text = "李四"
        title = doc.add_table(rows=1, cols=2)
        title.cell(0, 0).text = "论文作者："
        title.cell(0, 1).text = "张三"

        cover_label = cover.cell(0, 0).paragraphs[0]
        run = cover_label.runs[0]
        run.font.name = "SimSun"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")
        cover_label.alignment = WD_ALIGN_PARAGRAPH.DISTRIBUTE
        paragraphs = apply_format_spec._content_instance_paragraphs(doc)
        index = next(i for i, item in enumerate(paragraphs) if item._p is cover_label._p)
        table_sha256 = apply_format_spec._serialized_table_context_sha256(cover)
        instance_audit = [{
            "instance_id": "I-COVER-AUTHOR",
            "role": "cover_field_label",
            "text": "论文作者",
            "target_locators": [{
                "inventory_index": index,
                "text": "论文作者：",
                "text_sha256": hashlib.sha256("论文作者：".encode("utf-8")).hexdigest(),
                "source_location": {
                    "part": "document", "table_child_index": 1,
                    "row": 0, "column": 0, "paragraph": 0,
                },
                "source_kind": "table_cell",
                "table_context_sha256": table_sha256,
            }],
        }]
        actual = apply_format_spec.audit_serialized_content_instance_targets(doc, instance_audit)
        self.assertEqual(instance_audit[0]["serialized_target_status"], "verified")
        self.assertEqual(actual["I-COVER-AUTHOR"]["font"]["cjk"], "SimSun")
        self.assertEqual(actual["I-COVER-AUTHOR"]["paragraph"]["alignment"], "distributed")

        duplicate_doc = Document()
        for _ in range(2):
            duplicate = duplicate_doc.add_table(rows=2, cols=2)
            duplicate.cell(0, 0).text = "论文作者："
            duplicate.cell(0, 1).text = "张三"
            duplicate.cell(1, 0).text = "指导教师："
            duplicate.cell(1, 1).text = "李四"
        duplicate_sha256 = apply_format_spec._serialized_table_context_sha256(
            duplicate_doc.tables[0],
        )
        ambiguous_audit = [{
            "instance_id": "I-AMBIGUOUS",
            "role": "cover_field_label",
            "text": "论文作者",
            "target_locators": [{
                **instance_audit[0]["target_locators"][0],
                "table_context_sha256": duplicate_sha256,
            }],
        }]
        ambiguous_actual = apply_format_spec.audit_serialized_content_instance_targets(
            duplicate_doc, ambiguous_audit,
        )
        self.assertEqual(ambiguous_actual, {})
        self.assertEqual(ambiguous_audit[0]["serialized_target_status"], "unverified")
        self.assertIn(
            "serialized_target_table_ambiguous",
            ambiguous_audit[0]["serialized_target_errors"],
        )

    def test_document_evidence_propagates_explicit_cover_and_title_page_markers(self) -> None:
        with tempfile.TemporaryDirectory(prefix="source-section-marker-") as directory:
            path = Path(directory) / "sections.docx"
            doc = Document()
            doc.add_paragraph("（一）封面")
            cover = doc.add_table(rows=2, cols=2)
            cover.cell(0, 0).text = "专业名称："
            cover.cell(0, 1).text = "应用数学"
            cover.cell(1, 0).text = "作者学号："
            cover.cell(1, 1).text = "12345"
            doc.add_paragraph("（二）扉页")
            title_page = doc.add_table(rows=1, cols=2)
            title_page.cell(0, 0).text = "论文作者："
            title_page.cell(0, 1).text = "张三"
            doc.save(path)

            evidence = requirements_engine.extract_document_evidence(path)
            tables: dict[int, set[str]] = {}
            for item in evidence["evidence"]:
                location = item.get("location") or {}
                table_index = location.get("table_child_index")
                if type(table_index) is int:
                    tables.setdefault(table_index, set()).add(item.get("source_section_marker"))
            self.assertEqual(list(tables.values()), [{"cover"}, {"title_page"}])

    def test_scope_requires_explicit_anchor_array_even_for_an_explicit_clause_set(self) -> None:
        source = "字段"
        location = {"part": "document", "child_index": 0, "paragraph": 0}
        clauses = [clause("C1", source, source, evidence_id="E1", location=location)]
        errors, projection = validate_source_scope_binding({
            "kind": "explicit_clause_set", "target_clause_ids": ["C1"],
        }, clauses, cited_clause_ids={"C1"})
        self.assertIsNone(projection)
        self.assertEqual(errors, ["source_scope_anchor_clause_ids_required"])

    def test_duplicate_ordinal_reference_is_reported_as_ambiguity(self) -> None:
        first = {"id": "C1", "text": "第一条 甲项", "evidence_ids": ["E1"]}
        second = {"id": "C2", "text": "第一条 乙项", "evidence_ids": ["E2"]}
        instruction = {"id": "C3", "text": "第一条中导师姓名使用五号宋体", "evidence_ids": ["E3"]}
        result = diagnose_source_scope_reference(instruction, [first, second, instruction])
        self.assertIsNotNone(result)
        self.assertEqual(result["status"], "pending_input")
        self.assertEqual(result["classification"], "semantic_ambiguity")
        self.assertEqual(result["candidate_clause_ids"], ["C1", "C2"])
        self.assertEqual(result["resolution_policy"], "human_selects_exact_clause_ids_before_execution")

    def test_rule_compiler_keeps_scoped_style_pending_instead_of_role_wide(self) -> None:
        clauses = []
        for index, source_text in enumerate(("第一条 甲项", "第一条 乙项"), 1):
            clauses.append({
                "id": f"C{index}", "text": source_text, "evidence_ids": [f"E{index}"],
                "context_before": [], "context_after": [],
            })
        instruction = {
            "id": "C3", "text": "第一条中导师姓名使用五号宋体。", "evidence_ids": ["E3"],
            "context_before": ["第一条 甲项", "第一条 乙项"], "context_after": [],
        }
        clauses.append(instruction)
        spec, questions, conflicts = requirements_engine.build_rule_result(
            Path("source.docx"), clauses,
        )
        self.assertEqual(conflicts, [])
        self.assertFalse(spec["requirements"])
        self.assertNotIn("cover_field_label", spec["roles"])
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0]["source_scope_diagnosis"]["classification"], "semantic_ambiguity")
        self.assertEqual(questions[0]["source_scope_diagnosis"]["candidate_clause_ids"], ["C1", "C2"])

    def test_anaphora_lists_candidates_but_never_resolves_by_proximity(self) -> None:
        prior_a = {"id": "C1", "text": "目录项目格式", "evidence_ids": ["E1"]}
        prior_b = {"id": "C2", "text": "第一条标题格式", "evidence_ids": ["E2"]}
        instruction = {"id": "C3", "text": "上述规定适用五号宋体", "evidence_ids": ["E3"]}
        result = diagnose_source_scope_reference(instruction, [prior_a, prior_b, instruction])
        self.assertEqual(result["classification"], "semantic_ambiguity")
        self.assertEqual(result["candidate_clause_ids"], ["C1", "C2"])

    def test_forward_table_reference_without_structural_candidate_is_capability_gap(self) -> None:
        instruction = {"id": "C1", "text": "字段按下表规定填写", "evidence_ids": ["E1"]}
        result = diagnose_source_scope_reference(instruction, [instruction])
        self.assertEqual(result["status"], "pending_input")
        self.assertEqual(result["classification"], "parser_capability_gap")
        self.assertEqual(result["reason"], "forward_target_not_structurally_observable")

    def test_same_paragraph_range_binds_exact_contiguous_clause_set(self) -> None:
        source = "甲标签乙标签丙"
        location = {"part": "document", "child_index": 5, "paragraph": 0}
        clauses = [
            clause("C1", "甲标签", source, evidence_id="E1", location=location, start=0),
            clause("C2", "乙标签", source, evidence_id="E1", location=location, start=3),
            clause("C3", "丙", source, evidence_id="E1", location=location, start=6),
        ]
        errors, projection = validate_source_scope_binding({
            "kind": "same_paragraph_range",
            "anchor_clause_ids": ["C1", "C3"],
            "target_clause_ids": ["C1", "C2", "C3"],
        }, clauses, cited_clause_ids={"C1", "C2", "C3"})
        self.assertEqual(errors, [])
        self.assertEqual(projection["target_clause_ids"], ["C1", "C2", "C3"])
        self.assertEqual([item["start_offset"] for item in projection["target_spans"]], [0, 3, 6])

    def test_same_paragraph_range_rejects_omitted_middle_target_and_bad_source_hash(self) -> None:
        source = "甲标签乙标签丙"
        location = {"part": "document", "child_index": 5, "paragraph": 0}
        clauses = [
            clause("C1", "甲标签", source, evidence_id="E1", location=location, start=0),
            clause("C2", "乙标签", source, evidence_id="E1", location=location, start=3),
            clause("C3", "丙", source, evidence_id="E1", location=location, start=6),
        ]
        errors, projection = validate_source_scope_binding({
            "kind": "same_paragraph_range", "anchor_clause_ids": ["C1", "C3"],
            "target_clause_ids": ["C1", "C3"],
        }, clauses, cited_clause_ids={"C1", "C3"})
        self.assertIsNone(projection)
        self.assertIn("source_scope_paragraph_range_target_set_mismatch", errors)
        clauses[1]["source_span"]["source_sha256"] = "0" * 64
        errors, projection = validate_source_scope_binding({
            "kind": "same_paragraph_range", "anchor_clause_ids": ["C1", "C3"],
            "target_clause_ids": ["C1", "C2", "C3"],
        }, clauses, cited_clause_ids={"C1", "C2", "C3"})
        self.assertIsNone(projection)
        self.assertTrue(any("source_scope_clause_span_unverifiable:C2" in error for error in errors))

    def test_table_row_range_uses_exact_cell_coordinates_and_order(self) -> None:
        clauses = [
            clause(f"C{column + 1}", f"字段{column + 1}", f"字段{column + 1}",
                   evidence_id=f"E{column + 1}",
                   location={"part": "document", "table_child_index": 3, "row": 2,
                             "column": column, "paragraph": 0})
            for column in range(3)
        ]
        errors, projection = validate_source_scope_binding({
            "kind": "same_table_row_range",
            "anchor_clause_ids": ["C1", "C3"],
            "target_clause_ids": ["C1", "C2", "C3"],
        }, clauses, cited_clause_ids={"C1", "C2", "C3"})
        self.assertEqual(errors, [])
        self.assertEqual(projection["target_clause_ids"], ["C1", "C2", "C3"])
        clauses[1]["location"]["row"] = 4
        clauses[1]["source_span"]["location"]["row"] = 4
        errors, projection = validate_source_scope_binding({
            "kind": "same_table_row_range",
            "anchor_clause_ids": ["C1", "C3"],
            "target_clause_ids": ["C1", "C2", "C3"],
        }, clauses, cited_clause_ids={"C1", "C2", "C3"})
        self.assertIsNone(projection)
        self.assertIn("source_scope_table_row_range_target_set_mismatch", errors)

    def test_table_row_range_rejects_clause_crossing_the_end_anchor(self) -> None:
        source = "甲乙丙丁"
        location = {
            "part": "document", "table_child_index": 3, "row": 2,
            "column": 0, "paragraph": 0,
        }
        clauses = [
            clause("C1", "甲", source, evidence_id="E1", location=location, start=0),
            clause("C2", "乙", source, evidence_id="E1", location=location, start=1),
            clause("C3", "乙丙丁", source, evidence_id="E1", location=location, start=1),
        ]
        errors, projection = validate_source_scope_binding({
            "kind": "same_table_row_range",
            "anchor_clause_ids": ["C1", "C2"],
            "target_clause_ids": ["C1", "C2", "C3"],
        }, clauses, cited_clause_ids={"C1", "C2", "C3"})
        self.assertIsNone(projection)
        self.assertIn("source_scope_table_row_range_target_set_mismatch", errors)

    def test_scope_cannot_reference_uncited_or_unavailable_clause(self) -> None:
        source = "字段"
        loc = {"part": "document", "child_index": 0, "paragraph": 0}
        clauses = [clause("C1", source, source, evidence_id="E1", location=loc)]
        errors, projection = validate_source_scope_binding({
            "kind": "explicit_clause_set", "anchor_clause_ids": [],
            "target_clause_ids": ["C1", "C2"],
        }, clauses, cited_clause_ids={"C1"})
        self.assertIsNone(projection)
        self.assertEqual(errors, [
            "source_scope_references_unknown_clause",
            "source_scope_references_uncited_clause",
        ])

    def _scoped_merge_fixture(self, include_second_instance: bool = True):
        source = "甲标签乙标签"
        location = {"part": "document", "child_index": 5, "paragraph": 0}
        clauses = [
            clause("C1", "甲标签", source, evidence_id="E1", location=location, start=0),
            clause("C2", "乙标签", source, evidence_id="E1", location=location, start=3),
        ]
        instances = []
        for target in clauses[:2 if include_second_instance else 1]:
            instance_id, error = requirements_engine._register_content_instance(
                instances, {"field_key": None}, "cover_field_label",
                {"text": target["text"]}, {target["id"]}, {"E1"},
                {item["id"]: item for item in clauses}, "exact current source label",
            )
            self.assertIsNotNone(instance_id)
            self.assertIsNone(error)
        response = {
            "contract_version": "2.1",
            "requirements": [{
                "role": "cover_field_label",
                "properties": {"font": {"size_pt": 16}},
                "source_scope": {
                    "kind": "same_paragraph_range",
                    "anchor_clause_ids": ["C1", "C2"],
                    "target_clause_ids": ["C1", "C2"],
                },
                "clause_ids": ["C1", "C2"], "evidence_ids": ["E1"],
                "confidence": .99, "reason": "将已核实的同段来源范围绑定到具体字段。",
            }],
            "clause_reviews": [
                {"clause_id": clause_id, "classification": "executable",
                 "requirement_indexes": [0], "reason": "来源范围明确。"}
                for clause_id in ("C1", "C2")
            ],
            "unsupported_items": [], "reported_conflicts": [],
        }
        rule_spec = {
            "schema_version": "1.0", "source_document": "synthetic-source",
            "roles": {}, "requirements": [], "content_instances": instances,
        }
        return source, location, clauses, response, rule_spec

    def test_merge_binds_scoped_style_only_to_exact_field_instances(self) -> None:
        _source, _location, clauses, response, rule_spec = self._scoped_merge_fixture()
        spec, conflicts, audit = requirements_engine.merge_llm_primary(
            Path("synthetic-source"), rule_spec, clauses, response, {"E1"},
        )
        self.assertFalse(any(item.get("type") in {"llm_contract", "source_scope_binding"}
                             for item in conflicts), conflicts)
        requirement = spec["requirements"][0]
        self.assertEqual(
            set(requirement["field_instance_ids"]),
            {item["id"] for item in spec["content_instances"]},
        )
        self.assertEqual(requirement["source_scope"]["target_clause_ids"], ["C1", "C2"])
        self.assertEqual(spec["roles"], {})
        self.assertTrue(all(item["style_properties"]["font"]["size_pt"] == 16
                            for item in spec["content_instances"]))
        self.assertTrue(any(item.get("accepted") is True for item in audit))

    def test_merge_rejects_scoped_style_when_any_target_instance_is_missing(self) -> None:
        _source, _location, clauses, response, rule_spec = self._scoped_merge_fixture(
            include_second_instance=False,
        )
        spec, conflicts, audit = requirements_engine.merge_llm_primary(
            Path("synthetic-source"), rule_spec, clauses, response, {"E1"},
        )
        self.assertEqual(spec["requirements"], [])
        self.assertEqual(spec["roles"], {})
        self.assertTrue(any(item.get("type") == "source_scope_binding"
                            and "source_scope_target_instance_count:C2:0" in str(item.get("reason"))
                            for item in conflicts), conflicts)
        self.assertTrue(any(item.get("accepted") is False
                            and item.get("reason") == "source_scope_target_instance_unresolved"
                            for item in audit))

    def test_merge_rejects_target_instance_that_also_covers_unselected_source_clause(self) -> None:
        _source, _location, clauses, response, rule_spec = self._scoped_merge_fixture()
        rule_spec["content_instances"][0]["clause_ids"].append("C99")
        spec, conflicts, audit = requirements_engine.merge_llm_primary(
            Path("synthetic-source"), rule_spec, clauses, response, {"E1"},
        )
        self.assertEqual(spec["requirements"], [])
        self.assertFalse(spec["roles"])
        self.assertTrue(any(
            item.get("type") == "source_scope_binding"
            and "source_scope_target_instance_contains_unselected_clause:C1" in str(item.get("reason"))
            for item in conflicts
        ), conflicts)
        self.assertTrue(any(item.get("accepted") is False for item in audit))

    def test_scoped_request_text_cannot_create_or_style_an_out_of_scope_instance(self) -> None:
        _source, _location, clauses, response, rule_spec = self._scoped_merge_fixture()
        requirement = response["requirements"][0]
        requirement["properties"] = {"text": "甲标签", "font": {"size_pt": 16}}
        requirement["source_scope"] = {
            "kind": "explicit_clause_set", "anchor_clause_ids": [],
            "target_clause_ids": ["C2"],
        }
        requirement["clause_ids"] = ["C1", "C2"]
        requirement["evidence_ids"] = ["E1"]
        for review in response["clause_reviews"]:
            review["requirement_indexes"] = [0]
        spec, conflicts, _audit = requirements_engine.merge_llm_primary(
            Path("synthetic-source"), rule_spec, clauses, response, {"E1"},
        )
        self.assertFalse(any(item.get("type") in {"llm_contract", "source_scope_binding"}
                             for item in conflicts), conflicts)
        by_clause = {
            clause_id: instance for instance in spec["content_instances"]
            for clause_id in instance["clause_ids"]
        }
        self.assertNotIn("font", by_clause["C1"].get("style_properties", {}))
        self.assertEqual(by_clause["C2"]["style_properties"]["font"]["size_pt"], 16)
        self.assertEqual(spec["requirements"][0]["field_instance_ids"], [by_clause["C2"]["id"]])

    def test_referenced_content_text_without_explicit_scope_is_not_accepted(self) -> None:
        _source, _location, clauses, _response, rule_spec = self._scoped_merge_fixture()
        instruction_text = "第一条中导师姓名使用五号宋体"
        clauses.append(clause(
            "C3", instruction_text, instruction_text, evidence_id="E2",
            location={"part": "document", "child_index": 6, "paragraph": 0},
        ))
        response = {
            "contract_version": "2.1",
            "requirements": [{
                "role": "cover_field_label",
                "properties": {"text": "导师姓名", "font": {"size_pt": 10}},
                "clause_ids": ["C3"], "evidence_ids": ["E2"],
                "confidence": .99, "reason": "原文要求调整第一条中的导师姓名。",
            }],
            "clause_reviews": [{
                "clause_id": "C3", "classification": "covered",
                "requirement_indexes": [0], "reason": "引用特定条目。",
            }],
            "unsupported_items": [], "reported_conflicts": [],
        }
        spec, conflicts, _audit = requirements_engine.merge_llm_primary(
            Path("synthetic-source"), rule_spec, clauses, response, {"E1", "E2"},
        )
        self.assertTrue(any(
            item.get("type") == "llm_contract"
            and "source_scope_binding_required_for_referenced_content" in str(item.get("reason"))
            for item in conflicts
        ), conflicts)
        self.assertFalse(any(
            item.get("role") == "cover_field_label" and item.get("properties", {}).get("font")
            for item in spec["requirements"]
        ))

    def test_build_request_exposes_only_current_clause_ids_for_scope_anchors(self) -> None:
        _source, _location, clauses, _response, rule_spec = self._scoped_merge_fixture()
        request = requirements_engine.build_llm_request(
            [], clauses,
            evidence_doc={"evidence": [{
                "id": "E1", "text": "甲标签乙标签", "kind": "paragraph",
                "location": {"part": "document", "child_index": 5, "paragraph": 0},
            }]},
            rule_spec=rule_spec,
            mode="full",
            contract_version="2.1",
        )
        item = next(
            branch for branch in request["response_schema"]["properties"]["requirements"]["items"]["anyOf"]
            if branch["properties"]["role"]["enum"] == ["cover_field_label"]
        )
        scope = item["properties"]["source_scope"]["$ref"]
        scope_schema = request["response_schema"]["$defs"][scope.removeprefix("#/$defs/")]
        ids = scope_schema["properties"]["target_clause_ids"]["items"]["enum"]
        self.assertEqual(ids, ["C1", "C2"])


if __name__ == "__main__":
    unittest.main()
