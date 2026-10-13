from __future__ import annotations

import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path
import shutil

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from source_table_contract import (  # noqa: E402
    compile_source_table_contract,
    resolve_bound_requirement_docx,
    sha256_file,
)


THREE_LINE = "表的编排须采用国际通用的三线表。"
BORDER_WIDTHS = "说明4：表格的外边框线粗1.5磅，内边框线粗0.5磅。"


def clause(clause_id: str, text: str, evidence_id: str, child_index: int,
           span_text: str) -> dict:
    start = text.index(span_text)
    return {
        "id": clause_id,
        "text": text.rstrip("。"),
        "source_text_full": text,
        "evidence_ids": [evidence_id],
        "source_span": {
            "evidence_id": evidence_id,
            "start_offset": start,
            "end_offset": start + len(span_text),
            "text": span_text,
            "source_sha256": hashlib.sha256(text.encode()).hexdigest(),
            "location": {"part": "document", "child_index": child_index},
        },
    }


def make_bundle(root: Path) -> tuple[Path, Path, Path, Path, list[dict]]:
    source = root / "requirements.docx"
    doc = Document()
    doc.add_paragraph(THREE_LINE)
    doc.add_paragraph(BORDER_WIDTHS)
    doc.save(source)
    clauses = [
        clause("C-three", THREE_LINE, "E-three", 0, THREE_LINE[:-1]),
        clause("C-widths", BORDER_WIDTHS, "E-widths", 1, BORDER_WIDTHS[:-1]),
    ]
    clauses_path = root / "requirement-clauses.json"
    clauses_path.write_text(json.dumps(clauses, ensure_ascii=False), encoding="utf-8")
    spec_path = root / "format-spec.json"
    spec = {
        "source_document": str(source.resolve()),
        "run_id": "test-run",
        "roles": {"table_caption": {"position": "above"}},
        "semantic_review_provenance": {
            "run_id": "test-run", "source_sha256": sha256_file(source),
        },
        "clause_compliance": [
            {"clause_id": item["id"], "status": "unsupported_backend", "requirement_ids": []}
            for item in clauses
        ],
        "requirements": [],
    }
    spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    manifest_path = root / "extraction-manifest.json"
    manifest = {
        "status": "completed", "run_id": "test-run",
        "normalized_source_document": str(source.resolve()),
        "normalized_source_sha256": sha256_file(source),
        "requirement_clauses_sha256": sha256_file(clauses_path),
    }
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    return source, clauses_path, spec_path, manifest_path, clauses


class SourceTableContractTests(unittest.TestCase):
    def test_compiles_exact_bound_three_line_and_border_width_clauses(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, clauses_path, spec_path, manifest_path, clauses = make_bundle(root)
            bound_source, binding = resolve_bound_requirement_docx(manifest_path, clauses_path, spec_path)
            self.assertEqual(bound_source, source)
            self.assertEqual(binding["status"], "verified")
            spec = json.loads(spec_path.read_text())
            result = compile_source_table_contract(source, clauses, spec, source_binding=binding)
            self.assertEqual(result["status"], "compiled", result["findings"])
            self.assertEqual(result["contract"]["scope"], "captioned_tables")
            self.assertEqual(result["contract"]["border_widths_pt"], {
                "top": 1.5, "header": 0.5, "bottom": 1.5,
                "left": 0, "right": 0, "inside_h": 0, "inside_v": 0,
            })
            self.assertEqual([item["status"] for item in spec["clause_compliance"]],
                             ["pending_execution", "pending_execution"])

    def test_manifest_source_or_clause_digest_substitution_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, clauses_path, spec_path, manifest_path, _ = make_bundle(root)
            manifest = json.loads(manifest_path.read_text())
            manifest["normalized_source_sha256"] = "0" * 64
            manifest_path.write_text(json.dumps(manifest))
            result, binding = resolve_bound_requirement_docx(manifest_path, clauses_path, spec_path)
            self.assertIsNone(result)
            self.assertIn("requirements_extraction_binding_mismatch",
                          {item["code"] for item in binding["findings"]})

    def test_rule_only_format_spec_uses_manifest_bound_bytes_without_model_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, clauses_path, spec_path, manifest_path, _ = make_bundle(root)
            spec = json.loads(spec_path.read_text())
            spec.pop("semantic_review_provenance")
            spec["analysis_mode"] = "rule_only"
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            manifest = json.loads(manifest_path.read_text())
            manifest["format_spec_sha256"] = sha256_file(spec_path)
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

            bound_source, binding = resolve_bound_requirement_docx(
                manifest_path, clauses_path, spec_path,
            )

            self.assertEqual(bound_source, source)
            self.assertEqual(binding["status"], "verified")
            self.assertEqual(
                binding["format_spec_source_lineage"],
                "verified_run_source_path_and_sha256",
            )

    def test_normalized_source_keeps_original_path_lineage_and_checks_both_hashes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, clauses_path, spec_path, manifest_path, _ = make_bundle(root)
            normalized_dir = root / "normalized"
            normalized_dir.mkdir()
            normalized_source = normalized_dir / "requirements.docx"
            shutil.copyfile(source, normalized_source)

            spec = json.loads(spec_path.read_text())
            spec.pop("semantic_review_provenance")
            spec["analysis_mode"] = "rule_only"
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")

            manifest = json.loads(manifest_path.read_text())
            manifest.update({
                "source_document": str(source.resolve()),
                "source_bytes": source.stat().st_size,
                "source_sha256": sha256_file(source),
                "normalized_source_document": str(normalized_source.resolve()),
                "normalized_source_bytes": normalized_source.stat().st_size,
                "normalized_source_sha256": sha256_file(normalized_source),
                "format_spec_sha256": sha256_file(spec_path),
            })
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

            bound_source, binding = resolve_bound_requirement_docx(
                manifest_path, clauses_path, spec_path,
            )
            self.assertEqual(bound_source, normalized_source)
            self.assertEqual(binding["status"], "verified")

            wrong_original = root / "same-bytes-different-original-path.docx"
            shutil.copyfile(source, wrong_original)
            substituted_path = copy.deepcopy(manifest)
            substituted_path["source_document"] = str(wrong_original.resolve())
            manifest_path.write_text(json.dumps(substituted_path), encoding="utf-8")
            rejected_source, rejected = resolve_bound_requirement_docx(
                manifest_path, clauses_path, spec_path,
            )
            self.assertIsNone(rejected_source)
            self.assertIn("format_spec_source_lineage_mismatch",
                          {item["code"] for item in rejected["findings"]})

            substituted_hash = copy.deepcopy(manifest)
            substituted_hash["source_sha256"] = "0" * 64
            manifest_path.write_text(json.dumps(substituted_hash), encoding="utf-8")
            rejected_source, rejected = resolve_bound_requirement_docx(
                manifest_path, clauses_path, spec_path,
            )
            self.assertIsNone(rejected_source)
            self.assertIn("requirements_original_source_binding_mismatch",
                          {item["code"] for item in rejected["findings"]})

    def test_pipeline_manifest_can_bind_pipeline_owned_spec_enrichment(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, clauses_path, spec_path, manifest_path, _ = make_bundle(root)
            spec = json.loads(spec_path.read_text())
            spec.pop("semantic_review_provenance")
            spec["analysis_mode"] = "rule_only"
            parent_spec_sha = sha256_file(spec_path)
            spec["thesis_profile"] = {"metadata_status": "complete", "pending_fields": []}
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")

            manifest = json.loads(manifest_path.read_text())
            manifest.update({
                "source_document": str(source.resolve()),
                "source_bytes": source.stat().st_size,
                "source_sha256": sha256_file(source),
                "format_spec_sha256": parent_spec_sha,
            })
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

            def record(path: Path) -> dict:
                return {"path": str(path.resolve()), "bytes": path.stat().st_size,
                        "sha256": sha256_file(path)}

            pipeline_path = root / "pipeline-manifest.json"
            pipeline = {
                "run_id": "test-run",
                "inputs": {"requirements": record(source)},
                "format_spec_record": record(spec_path),
                "source_clause_record": record(clauses_path),
            }
            pipeline_path.write_text(json.dumps(pipeline), encoding="utf-8")

            bound_source, binding = resolve_bound_requirement_docx(
                manifest_path, clauses_path, spec_path, pipeline_path,
            )
            self.assertEqual(bound_source, source)
            self.assertEqual(binding["status"], "verified")
            self.assertEqual(
                binding["format_spec_source_lineage"],
                "verified_pipeline_manifest_current_spec",
            )

            pipeline["format_spec_record"]["sha256"] = "0" * 64
            pipeline_path.write_text(json.dumps(pipeline), encoding="utf-8")
            rejected_source, rejected = resolve_bound_requirement_docx(
                manifest_path, clauses_path, spec_path, pipeline_path,
            )
            self.assertIsNone(rejected_source)
            self.assertIn("format_spec_source_lineage_mismatch",
                          {item["code"] for item in rejected["findings"]})

    def test_deterministic_role_projection_replays_parent_and_rejects_tampering(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, clauses_path, parent_path, manifest_path, _ = make_bundle(root)
            parent = json.loads(parent_path.read_text())
            parent.pop("semantic_review_provenance")
            parent["analysis_mode"] = "rule_only"
            parent["requirements"] = [{
                "id": "R-cover-value", "role": "cover_field_value",
                "properties": {"font": {"cjk": "SimSun"}},
                "clause_ids": ["C-cover-value"], "evidence_ids": ["E-cover-value"],
                "field_instance_ids": ["I-cover-value"],
                "source_scope": {"kind": "explicit_clause_set", "anchor_clause_ids": [],
                                 "target_clause_ids": ["C-target-value"]},
            }]
            parent["content_instances"] = [{
                "id": "I-cover-value", "role": "cover_field_value", "text": "甲",
                "clause_ids": ["C-target-value"], "evidence_ids": ["E-target-value"],
            }]
            parent_path.write_text(json.dumps(parent, ensure_ascii=False), encoding="utf-8")
            relocated_parent_path = root / "same-bytes-other-path.json"
            relocated_parent_path.write_bytes(parent_path.read_bytes())
            manifest = json.loads(manifest_path.read_text())
            manifest["format_spec_sha256"] = sha256_file(parent_path)
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            projection = {
                "source_document": parent["source_document"], "run_id": parent["run_id"],
                "analysis_mode": "rule_only", "roles": {}, "status": "rule_resolved",
                "requirements": copy.deepcopy(parent["requirements"]),
                "content_instances": copy.deepcopy(parent["content_instances"]),
                "source_projection": {
                    "protocol": "deterministic_requirement_subset_v1",
                    "parent_format_spec": {"path": str(parent_path.resolve()),
                                            "sha256": sha256_file(parent_path)},
                    "selected_roles": ["cover_field_value"],
                    "projected_payload_sha256": hashlib.sha256(json.dumps(
                        {"requirements": parent["requirements"],
                         "content_instances": parent["content_instances"]},
                        ensure_ascii=False, sort_keys=True, separators=(",", ":"),
                    ).encode("utf-8")).hexdigest(),
                    "projection_code_sha256": sha256_file(ROOT / "scripts/source_table_contract.py"),
                },
            }
            projection_path = root / "focused-format-spec.json"
            projection_path.write_text(json.dumps(projection, ensure_ascii=False), encoding="utf-8")

            bound_source, binding = resolve_bound_requirement_docx(
                manifest_path, clauses_path, projection_path,
            )
            self.assertEqual(bound_source, source)
            self.assertEqual(binding["status"], "verified")
            self.assertEqual(
                binding["format_spec_source_lineage"],
                "verified_deterministic_subset_parent_projection",
            )

            for mutated_projection in (
                {**projection, "source_projection": {
                    **projection["source_projection"],
                    "parent_format_spec": {
                        **projection["source_projection"]["parent_format_spec"],
                        "path": str(relocated_parent_path.resolve()),
                    },
                }},
                {**projection, "source_projection": {
                    **projection["source_projection"],
                    "parent_format_spec": {
                        **projection["source_projection"]["parent_format_spec"],
                        "sha256": "0" * 64,
                    },
                }},
                {**projection, "source_projection": {
                    **projection["source_projection"], "projected_payload_sha256": "0" * 64,
                }},
                {**projection, "source_projection": {
                    **projection["source_projection"], "projection_code_sha256": "0" * 64,
                }},
            ):
                projection_path.write_text(
                    json.dumps(mutated_projection, ensure_ascii=False), encoding="utf-8",
                )
                rejected_source, rejected = resolve_bound_requirement_docx(
                    manifest_path, clauses_path, projection_path,
                )
                self.assertIsNone(rejected_source)
                self.assertIn("format_spec_source_lineage_mismatch",
                              {item["code"] for item in rejected["findings"]})

            without_audit = copy.deepcopy(projection)
            without_audit.pop("source_projection")
            projection_path.write_text(
                json.dumps(without_audit, ensure_ascii=False), encoding="utf-8",
            )
            rejected_source, rejected = resolve_bound_requirement_docx(
                manifest_path, clauses_path, projection_path,
            )
            self.assertIsNone(rejected_source)
            self.assertIn("format_spec_source_lineage_mismatch",
                          {item["code"] for item in rejected["findings"]})

            tampered = copy.deepcopy(projection)
            tampered["requirements"][0]["properties"]["font"]["cjk"] = "Arial"
            tampered["source_projection"]["projected_payload_sha256"] = hashlib.sha256(
                json.dumps(
                    {"requirements": tampered["requirements"],
                     "content_instances": tampered["content_instances"]},
                    ensure_ascii=False, sort_keys=True, separators=(",", ":"),
                ).encode("utf-8")
            ).hexdigest()
            projection_path.write_text(json.dumps(tampered, ensure_ascii=False), encoding="utf-8")
            rejected_source, rejected = resolve_bound_requirement_docx(
                manifest_path, clauses_path, projection_path,
            )
            self.assertIsNone(rejected_source)
            self.assertIn("format_spec_source_lineage_mismatch",
                          {item["code"] for item in rejected["findings"]})

    def test_clause_offset_or_source_location_tamper_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, _, _, _, clauses = make_bundle(root)
            bad_offset = copy.deepcopy(clauses)
            bad_offset[0]["source_span"]["start_offset"] = 1
            spec = {"requirements": [], "clause_compliance": []}
            spec["roles"] = {"table_caption": {"position": "above"}}
            result = compile_source_table_contract(source, bad_offset, spec)
            self.assertEqual(result["status"], "blocked")
            self.assertEqual(result["findings"][0]["reason"],
                             "source_span_offsets_do_not_match_source_paragraph")

            bad_location = copy.deepcopy(clauses)
            bad_location[1]["source_span"]["location"]["child_index"] = 9
            result = compile_source_table_contract(source, bad_location, spec)
            self.assertEqual(result["status"], "blocked")
            self.assertIn("source_document_locator_not_a_paragraph",
                          {item["reason"] for item in result["findings"]})

    def test_duplicate_topology_or_conflicting_spec_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, _, _, _, clauses = make_bundle(root)
            duplicate = copy.deepcopy(clauses)
            duplicate.append(copy.deepcopy(clauses[0]) | {"id": "C-three-duplicate"})
            result = compile_source_table_contract(source, duplicate,
                                                   {"requirements": [], "clause_compliance": []})
            self.assertEqual(result["status"], "blocked")
            self.assertEqual(result["findings"][0]["code"], "source_table_topology_ambiguous")

            conflict_spec = {"tables": {"style": "other"}, "requirements": [],
                             "clause_compliance": [],
                             "roles": {"table_caption": {"position": "above"}}}
            result = compile_source_table_contract(source, clauses, conflict_spec)
            self.assertEqual(result["status"], "blocked")
            self.assertEqual(result["findings"][0]["code"],
                             "source_table_contract_conflicts_with_format_spec")


if __name__ == "__main__":
    unittest.main()
