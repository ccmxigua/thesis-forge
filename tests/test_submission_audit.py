from __future__ import annotations

import hashlib
import json
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import submission_audit
from submission_audit import audit_docx
from render_attestation import sign

TEST_ATTESTATION_KEY = bytes.fromhex("11" * 32)


def add_field(paragraph, instruction: str, cached: str = "1") -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = cached
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def make_docx(path: Path, body_text: str = "第1章 引言\n正文", *, page: bool = True,
              styleref: bool = True, generic_header: bool = False) -> None:
    doc = Document()
    for line in body_text.split("\n"):
        doc.add_paragraph(line)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), "decimal")
    doc.sections[0]._sectPr.append(pg)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if page:
        add_field(footer, "PAGE \\* ARABIC")
    header = doc.sections[0].header.paragraphs[0]
    header.add_run("某大学博士/硕士学位论文" if generic_header else "University Doctoral Thesis")
    if styleref:
        add_field(header, 'STYLEREF "Heading 1" \\* MERGEFORMAT', "Heading 1")
    doc.save(path)


SPEC = {
    "page": {"page_number": {"body_format": "decimal", "alignment": "center"}},
    "roles": {"header": {"header_content": {"left_text": "University Doctoral Thesis",
                                                 "right_field": "styleref_heading_1"}}},
}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_pdf(path: Path, text: str = "Rendered thesis page 1", *,
             width: float = 595.28, height: float = 841.89,
             page_number: str | None = "1", header_text: str | None = "University Doctoral Thesis") -> dict[str, object]:
    import fitz

    document = fitz.open()
    page = document.new_page(width=width, height=height)
    if header_text is not None:
        page.insert_text((72, 36), header_text)
    page.insert_text((72, 100), text)
    if page_number is not None:
        point = fitz.Point(width / 2 - 3 * len(page_number), height - 36)
        page.insert_text(point, page_number)
    document.save(path)
    document.close()
    try:
        from pypdf import PdfReader
    except ImportError:
        from PyPDF2 import PdfReader
    extracted = "\n\f\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    return {
        "path": str(path.resolve()),
        "sha256": digest(path),
        "page_count": 1,
        "text_sha256": hashlib.sha256(extracted.encode("utf-8")).hexdigest(),
    }


def render_report(docx: Path, pdf: Path, text: str = "Rendered thesis page 1") -> dict[str, object]:
    report = {
        "schema_version": "1.0",
        "evidence_type": "microsoft_word_pdf_render",
        "renderer": {"name": "Microsoft Word", "version": "test"},
        "source_docx": {"path": str(docx.resolve()), "sha256": digest(docx)},
        "rendered_pdf": make_pdf(pdf, text),
        "word_export": {"story_count": 1, "field_count": 1, "updated_count": 1,
                        "failed_count": 0, "toc_count": 0},
    }
    report["attestation"] = {"algorithm": "HMAC-SHA256", "scope": "local_word_export_v1",
                             "signature": sign(report, TEST_ATTESTATION_KEY)}
    return report


def attest(report: dict[str, object]) -> dict[str, object]:
    report["word_export"] = {"story_count": 1, "field_count": 1, "updated_count": 1,
                             "failed_count": 0, "toc_count": 0}
    report["attestation"] = {"algorithm": "HMAC-SHA256", "scope": "local_word_export_v1",
                             "signature": sign(report, TEST_ATTESTATION_KEY)}
    return report


class SubmissionAuditTest(unittest.TestCase):
    def test_sidecar_loader_rejects_duplicate_json_keys(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "sidecar.json"
            path.write_text('{"submission_ready":true,"submission_ready":false}', encoding="utf-8")
            with self.assertRaises(ValueError):
                submission_audit._json(path)

    def test_placeholders_and_unresolved_reference_residue_block_submission(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); docx = td / "bad.docx"
            make_docx(docx, "研究方向：待确认\n见 推论[std_corollary) 的证明。")
            report = audit_docx(docx, SPEC)
            codes = {issue["code"] for issue in report["issues"]}
            self.assertIn("submission_placeholder_unconfirmed_metadata", codes)
            self.assertIn("unresolved_source_marker_unresolved_reference_fallback", codes)
            self.assertFalse(report["submission_ready"])

    def test_raw_citation_keys_without_bibliography_block_submission(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); docx = td / "missing-bibliography.docx"
            make_docx(
                docx,
                "第1章 引言\n已有研究[heston1993closedform]、[leland1985option]和[shreve2010stochastic]。",
            )
            report = audit_docx(docx, SPEC)
            codes = {issue["code"] for issue in report["issues"]}
            self.assertIn("unresolved_raw_citation_keys", codes)
            self.assertIn("missing_bibliography_section", codes)
            self.assertFalse(report["submission_ready"])

    def test_bibliography_and_numeric_citations_do_not_trigger_raw_key_audit(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); docx = td / "bibliography.docx"
            make_docx(
                docx,
                "第1章 引言\n已有研究[1-3]。\n参考文献\n[1] HESTON S L. A Closed-Form Solution[J].",
            )
            report = audit_docx(docx, SPEC)
            codes = {issue["code"] for issue in report["issues"]}
            self.assertNotIn("unresolved_raw_citation_keys", codes)
            self.assertNotIn("missing_bibliography_section", codes)

    def test_pdf_layout_accepts_footer_glyph_whose_box_crosses_footer_threshold(self) -> None:
        import fitz

        with tempfile.TemporaryDirectory() as td:
            pdf = Path(td) / "word-footer-position.pdf"
            document = fitz.open()
            page = document.new_page(width=595.2, height=841.92)
            # A 12pt glyph inserted at this Word-like baseline has y0 just above
            # the final 10% boundary while y1 extends below it.
            page.insert_text((310, 766), "10", fontsize=12)
            document.save(pdf)
            document.close()

            layout = submission_audit._pdf_layout_evidence(pdf)

        self.assertEqual(layout["footer_page_number_candidates"]["1"][0]["token"], "10")
        self.assertEqual(layout["page_numbering_runs"][0]["start"], 10)

    def test_rendered_number_coverage_allows_cover_zone_and_blank_versos(self) -> None:
        import fitz

        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"
            pdf = Path(td) / "cover-and-blank-versos.pdf"
            make_docx(path)
            document = fitz.open()
            page = document.new_page(width=595.28, height=841.89)
            page.insert_text((72, 72), "Unnumbered cover")
            page = document.new_page(width=595.28, height=841.89)
            page.insert_text((72, 72), "Front matter")
            page.insert_text((294, 806), "I")
            document.new_page(width=595.28, height=841.89)  # intentional blank verso
            page = document.new_page(width=595.28, height=841.89)
            page.insert_text((72, 72), "Body")
            page.insert_text((294, 806), "1")
            document.save(pdf)
            document.close()
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
            extracted = "\n\f\n".join((page.extract_text() or "") for page in PdfReader(str(pdf)).pages)
            report = attest({
                "schema_version": "1.0", "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": {
                    "path": str(pdf.resolve()), "sha256": digest(pdf), "page_count": 4,
                    "text_sha256": hashlib.sha256(extracted.encode("utf-8")).hexdigest(),
                },
            })
            audit = audit_docx(path, SPEC, report)

        failure_codes = {item["code"] for item in audit["render_validation"]["failures"]}
        self.assertNotIn("rendered_pdf_page_number_coverage_or_ambiguity", failure_codes)

    def setUp(self) -> None:
        self.key_patch = patch.object(submission_audit, "load_key", return_value=TEST_ATTESTATION_KEY)
        self.key_patch.start()

    def tearDown(self) -> None:
        self.key_patch.stop()

    def test_clean_serialized_docx_still_requires_render_proof(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"
            pdf = Path(td) / "clean.pdf"
            make_docx(path)
            audit = audit_docx(path, SPEC)
            self.assertTrue(audit["serialized_docx_valid"], audit["issues"])
            self.assertFalse(audit["submission_ready"])
            self.assertEqual(audit["render_validation"]["status"], "not_run")
            spoofed = audit_docx(path, SPEC, {"passed": True, "evidence": ["rendered.pdf"]})
            self.assertFalse(spoofed["submission_ready"])
            self.assertIn("unsupported_render_report_schema",
                          {x["code"] for x in spoofed["render_validation"]["failures"]})
            rendered = audit_docx(path, SPEC, render_report(path, pdf))
            self.assertTrue(rendered["submission_ready"])

    def test_pending_thesis_profile_blocks_submission_without_invalidating_serialization(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "pending-profile.docx"
            make_docx(path)
            audit = audit_docx(
                path,
                SPEC,
                thesis_profile={"metadata_status": "pending", "pending_fields": ["degree_category"]},
            )
            self.assertTrue(audit["serialized_docx_valid"], audit["issues"])
            self.assertFalse(audit["submission_ready"])
            self.assertIn(
                "incomplete_thesis_profile",
                {item["code"] for item in audit["submission_blockers"]},
            )
            self.assertNotIn(
                "incomplete_thesis_profile",
                {item["code"] for item in audit["issues"]},
            )

    def test_unresolved_docx_comments_block_submission(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "commented.docx"
            pdf = Path(td) / "commented.pdf"
            make_docx(path)
            comments_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="{submission_audit.W_NS}">
  <w:comment w:id="0" w:author="Template Author"><w:p><w:r><w:t>Remove before submission</w:t></w:r></w:p></w:comment>
</w:comments>'''.encode()
            with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("word/comments.xml", comments_xml)
            audit = audit_docx(path, SPEC, render_report(path, pdf))
            self.assertFalse(audit["submission_ready"])
            issue = next(item for item in audit["issues"]
                         if item["code"] == "unresolved_document_comments")
            self.assertEqual(issue["evidence"]["comment_count"], 1)
            self.assertEqual(issue["evidence"]["comments"][0]["author"], "Template Author")

    def test_render_pdf_with_word_error_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"
            pdf = Path(td) / "bad-render.pdf"
            make_docx(path)
            report = render_report(path, pdf, "Error! Reference source not found.")
            audit = audit_docx(path, SPEC, report)
            self.assertFalse(audit["submission_ready"])
            failures = {x["code"] for x in audit["render_validation"]["failures"]}
            self.assertIn("rendered_pdf_contains_word_errors", failures)

    def test_render_report_is_bound_to_exact_pdf_bytes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"
            pdf = Path(td) / "render.pdf"
            make_docx(path)
            report = render_report(path, pdf)
            pdf.write_bytes(pdf.read_bytes() + b"tampered")
            audit = audit_docx(path, SPEC, report)
            self.assertFalse(audit["submission_ready"])
            failures = {x["code"] for x in audit["render_validation"]["failures"]}
            self.assertIn("rendered_pdf_hash_mismatch", failures)

    def test_render_report_is_bound_to_exact_docx_bytes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"
            pdf = Path(td) / "render.pdf"
            make_docx(path)
            report = render_report(path, pdf)
            with path.open("ab") as handle:
                handle.write(b"tampered")
            audit = audit_docx(path, SPEC, report)
            self.assertFalse(audit["submission_ready"])
            failures = {x["code"] for x in audit["render_validation"]["failures"]}
            self.assertIn("source_docx_hash_mismatch", failures)

    def test_explicit_pdf_page_size_and_orientation_are_enforced(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"
            pdf = Path(td) / "wrong-size.pdf"
            make_docx(path)
            spec = json.loads(json.dumps(SPEC))
            spec["page"].update({"size": "A4", "orientation": "portrait"})
            report = attest({
                "schema_version": "1.0",
                "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": make_pdf(pdf, width=500, height=400),
            })
            audit = audit_docx(path, spec, report)
            failures = {x["code"] for x in audit["render_validation"]["failures"]}
            self.assertIn("rendered_pdf_page_size_mismatch", failures)
            self.assertIn("rendered_pdf_page_orientation_mismatch", failures)
            self.assertFalse(audit["submission_ready"])

    def test_pdf_layout_observations_do_not_create_unstated_school_rules(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"
            pdf = Path(td) / "layout.pdf"
            make_docx(path)
            import fitz
            document = fitz.open()
            for index, text in enumerate(("Repeated thesis page content long enough for duplicate detection.",
                                          "Repeated thesis page content long enough for duplicate detection.", ""), 1):
                page = document.new_page(width=595.28, height=841.89)
                page.insert_text((72, 36), "University Doctoral Thesis")
                if text:
                    page.insert_text((72, 100), text)
                page.insert_text((294, 806), str(index))
            document.save(pdf)
            document.close()
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
            extracted = "\n\f\n".join((page.extract_text() or "") for page in PdfReader(str(pdf)).pages)
            report = attest({
                "schema_version": "1.0",
                "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": {
                    "path": str(pdf.resolve()), "sha256": digest(pdf), "page_count": 3,
                    "text_sha256": hashlib.sha256(extracted.encode("utf-8")).hexdigest(),
                },
            })
            audit = audit_docx(path, SPEC, report)
            observations = {x["code"] for x in audit["render_validation"]["observations"]}
            self.assertIn("rendered_pdf_near_blank_page_candidates", observations)
            self.assertIn("rendered_pdf_duplicate_text_page_candidates", observations)
            self.assertTrue(audit["submission_ready"], audit["render_validation"])

    def test_self_asserted_word_renderer_without_attestation_cannot_pass(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"; pdf = Path(td) / "synthetic.pdf"
            make_docx(path)
            report = {
                "schema_version": "1.0", "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "self-asserted"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": make_pdf(pdf),
            }
            audit = audit_docx(path, SPEC, report)
            self.assertFalse(audit["submission_ready"])
            self.assertIn("word_export_attestation_invalid_or_missing",
                          {x["code"] for x in audit["render_validation"]["failures"]})

    def test_attestation_protocol_or_signed_payload_tampering_fails(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"; pdf = Path(td) / "clean.pdf"
            make_docx(path)
            report = render_report(path, pdf)
            report["attestation"]["scope"] = "untrusted_scope"
            scoped = audit_docx(path, SPEC, report)
            self.assertIn("word_export_attestation_invalid_or_missing",
                          {x["code"] for x in scoped["render_validation"]["failures"]})
            report = render_report(path, pdf)
            report["renderer"]["version"] = "tampered"
            tampered = audit_docx(path, SPEC, report)
            self.assertIn("word_export_attestation_invalid_or_missing",
                          {x["code"] for x in tampered["render_validation"]["failures"]})

    def test_missing_or_discontinuous_rendered_page_numbers_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"
            make_docx(path)
            missing_pdf = Path(td) / "missing.pdf"
            missing_audit = audit_docx(path, SPEC, {
                "schema_version": "1.0", "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": make_pdf(missing_pdf, page_number=None),
            })
            missing_codes = {x["code"] for x in missing_audit["render_validation"]["failures"]}
            self.assertIn("rendered_pdf_no_footer_page_number_candidates", missing_codes)

            import fitz
            sequence_pdf = Path(td) / "sequence.pdf"
            document = fitz.open()
            for page_number in (1, 3):
                page = document.new_page(width=595.28, height=841.89)
                page.insert_text((72, 72), f"Rendered thesis page {page_number}")
                page.insert_text((294, 806), str(page_number))
            document.save(sequence_pdf)
            document.close()
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
            extracted = "\n\f\n".join((page.extract_text() or "") for page in PdfReader(str(sequence_pdf)).pages)
            sequence_report = {
                "schema_version": "1.0", "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": {
                    "path": str(sequence_pdf.resolve()), "sha256": digest(sequence_pdf), "page_count": 2,
                    "text_sha256": hashlib.sha256(extracted.encode("utf-8")).hexdigest(),
                },
            }
            sequence_audit = audit_docx(path, SPEC, sequence_report)
            sequence_codes = {x["code"] for x in sequence_audit["render_validation"]["failures"]}
            self.assertIn("rendered_pdf_page_number_sequence_discontinuous", sequence_codes)

            gap_pdf = Path(td) / "gap.pdf"
            document = fitz.open()
            for index, footer in enumerate(("1", None, "3"), 1):
                page = document.new_page(width=595.28, height=841.89)
                # Keep content inside the body band so the missing middle footer
                # cannot be mistaken for an intentional blank verso.
                page.insert_text((72, 200), f"Rendered thesis page {index}")
                if footer is not None:
                    page.insert_text((294, 806), footer)
            document.save(gap_pdf); document.close()
            extracted = "\n\f\n".join((page.extract_text() or "") for page in PdfReader(str(gap_pdf)).pages)
            gap_report = {
                "schema_version": "1.0", "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": {"path": str(gap_pdf.resolve()), "sha256": digest(gap_pdf),
                                 "page_count": 3,
                                 "text_sha256": hashlib.sha256(extracted.encode("utf-8")).hexdigest()},
            }
            gap_audit = audit_docx(path, SPEC, gap_report)
            self.assertIn("rendered_pdf_page_number_coverage_or_ambiguity",
                          {x["code"] for x in gap_audit["render_validation"]["failures"]})

            roman_pdf = Path(td) / "roman.pdf"
            roman_report = {
                "schema_version": "1.0", "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": make_pdf(roman_pdf, page_number="I"),
            }
            roman_audit = audit_docx(path, SPEC, roman_report)
            roman_codes = {x["code"] for x in roman_audit["render_validation"]["failures"]}
            self.assertIn("rendered_pdf_page_number_format_mismatch", roman_codes)

    def test_rendered_header_and_toc_rules_use_final_pdf_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "clean.docx"
            make_docx(path)
            missing_header_pdf = Path(td) / "missing-header.pdf"
            missing_header_report = {
                "schema_version": "1.0", "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": make_pdf(missing_header_pdf, header_text=None),
            }
            missing_header = audit_docx(path, SPEC, missing_header_report)
            self.assertIn("rendered_pdf_fixed_header_text_missing",
                          {x["code"] for x in missing_header["render_validation"]["failures"]})

            toc_spec = json.loads(json.dumps(SPEC))
            toc_spec["document_structure"] = {"toc_depth": 3}
            missing_toc_pdf = Path(td) / "missing-toc.pdf"
            missing_toc = audit_docx(path, toc_spec, {
                "schema_version": "1.0", "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": make_pdf(missing_toc_pdf),
            })
            self.assertIn("rendered_pdf_toc_page_missing",
                          {x["code"] for x in missing_toc["render_validation"]["failures"]})

            import fitz
            toc_pdf = Path(td) / "toc.pdf"
            document = fitz.open()
            page = document.new_page(width=595.28, height=841.89)
            page.insert_text((72, 36), "University Doctoral Thesis")
            page.insert_text((270, 100), "CONTENTS")
            page.insert_text((72, 145), "Chapter 1 Introduction ........ 1")
            page.insert_text((294, 806), "1")
            document.save(toc_pdf)
            document.close()
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
            extracted = "\n\f\n".join((page.extract_text() or "") for page in PdfReader(str(toc_pdf)).pages)
            toc_report = attest({
                "schema_version": "1.0", "evidence_type": "microsoft_word_pdf_render",
                "renderer": {"name": "Microsoft Word", "version": "test"},
                "source_docx": {"path": str(path.resolve()), "sha256": digest(path)},
                "rendered_pdf": {
                    "path": str(toc_pdf.resolve()), "sha256": digest(toc_pdf), "page_count": 1,
                    "text_sha256": hashlib.sha256(extracted.encode("utf-8")).hexdigest(),
                },
            })
            passing = audit_docx(path, toc_spec, toc_report)
            self.assertTrue(passing["render_validation"]["rendered_verified"], passing["render_validation"])

    def test_fault_case_manifest_is_synthetic_and_tracks_regression_codes(self) -> None:
        manifest = json.loads((ROOT / "tests/golden/submission-audit-fault-cases.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["schema_version"], "1.0")
        cases = manifest["cases"]
        self.assertGreaterEqual(len(cases), 8)
        self.assertEqual(len({case["id"] for case in cases}), len(cases))
        expected_codes = {
            code for case in cases
            for key in ("expected_issue_codes", "expected_render_failure_codes")
            for code in case.get(key, [])
        }
        regression_codes = {
            "missing_serialized_page_fields", "page_field_section_coverage_or_format_mismatch",
            "missing_serialized_styleref", "unresolved_degree_header_placeholder",
            "unresolved_source_marker_bracket_reference", "unresolved_source_marker_latex_ref",
            "duplicate_equation_reference_prefix", "chapter_count_claim_mismatch",
            "rendered_pdf_contains_word_errors", "rendered_pdf_no_footer_page_number_candidates",
            "rendered_pdf_page_number_sequence_discontinuous", "rendered_pdf_page_size_mismatch",
            "rendered_pdf_page_orientation_mismatch",
        }
        self.assertEqual(expected_codes, regression_codes)

    def test_missing_page_field_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "missing-page.docx"
            make_docx(path, page=False)
            audit = audit_docx(path, SPEC)
            codes = {issue["code"] for issue in audit["issues"]}
            self.assertIn("missing_serialized_page_fields", codes)
            self.assertIn("page_field_section_coverage_or_format_mismatch", codes)
            self.assertFalse(audit["serialized_docx_valid"])

    def test_missing_styleref_and_generic_degree_placeholder_fail(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "bad-header.docx"
            make_docx(path, styleref=False, generic_header=True)
            audit = audit_docx(path, SPEC)
            codes = {issue["code"] for issue in audit["issues"]}
            self.assertIn("missing_serialized_styleref", codes)
            self.assertIn("unresolved_degree_header_placeholder", codes)

    def test_unresolved_markers_and_duplicate_equation_prefix_fail(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "markers.docx"
            make_docx(path, "第1章 引言\n见图 [fig:gcn) 和 \\ref{tab:a}。式 式（3.2）如下。")
            audit = audit_docx(path, SPEC)
            codes = {issue["code"] for issue in audit["issues"]}
            self.assertIn("unresolved_source_marker_bracket_reference", codes)
            self.assertIn("unresolved_source_marker_latex_ref", codes)
            self.assertIn("duplicate_equation_reference_prefix", codes)

    def test_chapter_claim_mismatch_fails(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "chapters.docx"
            make_docx(path, "全文共六章。\n第1章 引言\n第2章 方法\n第3章 实验")
            audit = audit_docx(path, SPEC)
            self.assertIn("chapter_count_claim_mismatch", {x["code"] for x in audit["issues"]})

    def test_benwen_chapter_claim_mismatch_fails(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "benwen-chapters.docx"
            make_docx(path, "本文共分为六章。\n第1章 引言\n第2章 方法\n第3章 实验")
            audit = audit_docx(path, SPEC)
            self.assertIn("chapter_count_claim_mismatch", {x["code"] for x in audit["issues"]})

    def test_localized_heading_1_is_counted_but_toc_1_is_not(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "localized-headings.docx"
            doc = Document()
            toc_style = doc.styles.add_style("toc 1", 1)
            toc_style.style_id = "TOC1"
            # Simulate a cached TOC result whose hidden PAGEREF page number is
            # exposed by raw OOXML text extraction.
            doc.add_paragraph("第一章  示例1", style=toc_style)
            doc.add_paragraph("第二章  示例3", style=toc_style)
            first = doc.add_heading("第一章  示例", level=1)
            pict = OxmlElement("w:pict")
            textbox = etree.Element("{urn:schemas-microsoft-com:vml}textbox")
            content = OxmlElement("w:txbxContent")
            note = OxmlElement("w:p")
            note_run = OxmlElement("w:r")
            note_text = OxmlElement("w:t")
            note_text.text = "黑体三号，居中，固定行距20磅"
            note_run.append(note_text)
            note.append(note_run)
            content.append(note)
            textbox.append(content)
            pict.append(textbox)
            first._p.insert(1, pict)
            doc.add_heading("第二章  示例", level=1)
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:fmt"), "decimal")
            doc.sections[0]._sectPr.append(pg)
            footer = doc.sections[0].footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_field(footer, "PAGE \\* ARABIC")
            doc.save(path)

            audit = audit_docx(path, {"page": {"page_number": {"body_format": "decimal", "alignment": "center"}}})
            self.assertNotIn("non_contiguous_chapter_numbering", {x["code"] for x in audit["issues"]})

    def test_mixed_plain_and_prefixed_chapter_headings_are_contiguous(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "mixed-headings.docx"
            doc = Document()
            doc.add_heading("1 引言", level=1)
            doc.add_heading("第2章 方法", level=1)
            doc.add_heading("第3章 实验", level=1)
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:fmt"), "decimal")
            doc.sections[0]._sectPr.append(pg)
            footer = doc.sections[0].footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_field(footer, "PAGE \\* ARABIC")
            doc.save(path)

            audit = audit_docx(path, {"page": {"page_number": {
                "body_format": "decimal", "alignment": "center",
            }}})
            self.assertNotIn("non_contiguous_chapter_numbering", {x["code"] for x in audit["issues"]})

    def test_references_heading_does_not_crash_without_ordering_policy(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "references.docx"
            make_docx(path, "第1章 引言\n参考文献")
            audit = audit_docx(path, SPEC)
            self.assertIsInstance(audit, dict)
            self.assertNotIn("docx_parse_failed", {x["code"] for x in audit["issues"]})

    def test_strict_wordprocessingml_namespace_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "strict.docx"
            make_docx(path)
            with zipfile.ZipFile(path, "a") as zf:
                xml = zf.read("word/document.xml").decode("utf-8")
                xml = xml.replace(
                    "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                    "http://purl.oclc.org/ooxml/wordprocessingml/main",
                )
                zf.writestr("word/document.xml", xml)
            audit = audit_docx(path, SPEC)
            self.assertFalse(audit["serialized_docx_valid"])
            self.assertIn("unsupported_wordprocessingml_namespace", {x["code"] for x in audit["issues"]})

    def test_page_field_switch_must_match_section_format(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "wrong-switch.docx"
            make_docx(path)
            spec = json.loads(json.dumps(SPEC))
            spec["page"]["page_number"]["body_format"] = "roman_upper"
            # Section properties are what the audit treats as authoritative for
            # each section; make them upper-Roman while leaving the field ARABIC.
            doc = Document(path)
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:fmt"), "upperRoman")
            doc.sections[0]._sectPr.append(pg)
            doc.save(path)
            audit = audit_docx(path, spec)
            issue = next(x for x in audit["issues"] if x["code"] == "page_field_section_coverage_or_format_mismatch")
            self.assertEqual(issue["evidence"][0]["expected_field_switch"], "ROMAN")

    def test_linked_section_inherits_previous_footer_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "linked.docx"
            doc = Document(); doc.add_heading("第1章 引言", 1)
            add_field(doc.sections[0].footer.paragraphs[0], "PAGE \\* ARABIC", "1")
            second = doc.add_section(WD_SECTION.NEW_PAGE)
            second.footer.is_linked_to_previous = True
            doc.save(path)
            spec = json.loads(json.dumps(SPEC))
            spec["page"]["page_number"]["body_start_selector"] = {
                "strategy": "section_index", "section_index": 1
            }
            audit = audit_docx(path, spec)
            codes = {x["code"] for x in audit["issues"]}
            self.assertNotIn("page_field_section_coverage_or_format_mismatch", codes)
            self.assertTrue(audit["evidence"]["sections"][1].get("footer_inherited", {}).get("default"))

    def test_semantic_body_selector_audits_front_and_body_formats_by_actual_heading_section(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "semantic-body-start.docx"
            doc = Document()
            doc.add_paragraph("摘 要")
            front = doc.sections[0]
            pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "lowerRoman")
            front._sectPr.append(pg)
            add_field(front.footer.paragraphs[0], "PAGE \\* roman", "i")
            body = doc.add_section(WD_SECTION.NEW_PAGE)
            body.footer.is_linked_to_previous = False
            inherited_pg = body._sectPr.find(qn("w:pgNumType"))
            if inherited_pg is not None:
                body._sectPr.remove(inherited_pg)
            pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "decimal")
            body._sectPr.append(pg)
            add_field(body.footer.paragraphs[0], "PAGE \\* ARABIC", "1")
            doc.add_heading("第1章 引言", 1)
            doc.save(path)
            spec = json.loads(json.dumps(SPEC))
            spec["page"]["page_number"].update({
                "front_matter_format": "roman",
                "body_format": "decimal",
                "body_start_selector": {"strategy": "first_heading_1"},
            })
            audit = audit_docx(path, spec)
            self.assertNotIn(
                "page_field_section_coverage_or_format_mismatch",
                {item["code"] for item in audit["issues"]},
            )

    def test_heading_text_body_selector_ignores_cached_toc_heading(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "toc-and-body-heading.docx"
            doc = Document()
            toc_style = doc.styles.add_style("toc 1", 1)
            toc_style.style_id = "TOC1"
            doc.add_paragraph("第1章  引言1", style=toc_style)

            front = doc.sections[0]
            pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "lowerRoman")
            front._sectPr.append(pg)
            add_field(front.footer.paragraphs[0], "PAGE \\* roman", "i")

            body = doc.add_section(WD_SECTION.NEW_PAGE)
            body.footer.is_linked_to_previous = False
            inherited_pg = body._sectPr.find(qn("w:pgNumType"))
            if inherited_pg is not None:
                body._sectPr.remove(inherited_pg)
            pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "decimal")
            body._sectPr.append(pg)
            add_field(body.footer.paragraphs[0], "PAGE \\* ARABIC", "1")
            doc.add_heading("第1章  引言", 1)
            doc.save(path)

            spec = json.loads(json.dumps(SPEC))
            spec["page"]["page_number"].update({
                "front_matter_format": "roman",
                "body_format": "decimal",
                "body_start_selector": {
                    "strategy": "heading_text",
                    "heading_text_pattern": r"^第\s*1\s*章",
                },
            })
            audit = audit_docx(path, spec)
            codes = {item["code"] for item in audit["issues"]}
            self.assertNotIn("ambiguous_body_start_selector", codes)
            self.assertNotIn("page_field_section_coverage_or_format_mismatch", codes)

    def test_explicit_front_restart_allows_intentionally_unnumbered_leading_section(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "unnumbered-leading-matter.docx"
            doc = Document(); doc.add_paragraph("题名页")
            front = doc.sections[0]
            pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "lowerRoman")
            front._sectPr.append(pg)
            abstract = doc.add_section(WD_SECTION.NEW_PAGE)
            abstract.footer.is_linked_to_previous = False
            pg = abstract._sectPr.find(qn("w:pgNumType"))
            if pg is None:
                pg = OxmlElement("w:pgNumType"); abstract._sectPr.append(pg)
            pg.set(qn("w:fmt"), "lowerRoman"); pg.set(qn("w:start"), "1")
            add_field(abstract.footer.paragraphs[0], "PAGE \\* roman", "i")
            doc.add_paragraph("摘 要")
            body = doc.add_section(WD_SECTION.NEW_PAGE)
            body.footer.is_linked_to_previous = False
            inherited_pg = body._sectPr.find(qn("w:pgNumType"))
            if inherited_pg is not None:
                body._sectPr.remove(inherited_pg)
            pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "decimal"); pg.set(qn("w:start"), "1")
            body._sectPr.append(pg); add_field(body.footer.paragraphs[0], "PAGE \\* ARABIC", "1")
            doc.add_heading("第1章 引言", 1); doc.save(path)
            spec = json.loads(json.dumps(SPEC))
            spec["page"]["page_number"].update({
                "front_matter_format": "roman", "body_format": "decimal",
                "front_matter_start": 1, "body_start": 1,
                "body_start_selector": {"strategy": "first_heading_1"},
            })
            audit = audit_docx(path, spec)
            self.assertNotIn(
                "page_field_section_coverage_or_format_mismatch",
                {item["code"] for item in audit["issues"]},
            )

    def test_template_profile_without_render_report_does_not_crash(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "no-render.docx"
            make_docx(path)
            with patch("template_audit.audit_template", return_value={
                "official_template_structure_valid": True,
                "manual_required_fields": [],
                "failures": [],
            }):
                audit = audit_docx(path, SPEC, template_profile_path=Path(td) / "profile.json")
            self.assertEqual(audit["render_validation"]["status"], "not_run")
            self.assertEqual(audit["render_validation"]["evidence"], {})
            self.assertFalse(audit["submission_ready"])

    def test_cli_rejects_invalid_spec_and_output_alias_before_audit(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); docx = root / "input.docx"; spec = root / "invalid.json"
            make_docx(docx); original = docx.read_bytes()
            spec.write_text(json.dumps({"page": {"page_number": {"body_format": "binary"}}}))
            with self.assertRaises(SystemExit) as invalid:
                submission_audit.main([str(docx), "--format-spec", str(spec), "--out", str(root / "out.json")])
            self.assertEqual(invalid.exception.code, 2)
            with self.assertRaises(SystemExit) as alias:
                submission_audit.main([str(docx), "--out", str(docx)])
            self.assertEqual(alias.exception.code, 2)
            self.assertEqual(docx.read_bytes(), original)

    def test_cli_refuses_hard_link_output_aliasing_docx(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); docx = root / "input.docx"; output = root / "audit.json"
            make_docx(docx); os.link(docx, output)
            with self.assertRaises(SystemExit) as raised:
                submission_audit.main([str(docx), "--out", str(output)])
            self.assertEqual(raised.exception.code, 2)
            self.assertTrue(docx.read_bytes().startswith(b"PK"))


if __name__ == "__main__":
    unittest.main()
