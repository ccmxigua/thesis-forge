"""Current table geometry may authorize rereview, never semantic repair or a pass."""
from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))
import host_agent_bridge as bridge
import native_semantic_review as native
import requirements_engine as engine
import thesis_format_pipeline as pipeline
import test_authoring_correction_routing as authoring_tests
from semantic_contract import sha256_json, sha256_file, evidence_payload, request_body_sha256
from semantic_contract import attach_request_provenance
from semantic_source_references import (
    REFERENCE_PROTOCOL, build_source_reference_packet, compile_source_reference_response,
    bind_validated_source_reference_selections,
)
from table_source_context import (
    build_table_structure_context, table_context_retry_is_source_bound, table_retry_feedback_is_source_bound, text_sha256,
)


class TableStructureReviewTests(unittest.TestCase):
    @staticmethod
    def fixture(label="批准日期", mode=None):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = label
        table.cell(0, 1).text = "年    月    日"
        table.cell(1, 0).text = "其他日期"
        table.cell(1, 1).text = "年    月    日"
        if mode == "multiple_paragraphs":
            table.cell(0, 0).add_paragraph("另一个字段")
        elif mode == "blank":
            table.cell(0, 0).text = ""
        elif mode in {"vertical", "horizontal", "span"}:
            cell_pr = table.cell(0, 1)._tc.get_or_add_tcPr()
            element = OxmlElement({"vertical": "w:vMerge", "horizontal": "w:hMerge", "span": "w:gridSpan"}[mode])
            element.set(qn("w:val"), "2" if mode == "span" else "restart")
            cell_pr.append(element)
        elif mode == "nested":
            table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套标签"
        elif mode == "grid_before":
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            element = OxmlElement("w:gridBefore")
            element.set(qn("w:val"), "1")
            tr_pr.append(element)
        elif mode == "rtl":
            element = OxmlElement("w:bidiVisual")
            table._tbl.tblPr.append(element)
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "current-source.docx"
            doc.save(path)
            evidence_doc = engine.extract_document_evidence(path)
        evidence = {e["id"]: e for e in evidence_doc["evidence"]}
        target = next(e for e in evidence.values() if e["location"].get("row") == 0
                      and e["location"].get("column") == 1)
        source = target["text"]
        clause = {"id": "arbitrary-clause", "text": "年 月 日", "source_kind": "table_cell",
                  "evidence_ids": [target["id"]], "source_span": {
                      "evidence_id": target["id"], "text": source, "source_sha256": text_sha256(source),
                      "start_offset": 0, "end_offset": len(source)}}
        provenance = {"run_id": "new-run", "source_sha256": "a" * 64,
                      "clause_sha256": "b" * 64, "evidence_sha256": sha256_json(evidence),
                      "request_sha256": "d" * 64}
        chunk = {"clauses": [clause], "evidence_context": evidence,
                 "provenance": provenance, "case_id": "arbitrary-case"}
        candidate = {"contract_version": "3.0", "provenance": provenance,
                     "requirements": [{"role": "cover", "properties": {"field": "candidate-proposal"},
                                       "clause_ids": [clause["id"]], "evidence_ids": [target["id"]]}],
                     "clause_reviews": [{"clause_id": clause["id"], "classification": "executable",
                                         "reason": "Structural proposal, not code inference."}]}
        return chunk, candidate

    @staticmethod
    def request(chunk, candidate):
        return native.build_obligation_coverage_request(candidate, chunk, run_id="new-run", chunk_index=1)

    @staticmethod
    def uncertain(check):
        source = check["document_text"]
        return {"results": [{"check_id": check["check_id"], "verdict": "uncertain",
                             "rationale": "The isolated date lacks a unique target.",
                             "identified_obligations": [{"source_quote": source, "disposition": "ambiguous",
                                                         "obligation_summary": "Date target", "requirement_refs": []}],
                             "evidence_quotes": [source], "machine_obligation_ids": []}]}

    @staticmethod
    def run_receipt_review(candidate, chunk, output, responses, calls):
        """No provider: exercise the real source compiler and receipt reconstruction."""
        def reviewer(request, **kwargs):
            calls.append(copy.deepcopy(request))
            packet = build_source_reference_packet(request)
            check = packet["checks"][0]
            span = next(s for s in check["source_spans"] if s["text"] == check["document_text"])
            result = copy.deepcopy(responses[len(calls) - 1]["results"][0])
            result.pop("evidence_quotes")
            result.pop("machine_obligation_ids")
            result["evidence_refs"] = [span["ref_id"]]
            for obligation in result["identified_obligations"]:
                obligation.pop("source_quote")
                obligation["source_ref"] = span["ref_id"]
            raw = {"results": [result]}
            compiled, compilation = compile_source_reference_response(
                raw, request, native.OBLIGATION_COVERAGE_SCHEMA, coverage=True,
                provider_nullable_optionals=True)
            original_compiled = copy.deepcopy(compiled)
            out = kwargs["output_dir"]
            out.mkdir(parents=True, exist_ok=True)
            paths = {"request_path": out / "request.json", "source_reference_packet_path": out / "source-reference-packet.json",
                     "raw_response_path": out / "raw-response.json", "compiled_response_path": out / "compiled-response.json",
                     "response_path": out / "response.json", "source_reference_compilation_path": out / "source-reference-compilation.json"}
            for key, data in (("request_path", request), ("source_reference_packet_path", packet),
                              ("raw_response_path", raw), ("compiled_response_path", original_compiled)):
                bridge._write_json(paths[key], data)
            results = native.validate_obligation_coverage_response(compiled, request["checks"])
            compilation = bind_validated_source_reference_selections(compilation, original_compiled, compiled, request)
            bridge._write_json(paths["response_path"], compiled)
            bridge._write_json(paths["source_reference_compilation_path"], compilation)
            return {"status": "completed", "protocol": native.OBLIGATION_COVERAGE_PROTOCOL,
                    "adapter_id": "codex", "host_runtime": "codex", "results": results,
                    "request_sha256": sha256_json(request), "response_sha256": sha256_file(paths["response_path"]),
                    "canonical_response_sha256": sha256_json(compiled), "source_reference_protocol": REFERENCE_PROTOCOL,
                    **{k: str(v.resolve()) for k, v in paths.items()},
                    "source_reference_packet_sha256": sha256_file(paths["source_reference_packet_path"]),
                    "raw_response_file_sha256": sha256_file(paths["raw_response_path"]),
                    "compiled_response_sha256": sha256_file(paths["compiled_response_path"]),
                    "source_reference_compilation_sha256": sha256_file(paths["source_reference_compilation_path"])}
        with patch.object(bridge, "run_native_semantic_review", side_effect=reviewer), patch.object(bridge.time, "sleep"):
            return bridge._run_independent_obligation_coverage_review(
                candidate, chunk, review_dir=Path(output), run_id="new-run", chunk_index=1,
                attempt=1, host_runtime="codex", model="gpt-5.6-luna", timeout=5,
                agent_id="main", runner="exec", binary="codex", config_path=None,
                controller=bridge.RunController())

    def test_current_row_facts_and_exact_source_catalog_are_separate(self):
        chunk, candidate = self.fixture()
        request = self.request(chunk, candidate)
        check = request["checks"][0]
        context = check["review_context"]["table_structure_context"]
        self.assertEqual(context["relationship"], "same_row_immediate_left_unmerged")
        self.assertEqual(context["source_row"]["cells"][0]["paragraphs"][0]["text"], "批准日期")
        self.assertEqual(check["document_text"], "年    月    日")
        self.assertTrue(table_context_retry_is_source_bound(check))
        packet = build_source_reference_packet(request)
        self.assertTrue(all(s["text"] == "年    月    日"
                            for s in packet["checks"][0]["source_spans"]))
        self.assertNotIn("批准日期", json.dumps(packet["checks"][0]["source_spans"], ensure_ascii=False))

    def test_arbitrary_left_label_does_not_compile_or_choose_a_semantic_field(self):
        for label in ("完成时间", "未知字段", "批准日期"):
            chunk, candidate = self.fixture(label)
            before = copy.deepcopy(candidate)
            self.assertTrue(table_context_retry_is_source_bound(self.request(chunk, candidate)["checks"][0]))
            self.assertEqual(candidate, before)
            self.assertEqual(candidate["requirements"][0]["properties"], {"field": "candidate-proposal"})

    def test_neighbor_retained_across_chunk_boundary_and_model_compaction(self):
        chunk, candidate = self.fixture()
        eid = chunk["clauses"][0]["source_span"]["evidence_id"]
        chunk["evidence_context"] = {eid: chunk["evidence_context"][eid]}
        check = self.request(chunk, candidate)["checks"][0]
        self.assertTrue(table_context_retry_is_source_bound(check))
        compact = bridge.compact_model_packet(chunk)
        self.assertEqual(compact["evidence_context"][eid]["table_row_context"],
                         chunk["evidence_context"][eid]["table_row_context"])

    def test_self_consistent_changed_row_cannot_replace_canonical_chunk_source(self):
        chunk, _candidate = self.fixture()
        evidence_doc = {"evidence": list(chunk["evidence_context"].values())}
        full = engine.build_llm_request([], chunk["clauses"], evidence_doc, {}, "full",
                                       contract_version="3.0")
        full = attach_request_provenance(full, source_sha256="a" * 64,
            evidence_doc=evidence_doc, clauses=chunk["clauses"], run_id="new-run")
        with tempfile.TemporaryDirectory() as td:
            manifest = engine.prepare_host_agent_review_packets(full, chunk["clauses"], evidence_doc,
                "a" * 64, Path(td), chunk_size=1)
            packets = json.loads((Path(td) / "llm-request-chunks.json").read_text())
        engine.validate_host_review_chunk_source_projection(full, packets, manifest)
        target = next(iter(packets[0]["evidence_context"].values()))
        row = target["table_row_context"]
        neighbor = row["cells"][0]["paragraphs"][0]
        neighbor["text"] = "伪造的字段"
        neighbor["text_sha256"] = text_sha256(neighbor["text"])
        row["row_sha256"] = sha256_json({k: v for k, v in row.items() if k != "row_sha256"})
        packets[0]["provenance"]["evidence_sha256"] = sha256_json(evidence_payload({
            "evidence": list(packets[0]["evidence_context"].values()),
            "page_evidence": packets[0].get("page_evidence", {}),
            "structure_evidence": packets[0].get("document_structure", {}),
        }))
        packets[0]["provenance"]["request_sha256"] = request_body_sha256(packets[0])
        manifest["chunk_projection_sha256"] = sha256_json(engine._chunk_projection(packets))
        with self.assertRaises(ValueError):
            engine.validate_host_review_chunk_source_projection(full, packets, manifest)

    def test_merged_blank_nested_and_nonunique_geometry_never_authorizes_retry(self):
        for mode in ("multiple_paragraphs", "blank", "vertical", "horizontal", "span", "nested", "grid_before", "rtl"):
            with self.subTest(mode=mode):
                chunk, candidate = self.fixture(mode=mode)
                check = self.request(chunk, candidate)["checks"][0]
                self.assertFalse(table_context_retry_is_source_bound(check))
                with self.assertRaises(native.NativeSemanticReviewError) as caught:
                    native.validate_obligation_coverage_response(self.uncertain(check), [check])
                self.assertNotIsInstance(caught.exception, native.TableContextUncertaintyError)

    def test_legacy_evidence_is_not_upgraded_from_free_neighbor_words(self):
        chunk, candidate = self.fixture()
        for evidence in chunk["evidence_context"].values():
            evidence.pop("table_row_context", None)
        check = self.request(chunk, candidate)["checks"][0]
        self.assertNotIn("table_structure_context", check["review_context"])
        self.assertFalse(table_context_retry_is_source_bound(check))

    def test_old_hash_cross_row_cross_part_and_forged_context_are_rejected(self):
        for mutation in ("hash", "row", "part", "target_text", "available_neighbor"):
            chunk, candidate = self.fixture()
            target = chunk["evidence_context"][chunk["clauses"][0]["source_span"]["evidence_id"]]
            row = target["table_row_context"]
            if mutation == "hash":
                row["row_sha256"] = "0" * 64
            elif mutation in {"row", "part"}:
                row[mutation] = 20 if mutation == "row" else "word/header1.xml"
                row["row_sha256"] = sha256_json({k: v for k, v in row.items() if k != "row_sha256"})
            elif mutation == "target_text":
                row["cells"][1]["paragraphs"][0]["text"] = "旧文本"
                row["row_sha256"] = sha256_json({k: v for k, v in row.items() if k != "row_sha256"})
            else:
                neighbor = row["cells"][0]["paragraphs"][0]["evidence_id"]
                chunk["evidence_context"][neighbor]["location"]["row"] = 10
            with self.subTest(mutation=mutation), self.assertRaises(native.NativeSemanticReviewError):
                self.request(chunk, candidate)
        chunk, candidate = self.fixture()
        check = self.request(chunk, candidate)["checks"][0]
        check["review_context"]["table_structure_context"]["immediate_left_column"] = 999
        self.assertFalse(table_context_retry_is_source_bound(check))

    def test_uncertainty_never_passes_and_does_not_mask_unrepresented_work(self):
        chunk, candidate = self.fixture()
        check = self.request(chunk, candidate)["checks"][0]
        with self.assertRaises(native.TableContextUncertaintyError):
            native.validate_obligation_coverage_response(self.uncertain(check), [check])
        mixed = self.uncertain(check)
        mixed["results"][0]["identified_obligations"].append({"source_quote": check["document_text"],
            "disposition": "unrepresented", "requirement_refs": []})
        with self.assertRaises(native.NativeSemanticReviewError) as caught:
            native.validate_obligation_coverage_response(mixed, [check])
        self.assertNotIsInstance(caught.exception, native.TableContextUncertaintyError)

    def test_bounded_rereview_keeps_candidate_and_second_uncertainty_fails_closed(self):
        chunk, candidate = self.fixture()
        check = self.request(chunk, candidate)["checks"][0]
        before = copy.deepcopy((chunk, candidate))
        calls = []
        with tempfile.TemporaryDirectory() as td:
            with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                authoring_tests.AuthoringCorrectionRoutingTests.run_review(
                    candidate, chunk, td, self.uncertain(check), calls)
            self.assertEqual(len(calls), 2)
            self.assertFalse(caught.exception.retryable)
            self.assertEqual(calls[1]["retry_feedback"]["code"], native.TableContextUncertaintyError.code)
            self.assertEqual(calls[0]["checks"], calls[1]["checks"])
            self.assertIn("source_bound_retry_authorized", next(Path(td).rglob("coverage-audit.json")).read_text())
        self.assertEqual((chunk, candidate), before)

    def test_corrective_review_can_accept_only_valid_independently_represented_result(self):
        chunk, candidate = self.fixture()
        check = self.request(chunk, candidate)["checks"][0]
        corrected = self.uncertain(check)
        corrected["results"][0]["verdict"] = "consistent"
        obligation = corrected["results"][0]["identified_obligations"][0]
        obligation["disposition"] = "represented"
        obligation["requirement_refs"] = [check["review_context"]["linked_requirements"][0]["requirement_ref"]]
        calls = []
        before = copy.deepcopy(candidate)
        with tempfile.TemporaryDirectory() as td:
            pointer = self.run_receipt_review(
                candidate, chunk, td, [self.uncertain(check), corrected], calls)
            self.assertEqual(len(calls), 2)
            self.assertEqual(pointer["status"], "completed")
            ledger = json.loads(next(Path(td).rglob("obligation-analysis-ledger.json")).read_text())
            self.assertFalse(ledger["submission_ready"])
            self.assertEqual(ledger["candidate_response_sha256"], bridge._response_sha256(candidate))
            envelope = json.loads((Path(td) / pointer["audit_path"]).read_text())
            bridge._validate_completed_obligation_ledger_chain(
                Path(td), envelope, pointer, candidate, chunk, chunk_index=1, attempt=1,
            )
        self.assertEqual(candidate, before)

    def test_retry_feedback_receipt_requires_current_target_and_exact_second_attempt(self):
        chunk, candidate = self.fixture()
        request = self.request(chunk, candidate)
        request["provider_attempt"] = 2
        request["retry_feedback"] = {"code": native.TableContextUncertaintyError.code,
                                     "clause_ids": ["arbitrary-clause"]}
        self.assertTrue(table_retry_feedback_is_source_bound(request))
        for attempt in (1, 3, True, "2"):
            changed = {**request, "provider_attempt": attempt}
            self.assertFalse(table_retry_feedback_is_source_bound(changed))
        for ids in ([], ["unknown"], ["arbitrary-clause", "arbitrary-clause"], [None]):
            changed = {**request, "retry_feedback": {**request["retry_feedback"], "clause_ids": ids}}
            self.assertFalse(table_retry_feedback_is_source_bound(changed))
        changed = copy.deepcopy(request)
        changed["checks"][0]["review_context"].pop("table_structure_context")
        self.assertFalse(table_retry_feedback_is_source_bound(changed))

    def test_table_corrective_receipt_reconstructs_and_resealed_neighbor_tampering_is_rejected(self):
        chunk, candidate = self.fixture()
        evidence_doc = {"evidence": list(chunk["evidence_context"].values())}
        full = engine.build_llm_request([], chunk["clauses"], evidence_doc, {}, "full", contract_version="3.0")
        full["case_id"] = chunk["case_id"]
        full = attach_request_provenance(full, source_sha256="a" * 64,
            evidence_doc=evidence_doc, clauses=chunk["clauses"], run_id="new-run")
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            engine.prepare_host_agent_review_packets(full, chunk["clauses"], evidence_doc, "a" * 64,
                                                     root, chunk_size=1)
            chunk = json.loads((root / "llm-request-chunks.json").read_text())[0]
            candidate["provenance"] = copy.deepcopy(chunk["provenance"])
            check = self.request(chunk, candidate)["checks"][0]
            corrected = self.uncertain(check)
            corrected["results"][0]["verdict"] = "consistent"
            item = corrected["results"][0]["identified_obligations"][0]
            item["disposition"] = "represented"
            item["requirement_refs"] = [check["review_context"]["linked_requirements"][0]["requirement_ref"]]
            pointer = self.run_receipt_review(candidate, chunk, td, [self.uncertain(check), corrected], [])
            candidate_path = root / "accepted-candidate.json"
            bridge._write_json(candidate_path, candidate)
            audit = {"chunk_count": 1, "adapter_id": "codex", "host_runtime": "codex",
                     "chunk_lifecycle": [{"chunk_index": 1, "status": "completed", "remote_operation_state": "completed"}],
                     "chunk_runs": [{"chunk_index": 1, "response_path": str(candidate_path),
                                     "accepted_response_sha256": sha256_json(candidate),
                                     "independent_obligation_review": pointer}]}
            validate = lambda: pipeline._validate_independent_obligation_receipts(
                audit=audit, review_root=root, expected_run_id="new-run",
                expected_request_body_sha=request_body_sha256(full),
                expected_request_envelope_sha=None, expected_request_file_sha=None)
            self.assertEqual(len(validate()), 1)
            envelope_path = root / pointer["audit_path"]
            envelope = json.loads(envelope_path.read_text())
            request_path = Path(envelope["review_audit"]["request_path"])
            request = json.loads(request_path.read_text())
            context = request["checks"][0]["review_context"]
            table = context["table_structure_context"]
            row = table["source_row"]
            row["cells"][0]["paragraphs"][0]["text"] = "伪造字段"
            row["cells"][0]["paragraphs"][0]["text_sha256"] = text_sha256("伪造字段")
            row["row_sha256"] = sha256_json({k: v for k, v in row.items() if k != "row_sha256"})
            context["cited_evidence"][table["target_evidence_id"]]["table_row_context"] = copy.deepcopy(row)
            self.assertTrue(table_retry_feedback_is_source_bound(request))
            bridge._write_json(request_path, request)
            request_sha = sha256_json(request)
            envelope["review_audit"]["request_sha256"] = request_sha
            envelope["review_request_sha256"] = pointer["review_request_sha256"] = request_sha
            bridge._write_json(envelope_path, envelope)
            pointer["audit_sha256"] = sha256_file(envelope_path)
            with self.assertRaisesRegex(ValueError, "not reconstructed from the canonical source packet"):
                validate()

    def test_forged_typed_error_cannot_authorize_a_retry(self):
        chunk, candidate = self.fixture(mode="blank")
        with tempfile.TemporaryDirectory() as td, patch.object(bridge, "run_native_semantic_review",
            side_effect=native.TableContextUncertaintyError(["arbitrary-clause"])) as reviewer:
            with self.assertRaises(bridge.IndependentObligationReviewError) as caught:
                bridge._run_independent_obligation_coverage_review(
                    candidate, chunk, review_dir=Path(td), run_id="new-run", chunk_index=1,
                    attempt=1, host_runtime="codex", model="gpt-5.6-luna", timeout=5,
                    agent_id="main", runner="exec", binary="codex", config_path=None,
                    controller=bridge.RunController())
            self.assertEqual(reviewer.call_count, 1)
            self.assertFalse(caught.exception.retryable)

    def test_prompt_explains_source_geometry_without_ordering_a_pass(self):
        chunk, candidate = self.fixture()
        request = self.request(chunk, candidate)
        request["retry_feedback"] = {"code": native.TableContextUncertaintyError.code,
                                     "clause_ids": ["arbitrary-clause"]}
        prompt = native._prompt(request)
        self.assertIn("not a guessed field meaning", prompt)
        self.assertIn("If ambiguity or misrepresentation remains, keep it", prompt)
        self.assertIn("source_refs still select only", prompt)
