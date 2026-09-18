from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from template_builder import build_template  # noqa: E402
from template_audit import _page_role_map, _section_numbering, audit_template  # noqa: E402
from template_profile import NS, compile_profile, load_profile, sha256  # noqa: E402

PROFILE = ROOT / "template_profiles" / "cau-graduate-thesis-2025-winter" / "profile.json"
BNU_PROFILE = ROOT / "template_profiles" / "bnu-degree-thesis-2026" / "profile.json"


class TemplateProfileBuilderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.compiled = compile_profile(PROFILE)

    def test_compile_emits_stable_ooxml_locator_contract(self):
        locator = self.compiled["build"]["locators"]["originality_fixed"]["locator"]
        self.assertEqual(locator["part"], "word/document.xml")
        self.assertEqual(locator["kind"], "paragraph")
        self.assertTrue(locator["xml_path"].startswith("/w:document/w:body/"))
        self.assertIsInstance(locator["body_child_index"], int)
        self.assertRegex(locator["normalized_text_sha256"], r"^[0-9a-f]{64}$")
        self.assertIn("style_id", locator)
        self.assertIn("table_cell_path", locator)
        self.assertIn("bookmarks", locator)
        self.assertIn("content_controls", locator)
        self.assertEqual(locator["section"]["index"], 1)
        table = self.compiled["build"]["locators"]["spine_example_table"]["locator"]
        self.assertEqual(table["kind"], "table")
        self.assertEqual(table["table_cell_path"]["table_index"], 0)
        body = self.compiled["role_evidence"]["body"]
        self.assertEqual(body["locator"]["style_id"], "2")
        self.assertIn("MACROBUTTON", body["selector"]["text"])
        self.assertEqual(len(self.compiled["role_evidence"]["cover"]["selector_variants"]), 2)

    def test_section_numbering_treats_omitted_word_format_as_inherited_or_decimal(self):
        root = etree.fromstring(f'''<w:document xmlns:w="{NS['w']}"><w:body>
          <w:p><w:pPr><w:sectPr><w:pgNumType w:fmt="upperRoman" w:start="1"/></w:sectPr></w:pPr></w:p>
          <w:p><w:pPr><w:sectPr><w:pgNumType/></w:sectPr></w:pPr></w:p>
          <w:sectPr><w:pgNumType w:fmt="decimal" w:start="1"/></w:sectPr>
        </w:body></w:document>'''.encode())
        numbering = _section_numbering(root)
        self.assertEqual(numbering[0]["format"], "upperRoman")
        self.assertEqual(numbering[1]["format"], "decimal")
        self.assertEqual(numbering[2], {"section_index": 2, "format": "decimal", "start": 1,
                                        "explicit_format": "decimal",
                                        "source": "word/document.xml:w:sectPr/w:pgNumType"})

    def test_rendered_role_map_does_not_treat_repeated_header_as_cover(self):
        profile = {
            "structure": {"ordered_roles": [
                {"role": "cover", "anchors": ["某大学硕士学位论文"]},
                {"role": "body", "anchors": ["第一章"]},
            ]},
            "render_rules": {"header_markers": ["某大学硕士学位论文"]},
        }
        pages = [
            {"page": 1, "full_text": "某大学硕士学位论文\n论文题目"},
            {"page": 2, "full_text": "某大学硕士学位论文\n第一章 引言\n正文"},
        ]
        role_pages, _ = _page_role_map(pages, profile)
        self.assertEqual(role_pages["cover"], [])
        self.assertEqual(role_pages["body"], [2])

    def test_rendered_role_map_uses_heading_lines_not_toc_or_prose_mentions(self):
        profile = {
            "structure": {"ordered_roles": [
                {"role": "acknowledgments", "anchors": ["致 谢"],
                 "selector": {"text": "致 谢", "match": "exact"}},
                {"role": "references", "anchors": ["参考文献"],
                 "selector": {"text": "参考文献", "match": "exact"}},
            ]},
            "render_rules": {},
        }
        pages = [
            {"page": 1, "full_text": "目录\n致 谢................7\n参考文献............8"},
            {"page": 2, "full_text": "本文对参考文献和致谢安排作出说明。"},
            {"page": 3, "full_text": "致 谢\n感谢导师。"},
            {"page": 4, "full_text": "参考文献\nAuthor. Title."},
        ]
        role_pages, _ = _page_role_map(pages, profile)
        self.assertEqual(role_pages["acknowledgments"], [3])
        self.assertEqual(role_pages["references"], [4])

    def test_ambiguous_selector_fails_closed(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            profile = json.loads(PROFILE.read_text(encoding="utf-8"))
            profile["resources"] = [{**item, "path": str((PROFILE.parent / item["path"]).resolve())}
                                    for item in profile["resources"]]
            profile["build"]["locators"]["ambiguous"] = {
                "kind": "paragraph", "text": "硕士/博士学位论文", "match": "exact"
            }
            path = temp / "profile.json"
            path.write_text(json.dumps(profile, ensure_ascii=False), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "ambiguous"):
                compile_profile(path)

    def test_builder_crops_declared_range_and_preserves_fixed_ooxml(self):
        with tempfile.TemporaryDirectory() as raw:
            output = Path(raw) / "built.docx"
            result = build_template(PROFILE, output)
            self.assertEqual(result["operations_applied"], 1)
            document = Document(output)
            self.assertFalse(any("书 脊 的 书 写" in p.text for p in document.paragraphs))
            self.assertTrue(any("独 创 性 声 明" in p.text for p in document.paragraphs))
            self.assertEqual(len(document.tables), 1)
            with zipfile.ZipFile(output) as archive:
                root = etree.fromstring(archive.read("word/document.xml"))
            for fixed in self.compiled["fixed_text_evidence"].values():
                expected = fixed["locators"][0]["ooxml_sha256"]
                hashes = [__import__("hashlib").sha256(etree.tostring(n, method="c14n")).hexdigest()
                          for n in root.xpath("//w:p | //w:tbl", namespaces=NS)]
                self.assertEqual(hashes.count(expected), 1)

    def test_replace_range_from_docx_closes_the_template_content_loop(self):
        compiled = copy.deepcopy(self.compiled)
        compiled["build"]["operations"] = [{
            "op": "replace_range_from_docx", "start": "spine_start",
            "end": "spine_example_table", "input": "body"
        }]
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            replacement = temp / "replacement.docx"
            incoming = Document()
            incoming.add_heading("第1章 可验证替换内容", level=1)
            incoming.add_paragraph("这是由外部 DOCX 提供的正文范围。")
            incoming.save(replacement)
            compiled_path = temp / "compiled.json"
            compiled_path.write_text(json.dumps(compiled), encoding="utf-8")
            output = temp / "out.docx"
            build_template(PROFILE, output, inputs={"body": replacement}, compiled_path=compiled_path)
            text = "\n".join(p.text for p in Document(output).paragraphs)
            self.assertIn("第1章 可验证替换内容", text)
            self.assertNotIn("书 脊 的 书 写", text)
            self.assertIn("独 创 性 声 明", text)

    def test_unknown_operator_fails_closed(self):
        compiled = copy.deepcopy(self.compiled)
        compiled["build"]["operations"][0]["op"] = "rewrite_everything"
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            compiled_path = temp / "compiled.json"
            output = temp / "out.docx"
            compiled_path.write_text(json.dumps(compiled), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "unknown template build operator"):
                build_template(PROFILE, output, compiled_path=compiled_path)

    def test_compiled_hash_binding_fails_closed(self):
        compiled = copy.deepcopy(self.compiled)
        compiled["official_docx_sha256"] = "0" * 64
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            compiled_path = temp / "compiled.json"
            compiled_path.write_text(json.dumps(compiled), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "not bound"):
                build_template(PROFILE, temp / "out.docx", compiled_path=compiled_path)

    @staticmethod
    def _complete_thesis_profile(degree_category="academic"):
        return {
            "schema_version": "1.0", "degree_level": "doctor", "degree_category": degree_category,
            "writing_language": "zh", "has_appendices": True,
            "cover_metadata": {
                "trust": {"source": "user_confirmed", "confirmed": True},
                "classification_number": "S123", "title_zh": "题目", "title_en": "Title",
                "author_name": "作者", "student_id": "20260001", "supervisor_name": "导师",
                "degree_discipline": "农学博士", "program_name": "作物学",
                "research_direction": "方向", "college_name": "学院", "completion_date": "2026-07",
            },
        }

    def test_trimmed_official_master_passes_structural_audit_with_complete_profile(self):
        with tempfile.TemporaryDirectory() as raw:
            output = Path(raw) / "built.docx"
            build_template(PROFILE, output)
            result = audit_template(output, PROFILE, thesis_profile=self._complete_thesis_profile())
            self.assertTrue(result["official_template_structure_valid"], result["failures"])
            self.assertEqual(result["evidence"]["role_hits"]["references"][0]["style_id"], "2")
            self.assertIn("MACROBUTTON Title1", result["evidence"]["role_hits"]["body"][0]["text"])

    @unittest.skipUnless(BNU_PROFILE.is_file(), "BNU 2026 profile is not present")
    def test_bnu_2026_profile_compiles_builds_and_preserves_official_structure(self):
        compiled = compile_profile(BNU_PROFILE)
        self.assertEqual(compiled["profile_id"], "bnu-degree-thesis-2026")
        self.assertEqual(compiled["build"]["operations"], [])
        self.assertIn("originality_declaration_2026", compiled["fixed_text_evidence"])
        self.assertIn("authorization_2026", compiled["fixed_text_evidence"])
        with tempfile.TemporaryDirectory() as raw:
            output = Path(raw) / "bnu-master.docx"
            built = build_template(BNU_PROFILE, output)
            self.assertEqual(built["operations_applied"], 0)
            result = audit_template(output, BNU_PROFILE, thesis_profile={
                "degree_level": "master", "degree_category": "academic",
                "cover_metadata": {
                    "title_zh": "题目", "title_en": "Title", "author_name": "作者",
                    "student_id": "20260001", "supervisor_name": "导师",
                    "degree_discipline": "教育学", "college_name": "学院",
                    "completion_date": "2026年6月",
                },
            })
            self.assertTrue(result["official_template_structure_valid"], result["failures"])
            self.assertEqual(result["evidence"]["role_hits"]["body"][0]["style_id"], "1")
            self.assertEqual(result["evidence"]["role_hits"]["defense_committee"][0]["section_index"], 17)

    def test_audit_reads_instrtext_and_required_selector_fails_closed(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            docx = temp / "field.docx"
            document = Document(); document.add_paragraph("visible")
            document.save(docx)
            with zipfile.ZipFile(docx, "r") as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            root = etree.fromstring(members["word/document.xml"])
            paragraph = root.xpath("/w:document/w:body/w:p[1]", namespaces=NS)[0]
            run = paragraph.find("w:r", NS)
            text = run.find("w:t", NS)
            text.getparent().remove(text)
            instr = etree.SubElement(run, f"{{{NS['w']}}}instrText")
            instr.text = " MACROBUTTON Demo [Field target]"
            members["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            with zipfile.ZipFile(docx, "w", zipfile.ZIP_DEFLATED) as archive:
                for name, payload in members.items():
                    archive.writestr(name, payload)

            profile = {
                "schema_version": "1.0", "profile_id": "test-field-selector",
                "authority": {"organization": "test", "source_url": "https://example.test",
                              "effective_version": "1"},
                "resources": [{"id": "doc", "path": str(docx), "sha256": sha256(docx),
                               "kind": "official_docx"}],
                "structure": {"ordered_roles": [{"role": "body", "required": True,
                    "selector": {"kind": "paragraph", "text": "MACROBUTTON Demo [Field target]",
                                 "match": "exact"}}]},
                "render_rules": {},
            }
            profile_path = temp / "profile.json"
            profile_path.write_text(json.dumps(profile), encoding="utf-8")
            passed = audit_template(docx, profile_path)
            self.assertTrue(passed["official_template_structure_valid"], passed["failures"])

            profile["structure"]["ordered_roles"][0]["selector"]["text"] = "missing"
            profile_path.write_text(json.dumps(profile), encoding="utf-8")
            failed = audit_template(docx, profile_path)
            self.assertFalse(failed["official_template_structure_valid"])
            self.assertTrue(any(item["code"] == "template_required_role_selector_not_unique"
                                for item in failed["failures"]))

            profile["structure"]["ordered_roles"][0]["selector"]["text"] = "MACROBUTTON Demo [Field target]"
            with zipfile.ZipFile(docx, "r") as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            root = etree.fromstring(members["word/document.xml"])
            body = root.find("w:body", NS)
            body.insert(1, copy.deepcopy(body[0]))
            members["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            with zipfile.ZipFile(docx, "w", zipfile.ZIP_DEFLATED) as archive:
                for name, payload in members.items():
                    archive.writestr(name, payload)
            profile["resources"][0]["sha256"] = sha256(docx)
            profile_path.write_text(json.dumps(profile), encoding="utf-8")
            ambiguous = audit_template(docx, profile_path)
            self.assertFalse(ambiguous["official_template_structure_valid"])
            finding = next(item for item in ambiguous["failures"]
                           if item["code"] == "template_required_role_selector_not_unique")
            self.assertEqual(finding["actual"], 2)
            diagnostics = finding["selector_diagnostics"]
            self.assertEqual(diagnostics["contract"], "selector-repair-evidence-v1")
            self.assertTrue(diagnostics["policy"]["llm_must_choose_listed_candidate"])
            self.assertTrue(diagnostics["candidates"])
            self.assertIn("candidate_selector", diagnostics["candidates"][0])
            self.assertIn("context", diagnostics["candidates"][0])

    def test_role_audit_ignores_word_toc_cache_entries(self):
        with tempfile.TemporaryDirectory() as raw:
            temp = Path(raw)
            docx = temp / "toc-cache.docx"
            document = Document()
            toc_style = document.styles.add_style("toc 1", 1)
            toc_style.style_id = "TOC1"
            cached = document.add_paragraph(style=toc_style)
            cached.add_run("附录A  标题 ")
            field = etree.Element(f"{{{NS['w']}}}fldSimple")
            field.set(f"{{{NS['w']}}}instr", " PAGEREF _Toc123 \\h ")
            run = etree.SubElement(field, f"{{{NS['w']}}}r")
            text = etree.SubElement(run, f"{{{NS['w']}}}t")
            text.text = "17"
            cached._p.append(field)
            document.add_paragraph("参考文献")
            document.add_paragraph("附录A 标题")
            document.add_paragraph("攻读硕士学位期间的学术成果")
            document.save(docx)

            profile = {
                "schema_version": "1.0", "profile_id": "test-toc-cache",
                "authority": {"organization": "test", "source_url": "https://example.test",
                              "effective_version": "1"},
                "resources": [{"id": "doc", "path": str(docx), "sha256": sha256(docx),
                               "kind": "official_docx"}],
                "structure": {"ordered_roles": [
                    {"role": "references", "required": True, "selector": {
                        "kind": "paragraph", "text": "参考文献", "match": "exact"}},
                    {"role": "appendices", "required": False, "min_occurrences": 0,
                     "max_occurrences": 1, "selector": {"kind": "paragraph",
                        "text": "附录A 标题", "match": "starts_with"}},
                    {"role": "academic_outputs", "required": True, "selector": {
                        "kind": "paragraph", "text": "攻读硕士学位期间的学术成果", "match": "exact"}},
                ]},
                "render_rules": {},
            }
            profile_path = temp / "profile.json"
            profile_path.write_text(json.dumps(profile), encoding="utf-8")
            result = audit_template(docx, profile_path)
            self.assertTrue(result["official_template_structure_valid"], result["failures"])
            self.assertEqual(len(result["evidence"]["role_hits"]["appendices"]), 1)
            self.assertEqual(result["evidence"]["role_hits"]["appendices"][0]["text"], "附录A 标题")

    @unittest.skipUnless(
        (ROOT / "build" / "bnu-template-engine" / "bnu-editing-compatibility.docx").is_file(),
        "historical BNU compatibility artifact is not part of the fresh-run inputs",
    )
    def test_selector_accepts_declared_staticized_field_text_without_relaxing_locator(self):
        profile = json.loads(BNU_PROFILE.read_text(encoding="utf-8"))
        body = next(item for item in profile["structure"]["ordered_roles"] if item["role"] == "body")
        self.assertEqual(body["selector"]["accepted_texts"], ["请输入章节名称"])
        result = audit_template(
            ROOT / "build" / "bnu-template-engine" / "bnu-editing-compatibility.docx",
            BNU_PROFILE,
            thesis_profile={
                "degree_level": "master",
                "cover_metadata": {
                    "title_zh": "测试题目", "title_en": "Test Title", "author_name": "测试作者",
                    "student_id": "20260000", "supervisor_name": "测试导师",
                    "degree_discipline": "教育学", "college_name": "测试学院",
                    "completion_date": "2026年6月",
                },
            },
        )
        self.assertTrue(result["official_template_structure_valid"], result["failures"])

    @unittest.skipUnless(
        (ROOT / "build" / "cau-template-engine" / "cau-wps-word-compatible-editing.docx").is_file(),
        "historical CAU compatibility artifact is not part of the fresh-run inputs",
    )
    def test_cau_selector_accepts_declared_staticized_field_text_without_relaxing_locator(self):
        profile = json.loads(PROFILE.read_text(encoding="utf-8"))
        body = next(item for item in profile["structure"]["ordered_roles"] if item["role"] == "body")
        self.assertEqual(
            body["selector"]["accepted_texts"],
            ["单击并输入第一级标题，对应样式：标题1"],
        )
        self.assertEqual(body["selector"]["section_index"], 7)
        self.assertEqual(body["selector"]["style_id"], "2")
        self.assertEqual(body["selector"]["accepted_style_ids"], ["1"])
        result = audit_template(
            ROOT / "build" / "cau-template-engine" / "cau-wps-word-compatible-editing.docx",
            PROFILE,
            thesis_profile=self._complete_thesis_profile(),
        )
        self.assertTrue(result["official_template_structure_valid"], result["failures"])

    @unittest.skipUnless(
        (ROOT / "build" / "cau-template-engine" / "final-audit" / "cau-word-final.docx").is_file(),
        "historical CAU final artifact is not part of the fresh-run inputs",
    )
    def test_cau_selector_accepts_word_normalized_heading_style_ids(self):
        result = audit_template(
            ROOT / "build" / "cau-template-engine" / "final-audit" / "cau-word-final.docx",
            PROFILE,
            thesis_profile=self._complete_thesis_profile(),
        )
        self.assertTrue(result["official_template_structure_valid"], result["failures"])


if __name__ == "__main__":
    unittest.main()
