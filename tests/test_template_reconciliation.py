from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from requirements_engine import merge_llm_primary
from format_spec_validation import load_and_validate
from resource_registry import materialize_declaration_resources
from template_reconciliation import extract_template_evidence, reconcile_template


def _source() -> dict:
    return {"kind": "official_template_docx", "path": "/official.docx",
            "bytes": 1, "sha256": "a" * 64}


def _item(item_id: str, role: str, properties: dict) -> dict:
    return {
        "id": item_id, "category": "semantic_style", "semantic_role": role,
        "property": "style_properties", "value": properties, "direct_semantic": True,
        "source_location": {"part": "word/styles.xml", "style_name": role},
        "provenance": {"source_kind": "official_template_docx",
                       "source_path": "/official.docx", "source_sha256": "a" * 64,
                       "extractor": "test", "extractor_version": "1.0"},
    }


def _evidence(*items: dict) -> dict:
    return {"schema_version": "1.0", "status": "extracted", "source": _source(),
            "items": list(items), "summary": {"item_count": len(items)}}


def _sources() -> dict:
    return {
        "requirements_source": {"kind": "docx", "path": "/requirements.docx",
                                "bytes": 1, "sha256": "b" * 64},
        "official_template_source": _source(),
        "structure_source": {"kind": "target_thesis_structure_docx",
                             "path": "/target.docx", "bytes": 1, "sha256": "c" * 64},
    }


class TemplateEvidenceTest(unittest.TestCase):
    def test_artifact_covers_page_field_structure_style_and_fixed_body_with_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "official.docx"
            doc = Document()
            style = doc.styles.add_style("Official Heading", WD_STYLE_TYPE.PARAGRAPH)
            style.font.name = "Times New Roman"
            style.font.size = Pt(16)
            style.font.bold = True
            doc.add_paragraph("第1章 绪论", style)
            doc.add_paragraph("学位论文原创性声明")
            doc.add_paragraph("本人郑重声明：本论文由本人独立完成。")
            doc.add_paragraph("作者签名：        日期：    年  月  日")
            doc.add_paragraph("摘 要")
            footer = doc.sections[0].footer.paragraphs[0]
            begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
            instruction = OxmlElement("w:instrText"); instruction.text = " PAGE "
            end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
            footer.add_run()._r.append(begin)
            footer.add_run()._r.append(instruction)
            footer.add_run()._r.append(end)
            doc.save(path)

            artifact = extract_template_evidence(path)
            categories = {item["category"] for item in artifact["items"]}
            self.assertTrue({"page_setting", "page_field", "semantic_style", "structure",
                             "fixed_text", "fixed_body"} <= categories)
            self.assertTrue(all(item.get("source_location") and item.get("provenance")
                                for item in artifact["items"]))
            declaration = next(item for item in artifact["items"] if item["category"] == "fixed_body")
            self.assertEqual(declaration["value"]["heading"], "学位论文原创性声明")
            self.assertEqual(declaration["value"]["body_parts"], ["本人郑重声明：本论文由本人独立完成。"])
            self.assertEqual(declaration["value"]["before_role"], "abstract_title_zh")
            self.assertRegex(artifact["source"]["sha256"], r"^[0-9a-f]{64}$")


class TemplateReconciliationTest(unittest.TestCase):
    def test_agreement_uniquely_resolves_question(self) -> None:
        clause = {"id": "C1", "text": "一级标题使用黑体三号", "evidence_ids": ["E1"]}
        spec = {
            "schema_version": "1.0", "source_document": "requirements.docx",
            "analysis_mode": "llm_primary", "status": "needs_clarification", "roles": {},
            "requirements": [],
            "completeness": {"reviewed_by": "llm", "covered_clause_ids": [],
                             "ignored_clause_ids": [], "unresolved_clause_ids": ["C1"],
                             "missing_clause_ids": [], "unsupported_items": []},
            "clause_compliance": [{"clause_id": "C1", "evidence_ids": ["E1"],
                                   "scope": "docx", "status": "unresolved",
                                   "requirement_ids": [], "reason": "ambiguous role"}],
        }
        question = {"id": "Q1", "clause_id": "C1", "source_text": clause["text"],
                    "candidate_roles": ["heading_1"], "evidence_ids": ["E1"]}
        evidence = _evidence(_item("TE1", "heading_1", {
            "font": {"cjk": "SimHei", "size_pt": 16},
        }))
        reconciled, questions, conflicts, report = reconcile_template(
            spec, [clause], [question], evidence, sources=_sources(),
            property_parser=lambda text: {"font": {"cjk": "SimHei", "size_pt": 16}},
            page_property_parser=lambda text: {}, fill_silent_values=True,
        )
        self.assertEqual(questions, [])
        self.assertEqual(conflicts, [])
        requirement = reconciled["requirements"][0]
        self.assertEqual(requirement["resolved_by"], "template")
        self.assertTrue(requirement["properties"])
        self.assertEqual(reconciled["clause_compliance"][0]["status"], "pending_execution")
        self.assertTrue(any(item["action"] == "resolved_by_template" for item in report["results"]))

    def test_direct_conflict_requires_clarification(self) -> None:
        spec = {
            "schema_version": "1.0", "source_document": "requirements.docx",
            "status": "semantic_resolved", "roles": {"body_text": {"font": {"size_pt": 12}}},
            "requirements": [{"id": "R1", "role": "body_text",
                              "properties": {"font": {"size_pt": 12}},
                              "evidence_ids": ["E1"], "clause_ids": ["C1"],
                              "resolved_by": "llm", "confidence": 1.0}],
        }
        evidence = _evidence(_item("TE1", "body_text", {"font": {"size_pt": 10.5}}))
        reconciled, questions, conflicts, report = reconcile_template(
            spec, [{"id": "C1", "text": "正文小四号", "evidence_ids": ["E1"]}], [],
            evidence, sources=_sources(), fill_silent_values=True,
        )
        self.assertEqual(reconciled["status"], "needs_clarification")
        self.assertEqual(len(conflicts), 1)
        self.assertEqual(conflicts[0]["type"], "official_template_conflict")
        self.assertEqual(len(questions), 1)
        self.assertEqual(report["summary"]["genuine_conflicts"], 1)

    def test_absence_is_insufficient_not_a_prohibition_or_conflict(self) -> None:
        spec = {
            "schema_version": "1.0", "source_document": "requirements.docx",
            "status": "semantic_resolved", "roles": {"header": {"font": {"size_pt": 10.5}}},
            "requirements": [{"id": "R1", "role": "header",
                              "properties": {"font": {"size_pt": 10.5}},
                              "evidence_ids": ["E1"], "clause_ids": ["C1"],
                              "resolved_by": "llm", "confidence": 1.0}],
        }
        reconciled, questions, conflicts, report = reconcile_template(
            spec, [{"id": "C1", "text": "页眉五号", "evidence_ids": ["E1"]}], [],
            _evidence(), sources=_sources(), fill_silent_values=True,
        )
        self.assertEqual(reconciled["status"], "semantic_resolved")
        self.assertEqual(questions, [])
        self.assertEqual(conflicts, [])
        self.assertTrue(any(item["result"] == "insufficient" for item in report["results"]))
        self.assertFalse(report["policy"]["sample_absence_is_prohibition"])

    def test_not_applicable_clause_is_reported_without_template_inference(self) -> None:
        spec = {
            "schema_version": "1.0", "source_document": "requirements.docx",
            "status": "semantic_resolved", "roles": {}, "requirements": [],
            "clause_compliance": [{"clause_id": "C-na", "evidence_ids": ["E-na"],
                                   "scope": "docx", "status": "not_applicable",
                                   "requirement_ids": [], "reason": "Doctoral-only clause; current thesis is master."}],
        }
        _, questions, conflicts, report = reconcile_template(
            spec, [{"id": "C-na", "text": "仅博士论文适用", "evidence_ids": ["E-na"]}], [],
            _evidence(), sources=_sources(),
        )
        self.assertEqual(questions, [])
        self.assertEqual(conflicts, [])
        item = next(item for item in report["results"] if item["clause_id"] == "C-na")
        self.assertEqual(item["result"], "not_applicable")

    def test_structure_order_uses_observed_relative_order_not_sample_completeness(self) -> None:
        requirement = {
            "id": "R-structure", "role": "document_structure",
            "properties": {"ordered_roles": ["abstract_title_zh", "heading_1", "heading_references"]},
            "evidence_ids": ["E1"], "clause_ids": ["C1"],
            "resolved_by": "llm", "confidence": 1.0,
        }
        spec = {"schema_version": "1.0", "source_document": "requirements.docx",
                "status": "semantic_resolved", "roles": {}, "requirements": [requirement]}
        structure = {
            "id": "TE-order", "category": "structure", "semantic_role": "document_structure",
            "property": "ordered_roles",
            "value": ["abstract_title_zh", "toc", "heading_1", "heading_references"],
            "source_location": {"part": "word/document.xml", "observations": []},
            "provenance": {"source_kind": "official_template_docx", "source_path": "/official.docx",
                           "source_sha256": "a" * 64, "extractor": "test", "extractor_version": "1.0"},
        }
        _, _, conflicts, report = reconcile_template(
            spec, [{"id": "C1", "text": "摘要、正文、参考文献依次排列", "evidence_ids": ["E1"]}], [],
            _evidence(structure), sources=_sources(),
        )
        self.assertEqual(conflicts, [])
        result = next(item for item in report["results"] if item.get("property") == "ordered_roles")
        self.assertEqual(result["result"], "agree")

    def test_declaration_body_is_exactly_materialized_from_official_template(self) -> None:
        clause = {"id": "C-decl", "text": "原创性声明见模板例文", "evidence_ids": ["E-decl"]}
        spec = {
            "schema_version": "1.0", "source_document": "requirements.docx",
            "analysis_mode": "llm_primary", "status": "needs_clarification", "roles": {},
            "requirements": [],
            "completeness": {"reviewed_by": "llm", "covered_clause_ids": [],
                             "ignored_clause_ids": [], "unresolved_clause_ids": ["C-decl"],
                             "missing_clause_ids": [], "unsupported_items": []},
            "clause_compliance": [{"clause_id": "C-decl", "evidence_ids": ["E-decl"],
                                   "scope": "docx", "status": "requires_source_content",
                                   "requirement_ids": [], "reason": "example body missing"}],
        }
        question = {"id": "Q1", "clause_id": "C-decl", "source_text": clause["text"],
                    "candidate_roles": ["unknown"], "evidence_ids": ["E-decl"]}
        with tempfile.TemporaryDirectory() as td:
            official = Path(td) / "official.docx"
            document = Document()
            document.add_paragraph("原创性声明")
            document.add_paragraph("这是官方模板的精确正文。")
            document.add_paragraph("作者签名：        日期：    年  月  日")
            document.add_paragraph("摘 要")
            document.save(official)
            evidence = extract_template_evidence(official)
            block = next(item for item in evidence["items"] if item["category"] == "fixed_body")
            reconciled, questions, conflicts, report = reconcile_template(
                spec, [clause], [question], evidence, sources=_sources(),
                property_parser=lambda text: {}, page_property_parser=lambda text: {},
            )
            self.assertEqual(questions, [])
            self.assertEqual(conflicts, [])
            materialized = materialize_declaration_resources(reconciled, "run-declaration")
            declaration = materialized["declarations"]["items"][0]
            resource = materialized["resource_registry"]["items"][declaration["resource_id"]]
            self.assertEqual(resource["heading"], "原创性声明")
            self.assertEqual(resource["body_parts"], ["这是官方模板的精确正文。"])
            self.assertEqual(resource["source_evidence_ids"], [block["id"]])
            self.assertEqual(load_and_validate(
                materialized, ROOT / "schema" / "format-spec.schema.json"
            ), [])
            self.assertTrue(any(item["action"] == "resource_from_template" for item in report["results"]))

    def test_existing_declaration_body_is_compared_exactly(self) -> None:
        properties = {"before_role": "document_start", "items": [{
            "id": "originality", "heading": "原创性声明",
            "body_parts": ["书面要求中的正文。"], "source_evidence_ids": ["E1"],
            "signature_placeholders": [],
        }]}
        spec = {
            "schema_version": "1.0", "source_document": "requirements.docx",
            "status": "semantic_resolved", "roles": {}, "requirements": [{
                "id": "R1", "role": "declarations", "properties": properties,
                "evidence_ids": ["E1"], "clause_ids": ["C1"],
                "resolved_by": "llm", "confidence": 1.0,
            }],
        }
        block = {
            "id": "TE1", "category": "fixed_body", "semantic_role": "declarations",
            "property": "items",
            "value": {"heading": "原创性声明", "body_parts": ["官方模板中的不同正文。"],
                      "signature_placeholders": [], "before_role": "document_start"},
            "source_location": {"part": "word/document.xml", "paragraphs": []},
            "provenance": {"source_kind": "official_template_docx", "source_path": "/official.docx",
                           "source_sha256": "a" * 64, "extractor": "test", "extractor_version": "1.0"},
        }
        _, questions, conflicts, report = reconcile_template(
            spec, [{"id": "C1", "text": "声明正文如下", "evidence_ids": ["E1"]}], [],
            _evidence(block), sources=_sources(),
        )
        self.assertEqual(len(conflicts), 1)
        self.assertEqual(conflicts[0]["property"], "items.originality.body_parts")
        self.assertEqual(len(questions), 1)
        self.assertTrue(any(item["result"] == "conflict" and item["semantic_role"] == "declarations"
                            for item in report["results"]))

    def test_declaration_conflict_question_targets_the_conflicting_item_clause(self) -> None:
        properties = {"items": [
            {"id": "originality", "heading": "原创性声明",
             "body_parts": ["原创性正文"], "source_evidence_ids": ["E0"],
             "signature_placeholders": []},
            {"id": "authorization", "heading": "授权书",
             "body_parts": ["书面正文"], "source_evidence_ids": ["E2"],
             "signature_placeholders": []},
        ]}
        clause_ids = ["C0", "C1", "C2", "C3"]
        spec = {
            "schema_version": "1.0", "source_document": "requirements.docx",
            "status": "semantic_resolved", "roles": {}, "requirements": [{
                "id": "R1", "role": "declarations", "properties": properties,
                "evidence_ids": ["E0", "E1", "E2", "E3"], "clause_ids": clause_ids,
                "resolved_by": "llm", "confidence": 1.0,
            }],
            "clause_compliance": [{
                "clause_id": clause_id, "evidence_ids": [evidence_id],
                "scope": "docx", "status": "pending_execution",
                "requirement_ids": ["R1"], "reason": "fixed declaration text",
            } for clause_id, evidence_id in zip(clause_ids, ["E0", "E1", "E2", "E3"])],
        }
        clauses = [
            {"id": "C0", "text": "原创性声明", "evidence_ids": ["E0"]},
            {"id": "C1", "text": "原创性正文", "evidence_ids": ["E1"]},
            {"id": "C2", "text": "授权书", "evidence_ids": ["E2"]},
            {"id": "C3", "text": "书面正文", "evidence_ids": ["E3"]},
        ]
        fixed_body = {
            "id": "TE-auth", "category": "fixed_body", "semantic_role": "declarations",
            "property": "items",
            "value": {"heading": "授权书", "body_parts": ["官方正文"],
                      "signature_placeholders": []},
            "source_location": {"part": "word/document.xml", "paragraphs": []},
            "provenance": {"source_kind": "official_template_docx", "source_path": "/official.docx",
                           "source_sha256": "a" * 64, "extractor": "test", "extractor_version": "1.0"},
        }
        reconciled, questions, conflicts, report = reconcile_template(
            spec, clauses, [], _evidence(fixed_body), sources=_sources(),
        )
        self.assertEqual(len(conflicts), 1)
        self.assertEqual(questions[0]["clause_id"], "C2")
        c0 = next(item for item in reconciled["clause_compliance"] if item["clause_id"] == "C0")
        self.assertEqual(c0["status"], "pending_execution")

    def test_c00023_manual_empty_requirement_is_not_executable(self) -> None:
        clause = {"id": "C00023", "text": "论文作者 指导教师", "evidence_ids": ["E1"]}
        response = {
            "contract_version": "2.1",
            "requirements": [{
                "role": "cover_field_label", "properties": {}, "clause_ids": ["C00023"],
                "evidence_ids": ["E1"], "confidence": 1.0, "reason": "labels require manual check",
                "verification": {"mode": "manual", "checks": ["inspect cover labels"]},
            }],
            "clause_reviews": [{"clause_id": "C00023", "classification": "executable",
                                "requirement_indexes": [0], "reason": "manual label check"}],
            "unsupported_items": [], "reported_conflicts": [],
        }
        spec, conflicts, _ = merge_llm_primary(
            Path("requirements.docx"),
            {"schema_version": "1.0", "source_document": "requirements.docx",
             "status": "rule_resolved", "roles": {}, "requirements": []},
            [clause], response, {"E1"}, require_provenance=False,
        )
        self.assertEqual(spec["requirements"], [])
        self.assertEqual(spec["clause_compliance"][0]["status"], "unverifiable")
        self.assertFalse(any(item.get("type") == "llm_contract" for item in conflicts))
        fixed = [
            {"id": "TE-author", "category": "fixed_text", "semantic_role": "fixed_text",
             "property": "text", "value": "论文作者", "normalized_text": "论文作者",
             "source_location": {"part": "word/document.xml", "paragraph_index": 1},
             "provenance": {"source_kind": "official_template_docx", "source_path": "/official.docx",
                            "source_sha256": "a" * 64, "extractor": "test", "extractor_version": "1.0"}},
            {"id": "TE-supervisor", "category": "fixed_text", "semantic_role": "fixed_text",
             "property": "text", "value": "指导教师", "normalized_text": "指导教师",
             "source_location": {"part": "word/document.xml", "paragraph_index": 2},
             "provenance": {"source_kind": "official_template_docx", "source_path": "/official.docx",
                            "source_sha256": "a" * 64, "extractor": "test", "extractor_version": "1.0"}},
        ]
        reconciled, _, genuine, report = reconcile_template(
            spec, [clause], [], _evidence(*fixed), sources=_sources()
        )
        self.assertEqual(reconciled["requirements"], [])
        self.assertEqual(genuine, [])
        self.assertTrue(any(item["action"] == "corroborated_manual" for item in report["results"]))

        invalid = {
            "schema_version": "1.0", "source_document": "requirements.docx",
            "status": "semantic_resolved", "roles": {},
            "requirements": [{"id": "R-empty", "role": "cover_field_label",
                              "properties": {}, "evidence_ids": ["E1"],
                              "resolved_by": "llm", "confidence": 1.0}],
        }
        errors = load_and_validate(invalid, ROOT / "schema" / "format-spec.schema.json")
        self.assertTrue(any("requires at least 1 properties" in error for error in errors))


class RequirementsEngineSourceSeparationTest(unittest.TestCase):
    def _run(self, *args: str) -> subprocess.CompletedProcess[str]:
        return subprocess.run([sys.executable, str(ROOT / "scripts" / "requirements_engine.py"), *args],
                              cwd=ROOT, text=True, capture_output=True)

    def test_structure_and_official_template_have_separate_identity_and_hashes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); requirements = td / "requirements.docx"
            target = td / "target.docx"; official = td / "official.docx"; out = td / "out"
            doc = Document(); doc.add_paragraph("正文使用小四号宋体。") ; doc.save(requirements)
            doc = Document(); doc.add_paragraph("目标论文第1章"); doc.save(target)
            doc = Document(); doc.add_paragraph("官方模板第1章", "Heading 1"); doc.save(official)
            result = self._run(str(requirements), "--out", str(out), "--analysis-mode", "rule_only",
                               "--structure-docx", str(target),
                               "--official-template-evidence-docx", str(official))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            document_evidence = json.loads((out / "document-evidence.json").read_text())
            template_evidence = json.loads((out / "template-evidence.json").read_text())
            manifest = json.loads((out / "extraction-manifest.json").read_text())
            self.assertEqual(Path(document_evidence["structure_source_document"]), target.resolve())
            self.assertEqual(Path(template_evidence["source"]["path"]), official.resolve())
            self.assertEqual(manifest["sources"]["structure_source"]["sha256"],
                             hashlib.sha256(target.read_bytes()).hexdigest())
            self.assertEqual(manifest["sources"]["official_template_source"]["sha256"],
                             hashlib.sha256(official.read_bytes()).hexdigest())
            self.assertNotEqual(manifest["sources"]["structure_source"]["path"],
                                manifest["sources"]["official_template_source"]["path"])

    def test_no_official_template_is_backward_compatible_and_does_not_use_structure_as_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); requirements = td / "requirements.docx"
            target = td / "target.docx"; out = td / "out"
            doc = Document(); doc.add_paragraph("正文使用小四号宋体。") ; doc.save(requirements)
            doc = Document(); doc.sections[0].orientation = 1; doc.add_paragraph("目标论文"); doc.save(target)
            result = self._run(str(requirements), "--out", str(out), "--analysis-mode", "rule_only",
                               "--structure-docx", str(target))
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            reconciliation = json.loads((out / "template-reconciliation.json").read_text())
            template_evidence = json.loads((out / "template-evidence.json").read_text())
            spec = json.loads((out / "format-spec.json").read_text())
            self.assertEqual(reconciliation["status"], "not_applicable")
            self.assertEqual(template_evidence["status"], "not_supplied")
            self.assertIsNone(reconciliation["sources"]["official_template_source"])
            self.assertNotIn("page", spec)

    def test_pipeline_routes_style_template_to_official_evidence_path(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); requirements = td / "requirements.docx"
            target = td / "target.docx"; official = td / "official.docx"
            output = td / "unused.docx"; work = td / "work"
            doc = Document(); doc.add_paragraph("正文使用小四号宋体。") ; doc.save(requirements)
            doc = Document(); doc.add_paragraph("目标论文"); doc.save(target)
            doc = Document(); doc.add_paragraph("官方模板", "Heading 1"); doc.save(official)
            result = subprocess.run([
                sys.executable, str(ROOT / "scripts" / "thesis_format_pipeline.py"),
                str(requirements), str(target), str(output), "--work-dir", str(work),
                "--style-template", str(official), "--analysis-mode", "llm_primary",
                "--compliance-mode", "full", "--prepare-host-review",
            ], cwd=ROOT, text=True, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            manifest = json.loads((work / "pipeline-manifest.json").read_text())
            command = next(item["command"] for item in manifest["steps"]
                           if item["name"] == "requirements")
            structure_index = command.index("--structure-docx")
            official_index = command.index("--official-template-evidence-docx")
            self.assertEqual(Path(command[structure_index + 1]), target.resolve())
            self.assertEqual(Path(command[official_index + 1]), official.resolve())
            self.assertEqual(Path(manifest["inputs"]["structure_source"]["path"]), target.resolve())
            self.assertEqual(Path(manifest["inputs"]["official_template_evidence"]["path"]), official.resolve())
            self.assertEqual(manifest["template_reconciliation"]["status"], "pending_semantic_review")


if __name__ == "__main__":
    unittest.main()
