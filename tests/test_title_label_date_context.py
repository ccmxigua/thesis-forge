"""Source-bound labels and ambiguous date evidence are separate responsibilities."""
from __future__ import annotations

import copy
import hashlib
import io
import json
from pathlib import Path
import sys
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import host_agent_bridge as bridge
import requirements_engine as engine
from apply_format_spec import _cover_field_text, apply_cover
from format_spec_validation import validate_instance
from semantic_contract import attach_request_provenance, sha256_json
from source_quote_reassessment import quote_context_reassessment, RULE_ID


def incident():
    data = json.loads((ROOT / "tests/fixtures/title-label-date-context-incident.json").read_text())
    source = data["source"]
    evidence = {"evidence": list(source["evidence_context"].values())}
    request = engine.build_llm_request([], source["clauses"], evidence, {}, "full", contract_version="3.0")
    request.update(source)
    request.update(case_id="offline-title-context", batch={"index": 1},
                   runtime_context={"code_fingerprint_sha256": sha256_json("offline-current-code")})
    request = attach_request_provenance(request, source_sha256=sha256_json(evidence),
        evidence_doc=evidence, clauses=source["clauses"], run_id="offline-title-context")
    return data["baseline"], request


def proposal(old, chunk):
    # Explicit offline primary proposal, never an automatic production repair
    # or a claim that the captured model selected this context.
    new = copy.deepcopy(old)
    by_id = {c["id"]: c for c in chunk["clauses"]}
    for review in new["clause_reviews"]:
        for atom in review.get("obligations", []):
            if atom.get("source_quote") == "年 月 日" and review["clause_id"] in by_id:
                atom["source_quote"] = by_id[review["clause_id"]]["source_span"]["text"]
    return new


class TitleLabelDateContextTests(unittest.TestCase):
    def test_captured_source_label_is_not_restricted_by_metadata_field_name(self):
        old, chunk = incident()
        errors = bridge.validate_host_agent_response(old, chunk)
        self.assertEqual(errors, ["$.clause_reviews[7].obligations[0].source_quote: must_equal_current_source_subspan"])
        self.assertEqual(old["requirements"][0]["properties"]["fields"][0]["label_display_policy"], "always")
        field = old["requirements"][0]["properties"]["fields"][0]
        schema = json.loads((ROOT / "schema/format-spec.schema.json").read_text())
        self.assertEqual(validate_instance(field, schema["$defs"]["coverField"], schema), [])

    def test_actual_ambiguous_quote_needs_explicit_full_context_and_review(self):
        old, chunk = incident(); frozen = copy.deepcopy(old)
        new = proposal(old, chunk)
        records = bridge.contract_error_records(bridge.validate_host_agent_response(old, chunk), response=old, chunk=chunk)
        paths = bridge._retry_change_paths(old, new)
        proofs = quote_context_reassessment(old, new, records, paths, chunk, validate=bridge.validate_host_agent_response)
        self.assertEqual(len(proofs), 1)
        self.assertEqual(proofs[0]["rule_id"], RULE_ID)
        self.assertTrue(proofs[0]["independent_review_required"])
        self.assertFalse(proofs[0]["mechanical_equivalence_claimed"])
        self.assertEqual(proofs[0]["ambiguous_match_count"], 2)
        self.assertEqual(bridge.validate_host_agent_response(new, chunk), [])
        ledger = []
        error, _ = bridge._retry_semantic_change_error(old, new, records,
            contract_version="3.0", chunk=chunk, authorization_out=ledger)
        self.assertIsNone(error)
        self.assertEqual(len(ledger), 1)
        self.assertTrue(ledger[0]["semantic_review_required"])
        self.assertEqual(old, frozen)

    def test_quote_proposal_cannot_silently_change_the_valid_label_policy(self):
        old, chunk = incident(); new = proposal(old, chunk)
        records = bridge.contract_error_records(bridge.validate_host_agent_response(old, chunk), response=old, chunk=chunk)
        new["requirements"][0]["properties"]["fields"][0]["label_display_policy"] = "with_value"
        error, changes = bridge._retry_semantic_change_error(old, new, records, contract_version="3.0", chunk=chunk)
        self.assertIsNotNone(error)
        self.assertTrue(any(path.endswith("label_display_policy") for path in changes))

    def test_always_title_label_requires_current_exact_linked_source(self):
        for change in ("invented", "missing_edge", "stale_source", "wrong_evidence", "two_occurrences"):
            old, chunk = incident(); current = proposal(old, chunk)
            requirement = current["requirements"][0]
            field = requirement["properties"]["fields"][0]
            clause = next(c for c in chunk["clauses"] if c["id"] == "C00042")
            if change == "invented": field["label"] = "虚构字段"
            elif change == "missing_edge": requirement["clause_ids"].remove(clause["id"])
            elif change == "stale_source": clause["source_span"]["source_sha256"] = "0" * 64
            elif change == "wrong_evidence": requirement["evidence_ids"].remove(clause["evidence_ids"][0])
            else:
                other = copy.deepcopy(clause); other["id"] = "another-current-title-label"
                chunk["clauses"].append(other); requirement["clause_ids"].append(other["id"])
            with self.subTest(change=change):
                errors = bridge.validate_host_agent_response(current, chunk)
                # Stale source identities fail at the earlier span gate,
                # before any field-specific check is allowed to run.
                marker = ("source_sha256" if change == "stale_source"
                    else "label_display_policy: must_be_bound_to_exact_linked_source_clause")
                self.assertTrue(any(marker in e for e in errors), errors)

    def test_rendered_title_policy_preserves_source_label_and_legacy_default(self):
        field = {"id": "title_zh", "label": "论文题目", "order": 1,
            "value_from": "thesis_profile.cover_metadata.title_zh", "display_policy": "required",
            "label_display_policy": "always"}
        self.assertEqual(_cover_field_text(field, "测试题目"), "论文题目：测试题目")
        self.assertEqual(_cover_field_text(field, ""), "论文题目：")
        for value in ("with_value", None):
            legacy = copy.deepcopy(field)
            if value is None: legacy.pop("label_display_policy")
            else: legacy["label_display_policy"] = value
            self.assertEqual(_cover_field_text(legacy, "测试题目"), "测试题目")
            self.assertIsNone(_cover_field_text(legacy, ""))
        doc = Document(); doc.add_paragraph("正文")
        cover = {"institution": "示例学校", "fields": [field]}
        counts = apply_cover(doc, cover, {})
        buffer = io.BytesIO(); doc.save(buffer); buffer.seek(0)
        self.assertIn("论文题目：——", [p.text for p in Document(buffer).paragraphs])
        self.assertEqual(counts["placeholder_fields_written"], 1)

    def test_literal_label_punctuation_and_whitespace_are_not_normalized_away(self):
        for suffix in ("：", ":", "：：", " "):
            old, chunk = incident(); current = proposal(old, chunk)
            current["requirements"][0]["properties"]["fields"][0]["label"] += suffix
            with self.subTest(suffix=suffix):
                self.assertTrue(any("label_display_policy: must_be_bound_to_exact_linked_source_clause" in e
                    for e in bridge.validate_host_agent_response(current, chunk)))
        old, chunk = incident(); current = proposal(old, chunk)
        clause = next(c for c in chunk["clauses"] if c["id"] == "C00042")
        evidence = chunk["evidence_context"][clause["evidence_ids"][0]]
        evidence["text"] += "："
        clause["source_text_full"] = evidence["text"]
        clause["source_span"]["source_sha256"] = hashlib.sha256(evidence["text"].encode()).hexdigest()
        current["requirements"][0]["properties"]["fields"][0]["label"] += "："
        # Segmentation can leave the original delimiter immediately outside
        # the span. Only the existing source-bound label compositor restores it.
        self.assertEqual(bridge.validate_host_agent_response(current, chunk), [])

    def test_unrelated_schema_constraints_still_reject_invalid_payload(self):
        old, chunk = incident(); current = proposal(old, chunk)
        current["requirements"][0]["properties"]["fields"][0]["display_policy"] = "invented"
        self.assertTrue(bridge.validate_host_agent_response(current, chunk))


if __name__ == "__main__":
    unittest.main()
