from __future__ import annotations

import sys
import tempfile
import unittest
import os
import subprocess
from zipfile import ZIP_DEFLATED, ZipFile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import word_render_export
import word_render_validate


class WordRenderExportTest(unittest.TestCase):
    def test_wait_for_document_accepts_macos_tmp_alias_and_returns_word_path(self) -> None:
        with patch.object(
            word_render_export,
            "active_document_path",
            return_value="/tmp/render/.final.docx.abc.docx",
        ):
            opened = word_render_export.wait_for_document(
                "/private/tmp/render/.final.docx.abc.docx", 1
            )
        self.assertEqual(opened, "/tmp/render/.final.docx.abc.docx")

    def test_repairs_stale_parallel_toc_targets_by_entry_order(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            docx = Path(td) / "parallel-toc.docx"
            Document().save(docx)
            with ZipFile(docx) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            w = f"{{{ns}}}"
            root = etree.fromstring(members["word/document.xml"])
            body = root.find(w + "body")
            for bookmark_id, target in enumerate(("_TocA", "_TocB"), 1):
                paragraph = etree.SubElement(body, w + "p")
                etree.SubElement(paragraph, w + "bookmarkStart", {
                    w + "id": str(bookmark_id), w + "name": target,
                })
                etree.SubElement(paragraph, w + "bookmarkEnd", {w + "id": str(bookmark_id)})
            for instruction in (
                "HYPERLINK \\l _TocA", "HYPERLINK \\l _TocB",
                "HYPERLINK \\l _StaleA", "HYPERLINK \\l _StaleB",
                "PAGEREF _TocA \\h", "PAGEREF _TocB \\h",
                "PAGEREF _StaleA \\h", "PAGEREF _StaleB \\h",
            ):
                paragraph = etree.SubElement(body, w + "p")
                run = etree.SubElement(paragraph, w + "r")
                etree.SubElement(run, w + "instrText").text = instruction
            members["word/document.xml"] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            with ZipFile(docx, "w", ZIP_DEFLATED) as archive:
                for name, payload in members.items():
                    archive.writestr(name, payload)

            self.assertEqual(word_render_export.repair_parallel_toc_targets(docx), 4)
            with ZipFile(docx) as archive:
                repaired = etree.fromstring(archive.read("word/document.xml"))
            instructions = ["".join(node.itertext()) for node in repaired.xpath(
                ".//w:instrText", namespaces={"w": ns}
            )]
            self.assertNotIn("_StaleA", " ".join(instructions))
            self.assertNotIn("_StaleB", " ".join(instructions))
            self.assertEqual(sum("_TocA" in item for item in instructions), 4)
            self.assertEqual(sum("_TocB" in item for item in instructions), 4)

    def test_cli_refuses_hard_link_aliases(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; final = root / "final.docx"
            report = root / "report.json"; pdf = root / "final.pdf"
            Document().save(source); Document().save(final); os.link(final, report)
            argv = ["word_render_export.py", str(source), str(final), str(pdf), "--report", str(report)]
            with patch.object(sys, "argv", argv), patch.object(word_render_export.platform, "system", return_value="Darwin"), self.assertRaises(SystemExit) as raised:
                word_render_export.main()
            self.assertEqual(raised.exception.code, 2)
            self.assertTrue(final.read_bytes().startswith(b"PK"))

    def test_failed_open_preserves_all_existing_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; final = root / "final.docx"
            pdf = root / "final.pdf"; report = root / "report.json"
            Document().save(source); final.write_bytes(b"old-docx"); pdf.write_bytes(b"old-pdf"); report.write_bytes(b"old-report")
            failed = subprocess.CompletedProcess([], 1, "", "open failed")
            argv = ["word_render_export.py", str(source), str(final), str(pdf), "--report", str(report)]
            with patch.object(sys, "argv", argv), patch.object(word_render_export.platform, "system", return_value="Darwin"), patch.object(word_render_export, "run", return_value=failed), self.assertRaises(SystemExit):
                word_render_export.main()
            self.assertEqual(final.read_bytes(), b"old-docx")
            self.assertEqual(pdf.read_bytes(), b"old-pdf")
            self.assertEqual(report.read_bytes(), b"old-report")

    def test_cli_refuses_hard_link_aliases(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; final = root / "final.docx"
            report = root / "report.json"; pdf = root / "final.pdf"
            Document().save(source); Document().save(final); os.link(final, report)
            argv = ["word_render_export.py", str(source), str(final), str(pdf), "--report", str(report)]
            with patch.object(sys, "argv", argv), patch.object(word_render_export.platform, "system", return_value="Darwin"), self.assertRaises(SystemExit) as raised:
                word_render_export.main()
            self.assertEqual(raised.exception.code, 2)
            self.assertTrue(final.read_bytes().startswith(b"PK"))

    def test_failed_open_preserves_all_existing_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; final = root / "final.docx"
            pdf = root / "final.pdf"; report = root / "report.json"
            Document().save(source); final.write_bytes(b"old-docx"); pdf.write_bytes(b"old-pdf"); report.write_bytes(b"old-report")
            failed = subprocess.CompletedProcess([], 1, "", "open failed")
            argv = ["word_render_export.py", str(source), str(final), str(pdf), "--report", str(report)]
            with patch.object(sys, "argv", argv), patch.object(word_render_export.platform, "system", return_value="Darwin"), patch.object(word_render_export, "run", return_value=failed), self.assertRaises(SystemExit):
                word_render_export.main()
            self.assertEqual(final.read_bytes(), b"old-docx")
            self.assertEqual(pdf.read_bytes(), b"old-pdf")
            self.assertEqual(report.read_bytes(), b"old-report")

    def test_word_script_uses_absolute_path_bounded_story_ranges_and_fail_closed_counts(self) -> None:
        script = word_render_export.APPLE_SCRIPT
        self.assertIn("POSIX full name of d is not expectedPath", script)
        # Main-story TOC/PAGEREF fields are deliberately updated once through
        # Word's native TOC API; updating each of them separately repeatedly
        # repaginates a thesis and can make automation appear to hang.
        self.assertNotIn("set end of storyTypes to main text story", script)
        for story in ("text frame story", "primary header story", "primary footer story",
                      "first page header story", "first page footer story"):
            self.assertIn(story, script)
        self.assertIn("update t", script)
        self.assertIn("update page numbers t", script)
        self.assertIn("storyChainCount > 64", script)
        self.assertIn("if failedCount is not 0 then error", script)
        self.assertIn("if storyAvailable then", script)
        self.assertIn("with timeout of 600 seconds", script)
        self.assertIn("save as d file name outPdf file format format PDF", script)
        self.assertNotIn("save as d file name (POSIX file outPdf)", script)
        self.assertIn("save as d file name outPdf file format format PDF",
                      word_render_export.EXPORT_ONLY_SCRIPT)

    def test_failed_validation_preserves_diagnostics(self) -> None:
        source = Path(word_render_export.__file__).read_text(encoding="utf-8")
        self.assertIn('diagnostics / "latest-failed.docx"', source)
        self.assertIn('diagnostics / "latest-failed.pdf"', source)
        self.assertIn('diagnostics / "latest-failed-validation.json"', source)
        self.assertIn("render error hits:", source)

    def test_word_render_stages_in_word_container_to_avoid_file_access_dialog(self) -> None:
        source = Path(word_render_export.__file__).read_text(encoding="utf-8")
        self.assertIn('"Containers" / "com.microsoft.Word" / "Data" / "tmp"', source)
        self.assertIn('staging_parent.relative_to(word_container_tmp)', source)
        self.assertIn('tempfile.TemporaryDirectory(prefix="thesis-word-render-", dir=staging_parent)', source)
        self.assertEqual(source.count('["open", "-g", "-a", "Microsoft Word", str(staged_docx)]'), 2)
        self.assertNotIn('["open", "-a", "Microsoft Word", str(staged_docx)]', source)
        self.assertIn('staging_root / "final-word.docx"', source)
        self.assertNotIn('tempfile.TemporaryDirectory(prefix="thesis-word-render-", dir="/tmp")', source)
        self.assertNotIn('staged_docx = sibling_temp(final_docx)', source)

    def test_cli_refuses_staging_outside_word_container_before_word_access(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "source.docx"
            Document().save(source)
            argv = [
                "word_render_export.py",
                str(source),
                str(root / "final.docx"),
                str(root / "final.pdf"),
                "--report",
                str(root / "render.json"),
                "--staging-dir",
                str(root),
            ]
            with patch.object(sys, "argv", argv), patch.object(
                word_render_export.platform, "system", return_value="Darwin"
            ), patch.object(word_render_export, "run") as run_mock, self.assertRaises(
                SystemExit
            ) as raised:
                word_render_export.main()
            self.assertEqual(raised.exception.code, 2)
            run_mock.assert_not_called()

    def test_post_update_pageref_repair_copies_target_and_cached_page(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            docx = Path(td) / "post-update.docx"
            Document().save(docx)
            with ZipFile(docx) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            w = f"{{{ns}}}"
            root = etree.fromstring(members["word/document.xml"])
            body = root.find(w + "body")
            for bookmark_id, target in enumerate(("_NewA", "_NewB"), 1):
                paragraph = etree.SubElement(body, w + "p")
                etree.SubElement(paragraph, w + "bookmarkStart", {
                    w + "id": str(bookmark_id), w + "name": target,
                })
                etree.SubElement(paragraph, w + "bookmarkEnd", {w + "id": str(bookmark_id)})
            for target, cached in (("_NewA", "7"), ("_NewB", "9"),
                                   ("_OldA", "错误!未定义书签。"),
                                   ("_OldB", "错误!未定义书签。")):
                paragraph = etree.SubElement(body, w + "p")
                field_parent = etree.SubElement(paragraph, w + "hyperlink")
                for marker in ("begin", None, "separate", None, "end"):
                    run = etree.SubElement(field_parent, w + "r")
                    if marker in ("begin", "separate", "end"):
                        etree.SubElement(run, w + "fldChar", {w + "fldCharType": marker})
                    elif marker is None and len(field_parent) == 2:
                        etree.SubElement(run, w + "instrText").text = f"PAGEREF {target} \\h"
                    else:
                        etree.SubElement(run, w + "t").text = cached
            members["word/document.xml"] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            with ZipFile(docx, "w", ZIP_DEFLATED) as archive:
                for name, payload in members.items():
                    archive.writestr(name, payload)

            self.assertEqual(word_render_export.repair_post_update_pageref_targets(docx), 2)
            with ZipFile(docx) as archive:
                repaired = etree.fromstring(archive.read("word/document.xml"))
            text = " ".join(repaired.xpath(".//w:instrText/text() | .//w:t/text()", namespaces={"w": ns}))
            self.assertNotIn("_OldA", text)
            self.assertNotIn("_OldB", text)
            self.assertNotIn("错误!未定义书签", text)
            self.assertEqual(text.count("7"), 2)
            self.assertEqual(text.count("9"), 2)

    def test_generic_validator_has_no_attestation_cli_switch(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); docx = root / "a.docx"; pdf = root / "a.pdf"
            Document().save(docx); pdf.write_bytes(b"not needed: parser must reject the switch first")
            argv = ["word_render_validate.py", str(docx), str(pdf), "--out", str(root / "r.json"),
                    "--attest-word-export"]
            with patch.object(sys, "argv", argv), self.assertRaises(SystemExit) as raised:
                word_render_validate.main()
            self.assertEqual(raised.exception.code, 2)

    def test_refuses_in_place_word_update_before_platform_or_word_access(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "source.docx"
            Document().save(source)
            argv = ["word_render_export.py", str(source), str(source), str(source.with_suffix(".pdf")),
                    "--report", str(source.with_suffix(".json"))]
            with patch.object(sys, "argv", argv), patch.object(
                    word_render_export.platform, "system", return_value="Darwin"):
                with self.assertRaises(SystemExit) as raised:
                    word_render_export.main()
            self.assertEqual(raised.exception.code, 2)

    def test_refuses_pdf_path_aliasing_final_docx(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "source.docx"; final = root / "final.docx"
            Document().save(source)
            argv = ["word_render_export.py", str(source), str(final), str(final),
                    "--report", str(root / "report.json")]
            with patch.object(sys, "argv", argv), patch.object(
                    word_render_export.platform, "system", return_value="Darwin"):
                with self.assertRaises(SystemExit) as raised:
                    word_render_export.main()
            self.assertEqual(raised.exception.code, 2)
            self.assertFalse(final.exists())


if __name__ == "__main__":
    unittest.main()
